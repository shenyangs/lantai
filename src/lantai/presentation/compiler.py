# -*- coding: utf-8 -*-
"""Compile DeckIR into editable PPTX using deterministic semantic layouts."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_CONNECTOR_TYPE, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from .layouts import check_capacity, get_layout
from .assets import resolve_deck_assets
from .models import (
    ChartObject,
    DeckIR,
    DeckObject,
    DeckSlide,
    ImageObject,
    MetricObject,
    ShapeObject,
    TableObject,
    TextObject,
)
from .theme import ThemeTokens, get_theme


SLIDE_WIDTH = 13.333
SLIDE_HEIGHT = 7.5
_SURFACE_LAYOUTS = {
    "two_column",
    "three_cards",
    "comparison",
    "process",
    "matrix",
    "hierarchy",
}


@dataclass(frozen=True)
class Frame:
    x: float
    y: float
    w: float
    h: float


@dataclass
class CompilationResult:
    output_path: Path
    deck_id: str
    slide_count: int
    editable_object_count: int
    degradations: list[str] = field(default_factory=list)
    native_repairs: list[str] = field(default_factory=list)


def _layout_frames(layout_id: str) -> dict[str, Frame]:
    title = Frame(0.75, 0.48, 11.85, 0.72)
    layouts: dict[str, dict[str, Frame]] = {
        "cover": {
            "title": Frame(0.9, 2.0, 6.85, 1.35),
            "subtitle": Frame(0.92, 3.55, 6.6, 0.8),
            "meta": Frame(0.92, 4.48, 6.6, 0.52),
            "cover_visual": Frame(8.05, 0.75, 4.7, 5.95),
        },
        "section": {
            "title": Frame(0.9, 2.25, 10.7, 1.0),
            "body": Frame(0.92, 3.45, 9.7, 1.1),
        },
        "key_judgment": {
            "title": title,
            "body": Frame(0.82, 1.55, 11.65, 4.95),
        },
        "two_column": {
            "title": title,
            "left": Frame(0.75, 1.55, 5.78, 5.15),
            "right": Frame(6.8, 1.55, 5.78, 5.15),
        },
        "three_cards": {
            "title": title,
            "card_1": Frame(0.75, 1.7, 3.82, 4.85),
            "card_2": Frame(4.76, 1.7, 3.82, 4.85),
            "card_3": Frame(8.77, 1.7, 3.82, 4.85),
        },
        "comparison": {
            "title": title,
            "left": Frame(0.75, 1.58, 5.78, 5.05),
            "right": Frame(6.8, 1.58, 5.78, 5.05),
        },
        "process": {
            "title": title,
            **{
                f"step_{i}": Frame(0.75 + (i - 1) * 2.42, 2.05, 2.14, 3.75)
                for i in range(1, 6)
            },
        },
        "data_conclusion": {
            "title": title,
            "chart": Frame(0.75, 1.6, 7.45, 4.7),
            "table": Frame(0.75, 1.6, 7.45, 4.7),
            "metric_1": Frame(8.55, 1.65, 3.8, 1.42),
            "metric_2": Frame(8.55, 3.25, 3.8, 1.42),
            "conclusion": Frame(8.55, 4.95, 3.8, 1.35),
        },
        "agenda": {"title": title, "items": Frame(0.95, 1.55, 11.2, 5.2)},
        "matrix": {
            "title": title,
            "cell_1": Frame(0.75, 1.55, 5.75, 2.35),
            "cell_2": Frame(6.82, 1.55, 5.75, 2.35),
            "cell_3": Frame(0.75, 4.15, 5.75, 2.35),
            "cell_4": Frame(6.82, 4.15, 5.75, 2.35),
        },
        "hierarchy": {
            "title": title,
            "root": Frame(0.75, 1.48, 11.85, 1.05),
            "node_1": Frame(0.75, 2.88, 5.78, 1.62),
            "node_2": Frame(6.8, 2.88, 5.78, 1.62),
            "node_3": Frame(0.75, 4.78, 5.78, 1.62),
            "node_4": Frame(6.8, 4.78, 5.78, 1.62),
        },
        "table": {
            "title": title,
            "table": Frame(0.75, 1.45, 11.85, 4.8),
            "conclusion": Frame(0.75, 6.35, 11.85, 0.5),
        },
        "chart": {
            "title": title,
            "chart": Frame(0.75, 1.45, 9.0, 5.2),
            "conclusion": Frame(10.0, 1.55, 2.55, 5.0),
        },
        "image_focus": {
            "title": title,
            "hero": Frame(0.75, 1.35, 11.85, 4.95),
            "caption": Frame(0.75, 6.38, 11.85, 0.48),
        },
        "quote": {
            "quote_visual": Frame(0.0, 0.0, 4.2, 7.5),
            "quote": Frame(4.75, 1.55, 7.55, 3.8),
            "attribution": Frame(7.1, 5.55, 5.0, 0.55),
        },
        "closing": {
            "title": Frame(1.0, 1.65, 10.8, 1.1),
            "body": Frame(1.0, 3.0, 10.5, 1.55),
            "contact": Frame(1.0, 5.25, 10.5, 0.65),
        },
    }
    return layouts[layout_id]


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


class PresentationCompiler:
    """A bounded compiler: models choose semantics, this class owns coordinates."""

    def compile(
        self,
        deck: DeckIR,
        output_path: str | Path,
    ) -> CompilationResult:
        deck = resolve_deck_assets(deck)
        get_theme(deck.theme_id)
        output = Path(output_path)
        output.parent.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        presentation.slide_width = Inches(SLIDE_WIDTH)
        presentation.slide_height = Inches(SLIDE_HEIGHT)
        editable_count = 0
        degradations: list[str] = []

        for index, deck_slide in enumerate(deck.slides, start=1):
            get_layout(deck_slide.layout_id)
            violations = check_capacity(deck_slide)
            for violation in violations:
                degradations.append(
                    f"{deck_slide.id}:{violation.slot}:{violation.message}",
                )
            slide = presentation.slides.add_slide(
                presentation.slide_layouts[6],
            )
            self._paint_background(
                slide,
                get_theme(deck.theme_id),
                deck_slide.layout_id,
            )
            if deck_slide.layout_id == "freeform":
                for obj in deck_slide.objects:
                    if obj.frame is None:
                        continue
                    frame = Frame(
                        obj.frame.x,
                        obj.frame.y,
                        obj.frame.w,
                        obj.frame.h,
                    )
                    editable_count += self._render_slot(
                        slide,
                        frame,
                        [obj],
                        get_theme(deck.theme_id),
                        deck_slide.layout_id,
                    )
                self._add_footer(
                    slide,
                    deck,
                    deck_slide,
                    index,
                    get_theme(deck.theme_id),
                )
                continue
            frames = _layout_frames(deck_slide.layout_id)
            grouped: dict[str, list[DeckObject]] = {}
            for obj in deck_slide.objects:
                grouped.setdefault(obj.slot, []).append(obj)
            for slot, objects in grouped.items():
                if slot not in frames:
                    continue
                editable_count += self._render_slot(
                    slide,
                    frames[slot],
                    objects,
                    get_theme(deck.theme_id),
                    deck_slide.layout_id,
                )
            self._add_footer(
                slide,
                deck,
                deck_slide,
                index,
                get_theme(deck.theme_id),
            )

        presentation.save(output)
        return CompilationResult(
            output,
            deck.id,
            len(deck.slides),
            editable_count,
            degradations,
        )

    def _paint_background(
        self,
        slide,
        theme: ThemeTokens,
        layout_id: str,
    ) -> None:
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = _rgb(
            self._cover_background(theme)
            if layout_id == "cover"
            else theme.background,
        )
        if layout_id == "cover":
            band = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(0),
                Inches(0),
                Inches(0.24),
                Inches(SLIDE_HEIGHT),
            )
            band.name = "lantai-theme-accent"
            band.fill.solid()
            band.fill.fore_color.rgb = _rgb(theme.accent)
            band.line.fill.background()

    def _render_slot(
        self,
        slide,
        frame: Frame,
        objects: list[DeckObject],
        theme: ThemeTokens,
        layout_id: str,
    ) -> int:
        if len(objects) > 1 and all(
            isinstance(obj, (TextObject, MetricObject)) for obj in objects
        ):
            each_height = frame.h / len(objects)
            for index, obj in enumerate(objects):
                item_frame = Frame(
                    frame.x,
                    frame.y + index * each_height,
                    frame.w,
                    each_height - 0.08,
                )
                if isinstance(obj, TextObject):
                    self._add_text(
                        slide,
                        item_frame,
                        obj,
                        theme,
                        layout_id,
                    )
                else:
                    self._add_metric(slide, item_frame, obj, theme)
            return len(objects)
        count = 0
        for obj in objects:
            if isinstance(obj, TextObject):
                self._add_text(slide, frame, obj, theme, layout_id)
            elif isinstance(obj, MetricObject):
                self._add_metric(slide, frame, obj, theme)
            elif isinstance(obj, ImageObject):
                self._add_image(slide, frame, obj, theme)
            elif isinstance(obj, TableObject):
                self._add_table(slide, frame, obj, theme)
            elif isinstance(obj, ChartObject):
                self._add_chart(slide, frame, obj, theme)
            elif isinstance(obj, ShapeObject):
                self._add_shape(slide, frame, obj, theme)
            count += 1
        return count

    def _add_shape(
        self,
        slide,
        frame: Frame,
        obj: ShapeObject,
        theme: ThemeTokens,
    ) -> None:
        if obj.shape_type == "line":
            rendered = slide.shapes.add_connector(
                MSO_CONNECTOR_TYPE.STRAIGHT,
                Inches(frame.x),
                Inches(frame.y),
                Inches(frame.x + frame.w),
                Inches(frame.y + frame.h),
            )
            rendered.name = obj.id
            rendered.line.color.rgb = _rgb(obj.line_color or theme.accent)
            rendered.line.width = Pt(obj.line_width)
            return
        shape_type = {
            "rect": MSO_SHAPE.RECTANGLE,
            "round_rect": MSO_SHAPE.ROUNDED_RECTANGLE,
            "ellipse": MSO_SHAPE.OVAL,
            "chevron": MSO_SHAPE.CHEVRON,
        }[obj.shape_type]
        rendered = slide.shapes.add_shape(
            shape_type,
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
        )
        rendered.name = obj.id
        rendered.rotation = obj.rotate
        rendered.fill.solid()
        rendered.fill.fore_color.rgb = _rgb(obj.fill_color or theme.surface)
        rendered.line.color.rgb = _rgb(obj.line_color or theme.accent)
        rendered.line.width = Pt(obj.line_width)
        if obj.text:
            rendered.text_frame.clear()
            rendered.text_frame.word_wrap = True
            rendered.text_frame.margin_left = Inches(0.12)
            rendered.text_frame.margin_right = Inches(0.12)
            rendered.text_frame.margin_top = Inches(0.08)
            rendered.text_frame.margin_bottom = Inches(0.08)
            paragraph = rendered.text_frame.paragraphs[0]
            paragraph.text = obj.text
            paragraph.font.name = theme.body_font
            paragraph.font.size = Pt(obj.font_size)
            paragraph.font.color.rgb = _rgb(obj.text_color or theme.text)
            paragraph.alignment = PP_ALIGN.CENTER
            rendered.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    def _add_text(
        self,
        slide,
        frame: Frame,
        obj: TextObject,
        theme: ThemeTokens,
        layout_id: str,
    ) -> None:
        surfaced = layout_id in _SURFACE_LAYOUTS and obj.slot != "title"
        if surfaced:
            self._add_text_surface(slide, frame, obj, theme)
        box = slide.shapes.add_textbox(
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
        )
        box.name = obj.id
        box.text_frame.clear()
        box.text_frame.word_wrap = True
        horizontal_margin = 0.2 if surfaced else 0.08
        vertical_margin = 0.16 if surfaced else 0.06
        box.text_frame.margin_left = Inches(horizontal_margin)
        box.text_frame.margin_right = Inches(horizontal_margin)
        box.text_frame.margin_top = Inches(vertical_margin)
        box.text_frame.margin_bottom = Inches(vertical_margin)
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = obj.text
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.line_spacing = 1.12
        paragraph.font.name = (
            theme.title_font
            if obj.role in {"title", "heading"}
            else theme.body_font
        )
        paragraph.font.color.rgb = _rgb(
            self._cover_foreground(theme)
            if layout_id == "cover"
            else theme.primary
            if obj.role in {"title", "heading"}
            else theme.text,
        )
        paragraph.font.bold = obj.emphasis or obj.role in {"title", "heading"}
        paragraph.font.italic = obj.role == "quote"
        paragraph.font.size = Pt(self._font_size(obj, frame, layout_id))
        box.text_frame.vertical_anchor = (
            MSO_ANCHOR.MIDDLE
            if layout_id
            in {
                "cover",
                "section",
                "quote",
                "closing",
                "three_cards",
                "process",
                "matrix",
                "hierarchy",
            }
            else MSO_ANCHOR.TOP
        )

    def _add_text_surface(
        self,
        slide,
        frame: Frame,
        obj: TextObject,
        theme: ThemeTokens,
    ) -> None:
        surface = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
        )
        surface.name = f"{obj.id}-surface"
        surface.fill.solid()
        surface.fill.fore_color.rgb = _rgb(theme.surface)
        surface.line.color.rgb = _rgb(theme.accent)
        surface.line.width = Pt(1.1)
        if obj.role == "quote":
            mark = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(frame.x + 0.16),
                Inches(frame.y + 0.2),
                Inches(0.07),
                Inches(max(0.35, frame.h - 0.4)),
            )
            mark.name = f"{obj.id}-quote-mark"
            mark.fill.solid()
            mark.fill.fore_color.rgb = _rgb(theme.accent)
            mark.line.fill.background()

    def _font_size(
        self,
        obj: TextObject,
        frame: Frame,
        layout_id: str,
    ) -> float:
        if obj.role == "title":
            base = 42 if layout_id == "cover" else 32
            minimum = base
        elif obj.role == "heading":
            base, minimum = 22, 18
        elif obj.role == "quote":
            base, minimum = 20, 17
        else:
            base, minimum = 20, 17
        capacity = max(
            1.0,
            frame.w * frame.h * (8.0 if obj.role == "title" else 20.0),
        )
        pressure = len(obj.text) / capacity
        return max(minimum, base - max(0.0, pressure - 1.0) * 5.0)

    def _cover_background(self, theme: ThemeTokens) -> str:
        return (
            theme.primary
            if self._luminance(theme.background)
            > self._luminance(theme.primary)
            else theme.background
        )

    def _cover_foreground(self, theme: ThemeTokens) -> str:
        background = self._cover_background(theme)
        candidates = {
            theme.surface,
            theme.primary,
            theme.text,
            theme.background,
        }
        return max(
            candidates,
            key=lambda color: abs(
                self._luminance(color) - self._luminance(background),
            ),
        )

    @staticmethod
    def _luminance(color: str) -> float:
        red, green, blue = (
            int(color[index : index + 2], 16) / 255 for index in (0, 2, 4)
        )
        return 0.2126 * red + 0.7152 * green + 0.0722 * blue

    def _add_metric(
        self,
        slide,
        frame: Frame,
        obj: MetricObject,
        theme: ThemeTokens,
    ) -> None:
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
        )
        card.name = obj.id
        card.fill.solid()
        card.fill.fore_color.rgb = _rgb(theme.surface)
        card.line.color.rgb = _rgb(theme.accent)
        text_frame = card.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.margin_left = Inches(0.16)
        text_frame.margin_right = Inches(0.16)
        text_frame.margin_top = Inches(0.12)
        text_frame.margin_bottom = Inches(0.1)
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        label = text_frame.paragraphs[0]
        label.text = obj.label
        label.font.name = theme.body_font
        label.font.size = Pt(15)
        label.font.bold = True
        label.font.color.rgb = _rgb(theme.primary)
        label.alignment = PP_ALIGN.LEFT
        value = text_frame.add_paragraph()
        value.text = obj.value
        value.font.name = theme.title_font
        value.font.size = Pt(24)
        value.font.bold = True
        value.font.color.rgb = _rgb(theme.accent)
        value.alignment = PP_ALIGN.LEFT
        if obj.detail:
            detail = text_frame.add_paragraph()
            detail.text = obj.detail
            detail.font.name = theme.body_font
            detail.font.size = Pt(15)
            detail.font.color.rgb = _rgb(theme.muted)
            detail.alignment = PP_ALIGN.LEFT

    def _add_image(
        self,
        slide,
        frame: Frame,
        obj: ImageObject,
        theme: ThemeTokens,
    ) -> None:
        if not obj.path:
            raise ValueError(
                f"image {obj.id} was not resolved to an approved local path",
            )
        path = Path(obj.path)
        if not path.is_file():
            placeholder = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(frame.x),
                Inches(frame.y),
                Inches(frame.w),
                Inches(frame.h),
            )
            placeholder.name = obj.id
            placeholder.fill.solid()
            placeholder.fill.fore_color.rgb = _rgb(theme.surface)
            placeholder.line.color.rgb = _rgb(theme.muted)
            paragraph = placeholder.text_frame.paragraphs[0]
            paragraph.text = obj.alt or "图片暂不可用"
            paragraph.font.name = theme.body_font
            paragraph.font.size = Pt(12)
            paragraph.font.color.rgb = _rgb(theme.muted)
            paragraph.alignment = PP_ALIGN.CENTER
            return
        shape = slide.shapes.add_picture(
            str(path),
            Inches(frame.x),
            Inches(frame.y),
            width=Inches(frame.w),
            height=Inches(frame.h),
        )
        shape.name = obj.id

    def _add_table(
        self,
        slide,
        frame: Frame,
        obj: TableObject,
        theme: ThemeTokens,
    ) -> None:
        table_shape = slide.shapes.add_table(
            len(obj.rows) + 1,
            len(obj.headers),
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
        )
        table_shape.name = obj.id
        table = table_shape.table
        for column, header in enumerate(obj.headers):
            table.cell(0, column).text = header
        for row_index, row in enumerate(obj.rows, start=1):
            for column, value in enumerate(row):
                table.cell(row_index, column).text = value
        for row_index, row in enumerate(table.rows):
            for cell in row.cells:
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(
                    theme.primary if row_index == 0 else theme.surface,
                )
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.name = theme.body_font
                    paragraph.font.size = Pt(10)
                    paragraph.font.color.rgb = _rgb(
                        theme.background if row_index == 0 else theme.text,
                    )

    def _add_chart(
        self,
        slide,
        frame: Frame,
        obj: ChartObject,
        theme: ThemeTokens,
    ) -> None:
        data = ChartData()
        data.categories = obj.categories
        for name, values in obj.series.items():
            data.add_series(name, values)
        chart_types = {
            "bar": XL_CHART_TYPE.BAR_CLUSTERED,
            "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
            "line": XL_CHART_TYPE.LINE,
            "pie": XL_CHART_TYPE.PIE,
        }
        chart = slide.shapes.add_chart(
            chart_types[obj.chart_type],
            Inches(frame.x),
            Inches(frame.y),
            Inches(frame.w),
            Inches(frame.h),
            data,
        )
        chart.name = obj.id
        chart.chart.has_legend = (
            obj.show_legend
            if obj.show_legend is not None
            else len(obj.series) > 1
        )
        chart.chart.has_title = False
        if chart.chart.series:
            series = chart.chart.series[0]
            if obj.chart_type == "line":
                series.format.line.color.rgb = _rgb(theme.accent)
                series.format.line.width = Pt(2.25)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = _rgb(theme.accent)

    def _add_footer(
        self,
        slide,
        deck: DeckIR,
        deck_slide: DeckSlide,
        index: int,
        theme: ThemeTokens,
    ) -> None:
        footer = slide.shapes.add_textbox(
            Inches(0.75),
            Inches(7.02),
            Inches(11.85),
            Inches(0.22),
        )
        footer.name = f"meta:{deck.id}:{deck_slide.id}"
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"{deck.title}  ·  {index:02d}/{len(deck.slides):02d}"
        paragraph.font.name = theme.body_font
        paragraph.font.size = Pt(7.5)
        paragraph.font.color.rgb = _rgb(
            self._cover_foreground(theme)
            if deck_slide.layout_id == "cover"
            else theme.muted,
        )
        paragraph.alignment = PP_ALIGN.RIGHT
