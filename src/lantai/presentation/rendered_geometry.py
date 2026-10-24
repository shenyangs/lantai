"""Conservative native foreground-divider QA using rendered PDF glyph bounds.

Text-frame whitespace is not ink. Inspect tight, actually shaped PDF glyph
bounds against opaque axis-aligned rectangles in the final PPTX, preserving
native z-order. This is a narrow geometry check, not an aesthetic judgment or
a substitute for the visual critic, Office validation, or human acceptance.
"""

from __future__ import annotations

from dataclasses import dataclass
from math import isfinite
from pathlib import Path

import pypdfium2 as pdfium
from pptx import Presentation
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..utils.operation_control import check_operation
from .models import DeckIR, QualityIssue


GLYPH_DIVIDER_RULESET = "native-foreground-glyph-dividers/v1"
_POINTS_PER_EMU = 72 / 914400
_MAX_RULE_THICKNESS = 0.06 * 72
_MIN_RULE_LENGTH = 0.5 * 72
_GLYPH_CLEARANCE = 0.25
_MIN_CROSSING_LENGTH = 0.5
_FRAME_TOLERANCE = 0.5
_Box = tuple[float, float, float, float]


@dataclass(frozen=True)
class _Candidate:
    slide_index: int
    text_name: str
    divider_name: str
    text_frame: _Box
    divider_frame: _Box
    vertical: bool
    later_masks: tuple[_Box, ...]


def _frame(shape) -> _Box:
    return tuple(
        value * _POINTS_PER_EMU
        for value in (
            shape.left, shape.top,
            shape.left + shape.width, shape.top + shape.height,
        )
    )


def _opaque_rect(shape) -> bool:
    if (
        shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE
        or shape._element.xpath("./p:spPr/a:prstGeom/@prst") != ["rect"]
        or shape.rotation != 0
        or shape.fill.type != MSO_FILL_TYPE.SOLID
    ):
        return False
    fill = shape._element.xpath("./p:spPr/a:solidFill")
    return bool(fill) and not (
        any(value != "100000" for value in fill[0].xpath(".//a:alpha/@val"))
        or fill[0].xpath(".//a:alphaMod | .//a:alphaOff")
    )


def _overlap(first: _Box, second: _Box) -> bool:
    return (
        max(first[0], second[0]) < min(first[2], second[2])
        and max(first[1], second[1]) < min(first[3], second[3])
    )


def _candidates(output: Path) -> tuple[float, float, int, list[_Candidate]]:
    native = Presentation(output)
    candidates: list[_Candidate] = []
    for slide_index, slide in enumerate(native.slides):
        check_operation()
        shapes = list(slide.shapes)
        for index, divider in enumerate(shapes):
            if not _opaque_rect(divider) or divider.text.strip():
                continue
            box = _frame(divider)
            width, height = box[2] - box[0], box[3] - box[1]
            vertical = 0 < width <= _MAX_RULE_THICKNESS and height >= _MIN_RULE_LENGTH
            horizontal = 0 < height <= _MAX_RULE_THICKNESS and width >= _MIN_RULE_LENGTH
            if not (vertical or horizontal):
                continue
            masks = tuple(_frame(shape) for shape in shapes[index + 1:] if _opaque_rect(shape))
            for text in shapes[:index]:
                if (
                    not text.has_text_frame or not text.text.strip()
                    or text.rotation != 0 or not _overlap(_frame(text), box)
                ):
                    continue
                candidates.append(_Candidate(
                    slide_index, text.name, divider.name, _frame(text), box,
                    vertical, masks,
                ))
    return (
        native.slide_width * _POINTS_PER_EMU,
        native.slide_height * _POINTS_PER_EMU,
        len(native.slides), candidates,
    )


def has_glyph_divider_candidates(output: Path) -> bool:
    return bool(_candidates(output)[3])


def _inside_text_frame(glyph: _Box, frame: _Box) -> bool:
    x, y = (glyph[0] + glyph[2]) / 2, (glyph[1] + glyph[3]) / 2
    return (
        frame[0] - _FRAME_TOLERANCE <= x <= frame[2] + _FRAME_TOLERANCE
        and frame[1] - _FRAME_TOLERANCE <= y <= frame[3] + _FRAME_TOLERANCE
    )


def _crosses(glyph: _Box, candidate: _Candidate) -> bool:
    if not _inside_text_frame(glyph, candidate.text_frame):
        return False
    rule = candidate.divider_frame
    if candidate.vertical:
        crosses = (
            glyph[0] < rule[0] - _GLYPH_CLEARANCE
            and glyph[2] > rule[2] + _GLYPH_CLEARANCE
            and min(glyph[3], rule[3]) - max(glyph[1], rule[1]) > _MIN_CROSSING_LENGTH
        )
    else:
        crosses = (
            glyph[1] < rule[1] - _GLYPH_CLEARANCE
            and glyph[3] > rule[3] + _GLYPH_CLEARANCE
            and min(glyph[2], rule[2]) - max(glyph[0], rule[0]) > _MIN_CROSSING_LENGTH
        )
    # A later opaque rectangle hides the divider over the whole glyph. Do not
    # call an invisible rule a visible foreground collision.
    return crosses and not any(
        mask[0] <= glyph[0] and mask[1] <= glyph[1]
        and mask[2] >= glyph[2] and mask[3] >= glyph[3]
        for mask in candidate.later_masks
    )


def _object_id(deck: DeckIR, candidate: _Candidate) -> str | None:
    slide = deck.slides[candidate.slide_index]
    return next((
        obj.id for obj in slide.objects
        if candidate.text_name in {obj.id, f"{slide.id}::{obj.id}"}
    ), None)


def inspect_rendered_glyph_dividers(
    deck: DeckIR, output: Path, pdf_path: Path | None,
) -> list[QualityIssue]:
    """Return repairable collisions; absent/ambiguous geometry cannot pass."""

    document = None
    try:
        native_width, native_height, slide_count, candidates = _candidates(output)
        if not candidates:
            return []
        if pdf_path is None or not pdf_path.is_file():
            raise ValueError("missing rendered PDF")
        document = pdfium.PdfDocument(pdf_path)
        if len(document) != slide_count or slide_count != len(deck.slides):
            raise ValueError("rendered PDF page count does not match final PPTX")
        issues: list[QualityIssue] = []
        for page_index in sorted({candidate.slide_index for candidate in candidates}):
            check_operation()
            page = document[page_index]
            text_page = None
            try:
                width, height = page.get_size()
                if (
                    page.get_rotation() != 0
                    or not all(isfinite(value) and value > 0 for value in (width, height))
                    or abs(width / height - native_width / native_height) > 0.005
                ):
                    raise ValueError("rendered PDF geometry does not match final PPTX")
                text_page = page.get_textpage()
                glyphs: list[_Box] = []
                for char_index in range(text_page.count_chars()):
                    if char_index % 128 == 0:
                        check_operation()
                    code = pdfium.raw.FPDFText_GetUnicode(text_page, char_index)
                    if code and chr(code).isspace():
                        continue
                    # PDFium's default is the tight glyph box, not the loose
                    # line/text-frame bounds. Unknown Unicode still has ink.
                    left, bottom, right, top = text_page.get_charbox(char_index, loose=False)
                    if (
                        not all(isfinite(value) for value in (left, bottom, right, top))
                        or right <= left or top <= bottom
                    ):
                        continue
                    glyphs.append((
                        left * native_width / width,
                        (height - top) * native_height / height,
                        right * native_width / width,
                        (height - bottom) * native_height / height,
                    ))
                for candidate in candidates:
                    if candidate.slide_index != page_index:
                        continue
                    if not any(_inside_text_frame(glyph, candidate.text_frame) for glyph in glyphs):
                        raise ValueError("rendered text has no usable glyph geometry")
                    count = sum(_crosses(glyph, candidate) for glyph in glyphs)
                    if count:
                        issues.append(QualityIssue(
                            code="layout.glyph_divider_collision", severity="error",
                            slide_id=deck.slides[page_index].id,
                            object_id=_object_id(deck, candidate), repairable=True,
                            message=(
                                f"{count} rendered glyph boxes cross the opaque foreground "
                                f"divider {candidate.divider_name} on PDF page {page_index + 1}; "
                                "repair layout without losing text"
                            ),
                        ))
            finally:
                if text_page is not None:
                    text_page.close()
                page.close()
        return issues
    except Exception as exc:
        check_operation()  # Preserve cancellation/deadline failures.
        return [QualityIssue(
            code="render.glyph_geometry_unavailable", severity="error",
            message=f"foreground-divider glyph geometry could not be verified ({type(exc).__name__})",
        )]
    finally:
        if document is not None:
            document.close()
