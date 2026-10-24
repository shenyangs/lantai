"""Transactional edits for generated presentations.

Text-only edits preserve the existing native objects.  A semantic redesign is
compiled into an isolated candidate and must preserve deck/slide scope and
reopen with every declared target object before it can be committed.  Both
paths leave the authoritative PPTX untouched until the final hash-guarded
commit.
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
from pathlib import Path
from typing import Protocol

from pptx import Presentation

from .slide_package import graft_compiled_slide
from .notes_editing import replace_native_speaker_notes
from .models import (
    DeckIR,
    PresentationEditOperation,
    PresentationEditPreview,
    PresentationEditReceipt,
    TextObject,
    TextReplacement,
)
from .text_editing import replace_native_text_object
from .table_editing import is_table_text_only_edit


class DeckCompiler(Protocol):
    def compile(self, deck: DeckIR, output_path: str | Path):
        ...


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _atomic_copy(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(f".{destination.name}.tmp")
    shutil.copy2(source, temporary)
    os.replace(temporary, destination)


def _replace_text_preserving_format(shape, new_text: str, *, obj: TextObject) -> None:
    """Replace native text without discarding the existing run formatting.

    Assigning to ``TextFrame.text`` rebuilds paragraphs and runs.  That makes a
    seemingly harmless natural-language edit silently lose the compiler's font
    sizes and other native formatting. Known metric label-only edits retain
    every native paragraph. A bounded single-run edit retains its exact native
    carriers. Other whole-body replacements keep the original first-run carrier
    behavior without claiming caption provenance.
    """

    text_frame = shape.text_frame
    paragraphs = list(text_frame.paragraphs)
    if not paragraphs:
        raise ValueError(f"native text object has no paragraph: {shape.name}")
    edited = replace_native_text_object(obj, new_text)
    if (
        shape.text == obj.text
        and edited.native_metric_text is not None
        and edited.runs
        and edited.text == "".join(run.text for run in edited.runs)
    ):
        parts = new_text.split("\n")
        if len(parts) != len(paragraphs):
            raise ValueError(f"native metric paragraph structure changed: {shape.name}")
        for paragraph, text in zip(paragraphs, parts, strict=True):
            if paragraph.text == text:
                continue
            carriers = [run for run in paragraph.runs if run.text]
            if len(carriers) != 1:
                raise ValueError(f"native metric label format is ambiguous: {shape.name}")
            carriers[0].text = text
        return
    if edited.runs and edited.native_metric_text is None:
        native_runs = [run for paragraph in paragraphs for run in paragraph.runs]
        before_runs = [run for run in obj.runs if not run.paragraph_break]
        after_runs = [run for run in edited.runs if not run.paragraph_break]
        if (
            shape.text != obj.text
            or [run.text for run in native_runs] != [run.text for run in before_runs]
            or len(after_runs) != len(before_runs)
        ):
            raise ValueError(f"native rich-text run carrier is stale: {shape.name}")
        for native, before, after in zip(native_runs, before_runs, after_runs, strict=True):
            if before.text != after.text:
                native.text = after.text
        return
    first_paragraph = paragraphs[0]
    first_runs = list(first_paragraph.runs)
    if first_runs:
        first_runs[0].text = new_text.replace("\n", "\v")
        for run in first_runs[1:]:
            run.text = ""
    else:
        run = first_paragraph.add_run()
        run.text = new_text.replace("\n", "\v")
    for paragraph in paragraphs[1:]:
        runs = list(paragraph.runs)
        if runs:
            for run in runs:
                run.text = ""
        else:
            paragraph.text = ""


class PresentationEditTransactionService:
    def __init__(self, root: str | Path) -> None:
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def prepare_speaker_notes(
        self, source_pptx: str | Path, original_deck: DeckIR, changed_deck: DeckIR,
        operation: PresentationEditOperation,
    ) -> PresentationEditPreview:
        """Notes-only preview; no visible slide, object or dependency rewrite."""
        source = Path(source_pptx).resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        if operation.allowed_content_types != {"speaker_notes"} or operation.object_ids:
            raise PermissionError("native notes-only edits require explicit notes authority")
        if (operation.deck_id != original_deck.id
                or original_deck.model_dump(exclude={"slides"}) != changed_deck.model_dump(exclude={"slides"})):
            raise PermissionError("native notes edit changed deck identity or metadata")
        if [slide.id for slide in original_deck.slides] != [slide.id for slide in changed_deck.slides]:
            raise PermissionError("native notes edit cannot add, remove or reorder slides")
        targets = operation.slide_ids
        if not targets or len(targets) != len(set(targets)):
            raise ValueError("native notes require unique explicit target slides")
        if set(targets) - {slide.id for slide in original_deck.slides}:
            raise ValueError("native notes target slide is unknown")
        changed_ids = []
        for original, changed in zip(original_deck.slides, changed_deck.slides, strict=True):
            if original.id not in targets:
                if original != changed:
                    raise PermissionError("native notes edit changed an unapproved slide")
                continue
            if original.model_dump(exclude={"speaker_notes", "revision", "scene_hash"}) != (
                changed.model_dump(exclude={"speaker_notes", "revision", "scene_hash"})
            ):
                raise PermissionError("native notes edit changed visible content or page metadata")
            if original.speaker_notes != changed.speaker_notes:
                changed_ids.append(original.id)
        if not changed_ids:
            raise ValueError("native notes transaction produced no material change")
        transaction_dir = self.root / operation.operation_id
        if transaction_dir.exists():
            raise FileExistsError(f"operation already exists: {operation.operation_id}")
        transaction_dir.mkdir(parents=True)
        snapshot = transaction_dir / "source.snapshot.pptx"
        candidate = transaction_dir / "candidate.pptx"
        source_sha256 = _sha256(source)
        shutil.copy2(source, snapshot)
        if _sha256(snapshot) != source_sha256:
            raise RuntimeError("source changed while taking the notes edit snapshot")
        replace_native_speaker_notes(snapshot, candidate, original_deck, changed_deck, changed_ids)
        reopened = Presentation(str(candidate))
        if len(reopened.slides) != len(changed_deck.slides):
            raise ValueError("native notes candidate changed slide count")
        for index, changed in enumerate(changed_deck.slides):
            if changed.id in changed_ids:
                notes_frame = reopened.slides[index].notes_slide.notes_text_frame
                if notes_frame is None or notes_frame.text != changed.speaker_notes:
                    raise ValueError("native notes verification failed")
        if _sha256(source) != source_sha256:
            raise RuntimeError("source changed while preparing notes preview; refusing a stale edit")
        preview = PresentationEditPreview(
            operation=operation, source_path=str(source), candidate_path=str(candidate),
            snapshot_path=str(snapshot), source_sha256=source_sha256,
            changed_object_ids=[], change_summary=[f"{slide_id}: speaker notes updated" for slide_id in changed_ids],
            verified=True,
        )
        (transaction_dir / "preview.json").write_text(preview.model_dump_json(indent=2), encoding="utf-8")
        return preview

    def prepare(
        self,
        source_pptx: str | Path,
        deck_ir: DeckIR,
        operation: PresentationEditOperation,
        replacements: list[TextReplacement],
    ) -> PresentationEditPreview:
        source = Path(source_pptx).resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        if operation.deck_id != deck_ir.id:
            raise ValueError("edit operation targets a different deck")
        if operation.allowed_content_types != {"text"}:
            raise ValueError("V1 transactional edits allow text replacements only")
        if not replacements:
            raise ValueError("an edit preview requires at least one replacement")

        slide_by_object: dict[str, str] = {}
        text_by_object: dict[str, TextObject] = {}
        slide_index: dict[str, int] = {}
        for index, slide in enumerate(deck_ir.slides):
            slide_index[slide.id] = index
            for obj in slide.objects:
                slide_by_object[obj.id] = slide.id
                if isinstance(obj, TextObject):
                    text_by_object[obj.id] = obj

        replacement_ids = [item.object_id for item in replacements]
        if len(set(replacement_ids)) != len(replacement_ids):
            raise ValueError("replacement object ids must be unique")
        if operation.object_ids and not set(replacement_ids).issubset(operation.object_ids):
            raise PermissionError("replacement exceeds the approved object scope")
        unknown = set(replacement_ids) - set(text_by_object)
        if unknown:
            raise ValueError(f"unknown or non-text object ids: {sorted(unknown)}")
        if operation.slide_ids:
            outside = {item for item in replacement_ids if slide_by_object[item] not in operation.slide_ids}
            if outside:
                raise PermissionError("replacement exceeds the approved slide scope")

        transaction_dir = self.root / operation.operation_id
        if transaction_dir.exists():
            raise FileExistsError(f"operation already exists: {operation.operation_id}")
        transaction_dir.mkdir(parents=True)
        snapshot = transaction_dir / "source.snapshot.pptx"
        candidate = transaction_dir / "candidate.pptx"
        source_sha256 = _sha256(source)
        shutil.copy2(source, snapshot)
        if _sha256(snapshot) != source_sha256:
            raise RuntimeError("source changed while taking the edit snapshot")
        shutil.copy2(snapshot, candidate)

        presentation = Presentation(str(candidate))
        replacement_map = {item.object_id: item.text for item in replacements}
        found: set[str] = set()
        for object_id, new_text in replacement_map.items():
            ppt_slide = presentation.slides[slide_index[slide_by_object[object_id]]]
            for shape in ppt_slide.shapes:
                if shape.name != object_id:
                    continue
                if not getattr(shape, "has_text_frame", False):
                    raise ValueError(f"native object is not editable text: {object_id}")
                _replace_text_preserving_format(
                    shape, new_text, obj=text_by_object[object_id],
                )
                found.add(object_id)
                break
        if found != set(replacement_map):
            missing = sorted(set(replacement_map) - found)
            raise ValueError(f"native PPTX objects are missing: {missing}")
        presentation.save(str(candidate))
        self._verify_text(candidate, deck_ir, replacement_map)

        if _sha256(source) != source_sha256:
            raise RuntimeError("source changed while preparing preview; refusing a stale edit")
        preview = PresentationEditPreview(
            operation=operation,
            source_path=str(source),
            candidate_path=str(candidate),
            snapshot_path=str(snapshot),
            source_sha256=source_sha256,
            changed_object_ids=replacement_ids,
            change_summary=[
                f"{text_by_object[item.object_id].text[:80]} → {item.text[:80]}" for item in replacements
            ],
            verified=True,
        )
        (transaction_dir / "preview.json").write_text(
            preview.model_dump_json(indent=2), encoding="utf-8"
        )
        return preview

    def prepare_compiled_redesign(
        self,
        source_pptx: str | Path,
        original_deck: DeckIR,
        changed_deck: DeckIR,
        operation: PresentationEditOperation,
        compiler: DeckCompiler,
        *,
        compiled_slide_id: str | None = None,
        compiled_slide_ids: list[str] | None = None,
    ) -> PresentationEditPreview:
        """Prepare explicitly scoped whole-slide redesigns without mutation.

        The reasoning layer may replace the target slides' layouts and native
        objects, but it may not change deck identity, slide order, or any
        untargeted slide. Native object IDs are verified in the reopened PPTX
        so a compiler degradation cannot be promoted as a successful edit.
        """

        source = Path(source_pptx).resolve()
        if not source.is_file():
            raise FileNotFoundError(source)
        if (
            operation.deck_id != original_deck.id
            or changed_deck.id != original_deck.id
        ):
            raise ValueError("slide redesign targets a different deck")
        if compiled_slide_id is not None and compiled_slide_ids is not None:
            raise ValueError("compiled redesign accepts one target form")
        if compiled_slide_ids is not None:
            target_ids = list(compiled_slide_ids)
        elif compiled_slide_id is not None:
            target_ids = [compiled_slide_id]
        else:
            if len(operation.slide_ids) != 1:
                raise ValueError("compiled redesign requires exactly one target slide")
            target_ids = [operation.slide_ids[0]]
        if not target_ids:
            raise ValueError("compiled redesign requires at least one target slide")
        if len(set(target_ids)) != len(target_ids):
            raise ValueError("compiled redesign target slides must be unique")
        if not set(target_ids).issubset(operation.slide_ids):
            raise PermissionError(
                "compiled slide exceeds the approved slide scope",
            )
        if not operation.allowed_content_types:
            raise ValueError("compiled redesign requires explicit content types")
        before_ids = [slide.id for slide in original_deck.slides]
        after_ids = [slide.id for slide in changed_deck.slides]
        if after_ids != before_ids:
            raise ValueError("slide redesign cannot add, remove, or reorder slides")
        original_by_id = {slide.id: slide for slide in original_deck.slides}
        changed_by_id = {slide.id: slide for slide in changed_deck.slides}
        unknown_targets = [
            target_id for target_id in target_ids if target_id not in original_by_id
        ]
        if unknown_targets:
            raise ValueError(f"unknown target slide: {unknown_targets[0]}")
        target_id_set = set(target_ids)
        for before, after in zip(
            original_deck.slides,
            changed_deck.slides,
            strict=True,
        ):
            if before.id not in target_id_set and before != after:
                raise PermissionError("slide redesign changed content outside its scope")
        for target_id in target_ids:
            original_slide = original_by_id[target_id]
            changed_slide = changed_by_id[target_id]
            if original_slide == changed_slide:
                raise ValueError(
                    "slide redesign must materially change every target slide",
                )
            changed_kinds = {obj.kind for obj in changed_slide.objects}
            if not changed_kinds.issubset(operation.allowed_content_types):
                raise PermissionError(
                    "slide redesign introduced an unapproved content type",
                )

        transaction_dir = self.root / operation.operation_id
        if transaction_dir.exists():
            raise FileExistsError(f"operation already exists: {operation.operation_id}")
        transaction_dir.mkdir(parents=True)
        snapshot = transaction_dir / "source.snapshot.pptx"
        candidate = transaction_dir / "candidate.pptx"
        compiled = transaction_dir / "compiled.full.pptx"
        source_sha256 = _sha256(source)
        shutil.copy2(source, snapshot)
        if _sha256(snapshot) != source_sha256:
            raise RuntimeError("source changed while taking the edit snapshot")
        compiler.compile(changed_deck, compiled)
        original_native = Presentation(str(snapshot))
        graft_source = snapshot
        stage_paths: list[Path] = []
        identity_repairs: list[str] = []
        try:
            for index, target_id in enumerate(target_ids):
                stage = transaction_dir / f"candidate.stage-{index + 1}.pptx"
                stage_paths.append(stage)
                shutil.copy2(compiled, stage)
                original_slide = original_by_id[target_id]
                changed_slide = changed_by_id[target_id]
                # Explicit objects can keep their native carriers only when
                # inherited styling/assets and the page's geometry authority
                # have not changed. A genuine redesign must still execute its
                # new layout/theme rather than freezing equal-looking IR.
                same_context = all(
                    getattr(original_deck, field) == getattr(changed_deck, field)
                    for field in (
                        "title", "theme_id", "language", "aspect_ratio", "composition_tokens",
                        "design_profile", "design_context", "visual_direction",
                        "assets", "template_binding",
                    )
                ) and all(
                    getattr(original_slide, field) == getattr(changed_slide, field)
                    for field in ("layout_id", "composition_authority", "template_anchors")
                )
                before_objects = {obj.id: obj for obj in original_slide.objects}
                preserved_ids = {
                    obj.id for obj in changed_slide.objects
                    if same_context and obj.frame is not None
                    and obj == before_objects.get(obj.id)
                }
                table_text_edits = {
                    obj.id: (before_objects[obj.id], obj)
                    for obj in changed_slide.objects
                    if same_context and is_table_text_only_edit(before_objects.get(obj.id), obj)
                }
                marker_id = f"meta:{original_deck.id}:{target_id}"
                if same_context and any(
                    shape.name == marker_id
                    for shape in original_native.slides[before_ids.index(target_id)].shapes
                ):
                    preserved_ids.add(marker_id)
                repairs = graft_compiled_slide(
                    graft_source,
                    stage,
                    before_ids.index(target_id),
                    preserved_object_ids=preserved_ids,
                    table_text_edits=table_text_edits,
                )
                identity_repairs.extend(f"{target_id}: {repair}" for repair in repairs or [])
                graft_source = stage
            shutil.copy2(graft_source, candidate)
        finally:
            compiled.unlink(missing_ok=True)
            for stage in stage_paths:
                stage.unlink(missing_ok=True)
        for target_id in target_ids:
            self._verify_compiled_redesign(candidate, changed_deck, target_id)

        changed_object_ids: set[str] = set()
        change_summary: list[str] = list(identity_repairs)
        for target_id in target_ids:
            original_slide = original_by_id[target_id]
            changed_slide = changed_by_id[target_id]
            before_object_ids = {obj.id for obj in original_slide.objects}
            after_object_ids = {obj.id for obj in changed_slide.objects}
            changed_object_ids.update(before_object_ids | after_object_ids)
            change_summary.extend([
                f"{target_id}: {original_slide.layout_id} → {changed_slide.layout_id}",
                f"{target_id} native objects: "
                f"{len(before_object_ids)} → {len(after_object_ids)}",
            ])
        if _sha256(source) != source_sha256:
            raise RuntimeError("source changed while preparing preview; refusing a stale edit")
        preview = PresentationEditPreview(
            operation=operation,
            source_path=str(source),
            candidate_path=str(candidate),
            snapshot_path=str(snapshot),
            source_sha256=source_sha256,
            changed_object_ids=sorted(changed_object_ids),
            change_summary=change_summary,
            verified=True,
        )
        (transaction_dir / "preview.json").write_text(
            preview.model_dump_json(indent=2), encoding="utf-8"
        )
        return preview

    def commit(self, preview: PresentationEditPreview) -> PresentationEditReceipt:
        source = Path(preview.source_path).resolve()
        candidate = Path(preview.candidate_path).resolve()
        snapshot = Path(preview.snapshot_path).resolve()
        if not preview.verified or not candidate.is_file() or not snapshot.is_file():
            raise ValueError("only a verified preview can be committed")
        if _sha256(source) != preview.source_sha256:
            raise RuntimeError("source changed after preview; refusing a stale commit")

        _atomic_copy(candidate, source)
        Presentation(str(source))
        receipt = PresentationEditReceipt(
            operation_id=preview.operation.operation_id,
            source_path=str(source),
            snapshot_path=str(snapshot),
            committed_sha256=_sha256(source),
            changed_object_ids=preview.changed_object_ids,
            verified=True,
        )
        receipt_path = candidate.parent / "receipt.json"
        receipt_path.write_text(receipt.model_dump_json(indent=2), encoding="utf-8")
        return receipt

    def rollback(self, receipt: PresentationEditReceipt) -> None:
        source = Path(receipt.source_path).resolve()
        snapshot = Path(receipt.snapshot_path).resolve()
        if not snapshot.is_file():
            raise FileNotFoundError("transaction snapshot is unavailable")
        if _sha256(source) != receipt.committed_sha256:
            raise RuntimeError("source changed after commit; refusing a stale rollback")
        _atomic_copy(snapshot, source)
        Presentation(str(source))

    undo = rollback

    @staticmethod
    def _verify_text(path: Path, deck_ir: DeckIR, expected: dict[str, str]) -> None:
        presentation = Presentation(str(path))
        slide_index = {
            obj.id: index
            for index, slide in enumerate(deck_ir.slides)
            for obj in slide.objects
            if obj.id in expected
        }
        for object_id, text in expected.items():
            shapes = presentation.slides[slide_index[object_id]].shapes
            match = next((shape for shape in shapes if shape.name == object_id), None)
            if match is None or not getattr(match, "has_text_frame", False):
                raise ValueError(f"cannot reopen edited object: {object_id}")
            if match.text != text:
                raise ValueError(f"text verification failed for {object_id}")

    @staticmethod
    def _verify_compiled_redesign(
        path: Path,
        deck_ir: DeckIR,
        target_slide_id: str,
    ) -> None:
        presentation = Presentation(str(path))
        if len(presentation.slides) != len(deck_ir.slides):
            raise ValueError("compiled redesign changed the slide count")
        slide_index = next(
            index
            for index, slide in enumerate(deck_ir.slides)
            if slide.id == target_slide_id
        )
        expected = {
            obj.id
            for obj in deck_ir.slides[slide_index].objects
        }
        present = {shape.name for shape in presentation.slides[slide_index].shapes}
        missing = sorted(expected - present)
        if missing:
            raise ValueError(
                "compiled redesign is missing native target objects: "
                + ", ".join(missing)
            )
