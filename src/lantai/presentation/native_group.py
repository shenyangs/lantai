"""Lossless, read-only carriers for bounded native text/shape groups.

No flattening, authored XML, nested groups, or relationship transplantation is
allowed. A carrier retains its original coordinates and inherited style context;
the compiler replaces a private placeholder before atomically publishing output.
"""

from __future__ import annotations

from copy import deepcopy
import hashlib
import json
import os
from pathlib import Path
import re
import tempfile
from typing import TYPE_CHECKING
from zipfile import ZipFile

from lxml import etree as XML
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

if TYPE_CHECKING:
    from .models import DeckIR, NativeGroupObject

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MAC = "http://schemas.microsoft.com/office/mac/drawingml/2011/main"
THM15 = "http://schemas.microsoft.com/office/thememl/2012/main"
_THEME_FAMILY_URI = "{05A4C25C-085E-4340-85A3-A5531E510DB2}"
_GUID = re.compile(r"\{[0-9A-Fa-f]{8}(?:-[0-9A-Fa-f]{4}){3}-[0-9A-Fa-f]{12}\}")
NS = {"p": P, "a": A}
NATIVE_CANVAS = (12192000, 6858000)
EMU = 914400
_CONTEXT_ROOTS = {
    "theme_xml": f"{{{A}}}theme",
    "master_other_style_xml": f"{{{P}}}otherStyle",
    "master_color_map_xml": f"{{{P}}}clrMap",
    "presentation_default_text_style_xml": f"{{{P}}}defaultTextStyle",
}


def _reject() -> None:
    # Never echo XML, source text, or a user-controlled identifier into errors.
    raise ValueError("unsupported native group; authoritative PPTX was not changed")


def _validate_theme_family(root, node) -> None:
    # Office saves may add this identity-only leaf. Do not accept its optional
    # extension list, other ThemeML objects, or any relationship transplantation.
    extension = node.getparent()
    extensions = extension.getparent() if extension is not None else None
    if (root.tag != f"{{{A}}}theme" or node.tag != f"{{{THM15}}}themeFamily"
            or extension is None or extension.tag != f"{{{A}}}ext"
            or dict(extension.attrib) != {"uri": _THEME_FAMILY_URI}
            or len(extension) != 1 or (extension.text or "").strip()
            or extensions is None or extensions.tag != f"{{{A}}}extLst"
            or extensions.getparent() is not root or extensions.attrib
            or len(root.findall(f"{{{A}}}extLst")) != 1
            or len(node) or (node.text or "").strip() or (node.tail or "").strip()
            or set(node.attrib) != {"name", "id", "vid"}
            or any(_GUID.fullmatch(node.get(key, "")) is None for key in ("id", "vid"))):
        _reject()


def parse_native_xml(payload: str | bytes, *, limit: int = 524288,
                     context_kind: str | None = None):
    blob = payload.encode("utf-8") if isinstance(payload, str) else payload
    if not blob or len(blob) > limit or b"<!DOCTYPE" in blob.upper() or b"<!ENTITY" in blob.upper():
        _reject()
    try:
        root = XML.fromstring(blob, parser=XML.XMLParser(
            resolve_entities=False, load_dtd=False, no_network=True,
        ))
    except (XML.XMLSyntaxError, ValueError):
        _reject()
    if context_kind is not None and root.tag != _CONTEXT_ROOTS.get(context_kind):
        _reject()
    nodes = list(root.iter())
    if len(nodes) > 4096:
        _reject()
    theme_families = 0
    for node in nodes:
        if not isinstance(node.tag, str):
            _reject()
        tag = XML.QName(node)
        if tag.namespace == THM15 and context_kind == "theme_xml":
            _validate_theme_family(root, node)
            theme_families += 1
            if theme_families > 1:
                _reject()
            continue
        if tag.namespace not in {P, A, MAC}:
            _reject()
        if tag.namespace == MAC and tag.localname != "wrappingTextBoxFlag":
            _reject()
        if tag.localname in {
            "hlinkClick", "hlinkHover", "blip", "blipFill", "videoFile",
            "audioFile", "oleObj", "graphicFrame", "contentPart", "ph",
        }:
            _reject()
        for attribute in node.attrib:
            name = XML.QName(attribute)
            if name.namespace is not None or name.localname in {"embed", "link", "href"}:
                _reject()
    return root


def canonical_native_xml(payload: str | bytes, *, context_kind: str | None = None) -> bytes:
    return XML.tostring(parse_native_xml(payload, context_kind=context_kind),
                        method="c14n", exclusive=True)


def native_context_sha256(context) -> str:
    parts = {
        key: (canonical_native_xml(getattr(context, key), context_kind=key).decode("utf-8")
              if getattr(context, key) else "")
        for key in _CONTEXT_ROOTS
    }
    return hashlib.sha256(json.dumps(parts, sort_keys=True, separators=(",", ":"),
                                     ensure_ascii=False).encode("utf-8")).hexdigest()


def validate_native_context(context) -> None:
    for key, expected in _CONTEXT_ROOTS.items():
        payload = getattr(context, key)
        if not payload:
            if key in {"theme_xml", "master_color_map_xml"}:
                _reject()
            continue
        root = parse_native_xml(payload, limit=8192 if key == "master_color_map_xml" else 524288,
                                context_kind=key)
        if root.tag != expected:
            _reject()
    if context.sha256 != native_context_sha256(context):
        _reject()


def inspect_native_group(payload: str | bytes, canvas: tuple[int, int]) -> dict:
    if canvas != NATIVE_CANVAS:
        _reject()
    root = parse_native_xml(payload)
    if root.tag != f"{{{P}}}grpSp":
        _reject()
    children = list(root)
    if ([node.tag for node in children[:2]] != [f"{{{P}}}nvGrpSpPr", f"{{{P}}}grpSpPr"]
            or not 1 <= len(children[2:]) <= 64
            or any(node.tag != f"{{{P}}}sp" for node in children[2:])):
        _reject()
    if len(root.findall(f".//{{{P}}}grpSp")):
        _reject()
    # All descendants must remain pure text/shape primitives. Unknown p:sp
    # extensions are not a license to import another object/relationship type.
    allowed_p = {"grpSp", "nvGrpSpPr", "cNvPr", "cNvGrpSpPr", "nvPr",
                 "grpSpPr", "sp", "nvSpPr", "cNvSpPr", "spPr", "txBody",
                 "style", "extLst", "ext"}
    if any(XML.QName(node).namespace == P and XML.QName(node).localname not in allowed_p
           for node in root.iter()):
        _reject()
    nv = root.find("p:nvGrpSpPr/p:cNvPr", NS)
    xfrm = root.find("p:grpSpPr/a:xfrm", NS)
    if nv is None or xfrm is None or xfrm.attrib:
        _reject()  # rotations/flips of the group are outside this contract
    values = {}
    try:
        for tag, keys in {"off": ("x", "y"), "ext": ("cx", "cy"),
                          "chOff": ("x", "y"), "chExt": ("cx", "cy")}.items():
            nodes = xfrm.findall(f"a:{tag}", NS)
            if len(nodes) != 1 or set(nodes[0].attrib) != set(keys):
                _reject()
            values[tag] = tuple(int(nodes[0].get(key)) for key in keys)
        ids = [int(value) for value in root.xpath(".//p:cNvPr/@id", namespaces=NS)]
        group_id = int(nv.get("id"))
    except (TypeError, ValueError):
        _reject()
    if (len(ids) != len(children[2:]) + 1 or min(ids) < 1 or len(set(ids)) != len(ids)
            or not nv.get("name") or any(v <= 0 for v in (*values["ext"], *values["chExt"]))):
        _reject()
    x, y = values["off"]
    w, h = values["ext"]
    if x < 0 or y < 0 or x + w > canvas[0] or y + h > canvas[1]:
        _reject()
    canonical = XML.tostring(root, method="c14n", exclusive=True)
    return {"root": root, "ids": ids, "shape_id": group_id, "name": nv.get("name"),
            "frame": (x / EMU, y / EMU, w / EMU, h / EMU),
            "sha256": hashlib.sha256(canonical).hexdigest(), "child_count": len(children[2:])}


def validate_native_carrier(group) -> dict:
    info = inspect_native_group(group.native_xml, (group.source_width_emu, group.source_height_emu))
    validate_native_context(group.style_context)
    if (info["sha256"] != group.native_sha256 or info["shape_id"] != group.native_shape_id
            or info["name"] != group.native_name
            or any(abs(getattr(group.frame, key) - expected) * EMU > .01
                   for key, expected in zip(("x", "y", "w", "h"), info["frame"]))):
        _reject()
    return info


def native_group_read_only_summary(group) -> dict:
    """Give intelligence visible content, not an XML authoring capability."""
    info = validate_native_carrier(group)
    return {
        "kind": "native_group", "id": group.id, "slot": group.slot,
        "frame": group.frame.model_dump(mode="json"), "read_only": True,
        "native_name": group.native_name, "native_sha256": group.native_sha256,
        "native_texts": info["root"].xpath(".//a:t/text()", namespaces=NS),
    }


def capture_native_context(presentation, slide):
    from .models import NativeGroupStyleContext

    master = slide.slide_layout.slide_master
    for owner in (slide._element, slide.slide_layout._element):
        override = owner.find("p:clrMapOvr", NS)
        if override is not None and (len(override) != 1 or override[0].tag != f"{{{A}}}masterClrMapping"):
            _reject()
    def text(node):
        return XML.tostring(node, encoding="unicode") if node is not None else ""
    try:
        theme = master.part.part_related_by(RT.THEME).blob.decode("utf-8")
    except (KeyError, ValueError, UnicodeDecodeError):
        _reject()
    values = {
        "theme_xml": theme,
        "master_other_style_xml": text(master._element.find("p:txStyles/p:otherStyle", NS)),
        "master_color_map_xml": text(master._element.find("p:clrMap", NS)),
        "presentation_default_text_style_xml": text(presentation._element.find("p:defaultTextStyle", NS)),
    }
    # Hash canonical elements, not unused namespace declarations or formatting.
    class Context:
        pass
    context = Context()
    for key, value in values.items():
        setattr(context, key, value)
    return NativeGroupStyleContext(**values, sha256=native_context_sha256(context))


def validate_native_groups(deck: DeckIR) -> list[NativeGroupObject]:
    from .models import NativeGroupObject

    groups = [obj for slide in deck.slides for obj in slide.objects if isinstance(obj, NativeGroupObject)]
    if not groups:
        return []
    if (deck.template_binding is not None or len(groups) > 32
            or sum(len(obj.native_xml.encode("utf-8")) for obj in groups) > 2097152
            or len({obj.style_context.sha256 for obj in groups}) != 1):
        _reject()
    for slide in deck.slides:
        reserved = set()
        names = set()
        for obj in slide.objects:
            if not isinstance(obj, NativeGroupObject):
                continue
            info = validate_native_carrier(obj)
            if reserved.intersection(info["ids"]) or info["name"] in names:
                _reject()
            reserved.update(info["ids"])
            names.add(info["name"])
        if any(obj.id in names for obj in slide.objects if not isinstance(obj, NativeGroupObject)):
            _reject()
    return groups


def native_group_context_matches(presentation, slide, group) -> bool:
    try:
        return capture_native_context(presentation, slide).sha256 == group.style_context.sha256
    except ValueError:
        return False


def restore_native_groups(deck: DeckIR, output: Path) -> int:
    from .models import NativeGroupObject

    groups = validate_native_groups(deck)
    if not groups:
        return 0
    replacements: dict[str, bytes] = {}
    child_count = 0
    with ZipFile(output) as archive:
        # This path owns a single generated master/theme; mixed master/template
        # packages must never silently receive one group's inheritance context.
        if ({name for name in archive.namelist() if name.startswith("ppt/slideMasters/") and name.endswith(".xml")}
                != {"ppt/slideMasters/slideMaster1.xml"}
                or {name for name in archive.namelist() if name.startswith("ppt/theme/") and name.endswith(".xml")}
                != {"ppt/theme/theme1.xml"}):
            _reject()
        for index, slide in enumerate(deck.slides, start=1):
            on_slide = [obj for obj in slide.objects if isinstance(obj, NativeGroupObject)]
            if not on_slide:
                continue
            part = f"ppt/slides/slide{index}.xml"
            # Generated slides contain r: attributes for ordinary pictures;
            # they are not parsed by the relationship-free carrier parser.
            root = XML.fromstring(archive.read(part), parser=XML.XMLParser(resolve_entities=False, no_network=True))
            tree = root.find("p:cSld/p:spTree", NS)
            if tree is None:
                _reject()
            reserved = {value for obj in on_slide for value in validate_native_carrier(obj)["ids"]}
            all_nv = tree.findall(".//p:cNvPr", NS)
            ids = [int(nv.get("id")) for nv in all_nv]
            if len(ids) != len(set(ids)):
                _reject()
            next_id = max([*ids, *reserved]) + 1
            remap = {}
            for nv in all_nv:
                old = int(nv.get("id"))
                if old in reserved:
                    remap[old] = next_id
                    nv.set("id", str(next_id))
                    next_id += 1
            for tag in ("stCxn", "endCxn"):
                for node in tree.findall(f".//a:{tag}", NS):
                    old = int(node.get("id"))
                    if old in remap:
                        node.set("id", str(remap[old]))
            for obj in on_slide:
                matches = [node for node in tree.findall("p:sp", NS)
                           if node.find("p:nvSpPr/p:cNvPr", NS).get("name") == obj.id]
                if len(matches) != 1:
                    _reject()
                info = validate_native_carrier(obj)
                tree.replace(matches[0], deepcopy(info["root"]))
                child_count += info["child_count"]
            final_ids = tree.xpath(".//p:cNvPr/@id", namespaces=NS)
            if len(final_ids) != len(set(final_ids)):
                _reject()
            replacements[part] = XML.tostring(root, encoding="utf-8", xml_declaration=True)
        context = groups[0].style_context
        replacements["ppt/theme/theme1.xml"] = context.theme_xml.encode("utf-8")
        master = XML.fromstring(archive.read("ppt/slideMasters/slideMaster1.xml"))
        presentation = XML.fromstring(archive.read("ppt/presentation.xml"))
        for owner, path, payload in (
            (master, "p:txStyles/p:otherStyle", context.master_other_style_xml),
            (master, "p:clrMap", context.master_color_map_xml),
            (presentation, "p:defaultTextStyle", context.presentation_default_text_style_xml),
        ):
            existing = owner.find(path, NS)
            if existing is not None:
                parent = existing.getparent()
                if payload:
                    parent.replace(existing, parse_native_xml(payload))
                else:
                    parent.remove(existing)
            elif payload:
                _reject()  # generated package contract changed; no guessed order
        replacements["ppt/slideMasters/slideMaster1.xml"] = XML.tostring(master, encoding="utf-8", xml_declaration=True)
        replacements["ppt/presentation.xml"] = XML.tostring(presentation, encoding="utf-8", xml_declaration=True)
    fd, name = tempfile.mkstemp(prefix=f".{output.stem}-groups-", suffix=".pptx", dir=output.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        with ZipFile(output) as source, ZipFile(temporary, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, replacements.get(entry.filename, source.read(entry.filename)))
        from pptx import Presentation
        restored = Presentation(temporary)
        if len(restored.slides) != len(deck.slides):
            _reject()
        for expected, native in zip(deck.slides, restored.slides):
            for obj in expected.objects:
                if not isinstance(obj, NativeGroupObject):
                    continue
                matches = [shape for shape in native.shapes if shape.shape_id == obj.native_shape_id
                           and shape.name == obj.native_name]
                if (len(matches) != 1 or hashlib.sha256(canonical_native_xml(XML.tostring(matches[0]._element))).hexdigest()
                        != obj.native_sha256 or not native_group_context_matches(restored, native, obj)):
                    _reject()
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)
    return child_count
