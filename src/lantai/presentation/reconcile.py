# -*- coding: utf-8 -*-
"""Re-import an authoritative GenOffice save into DeckIR.

This module is deliberately independent of the retired Office Skills.  It
uses the public python-pptx object model to preserve stable identities, native
geometry and user-authored objects before any later AI mutation is allowed.
"""

from __future__ import annotations

import hashlib
from math import isfinite
from pathlib import Path
from typing import Iterable

from lxml import etree as XML

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn

from .models import (
    AssetRecord,
    ChartAxisStyle,
    ChartObject,
    ChartPlotLayout,
    CompositionAuthority,
    DeckIR,
    DeckObject,
    DeckSlide,
    ImageObject,
    MetricObject,
    NativeGroupObject,
    NativeMetricText,
    NativeShapeStyle,
    ObjectFrame,
    RichTextRun,
    ShapeObject,
    TableCellStyle,
    TableObject,
    TextObject,
)
from .quality import file_sha256
from .native_group import capture_native_context, inspect_native_group, validate_native_groups
from .native_fonts import NativeFontResolver


_CANVAS_WIDTH = 13.333
_CANVAS_HEIGHT = 7.5
_DERIVATIVE_SEPARATOR = "::"


def _bounded_frame(shape, slide_width: int, slide_height: int) -> ObjectFrame:
    # PptxGenJS LAYOUT_WIDE is 13 1/3 inches; the IR canvas is rounded
    # to 13.333. Rescaling an already-native page by that difference shrinks
    # every textbox and can move a character onto the next line.
    emu_per_inch = 914400
    native_canvas = (
        abs(slide_width / emu_per_inch - _CANVAS_WIDTH) < 0.001
        and abs(slide_height / emu_per_inch - _CANVAS_HEIGHT) < 0.001
    )
    x_scale = 1 / emu_per_inch if native_canvas else _CANVAS_WIDTH / slide_width
    y_scale = 1 / emu_per_inch if native_canvas else _CANVAS_HEIGHT / slide_height
    x = max(0.0, min(_CANVAS_WIDTH, shape.left * x_scale))
    y = max(0.0, min(_CANVAS_HEIGHT, shape.top * y_scale))
    w = max(0.001, shape.width * x_scale)
    h = max(0.0, shape.height * y_scale)
    w = min(w, _CANVAS_WIDTH - x)
    h = min(h, _CANVAS_HEIGHT - y)
    return ObjectFrame(x=x, y=y, w=max(0.001, w), h=h)


def _stable_id(slide_id: str, shape) -> str:
    fingerprint = f"{shape.shape_id}:{shape.name}:{shape.shape_type}"
    digest = hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()[:12]
    return f"imported-{slide_id}-{digest}"


def _shape_text(shape, *, preserve_whitespace: bool = False) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    # Shape labels are editable values, including blank labels and trailing
    # line breaks. Keep textbox import's existing normalization separate.
    return shape.text if preserve_whitespace else shape.text.strip()


def _font_size(shape) -> float | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    fonts = NativeFontResolver(shape)
    if any(size is None for size in fonts.visible_sizes()):
        raise ValueError(f"unresolved native font size: {shape.name}")
    return fonts.base_size()


def _paragraph_runs(
    text_frame, fonts: NativeFontResolver, color_fallback=None, *, preserve_soft_breaks=False,
) -> list[RichTextRun]:
    result: list[RichTextRun] = []
    for index, paragraph in enumerate(text_frame.paragraphs):
        if index:
            result.append(RichTextRun(text="\n", paragraph_break=True))
        native_runs = iter(paragraph.runs)
        for child in paragraph._p:
            if child.tag == qn("a:br") and preserve_soft_breaks:
                result.append(RichTextRun(text="\v", paragraph_break=True))
                continue
            if child.tag != qn("a:r"):
                continue
            run = next(native_runs)
            size = fonts.run_size(paragraph, run)
            color = fonts.run_rgb(paragraph, run)
            result.append(RichTextRun(
                text=run.text,
                # Office formatting is authoritative. Semantic roles only
                # style newly authored text, never an existing native save.
                bold=fonts.run_flag(paragraph, run, "b"),
                italic=fonts.run_flag(paragraph, run, "i"),
                color=str(color) if color is not None else color_fallback,
                font_size=size,
            ))
    return result


def _text_runs(shape, obj: TextObject) -> list[RichTextRun]:
    """Import the supported inline formatting from the authoritative save."""
    result = _paragraph_runs(shape.text_frame, NativeFontResolver(shape), obj.color)
    # Soft breaks and other text constructs not represented by paragraph.runs
    # must never disappear merely to preserve supported inline formatting.
    return result if "".join(run.text for run in result).strip() == _shape_text(shape) else []


def _shape_text_style(shape, color_fallback=None) -> dict:
    """Keep existing captions in bounded runs, not the body-authoring field."""
    text = _shape_text(shape, preserve_whitespace=True)
    size = _font_size(shape) if text.strip() else None
    runs = _paragraph_runs(
        shape.text_frame, NativeFontResolver(shape), color_fallback, preserve_soft_breaks=True,
    ) if text else []
    if "".join(run.text for run in runs) != text:
        raise ValueError("unsupported native shape text; authoritative PPTX was not changed")
    return {"runs": runs, "font_size": size if size is not None and 17 <= size <= 72 else 17}


def _rgb(fill_or_line) -> str | None:
    try:
        color_format = getattr(fill_or_line, "color", None)
        if color_format is None:
            color_format = fill_or_line.fore_color
        color = color_format.rgb
    except (AttributeError, TypeError, ValueError):
        return None
    return str(color) if color is not None else None


def _native_shape_style(shape) -> NativeShapeStyle:
    """Read native paint without LineFormat.color's solid-fill mutation."""
    properties = shape._element.spPr

    def invisible(node):
        return node is not None and (
            node.find(qn("a:noFill")) is not None
            or node.xpath("./a:solidFill/*/a:alpha/@val") == ["0"]
        )

    line = properties.find(qn("a:ln"))
    # Native caption shapes often have no line paint at all, not an explicit
    # noFill. Only treat that as absent when there is no inherited line style;
    # authored or theme-backed shapes must not silently lose their outline.
    line_paint = line is not None and any(
        line.find(qn("a:" + name)) is not None
        for name in ("noFill", "solidFill", "gradFill", "pattFill", "grpFill")
    )
    inherited_line = bool(shape._element.xpath("./p:style/a:lnRef"))
    values = {"fill_none": invisible(properties),
              "line_none": invisible(line) or (not line_paint and not inherited_line)}
    if _shape_text(shape, preserve_whitespace=True):
        frame = shape.text_frame
        alignments = {
            {None: "left", PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}.get(p.alignment)
            for p in frame.paragraphs
        }
        if None in alignments or len(alignments) != 1:
            raise ValueError("unsupported native shape paragraph alignment; authoritative PPTX was not changed")
        vertical = {None: "top", MSO_ANCHOR.TOP: "top", MSO_ANCHOR.MIDDLE: "mid", MSO_ANCHOR.BOTTOM: "bottom"}.get(frame.vertical_anchor)
        if vertical is None:
            raise ValueError("unsupported native shape vertical alignment; authoritative PPTX was not changed")
        values.update(align=alignments.pop(), valign=vertical)
        values.update({"margin_" + side: getattr(frame, "margin_" + side).pt
                       for side in ("top", "right", "bottom", "left")})
    return NativeShapeStyle(**values)


def _native_shape_line_width(shape, fallback=1) -> float:
    line = shape._element.spPr.find(qn("a:ln"))
    width = float(line.get("w")) / 12700 if line is not None and line.get("w") is not None else fallback
    if not isfinite(width) or not 0 <= width <= 12:
        raise ValueError("unsupported native shape line width; authoritative PPTX was not changed")
    return width


def _native_shape_colors(shape, style: NativeShapeStyle, fill_fallback=None, line_fallback=None) -> dict:
    fill = getattr(shape, "fill", None)
    line = shape._element.spPr.find(qn("a:ln"))
    # Do not read LineFormat.color: it converts noFill into a solid fill.
    line_fill = shape.line.fill if line is not None else None
    return {
        "fill_color": None if style.fill_none else _rgb(fill) or fill_fallback,
        "line_color": None if style.line_none else _rgb(line_fill) or line_fallback,
    }


def _chart_data(shape) -> tuple[list[str], dict[str, list[float]]]:
    try:
        plot = shape.chart.plots[0]
        categories = [str(item.label) for item in plot.categories]
        series = {
            str(item.name or f"series-{index}"): [float(value) for value in item.values]
            for index, item in enumerate(shape.chart.series, start=1)
        }
    except (AttributeError, IndexError, TypeError, ValueError):
        return [], {}
    return categories, series


def _chart_type(shape) -> str:
    supported = {
        XL_CHART_TYPE.BAR_CLUSTERED: "bar",
        XL_CHART_TYPE.COLUMN_CLUSTERED: "column",
        XL_CHART_TYPE.LINE: "line",
        # PptxGenJS serializes its native line chart with markers.  It is the
        # same editable line semantics supported by DeckIR, not a combo or
        # stacked chart, so it must survive the product's own round trip.
        XL_CHART_TYPE.LINE_MARKERS: "line",
        XL_CHART_TYPE.PIE: "pie",
    }
    chart = shape.chart
    if len(chart.plots) != 1 or chart.chart_type not in supported:
        raise ValueError(f"unsupported native chart ({shape.name}); authoritative PPTX was not changed")
    categories, series = _chart_data(shape)
    if not categories or not series or len(series) != len(chart.series):
        raise ValueError(f"unreadable or duplicate chart series ({shape.name}); authoritative PPTX was not changed")
    return supported[chart.chart_type]


def _chart_axes(shape) -> dict:
    manual = shape.chart.part._element.find("c:chart/c:plotArea/c:layout/c:manualLayout",
                                           namespaces={"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"})
    layout = None
    if manual is not None:
        values = {XML.QName(node).localname: node.get("val") for node in manual}
        if len(values) != len(manual) or manual.attrib:
            raise ValueError("unsupported native chart plot layout; authoritative PPTX was not changed")
        layout = ChartPlotLayout(**values)
    if _chart_type(shape) == "pie":
        return {"category_axis": None, "value_axis": None, "plot_layout": layout}
    positions = {
        XL_TICK_LABEL_POSITION.NONE: "none", XL_TICK_LABEL_POSITION.LOW: "low",
        XL_TICK_LABEL_POSITION.HIGH: "high", XL_TICK_LABEL_POSITION.NEXT_TO_AXIS: "nextTo",
    }
    result = {"plot_layout": layout}
    for name, axis in (("category_axis", shape.chart.category_axis),
                       ("value_axis", shape.chart.value_axis)):
        result[name] = ChartAxisStyle(
            orientation="maxMin" if axis.reverse_order else "minMax",
            minimum=axis.minimum_scale, maximum=axis.maximum_scale,
            major_unit=getattr(axis, "major_unit", None),
            minor_unit=getattr(axis, "minor_unit", None),
            position=axis._element.find(qn("c:axPos")).get("val"),
            label_position=positions[axis.tick_label_position],
            number_format=axis.tick_labels.number_format or "General",
        )
    return result


def _table_data(shape) -> tuple[list[str], list[list[str]]]:
    # Cell whitespace is authored native content, not an import delimiter.
    # Trimming it here makes a text-only edit lossy on its next round trip.
    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
    return (rows[0], rows[1:]) if rows else ([], [])


def _table_cell_valign(cell) -> str:
    # Older product exports leaked the IR's "mid" into a:tcPr/@anchor.
    # Read that exact legacy spelling without rewriting authoritative XML;
    # all other values still go through python-pptx's validating enum reader.
    properties = cell._tc.tcPr
    if properties is not None and properties.get("anchor") == "mid":
        return "mid"
    return {MSO_ANCHOR.MIDDLE: "mid", MSO_ANCHOR.BOTTOM: "bottom"}.get(cell.vertical_anchor, "top")


def _table_styles(shape) -> list[list[TableCellStyle]]:
    styles = []
    for row in shape.table.rows:
        cells = []
        for cell in row.cells:
            fonts = NativeFontResolver(cell)
            paragraph = cell.text_frame.paragraphs[0]
            align = {PP_ALIGN.LEFT: "left", PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}
            if paragraph.alignment not in {None, *align}:
                raise ValueError("unsupported native table alignment; authoritative PPTX was not changed")
            runs = _paragraph_runs(cell.text_frame, fonts)
            if "".join(run.text for run in runs).strip() != cell.text.strip():
                raise ValueError("unsupported native table text; authoritative PPTX was not changed")
            cells.append(TableCellStyle(
                align=align.get(paragraph.alignment, "left"),
                valign=_table_cell_valign(cell),
                font_size=fonts.base_size(), fill_color=_rgb(cell.fill), runs=runs,
                borders={side: XML.tostring(node, encoding="unicode")
                         for side in ("lnL", "lnR", "lnT", "lnB")
                         if (node := cell._tc.get_or_add_tcPr().find(qn("a:" + side))) is not None} or None,
            ))
        styles.append(cells)
    return styles


def _shape_type(shape) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    # PptxGenJS writes a line preset as p:sp rather than p:cxnSp.
    # python-pptx classifies this as AUTO_SHAPE but its enum has no line entry.
    presets = shape._element.xpath("./p:spPr/a:prstGeom")
    if presets and presets[0].get("prst") == "line":
        return "line"
    supported = {
        MSO_AUTO_SHAPE_TYPE.RECTANGLE: "rect",
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE: "round_rect",
        MSO_AUTO_SHAPE_TYPE.OVAL: "ellipse",
        MSO_AUTO_SHAPE_TYPE.CHEVRON: "chevron",
    }
    try:
        return supported[shape.auto_shape_type]
    except (AttributeError, ValueError, KeyError) as exc:
        raise ValueError(f"unsupported native shape ({shape.name}); authoritative PPTX was not changed") from exc


def _extract_picture(
    shape,
    *,
    assets_dir: Path,
    slide_id: str,
    object_id: str,
    source_ref: str,
) -> tuple[ImageObject, AssetRecord]:
    blob = shape.image.blob
    digest = hashlib.sha256(blob).hexdigest()
    extension = shape.image.ext or "bin"
    destination = assets_dir / f"{digest[:16]}.{extension}"
    assets_dir.mkdir(parents=True, exist_ok=True)
    if not destination.is_file():
        destination.write_bytes(blob)
    asset = AssetRecord(
        asset_id=f"asset-{digest[:24]}",
        local_path=str(destination),
        source_ref=source_ref,
        rights_status="user_provided",
        sha256=digest,
        usage_role="explanation",
    )
    image = ImageObject(
        id=object_id,
        path=str(destination),
        asset_id=asset.asset_id,
        alt=shape.name or f"Image on {slide_id}",
        frame=None,
    )
    return image, asset


def _update_known_object(
    obj: DeckObject,
    shape,
    frame: ObjectFrame,
    *,
    assets_dir: Path,
    slide_id: str,
    source_ref: str,
) -> tuple[DeckObject | None, AssetRecord | None]:
    text = _shape_text(shape)
    if isinstance(obj, TextObject):
        if not text:
            return None, None
        metric_text = obj.native_metric_text
        if (obj.role != "metric" or metric_text is None or not metric_text.caption_paragraphs(
                [paragraph.text for paragraph in shape.text_frame.paragraphs])):
            metric_text = None
        return obj.model_copy(
            update={
                "text": text, "frame": frame, "font_size": _font_size(shape),
                # The writer prioritizes runs over text. Re-read formatting
                # even when wording is unchanged; old IR runs are not authority.
                "runs": _text_runs(shape, obj),
                "native_metric_text": metric_text,
            },
        ), None
    if isinstance(obj, MetricObject):
        if not text:
            return None, None
        # The backend writes value, label and detail into one native textbox.
        # Its edited text cannot safely be assigned wholesale to `value`,
        # which would duplicate the label/detail on the next export.
        imported = TextObject(
            id=obj.id, role="metric", slot=obj.slot, text=text,
            frame=frame, font_size=_font_size(shape), valign="mid",
        )
        metric_text = NativeMetricText(value=obj.value, label=obj.label, detail=obj.detail)
        if not metric_text.caption_paragraphs(
                [paragraph.text for paragraph in shape.text_frame.paragraphs]):
            metric_text = None
        return imported.model_copy(update={
            "runs": _text_runs(shape, imported), "native_metric_text": metric_text,
        }), None
    if isinstance(obj, TableObject) and getattr(shape, "has_table", False):
        headers, rows = _table_data(shape)
        if headers:
            return obj.model_copy(update={"headers": headers, "rows": rows, "frame": frame,
                                          "cell_styles": _table_styles(shape)}), None
        return obj.model_copy(update={"frame": frame}), None
    if isinstance(obj, ChartObject) and getattr(shape, "has_chart", False):
        categories, series = _chart_data(shape)
        update = {"frame": frame, "chart_type": _chart_type(shape), **_chart_axes(shape)}
        if categories and series:
            update.update({"categories": categories, "series": series})
        return obj.model_copy(update=update), None
    if isinstance(obj, ImageObject) and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        imported, asset = _extract_picture(
            shape,
            assets_dir=assets_dir,
            slide_id=slide_id,
            object_id=obj.id,
            source_ref=source_ref,
        )
        return imported.model_copy(update={"frame": frame, "alt": obj.alt}), asset
    if isinstance(obj, ShapeObject):
        # Native connectors have a stroke but no fill; a generated header
        # rule must remain reimportable alongside its page provenance.
        style = _native_shape_style(shape)
        return obj.model_copy(
            update={
                "frame": frame,
                "text": _shape_text(shape, preserve_whitespace=True),
                **_shape_text_style(shape, obj.text_color),
                "rotate": shape.rotation,
                "native_style": style,
                "line_width": _native_shape_line_width(shape, obj.line_width),
                **_native_shape_colors(shape, style, obj.fill_color, obj.line_color),
                "shape_type": _shape_type(shape),
            },
        ), None
    return obj.model_copy(update={"frame": frame}), None


def _import_unknown_object(
    shape,
    *,
    frame: ObjectFrame,
    assets_dir: Path,
    slide_id: str,
    source_ref: str,
) -> tuple[DeckObject | None, AssetRecord | None]:
    object_id = _stable_id(slide_id, shape)
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        image, asset = _extract_picture(
            shape,
            assets_dir=assets_dir,
            slide_id=slide_id,
            object_id=object_id,
            source_ref=source_ref,
        )
        return image.model_copy(update={"frame": frame}), asset
    if getattr(shape, "has_table", False):
        headers, rows = _table_data(shape)
        if headers:
            return TableObject(id=object_id, headers=headers, rows=rows, frame=frame,
                               cell_styles=_table_styles(shape)), None
    if getattr(shape, "has_chart", False):
        categories, series = _chart_data(shape)
        if categories and series:
            return ChartObject(
                id=object_id,
                chart_type=_chart_type(shape),
                categories=categories,
                series=series,
                unit="原文件标注",
                derivation="从用户保存的权威 PPTX 原样重导入",
                source_refs=[source_ref],
                frame=frame,
                **_chart_axes(shape),
            ), None
    text = _shape_text(shape)
    if shape.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE}:
        style = _native_shape_style(shape)
        return ShapeObject(
            id=object_id,
            shape_type=_shape_type(shape),
            frame=frame,
            text=_shape_text(shape, preserve_whitespace=True),
            rotate=shape.rotation,
            **_shape_text_style(shape),
            native_style=style,
            line_width=_native_shape_line_width(shape),
            **_native_shape_colors(shape, style),
        ), None
    if text:
        imported = TextObject(
            id=object_id,
            role="body",
            text=text,
            frame=frame,
            font_size=_font_size(shape),
        )
        return imported.model_copy(update={"runs": _text_runs(shape, imported)}), None
    return None, None


def _slide_id_from_marker(shape_names: Iterable[str], deck_id: str) -> str | None:
    prefix = f"meta:{deck_id}:"
    for name in shape_names:
        if name.startswith(prefix):
            return name[len(prefix):]
    return None


def reconcile_pptx(
    deck: DeckIR,
    pptx_path: str | Path,
    *,
    assets_dir: str | Path,
) -> DeckIR:
    """Return DeckIR reconciled to the current native PPTX.

    The PPTX remains authoritative and is never rewritten here.  Stable names
    win; slides without a v3 marker fall back to positional identity so legacy
    decks can migrate without loading any retired Skill implementation.
    """

    source = Path(pptx_path).resolve()
    presentation = Presentation(str(source))
    source_ref = f"genoffice-save:{file_sha256(source)}"
    old_by_id = {slide.id: slide for slide in deck.slides}
    marker_ids = [
        _slide_id_from_marker((shape.name for shape in slide.shapes), deck.id)
        for slide in presentation.slides
    ]
    claimed_ids: set[str] = set()
    identities = []
    for index, ppt_slide in enumerate(presentation.slides):
        previous = old_by_id.get(marker_ids[index] or "")
        if previous is None and not any(marker_ids) and index < len(deck.slides):
            previous = deck.slides[index]
        if previous is not None and previous.id in claimed_ids:
            previous = None
        slide_id = previous.id if previous else f"imported-slide-{index + 1:02d}"
        while slide_id in claimed_ids or (previous is None and slide_id in old_by_id):
            slide_id += "-copy"
        claimed_ids.add(slide_id)
        identities.append((previous, slide_id, {obj.id: obj for obj in previous.objects} if previous else {}))
    # Validate ALL groups and inheritance before the first asset write. The
    # supported subset remains native XML, never flattened into guessed frames.
    group_contexts = {}
    group_infos = {}
    native_carriers = {}
    carrier_slides = []
    for slide_index, ppt_slide in enumerate(presentation.slides, start=1):
        previous, slide_id, known = identities[slide_index - 1]
        carriers = []
        for shape in ppt_slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                _shape_type(shape)
            elif shape.shape_type not in {
                MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.PICTURE,
                MSO_SHAPE_TYPE.LINE, MSO_SHAPE_TYPE.TEXT_BOX,
            } and not (getattr(shape, "has_chart", False)
                       or getattr(shape, "has_table", False)
                       or (shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER
                           and getattr(shape, "has_text_frame", False))):
                raise ValueError(
                    f"slide {slide_index} contains an unsupported native object "
                    f"({shape.name}); authoritative PPTX was not changed"
                )
            if getattr(shape, "has_chart", False):
                _chart_type(shape)
            if (shape.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE}
                and (shape.name not in known or isinstance(known[shape.name], ShapeObject))):
                _native_shape_style(shape)
                _native_shape_line_width(shape)
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                ids = ppt_slide._element.xpath(".//p:cNvPr/@id")
                if len(ids) != len(set(ids)):
                    raise ValueError("unsupported native group; duplicate native shape ids")
                payload = XML.tostring(shape._element, encoding="unicode")
                group_infos[(slide_index - 1, shape.shape_id)] = inspect_native_group(
                    payload, (presentation.slide_width, presentation.slide_height),
                )
                group_contexts[slide_index - 1] = capture_native_context(presentation, ppt_slide)
                info = group_infos[(slide_index - 1, shape.shape_id)]
                matches = [obj for obj in known.values() if isinstance(obj, NativeGroupObject)
                           and obj.native_shape_id == shape.shape_id and obj.native_name == shape.name]
                if len(matches) > 1:
                    raise ValueError("unsupported native group; ambiguous carrier identity")
                carrier = NativeGroupObject(
                    id=matches[0].id if matches else _stable_id(slide_id, shape),
                    native_xml=payload, native_sha256=info["sha256"],
                    native_shape_id=shape.shape_id, native_name=shape.name,
                    frame=ObjectFrame(**dict(zip(("x", "y", "w", "h"), info["frame"]))),
                    source_ref=source_ref, style_context=group_contexts[slide_index - 1],
                )
                carriers.append(carrier)
                native_carriers[(slide_index - 1, shape.shape_id)] = carrier
        native_names = {obj.native_name for obj in carriers}
        carrier_ids = {obj.id for obj in carriers}
        for shape in ppt_slide.shapes:
            name = shape.name or ""
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP or name.startswith("meta:"):
                continue
            if _DERIVATIVE_SEPARATOR in name and name.split(_DERIVATIVE_SEPARATOR, 1)[0] in known:
                continue
            object_id = name if name in known else _stable_id(slide_id, shape)
            if object_id in native_names or object_id in carrier_ids:
                raise ValueError("unsupported native group; ambiguous native name")
        carrier_slides.append(DeckSlide(id=slide_id, layout_id="freeform", objects=carriers))
    if len({context.sha256 for context in group_contexts.values()}) > 1:
        raise ValueError("unsupported native group; mixed inheritance contexts")
    validate_native_groups(deck.model_copy(update={"slides": carrier_slides}))
    asset_root = Path(assets_dir).resolve()
    updated_slides: list[DeckSlide] = []
    assets_by_id = {asset.asset_id: asset for asset in deck.assets}

    for index, ppt_slide in enumerate(presentation.slides):
        previous, slide_id, known = identities[index]
        objects: list[DeckObject] = []
        slide_assets: set[str] = set()

        for shape in ppt_slide.shapes:
            name = shape.name or ""
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                objects.append(native_carriers[(index, shape.shape_id)])
                continue
            if name.startswith("meta:"):
                continue
            if _DERIVATIVE_SEPARATOR in name and name.split(_DERIVATIVE_SEPARATOR, 1)[0] in known:
                continue
            frame = _bounded_frame(shape, presentation.slide_width, presentation.slide_height)
            asset: AssetRecord | None = None
            if name in known:
                imported, asset = _update_known_object(
                    known[name],
                    shape,
                    frame,
                    assets_dir=asset_root,
                    slide_id=slide_id,
                    source_ref=source_ref,
                )
            else:
                imported, asset = _import_unknown_object(
                    shape,
                    frame=frame,
                    assets_dir=asset_root,
                    slide_id=slide_id,
                    source_ref=source_ref,
                )
            if imported is not None:
                objects.append(imported)
                if isinstance(imported, ImageObject) and imported.asset_id:
                    slide_assets.add(imported.asset_id)
            if asset is not None:
                assets_by_id[asset.asset_id] = asset

        surviving_ids = {obj.id for obj in objects}
        groups = [
            group.model_copy(update={"object_ids": [item for item in group.object_ids if item in surviving_ids]})
            for group in (previous.object_groups if previous else [])
            if any(item in surviving_ids for item in group.object_ids)
        ]
        relationships = [
            relation
            for relation in (previous.relationships if previous else [])
            if set(relation.source_object_ids + relation.target_object_ids).issubset(surviving_ids)
        ]
        updated_slides.append(
            DeckSlide(
                id=slide_id,
                layout_id="freeform",
                semantic_layout=previous.provenance_layout if previous else None,
                intent_id=previous.intent_id if previous else None,
                # Native office geometry is the authoring authority after an
                # explicit save/import.  Losing that fact here makes the
                # downstream composer treat fully framed text/image pages as
                # AUTO and silently add or rearrange objects on the next edit.
                composition_authority=CompositionAuthority.AUTHOR,
                objects=objects,
                speaker_notes=(ppt_slide.notes_slide.notes_text_frame.text
                               if ppt_slide.has_notes_slide
                               and ppt_slide.notes_slide.notes_text_frame is not None else ""),
                revision=(previous.revision + 1 if previous else 1),
                editability_mode=previous.editability_mode if previous else "native-editable",
                source_fact_ids=previous.source_fact_ids if previous else [],
                asset_ids=sorted(slide_assets),
                object_groups=groups,
                relationships=relationships,
            ),
        )

    if not updated_slides:
        raise ValueError("authoritative PPTX contains no slides")
    reconciled = deck.model_copy(
        update={
            "slides": updated_slides,
            "assets": list(assets_by_id.values()),
        },
    )
    # Preflight rejects template-bound groups; never guess a context migration.
    validate_native_groups(reconciled)
    return reconciled
