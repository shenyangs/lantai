"""Preserve bounded, relationship-free native cell border XML in staging."""

import os
from pathlib import Path
import re
import tempfile
import zipfile

from lxml import etree as XML
from pptx import Presentation
from pptx.oxml.ns import qn

from .native_group import parse_native_xml


SIDES = ("lnL", "lnR", "lnT", "lnB")
_ATTRS = {
    **{side: {"w", "cap", "cmpd", "algn"} for side in SIDES},
    "solidFill": set(), "srgbClr": {"val"}, "noFill": set(),
    "prstDash": {"val"}, "round": set(), "bevel": set(), "miter": {"lim"},
    "headEnd": {"type", "w", "len"}, "tailEnd": {"type", "w", "len"},
}
_ENUMS = {
    "cap": {"rnd", "sq", "flat"},
    "cmpd": {"sng", "dbl", "thickThin", "thinThick", "tri"},
    "algn": {"ctr", "in"},
    "prstDash": {"solid", "dot", "dash", "lgDash", "dashDot", "lgDashDot",
                 "lgDashDotDot", "sysDash", "sysDot", "sysDashDot", "sysDashDotDot"},
    "type": {"none", "triangle", "stealth", "diamond", "oval", "arrow"},
    "end_size": {"sm", "med", "lg"},
}


def parse_table_border(side: str, payload: str):
    root = parse_native_xml(payload, limit=8192)
    if side not in SIDES or root.tag != qn("a:" + side):
        raise ValueError("unsupported native table border")
    seen = set()
    for node in root.iter():
        name = XML.QName(node).localname
        if node.tag != qn("a:" + name) or name not in _ATTRS or set(node.attrib) - _ATTRS[name]:
            raise ValueError("unsupported native table border")
        if (node.text and node.text.strip()) or (node.tail and node.tail.strip()):
            raise ValueError("unsupported native table border text")
        if node is not root:
            if name in SIDES or (node.getparent() is not root and not
                                 (name == "srgbClr" and node.getparent().tag == qn("a:solidFill"))):
                raise ValueError("unsupported native table border structure")
            if node.getparent() is root:
                if name in seen:
                    raise ValueError("duplicate native table border style")
                seen.add(name)
        for attribute, value in node.attrib.items():
            if not re.fullmatch(r"[0-9A-Za-z]{1,16}", value):
                raise ValueError("unsupported native table border value")
            enum = ("end_size" if name in {"headEnd", "tailEnd"} and attribute in {"w", "len"}
                    else "prstDash" if name == "prstDash" and attribute == "val" else attribute)
            if enum in _ENUMS and value not in _ENUMS[enum]:
                raise ValueError("unsupported native table border enum")
        if name == "srgbClr" and not re.fullmatch(r"[0-9A-Fa-f]{6}", node.get("val", "")):
            raise ValueError("unsupported native table border color")
        if name in SIDES and "w" in node.attrib:
            if not node.get("w").isdigit() or not 0 <= int(node.get("w")) <= 20116800:
                raise ValueError("unsupported native table border width")
        if name == "miter" and "lim" in node.attrib and not node.get("lim").isdigit():
            raise ValueError("unsupported native table border miter limit")
        if name == "solidFill" and (len(node) != 1 or node[0].tag != qn("a:srgbClr")):
            raise ValueError("unsupported native table border fill")
        if name not in {*SIDES, "solidFill"} and len(node):
            raise ValueError("unsupported native table border structure")
    if {"solidFill", "noFill"} <= seen:
        raise ValueError("conflicting native table border fills")
    if len(seen & {"round", "bevel", "miter"}) > 1:
        raise ValueError("conflicting native table border joins")
    return root


def restore_table_borders(deck, output: Path) -> None:
    from .models import TableObject
    if not any(isinstance(obj, TableObject) and any(cell.borders is not None for row in obj.cell_styles for cell in row)
               for slide in deck.slides for obj in slide.objects):
        return
    native = Presentation(output)
    if len(native.slides) != len(deck.slides):
        raise RuntimeError("native table border slide binding is inconsistent")
    replacements = {}
    for authored, slide in zip(deck.slides, native.slides):
        changed = False
        for obj in authored.objects:
            if not isinstance(obj, TableObject) or not obj.cell_styles:
                continue
            matches = [shape.table for shape in slide.shapes if shape.has_table and shape.name == obj.id]
            if len(matches) != 1:
                raise RuntimeError("native table borders have no unique authored binding")
            table = matches[0]
            if len(table.rows) != len(obj.cell_styles) or len(table.columns) != len(obj.headers):
                raise RuntimeError("native table border grid binding is inconsistent")
            for ri, row in enumerate(obj.cell_styles):
                for ci, cell in enumerate(row):
                    if cell.borders is None:
                        continue
                    props = table.cell(ri, ci)._tc.get_or_add_tcPr()
                    for node in list(props):
                        if node.tag in {qn("a:" + side) for side in SIDES}:
                            props.remove(node)
                    for side in reversed(SIDES):
                        if side in cell.borders:
                            props.insert(0, parse_table_border(side, cell.borders[side]))
                    changed = True
        if changed:
            replacements[str(slide.part.partname).lstrip("/")] = XML.tostring(slide._element, encoding="utf-8", xml_declaration=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{output.stem}-borders-", suffix=".pptx", dir=output.parent)
    os.close(fd)
    staged = Path(temporary)
    try:
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(staged, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, replacements.get(entry.filename, source.read(entry.filename)))
        os.replace(staged, output)
    finally:
        staged.unlink(missing_ok=True)
