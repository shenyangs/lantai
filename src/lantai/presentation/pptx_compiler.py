# -*- coding: utf-8 -*-
"""Generic native PPTX compiler boundary for Lantai presentations."""

from __future__ import annotations

import json
from io import BytesIO
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from xml.etree import ElementTree
import zipfile

from lxml import etree as XML
from ..utils import owned_subprocess as subprocess
from ..utils.operation_control import (
    check_operation,
    publication_boundary,
    remaining_timeout,
)

from .compiler import CompilationResult
from .design_system import prepare_deck_for_composition
from .models import ChartObject, DeckIR
from .native_group import restore_native_groups, validate_native_groups
from .native_chart_axes import restore_chart_axes
from .native_table_borders import restore_table_borders
from .native_table_identity import restore_generated_table_ids
from .template_binding import apply_template_binding, validate_template_binding


class NativePptxCompilerUnavailable(RuntimeError):
    """Raised when the native composition runtime cannot be started."""


_CONTENT_TYPES_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/content-types"
)
_SLIDE_MASTER_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml."
    "slideMaster+xml"
)


def _repair_backend_content_types(output: Path) -> None:
    """Repair a bounded PptxGenJS 4.0.x package-manifest defect.

    PptxGenJS 4.0.1 emits one slide-master override per slide while writing
    only ``slideMaster1.xml``.  Office tolerates the dangling overrides, but
    they make the OOXML package internally inconsistent.  This compiler
    boundary removes only missing slide-master overrides, then rejects every
    other missing override instead of masking a broader backend failure.
    """

    with zipfile.ZipFile(output, "r") as package:
        members = set(package.namelist())
        try:
            content_types = package.read("[Content_Types].xml")
        except KeyError as exc:
            raise RuntimeError(
                "native PPTX package has no content types",
            ) from exc

    try:
        root = ElementTree.fromstring(content_types)
    except ElementTree.ParseError as exc:
        raise RuntimeError(
            "native PPTX content types are invalid XML",
        ) from exc

    override_tag = f"{{{_CONTENT_TYPES_NAMESPACE}}}Override"
    removed = False
    for override in list(root.findall(override_tag)):
        part_name = str(override.attrib.get("PartName", "")).lstrip("/")
        if not part_name or part_name in members:
            continue
        if override.attrib.get("ContentType") == _SLIDE_MASTER_CONTENT_TYPE:
            root.remove(override)
            removed = True
            continue
        raise RuntimeError(
            "native PPTX content type targets missing package part: "
            f"{part_name}",
        )

    if not removed:
        return

    ElementTree.register_namespace("", _CONTENT_TYPES_NAMESPACE)
    repaired_content_types = ElementTree.tostring(
        root,
        encoding="utf-8",
        xml_declaration=True,
    )
    temp_fd, temp_name = tempfile.mkstemp(
        prefix=f".{output.stem}-",
        suffix=".pptx",
        dir=output.parent,
    )
    os.close(temp_fd)
    temp_output = Path(temp_name)
    try:
        with (
            zipfile.ZipFile(output, "r") as source,
            zipfile.ZipFile(temp_output, "w") as target,
        ):
            for entry in source.infolist():
                payload = (
                    repaired_content_types
                    if entry.filename == "[Content_Types].xml"
                    else source.read(entry.filename)
                )
                target.writestr(entry, payload)
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)

    with zipfile.ZipFile(output, "r") as package:
        members = set(package.namelist())
        repaired_root = ElementTree.fromstring(
            package.read("[Content_Types].xml"),
        )
    missing = [
        str(override.attrib.get("PartName", "")).lstrip("/")
        for override in repaired_root.findall(override_tag)
        if str(override.attrib.get("PartName", "")).lstrip("/") not in members
    ]
    if missing:
        raise RuntimeError(
            "native PPTX package still has missing content type targets: "
            + ", ".join(sorted(missing)),
        )


def _repair_backend_chart_axes(output: Path) -> None:
    """Remove only PptxGenJS's known orphan series axis in 2-D charts.

    Version 4.0.1 references its primary series axis in bar/line charts,
    although it declares that axis only for 3-D charts. Do not remove a
    declared series axis, repair another missing axis, or change chart data.
    This operates on the private compiler stage, never an existing candidate.
    """
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    def tag(name: str) -> str:
        return f"{{{namespace}}}{name}"

    axis_tags = {tag(name) for name in ("catAx", "valAx", "dateAx", "serAx")}
    repaired_parts: dict[str, bytes] = {}
    with zipfile.ZipFile(output) as package:
        for part in package.namelist():
            if not part.startswith("ppt/charts/") or not part.endswith(".xml"):
                continue
            payload = package.read(part)
            if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
                raise RuntimeError("native PPTX chart contains forbidden XML entities")
            try:
                root = XML.fromstring(payload, parser=XML.XMLParser(
                    resolve_entities=False, load_dtd=False, no_network=True,
                ))
            except XML.XMLSyntaxError:
                raise RuntimeError("native PPTX chart is invalid XML") from None
            plot_area = root.find(f".//{tag('plotArea')}")
            if plot_area is None:
                raise RuntimeError("native PPTX chart has no plot area")
            declared = {
                axis_id.get("val")
                for axis in plot_area if axis.tag in axis_tags
                for axis_id in axis.findall(tag("axId"))
            }
            changed = False
            for axis_id in plot_area.iter(tag("axId")):
                parent = axis_id.getparent()
                value = axis_id.get("val")
                if parent.tag in axis_tags or value in declared:
                    continue
                if (
                    value == "2094734556"
                    and parent.tag in {tag("barChart"), tag("lineChart")}
                    and {"2094734552", "2094734554"}.issubset(declared)
                ):
                    parent.remove(axis_id)
                    changed = True
                    continue
                raise RuntimeError("native PPTX chart references an undeclared axis")
            if any(axis.get("val") not in declared for axis in plot_area.iter(tag("crossAx"))):
                raise RuntimeError("native PPTX chart crosses an undeclared axis")
            if changed:
                repaired_parts[part] = XML.tostring(
                    root, encoding="utf-8", xml_declaration=True,
                )
    if not repaired_parts:
        return
    temp_fd, temp_name = tempfile.mkstemp(
        prefix=f".{output.stem}-", suffix=".pptx", dir=output.parent,
    )
    os.close(temp_fd)
    temp_output = Path(temp_name)
    try:
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(temp_output, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, repaired_parts.get(entry.filename, source.read(entry.filename)))
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)


def _chart_workbook_facts(output: Path, deck: DeckIR) -> dict[str, dict[str, float]]:
    """Bind generated workbooks to explicit authored chart IDs and values."""
    charts = [obj for slide in deck.slides for obj in slide.objects
              if isinstance(obj, ChartObject)]
    if not charts:
        return {}
    from pptx import Presentation

    presentation = Presentation(output)
    if len(presentation.slides) != len(deck.slides):
        raise RuntimeError("native chart workbook slide binding is inconsistent")
    facts: dict[str, dict[str, float]] = {}
    for authored_slide, native_slide in zip(deck.slides, presentation.slides):
        for obj in authored_slide.objects:
            if not isinstance(obj, ChartObject):
                continue
            matches = [shape for shape in native_slide.shapes
                       if shape.has_chart and shape.name == obj.id]
            if len(matches) != 1:
                raise RuntimeError("native chart workbook has no unique authored binding")
            workbook = matches[0].chart.part.chart_workbook.xlsx_part
            if workbook is None:
                raise RuntimeError("native chart has no editable workbook")
            part_name = str(workbook.partname).lstrip("/")
            if part_name in facts:
                raise RuntimeError("native charts unexpectedly share a data workbook")
            cells: dict[str, float] = {}
            for column_index, values in enumerate(obj.series.values(), start=2):
                letters = ""
                remainder = column_index
                while remainder:
                    remainder, digit = divmod(remainder - 1, 26)
                    letters = chr(ord("A") + digit) + letters
                for row_index, value in enumerate(values, start=2):
                    cells[f"{letters}{row_index}"] = value
            facts[part_name] = cells
    return facts


def _repair_backend_chart_workbooks(output: Path, deck: DeckIR) -> None:
    """Fix two bounded workbook defects against explicit authored chart facts.

    PptxGenJS 4.0.1 writes ``A1:C3'`` instead of ``A1:C3`` in category
    chart workbooks. Apple importers consequently discard the native chart.
    Pin the repair to the actual sheet dimension and table column count;
    reject other inconsistent ranges. It also serializes numeric zero as an
    empty cell. Repair only those empty numeric cells whose authored value is
    exactly zero; reject any other changed fact and preserve all other parts.
    """
    authored_facts = _chart_workbook_facts(output, deck)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    def tag(name: str) -> str:
        return f"{{{namespace}}}{name}"

    def parse(payload: bytes):
        if b"<!DOCTYPE" in payload.upper() or b"<!ENTITY" in payload.upper():
            raise RuntimeError("native chart workbook contains forbidden XML entities")
        try:
            return XML.fromstring(payload, parser=XML.XMLParser(
                resolve_entities=False, load_dtd=False, no_network=True,
            ))
        except XML.XMLSyntaxError:
            raise RuntimeError("native chart workbook is invalid XML") from None

    repaired_parts: dict[str, bytes] = {}
    with zipfile.ZipFile(output) as package:
        for part in package.namelist():
            if not part.startswith("ppt/embeddings/Microsoft_Excel_Worksheet") or not part.endswith(".xlsx"):
                continue
            with zipfile.ZipFile(BytesIO(package.read(part))) as workbook:
                try:
                    worksheet = parse(workbook.read("xl/worksheets/sheet1.xml"))
                    table = parse(workbook.read("xl/tables/table1.xml"))
                except KeyError:
                    raise RuntimeError("native chart workbook has no data sheet/table") from None
                dimension = worksheet.find(tag("dimension"))
                reference = dimension.get("ref", "") if dimension is not None else ""
                bounds = re.fullmatch(r"A1:([A-Z]{1,3})([1-9][0-9]{0,6})", reference)
                if bounds is None:
                    raise RuntimeError("native chart workbook has an invalid data range")
                column_count = 0
                for character in bounds.group(1):
                    column_count = column_count * 26 + ord(character) - ord("A") + 1
                columns = table.find(tag("tableColumns"))
                if (
                    column_count > 16384 or int(bounds.group(2)) > 1048576
                    or columns is None or columns.get("count") != str(column_count)
                    or len(columns.findall(tag("tableColumn"))) != column_count
                ):
                    raise RuntimeError("native chart workbook has inconsistent table bounds")
                table_reference = table.get("ref")
                if table_reference not in {reference, reference + "'"}:
                    raise RuntimeError("native chart workbook has an inconsistent table range")
                replacements: dict[str, bytes] = {}
                if table_reference != reference:
                    table.set("ref", reference)
                    replacements["xl/tables/table1.xml"] = XML.tostring(
                        table, encoding="utf-8", xml_declaration=True,
                    )
                cells = {cell.get("r"): cell for cell in worksheet.iter(tag("c"))}
                changed = False
                for address, expected in authored_facts.get(part, {}).items():
                    cell = cells.get(address)
                    value = cell.find(tag("v")) if cell is not None else None
                    if cell is None or cell.get("t", "n") != "n" or value is None:
                        raise RuntimeError("native chart workbook has a missing numeric fact")
                    if value.text in (None, "") and expected == 0:
                        value.text = "0"
                        changed = True
                        continue
                    try:
                        actual = float(value.text)
                    except (TypeError, ValueError):
                        raise RuntimeError("native chart workbook has an invalid numeric fact") from None
                    if actual != expected:
                        raise RuntimeError("native chart workbook changed an authored numeric fact")
                if changed:
                    replacements["xl/worksheets/sheet1.xml"] = XML.tostring(
                        worksheet, encoding="utf-8", xml_declaration=True,
                    )
                if not replacements:
                    continue
                repaired_workbook = BytesIO()
                with zipfile.ZipFile(repaired_workbook, "w") as target:
                    for entry in workbook.infolist():
                        payload = replacements.get(entry.filename, workbook.read(entry.filename))
                        target.writestr(entry, payload)
                repaired_parts[part] = repaired_workbook.getvalue()
    if not repaired_parts:
        return
    temp_fd, temp_name = tempfile.mkstemp(
        prefix=f".{output.stem}-", suffix=".pptx", dir=output.parent,
    )
    os.close(temp_fd)
    temp_output = Path(temp_name)
    try:
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(temp_output, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, repaired_parts.get(entry.filename, source.read(entry.filename)))
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)


class NativePptxCompiler:
    """Compile an explicitly composed DeckIR with one authoritative writer.

    The public boundary is independent from the currently selected JavaScript
    backend. Style and composition decisions must already exist in DeckIR.
    """

    engine_id = "native_pptx"

    def __init__(
        self,
        *,
        node_binary: str | Path | None = None,
        backend_module_path: str | Path | None = None,
        module_path: str | Path | None = None,
    ) -> None:
        self.node_binary = self._resolve_node(node_binary)
        # ``module_path`` is retained as a source-compatible constructor alias
        # for callers pinned before the compiler/backend boundary was named.
        self.backend_module_path = self._resolve_backend_module(
            backend_module_path or module_path,
        )
        self.script_path = Path(__file__).with_name("pptx_compiler.mjs")

    @staticmethod
    def _runtime_roots() -> list[Path]:
        configured = os.environ.get("LANTAI_PRESENTATION_RUNTIME")
        roots = [Path(configured).expanduser()] if configured else []
        if getattr(sys, "frozen", False):
            executable = Path(sys.executable)
            resources = next(
                (
                    parent
                    for parent in executable.parents
                    if parent.name == "Resources"
                ),
                executable.parent.parent / "Resources",
            )
            roots.append(resources / "presentation-runtime")
        roots.append(Path(__file__).with_name("runtime_assets"))
        return roots

    @staticmethod
    def _resolve_node(explicit: str | Path | None) -> Path:
        candidates = [
            explicit,
            os.environ.get("LANTAI_NODE_BIN"),
            os.environ.get("PRESENTATION_NODE_EXECUTABLE"),
            *(
                root / "node" / "bin" / "node"
                for root in NativePptxCompiler._runtime_roots()
            ),
            shutil.which("node"),
        ]
        for candidate in candidates:
            if candidate and Path(candidate).is_file():
                return Path(candidate)
        raise NativePptxCompilerUnavailable(
            (
                "Node.js runtime is unavailable; configure "
                "LANTAI_PRESENTATION_RUNTIME or LANTAI_NODE_BIN"
            ),
        )

    @staticmethod
    def _resolve_backend_module(explicit: str | Path | None) -> Path:
        node_modules = os.environ.get("LANTAI_NODE_MODULES") or os.environ.get(
            "PRESENTATION_NODE_MODULES",
        )
        candidates = [
            explicit,
            os.environ.get("LANTAI_PPTXGENJS_MODULE"),
            Path(node_modules) / "pptxgenjs" if node_modules else None,
            *(
                root / "node" / "node_modules" / "pptxgenjs"
                for root in NativePptxCompiler._runtime_roots()
            ),
        ]
        for candidate in candidates:
            if candidate and Path(candidate).exists():
                return Path(candidate)
        raise NativePptxCompilerUnavailable(
            (
                "native PPTX backend module is unavailable; "
                "configure LANTAI_PRESENTATION_RUNTIME or "
                "LANTAI_PPTXGENJS_MODULE"
            ),
        )

    def compile(
        self,
        deck: DeckIR,
        output_path: str | Path,
    ) -> CompilationResult:
        check_operation()
        # The native backend silently drops additional pie series. A chart
        # type change is an author's decision, not a compiler repair. Reject
        # this unsupported request before creating or replacing any artifact.
        if any(isinstance(obj, ChartObject) and obj.chart_type == "pie"
               and len(obj.series) != 1
               for slide in deck.slides for obj in slide.objects):
            raise RuntimeError("native pie chart requires exactly one data series")
        validate_native_groups(deck)
        template_snapshot = validate_template_binding(deck)
        deck = prepare_deck_for_composition(deck)
        output = Path(output_path).resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        if not self.script_path.is_file():
            raise NativePptxCompilerUnavailable(
                f"native composition script is missing: {self.script_path}",
            )
        artifact_fd, artifact_name = tempfile.mkstemp(
            prefix=f".{output.stem}-",
            suffix=".pptx",
            dir=output.parent,
        )
        os.close(artifact_fd)
        staged_output = Path(artifact_name)
        try:
            with tempfile.TemporaryDirectory(
                prefix="lantai-native-pptx-",
            ) as temp:
                spec_path = Path(temp) / "deck.json"
                spec_path.write_text(
                    deck.model_dump_json(indent=2),
                    encoding="utf-8",
                )
                child_env = os.environ.copy()
                if os.environ.get(
                    "PRESENTATION_NODE_MODE",
                    "",
                ).strip().lower() == "electron" and self.node_binary == Path(
                    os.environ.get("PRESENTATION_NODE_EXECUTABLE", ""),
                ):
                    child_env["ELECTRON_RUN_AS_NODE"] = "1"
                completed = subprocess.run(
                    [
                        str(self.node_binary),
                        str(self.script_path),
                        str(spec_path),
                        str(staged_output),
                        str(self.backend_module_path),
                    ],
                    env=child_env,
                    capture_output=True,
                    text=True,
                    timeout=remaining_timeout(180),
                    check=False,
                )
            if completed.returncode != 0:
                detail = (completed.stderr or completed.stdout).strip()
                raise RuntimeError(f"native PPTX composition failed: {detail}")
            if (
                not staged_output.is_file()
                or staged_output.stat().st_size == 0
            ):
                raise RuntimeError("native PPTX compiler produced no artifact")
            _repair_backend_content_types(staged_output)
            _repair_backend_chart_axes(staged_output)
            _repair_backend_chart_workbooks(staged_output, deck)
            restore_chart_axes(deck, staged_output)
            restore_table_borders(deck, staged_output)
            table_identity_repairs = restore_generated_table_ids(deck, staged_output)
            try:
                payload = json.loads(completed.stdout)
            except json.JSONDecodeError as exc:
                raise RuntimeError(
                    "native PPTX compiler returned invalid execution evidence",
                ) from exc
            if payload.get("engine") != self.engine_id:
                raise RuntimeError("native PPTX compiler identity mismatch")
            editable_object_count = int(payload["editable_object_count"])
            degradations = list(payload.get("degradations", []))
            editable_object_count += apply_template_binding(template_snapshot, staged_output)
            editable_object_count += restore_native_groups(deck, staged_output)
            # The backend never writes the authoritative candidate directly.
            # Only a complete, package-valid artifact with matching execution
            # evidence may replace the previous candidate, and the rename is
            # atomic because the staged file lives in the same directory.
            with publication_boundary():
                os.replace(staged_output, output)
        finally:
            staged_output.unlink(missing_ok=True)
        return CompilationResult(
            output_path=output,
            deck_id=deck.id,
            slide_count=len(deck.slides),
            editable_object_count=editable_object_count,
            degradations=degradations,
            native_repairs=table_identity_repairs,
        )
