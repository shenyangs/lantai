# -*- coding: utf-8 -*-
"""Structural and editability QA for compiled presentations."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from pathlib import Path

from PIL import Image, ImageStat
from lxml import etree as XML
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.dml import MSO_FILL_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_SHAPE_TYPE
from pptx.shapes.base import BaseShape

from ..utils.operation_control import check_operation
from .compiler import SLIDE_HEIGHT, SLIDE_WIDTH
from .design_system import measure_text
from .layouts import check_capacity
from .models import (
    ChartObject,
    DeckIR,
    EvidenceState,
    MetricObject,
    NativeGroupObject,
    QualityIssue,
    QualityReport,
    RelationshipType,
    ShapeObject,
)
from .native_group import canonical_native_xml, native_group_context_matches, validate_native_carrier
from .native_fonts import NativeFontResolver
from .rendered_geometry import GLYPH_DIVIDER_RULESET, has_glyph_divider_candidates


_PROCESS_TITLE_RE = re.compile(r"路线|流程|阶段|步骤|路径|时间线|进程")
_COMPARISON_TITLE_RE = re.compile(
    r"对比|比较|差异|区别|差距|中美|ToB\s*/\s*ToC|业务定位|\bvs\.?\b",
    re.IGNORECASE,
)


def delivery_status_for(
    issues: list[QualityIssue],
    *,
    content_complete: bool,
    editable: bool,
) -> str:
    """Only an error-free candidate may be considered deliverable."""

    errors = [issue for issue in issues if issue.severity == "error"]
    if not errors:
        return "passed"
    return "failed"


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _normalized_text(value: str) -> str:
    return re.sub(r"\s+", "", value)


def _native_text_overflows_fixed_frame(shape, text: str, sizes: list[float | None]) -> bool:
    """Detect severe horizontal overflow without claiming exact font shaping.

    Use the smallest effective native font, omit leading, and allow one whole
    estimated line of wrapping uncertainty. Growing, rotated, vertical, and
    multicolumn frames require Office-aware layout evidence instead.
    """
    if not text or not sizes or any(size is None for size in sizes):
        return False
    properties = shape._element.xpath("./p:txBody/a:bodyPr")
    if not properties:
        return False
    body = properties[0]
    if (
        body.xpath("./a:spAutoFit")
        or body.get("vert", "horz") != "horz"
        or body.get("numCol", "1") != "1"
        or shape.rotation != 0
        or shape.text_frame.word_wrap is False
        or any(
            int(properties.get("spc", "0")) < 0
            for properties in shape._element.xpath(".//a:rPr[@spc] | .//a:defRPr[@spc]")
        )
    ):
        return False
    frame = shape.text_frame
    width = (shape.width - frame.margin_left - frame.margin_right) / 914400
    height = (shape.height - frame.margin_top - frame.margin_bottom) / 914400
    if width <= 0 or height <= 0:
        return False
    font_size = min(size for size in sizes if size is not None)
    measured = measure_text(
        text.replace("\v", "\n"),
        font_size=font_size,
        frame_width=width,
        frame_height=height,
    )
    lower_bound_height = (measured.estimated_lines - 1) * font_size / 72
    return measured.estimated_lines >= 3 and lower_bound_height > height + 0.01


def _has_comparison_labels(labels: list[str]) -> bool:
    return len({_normalized_text(label) for label in labels if label.strip()}) >= 2


def _design_issues(
    deck: DeckIR,
    *,
    native_comparison_chart_ids: set[tuple[str, str]],
) -> list[QualityIssue]:
    """Reject document dumps that are mechanically valid but not slides."""
    issues: list[QualityIssue] = []
    content_slides = [
        slide
        for slide in deck.slides
        if slide.layout_id not in {"cover", "section", "closing"}
    ]
    text_dump_slides = 0
    body_signatures: list[tuple[str, frozenset[str]]] = []

    for slide in content_slides:
        relationship_types = {
            relationship.relationship_type
            for relationship in slide.relationships
        }
        text_objects = [obj for obj in slide.objects if obj.kind == "text"]
        title = next(
            (obj.text for obj in text_objects if obj.slot == "title"),
            "",
        )
        body_text = "\n".join(
            obj.text for obj in text_objects if obj.slot != "title"
        ).strip()
        body_lines = [line for line in body_text.splitlines() if line.strip()]
        body_signature = frozenset(
            normalized
            for obj in text_objects
            if obj.slot != "title"
            if len(normalized := _normalized_text(obj.text)) >= 6
        )
        duplicated_from = ""
        duplicated_count = 0
        duplicated_ratio = 0.0
        if len(body_signature) >= 3:
            for prior_slide_id, prior_signature in body_signatures:
                shared = body_signature & prior_signature
                shorter_size = min(len(body_signature), len(prior_signature))
                overlap_ratio = len(shared) / shorter_size
                if len(shared) >= 3 and overlap_ratio >= 0.75:
                    duplicated_from = prior_slide_id
                    duplicated_count = len(shared)
                    duplicated_ratio = overlap_ratio
                    break
        if duplicated_from:
            issues.append(
                QualityIssue(
                    code="design.repetition",
                    severity="error",
                    slide_id=slide.id,
                    message=(
                        f"{duplicated_count} body items ({duplicated_ratio:.0%} of "
                        "the shorter slide) duplicate slide "
                        f"{duplicated_from!r}; redesign one slide instead of "
                        "reusing the same information in a new shell"
                    ),
                ),
            )
        elif len(body_signature) >= 3:
            body_signatures.append((slide.id, body_signature))

        process_semantics = bool(
            relationship_types
            & {
                RelationshipType.PROCESS,
                RelationshipType.CHRONOLOGY,
                RelationshipType.JOURNEY,
                RelationshipType.ROADMAP,
                RelationshipType.CAUSE_EFFECT,
                RelationshipType.HIERARCHY,
            },
        ) or slide.layout_id in {"process", "hierarchy", "chart", "data_conclusion"}
        ambiguous_semantics = not relationship_types or relationship_types <= {
            RelationshipType.COMPOSITION,
            RelationshipType.MIXED,
        }
        if (
            _PROCESS_TITLE_RE.search(title)
            and not process_semantics
            and ambiguous_semantics
        ):
            issues.append(
                QualityIssue(
                    code="design.semantic_layout_mismatch",
                    severity="error",
                    slide_id=slide.id,
                    message="a route/process title requires a process or hierarchy composition",
                    repairable=True,
                ),
            )
        comparison_semantics = bool(
            relationship_types
            & {
                RelationshipType.COMPARISON,
                RelationshipType.BEFORE_AFTER,
                RelationshipType.DATA_TREND,
            },
        ) or slide.layout_id in {"comparison", "two_column", "chart", "data_conclusion"} or any(
            isinstance(obj, ChartObject)
            and obj.chart_type in {"bar", "column"}
            and (slide.id, obj.id) in native_comparison_chart_ids
            and (
                _has_comparison_labels(obj.categories)
                or _has_comparison_labels(list(obj.series))
            )
            for obj in slide.objects
        )
        if (
            _COMPARISON_TITLE_RE.search(title)
            and not comparison_semantics
            and ambiguous_semantics
        ):
            issues.append(
                QualityIssue(
                    code="design.semantic_layout_mismatch",
                    severity="error",
                    slide_id=slide.id,
                    message="a comparison title requires a comparison composition",
                    repairable=True,
                ),
            )

        prose_objects = [
            obj for obj in text_objects
            if obj.slot != "title" and len(obj.text.strip()) >= 80
        ]
        largest_prose = max((len(obj.text) for obj in prose_objects), default=0)
        # Several compact proof points are a legitimate judgment composition.
        # A text wall means one or two prose objects carry document-scale copy.
        judgment_like = slide.layout_id == "key_judgment" or bool(
            relationship_types
            & {
                RelationshipType.THESIS,
                RelationshipType.HERO_STATEMENT,
                RelationshipType.COMPOSITION,
                RelationshipType.MIXED,
            },
        )
        is_text_dump = judgment_like and (
            largest_prose > 140
            or (
                len(text_objects) <= 3
                and (len(body_text) > 180 or len(body_lines) > 5)
            )
        )
        if is_text_dump:
            text_dump_slides += 1
        if judgment_like and (
            largest_prose > 260
            or (
                len(text_objects) <= 3
                and len(body_text) > 320
                and len(body_lines) > 5
            )
        ):
            issues.append(
                QualityIssue(
                    code="design.text_wall",
                    severity="error",
                    slide_id=slide.id,
                    message="key-judgment slide contains a document-scale text wall",
                ),
            )

    if (
        len(content_slides) >= 3
        and text_dump_slides >= 3
        and text_dump_slides / len(content_slides) >= 0.6
    ):
        issues.append(
            QualityIssue(
                code="design.repetitive_text_dump",
                severity="error",
                message="most content slides repeat the same text-only document-dump composition",
            ),
        )
    if len(content_slides) >= 5:
        semantic_families = [
            "/".join(
                sorted(
                    relationship.relationship_type.value
                    for relationship in slide.relationships
                ),
            )
            or slide.layout_id
            for slide in content_slides
        ]
        dominant_layout, dominant_count = Counter(semantic_families).most_common(1)[0]
        if dominant_count / len(content_slides) >= 0.6:
            issues.append(
                QualityIssue(
                    code="design.repetitive_layout",
                    severity="warning",
                    message=(
                        f"layout {dominant_layout!r} dominates the content deck; "
                        "semantic visual review should confirm intentional variety"
                    ),
                ),
            )
    return issues


def _chart_semantic_issues(deck: DeckIR) -> list[QualityIssue]:
    """Reject charts that cannot be interpreted or audited downstream."""

    issues: list[QualityIssue] = []
    allowed_sources = set(deck.source_refs)
    for slide in deck.slides:
        relationship_types = {
            relationship.relationship_type
            for relationship in slide.relationships
        }
        for obj in slide.objects:
            if not isinstance(obj, ChartObject):
                continue
            if not obj.unit.strip():
                issues.append(
                    QualityIssue(
                        code="chart.unit_missing",
                        severity="error",
                        slide_id=slide.id,
                        object_id=obj.id,
                        message="native chart has no declared measurement unit",
                    ),
                )
            elif obj.unit.strip().lower() in {
                "benchmark score",
                "demo score",
                "sample score",
                "placeholder",
            }:
                issues.append(
                    QualityIssue(
                        code="chart.unit_placeholder",
                        severity="error",
                        slide_id=slide.id,
                        object_id=obj.id,
                        message="native chart uses a fixture/demo placeholder as its measurement unit",
                    ),
                )
            if not obj.derivation.strip():
                issues.append(
                    QualityIssue(
                        code="chart.derivation_missing",
                        severity="error",
                        slide_id=slide.id,
                        object_id=obj.id,
                        message="native chart has no derivation statement",
                    ),
                )
            if not obj.source_refs:
                issues.append(
                    QualityIssue(
                        code="chart.source_missing",
                        severity="error",
                        slide_id=slide.id,
                        object_id=obj.id,
                        message="native chart is not bound to any source",
                    ),
                )
            elif allowed_sources and set(obj.source_refs) - allowed_sources:
                issues.append(
                    QualityIssue(
                        code="chart.source_invalid",
                        severity="error",
                        slide_id=slide.id,
                        object_id=obj.id,
                        message="native chart references an undeclared source",
                    ),
                )
        if relationship_types & {
            RelationshipType.COMPOSITION,
            RelationshipType.COMPARISON,
        }:
            for metric in (
                obj for obj in slide.objects if isinstance(obj, MetricObject)
            ):
                if metric.label.strip() in {"起点", "当前", "开始", "现在"}:
                    issues.append(
                        QualityIssue(
                            code="chart.metric_relationship_mismatch",
                            severity="error",
                            slide_id=slide.id,
                            object_id=metric.id,
                            message=(
                                "composition/comparison slide labels a category as a temporal endpoint"
                            ),
                        ),
                    )
    return issues


def inspect_rendered_slides(
    paths: list[Path] | tuple[Path, ...],
) -> list[QualityIssue]:
    """Run deterministic image-level checks without interpreting slide semantics."""
    issues: list[QualityIssue] = []
    for index, path in enumerate(paths):
        try:
            with Image.open(path) as image:
                rgb = image.convert("RGB")
                width, height = rgb.size
                if width < 640 or height < 360:
                    issues.append(
                        QualityIssue(
                            code="render.resolution",
                            severity="error",
                            slide_id=f"slide-{index + 1}",
                            message="rendered slide resolution is too small",
                            repairable=True,
                        ),
                    )
                ratio = width / max(1, height)
                if abs(ratio - (16 / 9)) > 0.03:
                    issues.append(
                        QualityIssue(
                            code="render.aspect_ratio",
                            severity="error",
                            slide_id=f"slide-{index + 1}",
                            message="rendered slide is not 16:9",
                            repairable=True,
                        ),
                    )
                grayscale = rgb.convert("L").resize((160, 90))
                stat = ImageStat.Stat(grayscale)
                if stat.var[0] < 2.0:
                    issues.append(
                        QualityIssue(
                            code="render.blank_or_flat",
                            severity="error",
                            slide_id=f"slide-{index + 1}",
                            message="rendered slide appears blank or visually flat",
                            repairable=True,
                        ),
                    )
        except Exception as exc:
            issues.append(
                QualityIssue(
                    code="render.image_unreadable",
                    severity="error",
                    slide_id=f"slide-{index + 1}",
                    message=str(exc),
                    repairable=True,
                ),
            )
    return issues


def delivery_evidence_issues(
    deck: DeckIR,
    report: QualityReport,
    output_path: str | Path,
    rendered_paths: list[Path] | tuple[Path, ...],
) -> list[QualityIssue]:
    """Verify that a QA report is cryptographically bound to final artifacts."""

    issues: list[QualityIssue] = []
    output = Path(output_path)
    if not output.is_file():
        issues.append(
            QualityIssue(
                code="evidence.output_missing",
                severity="error",
                message="the final PPTX referenced by the report is missing",
            ),
        )
    elif not report.output_sha256 or report.output_sha256 != file_sha256(
        output,
    ):
        issues.append(
            QualityIssue(
                code="evidence.output_hash_mismatch",
                severity="error",
                message="the final PPTX hash does not match the quality report",
            ),
        )
    if output.is_file() and report.rendered_glyph_check_version != GLYPH_DIVIDER_RULESET:
        try:
            needs_glyph_evidence = has_glyph_divider_candidates(output)
        except Exception:
            check_operation()
            needs_glyph_evidence = True
        if needs_glyph_evidence:
            issues.append(QualityIssue(
                code="evidence.glyph_geometry_missing", severity="error",
                message="native foreground dividers lack current rendered-glyph QA evidence",
            ))
    if not report.renderer_id or not report.renderer_version:
        issues.append(
            QualityIssue(
                code="evidence.renderer_missing",
                severity="error",
                message="renderer identity and version are required",
            ),
        )
    if not report.font_manifest:
        issues.append(
            QualityIssue(
                code="evidence.font_manifest_missing",
                severity="error",
                message="the render evidence has no font manifest",
            ),
        )
    if report.rendered_slide_count != len(deck.slides):
        issues.append(
            QualityIssue(
                code="evidence.render_count_mismatch",
                severity="error",
                message="the report render count does not match DeckIR",
            ),
        )
    if len(rendered_paths) != len(deck.slides):
        issues.append(
            QualityIssue(
                code="evidence.render_files_missing",
                severity="error",
                message="the final render set does not contain every slide",
            ),
        )
    if len(report.rendered_slides) != len(deck.slides):
        issues.append(
            QualityIssue(
                code="evidence.slide_manifest_mismatch",
                severity="error",
                message="the per-slide evidence manifest is incomplete",
            ),
        )

    for index, (slide, evidence, image_path) in enumerate(
        zip(
            deck.slides,
            report.rendered_slides,
            rendered_paths,
            strict=False,
        ),
        start=1,
    ):
        if evidence.slide_id != slide.id or evidence.slide_index != index:
            issues.append(
                QualityIssue(
                    code="evidence.slide_identity_mismatch",
                    severity="error",
                    slide_id=slide.id,
                    message="render evidence is bound to the wrong slide",
                ),
            )
            continue
        try:
            digest = file_sha256(image_path)
            with Image.open(image_path) as image:
                width, height = image.size
        except (OSError, ValueError) as exc:
            issues.append(
                QualityIssue(
                    code="evidence.render_unreadable",
                    severity="error",
                    slide_id=slide.id,
                    message=str(exc),
                ),
            )
            continue
        if evidence.sha256 != digest:
            issues.append(
                QualityIssue(
                    code="evidence.render_hash_mismatch",
                    severity="error",
                    slide_id=slide.id,
                    message="render image hash does not match the report",
                ),
            )
        if evidence.width != width or evidence.height != height:
            issues.append(
                QualityIssue(
                    code="evidence.render_dimensions_mismatch",
                    severity="error",
                    slide_id=slide.id,
                    message="render image dimensions do not match the report",
                ),
            )
    issues.extend(inspect_rendered_slides(rendered_paths))
    return issues


def is_verified_delivery(
    deck: DeckIR,
    report: QualityReport,
    output_path: str | Path,
    rendered_paths: list[Path] | tuple[Path, ...],
) -> bool:
    """Return true only for an error-free PPTX bound to its QA evidence."""

    if not (
        report.passed
        and report.delivery_status == "passed"
        and report.render_verified
        and report.content_evidence == EvidenceState.PASSED
        and report.structure_evidence == EvidenceState.PASSED
        and report.render_evidence == EvidenceState.PASSED
        and report.editability_evidence == EvidenceState.PASSED
        and report.provenance_evidence == EvidenceState.PASSED
        and not report.blocking_gaps
        and (
            (
                report.visual_verified
                and report.visual_evidence == EvidenceState.PASSED
            )
            or not report.vision_critique_required
        )
        and report.deck_id == deck.id
        and not delivery_evidence_issues(
            deck,
            report,
            output_path,
            rendered_paths,
        )
    ):
        return False
    rechecked = inspect_pptx(
        deck,
        output_path,
        rendered_slide_count=len(rendered_paths),
        repair_rounds=report.repair_rounds,
    )
    return (
        rechecked.passed
        and rechecked.delivery_status == "passed"
        and rechecked.output_sha256 == report.output_sha256
        and rechecked.render_verified
    )


def inspect_pptx(
    deck: DeckIR,
    path: str | Path,
    *,
    rendered_slide_count: int | None = None,
    repair_rounds: int = 0,
) -> QualityReport:
    source = Path(path)
    issues: list[QualityIssue] = []
    try:
        presentation = Presentation(source)
    except Exception as exc:
        return QualityReport(
            deck_id=deck.id,
            passed=False,
            delivery_status="failed",
            content_complete=False,
            editable=False,
            content_evidence=EvidenceState.FAILED,
            structure_evidence=EvidenceState.FAILED,
            editability_evidence=EvidenceState.FAILED,
            provenance_evidence=EvidenceState.FAILED,
            blocking_gaps=[
                "content evidence",
                "structure evidence",
                "editability evidence",
                "provenance evidence",
            ],
            issues=[
                QualityIssue(
                    code="pptx.reopen_failed",
                    severity="error",
                    message=str(exc),
                ),
            ],
        )
    if len(presentation.slides) != len(deck.slides):
        issues.append(
            QualityIssue(
                code="deck.slide_count",
                severity="error",
                message="compiled slide count differs from DeckIR",
            ),
        )

    editable_names: set[str] = set()
    native_comparison_chart_ids: set[tuple[str, str]] = set()
    comparison_chart_types = {
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    }
    visible_text: list[str] = []
    expected_objects = {
        obj.id: (deck_slide, obj)
        for deck_slide in deck.slides
        for obj in deck_slide.objects
    }
    design_profile = deck.design_context.profile if deck.design_context else deck.design_profile
    caption_minimum = max(
        12.0, design_profile.quality_floor.minimum_caption_font_size if design_profile else 12.0,
    )
    for slide_index, slide in enumerate(presentation.slides):
        slide_id = (
            deck.slides[slide_index].id
            if slide_index < len(deck.slides)
            else None
        )
        expected_on_slide = (
            {obj.id for obj in deck.slides[slide_index].objects}
            if slide_index < len(deck.slides)
            else set()
        )
        native_bounds: list[tuple[str, str, float, float, float, float]] = []
        native_layers: dict[str, tuple[int, BaseShape]] = {}
        for native_index, shape in enumerate(slide.shapes):
            bound_name = shape.name
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP and slide_index < len(deck.slides):
                matches = [obj for obj in deck.slides[slide_index].objects
                           if isinstance(obj, NativeGroupObject)
                           and obj.native_shape_id == shape.shape_id and obj.native_name == shape.name]
                if len(matches) == 1:
                    carrier = matches[0]
                    unchanged = False
                    try:
                        validate_native_carrier(carrier)
                        unchanged = (hashlib.sha256(canonical_native_xml(
                            XML.tostring(shape._element))).hexdigest() == carrier.native_sha256)
                    except ValueError:
                        pass
                    if unchanged:
                        bound_name = carrier.id
                        editable_names.add(carrier.id)
                    else:
                        issues.append(QualityIssue(
                            code="editability.native_group_changed", severity="error",
                            slide_id=slide_id, object_id=carrier.id,
                            message="native group XML differs from the imported read-only carrier",
                        ))
                    if not native_group_context_matches(presentation, slide, carrier):
                        issues.append(QualityIssue(
                            code="editability.native_group_context_changed", severity="error",
                            slide_id=slide_id, object_id=carrier.id,
                            message="native group inherited style context differs from the imported carrier",
                        ))
                # Grouped text is visible native content, not an empty frame.
                visible_text.extend(shape._element.xpath(".//a:t/text()"))
            editable_names.add(shape.name)
            if shape.has_chart and bound_name in expected_on_slide:
                expected_chart = expected_objects[bound_name][1]
                if (
                    isinstance(expected_chart, ChartObject)
                    and shape.chart.chart_type == comparison_chart_types.get(expected_chart.chart_type)
                    and (
                        _has_comparison_labels([
                            category.label for category in shape.chart.plots[0].categories
                        ])
                        or _has_comparison_labels([series.name for series in shape.chart.series])
                    )
                ):
                    native_comparison_chart_ids.add((slide_id, bound_name))
            left = shape.left / 914400
            top = shape.top / 914400
            right = left + shape.width / 914400
            bottom = top + shape.height / 914400
            if (
                left < -0.01
                or top < -0.01
                or right > SLIDE_WIDTH + 0.01
                or bottom > SLIDE_HEIGHT + 0.01
            ):
                issues.append(
                    QualityIssue(
                        code="layout.out_of_bounds",
                        severity="error",
                        slide_id=slide_id,
                        object_id=shape.name,
                        message="shape is outside slide bounds",
                        repairable=True,
                    ),
                )
            if bound_name in expected_on_slide:
                expected_kind = expected_objects[bound_name][1].kind
                native_bounds.append(
                    (bound_name, expected_kind, left, top, right, bottom),
                )
                native_layers[bound_name] = (native_index, shape)
            if getattr(shape, "has_table", False):
                visible_text.extend(
                    text
                    for row in shape.table.rows
                    for cell in row.cells
                    if (text := cell.text.strip())
                )
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    visible_text.append(text)
                    area = max(0.1, (right - left) * (bottom - top))
                    pressure = len(text) / (area * 30.0)
                    if pressure > 1.8:
                        issues.append(
                            QualityIssue(
                                code="layout.text_overflow_risk",
                                severity="error",
                                slide_id=slide_id,
                                object_id=shape.name,
                                message="text density is likely to overflow its frame",
                                repairable=True,
                            ),
                        )
                    elif pressure > 1.25:
                        issues.append(
                            QualityIssue(
                                code="layout.text_density",
                                severity="warning",
                                slide_id=slide_id,
                                object_id=shape.name,
                                message="text density is high",
                            ),
                        )
                expected = expected_objects.get(shape.name)
                if expected and expected[1].kind == "text":
                    expected_slide, expected_text = expected
                    paragraph_sizes = NativeFontResolver(shape).visible_sizes_by_paragraph()
                    sizes = [size for paragraph in paragraph_sizes for size in paragraph]
                    if (
                        text
                        and pressure <= 1.8
                        and _native_text_overflows_fixed_frame(shape, text, sizes)
                    ):
                        issues.append(
                            QualityIssue(
                                code="layout.text_overflow_risk",
                                severity="error",
                                slide_id=slide_id,
                                object_id=shape.name,
                                message="effective native font requires more lines than the fixed text frame can hold",
                                repairable=True,
                            ),
                        )
                    minimum = (
                        42.0
                        if expected_text.role == "title"
                        and expected_slide.layout_id == "cover"
                        else 32.0
                        if expected_text.role == "title"
                        else 17.0
                    )
                    caption_paragraphs = (
                        expected_text.native_metric_text.caption_paragraphs(
                            [paragraph.text for paragraph in shape.text_frame.paragraphs],
                        )
                        if expected_text.role == "metric" and expected_text.native_metric_text
                        and expected_text.text == text
                        else frozenset()
                    )
                    size_requirements = (
                        (size, caption_minimum if index in caption_paragraphs else minimum)
                        for index, paragraph in enumerate(paragraph_sizes)
                        for size in paragraph
                    )
                    failed_minimum = next((
                        floor for size, floor in size_requirements
                        if size is None or size < floor
                    ), None)
                    if not sizes or failed_minimum is not None:
                        issues.append(
                            QualityIssue(
                                code="design.font_too_small",
                                severity="error",
                                slide_id=slide_id,
                                object_id=shape.name,
                                message=(
                                    f"text is below the {failed_minimum if failed_minimum is not None else minimum:g} pt "
                                    "presentation readability floor"
                                ),
                                repairable=True,
                            ),
                        )
        for index, first in enumerate(native_bounds):
            for second in native_bounds[index + 1 :]:
                # A full-bleed or documentary image is commonly the intentional
                # visual ground for native text. It must remain editable and is
                # therefore not an accidental object-on-object collision.
                if "image" in {first[1], second[1]}:
                    continue
                # Relationship connectors are expected to cross or terminate
                # inside native diagram nodes.  Treating their thin bounding
                # boxes as filled objects creates false collision failures.
                first_expected = expected_objects.get(first[0])
                second_expected = expected_objects.get(second[0])
                if any(
                    isinstance(expected[1], ShapeObject)
                    and expected[1].shape_type == "line"
                    for expected in (first_expected, second_expected)
                    if expected is not None
                ):
                    continue
                first_area = (first[4] - first[2]) * (first[5] - first[3])
                second_area = (second[4] - second[2]) * (second[5] - second[3])

                def intentionally_contained(
                    outer: tuple[str, str, float, float, float, float],
                    inner: tuple[str, str, float, float, float, float],
                    *,
                    outer_area: float,
                    inner_area: float,
                ) -> bool:
                    """Recognize an editable card/background containing its content.

                    Native presentations intentionally layer text, icons and diagrams
                    inside editable shape containers.  That is not an accidental
                    collision.  Near-identical duplicate shapes remain errors.
                    """

                    if outer[1] != "shape" or outer_area <= 0 or inner_area <= 0:
                        return False
                    tolerance = 0.03
                    contained = (
                        inner[2] >= outer[2] - tolerance
                        and inner[3] >= outer[3] - tolerance
                        and inner[4] <= outer[4] + tolerance
                        and inner[5] <= outer[5] + tolerance
                    )
                    if not contained:
                        return False
                    outer_index, outer_shape = native_layers[outer[0]]
                    inner_index, inner_shape = native_layers[inner[0]]
                    if (
                        outer_shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                        and outer_shape.auto_shape_type in {
                            MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
                        }
                        and outer_shape.rotation == inner_shape.rotation == 0
                        and outer_shape.fill.type == MSO_FILL_TYPE.SOLID
                    ):
                        if outer_index > inner_index:
                            # An opaque foreground card is not a background:
                            # charts, tables and native shapes can be hidden
                            # just as completely as a transparent text label.
                            alpha = outer_shape._element.xpath(
                                "./p:spPr/a:solidFill/*/a:alpha/@val",
                            )
                            alpha_adjustments = outer_shape._element.xpath(
                                "./p:spPr/a:solidFill/*/a:alphaMod | "
                                "./p:spPr/a:solidFill/*/a:alphaOff",
                            )
                            if not alpha_adjustments and all(
                                value == "100000" for value in alpha
                            ):
                                return False
                        if (
                            inner[1] == "text"
                            and inner_shape.has_text_frame
                            and inner_shape.text.strip()
                        ):
                            # A transparent native label on a blank, earlier
                            # rectangular fill is intentional even when its
                            # editable text frame nearly fills the background.
                            if (
                                outer_index < inner_index
                                and not outer_shape.text.strip()
                                and inner_shape.fill.type == MSO_FILL_TYPE.BACKGROUND
                            ):
                                return True
                    if inner_area / outer_area >= 0.9:
                        return False
                    return True

                if intentionally_contained(
                    first,
                    second,
                    outer_area=first_area,
                    inner_area=second_area,
                ) or intentionally_contained(
                    second,
                    first,
                    outer_area=second_area,
                    inner_area=first_area,
                ):
                    continue
                overlap_width = max(
                    0.0,
                    min(first[4], second[4]) - max(first[2], second[2]),
                )
                overlap_height = max(
                    0.0,
                    min(first[5], second[5]) - max(first[3], second[3]),
                )
                overlap = overlap_width * overlap_height
                smaller = min(first_area, second_area)
                if smaller > 0 and overlap / smaller > 0.9:
                    issues.append(
                        QualityIssue(
                            code="layout.complete_overlap",
                            severity="error",
                            slide_id=slide_id,
                            object_id=second[0],
                            message=f"object substantially overlaps {first[0]}",
                            repairable=True,
                        ),
                    )

    for slide in deck.slides:
        for violation in check_capacity(slide):
            issues.append(
                QualityIssue(
                    code="layout.capacity",
                    severity="error",
                    slide_id=slide.id,
                    message=violation.message,
                    repairable=True,
                ),
            )

    issues.extend(_design_issues(deck, native_comparison_chart_ids=native_comparison_chart_ids))
    issues.extend(_chart_semantic_issues(deck))

    expected_ids = {obj.id for slide in deck.slides for obj in slide.objects}
    missing_ids = expected_ids - editable_names
    for object_id in sorted(missing_ids):
        issues.append(
            QualityIssue(
                code="editability.object_missing",
                severity="error",
                object_id=object_id,
                message="DeckIR object is not represented as a native named shape",
            ),
        )
    joined_text = "\n".join(visible_text)
    normalized_joined_text = _normalized_text(joined_text)
    if deck.source_refs and not deck.required_facts:
        issues.append(
            QualityIssue(
                code="content.source_unverified",
                severity="error",
                message="source-backed deck has no required facts to verify",
            ),
        )
    for fact in deck.required_facts:
        if _normalized_text(fact) not in normalized_joined_text:
            issues.append(
                QualityIssue(
                    code="content.required_fact_missing",
                    severity="error",
                    message=f"required source fact missing: {fact}",
                ),
            )
    for immutable in deck.immutable_data:
        if _normalized_text(immutable) not in normalized_joined_text:
            issues.append(
                QualityIssue(
                    code="content.immutable_missing",
                    severity="error",
                    message=f"immutable data missing: {immutable}",
                ),
            )
    provenance_errors: list[QualityIssue] = []
    for material_coverage in deck.source_coverage:
        if material_coverage.status == "converted_doc_unverified":
            provenance_errors.append(
                QualityIssue(
                    code="material.conversion_unverified",
                    severity="error",
                    message=(
                        f"legacy DOC {material_coverage.source_name} conversion is not "
                        "verified against the original; formal delivery is blocked"
                    ),
                )
            )
    source_refs = set(deck.source_refs)
    fact_ids = {fact.fact_id for fact in deck.facts}
    if source_refs and not deck.facts:
        provenance_errors.append(
            QualityIssue(
                code="provenance.fact_catalog_missing",
                severity="error",
                message="source-backed deck has no fact provenance catalog",
            ),
        )
    for fact in deck.facts:
        if not fact.source_refs or set(fact.source_refs) - source_refs:
            provenance_errors.append(
                QualityIssue(
                    code="provenance.fact_source_invalid",
                    severity="error",
                    message=f"fact {fact.fact_id} is not bound to an allowed source",
                ),
            )
    for slide in deck.slides:
        if source_refs and slide.provenance_layout not in {"cover", "closing"}:
            if not slide.source_fact_ids:
                provenance_errors.append(
                    QualityIssue(
                        code="provenance.slide_fact_missing",
                        severity="error",
                        slide_id=slide.id,
                        message="source-backed slide has no fact provenance binding",
                    ),
                )
            elif set(slide.source_fact_ids) - fact_ids:
                provenance_errors.append(
                    QualityIssue(
                        code="provenance.slide_fact_invalid",
                        severity="error",
                        slide_id=slide.id,
                        message="slide references an unknown fact",
                    ),
                )
    for asset in deck.assets:
        asset_path = Path(asset.local_path)
        if (
            asset.rights_status == "review"
            or asset.ingestion_status not in {"ready", "recovered"}
            or not asset_path.is_file()
            or file_sha256(asset_path) != asset.sha256
        ):
            provenance_errors.append(
                QualityIssue(
                    code="provenance.asset_invalid",
                    severity="error",
                    message=(
                        f"asset {asset.asset_id} lacks verified local rights, "
                        "ingestion state, or content hash"
                    ),
                ),
            )
    issues.extend(provenance_errors)

    actual_rendered = rendered_slide_count or 0
    if rendered_slide_count is not None and actual_rendered != len(
        deck.slides,
    ):
        issues.append(
            QualityIssue(
                code="render.slide_count",
                severity="error",
                message="rendered slide count differs from DeckIR",
                repairable=True,
            ),
        )
    passed = not any(issue.severity == "error" for issue in issues)
    content_complete = not any(
        issue.severity == "error" and issue.code.startswith("content.")
        for issue in issues
    )
    delivery_status = delivery_status_for(
        issues,
        content_complete=content_complete,
        editable=not missing_ids,
    )
    content_evidence = (
        EvidenceState.FAILED
        if any(
            issue.severity == "error" and issue.code.startswith("content.")
            for issue in issues
        )
        else EvidenceState.PASSED
    )
    structure_evidence = (
        EvidenceState.FAILED
        if any(
            issue.severity == "error"
            and issue.code.startswith(
                ("pptx.", "deck.", "layout.", "design.", "chart."),
            )
            for issue in issues
        )
        else EvidenceState.PASSED
    )
    editability_evidence = (
        EvidenceState.FAILED if missing_ids else EvidenceState.PASSED
    )
    provenance_evidence = (
        EvidenceState.FAILED if provenance_errors else EvidenceState.PASSED
    )
    render_evidence = (
        EvidenceState.NOT_RUN
        if rendered_slide_count is None
        else EvidenceState.PASSED
        if actual_rendered == len(deck.slides)
        else EvidenceState.FAILED
    )
    blocking_gaps = [
        label
        for label, state in (
            ("content evidence", content_evidence),
            ("structure evidence", structure_evidence),
            ("editability evidence", editability_evidence),
            ("provenance evidence", provenance_evidence),
        )
        if state != EvidenceState.PASSED
    ]
    if render_evidence == EvidenceState.FAILED:
        blocking_gaps.append("render evidence")
    if not passed or blocking_gaps:
        delivery_status = "failed"
    return QualityReport(
        deck_id=deck.id,
        passed=passed,
        delivery_status=delivery_status,
        content_complete=content_complete,
        editable=not missing_ids,
        output_sha256=file_sha256(source),
        rendered_slide_count=actual_rendered,
        render_verified=(
            rendered_slide_count is not None
            and rendered_slide_count == len(deck.slides)
        ),
        # Opening and rendering are deterministic evidence checks.  They do
        # not establish hierarchy, composition, chart semantics, or beauty.
        visual_verified=False,
        content_evidence=content_evidence,
        structure_evidence=structure_evidence,
        render_evidence=render_evidence,
        visual_evidence=EvidenceState.NOT_RUN,
        editability_evidence=editability_evidence,
        provenance_evidence=provenance_evidence,
        blocking_gaps=blocking_gaps,
        issues=issues,
        repair_rounds=repair_rounds,
    )
