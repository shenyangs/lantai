"""Resolve native OOXML typography without semantic guesses or size clamping."""

from pptx.oxml.ns import qn


def _size(properties) -> float | None:
    if properties is None or properties.get("sz") is None:
        return None
    value = int(properties.get("sz")) / 100
    return value if value > 0 else None


def _placeholder_type(shape) -> str:
    placeholder = shape._element.xpath("./p:nvSpPr/p:nvPr/p:ph")
    return placeholder[0].get("type", "obj") if placeholder else "obj"


def _master_type(kind: str) -> str:
    if kind in {"title", "ctrTitle"}:
        return "title"
    if kind in {"body", "subTitle", "obj", "chart", "tbl", "dgm", "pic", "clipArt", "media"}:
        return "body"
    return kind


def _style_size(style, level: int) -> float | None:
    if style is None:
        return None
    if style.tag == qn("p:txBody"):
        for paragraph in style.findall(qn("a:p")):
            properties = paragraph.find(qn("a:pPr"))
            if properties is not None and int(properties.get("lvl", "0")) == level:
                size = _size(properties.find(qn("a:defRPr")))
                if size is not None:
                    return size
        style = style.find(qn("a:lstStyle"))
        if style is None:
            return None
    for name in (f"a:lvl{level + 1}pPr", "a:defPPr"):
        rule = style.find(qn(name))
        if rule is not None and (size := _size(rule.find(qn("a:defRPr")))) is not None:
            return size
    return None


def _style_run_properties(style, level: int):
    """Yield nearest-first rules, resolving each character property separately."""
    if style is None:
        return
    if style.tag == qn("p:txBody"):
        for paragraph in style.findall(qn("a:p")):
            properties = paragraph.find(qn("a:pPr"))
            if properties is not None and int(properties.get("lvl", "0")) == level:
                rule = properties.find(qn("a:defRPr"))
                if rule is not None:
                    yield rule
        style = style.find(qn("a:lstStyle"))
        if style is None:
            return
    for name in (f"a:lvl{level + 1}pPr", "a:defPPr"):
        properties = style.find(qn(name))
        if properties is not None:
            rule = properties.find(qn("a:defRPr"))
            if rule is not None:
                yield rule


class NativeFontResolver:
    """Shared native size authority plus independently inherited run properties."""

    def __init__(self, shape):
        self.shape = shape
        # Table cells have the same DrawingML text body but are not shapes.
        root = getattr(shape, "_element", None)
        if root is None:
            root = shape._tc
        body = root.find(qn("p:txBody"))
        if body is None:
            body = root.find(qn("a:txBody"))
        fit = body.xpath("./a:bodyPr/a:normAutofit") if body is not None else []
        self.scale = int(fit[0].get("fontScale", "100000")) / 100000 if fit else 1.0
        if self.scale <= 0:
            raise ValueError("invalid native font scale")
        self.styles = body.xpath("./a:lstStyle") if body is not None else []
        if getattr(shape, "is_placeholder", False):
            layout = shape.part.slide_layout.placeholders.get(idx=shape.placeholder_format.idx)
            master = shape.part.slide_layout.slide_master
            kind = _master_type(_placeholder_type(layout or shape))
            if layout is not None:
                self.styles.extend(layout._element.xpath("./p:txBody"))
            placeholders = [candidate for candidate in master.placeholders
                            if _placeholder_type(candidate) == kind]
            if len(placeholders) > 1:
                raise ValueError("ambiguous master placeholder font inheritance")
            if placeholders:
                self.styles.extend(placeholders[0]._element.xpath("./p:txBody"))
            group = "titleStyle" if kind == "title" else "bodyStyle" if kind == "body" else "otherStyle"
            self.styles.extend(master._element.xpath(f"./p:txStyles/p:{group}"))
        defaults = shape.part.package.presentation_part._element.find(qn("p:defaultTextStyle"))
        if defaults is not None:
            self.styles.append(defaults)

    def paragraph_size(self, paragraph) -> float | None:
        properties = paragraph._p.find(qn("a:pPr"))
        if properties is not None:
            size = _size(properties.find(qn("a:defRPr")))
            if size is not None:
                return size
        level = int(properties.get("lvl", "0")) if properties is not None else 0
        for style in self.styles:
            if (size := _style_size(style, level)) is not None:
                return size
        return None

    def run_size(self, paragraph, run) -> float | None:
        size = _size(run._r.find(qn("a:rPr")))
        size = size if size is not None else self.paragraph_size(paragraph)
        return size * self.scale if size is not None else None

    def _run_properties(self, paragraph, run):
        explicit = run._r.find(qn("a:rPr"))
        if explicit is not None:
            yield explicit
        properties = paragraph._p.find(qn("a:pPr"))
        if properties is not None:
            default = properties.find(qn("a:defRPr"))
            if default is not None:
                yield default
        level = int(properties.get("lvl", "0")) if properties is not None else 0
        for style in self.styles:
            yield from _style_run_properties(style, level)

    def run_flag(self, paragraph, run, flag: str) -> bool:
        if flag not in {"b", "i"}:
            raise ValueError("unsupported native character flag")
        for properties in self._run_properties(paragraph, run):
            value = properties.get(flag)
            if value is not None:
                if value not in {"0", "1", "true", "false"}:
                    raise ValueError("invalid native character flag")
                return value in {"1", "true"}
        return False

    def run_rgb(self, paragraph, run) -> str | None:
        for properties in self._run_properties(paragraph, run):
            fill = properties.find(qn("a:solidFill"))
            if fill is not None:
                color = fill.find(qn("a:srgbClr"))
                return color.get("val") if color is not None else None
        return None

    def base_size(self) -> float | None:
        for paragraph in self.shape.text_frame.paragraphs:
            ending = _size(paragraph._p.find(qn("a:endParaRPr")))
            if ending is not None:
                return ending * self.scale
            size = self.paragraph_size(paragraph)
            if size is not None:
                return size * self.scale
            for run in paragraph.runs:
                if (size := self.run_size(paragraph, run)) is not None:
                    return size
        return None

    def visible_sizes_by_paragraph(self) -> list[list[float | None]]:
        """Resolve every visible run/field without losing its paragraph index."""
        paragraphs = []
        for paragraph in self.shape.text_frame.paragraphs:
            sizes = []
            for run in paragraph._p:
                if run.tag not in {qn("a:r"), qn("a:fld")}:
                    continue
                text = run.find(qn("a:t"))
                if text is not None and (text.text or "").strip():
                    size = _size(run.find(qn("a:rPr")))
                    size = size if size is not None else self.paragraph_size(paragraph)
                    sizes.append(size * self.scale if size is not None else None)
            paragraphs.append(sizes)
        return paragraphs

    def visible_sizes(self) -> list[float | None]:
        return [size for paragraph in self.visible_sizes_by_paragraph() for size in paragraph]
