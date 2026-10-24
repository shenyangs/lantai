# -*- coding: utf-8 -*-
"""Object-level PPTX measurement for reference learning.

Only normalized measurements and aggregate style signals leave this module.
The reference package itself is never copied into the product tree.
"""

from __future__ import annotations

from collections import Counter
from hashlib import sha256
from pathlib import Path
from statistics import median
from typing import Any, Iterable, Literal
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from .models import (
    ReferenceDeckMeasurement,
    ReferenceObjectMeasurement,
    ReferenceSlideMeasurement,
)


def _inches(value: int | Emu | None) -> float:
    if value is None:
        return 0.0
    return round(float(value) / 914400.0, 4)


def _hex_color(color: Any) -> str | None:
    try:
        rgb = color.rgb
    except (AttributeError, ValueError, TypeError):
        return None
    return str(rgb).upper() if rgb is not None else None


def _unique(values: Iterable[str | None]) -> list[str]:
    return list(dict.fromkeys(value for value in values if value))


def _alpha_from_xml(element: Any) -> float | None:
    """Read DrawingML alpha as transparency (0 opaque, 1 transparent)."""

    if element is None:
        return None
    alpha = element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}alpha")
    if alpha is None:
        return None
    try:
        opacity = int(alpha.get("val")) / 100000.0
    except (TypeError, ValueError):
        return None
    return round(max(0.0, min(1.0, 1.0 - opacity)), 4)


def _has_drawingml(shape: Any, local_name: str) -> bool:
    element = getattr(shape, "_element", None)
    if element is None:
        return False
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    return element.find(f".//{{{namespace}}}{local_name}") is not None


def _color_from_fill(fill: Any) -> str | None:
    try:
        if fill.type is None:
            return None
        return _hex_color(fill.fore_color)
    except (AttributeError, ValueError, TypeError):
        return None


def _shape_kind(shape: Any) -> str:
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if getattr(shape, "has_text_frame", False):
        if shape.is_placeholder:
            placeholder = str(shape.placeholder_format.type).casefold()
            if "title" in placeholder:
                return "title"
        return "text"
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    return "shape"


def _role_hint(shape: Any, kind: str) -> str:
    if kind == "title":
        return "title"
    if shape.is_placeholder:
        placeholder = str(shape.placeholder_format.type).casefold()
        if "footer" in placeholder:
            return "footer"
        if "slide_number" in placeholder or "slide number" in placeholder:
            return "page_number"
        if "date" in placeholder:
            return "date"
    name = str(getattr(shape, "name", "")).casefold()
    if "logo" in name:
        return "logo"
    if kind == "image":
        return "image"
    if kind in {"chart", "table"}:
        return "evidence"
    return "content"


def _text_measurements(shape: Any) -> dict[str, Any]:
    if not getattr(shape, "has_text_frame", False):
        return {
            "text_length": 0,
            "line_count": 0,
            "font_families": [],
            "font_sizes": [],
            "text_colors": [],
            "alignments": [],
            "margins": {},
        }
    frame = shape.text_frame
    fonts: list[str] = []
    sizes: list[float] = []
    colors: list[str] = []
    alignments: list[str] = []
    texts: list[str] = []
    for paragraph in frame.paragraphs:
        if paragraph.alignment is not None:
            alignments.append(str(paragraph.alignment).split(" (")[0].casefold())
        for run in paragraph.runs:
            texts.append(run.text)
            if run.font.name:
                fonts.append(str(run.font.name))
            if run.font.size:
                sizes.append(round(float(run.font.size.pt), 2))
            color = _hex_color(run.font.color)
            if color:
                colors.append(color)
    text = "\n".join(paragraph.text for paragraph in frame.paragraphs)
    return {
        "text_length": len(text.strip()),
        "line_count": len([line for line in text.splitlines() if line.strip()]),
        "font_families": _unique(fonts),
        "font_sizes": sorted(set(sizes)),
        "text_colors": _unique(colors),
        "alignments": _unique(alignments),
        "margins": {
            "left": _inches(frame.margin_left),
            "right": _inches(frame.margin_right),
            "top": _inches(frame.margin_top),
            "bottom": _inches(frame.margin_bottom),
        },
    }


def _measure_shape(
    shape: Any,
    slide_number: int,
    z_order: int,
    *,
    source_layer: Literal["slide", "layout", "master"] = "slide",
) -> ReferenceObjectMeasurement:
    kind = _shape_kind(shape)
    text = _text_measurements(shape)
    fill_color = _color_from_fill(getattr(shape, "fill", None))
    line_color = _hex_color(getattr(getattr(shape, "line", None), "color", None))
    image_aspect: float | None = None
    displayed_aspect: float | None = None
    image_crop: dict[str, float] = {}
    if kind == "image":
        try:
            image_width, image_height = shape.image.size
            image_aspect = round(image_width / image_height, 4) if image_height else None
            displayed_aspect = round(float(shape.width) / float(shape.height), 4) if shape.height else None
            image_crop = {
                "left": round(float(shape.crop_left), 4),
                "right": round(float(shape.crop_right), 4),
                "top": round(float(shape.crop_top), 4),
                "bottom": round(float(shape.crop_bottom), 4),
            }
        except (AttributeError, TypeError, ValueError, ZeroDivisionError):
            pass
    chart_type = ""
    if kind == "chart":
        try:
            chart_type = str(shape.chart.chart_type).split(" (")[0].casefold()
        except (AttributeError, ValueError):
            chart_type = "unknown"
    table_shape = None
    if kind == "table":
        try:
            table_shape = (len(shape.table.rows), len(shape.table.columns))
        except (AttributeError, TypeError):
            table_shape = None
    placeholder_type = ""
    if shape.is_placeholder:
        placeholder_type = str(shape.placeholder_format.type).split(" (")[0].casefold()
    layer_prefix = "" if source_layer == "slide" else f"{source_layer}-"
    return ReferenceObjectMeasurement(
        object_id=f"slide-{slide_number}-{layer_prefix}object-{z_order + 1}",
        slide_number=slide_number,
        z_order=z_order,
        source_layer=source_layer,
        name=str(getattr(shape, "name", "")),
        kind=kind,
        role_hint=_role_hint(shape, kind),
        x=_inches(shape.left),
        y=_inches(shape.top),
        width=_inches(shape.width),
        height=_inches(shape.height),
        rotation=round(float(getattr(shape, "rotation", 0) or 0), 2),
        fill_colors=[fill_color] if fill_color else [],
        line_colors=[line_color] if line_color else [],
        fill_transparency=_alpha_from_xml(getattr(getattr(shape, "fill", None), "_xPr", None)),
        line_transparency=_alpha_from_xml(getattr(getattr(shape, "line", None), "_get_or_add_ln", lambda: None)()),
        has_shadow=_has_drawingml(shape, "outerShdw") or _has_drawingml(shape, "innerShdw"),
        has_gradient=_has_drawingml(shape, "gradFill"),
        has_bevel=_has_drawingml(shape, "bevelT") or _has_drawingml(shape, "bevelB"),
        is_placeholder=bool(shape.is_placeholder),
        placeholder_type=placeholder_type,
        image_aspect_ratio=image_aspect,
        displayed_aspect_ratio=displayed_aspect,
        image_crop=image_crop,
        chart_type=chart_type,
        table_shape=table_shape,
        group_children=len(shape.shapes) if kind == "group" else 0,
        **text,
    )


def _shows_inherited_shapes(owner: Any) -> bool:
    return str(owner._element.get("showMasterSp", "true")).casefold() not in {"0", "false"}


def _hidden_shape(shape: Any) -> bool:
    namespace = "http://schemas.openxmlformats.org/presentationml/2006/main"
    # Inspect this shape's own non-visual properties, not hidden children of
    # an otherwise visible group.
    return any(
        str(item.get("hidden", "false")).casefold() in {"1", "true"}
        for item in shape._element.findall(f"./*/{{{namespace}}}cNvPr")
    )


def _inherited_objects(slide: Any, slide_number: int) -> list[ReferenceObjectMeasurement]:
    """Measure visible identity on this slide's active layout/master chain.

    Placeholder samples in a layout or master are style definitions, not
    independent rendered objects. Page-authored counts remain unchanged.
    """

    if not _shows_inherited_shapes(slide):
        return []
    layout = slide.slide_layout
    owners: list[tuple[Literal["layout", "master"], Any]] = []
    if _shows_inherited_shapes(layout):
        owners.append(("master", layout.slide_master))
    owners.append(("layout", layout))
    return [
        _measure_shape(shape, slide_number, z_order, source_layer=layer)
        for layer, owner in owners
        for z_order, shape in enumerate(owner.shapes)
        if not shape.is_placeholder and not _hidden_shape(shape)
    ]


def _clusters(values: Iterable[float], tolerance: float = 0.08) -> list[float]:
    clusters: list[list[float]] = []
    for value in sorted(values):
        if not clusters or abs(value - median(clusters[-1])) > tolerance:
            clusters.append([value])
        else:
            clusters[-1].append(value)
    return [round(float(median(cluster)), 3) for cluster in clusters]


def _slide_title(slide: Any, objects: list[ReferenceObjectMeasurement]) -> str:
    try:
        if slide.shapes.title and slide.shapes.title.text.strip():
            return slide.shapes.title.text.strip()[:160]
    except AttributeError:
        pass
    candidates = [
        (max(item.font_sizes, default=0), -item.y, shape.text.strip())
        for shape, item in zip(slide.shapes, objects)
        if getattr(shape, "has_text_frame", False) and shape.text.strip()
    ]
    return max(candidates, default=(0, 0, ""))[2][:160]


def _background_colors(slide: Any) -> list[str]:
    colors: list[str] = []
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        try:
            value = _color_from_fill(owner.background.fill)
        except (AttributeError, ValueError):
            value = None
        if value:
            colors.append(value)
    return _unique(colors)


def _theme_metadata(path: Path) -> tuple[dict[str, str], list[str]]:
    fonts: dict[str, str] = {}
    colors: list[str] = []
    try:
        with ZipFile(path) as archive:
            names = [name for name in archive.namelist() if name.startswith("ppt/theme/theme") and name.endswith(".xml")]
            if not names:
                return fonts, colors
            from xml.etree import ElementTree as ET

            root = ET.fromstring(archive.read(names[0]))
            namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
            for key, xpath in {
                "major_latin": ".//a:majorFont/a:latin",
                "major_east_asian": ".//a:majorFont/a:ea",
                "minor_latin": ".//a:minorFont/a:latin",
                "minor_east_asian": ".//a:minorFont/a:ea",
            }.items():
                element = root.find(xpath, namespace)
                if element is not None and element.get("typeface"):
                    fonts[key] = str(element.get("typeface"))
            for element in root.findall(".//a:clrScheme/*", namespace):
                child = next(iter(element), None)
                if child is not None:
                    value = child.get("val") or child.get("lastClr")
                    if value and len(value) == 6:
                        colors.append(value.upper())
    except (OSError, KeyError, ValueError):
        pass
    return fonts, _unique(colors)


class PptxObjectParser:
    """Parse a PPTX into normalized object measurements."""

    def parse(self, path: str | Path) -> ReferenceDeckMeasurement:
        source = Path(path).expanduser().resolve()
        if source.suffix.casefold() != ".pptx":
            raise ValueError("PptxObjectParser accepts .pptx files only")
        if not source.is_file():
            raise FileNotFoundError(source)
        presentation = Presentation(str(source))
        width = _inches(presentation.slide_width)
        height = _inches(presentation.slide_height)
        slide_records: list[ReferenceSlideMeasurement] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            objects = [
                _measure_shape(shape, slide_number, z_order)
                for z_order, shape in enumerate(slide.shapes)
            ]
            occupied = min(
                1.0,
                sum(item.width * item.height for item in objects) / max(width * height, 0.001),
            )
            text_chars = sum(item.text_length for item in objects)
            density_score = occupied * 0.65 + min(1.0, text_chars / 900) * 0.35
            density = "high" if density_score >= 0.58 else "low" if density_score <= 0.28 else "medium"
            slide_records.append(
                ReferenceSlideMeasurement(
                    slide_number=slide_number,
                    layout_name=str(getattr(slide.slide_layout, "name", "")),
                    master_name=str(getattr(slide.slide_layout.slide_master, "name", "")),
                    background_colors=_background_colors(slide),
                    object_count=len(objects),
                    text_object_count=sum(item.kind in {"text", "title"} for item in objects),
                    image_count=sum(item.kind == "image" for item in objects),
                    chart_count=sum(item.kind == "chart" for item in objects),
                    table_count=sum(item.kind == "table" for item in objects),
                    shape_count=sum(item.kind in {"shape", "line", "group"} for item in objects),
                    text_characters=text_chars,
                    title_text=_slide_title(slide, objects),
                    left_edge_clusters=_clusters(item.x for item in objects if item.width > 0),
                    top_edge_clusters=_clusters(item.y for item in objects if item.height > 0),
                    estimated_whitespace=round(max(0.0, 1.0 - occupied), 4),
                    density_band=density,
                    objects=objects,
                    inherited_objects=_inherited_objects(slide, slide_number),
                ),
            )
        theme_fonts, theme_colors = _theme_metadata(source)
        masters = _unique(str(getattr(master, "name", "")) for master in presentation.slide_masters)
        layouts = _unique(
            str(getattr(layout, "name", ""))
            for master in presentation.slide_masters
            for layout in master.slide_layouts
        )
        return ReferenceDeckMeasurement(
            source_name=source.name,
            source_format="pptx",
            source_sha256=sha256(source.read_bytes()).hexdigest(),
            slide_width=width,
            slide_height=height,
            slide_count=len(slide_records),
            aspect_ratio=round(width / height, 4),
            master_names=masters,
            layout_names=layouts,
            theme_fonts=theme_fonts,
            theme_colors=theme_colors,
            slides=slide_records,
        )
