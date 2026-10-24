"""Restore bounded native axis properties unsupported by PptxGenJS's API."""

import os
from pathlib import Path
import tempfile
import zipfile

from lxml import etree as XML
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from .models import ChartObject, DeckIR


def restore_chart_axes(deck: DeckIR, output: Path) -> None:
    """Modify only bound chart parts in the compiler's private staging file.

    Using catLabelFormatCode in PptxGenJS converts a category axis into a date
    axis. Preserve the original number format without changing the axis type.
    Its public API also lacks minor units and explicit axis positions.
    """
    if not any(isinstance(obj, ChartObject) and (obj.category_axis or obj.value_axis or obj.plot_layout)
               for slide in deck.slides for obj in slide.objects):
        return
    presentation = Presentation(output)
    if len(presentation.slides) != len(deck.slides):
        raise RuntimeError("native chart axis slide binding is inconsistent")
    replacements = {}
    for authored, native in zip(deck.slides, presentation.slides):
        for obj in authored.objects:
            if not isinstance(obj, ChartObject) or not (obj.category_axis or obj.value_axis or obj.plot_layout):
                continue
            matches = [shape.chart for shape in native.shapes if shape.has_chart and shape.name == obj.id]
            if len(matches) != 1:
                raise RuntimeError("native chart axes have no unique authored binding")
            chart = matches[0]
            part = str(chart.part.partname).lstrip("/")
            if part in replacements:
                raise RuntimeError("native chart axes unexpectedly share a chart part")
            if obj.plot_layout is not None:
                layout = chart.part._element.find("c:chart/c:plotArea/c:layout",
                                                 namespaces={"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"})
                if layout is None:
                    raise RuntimeError("native chart has no plot layout binding")
                for old in list(layout):
                    if old.tag == qn("c:manualLayout"):
                        layout.remove(old)
                manual = OxmlElement("c:manualLayout")
                for name, value in obj.plot_layout.model_dump(exclude_none=True).items():
                    node = OxmlElement("c:" + name)
                    node.set("val", str(value))
                    manual.append(node)
                layout.insert(0, manual)
            for name in ("category_axis", "value_axis"):
                style = getattr(obj, name)
                if style is None:
                    continue
                axis = getattr(chart, name)
                axis.tick_labels.number_format = style.number_format
                if name == "value_axis":
                    axis.minor_unit = style.minor_unit
                if style.position is not None:
                    position = axis._element.find(qn("c:axPos"))
                    if position is None:
                        raise RuntimeError("native chart axis has no position")
                    position.set("val", style.position)
            replacements[part] = XML.tostring(chart.part._element, encoding="utf-8", xml_declaration=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{output.stem}-axes-", suffix=".pptx", dir=output.parent)
    os.close(fd)
    staged = Path(temporary)
    try:
        with zipfile.ZipFile(output) as source, zipfile.ZipFile(staged, "w") as target:
            for entry in source.infolist():
                target.writestr(entry, replacements.get(entry.filename, source.read(entry.filename)))
        os.replace(staged, output)
    finally:
        staged.unlink(missing_ok=True)
