"""Version-bound report exports and byte receipts shared by UI and chat tools."""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import json
from pathlib import Path
import re
from uuid import uuid4

from .citations import format_citations, STYLES
from .evidence import _digest, _within
from .library import LibraryService
from .store import canonical, RevisionConflict
from .figures import FIGURE_LINE, FigureService, figure_lines, portable_markdown

FORMATS = {'markdown': ('md', 'text/markdown; charset=utf-8'),
           'docx': ('docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'),
           'pdf': ('pdf', 'application/pdf')}
STATUS_LABELS = {'reviewed': '核验完成', 'not_reviewed': '草稿，待核验', 'needs_revision': '需要修订',
                 'source_changed': '资料已变化，需要重新核验', 'metadata_changed': '引用元数据已变化，需要重新核验'}

FONT_NAME = 'Lantai Research Sans'
FONT_FILE = 'LantaiResearchSans-Regular.ttf'


def _font_data():
    directory = Path(__file__).with_name('data')
    data = (directory / FONT_FILE).read_bytes()
    manifest = json.loads((directory / 'manifest.json').read_text(encoding='utf-8'))
    if _digest(data) != manifest['files'][FONT_FILE]['sha256']:
        raise ValueError('Report font integrity mismatch')
    return data


def _report_font(text):
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    if FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(FONT_NAME, BytesIO(_font_data())))
    font = pdfmetrics.getFont(FONT_NAME)
    missing = sorted({ord(c) for c in text if not c.isspace() and ord(c) not in font.face.charToGlyph})
    if missing:
        raise ValueError('Report font does not cover characters: ' + ', '.join(f'U+{c:04X}' for c in missing[:12]))
    return font


def _embed_docx_font(document):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.xmlchemy import serialize_for_reading
    from docx.oxml.ns import qn
    key = uuid4()
    data = bytearray(_font_data())
    mask = key.bytes[::-1]
    for index in range(32):
        data[index] ^= mask[index % 16]
    table = document.part.part_related_by(RT.FONT_TABLE)
    part = Part(PackURI('/word/fonts/research.odttf'),
                'application/vnd.openxmlformats-officedocument.obfuscatedFont', bytes(data), document.part.package)
    relationship = table.relate_to(part, RT.FONT)
    font = OxmlElement('w:font'); font.set(qn('w:name'), FONT_NAME)
    embedded = OxmlElement('w:embedRegular')
    embedded.set(qn('r:id'), relationship)
    embedded.set(qn('w:fontKey'), '{' + str(key).upper() + '}')
    embedded.set(qn('w:subsetted'), 'false')
    font.append(embedded)
    root = parse_xml(table.blob); root.append(font)
    table._blob = serialize_for_reading(root).encode('utf-8')
    settings = document.settings.element
    for name in ('embedTrueTypeFonts', 'saveSubsetFonts'):
        element = OxmlElement('w:' + name)
        element.set(qn('w:val'), 'true' if name == 'embedTrueTypeFonts' else 'false')
        settings.append(element)


def _bibliography(service, project, db):
    sources, editions = [], {}
    for identifier in project['source_versions']:
        source = service.store.get(identifier, db, 'source')
        sources.append({k: source[k] for k in ('id', 'path', 'raw_hash', 'text_hash', 'coverage', 'evidence_level')})
        link = db.execute('SELECT * FROM source_works WHERE source=?', (identifier,)).fetchone()
        if link:
            record = service.store.get(link['version'], db, 'work_version')
            selected = db.execute('SELECT selected_version_id FROM works WHERE id=?', (link['work'],)).fetchone()[0]
            editions[identifier] = {'id': record['id'], 'work_id': link['work'], 'version_id': record['id'],
                'metadata_hash': record['metadata_hash'], 'selected_version_id': selected, 'csl': record['csl']}
        else:
            csl = {'title': Path(source['path']).stem, 'type': 'document'}
            editions[identifier] = {'id': identifier, 'work_id': None, 'version_id': identifier,
                                     'metadata_hash': _digest(canonical(csl).encode()), 'csl': csl}
    return sources, editions


def _citation_audit(report, citation, review=None):
    audit = dict(citation['audit']) if citation else {
        'schema_version': 1, 'status': 'no_structured_citations', 'scope': 'structured_citations_only',
        'clusters': [], 'bibliography_ids': [], 'metadata_warnings': []}
    positions = sorted({claim['end'] for claim in report['claims'] if claim['evidence']})
    audit['nodes'] = [{'report_end': position, 'bibliography_ids': cluster}
                      for position, cluster in zip(positions, audit['clusters'])]
    # Recognize explicit markup only; this is not semantic figure/reference validation.
    images = re.findall(r'!\[[^\]]*\]\([^\n]*\)|<img\b[^>]*>', report['text'], re.IGNORECASE)
    audit['figure_rendering'] = 'markup_only' if images else 'no_image_markup_detected'
    audit['image_markup_count'] = len(images)
    audit['embedded_figure_count'] = len(list(figure_lines(report['text'])))
    audit['figures'] = report.get('figures', [])
    if audit['embedded_figure_count']:
        audit['figure_rendering'] = 'project_figures_embedded'
    audit['table_figure_reference_check'] = 'not_semantically_verified'
    reference_review = (review or {}).get('reference_review')
    if reference_review and reference_review.get('report_hash') == report['text_hash']:
        audit['table_figure_reference_check'] = reference_review['status']
        audit['reference_review'] = reference_review
    return audit


def _markdown(service, project, report, review, sources, editions, style, human_reviews=()):
    groups = defaultdict(list)
    used = {}
    for claim in sorted(report['claims'], key=lambda c: (c['end'], c['start'])):
        for evidence in claim['evidence']:
            item = editions[evidence['source_version_id']]
            used[item['id']] = item
            if not any(c['id'] == item['id'] for c in groups[claim['end']]):
                source = service.source_manifest(evidence['source_version_id'])
                page = next((p['page'] for p in source['pages'] if p['start'] <= evidence['start'] < p['end']), None)
                cite = {'id': item['id']}
                if page and source['parser'].startswith('pypdf'):
                    cite.update(locator=str(page), label='page')
                groups[claim['end']].append(cite)
    positions = sorted(groups)
    citation = format_citations([{**item['csl'], 'id': identifier} for identifier, item in used.items()],
                               [groups[pos] for pos in positions], style) if positions else None
    text = report['text']
    if citation:
        for pos, value in reversed(list(zip(positions, citation['citations']))):
            text = text[:pos] + ' ' + value + text[pos:]
    title = project['question'].replace('\n', ' ')
    status = project['review_status']
    lines = ['# ' + title, '', STATUS_LABELS.get(status, status), '', text]
    if citation:
        lines += ['', '## 参考文献', ''] + citation['bibliography']
    lines += ['', '## 研究范围与限制', '', f"研究修订：{project['revision']}。纳入 {len(sources)} 份固定版本资料。"]
    if project['scope']:
        lines.append('范围：' + project['scope'])
    if project['exclusions']:
        lines.append('排除条件：' + project['exclusions'])
    if not citation:
        lines.append('正文尚未连接文献引用；此导出保留草稿，不代表证据核验完成。')
    for source in sources:
        original = service.source_manifest(source['id'])
        if source['coverage'] != 'complete':
            lines.append(f"资料 {Path(source['path']).name} 存在提取缺口：" + '；'.join(original['issues']))
        if source['evidence_level'] == 'abstract':
            lines.append(f"资料 {Path(source['path']).name} 仅包含摘要，不等同于全文。")
        if source['evidence_level'] == 'excerpt':
            lines.append(f"资料 {Path(source['path']).name} 仅包含部分摘录，不能代表全文覆盖。")
    audit = _citation_audit(report, citation, review)
    if audit.get('reference_review'):
        check = audit['reference_review']
        label = '完成' if check['status'] == 'text_references_reviewed' else '存在待修订项'
        lines.append(f'表图文字引用检查：{label}；范围为表格文字与图片标题，图像像素内容未核验。')
    labels = {'author': '作者', 'issued': '出版日期', 'container-title': '期刊名'}
    for warning in audit['metadata_warnings']:
        title = used[warning['id']]['csl']['title']
        fields = '、'.join(labels[field] for field in warning['missing_fields'])
        lines.append(f'引用元数据待补充：《{title}》缺少{fields}；导出未补造这些字段。')
    if audit['embedded_figure_count']:
        scope = '图片内容' if audit.get('reference_review') else '图片内容与图表交叉引用'
        lines.append(f'项目图片已嵌入导出；{scope}仍需人工核对。外部图片标记仅保留文本。')
    if audit['figure_rendering'] == 'markup_only':
        lines.append('图片引用保留为原始标记，Word/PDF 未嵌入图像；图像内容与图表交叉引用仍需人工核对。')
    if project.get('stale_sources'):
        lines.append('原始资料已经变化；本报告引用历史快照，需要更新资料后重新核验。')
    if project.get('stale_bibliography'):
        lines.append('文献所选元数据已经变化；本导出保留引用所用版本，需要重新核对书目。')
    if review:
        lines.append(f"核验模型：{review['identity']['provider']} / {review['identity']['model']}。核验覆盖 {len(review['coverage'])} 段正文。")
        lines.extend('核验待处理项：' + issue for issue in review['issues'])
        from .review import claim_is_resolved
        secondary = review.get('critical_review', {})
        if secondary.get('required_claims'):
            lines.append(f"重要论点独立复核：{secondary['required_claims']} 项，{len(secondary['checkpoint_ids'])} 次模型调用。")
        for claim in review['claims'] + secondary.get('claims', []):
            if claim.get('disposition') == 'disclosed':
                lines.append(f"已披露的证据分歧（{claim['assessment']}）：{claim['text']}。{claim['reason']} 披露位置：{claim['disclosure']['text']}")
            elif not claim_is_resolved(claim):
                lines.append(f"待处理结论：{claim['text']}。{claim['reason']}")
    if human_reviews:
        lines += ['', '## 人工复核意见', '', '以下为用户录入的独立意见，保留原模型核验；不覆盖结构错误或自动提升核验状态。']
        labels = {'agree': '认同', 'disagree': '不同意', 'question': '待澄清'}
        for opinion in human_reviews:
            lines.append(f"{labels[opinion['verdict']]}：{opinion['text']}")
            if opinion['anchor']:
                lines.append('对应正文：' + opinion['anchor']['text'])
            lines.append('绑定审核：' + (opinion['review_id'] or '尚无模型审核') + '；记录：' + opinion['id'])
    if report['claims']:
        lines += ['', '## 原文证据记录', '']
        for index, claim in enumerate(report['claims'], 1):
            lines += [f"### 结论 {index}", '', report['text'][claim['start']:claim['end']], '']
            for evidence in claim['evidence']:
                source = service.source_manifest(evidence['source_version_id'])
                page = next((p['page'] for p in source['pages'] if p['start'] <= evidence['start'] < p['end']), None)
                lines.append('\n'.join('> ' + line for line in evidence['quote'].splitlines()))
                lines += ['', f"来源：{Path(source['path']).name}，文件页 {page or '未知'}，字符 {evidence['start']}-{evidence['end']}；版本 {source['raw_hash'][:12]}。", '']
    if citation:
        lines += ['', citation['provenance']['attribution']]
    return '\n'.join(lines) + '\n', citation, list(used.values())


def _table_cells(line):
    # Only unescaped pipes separate cells; literal escaped pipes remain text.
    parts, current, escaped = [], [], False
    for char in line.strip():
        if escaped:
            current.append(char if char == '|' else '\\' + char)
            escaped = False
        elif char == '\\':
            escaped = True
        elif char == '|':
            parts.append(''.join(current).strip()); current = []
        else:
            current.append(char)
    if escaped:
        current.append('\\')
    parts.append(''.join(current).strip())
    if len(parts) == 1:
        return None
    if not parts[0]:
        parts.pop(0)
    if parts and not parts[-1]:
        parts.pop()
    return parts


def _blocks(text, figures=()):
    lines = text.splitlines()
    index, fence = 0, None
    while index < len(lines):
        line = lines[index]
        marker = re.match(r'^\s*(`{3,}|~{3,})', line)
        if marker:
            token = marker.group(1)
            if fence is None:
                fence = token
            elif token[0] == fence[0] and len(token) >= len(fence):
                fence = None
            yield 'paragraph', (0, line)
            index += 1
            continue
        figure = FIGURE_LINE.fullmatch(line) if fence is None else None
        if figure and figure.group(2) in figures:
            yield 'figure', (figure.group(1), figure.group(2))
            index += 1
            continue
        cells = _table_cells(line) if fence is None else None
        separators = _table_cells(lines[index + 1]) if cells and index + 1 < len(lines) else None
        if separators and len(cells) == len(separators) and all(re.fullmatch(r':?-{3,}:?', c) for c in separators):
            rows = [cells]
            index += 2
            while index < len(lines):
                row = _table_cells(lines[index])
                if not row or len(row) != len(cells):
                    break
                rows.append(row)
                index += 1
            yield 'table', rows
            continue
        if line.strip():
            heading = re.match(r'^(#{1,6})\s+(.*)$', line) if fence is None else None
            yield 'paragraph', ((len(heading.group(1)), heading.group(2)) if heading else
                                (0, line[2:] if line.startswith('> ') and fence is None else line))
        index += 1


def _docx(text, figures=None):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    _report_font(text)
    document = Document()
    _embed_docx_font(document)
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.9)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Heading 4', 'Footer'):
        style = document.styles[name]
        style.font.name = FONT_NAME
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attribute), FONT_NAME)
    document.styles['Normal'].font.size = Pt(11)
    for name, size in (('Heading 1', 20), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 11), ('Footer', 9)):
        document.styles[name].font.size = Pt(size)
    document.styles['Normal'].paragraph_format.space_after = Pt(8)
    for kind, content in _blocks(text, figures or {}):
        if kind == 'figure':
            import base64
            caption, identifier = content
            record = figures[identifier]
            scale = min(6.4 / record['width'], 7.5 / record['height'])
            shape = document.add_picture(BytesIO(base64.b64decode(record['data'])),
                width=Inches(record['width'] * scale), height=Inches(record['height'] * scale))
            shape._inline.docPr.set('descr', caption)
            paragraph = document.paragraphs[-1]
            paragraph.alignment = 1
            paragraph.paragraph_format.keep_with_next = bool(caption)
            if caption:
                document.add_paragraph(caption).alignment = 1
            continue
        if kind == 'table':
            table = document.add_table(rows=0, cols=len(content[0]))
            table.style = 'Table Grid'
            for number, values in enumerate(content):
                row = table.add_row()
                row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                for cell, value in zip(row.cells, values):
                    cell.text = value
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.space_after = Pt(4)
                        for run in paragraph.runs:
                            run.bold = number == 0
                if number == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            document.add_paragraph()
            continue
        level, line = content
        if level:
            document.add_heading(line, min(level, 4))
        else:
            document.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    output = BytesIO(); document.save(output)
    return output.getvalue()


def _pdf(text, figures=None):
    from html import escape
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Image, KeepTogether
    _report_font(text)
    normal = ParagraphStyle('ResearchBody', fontName=FONT_NAME, fontSize=10.5, leading=17,
                            alignment=TA_LEFT, wordWrap='CJK', spaceAfter=8, textColor=colors.HexColor('#191919'))
    headings = {i: ParagraphStyle(f'ResearchHeading{i}', parent=normal, fontSize=20 if i == 1 else 14 if i == 2 else 12,
                                 leading=27 if i == 1 else 21, spaceBefore=12, spaceAfter=8, keepWithNext=True) for i in range(1, 7)}
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=48, bottomMargin=48,
                                 title='澜台研究报告', author='澜台')
    story = []
    for kind, content in _blocks(text, figures or {}):
        if kind == 'figure':
            import base64
            caption, identifier = content
            record = figures[identifier]
            scale = min(document.width / record['width'], (document.height - 100) / record['height'])
            picture = Image(BytesIO(base64.b64decode(record['data'])),
                width=record['width'] * scale, height=record['height'] * scale)
            story.append(KeepTogether([picture, Paragraph(escape(caption), normal)] if caption else [picture]))
        elif kind == 'table':
            rows = [[Paragraph(escape(cell), normal) for cell in row] for row in content]
            table = LongTable(rows, colWidths=[document.width / len(rows[0])] * len(rows[0]),
                              repeatRows=1, splitInRow=1, hAlign='LEFT', spaceAfter=12)
            table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
                ('GRID', (0, 0), (-1, -1), .4, colors.HexColor('#aaaaaa')),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]))
            story.append(table)
        else:
            level, line = content
            story.append(Paragraph(escape(line), headings[level] if level else normal))
    if not story:
        story = [Spacer(1, 1)]
    def page(canvas, doc):
        canvas.saveState(); canvas.setFont(FONT_NAME, 9)
        canvas.drawRightString(A4[0] - 50, 28, str(doc.page)); canvas.restoreState()
    document.build(story, onFirstPage=page, onLaterPages=page)
    return output.getvalue()


def export_report(service, project_id, revision, report_id, format, style, idempotency_key):
    if format not in FORMATS or style not in STYLES:
        raise ValueError('Choose markdown/docx/pdf and apa/numeric/gb7714')
    LibraryService(service.workspace)
    payload = dict(project_id=project_id, revision=revision, report_id=report_id, format=format, style=style)
    with service.store.transaction() as db:
        existing = service.store.request(db, 'export:' + idempotency_key, payload)
        if existing:
            return existing
        current = service.store.expect(project_id, revision, db)
        if current['report_id'] != report_id:
            raise RevisionConflict('Report changed; reload before exporting')
        report = service.store.get(report_id, db, 'report')
        figure_service = FigureService(service.workspace, store=service.store)
        figures = {f['id']: figure_service.load(project_id, f['id'], db) for f in report.get('figures', [])}
        for item in report.get('figures', []):
            if figures[item['id']]['sha256'] != item['sha256']:
                raise ValueError('Report figure integrity mismatch')
        sources, editions = _bibliography(service, current, db)
        review = service.store.get(current['review_id'], db, 'review') if current['review_id'] else None
    project = service.project(project_id)
    from ..runtime.research_delivery import project_research_evidence_graph
    evidence_graph_hash = project_research_evidence_graph(service, project_id).graph_hash
    human_reviews = [r for r in service.store.list('human_review', project_id) if r['report_id'] == report_id]
    text, citation, bibliography = _markdown(service, project, report, review, sources, editions, style, human_reviews)
    if format == 'docx':
        text += '\n## 引用编辑说明\n\n引用为静态文本，不含 Zotero 动态引用字段。修改文献或引用格式后，请在澜台重新导出。\n'
    data = portable_markdown(text, figures).encode('utf-8') if format == 'markdown' else _docx(text, figures) if format == 'docx' else _pdf(text, figures)
    artifact = '.lantai/research/exports/' + uuid4().hex + '.' + FORMATS[format][0]
    path = _within(service.workspace, artifact)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with service.store.transaction() as db:
            current = service.store.expect(project_id, revision, db)
            if current['report_id'] != report_id or current['review_id'] != project['review_id']:
                raise RevisionConflict('Report or review changed during export')
            _, now = _bibliography(service, current, db)
            if canonical(now) != canonical(editions):
                raise RevisionConflict('Bibliography changed during export')
            from .sources import source_is_current
            stale = [s for s in current['source_versions'] if not source_is_current(service.workspace, service.store.get(s, db, 'source'))]
            if stale != project['stale_sources']:
                raise RevisionConflict('Sources changed during export')
            # Recheck idempotency after the relatively expensive CSL/Office work.
            existing = service.store.request(db, 'export:' + idempotency_key, payload)
            if existing:
                return existing
            temp = path.with_suffix(path.suffix + '.tmp'); temp.write_bytes(data); temp.replace(path)
            result = service.store.put(db, 'export', {**payload, 'status': project['review_status'],
                'artifact_path': artifact, 'artifact_hash': _digest(data), 'artifact_bytes': len(data),
                'media_type': FORMATS[format][1], 'filename': f"研究报告-修订{revision}.{FORMATS[format][0]}",
                'report_hash': report['text_hash'], 'review_id': current['review_id'], 'sources': sources,
                'evidence_graph_hash': evidence_graph_hash, 'citation_mode': 'static_text',
                'human_review_ids': [r['id'] for r in human_reviews],
                'bibliography_versions': [{k: v for k, v in item.items() if k != 'csl'} for item in bibliography],
                'citation_provenance': citation['provenance'] if citation else None,
                'citation_audit': _citation_audit(report, citation, review)}, project_id, revision)
            service.store.remember(db, 'export:' + idempotency_key, payload, result)
        return result
    except BaseException:
        path.unlink(missing_ok=True)
        path.with_suffix(path.suffix + '.tmp').unlink(missing_ok=True)
        raise


def export_file(service, export_id):
    record = service.store.get(export_id, kind='export')
    path = _within(service.workspace, record['artifact_path'])
    if not path.is_file():
        raise FileNotFoundError('Export file is unavailable')
    if _digest(path.read_bytes()) != record['artifact_hash']:
        raise ValueError('Export file integrity mismatch')
    return path
