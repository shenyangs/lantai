"""Immutable table snapshots and deterministic, restricted aggregations.

This module deliberately has no Python/SQL execution surface. The model may
choose a structured analysis plan; the runtime validates and executes it.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import shutil
import tempfile
import threading
import zipfile
from xml.etree import ElementTree
from collections import defaultdict
from contextlib import contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, localcontext
from functools import wraps
from pathlib import Path
from typing import Any
from weakref import WeakValueDictionary

MAX_SOURCE_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
MAX_SCANNED_ROWS = 250_000
MAX_COLUMNS = 100
MAX_CELLS = 2_000_000
MAX_OUTPUT_ROWS = 500
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_PACKAGE_PARTS = 4096
MAX_DELIVERY_BYTES = 100 * 1024 * 1024
MAX_RECORD_BYTES = 24 * 1024 * 1024
MAX_RESULT_RECORDS_PER_SCOPE = 1000
MAX_CELL_CHARS = 1000
_ID_RE = re.compile(r"^[0-9a-f]{64}$")
_PLAN_KEYS = frozenset({"filters", "date_buckets", "deduplicate_by", "group_by", "metrics", "ratios", "comparisons", "shares", "ranks", "pivot", "sort", "limit"})
SCHEMA_VERSION = 1
PARSER_VERSION = "csv-xlsx-v5"
ENGINE_VERSION = "restricted-aggregation-v10"
_TOTAL_MARKERS = frozenset({"总计", "合计", "小计", "total", "grand total", "subtotal"})
_CURRENCY_COLUMNS = frozenset({"currency", "currency_code", "币种", "货币", "结算币种"})
_XML_INVALID_TEXT = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_SCOPE_LOCKS: WeakValueDictionary[Path, Any] = WeakValueDictionary()
_SCOPE_LOCKS_GUARD = threading.Lock()


class AnalysisError(ValueError):
    """A user-facing analysis boundary or validation failure."""


def _synchronized_scope(method):
    @wraps(method)
    def wrapped(self, *args, **kwargs):
        with self.active_scope():
            return method(self, *args, **kwargs)
    return wrapped


def _canonical(value: Any) -> bytes:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False).encode("utf-8")


def _digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _read_bounded(path: Path, *, limit: int, label: str) -> bytes:
    """Cap allocation even if a file grows after its size was checked."""

    with path.open("rb") as stream:
        blob = stream.read(limit + 1)
    if len(blob) > limit:
        raise AnalysisError(f"{label} exceeds the safety limit")
    return blob


def _atomic_json(path: Path, value: dict[str, Any]) -> dict[str, Any]:
    sealed = {**value, "record_sha256": _digest(_canonical(value))}
    encoded = _canonical(sealed)
    if len(encoded) > MAX_RECORD_BYTES:
        raise AnalysisError("Analysis record exceeds the safety limit")
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile("wb", dir=path.parent, prefix=".analysis-", delete=False) as stream:
        temporary = Path(stream.name)
        stream.write(encoded)
        stream.flush()
        os.fsync(stream.fileno())
    try:
        os.link(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)
    return sealed


def _cell(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    rendered = str(value)
    if len(rendered) > MAX_CELL_CHARS:
        raise AnalysisError("Table cell length exceeds the safety limit")
    _require_xml_text(rendered)
    return rendered


def _require_xml_text(value: str) -> None:
    if _XML_INVALID_TEXT.search(value):
        raise AnalysisError("Analysis input contains an XML-incompatible character; clean it before analysis")


def _require_xml_data_text(value: Any) -> None:
    """Keep plan text safe for DOCX/PPTX/SVG evidence without rewriting it."""

    pending = [value]
    while pending:
        item = pending.pop()
        if isinstance(item, str):
            _require_xml_text(item)
        elif isinstance(item, dict):
            pending.extend(item.keys())
            pending.extend(item.values())
        elif isinstance(item, list):
            pending.extend(item)


def _reject_ambiguous_xlsx_features(archive: zipfile.ZipFile, members: list[zipfile.ZipInfo]) -> None:
    if any(item.filename.lower().startswith("xl/externallinks/") or item.filename.lower().endswith(".bin") for item in members):
        raise AnalysisError("XLSX external links or embedded binaries are not accepted")
    for item in members:
        name = item.filename
        # A workbook relationship may reference a worksheet at any ZIP path or
        # with any extension. The archive's total expanded size is capped above;
        # inspect each part's root so hidden data cannot evade the check by
        # moving or renaming a worksheet part.
        if name.endswith("/"):
            continue
        with archive.open(item) as stream:
            parser = ElementTree.iterparse(stream, events=("start", "end"))
            try:
                _, root = next(parser)
            except (StopIteration, ElementTree.ParseError):
                # Non-XML package parts (images, metadata, etc.) are not sheets.
                continue
            if root.tag.rsplit("}", 1)[-1] != "worksheet":
                continue
            for event, element in parser:
                tag = element.tag.rsplit("}", 1)[-1]
                if event == "start":
                    if tag in {"row", "col"} and element.attrib.get("hidden", "").lower() in {"1", "true"}:
                        raise AnalysisError("XLSX hidden rows or columns need an explicit exported-value table")
                    if tag in {"autoFilter", "mergeCells"}:
                        raise AnalysisError("XLSX filters or merged cells need an explicit exported-value table")
                elif tag == "row":
                    root.clear()


def _parse_table(blob: bytes, suffix: str, sheet_name: str | None) -> tuple[list[str], list[dict[str, str | None]], str | None]:
    if suffix == ".csv":
        if sheet_name:
            raise AnalysisError("CSV does not have sheets")
        try:
            text = blob.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise AnalysisError("CSV must be UTF-8 or UTF-8 with BOM") from exc
        reader = csv.reader(io.StringIO(text, newline=""), strict=True)
        try:
            records = iter(reader)
            header = next(records)
        except (StopIteration, csv.Error) as exc:
            raise AnalysisError("CSV is empty or malformed") from exc
        chosen_sheet = None
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook

            with zipfile.ZipFile(io.BytesIO(blob)) as archive:
                members = archive.infolist()
                if len(members) > MAX_XLSX_PACKAGE_PARTS:
                    raise AnalysisError("XLSX package part count exceeds the safety limit")
                if sum(item.file_size for item in members) > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise AnalysisError("XLSX uncompressed size exceeds the safety limit")
                _reject_ambiguous_xlsx_features(archive, members)
            workbook = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
        except AnalysisError:
            raise
        except (OSError, ValueError, zipfile.BadZipFile) as exc:
            raise AnalysisError("XLSX is invalid") from exc
        if sheet_name is None:
            if len(workbook.sheetnames) != 1:
                workbook.close()
                raise AnalysisError("Choose an XLSX sheet explicitly")
            sheet_name = workbook.sheetnames[0]
        if sheet_name not in workbook.sheetnames:
            workbook.close()
            raise AnalysisError("XLSX sheet does not exist")
        sheet = workbook[sheet_name]
        if sheet.sheet_state != "visible":
            workbook.close()
            raise AnalysisError("XLSX hidden sheets need an explicitly exported visible-value table")
        records = iter(sheet.iter_rows(values_only=False))
        try:
            header_cells = next(records)
            if any(cell.data_type == "f" for cell in header_cells):
                raise AnalysisError("XLSX formula cells are not accepted; export values first")
            header = [_cell(cell.value) for cell in header_cells]
        except StopIteration as exc:
            workbook.close()
            raise AnalysisError("XLSX sheet is empty") from exc
        except Exception:
            workbook.close()
            raise
        chosen_sheet = sheet_name
    else:
        raise AnalysisError("Only CSV and XLSX are supported")

    try:
        columns = [str(item).strip() if item is not None else "" for item in header]
        for column in columns:
            _require_xml_text(column)
        if not columns or len(columns) > MAX_COLUMNS or not all(columns) or any(len(column) > 120 for column in columns):
            raise AnalysisError("Table headers must be nonempty and within the column limit")
        if len(set(columns)) != len(columns):
            raise AnalysisError("Table has duplicate column headers")
        rows: list[dict[str, str | None]] = []
        for scanned_rows, record in enumerate(records, start=1):
            if scanned_rows > MAX_SCANNED_ROWS:
                raise AnalysisError("Table scanned row count exceeds the safety limit")
            if suffix == ".xlsx":
                if any(cell.data_type == "f" for cell in record):
                    raise AnalysisError("XLSX formula cells are not accepted; export values first")
                values = [_cell(cell.value) for cell in record]
            else:
                values = [_cell(value) for value in record]
            if len(values) != len(columns):
                raise AnalysisError("Table contains a row with a different column count")
            if not any(value is not None for value in values):
                continue
            rows.append(dict(zip(columns, values)))
            if len(rows) > MAX_ROWS or len(rows) * len(columns) > MAX_CELLS:
                raise AnalysisError("Table exceeds the row or cell safety limit")
        if not rows:
            raise AnalysisError("Table has no data rows")
        return columns, rows, chosen_sheet
    except csv.Error as exc:
        raise AnalysisError("CSV is malformed") from exc
    finally:
        if suffix == ".xlsx":
            workbook.close()


def _number(value: str | None, *, column: str) -> Decimal | None:
    if value is None:
        return None
    try:
        parsed = Decimal(value)
    except InvalidOperation as exc:
        raise AnalysisError(f"column {column!r} contains nonnumeric data") from exc
    if not parsed.is_finite():
        raise AnalysisError(f"column {column!r} contains a non-finite number")
    if parsed and (parsed.adjusted() > 30 or parsed.as_tuple().exponent < -12):
        raise AnalysisError(f"column {column!r} exceeds supported numeric precision")
    return parsed


def _as_text(value: Decimal) -> str:
    return format(value, "f")


def _profile(columns: list[str], rows: list[dict[str, str | None]]) -> dict[str, Any]:
    """Small, non-sensitive quality summary; never return raw sample rows."""

    candidate_types: dict[str, str] = {}
    date_ranges: dict[str, dict[str, str]] = {}
    for column in columns:
        present = [row[column] for row in rows if row[column] is not None]
        if present and all(_is_bounded_number(value) for value in present):
            candidate_types[column] = "number"
        elif present and all(_is_iso_date(value, column=column) for value in present):
            candidate_types[column] = "date"
            parsed = [_iso_date(value, column=column) for value in present]
            date_ranges[column] = {"min": min(parsed).isoformat(), "max": max(parsed).isoformat()}
        else:
            candidate_types[column] = "text"
    possible_total_count = sum(_possible_total_row(row, columns) for row in rows)
    notes = ["Values and business definitions are not independently verified"]
    if possible_total_count:
        notes.append("Possible total or subtotal rows must be explicitly excluded or inspected before aggregation")
    if date_ranges:
        notes.append("Date coverage is observed, not proven complete; a grouped period may be incomplete")
    return {
        "null_counts": {column: sum(row[column] is None for row in rows) for column in columns},
        "candidate_types": candidate_types,
        "date_ranges": date_ranges,
        "possible_total_row_count": possible_total_count,
        "grain": "unconfirmed",
        "unit": "unconfirmed",
        "notes": notes,
    }


def _is_bounded_number(value: str) -> bool:
    try:
        return _number(value, column="profile") is not None
    except AnalysisError:
        return False


def _is_iso_date(value: str, *, column: str) -> bool:
    try:
        _iso_date(value, column=column)
        return True
    except AnalysisError:
        return False


def _possible_total_row(row: dict[str, str | None], columns: list[str]) -> bool:
    return any(
        isinstance(row[column], str) and row[column].strip().casefold() in _TOTAL_MARKERS
        for column in columns
    )


def _iso_date(value: str, *, column: str) -> date:
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}(?:\.\d{1,6})?", value):
        try:
            return datetime.fromisoformat(value).date()
        except ValueError as exc:
            raise AnalysisError(f"column {column!r} contains an invalid ISO datetime") from exc
    if not re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?", value):
        raise AnalysisError(f"column {column!r} requires unambiguous ISO dates or timezone-free ISO datetimes")
    try:
        return date.fromisoformat(value + "-01" if len(value) == 7 else value)
    except ValueError as exc:
        raise AnalysisError(f"column {column!r} contains an invalid ISO date") from exc


def _validate_filters(items: Any, columns: list[str], *, limit: int) -> list[dict[str, Any]]:
    if not isinstance(items, list) or len(items) > limit:
        raise AnalysisError("Too many filters")
    normalized = []
    for item in items:
        if not isinstance(item, dict) or set(item) != {"column", "op", "value"}:
            raise AnalysisError("Invalid filter operation")
        if not isinstance(item["column"], str) or item["column"] not in columns or not isinstance(item["op"], str) or item["op"] not in {"eq", "ne", "gt", "gte", "lt", "lte", "in"}:
            raise AnalysisError("Invalid filter column or operation")
        if item["op"] == "in":
            if not isinstance(item["value"], list) or len(item["value"]) > 100 or any(not isinstance(value, (str, int, float, bool, type(None))) or len(str(value)) > 200 for value in item["value"]):
                raise AnalysisError("Invalid filter list")
        elif not isinstance(item["value"], (str, int, float, bool, type(None))) or len(str(item["value"])) > 200:
            raise AnalysisError("Invalid filter value")
        normalized.append(item)
    return normalized


def _validate_plan(plan: dict[str, Any], columns: list[str]) -> dict[str, Any]:
    if not isinstance(plan, dict) or set(plan) - _PLAN_KEYS:
        raise AnalysisError("Unsupported analysis operation")
    _require_xml_data_text(plan)
    group_by = plan.get("group_by", [])
    metrics = plan.get("metrics", [])
    filters = plan.get("filters", [])
    date_buckets = plan.get("date_buckets", [])
    ratios = plan.get("ratios", [])
    comparisons = plan.get("comparisons", [])
    shares = plan.get("shares", [])
    ranks = plan.get("ranks", [])
    pivot = plan.get("pivot")
    deduplicate_by = plan.get("deduplicate_by", [])
    sort = plan.get("sort", [])
    limit = plan.get("limit", MAX_OUTPUT_ROWS)
    if not isinstance(date_buckets, list) or len(date_buckets) > 2:
        raise AnalysisError("Too many date buckets")
    normalized_buckets = []
    available_columns = list(columns)
    for bucket in date_buckets:
        if not isinstance(bucket, dict) or set(bucket) != {"column", "grain", "as"} or not isinstance(bucket["column"], str) or bucket["column"] not in columns or not isinstance(bucket["grain"], str) or bucket["grain"] not in {"year", "quarter", "month"} or not isinstance(bucket["as"], str) or not bucket["as"] or len(bucket["as"]) > 80 or bucket["as"] in available_columns:
            raise AnalysisError("Invalid date bucket")
        available_columns.append(bucket["as"])
        normalized_buckets.append(bucket)
    if not isinstance(group_by, list) or len(group_by) > 4 or any(not isinstance(c, str) or c not in available_columns for c in group_by) or len(set(group_by)) != len(group_by):
        raise AnalysisError("Invalid group_by column")
    if not isinstance(deduplicate_by, list) or len(deduplicate_by) > 4 or any(not isinstance(c, str) or c not in columns for c in deduplicate_by) or len(set(deduplicate_by)) != len(deduplicate_by):
        raise AnalysisError("Invalid deduplication key")
    if not isinstance(metrics, list) or not (1 <= len(metrics) <= 12):
        raise AnalysisError("Select 1 to 12 metrics")
    normalized_filters = _validate_filters(filters, columns, limit=12)
    if not isinstance(ratios, list) or len(ratios) > 6:
        raise AnalysisError("Too many ratios")
    if not isinstance(comparisons, list) or len(comparisons) > 6:
        raise AnalysisError("Too many comparisons")
    if not isinstance(shares, list) or len(shares) > 6:
        raise AnalysisError("Too many shares")
    if not isinstance(ranks, list) or len(ranks) > 6:
        raise AnalysisError("Too many ranks")
    if not isinstance(sort, list) or len(sort) > 4:
        raise AnalysisError("Too many sort keys")
    if not isinstance(limit, int) or isinstance(limit, bool) or not 1 <= limit <= MAX_OUTPUT_ROWS:
        raise AnalysisError("Invalid output limit")
    names = set(group_by)
    metric_names: set[str] = set()
    normalized_metrics = []
    for metric in metrics:
        if not isinstance(metric, dict) or set(metric) - {"column", "op", "as", "where", "unit"}:
            raise AnalysisError("Invalid metric operation")
        op = metric.get("op")
        column = metric.get("column")
        if not isinstance(op, str) or op not in {"sum", "avg", "min", "max", "count", "count_distinct"}:
            raise AnalysisError("Unsupported metric operation")
        if op != "count" and (not isinstance(column, str) or column not in columns):
            raise AnalysisError("Invalid metric column")
        if op == "count" and column is not None and (not isinstance(column, str) or column not in columns):
            raise AnalysisError("Invalid metric column")
        unit = metric.get("unit")
        if unit is not None and (
            op not in {"sum", "avg", "min", "max"}
            or not isinstance(unit, dict)
            or set(unit) != {"kind", "name"}
            or unit.get("kind") != "quantity"
            or not isinstance(unit.get("name"), str)
            or not 1 <= len(unit["name"]) <= 40
            or unit["name"] != unit["name"].strip()
            or re.search(r"[\x00-\x1f\x7f]", unit["name"])
        ):
            raise AnalysisError("Invalid metric unit declaration")
        alias = metric.get("as") or (f"{column}_{op}" if column else "row_count")
        if not isinstance(alias, str) or not alias or len(alias) > 80 or alias in names:
            raise AnalysisError("Metric aliases must be unique")
        names.add(alias)
        metric_names.add(alias)
        normalized_metrics.append({"column": column, "op": op, "as": alias, "where": _validate_filters(metric.get("where", []), columns, limit=8), "unit": unit})
    normalized_ratios = []
    for item in ratios:
        if not isinstance(item, dict) or set(item) != {"numerator", "denominator", "as"}:
            raise AnalysisError("Invalid ratio operation")
        if not isinstance(item["numerator"], str) or not isinstance(item["denominator"], str) or item["numerator"] not in metric_names or item["denominator"] not in metric_names:
            raise AnalysisError("Ratio references an unknown metric")
        if not isinstance(item["as"], str) or not item["as"] or len(item["as"]) > 80 or item["as"] in names:
            raise AnalysisError("Ratio alias must be unique")
        names.add(item["as"])
        normalized_ratios.append(item)
    normalized_comparisons = []
    comparable_names = set(metric_names) | {item["as"] for item in normalized_ratios}
    for item in comparisons:
        if not isinstance(item, dict) or set(item) != {"current", "baseline", "op", "as"} or not isinstance(item.get("op"), str) or item["op"] not in {"difference", "growth_rate"}:
            raise AnalysisError("Invalid comparison operation")
        if not isinstance(item["current"], str) or not isinstance(item["baseline"], str) or item["current"] not in comparable_names or item["baseline"] not in comparable_names:
            raise AnalysisError("Comparison references an unknown metric")
        if not isinstance(item["as"], str) or not item["as"] or len(item["as"]) > 80 or item["as"] in names:
            raise AnalysisError("Comparison alias must be unique")
        names.add(item["as"])
        normalized_comparisons.append(item)
    normalized_shares = []
    for item in shares:
        if not isinstance(item, dict) or set(item) != {"metric", "as"} or not isinstance(item.get("metric"), str) or item["metric"] not in metric_names:
            raise AnalysisError("Share references an unknown metric")
        if not isinstance(item.get("as"), str) or not item["as"] or len(item["as"]) > 80 or item["as"] in names:
            raise AnalysisError("Share alias must be unique")
        names.add(item["as"])
        normalized_shares.append(item)
    normalized_ranks = []
    for item in ranks:
        if not isinstance(item, dict) or set(item) != {"metric", "as", "descending", "method"} or not isinstance(item.get("metric"), str) or item["metric"] not in metric_names or not isinstance(item.get("descending"), bool) or item.get("method") not in {"competition", "dense"}:
            raise AnalysisError("Invalid rank operation")
        if not isinstance(item.get("as"), str) or not item["as"] or len(item["as"]) > 80 or item["as"] in names:
            raise AnalysisError("Rank alias must be unique")
        names.add(item["as"])
        normalized_ranks.append(item)
    for item in sort:
        if not isinstance(item, dict) or set(item) - {"column", "descending"} or not isinstance(item.get("column"), str) or item.get("column") not in names:
            raise AnalysisError("Invalid sort column")
        if not isinstance(item.get("descending", False), bool):
            raise AnalysisError("Invalid sort direction")
    if pivot is not None:
        if (
            not isinstance(pivot, dict)
            or set(pivot) != {"column", "metric", "as_prefix"}
            or not isinstance(pivot.get("column"), str)
            or pivot["column"] not in group_by
            or not isinstance(pivot.get("metric"), str)
            or len(normalized_metrics) != 1
            or pivot["metric"] != normalized_metrics[0]["as"]
            or not isinstance(pivot.get("as_prefix"), str)
            or not 1 <= len(pivot["as_prefix"]) <= 40
            or re.search(r"[\x00-\x1f\x7f]", pivot["as_prefix"])
            or normalized_ratios or normalized_comparisons or normalized_shares or normalized_ranks
            or any(item["column"] not in group_by or item["column"] == pivot["column"] for item in sort)
        ):
            raise AnalysisError("Pivot requires one grouped metric, a grouped pivot column, and sort keys from remaining dimensions")
    return {"filters": normalized_filters, "date_buckets": normalized_buckets, "deduplicate_by": deduplicate_by, "group_by": group_by, "metrics": normalized_metrics, "ratios": normalized_ratios, "comparisons": normalized_comparisons, "shares": normalized_shares, "ranks": normalized_ranks, "pivot": pivot, "sort": sort, "limit": limit}


def _matches(row: dict[str, str | None], filters: list[dict[str, Any]]) -> bool:
    for item in filters:
        actual = row[item["column"]]
        expected = item["value"]
        op = item["op"]
        if op == "in":
            matched = actual in {None if value is None else str(value) for value in expected}
        elif op in {"eq", "ne"}:
            matched = actual == (None if expected is None else str(expected))
            if op == "ne":
                matched = not matched
        else:
            if actual is None or expected is None:
                matched = False
            else:
                if isinstance(expected, str) and re.fullmatch(r"\d{4}-\d{2}(?:-\d{2})?", expected):
                    left, right = _iso_date(actual, column=item["column"]), _iso_date(expected, column=item["column"])
                else:
                    left, right = _number(actual, column=item["column"]), _number(str(expected), column=item["column"])
                matched = {"gt": left > right, "gte": left >= right, "lt": left < right, "lte": left <= right}[op]
        if not matched:
            return False
    return True


def _aggregate(rows: list[dict[str, str | None]], plan: dict[str, Any]) -> tuple[list[dict[str, Any]], list[str]]:
    selected = []
    for row in rows:
        if not _matches(row, plan["filters"]):
            continue
        enriched = row.copy()
        for bucket in plan["date_buckets"]:
            raw = row[bucket["column"]]
            parsed = _iso_date(raw, column=bucket["column"]) if raw is not None else None
            grain = bucket["grain"]
            enriched[bucket["as"]] = (
                None if parsed is None else
                f"{parsed.year:04d}" if grain == "year" else
                f"{parsed.year:04d}-Q{(parsed.month - 1) // 3 + 1}" if grain == "quarter" else
                f"{parsed.year:04d}-{parsed.month:02d}"
            )
        selected.append(enriched)
    warnings: list[str] = []
    if plan["deduplicate_by"]:
        unique: dict[tuple[str, ...], dict[str, str | None]] = {}
        deduplicated = []
        removed = 0
        for row in selected:
            key_values = [row[column] for column in plan["deduplicate_by"]]
            if any(value is None for value in key_values):
                raise AnalysisError("Deduplication key contains a missing value")
            key = tuple(str(value) for value in key_values)
            previous = unique.get(key)
            if previous is not None:
                if previous != row:
                    raise AnalysisError("Deduplication key has conflicting rows; choose a finer key or resolve the source")
                removed += 1
            else:
                unique[key] = row
                deduplicated.append(row)
        selected = deduplicated
        if removed:
            warnings.append(f"{removed} exact duplicate rows removed using the explicit key {plan['deduplicate_by']}")
    columns = list(rows[0]) if rows else []
    currency_columns = [column for column in columns if column.strip().casefold() in _CURRENCY_COLUMNS]
    if len(selected) > 1 and any(_possible_total_row(row, columns) for row in selected):
        raise AnalysisError("Possible total or subtotal row mixed with detail; inspect it or explicitly filter it before aggregation")
    groups: dict[tuple[str | None, ...], list[dict[str, str | None]]] = defaultdict(list)
    for row in selected:
        groups[tuple(row[column] for column in plan["group_by"])].append(row)
    if not plan["group_by"] and not groups:
        groups[()] = []
    if len(groups) > 10_000:
        raise AnalysisError("Analysis produced too many groups; add filters or grouping")
    output = []
    currencies_by_group: list[dict[str, dict[str, str]]] = []
    if selected and plan["date_buckets"]:
        warnings.append("Date coverage is observed, not proven complete; a grouped period may be incomplete. Do not claim a full-period comparison without checking coverage")
    for key, members in groups.items():
        record: dict[str, Any] = dict(zip(plan["group_by"], key))
        metric_currencies: dict[str, dict[str, str]] = {}
        known_units: dict[str, tuple[Any, ...] | None] = {}
        for metric in plan["metrics"]:
            op, column, alias = metric["op"], metric["column"], metric["as"]
            metric_members = [member for member in members if _matches(member, metric["where"])]
            if metric["unit"] is not None:
                warnings.append(f"{alias}: quantity unit {metric['unit']['name']!r} was declared in the analysis plan, not verified from source metadata")
            if op in {"sum", "avg", "min", "max"} and metric["unit"] is None:
                resolved: dict[str, str] = {}
                for currency_column in currency_columns:
                    units = {member[currency_column] for member in metric_members if member[column] is not None}
                    if any(unit is None or not unit.strip() for unit in units):
                        raise AnalysisError(f"{alias}: missing currency values cannot support monetary aggregation")
                    if len(units) > 1:
                        raise AnalysisError(f"{alias}: mixed or missing currency values cannot be aggregated together")
                    if units:
                        unit = next(iter(units))
                        if unit is not None:
                            resolved[currency_column] = unit
                metric_currencies[alias] = resolved
            if op in {"count", "count_distinct"}:
                known_units[alias] = ("count",)
            elif metric["unit"] is not None:
                known_units[alias] = ("quantity", metric["unit"]["name"])
            elif metric_currencies.get(alias):
                known_units[alias] = ("currency", tuple(sorted(metric_currencies[alias].items())))
            else:
                known_units[alias] = None
            values = [member[column] for member in metric_members] if column else []
            missing = sum(value is None for value in values)
            if missing:
                warnings.append(f"{alias}: {missing} missing values excluded from this group")
            if op == "count":
                record[alias] = len([value for value in values if value is not None]) if column else len(metric_members)
            elif op == "count_distinct":
                record[alias] = len({value for value in values if value is not None})
            else:
                numbers = [_number(value, column=column) for value in values]
                nonnull = [value for value in numbers if value is not None]
                if not nonnull:
                    record[alias] = None
                    warnings.append(f"{alias}: no non-null numeric values")
                elif op == "sum":
                    with localcontext() as context:
                        context.prec = 64
                        record[alias] = _as_text(sum(nonnull, Decimal(0)))
                elif op == "avg":
                    with localcontext() as context:
                        context.prec = 64
                        record[alias] = _as_text(sum(nonnull, Decimal(0)) / len(nonnull))
                elif op == "min":
                    record[alias] = _as_text(min(nonnull))
                else:
                    record[alias] = _as_text(max(nonnull))
        for ratio in plan["ratios"]:
            left_currency = metric_currencies.get(ratio["numerator"], {})
            right_currency = metric_currencies.get(ratio["denominator"], {})
            if any(left_currency.get(column) != right_currency.get(column) for column in left_currency.keys() & right_currency.keys()):
                raise AnalysisError(f"{ratio['as']}: cross-currency ratio requires an explicit conversion basis")
            left_unit = known_units[ratio["numerator"]]
            right_unit = known_units[ratio["denominator"]]
            known_units[ratio["as"]] = (
                None if left_unit is None or right_unit is None else
                ("dimensionless",) if left_unit == right_unit else
                ("rate", left_unit, right_unit)
            )
            numerator = record[ratio["numerator"]]
            denominator = record[ratio["denominator"]]
            if numerator is None or denominator is None or Decimal(str(denominator)) == 0:
                record[ratio["as"]] = None
                warnings.append(f"{ratio['as']}: zero or missing denominator")
            else:
                with localcontext() as context:
                    context.prec = 28
                    record[ratio["as"]] = _as_text(Decimal(str(numerator)) / Decimal(str(denominator)))
        for comparison in plan["comparisons"]:
            left_currency = metric_currencies.get(comparison["current"], {})
            right_currency = metric_currencies.get(comparison["baseline"], {})
            if any(left_currency.get(column) != right_currency.get(column) for column in left_currency.keys() & right_currency.keys()):
                raise AnalysisError(f"{comparison['as']}: cross-currency comparison requires an explicit conversion basis")
            left_unit = known_units[comparison["current"]]
            right_unit = known_units[comparison["baseline"]]
            if left_unit is not None and right_unit is not None and left_unit != right_unit:
                raise AnalysisError(f"{comparison['as']}: incompatible units cannot support a comparison")
            current = record[comparison["current"]]
            baseline = record[comparison["baseline"]]
            alias = comparison["as"]
            if current is None or baseline is None:
                record[alias] = None
                warnings.append(f"{alias}: missing current or baseline value")
                continue
            with localcontext() as context:
                context.prec = 28
                delta = Decimal(str(current)) - Decimal(str(baseline))
                if comparison["op"] == "difference":
                    record[alias] = _as_text(delta)
                elif Decimal(str(baseline)) <= 0:
                    record[alias] = None
                    warnings.append(f"{alias}: zero or negative baseline cannot support a conventional growth rate")
                else:
                    record[alias] = _as_text(delta / Decimal(str(baseline)))
        output.append(record)
        currencies_by_group.append(metric_currencies)
    for share in plan["shares"]:
        metric, alias = share["metric"], share["as"]
        units = {tuple(sorted(group.get(metric, {}).items())) for group in currencies_by_group}
        if len(units) > 1:
            raise AnalysisError(f"{alias}: cross-currency share requires an explicit conversion basis")
        values = [row[metric] for row in output]
        if any(value is None or Decimal(str(value)) < 0 for value in values):
            raise AnalysisError(f"{alias}: share needs nonnegative, nonmissing values and a positive denominator")
        with localcontext() as context:
            context.prec = 64
            denominator = sum((Decimal(str(value)) for value in values), Decimal(0))
            if denominator <= 0:
                raise AnalysisError(f"{alias}: share needs a positive denominator")
            for row, value in zip(output, values):
                row[alias] = _as_text(Decimal(str(value)) / denominator)
    for rank in plan["ranks"]:
        metric, alias = rank["metric"], rank["as"]
        units = {tuple(sorted(group.get(metric, {}).items())) for group in currencies_by_group}
        if len(units) > 1:
            raise AnalysisError(f"{alias}: cross-currency rank requires an explicit conversion basis")
        ordered = sorted(
            (Decimal(str(row[metric])) for row in output if row[metric] is not None),
            reverse=rank["descending"],
        )
        positions: dict[Decimal, int] = {}
        for position, value in enumerate(ordered, start=1):
            if value not in positions:
                positions[value] = position if rank["method"] == "competition" else len(positions) + 1
        for row in output:
            row[alias] = None if row[metric] is None else positions[Decimal(str(row[metric]))]
        if any(row[alias] is None for row in output):
            warnings.append(f"{alias}: groups with a missing metric are unranked")
    if plan["pivot"] is not None:
        pivot = plan["pivot"]
        column, metric, prefix = pivot["column"], pivot["metric"], pivot["as_prefix"]
        dimensions = [name for name in plan["group_by"] if name != column]
        categories = sorted({row[column] for row in output}, key=lambda value: "" if value is None else str(value))
        if len(categories) > 20:
            raise AnalysisError("Too many pivot columns; filter the data or keep a long-form table")
        if any(
            category is None or not str(category).strip()
            or len(prefix + str(category)) > 120
            or re.search(r"[\x00-\x1f\x7f]", str(category))
            or prefix + str(category) in dimensions
            for category in categories
        ):
            raise AnalysisError("Ambiguous pivot columns; choose a different pivot or prefix")
        wide: dict[tuple[str | None, ...], dict[str, Any]] = {}
        for row in output:
            key = tuple(row[name] for name in dimensions)
            if key not in wide:
                wide[key] = {**dict(zip(dimensions, key)), **{prefix + str(category): None for category in categories}}
            wide[key][prefix + str(row[column])] = row[metric]
        output = list(wide.values())
    for key in reversed(plan["sort"]):
        column = key["column"]
        reverse = key.get("descending", False)
        numeric = column not in plan["group_by"]
        output.sort(
            key=lambda row: (Decimal(str(row[column])) if row[column] is not None else Decimal(0)) if numeric else str(row[column] or ""),
            reverse=reverse,
        )
        output.sort(key=lambda row: row[column] is None)
    if len(output) > plan["limit"]:
        raise AnalysisError("Result exceeds output limit; add filters or increase limit")
    return output, list(dict.fromkeys(warnings))


class AnalysisService:
    """Scope-local, immutable evidence repository for one chat/workspace."""

    def __init__(self, state_root: Path, *, scope: str):
        if not scope or len(scope) > 512:
            raise AnalysisError("A bounded analysis scope is required")
        self.state_root = Path(state_root).expanduser().resolve()
        self.root = self.state_root / "analysis" / _digest(scope.encode("utf-8"))
        self.retired_marker = self.root.parent / f".{self.root.name}.retired"

    def _assert_safe_scope(self, *, allow_retired: bool = False) -> None:
        if self.state_root.is_symlink() or self.root.parent.is_symlink() or self.root.is_symlink() or self.retired_marker.is_symlink():
            raise AnalysisError("Analysis scope path cannot contain a symbolic link")
        if self.retired_marker.exists() and not allow_retired:
            raise AnalysisError("Analysis scope was retired with its deleted conversation")

    @contextmanager
    def active_scope(self):
        """Hold the scope lock while another component stores linked metadata."""

        with _SCOPE_LOCKS_GUARD:
            lock = _SCOPE_LOCKS.setdefault(self.root, threading.RLock())
        with lock:
            self._assert_safe_scope()
            yield

    def _record_path(self, kind: str, identifier: str) -> Path:
        self._assert_safe_scope()
        if not _ID_RE.fullmatch(identifier):
            raise AnalysisError(f"Invalid {kind} id")
        return self.root / kind / f"{identifier}.json"

    def _load(self, kind: str, identifier: str) -> dict[str, Any]:
        path = self._record_path(kind, identifier)
        label = {"datasets": "dataset", "results": "result", "deliveries": "delivery"}[kind]
        try:
            record = json.loads(_read_bounded(path, limit=MAX_RECORD_BYTES, label="Analysis record"))
        except (FileNotFoundError, UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise AnalysisError(f"Unknown {label} id") from exc
        identity_key = {"datasets": "dataset_id", "results": "analysis_id", "deliveries": "delivery_id"}[kind]
        if not isinstance(record, dict) or record.get("schema_version") != SCHEMA_VERSION or record.get(identity_key) != identifier:
            raise AnalysisError(f"{label} record failed integrity check")
        checksum = record.pop("record_sha256", None)
        if checksum != _digest(_canonical(record)):
            raise AnalysisError(f"{label} record failed integrity check")
        record["record_sha256"] = checksum
        return record

    @_synchronized_scope
    def ingest(self, source_path: Path, *, sheet_name: str | None = None, cancel_event: threading.Event | None = None) -> dict[str, Any]:
        source = Path(source_path).expanduser().resolve(strict=True)
        _require_xml_text(str(source))
        if not source.is_file() or source.suffix.lower() not in {".csv", ".xlsx"}:
            raise AnalysisError("Source must be a CSV or XLSX file")
        stat_before = source.stat()
        if stat_before.st_size > MAX_SOURCE_BYTES:
            raise AnalysisError("Source exceeds the 20 MB safety limit")
        blob = _read_bounded(source, limit=MAX_SOURCE_BYTES, label="Source")
        stat_after = source.stat()
        if len(blob) > MAX_SOURCE_BYTES or (stat_before.st_size, stat_before.st_mtime_ns) != (stat_after.st_size, stat_after.st_mtime_ns):
            raise AnalysisError("Source changed during import or exceeds the safety limit")
        from .worker import run_worker

        parsed = run_worker("ingest", blob=blob, suffix=source.suffix.lower(), sheet_name=sheet_name, cancel_event=cancel_event)
        if cancel_event is not None and cancel_event.is_set():
            raise AnalysisError("Analysis import cancelled; the dataset was not saved")
        columns, chosen_sheet = parsed["columns"], parsed["sheet_name"]
        source_hash = _digest(blob)
        identity = {"source_path": str(source), "source_sha256": source_hash, "sheet_name": chosen_sheet, "parser_version": PARSER_VERSION}
        identifier = _digest(_canonical(identity))
        snapshot = self.root / "snapshots" / f"{identifier}{source.suffix.lower()}"
        snapshot.parent.mkdir(parents=True, exist_ok=True)
        if not snapshot.exists():
            with tempfile.NamedTemporaryFile("wb", dir=snapshot.parent, prefix=".snapshot-", delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(blob)
                stream.flush()
                os.fsync(stream.fileno())
            try:
                os.link(temporary, snapshot)
            except FileExistsError:
                pass
            finally:
                temporary.unlink(missing_ok=True)
        if _digest(_read_bounded(snapshot, limit=MAX_SOURCE_BYTES, label="Dataset snapshot")) != source_hash:
            raise AnalysisError("Dataset snapshot failed integrity check")
        record = {"schema_version": SCHEMA_VERSION, "dataset_id": identifier, **identity, "columns": columns, "row_count": parsed["row_count"], "profile": parsed["profile"], "snapshot_path": str(snapshot), "imported_at": datetime.now(timezone.utc).isoformat()}
        path = self._record_path("datasets", identifier)
        if not path.exists():
            try:
                record = _atomic_json(path, record)
            except FileExistsError:
                record = self._load("datasets", identifier)
        else:
            record = self._load("datasets", identifier)
        return {key: record[key] for key in ("schema_version", "dataset_id", "source_path", "source_sha256", "sheet_name", "columns", "row_count", "profile", "imported_at")}

    def get_dataset(self, dataset_id: str) -> dict[str, Any]:
        return self._load("datasets", dataset_id)

    def source_status(self, dataset_id: str) -> dict[str, str]:
        dataset = self.get_dataset(dataset_id)
        source = Path(dataset["source_path"])
        try:
            if source.stat().st_size > MAX_SOURCE_BYTES:
                return {"status": "changed", "source_sha256": dataset["source_sha256"]}
            current = _digest(_read_bounded(source, limit=MAX_SOURCE_BYTES, label="Source"))
        except AnalysisError:
            return {"status": "changed", "source_sha256": dataset["source_sha256"]}
        except (FileNotFoundError, PermissionError):
            return {"status": "unavailable", "source_sha256": dataset["source_sha256"]}
        return {"status": "current" if current == dataset["source_sha256"] else "changed", "source_sha256": dataset["source_sha256"]}

    @_synchronized_scope
    def run(self, dataset_id: str, plan: dict[str, Any], *, cancel_event: threading.Event | None = None) -> dict[str, Any]:
        dataset = self.get_dataset(dataset_id)
        snapshot = Path(dataset["snapshot_path"])
        expected = self.root / "snapshots" / f"{dataset_id}{Path(dataset['source_path']).suffix.lower()}"
        if snapshot != expected:
            raise AnalysisError("Dataset snapshot path failed integrity check")
        blob = _read_bounded(snapshot, limit=MAX_SOURCE_BYTES, label="Dataset snapshot")
        if _digest(blob) != dataset["source_sha256"]:
            raise AnalysisError("Dataset snapshot failed integrity check")
        from .worker import run_worker

        computed = run_worker("compute", blob=blob, suffix=snapshot.suffix.lower(), sheet_name=dataset["sheet_name"], plan=plan, cancel_event=cancel_event)
        if cancel_event is not None and cancel_event.is_set():
            raise AnalysisError("Analysis calculation cancelled; the result was not saved")
        columns = computed["columns"]
        if columns != dataset["columns"] or computed["row_count"] != dataset["row_count"]:
            raise AnalysisError("Dataset snapshot schema changed")
        normalized, output, warnings = computed["plan"], computed["rows"], computed["warnings"]
        identifier = _digest(_canonical({"dataset_id": dataset_id, "plan": normalized, "engine_version": ENGINE_VERSION}))
        record = {
            "schema_version": SCHEMA_VERSION,
            "engine_version": ENGINE_VERSION,
            "analysis_id": identifier,
            "dataset_id": dataset_id,
            "plan": normalized,
            "rows": output,
            "warnings": warnings,
            "verification": {
                "source_integrity": {"status": "passed", "evidence": "Snapshot SHA-256 matched the recorded source hash"},
                "calculation_validity": {"status": "passed", "evidence": "Validated restricted plan executed by the versioned aggregation engine"},
                "semantic_review": {"status": "pending", "reason": "Business grain, units, filters, and interpretation require contextual review"},
                "delivery_consistency": {"status": "not_applicable", "reason": "No delivery artifact has been generated by this result"},
            },
            "evidence": {
                "source_path": dataset["source_path"],
                "source_sha256": dataset["source_sha256"],
                "source_row_count": dataset["row_count"],
                "source_columns": columns,
                "dataset_profile": dataset["profile"],
                "sheet_name": dataset["sheet_name"],
                "analysis_plan": normalized,
                "analysis_id": identifier,
                "engine_version": ENGINE_VERSION,
                "parser_version": dataset["parser_version"],
                "calculated_at": datetime.now(timezone.utc).isoformat(),
            },
        }
        path = self._record_path("results", identifier)
        if not path.exists():
            try:
                record = _atomic_json(path, record)
            except FileExistsError:
                record = self._load("results", identifier)
        else:
            record = self._load("results", identifier)
        return record

    def get_result(self, analysis_id: str) -> dict[str, Any]:
        return self._load("results", analysis_id)

    @_synchronized_scope
    def latest_result(self) -> dict[str, Any]:
        """Return the newest immutable result in this conversation scope.

        Result discovery stays inside the already hashed scope. Every candidate
        is loaded through the normal checksum gate before its recorded UTC
        calculation time is considered; corrupt or excessive state fails
        closed instead of silently selecting a different result.
        """

        results_dir = self.root / "results"
        if not results_dir.exists():
            raise AnalysisError("No analysis result is available in this conversation")
        if results_dir.is_symlink() or not results_dir.is_dir():
            raise AnalysisError("Analysis results path failed integrity check")
        candidates: list[tuple[datetime, str, dict[str, Any]]] = []
        entry_count = 0
        for path in results_dir.iterdir():
            entry_count += 1
            if entry_count > MAX_RESULT_RECORDS_PER_SCOPE:
                raise AnalysisError("Analysis result count exceeds the safety limit")
            if path.is_symlink():
                raise AnalysisError("Analysis result record failed integrity check")
            if not path.is_file() or path.suffix != ".json" or not _ID_RE.fullmatch(path.stem):
                raise AnalysisError("Analysis result record failed integrity check")
            result = self._load("results", path.stem)
            evidence = result.get("evidence")
            calculated_at = evidence.get("calculated_at") if isinstance(evidence, dict) else None
            try:
                timestamp = datetime.fromisoformat(calculated_at)
            except (TypeError, ValueError) as exc:
                raise AnalysisError("Analysis result record failed integrity check") from exc
            if timestamp.tzinfo is None:
                raise AnalysisError("Analysis result record failed integrity check")
            candidates.append((timestamp.astimezone(timezone.utc), result["analysis_id"], result))
        if not candidates:
            raise AnalysisError("No analysis result is available in this conversation")
        return max(candidates, key=lambda item: (item[0], item[1]))[2]

    @_synchronized_scope
    def refresh_result(self, analysis_id: str, *, cancel_event: threading.Event | None = None) -> dict[str, Any]:
        """Re-run the same validated plan against a newly captured source version."""

        previous = self.get_result(analysis_id)
        previous_dataset_id = previous["dataset_id"]
        status = self.source_status(previous_dataset_id)["status"]
        if status == "current":
            return {"status": "current", "analysis_id": analysis_id, "dataset_id": previous_dataset_id}
        if status == "unavailable":
            return {"status": "unavailable", "analysis_id": analysis_id, "dataset_id": previous_dataset_id}
        previous_dataset = self.get_dataset(previous_dataset_id)
        replacement = self.ingest(
            Path(previous_dataset["source_path"]),
            sheet_name=previous_dataset["sheet_name"], cancel_event=cancel_event,
        )
        if cancel_event is not None and cancel_event.is_set():
            raise AnalysisError("Analysis refresh cancelled; the result was not saved")
        refreshed = self.run(replacement["dataset_id"], previous["plan"], cancel_event=cancel_event)
        return {
            "status": "changed",
            "previous_analysis_id": analysis_id,
            "previous_dataset_id": previous_dataset_id,
            "analysis_id": refreshed["analysis_id"],
            "dataset_id": replacement["dataset_id"],
            "result": refreshed,
        }

    @_synchronized_scope
    def bind_delivery(
        self,
        analysis_id: str,
        artifact_path: Path,
        *,
        kind: str,
        bound_elements: list[str],
        verification_evidence: str | None = None,
    ) -> dict[str, Any]:
        """Seal a file; only callers that reopened it may attest consistency."""

        self.get_result(analysis_id)
        if kind not in {"docx", "pptx", "svg", "evidence"}:
            raise AnalysisError("Unsupported analysis delivery kind")
        if not isinstance(bound_elements, list) or not bound_elements or len(bound_elements) > 100 or any(
            not isinstance(item, str) or not item or len(item) > 200 for item in bound_elements
        ):
            raise AnalysisError("Delivery needs bounded result element bindings")
        if verification_evidence is not None and (
            not isinstance(verification_evidence, str)
            or not verification_evidence.strip()
            or len(verification_evidence) > 300
        ):
            raise AnalysisError("Delivery verification evidence is invalid")
        artifact = Path(artifact_path).expanduser().resolve(strict=True)
        if not artifact.is_file() or artifact.stat().st_size > MAX_DELIVERY_BYTES:
            raise AnalysisError("Analysis delivery is missing or too large")
        content_hash = _digest(_read_bounded(artifact, limit=MAX_DELIVERY_BYTES, label="Analysis delivery"))
        identity = {
            "analysis_id": analysis_id,
            "artifact_path": str(artifact),
            "content_sha256": content_hash,
            "kind": kind,
        }
        delivery_id = _digest(_canonical(identity))
        record = {
            "schema_version": SCHEMA_VERSION,
            "delivery_id": delivery_id,
            **identity,
            "bound_elements": bound_elements,
            "verification": {
                "delivery_consistency": {
                    "status": "passed" if verification_evidence is not None else "pending",
                    **({"evidence": verification_evidence} if verification_evidence is not None else {"reason": "Final-file content was not checked against the bound result"}),
                },
                "semantic_review": {"status": "pending", "reason": "Business interpretation requires contextual review"},
            },
            "bound_at": datetime.now(timezone.utc).isoformat(),
        }
        path = self._record_path("deliveries", delivery_id)
        if not path.exists():
            try:
                return _atomic_json(path, record)
            except FileExistsError:
                pass
        return self._load("deliveries", delivery_id)

    def delivery_status(self, delivery_id: str) -> dict[str, Any]:
        """An edit never inherits the prior successful final-file verification."""

        record = self._load("deliveries", delivery_id)
        result = self.get_result(record["analysis_id"])
        artifact = Path(record["artifact_path"])
        try:
            if not artifact.is_file() or artifact.stat().st_size > MAX_DELIVERY_BYTES:
                status = "unavailable_unverified"
            else:
                status = "current" if _digest(_read_bounded(artifact, limit=MAX_DELIVERY_BYTES, label="Analysis delivery")) == record["content_sha256"] else "modified_unverified"
        except (OSError, PermissionError, AnalysisError):
            status = "unavailable_unverified"
        return {
            "delivery_id": delivery_id,
            "analysis_id": record["analysis_id"],
            "artifact_path": str(artifact),
            "kind": record["kind"],
            "content_sha256": record["content_sha256"],
            "status": status,
            "source_status": self.source_status(result["dataset_id"])["status"],
            "bound_elements": record["bound_elements"],
            "verification": record["verification"],
        }

    def get_delivery(self, delivery_id: str) -> dict[str, Any]:
        """Read a sealed delivery receipt without probing its source path."""

        return self._load("deliveries", delivery_id)

    def _clear_scope_records(self) -> dict[str, Any]:
        counts = {
            f"{kind}_removed": len(list((self.root / kind).glob("*.json")))
            for kind in ("datasets", "results", "deliveries")
        }
        counts["snapshots_removed"] = len(list((self.root / "snapshots").iterdir())) if (self.root / "snapshots").is_dir() else 0
        if self.root.exists():
            shutil.rmtree(self.root)
        return {**counts, "external_files_removed": False, "managed_records_recoverable": False, "source_files_untouched": True}

    @_synchronized_scope
    def clear(self) -> dict[str, Any]:
        """Forget this conversation's analysis while keeping the chat usable."""

        return self._clear_scope_records()

    def retire(self) -> dict[str, Any]:
        """Delete managed records and reject late writes from a deleted chat.

        The marker is outside the deleted scope so in-flight calls cannot
        repopulate it after the conversation deletion. Repeating retirement
        completes an interrupted cleanup without touching external files.
        """

        with _SCOPE_LOCKS_GUARD:
            lock = _SCOPE_LOCKS.setdefault(self.root, threading.RLock())
        with lock:
            self._assert_safe_scope(allow_retired=True)
            self.root.parent.mkdir(parents=True, exist_ok=True)
            if not self.retired_marker.exists():
                with tempfile.NamedTemporaryFile("wb", dir=self.root.parent, prefix=".retired-", delete=False) as stream:
                    temporary = Path(stream.name)
                    stream.write(b'{"analysis_scope_retired":true}')
                    stream.flush()
                    os.fsync(stream.fileno())
                try:
                    try:
                        os.link(temporary, self.retired_marker)
                    except FileExistsError:
                        pass
                finally:
                    temporary.unlink(missing_ok=True)
            return self._clear_scope_records()
