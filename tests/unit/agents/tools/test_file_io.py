# -*- coding: utf-8 -*-
from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook
from docx import Document
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

from lantai.agents.tools.utils import read_file_safe
from lantai.agents.tools.file_io import edit_file, read_file, write_file
import lantai.agents.tools.file_io as file_io_module
from lantai.config.context import (
    WorkspaceContext,
    set_current_recent_max_bytes,
    set_current_workspace_context,
    set_current_workspace_dir,
)
from lantai.constant import TRUNCATION_NOTICE_MARKER
from lantai.security.file_access_grants import (
    issue_external_file_grant,
    set_current_external_file_grants,
)
from lantai.security.tool_policy import _detect_distribution


def _issue_test_external_file_grant(path, *, operations):
    return issue_external_file_grant(
        path,
        operations=operations,
        security_scoped_bookmark=_detect_distribution() == "app_store",
    )


@pytest.fixture
def workspace(tmp_path):
    root = tmp_path / "workspace"
    root.mkdir()
    set_current_workspace_dir(root)
    try:
        yield root
    finally:
        set_current_recent_max_bytes(None)
        set_current_workspace_dir(None)
        set_current_external_file_grants(None)


@pytest.mark.asyncio
async def test_file_tools_allow_absolute_paths_inside_workspace(workspace):
    target = workspace / "notes.txt"

    write_response = await write_file(str(target), "hello")
    read_response = await read_file(str(target))

    assert "Wrote" in write_response.content[0]["text"]
    assert read_response.content[0]["text"] == "hello"


@pytest.mark.asyncio
async def test_write_file_creates_missing_parent_directories_inside_workspace(
    workspace,
):
    target = workspace / "coffee-revenue-dataset" / "CLAUDE.md"

    write_response = await write_file(str(target), "dataset constraints")
    read_response = await read_file(str(target))

    assert write_response.is_error is False
    assert "Wrote" in write_response.content[0]["text"]
    assert read_response.content[0]["text"] == "dataset constraints"


@pytest.mark.asyncio
async def test_read_file_recovers_workspace_directory_as_file_overview(
    workspace,
):
    (workspace / ".lantai").mkdir()
    (workspace / ".lantai" / "runtime.json").write_text(
        "internal",
        encoding="utf-8",
    )
    (workspace / "briefs").mkdir()
    (workspace / "briefs" / "launch.md").write_text(
        "launch brief",
        encoding="utf-8",
    )

    response = await read_file(str(workspace))
    text = response.content[0]["text"]

    assert response.is_error is False
    assert "directory, not a file" in text
    assert "briefs/launch.md" in text
    assert ".lantai" not in text
    assert "read_file" in text


@pytest.mark.asyncio
async def test_read_file_recovers_empty_workspace_directory_with_create_hint(
    workspace,
):
    (workspace / ".lantai-topic-workspace.json").write_text(
        "{}",
        encoding="utf-8",
    )

    response = await read_file(str(workspace))
    text = response.content[0]["text"]

    assert response.is_error is False
    assert "No user files" in text
    assert "write_file" in text
    assert "concrete file path" in text


@pytest.mark.asyncio
async def test_read_file_extracts_all_xlsx_sheets_without_external_python(
    tmp_path,
    monkeypatch,
) -> None:
    from lantai.app import office_documents

    workbook_path = tmp_path / "media-budget.xlsx"
    workbook = Workbook()
    summary = workbook.active
    summary.title = "总表"
    summary.append(["专项", "金额"])
    summary.append(["媒体", 300])
    detail = workbook.create_sheet("媒体合作明细")
    detail.append(["媒体名称", "折后价"])
    detail.append(["36氪", 30])
    workbook.save(workbook_path)

    monkeypatch.setattr(office_documents, "get_officecli_command", lambda: None)
    set_current_workspace_dir(tmp_path)
    try:
        response = await read_file(workbook_path.name)
    finally:
        set_current_workspace_dir(None)

    text = response.content[0]["text"]
    assert response.is_error is False
    assert "Sheet: 总表" in text
    assert "媒体\t300" in text
    assert "Sheet: 媒体合作明细" in text
    assert "36氪\t30" in text


@pytest.mark.asyncio
async def test_read_file_extracts_pdf_text_instead_of_binary_bytes(
    workspace,
) -> None:
    pdf_path = workspace / "report.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "Lantai readable PDF contract")
    pdf.save()

    response = await read_file(pdf_path.name)

    text = response.content[0]["text"]
    assert response.is_error is False
    assert "Lantai readable PDF contract" in text
    assert "%PDF" not in text
    assert response.metadata["source_format"] == "pdf"


@pytest.mark.asyncio
async def test_read_file_preserves_pdf_page_boundaries(workspace) -> None:
    pdf_path = workspace / "pages.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "First page contract")
    pdf.showPage()
    pdf.drawString(72, 720, "Second page contract")
    pdf.save()

    response = await read_file(pdf_path.name)
    text = response.content[0]["text"]

    assert "## Page 1" in text
    assert "First page contract" in text
    assert "## Page 2" in text
    assert "Second page contract" in text
    assert response.metadata["page_count"] == 2


@pytest.mark.asyncio
async def test_read_file_detects_pdf_magic_even_with_wrong_extension(workspace) -> None:
    disguised = workspace / "report.bin"
    pdf = canvas.Canvas(str(disguised))
    pdf.drawString(72, 720, "Magic signature contract")
    pdf.save()

    response = await read_file(disguised.name)

    assert "Magic signature contract" in response.content[0]["text"]
    assert "%PDF" not in response.content[0]["text"]
    assert response.metadata["source_format"] == "pdf"


@pytest.mark.asyncio
async def test_read_file_marks_image_only_pdf_as_scanned(workspace) -> None:
    pdf_path = workspace / "scan.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.rect(72, 600, 200, 100, fill=1)
    pdf.save()

    response = await read_file(pdf_path.name)

    assert response.is_error is False
    assert response.metadata["extraction_status"] == "no_text"
    assert response.metadata["scanned_document"] is True
    assert "OCR" in response.content[0]["text"]
    assert "%PDF" not in response.content[0]["text"]


@pytest.mark.asyncio
async def test_read_file_reports_encrypted_pdf_without_binary(workspace) -> None:
    plain = workspace / "plain.pdf"
    encrypted = workspace / "encrypted.pdf"
    pdf = canvas.Canvas(str(plain))
    pdf.drawString(72, 720, "Secret contract")
    pdf.save()
    reader = PdfReader(str(plain))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt("password")
    with encrypted.open("wb") as file:
        writer.write(file)

    response = await read_file(encrypted.name)

    assert response.is_error is True
    assert response.metadata["extraction_status"] == "encrypted"
    assert "加密" in response.content[0]["text"]
    assert "%PDF" not in response.content[0]["text"]


@pytest.mark.asyncio
async def test_read_file_reports_corrupt_pdf_without_object_stream(workspace) -> None:
    pdf_path = workspace / "damaged.pdf"
    pdf_path.write_bytes(b"%PDF-1.7\n1 0 obj\ntruncated")

    response = await read_file(pdf_path.name)

    assert response.is_error is True
    assert response.metadata["extraction_status"] == "corrupt"
    assert "损坏" in response.content[0]["text"]
    assert "1 0 obj" not in response.content[0]["text"]


@pytest.mark.asyncio
async def test_read_file_long_pdf_is_text_only_and_continuable(workspace) -> None:
    pdf_path = workspace / "long-report.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    for page_number in range(1, 106):
        pdf.drawString(72, 720, f"Long report page {page_number}")
        pdf.showPage()
    pdf.save()
    set_current_recent_max_bytes(800)

    response = await read_file(pdf_path.name)
    text = response.content[0]["text"]

    assert response.metadata["page_count"] == 105
    assert "Long report page 1" in text
    assert TRUNCATION_NOTICE_MARKER in text
    assert "start_line=" in text
    assert "%PDF" not in text


@pytest.mark.asyncio
async def test_read_file_preserves_docx_table_boundaries(workspace) -> None:
    docx_path = workspace / "contract.docx"
    document = Document()
    document.add_paragraph("合同摘要")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "金额"
    table.cell(1, 0).text = "年度服务"
    table.cell(1, 1).text = "120000"
    document.save(docx_path)

    response = await read_file(docx_path.name)

    text = response.content[0]["text"]
    assert "[Table 1]" in text
    assert "项目 | 金额" in text
    assert "年度服务 | 120000" in text


@pytest.mark.asyncio
async def test_read_file_preserves_pptx_slide_layout_and_speaker_notes(
    workspace,
) -> None:
    pptx_path = workspace / "briefing.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "季度汇报"
    slide.placeholders[1].text = "关键进展"
    slide.notes_slide.notes_text_frame.text = "演讲人备注：强调风险"
    presentation.save(pptx_path)

    response = await read_file(pptx_path.name)

    text = response.content[0]["text"]
    assert "[Slide 1]" in text
    assert "[Layout:" in text
    assert "季度汇报" in text
    assert "[Speaker notes]" in text
    assert "演讲人备注：强调风险" in text


@pytest.mark.asyncio
async def test_read_file_preserves_xlsx_formula_without_cached_value(
    workspace,
) -> None:
    workbook_path = workspace / "forecast.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["项目", "单价", "数量", "合计"])
    sheet.append(["服务", 20, 3, "=B2*C2"])
    workbook.save(workbook_path)

    response = await read_file(workbook_path.name)

    text = response.content[0]["text"]
    assert "=B2*C2" in text


@pytest.mark.asyncio
async def test_read_file_falls_back_to_agent_product_guide_from_topic_workspace(
    tmp_path,
):
    agent_home = tmp_path / "agent-home"
    topic_workspace = agent_home / "topic-workspaces" / "chat-1"
    topic_workspace.mkdir(parents=True)
    (agent_home / "PRODUCT_GUIDE.md").write_text(
        "澜台的设计哲学是将主动性交还给用户。",
        encoding="utf-8",
    )
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=topic_workspace,
            project_root=topic_workspace,
        ),
    )

    try:
        response = await read_file("PRODUCT_GUIDE.md")
    finally:
        set_current_workspace_context(None)

    assert response.content[0]["text"] == (
        "澜台的设计哲学是将主动性交还给用户。"
    )


@pytest.mark.asyncio
async def test_read_file_prefers_topic_product_guide_over_agent_fallback(
    tmp_path,
):
    agent_home = tmp_path / "agent-home"
    topic_workspace = agent_home / "topic-workspaces" / "chat-1"
    topic_workspace.mkdir(parents=True)
    (agent_home / "PRODUCT_GUIDE.md").write_text("agent", encoding="utf-8")
    (topic_workspace / "PRODUCT_GUIDE.md").write_text("topic", encoding="utf-8")
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=topic_workspace,
            project_root=topic_workspace,
        ),
    )

    try:
        response = await read_file("PRODUCT_GUIDE.md")
    finally:
        set_current_workspace_context(None)

    assert response.content[0]["text"] == "topic"


@pytest.mark.asyncio
async def test_read_file_does_not_fallback_to_other_agent_home_files(tmp_path):
    agent_home = tmp_path / "agent-home"
    topic_workspace = agent_home / "topic-workspaces" / "chat-1"
    topic_workspace.mkdir(parents=True)
    (agent_home / "SOUL.md").write_text("root-only persona", encoding="utf-8")
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=topic_workspace,
            project_root=topic_workspace,
        ),
    )

    try:
        response = await read_file("SOUL.md")
    finally:
        set_current_workspace_context(None)

    text = response.content[0]["text"]
    assert str(topic_workspace / "SOUL.md") in text
    assert "does not exist" in text
    assert "root-only persona" not in text


@pytest.mark.asyncio
async def test_read_file_redirects_topic_local_skill_paths_to_read_skill(
    tmp_path,
):
    agent_home = tmp_path / "agent-home"
    topic_workspace = agent_home / "topic-workspaces" / "chat-1"
    topic_workspace.mkdir(parents=True)
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=topic_workspace,
            project_root=topic_workspace,
        ),
    )

    try:
        response = await read_file("skills/pptx/SKILL.md")
    finally:
        set_current_workspace_context(None)

    text = response.content[0]["text"]
    assert response.is_error is True
    assert 'read_skill(skill_name="pptx", file_path="SKILL.md")' in text


@pytest.mark.asyncio
async def test_write_file_never_writes_through_agent_reference_fallback(
    tmp_path,
):
    agent_home = tmp_path / "agent-home"
    topic_workspace = agent_home / "topic-workspaces" / "chat-1"
    topic_workspace.mkdir(parents=True)
    agent_guide = agent_home / "PRODUCT_GUIDE.md"
    agent_guide.write_text("agent", encoding="utf-8")
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=topic_workspace,
            project_root=topic_workspace,
        ),
    )

    try:
        response = await write_file("PRODUCT_GUIDE.md", "topic")
    finally:
        set_current_workspace_context(None)

    assert "Wrote" in response.content[0]["text"]
    assert agent_guide.read_text(encoding="utf-8") == "agent"
    assert (topic_workspace / "PRODUCT_GUIDE.md").read_text(encoding="utf-8") == (
        "topic"
    )


@pytest.mark.asyncio
async def test_file_tools_reject_absolute_paths_outside_workspace(tmp_path, workspace):
    outside = tmp_path / "outside.md"
    outside.write_text("secret", encoding="utf-8")

    read_response = await read_file(str(outside))
    write_response = await write_file(str(outside), "overwrite")

    assert (
        "file path must stay inside workspace"
        in read_response.content[0]["text"]
    )
    assert (
        "file path must stay inside workspace"
        in write_response.content[0]["text"]
    )
    assert outside.read_text(encoding="utf-8") == "secret"


@pytest.mark.asyncio
async def test_read_file_explains_when_an_external_directory_was_selected(
    tmp_path,
    workspace,
):
    outside_directory = tmp_path / "another-topic-workspace"
    outside_directory.mkdir()

    response = await read_file(str(outside_directory))
    text = response.content[0]["text"]

    assert response.is_error is True
    assert "path is a directory" in text
    assert "glob_search" in text
    assert "external grant rejected" not in text


@pytest.mark.asyncio
async def test_file_tools_allow_external_absolute_read_with_grant(tmp_path, workspace):
    outside = tmp_path / "outside.txt"
    outside.write_text("shared", encoding="utf-8")
    grant = _issue_test_external_file_grant(outside, operations={"read"})
    set_current_external_file_grants([grant])

    response = await read_file(str(outside))

    assert response.content[0]["text"] == "shared"


@pytest.mark.asyncio
async def test_external_read_grant_does_not_allow_write(tmp_path, workspace):
    outside = tmp_path / "outside.txt"
    outside.write_text("shared", encoding="utf-8")
    grant = _issue_test_external_file_grant(outside, operations={"read"})
    set_current_external_file_grants([grant])

    response = await write_file(str(outside), "overwrite")

    assert "does not allow write access" in response.content[0]["text"]
    assert outside.read_text(encoding="utf-8") == "shared"


@pytest.mark.asyncio
async def test_external_write_grant_allows_edit_without_general_read(
    tmp_path,
    workspace,
):
    outside = tmp_path / "outside.md"
    outside.write_text("old", encoding="utf-8")
    grant = _issue_test_external_file_grant(outside, operations={"write"})
    set_current_external_file_grants([grant])

    read_response = await read_file(str(outside))
    edit_response = await edit_file(str(outside), "old", "new")

    assert "does not allow read access" in read_response.content[0]["text"]
    assert "Successfully replaced" in edit_response.content[0]["text"]
    assert outside.read_text(encoding="utf-8") == "new"


@pytest.mark.asyncio
async def test_external_file_grant_rejects_path_replacement(tmp_path, workspace):
    outside = tmp_path / "outside.txt"
    outside.write_text("first", encoding="utf-8")
    grant = _issue_test_external_file_grant(outside, operations={"read"})
    set_current_external_file_grants([grant])
    outside.unlink()
    outside.write_text("secret replacement", encoding="utf-8")

    response = await read_file(str(outside))
    text = response.content[0]["text"]

    assert "external file grant no longer matches the file" in text
    assert "secret replacement" not in text


def test_app_store_external_file_grant_requires_bookmark(
    monkeypatch,
    tmp_path,
):
    outside = tmp_path / "outside.txt"
    outside.write_text("shared", encoding="utf-8")
    monkeypatch.setenv("LANTAI_APP_STORE", "1")

    with pytest.raises(ValueError, match="security-scoped bookmark"):
        issue_external_file_grant(outside, operations={"read"})

    grant = issue_external_file_grant(
        outside,
        operations={"read"},
        security_scoped_bookmark=True,
    )

    assert grant.security_scoped_bookmark is True


@pytest.mark.asyncio
async def test_file_tools_reject_symlink_escape(tmp_path, workspace):
    outside = tmp_path / "outside.txt"
    outside.write_text("secret", encoding="utf-8")
    (workspace / "link.txt").symlink_to(outside)

    response = await read_file("link.txt")

    assert (
        "file path must stay inside workspace" in response.content[0]["text"]
    )


@pytest.mark.asyncio
async def test_read_file_rejects_symlink_swap_after_validation(
    monkeypatch,
    tmp_path,
    workspace,
):
    target = workspace / "notes.md"
    target.write_text("safe", encoding="utf-8")
    outside = tmp_path / "outside.md"
    outside.write_text("secret", encoding="utf-8")
    real_isfile = file_io_module.os.path.isfile
    swapped = False

    def swap_after_isfile(path):
        nonlocal swapped
        result = real_isfile(path)
        if Path(path) == target and result and not swapped:
            target.unlink()
            target.symlink_to(outside)
            swapped = True
        return result

    monkeypatch.setattr(file_io_module.os.path, "isfile", swap_after_isfile)

    response = await read_file(str(target))
    text = response.content[0]["text"]

    assert "symlink" in text or "regular file" in text
    assert "secret" not in text


@pytest.mark.asyncio
async def test_read_file_safe_enforces_byte_budget_for_multibyte_text(
    tmp_path,
):
    path = tmp_path / "multibyte.txt"
    path.write_text("é" * 4, encoding="utf-8")

    content = await read_file_safe(str(path), max_bytes=5)

    assert len(content.encode("utf-8")) <= 5


@pytest.mark.asyncio
async def test_read_file_long_single_line_uses_byte_continuation(workspace):
    path = workspace / "long.txt"
    path.write_text("a" * 25, encoding="utf-8")
    set_current_recent_max_bytes(10)

    first_response = await read_file(str(path))
    first_text = first_response.content[0]["text"]

    assert first_text.startswith("a" * 10)
    assert TRUNCATION_NOTICE_MARKER in first_text
    assert "start_byte=10" in first_text

    second_response = await read_file(str(path), start_byte=10)
    second_text = second_response.content[0]["text"]

    assert second_text.startswith("a" * 10)
    assert "start_byte=20" in second_text


@pytest.mark.asyncio
async def test_write_file_replace_failure_keeps_original_file(monkeypatch, workspace):
    path = workspace / "notes.txt"
    path.write_text("old", encoding="utf-8")

    def fail_replace(src, dst):  # noqa: ARG001
        raise OSError("replace failed")

    monkeypatch.setattr(file_io_module.os, "replace", fail_replace)

    response = await write_file(str(path), "new")

    assert "replace failed" in response.content[0]["text"]
    assert path.read_text(encoding="utf-8") == "old"
    assert not list(workspace.glob(".notes.txt.*.tmp"))


@pytest.mark.asyncio
async def test_edit_file_rejects_stale_expected_sha256(workspace):
    path = workspace / "notes.md"
    path.write_text("old", encoding="utf-8")

    response = await edit_file(
        str(path),
        old_text="old",
        new_text="new",
        expected_sha256="0" * 64,
    )

    text = response.content[0]["text"]
    assert "hash mismatch" in text
    assert path.read_text(encoding="utf-8") == "old"


@pytest.mark.asyncio
async def test_edit_file_accepts_matching_expected_sha256(workspace):
    path = workspace / "notes.md"
    path.write_text("old", encoding="utf-8")
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    response = await edit_file(
        str(path),
        old_text="old",
        new_text="new",
        expected_sha256=digest,
    )

    assert "Successfully replaced" in response.content[0]["text"]
    assert path.read_text(encoding="utf-8") == "new"
