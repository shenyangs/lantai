from __future__ import annotations

import asyncio
import hashlib
import json
import threading
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from jsonschema import Draft202012Validator
from jsonschema.exceptions import ValidationError as JsonSchemaValidationError
from PIL import Image
from pptx import Presentation

from lantai.agents.tools.image_generation import (
    ImageSource,
    _generated_image_filename,
)
from lantai.agents.tools.presentation import (
    CREATE_PRESENTATION_INPUT_SCHEMA,
    _build_authored_storyboard_plan,
    _visual_verification_attestation,
    create_presentation,
    finalize_presentation_generation,
    get_presentation_generation_status,
    retry_presentation_generation,
)
from lantai.config.context import (
    WorkspaceContext,
    set_current_workspace_context,
)
from lantai.exceptions import ProviderError
from lantai.presentation.jobs import (
    PresentationJobRecord,
    PresentationJobStore,
    cancel_active_jobs_for_runtime_run,
)
from lantai.presentation.models import (
    ChartObject,
    DeckIR,
    DeckSlide,
    InputMaterial,
    JobStage,
    JobStatus,
    ObjectFrame,
    PresentationJobRequest,
    PresentationStageMetric,
    VisualCriticResult,
    VisualIssue,
    VisualReviewCheckpoint,
)
from lantai.presentation.render import (
    PresentationRenderUnavailable,
    RenderResult,
)
from lantai.presentation.service import PresentationService
from lantai.presentation.theme import normalize_theme_id
from lantai.runtime.effect_context import (
    RuntimeEffectContext,
    bind_runtime_effect_context,
)
from lantai.runtime.tool_response import TextBlock, ToolResponse


class _PassingVisualCritic:
    def review(self, _slides):
        return VisualCriticResult(critic_id="test-visual-critic", issues=[])


class _DocumentLikeVisualCritic:
    def review(self, slides):
        return VisualCriticResult(
            critic_id="test-document-like-critic",
            issues=[
                VisualIssue(
                    slide_id=slides[0].slide_id,
                    severity="error",
                    issue_type="aesthetic_staleness",
                    description=(
                        "The slide resembles a Word document export rather "
                        "than a presentation with deliberate visual hierarchy."
                    ),
                    proposed_action="Redesign the slide before delivery.",
                )
            ],
        )


@pytest.fixture(autouse=True)
def _stub_presentation_intelligence(monkeypatch) -> None:
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.UtilityPresentationVisualCritic",
        lambda _agent_id: _PassingVisualCritic(),
    )


def _context(tmp_path: Path) -> WorkspaceContext:
    project_root = tmp_path / "workspace"
    state_root = tmp_path / "state"
    agent_home = tmp_path / "agent-home"
    project_root.mkdir()
    state_root.mkdir()
    agent_home.mkdir()
    return WorkspaceContext(
        agent_home=agent_home,
        state_root=state_root,
        project_root=project_root,
        workspace_grant_id="conversation-ppt",
        trust_state="trusted",
    )


def _slides() -> list[dict[str, object]]:
    return [
        {
            "layout_id": "cover",
            "objects": [
                {
                    "kind": "text",
                    "role": "title",
                    "text": "结构化演示文稿",
                    "slot": "title",
                },
                {
                    "kind": "text",
                    "role": "body",
                    "text": "可编辑、可重开、可核验",
                    "slot": "subtitle",
                },
            ],
        }
    ]


def _render_one(_source, output_dir, *, timeout=120) -> RenderResult:
    _ = timeout
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"pdf")
    image_path = destination / "slide-1.png"
    image = Image.new("RGB", (640, 360), "white")
    image.paste("black", (0, 0, 320, 360))
    image.save(image_path)
    return RenderResult(pdf_path=pdf_path, slide_images=(image_path,))


def test_visual_verification_attestation_seals_receipt_bytes(tmp_path: Path) -> None:
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"deck")
    preview = tmp_path / "slide-1.png"
    Image.new("RGB", (64, 36), "white").save(preview)
    quality = {
        "schema_version": "1.0",
        "passed": True,
        "delivery_status": "passed",
        "render_verified": True,
        "vision_critique_required": True,
        "visual_verified": True,
        "render_evidence": "passed",
        "visual_evidence": "passed",
        "visual_critic_id": "test-visual-critic",
        "rendered_slide_count": 1,
    }
    receipt = tmp_path / "receipt.json"
    receipt.write_text(
        json.dumps(
            {
                "output_path": str(deck),
                "output_sha256": hashlib.sha256(deck.read_bytes()).hexdigest(),
                "render_evidence": [{"slide_index": 1, "path": str(preview)}],
            }
        ),
        encoding="utf-8",
    )

    attestation = _visual_verification_attestation(deck, receipt, quality)

    assert attestation["receipt_sha256"] == (
        "sha256:" + hashlib.sha256(receipt.read_bytes()).hexdigest()
    )


def test_generated_image_without_remote_suffix_gets_mime_extension() -> None:
    filename = _generated_image_filename(
        ImageSource(
            "url", "https://example.test/generated?id=42", "image/webp"
        ),
        index=2,
    )

    assert filename.endswith("_generated.webp")


def test_model_authored_presentation_theme_is_safely_normalized() -> None:
    assert normalize_theme_id("corporate_blue") == "lantai-ink"
    assert normalize_theme_id("future-brand-theme") == "lantai-ink"
    assert normalize_theme_id("lantai-slate") == "lantai-slate"


def test_create_presentation_schema_rejects_slots_outside_selected_layout() -> (
    None
):
    validator = Draft202012Validator(CREATE_PRESENTATION_INPUT_SCHEMA)
    invalid = {
        "title": "版式契约",
        "slides": [
            {
                "layout_id": "key_judgment",
                "objects": [
                    {
                        "kind": "text",
                        "role": "body",
                        "text": "不应进入渲染阶段",
                        "slot": "point_1",
                    }
                ],
            }
        ],
    }

    with pytest.raises(JsonSchemaValidationError):
        validator.validate(invalid)


def test_create_presentation_schema_requires_frames_for_freeform_objects() -> (
    None
):
    validator = Draft202012Validator(CREATE_PRESENTATION_INPUT_SCHEMA)
    invalid = {
        "title": "自由画布契约",
        "slides": [
            {
                "layout_id": "freeform",
                "objects": [
                    {
                        "kind": "text",
                        "role": "title",
                        "text": "没有边界的对象",
                        "slot": "title",
                    }
                ],
            }
        ],
    }

    with pytest.raises(JsonSchemaValidationError):
        validator.validate(invalid)


@pytest.mark.parametrize(
    "partial",
    [
        {"narrative_arc": "先给判断，再给证据。"},
        {
            "storyboard": [
                {
                    "purpose": "给出判断",
                    "key_message": "证据支持判断",
                    "evidence": [],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [],
                }
            ]
        },
        {"source_paths": ["/trusted/topic/attachments/source.pdf"]},
    ],
)
def test_create_presentation_schema_rejects_incomplete_source_story_contract(
    partial,
) -> None:
    validator = Draft202012Validator(CREATE_PRESENTATION_INPUT_SCHEMA)
    payload = {
        "title": "来源契约",
        "slides": _slides(),
        **partial,
    }

    with pytest.raises(JsonSchemaValidationError):
        validator.validate(payload)


def test_create_presentation_schema_tells_model_to_use_trusted_canonical_path(
) -> None:
    source_description = CREATE_PRESENTATION_INPUT_SCHEMA["properties"][
        "source_paths"
    ]["description"]

    assert "canonical_path" in source_description
    assert "display filename" in source_description
    assert "do not search" in source_description


def test_storyboard_schema_explains_exact_deck_fact_binding() -> None:
    properties = CREATE_PRESENTATION_INPUT_SCHEMA["properties"]["storyboard"][
        "items"
    ]["properties"]
    for page_field, deck_field in (
        ("required_facts", "required_facts"),
        ("immutable_facts", "immutable_data"),
    ):
        description = properties[page_field].get("description", "")
        assert "exact" in description
        assert "top-level " + deck_field in description
        assert "evidence" in description


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "invalid_field", ["required_facts", "immutable_facts"]
)
@pytest.mark.parametrize("with_source", [False, True])
async def test_presentation_fact_preflight_rejection_is_safe_to_correct(
    tmp_path,
    monkeypatch,
    invalid_field,
    with_source,
) -> None:
    def should_not_create_job(*_args, **_kwargs):
        raise AssertionError("fact rejection must precede any durable job")

    def should_not_extract_material(*_args, **_kwargs):
        raise AssertionError("fact rejection must precede source extraction")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.create_job",
        should_not_create_job,
    )
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.build_brief",
        should_not_extract_material,
    )
    context = _context(tmp_path)
    source_paths = []
    if with_source:
        source = context.project_root / "source.txt"
        source.write_text("可编辑的演示文稿资料", encoding="utf-8")
        source_paths = [str(source)]
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="结构化演示文稿",
            slides=_slides(),
            required_facts=["可编辑"],
            immutable_data=["一页"],
            source_paths=source_paths,
            narrative_arc="从主题到诗句",
            storyboard=[
                {
                    "purpose": "呈现主题",
                    "key_message": "结构化演示文稿",
                    "evidence": ["自由创作的诗句"],
                    "relationship_type": "thesis",
                    "visual_role": "cover",
                    "source_indexes": [0] if with_source else [],
                    invalid_field: ["不是全稿事实的诗句"],
                },
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["side_effect_outcome"] == "known_not_applied"
    assert response.metadata["error_code"] == "invalid_presentation_arguments"
    assert "exact" in response.content[0]["text"]
    assert "evidence" in response.content[0]["text"]
    assert not (context.state_root / "presentation-jobs").exists()
    assert not (context.project_root / "presentation.pptx").exists()


@pytest.mark.asyncio
async def test_presentation_extraction_failure_cannot_claim_unapplied_effects(
    tmp_path,
    monkeypatch,
) -> None:
    context = _context(tmp_path)
    marker = context.state_root / "extraction-started"

    def extract_and_fail(_request):
        marker.write_text(
            "conversion or extraction has started", encoding="utf-8"
        )
        raise ValueError("source extraction failed after starting")

    def should_not_create_job(*_args, **_kwargs):
        raise AssertionError("extraction failure must precede the durable job")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.build_brief",
        extract_and_fail,
    )
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.create_job",
        should_not_create_job,
    )
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="结构化演示文稿", slides=_slides()
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert marker.is_file()
    assert response.metadata.get("side_effect_outcome") != "known_not_applied"
    assert not (context.state_root / "presentation-jobs").exists()


@pytest.mark.asyncio
async def test_create_presentation_rejects_layout_contract_before_job_start(
    tmp_path,
    monkeypatch,
) -> None:
    def should_not_create_job(*_args, **_kwargs):
        raise AssertionError("invalid layout input must not start a job")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.create_job",
        should_not_create_job,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="版式契约",
            slides=[
                {
                    "layout_id": "key_judgment",
                    "objects": [
                        {
                            "kind": "text",
                            "role": "body",
                            "text": "不应进入渲染阶段",
                            "slot": "point_1",
                        }
                    ],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata == {
        "error_code": "invalid_presentation_arguments",
        "deterministic_failure": True,
        "side_effect_outcome": "known_not_applied",
    }
    assert "slide 1 (key_judgment)" in response.content[0]["text"]
    assert "slot 'point_1' is not defined" in response.content[0]["text"]


@pytest.mark.asyncio
async def test_create_presentation_delivers_verified_pptx_idempotently(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    set_current_workspace_context(_context(tmp_path))
    try:
        first = await create_presentation(
            title="结构化演示文稿",
            slides=_slides(),
            output_filename="deck.pptx",
        )
        second = await create_presentation(
            title="结构化演示文稿",
            slides=_slides(),
            output_filename="deck.pptx",
        )
    finally:
        set_current_workspace_context(None)

    output = tmp_path / "workspace" / "deck.pptx"
    assert first.is_error is False
    assert first.content[0]["type"] == "file"
    assert first.metadata["verification_status"] == "visual_verified"
    assert first.metadata["reused"] is False
    assert second.metadata["reused"] is True
    assert second.metadata["job_id"] == first.metadata["job_id"]
    assert len(render_calls) == 1
    assert output.is_file()
    assert len(Presentation(output).slides) == 1


@pytest.mark.asyncio
async def test_create_presentation_preserves_unowned_destination_before_generation(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    context = _context(tmp_path)
    destination = context.project_root / "existing-author-work.pptx"
    destination.write_bytes(b"author-owned-pptx-bytes")
    original = destination.read_bytes()
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="不能覆盖已有作品",
            slides=_slides(),
            output_filename=destination.name,
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_delivery_conflict"
    assert response.metadata["side_effect_outcome"] == "none"
    assert response.metadata["publish_ready"] is False
    assert response.metadata["requires_publish"] is False
    assert destination.read_bytes() == original
    assert render_calls == []
    jobs_root = context.state_root / "presentation-jobs"
    assert not jobs_root.exists()


@pytest.mark.asyncio
async def test_create_presentation_does_not_replace_newer_author_save(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    context = _context(tmp_path)
    destination = context.project_root / "author-latest.pptx"
    set_current_workspace_context(context)
    try:
        first = await create_presentation(
            title="作者初版",
            slides=_slides(),
            output_filename=destination.name,
        )
        destination.write_bytes(destination.read_bytes() + b"author-save")
        authored_bytes = destination.read_bytes()

        revised = await create_presentation(
            title="模型局部修改",
            slides=_slides(),
            output_filename=destination.name,
        )
    finally:
        set_current_workspace_context(None)

    assert first.is_error is False
    assert revised.is_error is True
    assert revised.metadata["error_code"] == "presentation_delivery_conflict"
    assert revised.metadata["side_effect_outcome"] == "none"
    assert revised.metadata["publish_ready"] is False
    assert revised.metadata["requires_publish"] is False
    assert destination.read_bytes() == authored_bytes
    assert len(render_calls) == 1
    store = PresentationJobStore(context.state_root / "presentation-jobs")
    assert len(store.list_records(context.workspace_grant_id)) == 1


@pytest.mark.asyncio
async def test_presentation_job_reuses_native_runtime_owner(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    workspace_context = _context(tmp_path)
    set_current_workspace_context(workspace_context)
    try:
        with bind_runtime_effect_context(
            RuntimeEffectContext(
                run_id="native-run-ppt-1",
                attempt_id="native-run-ppt-1:attempt-1",
                effect_id="effect-create-presentation",
                state_workspace_dir=str(workspace_context.state_root),
            )
        ):
            response = await create_presentation(
                title="Native owner 复用",
                slides=_slides(),
                output_filename="native-owner.pptx",
            )
    finally:
        set_current_workspace_context(None)

    record = PresentationService(
        workspace_context.state_root / "presentation-jobs"
    ).store.get(response.metadata["job_id"])
    assert record.runtime_run_id == "native-run-ppt-1"


@pytest.mark.asyncio
async def test_create_presentation_rebuilds_when_image_content_changes(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    context = _context(tmp_path)
    image_path = context.project_root / "hero.png"
    Image.new("RGB", (320, 180), "red").save(image_path)
    slides = [
        {
            "layout_id": "image_focus",
            "objects": [
                {
                    "kind": "text",
                    "role": "title",
                    "text": "图片内容可追踪",
                    "slot": "title",
                },
                {
                    "kind": "image",
                    "path": "hero.png",
                    "alt": "主视觉",
                    "slot": "hero",
                },
            ],
        }
    ]
    set_current_workspace_context(context)
    try:
        first = await create_presentation(title="图片指纹", slides=slides)
        Image.new("RGB", (320, 180), "blue").save(image_path)
        second = await create_presentation(title="图片指纹", slides=slides)
    finally:
        set_current_workspace_context(None)

    assert first.is_error is False
    assert second.is_error is False
    assert second.metadata["reused"] is False
    assert second.metadata["job_id"] != first.metadata["job_id"]
    assert len(render_calls) == 2


@pytest.mark.asyncio
async def test_create_presentation_promotes_semantic_slides_to_designed_deck(
    tmp_path,
    monkeypatch,
) -> None:
    captured: list[object] = []

    def capture_compose(_service, _job_id, deck, *, runtime_run_id=None):
        assert runtime_run_id is None
        captured.append(deck)
        raise RuntimeError("stop after design boundary")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        capture_compose,
    )
    set_current_workspace_context(_context(tmp_path))
    try:
        response = await create_presentation(
            title="设计编排不可旁路",
            slides=_slides(),
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert len(captured) == 1
    deck = captured[0]
    assert deck.design_context is not None
    assert deck.design_profile is not None
    assert deck.composition_tokens is not None
    assert deck.visual_plan is not None
    assert deck.slides[0].layout_id == "freeform"
    assert all(obj.frame is not None for obj in deck.slides[0].objects)


@pytest.mark.asyncio
async def test_create_presentation_reports_durable_provider_recovery_route(
    tmp_path,
    monkeypatch,
) -> None:
    def fail_compose(_service, _job_id, _deck, *, runtime_run_id=None):
        assert runtime_run_id is None
        raise ProviderError("critic provider is unavailable")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        fail_compose,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="模型故障按阶段恢复",
            slides=_slides(),
        )
        status = await get_presentation_generation_status(
            response.metadata["job_id"]
        )
        retried = await retry_presentation_generation(
            response.metadata["job_id"]
        )
        service = PresentationService(context.state_root / "presentation-jobs")
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata.get("side_effect_outcome") != "known_not_applied"
    assert response.metadata["failure_categories"] == ["model_service"]
    assert response.metadata["recovery_actions"] == ["retry_model_service"]
    assert response.metadata["automatic_retry_allowed"] is True
    assert response.metadata["requires_human"] is False
    assert status.metadata["failure_routes"] == response.metadata[
        "failure_routes"
    ]
    assert status.metadata["failure_routes"][0]["retry_stage"] == "compose"
    assert retried.metadata["job_id"] == response.metadata["job_id"]
    assert retried.metadata["recovery_attempt_count"] == 1
    assert len(service.store.list_records("conversation-ppt")) == 1
    assert not (tmp_path / "presentation.pptx").exists()


@pytest.mark.asyncio
async def test_create_presentation_keeps_explicit_stop_out_of_failure_routes(
    tmp_path,
    monkeypatch,
) -> None:
    def cancel_compose(service, job_id, _deck, *, runtime_run_id=None):
        assert runtime_run_id is None
        record = service.store.get(job_id)
        service.store.cancel(job_id, record.conversation_id)
        raise RuntimeError("worker observed explicit stop")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        cancel_compose,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="用户停止保持独立状态",
            slides=_slides(),
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_job_cancelled"
    assert response.metadata["job_status"] == "cancelled"
    assert "failure_routes" not in response.metadata


@pytest.mark.asyncio
async def test_create_presentation_passes_authored_storyboard_to_visual_review(
    tmp_path,
    monkeypatch,
) -> None:
    captured: list[tuple[object, object]] = []
    design_plans = []
    context = _context(tmp_path)
    source = context.project_root / "source.md"
    source.write_text("任务状态与来源核对必须分开", encoding="utf-8")
    authored_direction = {
        "visual_concept": "新品发布会风格，突出一个主张",
        "palette": ["#10131A", "#F8FAFC", "#C73542"],
        "typography": "Noto Sans SC / Noto Sans SC",
        "density": "low",
        "image_language": "真实产品特写",
        "chart_language": "原生图表突出结论",
        "shape_language": "简洁几何",
        "grid": "12 columns",
        "recurring_motif": "页码和细线",
    }

    def capture_compose(service, job_id, deck, *, runtime_run_id=None):
        from lantai.presentation.models import DeckPlan, PresentationBrief

        brief_path = service.store.checkpoint_path(job_id, "brief.json")
        brief_checkpoint = PresentationBrief.model_validate_json(
            brief_path.read_text(encoding="utf-8")
        )
        assert brief_checkpoint.source_coverage[0].material_id == deck.source_refs[0]
        assert brief_checkpoint.source_coverage[0].source_sha256 == hashlib.sha256(
            source.read_bytes()
        ).hexdigest()
        plan_path = service.store.checkpoint_path(job_id, "deck-plan.json")
        plan = DeckPlan.model_validate_json(plan_path.read_text(encoding="utf-8"))
        captured.append((deck, plan))
        raise RuntimeError("stop after storyboard boundary")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        capture_compose,
    )
    from lantai.presentation.design_intelligence import (
        compose_deck_ir_with_design_context as compose_with_context,
    )

    def capture_design_plan(deck, request, brief, *, plan=None, author_direction=None):
        design_plans.append(plan)
        return compose_with_context(
            deck, request, brief, plan=plan,
            author_direction=author_direction,
        )

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.compose_deck_ir_with_design_context",
        capture_design_plan,
    )
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="有判断的复盘",
            brief="给业务负责人解释任务完成与数据可用不是同一件事。",
            slides=_slides(),
            source_paths=["source.md"],
            visual_direction=authored_direction,
            narrative_arc="先交代判断，再展示证据，最后给出处置。",
            storyboard=[
                {
                    "purpose": "说明核心判断",
                    "key_message": "完成任务不代表数据已经可用于决策",
                    "evidence": ["任务状态与来源核对必须分开"],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert len(captured) == 1
    deck, plan = captured[0]
    assert deck.slides[0].intent_id == plan.slides[0].intent.id
    assert plan.slides[0].intent.key_message == "完成任务不代表数据已经可用于决策"
    assert plan.slides[0].intent.source_refs == deck.source_refs
    assert plan.narrative_arc == "先交代判断，再展示证据，最后给出处置。"
    assert plan.design_context == deck.design_context
    assert deck.design_profile.identity.startswith("autonomous:product-launch:")
    assert deck.composition_tokens.accent == "C73542"
    assert deck.visual_direction.model_dump() == authored_direction
    assert len(design_plans) == 1
    assert design_plans[0].slides[0].intent.id == deck.slides[0].intent_id


@pytest.mark.asyncio
async def test_create_presentation_records_docx_source_coverage(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    source = context.project_root / "source.docx"
    document = Document()
    document.add_heading("来源说明", level=1)
    document.add_paragraph("仅 11 位受访者，不应外推。")
    document.save(source)
    slides = _slides()
    slides[0]["objects"][1]["text"] = "仅 11 位受访者，不应外推。"

    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="材料追溯",
            slides=slides,
            source_paths=["source.docx"],
            required_facts=["仅 11 位受访者，不应外推。"],
            narrative_arc="先交代样本边界，再给出结论。",
            storyboard=[
                {
                    "purpose": "交代样本边界",
                    "key_message": "仅 11 位受访者，不应外推。",
                    "evidence": ["仅 11 位受访者，不应外推。"],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                    "required_facts": ["仅 11 位受访者，不应外推。"],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False, response.content
    service = PresentationService(context.state_root / "presentation-jobs")
    record = service.store.get(response.metadata["job_id"])
    brief_path = service.store.checkpoint_path(record.job_id, "brief.json")
    assert record.checkpoints["brief"] == str(brief_path)
    saved = json.loads(brief_path.read_text(encoding="utf-8"))
    assert saved["source_coverage"][0]["method"] == "docx_ooxml"
    assert saved["source_coverage"][0]["source_sha256"] == hashlib.sha256(
        source.read_bytes()
    ).hexdigest()


@pytest.mark.asyncio
async def test_create_presentation_does_not_deliver_unverified_legacy_doc(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    source = context.project_root / "legacy.doc"
    source.write_bytes(b"legacy binary input")
    converted = tmp_path / "converted.docx"
    document = Document()
    document.add_paragraph("旧稿中的关键事实 28。")
    document.save(converted)

    def fake_convert(input_path, extension, output_dir):
        assert input_path.name == source.name
        assert input_path.read_bytes() == source.read_bytes()
        assert input_path != source  # Conversion consumes the immutable snapshot.
        assert extension == "docx"
        assert output_dir
        return SimpleNamespace(output_path=converted)

    def fake_validate(input_path):
        assert input_path.name == source.name
        assert input_path.read_bytes() == source.read_bytes()
        assert input_path != source

    monkeypatch.setattr(
        "lantai.app.office_conversion.convert_office_document", fake_convert
    )
    monkeypatch.setattr(
        "lantai.presentation.material_extraction._validate_legacy_doc",
        fake_validate,
    )
    slides = _slides()
    slides[0]["objects"][1]["text"] = "旧稿中的关键事实 28。"

    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="旧稿来源门禁",
            slides=slides,
            source_paths=["legacy.doc"],
            output_filename="legacy-result.pptx",
            required_facts=["旧稿中的关键事实 28。"],
            narrative_arc="先给出旧稿事实，再说明转换核验状态。",
            storyboard=[
                {
                    "purpose": "交代旧稿事实",
                    "key_message": "旧稿中的关键事实 28。",
                    "evidence": ["旧稿中的关键事实 28。"],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                    "required_facts": ["旧稿中的关键事实 28。"],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_quality_failed"
    assert "material.conversion_unverified" in response.content[0]["text"]
    assert not any(block["type"] == "file" for block in response.content)
    assert not (context.project_root / "legacy-result.pptx").exists()
    service = PresentationService(context.state_root / "presentation-jobs")
    record = service.store.get(response.metadata["job_id"])
    assert record.status.value == "failed"
    brief_path = service.store.checkpoint_path(record.job_id, "brief.json")
    assert json.loads(brief_path.read_text(encoding="utf-8"))[
        "source_coverage"
    ][0]["status"] == "converted_doc_unverified"


@pytest.mark.asyncio
async def test_create_presentation_rejects_source_without_authored_storyboard_early(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    source = context.project_root / "source.docx"
    document = Document()
    document.add_paragraph("样本只有 11 人。")
    document.save(source)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="材料汇报",
            slides=_slides(),
            source_paths=["source.docx"],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert "source-backed presentation requires" in response.content[0]["text"]
    service = PresentationService(context.state_root / "presentation-jobs")
    assert service.store.list_records(context.workspace_grant_id) == []


@pytest.mark.asyncio
async def test_create_presentation_missing_source_error_requires_canonical_path(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="材料路径契约",
            slides=_slides(),
            source_paths=["display-name-only.pdf"],
            narrative_arc="先给判断，再给证据。",
            storyboard=[
                {
                    "purpose": "说明来源",
                    "key_message": "只使用可信来源",
                    "evidence": [],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "invalid_presentation_arguments"
    assert "canonical_path" in response.content[0]["text"]
    assert "display filename" in response.content[0]["text"]


@pytest.mark.asyncio
async def test_create_presentation_rejects_storyboard_evidence_absent_from_docx(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    source = context.project_root / "source.docx"
    document = Document()
    document.add_paragraph("实际样本只有 11 人。")
    document.save(source)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="虚构证据不得入册",
            slides=_slides(),
            source_paths=["source.docx"],
            required_facts=["样本有 1000 人。"],
            narrative_arc="先给判断，再给证据。",
            storyboard=[
                {
                    "purpose": "解释样本",
                    "key_message": "样本足够大",
                    "evidence": ["样本有 1000 人。"],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                    "required_facts": ["样本有 1000 人。"],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "invalid_presentation_arguments"
    assert "not present in its declared source" in response.content[0]["text"]
    service = PresentationService(context.state_root / "presentation-jobs")
    assert service.store.list_records(context.workspace_grant_id) == []


@pytest.mark.asyncio
async def test_create_presentation_requires_source_binding_for_immutable_data(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    source = context.project_root / "source.docx"
    document = Document()
    document.add_paragraph("基准值为 11。")
    document.save(source)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="不可改写的数据",
            slides=_slides(),
            source_paths=["source.docx"],
            immutable_data=["11"],
            narrative_arc="先交代基准值，再解释变化。",
            storyboard=[
                {
                    "purpose": "交代基准值",
                    "key_message": "基准值为 11",
                    "evidence": ["基准值为 11。"],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [0],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert "immutable facts lack source-bound" in response.content[0]["text"]
    service = PresentationService(context.state_root / "presentation-jobs")
    assert service.store.list_records(context.workspace_grant_id) == []


@pytest.mark.asyncio
async def test_create_presentation_critic_receives_authored_claim_and_global_facts(
    tmp_path,
    monkeypatch,
) -> None:
    captured = []

    class CapturingVisualCritic:
        def review(self, slides):
            captured.extend(slides)
            return VisualCriticResult(critic_id="capturing-critic", issues=[])

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.UtilityPresentationVisualCritic",
        lambda _agent_id: CapturingVisualCritic(),
    )
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    authored_slides = _slides()
    authored_slides[0]["objects"][1]["text"] = "全局事实不得丢失"
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="经编辑的故事板",
            slides=authored_slides,
            required_facts=["全局事实不得丢失"],
            narrative_arc="先给判断，再限定证据边界。",
            storyboard=[
                {
                    "purpose": "交代结论",
                    "key_message": "这页主张来自故事板而非页标题",
                    "evidence": [],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": [],
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False
    assert len(captured) == 1
    assert captured[0].key_message == "这页主张来自故事板而非页标题"
    assert captured[0].intent.key_message == captured[0].key_message
    assert "全局事实不得丢失" in captured[0].factual_constraints
    assert captured[0].full_deck_context == [
        "Narrative arc: 先给判断，再限定证据边界。",
        "1. 这页主张来自故事板而非页标题",
    ]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("source_indexes", "page_facts", "expected"),
    [
        ([0], [], "invalid source_indexes"),
        ([], ["无来源的新事实"], "unknown required_facts"),
    ],
)
async def test_create_presentation_rejects_unbound_storyboard_evidence(
    tmp_path,
    source_indexes,
    page_facts,
    expected,
) -> None:
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="来源边界",
            slides=_slides(),
            narrative_arc="一页判断",
            storyboard=[
                {
                    "purpose": "判断",
                    "key_message": "只允许已有事实",
                    "evidence": [],
                    "relationship_type": "thesis",
                    "visual_role": "thesis",
                    "source_indexes": source_indexes,
                    "required_facts": page_facts,
                }
            ],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert expected in response.content[0]["text"]
    assert PresentationService(
        context.state_root / "presentation-jobs"
    ).store.list_records(context.workspace_grant_id) == []


def test_authored_storyboard_does_not_limit_native_charts_to_planner_quota() -> None:
    charts = [
        ChartObject(
            id=f"chart-{index}",
            categories=["A", "B"],
            series={"值": [index, index + 1]},
            frame=ObjectFrame(x=0.5 + index * 4, y=1, w=3.5, h=3),
        )
        for index in range(3)
    ]
    deck = DeckIR(
        title="三项证据",
        slides=[
            DeckSlide(
                id="slide-1",
                layout_id="freeform",
                composition_authority="author",
                objects=charts,
            )
        ],
    )
    bound_deck, plan = _build_authored_storyboard_plan(
        deck,
        [],
        [{
            "purpose": "分别展示三组数据",
            "key_message": "三组数据共同支撑判断",
            "evidence": [],
            "relationship_type": "data_trend",
            "visual_role": "evidence",
            "source_indexes": [],
        }],
        "按三组数据作判断",
    )

    assert plan is not None
    assert plan.slides[0].charts == []
    assert [obj.id for obj in bound_deck.slides[0].objects] == [
        "chart-0", "chart-1", "chart-2"
    ]


def test_authored_storyboard_requires_every_input_material_to_have_page_destination() -> None:
    deck = DeckIR(
        title="素材覆盖", slides=[DeckSlide(id="slide-1", layout_id="freeform")]
    )
    materials = [
        InputMaterial(id="source-a", name="A", content="事实 A"),
        InputMaterial(id="source-b", name="B", content="事实 B"),
    ]
    storyboard = [{
        "purpose": "形成判断",
        "key_message": "只映射了 A",
        "evidence": ["事实 A"],
        "relationship_type": "thesis",
        "visual_role": "thesis",
        "source_indexes": [0],
    }]

    with pytest.raises(ValueError, match="omits source_paths indexes: 1"):
        _build_authored_storyboard_plan(deck, materials, storyboard, "先给判断")


@pytest.mark.asyncio
async def test_create_presentation_returns_durable_job_when_wait_is_cancelled(
    tmp_path,
    monkeypatch,
) -> None:
    entered = threading.Event()
    release = threading.Event()

    def slow_compose(_service, _job_id, _deck, *, runtime_run_id=None):
        assert runtime_run_id is None
        entered.set()
        release.wait(timeout=5)
        raise RuntimeError("background test generation stopped")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        slow_compose,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    task = asyncio.create_task(
        create_presentation(title="可续跑生成", slides=_slides())
    )
    try:
        assert await asyncio.to_thread(entered.wait, 1)
        task.cancel()
        response = await task
        retried = await create_presentation(
            title="可续跑生成",
            slides=_slides(),
        )
        service = PresentationService(
            context.state_root / "presentation-jobs"
        )
    finally:
        release.set()
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.content[0]["type"] == "text"
    assert "continues in the background" in response.content[0]["text"]
    assert response.metadata["error_code"] == "presentation_generation_pending"
    assert response.metadata["job_id"]
    assert response.metadata["job_status"] in {"pending", "queued", "running"}
    assert response.metadata["verification_status"] == "pending"
    assert response.metadata["requires_probe"] is True
    assert response.metadata["requires_human"] is False
    assert response.metadata["automatic_retry_allowed"] is False
    assert response.metadata["side_effect_outcome"] == "unknown"
    assert "candidate_path" not in response.metadata
    assert retried.is_error is False
    assert retried.metadata["job_id"] == response.metadata["job_id"]
    assert retried.metadata["verification_status"] == "pending"
    assert len(service.store.list_records("conversation-ppt")) == 1


@pytest.mark.asyncio
async def test_create_presentation_reports_explicit_stop_when_wait_is_cancelled(
    tmp_path,
    monkeypatch,
) -> None:
    entered = threading.Event()
    release = threading.Event()

    def slow_compose(_service, _job_id, _deck, *, runtime_run_id=None):
        assert runtime_run_id == "native-run-explicit-stop"
        entered.set()
        release.wait(timeout=5)
        raise RuntimeError("stopped worker returned after foreground cancellation")

    monkeypatch.setattr(
        "lantai.agents.tools.presentation.PresentationService.compose",
        slow_compose,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        with bind_runtime_effect_context(
            RuntimeEffectContext(
                run_id="native-run-explicit-stop",
                attempt_id="native-run-explicit-stop:attempt-1",
                effect_id="effect-create-presentation-stop",
                state_workspace_dir=str(context.state_root),
            )
        ):
            task = asyncio.create_task(
                create_presentation(
                    title="显式停止生成",
                    slides=_slides(),
                )
            )
        assert await asyncio.to_thread(entered.wait, 1)
        assert cancel_active_jobs_for_runtime_run(
            context.state_root,
            "native-run-explicit-stop",
        )
        task.cancel()
        response = await task
    finally:
        release.set()
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_job_cancelled"
    assert response.metadata["job_status"] == "cancelled"
    assert response.metadata["verification_status"] == "not_delivered"
    assert response.metadata["requires_probe"] is False
    assert response.metadata["publish_ready"] is False
    assert response.metadata["requires_publish"] is False
    assert "continues in the background" not in response.content[0]["text"]


@pytest.mark.asyncio
async def test_get_presentation_generation_status_does_not_duplicate_running_job(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    service = PresentationService(context.state_root / "presentation-jobs")
    record = service.create_job(
        PresentationJobRequest(
            conversation_id=context.workspace_grant_id,
            user_request="后台继续生成",
            slide_count=1,
        )
    )

    set_current_workspace_context(context)
    try:
        response = await get_presentation_generation_status(record.job_id)
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False
    assert not any(block["type"] == "file" for block in response.content)
    assert response.metadata["job_id"] == record.job_id
    assert response.metadata["job_status"] == "pending"
    assert response.metadata["verification_status"] == "pending"
    assert response.metadata["requires_probe"] is True
    assert response.metadata["progress"] == {
        "requested_slide_count": 1,
        "planned_slide_count": None,
        "authored_slide_count": 0,
        "rendered_slide_count": 0,
        "reviewed_slide_count": 0,
        "deck_review_completed": False,
        "latest_repair_round": 0,
        "repair_attempts_completed": 0,
    }
    assert len(service.store.list_records(context.workspace_grant_id)) == 1


@pytest.mark.asyncio
async def test_status_reports_only_durable_per_page_progress(
    tmp_path: Path,
) -> None:
    context = _context(tmp_path)
    service = PresentationService(context.state_root / "presentation-jobs")
    record = service.create_job(
        PresentationJobRequest(
            conversation_id=context.workspace_grant_id,
            user_request="后台生成三页",
            slide_count=3,
        )
    )
    for index in (1, 2):
        slide = DeckSlide(
            id=f"slide-{index}",
            layout_id="freeform",
            objects=[],
        )
        service.store.write_checkpoint(
            record.job_id,
            f"page-scene-{index:03d}.json",
            json.dumps({"slide": slide.model_dump(mode="json")}),
        )
    render_dir = service.store.job_dir(record.job_id) / "renders" / "round-0"
    render_dir.mkdir(parents=True)
    for index in (1, 2):
        Image.new("RGB", (640, 360), "white").save(
            render_dir / f"slide-{index}.png"
        )
    now = datetime.now(timezone.utc)
    service.store.record_stage_metric(
        record.job_id,
        PresentationStageMetric(
            stage=JobStage.QUALITY,
            operation="critic",
            repair_round=0,
            slide_id="slide-1",
            outcome="passed",
            duration_ms=25,
            started_at=now,
            finished_at=now,
        ),
    )
    service.store.write_quality_checkpoint(
        record.job_id,
        "slide-"
        + hashlib.sha256(b"slide-1").hexdigest()
        + ".json",
        VisualReviewCheckpoint(
            scope="slide",
            dependency_fingerprint="a" * 64,
            critic_cache_identity="test-critic",
            result=VisualCriticResult(
                critic_id="test-critic",
                issues=[],
            ),
            rendered_sha256=hashlib.sha256(
                (render_dir / "slide-1.png").read_bytes()
            ).hexdigest(),
        ).model_dump_json(),
    )
    service.store.append_event(
        record.job_id,
        stage=JobStage.QUALITY,
        status=JobStatus.RUNNING,
        message="正在逐页检查",
    )
    job_dir = service.store.job_dir(record.job_id)
    files_before = {
        str(path.relative_to(job_dir)): path.read_bytes()
        for path in job_dir.rglob("*")
        if path.is_file()
    }
    directories_before = {
        str(path.relative_to(job_dir))
        for path in job_dir.rglob("*")
        if path.is_dir()
    }

    set_current_workspace_context(context)
    try:
        response = await get_presentation_generation_status(record.job_id)
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False
    assert response.metadata["progress"] == {
        "requested_slide_count": 3,
        "planned_slide_count": None,
        "authored_slide_count": 2,
        "rendered_slide_count": 2,
        "reviewed_slide_count": 1,
        "deck_review_completed": False,
        "latest_repair_round": 0,
        "repair_attempts_completed": 0,
    }
    assert "Authored 2/3; rendered 2/3; reviewed 1/3" in response.content[0][
        "text"
    ]
    assert not any(block["type"] == "file" for block in response.content)
    assert {
        str(path.relative_to(job_dir)): path.read_bytes()
        for path in job_dir.rglob("*")
        if path.is_file()
    } == files_before
    assert {
        str(path.relative_to(job_dir))
        for path in job_dir.rglob("*")
        if path.is_dir()
    } == directories_before

    Image.new("RGB", (640, 360), "black").save(
        render_dir / "slide-1.png"
    )
    set_current_workspace_context(context)
    try:
        stale_checkpoint = await get_presentation_generation_status(
            record.job_id
        )
    finally:
        set_current_workspace_context(None)
    assert stale_checkpoint.metadata["progress"]["reviewed_slide_count"] == 0

    repaired_deck = DeckIR(
        title="拆页后的当前作品",
        slides=[
            DeckSlide(
                id=f"slide-{index}",
                layout_id="freeform",
                objects=[],
            )
            for index in (1, 2, 3)
        ],
    )
    service.store.write_checkpoint(
        record.job_id,
        "deck-ir-redesign-1.json",
        repaired_deck.model_dump_json(),
    )
    repaired_render_dir = job_dir / "renders" / "round-1"
    repaired_render_dir.mkdir(parents=True)
    for index in (1, 2, 3):
        Image.new(
            "RGB",
            (640, 360),
            f"#{index}{index}{index}{index}{index}{index}",
        ).save(
            repaired_render_dir / f"slide-{index}.png",
        )
    service.store.record_stage_metric(
        record.job_id,
        PresentationStageMetric(
            stage=JobStage.COMPOSE,
            operation="budget_reservation",
            repair_round=1,
            outcome="passed",
            duration_ms=0,
            started_at=now,
            finished_at=now,
        ),
    )
    set_current_workspace_context(context)
    try:
        repaired = await get_presentation_generation_status(record.job_id)
    finally:
        set_current_workspace_context(None)
    assert repaired.metadata["progress"] == {
        "requested_slide_count": 3,
        "planned_slide_count": None,
        "authored_slide_count": 3,
        "rendered_slide_count": 3,
        "reviewed_slide_count": 0,
        "deck_review_completed": False,
        "latest_repair_round": 1,
        "repair_attempts_completed": 0,
    }


@pytest.mark.asyncio
async def test_get_presentation_generation_status_does_not_materialize_unknown_job(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    presentation_root = context.state_root / "presentation-jobs"

    set_current_workspace_context(context)
    try:
        response = await get_presentation_generation_status(
            "presentation-does-not-exist",
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_job_not_found"
    assert not presentation_root.exists()


@pytest.mark.asyncio
async def test_get_presentation_generation_status_does_not_mutate_bare_job(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    presentation_root = context.state_root / "presentation-jobs"
    job_id = "presentation-read-only-probe"
    job_dir = presentation_root / job_id
    job_dir.mkdir(parents=True)
    record = PresentationJobRecord(
        job_id=job_id,
        conversation_id=context.workspace_grant_id,
        request=PresentationJobRequest(
            conversation_id=context.workspace_grant_id,
            user_request="只读状态查询",
            slide_count=3,
        ),
    )
    (job_dir / "manifest.json").write_text(
        record.model_dump_json(indent=2),
        encoding="utf-8",
    )
    entries_before = {
        str(path.relative_to(presentation_root))
        for path in presentation_root.rglob("*")
    }

    set_current_workspace_context(context)
    try:
        response = await get_presentation_generation_status(job_id)
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False
    assert response.metadata["job_status"] == "pending"
    assert {
        str(path.relative_to(presentation_root))
        for path in presentation_root.rglob("*")
    } == entries_before


@pytest.mark.asyncio
async def test_get_presentation_generation_status_is_read_only_for_completed_job(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        generated = await create_presentation(
            title="后台生成已完成",
            slides=_slides(),
            output_filename="first-delivery.pptx",
        )
        recovered = await get_presentation_generation_status(
            generated.metadata["job_id"],
        )
    finally:
        set_current_workspace_context(None)

    assert generated.is_error is False
    assert recovered.is_error is False
    assert recovered.metadata["job_id"] == generated.metadata["job_id"]
    assert recovered.metadata["verification_status"] == "visual_verified"
    assert recovered.metadata["publish_ready"] is True
    assert recovered.metadata["requires_publish"] is True
    assert not any(block["type"] == "file" for block in recovered.content)
    assert not (context.project_root / "recovered-delivery.pptx").exists()
    service = PresentationService(context.state_root / "presentation-jobs")
    assert len(service.store.list_records(context.workspace_grant_id)) == 1


@pytest.mark.asyncio
async def test_finalize_presentation_generation_explicitly_publishes_completed_job(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        generated = await create_presentation(
            title="后台产物显式发布",
            slides=_slides(),
            output_filename="initial.pptx",
        )
        published = await finalize_presentation_generation(
            generated.metadata["job_id"],
            output_filename="published.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert generated.is_error is False
    assert published.is_error is False
    assert published.metadata["job_id"] == generated.metadata["job_id"]
    assert published.metadata["verification_status"] == "visual_verified"
    assert published.content[0]["type"] == "file"
    assert (context.project_root / "published.pptx").is_file()


@pytest.mark.asyncio
async def test_finalize_does_not_overwrite_user_edited_delivery(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    context = _context(tmp_path)
    destination = context.project_root / "author-edited.pptx"
    set_current_workspace_context(context)
    try:
        generated = await create_presentation(
            title="保留作者保存",
            slides=_slides(),
            output_filename=destination.name,
        )
        destination.write_bytes(destination.read_bytes() + b"author-save")
        authored_bytes = destination.read_bytes()

        replayed = await finalize_presentation_generation(
            generated.metadata["job_id"],
            output_filename=destination.name,
        )
    finally:
        set_current_workspace_context(None)

    assert replayed.is_error is True
    assert replayed.metadata["error_code"] == "presentation_delivery_conflict"
    assert replayed.metadata["job_id"] == generated.metadata["job_id"]
    assert replayed.metadata["side_effect_outcome"] == "none"
    assert replayed.metadata["requires_human"] is True
    assert replayed.metadata["publish_ready"] is False
    assert replayed.metadata["requires_publish"] is False
    assert destination.read_bytes() == authored_bytes


@pytest.mark.asyncio
async def test_get_presentation_generation_status_does_not_leak_other_conversation(
    tmp_path,
) -> None:
    context = _context(tmp_path)
    service = PresentationService(context.state_root / "presentation-jobs")
    record = service.create_job(
        PresentationJobRequest(
            conversation_id="another-conversation",
            user_request="不应泄露",
            slide_count=1,
        )
    )

    set_current_workspace_context(context)
    try:
        response = await get_presentation_generation_status(record.job_id)
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata["error_code"] == "presentation_job_not_found"
    assert record.job_id not in response.content[0]["text"]
    assert "candidate_path" not in response.metadata


@pytest.mark.asyncio
async def test_create_presentation_rebuilds_when_render_evidence_changes(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    set_current_workspace_context(_context(tmp_path))
    try:
        first = await create_presentation(
            title="证据不可篡改",
            slides=_slides(),
            output_filename="evidence.pptx",
        )
        receipt = json.loads(
            Path(first.metadata["verification_receipt"]).read_text(
                encoding="utf-8"
            )
        )
        render_path = Path(receipt["render_evidence"][0]["path"])
        Image.new("RGB", (640, 360), "white").save(render_path)
        second = await create_presentation(
            title="证据不可篡改",
            slides=_slides(),
            output_filename="evidence.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert first.is_error is False
    assert second.is_error is False
    assert second.metadata["reused"] is False
    assert second.metadata["job_id"] != first.metadata["job_id"]
    assert len(render_calls) == 2


@pytest.mark.asyncio
async def test_create_presentation_resumes_interrupted_delivery_without_regeneration(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    context = _context(tmp_path)
    destination = context.project_root / "interrupted-delivery.pptx"
    receipts_root = context.state_root / "presentation-tool-receipts"
    original_write_text = Path.write_text

    def interrupt_receipt_write(path, data, *args, **kwargs):
        if path.parent == receipts_root and path.name.endswith(".tmp"):
            raise OSError("simulated receipt write interruption")
        return original_write_text(path, data, *args, **kwargs)

    set_current_workspace_context(context)
    try:
        with monkeypatch.context() as interrupted:
            interrupted.setattr(Path, "write_text", interrupt_receipt_write)
            interrupted_response = await create_presentation(
                title="中断后续传",
                slides=_slides(),
                output_filename=destination.name,
            )

        store = PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        )
        records = store.list_records(context.workspace_grant_id)
        assert len(records) == 1
        interrupted_job_id = records[0].job_id
        assert records[0].status == JobStatus.COMPLETED
        assert destination.is_file()
        assert list(receipts_root.glob("*.json")) == []
        assert interrupted_response.is_error is True
        assert (
            interrupted_response.metadata["error_code"]
            == "presentation_delivery_file_error"
        )
        assert interrupted_response.metadata["job_id"] == interrupted_job_id
        assert interrupted_response.metadata["side_effect_outcome"] == "partial"
        assert interrupted_response.metadata["automatic_retry_allowed"] is True
        assert interrupted_response.metadata["requires_probe"] is False
        assert interrupted_response.metadata["requires_human"] is False
        assert interrupted_response.metadata["publish_ready"] is False
        assert interrupted_response.metadata["requires_publish"] is False
        assert interrupted_response.metadata["failure_categories"] == ["file"]
        assert interrupted_response.metadata["recovery_actions"] == [
            "retry_file_stage"
        ]

        resumed = await create_presentation(
            title="中断后续传",
            slides=_slides(),
            output_filename=destination.name,
        )
    finally:
        set_current_workspace_context(None)

    assert resumed.is_error is False
    assert resumed.metadata["job_id"] == interrupted_job_id
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert len(
        PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
    ) == 1
    assert Path(resumed.metadata["verification_receipt"]).is_file()


@pytest.mark.asyncio
async def test_create_presentation_retries_failed_final_delivery_verification(
    tmp_path,
    monkeypatch,
) -> None:
    from lantai.presentation.quality import (
        is_verified_delivery as actual_is_verified_delivery,
    )

    verification_calls = 0
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    def fail_first_copied_delivery(*args, **kwargs):
        nonlocal verification_calls
        verification_calls += 1
        if verification_calls == 2:
            return False
        return actual_is_verified_delivery(*args, **kwargs)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.is_verified_delivery",
        fail_first_copied_delivery,
    )
    context = _context(tmp_path)
    destination = context.project_root / "verification-interrupted.pptx"
    set_current_workspace_context(context)
    try:
        interrupted = await create_presentation(
            title="最终核验中断后续传",
            slides=_slides(),
            output_filename=destination.name,
        )
        resumed = await create_presentation(
            title="最终核验中断后续传",
            slides=_slides(),
            output_filename=destination.name,
        )
    finally:
        set_current_workspace_context(None)

    assert interrupted.is_error is True
    assert (
        interrupted.metadata["error_code"]
        == "presentation_delivery_verification_failed"
    )
    assert interrupted.metadata["job_id"] == resumed.metadata["job_id"]
    assert interrupted.metadata["side_effect_outcome"] == "partial"
    assert interrupted.metadata["automatic_retry_allowed"] is True
    assert interrupted.metadata["requires_probe"] is False
    assert interrupted.metadata["requires_human"] is False
    assert interrupted.metadata["publish_ready"] is False
    assert interrupted.metadata["requires_publish"] is False
    assert interrupted.metadata["failure_categories"] == ["file"]
    assert interrupted.metadata["recovery_actions"] == ["retry_file_stage"]
    assert destination.is_file()
    assert resumed.is_error is False
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert len(
        PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
    ) == 1


@pytest.mark.asyncio
async def test_create_presentation_retries_failed_file_response_without_regeneration(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []
    send_calls: list[str] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    from lantai.agents.tools.send_file import send_file_to_user as actual_send

    async def interrupt_first_response(file_path: str) -> ToolResponse:
        send_calls.append(file_path)
        if len(send_calls) == 1:
            return ToolResponse(
                content=[TextBlock(text="Error: simulated file response failure")],
                is_error=True,
            )
        return await actual_send(file_path)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.send_file_to_user",
        interrupt_first_response,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        interrupted = await create_presentation(
            title="文件响应中断后续发",
            slides=_slides(),
            output_filename="response-interrupted.pptx",
        )
        resumed = await create_presentation(
            title="文件响应中断后续发",
            slides=_slides(),
            output_filename="response-interrupted.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert interrupted.is_error is True
    assert (
        interrupted.metadata["error_code"]
        == "presentation_delivery_response_error"
    )
    assert interrupted.metadata["side_effect_outcome"] == "partial"
    assert interrupted.metadata["automatic_retry_allowed"] is True
    assert interrupted.metadata["requires_probe"] is False
    assert interrupted.metadata["requires_human"] is False
    assert interrupted.metadata["publish_ready"] is True
    assert interrupted.metadata["requires_publish"] is True
    assert Path(interrupted.metadata["verification_receipt"]).is_file()
    assert resumed.is_error is False
    assert resumed.metadata["job_id"] == interrupted.metadata["job_id"]
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert len(send_calls) == 2
    assert len(
        PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
    ) == 1


@pytest.mark.asyncio
async def test_create_presentation_retries_raised_file_response_without_regeneration(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []
    send_calls: list[str] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    from lantai.agents.tools.send_file import send_file_to_user as actual_send

    async def raise_first_response(file_path: str) -> ToolResponse:
        send_calls.append(file_path)
        if len(send_calls) == 1:
            raise RuntimeError("simulated file response exception")
        return await actual_send(file_path)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.send_file_to_user",
        raise_first_response,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        interrupted = await create_presentation(
            title="文件响应异常后续发",
            slides=_slides(),
            output_filename="response-exception.pptx",
        )
        resumed = await create_presentation(
            title="文件响应异常后续发",
            slides=_slides(),
            output_filename="response-exception.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert interrupted.is_error is True
    assert (
        interrupted.metadata["error_code"]
        == "presentation_delivery_response_error"
    )
    assert interrupted.metadata["job_id"] == resumed.metadata["job_id"]
    assert interrupted.metadata["side_effect_outcome"] == "partial"
    assert interrupted.metadata["automatic_retry_allowed"] is True
    assert Path(interrupted.metadata["verification_receipt"]).is_file()
    assert resumed.is_error is False
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert len(send_calls) == 2
    assert len(
        PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
    ) == 1


@pytest.mark.asyncio
async def test_create_presentation_does_not_swallow_cancelled_file_response(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls: list[Path] = []
    send_calls: list[str] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    from lantai.agents.tools.send_file import send_file_to_user as actual_send

    async def cancel_first_response(file_path: str) -> ToolResponse:
        send_calls.append(file_path)
        if len(send_calls) == 1:
            raise asyncio.CancelledError
        return await actual_send(file_path)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.send_file_to_user",
        cancel_first_response,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        with pytest.raises(asyncio.CancelledError):
            await create_presentation(
                title="用户停止文件响应",
                slides=_slides(),
                output_filename="response-cancelled.pptx",
            )
        records_after_stop = PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
        resumed = await create_presentation(
            title="用户停止文件响应",
            slides=_slides(),
            output_filename="response-cancelled.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert len(records_after_stop) == 1
    assert resumed.is_error is False
    assert resumed.metadata["job_id"] == records_after_stop[0].job_id
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert len(send_calls) == 2


@pytest.mark.asyncio
async def test_create_presentation_retries_failed_delivery_attestation(
    tmp_path,
    monkeypatch,
) -> None:
    from lantai.agents.tools.presentation import (
        _visual_verification_attestation as actual_attestation,
    )

    attestation_calls = 0
    render_calls: list[Path] = []

    def render(source, output_dir, *, timeout=120):
        render_calls.append(Path(source))
        return _render_one(source, output_dir, timeout=timeout)

    def interrupt_first_attestation(destination, receipt_path, quality):
        nonlocal attestation_calls
        attestation_calls += 1
        if attestation_calls == 1:
            raise ValueError("simulated attestation interruption")
        return actual_attestation(destination, receipt_path, quality)

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation._visual_verification_attestation",
        interrupt_first_attestation,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        interrupted = await create_presentation(
            title="认证中断后续传",
            slides=_slides(),
            output_filename="attestation-interrupted.pptx",
        )
        resumed = await create_presentation(
            title="认证中断后续传",
            slides=_slides(),
            output_filename="attestation-interrupted.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert interrupted.is_error is True
    assert (
        interrupted.metadata["error_code"]
        == "presentation_verification_attestation_failed"
    )
    assert interrupted.metadata["job_id"] == resumed.metadata["job_id"]
    assert interrupted.metadata["side_effect_outcome"] == "partial"
    assert interrupted.metadata["automatic_retry_allowed"] is True
    assert interrupted.metadata["requires_probe"] is False
    assert interrupted.metadata["requires_human"] is False
    assert interrupted.metadata["publish_ready"] is False
    assert interrupted.metadata["requires_publish"] is False
    assert interrupted.metadata["failure_categories"] == ["file"]
    assert interrupted.metadata["recovery_actions"] == ["retry_file_stage"]
    assert Path(interrupted.metadata["verification_receipt"]).is_file()
    assert resumed.is_error is False
    assert resumed.metadata["reused"] is True
    assert len(render_calls) == 1
    assert attestation_calls == 2
    assert len(
        PresentationJobStore(
            context.state_root / "presentation-jobs",
            read_only=True,
        ).list_records(context.workspace_grant_id)
    ) == 1


@pytest.mark.asyncio
async def test_create_presentation_blocks_delivery_without_visual_verification(
    tmp_path,
    monkeypatch,
) -> None:
    def unavailable(*_args, **_kwargs):
        raise PresentationRenderUnavailable("renderer unavailable")

    monkeypatch.setattr("lantai.presentation.service.render_pptx", unavailable)
    monkeypatch.setattr(
        "lantai.presentation.service.PresentationService.compiler_for",
        lambda service, _engine: service.compiler,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="未视觉验证",
            slides=_slides(),
            output_filename="unverified.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert not any(block["type"] == "file" for block in response.content)
    assert response.metadata["error_code"] == "presentation_quality_failed"
    assert response.metadata["verification_status"] == "failed"
    assert "render.unavailable" in response.content[0]["text"]
    assert response.metadata["failure_categories"] == ["file"]
    assert response.metadata["recovery_actions"] == ["retry_file_stage"]
    assert response.metadata["draft_preview"] is None
    assert not (context.project_root / "unverified.pptx").exists()


@pytest.mark.asyncio
async def test_create_presentation_quarantines_document_like_visual_failure(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_one)
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.UtilityPresentationVisualCritic",
        lambda _agent_id: _DocumentLikeVisualCritic(),
    )
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.UtilityPresentationRepairDesigner",
        lambda _agent_id: None,
    )
    context = _context(tmp_path)
    set_current_workspace_context(context)
    try:
        response = await create_presentation(
            title="文档搬运式 PPT 不得交付",
            slides=_slides(),
            output_filename="document-like.pptx",
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert not any(block["type"] == "file" for block in response.content)
    assert response.metadata["error_code"] == "presentation_quality_failed"
    assert response.metadata["verification_status"] == "failed"
    assert "vision.aesthetic_staleness" in response.content[0]["text"]
    assert response.metadata["failure_categories"] == ["aesthetic"]
    assert response.metadata["recovery_actions"] == ["redesign_slide"]
    assert "candidate_path" not in response.metadata
    assert not (context.project_root / "document-like.pptx").exists()
    preview = response.metadata["draft_preview"]
    assert preview["status"] == "unverified_draft"
    assert preview["slide_count"] == 1
    assert len(preview["render_paths"]) == 1
    assert Path(preview["render_paths"][0]).is_file()
    assert response.metadata["job_stage"] == "failed"
    assert response.metadata["stage_metrics"]
    assert response.metadata["side_effect_outcome"] == "none"
    assert response.metadata["publish_ready"] is False
    assert response.metadata["requires_publish"] is False

    set_current_workspace_context(context)
    try:
        status = await get_presentation_generation_status(
            response.metadata["job_id"]
        )
        assert status.metadata["draft_preview"] == preview
        render_path = Path(preview["render_paths"][0])
        original_render = render_path.read_bytes()
        Image.new("RGB", (640, 360), "white").save(render_path)
        tampered = await get_presentation_generation_status(
            response.metadata["job_id"]
        )
        render_path.write_bytes(original_render)
        service = PresentationService(context.state_root / "presentation-jobs")
        candidate = (
            service.store.job_dir(response.metadata["job_id"])
            / "candidates"
            / "round-0"
            / "presentation.pptx"
        )
        candidate.write_bytes(candidate.read_bytes() + b"tampered")
        tampered_candidate = await get_presentation_generation_status(
            response.metadata["job_id"]
        )
    finally:
        set_current_workspace_context(None)

    assert tampered.metadata["draft_preview"] is None
    assert tampered_candidate.metadata["draft_preview"] is None
    assert not any(block["type"] == "file" for block in status.content)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "output_filename",
    ["../outside.pptx", r"..\outside.pptx"],
)
async def test_create_presentation_rejects_unsafe_output_name(
    tmp_path,
    output_filename,
) -> None:
    set_current_workspace_context(_context(tmp_path))
    try:
        response = await create_presentation(
            title="unsafe",
            slides=_slides(),
            output_filename=output_filename,
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is True
    assert response.metadata == {
        "error_code": "invalid_presentation_arguments",
        "deterministic_failure": True,
        "side_effect_outcome": "known_not_applied",
    }
