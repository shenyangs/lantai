"""Exercise correction, real PPTX authoring, full-scope review and delivery.

Only the Office renderer and paid visual Provider are test doubles. The source
binding, serializer, receipt, identity/scope checks and Works gate are real.
"""
from __future__ import annotations

import json
import hashlib
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from lantai.agents.tools.presentation import create_presentation, _visual_verification_attestation
from lantai.app.runner import runner as runner_module
from lantai.app.runner.runner import AgentRunner
from lantai.app.experience.events import append_chat_event
from lantai.app.experience.runtime_runs import record_chat_runtime_run_started
from lantai.config.context import WorkspaceContext, set_current_workspace_context
from lantai.presentation.models import VisualCriticResult
from lantai.presentation.render import RenderResult
from lantai.runtime.deliverable_contract import bind_deliverable_contract
from lantai.runtime.effect_context import RuntimeEffectContext, bind_runtime_effect_context
from lantai.runtime.native_turn_support import _validate_tool_call_arguments
from lantai.runtime.runs import build_chat_task_contract
from lantai.runtime.store import RuntimeRunStore
from lantai.runtime.tools import LantaiToolRegistry, ToolCall


@pytest.mark.parametrize("path", [None, "", "   ", 123])
def test_attestation_cannot_supply_missing_or_invalid_preview_paths(tmp_path, path):
    deck = tmp_path / "deck.pptx"
    deck.write_bytes(b"candidate bytes")
    receipt = tmp_path / "receipt.json"
    receipt.write_text(json.dumps({
        "output_path": str(deck),
        "output_sha256": hashlib.sha256(deck.read_bytes()).hexdigest(),
        "render_evidence": [{"slide_index": 1, "path": path}],
    }), encoding="utf-8")
    quality = dict(passed=True, delivery_status="passed", render_verified=True,
                   vision_critique_required=True, visual_verified=True,
                   render_evidence="passed", visual_evidence="passed",
                   visual_critic_id="test-critic", rendered_slide_count=1)
    with pytest.raises(ValueError, match="render evidence paths"):
        _visual_verification_attestation(deck, receipt, quality)


@pytest.mark.asyncio
async def test_two_source_mapping_errors_recover_to_26_page_verified_delivery(
    tmp_path, monkeypatch,
):
    workspace, state = tmp_path / "workspace", tmp_path / "state"
    workspace.mkdir()
    state.mkdir()
    context = WorkspaceContext(
        agent_home=tmp_path / "home", state_root=state, project_root=workspace,
        workspace_grant_id="recovery-chain", trust_state="trusted",
    )
    evidence = [f"Evidence section {letter} was retained from the source."
                for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ"]
    source = workspace / "report.md"
    source.write_text("\n".join(evidence), encoding="utf-8")
    slides = [{"id": f"page-{index}", "layout_id": "cover", "objects": [
        {"kind": "text", "role": "title", "text": f"Section {index}", "slot": "title"},
        {"kind": "text", "role": "body", "text": fact, "slot": "subtitle"},
    ]} for index, fact in enumerate(evidence, start=1)]
    storyboard = [{"purpose": "Explain the source evidence", "key_message": fact,
                   "visual_role": "opening" if index == 0 else "evidence",
                   "relationship_type": "thesis", "evidence": [fact],
                   "source_indexes": [0]}
                  for index, fact in enumerate(evidence)]
    rendered, reviewed = [], []

    def render(deck_path, output_dir, *, timeout=120):
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        paths = []
        for index, _slide in enumerate(Presentation(deck_path).slides, start=1):
            path = destination / f"slide-{index}.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste((index * 7 % 255, 60, 110), (0, 0, 320, 360))
            image.save(path)
            paths.append(path)
        rendered.extend(paths)
        pdf = destination / "presentation.pdf"
        pdf.write_bytes(b"renderer test double; PNG evidence is decoded separately")
        return RenderResult(pdf_path=pdf, slide_images=tuple(paths))

    class Critic:
        def review(self, pages):
            reviewed.extend(page.slide_id for page in pages)
            return VisualCriticResult(critic_id="test-full-scope-critic", issues=[])

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    monkeypatch.setattr("lantai.agents.tools.presentation.UtilityPresentationVisualCritic",
                        lambda _agent: Critic())
    run_id = "run-recovery-chain"
    await record_chat_runtime_run_started(
        workspace_dir=state, agent_id="default", chat_id="chat-recovery-chain",
        session_id="console:chat-recovery-chain", run_id=run_id,
        effective_workspace_dir=str(workspace),
        task_contract=build_chat_task_contract(
            objective="Create a source-backed 26-page presentation",
            client_turn_id="client-recovery-chain", complexity="light",
            work_method_id="standard",
        ),
    )
    await bind_deliverable_contract(
        state_workspace_dir=state, run_id=run_id, kind="presentation",
        accepted_extensions=[".pptx"], visual_verification_required=True,
    )
    contract = await RuntimeRunStore(state).load_contract(run_id)
    assert contract is not None
    arguments = dict(title="Source-backed report", brief="Retain all source evidence.",
                     slides=slides, storyboard=storyboard,
                     narrative_arc="Introduce evidence, review all sections, then conclude.",
                     required_facts=evidence, output_filename="corrected-deck.pptx")
    original = workspace / "fallback.pptx"
    original.write_bytes(b"preserve the original unverified candidate")
    set_current_workspace_context(context)
    try:
        no_sources = await create_presentation(**arguments)
        assert no_sources.is_error
        assert no_sources.metadata["source_index_repair"]["source_count"] == 0
        erased_indexes = await create_presentation(**{
            **arguments, "source_paths": [str(source)],
            "storyboard": [{**page, "source_indexes": []} for page in storyboard],
        })
        assert erased_indexes.is_error
        assert erased_indexes.metadata["source_index_repair"]["missing_source_indexes"] == [0]
        assert rendered == reviewed == []
        with bind_runtime_effect_context(RuntimeEffectContext(
            run_id=run_id, attempt_id=contract.attempt_id,
            effect_id="effect-corrected-presentation", state_workspace_dir=str(state),
        )):
            generated = await create_presentation(**arguments, source_paths=[str(source)])
    finally:
        set_current_workspace_context(None)

    assert not generated.is_error, generated.content
    deck = workspace / "corrected-deck.pptx"
    assert len(Presentation(deck).slides) == 26
    assert len(rendered) == 26
    assert reviewed == [f"page-{index}" for index in range(1, 27)]
    attestation = generated.metadata["verification_attestation"]
    assert attestation["source_page_count"] == 26
    assert attestation["observed_page_indexes"] == list(range(1, 27))
    public_text = "\n".join(block.get("text", "") for block in generated.content)
    # The actual model-visible output must provide the next tool's path mapping,
    # not expect hidden quality metadata or invent a nonexistent render_paths key.
    inputs = json.loads(public_text.split("finalize_deliverable_inputs: ", 1)[1])
    assert inputs["primary_path"] == str(deck)
    receipt = json.loads(Path(generated.metadata["verification_receipt"]).read_text())
    assert inputs["verification_paths"] == [row["path"] for row in receipt["render_evidence"]]
    assert len(inputs["verification_paths"]) == 26
    for name in ("build_lantai_tool_registry", "build_memory_tool_registry", "build_mcp_client_registry"):
        monkeypatch.setattr(runner_module, name, lambda *_, **__: LantaiToolRegistry())
    registry = AgentRunner(agent_id="default", workspace_dir=state)._build_native_tool_registry(
        agent_config=SimpleNamespace(), mcp_clients=[], runtime_attempt_id=contract.attempt_id,
        request_context={"state_workspace_dir": str(state), "workspace_dir": str(workspace),
                         "works_workspace_dir": str(state), "run_id": run_id,
                         "chat_id": "chat-recovery-chain"},
    )
    spec, executor = registry.get_tool("finalize_deliverable")
    call = ToolCall(
        id="finalize-corrected-presentation", name=spec.name,
        arguments={**inputs, "title": arguments["title"]},
    )
    # Nothing except the task title may be added to the advertised mapping.
    # Validate and execute the real runner entry point, not just the inner gate.
    _validate_tool_call_arguments(call, specs_by_name={spec.name: spec})
    assert inputs["acceptance_evidence"] == []
    for event_type, status in (("tool.call.started", "running"), ("tool.call.completed", "completed")):
        await append_chat_event(
            workspace_dir=state, agent_id="default", chat_id="chat-recovery-chain",
            session_id="console:chat-recovery-chain", run_id=run_id,
            kind="chat.tool_call.finished" if status == "completed" else "chat.tool_call.started",
            type=event_type, source="tool", status=status,
            summary="Canonical presentation generation",
            payload={"name": "create_presentation", "call_id": "corrected-presentation",
                     "attempt_id": contract.attempt_id,
                     "canonical_tool": {"effect_attempt_id": contract.attempt_id},
                     **({"ok": True, "metadata": generated.metadata}
                        if status == "completed" else {})},
        )
    finalized = await executor(**call.arguments)
    assert finalized["delivery_status"] == "verified"
    assert finalized["artifact_id"]
    assert finalized["work_id"]
    assert all(check["passed"] for check in finalized["acceptance_results"])
    assert original.read_bytes() == b"preserve the original unverified candidate"
