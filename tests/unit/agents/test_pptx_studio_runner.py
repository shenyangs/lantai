from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from pptx import Presentation

from lantai.agents.tools.skill_resources import run_skill_script
from lantai.config.context import WorkspaceContext, set_current_workspace_context


PROJECT_ROOT = Path(__file__).resolve().parents[3]
SKILLS_ROOT = PROJECT_ROOT / "src" / "lantai" / "agents" / "skills"


@pytest.mark.parametrize("locale", ["pptx-zh", "pptx-en"])
def test_bundled_pptx_skill_exposes_only_v31_runtime_entrypoints(locale: str) -> None:
    skill_root = SKILLS_ROOT / locale
    skill = (skill_root / "SKILL.md").read_text(encoding="utf-8")
    boundary = (skill_root / "references" / "runtime-boundary.md").read_text(
        encoding="utf-8",
    )

    assert "version: 3.1.0" in skill
    assert "license: Apache-2.0" in skill
    assert "Presentation Skill v3" in skill
    assert "PptxGenJS" in skill
    assert "DeckIR" in skill
    assert "GenOffice" in skill
    assert "scripts/run_presentation_v3.py" in skill
    assert "scripts/preflight.py" in skill
    assert "lantai.presentation.render" in boundary
    for tool_name in (
        "create_presentation",
        "get_presentation_generation_status",
        "retry_presentation_generation",
        "finalize_presentation_generation",
        "resolve_pending_deliverable_choice",
    ):
        assert tool_name in skill
        assert tool_name in boundary
    for reference_name in (
        "editorial.md",
        "design.md",
        "template-assets.md",
        "native-author.md",
        "review-recovery.md",
    ):
        assert (skill_root / "references" / reference_name).is_file()
    for retired in (
        "scripts/generate_presentation.py",
        "scripts/render_pptxgenjs.py",
        "scripts/pptxgenjs_compat.cjs",
        "studio-workflow.md",
    ):
        assert not (skill_root / retired).exists()


def test_localized_pptx_entrypoints_share_one_runtime_implementation() -> None:
    localized_skills = [
        (SKILLS_ROOT / locale / "SKILL.md").read_text(encoding="utf-8")
        for locale in ("pptx-zh", "pptx-en")
    ]
    assert all("\nname: pptx\n" in skill for skill in localized_skills)
    assert not (SKILLS_ROOT / "pptx-flagship").exists()
    for script_name in ("run_presentation_v3.py", "preflight.py"):
        localized = [
            (SKILLS_ROOT / locale / "scripts" / script_name).read_text(
                encoding="utf-8",
            )
            for locale in ("pptx-zh", "pptx-en")
        ]
        assert localized[0] == localized[1]


@pytest.mark.asyncio
async def test_pptx_preflight_runs_through_logical_runtime_tool(
    tmp_path: Path,
) -> None:
    agent_home = tmp_path / "agent"
    project_root = agent_home / "topic-workspaces" / "chat-1"
    project_root.mkdir(parents=True)
    shutil.copytree(SKILLS_ROOT / "pptx-zh", agent_home / "skills" / "pptx")
    (agent_home / "skill.json").write_text(
        json.dumps(
            {
                "schema_version": "workspace-skill-manifest.v1",
                "version": 1,
                "skills": {
                    "pptx": {
                        "enabled": True,
                        "channels": ["all"],
                        "source": "builtin",
                    },
                },
            },
        ),
        encoding="utf-8",
    )
    deck = Presentation()
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(project_root / "sample.pptx")
    set_current_workspace_context(
        WorkspaceContext(
            agent_home=agent_home,
            state_root=project_root,
            project_root=project_root,
        ),
    )
    try:
        response = await run_skill_script(
            "pptx",
            "scripts/preflight.py",
            ["--path", "sample.pptx"],
        )
    finally:
        set_current_workspace_context(None)

    assert response.is_error is False, response.content[0]["text"]
    result = json.loads(response.content[0]["text"])
    assert response.metadata["skill_name"] == "pptx"
    assert result["passed"] is True
    assert result["sha256"]
