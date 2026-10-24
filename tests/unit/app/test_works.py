# -*- coding: utf-8 -*-
from __future__ import annotations

import base64
import hashlib
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from lantai.app.capabilities.artifacts import load_artifact_manifest
from lantai.app import office_conversion, work_covers, works as works_store
from lantai.app.office_conversion import OfficeConversionResult
from lantai.app.work_covers import generate_ai_cover
from lantai.app.works import (
    add_work,
    build_work_edit_context,
    create_work_edit_copy,
    derive_work_to_pdf,
    get_work_cover_file,
    infer_work_kind,
    list_works,
    remove_work,
    resolve_canonical_artifact_edit_target,
    set_work_cover,
)
from lantai.config.config import ImageGenerationConfig, ModelSlotConfig


def _write_pptx(path: Path, *, slides: int = 3) -> None:
    deck = Presentation()
    for number in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = f"slide {number + 1}"
    deck.save(path)


def test_work_library_copies_and_preserves_a_generated_html_artifact(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "games" / "tetris.html"
    source.parent.mkdir()
    source.write_text("<title>俄罗斯方块</title>", encoding="utf-8")
    (source.parent / "game.js").write_text(
        "window.ready = true",
        encoding="utf-8",
    )

    work = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="俄罗斯方块",
        source_chat_id="chat-1",
    )

    assert work.kind == "html"
    assert work.title == "俄罗斯方块"
    assert Path(work.path).read_text(encoding="utf-8") == source.read_text(
        encoding="utf-8",
    )
    assert Path(work.path) != source
    assert (Path(work.path).parent / "game.js").read_text(
        encoding="utf-8",
    ) == ("window.ready = true")
    assert list_works(workspace) == [work]
    manifest = load_artifact_manifest(workspace)
    assert manifest["artifacts"][0]["retention"] == "user_saved"
    assert manifest["artifacts"][0]["metadata"]["work_id"] == work.id


def test_work_edit_context_points_at_the_managed_bundle_not_the_old_chat(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "game" / "index.html"
    source.parent.mkdir()
    source.write_text("<title>Game</title>", encoding="utf-8")
    (source.parent / "game.js").write_text("play()", encoding="utf-8")
    work = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="迷你马里奥",
        source_chat_id="chat-old",
        source_run_id="run-old",
    )

    context = build_work_edit_context(
        work,
        work.path,
        source_chat_name="做一个小游戏",
    )

    assert f"主文件：{work.path}" in context
    assert f"工作目录：{Path(work.path).parent}" in context
    assert "独立可编辑副本已挂载" in context
    assert "收藏的来源版本保持不变" in context
    assert "来源会话：做一个小游戏" in context


def test_work_edit_branch_preserves_original_and_records_version_lineage(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "game" / "index.html"
    source.parent.mkdir()
    source.write_text("<title>Original</title>", encoding="utf-8")
    (source.parent / "game.js").write_text("play(1)", encoding="utf-8")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
    )
    original_bytes = Path(original.path).read_bytes()

    edit_entry, parent_digest = create_work_edit_copy(
        workspace,
        original,
        original.path,
    )
    edit_entry.write_text("<title>Version 2</title>", encoding="utf-8")
    (edit_entry.parent / "game.js").write_text("play(2)", encoding="utf-8")
    version_two = add_work(
        workspace,
        source_path=edit_entry,
        allowed_roots=[edit_entry.parent],
    )

    assert Path(original.path).read_bytes() == original_bytes
    assert (Path(original.path).parent / "game.js").read_text(
        encoding="utf-8",
    ) == "play(1)"
    assert version_two.parent_work_id == original.id
    assert version_two.parent_digest == parent_digest
    assert version_two.lineage_root_id == original.id
    assert version_two.version == 2
    assert version_two.source_chat_id == original.source_chat_id
    assert (
        Path(version_two.path).read_text(encoding="utf-8")
        == "<title>Version 2</title>"
    )
    assert not (
        Path(version_two.path).parent / ".lantai-work-lineage.json"
    ).exists()


def test_pptx_pdf_derivation_records_exact_cross_format_provenance(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    _write_pptx(source, slides=1)
    version_one = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="月度复盘",
        source_chat_id="chat-1",
        source_run_id="run-1",
    )
    edit_entry, _ = create_work_edit_copy(
        workspace,
        version_one,
        version_one.path,
    )
    _write_pptx(edit_entry, slides=2)
    version_two = add_work(
        workspace,
        source_path=edit_entry,
        allowed_roots=[edit_entry.parent],
        title="月度复盘",
        source_chat_id="chat-1",
    )
    calls = 0

    def fake_convert(
        source_path: str | Path,
        target_extension: str,
        output_dir: str | Path,
        **_kwargs,
    ) -> OfficeConversionResult:
        nonlocal calls
        calls += 1
        assert Path(source_path) == Path(version_two.path)
        assert target_extension == "pdf"
        output = Path(output_dir) / "deck.pptx.pdf"
        output.parent.mkdir(parents=True, exist_ok=True)
        payload = b"%PDF-1.4\nexact derived bytes\n%%EOF\n"
        output.write_bytes(payload)
        return OfficeConversionResult(
            output_path=output,
            sha256=f"sha256:{hashlib.sha256(payload).hexdigest()}",
            converter="libreoffice-test",
        )

    monkeypatch.setattr(
        office_conversion,
        "convert_office_document",
        fake_convert,
    )

    derived = derive_work_to_pdf(workspace, version_two.id)
    repeated = derive_work_to_pdf(workspace, version_two.id)

    assert repeated.id == derived.id
    assert calls == 1
    assert derived.kind == "pdf"
    assert derived.version == 1
    assert derived.parent_work_id is None
    assert derived.parent_digest is None
    assert derived.lineage_root_id == derived.id
    assert derived.source_chat_id == version_two.source_chat_id
    assert derived.source_run_id == version_two.source_run_id
    assert derived.derivation is not None
    assert derived.derivation.source_work_id == version_two.id
    assert derived.derivation.source_artifact_id == version_two.artifact_id
    assert derived.derivation.source_version == 2
    assert derived.derivation.source_digest == version_two.source_digest
    assert derived.derivation.output_digest == derived.source_digest
    assert derived.derivation.operation == "office_conversion"
    assert derived.derivation.converter == "libreoffice-test"
    manifest = load_artifact_manifest(workspace)
    artifact = next(
        item
        for item in manifest["artifacts"]
        if item["id"] == derived.artifact_id
    )
    assert artifact["metadata"]["derivation"] == derived.derivation.model_dump(
        mode="json",
    )


def test_pdf_derivation_refuses_output_changed_during_work_copy(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    _write_pptx(source, slides=1)
    presentation = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="月度复盘",
        source_chat_id="chat-1",
        source_run_id="run-1",
    )
    converted_output: Path | None = None

    def fake_convert(
        _source_path: str | Path,
        _target_extension: str,
        output_dir: str | Path,
        **_kwargs,
    ) -> OfficeConversionResult:
        nonlocal converted_output
        converted_output = Path(output_dir) / "deck.pptx.pdf"
        converted_output.parent.mkdir(parents=True, exist_ok=True)
        payload = b"%PDF-1.4\nverified conversion\n%%EOF\n"
        converted_output.write_bytes(payload)
        return OfficeConversionResult(
            output_path=converted_output,
            sha256=f"sha256:{hashlib.sha256(payload).hexdigest()}",
            converter="libreoffice-test",
        )

    real_copy2 = works_store.shutil.copy2
    changed = False

    def change_output_before_work_copy(
        source_path: str | Path,
        destination_path: str | Path,
        *args,
        **kwargs,
    ):
        nonlocal changed
        if Path(source_path) == converted_output and not changed:
            changed = True
            Path(source_path).write_bytes(
                b"%PDF-1.4\nchanged after conversion receipt\n%%EOF\n"
            )
        return real_copy2(source_path, destination_path, *args, **kwargs)

    monkeypatch.setattr(
        office_conversion,
        "convert_office_document",
        fake_convert,
    )
    monkeypatch.setattr(works_store.shutil, "copy2", change_output_before_work_copy)

    with pytest.raises(
        ValueError,
        match="derived Work output changed during registration",
    ):
        derive_work_to_pdf(workspace, presentation.id)

    assert changed is True
    assert list_works(workspace) == [presentation]
    artifacts = load_artifact_manifest(workspace)["artifacts"]
    assert len(artifacts) == 1
    assert artifacts[0]["id"] == presentation.artifact_id


def test_same_basename_pdf_is_not_inferred_as_a_pptx_derivative(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"presentation")
    presentation = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    unrelated_pdf = workspace / "deck.pdf"
    unrelated_pdf.write_bytes(b"%PDF unrelated")
    pdf = add_work(
        workspace,
        source_path=unrelated_pdf,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )

    assert pdf.derivation is None
    assert pdf.lineage_root_id == pdf.id
    assert pdf.lineage_root_id != presentation.lineage_root_id


def test_same_file_can_be_owned_by_two_canonical_chats(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"same immutable candidate")

    chat_one = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
        source_run_id="run-shared",
    )
    chat_two = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-2",
        source_run_id="run-shared",
    )

    assert chat_one.id != chat_two.id
    assert {item.source_chat_id for item in list_works(workspace)} == {
        "chat-1",
        "chat-2",
    }


def test_work_version_cannot_move_to_another_canonical_chat(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"version one")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)
    edit_entry.write_bytes(b"version two")

    with pytest.raises(ValueError, match="canonical chat does not match"):
        add_work(
            workspace,
            source_path=edit_entry,
            allowed_roots=[edit_entry.parent],
            source_chat_id="chat-2",
        )


def test_work_version_can_bind_a_legacy_unowned_parent_to_a_canonical_chat(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "draft.md"
    source.write_text("version one", encoding="utf-8")
    original = add_work(workspace, source_path=source, allowed_roots=[workspace])
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)
    edit_entry.write_text("version two", encoding="utf-8")

    version_two = add_work(
        workspace,
        source_path=edit_entry,
        allowed_roots=[edit_entry.parent],
        source_chat_id="chat-private-editor",
    )

    assert original.source_chat_id is None
    assert version_two.parent_work_id == original.id
    assert version_two.source_chat_id == "chat-private-editor"


def test_work_edit_branch_rejects_a_stale_parallel_save(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"version one")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
    )
    first_edit, _ = create_work_edit_copy(workspace, original, original.path)
    stale_edit, _ = create_work_edit_copy(workspace, original, original.path)

    first_edit.write_bytes(b"version two")
    version_two = add_work(
        workspace,
        source_path=first_edit,
        allowed_roots=[first_edit.parent],
    )
    assert version_two.version == 2

    stale_edit.write_bytes(b"conflicting version")
    with pytest.raises(ValueError, match="newer work version already exists"):
        add_work(
            workspace,
            source_path=stale_edit,
            allowed_roots=[stale_edit.parent],
        )


def test_unchanged_work_edit_branch_does_not_create_a_duplicate_version(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "notes.md"
    source.write_text("keep this", encoding="utf-8")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)

    result = add_work(
        workspace,
        source_path=edit_entry,
        allowed_roots=[edit_entry.parent],
    )

    assert result.id == original.id
    assert [item.id for item in list_works(workspace)] == [original.id]


def test_canonical_artifact_edit_target_binds_chat_revision_and_page(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    _write_pptx(source, slides=3)
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)

    target, context = resolve_canonical_artifact_edit_target(
        workspace,
        {
            "schema_version": 1,
            "chat_id": "chat-1",
            "asset_id": original.artifact_id,
            "work_id": original.id,
            "edit_path": str(edit_entry),
            "base_revision": {
                "version": original.version,
                "digest": original.source_digest,
            },
            "selection": {"kind": "page", "page": 3},
        },
        chat_id="chat-1",
    )

    assert target["edit_path"] == str(edit_entry.resolve())
    assert target["selection"] == {"kind": "page", "page": 3}
    assert "【当前会话成果修改目标】" in context
    assert "本轮选择：第 3 页" in context


def test_canonical_artifact_edit_target_rejects_page_outside_base_revision(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    _write_pptx(source, slides=2)
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)

    with pytest.raises(ValueError, match="out of range for 2 pages"):
        resolve_canonical_artifact_edit_target(
            workspace,
            {
                "schema_version": 1,
                "chat_id": "chat-1",
                "asset_id": original.artifact_id,
                "work_id": original.id,
                "edit_path": str(edit_entry),
                "base_revision": {
                    "version": original.version,
                    "digest": original.source_digest,
                },
                "selection": {"kind": "page", "page": 3},
            },
            chat_id="chat-1",
        )


def test_edit_target_rejects_edit_changed_before_confirmation(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "notes.md"
    source.write_text("base", encoding="utf-8")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)
    edit_entry.write_text("changed before send", encoding="utf-8")

    with pytest.raises(ValueError, match="changed before confirmation"):
        resolve_canonical_artifact_edit_target(
            workspace,
            {
                "schema_version": 1,
                "chat_id": "chat-1",
                "asset_id": original.artifact_id,
                "work_id": original.id,
                "edit_path": str(edit_entry),
                "base_revision": {
                    "version": original.version,
                    "digest": original.source_digest,
                },
            },
            chat_id="chat-1",
        )


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda target: target.update(chat_id="chat-2"), "another chat"),
        (
            lambda target: target["base_revision"].update(version=99),
            "base revision",
        ),
        (
            lambda target: target.update(edit_path="/tmp/forged.pptx"),
            "managed edits",
        ),
    ],
)
def test_canonical_artifact_edit_target_rejects_forged_binding(
    tmp_path: Path,
    mutate,
    message: str,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"version one")
    original = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
    )
    edit_entry, _ = create_work_edit_copy(workspace, original, original.path)
    target = {
        "schema_version": 1,
        "chat_id": "chat-1",
        "asset_id": original.artifact_id,
        "work_id": original.id,
        "edit_path": str(edit_entry),
        "base_revision": {
            "version": original.version,
            "digest": original.source_digest,
        },
    }
    mutate(target)

    with pytest.raises((ValueError, PermissionError), match=message):
        resolve_canonical_artifact_edit_target(
            workspace,
            target,
            chat_id="chat-1",
        )


def test_listing_marks_modified_html_bundle_without_rewriting_history(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "game" / "index.html"
    source.parent.mkdir()
    source.write_text("<title>Game</title>", encoding="utf-8")
    (source.parent / "game.js").write_text("play(1)", encoding="utf-8")
    work = add_work(workspace, source_path=source, allowed_roots=[workspace])
    managed_script = Path(work.path).parent / "game.js"
    future_timestamp = work.updated_at.timestamp() + 10
    managed_script.write_text("play(2)", encoding="utf-8")
    os.utime(managed_script, (future_timestamp, future_timestamp))

    refreshed = list_works(workspace)[0]

    assert refreshed.id == work.id
    assert refreshed.updated_at == work.updated_at
    assert refreshed.availability_status == "available"
    assert refreshed.integrity_status == "modified"
    artifact = load_artifact_manifest(workspace)["artifacts"][0]
    assert artifact["id"] == work.artifact_id


def test_work_listing_marks_missing_saved_copy(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "report.md"
    source.write_text("saved", encoding="utf-8")
    work = add_work(workspace, source_path=source, allowed_roots=[workspace])
    Path(work.path).unlink()

    refreshed = list_works(workspace)[0]

    assert refreshed.availability_status == "missing"
    assert refreshed.integrity_status == "unknown"


def test_work_edit_copy_rejects_modified_saved_version(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "report.md"
    source.write_text("saved", encoding="utf-8")
    work = add_work(workspace, source_path=source, allowed_roots=[workspace])
    Path(work.path).write_text("changed behind the library", encoding="utf-8")

    with pytest.raises(ValueError, match="saved work version was modified"):
        create_work_edit_copy(workspace, work, work.path)


def test_adding_the_same_source_is_idempotent(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "report.pdf"
    source.write_bytes(b"%PDF-test")

    first = add_work(workspace, source_path=source, allowed_roots=[workspace])
    second = add_work(workspace, source_path=source, allowed_roots=[workspace])

    assert first.id == second.id
    assert len(list_works(workspace)) == 1


def test_legacy_work_without_digest_does_not_claim_changed_source_bytes(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "report.pdf"
    source.write_bytes(b"%PDF-old")
    legacy = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
        source_run_id="run-old",
    )
    index_path = workspace / ".lantai" / "works" / "index.json"
    payload = json.loads(index_path.read_text(encoding="utf-8"))
    payload["works"][0].pop("source_digest")
    index_path.write_text(
        json.dumps(payload, ensure_ascii=False),
        encoding="utf-8",
    )
    source.write_bytes(b"%PDF-new")

    current = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        source_chat_id="chat-1",
        source_run_id="run-current",
    )

    assert current.id != legacy.id
    assert Path(current.path).read_bytes() == b"%PDF-new"
    assert current.source_digest != legacy.source_digest
    assert {item.id for item in list_works(workspace)} == {
        legacy.id,
        current.id,
    }


def test_same_run_and_digest_is_idempotent_across_reconnected_paths(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    first_source = workspace / "attempt-1" / "index.html"
    second_source = workspace / "attempt-2" / "index.html"
    first_source.parent.mkdir()
    second_source.parent.mkdir()
    content = "<!doctype html><html><body>same</body></html>"
    first_source.write_text(content, encoding="utf-8")
    second_source.write_text(content, encoding="utf-8")

    first = add_work(
        workspace,
        source_path=first_source,
        allowed_roots=[workspace],
        source_run_id="run-1",
    )
    second = add_work(
        workspace,
        source_path=second_source,
        allowed_roots=[workspace],
        source_run_id="run-1",
    )

    assert first.id == second.id
    assert first.source_digest == second.source_digest
    assert len(list_works(workspace)) == 1


def test_remove_work_deletes_only_the_managed_copy(tmp_path: Path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "deck.pptx"
    source.write_bytes(b"deck")
    work = add_work(workspace, source_path=source, allowed_roots=[workspace])

    removed = remove_work(workspace, work.id)

    assert removed == work
    assert source.read_bytes() == b"deck"
    assert not Path(work.path).exists()
    assert list_works(workspace) == []
    assert load_artifact_manifest(workspace)["artifacts"] == []


def test_work_cover_is_managed_with_the_work_and_round_trips_unicode(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "report.pdf"
    source.write_bytes(b"%PDF-test")
    work = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="市场研究报告",
    )
    generated = tmp_path / "generated.png"
    generated.write_bytes(b"png")

    covered = set_work_cover(
        workspace,
        work.id,
        source_path=generated,
        mime_type="image/png",
    )

    assert covered.title == "市场研究报告"
    assert covered.cover_source == "ai"
    managed = get_work_cover_file(workspace, work.id)
    assert managed is not None
    assert managed[1].read_bytes() == b"png"


def test_setting_cover_never_deletes_primary_work_named_cover_png(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "cover.png"
    source.write_bytes(b"primary-work")
    work = add_work(workspace, source_path=source, allowed_roots=[workspace])
    managed_work = Path(work.path)
    generated = workspace / "generated.png"
    generated.write_bytes(b"generated-cover")

    covered = set_work_cover(
        workspace,
        work.id,
        source_path=generated,
        mime_type="image/png",
    )

    assert managed_work.read_bytes() == b"primary-work"
    assert Path(covered.cover_path or "").read_bytes() == b"generated-cover"


async def test_ai_cover_uses_selected_model_and_keeps_title_out_of_image_bytes(
    tmp_path: Path,
    monkeypatch,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "game.html"
    source.write_text("<!doctype html><title>Game</title>", encoding="utf-8")
    work = add_work(
        workspace,
        source_path=source,
        allowed_roots=[workspace],
        title="俄罗斯方块",
    )
    image_buffer = io.BytesIO()
    Image.new("RGB", (16, 10), "navy").save(image_buffer, format="PNG")
    captured = {}

    class _Provider:
        models = [SimpleNamespace(id="cover-v1")]
        extra_models = []

        async def generate_image(self, **kwargs):
            captured.update(kwargs)
            return {
                "data": [
                    {
                        "b64_json": base64.b64encode(
                            image_buffer.getvalue(),
                        ).decode("ascii"),
                    },
                ],
            }

    provider = _Provider()
    manager = SimpleNamespace(
        get_provider=lambda _provider_id: provider,
        is_model_slot_available=lambda _slot: True,
    )
    monkeypatch.setattr(
        work_covers.ProviderManager,
        "get_instance",
        lambda: manager,
    )
    config = ImageGenerationConfig(
        models=[ModelSlotConfig(provider_id="image", model="cover-v1")],
        works_cover_enabled=True,
        works_cover_model=ModelSlotConfig(
            provider_id="image",
            model="cover-v1",
        ),
    )

    covered = await generate_ai_cover(workspace, work, config)

    assert covered.cover_source == "ai"
    assert Path(covered.cover_path or "").is_file()
    assert captured["model_id"] == "cover-v1"
    assert "Do not render any letters" in captured["prompt"]


def test_add_work_rejects_sources_outside_the_allowed_workspace(
    tmp_path: Path,
):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    outside = tmp_path / "outside.html"
    outside.write_text("outside", encoding="utf-8")

    with pytest.raises(PermissionError, match="outside the active workspace"):
        add_work(workspace, source_path=outside, allowed_roots=[workspace])


@pytest.mark.parametrize(
    ("filename", "kind"),
    [
        ("poster.png", "image"),
        ("report.pdf", "pdf"),
        ("slides.pptx", "presentation"),
        ("budget.xlsx", "spreadsheet"),
        ("game.html", "html"),
    ],
)
def test_infer_work_kind(filename: str, kind: str):
    assert infer_work_kind(filename) == kind
