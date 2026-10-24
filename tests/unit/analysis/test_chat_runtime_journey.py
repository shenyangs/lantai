"""A normal chat can discover and execute the bounded analysis tools."""

from __future__ import annotations

import asyncio
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from lantai.agents.tool_registry import build_lantai_tool_registry
from lantai.app.agent_context import current_session_id_scope
from lantai.config.context import WorkspaceContext, set_current_workspace_context
from lantai.runtime.adapters.native_runtime import NativeLantaiRuntime
from lantai.runtime.input import RuntimeInput, RuntimeRequestContext
from lantai.runtime.messages import RuntimeMessage
from lantai.runtime.models import ModelMessage, ModelResponse, ModelStreamEvent
from lantai.runtime.tools import ToolCall
from lantai.presentation.models import VisualCriticResult
from lantai.presentation.render import RenderResult
from lantai.security.tool_guard.approval import ApprovalDecision


class AnalysisChatModel:
    """Script only model decisions; execute real registry tools and worker."""

    provider_id = "test"
    model_name = "scripted-analysis"

    def __init__(self) -> None:
        self.requests = []
        self.dataset_id: str | None = None
        self.analysis_id: str | None = None

    async def stream(self, request):
        self.requests.append(request)
        round_number = len(self.requests)
        if round_number == 1:
            name, arguments = "equip_lantai_tools", {"files": True}
        elif round_number == 2:
            name, arguments = "import_data_table", {"file_path": "sales.csv"}
        elif round_number == 3:
            imported = json.loads(request.messages[-1].content[0].data["output"][0]["text"])
            self.dataset_id = imported["dataset_id"]
            name, arguments = "analyze_data_table", {
                "dataset_id": self.dataset_id,
                "plan": {
                    "group_by": ["region"],
                    "metrics": [{"column": "amount", "op": "sum", "as": "revenue"}],
                    "sort": [{"column": "region"}],
                },
            }
        elif round_number == 4:
            analyzed = json.loads(request.messages[-1].content[0].data["output"][0]["text"])
            assert analyzed["rows"] == [
                {"region": "East", "revenue": "10.20"},
                {"region": "West", "revenue": "5.30"},
            ]
            self.analysis_id = analyzed["analysis_id"]
            yield ModelStreamEvent(kind="text_delta", message_id="final", delta="已按区域汇总，来源与计算结果已保存。")
            yield ModelStreamEvent(
                kind="message_completed", message_id="final",
                response=ModelResponse(message=ModelMessage.text(
                    "assistant", "已按区域汇总，来源与计算结果已保存。",
                )),
            )
            return
        else:
            raise AssertionError("Unexpected extra model round")

        message_id = f"tool-{round_number}"
        yield ModelStreamEvent(
            kind="tool_call_completed", message_id=message_id,
            tool_call=ToolCall(id=f"call-{round_number}", name=name, arguments=arguments),
        )
        yield ModelStreamEvent(
            kind="message_completed", message_id=message_id,
            response=ModelResponse(message=ModelMessage(id=message_id, role="assistant")),
        )

    async def complete(self, _request):
        return ModelResponse(message=ModelMessage.text("assistant", "完成"))


class AnalysisFollowupModel:
    provider_id = "test"
    model_name = "scripted-analysis-followup"

    def __init__(self, analysis_id: str, source_sha256: str) -> None:
        self.analysis_id = analysis_id
        self.source_sha256 = source_sha256
        self.requests = []

    async def stream(self, request):
        self.requests.append(request)
        round_number = len(self.requests)
        if round_number <= 2:
            name = "equip_lantai_tools" if round_number == 1 else "get_analysis_result"
            arguments = {"files": True} if round_number == 1 else {
                "analysis_id": self.analysis_id,
            }
            message_id = f"followup-tool-{round_number}"
            yield ModelStreamEvent(
                kind="tool_call_completed", message_id=message_id,
                tool_call=ToolCall(id=f"followup-call-{round_number}", name=name, arguments=arguments),
            )
            yield ModelStreamEvent(
                kind="message_completed", message_id=message_id,
                response=ModelResponse(message=ModelMessage(id=message_id, role="assistant")),
            )
            return
        if round_number != 3:
            raise AssertionError("Unexpected extra follow-up round")
        result = json.loads(request.messages[-1].content[0].data["output"][0]["text"])
        assert result["analysis_id"] == self.analysis_id
        assert result["evidence"]["source_sha256"] == self.source_sha256
        assert result["rows"][0]["revenue"] == "10.20"
        yield ModelStreamEvent(kind="text_delta", message_id="followup-final", delta="East 为 10.20。")
        yield ModelStreamEvent(
            kind="message_completed", message_id="followup-final",
            response=ModelResponse(message=ModelMessage.text("assistant", "East 为 10.20。")),
        )

    async def complete(self, _request):
        return ModelResponse(message=ModelMessage.text("assistant", "完成"))


class AnalysisRefreshModel:
    provider_id = "test"
    model_name = "scripted-analysis-refresh"

    def __init__(self, analysis_id: str, source_sha256: str) -> None:
        self.analysis_id = analysis_id
        self.source_sha256 = source_sha256
        self.requests = []
        self.refreshed_analysis_id: str | None = None
        self.refreshed_source_sha256: str | None = None

    async def stream(self, request):
        self.requests.append(request)
        round_number = len(self.requests)
        if round_number == 1:
            name, arguments = "equip_lantai_tools", {"files": True}
        elif round_number == 2:
            name, arguments = "refresh_analysis_result", {"analysis_id": self.analysis_id}
        elif round_number == 3:
            refreshed = json.loads(request.messages[-1].content[0].data["output"][0]["text"])
            assert refreshed["status"] == "changed"
            assert refreshed["previous_analysis_id"] == self.analysis_id
            assert refreshed["result"]["rows"] == [
                {"region": "East", "revenue": "12.30"},
                {"region": "West", "revenue": "5.30"},
            ]
            self.refreshed_analysis_id = refreshed["analysis_id"]
            self.refreshed_source_sha256 = refreshed["result"]["evidence"]["source_sha256"]
            assert self.refreshed_analysis_id != self.analysis_id
            assert self.refreshed_source_sha256 != self.source_sha256
            name, arguments = "get_analysis_result", {"analysis_id": self.analysis_id}
        elif round_number == 4:
            previous = json.loads(request.messages[-1].content[0].data["output"][0]["text"])
            assert previous["analysis_id"] == self.analysis_id
            assert previous["evidence"]["source_sha256"] == self.source_sha256
            assert previous["rows"][0]["revenue"] == "10.20"
            answer = "源数据已更新；新版 East 为 12.30，旧版仍保留 10.20。"
            yield ModelStreamEvent(kind="text_delta", message_id="refresh-final", delta=answer)
            yield ModelStreamEvent(
                kind="message_completed", message_id="refresh-final",
                response=ModelResponse(message=ModelMessage.text("assistant", answer)),
            )
            return
        else:
            raise AssertionError("Unexpected extra refresh round")

        message_id = f"refresh-tool-{round_number}"
        yield ModelStreamEvent(
            kind="tool_call_completed", message_id=message_id,
            tool_call=ToolCall(id=f"refresh-call-{round_number}", name=name, arguments=arguments),
        )
        yield ModelStreamEvent(
            kind="message_completed", message_id=message_id,
            response=ModelResponse(message=ModelMessage(id=message_id, role="assistant")),
        )

    async def complete(self, _request):
        return ModelResponse(message=ModelMessage.text("assistant", "完成"))


def _tool_analysis_metadata(request) -> dict:
    for message in reversed(request.messages):
        for content in message.content:
            output = content.data.get("output")
            if not isinstance(output, list):
                continue
            for block in output:
                if isinstance(block, dict) and isinstance(block.get("text"), str):
                    try:
                        parsed = json.loads(block["text"])
                    except json.JSONDecodeError:
                        continue
                    if isinstance(parsed, dict) and "analysis_id" in parsed:
                        return parsed
    raise AssertionError("The model did not receive analysis delivery metadata")


class AnalysisDeliveryModel:
    provider_id = "test"
    model_name = "scripted-analysis-delivery"

    def __init__(self, analysis_id: str, source_sha256: str) -> None:
        self.analysis_id = analysis_id
        self.source_sha256 = source_sha256
        self.requests = []
        self.deck_metadata: dict | None = None

    async def stream(self, request):
        self.requests.append(request)
        round_number = len(self.requests)
        if round_number == 1:
            name, arguments = "equip_lantai_tools", {"files": True}
        elif round_number == 2:
            name, arguments = "create_analysis_chart", {
                "analysis_id": self.analysis_id,
                "category_column": "region",
                "value_column": "revenue",
                "output_filename": "chat-analysis-chart.svg",
            }
        elif round_number == 3:
            chart = _tool_analysis_metadata(request)
            assert chart["analysis_id"] == self.analysis_id
            assert chart["source_sha256"] == self.source_sha256
            assert chart["chat_chart_reference"] == {
                "analysis_id": self.analysis_id,
                "x_column": "region",
                "y_column": "revenue",
                "type": "bar",
            }
            name, arguments = "create_analysis_report", {
                "analysis_id": self.analysis_id,
                "title": "区域收入分析",
                "output_filename": "chat-analysis-report.docx",
                "sections": [{"heading": "关键发现", "paragraphs": [[
                    {"text": "East 收入为 "}, {"row": 0, "column": "revenue"}, {"text": "。"},
                ]]}],
            }
        elif round_number == 4:
            report = _tool_analysis_metadata(request)
            assert report["analysis_id"] == self.analysis_id
            assert report["source_sha256"] == self.source_sha256
            assert report["delivery_consistency"] == "passed"
            name, arguments = "create_analysis_presentation", {
                "analysis_id": self.analysis_id,
                "title": "区域收入分析",
                "output_filename": "chat-analysis-deck.pptx",
                "slides": [{"layout_id": "data_conclusion", "objects": [
                    {"kind": "text", "role": "title", "text": "区域收入", "slot": "title"},
                    {"kind": "chart", "slot": "chart", "chart_type": "bar",
                     "categories": ["placeholder"], "series": {"revenue": [999]},
                     "unit": "元", "derivation": "placeholder", "source_refs": ["placeholder"]},
                    {"kind": "metric", "label": "East 收入", "value": "999", "slot": "metric_1"},
                    {"kind": "text", "role": "body", "text": "East 高于 West。", "slot": "conclusion"},
                ]}],
                "fact_refs": [{"row": 0, "column": "revenue"}],
                "bindings": [
                    {"kind": "chart", "slide": 0, "object": 1,
                     "category_column": "region", "value_column": "revenue", "series_name": "revenue"},
                    {"kind": "metric", "slide": 0, "object": 2,
                     "row": 0, "column": "revenue"},
                ],
            }
        elif round_number == 5:
            deck = _tool_analysis_metadata(request)
            self.deck_metadata = deck
            assert deck["analysis_id"] == self.analysis_id
            assert deck["source_sha256"] == self.source_sha256
            assert deck["delivery_consistency"] == "passed"
            answer = "图表、报告与演示文稿已基于同一份分析证据生成。"
            yield ModelStreamEvent(kind="text_delta", message_id="delivery-final", delta=answer)
            yield ModelStreamEvent(
                kind="message_completed", message_id="delivery-final",
                response=ModelResponse(message=ModelMessage.text("assistant", answer)),
            )
            return
        else:
            raise AssertionError("Unexpected extra delivery round")

        message_id = f"delivery-tool-{round_number}"
        yield ModelStreamEvent(
            kind="tool_call_completed", message_id=message_id,
            tool_call=ToolCall(id=f"delivery-call-{round_number}", name=name, arguments=arguments),
        )
        yield ModelStreamEvent(
            kind="message_completed", message_id=message_id,
            response=ModelResponse(message=ModelMessage(id=message_id, role="assistant")),
        )

    async def complete(self, _request):
        return ModelResponse(message=ModelMessage.text("assistant", "完成"))


class ExplicitDeliveryApproval:
    """Model the user's approval UI response, without changing tool policy."""

    def __init__(self) -> None:
        self.requests: list[dict] = []

    async def create_pending(self, **kwargs):
        self.requests.append(kwargs)
        future = asyncio.get_running_loop().create_future()
        future.set_result(ApprovalDecision.APPROVED)
        return SimpleNamespace(
            request_id=f"analysis-delivery-{len(self.requests)}",
            future=future,
            timeout_seconds=1.0,
        )


@pytest.mark.asyncio
async def test_plain_chat_equips_files_then_imports_and_analyzes_source(
    tmp_path: Path, monkeypatch,
) -> None:
    class PassingCritic:
        def review(self, _slides):
            return VisualCriticResult(critic_id="analysis-chat-journey-critic", issues=[])

    def render_stub(_source, output_dir, *, timeout=120):
        _ = timeout
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        pdf_path = destination / "presentation.pdf"
        pdf_path.write_bytes(b"pdf")
        preview = destination / "slide-1.png"
        image = Image.new("RGB", (640, 360), "white")
        image.paste("black", (0, 0, 320, 360))
        image.save(preview)
        return RenderResult(pdf_path=pdf_path, slide_images=(preview,))

    # The offline journey exercises the real deck writer and final-file checks;
    # utility-model visual review and PDF rendering are separately gated.
    monkeypatch.setattr(
        "lantai.agents.tools.presentation.UtilityPresentationVisualCritic",
        lambda _agent_id: PassingCritic(),
    )
    monkeypatch.setattr("lantai.presentation.service.render_pptx", render_stub)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "sales.csv").write_text(
        "region,amount\nEast,10.20\nWest,5.30\n", encoding="utf-8",
    )
    context = WorkspaceContext(
        agent_home=tmp_path / "home", state_root=tmp_path / "state", project_root=workspace,
    )
    registry = build_lantai_tool_registry(
        agent_config=SimpleNamespace(tools=SimpleNamespace(builtin_tools={})),
        request_context={"distribution": "developer_id"},
    )
    model = AnalysisChatModel()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: model,
        tool_registry=registry,
        progressive_tools=True,
    )
    runtime_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default", session_id="analysis-chat", user_id="default",
            run_id="turn-1", workspace_dir=str(workspace),
        ),
        messages=[RuntimeMessage.text("user", "请按区域汇总 sales.csv 的收入")],
    )

    set_current_workspace_context(context)
    try:
        with current_session_id_scope("analysis-chat"):
            events = [event async for event in runtime.run_turn(runtime_input)]
    finally:
        set_current_workspace_context(None)

    assert len(model.requests) == 4
    assert "import_data_table" not in {tool.name for tool in model.requests[0].tools}
    assert {"import_data_table", "analyze_data_table"} <= {
        tool.name for tool in model.requests[1].tools
    }
    assert model.dataset_id and model.analysis_id
    assert sum(event.type == "tool.call.completed" for event in events) == 3
    assert events[-1].type == "turn.completed"

    source_sha256 = json.loads(
        model.requests[2].messages[-1].content[0].data["output"][0]["text"]
    )["source_sha256"]
    followup_model = AnalysisFollowupModel(model.analysis_id, source_sha256)
    followup_runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: followup_model,
        tool_registry=registry,
        progressive_tools=True,
    )
    followup_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default", session_id="analysis-chat", user_id="default",
            run_id="turn-2", workspace_dir=str(workspace),
        ),
        messages=[
            RuntimeMessage.text("user", "请按区域汇总 sales.csv 的收入"),
            RuntimeMessage.text("assistant", "已按区域汇总，来源与计算结果已保存。"),
            RuntimeMessage.text("user", f"追问：上次分析 {model.analysis_id} 中 East 是多少？"),
        ],
    )
    set_current_workspace_context(context)
    try:
        with current_session_id_scope("analysis-chat"):
            followup_events = [
                event async for event in followup_runtime.run_turn(followup_input)
            ]
    finally:
        set_current_workspace_context(None)

    assert len(followup_model.requests) == 3
    assert followup_events[-1].type == "turn.completed"

    delivery_model = AnalysisDeliveryModel(model.analysis_id, source_sha256)
    delivery_approval = ExplicitDeliveryApproval()
    delivery_runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: delivery_model,
        tool_registry=registry,
        progressive_tools=True,
        approval_service=delivery_approval,
    )
    delivery_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default", session_id="analysis-chat", user_id="default",
            run_id="turn-3", workspace_dir=str(workspace),
        ),
        messages=[
            RuntimeMessage.text("user", "请按区域汇总 sales.csv 的收入"),
            RuntimeMessage.text("assistant", "已按区域汇总，来源与计算结果已保存。"),
            RuntimeMessage.text("user", f"追问：上次分析 {model.analysis_id} 中 East 是多少？"),
            RuntimeMessage.text("assistant", "East 为 10.20。"),
            RuntimeMessage.text("user", "再用这份分析做图表、报告和 PPT。"),
        ],
    )
    set_current_workspace_context(context)
    try:
        with current_session_id_scope("analysis-chat"):
            delivery_events = [event async for event in delivery_runtime.run_turn(delivery_input)]
    finally:
        set_current_workspace_context(None)

    assert len(delivery_model.requests) == 5
    assert sum(event.type == "tool.call.completed" for event in delivery_events) == 4
    assert len(delivery_approval.requests) == 3
    assert sum(event.type == "approval.requested" for event in delivery_events) == 3
    assert sum(event.type == "approval.resolved" for event in delivery_events) == 3
    event_types = [event.type for event in delivery_events]
    assert event_types.index("approval.requested") < event_types.index("tool.started")
    assert [
        event.payload["approval_binding"]["tool_name"]
        for event in delivery_events if event.type == "approval.requested"
    ] == ["create_analysis_chart", "create_analysis_report", "create_analysis_presentation"]
    assert all(
        event.payload["decision"] == ApprovalDecision.APPROVED.value
        for event in delivery_events if event.type == "approval.resolved"
    )
    delivery_messages = [
        event.payload["message"]
        for event in delivery_events
        if event.type == "message.completed"
        and event.status == "completed"
        and event.source == "assistant"
    ]
    delivery_text = "\n".join(
        str(block.get("text") or "")
        for message in delivery_messages
        for block in message.get("content", [])
        if isinstance(block, dict) and block.get("type") == "text"
    )
    assert "```analysis-chart" in delivery_text
    assert f'"analysis_id":"{model.analysis_id}"' in delivery_text
    assert '"type":"bar"' in delivery_text
    assert '"x_column":"region"' in delivery_text
    assert '"y_column":"revenue"' in delivery_text
    assert delivery_events[-1].type == "turn.completed"
    assert delivery_model.deck_metadata is not None
    assert "10.20" in (workspace / "chat-analysis-chart.svg").read_text(encoding="utf-8")
    report = Document(workspace / "chat-analysis-report.docx")
    assert model.analysis_id in "\n".join(paragraph.text for paragraph in report.paragraphs)
    assert source_sha256 in "\n".join(paragraph.text for paragraph in report.paragraphs)
    assert "10.20" in "\n".join(paragraph.text for paragraph in report.paragraphs)
    assert (workspace / "chat-analysis-deck.pptx").exists()
    deck = Presentation(workspace / "chat-analysis-deck.pptx")
    deck_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "10.20" in deck_text
    original_artifacts = {
        filename: hashlib.sha256((workspace / filename).read_bytes()).hexdigest()
        for filename in (
            "chat-analysis-chart.svg",
            "chat-analysis-report.docx",
            "chat-analysis-deck.pptx",
        )
    }

    (workspace / "sales.csv").write_text(
        "region,amount\nEast,12.30\nWest,5.30\n", encoding="utf-8",
    )
    refresh_model = AnalysisRefreshModel(model.analysis_id, source_sha256)
    refresh_runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: refresh_model,
        tool_registry=registry,
        progressive_tools=True,
    )
    refresh_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default", session_id="analysis-chat", user_id="default",
            run_id="turn-4", workspace_dir=str(workspace),
        ),
        messages=[
            RuntimeMessage.text("user", "请按区域汇总 sales.csv 的收入"),
            RuntimeMessage.text("assistant", "已按区域汇总，来源与计算结果已保存。"),
            RuntimeMessage.text("user", f"追问：上次分析 {model.analysis_id} 中 East 是多少？"),
            RuntimeMessage.text("assistant", "East 为 10.20。"),
            RuntimeMessage.text("user", "再用这份分析做图表、报告和 PPT。"),
            RuntimeMessage.text("assistant", "图表、报告与演示文稿已基于同一份分析证据生成。"),
            RuntimeMessage.text("user", "sales.csv 更新了；请刷新上次分析，并保留旧版供追溯。"),
        ],
    )
    set_current_workspace_context(context)
    try:
        with current_session_id_scope("analysis-chat"):
            refresh_events = [event async for event in refresh_runtime.run_turn(refresh_input)]
    finally:
        set_current_workspace_context(None)

    assert len(refresh_model.requests) == 4
    assert refresh_model.refreshed_analysis_id
    assert refresh_model.refreshed_source_sha256
    assert sum(event.type == "tool.call.completed" for event in refresh_events) == 3
    assert refresh_events[-1].type == "turn.completed"
    assert {
        filename: hashlib.sha256((workspace / filename).read_bytes()).hexdigest()
        for filename in original_artifacts
    } == original_artifacts
    assert "10.20" in (workspace / "chat-analysis-chart.svg").read_text(encoding="utf-8")
    assert "12.30" not in (workspace / "chat-analysis-chart.svg").read_text(encoding="utf-8")
