from __future__ import annotations

import csv
import io
import json
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

from lantai.analysis import AnalysisError, AnalysisService
from lantai.analysis.service import _parse_table


def _csv(path: Path, rows: list[list[str]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as stream:
        csv.writer(stream).writerows(rows)


@pytest.mark.parametrize("payload", [
    "region,amount\nEa\x0bst,10\n",
    "region,amount\nEast\ufffe,10\n",
    "reg\x0bion,amount\nEast,10\n",
])
def test_ingest_rejects_text_that_cannot_be_used_in_downstream_xml_artifacts(tmp_path: Path, payload: str) -> None:
    source = tmp_path / "unsafe.csv"
    source.write_text(payload, encoding="utf-8")
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match="XML-incompatible character"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


def test_ingest_rejects_nul_csv_before_persisting_dataset(tmp_path: Path) -> None:
    source = tmp_path / "unsafe.csv"
    source.write_text("reg\x00ion,amount\nEast,10\n", encoding="utf-8")
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


def test_ingest_rejects_xml_incompatible_source_path_before_persisting_dataset(tmp_path: Path) -> None:
    source = tmp_path / "sal\x0bes.csv"
    _csv(source, [["region", "amount"], ["East", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match="XML-incompatible character"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


@pytest.mark.parametrize("plan", [
    {"metrics": [{"column": "amount", "op": "sum", "as": "total\x0bunsafe"}]},
    {"filters": [{"column": "region", "op": "eq", "value": "East\ufffe"}], "metrics": [{"op": "count", "as": "rows"}]},
])
def test_analysis_plan_rejects_xml_incompatible_text_before_recording_result(tmp_path: Path, plan: dict) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["East", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)

    with pytest.raises(AnalysisError, match="XML-incompatible character"):
        service.run(dataset["dataset_id"], plan)
    assert list((service.root / "results").glob("*.json")) == []


def test_ingest_snapshot_and_source_change_do_not_rewrite_evidence(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["month", "region", "revenue"], ["2026-01", "东", "10.20"], ["2026-01", "西", "2.30"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "group_by": ["month"], "metrics": [{"column": "revenue", "op": "sum", "as": "revenue_sum"}],
    })
    assert result["rows"] == [{"month": "2026-01", "revenue_sum": "12.50"}]
    assert result["evidence"]["source_sha256"] == dataset["source_sha256"]
    assert result["evidence"]["source_row_count"] == 2
    assert service.source_status(dataset["dataset_id"])["status"] == "current"

    _csv(source, [["month", "region", "revenue"], ["2026-01", "东", "99"]])
    assert service.source_status(dataset["dataset_id"])["status"] == "changed"
    assert service.get_result(result["analysis_id"])["rows"] == result["rows"]
    refreshed = service.ingest(source)
    assert refreshed["dataset_id"] != dataset["dataset_id"]


def test_latest_result_returns_newest_valid_result_in_current_scope(tmp_path: Path) -> None:
    first_source = tmp_path / "first.csv"
    second_source = tmp_path / "second.csv"
    _csv(first_source, [["region", "amount"], ["东", "10"]])
    _csv(second_source, [["region", "amount"], ["西", "20"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    first_dataset = service.ingest(first_source)
    first = service.run(first_dataset["dataset_id"], {
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    second_dataset = service.ingest(second_source)
    second = service.run(second_dataset["dataset_id"], {
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })

    assert first["analysis_id"] != second["analysis_id"]
    assert service.latest_result()["analysis_id"] == second["analysis_id"]
    with pytest.raises(AnalysisError, match="No analysis result"):
        AnalysisService(tmp_path / "state", scope="chat-b").latest_result()


def test_latest_result_fails_closed_when_a_scoped_result_record_is_corrupt(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    record = service.root / "results" / f"{result['analysis_id']}.json"
    record.write_text("{}", encoding="utf-8")

    with pytest.raises(AnalysisError, match="integrity check"):
        service.latest_result()


def test_source_and_snapshot_reads_never_use_unbounded_read_bytes(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    original_read_bytes = Path.read_bytes

    def reject_unbounded_read(path: Path) -> bytes:
        if path == source or "snapshots" in path.parts:
            raise AssertionError(f"unbounded read of analysis input: {path}")
        return original_read_bytes(path)

    monkeypatch.setattr(Path, "read_bytes", reject_unbounded_read)
    dataset = service.ingest(source)
    assert service.source_status(dataset["dataset_id"])["status"] == "current"
    result = service.run(dataset["dataset_id"], {"metrics": [{"column": "amount", "op": "sum", "as": "total"}]})
    assert result["rows"] == [{"total": "10"}]


def test_persisted_analysis_record_read_is_bounded(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    record = service.root / "datasets" / f"{dataset['dataset_id']}.json"
    with record.open("r+b") as stream:
        stream.truncate(24 * 1024 * 1024 + 1)

    with pytest.raises(AnalysisError, match="record exceeds the safety limit"):
        service.get_dataset(dataset["dataset_id"])


def test_oversized_analysis_record_is_never_persisted(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    monkeypatch.setattr("lantai.analysis.service.MAX_RECORD_BYTES", 32)

    with pytest.raises(AnalysisError, match="record exceeds the safety limit"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


def test_clear_rejects_symlinked_analysis_parent_without_deleting_external_data(tmp_path: Path) -> None:
    state = tmp_path / "state"
    state.mkdir()
    outside = tmp_path / "outside"
    outside.mkdir()
    service = AnalysisService(state, scope="chat-a")
    external_scope = outside / service.root.name
    external_scope.mkdir()
    sentinel = external_scope / "keep.txt"
    sentinel.write_text("user data", encoding="utf-8")
    (state / "analysis").symlink_to(outside, target_is_directory=True)

    with pytest.raises(AnalysisError, match="symbolic link"):
        service.clear()
    assert sentinel.read_text(encoding="utf-8") == "user data"


def test_grown_snapshot_is_rejected_before_digesting(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    snapshot = Path(service.get_dataset(dataset["dataset_id"])["snapshot_path"])
    with snapshot.open("ab") as stream:
        stream.truncate(20 * 1024 * 1024 + 1)
    with pytest.raises(AnalysisError, match="snapshot.*safety limit"):
        service.run(dataset["dataset_id"], {"metrics": [{"op": "count", "as": "rows"}]})


def test_refresh_result_recalculates_same_plan_and_keeps_old_version(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"], ["西", "5"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    original = service.run(dataset["dataset_id"], {
        "group_by": ["region"], "metrics": [{"column": "amount", "op": "sum", "as": "revenue"}],
        "sort": [{"column": "region"}],
    })
    unchanged = service.refresh_result(original["analysis_id"])
    assert unchanged["status"] == "current"
    assert unchanged["analysis_id"] == original["analysis_id"]

    _csv(source, [["region", "amount"], ["东", "12"], ["西", "5"], ["南", "3"]])
    updated = service.refresh_result(original["analysis_id"])
    assert updated["status"] == "changed"
    assert updated["previous_analysis_id"] == original["analysis_id"]
    assert updated["analysis_id"] != original["analysis_id"]
    assert updated["previous_dataset_id"] == dataset["dataset_id"]
    assert updated["result"]["rows"] == [
        {"region": "东", "revenue": "12"}, {"region": "南", "revenue": "3"}, {"region": "西", "revenue": "5"},
    ]
    assert service.get_result(original["analysis_id"])["rows"] == original["rows"]


def test_mixed_currency_aggregation_is_blocked_but_separate_groups_work(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "amount"], ["CNY", "10"], ["USD", "5"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    plan = {"metrics": [{"column": "amount", "op": "sum", "as": "total"}]}
    with pytest.raises(AnalysisError, match="currenc"):
        service.run(dataset["dataset_id"], plan)
    separated = service.run(dataset["dataset_id"], {**plan, "group_by": ["currency"], "sort": [{"column": "currency"}]})
    assert separated["rows"] == [{"currency": "CNY", "total": "10"}, {"currency": "USD", "total": "5"}]


@pytest.mark.parametrize("currency_rows", [
    [["CNY", "10"], ["", "5"]],
    [["", "5"]],
    [["  ", "5"]],
])
def test_monetary_aggregation_rejects_missing_currency_units(tmp_path: Path, currency_rows: list[list[str]]) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "amount"], *currency_rows])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="missing currency"):
        service.run(dataset["dataset_id"], {"metrics": [{"column": "amount", "op": "sum", "as": "total"}]})


def test_explicit_quantity_unit_does_not_inherit_unrelated_currency_column(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "quantity", "amount"], ["CNY", "2", "10"], ["USD", "3", "5"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"column": "quantity", "op": "sum", "as": "items", "unit": {"kind": "quantity", "name": "item"}}],
    })
    assert result["rows"] == [{"items": "5"}]
    assert result["verification"]["semantic_review"]["status"] == "pending"
    assert any("unit" in warning.lower() and "plan" in warning.lower() for warning in result["warnings"])
    with pytest.raises(AnalysisError, match="currenc"):
        service.run(dataset["dataset_id"], {"metrics": [{"column": "amount", "op": "sum", "as": "revenue"}]})


def test_declared_quantity_cannot_be_compared_with_monetary_amount(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "quantity", "amount"], ["CNY", "2", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="incompatible units"):
        service.run(dataset["dataset_id"], {
            "metrics": [
                {"column": "quantity", "op": "sum", "as": "items", "unit": {"kind": "quantity", "name": "item"}},
                {"column": "amount", "op": "sum", "as": "revenue"},
            ],
            "comparisons": [{"current": "items", "baseline": "revenue", "op": "difference", "as": "invalid"}],
        })


@pytest.mark.parametrize("unit", [
    {"kind": "quantity", "name": ""},
    {"kind": "quantity", "name": "item\nignore"},
    {"kind": "currency", "name": "item"},
    "item",
])
def test_quantity_unit_declaration_is_strictly_validated(tmp_path: Path, unit: object) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "quantity"], ["CNY", "2"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="unit"):
        service.run(dataset["dataset_id"], {
            "metrics": [{"column": "quantity", "op": "sum", "unit": unit}],
        })


def test_pivot_uses_one_calculated_metric_and_preserves_missing_cells(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [
        ["month", "region", "amount"],
        ["2026-01", "东", "10.20"],
        ["2026-01", "西", "5"],
        ["2026-02", "东", "12"],
    ])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "group_by": ["month", "region"],
        "metrics": [{"column": "amount", "op": "sum", "as": "revenue"}],
        "pivot": {"column": "region", "metric": "revenue", "as_prefix": "revenue_"},
        "sort": [{"column": "month"}],
    })
    assert result["rows"] == [
        {"month": "2026-01", "revenue_东": "10.20", "revenue_西": "5"},
        {"month": "2026-02", "revenue_东": "12", "revenue_西": None},
    ]


def test_pivot_rejects_unbounded_or_ambiguous_columns(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], *[[f"region-{index}", "1"] for index in range(21)]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    plan = {
        "group_by": ["region"],
        "metrics": [{"column": "amount", "op": "sum", "as": "revenue"}],
        "pivot": {"column": "region", "metric": "revenue", "as_prefix": "revenue_"},
    }
    with pytest.raises(AnalysisError, match="pivot columns"):
        service.run(dataset["dataset_id"], plan)
    with pytest.raises(AnalysisError, match="Pivot requires"):
        service.run(dataset["dataset_id"], {**plan, "group_by": []})


def test_cross_currency_ratio_is_blocked_even_when_metric_filters_separate_them(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "amount"], ["CNY", "10"], ["USD", "5"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    plan = {
        "metrics": [
            {"column": "amount", "op": "sum", "as": "cny", "where": [{"column": "currency", "op": "eq", "value": "CNY"}]},
            {"column": "amount", "op": "sum", "as": "usd", "where": [{"column": "currency", "op": "eq", "value": "USD"}]},
        ],
        "ratios": [{"numerator": "cny", "denominator": "usd", "as": "bad_ratio"}],
    }
    with pytest.raises(AnalysisError, match="currenc"):
        service.run(dataset["dataset_id"], plan)


@pytest.mark.parametrize("operation", ["difference", "growth_rate"])
def test_comparison_rejects_count_against_known_currency_amount(tmp_path: Path, operation: str) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "amount"], ["CNY", "10"], ["CNY", "20"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="incompatible units"):
        service.run(dataset["dataset_id"], {
            "metrics": [
                {"op": "count", "as": "orders"},
                {"column": "amount", "op": "sum", "as": "revenue"},
            ],
            "comparisons": [{"current": "orders", "baseline": "revenue", "op": operation, "as": "invalid"}],
        })


def test_comparison_rejects_money_per_order_against_total_money(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "amount"], ["CNY", "10"], ["CNY", "20"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    plan = {
        "metrics": [
            {"op": "count", "as": "orders"},
            {"column": "amount", "op": "sum", "as": "revenue"},
        ],
        "ratios": [{"numerator": "revenue", "denominator": "orders", "as": "per_order"}],
    }
    assert service.run(dataset["dataset_id"], plan)["rows"] == [
        {"orders": 2, "revenue": "30", "per_order": "15"},
    ]
    with pytest.raises(AnalysisError, match="incompatible units"):
        service.run(dataset["dataset_id"], {
            **plan,
            "comparisons": [{"current": "per_order", "baseline": "revenue", "op": "difference", "as": "invalid"}],
        })


def test_scope_isolation_and_invalid_plan(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "value"], ["A", "1"]])
    first = AnalysisService(tmp_path / "state", scope="chat-a")
    second = AnalysisService(tmp_path / "state", scope="chat-b")
    dataset = first.ingest(source)
    with pytest.raises(AnalysisError, match="dataset"):
        second.run(dataset["dataset_id"], {"group_by": [], "metrics": [{"column": "value", "op": "sum"}]})
    with pytest.raises(AnalysisError, match="operation"):
        first.run(dataset["dataset_id"], {"python": "print(1)"})
    with pytest.raises(AnalysisError, match="column"):
        first.run(dataset["dataset_id"], {"group_by": ["missing"], "metrics": [{"column": "value", "op": "sum"}]})


def test_explicit_analysis_clear_only_removes_current_scope(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "value"], ["A", "1"]])
    first = AnalysisService(tmp_path / "state", scope="chat-a")
    second = AnalysisService(tmp_path / "state", scope="chat-b")
    first_dataset = first.ingest(source)
    second_dataset = second.ingest(source)
    result = first.run(first_dataset["dataset_id"], {"metrics": [{"column": "value", "op": "sum"}]})
    receipt = first.clear()
    assert receipt["datasets_removed"] == 1
    assert receipt["results_removed"] == 1
    assert receipt["snapshots_removed"] == 1
    assert receipt["external_files_removed"] is False
    assert not first.root.exists()
    assert source.is_file()
    with pytest.raises(AnalysisError, match="result"):
        first.get_result(result["analysis_id"])
    assert second.get_dataset(second_dataset["dataset_id"])["dataset_id"] == second_dataset["dataset_id"]
    assert first.clear()["datasets_removed"] == 0


def test_clear_waits_for_in_flight_import_before_removing_scope(tmp_path: Path, monkeypatch) -> None:
    from lantai.analysis import worker

    source = tmp_path / "data.csv"
    _csv(source, [["team", "value"], ["A", "1"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    started, release = threading.Event(), threading.Event()
    original = worker.run_worker

    def paused_worker(operation, **kwargs):
        if operation == "ingest":
            started.set()
            assert release.wait(5)
        return original(operation, **kwargs)

    monkeypatch.setattr(worker, "run_worker", paused_worker)
    with ThreadPoolExecutor(max_workers=2) as pool:
        import_task = pool.submit(service.ingest, source)
        assert started.wait(5)
        clear_task = pool.submit(service.clear)
        assert not clear_task.done()
        release.set()
        assert import_task.result(timeout=10)["dataset_id"]
        assert clear_task.result(timeout=10)["datasets_removed"] == 1
    assert not service.root.exists()


def test_filters_counts_average_and_zero_denominator(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["region", "revenue", "orders"], ["东", "10", "2"], ["东", "20", "0"], ["西", "5", "1"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "filters": [{"column": "region", "op": "eq", "value": "东"}],
        "group_by": ["region"],
        "metrics": [
            {"column": "revenue", "op": "sum", "as": "revenue"},
            {"column": "orders", "op": "sum", "as": "orders"},
            {"column": "revenue", "op": "avg", "as": "average"},
            {"op": "count", "as": "rows"},
        ],
        "ratios": [{"numerator": "revenue", "denominator": "orders", "as": "ticket"}],
    })
    assert result["rows"] == [{"region": "东", "revenue": "30", "orders": "2", "average": "15", "rows": 2, "ticket": "15"}]
    zero = service.run(dataset["dataset_id"], {
        "filters": [{"column": "revenue", "op": "eq", "value": "20"}],
        "metrics": [{"column": "orders", "op": "sum", "as": "orders"}, {"column": "revenue", "op": "sum", "as": "revenue"}],
        "ratios": [{"numerator": "revenue", "denominator": "orders", "as": "ratio"}],
    })
    assert zero["rows"][0]["ratio"] is None
    assert "zero or missing denominator" in zero["warnings"][0]


def test_xlsx_formulas_are_rejected(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "value"])
    sheet.append(["A", "=1+2"])
    book.save(source)
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    with pytest.raises(AnalysisError, match="formula"):
        service.ingest(source)


def test_duplicate_headers_and_oversize_are_rejected(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["name", "name"], ["A", "B"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    with pytest.raises(AnalysisError, match="duplicate"):
        service.ingest(source)


def test_numeric_sort_and_empty_filter_are_explicit(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["region", "amount"], ["A", "2"], ["B", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "group_by": ["region"],
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
        "sort": [{"column": "total", "descending": True}],
    })
    assert [row["region"] for row in result["rows"]] == ["B", "A"]
    empty = service.run(dataset["dataset_id"], {
        "filters": [{"column": "region", "op": "eq", "value": "Z"}],
        "metrics": [{"op": "count", "as": "rows"}],
    })
    assert empty["rows"] == [{"rows": 0}]


def test_formula_header_and_unbounded_source_status(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["=1+1", "value"])
    sheet.append(["A", 2])
    source = tmp_path / "formula-header.xlsx"
    workbook.save(source)
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    with pytest.raises(AnalysisError, match="formula"):
        service.ingest(source)

    csv_source = tmp_path / "data.csv"
    _csv(csv_source, [["name", "amount"], ["A", "1"]])
    dataset = service.ingest(csv_source)
    csv_source.write_bytes(b"x" * (20 * 1024 * 1024 + 1))
    assert service.source_status(dataset["dataset_id"])["status"] == "changed"


@pytest.mark.parametrize("plan", [
    {"group_by": [["team"]], "metrics": [{"op": "count"}]},
    {"metrics": [{"op": "sum", "column": "amount", "as": ["bad"]}]},
    {"metrics": [{"op": "count"}], "filters": [{"column": ["team"], "op": "eq", "value": "A"}]},
    {"metrics": [{"op": "count"}], "filters": [{"column": "team", "op": "in", "value": [["A"]]}]},
    {"metrics": [{"op": "count"}], "ratios": [{"numerator": ["x"], "denominator": "row_count", "as": "x"}]},
    {"metrics": [{"op": "count"}], "sort": [{"column": ["row_count"]}]},
])
def test_malformed_structured_plans_fail_cleanly(tmp_path: Path, plan: dict) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "amount"], ["A", "1"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError):
        service.run(dataset["dataset_id"], plan)


def test_descending_sort_keeps_missing_values_last(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "amount"], ["A", "2"], ["B", ""], ["C", "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "group_by": ["team"], "metrics": [{"op": "sum", "column": "amount", "as": "total"}],
        "sort": [{"column": "total", "descending": True}],
    })
    assert [row["team"] for row in result["rows"]] == ["C", "A", "B"]


def test_quality_and_verification_do_not_claim_business_review(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "amount"], ["A", "2"], ["A", ""]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    assert dataset["profile"]["null_counts"] == {"team": 0, "amount": 1}
    assert dataset["profile"]["grain"] == "unconfirmed"
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"op": "sum", "column": "amount", "as": "total"}],
    })
    assert result["rows"] == [{"total": "2"}]
    assert any("1 missing" in warning for warning in result["warnings"])
    assert result["verification"]["source_integrity"]["status"] == "passed"
    assert result["verification"]["calculation_validity"]["status"] == "passed"
    assert result["verification"]["semantic_review"]["status"] == "pending"
    assert result["verification"]["delivery_consistency"]["status"] == "not_applicable"


def test_saved_result_rejects_tampered_numbers_and_identity(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["team", "amount"], ["A", "2"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"op": "sum", "column": "amount", "as": "total"}],
    })
    result_path = service.root / "results" / f"{result['analysis_id']}.json"
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    payload["rows"][0]["total"] = "999"
    result_path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(AnalysisError, match="integrity"):
        service.get_result(result["analysis_id"])


def test_long_cells_and_hidden_xlsx_rows_are_not_silently_accepted(tmp_path: Path) -> None:
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    source = tmp_path / "long.csv"
    _csv(source, [["name", "value"], ["A", "x" * 1001]])
    with pytest.raises(AnalysisError, match="length"):
        service.ingest(source)

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "value"])
    sheet.append(["A", 2])
    sheet.row_dimensions[2].hidden = True
    hidden = tmp_path / "hidden.xlsx"
    workbook.save(hidden)
    with pytest.raises(AnalysisError, match="hidden"):
        service.ingest(hidden)

    hidden_sheet_book = Workbook()
    hidden_sheet = hidden_sheet_book.active
    hidden_sheet.append(["name", "value"])
    hidden_sheet.append(["A", 2])
    hidden_sheet.sheet_state = "hidden"
    hidden_sheet_book.create_sheet("visible")
    hidden_sheet_file = tmp_path / "hidden-sheet.xlsx"
    hidden_sheet_book.save(hidden_sheet_file)
    with pytest.raises(AnalysisError, match="hidden sheets"):
        service.ingest(hidden_sheet_file, sheet_name=hidden_sheet.title)


@pytest.mark.parametrize("new_part", [
    "xl/custom/sheet1.xml",
    "custom/sheet1.xml",
    "xl/custom/sheet1.dat",
])
@pytest.mark.parametrize("feature, message", [
    ("hidden_row", "hidden"),
    ("hidden_column", "hidden"),
    ("filter", "filters"),
    ("merge", "merged"),
])
def test_ambiguous_features_in_nonstandard_xlsx_worksheet_part_are_rejected(
    tmp_path: Path, new_part: str, feature: str, message: str,
) -> None:
    """OOXML relationships can point outside xl/worksheets/; scan the referenced sheet."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "value"])
    sheet.append(["A", 2])
    if feature == "hidden_row":
        sheet.row_dimensions[2].hidden = True
    elif feature == "hidden_column":
        sheet.column_dimensions["B"].hidden = True
    elif feature == "filter":
        sheet.auto_filter.ref = "A1:B2"
    else:
        sheet.merge_cells("A2:B2")
    source_bytes = io.BytesIO()
    workbook.save(source_bytes)
    old_part = "xl/worksheets/sheet1.xml"
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source_bytes.getvalue())) as original:
        with zipfile.ZipFile(rewritten, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for item in original.infolist():
                content = original.read(item.filename)
                if item.filename in {"xl/_rels/workbook.xml.rels", "[Content_Types].xml"}:
                    content = content.replace(old_part.encode(), new_part.encode())
                archive.writestr(new_part if item.filename == old_part else item.filename, content)
    source = tmp_path / "nonstandard-sheet.xlsx"
    source.write_bytes(rewritten.getvalue())
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match=message):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


@pytest.mark.parametrize("member", [
    "xl/externalLinks/externalLink1.xml",
    "xl/ExternalLinks/externalLink1.xml",
    "xl/embeddings/oleObject1.BIN",
])
def test_xlsx_link_or_binary_part_is_rejected_before_persistence(tmp_path: Path, member: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "value"])
    sheet.append(["A", 2])
    source_bytes = io.BytesIO()
    workbook.save(source_bytes)
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source_bytes.getvalue())) as original:
        with zipfile.ZipFile(rewritten, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for item in original.infolist():
                archive.writestr(item, original.read(item.filename))
            archive.writestr(member, b"<externalLink/>")
    source = tmp_path / "linked.xlsx"
    source.write_bytes(rewritten.getvalue())
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match="external links or embedded binaries"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


def test_xlsx_with_excessive_zip_members_is_rejected_before_persistence(tmp_path: Path) -> None:
    """Small OOXML archives can still contain enough parts to exhaust a scan."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["name", "value"])
    sheet.append(["A", 2])
    source_bytes = io.BytesIO()
    workbook.save(source_bytes)
    rewritten = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source_bytes.getvalue())) as original:
        with zipfile.ZipFile(rewritten, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for item in original.infolist():
                archive.writestr(item, original.read(item.filename))
            for index in range(4096):
                archive.writestr(f"padding/{index:04d}.xml", b"<padding/>")
    source = tmp_path / "many-parts.xlsx"
    source.write_bytes(rewritten.getvalue())
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match="package part count exceeds the safety limit"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_blank_rows_count_toward_parse_scan_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, suffix: str,
) -> None:
    """Empty physical rows must not bypass the work budget for parsing."""

    monkeypatch.setattr("lantai.analysis.service.MAX_SCANNED_ROWS", 3, raising=False)
    source = tmp_path / f"sparse{suffix}"
    records = [["name", "value"], ["A", "1"], ["", ""], ["", ""], ["", ""], ["B", "2"]]
    if suffix == ".csv":
        _csv(source, records)
    else:
        workbook = Workbook()
        sheet = workbook.active
        for record in records:
            sheet.append(record)
        workbook.save(source)
    with pytest.raises(AnalysisError, match="scanned row count exceeds the safety limit"):
        _parse_table(source.read_bytes(), suffix, None)


def test_sparse_csv_scan_limit_rejects_worker_ingest_before_persistence(tmp_path: Path) -> None:
    source = tmp_path / "sparse.csv"
    source.write_text("name,value\nA,1\n" + ",\n" * 250_000 + "B,2\n", encoding="utf-8")
    service = AnalysisService(tmp_path / "state", scope="chat-a")

    with pytest.raises(AnalysisError, match="scanned row count exceeds the safety limit"):
        service.ingest(source)
    assert list((service.root / "datasets").glob("*.json")) == []


def test_date_bucket_conditional_metrics_and_period_change(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [
        ["date", "region", "status", "revenue"],
        ["2026-01-03", "东", "paid", "10.20"],
        ["2026-01-28", "东", "cancelled", "50"],
        ["2026-02-03", "东", "paid", "15.30"],
        ["2026-02-13", "西", "paid", "5"],
    ])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "date_buckets": [{"column": "date", "grain": "month", "as": "month"}],
        "group_by": ["month"],
        "metrics": [
            {"column": "revenue", "op": "sum", "as": "paid_revenue", "where": [{"column": "status", "op": "eq", "value": "paid"}]},
            {"op": "count", "as": "orders"},
        ],
        "sort": [{"column": "month"}],
    })
    assert result["rows"] == [
        {"month": "2026-01", "paid_revenue": "10.20", "orders": 2},
        {"month": "2026-02", "paid_revenue": "20.30", "orders": 2},
    ]
    comparison = service.run(dataset["dataset_id"], {
        "metrics": [
            {"column": "revenue", "op": "sum", "as": "current", "where": [{"column": "date", "op": "gte", "value": "2026-02-01"}]},
            {"column": "revenue", "op": "sum", "as": "baseline", "where": [{"column": "date", "op": "lt", "value": "2026-02-01"}]},
        ],
        "comparisons": [
            {"current": "current", "baseline": "baseline", "op": "difference", "as": "delta"},
            {"current": "current", "baseline": "baseline", "op": "growth_rate", "as": "growth"},
        ],
    })
    assert comparison["rows"] == [{"current": "20.30", "baseline": "60.20", "delta": "-39.90", "growth": "-0.6627906976744186046511627907"}]


def test_date_bucket_rejects_ambiguous_dates_and_negative_growth_baseline(tmp_path: Path) -> None:
    source = tmp_path / "data.csv"
    _csv(source, [["date", "amount"], ["03/04/2026", "-1"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="ISO"):
        service.run(dataset["dataset_id"], {
            "date_buckets": [{"column": "date", "grain": "month", "as": "month"}],
            "group_by": ["month"], "metrics": [{"op": "count"}],
        })
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"column": "amount", "op": "sum", "as": "base"}, {"op": "count", "as": "current"}],
        "comparisons": [{"current": "current", "baseline": "base", "op": "growth_rate", "as": "growth"}],
    })
    assert result["rows"] == [{"base": "-1", "current": 1, "growth": None}]
    assert any("negative baseline" in warning for warning in result["warnings"])


def test_xlsx_calendar_datetime_can_be_bucketed_without_timezone_guess(tmp_path: Path) -> None:
    from datetime import datetime

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["when", "amount"])
    sheet.append([datetime(2026, 2, 3, 9, 30), 12.5])
    source = tmp_path / "dates.xlsx"
    workbook.save(source)
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "date_buckets": [{"column": "when", "grain": "quarter", "as": "quarter"}],
        "group_by": ["quarter"],
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    assert result["rows"] == [{"quarter": "2026-Q1", "total": "12.5"}]


def test_delivery_receipt_detects_external_edit_and_is_scope_bound(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "12.50"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    other = AnalysisService(tmp_path / "state", scope="chat-b")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    artifact = tmp_path / "report.docx"
    artifact.write_bytes(b"verified report content")
    original_read_bytes = Path.read_bytes

    def reject_unbounded_artifact_read(path: Path) -> bytes:
        if path == artifact:
            raise AssertionError("unbounded read of analysis delivery")
        return original_read_bytes(path)

    monkeypatch.setattr(Path, "read_bytes", reject_unbounded_artifact_read)
    receipt = service.bind_delivery(result["analysis_id"], artifact, kind="docx", bound_elements=["result-table"])
    assert receipt["verification"]["delivery_consistency"]["status"] == "pending"
    assert service.delivery_status(receipt["delivery_id"])["status"] == "current"
    assert service.delivery_status(receipt["delivery_id"])["source_status"] == "current"
    with pytest.raises(AnalysisError, match="delivery"):
        other.delivery_status(receipt["delivery_id"])
    _csv(source, [["region", "amount"], ["东", "99"]])
    assert service.delivery_status(receipt["delivery_id"])["source_status"] == "changed"
    artifact.write_bytes(b"changed outside Lantai")
    assert service.delivery_status(receipt["delivery_id"])["status"] == "modified_unverified"
    assert service.get_result(result["analysis_id"])["rows"] == [{"total": "12.50"}]


def test_possible_total_row_blocks_double_count_until_explicitly_filtered(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["region", "amount"], ["东", "10"], ["西", "20"], ["总计", "30"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    assert dataset["profile"]["possible_total_row_count"] == 1
    with pytest.raises(AnalysisError, match="total or subtotal"):
        service.run(dataset["dataset_id"], {"metrics": [{"column": "amount", "op": "sum", "as": "total"}]})
    result = service.run(dataset["dataset_id"], {
        "filters": [{"column": "region", "op": "ne", "value": "总计"}],
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    assert result["rows"] == [{"total": "30"}]


def test_total_marker_in_later_dimension_still_blocks_double_count(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["month", "region", "amount"], ["2026-01", "东", "10"], ["2026-01", "西", "20"], ["2026-01", "总计", "30"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    assert dataset["profile"]["possible_total_row_count"] == 1
    with pytest.raises(AnalysisError, match="total or subtotal"):
        service.run(dataset["dataset_id"], {"metrics": [{"column": "amount", "op": "sum", "as": "total"}]})
    result = service.run(dataset["dataset_id"], {
        "filters": [{"column": "region", "op": "ne", "value": "总计"}],
        "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    assert result["rows"] == [{"total": "30"}]


def test_current_period_is_flagged_without_inventing_a_full_month(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    current_day = date.today().isoformat()
    _csv(source, [["date", "amount"], [current_day, "10"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    assert dataset["profile"]["candidate_types"]["date"] == "date"
    assert any("incomplete" in note for note in dataset["profile"]["notes"])
    result = service.run(dataset["dataset_id"], {
        "date_buckets": [{"column": "date", "grain": "month", "as": "month"}],
        "group_by": ["month"], "metrics": [{"column": "amount", "op": "sum", "as": "total"}],
    })
    assert result["rows"] == [{"month": current_day[:7], "total": "10"}]
    assert any("incomplete" in warning for warning in result["warnings"])


def test_share_rank_and_explicit_deduplication_are_deterministic(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["order", "region", "revenue"], ["1", "East", "10"], ["1", "East", "10"], ["2", "West", "5"], ["3", "North", "5"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    result = service.run(dataset["dataset_id"], {
        "deduplicate_by": ["order"],
        "group_by": ["region"],
        "metrics": [{"column": "revenue", "op": "sum", "as": "Revenue"}],
        "shares": [{"metric": "Revenue", "as": "Share"}],
        "ranks": [{"metric": "Revenue", "as": "Rank", "descending": True, "method": "competition"}],
        "sort": [{"column": "Rank"}],
    })
    assert result["rows"] == [
        {"region": "East", "Revenue": "10", "Share": "0.5", "Rank": 1},
        {"region": "West", "Revenue": "5", "Share": "0.25", "Rank": 2},
        {"region": "North", "Revenue": "5", "Share": "0.25", "Rank": 2},
    ]
    assert any("duplicate" in warning.lower() for warning in result["warnings"])


def test_deduplication_refuses_conflicting_rows_and_share_refuses_invalid_denominator(tmp_path: Path) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["order", "region", "revenue"], ["1", "East", "10"], ["1", "East", "20"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="conflicting"):
        service.run(dataset["dataset_id"], {
            "deduplicate_by": ["order"],
            "group_by": ["region"],
            "metrics": [{"column": "revenue", "op": "sum", "as": "Revenue"}],
        })
    with pytest.raises(AnalysisError, match="positive"):
        service.run(dataset["dataset_id"], {
            "group_by": ["region"],
            "metrics": [{"column": "revenue", "op": "sum", "as": "Revenue", "where": [
                {"column": "region", "op": "eq", "value": "Absent"},
            ]}],
            "shares": [{"metric": "Revenue", "as": "Share"}],
        })


@pytest.mark.parametrize("derived", [
    {"shares": [{"metric": "Revenue", "as": "Share"}]},
    {"ranks": [{"metric": "Revenue", "as": "Rank", "descending": True, "method": "dense"}]},
])
def test_cross_currency_derived_group_metrics_are_blocked(tmp_path: Path, derived: dict) -> None:
    source = tmp_path / "sales.csv"
    _csv(source, [["currency", "region", "revenue"], ["CNY", "East", "10"], ["USD", "West", "20"]])
    service = AnalysisService(tmp_path / "state", scope="chat-a")
    dataset = service.ingest(source)
    with pytest.raises(AnalysisError, match="cross-currency"):
        service.run(dataset["dataset_id"], {
            "group_by": ["region"],
            "metrics": [{"column": "revenue", "op": "sum", "as": "Revenue"}],
            **derived,
        })
