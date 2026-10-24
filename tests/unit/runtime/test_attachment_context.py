# -*- coding: utf-8 -*-
"""Behavior locks for model-readable uploaded document context."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from lantai.runtime.attachment_context import (
    MAX_ATTACHMENT_BYTES,
    MAX_ATTACHMENT_CONTEXT_CHARS,
    attachment_context_text,
)


def test_attachment_context_accepts_the_chat_upload_limit() -> None:
    assert MAX_ATTACHMENT_BYTES == 150 * 1024 * 1024


def _file_block(path: Path) -> dict:
    return {
        "filename": path.name,
        "source": {"type": "url", "url": str(path)},
    }


def test_text_attachment_is_readable_and_keeps_local_location(tmp_path: Path) -> None:
    path = tmp_path / "说明.md"
    path.write_text("# 澜台附件\n\n这是正文。", encoding="utf-8")

    context = attachment_context_text(_file_block(path))

    assert "# 澜台附件" in context
    assert "这是正文。" in context
    assert str(path.resolve()) in context


def test_modern_office_attachments_are_converted_to_readable_text(
    tmp_path: Path,
) -> None:
    docx_path = tmp_path / "合同.docx"
    document = Document()
    document.add_paragraph("Word 合同正文")
    document.save(docx_path)

    pptx_path = tmp_path / "汇报.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPT 汇报标题"
    presentation.save(pptx_path)

    xlsx_path = tmp_path / "数据.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "Excel 指标"
    workbook.active["B1"] = 42
    workbook.save(xlsx_path)

    assert "Word 合同正文" in attachment_context_text(_file_block(docx_path))
    assert "PPT 汇报标题" in attachment_context_text(_file_block(pptx_path))
    xlsx_context = attachment_context_text(_file_block(xlsx_path))
    assert "Excel 指标" in xlsx_context
    assert "42" in xlsx_context


def test_pdf_attachment_text_is_extracted(tmp_path: Path) -> None:
    path = tmp_path / "report.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Lantai PDF attachment content")
    pdf.save()

    context = attachment_context_text(_file_block(path))

    assert "Lantai PDF attachment content" in context


def test_zip_attachment_lists_files_and_previews_text_members(
    tmp_path: Path,
) -> None:
    path = tmp_path / "project.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("README.md", "# Project\nRun the demo.")
        archive.writestr("assets/logo.bin", b"\x00\x01")

    context = attachment_context_text(_file_block(path))

    assert "Archive contents:" in context
    assert "README.md" in context
    assert "# Project" in context
    assert "Run the demo." in context
    assert "assets/logo.bin" in context


def test_large_text_attachment_is_streamed_and_bounded(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = tmp_path / "large.txt"
    path.write_text("A" * (MAX_ATTACHMENT_CONTEXT_CHARS + 50_000), encoding="utf-8")

    def forbid_unbounded_read(_path: Path) -> bytes:
        raise AssertionError("attachment extraction must not call Path.read_bytes")

    monkeypatch.setattr(Path, "read_bytes", forbid_unbounded_read)
    context = attachment_context_text(_file_block(path))

    assert "[Attachment text truncated by Lantai.]" in context
    assert len(context) < MAX_ATTACHMENT_CONTEXT_CHARS + 500


def test_zip_attachment_opens_archive_path_without_loading_whole_file(
    tmp_path: Path,
    monkeypatch,
) -> None:
    path = tmp_path / "project-large.zip"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("README.md", "# Streamed archive")

    def forbid_unbounded_read(_path: Path) -> bytes:
        raise AssertionError("zip extraction must not call Path.read_bytes")

    monkeypatch.setattr(Path, "read_bytes", forbid_unbounded_read)
    context = attachment_context_text(_file_block(path))

    assert "# Streamed archive" in context
