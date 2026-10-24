"""Artifact scope proofs do not replace Provider, Office or human acceptance."""

from __future__ import annotations

import hashlib
import importlib.util
import json
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image
import pytest

from lantai.presentation.design_system import prepare_deck_for_composition
from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    AssetRecord, DeckIR, DeckSlide, ImageObject, ObjectFrame, PresentationEditOperation, TextObject,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


SCRIPT = Path(__file__).resolve().parents[3] / "scripts/verify_presentation_native_edit_scope.py"


def _module():
    spec = importlib.util.spec_from_file_location("native_edit_scope", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _replace_member(path: Path, name: str, payload: bytes):
    staged = path.with_suffix(".stage.pptx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(staged, "w") as target:
        for item in source.infolist():
            target.writestr(item, payload if item.filename == name else source.read(item.filename))
    staged.replace(path)


def _fixture(tmp_path: Path, *, with_image: bool = False):
    deck = DeckIR(id="scope-deck", title="Synthetic scope proof", slides=[
        DeckSlide(id=f"page-{n}", layout_id="freeform", objects=[TextObject(
            id=f"title-{n}", text=f"Synthetic page {n}", role="title", font_size=42,
            frame=ObjectFrame(x=1, y=1, w=10, h=1),
        )]) for n in range(1, 4)
    ])
    if with_image:
        asset_path = tmp_path / "synthetic-image.png"
        Image.new("RGB", (16, 8), "navy").save(asset_path)
        deck.assets = [AssetRecord(
            asset_id="synthetic-image", local_path=str(asset_path), source_ref="synthetic-fixture",
            rights_status="generated", sha256=hashlib.sha256(asset_path.read_bytes()).hexdigest(),
            dimensions=(16, 8),
        )]
        deck.slides[2].objects.append(ImageObject(
            id="image-3", asset_id="synthetic-image", path=str(asset_path), alt="Synthetic image",
            frame=ObjectFrame(x=1, y=3, w=4, h=2),
        ))
    before = tmp_path / "before.pptx"
    compiler = NativePptxCompiler()
    compiler.compile(deck, before)
    canonical = prepare_deck_for_composition(reconcile_pptx(deck, before, assets_dir=tmp_path / "assets"))
    changed = canonical.model_copy(deep=True)
    changed.slides[0].objects[0].text = "Revised synthetic page"
    # Match the production text-edit planner: old rich runs cannot override
    # the newly requested wording in the native compiler.
    changed.slides[0].objects[0].runs = []
    changed.slides[0].revision += 1
    changed = prepare_deck_for_composition(changed)
    operation = PresentationEditOperation(
        deck_id=deck.id, slide_ids=["page-1"], object_ids=["title-1"],
        allowed_content_types={"text"}, instruction="Synthetic scoped revision",
    )
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        before, canonical, changed, operation, compiler, compiled_slide_ids=["page-1"],
    )
    from pptx import Presentation
    assert any(shape.has_text_frame and shape.text == "Revised synthetic page"
               for shape in Presentation(preview.candidate_path).slides[0].shapes)
    before_ir, after_ir = tmp_path / "before.json", tmp_path / "after.json"
    before_ir.write_text(deck.model_dump_json(), encoding="utf-8")
    after_ir.write_text(changed.model_dump_json(), encoding="utf-8")
    old_renders, new_renders = tmp_path / "old-renders", tmp_path / "new-renders"
    old_renders.mkdir()
    new_renders.mkdir()
    for n in range(1, 4):
        # Tiny synthetic pixels test the verifier, not an Office render claim.
        Image.new("RGBA", (16, 8), (255, 255, 255, 255)).save(old_renders / f"slide-{n}.png")
        Image.new("RGBA", (16, 8), (0, 0, 255, 255) if n == 1 else (255, 255, 255, 255)).save(new_renders / f"slide-{n}.png")
    return dict(before_pptx=before, after_pptx=Path(preview.candidate_path),
                before_deck_ir=before_ir, after_deck_ir=after_ir,
                before_render_dir=old_renders, after_render_dir=new_renders,
                allowed_slide_ids=["page-1", "page-2"], required_slide_ids=["page-1"])


def test_scoped_native_edit_accepts_import_bookkeeping_without_forcing_all_allowed_pages(tmp_path):
    args = _fixture(tmp_path)
    before_files = {p: hashlib.sha256(p.read_bytes()).hexdigest()
                    for p in tmp_path.rglob("*") if p.is_file()}
    result = _module().verify_native_edit_scope(**args)
    assert result["ok"] is True
    assert result["non_target_count"] == 1
    assert result["required_slide_ids"] == ["page-1"]
    assert result["slides"][1]["canonical_ir_equal"] is True
    assert result["slides"][1]["package_dependencies_equal"] is True
    assert result["slides"][1]["render_pixels_equal"] is True
    assert result["full_case_or_office_app_human_acceptance_claimed"] is False
    assert result["provider_calls"] == 0
    assert {p: hashlib.sha256(p.read_bytes()).hexdigest()
            for p in tmp_path.rglob("*") if p.is_file()} == before_files


def test_content_types_cannot_hide_a_changed_existing_non_target_declaration(tmp_path):
    args = _fixture(tmp_path)
    with zipfile.ZipFile(args["after_pptx"]) as package:
        root = ET.fromstring(package.read("[Content_Types].xml"))
    item = next(n for n in root if n.get("PartName") == "/ppt/slides/slide3.xml")
    item.set("ContentType", "application/synthetic-wrong-type")
    _replace_member(args["after_pptx"], "[Content_Types].xml", ET.tostring(root))
    result = _module().verify_native_edit_scope(**args)
    assert result["ok"] is False
    assert result["content_types_scoped"] is False


@pytest.mark.parametrize("change", ["rgb", "alpha", "dimension", "ir", "shared_theme", "global_metadata"])
def test_non_target_changes_fail_even_when_other_evidence_is_unchanged(tmp_path, change):
    args = _fixture(tmp_path)
    if change in {"rgb", "alpha", "dimension"}:
        color = (255, 255, 255, 0) if change == "alpha" else (255, 0, 0, 255)
        size = (17, 8) if change == "dimension" else (16, 8)
        Image.new("RGBA", size, color).save(args["after_render_dir"] / "slide-3.png")
    elif change == "ir":
        after = DeckIR.model_validate_json(args["after_deck_ir"].read_text())
        after.slides[2].objects[0].text = "Unexpected unrelated revision"
        args["after_deck_ir"].write_text(after.model_dump_json(), encoding="utf-8")
    else:
        part = "ppt/theme/theme1.xml" if change == "shared_theme" else "docProps/core.xml"
        with zipfile.ZipFile(args["after_pptx"]) as package:
            payload = package.read(part) + b"\n"
        _replace_member(args["after_pptx"], part, payload)
    result = _module().verify_native_edit_scope(**args)
    assert result["ok"] is False
    row = result["slides"][2]
    if change == "ir":
        assert row["canonical_ir_equal"] is False
    elif change == "shared_theme":
        assert row["package_dependencies_equal"] is False
    elif change == "global_metadata":
        assert result["package_changes_scoped"] is False
    else:
        assert row["render_pixels_equal"] is False


def test_revision_only_changes_do_not_satisfy_a_required_page(tmp_path):
    args = _fixture(tmp_path)
    after = DeckIR.model_validate_json(args["after_deck_ir"].read_text())
    after.slides[1].revision += 1
    args["after_deck_ir"].write_text(after.model_dump_json(), encoding="utf-8")
    args["required_slide_ids"].append("page-2")
    result = _module().verify_native_edit_scope(**args)
    assert result["ok"] is False
    assert result["required_slides_materially_changed"] is False
    assert result["slides"][1]["material_ir_changed"] is False


@pytest.mark.parametrize("allowed,required", [([], []), (["page-4"], []),
    (["page-1", "page-1"], []), (["page-1"], ["page-2"]),
    (["page-1"], ["page-1", "page-1"])])
def test_invalid_scope_fails_closed(tmp_path, allowed, required):
    args = _fixture(tmp_path)
    args.update(allowed_slide_ids=allowed, required_slide_ids=required)
    with pytest.raises(ValueError):
        _module().verify_native_edit_scope(**args)


def test_every_render_input_is_hash_bound_without_copying_slide_content(tmp_path):
    args = _fixture(tmp_path)
    result = _module().verify_native_edit_scope(**args)
    for n, row in enumerate(result["slides"], 1):
        assert row["before_render_sha256"] == hashlib.sha256(
            (args["before_render_dir"] / f"slide-{n}.png").read_bytes()).hexdigest()
        assert row["after_render_sha256"] == hashlib.sha256(
            (args["after_render_dir"] / f"slide-{n}.png").read_bytes()).hexdigest()
    assert "Synthetic page" not in json.dumps(result)
    assert "Revised synthetic page" not in json.dumps(result)


def test_verified_image_locations_are_normalized_without_erasing_provenance(tmp_path):
    args = _fixture(tmp_path, with_image=True)
    result = _module().verify_native_edit_scope(**args)
    assert result["ok"] is True
    assert result["slides"][2]["canonical_ir_equal"] is True
    after = DeckIR.model_validate_json(args["after_deck_ir"].read_text())
    image = next(obj for obj in after.slides[2].objects if obj.kind == "image")
    next(asset for asset in after.assets if asset.asset_id == image.asset_id).source_ref = "altered-provenance"
    args["after_deck_ir"].write_text(after.model_dump_json(), encoding="utf-8")
    assert _module().verify_native_edit_scope(**args)["ok"] is False


def test_an_image_path_is_not_normalized_until_its_bytes_match_the_asset_hash(tmp_path):
    args = _fixture(tmp_path, with_image=True)
    after = DeckIR.model_validate_json(args["after_deck_ir"].read_text())
    image = next(obj for obj in after.slides[2].objects if obj.kind == "image")
    Image.new("RGB", (16, 8), "red").save(image.path)
    with pytest.raises(ValueError, match="bytes do not match"):
        _module().verify_native_edit_scope(**args)


def _cli_argv(args, output):
    command = [str(SCRIPT)]
    for name, value in args.items():
        flag = "--" + name.replace("_", "-")
        if name in {"allowed_slide_ids", "required_slide_ids"}:
            flag = flag[:-1]
            for slide in value:
                command.extend([flag, slide])
        else:
            command.extend([flag, str(value)])
    return [*command, "--output", str(output)]


def test_cli_does_not_write_evidence_into_a_render_input_directory(tmp_path, monkeypatch, capsys):
    args = _fixture(tmp_path)
    target = args["before_render_dir"] / "scope-proof.json"
    monkeypatch.setattr(sys, "argv", _cli_argv(args, target))
    assert _module().main() == 2
    assert not target.exists()
    assert json.loads(capsys.readouterr().out)["state"] == "invalid_scope_evidence"


def test_cli_does_not_overwrite_existing_input_or_print_invalid_ir_content(tmp_path, monkeypatch, capsys):
    args = _fixture(tmp_path)
    original = args["before_deck_ir"].read_bytes()
    monkeypatch.setattr(sys, "argv", _cli_argv(args, args["before_deck_ir"]))
    assert _module().main() == 2
    assert args["before_deck_ir"].read_bytes() == original
    assert json.loads(capsys.readouterr().out)["error_type"] in {"ValueError", "FileExistsError"}
    marker = "SYNTHETIC_PRIVATE_DIAGNOSTIC_NOT_TO_PRINT"
    args["after_deck_ir"].write_text(marker, encoding="utf-8")
    monkeypatch.setattr(sys, "argv", _cli_argv(args, tmp_path / "proof.json"))
    assert _module().main() == 2
    output = capsys.readouterr().out
    assert marker not in output
    assert len(json.loads(output)["error_sha256"]) == 64
