"""Local edits must not reserialize unrequested native objects."""

from hashlib import sha256
from zipfile import ZipFile

from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches
import pytest
from PIL import Image

from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    AddObjectAction, AssetRecord, ChartObject, CompositionAuthority, DeckIR, DeckSlide,
    ImageObject, NativeCompositionTokens, ObjectFrame,
    PresentationEditOperation, PresentationObjectEditPlan, RemoveObjectAction,
    ReplaceTextAction, ShapeObject, TextObject, UpdateChartAction,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _chart_package_parts(path):
    with ZipFile(path) as package:
        return {name: package.read(name) for name in package.namelist()
                if name.startswith(("ppt/charts/", "ppt/embeddings/"))}


@pytest.mark.parametrize("has_marker", [True, False])
def test_local_native_edit_keeps_untouched_xml_and_existing_shape_ids(tmp_path, has_marker):
    deck = DeckIR(id="native-preservation", title="合成局部编辑", slides=[DeckSlide(
        id="page", layout_id="freeform",
        composition_authority=CompositionAuthority.AUTHOR, objects=[
            ShapeObject(id="remove", frame=ObjectFrame(x=1, y=1, w=1, h=1)),
            TextObject(id="edit", role="body", text="旧标签", font_size=18,
                       frame=ObjectFrame(x=3, y=1, w=3, h=1)),
            TextObject(id="keep", role="body", text="保留原生段落\n\n保留原生格式",
                       font_size=18, frame=ObjectFrame(x=1, y=3, w=5, h=2)),
        ],
    )])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    if not has_marker:
        external = Presentation(source)
        marker = external.slides[0].shapes[-1]._element
        marker.getparent().remove(marker)
        external.save(source)
    original_bytes = source.read_bytes()
    original = {s.name: s for s in Presentation(source).slides[0].shapes}
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(
        slide_id="page", actions=[
            ReplaceTextAction(object_id="edit", text="新标签"),
            RemoveObjectAction(object_id="remove"),
            AddObjectAction(object=ShapeObject(
                id="added", frame=ObjectFrame(x=7, y=3, w=1, h=1),
            )),
        ],
    ))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, imported, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["page"], object_ids=ids,
            allowed_content_types={"text", "shape"}, instruction="只修改标签并替换方块",
        ), compiler,
    )

    reopened = Presentation(preview.candidate_path)
    actual = {s.name: s for s in reopened.slides[0].shapes}
    assert actual["keep"]._element.xml == original["keep"]._element.xml
    assert actual["edit"].shape_id == original["edit"].shape_id
    if has_marker:
        assert actual["meta:native-preservation:page"]._element.xml == original[
            "meta:native-preservation:page"
        ]._element.xml
    assert actual["edit"].text == "新标签"
    assert "remove" not in actual and "added" in actual
    native_ids = [s.shape_id for s in reopened.slides[0].shapes]
    assert len(native_ids) == len(set(native_ids))
    assert source.read_bytes() == original_bytes


def test_local_edit_preserves_picture_and_chart_relationships_when_ids_collide(tmp_path):
    pictures = []
    assets = []
    for name, color in (("removed", "red"), ("kept", "blue"), ("added", "green")):
        path = tmp_path / f"{name}.png"
        Image.new("RGB", (16, 16), color).save(path)
        assets.append(AssetRecord(
            asset_id=f"asset-{name}", local_path=str(path), source_ref="synthetic",
            rights_status="generated", sha256=sha256(path.read_bytes()).hexdigest(),
            dimensions=(16, 16),
        ))
        pictures.append(ImageObject(id=name, path=str(path), alt=name,
                                    frame=ObjectFrame(x=1, y=1, w=1, h=1)))
    chart = ChartObject(id="chart", categories=["A", "B"], series={"合成数量": [12, 34]},
                        unit="个", derivation="合成原始数据直接映射", source_refs=["synthetic"],
                        frame=ObjectFrame(x=3, y=1, w=6, h=4))
    deck = DeckIR(id="relationships", title="合成依赖", assets=assets, slides=[DeckSlide(
        id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
        objects=[*pictures[:2], chart],
    )])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    original_bytes = source.read_bytes()
    original = {s.name: s for s in Presentation(source).slides[0].shapes}
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(
        slide_id="page", actions=[RemoveObjectAction(object_id="removed"),
                                  AddObjectAction(object=pictures[2])],
    ))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, imported, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["page"], object_ids=ids,
            allowed_content_types={"image", "chart"}, instruction="只替换第一张图片",
        ), compiler,
    )
    actual = {s.name: s for s in Presentation(preview.candidate_path).slides[0].shapes}
    for name in ("kept", "chart"):
        assert actual[name]._element.xml == original[name]._element.xml
    assert actual["kept"].image.blob == original["kept"].image.blob
    assert actual["added"].image.blob == (tmp_path / "added.png").read_bytes()
    assert actual["chart"].chart.part.blob == original["chart"].chart.part.blob
    assert actual["chart"].chart.series[0].values == (12.0, 34.0)
    assert _chart_package_parts(preview.candidate_path) == _chart_package_parts(source)
    assert source.read_bytes() == original_bytes


@pytest.mark.parametrize("change_chart", [False, True])
def test_caption_edit_only_imports_chart_dependencies_when_chart_is_changed(tmp_path, change_chart):
    chart = ChartObject(
        id="chart", categories=["A", "B"], series={"合成数量": [12, 34]},
        unit="个", derivation="合成原始数据直接映射", source_refs=["synthetic"],
        frame=ObjectFrame(x=1, y=1, w=6, h=4),
    )
    deck = DeckIR(id="chart-caption", title="合成图表说明", slides=[
        DeckSlide(id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
                  objects=[chart, TextObject(
                      id="caption", role="body", text="旧说明", font_size=18,
                      frame=ObjectFrame(x=1, y=5, w=6, h=1),
                  )]),
        DeckSlide(id="untouched", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
                  objects=[TextObject(id="keep", role="body", text="未涉及的页面", font_size=18,
                                      frame=ObjectFrame(x=1, y=1, w=6, h=1))]),
    ])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    # Existing unreferenced native parts belong to the user, not to our edit.
    with ZipFile(source, "a") as package:
        workbook = next(name for name in package.namelist() if name.endswith(".xlsx"))
        package.writestr("ppt/embeddings/original-unused.xlsx", package.read(workbook))
    original_bytes = source.read_bytes()
    original_parts = _chart_package_parts(source)
    original_chart = next(s for s in Presentation(source).slides[0].shapes if s.has_chart)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    actions = [ReplaceTextAction(object_id="caption", text="新说明：单位为个")]
    if change_chart:
        actions.append(UpdateChartAction(
            object_id="chart", categories=["A", "B"], series={"合成数量": [22, 44]},
            unit="个", derivation="合成更新数据直接映射", source_refs=["synthetic"],
        ))
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(slide_id="page", actions=actions))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, imported, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["page"], object_ids=ids,
            allowed_content_types={"text", "chart"}, instruction="修改图表说明，仅显式要求时更新图表数据",
        ), compiler,
    )
    parts = _chart_package_parts(preview.candidate_path)
    assert all(parts[name] == payload for name, payload in original_parts.items())
    actual = {s.name: s for s in Presentation(preview.candidate_path).slides[0].shapes}
    assert actual["caption"].text == "新说明：单位为个"
    if change_chart:
        new_parts = set(parts) - original_parts.keys()
        assert len([name for name in new_parts if name.startswith("ppt/charts/") and name.endswith(".xml")]) == 1
        assert len([name for name in new_parts if name.endswith(".xlsx")]) == 1
        assert actual["chart"].chart.series[0].values == (22.0, 44.0)
    else:
        assert parts == original_parts
        assert actual["chart"]._element.xml == original_chart._element.xml
        assert actual["chart"].chart.series[0].values == (12.0, 34.0)
    with ZipFile(source) as before, ZipFile(preview.candidate_path) as after:
        changed_parts = {"[Content_Types].xml", "ppt/slides/slide1.xml", "ppt/slides/_rels/slide1.xml.rels"}
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in changed_parts)
    assert source.read_bytes() == original_bytes


def test_inherited_palette_redesign_does_not_freeze_equal_object_ir(tmp_path):
    deck = DeckIR(id="context", title="合成主题变更", slides=[DeckSlide(
        id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
        objects=[TextObject(id="text", text="原生文本", role="body", font_size=18,
                            frame=ObjectFrame(x=1, y=1, w=5, h=1))],
    )])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    before = Presentation(source).slides[0].shapes[0]
    changed = deck.model_copy(update={
        "slides": [deck.slides[0].model_copy(update={"revision": deck.slides[0].revision + 1})],
        "composition_tokens": NativeCompositionTokens(
        canvas="FFFFFF", ink="123456", muted="607582", accent="327A72", subtle="E8ECEF",
        title_font="Noto Sans SC", body_font="Noto Sans SC", title_size=28,
        body_size=18, caption_size=12,
    )})
    assert changed.slides[0].objects == deck.slides[0].objects
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, deck, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["page"], object_ids=["text"],
            allowed_content_types={"text"}, instruction="整页采用新的合成配色",
        ), compiler,
    )
    actual = Presentation(preview.candidate_path).slides[0].shapes[0]
    assert str(actual.text_frame.paragraphs[0].runs[0].font.color.rgb) == "123456"
    assert actual.shape_id == before.shape_id
    assert actual.text == before.text


@pytest.mark.parametrize("remove_endpoint", [False, True])
def test_native_connector_keeps_ids_or_rejects_removed_endpoint(tmp_path, remove_endpoint):
    deck = DeckIR(id="connections", title="合成原生连接", slides=[DeckSlide(
        id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
        objects=[
            ShapeObject(id="remove", frame=ObjectFrame(x=1, y=1, w=1, h=1)),
            TextObject(id="endpoint-a", role="body", text="旧标签", font_size=18,
                       frame=ObjectFrame(x=3, y=1, w=1, h=1)),
            ShapeObject(id="endpoint-b", frame=ObjectFrame(x=6, y=1, w=1, h=1)),
            ShapeObject(id="link", shape_type="line", frame=ObjectFrame(x=4, y=1, w=2, h=1)),
        ],
    )])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    native = Presentation(source)
    slide = native.slides[0]
    shapes = {s.name: s for s in slide.shapes}
    old_line = shapes["link"]._element
    old_line.getparent().remove(old_line)
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(4), Inches(1), Inches(6), Inches(2),
    )
    connector.name = "link"
    connector.begin_connect(shapes["endpoint-a"], 0)
    connector.end_connect(shapes["endpoint-b"], 0)
    native.save(source)
    original_bytes = source.read_bytes()
    before = {s.name: s for s in Presentation(source).slides[0].shapes}
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    actions = [RemoveObjectAction(object_id="remove"),
               AddObjectAction(object=ShapeObject(id="added", frame=ObjectFrame(x=8, y=1, w=1, h=1)))]
    if remove_endpoint:
        actions.append(RemoveObjectAction(object_id="endpoint-a"))
    else:
        actions.append(ReplaceTextAction(object_id="endpoint-a", text="新标签"))
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(slide_id="page", actions=actions))
    service = PresentationEditTransactionService(tmp_path / "edits")
    operation = PresentationEditOperation(
        deck_id=deck.id, slide_ids=["page"], object_ids=ids,
        allowed_content_types={"text", "shape"}, instruction="只替换标签与独立方块",
    )
    if remove_endpoint:
        with pytest.raises(ValueError, match="connector references a removed object"):
            service.prepare_compiled_redesign(source, imported, changed, operation, compiler)
    else:
        preview = service.prepare_compiled_redesign(source, imported, changed, operation, compiler)
        after = {s.name: s for s in Presentation(preview.candidate_path).slides[0].shapes}
        assert after["link"]._element.xml == before["link"]._element.xml
        assert after["endpoint-a"].shape_id == before["endpoint-a"].shape_id
        assert after["endpoint-a"].text == "新标签"
        assert after["added"].shape_id > max(s.shape_id for s in before.values())
    assert source.read_bytes() == original_bytes
