"""Native chart ordering and axis rules remain editable after an Office save."""

import pytest
from lxml import etree as XML
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_LABEL_POSITION
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from lantai.presentation.models import ChartObject, DeckIR, DeckSlide, ObjectFrame
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


@pytest.mark.parametrize("known", [True, False])
@pytest.mark.parametrize("automatic_max", [True, False])
def test_reversed_categories_and_explicit_value_scale_survive_native_export(tmp_path, known, automatic_max):
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    values = [4, 5, 9, 31]
    data = CategoryChartData()
    data.categories = ["合成一", "合成二", "合成三", "合成四"]
    data.add_series("数量", values)
    shape = slide.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(1), Inches(1),
                                  Inches(9), Inches(5), data)
    shape.name = "native-chart"
    chart = shape.chart
    chart.category_axis.reverse_order = True
    chart.value_axis.minimum_scale = 0
    chart.value_axis.maximum_scale = None if automatic_max else 40
    chart.value_axis.major_unit = 10
    chart.value_axis.minor_unit = 5
    chart.category_axis.tick_labels.number_format = "0"
    chart.value_axis.tick_labels.number_format = "0"
    chart.value_axis.tick_label_position = XL_TICK_LABEL_POSITION.HIGH
    chart.value_axis._element.find(qn("c:axPos")).set("val", "t")
    plot_area = chart.part._element.find(qn("c:chart")).find(qn("c:plotArea"))
    layout = plot_area.find(qn("c:layout"))
    if layout is None:
        layout = OxmlElement("c:layout")
        plot_area.insert(0, layout)
    manual = OxmlElement("c:manualLayout")
    for tag, value in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"),
                       ("x", "0.199798"), ("y", "0.0409448"), ("w", "0.786139"), ("h", "0.881395")):
        node = OxmlElement("c:" + tag)
        node.set("val", value)
        manual.append(node)
    layout.append(manual)
    expected_layout = XML.tostring(manual, method="c14n", exclusive=True)
    objects = [ChartObject(id=shape.name, chart_type="bar", categories=["合成一", "合成二", "合成三", "合成四"],
                           series={"数量": values}, unit="数量", derivation="合成测试中的原始数量，未派生",
                           source_refs=["synthetic:chart-axes"],
                           frame=ObjectFrame(x=1, y=1, w=9, h=5))] if known else []
    deck = DeckIR(id="chart-axes-deck", title="合成测试", slides=[DeckSlide(
        id="chart-axes-slide", layout_id="freeform", objects=objects,
    )])
    source = tmp_path / "source.pptx"
    presentation.save(source)
    before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    output = tmp_path / "exported.pptx"
    NativePptxCompiler().compile(imported, output)
    native = next(s.chart for s in Presentation(output).slides[0].shapes if s.has_chart)
    assert native.category_axis.reverse_order is True
    assert native.value_axis.minimum_scale == 0
    assert native.value_axis.maximum_scale == (None if automatic_max else 40)
    assert native.value_axis.major_unit == 10
    assert native.value_axis.minor_unit == 5
    assert native.value_axis._element.find(qn("c:axPos")).get("val") == "t"
    assert native.category_axis._element.tag == qn("c:catAx")
    assert native.category_axis.tick_labels.number_format == "0"
    assert native.value_axis.tick_labels.number_format == "0"
    assert native.value_axis.tick_label_position == XL_TICK_LABEL_POSITION.HIGH
    native_layout = native.part._element.find(qn("c:chart")).find(qn("c:plotArea")).find(qn("c:layout"))
    assert native_layout.find(qn("c:manualLayout")) is not None
    assert XML.tostring(native_layout.find(qn("c:manualLayout")), method="c14n", exclusive=True) == expected_layout
    assert list(native.series[0].values) == values
    assert [c.label for c in native.plots[0].categories] == ["合成一", "合成二", "合成三", "合成四"]
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "second-assets")
    assert second.slides[0].objects[0].category_axis == imported.slides[0].objects[0].category_axis
    assert second.slides[0].objects[0].value_axis == imported.slides[0].objects[0].value_axis
    assert second.slides[0].objects[0].plot_layout == imported.slides[0].objects[0].plot_layout
    assert source.read_bytes() == before
