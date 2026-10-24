# -*- coding: utf-8 -*-
from __future__ import annotations

import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from PIL import Image

from lantai.app.structured_utility import StructuredTaskValidationError
from lantai.app.utility_model import UtilityTaskTimeoutError
from lantai.presentation.assets import (
    PresentationAssetError,
    resolve_deck_assets,
)
from lantai.presentation.design_system import (
    apply_design_measurements,
    measure_text,
)
from lantai.presentation.jobs import (
    PresentationJobInterruptedError,
    PresentationJobStore,
    PresentationWriterBusyError,
)
from lantai.presentation.intelligence import UtilityPresentationPageAuthor
from lantai.presentation.models import (
    AddObjectAction,
    AssetRecord,
    ChartObject,
    CompositionAuthority,
    DeckIR,
    DeckPlan,
    DeckPlanSlide,
    DeckSlide,
    ImageObject,
    InputMaterial,
    JobStage,
    JobStatus,
    MetricObject,
    MoveResizeAction,
    ObjectFrame,
    ObjectGroup,
    PlanningMode,
    PresentationFailureCategory,
    PresentationFailureRoute,
    PresentationJobRequest,
    PresentationObjectEditPlan,
    PresentationProviderCostStatus,
    PresentationRecoveryAction,
    QualityReport,
    RemoveObjectAction,
    ReplaceImageAction,
    ReplaceTextAction,
    ShapeObject,
    SlideIntent,
    SlideIntentPatch,
    TemplateMode,
    TextObject,
    UpdateChartAction,
    VisualDirection,
)
from lantai.presentation.natural_edit import (
    apply_object_edit_plan,
    apply_slide_intent_patch,
)
from lantai.presentation.page_authoring import PresentationPageAuthorSceneError
from lantai.presentation.planner import (
    build_brief,
    build_deck_ir,
    build_deterministic_fixture_plan,
)
from lantai.presentation.pipeline_registry import (
    PipelineCapability,
    PresentationPipelineRegistry,
)
from lantai.presentation.quality import file_sha256
from lantai.presentation.runtime import PresentationJobRuntime
from lantai.presentation.service import PresentationService
from lantai.presentation.versioning import (
    PresentationVersionStore,
    StalePresentationVersionError,
)


def _request(**updates) -> PresentationJobRequest:
    request = PresentationJobRequest(
        conversation_id="conversation-v3-control",
        user_request="制作一份产品复盘演示",
        slide_count=3,
        visual_verification_required=False,
        vision_critique_required=False,
        idempotency_key="v3-control-plane-key",
    )
    return request.model_copy(update=updates)


def _deck(*objects, assets=None) -> DeckIR:
    return DeckIR(
        id="deck-v3-control",
        title="控制面验证",
        slides=[
            DeckSlide(
                id="slide-control",
                layout_id="key_judgment",
                objects=list(objects)
                or [
                    TextObject(
                        id="title-control",
                        role="title",
                        slot="title",
                        text="控制面验证",
                    ),
                ],
            ),
        ],
        assets=assets or [],
    )


class _FixturePageAuthor:
    def author(
        self,
        request,
        brief,
        plan,
        planned_slide,
        scaffold_slide,
        *,
        retry_directive=None,
    ):
        return scaffold_slide.model_copy(
            update={
                "composition_authority": CompositionAuthority.AUTHOR,
                "scene_hash": None,
                "speaker_notes": "",
                "objects": [
                    obj for obj in scaffold_slide.objects
                    if obj.slot not in {"relationship-marker", "page-marker"}
                ],
                "object_groups": [],
                "relationships": [],
            },
        )


def test_production_request_defaults_to_model_planning() -> None:
    assert _request().planning_mode == PlanningMode.MODEL


def test_runtime_uses_injected_semantic_planner_for_default_request(
    tmp_path,
    monkeypatch,
) -> None:
    calls: list[str] = []

    class Planner:
        def plan(self, request, brief):
            calls.append(request.planning_mode.value)
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path, semantic_planner=Planner(), page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            slide_count=1,
            materials=[
                InputMaterial(
                    name="模型规划输入.md",
                    content=(
                        "# 背景\n当前产品需要复盘。\n\n"
                        "# 判断\n默认生产链必须由语义规划器负责。\n\n"
                        "# 行动\n依据证据生成可编辑演示。"
                    ),
                ),
            ],
        ),
    )
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda *_args, **_kwargs: None,
    )
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        assert calls == ["model"]
        assert runtime.service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).is_file()
        deck = DeckIR.model_validate_json(
            runtime.service.store.checkpoint_path(
                record.job_id, "deck-ir.json",
            ).read_text(encoding="utf-8"),
        )
        assert runtime.service.store.get(record.job_id).cost_used_units == (
            10 + len(deck.slides)
        )
        provider_metrics = [
            metric
            for metric in runtime.service.store.get(record.job_id).stage_metrics
            if metric.cost_units > 0
        ]
        assert [metric.operation for metric in provider_metrics] == [
            "plan",
            *["author" for _slide in deck.slides],
        ]
        assert all(
            metric.provider_cost_status
            is PresentationProviderCostStatus.UNKNOWN
            for metric in provider_metrics
        )
        assert all(
            metric.provider_cost_micro_cny is None
            for metric in provider_metrics
        )
    finally:
        runtime.shutdown()


def test_reviewed_plan_retry_reinvokes_planner_instead_of_reusing_old_alias(
    tmp_path,
    monkeypatch,
) -> None:
    calls: list[int] = []

    class Planner:
        def plan(self, request, brief, *, design_context=None):
            calls.append(len(calls) + 1)
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=Planner(),
        page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            slide_count=1,
            materials=[InputMaterial(
                name="规划恢复.md",
                content="# 判断\n复核后的 PLAN 重试必须重做语义规划。",
            )],
        ),
    )
    monkeypatch.setattr(runtime.service, "compose", lambda *_args, **_kwargs: None)
    store = runtime.service.store
    try:
        with store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        old_plan = store.checkpoint_path(record.job_id, "deck-plan.json")
        old_deck = store.checkpoint_path(record.job_id, "deck-ir.json")
        assert calls == [1]
        assert old_plan.is_file() and old_deck.is_file()

        persisted_route = PresentationFailureRoute(
            category=PresentationFailureCategory.UNKNOWN,
            recovery_action=PresentationRecoveryAction.INSPECT_FAILURE,
            issue_codes=["required_fact_contract"],
            retry_stage=JobStage.PLAN,
            automatic_retry_allowed=False,
            requires_human=True,
        )
        store.append_event(
            record.job_id,
            stage=JobStage.FAILED,
            status=JobStatus.FAILED,
            message="required fact contract needs review",
            recoverable=True,
            failure_routes=[persisted_route],
        )
        reviewed_route = persisted_route.model_copy(update={
            "category": PresentationFailureCategory.CONTENT,
            "recovery_action": PresentationRecoveryAction.REVISE_CONTENT,
            "issue_codes": ["content.required_fact_contract"],
            "requires_human": False,
        })
        store.retry_failed_after_review(
            record.job_id,
            record.conversation_id,
            reviewed_routes=[reviewed_route],
            review_basis="required_fact_contract.v2",
        )
        assert not old_plan.exists()
        assert not old_deck.exists()

        with store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)

        assert calls == [1, 2]
        assert old_plan.is_file() and old_deck.is_file()
        history = (
            store.job_dir(record.job_id)
            / "checkpoints"
            / "history"
            / "recovery-001-from-plan"
        )
        assert (history / "deck-plan.json").is_file()
        assert (history / "deck-ir.json").is_file()
    finally:
        runtime.shutdown()


def test_runtime_reserves_every_staged_semantic_planner_call(
    tmp_path,
    monkeypatch,
) -> None:
    class Planner:
        def estimated_cost_units(self, request):
            assert request.slide_count == 1
            return 30

        def plan(self, request, brief):
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=Planner(),
        page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            slide_count=1,
            materials=[InputMaterial(
                name="模型规划输入.md",
                content="# 判断\n全部 staged utility 调用必须计入现有预算。",
            )],
        ),
    )
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda *_args, **_kwargs: None,
    )
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        saved = runtime.service.store.get(record.job_id)
        assert saved.cost_used_units == 31
        plan_metric = next(
            metric
            for metric in saved.stage_metrics
            if metric.operation == "plan"
        )
        assert plan_metric.cost_units == 30
    finally:
        runtime.shutdown()


def test_runtime_keeps_batched_claim_failure_without_persisting_model_text():
    assert PresentationJobRuntime._safe_semantic_contract_failures([
        {
            "stage": "claims:001", "attempt": 1,
            "violation_code": "output_truncated",
            "unsafe_model_text": "must never persist",
        },
        {
            "stage": "claims:private text", "attempt": 1,
            "violation_code": "output_truncated",
        },
    ]) == [{
        "stage": "claims:001", "attempt": 1,
        "violation_code": "output_truncated",
    }]


def test_runtime_checkpoints_safe_semantic_substage_diagnostics_on_failure(
    tmp_path,
) -> None:
    class Planner:
        last_stage_timings_ms = {"claims": 101, "outline": 120_003}
        last_stage_attempt_counts = {"claims": 1, "outline": 3}
        last_stage_output_token_budgets = {"claims": 1_024, "outline": 2_048}
        last_stage_contract_failures = [{
            "stage": "detail:001-001",
            "attempt": 3,
            "violation_code": "claim_grouping",
            "target_orders": [1],
            "expected_claim_count": 2,
            "content_block_count": 2,
            "claim_group_count": 1,
            "missing_claim_ids": ["claim-002"],
            "duplicate_claim_ids": [],
            "unknown_claim_ids": ["claim-999"],
            "unsafe_model_text": "must never persist",
        }]
        last_utility_call_count = 4
        last_output_token_budget = 7_168

        def estimated_cost_units(self, request):
            assert request.slide_count == 1
            return 30

        def plan(self, request, brief, *, design_context=None):
            raise UtilityTaskTimeoutError(
                phase="total",
                timeout_seconds=120,
            )

    runtime = PresentationJobRuntime(tmp_path, semantic_planner=Planner())
    record = runtime.service.create_job(
        _request(
            slide_count=1,
            materials=[InputMaterial(
                name="模型规划输入.md",
                content="# 判断\n记录安全的分阶段耗时，不保存模型正文。",
            )],
        ),
    )
    try:
        runtime._run(record.job_id)

        diagnostics_path = runtime.service.store.checkpoint_path(
            record.job_id,
            "semantic-plan-diagnostics-recovery-000.json",
        )
        diagnostics = json.loads(diagnostics_path.read_text(encoding="utf-8"))
        assert diagnostics == {
            "external_recovery_attempt_count": 0,
            "implementation_recovery_attempt_count": 0,
            "recovery_attempt_count": 0,
            "schema_version": "1.0",
            "stage_contract_failures": [{
                "attempt": 3,
                "claim_group_count": 1,
                "content_block_count": 2,
                "duplicate_claim_ids": [],
                "expected_claim_count": 2,
                "missing_claim_ids": ["claim-002"],
                "stage": "detail:001-001",
                "target_orders": [1],
                "unknown_claim_ids": ["claim-999"],
                "violation_code": "claim_grouping",
            }],
            "stage_attempt_counts": {"claims": 1, "outline": 3},
            "stage_output_token_budgets": {
                "claims": 1_024,
                "outline": 2_048,
            },
            "stage_timings_ms": {"claims": 101, "outline": 120_003},
            "total_output_token_budget": 7_168,
            "utility_call_count": 4,
        }
        assert "模型正文" not in diagnostics_path.read_text(encoding="utf-8")
        assert "must never persist" not in diagnostics_path.read_text(
            encoding="utf-8",
        )
        saved = runtime.service.store.get(record.job_id)
        assert saved.status == JobStatus.FAILED
        assert saved.events[-2].checkpoint == str(diagnostics_path)
        assert saved.events[-1].failure_routes[0].issue_codes == [
            "utilitytasktimeouterror",
        ]
    finally:
        runtime.shutdown()


def test_model_job_authors_and_checkpoints_each_native_page_before_compose(
    tmp_path,
    monkeypatch,
) -> None:
    calls: list[str] = []

    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    class PageAuthor:
        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            calls.append(scaffold_slide.id)
            return scaffold_slide.model_copy(
                update={
                    "composition_authority": CompositionAuthority.AUTHOR,
                    "scene_hash": None,
                    "speaker_notes": "",
                    "objects": [
                        obj for obj in scaffold_slide.objects
                        if obj.slot != "page-marker"
                    ],
                    "object_groups": [],
                    "relationships": [],
                },
            )

    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=Planner(),
        page_author=PageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            materials=[InputMaterial(name="判断.md", content="# 判断\n核心事实必须保留。")],
        ),
    )
    captured: list[DeckIR] = []
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda _job_id, deck, **_kwargs: captured.append(deck),
    )
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        assert len(calls) == len(captured[0].slides)
        assert all(
            slide.composition_authority == CompositionAuthority.AUTHOR
            for slide in captured[0].slides
        )
        assert all(
            runtime.service.store.checkpoint_path(
                record.job_id,
                f"page-scene-{index:03d}.json",
            ).is_file()
            for index in range(1, len(calls) + 1)
        )
        assert runtime.service.store.get(record.job_id).cost_used_units == (
            10 + len(calls)
        )
    finally:
        runtime.shutdown()


def test_model_job_fails_closed_without_page_author(tmp_path, monkeypatch) -> None:
    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(tmp_path, semantic_planner=Planner())
    record = runtime.service.create_job(
        _request(
            materials=[InputMaterial(name="判断.md", content="# 判断\n事实必须保留。")],
        ),
    )
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda *_args, **_kwargs: pytest.fail("compose must not run"),
    )
    try:
        runtime._run(record.job_id)
        failed = runtime.service.store.get(record.job_id)
        assert failed.status == JobStatus.FAILED
        assert "缺少页级原生作者" in failed.events[-1].message
        assert not runtime.service.store.checkpoint_path(
            record.job_id, "deck-ir.json",
        ).exists()
    finally:
        runtime.shutdown()


def test_model_job_resumes_from_validated_page_checkpoint(tmp_path, monkeypatch) -> None:
    calls: list[str] = []

    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    class IntermittentAuthor(_FixturePageAuthor):
        failed_once = False

        def author(self, request, brief, plan, planned_slide, scaffold_slide):
            calls.append(scaffold_slide.id)
            if planned_slide.order == 2 and not self.failed_once:
                self.failed_once = True
                raise RuntimeError("model service interrupted")
            return super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    runtime = PresentationJobRuntime(
        tmp_path, semantic_planner=Planner(), page_author=IntermittentAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            materials=[InputMaterial(
                name="判断.md",
                content=(
                    "# 材料\n## 背景\n事实必须保留。\n"
                    "## 判断\n第二页判断。\n## 行动\n第三页行动。"
                ),
            )],
        ),
    )
    composed: list[DeckIR] = []
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda _job_id, deck, **_kwargs: composed.append(deck),
    )
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            with pytest.raises(RuntimeError, match="model service interrupted"):
                runtime._run_with_lease(record.job_id, lease)
        first = runtime.service.store.checkpoint_path(
            record.job_id, "page-scene-001.json",
        )
        assert first.is_file()
        assert not runtime.service.store.checkpoint_path(
            record.job_id, "deck-ir.json",
        ).exists()
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        assert calls.count("slide-001") == 1
        assert calls.count("slide-002") == 2
        assert composed and all(
            slide.composition_authority == CompositionAuthority.AUTHOR
            for slide in composed[0].slides
        )
    finally:
        runtime.shutdown()


@pytest.mark.parametrize(
    "transient_error",
    [
        StructuredTaskValidationError(
            "structured utility task output failed schema validation",
            validation_issues=(
                "missing@$.objects[0].frame",
                "extra_forbidden@$.unexpected",
            ),
        ),
        UtilityTaskTimeoutError(phase="total", timeout_seconds=120),
    ],
)
def test_page_author_retries_transient_utility_failure_and_charges_every_call(
    tmp_path,
    transient_error: Exception,
) -> None:
    class TransientRetryAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.calls = 0
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.calls += 1
            self.retry_directives.append(retry_directive)
            if self.calls == 1:
                raise transient_error
            return super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    author = TransientRetryAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n结构化模型响应失败时只重试当前页。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        saved = runtime.service.store.get(record.job_id)
        author_metrics = [
            metric for metric in saved.stage_metrics
            if metric.operation == "author"
        ]
        assert author.calls == 2
        assert saved.cost_used_units == 2
        assert [metric.outcome for metric in author_metrics] == [
            "failed",
            "passed",
        ]
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
        if isinstance(transient_error, StructuredTaskValidationError):
            retry_message = next(
                event.message for event in saved.events
                if "schema" in event.message
            )
            assert "missing@$.objects[0].frame" in retry_message
            assert "extra_forbidden@$" in retry_message
            assert "unexpected" not in retry_message
            assert author.retry_directives[0] is None
            assert "native_scene_contract" in author.retry_directives[1]
            assert "TextObject.font_size 必须在 17 到 72 之间" in (
                author.retry_directives[1]
            )
            assert "missing@$.objects[0].frame" in author.retry_directives[1]
            assert "extra_forbidden@$" in author.retry_directives[1]
            assert "unexpected" not in author.retry_directives[1]
        assert runtime.service.store.checkpoint_path(
            record.job_id, "page-scene-001.json",
        ).is_file()
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("correct_on_retry", [True, False])
def test_public_page_author_anchor_retry_is_safe_and_bounded(
    tmp_path, monkeypatch, correct_on_retry: bool,
) -> None:
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="controlled.md", content="# Controlled evidence\nA verified fact.",
        )],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    planned = plan.slides[0]
    rejected_id = "private-rejected-anchor-sentinel"
    rejected_text = "private-rejected-content-sentinel"
    valid_scene = {
        "objects": [
            {
                "kind": "text", "id": "controlled-title", "role": "title",
                "text": planned.title, "font_size": 36,
                "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1},
            },
            {
                "kind": "text", "id": "controlled-body", "role": "body",
                "text": "\n".join(
                    planned.intent.required_facts + planned.intent.immutable_facts,
                ) or "A verified fact.",
                "font_size": 24,
                "frame": {"x": 1, "y": 2, "w": 10, "h": 4},
            },
        ],
    }
    invalid_scene = {
        "objects": [{
            **valid_scene["objects"][0],
            "id": rejected_id, "text": rejected_text,
        }],
        "template_anchors": {
            "source_slide_number": 1,
            "uses": [{
                "source_shape_id": 2, "purpose": "static_brand",
                "target_object_id": rejected_id,
            }],
        },
    }
    payloads: list[dict] = []

    async def transport(**kwargs):
        payload = json.loads(kwargs["user_prompt"])
        payloads.append(payload)
        directive = payload.get("retry_directive", "")
        corrected = correct_on_retry and "static_brand 和 page_number" in directive
        return SimpleNamespace(
            text=json.dumps(valid_scene if corrected else invalid_scene),
            model_label="approved-provider/controlled-model",
        )

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    monkeypatch.setattr(
        "lantai.presentation.intelligence._utility_route_cache_identity",
        lambda _agent_id: "approved-provider/controlled-model",
    )
    runtime = PresentationJobRuntime(
        tmp_path, page_author=UtilityPresentationPageAuthor("controlled-agent"),
    )
    record = runtime.service.create_job(request)
    try:
        if correct_on_retry:
            authored = runtime._author_pages(
                record.job_id, request, brief, plan, scaffold,
            )
            assert authored.slides[0].composition_authority == CompositionAuthority.AUTHOR
            assert runtime.service.store.checkpoint_path(
                record.job_id, "page-scene-001.json",
            ).is_file()
        else:
            with pytest.raises(StructuredTaskValidationError):
                runtime._author_pages(record.job_id, request, brief, plan, scaffold)
            assert not runtime.service.store.checkpoint_path(
                record.job_id, "page-scene-001.json",
            ).exists()
        saved = runtime.service.store.get(record.job_id)
        assert len(payloads) == saved.cost_used_units == (2 if correct_on_retry else 3)
        assert "retry_directive" not in payloads[0]
        for payload in payloads[1:]:
            directive = payload["retry_directive"]
            assert "value_error@$" in directive
            assert "objects 中的 id 必须唯一" in directive
            assert "title_frame 必须引用本次返回的唯一 title" in directive
            assert "static_brand 和 page_number" in directive
            assert "不在 objects 中的新 ID" in directive
            assert "不得修改事实或图表数据" in directive
            assert rejected_id not in directive
            assert rejected_text not in directive
        recorded = saved.model_dump_json()
        assert rejected_id not in recorded
        assert rejected_text not in recorded
        assert runtime._PAGE_AUTHOR_TRANSIENT_ATTEMPT_LIMIT == 5
        assert runtime._PAGE_AUTHOR_RETRY_LIMIT_PER_CAUSE == 2
    finally:
        runtime.shutdown()


def test_page_author_retries_grounding_failure_with_safe_directive(
    tmp_path,
) -> None:
    class ContentRetryAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.retry_directives.append(retry_directive)
            authored = super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )
            if len(self.retry_directives) == 1:
                return authored.model_copy(update={
                    "speaker_notes": "不要展示来源材料没有提供的99万元。",
                })
            return authored

    author = ContentRetryAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n事实必须保留，不能新增数字。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        saved = runtime.service.store.get(record.job_id)
        assert len(author.retry_directives) == 2
        assert author.retry_directives[0] is None
        assert "speaker_notes 必须返回空字符串" in (
            author.retry_directives[1]
        )
        assert "不要保留、改写或概括上一版备注" in (
            author.retry_directives[1]
        )
        assert saved.cost_used_units == 2
        assert any(
            "事实门禁" in event.message for event in saved.events
        )
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
    finally:
        runtime.shutdown()


def test_page_author_numeric_notes_retry_requires_empty_string(
    tmp_path,
) -> None:
    class NumericNotesAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.retry_directives.append(retry_directive)
            authored = super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )
            if "speaker_notes 必须返回空字符串" not in str(retry_directive):
                return authored.model_copy(update={
                    "speaker_notes": "第1步讲清本页判断。",
                })
            return authored.model_copy(update={"speaker_notes": ""})

    author = NumericNotesAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n事实必须保留，且备注不得添加数字。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        assert len(author.retry_directives) == 2
        assert "speaker_notes 必须返回空字符串" in str(
            author.retry_directives[1],
        )
        assert authored.slides[0].speaker_notes == ""
        assert runtime.service.store.checkpoint_path(
            record.job_id, "page-scene-001.json",
        ).is_file()
    finally:
        runtime.shutdown()


def test_page_author_retries_exact_missing_fact_path_without_fact_text(
    tmp_path,
) -> None:
    class MissingFactRetryAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.retry_directives.append(retry_directive)
            authored = super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )
            if len(self.retry_directives) == 1:
                missing_fact = planned_slide.intent.required_facts[0]
                authored = authored.model_copy(update={
                    "objects": [
                        obj.model_copy(update={"text": "只保留概括判断"})
                        if isinstance(obj, TextObject)
                        and obj.role != "title"
                        and missing_fact in obj.text
                        else obj
                        for obj in authored.objects
                    ],
                })
            return authored

    author = MissingFactRetryAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n必须逐字保留这条事实。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    missing_fact = plan.slides[0].intent.required_facts[0]
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        assert len(author.retry_directives) == 2
        assert "required_facts[0]" in str(author.retry_directives[1])
        assert missing_fact not in str(author.retry_directives[1])
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
        assert runtime.service.store.checkpoint_path(
            record.job_id, "page-scene-001.json",
        ).is_file()
    finally:
        runtime.shutdown()


def test_page_author_retains_missing_fact_constraint_after_numeric_retry(
    tmp_path,
) -> None:
    class OscillatingContentAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.retry_directives.append(retry_directive)
            authored = super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )
            missing_fact = planned_slide.intent.required_facts[0]
            if len(self.retry_directives) == 1:
                return authored.model_copy(update={
                    "objects": [
                        obj.model_copy(update={"text": "只保留概括判断"})
                        if isinstance(obj, TextObject)
                        and obj.role != "title"
                        and missing_fact in obj.text
                        else obj
                        for obj in authored.objects
                    ],
                })
            if len(self.retry_directives) == 2:
                assert "required_facts[0]" in str(retry_directive)
                return authored.model_copy(update={
                    "speaker_notes": "不要添加来源没有提供的99万元。",
                })
            assert "required_facts[0]" in str(retry_directive)
            assert "speaker_notes 必须返回空字符串" in str(retry_directive)
            assert missing_fact not in str(retry_directive)
            return authored

    author = OscillatingContentAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n必须逐字保留这条事实，且不能新增数字。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        assert len(author.retry_directives) == 3
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
        assert runtime.service.store.checkpoint_path(
            record.job_id, "page-scene-001.json",
        ).is_file()
    finally:
        runtime.shutdown()


def test_page_author_retries_dangling_object_reference_with_safe_directive(
    tmp_path,
) -> None:
    class SceneRetryAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.retry_directives.append(retry_directive)
            authored = super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )
            if len(self.retry_directives) == 1:
                return authored.model_copy(update={
                    "object_groups": [ObjectGroup(
                        group_id="evidence-group",
                        role="evidence",
                        object_ids=[authored.objects[0].id, "missing-object"],
                    )],
                })
            return authored

    author = SceneRetryAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n事实必须保留。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        saved = runtime.service.store.get(record.job_id)
        assert len(author.retry_directives) == 2
        assert author.retry_directives[0] is None
        assert "object_groups" in author.retry_directives[1]
        assert "relationships" in author.retry_directives[1]
        assert "objects" in author.retry_directives[1]
        assert "空列表" in author.retry_directives[1]
        assert "missing-object" not in author.retry_directives[1]
        assert saved.cost_used_units == 2
        assert any(
            "场景关系" in event.message for event in saved.events
        )
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
    finally:
        runtime.shutdown()


def test_page_author_gives_each_failure_cause_its_own_bounded_retry(tmp_path) -> None:
    class MixedFailureAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.calls = 0
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.calls += 1
            self.retry_directives.append(retry_directive)
            if self.calls <= 2:
                raise PresentationPageAuthorSceneError(
                    "authored scene references missing objects",
                    violation_code="dangling_object_reference",
                )
            if self.calls == 3:
                raise StructuredTaskValidationError(
                    "structured utility task output failed schema validation",
                    validation_issues=("greater_than@$.objects[0].shape.frame.w",),
                )
            return super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    author = MixedFailureAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n异类错误不能互相耗尽页级修复机会。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        saved = runtime.service.store.get(record.job_id)
        assert author.calls == 4
        assert saved.cost_used_units == 4
        assert len(author.retry_directives) == 4
        assert "空列表" in str(author.retry_directives[1])
        assert "空列表" in str(author.retry_directives[2])
        assert "native_scene_contract" in str(author.retry_directives[3])
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
    finally:
        runtime.shutdown()


def test_page_author_distinguishes_schema_fields_but_bounds_repeated_field(
    tmp_path,
) -> None:
    class DistinctSchemaFailureAuthor(_FixturePageAuthor):
        def __init__(self) -> None:
            self.calls = 0
            self.retry_directives: list[str | None] = []

        def author(
            self,
            request,
            brief,
            plan,
            planned_slide,
            scaffold_slide,
            *,
            retry_directive=None,
        ):
            self.calls += 1
            self.retry_directives.append(retry_directive)
            if self.calls == 1:
                raise StructuredTaskValidationError(
                    "structured utility task output failed schema validation",
                    validation_issues=(
                        "greater_than@$.objects[5].shape.frame.w",
                    ),
                )
            if self.calls == 2:
                raise StructuredTaskValidationError(
                    "structured utility task output failed schema validation",
                    validation_issues=(
                        "greater_than@$.objects[6].shape.frame.w",
                    ),
                )
            if self.calls == 3:
                raise StructuredTaskValidationError(
                    "structured utility task output failed schema validation",
                    validation_issues=(
                        "list_type@$.objects[2].text.runs",
                    ),
                )
            return super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    author = DistinctSchemaFailureAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(
        slide_count=1,
        materials=[InputMaterial(
            name="事实.md",
            content="# 判断\n不同 schema 字段应各自获得有界修复机会。",
        )],
    )
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        authored = runtime._author_pages(
            record.job_id, request, brief, plan, scaffold,
        )

        saved = runtime.service.store.get(record.job_id)
        assert author.calls == 4
        assert saved.cost_used_units == 4
        assert len(author.retry_directives) == 4
        assert "objects[5].shape.frame.w" in str(author.retry_directives[1])
        assert "objects[6].shape.frame.w" in str(author.retry_directives[2])
        assert "objects[2].text.runs" in str(author.retry_directives[3])
        assert authored.slides[0].composition_authority == (
            CompositionAuthority.AUTHOR
        )
    finally:
        runtime.shutdown()


def test_page_scene_checkpoint_invalidates_when_author_route_changes(tmp_path) -> None:
    class RoutedAuthor(_FixturePageAuthor):
        route = "provider/model-a"

        def __init__(self):
            self.calls: list[str] = []

        def cache_identity(self) -> str:
            return self.route

        def author(self, request, brief, plan, planned_slide, scaffold_slide):
            self.calls.append(scaffold_slide.id)
            return super().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    author = RoutedAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(materials=[InputMaterial(
        name="事实.md",
        content="# 背景\n事实一。\n# 判断\n事实二。\n# 行动\n事实三。",
    )])
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        runtime._author_pages(record.job_id, request, brief, plan, scaffold)
        assert len(author.calls) == len(plan.slides)
        runtime._author_pages(record.job_id, request, brief, plan, scaffold)
        assert len(author.calls) == len(plan.slides)
        author.route = "provider/model-b"
        runtime._author_pages(record.job_id, request, brief, plan, scaffold)
        assert len(author.calls) == 2 * len(plan.slides)
    finally:
        runtime.shutdown()


def test_utility_page_scene_checkpoint_only_invalidates_affected_page(tmp_path) -> None:
    class RecordingUtilityAuthor(UtilityPresentationPageAuthor):
        def __init__(self):
            super().__init__("agent-cache-test")
            self.calls: list[str] = []

        def cache_identity(self) -> str:
            return "provider/model-a"

        def author(self, request, brief, plan, planned_slide, scaffold_slide):
            self.calls.append(scaffold_slide.id)
            return _FixturePageAuthor().author(
                request, brief, plan, planned_slide, scaffold_slide,
            )

    author = RecordingUtilityAuthor()
    runtime = PresentationJobRuntime(tmp_path, page_author=author)
    request = _request(materials=[InputMaterial(
        name="事实.md",
        content="# 材料\n## 背景\n事实一。\n## 判断\n事实二。\n## 行动\n事实三。",
    )])
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan)
    try:
        runtime._author_pages(record.job_id, request, brief, plan, scaffold)
        assert author.calls == [slide.id for slide in scaffold.slides]
        runtime._author_pages(record.job_id, request, brief, plan, scaffold)
        assert len(author.calls) == len(scaffold.slides)

        changed_slide = plan.slides[1].model_copy(update={
            "content_blocks": [*plan.slides[1].content_blocks, "事实二。"],
        })
        changed_plan = plan.model_copy(update={
            "slides": [plan.slides[0], changed_slide, *plan.slides[2:]],
        })
        changed_scaffold = build_deck_ir(request, changed_plan)
        runtime._author_pages(
            record.job_id, request, brief, changed_plan, changed_scaffold,
        )
        assert author.calls == [
            *[slide.id for slide in scaffold.slides],
            changed_scaffold.slides[1].id,
        ]

        changed_direction = changed_plan.visual_direction.model_copy(update={
            "visual_concept": "共享视觉方向已更新",
        })
        shared_plan = changed_plan.model_copy(update={
            "visual_direction": changed_direction,
        })
        runtime._author_pages(
            record.job_id, request, brief, shared_plan,
            build_deck_ir(request, shared_plan),
        )
        assert len(author.calls) == 2 * len(scaffold.slides) + 1

        author.SYSTEM_PROMPT += "\n新增全套作者规则"
        runtime._author_pages(
            record.job_id, request, brief, shared_plan,
            build_deck_ir(request, shared_plan),
        )
        assert len(author.calls) == 3 * len(scaffold.slides) + 1
    finally:
        runtime.shutdown()


def test_old_model_deck_checkpoint_is_preserved_and_reauthored(
    tmp_path, monkeypatch,
) -> None:
    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path, semantic_planner=Planner(), page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(materials=[InputMaterial(
            name="原材料.md", content="# 判断\n原始判断必须保留。",
        )]),
    )
    brief = build_brief(record.request)
    plan = build_deterministic_fixture_plan(record.request, brief)
    old_deck = build_deck_ir(record.request, plan)
    store = runtime.service.store
    store.write_checkpoint(record.job_id, "brief.json", brief.model_dump_json())
    store.write_checkpoint(record.job_id, "deck-plan.json", plan.model_dump_json())
    original = old_deck.model_dump_json(indent=2)
    store.write_checkpoint(record.job_id, "deck-ir.json", original)
    composed: list[DeckIR] = []
    monkeypatch.setattr(
        runtime.service, "compose",
        lambda _job_id, deck, **_kwargs: composed.append(deck),
    )
    try:
        with store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        assert store.checkpoint_path(
            record.job_id, "deck-ir-legacy-before-page-author.json",
        ).read_text(encoding="utf-8") == original
        assert composed and all(
            slide.composition_authority == CompositionAuthority.AUTHOR
            for slide in composed[0].slides
        )
        assert all(slide.scene_hash for slide in composed[0].slides)
    finally:
        runtime.shutdown()


def test_tampered_authored_deck_checkpoint_fails_closed(
    tmp_path, monkeypatch,
) -> None:
    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path, semantic_planner=Planner(), page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(materials=[InputMaterial(
            name="原材料.md", content="# 判断\n原始判断必须保留。",
        )]),
    )
    calls: list[DeckIR] = []
    monkeypatch.setattr(
        runtime.service, "compose",
        lambda _job_id, deck, **_kwargs: calls.append(deck),
    )
    store = runtime.service.store
    try:
        with store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)
        assert len(calls) == 1
        path = store.checkpoint_path(record.job_id, "deck-ir.json")
        tampered = DeckIR.model_validate_json(path.read_text(encoding="utf-8"))
        tampered.slides[0].speaker_notes = "篡改后的备注"
        store.write_checkpoint(record.job_id, "deck-ir.json", tampered.model_dump_json())
        with store.writer_lease(record.job_id) as lease:
            with pytest.raises(ValueError, match="checkpoint hash"):
                runtime._run_with_lease(record.job_id, lease)
        assert len(calls) == 1
        assert not store.checkpoint_path(
            record.job_id, "deck-ir-legacy-before-page-author.json",
        ).exists()
    finally:
        runtime.shutdown()


def test_model_authored_page_reaches_real_native_pptx_writer(tmp_path) -> None:
    class Planner:
        def plan(self, request, brief, *, design_context=None):
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path, semantic_planner=Planner(), page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            slide_count=1,
            materials=[InputMaterial(
                name="原材料.md", content="# 判断\n原始判断必须保留。",
            )],
        ),
    )
    try:
        runtime._run(record.job_id)
        output = (
            runtime.service.store.job_dir(record.job_id)
            / "output" / "presentation.pptx"
        )
        assert output.is_file()
        assert output.read_bytes()[:2] == b"PK"
        assert runtime.service.store.get(record.job_id).status == JobStatus.COMPLETED
        deck = DeckIR.model_validate_json(
            runtime.service.store.checkpoint_path(
                record.job_id, "deck-ir.json",
            ).read_text(encoding="utf-8"),
        )
        assert deck.slides[0].composition_authority == CompositionAuthority.AUTHOR
    finally:
        runtime.shutdown()


def test_runtime_preserves_legacy_doc_coverage_through_model_plan_checkpoint(
    tmp_path,
    monkeypatch,
) -> None:
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"legacy binary input")
    converted = tmp_path / "converted.docx"
    document = Document()
    document.add_heading("旧稿事实", level=1)
    document.add_paragraph("基准值为 28。")
    document.save(converted)

    def fake_convert(input_path, extension, output_dir):
        assert input_path != source
        assert input_path.name == source.name
        assert input_path.read_bytes() == source.read_bytes()
        assert extension == "docx"
        assert output_dir
        return SimpleNamespace(output_path=converted)

    def fake_validate(input_path):
        assert input_path != source
        assert input_path.name == source.name
        assert input_path.read_bytes() == source.read_bytes()

    monkeypatch.setattr(
        "lantai.app.office_conversion.convert_office_document", fake_convert
    )
    monkeypatch.setattr(
        "lantai.presentation.material_extraction._validate_legacy_doc",
        fake_validate,
    )

    class Planner:
        def plan(self, request, brief):
            # Simulate a model that omits deterministic source metadata.
            return build_deterministic_fixture_plan(request, brief).model_copy(
                update={"source_coverage": []}
            )

    runtime = PresentationJobRuntime(
        tmp_path / "jobs", semantic_planner=Planner(),
        page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(materials=[InputMaterial(name=source.name, path=str(source))]),
    )
    monkeypatch.setattr(runtime.service, "compose", lambda *_args, **_kwargs: None)
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)

        for checkpoint in ("brief.json", "deck-plan.json", "deck-ir.json"):
            saved = json.loads(
                runtime.service.store.checkpoint_path(
                    record.job_id, checkpoint
                ).read_text(encoding="utf-8")
            )
            assert saved["source_coverage"][0]["status"] == (
                "converted_doc_unverified"
            )
    finally:
        runtime.shutdown()


def test_runtime_reference_learning_is_optional_and_shared_with_planner(
    tmp_path,
    monkeypatch,
) -> None:
    reference = tmp_path / "reference.png"
    Image.new("RGB", (1600, 900), "#f5f3ee").save(reference)
    received_contexts = []

    class Planner:
        def plan(self, request, brief, *, design_context=None):
            received_contexts.append(design_context)
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path / "jobs", semantic_planner=Planner(),
        page_author=_FixturePageAuthor(),
    )
    record = runtime.service.create_job(
        _request(
            user_request="只参考这份演示的感觉，用新内容重新设计",
            template_path=str(reference),
            materials=[
                InputMaterial(
                    name="判断.md",
                    content="# 判断\n引用输入只增强设计，不替代内容事实。",
                ),
            ],
        ),
    )
    monkeypatch.setattr(runtime.service, "compose", lambda *_args, **_kwargs: None)
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            runtime._run_with_lease(record.job_id, lease)

        assert len(received_contexts) == 1
        context = received_contexts[0]
        assert context is not None
        assert context.reference_input_present is True
        assert context.profile.template_mode == TemplateMode.STYLE_REFERENCE
        assert runtime.service.store.checkpoint_path(
            record.job_id,
            "reference-learning.json",
        ).is_file()
        plan = runtime.service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).read_text(encoding="utf-8")
        assert "custom-reference-" in plan
    finally:
        runtime.shutdown()


def test_runtime_reference_checkpoint_rejects_changed_template_bytes(
    tmp_path,
) -> None:
    reference = tmp_path / "reference.png"
    Image.new("RGB", (1600, 900), "#f5f3ee").save(reference)
    runtime = PresentationJobRuntime(
        tmp_path / "jobs", page_author=_FixturePageAuthor(),
    )
    request = _request(template_path=str(reference))
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    try:
        first = runtime._resolve_job_design_context(record.job_id, request, brief)
        context_path = runtime.service.store.checkpoint_path(
            record.job_id, "presentation-design-context.json",
        )
        original_context_bytes = context_path.read_bytes()
        learning = json.loads(
            runtime.service.store.checkpoint_path(
                record.job_id, "reference-learning.json",
            ).read_text(encoding="utf-8")
        )
        assert first.reference_artifact_id == (
            f"sha256:{learning['measurement']['source_sha256']}"
        )

        Image.new("RGB", (1600, 900), "#e1e5ef").save(reference)
        with pytest.raises(ValueError, match="reference.*changed"):
            runtime._resolve_job_design_context(record.job_id, request, brief)
        assert context_path.read_bytes() == original_context_bytes

        # Older jobs may have learning evidence but no final context checkpoint.
        context_path.unlink()
        with pytest.raises(ValueError, match="reference.*changed"):
            runtime._resolve_job_design_context(record.job_id, request, brief)
        assert not context_path.exists()

        reference.unlink()
        with pytest.raises(FileNotFoundError, match="reference does not exist"):
            runtime._resolve_job_design_context(record.job_id, request, brief)
        assert not context_path.exists()
    finally:
        runtime.shutdown()


def test_runtime_legacy_reference_context_reads_learning_identity(tmp_path) -> None:
    reference = tmp_path / "legacy-reference.png"
    Image.new("RGB", (1600, 900), "#f5f3ee").save(reference)
    runtime = PresentationJobRuntime(
        tmp_path / "jobs", page_author=_FixturePageAuthor(),
    )
    request = _request(template_path=str(reference))
    record = runtime.service.create_job(request)
    brief = build_brief(request)
    try:
        context = runtime._resolve_job_design_context(record.job_id, request, brief)
        context_path = runtime.service.store.checkpoint_path(
            record.job_id, "presentation-design-context.json",
        )
        legacy_context = context.model_copy(update={"reference_artifact_id": None})
        context_path.write_text(
            legacy_context.model_dump_json(indent=2), encoding="utf-8",
        )

        resumed = runtime._resolve_job_design_context(record.job_id, request, brief)
        assert resumed.reference_artifact_id == context.reference_artifact_id
        assert context_path.read_text(encoding="utf-8") == (
            legacy_context.model_dump_json(indent=2)
        )

        Image.new("RGB", (1600, 900), "#e1e5ef").save(reference)
        with pytest.raises(ValueError, match="reference.*changed"):
            runtime._resolve_job_design_context(record.job_id, request, brief)
    finally:
        runtime.shutdown()


def test_runtime_fails_closed_when_default_model_planner_is_unavailable(
    tmp_path,
) -> None:
    runtime = PresentationJobRuntime(tmp_path)
    record = runtime.service.create_job(_request())
    try:
        runtime._run(record.job_id)
        failed = runtime.service.store.get(record.job_id)
        assert failed.status == JobStatus.FAILED
        assert "禁止静默降级" in failed.events[-1].message
        assert not runtime.service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).exists()
    finally:
        runtime.shutdown()


def test_cancel_during_semantic_planning_does_not_publish_plan_checkpoint(
    tmp_path,
    monkeypatch,
) -> None:
    runtime = PresentationJobRuntime.__new__(PresentationJobRuntime)

    class CancellingPlanner:
        def plan(self, request, brief, *, design_context=None):
            runtime.service.store.cancel(
                record.job_id,
                record.conversation_id,
            )
            return build_deterministic_fixture_plan(request, brief)

    PresentationJobRuntime.__init__(
        runtime,
        tmp_path,
        semantic_planner=CancellingPlanner(),
    )
    record = runtime.service.create_job(
        _request(
            materials=[
                InputMaterial(
                    name="取消竞态输入.md",
                    content=(
                        "# 判断\n规划返回后仍必须尊重用户 Stop。\n\n"
                        "# 行动\n取消后不写入大纲检查点。"
                    ),
                ),
            ],
        ),
    )
    monkeypatch.setattr(runtime.service, "compose", lambda *_args, **_kwargs: None)
    try:
        with runtime.service.store.writer_lease(record.job_id) as lease:
            with pytest.raises(
                PresentationJobInterruptedError,
                match="cancelled",
            ):
                runtime._run_with_lease(record.job_id, lease)

        cancelled = runtime.service.store.get(record.job_id)
        assert cancelled.status == JobStatus.CANCELLED
        assert not runtime.service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).exists()
        assert any(
            metric.operation == "plan" and metric.outcome == "cancelled"
            for metric in cancelled.stage_metrics
        )
    finally:
        runtime.shutdown()


def test_deck_ir_v1_checkpoint_migrates_to_v2() -> None:
    deck = DeckIR.model_validate(
        {
            "schema_version": "1.0",
            "id": "legacy-deck",
            "title": "旧检查点",
            "slides": [
                {
                    "id": "legacy-slide",
                    "layout_id": "cover",
                    "objects": [
                        {
                            "kind": "text",
                            "id": "legacy-title",
                            "role": "title",
                            "slot": "title",
                            "text": "旧版仍可读取",
                        },
                    ],
                },
            ],
        },
    )
    assert deck.schema_version == "2.0"
    assert deck.slides[0].revision == 1
    assert deck.slides[0].object_groups == []
    assert deck.facts == []
    assert deck.assets == []


def test_deck_ir_v2_rejects_unknown_fact_asset_and_object_references() -> None:
    with pytest.raises(ValueError, match="unknown facts"):
        DeckIR(
            title="错误引用",
            slides=[
                DeckSlide(
                    layout_id="cover",
                    source_fact_ids=["missing-fact"],
                ),
            ],
        )
    with pytest.raises(ValueError, match="unknown asset"):
        DeckIR(
            title="错误引用",
            slides=[
                DeckSlide(
                    layout_id="cover",
                    objects=[
                        ImageObject(
                            id="image-missing",
                            asset_id="missing-asset",
                            alt="missing",
                        ),
                    ],
                ),
            ],
        )


def _asset(path: Path, asset_id: str, **updates) -> AssetRecord:
    values = {
        "asset_id": asset_id,
        "local_path": str(path),
        "source_ref": f"fixture:{asset_id}",
        "rights_status": "generated",
        "sha256": file_sha256(path),
        "focal_point": (0.5, 0.4),
        "upload_policy": "never",
    }
    values.update(updates)
    return AssetRecord(**values)


def test_asset_resolution_is_local_rights_checked_and_hash_bound(
    tmp_path,
) -> None:
    image = tmp_path / "asset.png"
    image.write_bytes(b"audited-image")
    asset = _asset(image, "asset-a")
    deck = _deck(
        ImageObject(id="image-a", asset_id=asset.asset_id, alt="audited"),
        assets=[asset],
    )
    resolved = resolve_deck_assets(deck)
    resolved_image = resolved.slides[0].objects[0]
    assert isinstance(resolved_image, ImageObject)
    assert resolved_image.path == str(image.resolve())
    assert resolved_image.focal_point == (0.5, 0.4)
    assert resolved.slides[0].asset_ids == ["asset-a"]
    assert resolved.assets[0].dedupe_key == resolved.assets[0].sha256
    assert resolved.assets[0].upload_policy == "never"

    blocked = deck.model_copy(
        update={
            "assets": [
                asset.model_copy(update={"rights_status": "review"}),
            ],
        },
    )
    with pytest.raises(PresentationAssetError, match="rights require review"):
        resolve_deck_assets(blocked)

    image.write_bytes(b"changed-after-ingestion")
    with pytest.raises(PresentationAssetError, match="hash mismatch"):
        resolve_deck_assets(deck)


def test_asset_contract_requires_source_rights_and_recovery_details(
    tmp_path,
) -> None:
    source = tmp_path / "asset.svg"
    source.write_text("<svg/>", encoding="utf-8")
    digest = file_sha256(source)
    with pytest.raises(ValueError, match="source_uri or source_ref"):
        AssetRecord(
            local_path=str(source),
            rights_status="generated",
            sha256=digest,
        )
    with pytest.raises(ValueError, match="license_name"):
        AssetRecord(
            local_path=str(source),
            source_uri="https://example.invalid/asset.svg",
            rights_status="licensed",
            sha256=digest,
        )
    with pytest.raises(ValueError, match="recovery_action"):
        AssetRecord(
            local_path=str(source),
            source_ref="fixture:recovered",
            rights_status="generated",
            sha256=digest,
            ingestion_status="recovered",
        )


def test_asset_dimensions_are_verified_against_local_bytes(tmp_path) -> None:
    source = tmp_path / "asset.png"
    Image.new("RGB", (320, 180), "navy").save(source)
    asset = _asset(source, "asset-dimensions", dimensions=(640, 360))
    deck = _deck(
        ImageObject(
            id="image-dimensions",
            asset_id=asset.asset_id,
            alt="尺寸校验",
        ),
        assets=[asset],
    )
    with pytest.raises(PresentationAssetError, match="dimensions mismatch"):
        resolve_deck_assets(deck)


def test_design_measurement_handles_cjk_and_validates_chart_semantics() -> None:
    cjk = measure_text(
        "中文测试",
        font_size=18,
        frame_width=2,
        frame_height=1,
    )
    latin = measure_text(
        "test",
        font_size=18,
        frame_width=2,
        frame_height=1,
    )
    assert cjk.weighted_characters > latin.weighted_characters

    measured = apply_design_measurements(
        _deck(
            TextObject(
                id="measured-title",
                role="title",
                text="中文与 Latin 混排标题",
                frame=ObjectFrame(x=0.5, y=0.5, w=4, h=0.8),
            ),
            ChartObject(
                id="audited-chart",
                categories=["Q1", "Q2"],
                series={"收入": [10, 12]},
                unit="万元",
                number_format="0.0",
                derivation="direct values from the cited fixture",
                source_refs=["fixture://chart-semantics"],
            ),
        ),
    )
    title = measured.slides[0].objects[0]
    assert isinstance(title, TextObject)
    assert title.font_size is not None and title.font_size >= 17
    assert measured.slides[0].scene_hash is not None

    with pytest.raises(ValueError, match="derivation"):
        apply_design_measurements(
            _deck(
                ChartObject(
                    id="unbound-chart",
                    categories=["Q1"],
                    series={"收入": [10]},
                    unit="万元",
                    source_refs=["fixture://chart-semantics"],
                ),
            ),
        )


def test_pipeline_registry_pins_old_jobs_and_rolls_back_new_default(
    tmp_path,
) -> None:
    registry = PresentationPipelineRegistry(tmp_path)
    pinned_v3 = registry.bind_new_request(_request())
    assert pinned_v3.pipeline_version == "lantai-presentation-v3"

    registry.register(
        PipelineCapability(
            version="lantai-presentation-v4-canary",
            composition_engine="pptxgenjs_v3",
        ),
    )
    registry.activate_for_new_jobs(
        "lantai-presentation-v4-canary",
        reason="controlled acceptance passed",
        metrics={"acceptance_pass_rate": 1.0},
    )
    new_default = registry.bind_new_request(
        _request(idempotency_key="v4-default-request"),
    )
    explicit_v3 = registry.bind_new_request(
        _request(
            pipeline_version="lantai-presentation-v3",
            idempotency_key="v3-pinned-request",
        ),
    )
    assert new_default.pipeline_version == "lantai-presentation-v4-canary"
    assert explicit_v3.pipeline_version == "lantai-presentation-v3"

    state = registry.rollback_default(
        "lantai-presentation-v3",
        reason="canary regression exceeded threshold",
        metrics={"visual_regression_rate": 0.12},
    )
    assert state.default_version == "lantai-presentation-v3"
    assert [item.action for item in state.transitions] == [
        "activate",
        "rollback",
    ]
    assert registry.bind_new_request(
        _request(idempotency_key="post-rollback-request"),
    ).pipeline_version == "lantai-presentation-v3"
    registry.assert_pinned_available(pinned_v3.pipeline_version)


def test_job_creation_is_idempotent_and_writer_lease_is_cross_store(
    tmp_path,
) -> None:
    service = PresentationService(tmp_path)
    first = service.create_job(_request())
    second = service.create_job(_request())
    assert second.job_id == first.job_id
    assert len(second.events) == 1

    other_process_view = PresentationJobStore(tmp_path)
    with service.store.writer_lease(first.job_id) as lease:
        with pytest.raises(PresentationWriterBusyError):
            with other_process_view.writer_lease(first.job_id):
                pass
        assert (
            service.store.get(first.job_id).active_lease_id == lease.lease_id
        )
    assert service.store.get(first.job_id).active_lease_id is None


def test_writer_lease_recovers_stale_attempt_and_enforces_budget_deadline(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(
        _request(
            idempotency_key="v3-control-recovery",
            cost_budget_units=5,
        ),
    )
    record.active_lease_id = "lease-from-dead-process"
    record.generation = 4
    store._write(record)

    restarted_runtime_id = "run-after-process-restart"
    with store.writer_lease(
        record.job_id,
        runtime_run_id=restarted_runtime_id,
    ) as lease:
        assert lease.recovered_lease_id == "lease-from-dead-process"
        assert lease.generation == 5
        assert lease.runtime_run_id == restarted_runtime_id
        assert store.get(record.job_id).runtime_run_id == restarted_runtime_id
        store.consume_cost(record.job_id, 5)
        with pytest.raises(RuntimeError, match="exhausted"):
            store.assert_execution_budget(record.job_id)

    expired = store.get(record.job_id)
    expired.cost_used_units = 0
    expired.deadline_at = datetime.now(timezone.utc) - timedelta(seconds=1)
    store._write(expired)
    with pytest.raises(TimeoutError, match="deadline"):
        store.assert_execution_budget(record.job_id)


def test_version_ledger_rejects_stale_base_and_non_monotonic_generation(
    tmp_path,
) -> None:
    output = tmp_path / "presentation.pptx"
    output.write_bytes(b"version-0")
    deck = _deck()
    store = PresentationVersionStore(tmp_path)
    initial = store.bootstrap(output, deck, writer="pptxgenjs_v3")
    assert initial.generation == 0

    output.write_bytes(b"version-1")
    with pytest.raises(StalePresentationVersionError, match="generation"):
        store.record(
            output,
            deck,
            writer="lantai_ai_edit",
            parent=initial,
            generation=0,
            base_pptx_sha256=initial.pptx_sha256,
        )
    with pytest.raises(StalePresentationVersionError, match="lease base"):
        store.record(
            output,
            deck,
            writer="lantai_ai_edit",
            parent=initial,
            generation=1,
            base_pptx_sha256="0" * 64,
        )


def test_object_edit_plan_updates_chart_image_geometry_and_object_set(
    tmp_path,
) -> None:
    first_image = tmp_path / "first.png"
    second_image = tmp_path / "second.png"
    first_image.write_bytes(b"first")
    second_image.write_bytes(b"second")
    assets = [
        _asset(first_image, "asset-first"),
        _asset(second_image, "asset-second"),
    ]
    deck = _deck(
        TextObject(id="title", role="title", slot="title", text="旧标题"),
        ChartObject(
            id="chart",
            categories=["旧"],
            series={"值": [1]},
            unit="count",
            derivation="direct fixture values",
            source_refs=["fixture://object-edit"],
        ),
        ImageObject(
            id="image",
            asset_id="asset-first",
            alt="旧图",
        ),
        ShapeObject(
            id="shape",
            shape_type="rect",
            frame=ObjectFrame(x=1, y=1, w=2, h=1),
        ),
        MetricObject(id="remove-me", label="旧指标", value="1"),
        assets=assets,
    )
    plan = PresentationObjectEditPlan(
        slide_id="slide-control",
        actions=[
            ReplaceTextAction(object_id="title", text="新标题"),
            UpdateChartAction(
                object_id="chart",
                chart_type="line",
                categories=["Q1", "Q2"],
                series={"收入": [2, 5]},
                show_legend=False,
                unit="万元",
                number_format="0.0",
                derivation="direct fixture values after approved edit",
                source_refs=["fixture://object-edit-updated"],
            ),
            ReplaceImageAction(
                object_id="image",
                asset_id="asset-second",
                alt="新图",
                crop="contain",
            ),
            MoveResizeAction(
                object_id="shape",
                frame=ObjectFrame(x=4, y=2, w=3, h=2),
            ),
            RemoveObjectAction(object_id="remove-me"),
            AddObjectAction(
                object=TextObject(
                    id="new-caption",
                    role="caption",
                    slot="caption",
                    text="新增说明",
                ),
            ),
        ],
        rationale="对象级、可审计修改",
    )
    updated, changed = apply_object_edit_plan(deck, plan)
    slide = updated.slides[0]
    objects = {obj.id: obj for obj in slide.objects}

    assert set(changed) == {
        "title",
        "chart",
        "image",
        "shape",
        "remove-me",
        "new-caption",
    }
    assert objects["title"].text == "新标题"
    assert objects["chart"].chart_type == "line"
    assert objects["chart"].categories == ["Q1", "Q2"]
    assert objects["chart"].show_legend is False
    assert objects["chart"].unit == "万元"
    assert objects["chart"].source_refs == [
        "fixture://object-edit-updated",
    ]
    assert objects["image"].asset_id == "asset-second"
    assert objects["image"].crop == "contain"
    assert objects["shape"].frame.x == 4
    assert "remove-me" not in objects
    assert "new-caption" in objects
    assert slide.revision == 2
    assert slide.scene_hash is not None
    assert slide.asset_ids == ["asset-second"]


def test_object_edit_plan_can_atomically_update_explicit_additional_slides() -> None:
    deck = _deck(
        TextObject(id="primary-title", role="title", slot="title", text="主页面"),
    )
    deck = deck.model_copy(
        update={
            "slides": [
                deck.slides[0],
                DeckSlide(
                    id="slide-other",
                    layout_id="freeform",
                    objects=[
                        TextObject(
                            id="other-title",
                            role="title",
                            slot="title",
                            text="关联页面",
                        ),
                    ],
                ),
            ],
        },
    )
    plan = PresentationObjectEditPlan(
        slide_id="slide-control",
        additional_slide_ids=["slide-other"],
        actions=[
            ReplaceTextAction(object_id="primary-title", text="主页面已修"),
            ReplaceTextAction(object_id="other-title", text="关联页面已修"),
        ],
        rationale="两个互相依赖的页面必须在一次整套审查中原子提交",
    )

    updated, changed = apply_object_edit_plan(deck, plan)

    assert changed == ["primary-title", "other-title"]
    assert updated.slides[0].objects[0].text == "主页面已修"
    assert updated.slides[1].objects[0].text == "关联页面已修"
    assert updated.slides[0].revision == deck.slides[0].revision + 1
    assert updated.slides[1].revision == deck.slides[1].revision + 1


def test_object_edit_plan_rejects_implicit_cross_slide_scope() -> None:
    deck = _deck(
        TextObject(id="primary-title", role="title", slot="title", text="主页面"),
    )
    deck = deck.model_copy(
        update={
            "slides": [
                deck.slides[0],
                DeckSlide(
                    id="slide-other",
                    layout_id="freeform",
                    objects=[
                        TextObject(
                            id="other-title",
                            role="title",
                            slot="title",
                            text="关联页面",
                        ),
                    ],
                ),
            ],
        },
    )
    plan = PresentationObjectEditPlan(
        slide_id="slide-control",
        actions=[ReplaceTextAction(object_id="other-title", text="越界修改")],
        rationale="未显式声明关联页面",
    )

    with pytest.raises(ValueError, match="unknown object id"):
        apply_object_edit_plan(deck, plan)


def test_object_edit_plan_can_patch_only_the_linked_slide_intent() -> None:
    deck = _deck(
        TextObject(id="title", role="title", slot="title", text="Q1 事实"),
    )
    deck = deck.model_copy(
        update={
            "slides": [
                deck.slides[0].model_copy(update={"intent_id": "intent-q1"}),
            ],
        },
    )
    plan = DeckPlan(
        title="季度复盘",
        narrative_arc="按季度核对经营事实",
        visual_direction=VisualDirection(
            visual_concept="克制经营复盘",
            palette=["#111111", "#FFFFFF"],
            typography="清晰",
            image_language="无装饰图",
            chart_language="原生图表",
            shape_language="少量分隔线",
            grid="十二列",
            recurring_motif="季度标记",
        ),
        slides=[
            DeckPlanSlide(
                order=1,
                title="Q1",
                intent=SlideIntent(
                    id="intent-q1",
                    purpose="核对 Q1",
                    key_message="Q1 客户数为 21",
                ),
            ),
        ],
    )
    edit = PresentationObjectEditPlan(
        slide_id="slide-control",
        intent_patch=SlideIntentPatch(key_message="Q1 客户数为 18"),
        rationale="按本页原始数据纠正主判断",
    )

    updated = apply_slide_intent_patch(deck, plan, edit)

    assert updated.slides[0].intent.key_message == "Q1 客户数为 18"
    assert plan.slides[0].intent.key_message == "Q1 客户数为 21"

    cross_deck = deck.model_copy(update={
        "slides": [
            *deck.slides,
            DeckSlide(
                id="slide-other",
                intent_id="intent-other",
                layout_id="freeform",
                objects=[TextObject(
                    id="other-title",
                    role="title",
                    slot="title",
                    text="其他页",
                )],
            ),
        ],
    })
    cross_plan = plan.model_copy(update={
        "slides": [
            *plan.slides,
            DeckPlanSlide(
                order=2,
                title="其他页",
                intent=SlideIntent(
                    id="intent-other",
                    purpose="核对其他页",
                    key_message="其他页保持不变",
                ),
            ),
        ],
    })
    cross_edit = PresentationObjectEditPlan(
        slide_id="slide-other",
        intent_patch=SlideIntentPatch(
            slide_id="slide-control",
            key_message="Q1 客户数为 18",
        ),
        rationale="与另一页几何修复一起做原子复审",
    )
    cross_updated = apply_slide_intent_patch(
        cross_deck,
        cross_plan,
        cross_edit,
    )
    assert cross_updated.slides[0].intent.key_message == "Q1 客户数为 18"
    assert cross_updated.slides[1].intent.key_message == "其他页保持不变"

    with pytest.raises(ValueError, match="no material change"):
        apply_slide_intent_patch(deck, updated, edit)
    with pytest.raises(ValueError, match="action or intent patch"):
        PresentationObjectEditPlan(slide_id="slide-control")


def test_unrun_evidence_cannot_be_reported_as_verified() -> None:
    report = QualityReport(
        deck_id="deck-v3-control",
        passed=False,
        delivery_status="failed",
        blocking_gaps=["semantic planner evidence", "visual critic evidence"],
    )
    assert report.render_evidence.value == "not_run"
    assert report.visual_evidence.value == "not_run"
    assert report.blocking_gaps == [
        "semantic planner evidence",
        "visual critic evidence",
    ]


def _external_save_runtime(tmp_path, *, grouped=False):
    from pptx import Presentation
    from lantai.presentation.compiler import PresentationCompiler
    from lantai.presentation.models import JobStage

    runtime = PresentationJobRuntime(tmp_path)
    record = runtime.service.create_job(_request())
    store = runtime.service.store
    job_dir = store.job_dir(record.job_id)
    output = job_dir / "output/presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store.write_checkpoint(record.job_id, "deck-ir.json", deck.model_dump_json())
    versions = PresentationVersionStore(job_dir)
    original = versions.record(output, deck, writer="native_pptx")
    presentation = Presentation(str(output))
    shape = next(s for s in presentation.slides[0].shapes if s.name == "title-control")
    shape.text = "人工保存的新标题"
    from pptx.util import Pt
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    if grouped:
        presentation.slides[0].shapes.add_group_shape([shape])
    presentation.save(str(output))
    versions.register_external_save(
        output, writer="genoffice_human", expected_parent_sha256=original.pptx_sha256,
    )
    store.append_event(record.job_id, stage=JobStage.COMPLETED,
                       status=JobStatus.COMPLETED, message="test fixture")
    return runtime, record, output, versions


def _reimport_report_fixture(deck, output, render_dir):
    """Real PPTX inspection with explicitly synthetic render/critic evidence."""
    from PIL import ImageDraw
    from lantai.presentation.models import RenderedSlideEvidence
    from lantai.presentation.quality import inspect_pptx
    render_dir.mkdir(parents=True, exist_ok=True)
    image_path = render_dir / "slide-1.png"
    image = Image.new("RGB", (1280, 720), "white")
    ImageDraw.Draw(image).rectangle((50, 50, 600, 150), fill="black")
    image.save(image_path)
    report = inspect_pptx(deck, output, rendered_slide_count=1)
    report.vision_critique_required = True
    report.visual_verified = True
    report.visual_evidence = "passed"
    report.visual_critic_id = "test-double"
    report.renderer_id = "synthetic-test-fixture"
    report.renderer_version = "1"
    report.font_manifest = ["test-fixture"]
    report.rendered_slides = [RenderedSlideEvidence(slide_id=deck.slides[0].id,
        slide_index=1, sha256=file_sha256(image_path), width=1280, height=720)]
    return report, (image_path,)


@pytest.mark.parametrize("failure", ["group", "concurrent_save", "critic_unavailable", "critic_failed", "critic_not_run", "version_failure"])
def test_editor_reimport_failure_preserves_semantic_checkpoint(tmp_path, monkeypatch, failure):
    from pptx import Presentation
    from lantai.presentation.models import QualityIssue

    runtime, record, output, versions = _external_save_runtime(tmp_path, grouped=failure == "group")
    checkpoint = runtime.service.store.checkpoint_path(record.job_id, "deck-ir.json")
    before_checkpoint = checkpoint.read_bytes()
    before_version = versions.current()
    base_hash = file_sha256(output)
    verified = []

    def verify(deck, *args, **kwargs):
        verified.append(True)
        args[1].mkdir(parents=True, exist_ok=True)
        if failure == "concurrent_save":
            presentation = Presentation(str(output))
            presentation.slides[0].shapes[0].text = "验证期间再次人工保存"
            presentation.save(str(output))
        issues = []
        if failure in {"critic_unavailable", "critic_failed"}:
            issues = [QualityIssue(code=f"vision.{failure}", severity="error", message="test failure")]
        report, images = _reimport_report_fixture(deck, output, args[1])
        report.output_sha256 = base_hash
        report.passed = report.passed and not issues
        report.visual_evidence = "not_run" if failure == "critic_not_run" else "unavailable" if issues else "passed"
        report.visual_verified = not issues and failure != "critic_not_run"
        report.issues.extend(issues)
        return report, images


    monkeypatch.setattr(runtime.service, "verify_artifact", verify)
    persistence_attempts = []
    if failure == "version_failure":
        def fail_record(*args, **kwargs):
            persistence_attempts.append(True)
            raise RuntimeError("simulated version persistence failure")
        monkeypatch.setattr(PresentationVersionStore, "record", fail_record)
    with pytest.raises((ValueError, RuntimeError)):
        runtime.reconcile_editor_save(record.job_id, record.conversation_id, base_output_sha256=base_hash)
    if failure == "version_failure":
        assert persistence_attempts == [True]
    assert checkpoint.read_bytes() == before_checkpoint
    assert versions.current() == before_version
    assert not (output.parent.parent / "quality/quality-report.json").exists()
    if failure == "group":
        assert not verified
    if failure != "concurrent_save":
        assert file_sha256(output) == base_hash
    else:
        assert "验证期间再次人工保存" in Presentation(str(output)).slides[0].shapes[0].text


@pytest.mark.parametrize("failure", ["passed", "delivery_status", "blocking_gaps", "general_error", "visual_failed"])
def test_editor_reimport_never_promotes_failed_report(tmp_path, monkeypatch, failure):
    from lantai.presentation.models import QualityIssue
    runtime, record, output, versions = _external_save_runtime(tmp_path)
    original = versions.current()
    digest = file_sha256(output)
    checkpoint = runtime.service.store.checkpoint_path(record.job_id, "deck-ir.json")
    checkpoint_bytes = checkpoint.read_bytes()
    def verify(deck, *args, **kwargs):
        args[1].mkdir(parents=True, exist_ok=True)
        report, images = _reimport_report_fixture(deck, output, args[1])
        assert report.passed and report.delivery_status == "passed"
        if failure == "passed": report.passed = False
        if failure == "delivery_status": report.delivery_status = "failed"
        if failure == "blocking_gaps": report.blocking_gaps = ["missing verification"]
        if failure == "general_error": report.issues = [QualityIssue(code="design.text_wall", severity="error", message="Too dense")]
        if failure == "visual_failed": report.visual_evidence = "failed"
        return report, images
    monkeypatch.setattr(runtime.service, "verify_artifact", verify)
    with pytest.raises(ValueError, match="未通过"):
        runtime.reconcile_editor_save(record.job_id, record.conversation_id, base_output_sha256=digest)
    assert runtime.get(record.job_id, record.conversation_id).status == JobStatus.FAILED
    assert versions.current() == original
    assert checkpoint.read_bytes() == checkpoint_bytes
    assert file_sha256(output) == digest


def test_editor_reimport_accepts_complete_bound_evidence(tmp_path, monkeypatch):
    runtime, record, output, versions = _external_save_runtime(tmp_path)
    monkeypatch.setattr(runtime.service, "verify_artifact",
        lambda deck, output, render_dir, **kwargs: _reimport_report_fixture(deck, output, render_dir))
    result = runtime.reconcile_editor_save(record.job_id, record.conversation_id,
        base_output_sha256=file_sha256(output))
    assert result.reconciled
    assert versions.current() == result
    assert runtime.get(record.job_id, record.conversation_id).status == JobStatus.COMPLETED
