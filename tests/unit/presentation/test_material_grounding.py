from __future__ import annotations

import hashlib
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement

from lantai.presentation.cli import _load_request
from lantai.presentation.models import (
    InputMaterial,
    MetricObject,
    PresentationJobRequest,
    TextObject,
)
from lantai.presentation.planner import (
    PresentationMaterialError,
    build_brief,
    build_deck_ir,
    build_plan,
)


def _semantic_objects(slide):
    """Exclude native design grammar while asserting authored semantics."""

    return [obj for obj in slide.objects if "-grammar-" not in obj.id]


def test_markdown_outline_drives_slide_titles_and_body_content(
    tmp_path,
) -> None:
    outline = tmp_path / "outline.md"
    outline.write_text(
        """# WPS for PAD 用户调研报告

## 第1页 封面
从平板适配型产品到平板生产力体验型产品。

## 第2页 核心结论
9/11 用户将 PAD 视为便携阅读和批注终端。
手写体验不佳覆盖 10/11 用户。

## 第3页 产品建议
优先级为基础体验、效率场景、差异化和 AI 探索。
""",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-material",
        user_request="将调研报告整理成 3 页汇报",
        audience="产品决策团队",
        slide_count=3,
        materials=[
            InputMaterial(
                id="source-outline",
                name=outline.name,
                path=str(outline),
                media_type="text/markdown",
            )
        ],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)
    deck = build_deck_ir(request, plan)
    visible_text = "\n".join(
        obj.text
        for slide in deck.slides
        for obj in slide.objects
        if isinstance(obj, TextObject)
    )

    assert [slide.title for slide in plan.slides] == [
        "封面",
        "核心结论",
        "产品建议",
    ]
    assert "9/11 用户将 PAD 视为便携阅读和批注终端" in visible_text
    assert "手写体验不佳覆盖 10/11 用户" in visible_text
    assert "基础体验、效率场景、差异化和 AI 探索" in visible_text
    assert "推进叙事" not in visible_text
    assert "先明确目标、证据与行动" not in visible_text
    assert deck.source_refs == ["source-outline"]
    assert all(fact in visible_text for fact in deck.required_facts)
    assert {"9/11", "10/11"}.issubset(set(deck.immutable_data))


def test_deck_immutable_data_comes_only_from_explicit_slide_intents(
    tmp_path,
) -> None:
    source = tmp_path / "immutable-source.md"
    source.write_text(
        "## 结论\n材料明确硬事实为 42%。",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-explicit-immutable",
        user_request="将材料整理成一页汇报",
        slide_count=1,
        materials=[InputMaterial(name=source.name, path=str(source))],
    )
    brief = build_brief(request)
    plan = build_plan(request, brief)
    planned_slide = plan.slides[0]
    plan = plan.model_copy(
        update={
            "slides": [
                planned_slide.model_copy(
                    update={
                        "content_blocks": [
                            f"{planned_slide.content_blocks[0]} 连接字段 314159",
                        ],
                    },
                ),
            ],
        },
    )

    deck = build_deck_ir(request, plan)

    assert deck.immutable_data == planned_slide.intent.immutable_facts
    assert "314159" not in deck.immutable_data


def test_two_different_materials_produce_different_deck_content(
    tmp_path,
) -> None:
    first = tmp_path / "first.md"
    second = tmp_path / "second.md"
    first.write_text("## 第1页 A\n独有事实 Alpha-314", encoding="utf-8")
    second.write_text("## 第1页 B\n独有事实 Beta-271", encoding="utf-8")

    def deck_text(path):
        request = PresentationJobRequest(
            conversation_id="conversation-differential",
            user_request="根据材料生成汇报",
            slide_count=1,
            materials=[InputMaterial(name=path.name, path=str(path))],
        )
        brief = build_brief(request)
        deck = build_deck_ir(request, build_plan(request, brief))
        return "\n".join(
            obj.text
            for slide in deck.slides
            for obj in slide.objects
            if isinstance(obj, TextObject)
        )

    assert deck_text(first) != deck_text(second)
    assert "Alpha-314" in deck_text(first)
    assert "Beta-271" in deck_text(second)


def test_merging_two_sources_into_one_page_keeps_both_provenance_keys(
    tmp_path,
) -> None:
    first = tmp_path / "first.md"
    second = tmp_path / "second.md"
    first.write_text("## 第一来源\n事实 A", encoding="utf-8")
    second.write_text("## 第二来源\n事实 B", encoding="utf-8")
    request = PresentationJobRequest(
        conversation_id="conversation-merged-sources",
        user_request="将两份材料合成一页",
        slide_count=1,
        materials=[
            InputMaterial(id="source-A", name=first.name, path=str(first)),
            InputMaterial(id="source-B", name=second.name, path=str(second)),
        ],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)
    deck = build_deck_ir(request, plan)

    assert plan.slides[0].intent.source_refs == ["source-A", "source-B"]
    assert deck.source_refs == ["source-A", "source-B"]
    assert len(brief.source_coverage) == 2


def test_cli_resolves_relative_material_path_from_request_file(
    tmp_path,
) -> None:
    (tmp_path / "outline.md").write_text("## 第1页 结论\n真实内容", encoding="utf-8")
    request_path = tmp_path / "request.json"
    request_path.write_text(
        """{
          "conversation_id": "conversation-relative",
          "user_request": "生成汇报",
          "slide_count": 1,
          "materials": [{"type": "file", "path": "outline.md"}]
        }""",
        encoding="utf-8",
    )

    request = _load_request(request_path)

    assert request.materials[0].path == str(
        (tmp_path / "outline.md").resolve()
    )
    assert build_brief(request).source_sections[0].body == "真实内容"


def test_cli_rejects_material_outside_request_workspace(tmp_path) -> None:
    request_dir = tmp_path / "request-workspace"
    request_dir.mkdir()
    outside = tmp_path / "outside.md"
    outside.write_text("不应被读取", encoding="utf-8")
    request_path = request_dir / "request.json"
    request_path.write_text(
        """{
          "conversation_id": "conversation-scope",
          "user_request": "生成汇报",
          "slide_count": 1,
          "materials": [{"name": "outside.md", "path": "../outside.md"}]
        }""",
        encoding="utf-8",
    )

    with pytest.raises(ValueError, match="must stay inside"):
        _load_request(request_path)


def test_declared_unreadable_material_fails_closed(tmp_path) -> None:
    request = PresentationJobRequest(
        conversation_id="conversation-missing-material",
        user_request="根据材料生成汇报",
        materials=[
            InputMaterial(
                name="missing.md", path=str(tmp_path / "missing.md")
            ),
        ],
    )

    with pytest.raises(PresentationMaterialError, match="does not exist"):
        build_brief(request)


def test_pdf_text_extraction_preserves_digest_and_declares_visual_limitations(
    tmp_path, monkeypatch
) -> None:
    source = tmp_path / "source.pdf"
    source.write_bytes(b"%PDF-1.4\ntext-fixture")
    extracted_text = "行业事实：样本覆盖 16 页，结论必须保留。"

    def fake_extract(path, *, detected_suffix=None, require_complete_pdf=False, pdf_max_chars=None):
        assert path != source
        assert path.name == source.name
        assert path.read_bytes() == source.read_bytes()
        assert detected_suffix == ".pdf"
        assert require_complete_pdf is True
        assert pdf_max_chars == 1_000_000
        return {
            "text": extracted_text,
            "status": "text_extracted",
            "method": "pdfkit",
            "message": "已抽取可检索正文。",
            "text_chars": len(extracted_text),
            "page_count": 16,
            "text_complete": True,
            "extracted_pages": 16,
            "analyzed_pages": 12,
            "layout_block_count": 15,
        }

    monkeypatch.setattr(
        "lantai.app.wiki.core.extract_source_text_detail",
        fake_extract,
    )
    request = PresentationJobRequest(
        conversation_id="conversation-pdf-source",
        user_request="根据 PDF 生成汇报",
        materials=[
            InputMaterial(
                id="pdf-source",
                name=source.name,
                path=str(source),
                media_type="application/pdf",
            )
        ],
    )

    brief = build_brief(request)

    coverage = brief.source_coverage[0]
    assert extracted_text in brief.source_sections[0].body
    assert coverage.material_id == "pdf-source"
    assert coverage.source_sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    assert coverage.converted_sha256 is None
    assert coverage.method == "pdf_pdfkit"
    assert coverage.status == "supported_content_extracted"
    assert coverage.extracted_chars == len(extracted_text)
    assert coverage.limitations == [
        "PDF text extraction covered 16 pages; charts, images, and visual layout "
        "were not treated as grounded facts. Only extracted text may be used."
    ]


@pytest.mark.parametrize(
    "status",
    ["ocr_extracted", "no_text", "encrypted", "corrupt", "failed"],
)
def test_pdf_non_text_or_failed_extraction_fails_closed_without_leaking_content(
    tmp_path, monkeypatch, status
) -> None:
    source = tmp_path / f"{status}.pdf"
    source.write_bytes(b"%PDF-1.4\nunsafe-fixture")

    monkeypatch.setattr(
        "lantai.app.wiki.core.extract_source_text_detail",
        lambda _path, **kwargs: {
            "text": "DO-NOT-ECHO-UNVERIFIED-PDF-CONTENT",
            "status": status,
            "method": "fixture",
            "message": "DO-NOT-ECHO-EXTRACTOR-DETAIL",
            "page_count": 3,
        },
    )
    request = PresentationJobRequest(
        conversation_id=f"conversation-pdf-{status}",
        user_request="根据 PDF 生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError) as exc_info:
        build_brief(request)

    message = str(exc_info.value)
    assert status in message
    assert "DO-NOT-ECHO-UNVERIFIED-PDF-CONTENT" not in message
    assert "DO-NOT-ECHO-EXTRACTOR-DETAIL" not in message


def test_docx_body_and_tables_keep_order_and_source_coverage(tmp_path) -> None:
    source = tmp_path / "研究记录.docx"
    document = Document()
    document.add_heading("研究记录", level=1)
    document.add_heading("核心判断", level=2)
    document.add_paragraph("样本只有 11 人，不能外推到全部用户。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "反馈人数"
    table.cell(1, 1).text = "9/11"
    document.add_paragraph("下一步先核实分母。")
    document.save(source)
    request = PresentationJobRequest(
        conversation_id="conversation-docx-body",
        user_request="根据材料生成汇报",
        materials=[InputMaterial(id="word-source", name=source.name, path=str(source))],
    )

    brief = build_brief(request)

    assert brief.source_title == "研究记录"
    assert len(brief.source_sections) == 1
    body = brief.source_sections[0].body
    assert body.index("样本只有 11 人") < body.index("反馈人数") < body.index("下一步")
    assert "反馈人数 | 9/11" in body
    assert brief.source_coverage[0].material_id == "word-source"
    assert brief.source_coverage[0].status == "supported_content_extracted"
    assert brief.source_coverage[0].method == "docx_ooxml"
    assert brief.source_coverage[0].source_sha256


def test_docx_intro_before_first_section_is_not_dropped(tmp_path) -> None:
    source = tmp_path / "intro.docx"
    document = Document()
    document.add_heading("研究记录", level=1)
    document.add_paragraph("范围说明：只覆盖 11 位受访者。")
    document.add_heading("核心判断", level=2)
    document.add_paragraph("9 位提到了批注。")
    document.save(source)
    request = PresentationJobRequest(
        conversation_id="conversation-docx-intro",
        user_request="根据材料汇报",
        slide_count=1,
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)

    assert "只覆盖 11 位受访者" in "\n".join(
        section.body for section in brief.source_sections
    )
    assert "只覆盖 11 位受访者" in "\n".join(
        slide.intent.required_facts[0] for slide in plan.slides
    )


@pytest.mark.parametrize(
    ("part_name", "message"),
    [
        ("word/comments.xml", "comments"),
        ("word/footnotes.xml", "footnotes"),
    ],
)
def test_docx_with_unreadable_annotation_part_fails_closed(
    tmp_path, part_name, message
) -> None:
    source = tmp_path / "annotated.docx"
    document = Document()
    document.add_paragraph("正文事实 42")
    document.save(source)
    with ZipFile(source, "a") as archive:
        archive.writestr(part_name, "<root>unread annotation</root>")
    request = PresentationJobRequest(
        conversation_id="conversation-docx-annotation",
        user_request="生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match=message):
        build_brief(request)


def test_docx_with_header_part_fails_closed(tmp_path) -> None:
    source = tmp_path / "header.docx"
    document = Document()
    document.add_paragraph("正文事实 42")
    document.save(source)
    with ZipFile(source, "a") as archive:
        archive.writestr("word/header1.xml", "<root/>")
    request = PresentationJobRequest(
        conversation_id="conversation-docx-header",
        user_request="生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match="header/footer"):
        build_brief(request)


def test_docx_with_drawing_fails_instead_of_silently_dropping_chart(
    tmp_path,
) -> None:
    source = tmp_path / "chart.docx"
    document = Document()
    document.add_paragraph("图表中的数值必须核验")
    document.save(source)
    updated = BytesIO()
    with ZipFile(source) as archive, ZipFile(updated, "w") as destination:
        original = archive.read("word/document.xml")
        for name in archive.namelist():
            payload = archive.read(name)
            if name == "word/document.xml":
                payload = original.replace(
                    b"</w:body>",
                    b"<w:p><w:r><w:drawing/></w:r></w:p></w:body>",
                )
            destination.writestr(name, payload)
    source.write_bytes(updated.getvalue())
    request = PresentationJobRequest(
        conversation_id="conversation-docx-chart",
        user_request="生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match="drawing"):
        build_brief(request)


def test_docx_with_equation_fails_instead_of_silently_dropping_formula(
    tmp_path,
) -> None:
    source = tmp_path / "formula.docx"
    document = Document()
    document.add_paragraph("公式下方的文字")
    document.save(source)
    updated = BytesIO()
    with ZipFile(source) as archive, ZipFile(updated, "w") as destination:
        for name in archive.namelist():
            payload = archive.read(name)
            if name == "word/document.xml":
                payload = payload.replace(
                    b"</w:body>",
                    (
                        b'<w:p><m:oMath xmlns:m="http://schemas.openxmlformats.org/'
                        b'officeDocument/2006/math"/></w:p></w:body>'
                    ),
                )
            destination.writestr(name, payload)
    source.write_bytes(updated.getvalue())
    request = PresentationJobRequest(
        conversation_id="conversation-docx-formula",
        user_request="生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match="equation"):
        build_brief(request)


def test_docx_with_merged_cells_fails_instead_of_flattening_table_meaning(
    tmp_path,
) -> None:
    source = tmp_path / "merged-table.docx"
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "两个列共用的标题"
    table.cell(1, 0).text = "11"
    table.cell(1, 1).text = "9"
    document.save(source)
    request = PresentationJobRequest(
        conversation_id="conversation-docx-merged",
        user_request="根据表格生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match="merged table cells"):
        build_brief(request)


def test_docx_with_inherited_style_numbering_fails_instead_of_losing_labels(
    tmp_path,
) -> None:
    source = tmp_path / "numbered-style.docx"
    document = Document()
    numbered = document.styles.add_style("FactualList", WD_STYLE_TYPE.PARAGRAPH)
    numbered._element.get_or_add_pPr().append(OxmlElement("w:numPr"))
    child = document.styles.add_style("FactualListChild", WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = numbered
    document.add_paragraph("结论", style="FactualListChild")
    document.save(source)
    request = PresentationJobRequest(
        conversation_id="conversation-docx-style-numbering",
        user_request="根据列表生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match="numbering"):
        build_brief(request)


@pytest.mark.parametrize(
    ("injected_xml", "reason"),
    [
        (b"<w:hyperlink><w:r><w:t>source</w:t></w:r></w:hyperlink>", "hyperlink"),
        (b"<w:pPr><w:numPr/></w:pPr><w:r><w:t>item</w:t></w:r>", "numbering"),
        (b"<w:r><w:fldChar w:fldCharType=\"begin\"/></w:r>", "field"),
    ],
)
def test_docx_rejects_semantics_not_preserved_by_plain_text_reader(
    tmp_path, injected_xml, reason
) -> None:
    source = tmp_path / "semantic.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(source)
    updated = BytesIO()
    with ZipFile(source) as archive, ZipFile(updated, "w") as destination:
        for name in archive.namelist():
            payload = archive.read(name)
            if name == "word/document.xml":
                payload = payload.replace(
                    b"</w:body>",
                    b"<w:p>" + injected_xml + b"</w:p></w:body>",
                )
            destination.writestr(name, payload)
    source.write_bytes(updated.getvalue())
    request = PresentationJobRequest(
        conversation_id="conversation-docx-semantics",
        user_request="生成汇报",
        materials=[InputMaterial(name=source.name, path=str(source))],
    )

    with pytest.raises(PresentationMaterialError, match=reason):
        build_brief(request)


def test_legacy_doc_conversion_is_marked_unverified(tmp_path, monkeypatch) -> None:
    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"old binary source")
    converted = tmp_path / "converted.docx"
    document = Document()
    document.add_paragraph("旧稿中的关键事实 28")
    document.save(converted)

    def fake_convert(source, extension, output_dir):
        assert source != legacy
        assert source.name == legacy.name
        assert source.read_bytes() == legacy.read_bytes()
        assert extension == "docx"
        assert output_dir
        return SimpleNamespace(output_path=converted)

    def fake_validate(source):
        assert source != legacy
        assert source.name == legacy.name
        assert source.read_bytes() == legacy.read_bytes()

    monkeypatch.setattr(
        "lantai.app.office_conversion.convert_office_document", fake_convert
    )
    monkeypatch.setattr(
        "lantai.presentation.material_extraction._validate_legacy_doc",
        fake_validate,
    )
    request = PresentationJobRequest(
        conversation_id="conversation-legacy-doc",
        user_request="根据旧稿汇报",
        materials=[InputMaterial(name=legacy.name, path=str(legacy))],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)
    deck = build_deck_ir(request, plan)

    assert "旧稿中的关键事实 28" in brief.source_sections[0].body
    assert brief.source_coverage[0].status == "converted_doc_unverified"
    assert brief.source_coverage[0].method == "libreoffice_doc_to_docx_ooxml"
    assert len(brief.source_coverage[0].converted_sha256 or "") == 64
    assert brief.source_coverage[0].limitations
    assert plan.source_coverage[0].status == "converted_doc_unverified"
    assert deck.source_coverage[0].status == "converted_doc_unverified"


def test_invalid_legacy_doc_is_rejected_before_conversion_can_claim_success(
    tmp_path,
) -> None:
    legacy = tmp_path / "corrupt.doc"
    legacy.write_bytes(
        bytes.fromhex("d0cf11e0a1b11ae1")
        + b"SRC01-CORRUPT-NOT-A-WORD-DOCUMENT"
    )
    request = PresentationJobRequest(
        conversation_id="conversation-corrupt-legacy-doc",
        user_request="根据旧稿汇报",
        materials=[InputMaterial(name=legacy.name, path=str(legacy))],
    )

    with pytest.raises(
        PresentationMaterialError,
        match="not a valid OLE compound document",
    ):
        build_brief(request)


def test_materialless_stable_planner_fails_closed_instead_of_inventing_content(
    ) -> None:
    request = PresentationJobRequest(
        conversation_id="conversation-no-material",
        user_request="生成一份 9 页的产品汇报",
        slide_count=9,
    )

    brief = build_brief(request)

    with pytest.raises(
        PresentationMaterialError,
        match="requires source material or model-authored slides",
    ):
        build_plan(request, brief)


def test_outline_count_mismatch_is_restructured_without_dropping_sections(
    tmp_path,
) -> None:
    outline = tmp_path / "three-slides.md"
    outline.write_text(
        "## 第1页 A\n事实 A\n## 第2页 B\n事实 B\n## 第3页 C\n事实 C",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-count-mismatch",
        user_request="生成两页汇报",
        slide_count=2,
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )
    brief = build_brief(request)

    plan = build_plan(request, brief)

    assert len(plan.slides) == 2
    combined = "\n".join(
        block for slide in plan.slides for block in slide.content_blocks
    )
    assert all(f"事实 {letter}" in combined for letter in ("A", "B", "C"))


def test_nested_markdown_headings_stay_inside_their_slide(tmp_path) -> None:
    outline = tmp_path / "nested.md"
    outline.write_text(
        "# 调研\n## 第1页 结论\n### 证据\n唯一事实 42",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-nested-heading",
        user_request="生成一页汇报",
        slide_count=1,
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )

    brief = build_brief(request)

    assert len(brief.source_sections) == 1
    assert brief.source_sections[0].title == "结论"
    assert brief.source_sections[0].body == "证据\n唯一事实 42"


def test_grounded_bullets_choose_semantic_layout_and_keep_object_ids_unique(
    tmp_path,
) -> None:
    outline = tmp_path / "route.md"
    outline.write_text(
        """# AGI 技术交流

## 第1页 封面
从单点能力走向通用智能。

## 第2页 AGI 技术路线图
① 阶段一：扩大基础模型与数据规模
- 阶段二：强化推理、工具使用与反馈
- 阶段三：建立长程任务规划能力
- 阶段四：提升环境交互与自主学习
- 阶段五：形成可验证的通用智能

## 第3页 总结
把可验证的能力进展作为下一阶段决策依据。
""",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-semantic-route",
        user_request="将材料整理成 3 页汇报",
        slide_count=3,
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)
    deck = build_deck_ir(request, plan)

    route_plan = plan.slides[1]
    route_slide = deck.slides[1]
    assert route_plan.intent.semantic_layout == "process"
    assert route_plan.content_blocks == [
        "阶段一：扩大基础模型与数据规模",
        "阶段二：强化推理、工具使用与反馈",
        "阶段三：建立长程任务规划能力",
        "阶段四：提升环境交互与自主学习",
        "阶段五：形成可验证的通用智能",
    ]
    assert [obj.slot for obj in _semantic_objects(route_slide)] == [
        "title",
        "step_1",
        "step_2",
        "step_3",
        "step_4",
        "step_5",
    ]
    object_ids = [obj.id for obj in route_slide.objects]
    assert len(object_ids) == len(set(object_ids))


def test_dense_outline_is_grouped_into_semantic_surfaces_without_losing_facts(
    tmp_path,
) -> None:
    outline = tmp_path / "dense.md"
    outline.write_text(
        """# 投资者交流会

## 第1页 封面
内部投研交流。

## 第2页 核心观点
- 观点一：愿景驱动组织
- 观点二：开源是长期选择
- 观点三：商业化保持克制
- 观点四：只赚取合理利润
- 观点五：AGI 是长期目标
- 观点六：算力是最大差距
- 观点七：国产芯片有历史机遇
- 观点八：团队稳定是核心利益

## 第3页 AGI 技术路线图
- 当前局限：AI 缺少持续学习能力
- 阶梯六步：语言模型 → CoT → Agent → 持续学习 → 自我迭代 → 具身智能
- 今年重点是 Agent
- 下一道坎是持续学习
- 希望下一版模型帮助开发下一代模型
- 相信 Scaling，中国还未摸到墙

## 第4页 总结
以可验证的能力进展作为决策依据。
""",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-dense-outline",
        user_request="将材料整理成 4 页汇报",
        slide_count=4,
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )

    brief = build_brief(request)
    plan = build_plan(request, brief)
    deck = build_deck_ir(request, plan)

    assert plan.slides[1].intent.semantic_layout == "matrix"
    assert plan.slides[2].intent.semantic_layout == "hierarchy"
    assert plan.slides[2].intent.relationship_type.value == "process"
    assert [obj.slot for obj in _semantic_objects(deck.slides[2])] == [
        "title",
        "step_1",
        "step_2",
        "step_3",
        "step_4",
        "step_5",
        "step_6",
    ]
    visible_text = "\n".join(
        obj.text
        for slide in deck.slides
        for obj in slide.objects
        if isinstance(obj, TextObject)
    )
    assert all(fact in visible_text for fact in deck.required_facts)
    object_ids = [obj.id for slide in deck.slides for obj in slide.objects]
    assert len(object_ids) == len(set(object_ids))


def test_grounded_outline_uses_real_cover_title_and_varied_layouts(
    tmp_path,
) -> None:
    outline = tmp_path / "varied.md"
    sections = [
        (
            "封面",
            [
                "深度求索投资者交流会纪要",
                "2026 年 7 月 23 日",
                "内部资料 · 注意保密",
            ],
        ),
        ("公司愿景与组织", [f"组织事实 {index}" for index in range(1, 8)]),
        ("开源战略与商业克制", [f"开源事实 {index}" for index in range(1, 7)]),
        ("定价与商业模式", [f"定价事实 {index}" for index in range(1, 6)]),
        (
            "AGI 技术路线图",
            [
                "阶梯六步：语言模型 → CoT → Agent → 持续学习",
                *[f"路线事实 {index}" for index in range(1, 6)],
            ],
        ),
        ("中美差距与算力策略", [f"对比事实 {index}" for index in range(1, 7)]),
        ("ToB / ToC 业务定位", [f"业务事实 {index}" for index in range(1, 6)]),
        ("关键数字", [f"数字事实 {index}" for index in range(1, 7)]),
        ("结束页", ["内部资料 · 注意保密", "仅限内部使用"]),
    ]
    outline.write_text(
        "# 深度求索投资者交流\n\n"
        + "\n\n".join(
            f"## 第{index}页 {title}\n"
            + "\n".join(f"- {fact}" for fact in facts)
            for index, (title, facts) in enumerate(sections, start=1)
        ),
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-varied-outline",
        user_request="生成专业投研汇报",
        slide_count=len(sections),
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )

    plan = build_plan(request, build_brief(request))
    deck = build_deck_ir(request, plan)

    assert [slide.intent.semantic_layout for slide in plan.slides] == [
        "cover",
        "matrix",
        "matrix",
        "matrix",
        "hierarchy",
        "comparison",
        "comparison",
        "three_cards",
        "closing",
    ]
    cover_text = {
        obj.slot: obj.text
        for obj in deck.slides[0].objects
        if isinstance(obj, TextObject)
    }
    assert cover_text["hero_statement"] == "深度求索投资者交流会纪要"
    assert "2026 年 7 月 23 日" in cover_text["context"]
    assert plan.slides[1].intent.relationship_type.value == "composition"
    assert [obj.slot for obj in _semantic_objects(deck.slides[1])] == [
        "title",
        "lead",
        "support_1",
        "support_2",
        "support_3",
        "support_4",
        "support_5",
        "support_6",
    ]
    assert [obj.slot for obj in _semantic_objects(deck.slides[2])] == [
        "title",
        "lead",
        "support_1",
        "support_2",
        "support_3",
        "support_4",
        "support_5",
    ]


def test_metric_and_quote_sections_use_distinct_visual_objects(
    tmp_path,
) -> None:
    outline = tmp_path / "metrics-and-quotes.md"
    outline.write_text(
        """# 投研摘要

## 第1页 封面
- 投研摘要
- 内部资料

## 第2页 关键数字
- 当前算力：约两万张 H 等效算力
- 华为卡：约一万六千张，只够训当前一代
- API 定价：十个月收回设备成本，约六倍利润
- 激活参数：国内几十 B vs 美国 800B
- 追赶：落后美国一到两年，只用二十分之一算力
- 华为卡对比英伟达：四张顶一张，落后约两年

## 第3页 金句摘录
- 「我们怀着对这个世界非常大的善意来做这个事情。」
- 「你越克制，你越有可能能够做成。」
- 「核心利益只有一个：保持团队的稳定性。」
- 「我们希望 AI 是负担得起的。」
""",
        encoding="utf-8",
    )
    request = PresentationJobRequest(
        conversation_id="conversation-metrics-quotes",
        user_request="生成投研汇报",
        slide_count=3,
        materials=[InputMaterial(name=outline.name, path=str(outline))],
    )

    deck = build_deck_ir(
        request,
        build_plan(request, build_brief(request)),
    )

    metrics = [
        obj for obj in deck.slides[1].objects if isinstance(obj, MetricObject)
    ]
    quotes = [
        obj
        for obj in _semantic_objects(deck.slides[2])
        if isinstance(obj, TextObject) and obj.slot != "title"
    ]
    assert len(metrics) == 6
    assert [metric.slot for metric in metrics] == [
        "lead_metric",
        "support_metric_2",
        "support_metric_3",
        "support_metric_4",
        "support_metric_5",
        "support_metric_6",
    ]
    assert metrics[0].label == "当前算力"
    assert metrics[0].value == "约两万张 H 等效算力"
    assert all(quote.role == "quote" for quote in quotes)
