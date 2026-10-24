from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.models import DeckIR, DeckSlide, ObjectFrame, ShapeObject, TextObject
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx


@pytest.fixture(params=[NativePptxCompiler, PresentationCompiler], ids=["native", "legacy"])
def compiler(request):
    return request.param()


def _label_deck(*, background_last: bool = False, compact: bool = False) -> DeckIR:
    background = ShapeObject(
        id="label-background",
        shape_type="rect",
        fill_color="1565C0",
        frame=ObjectFrame(x=8.4, y=1.95, w=4.18, h=1.1),
    )
    label = TextObject(
        id="label-text",
        role="body",
        text="Category A",
        color="FFFFFF",
        font_size=20,
        align="center",
        valign="mid",
        frame=(
            ObjectFrame(x=8.8, y=2.2, w=2.0, h=0.5) if compact
            else ObjectFrame(x=8.6, y=1.95, w=3.9, h=1.1)
        ),
    )
    return DeckIR(
        id="deck-colored-label",
        title="Native editable label",
        required_facts=["Category A"],
        slides=[
            DeckSlide(
                id="slide-label",
                layout_id="freeform",
                objects=[label, background] if background_last else [background, label],
            ),
        ],
    )


def _label_shapes(presentation):
    return [
        shape for shape in presentation.slides[0].shapes
        if shape.name in {"label-background", "label-text"}
    ]


def _assert_overlap_blocked(report) -> None:
    assert report.passed is False
    assert report.delivery_status == "failed"
    assert any(
        issue.code == "layout.complete_overlap" and issue.severity == "error"
        for issue in report.issues
    )


@pytest.mark.parametrize("shape_type", ["rect", "round_rect"])
@pytest.mark.parametrize("coverage", ["padded", "full"])
def test_editable_label_can_fill_its_native_rectangular_background(tmp_path, compiler, shape_type, coverage) -> None:
    deck = _label_deck()
    background, label = deck.slides[0].objects
    background.shape_type = shape_type
    if coverage == "full":
        label.frame = background.frame.model_copy()
    output = tmp_path / "colored-label.pptx"
    compiler.compile(deck, output)

    native_shapes = _label_shapes(Presentation(output))
    assert [shape.name for shape in native_shapes] == ["label-background", "label-text"]
    assert native_shapes[0].text == ""
    assert native_shapes[1].text == "Category A"
    assert native_shapes[1].width / native_shapes[0].width > 0.9

    report = inspect_pptx(deck, output)

    assert report.editable is True
    assert report.content_complete is True
    assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)


def test_foreground_native_rectangle_covering_a_label_is_blocked(tmp_path, compiler) -> None:
    deck = _label_deck(background_last=True, compact=True)
    output = tmp_path / "covered-label.pptx"
    compiler.compile(deck, output)
    native_shapes = _label_shapes(Presentation(output))
    assert [shape.name for shape in native_shapes] == ["label-text", "label-background"]
    assert native_shapes[0].width * native_shapes[0].height < (
        native_shapes[1].width * native_shapes[1].height * 0.9
    )

    report = inspect_pptx(deck, output)

    assert report.editable is True
    assert report.content_complete is True
    _assert_overlap_blocked(report)
    assert any(
        issue.code == "layout.complete_overlap" and issue.severity == "error"
        and issue.object_id == "label-background"
        for issue in report.issues
    )


def test_overlap_inspects_reopened_native_order_not_the_authored_order(tmp_path, compiler) -> None:
    deck = _label_deck(compact=True)
    original = tmp_path / "original-order.pptx"
    output = tmp_path / "reordered.pptx"
    compiler.compile(deck, original)
    assert not any(
        issue.code == "layout.complete_overlap"
        for issue in inspect_pptx(deck, original).issues
    )
    native = Presentation(original)
    background, label = _label_shapes(native)
    background._element.getparent().remove(background._element)
    label._element.addnext(background._element)
    native.save(output)

    assert [shape.name for shape in _label_shapes(Presentation(output))] == [
        "label-text", "label-background",
    ]
    assert [obj.id for obj in deck.slides[0].objects] == ["label-background", "label-text"]
    _assert_overlap_blocked(inspect_pptx(deck, output))


@pytest.mark.parametrize("shape_type,rotate,text", [
    ("ellipse", 0, ""),
    ("rect", 15, ""),
    ("rect", 0, "Another label"),
], ids=["nonrectangular", "rotated", "nonblank-background"])
def test_high_coverage_label_exception_is_narrow(tmp_path, compiler, shape_type, rotate, text) -> None:
    deck = _label_deck()
    background = deck.slides[0].objects[0]
    background.shape_type = shape_type
    background.rotate = rotate
    background.text = text
    output = tmp_path / "not-a-blank-rectangular-background.pptx"
    compiler.compile(deck, output)

    _assert_overlap_blocked(inspect_pptx(deck, output))


def test_reopened_opaque_label_is_not_exempted(tmp_path, compiler) -> None:
    deck = _label_deck()
    original = tmp_path / "transparent-label.pptx"
    output = tmp_path / "opaque-label.pptx"
    compiler.compile(deck, original)
    native = Presentation(original)
    label = _label_shapes(native)[1]
    label.fill.solid()
    label.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    native.save(output)

    _assert_overlap_blocked(inspect_pptx(deck, output))


def test_near_identical_native_shapes_remain_blocked(tmp_path, compiler) -> None:
    deck = _label_deck()
    deck.required_facts = []
    deck.slides[0].objects[1] = ShapeObject(
        id="duplicate-background",
        shape_type="rect",
        fill_color="FFFFFF",
        frame=ObjectFrame(x=8.6, y=1.95, w=3.9, h=1.1),
    )
    output = tmp_path / "duplicate-shapes.pptx"
    compiler.compile(deck, output)

    _assert_overlap_blocked(inspect_pptx(deck, output))


@pytest.mark.parametrize("alpha", ["0", "50000", "100000"], ids=["transparent", "translucent", "opaque"])
def test_reopened_native_alpha_is_respected_for_compact_containment(tmp_path, compiler, alpha) -> None:
    deck = _label_deck(background_last=True, compact=True)
    original = tmp_path / "solid-background.pptx"
    output = tmp_path / "nonopaque-background.pptx"
    compiler.compile(deck, original)
    native = Presentation(original)
    background = _label_shapes(native)[1]
    color = background._element.xpath("./p:spPr/a:solidFill/*")[0]
    for transform in list(color):
        if transform.tag.endswith("}alpha"):
            color.remove(transform)
    opacity = OxmlElement("a:alpha")
    opacity.set("val", alpha)
    color.append(opacity)
    native.save(output)

    assert _label_shapes(Presentation(output))[1]._element.xpath(
        "./p:spPr/a:solidFill/*/a:alpha/@val",
    ) == [alpha]
    report = inspect_pptx(deck, output)
    if alpha == "100000":
        _assert_overlap_blocked(report)
    else:
        assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)


def test_high_coverage_label_extending_outside_background_remains_blocked(tmp_path, compiler) -> None:
    deck = _label_deck()
    deck.slides[0].objects[1].frame = ObjectFrame(x=8.6, y=1.95, w=4.18, h=1.1)
    output = tmp_path / "outside-background.pptx"
    compiler.compile(deck, output)

    _assert_overlap_blocked(inspect_pptx(deck, output))
