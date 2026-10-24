"""Public Office round trips preserve cell alignment and native typography."""

import pytest
from lxml import etree as XML
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from lantai.presentation.models import DeckIR, DeckSlide, ObjectFrame, TableObject
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


@pytest.mark.parametrize("known", [True, False])
@pytest.mark.parametrize("anchor", ["ctr", "mid"])
def test_native_table_does_not_realign_numbers_or_bolden_plain_headers(tmp_path, known, anchor):
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(9), Inches(2))
    shape.name = "native-table"
    values = [["合成阶段", "数量", "数值"], ["本地验收", "2", "31%"]]
    for ri, row in enumerate(values):
        for ci, value in enumerate(row):
            cell = shape.table.cell(ri, ci)
            cell.text = value
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            # Legacy PptxGenJS exports leaked the IR spelling into OOXML.
            # Import must read these saves without mutating their native XML.
            cell._tc.get_or_add_tcPr().set("anchor", anchor)
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.font.size = Pt(18)
            paragraph.font.bold = False
            paragraph.font.italic = False
            paragraph.runs[0].font.color.rgb = RGBColor.from_string("315F9A")
            for side in ("lnL", "lnR", "lnT", "lnB"):
                border = OxmlElement("a:" + side)
                if ri == 1:
                    border.set("w", "25400")
                fill = OxmlElement("a:solidFill")
                rgb = OxmlElement("a:srgbClr")
                rgb.set("val", "D4DDE1")
                fill.append(rgb)
                border.append(fill)
                cell._tc.get_or_add_tcPr().append(border)
    objects = [TableObject(id=shape.name, headers=values[0], rows=values[1:],
                           frame=ObjectFrame(x=1, y=1, w=9, h=2))] if known else []
    deck = DeckIR(id="table-style-deck", title="合成测试", slides=[DeckSlide(
        id="table-style-slide", layout_id="freeform", objects=objects,
    )])
    source = tmp_path / "source.pptx"
    presentation.save(source)
    before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    imported_table = next(obj for obj in imported.slides[0].objects if isinstance(obj, TableObject))
    assert all(style.valign == "mid" for row in imported_table.cell_styles for style in row)
    output = tmp_path / "exported.pptx"
    NativePptxCompiler().compile(imported, output)
    native = next(s for s in Presentation(output).slides[0].shapes if s.has_table)
    assert [[c.text for c in r.cells] for r in native.table.rows] == values
    original_table = shape.table
    for ri, row in enumerate(native.table.rows):
        for ci, cell in enumerate(row.cells):
            for side in ("lnL", "lnR", "lnT", "lnB"):
                expected = original_table.cell(ri, ci)._tc.get_or_add_tcPr().find(qn("a:" + side))
                actual = cell._tc.get_or_add_tcPr().find(qn("a:" + side))
                assert XML.tostring(actual, method="c14n", exclusive=True) == XML.tostring(expected, method="c14n", exclusive=True)
    for row in native.table.rows:
        for cell in row.cells:
            assert cell.vertical_anchor == MSO_ANCHOR.MIDDLE
            assert cell._tc.get_or_add_tcPr().get("anchor") == "ctr"
            paragraph = cell.text_frame.paragraphs[0]
            assert paragraph.alignment == PP_ALIGN.LEFT
            assert all(run.font.bold is not True and run.font.italic is not True
                       for run in paragraph.runs)
            assert all(str(run.font.color.rgb) == "315F9A" for run in paragraph.runs)
            assert all(run.font.size.pt == 18 for run in paragraph.runs)
    assert source.read_bytes() == before


@pytest.mark.parametrize("known", [True, False])
@pytest.mark.parametrize("anchor", ["middle", "invalid-anchor"])
def test_unknown_native_table_anchor_is_rejected_without_rewriting_source(tmp_path, known, anchor):
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_table(2, 1, Inches(1), Inches(1), Inches(4), Inches(2))
    shape.name = "unknown-anchor-table"
    for row, value in zip(shape.table.rows, ["合成阶段", "本地验收"]):
        cell = row.cells[0]
        cell.text = value
        cell.text_frame.paragraphs[0].font.size = Pt(18)
        cell._tc.get_or_add_tcPr().set("anchor", anchor)
    objects = [TableObject(id=shape.name, headers=["合成阶段"], rows=[["本地验收"]],
                           frame=ObjectFrame(x=1, y=1, w=4, h=2))] if known else []
    deck = DeckIR(id="unknown-anchor-deck", title="合成测试", slides=[DeckSlide(
        id="unknown-anchor-slide", layout_id="freeform", objects=objects,
    )])
    source = tmp_path / "unknown-anchor.pptx"
    presentation.save(source)
    before = source.read_bytes()

    with pytest.raises(ValueError):
        reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")

    assert source.read_bytes() == before
