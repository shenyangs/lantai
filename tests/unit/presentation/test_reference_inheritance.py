from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from lantai.presentation.models import ConstraintLevel, TemplateMode
from lantai.presentation.reference_deck.models import ReferenceVisualMeasurement
from lantai.presentation.reference_deck.object_parser import PptxObjectParser
from lantai.presentation.reference_deck.profile_builder import ReferenceProfileBuilder


def _fixture(
    tmp_path: Path,
    *,
    slide_count: int = 3,
    slide_show: str | None = None,
    layout_show: str | None = None,
    hidden_logo: str | None = None,
) -> Path:
    """Build controlled native inheritance test data, not a product candidate."""

    presentation = Presentation()
    layout = presentation.slide_layouts[6]
    for number in range(1, slide_count + 1):
        slide = presentation.slides.add_slide(layout)
        if slide_show is not None:
            slide._element.set("showMasterSp", slide_show)
        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        title.text = f"Controlled page {number}"
        title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
        body = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        body.text = "Controlled content"
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    if layout_show is not None:
        layout._element.set("showMasterSp", layout_show)

    scratch = presentation.slides[0]
    image_bytes = BytesIO()
    Image.new("RGB", (16, 8), "#683FA5").save(image_bytes, format="PNG")
    image_bytes.seek(0)
    picture = scratch.shapes.add_picture(
        image_bytes, Inches(8), Inches(0.1), width=Inches(1), height=Inches(0.5),
    )
    picture.name = "master-logo"
    master_picture = deepcopy(picture._element)
    old_rid = master_picture.xpath(".//a:blip")[0].get(qn("r:embed"))
    new_rid = layout.slide_master.part.relate_to(scratch.part.related_part(old_rid), RT.IMAGE)
    master_picture.xpath(".//a:blip")[0].set(qn("r:embed"), new_rid)
    if hidden_logo is not None:
        master_picture.xpath("./p:nvPicPr/p:cNvPr")[0].set("hidden", hidden_logo)
    layout.slide_master.shapes._spTree.append(master_picture)
    scratch.shapes._spTree.remove(picture._element)
    scratch.part.drop_rel(old_rid)

    wordmark = scratch.shapes.add_textbox(Inches(6), Inches(0.1), Inches(1.5), Inches(0.4))
    wordmark.name = "layout-wordmark"
    wordmark.text = "Controlled Brand"
    run = wordmark.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string("683FA5")
    layout.shapes._spTree.append(deepcopy(wordmark._element))
    scratch.shapes._spTree.remove(wordmark._element)

    # This shape is in a valid, but unused layout. It is not visible evidence.
    unused = deepcopy(master_picture)
    unused.xpath("./p:nvPicPr/p:cNvPr")[0].set("name", "unused-logo")
    unused_layout = presentation.slide_layouts[0]
    unused_rid = unused_layout.part.relate_to(layout.slide_master.part.related_part(new_rid), RT.IMAGE)
    unused.xpath(".//a:blip")[0].set(qn("r:embed"), unused_rid)
    unused_layout.shapes._spTree.append(unused)

    path = tmp_path / "controlled-reference.pptx"
    presentation.save(path)
    return path


def _build(path: Path):
    return ReferenceProfileBuilder().build(
        PptxObjectParser().parse(path), ReferenceVisualMeasurement(),
        profile_id="controlled-inheritance", template_mode=TemplateMode.STRICT,
    )


def test_active_inheritance_is_measured_without_changing_page_content_counts(tmp_path: Path) -> None:
    path = _fixture(tmp_path)
    original = sha256(path.read_bytes()).hexdigest()
    deck = PptxObjectParser().parse(path)

    assert sha256(path.read_bytes()).hexdigest() == original == deck.source_sha256
    all_ids = []
    for slide in deck.slides:
        assert slide.object_count == slide.text_object_count == 2
        assert slide.image_count == 0
        assert len(slide.objects) == 2
        assert all(item.source_layer == "slide" for item in slide.objects)
        assert {(item.name, item.source_layer) for item in slide.inherited_objects} == {
            ("master-logo", "master"), ("layout-wordmark", "layout"),
        }
        logo = next(item for item in slide.inherited_objects if item.role_hint == "logo")
        assert logo.kind == "image"
        assert logo.image_aspect_ratio == logo.displayed_aspect_ratio == 2
        all_ids.extend(item.object_id for item in [*slide.objects, *slide.inherited_objects])
    assert len(all_ids) == len(set(all_ids))
    assert "Controlled Brand" not in deck.model_dump_json()
    assert "unused-logo" not in deck.model_dump_json()


@pytest.mark.parametrize("flag", ["0", "false"])
def test_slide_hiding_inheritance_does_not_supply_brand_evidence(tmp_path: Path, flag: str) -> None:
    result = _build(_fixture(tmp_path, slide_show=flag))
    assert all(not slide.inherited_objects for slide in result.measurement.slides)
    assert not any(item.category == "logo" for item in result.design_context.profile.hard_constraints)


@pytest.mark.parametrize("flag", ["0", "false"])
def test_layout_hiding_master_keeps_only_its_own_visible_objects(tmp_path: Path, flag: str) -> None:
    result = _build(_fixture(tmp_path, layout_show=flag))
    assert all(
        [(item.name, item.source_layer) for item in slide.inherited_objects]
        == [("layout-wordmark", "layout")]
        for slide in result.measurement.slides
    )
    assert not any(item.category == "logo" for item in result.design_context.profile.hard_constraints)


@pytest.mark.parametrize("flag", ["1", "true"])
def test_hidden_native_logo_is_not_a_preservation_requirement(tmp_path: Path, flag: str) -> None:
    result = _build(_fixture(tmp_path, hidden_logo=flag))
    assert all(
        all(item.name != "master-logo" for item in slide.inherited_objects)
        for slide in result.measurement.slides
    )
    assert not any(item.category == "logo" for item in result.design_context.profile.hard_constraints)


def test_visible_master_logo_reaches_hard_brand_constraints_and_observed_palette(tmp_path: Path) -> None:
    result = _build(_fixture(tmp_path))
    logos = [item for item in result.design_context.profile.hard_constraints if item.category == "logo"]
    assert logos
    inherited_ids = {
        item.object_id for slide in result.measurement.slides for item in slide.inherited_objects
        if item.role_hint == "logo"
    }
    assert inherited_ids == {object_id for item in logos for object_id in item.object_ids}
    assert all(item.level == ConstraintLevel.HARD for item in logos)
    assert "683FA5" in result.style_profile.palette
    assert result.style_profile.typography["body_size"] >= 17


def test_one_slide_visible_declared_logo_is_not_lost_to_repetition_threshold(tmp_path: Path) -> None:
    result = _build(_fixture(tmp_path, slide_count=1))
    logos = [item for item in result.design_context.profile.brand_constraints if item.category == "logo"]
    assert logos
    assert all(item.level == ConstraintLevel.HARD for item in logos)
