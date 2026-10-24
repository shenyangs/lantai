"""Slide-local exports retain the caller's lifecycle and publication fence."""
from threading import Barrier, Event, Thread
from types import SimpleNamespace
import json
import time

from pptx import Presentation

from lantai.app import office_conversion
from lantai.presentation.render import render_pptx
from lantai.utils.operation_control import (
    OperationControl, bind_operation, check_operation, current_operation,
)


def test_failed_page_settles_siblings_before_return_and_never_publishes(tmp_path, monkeypatch):
    source = tmp_path / 'source.pptx'
    deck = Presentation()
    for _ in range(8):
        deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    output = tmp_path / 'render'
    output.mkdir()
    old_pdf = output / 'presentation.pdf'
    old_pdf.write_bytes(b'previous verified PDF')
    ready = Barrier(4)
    failure = Event()
    emergency_release = Event()
    finished = Event()
    admitted = []
    controls = []
    errors = []
    parent = OperationControl(
        owner_key=('job-id', 'lease-id'), identity=('original generation',),
        ensure_active=lambda: None, remaining_seconds=lambda: 60,
    )
    monkeypatch.setattr(office_conversion, 'find_soffice', lambda: '/controlled/soffice')

    def converter(args, **kwargs):
        options = json.loads(args[args.index('--convert-to') + 1].split(':', 2)[2])
        number = int(options['PageRange']['value'])
        controls.append(current_operation())
        admitted.append(number)
        if number <= 4:
            ready.wait(timeout=3)
        if number == 1:
            failure.set()
            return SimpleNamespace(returncode=1, stdout='', stderr='controlled export failure')
        while not emergency_release.is_set():
            check_operation()
            time.sleep(.005)
        raise RuntimeError('emergency test cleanup')

    monkeypatch.setattr(office_conversion.subprocess, 'run', converter)

    def render():
        try:
            with bind_operation(parent):
                render_pptx(source, output, timeout=10)
        except BaseException as exc:
            errors.append(exc)
        finally:
            finished.set()

    thread = Thread(target=render)
    thread.start()
    try:
        assert failure.wait(3)
        assert finished.wait(2), 'failed page left sibling conversions running'
        assert len(errors) == 1
        assert isinstance(errors[0], RuntimeError)
        assert 'controlled export failure' in str(errors[0])
        assert set(admitted) == {1, 2, 3, 4}
        assert all(control.owner_key == parent.owner_key for control in controls)
        assert all(control.identity == parent.identity for control in controls)
        assert all(control.resources is parent.resources for control in controls)
        assert old_pdf.read_bytes() == b'previous verified PDF'
        assert tuple(output.glob('slide-*.png')) == ()
        assert tuple(output.glob('.slide-local-render-*')) == ()
    finally:
        emergency_release.set()
        thread.join(5)
        assert not thread.is_alive()
