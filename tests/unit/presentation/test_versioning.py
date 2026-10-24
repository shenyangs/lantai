from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

import lantai.presentation.versioning as versioning_module
from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.models import (
    ChartObject, DeckIR, DeckSlide, NativeShapeStyle, ObjectFrame, ShapeObject, TableObject,
    TemplateAnchorUse, TemplatePackageBinding, TemplateSlideAnchors, TextObject,
)
from lantai.presentation.quality import file_sha256, is_verified_delivery
from lantai.presentation.render import RenderResult, slide_image_paths
from lantai.presentation.service import PresentationService
from lantai.presentation.versioning import (
    PresentationVersionStore,
    StalePresentationVersionError,
    deck_ir_version,
)


def _deck(title: str = "版本所有权") -> DeckIR:
    return DeckIR(
        id="versioned-deck",
        title=title,
        slides=[
            DeckSlide(
                id="slide-1",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="title-1",
                        role="title",
                        slot="title",
                        text=title,
                    )
                ],
            )
        ],
    )


def test_version_hash_ignores_only_absent_legacy_chart_legend_policy() -> None:
    deck = DeckIR(
        id="legacy-chart-deck",
        title="旧图表快照",
        slides=[
            DeckSlide(
                id="slide-chart",
                layout_id="chart",
                objects=[
                    ChartObject(
                        id="chart-1",
                        categories=["Q1", "Q2"],
                        series={"收入": [10.0, 12.0]},
                    )
                ],
            )
        ],
    )
    legacy_payload = deck.model_dump(mode="json")
    legacy_payload["slides"][0]["objects"][0].pop("show_legend")
    legacy_payload["slides"][0]["objects"][0].pop("calculation_proofs")
    legacy_hash = hashlib.sha256(
        json.dumps(
            legacy_payload,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()

    assert deck_ir_version(deck) == legacy_hash
    assert deck_ir_version(
        deck.model_copy(
            update={
                "slides": [
                    deck.slides[0].model_copy(
                        update={
                            "objects": [
                                deck.slides[0].objects[0].model_copy(
                                    update={"show_legend": False}
                                )
                            ]
                        }
                    )
                ]
            }
        )
    ) != legacy_hash


def test_version_hash_preserves_pre_calculation_chart_ownership():
    deck = _deck()
    deck.slides[0].objects.append(ChartObject(
        id="owned-chart", categories=["Q2"], series={"利润": [37]}, unit="万元",
    ))
    payload = deck.model_dump(mode="json")
    chart = payload["slides"][0]["objects"][1]
    chart.pop("show_legend")
    chart.pop("calculation_proofs")
    legacy_hash = hashlib.sha256(json.dumps(
        payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    loaded = DeckIR.model_validate(payload)
    assert deck_ir_version(loaded) == legacy_hash
    # An explicitly persisted empty proof list still carries no calculation.
    assert deck_ir_version(DeckIR.model_validate_json(loaded.model_dump_json())) == legacy_hash
    chart["calculation_proofs"] = [{
        "terms": [
            {"source_ref": "data", "source_excerpt": "收入145万元、成本108万元",
             "value": "145", "unit": "万元", "coefficient": 1},
            {"source_ref": "data", "source_excerpt": "收入145万元、成本108万元",
             "value": "108", "unit": "万元", "coefficient": -1},
        ],
        "result": "37", "unit": "万元", "authorization_source_ref": "data",
        "authorization_excerpt": "利润按收入减成本计算",
    }]
    assert deck_ir_version(DeckIR.model_validate(payload)) != legacy_hash


def test_version_ledger_reads_do_not_materialize_absent_state(
    tmp_path: Path,
) -> None:
    store = PresentationVersionStore(tmp_path)

    assert store.current() is None
    assert store.history() == []
    assert not store.root.exists()


def test_writer_accepts_absent_native_defaults_without_rewriting_history(
    tmp_path: Path,
) -> None:
    deck = _deck()
    deck.slides[0].objects.extend([
        ShapeObject(id="legacy-shape", frame=ObjectFrame(x=1, y=2, w=3, h=1)),
        ChartObject(id="legacy-chart", categories=["Q1"], series={"收入": [10.0]}),
        TableObject(id="legacy-table", headers=["季度"], rows=[["Q1"]]),
    ])
    legacy_payload = deck.model_dump(mode="json")
    legacy_payload.pop("template_binding")
    for slide in legacy_payload["slides"]:
        slide.pop("semantic_layout")
        slide.pop("template_anchors")
        for obj in slide["objects"]:
            if obj["kind"] == "shape":
                obj.pop("runs")
                obj.pop("native_style")
            elif obj["kind"] == "chart":
                for field in ("show_legend", "calculation_proofs", "category_axis", "value_axis", "plot_layout"):
                    obj.pop(field)
            elif obj["kind"] == "table":
                obj.pop("cell_styles")
    legacy_hash = hashlib.sha256(json.dumps(
        legacy_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    loaded = DeckIR.model_validate(legacy_payload)
    canonical_hash = deck_ir_version(loaded)
    assert canonical_hash != legacy_hash
    output = tmp_path / "presentation.pptx"
    PresentationCompiler().compile(loaded, output)
    store = PresentationVersionStore(tmp_path)
    current = store.bootstrap(output, loaded, writer="legacy_python_pptx")
    legacy = current.model_copy(update={"deck_ir_version": legacy_hash})
    store.current_path.write_text(legacy.model_dump_json(indent=2), encoding="utf-8")
    store.history_path.write_text(legacy.model_dump_json() + "\n", encoding="utf-8")
    before = (store.current_path.read_bytes(), store.history_path.read_bytes())

    assert store.assert_writer_base(output, loaded) == legacy
    assert (store.current_path.read_bytes(), store.history_path.read_bytes()) == before
    assert deck_ir_version(loaded) == canonical_hash
    # Neither real semantics nor explicitly persisted neutral fields may be
    # erased by the compatibility reader in order to match an older hash.
    changes = [
        (("slides", 0, "semantic_layout"), "cover"),
        (("slides", 0, "semantic_layout"), None),
        (("slides", 0, "objects", 0, "text"), "真正改变的标题"),
        (("slides", 0, "objects", 1, "runs"), [{"text": "新标签"}]),
        (("slides", 0, "objects", 1, "runs"), []),
        (("slides", 0, "objects", 1, "native_style"), {"fill_none": True}),
        (("slides", 0, "objects", 1, "native_style"), None),
        (("slides", 0, "objects", 1, "frame", "x"), 2),
        (("slides", 0, "objects", 2, "series", "收入"), [11.0]),
        (("slides", 0, "objects", 2, "category_axis"), {"orientation": "maxMin"}),
        (("slides", 0, "objects", 2, "value_axis"), {"minimum": 0, "maximum": 20}),
        (("slides", 0, "objects", 2, "plot_layout"), {"x": 0.1}),
        (("slides", 0, "objects", 2, "value_axis"), None),
        (("slides", 0, "objects", 3, "cell_styles"), [[{"align": "right"}], [{}]]),
        (("slides", 0, "objects", 3, "cell_styles"), []),
        (("slides", 0, "objects", 3, "rows"), [["Q2"]]),
    ]
    for path, value in changes:
        changed_payload = json.loads(json.dumps(legacy_payload))
        target = changed_payload
        for key in path[:-1]:
            target = target[key]
        target[path[-1]] = value
        with pytest.raises(StalePresentationVersionError, match="DeckIR base hash"):
            store.assert_writer_base(output, DeckIR.model_validate(changed_payload))
        assert (store.current_path.read_bytes(), store.history_path.read_bytes()) == before


def test_writer_accepts_absent_native_caption_defaults_but_keeps_paint_semantics(
    tmp_path: Path,
) -> None:
    deck = _deck()
    deck.slides[0].objects.append(ShapeObject(
        id="native-caption", frame=ObjectFrame(x=1, y=2, w=3, h=1),
        native_style=NativeShapeStyle(fill_none=True, line_none=True),
    ))
    legacy_payload = deck.model_dump(mode="json")
    caption = legacy_payload["slides"][0]["objects"][1]["native_style"]
    for field in ("align", "valign", "margin_top", "margin_right", "margin_bottom", "margin_left"):
        caption.pop(field)
    legacy_hash = hashlib.sha256(json.dumps(
        legacy_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    loaded = DeckIR.model_validate(legacy_payload)
    output = tmp_path / "presentation.pptx"
    PresentationCompiler().compile(loaded, output)
    store = PresentationVersionStore(tmp_path)
    current = store.bootstrap(output, loaded, writer="legacy_python_pptx")
    legacy = current.model_copy(update={"deck_ir_version": legacy_hash})
    store.current_path.write_text(legacy.model_dump_json(indent=2), encoding="utf-8")
    store.history_path.write_text(legacy.model_dump_json() + "\n", encoding="utf-8")
    before = (store.current_path.read_bytes(), store.history_path.read_bytes())

    assert store.assert_writer_base(output, loaded) == legacy
    assert deck_ir_version(loaded) == current.deck_ir_version != legacy_hash
    for style_change in (
        {"fill_none": False}, {"line_none": False}, {"align": "right"},
        {"valign": "bottom"}, {"align": None},
        {"margin_top": 0, "margin_right": 0, "margin_bottom": 0, "margin_left": 0},
    ):
        changed_payload = json.loads(json.dumps(legacy_payload))
        changed_payload["slides"][0]["objects"][1]["native_style"].update(style_change)
        with pytest.raises(StalePresentationVersionError, match="DeckIR base hash"):
            store.assert_writer_base(output, DeckIR.model_validate(changed_payload))
    assert (store.current_path.read_bytes(), store.history_path.read_bytes()) == before


@pytest.mark.parametrize("omitted", [(True, False), (False, True), (True, True)])
def test_writer_accepts_legacy_absent_template_defaults_without_rewriting_history(
    tmp_path: Path, omitted: tuple[bool, bool],
) -> None:
    deck = _deck()
    output = tmp_path / "presentation.pptx"
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    current = store.bootstrap(output, deck, writer="legacy_python_pptx")
    legacy_payload = deck.model_dump(mode="json")
    if omitted[0]:
        legacy_payload.pop("template_binding")
    if omitted[1]:
        for slide in legacy_payload["slides"]:
            slide.pop("template_anchors")
    legacy_hash = hashlib.sha256(
        json.dumps(
            legacy_payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
        ).encode("utf-8")
    ).hexdigest()
    # A persisted pre-template ledger is the external input under test. Never
    # migrate that ledger in order to make the new reader accept it.
    legacy = current.model_copy(update={"deck_ir_version": legacy_hash})
    store.current_path.write_text(legacy.model_dump_json(indent=2), encoding="utf-8")
    store.history_path.write_text(legacy.model_dump_json() + "\n", encoding="utf-8")
    before = (store.current_path.read_bytes(), store.history_path.read_bytes())

    assert store.assert_writer_base(output, deck) == legacy
    assert (store.current_path.read_bytes(), store.history_path.read_bytes()) == before
    assert deck_ir_version(deck) == current.deck_ir_version
    with pytest.raises(StalePresentationVersionError, match="DeckIR base hash"):
        store.assert_writer_base(output, _deck("真正改变的语义"))
    bound = deck.model_copy(update={"template_binding": TemplatePackageBinding(
        source_path="template.pptx", source_sha256="a" * 64, source_slide_numbers=[1],
    )})
    anchored = deck.model_copy(update={"slides": [deck.slides[0].model_copy(update={
        "template_anchors": TemplateSlideAnchors(source_slide_number=1, uses=[
            TemplateAnchorUse(source_shape_id=2, purpose="title_frame", target_object_id="title-1"),
        ]),
    })]})
    for changed in (bound, anchored):
        with pytest.raises(StalePresentationVersionError, match="DeckIR base hash"):
            store.assert_writer_base(output, changed)
    assert (store.current_path.read_bytes(), store.history_path.read_bytes()) == before


def test_version_ledger_builds_an_explicit_parent_chain(tmp_path: Path) -> None:
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)

    first = store.bootstrap(output, deck, writer="legacy_python_pptx")
    base = store.assert_writer_base(output, deck)
    assert base == first

    changed = _deck("新的可验证标题")
    PresentationCompiler().compile(changed, output)
    second = store.record(output, changed, writer="lantai_ai_edit", parent=base)

    assert second.parent_version_id == first.version_id
    assert second.deck_ir_version == deck_ir_version(changed)
    assert store.current() == second
    assert len(store.history_path.read_text(encoding="utf-8").splitlines()) == 2


@pytest.mark.parametrize("damage", [
    None, "legacy_defaults", "png_hash", "missing_page", "additional_page",
    "unknown_preview",
    "deck_semantics", "report_output", "failed_quality", "stale_base",
    "non_new_generation", "unreconciled_target",
])
def test_rollback_restores_nineteen_evidence_bound_pages_in_numeric_order(
    tmp_path: Path, monkeypatch, damage: str | None,
) -> None:
    def render_fixture(source, output_dir, *, timeout=120):
        # Only the external renderer is replaced. Production compilation,
        # evidence binding, quality verification and rollback all run normally.
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        pdf = destination / "presentation.pdf"
        pdf.write_bytes(b"synthetic renderer fixture")
        images = []
        for index in range(1, len(Presentation(str(source)).slides) + 1):
            path = destination / f"slide-{index}.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste("black", (0, 0, 280 + index, 360))
            image.save(path)
            images.append(path)
        return RenderResult(pdf, tuple(images), "version-fixture", "1.0")

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render_fixture)
    deck = DeckIR(id="nineteen-page-version", title="原版本", slides=[
        DeckSlide(id=f"page-{index}", layout_id="cover", objects=[
            TextObject(id=f"title-{index}", role="title", text=f"原版本第 {index} 页", slot="title"),
        ]) for index in range(1, 20)
    ])
    output = tmp_path / "presentation.pptx"
    deck_path = tmp_path / "deck-ir.json"
    quality_path = tmp_path / "quality-report.json"
    render_root = tmp_path / "renders"
    compiler = PresentationCompiler()
    service = PresentationService(tmp_path / "jobs")
    compiler.compile(deck, output)
    report, images = service.verify_artifact(
        deck, output, render_root / "round-0", render_set_id="round-0",
    )
    assert is_verified_delivery(deck, report, output, images), [
        (issue.code, issue.slide_id) for issue in report.issues
    ]
    deck_path.write_text(deck.model_dump_json(indent=2), encoding="utf-8")
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    original = {path: path.read_bytes() for path in (output, deck_path, quality_path)}
    store = PresentationVersionStore(tmp_path)
    target = store.bootstrap(
        output, deck, writer="legacy_python_pptx", quality_report_path=quality_path,
        render_dir=render_root / "round-0",
    )
    changed = deck.model_copy(deep=True)
    changed.slides[0].objects[0].text = "更新后的第一页"
    compiler.compile(changed, output)
    changed_report, changed_images = service.verify_artifact(
        changed, output, render_root / "round-1", render_set_id="round-1",
    )
    assert is_verified_delivery(changed, changed_report, output, changed_images)
    deck_path.write_text(changed.model_dump_json(indent=2), encoding="utf-8")
    quality_path.write_text(changed_report.model_dump_json(indent=2), encoding="utf-8")
    edited = store.record(
        output, changed, writer="lantai_ai_edit", parent=target, generation=1,
        quality_report_path=quality_path, render_dir=render_root / "round-1",
    )
    bundle = store.root / "artifacts" / target.version_id
    if damage == "legacy_defaults":
        legacy_payload = json.loads((bundle / "deck-ir.json").read_text(encoding="utf-8"))
        legacy_payload.pop("template_binding")
        for slide in legacy_payload["slides"]:
            slide.pop("template_anchors")
            slide.pop("semantic_layout")
        legacy_bytes = json.dumps(legacy_payload, ensure_ascii=False, sort_keys=True,
                                  separators=(",", ":")).encode("utf-8")
        (bundle / "deck-ir.json").write_bytes(legacy_bytes)
        target = target.model_copy(update={"deck_ir_version": hashlib.sha256(legacy_bytes).hexdigest()})
        original[deck_path] = legacy_bytes
    elif damage == "png_hash":
        image = bundle / "renders/slide-10.png"
        image.write_bytes(image.read_bytes() + b"changed-snapshot")
    elif damage == "missing_page":
        (bundle / "renders/slide-10.png").rename(bundle / "renders/slide-10.missing")
    elif damage == "additional_page":
        (bundle / "renders/slide-20.png").write_bytes((bundle / "renders/slide-19.png").read_bytes())
    elif damage == "unknown_preview":
        (bundle / "renders/unrecognized.png").write_bytes((bundle / "renders/slide-19.png").read_bytes())
    elif damage == "deck_semantics":
        changed_payload = json.loads((bundle / "deck-ir.json").read_text(encoding="utf-8"))
        changed_payload["slides"][0]["objects"][0]["text"] = "快照被更改"
        (bundle / "deck-ir.json").write_text(json.dumps(changed_payload), encoding="utf-8")
    elif damage in {"report_output", "failed_quality"}:
        altered_report = report.model_copy(update=(
            {"output_sha256": "0" * 64} if damage == "report_output"
            else {"passed": False, "delivery_status": "failed"}
        ))
        (bundle / "quality-report.json").write_text(altered_report.model_dump_json(), encoding="utf-8")
    elif damage == "unreconciled_target":
        target = target.model_copy(update={"reconciled": False})
    if damage in {"legacy_defaults", "unreconciled_target"}:
        store.history_path.write_text(target.model_dump_json() + "\n" + edited.model_dump_json() + "\n",
                                      encoding="utf-8")
    snapshots = {path: path.read_bytes() for path in bundle.rglob("*") if path.is_file()}
    current_files = {path: path.read_bytes() for path in (
        output, deck_path, quality_path, store.current_path, store.history_path,
    )}
    rollback_args = dict(
        pptx_path=output, deck_ir_path=deck_path, quality_report_path=quality_path,
        render_root=render_root, generation=(1 if damage == "non_new_generation" else 2),
        base_pptx_sha256=(target.pptx_sha256 if damage == "stale_base" else edited.pptx_sha256),
    )
    if damage not in {None, "legacy_defaults"}:
        with pytest.raises(StalePresentationVersionError):
            store.rollback_to(target.version_id, **rollback_args)
        assert {path: path.read_bytes() for path in current_files} == current_files
        assert {path: path.read_bytes() for path in snapshots} == snapshots
        assert store.current() == edited and store.history() == [target, edited]
        assert not list(store.root.glob(".rollback-*"))
        return

    restored, restored_deck, restored_report = store.rollback_to(target.version_id, **rollback_args)

    assert {path: path.read_bytes() for path in original} == original
    assert {path: path.read_bytes() for path in snapshots} == snapshots
    assert restored_deck == deck and restored_report == report
    assert restored.writer == "version_rollback"
    assert restored.parent_version_id == edited.version_id and restored.generation == 2
    assert store.history() == [target, edited, restored]
    restored_images = slide_image_paths(render_root / report.render_set_id)
    assert [path.name for path in restored_images] == [f"slide-{index}.png" for index in range(1, 20)]
    assert is_verified_delivery(restored_deck, restored_report, output, restored_images)


def test_external_or_human_changes_fail_closed_until_reimport(tmp_path: Path) -> None:
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    store.bootstrap(output, deck, writer="legacy_python_pptx")

    output.write_bytes(output.read_bytes() + b"external-change")
    with pytest.raises(StalePresentationVersionError, match="outside"):
        store.assert_writer_base(output, deck)

    human = store.register_external_save(output, writer="genoffice_human")
    assert human.reconciled is False
    with pytest.raises(StalePresentationVersionError, match="unreconciled"):
        store.assert_writer_base(output, deck)

    reconciled_deck = _deck("从人工文件重建的语义版本")
    reconciled = store.record(
        output,
        reconciled_deck,
        writer="external_reconcile",
        parent=human,
    )
    assert reconciled.reconciled is True
    assert store.assert_writer_base(output, reconciled_deck) == reconciled


def test_deck_ir_cannot_overwrite_a_newer_pptx_base(tmp_path: Path) -> None:
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    store.bootstrap(output, deck, writer="legacy_python_pptx")

    with pytest.raises(StalePresentationVersionError, match="DeckIR base hash"):
        store.assert_writer_base(output, _deck("stale semantic checkpoint"))


def test_record_rejects_bytes_that_no_longer_match_verified_result(
    tmp_path: Path,
) -> None:
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    base = store.bootstrap(output, deck, writer="legacy_python_pptx")

    candidate = _deck("AI 已验证的候选标题")
    PresentationCompiler().compile(candidate, output)
    candidate_sha256 = file_sha256(output)
    PresentationCompiler().compile(_deck("用户随后保存的新标题"), output)

    with pytest.raises(
        StalePresentationVersionError,
        match="changed before version record",
    ):
        store.record(
            output,
            candidate,
            writer="lantai_ai_edit",
            parent=base,
            expected_pptx_sha256=candidate_sha256,
        )

    assert store.current() == base
    assert store.history() == [base]


def test_record_current_replace_failure_leaves_no_partial_history(tmp_path, monkeypatch):
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    base = store.bootstrap(output, deck, writer="legacy_python_pptx")
    before_current = store.current_path.read_bytes()
    before_history = store.history_path.read_bytes()
    candidate = _deck("未发表候选")
    PresentationCompiler().compile(candidate, output)
    original_replace = versioning_module.os.replace
    failed_once = False

    def fail_current(source, destination):
        nonlocal failed_once
        if Path(destination) == store.current_path and not failed_once:
            failed_once = True
            raise OSError("synthetic current replace error")
        return original_replace(source, destination)

    monkeypatch.setattr(versioning_module.os, "replace", fail_current)
    with pytest.raises(OSError, match="current replace"):
        store.record(output, candidate, writer="lantai_ai_edit", parent=base)
    assert store.current_path.read_bytes() == before_current
    assert store.history_path.read_bytes() == before_history


def test_unpublished_record_rollback_cannot_erase_a_newer_external_version(tmp_path):
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    base = store.bootstrap(output, deck, writer="legacy_python_pptx")
    checkpoint = store.publication_checkpoint(parent=base)
    candidate = _deck("AI 候选")
    PresentationCompiler().compile(candidate, output)
    candidate_sha = file_sha256(output)
    store.record(output, candidate, writer="lantai_ai_edit", parent=base, generation=1)
    PresentationCompiler().compile(_deck("更新的用户文件"), output)
    human = store.register_external_save(output, writer="genoffice_human")
    before_current = store.current_path.read_bytes()
    before_history = store.history_path.read_bytes()

    assert not store.rollback_unpublished_record(
        checkpoint, writer="lantai_ai_edit", generation=1,
        base_pptx_sha256=base.pptx_sha256, expected_pptx_sha256=candidate_sha,
    )
    assert store.current() == human
    assert store.current_path.read_bytes() == before_current
    assert store.history_path.read_bytes() == before_history


def test_record_aborts_if_source_changes_while_snapshotting(
    tmp_path: Path,
    monkeypatch,
) -> None:
    output = tmp_path / "presentation.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, output)
    store = PresentationVersionStore(tmp_path)
    base = store.bootstrap(output, deck, writer="legacy_python_pptx")
    artifacts = store.root / "artifacts"
    original_bundles = set(artifacts.iterdir())

    candidate = _deck("快照中的 AI 候选标题")
    PresentationCompiler().compile(candidate, output)
    candidate_sha256 = file_sha256(output)
    original_copy2 = versioning_module.shutil.copy2

    def copy_then_change_source(source, destination, *args, **kwargs):
        result = original_copy2(source, destination, *args, **kwargs)
        if (
            Path(source).resolve() == output.resolve()
            and Path(destination).name == "presentation.pptx"
        ):
            output.write_bytes(output.read_bytes() + b"newer-editor-save")
        return result

    monkeypatch.setattr(versioning_module.shutil, "copy2", copy_then_change_source)

    with pytest.raises(
        StalePresentationVersionError,
        match="changed before version snapshot was committed",
    ):
        store.record(
            output,
            candidate,
            writer="lantai_ai_edit",
            parent=base,
            expected_pptx_sha256=candidate_sha256,
        )

    assert file_sha256(output) != candidate_sha256
    assert store.current() == base
    assert store.history() == [base]
    assert set(artifacts.iterdir()) == original_bundles
