"""Visual review must receive the same shape captions as the native writer."""

from __future__ import annotations

import hashlib
import json
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from lantai.presentation.intelligence import UtilityPresentationVisualCritic
from lantai.presentation.models import (
    ConstraintLevel, DeckIR, DeckSlide, FactRecord, ObjectFrame,
    PresentationBrief, PresentationSourceSection, RichTextRun, ShapeObject,
    TemplateMode, TextObject, VisualCriticSlideInput, VisualIssue,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from tests.unit.presentation.visual_review_fixtures import synthetic_review_response


def _shape(*, shape_type="round_rect", text="Current synthetic pending decision", runs=None):
    return ShapeObject(
        id="page-a::caption", shape_type=shape_type, text=text,
        runs=runs if runs is not None else [
            RichTextRun(text="Old synthetic approved decision", bold=True,
                        font_size=24, color="123456"),
        ],
        frame=ObjectFrame(x=1, y=2, w=9, h=1.4), font_size=24,
        fill_color="FFFFFF", line_color="123456",
    )


def _slide(tmp_path, obj):
    rendered = tmp_path / "synthetic.png"
    Image.new("RGB", (320, 180), "white").save(rendered)
    brief = PresentationBrief(
        objective="Keep the decision pending", audience="Synthetic reviewers",
        narrative="Separate observations from decisions",
        required_facts=["The decision remains pending."],
        source_sections=[PresentationSourceSection(
            title="Complete synthetic original", source_ref="synthetic-source",
            body="The owner has a preference, but the decision remains pending.",
        )],
    )
    return VisualCriticSlideInput(
        slide_id="page-a", slide_index=1, rendered_path=str(rendered),
        rendered_sha256=hashlib.sha256(rendered.read_bytes()).hexdigest(),
        original_user_request="Retain the pending state and every source qualification.",
        source_brief=brief, key_message="The decision remains pending.",
        factual_constraints=["The decision remains pending."],
        page_factual_constraints=["The decision remains pending."],
        source_facts=[FactRecord(
            fact_id="pending-state", text="The decision remains pending.",
            source_refs=["synthetic-source"], immutable=True,
        )],
        object_ids=[obj.id], slide_objects=[obj], template_mode=TemplateMode.STRICT,
        constraint_levels_by_object={obj.id: ConstraintLevel.HARD},
    )


def _review(monkeypatch, slide, scope, severity="error"):
    captured = {}
    finding = VisualIssue(
        slide_id=slide.slide_id, object_ids=slide.object_ids, severity=severity,
        issue_type="content_accuracy", description="A synthetic finding must survive",
        proposed_action="Correct the candidate using the complete original evidence",
    )

    async def transport(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(
            text=json.dumps(synthetic_review_response(kwargs, [finding.model_dump(mode="json")])),
            model_label="synthetic-vision",
        )

    # Replace only external transport: real context serialization, parser, schema and
    # issue-reference validation still run at the public critic boundary.
    monkeypatch.setattr("lantai.app.structured_utility.run_utility_content_task", transport)
    monkeypatch.setattr(
        "lantai.presentation.intelligence._utility_route_cache_identity", lambda _: "test-route",
    )
    critic = UtilityPresentationVisualCritic("test-agent")
    result = critic.review_slide(slide) if scope == "slide" else critic.review_deck([slide])
    context = json.loads(captured["user_content"][1]["text"])
    if scope == "deck":
        context = context["ordered_slides"][0]
    assert result.issues == [finding], "Context projection must never filter or downgrade findings"
    return context


@pytest.mark.parametrize("scope", ["slide", "deck"])
@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron"])
def test_public_critic_caption_matches_actual_native_export_not_immutable_old_runs(
    tmp_path, monkeypatch, scope, shape_type,
):
    shape = _shape(shape_type=shape_type)
    slide = _slide(tmp_path, shape)
    before = slide.model_dump(mode="json")
    deck = DeckIR(title="Synthetic caption authority", slides=[DeckSlide(
        id=slide.slide_id, layout_id="freeform", objects=[shape],
        speaker_notes="Preserved synthetic notes",
    )])
    before_deck = deck.model_dump(mode="json")
    output = tmp_path / "synthetic-native.pptx"
    NativePptxCompiler().compile(deck, output)
    native = next(item for item in Presentation(str(output)).slides[0].shapes
                  if item.name == shape.id)
    assert native.text == shape.text
    context = _review(monkeypatch, slide, scope)
    caption = context["slide_objects"][0]
    assert caption["text"] == native.text
    assert "".join(run["text"] for run in caption["runs"]) == native.text
    assert caption["runs"] == [{
        "text": shape.text, "bold": True, "font_size": 24.0, "color": "123456",
    }]
    assert caption["frame"] == shape.frame.model_dump(mode="json")
    assert caption["fill_color"] == shape.fill_color and caption["line_color"] == shape.line_color
    assert "Old synthetic approved decision" not in json.dumps(caption)
    assert slide.model_dump(mode="json") == before
    assert deck.model_dump(mode="json") == before_deck


@pytest.mark.parametrize("scope", ["slide", "deck"])
@pytest.mark.parametrize("severity", ["info", "warning", "error"])
def test_projection_keeps_primary_sources_fact_catalog_and_hard_findings(
    tmp_path, monkeypatch, scope, severity,
):
    slide = _slide(tmp_path, _shape())
    before = slide.model_dump(mode="json")
    context = _review(monkeypatch, slide, scope, severity)
    if scope == "slide":
        assert context["source_brief"] == before["source_brief"]
        assert context["source_facts"] == before["source_facts"]
        assert context["original_user_request"] == before["original_user_request"]
    assert context["factual_constraints"] == before["factual_constraints"]
    assert context["page_factual_constraints"] == before["page_factual_constraints"]
    assert context["constraint_levels_by_object"] == {slide.object_ids[0]: "hard"}
    assert slide.model_dump(mode="json") == before


@pytest.mark.parametrize("text,shape_type", [("", "rect"), ("", "line"), ("Ignored", "line")])
def test_unrendered_shape_labels_are_not_sent_as_visible_captions(tmp_path, text, shape_type):
    shape = _shape(text=text, shape_type=shape_type)
    slide = _slide(tmp_path, shape)
    before = slide.model_dump(mode="json")
    context = UtilityPresentationVisualCritic._slide_context(slide)
    caption = context["slide_objects"][0]
    assert caption["text"] == "" and caption["runs"] == []
    assert slide.model_dump(mode="json") == before


def test_matching_rich_runs_remain_lossless_including_paragraph_breaks(tmp_path):
    runs = [RichTextRun(text="One", bold=True),
            RichTextRun(text="\n", paragraph_break=True),
            RichTextRun(text="Two", italic=True)]
    shape = _shape(text="One\nTwo", runs=runs)
    slide = _slide(tmp_path, shape)
    caption = UtilityPresentationVisualCritic._slide_context(slide)["slide_objects"][0]
    assert caption == {"kind": shape.kind, **shape.model_dump(
        mode="json", exclude_none=True, exclude_defaults=True,
    )}


@pytest.mark.parametrize("runs", [[], [RichTextRun(text="\n", paragraph_break=True)]])
def test_absent_caption_style_falls_back_to_current_text_without_old_runs(tmp_path, runs):
    slide = _slide(tmp_path, _shape(runs=runs))
    caption = UtilityPresentationVisualCritic._slide_context(slide)["slide_objects"][0]
    assert caption["text"] == slide.slide_objects[0].text
    assert caption.get("runs", []) == []


def test_replacement_uses_first_nonparagraph_style_and_clears_break(tmp_path):
    shape = _shape(runs=[
        RichTextRun(text="\n", paragraph_break=True, color="FFFFFF"),
        RichTextRun(text="Old", color="123456", italic=True),
        RichTextRun(text=" text", bold=True),
    ])
    slide = _slide(tmp_path, shape)
    caption = UtilityPresentationVisualCritic._slide_context(slide)["slide_objects"][0]
    assert caption["runs"] == [{"text": shape.text, "color": "123456", "italic": True}]


def test_text_objects_keep_their_different_rich_text_authority_unchanged(tmp_path):
    obj = TextObject(id="page-a::caption", role="body", text="Plain fallback",
                     runs=[RichTextRun(text="Visible rich text", bold=True)])
    slide = _slide(tmp_path, obj)
    caption = UtilityPresentationVisualCritic._slide_context(slide)["slide_objects"][0]
    assert caption == {"kind": obj.kind, **obj.model_dump(
        mode="json", exclude_none=True, exclude_defaults=True,
    )}


@pytest.mark.parametrize("scope,version", [("slide", "v6"), ("deck", "v7")])
def test_native_caption_context_invalidates_both_previous_review_cache_identities(
    monkeypatch, scope, version,
):
    monkeypatch.setattr(
        "lantai.presentation.intelligence._utility_route_cache_identity", lambda _: "test-route",
    )
    critic = UtilityPresentationVisualCritic("test-agent")
    old = (f"utility-vision/test-agent/page-and-deck-{version}/source-fact-catalog-v1/"
           "primary-sources-first-v1/image-capable-foreground-v1/test-route")
    current = critic.slide_cache_identity() if scope == "slide" else critic.deck_cache_identity()
    assert current != old
    assert "/primary-sources-first-v1/native-caption-context-v1/" in current
