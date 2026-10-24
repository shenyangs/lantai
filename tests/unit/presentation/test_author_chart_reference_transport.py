"""Model-selected exact native chart transport through the public author."""

from __future__ import annotations

import json
import time
from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from pptx import Presentation

from lantai.app.structured_utility import StructuredTaskValidationError
from lantai.presentation.intelligence import UtilityPresentationPageAuthor
from lantai.presentation.models import (
    ChartAxisStyle, ChartObject, ChartPlotLayout, InputMaterial, JobStatus, ObjectFrame,
    PresentationJobRequest, SourceCalculationProof,
)
from lantai.presentation.page_authoring import (
    PresentationPageAuthorContentError, grounded_source_text, validate_authored_page,
)
from lantai.presentation.planner import build_brief, build_deck_ir, build_deterministic_fixture_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.runtime import PresentationJobRuntime


def chart_case():
    request = PresentationJobRequest(
        conversation_id="synthetic-chart-reference-transport",
        user_request="制作一页收入判断，保留原始图表数据",
        slide_count=1,
        materials=[InputMaterial(name="synthetic.md", content="甲为42万元；乙为38万元。")],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    chart = ChartObject(
        id="planned-chart", categories=["甲", "乙"], series={"收入": [42, 38]},
        show_legend=False, unit="万元", number_format="0.0",
        derivation="原始记录", source_refs=[brief.source_sections[0].source_ref],
        frame=ObjectFrame(x=5, y=2, w=7, h=4),
    )
    planned = plan.slides[0].model_copy(update={
        "title": "先核验证据，再判断收入", "charts": [chart],
        "intent": plan.slides[0].intent.model_copy(update={
            "required_facts": ["甲为42万元；乙为38万元。"],
            "immutable_facts": ["甲为42万元；乙为38万元。"],
        }),
    })
    plan = plan.model_copy(update={"slides": [planned]})
    deck = build_deck_ir(request, plan)
    scene = {
        "objects": [
            {"kind": "text", "id": "title", "role": "title",
             "text_refs": ["storyboard.title"], "font_size": 36,
             "frame": {"x": 0.8, "y": 0.5, "w": 11.5, "h": 1}},
            {"kind": "text", "id": "facts", "role": "body",
             "text_refs": ["required_facts[0]", "immutable_facts[0]"], "font_size": 24,
             "frame": {"x": 0.8, "y": 2, "w": 3.5, "h": 3}},
            {"kind": "chart", "id": "author-chosen-chart", "slot": "evidence",
             "chart_ref": "storyboard.charts[0]",
             "frame": {"x": 4.8, "y": 2.2, "w": 7.2, "h": 4}},
        ],
        "speaker_notes": "", "object_groups": [], "relationships": [],
    }
    return request, brief, plan, deck, scene


def public_author(monkeypatch, case, calls):
    request, brief, plan, deck, scene = case

    async def transport(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="synthetic/test")

    # Replace only the external model transport, not schema parsing or guards.
    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    return UtilityPresentationPageAuthor("synthetic-agent").author(
        request, brief, plan, plan.slides[0], deck.slides[0],
    )


def test_public_author_preserves_selected_chart_data_and_author_geometry(monkeypatch):
    case = chart_case()
    request, brief, plan, deck, scene = case
    calls = []
    authored = public_author(monkeypatch, case, calls)
    chart = next(obj for obj in authored.objects if isinstance(obj, ChartObject))
    assert type(chart) is ChartObject
    immutable = {"id", "slot", "frame"}
    assert chart.model_dump(mode="json", exclude=immutable) == plan.slides[0].charts[0].model_dump(
        mode="json", exclude=immutable,
    )
    decision = scene["objects"][-1]
    assert chart.id == decision["id"] and chart.slot == decision["slot"]
    assert chart.frame.model_dump(mode="json") == decision["frame"]
    assert "chart_ref" not in chart.model_dump(mode="json")
    checked = validate_authored_page(
        authored, plan.slides[0], deck.slides[0],
        source_text=grounded_source_text(request, brief), source_brief=brief,
    )
    assert checked.scene_hash
    transport = json.loads(calls[0]["user_prompt"])["chart_reference_transport"]
    assert transport["references"] == [{"ref": "storyboard.charts[0]"}]
    assert transport["missing_charts_are_not_filled"] is True
    assert transport["same_page_only"] is True


def advertised_validator(monkeypatch):
    case = chart_case()
    calls = []
    public_author(monkeypatch, case, calls)
    schema = calls[0]["structured_output_schema"]
    assert schema == json.loads(calls[0]["system_prompt"].split("JSON Schema:\n", 1)[1])
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    assert validator.is_valid(case[-1])
    return validator


@pytest.mark.parametrize("field", sorted(set(ChartObject.model_fields) - {"kind", "id", "slot", "frame"}))
@pytest.mark.parametrize("use_null", [False, True])
def test_reference_wire_and_parser_reject_every_explicit_contract_override(
    monkeypatch, field, use_null,
):
    validator = advertised_validator(monkeypatch)
    case = chart_case()
    case[-1]["objects"][-1][field] = (
        None if use_null else case[2].slides[0].charts[0].model_dump(mode="json")[field]
    )
    assert not validator.is_valid(case[-1])
    with pytest.raises(StructuredTaskValidationError):
        public_author(monkeypatch, case, [])


@pytest.mark.parametrize("reference", [
    "storyboard.charts[1]", "surrounding_storyboard[1].charts[0]",
    "brief.charts[0]", "/private/SYNTHETIC_CHART_REFERENCE", "storyboard.charts[99999999]",
])
def test_public_author_rejects_out_of_page_chart_reference_without_echo(monkeypatch, reference):
    case = chart_case()
    case[-1]["objects"][-1]["chart_ref"] = reference
    with pytest.raises(PresentationPageAuthorContentError) as failure:
        public_author(monkeypatch, case, [])
    assert failure.value.violation_code == "invalid_chart_reference"
    assert reference not in str(failure.value)


@pytest.mark.parametrize("mutation", ["missing", "duplicate"])
def test_existing_exact_chart_gate_still_rejects_missing_or_duplicate_selection(monkeypatch, mutation):
    case = chart_case()
    if mutation == "missing":
        case[-1]["objects"].pop()
    else:
        duplicate = deepcopy(case[-1]["objects"][-1])
        duplicate["id"] = "distinct-duplicate-selection"
        case[-1]["objects"].append(duplicate)
    authored = public_author(monkeypatch, case, [])
    assert sum(isinstance(obj, ChartObject) for obj in authored.objects) == (0 if mutation == "missing" else 2)
    with pytest.raises(PresentationPageAuthorContentError) as failure:
        validate_authored_page(
            authored, case[2].slides[0], case[3].slides[0],
            source_text=grounded_source_text(case[0], case[1]), source_brief=case[1],
        )
    assert failure.value.violation_code == "planned_chart_data_changed"


def test_reference_without_a_chart_on_the_page_cannot_fill_from_another_page(monkeypatch):
    case = chart_case()
    planned = case[2].slides[0].model_copy(update={"charts": []})
    plan = case[2].model_copy(update={"slides": [planned]})
    case = (case[0], case[1], plan, case[3], case[-1])
    calls = []
    with pytest.raises(PresentationPageAuthorContentError):
        public_author(monkeypatch, case, calls)
    assert json.loads(calls[0]["user_prompt"])["chart_reference_transport"]["references"] == []


@pytest.mark.parametrize("mutation", [
    "neither", "empty_reference", "missing_categories", "missing_series",
    "empty_categories", "empty_series", "short_series", "long_series",
])
def test_reference_transport_does_not_weaken_literal_native_chart_validation(monkeypatch, mutation):
    case = chart_case()
    literal = case[2].slides[0].charts[0].model_dump(mode="json")
    literal.update({"id": "author-chosen-chart", "frame": case[-1]["objects"][-1]["frame"]})
    if mutation == "neither":
        literal = {key: value for key, value in literal.items() if key in {"kind", "id", "slot", "frame"}}
    elif mutation == "empty_reference":
        literal = deepcopy(case[-1]["objects"][-1])
        literal["chart_ref"] = ""
    elif mutation == "missing_categories":
        literal.pop("categories")
    elif mutation == "missing_series":
        literal.pop("series")
    elif mutation == "empty_categories":
        literal["categories"] = []
    elif mutation == "empty_series":
        literal["series"] = {}
    else:
        literal["series"] = {"收入": [42] if mutation == "short_series" else [42, 38, 20]}
    case[-1]["objects"][-1] = literal
    with pytest.raises(StructuredTaskValidationError):
        public_author(monkeypatch, case, [])


def test_selected_chart_preserves_source_authorized_calculation_proof(monkeypatch):
    case = chart_case()
    source = "甲为42万元；乙为38万元。\n差额按甲减乙计算。"
    request = case[0].model_copy(update={"materials": [InputMaterial(name="synthetic.md", content=source)]})
    brief = build_brief(request)
    source_ref = brief.source_sections[0].source_ref
    proof = SourceCalculationProof.model_validate({
        "terms": [
            {"source_ref": source_ref, "source_excerpt": "甲为42万元；乙为38万元。",
             "value": "42", "unit": "万元", "coefficient": 1},
            {"source_ref": source_ref, "source_excerpt": "甲为42万元；乙为38万元。",
             "value": "38", "unit": "万元", "coefficient": -1},
        ],
        "result": "4", "unit": "万元", "authorization_source_ref": source_ref,
        "authorization_excerpt": "差额按甲减乙计算。",
    })
    chart = case[2].slides[0].charts[0].model_copy(update={
        "categories": ["差额"], "series": {"差额": [4]}, "source_refs": [source_ref],
        "derivation": "42万元−38万元＝4万元", "calculation_proofs": [proof],
    })
    planned = case[2].slides[0].model_copy(update={
        "charts": [chart],
        "intent": case[2].slides[0].intent.model_copy(update={"source_refs": [source_ref]}),
    })
    plan = case[2].model_copy(update={"slides": [planned]})
    authored = public_author(monkeypatch, (request, brief, plan, case[3], case[-1]), [])
    selected = next(obj for obj in authored.objects if isinstance(obj, ChartObject))
    assert selected.calculation_proofs == [proof]
    assert selected.source_refs == [source_ref] and selected.series == {"差额": [4]}
    assert validate_authored_page(
        authored, planned, case[3].slides[0], source_text=source, source_brief=brief,
    ).scene_hash
    stripped = authored.model_copy(update={"objects": [
        obj.model_copy(update={"calculation_proofs": []}) if isinstance(obj, ChartObject) else obj
        for obj in authored.objects
    ]})
    with pytest.raises(PresentationPageAuthorContentError) as failure:
        validate_authored_page(stripped, planned, case[3].slides[0], source_text=source, source_brief=brief)
    assert failure.value.violation_code == "planned_chart_data_changed"


def test_reference_native_chart_keeps_axes_data_and_embedded_editable_workbook(monkeypatch, tmp_path):
    case = chart_case()
    chart = case[2].slides[0].charts[0].model_copy(update={
        "chart_type": "column",
        "category_axis": ChartAxisStyle(orientation="maxMin", label_position="high"),
        "value_axis": ChartAxisStyle(minimum=0, maximum=80, major_unit=20, number_format="0.0"),
        "plot_layout": ChartPlotLayout(
            layoutTarget="inner", xMode="factor", yMode="factor", wMode="factor", hMode="factor",
            x=0.1, y=0.1, w=0.8, h=0.8,
        ),
    })
    plan = case[2].model_copy(update={"slides": [case[2].slides[0].model_copy(update={"charts": [chart]})]})
    case = (case[0], case[1], plan, case[3], case[-1])
    authored = public_author(monkeypatch, case, [])
    accepted = validate_authored_page(
        authored, plan.slides[0], case[3].slides[0],
        source_text=grounded_source_text(case[0], case[1]), source_brief=case[1],
    )
    authored_chart = next(obj for obj in accepted.objects if isinstance(obj, ChartObject))
    assert authored_chart.model_dump(exclude={"id", "slot", "frame"}) == chart.model_dump(
        exclude={"id", "slot", "frame"},
    )
    deck = case[3].model_copy(update={"slides": [accepted]})
    reference_output = tmp_path / "selected-native-chart.pptx"
    NativePptxCompiler().compile(deck, reference_output)
    literal_scene = deepcopy(case[-1])
    literal_scene["objects"][-1] = authored_chart.model_dump(mode="json")
    literal = public_author(monkeypatch, (*case[:-1], literal_scene), [])
    literal_output = tmp_path / "literal-native-chart.pptx"
    NativePptxCompiler().compile(deck.model_copy(update={"slides": [literal]}), literal_output)
    shape = next(shape for shape in Presentation(reference_output).slides[0].shapes if shape.has_chart)
    assert shape.name == authored_chart.id
    assert [category.label for category in shape.chart.plots[0].categories] == ["甲", "乙"]
    assert list(shape.chart.series[0].values) == [42.0, 38.0]
    assert shape.chart.has_legend is False
    assert shape.chart.value_axis.minimum_scale == 0 and shape.chart.value_axis.maximum_scale == 80
    with ZipFile(reference_output) as reference_zip, ZipFile(literal_output) as literal_zip:
        assert reference_zip.read("ppt/charts/chart1.xml") == literal_zip.read("ppt/charts/chart1.xml")
        workbooks = [name for name in reference_zip.namelist() if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]
        assert len(workbooks) == 1
        workbook = load_workbook(BytesIO(reference_zip.read(workbooks[0])), data_only=True)
        # The established native writer uses an empty category-header cell,
        # not a missing cell. This transport must preserve that representation.
        assert list(workbook.active.values) == [("", "收入"), ("甲", 42), ("乙", 38)]
        literal_workbook = load_workbook(BytesIO(literal_zip.read(workbooks[0])), data_only=True)
        assert list(workbook.active.values) == list(literal_workbook.active.values)


def settled_failure(runtime, job_id, conversation_id):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        record = runtime.get(job_id, conversation_id)
        if record.status == JobStatus.FAILED and record.active_lease_id is None:
            assert not record.resource_settlement_pending
            return record
        time.sleep(0.01)
    raise AssertionError("synthetic public chart job did not settle")


def test_public_runtime_retries_chart_selection_and_compiles_native_output(tmp_path, monkeypatch):
    case = chart_case()
    request = case[0].model_copy(update={
        "visual_verification_required": False, "vision_critique_required": False,
    })
    calls = []

    async def transport(**kwargs):
        calls.append(json.loads(kwargs["user_prompt"]))
        scene = deepcopy(case[-1])
        if len(calls) == 1:
            scene["objects"][-1]["chart_ref"] = "storyboard.charts[1]"
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    runtime = PresentationJobRuntime(tmp_path, page_author=UtilityPresentationPageAuthor("default"))
    try:
        record = runtime.service.create_job(request)
        runtime.service.store.write_checkpoint(record.job_id, "brief.json", case[1].model_dump_json())
        runtime.service.store.write_checkpoint(record.job_id, "deck-plan.json", case[2].model_dump_json())
        runtime.pause(record.job_id, request.conversation_id)
        runtime.resume(record.job_id, request.conversation_id)
        deadline = time.monotonic() + 15
        while time.monotonic() < deadline:
            saved = runtime.get(record.job_id, request.conversation_id)
            if saved.status in {JobStatus.COMPLETED, JobStatus.FAILED} and saved.active_lease_id is None:
                break
            time.sleep(0.01)
        assert saved.status == JobStatus.COMPLETED
        assert not saved.resource_settlement_pending and len(calls) == 2
        assert saved.cost_used_units == 10 == sum(metric.cost_units for metric in saved.stage_metrics)
        assert saved.deadline_at == record.deadline_at and saved.attempt_count == 1
        assert "chart_reference_transport.references" in calls[1]["retry_directive"]
        authored_metrics = [metric for metric in saved.stage_metrics if metric.operation == "author" and metric.slide_id]
        assert [metric.outcome for metric in authored_metrics] == ["failed", "passed"]
        output = runtime.output_path(record.job_id, request.conversation_id)
        charts = [shape.chart for slide in Presentation(output).slides for shape in slide.shapes if shape.has_chart]
        assert len(charts) == 1 and list(charts[0].series[0].values) == [42.0, 38.0]
        assert [category.label for category in charts[0].plots[0].categories] == ["甲", "乙"]
        accepted = json.loads(runtime.service.store.checkpoint_path(record.job_id, "page-scene-001.json").read_text())
        assert all("chart_ref" not in obj for obj in accepted["slide"]["objects"])
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("failure_kind", ["reference", "literal_style"])
def test_public_reviewed_chart_recovery_keeps_owned_catalog_feedback(
    tmp_path, monkeypatch, failure_kind,
):
    case = chart_case()
    request, brief, plan, _deck, scene = case
    private_marker = "SYNTHETIC_PRIVATE_CHART_REFERENCE"
    expected_code = "invalid_chart_reference"
    if failure_kind == "reference":
        scene["objects"][-1]["chart_ref"] = "/private/" + private_marker
    else:
        chart = plan.slides[0].charts[0].model_dump(mode="json")
        chart.update({
            "id": scene["objects"][-1]["id"], "frame": scene["objects"][-1]["frame"],
            "show_legend": True,
        })
        scene["objects"][-1] = chart
        expected_code = "planned_chart_data_changed"
    calls = []

    async def transport(**kwargs):
        calls.append(json.loads(kwargs["user_prompt"]))
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    runtime = PresentationJobRuntime(tmp_path, page_author=UtilityPresentationPageAuthor("default"))
    try:
        record = runtime.service.create_job(request)
        runtime.service.store.write_checkpoint(record.job_id, "brief.json", brief.model_dump_json())
        runtime.service.store.write_checkpoint(record.job_id, "deck-plan.json", plan.model_dump_json())
        runtime.pause(record.job_id, request.conversation_id)
        runtime.resume(record.job_id, request.conversation_id)
        failed = settled_failure(runtime, record.job_id, request.conversation_id)
        assert len(calls) == failed.cost_used_units == 3
        assert calls[0].get("retry_directive") is None
        assert "chart_reference_transport.references" in calls[1]["retry_directive"]
        assert "chart_ref" in calls[2]["retry_directive"]
        routes = failed.events[-1].failure_routes
        assert routes[0].issue_codes == ["content." + expected_code]
        original_deadline, original_request, original_events = failed.deadline_at, failed.request, failed.events
        feedback_path = runtime.service.store.checkpoint_path(record.job_id, "page-author-feedback-001.json")
    finally:
        runtime.shutdown()

    feedback = json.loads(feedback_path.read_text(encoding="utf-8"))
    assert "chart_contract" in feedback["parts"]
    # Files are untrusted. Recovery must reconstruct only owned instructions,
    # not replay an arbitrary directive or an out-of-catalog chart path.
    feedback["retry_directive"] = private_marker
    feedback["chart_reference"] = "/private/" + private_marker
    feedback_path.write_text(json.dumps(feedback), encoding="utf-8")
    restarted = PresentationJobRuntime(tmp_path, page_author=UtilityPresentationPageAuthor("default"))
    try:
        restarted.retry_failed_after_review(
            record.job_id, request.conversation_id, reviewed_routes=routes,
            review_basis="synthetic-chart-contract-review",
        )
        failed = settled_failure(restarted, record.job_id, request.conversation_id)
        assert len(calls) == failed.cost_used_units == 6
        assert failed.deadline_at == original_deadline and failed.request == original_request
        assert failed.events[:len(original_events)] == original_events
        assert failed.recovery_attempt_count == 1
        directive = calls[3].get("retry_directive")
        assert directive and "chart_reference_transport.references" in directive
        assert "不会" in directive and "chart_ref" in directive
        assert private_marker not in json.dumps(calls, ensure_ascii=False)
        assert all(payload["chart_reference_transport"]["references"] == [{"ref": "storyboard.charts[0]"}] for payload in calls)
        assert not list(restarted.service.store.job_dir(record.job_id).rglob("*.pptx"))
    finally:
        restarted.shutdown()
