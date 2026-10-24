from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from lantai.presentation.models import (
    CompositionAuthority, CompositionEngine, DeckIR, DeckSlide, InputMaterial, ObjectFrame,
    PlanningMode, PresentationJobRequest, TemplateMode, TextObject,
)
from lantai.app.structured_utility import StructuredTaskValidationError
from lantai.presentation.intelligence import UtilityPresentationPageAuthor
from lantai.presentation.page_authoring import scope_authored_page_object_ids
from lantai.presentation.planner import build_brief, build_deck_ir, build_deterministic_fixture_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reference_deck.analyzer import ReferenceDeckAnalyzer
from lantai.presentation.reference_deck.models import ReferenceVisualMeasurement
from lantai.presentation.reference_deck.object_parser import PptxObjectParser
from lantai.presentation.reference_deck.profile_builder import ReferenceProfileBuilder
from lantai.presentation.service import PresentationService
from lantai.presentation.template_binding import (
    TemplateBindingError, bind_authorized_template,
)
from scripts.presentation_flagship.validate_d06_brand_preservation import (
    inspect_logo_inheritance,
)


LOGO_FRAME = (7315200, 91440, 914400, 457200)


def _template(tmp_path: Path) -> tuple[Path, str]:
    """Controlled OOXML fixture, never a user-facing product candidate."""
    source = Presentation()
    source.slide_width, source.slide_height = 12192000, 6858000
    layout = source.slide_layouts[6]
    for _ in range(2):
        page = source.slides.add_slide(layout)
        body = page.shapes.add_textbox(Inches(1), Inches(2), Inches(5), Inches(1))
        body.text = "SOURCE BODY MUST NOT BE COPIED"
    scratch = source.slides[0]
    pixels = BytesIO()
    Image.new("RGB", (16, 8), "#683FA5").save(pixels, format="PNG")
    logo_sha = sha256(pixels.getvalue()).hexdigest()
    pixels.seek(0)
    logo = scratch.shapes.add_picture(
        pixels, Inches(8), Inches(0.1), width=Inches(1), height=Inches(0.5),
    )
    picture = deepcopy(logo._element)
    identity = picture.xpath("./p:nvPicPr/p:cNvPr")[0]
    identity.set("id", "100")
    identity.set("name", "controlled-native-logo")
    rid = picture.xpath(".//a:blip")[0].get(qn("r:embed"))
    master_rid = layout.slide_master.part.relate_to(scratch.part.related_part(rid), RT.IMAGE)
    picture.xpath(".//a:blip")[0].set(qn("r:embed"), master_rid)
    layout.slide_master.shapes._spTree.append(picture)
    scratch.shapes._spTree.remove(logo._element)
    scratch.part.drop_rel(rid)
    wordmark = scratch.shapes.add_textbox(Inches(6), Inches(0.1), Inches(1.5), Inches(0.4))
    wordmark.text = "Controlled Brand"
    wordmark.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    wordmark.text_frame.paragraphs[0].runs[0].font.name = "Noto Sans CJK SC"
    copied = deepcopy(wordmark._element)
    copied.xpath("./p:nvSpPr/p:cNvPr")[0].set("id", "200")
    copied.xpath("./p:nvSpPr/p:cNvPr")[0].set("name", "controlled-layout-wordmark")
    layout.shapes._spTree.append(copied)
    scratch.shapes._spTree.remove(wordmark._element)
    unused = deepcopy(picture)
    unused.xpath("./p:nvPicPr/p:cNvPr")[0].set("name", "unused-layout-logo")
    unused_layout = source.slide_layouts[0]
    unused_rid = unused_layout.part.relate_to(layout.slide_master.part.related_part(master_rid), RT.IMAGE)
    unused.xpath(".//a:blip")[0].set(qn("r:embed"), unused_rid)
    unused_layout.shapes._spTree.append(unused)
    path = tmp_path / "authorized-template.pptx"
    source.save(path)
    return path, logo_sha


def _deck(count: int = 3) -> DeckIR:
    return DeckIR(
        id="template-binding-test", title="Template binding",
        slides=[
            DeckSlide(
                id=f"slide-{number}", layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=f"Source note {number}",
                objects=[TextObject(
                    id=f"title-{number}", role="title", text=f"New content {number}",
                    frame=ObjectFrame(x=1, y=1, w=8, h=0.8), font_size=32,
                )],
            ) for number in range(1, count + 1)
        ],
    )


def test_native_compilation_preserves_authorized_master_logo_without_copying_source_body(tmp_path: Path) -> None:
    template, logo_sha = _template(tmp_path)
    source_sha = sha256(template.read_bytes()).hexdigest()
    assert inspect_logo_inheritance(
        template, logo_sha256=logo_sha, frame_emu=LOGO_FRAME, expected_slide_count=2,
    )["passed"]
    deck = bind_authorized_template(
        _deck(), template_path=template, source_sha256=source_sha,
        template_mode=TemplateMode.STRICT,
    )
    output = tmp_path / "compiled.pptx"
    result = NativePptxCompiler().compile(deck, output)

    report = inspect_logo_inheritance(
        output, logo_sha256=logo_sha, frame_emu=LOGO_FRAME, expected_slide_count=3,
    )
    assert report["passed"], report["findings"]
    assert result.editable_object_count == 3
    assert sha256(template.read_bytes()).hexdigest() == source_sha
    reopened = Presentation(output)
    for number, page in enumerate(reopened.slides, 1):
        assert [shape.text for shape in page.shapes if shape.has_text_frame and not shape.name.startswith("meta:")] == [f"New content {number}"]
        assert any(shape.name == f"meta:template-binding-test:slide-{number}" for shape in page.shapes)
        assert f"Source note {number}" in page.notes_slide.notes_text_frame.text
        wordmark = next(shape for shape in page.slide_layout.shapes if shape.name == "controlled-layout-wordmark")
        run = wordmark.text_frame.paragraphs[0].runs[0]
        assert run.text == "Controlled Brand"
        assert run.font.size == Pt(14)
        assert run.font.name == "Noto Sans CJK SC"
    assert len(reopened.slide_masters) == 1
    assert len(reopened.slide_masters[0].slide_layouts) == 1
    assert not any(shape.name == "unused-layout-logo" for layout in reopened.slide_layouts for shape in layout.shapes)


def test_public_author_rejects_omitted_required_anchor_without_copying_template(tmp_path: Path, monkeypatch) -> None:
    import json
    from types import SimpleNamespace

    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.25))
    footer.text = "Controlled Brand"
    number = source.slides[0].shapes.add_textbox(Inches(12), Inches(7), Inches(0.5), Inches(0.25))
    number.text = "1"
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    requirements = {
        "source_sha256": source_sha,
        "pages": [{
            "target_slide_number": 1, "source_slide_number": 1,
            "anchors": [
                {"source_shape_id": footer.shape_id, "purpose": "static_brand"},
                {"source_shape_id": number.shape_id, "purpose": "page_number"},
            ],
        }],
    }
    request = PresentationJobRequest(
        conversation_id="controlled-required-anchor-author",
        user_request="Preserve the explicitly required native brand anchors.",
        slide_count=1, template_path=str(template), template_mode=TemplateMode.STRICT,
        template_anchor_requirements=requirements,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        materials=[InputMaterial(name="controlled.md", content="# Evidence\nA verified fact.")],
    )
    learned = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-required-author", template_mode=TemplateMode.STRICT,
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief).model_copy(update={"design_context": learned.design_context})
    scaffold = build_deck_ir(request, plan).slides[0]
    calls = []

    async def transport(**kwargs):
        payload = json.loads(kwargs["user_prompt"])
        calls.append(payload)
        assert payload["template_anchor_requirements"] == {
            "source_sha256": source_sha, "page": requirements["pages"][0],
        }
        return SimpleNamespace(text=json.dumps({
            "objects": [{
                "kind": "text", "id": "title", "role": "title", "text": plan.slides[0].title,
                "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1}, "font_size": 36,
            }],
            "template_anchors": {
                "source_slide_number": 1,
                "uses": [{"source_shape_id": footer.shape_id, "purpose": "static_brand", "target_object_id": "brand-footer"}],
            },
        }), model_label="approved-provider/controlled-model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    author = UtilityPresentationPageAuthor("controlled-agent")
    with pytest.raises(TemplateBindingError, match="^template_anchor_required_coverage_missing$") as captured:
        author.author(request, brief, plan, plan.slides[0], scaffold)

    assert captured.value.missing_anchors == ((number.shape_id, "page_number"),)
    assert len(calls) == 1
    assert scaffold.template_anchors is None
    assert sha256(template.read_bytes()).hexdigest() == source_sha
    assert author.cache_dependencies(request, brief, plan, plan.slides[0], scaffold)["user_payload"]["template_anchor_requirements"] == calls[0]["template_anchor_requirements"]


def test_author_selected_template_title_retains_native_frame_without_copying_source_title(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    title = source.slides[0].shapes.add_textbox(914401, 457201, 7315201, 731521)
    title.text = "SOURCE TITLE MUST NOT BE COPIED"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].runs[0].font.name = "Noto Sans CJK SC"
    source.save(template)
    raw = _deck(count=1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 1,
        "uses": [{
            "source_shape_id": title.shape_id,
            "purpose": "title_frame",
            "target_object_id": "title-1",
        }],
    }
    raw["slides"][0]["objects"][0]["frame"] = {
        "x": title.left / 914400, "y": title.top / 914400,
        "w": title.width / 914400, "h": title.height / 914400,
    }
    raw["slides"][0]["objects"][0]["font_size"] = 36
    deck = bind_authorized_template(
        DeckIR.model_validate(raw), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )

    output = tmp_path / "explicit-title-anchor.pptx"
    result = NativePptxCompiler().compile(deck, output)

    page = Presentation(output).slides[0]
    written_title = next(shape for shape in page.shapes if shape.name == "title-1")
    assert written_title.text == "New content 1"
    assert (written_title.left, written_title.top, written_title.width, written_title.height) == (
        title.left, title.top, title.width, title.height,
    )
    assert result.editable_object_count == 1
    assert not any("SOURCE" in shape.text for shape in page.shapes if shape.has_text_frame)


def test_author_selected_brand_footer_retains_small_native_text_without_copying_other_source_content(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(714376, 6410838, 7315200, 228600)
    footer.text = "Controlled Brand · Confidential"
    footer.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    footer.text_frame.paragraphs[0].runs[0].font.name = "Noto Sans CJK SC"
    source.save(template)
    raw = _deck(count=1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 1,
        "uses": [{
            "source_shape_id": footer.shape_id,
            "purpose": "static_brand",
            "target_object_id": "controlled-footer",
        }],
    }
    deck = bind_authorized_template(
        DeckIR.model_validate(raw), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )

    output = tmp_path / "explicit-footer-anchor.pptx"
    result = NativePptxCompiler().compile(deck, output)

    page = Presentation(output).slides[0]
    written = next(shape for shape in page.shapes if shape.name == "controlled-footer")
    assert written.text == footer.text
    assert (written.left, written.top, written.width, written.height) == (
        footer.left, footer.top, footer.width, footer.height,
    )
    assert written.text_frame.paragraphs[0].runs[0].font.size == Pt(11)
    assert written.text_frame.paragraphs[0].runs[0].font.name == "Noto Sans CJK SC"
    assert result.editable_object_count == 2
    assert not any("SOURCE" in shape.text for shape in page.shapes if shape.has_text_frame)


def test_author_selected_page_number_uses_output_order_and_preserves_native_format(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    number = source.slides[0].shapes.add_textbox(10972800, 6410838, 365760, 228600)
    number.text = "42"
    number.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    source.save(template)
    raw = _deck().model_dump(mode="python")
    for page in raw["slides"]:
        page["template_anchors"] = {
            "source_slide_number": 1,
            "uses": [{
                "source_shape_id": number.shape_id,
                "purpose": "page_number",
                "target_object_id": "native-page-number",
            }],
        }
    deck = bind_authorized_template(
        DeckIR.model_validate(raw), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )

    output = tmp_path / "explicit-page-number-anchor.pptx"
    result = NativePptxCompiler().compile(deck, output)

    for index, page in enumerate(Presentation(output).slides, 1):
        written = next(shape for shape in page.shapes if shape.name == "native-page-number")
        assert written.text == str(index)
        assert (written.left, written.top, written.width, written.height) == (
            number.left, number.top, number.width, number.height,
        )
        assert written.text_frame.paragraphs[0].runs[0].font.size == Pt(11)
    assert result.editable_object_count == 6


def test_author_selected_text_free_brand_rule_retains_native_geometry_and_color(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    rule = source.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, 594361, 512064, 82296, 630936)
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor.from_string("683FA5")
    rule._element.remove(rule._element.find(qn("p:txBody")))
    source.save(template)
    raw = _deck(count=1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 1,
        "uses": [{
            "source_shape_id": rule.shape_id,
            "purpose": "static_brand",
            "target_object_id": "native-brand-rule",
        }],
    }
    deck = bind_authorized_template(
        DeckIR.model_validate(raw), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )

    output = tmp_path / "explicit-rule-anchor.pptx"
    result = NativePptxCompiler().compile(deck, output)

    written = next(shape for shape in Presentation(output).slides[0].shapes if shape.name == "native-brand-rule")
    assert (written.left, written.top, written.width, written.height) == (
        rule.left, rule.top, rule.width, rule.height,
    )
    assert written.fill.fore_color.rgb == rule.fill.fore_color.rgb
    assert written._element.find(qn("p:txBody")) is None
    assert result.editable_object_count == 2


def test_utility_author_selects_pinned_native_footer_through_actual_scene_schema(tmp_path: Path, monkeypatch) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(714376, 6410838, 7315200, 228600)
    footer.text = "Controlled Brand · Confidential"
    footer.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    learned = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-author", template_mode=TemplateMode.STRICT,
    )
    request = PresentationJobRequest(
        conversation_id="controlled-author-template",
        user_request="Create a new controlled title while preserving the brand footer.",
        slide_count=1, template_path=str(template), template_mode=TemplateMode.STRICT,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        materials=[InputMaterial(name="controlled-facts.md", content="# New title\nControlled source evidence.")],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief).model_copy(update={"design_context": learned.design_context})
    scaffold = build_deck_ir(request, plan).slides[0]
    calls = []

    async def transport(**kwargs):
        import json
        from types import SimpleNamespace
        calls.append(kwargs)
        payload = json.loads(kwargs["user_prompt"])
        catalog = payload["template_anchor_catalog"]
        assert catalog["source_sha256"] == source_sha
        selected = next(shape for shape in catalog["pages"][0]["native_shapes"] if shape["source_shape_id"] == footer.shape_id)
        assert selected["native_text"] == footer.text
        assert selected["frame_inches"] == {
            "x": footer.left / 914400, "y": footer.top / 914400,
            "w": footer.width / 914400, "h": footer.height / 914400,
        }
        assert str(template) not in json.dumps(catalog)
        return SimpleNamespace(text=json.dumps({
            "objects": [{
                "kind": "text", "id": "title", "role": "title", "text": plan.slides[0].title,
                "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1}, "font_size": 36,
            }],
            "template_anchors": {
                "source_slide_number": 1,
                "uses": [{"source_shape_id": footer.shape_id, "purpose": "static_brand", "target_object_id": "brand-footer"}],
            },
        }), model_label="approved-provider/controlled-model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    author = UtilityPresentationPageAuthor("controlled-agent")
    authored = author.author(request, brief, plan, plan.slides[0], scaffold)
    assert authored.template_anchors is not None
    assert author.cache_dependencies(request, brief, plan, plan.slides[0], scaffold)["user_payload"]["template_anchor_catalog"]["source_sha256"] == source_sha
    bound = bind_authorized_template(
        build_deck_ir(request, plan).model_copy(update={"slides": [scope_authored_page_object_ids(authored)]}),
        template_path=template, source_sha256=source_sha, template_mode=TemplateMode.STRICT,
    )
    output = tmp_path / "actual-authored-native-footer.pptx"
    NativePptxCompiler().compile(bound, output)
    assert any(shape.text == footer.text for shape in Presentation(output).slides[0].shapes if shape.has_text_frame)
    assert len(calls) == 1


def test_page_scoping_updates_generated_and_preserved_template_anchor_ids_idempotently() -> None:
    raw = _deck(count=1).slides[0].model_dump(mode="python")
    raw["template_anchors"] = {
        "source_slide_number": 2,
        "uses": [
            {"source_shape_id": 3, "purpose": "title_frame", "target_object_id": "title-1"},
            {"source_shape_id": 5, "purpose": "static_brand", "target_object_id": "brand-footer"},
        ],
    }
    original = DeckSlide.model_validate(raw)

    scoped = scope_authored_page_object_ids(original)

    assert [use.target_object_id for use in scoped.template_anchors.uses] == [
        "slide-1::title-1", "slide-1::brand-footer",
    ]
    assert scoped.objects[0].id == "slide-1::title-1"
    assert DeckSlide.model_validate(scoped.model_dump(mode="python")) == scoped
    assert scope_authored_page_object_ids(scoped) == scoped


def test_native_service_binds_the_reference_pin_before_compilation(tmp_path: Path) -> None:
    template, logo_sha = _template(tmp_path)
    learned = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-service", template_mode=TemplateMode.STRICT,
    )
    service = PresentationService(tmp_path / "service")
    job = service.create_job(PresentationJobRequest(
        conversation_id="controlled-template-binding",
        user_request="Preserve the authorized template identity.",
        slide_count=3, template_path=str(template), template_mode=TemplateMode.STRICT,
        composition_engine=CompositionEngine.NATIVE_PPTX,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        visual_verification_required=False, vision_critique_required=False,
        cost_budget_units=1000,
    ))
    deck = _deck().model_copy(update={"design_context": learned.design_context})

    result, _ = service.compose(job.job_id, deck)

    report = inspect_logo_inheritance(
        result.output_path, logo_sha256=logo_sha, frame_emu=LOGO_FRAME,
        expected_slide_count=3,
    )
    assert report["passed"], report["findings"]
    checkpoint = DeckIR.model_validate_json(
        service.store.checkpoint_path(job.job_id, "deck-ir.json").read_text(),
    )
    assert checkpoint.template_binding is not None
    assert checkpoint.template_binding.source_sha256 == learned.measurement.source_sha256


def test_style_reference_does_not_read_or_copy_original_template_assets(tmp_path: Path) -> None:
    template, logo_sha = _template(tmp_path)
    preserve = bind_authorized_template(
        _deck(), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )
    # A stale native binding must not survive a switch to abstract style transfer.
    transferred = bind_authorized_template(
        preserve, template_path=tmp_path / "does-not-exist.pptx",
        source_sha256="0" * 64, template_mode=TemplateMode.STYLE_REFERENCE,
    )
    assert transferred.template_binding is None
    output = tmp_path / "style-only.pptx"
    NativePptxCompiler().compile(transferred, output)
    report = inspect_logo_inheritance(
        output, logo_sha256=logo_sha, frame_emu=LOGO_FRAME, expected_slide_count=3,
    )
    assert not report["passed"]
    reopened = Presentation(output)
    assert not any(
        shape.name == "controlled-layout-wordmark"
        for layout in reopened.slide_layouts for shape in layout.shapes
    )


def test_switching_to_style_reference_discards_explicit_native_anchors_without_source_io(tmp_path: Path) -> None:
    raw = _deck(count=1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 1,
        "uses": [{"source_shape_id": 5, "purpose": "static_brand", "target_object_id": "old-brand-footer"}],
    }
    original = DeckIR.model_validate(raw)

    transferred = bind_authorized_template(
        original, template_path=tmp_path / "intentionally-unavailable.pptx",
        source_sha256="0" * 64, template_mode=TemplateMode.STYLE_REFERENCE,
    )

    assert transferred.template_binding is None
    assert transferred.slides[0].template_anchors is None
    assert transferred.slides[0].objects == original.slides[0].objects
    output = tmp_path / "style-without-original-anchors.pptx"
    result = NativePptxCompiler().compile(transferred, output)
    assert result.editable_object_count == 1
    assert not any(shape.name == "old-brand-footer" for shape in Presentation(output).slides[0].shapes)


def test_changed_source_pin_is_rejected_without_overwriting_existing_output(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    original = sha256(template.read_bytes()).hexdigest()
    bound = bind_authorized_template(
        _deck(), template_path=template, source_sha256=original,
        template_mode=TemplateMode.STRICT,
    )
    output = tmp_path / "authoritative.pptx"
    NativePptxCompiler().compile(bound, output)
    previous = output.read_bytes()
    changed = Presentation(template)
    changed.slides[0].shapes[0].text = "Different source content"
    changed.save(template)

    with pytest.raises(TemplateBindingError, match="^template_source_pin_mismatch$"):
        NativePptxCompiler().compile(bound, output)

    assert output.read_bytes() == previous
    assert not list(tmp_path.glob(".template-binding-*.pptx"))


def test_broken_native_image_relationship_is_rejected_before_compilation(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    damaged = Presentation(template)
    picture = next(shape for shape in damaged.slide_masters[0].shapes if shape.name == "controlled-native-logo")
    picture._element.xpath(".//a:blip")[0].set(qn("r:embed"), "rIdMissingImage")
    damaged.save(template)

    with pytest.raises(TemplateBindingError, match="^template_xml_relationship_missing$"):
        bind_authorized_template(
            _deck(), template_path=template,
            source_sha256=sha256(template.read_bytes()).hexdigest(),
            template_mode=TemplateMode.STRICT,
        )


def test_image_reference_keeps_its_existing_abstract_design_path(tmp_path: Path) -> None:
    image_path = tmp_path / "authorized-reference.png"
    Image.new("RGB", (1280, 720), "#683FA5").save(image_path)
    original_sha = sha256(image_path.read_bytes()).hexdigest()
    learned = ReferenceDeckAnalyzer().analyze(image_path)
    service = PresentationService(tmp_path / "image-reference-service")
    job = service.create_job(PresentationJobRequest(
        conversation_id="controlled-image-reference",
        user_request="Use this image as design reference for new content.",
        slide_count=3, template_path=str(image_path),
        composition_engine=CompositionEngine.NATIVE_PPTX,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        visual_verification_required=False, vision_critique_required=False,
        cost_budget_units=1000,
    ))
    result, _ = service.compose(
        job.job_id, _deck().model_copy(update={"design_context": learned.design_context}),
    )

    assert len(Presentation(result.output_path).slides) == 3
    checkpoint = DeckIR.model_validate_json(
        service.store.checkpoint_path(job.job_id, "deck-ir.json").read_text(),
    )
    assert checkpoint.template_binding is None
    assert sha256(image_path.read_bytes()).hexdigest() == original_sha


def test_different_native_contexts_require_and_execute_an_author_selected_mapping(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    for page, color in zip(source.slides, ("F7F3EE", "DCECE0"), strict=True):
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(color)
    source.slides[1]._element.set("showMasterSp", "0")
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    with pytest.raises(TemplateBindingError, match="^template_context_mapping_required$"):
        bind_authorized_template(
            _deck(), template_path=template, source_sha256=source_sha,
            template_mode=TemplateMode.STRICT,
        )

    deck = bind_authorized_template(
        _deck(), template_path=template, source_sha256=source_sha,
        template_mode=TemplateMode.STRICT, source_slide_numbers=[2, 1, 2],
    )
    output = tmp_path / "explicit-context-mapping.pptx"
    NativePptxCompiler().compile(deck, output)

    reopened = Presentation(output)
    assert [str(page.background.fill.fore_color.rgb) for page in reopened.slides] == [
        "DCECE0", "F7F3EE", "DCECE0",
    ]
    assert [page._element.get("showMasterSp") for page in reopened.slides] == ["0", None, "0"]
    assert [page.shapes[0].text for page in reopened.slides] == [
        "New content 1", "New content 2", "New content 3",
    ]


def test_service_does_not_silently_repin_an_unidentified_reference(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    learned = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-unpinned", template_mode=TemplateMode.STRICT,
    )
    context = learned.design_context.model_copy(update={"reference_artifact_id": None})
    service = PresentationService(tmp_path / "unpinned-service")
    job = service.create_job(PresentationJobRequest(
        conversation_id="controlled-unpinned-reference",
        user_request="Preserve the confirmed template.", slide_count=3,
        template_path=str(template), template_mode=TemplateMode.STRICT,
        composition_engine=CompositionEngine.NATIVE_PPTX,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        visual_verification_required=False, vision_critique_required=False,
    ))
    with pytest.raises(TemplateBindingError, match="^template_reference_pin_required$"):
        service.compose(job.job_id, _deck().model_copy(update={"design_context": context}))

    current = service.store.get(job.job_id)
    assert current.cost_used_units == job.cost_used_units
    assert not service.store.checkpoint_path(job.job_id, "deck-ir.json").exists()
    assert not any((service.store.job_dir(job.job_id) / "candidates").iterdir())


def test_author_selected_anchors_execute_different_contexts_without_a_second_page_mapping(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer_ids = []
    for number, (page, color) in enumerate(zip(source.slides, ("F7F3EE", "DCECE0"), strict=True), 1):
        page.background.fill.solid()
        page.background.fill.fore_color.rgb = RGBColor.from_string(color)
        footer = page.shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.3))
        footer.text = f"Controlled Brand {number}"
        footer_ids.append(footer.shape_id)
    source.save(template)
    raw = _deck().model_dump(mode="python")
    for page, source_number in zip(raw["slides"], (2, 1, 2), strict=True):
        page["template_anchors"] = {
            "source_slide_number": source_number,
            "uses": [{
                "source_shape_id": footer_ids[source_number - 1],
                "purpose": "static_brand", "target_object_id": "selected-footer",
            }],
        }

    deck = bind_authorized_template(
        DeckIR.model_validate(raw), template_path=template,
        source_sha256=sha256(template.read_bytes()).hexdigest(),
        template_mode=TemplateMode.STRICT,
    )
    output = tmp_path / "author-selected-contexts.pptx"
    NativePptxCompiler().compile(deck, output)

    reopened = Presentation(output)
    assert deck.template_binding.source_slide_numbers == [2, 1, 2]
    assert [str(page.background.fill.fore_color.rgb) for page in reopened.slides] == ["DCECE0", "F7F3EE", "DCECE0"]
    assert [next(shape.text for shape in page.shapes if shape.name == "selected-footer") for page in reopened.slides] == [
        "Controlled Brand 2", "Controlled Brand 1", "Controlled Brand 2",
    ]
    assert not any("SOURCE" in shape.text for page in reopened.slides for shape in page.shapes if shape.has_text_frame)


@pytest.mark.parametrize("failure", [
    "static_brand_collision", "page_number_collision", "missing_title",
    "non_title_target", "duplicate_object_id",
])
def test_public_page_author_reports_anchor_collision_as_safe_schema_failure(monkeypatch, failure: str) -> None:
    import json
    from types import SimpleNamespace

    request = PresentationJobRequest(
        conversation_id="controlled-anchor-boundary",
        user_request="Create a controlled presentation.", slide_count=1,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        materials=[InputMaterial(name="controlled.md", content="# Controlled evidence\nA verified fact.")],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan).slides[0]
    rejected_id = "private-rejected-object-sentinel"
    rejected_text = "private-rejected-text-sentinel"
    objects = [{
        "kind": "text", "id": rejected_id, "role": "title", "text": rejected_text,
        "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1}, "font_size": 36,
    }]
    purpose = "static_brand"
    target = rejected_id
    if failure == "page_number_collision":
        purpose = "page_number"
    elif failure == "missing_title":
        purpose, target = "title_frame", "private-missing-title-sentinel"
    elif failure == "non_title_target":
        purpose = "title_frame"
        objects[0]["role"] = "body"
    elif failure == "duplicate_object_id":
        objects.append({**objects[0], "role": "body"})
        target = "unused-controlled-anchor"

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps({
            "objects": objects,
            "template_anchors": {
                "source_slide_number": 1,
                "uses": [{"source_shape_id": 2, "purpose": purpose, "target_object_id": target}],
            },
        }), model_label="approved-provider/controlled-model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)

    with pytest.raises(StructuredTaskValidationError) as captured:
        UtilityPresentationPageAuthor("controlled-agent").author(request, brief, plan, plan.slides[0], scaffold)

    assert captured.value.validation_issues == ("value_error@$",)
    assert rejected_id not in str(captured.value)
    assert rejected_text not in str(captured.value)
    assert target not in str(captured.value)


def test_public_page_author_preserves_valid_anchor_selection(monkeypatch) -> None:
    import json
    from types import SimpleNamespace

    request = PresentationJobRequest(
        conversation_id="controlled-valid-anchor-boundary",
        user_request="Create a controlled presentation.", slide_count=1,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        materials=[InputMaterial(name="controlled.md", content="# Evidence\nA verified fact.")],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    scaffold = build_deck_ir(request, plan).slides[0]
    anchors = {
        "source_slide_number": 2,
        "uses": [
            {"source_shape_id": 4, "purpose": "title_frame", "target_object_id": "controlled-title"},
            {"source_shape_id": 5, "purpose": "static_brand", "target_object_id": "controlled-brand"},
            {"source_shape_id": 6, "purpose": "page_number", "target_object_id": "controlled-page"},
        ],
    }

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps({
            "objects": [{
                "kind": "text", "id": "controlled-title", "role": "title", "text": plan.slides[0].title,
                "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1}, "font_size": 36,
            }],
            "template_anchors": anchors,
        }), model_label="approved-provider/controlled-model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    authored = UtilityPresentationPageAuthor("controlled-agent").author(request, brief, plan, plan.slides[0], scaffold)
    assert authored.template_anchors.model_dump(mode="json") == anchors
    assert authored.objects[0].id == "controlled-title"


def test_public_binding_rejects_missing_explicitly_required_brand_anchors(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    page = source.slides[0]
    footer = page.shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.25))
    footer.text = "Controlled Brand"
    rule = page.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(6.9), Inches(10), Inches(0.02))
    number = page.shapes.add_textbox(Inches(12), Inches(7), Inches(0.5), Inches(0.25))
    number.text = "1"
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    raw = _deck(count=1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 1,
        "uses": [{"source_shape_id": footer.shape_id, "purpose": "static_brand", "target_object_id": "footer"}],
    }
    deck = DeckIR.model_validate(raw)
    unchanged = deck.model_dump(mode="json")
    requirements = {
        "source_sha256": source_sha,
        "pages": [{
            "target_slide_number": 1, "source_slide_number": 1,
            "anchors": [
                {"source_shape_id": footer.shape_id, "purpose": "static_brand"},
                {"source_shape_id": rule.shape_id, "purpose": "static_brand"},
                {"source_shape_id": number.shape_id, "purpose": "page_number"},
            ],
        }],
    }

    with pytest.raises(TemplateBindingError, match="^template_anchor_required_coverage_missing$") as captured:
        bind_authorized_template(
            deck, template_path=template, source_sha256=source_sha,
            template_mode=TemplateMode.STRICT, required_anchors=requirements,
        )

    assert captured.value.missing_anchors == (
        (rule.shape_id, "static_brand"), (number.shape_id, "page_number"),
    )
    assert deck.model_dump(mode="json") == unchanged
    assert sha256(template.read_bytes()).hexdigest() == source_sha


@pytest.mark.parametrize("correct_on_retry", [True, False])
def test_required_anchor_feedback_uses_existing_finite_page_budget(tmp_path: Path, monkeypatch, correct_on_retry: bool) -> None:
    import json
    from types import SimpleNamespace
    from lantai.presentation.runtime import PresentationJobRuntime

    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.25))
    footer.text = "private-template-brand-sentinel"
    number = source.slides[0].shapes.add_textbox(Inches(12), Inches(7), Inches(0.5), Inches(0.25))
    number.text = "1"
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    request = PresentationJobRequest(
        conversation_id="controlled-required-anchor-feedback",
        user_request="Preserve the explicitly required native brand anchors.",
        slide_count=1, template_path=str(template), template_mode=TemplateMode.STRICT,
        template_anchor_requirements={
            "source_sha256": source_sha,
            "pages": [{
                "target_slide_number": 1, "source_slide_number": 1,
                "anchors": [
                    {"source_shape_id": footer.shape_id, "purpose": "static_brand"},
                    {"source_shape_id": number.shape_id, "purpose": "page_number"},
                ],
            }],
        },
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        materials=[InputMaterial(name="controlled.md", content="# Evidence\nA verified fact.")],
    )
    learned = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-required-feedback", template_mode=TemplateMode.STRICT,
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief).model_copy(update={"design_context": learned.design_context})
    scaffold = build_deck_ir(request, plan)
    calls = []

    async def transport(**kwargs):
        payload = json.loads(kwargs["user_prompt"])
        calls.append(payload)
        uses = [{"source_shape_id": footer.shape_id, "purpose": "static_brand", "target_object_id": "private-author-brand-target-sentinel"}]
        if correct_on_retry and "template_anchor_requirements" in payload.get("retry_directive", ""):
            uses.append({"source_shape_id": number.shape_id, "purpose": "page_number", "target_object_id": "page"})
        return SimpleNamespace(text=json.dumps({
            "objects": [
                {"kind": "text", "id": "title", "role": "title", "text": plan.slides[0].title,
                 "frame": {"x": 1, "y": 0.5, "w": 10, "h": 1}, "font_size": 36},
                {"kind": "text", "id": "body", "role": "body",
                 "text": "\n".join(plan.slides[0].intent.required_facts + plan.slides[0].intent.immutable_facts) or "A verified fact.",
                 "frame": {"x": 1, "y": 2, "w": 10, "h": 3}, "font_size": 24},
            ],
            "template_anchors": {"source_slide_number": 1, "uses": uses},
        }), model_label="approved-provider/controlled-model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    monkeypatch.setattr("lantai.presentation.intelligence._utility_route_cache_identity", lambda _: "approved-provider/controlled-model")
    runtime = PresentationJobRuntime(tmp_path / "service", page_author=UtilityPresentationPageAuthor("controlled-agent"))
    job = runtime.service.create_job(request)
    try:
        if correct_on_retry:
            authored = runtime._author_pages(job.job_id, request, brief, plan, scaffold)
            assert len(authored.slides[0].template_anchors.uses) == 2
            bound = bind_authorized_template(
                authored, template_path=template, source_sha256=source_sha,
                template_mode=TemplateMode.STRICT, required_anchors=request.template_anchor_requirements,
            )
            output = tmp_path / "required-anchor-verified.pptx"
            NativePptxCompiler().compile(bound, output)
            texts = [shape.text for shape in Presentation(output).slides[0].shapes if shape.has_text_frame]
            assert footer.text in texts and "1" in texts
            assert "SOURCE BODY MUST NOT BE COPIED" not in texts
        else:
            with pytest.raises(TemplateBindingError, match="^template_anchor_required_coverage_missing$"):
                runtime._author_pages(job.job_id, request, brief, plan, scaffold)
            assert not runtime.service.store.checkpoint_path(job.job_id, "page-scene-001.json").exists()
        saved = runtime.service.store.get(job.job_id)
        assert len(calls) == saved.cost_used_units == (2 if correct_on_retry else 3)
        for payload in calls[1:]:
            directive = payload["retry_directive"]
            assert "template_anchor_requirements" in directive
            assert f"source_shape_id={number.shape_id}/purpose=page_number" in directive
            assert "不得复制旧正文" in directive
            assert "不得修改事实或图表数据" in directive
            assert footer.text not in directive
            assert "private-author-brand-target-sentinel" not in directive
            assert str(template) not in directive
        assert footer.text not in "\n".join(event.message for event in saved.events)
        assert runtime._PAGE_AUTHOR_TRANSIENT_ATTEMPT_LIMIT == 5
        assert runtime._PAGE_AUTHOR_RETRY_LIMIT_PER_CAUSE == 2
        assert sha256(template.read_bytes()).hexdigest() == source_sha
    finally:
        runtime.shutdown()


def test_service_checks_required_anchor_coverage_before_checkpoint_or_output(tmp_path: Path) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.25))
    footer.text = "Controlled Brand"
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    context = ReferenceProfileBuilder().build(
        PptxObjectParser().parse(template), ReferenceVisualMeasurement(),
        profile_id="controlled-service-required", template_mode=TemplateMode.STRICT,
    ).design_context
    service = PresentationService(tmp_path / "required-service")
    job = service.create_job(PresentationJobRequest(
        conversation_id="controlled-required-service", user_request="Preserve required anchors.",
        slide_count=1, template_path=str(template), template_mode=TemplateMode.STRICT,
        template_anchor_requirements={
            "source_sha256": source_sha,
            "pages": [{"target_slide_number": 1, "source_slide_number": 1,
                       "anchors": [{"source_shape_id": footer.shape_id, "purpose": "static_brand"}]}],
        },
        composition_engine=CompositionEngine.NATIVE_PPTX,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        visual_verification_required=False, vision_critique_required=False,
    ))
    deck = _deck(1).model_copy(update={"design_context": context})
    unchanged = deck.model_dump(mode="json")

    with pytest.raises(TemplateBindingError, match="^template_anchor_required_coverage_missing$"):
        service.compose(job.job_id, deck)

    assert not service.store.checkpoint_path(job.job_id, "deck-ir.json").exists()
    assert not (service.store.job_dir(job.job_id) / "output" / "presentation.pptx").exists()
    assert deck.model_dump(mode="json") == unchanged
    assert sha256(template.read_bytes()).hexdigest() == source_sha
    assert service.store.get(job.job_id).cost_used_units == job.cost_used_units


def test_required_native_anchors_cannot_be_silently_ignored_by_legacy_engine() -> None:
    from pydantic import ValidationError
    with pytest.raises(ValidationError, match="required template anchors need a native composition engine"):
        PresentationJobRequest(
            conversation_id="controlled-legacy-contract", user_request="Keep required native anchors.",
            slide_count=1, template_path="controlled.pptx", template_mode=TemplateMode.STRICT,
            composition_engine=CompositionEngine.LEGACY_PYTHON_PPTX,
            template_anchor_requirements={
                "source_sha256": "a" * 64,
                "pages": [{"target_slide_number": 1, "source_slide_number": 1,
                           "anchors": [{"source_shape_id": 2, "purpose": "static_brand"}]}],
            },
        )


@pytest.mark.parametrize("failure,code", [
    ("pin", "template_anchor_requirements_pin_mismatch"),
    ("target_pages", "template_anchor_requirements_target_pages_mismatch"),
    ("source_page", "template_anchor_requirements_source_page_invalid"),
    ("source_shape", "template_anchor_requirements_source_shape_invalid"),
    ("duplicate_shape", "template_anchor_requirements_invalid"),
    ("wrong_role", "template_anchor_required_coverage_missing"),
    ("wrong_source", "template_anchor_required_coverage_missing"),
])
def test_required_anchor_contract_is_independently_pinned_and_not_author_selected(tmp_path: Path, failure: str, code: str) -> None:
    template, _ = _template(tmp_path)
    source = Presentation(template)
    footer = source.slides[0].shapes.add_textbox(Inches(1), Inches(7), Inches(5), Inches(0.25))
    footer.text = "Controlled Brand"
    source.save(template)
    source_sha = sha256(template.read_bytes()).hexdigest()
    required = {
        "source_sha256": source_sha,
        "pages": [{"target_slide_number": 1, "source_slide_number": 1,
                   "anchors": [{"source_shape_id": footer.shape_id, "purpose": "static_brand"}]}],
    }
    page = required["pages"][0]
    if failure == "pin":
        required["source_sha256"] = "0" * 64
    elif failure == "target_pages":
        page["target_slide_number"] = 2
    elif failure == "source_page":
        page["source_slide_number"] = 3
    elif failure == "source_shape":
        page["anchors"][0]["source_shape_id"] = 9999
    elif failure == "duplicate_shape":
        page["anchors"].append(dict(page["anchors"][0]))
    raw = _deck(1).model_dump(mode="python")
    raw["slides"][0]["template_anchors"] = {
        "source_slide_number": 2 if failure == "wrong_source" else 1,
        "uses": [{"source_shape_id": footer.shape_id,
                  "purpose": "page_number" if failure == "wrong_role" else "static_brand",
                  "target_object_id": "footer"}],
    }
    deck = DeckIR.model_validate(raw)
    unchanged = deck.model_dump(mode="json")
    with pytest.raises(TemplateBindingError, match=f"^{code}$"):
        bind_authorized_template(
            deck, template_path=template, source_sha256=source_sha,
            template_mode=TemplateMode.STRICT, required_anchors=required,
        )
    assert deck.model_dump(mode="json") == unchanged
    assert sha256(template.read_bytes()).hexdigest() == source_sha
