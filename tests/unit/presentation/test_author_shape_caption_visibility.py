"""New-page captions must be visible without changing native edit semantics."""

from __future__ import annotations

import json
import time
from types import SimpleNamespace

import pytest
from pptx import Presentation

from lantai.app.structured_utility import StructuredTaskValidationError
from lantai.presentation.intelligence import UtilityPresentationPageAuthor
from lantai.presentation.models import (
    InputMaterial, JobStatus, PresentationJobRequest,
    PresentationObjectEditPlan, ShapeObject,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.planner import (
    build_brief, build_deck_ir, build_deterministic_fixture_plan,
)
from lantai.presentation.runtime import PresentationJobRuntime
from lantai.presentation.pptx_compiler import NativePptxCompiler


def _case(shape):
    request = PresentationJobRequest(
        conversation_id="synthetic-authored-shape-caption",
        user_request="制作一页定性证据判断，不新增数字",
        slide_count=1,
        materials=[InputMaterial(name="合成来源.md", content="风险尚待核验。")],
    )
    brief = build_brief(request)
    original = build_deterministic_fixture_plan(request, brief)
    planned = original.slides[0].model_copy(update={
        "title": "先核验证据，再判断风险", "charts": [],
        "intent": original.slides[0].intent.model_copy(update={
            "required_facts": ["风险尚待核验。"], "immutable_facts": [],
        }),
    })
    plan = original.model_copy(update={"slides": [planned]})
    scaffold = build_deck_ir(request, plan).slides[0]
    scene = {
        "objects": [
            {
                "kind": "text", "id": "title", "role": "title",
                "text_refs": ["storyboard.title"], "font_size": 36,
                "frame": {"x": 0.8, "y": 0.7, "w": 11.5, "h": 0.9},
            },
            {
                "kind": "text", "id": "evidence", "role": "body",
                "text_refs": ["required_facts[0]"], "font_size": 24,
                "frame": {"x": 0.8, "y": 2, "w": 11.5, "h": 1},
            },
            {
                "kind": "shape", "id": "synthetic-caption",
                "frame": {"x": 0.8, "y": 4, "w": 5, "h": 1.5},
                "font_size": 24, **shape,
            },
        ],
        "speaker_notes": "", "object_groups": [], "relationships": [],
    }
    return request, brief, plan, scaffold, scene


def test_public_author_rejects_runs_only_shape_caption_before_it_disappears(
    monkeypatch,
):
    request, brief, plan, scaffold, scene = _case({
        "text": "", "runs": [{"text": "合成标签", "font_size": 24}],
    })

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps(scene), model_label="synthetic/test")

    # Replace only the external transport; parsing and authored validation are real.
    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    with pytest.raises(StructuredTaskValidationError) as captured:
        UtilityPresentationPageAuthor("default").author(
            request, brief, plan, plan.slides[0], scaffold,
        )
    assert "authored_shape_caption_text@$.objects[2].shape" in captured.value.validation_issues
    assert "合成标签" not in str(captured.value)


@pytest.mark.parametrize("line_caption", [False, True])
def test_public_runtime_explains_caption_authority_without_echoing_rejected_text(
    tmp_path, monkeypatch, line_caption,
):
    shape = {
        "text": "", "runs": [{"text": "合成私密标签", "font_size": 24}],
    }
    if line_caption:
        shape.update({
            "shape_type": "line", "text": "合成私密标签", "runs": [],
            "frame": {"x": 0.8, "y": 4, "w": 5, "h": 0},
        })
    issue = "authored_shape_line_caption" if line_caption else "authored_shape_caption_text"
    request, brief, plan, _, scene = _case(shape)
    calls = []

    async def transport(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(text=json.dumps(scene), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    runtime = PresentationJobRuntime(tmp_path, page_author=UtilityPresentationPageAuthor("default"))
    try:
        record = runtime.service.create_job(request)
        job_id = record.job_id
        runtime.service.store.write_checkpoint(job_id, "brief.json", brief.model_dump_json())
        runtime.service.store.write_checkpoint(job_id, "deck-plan.json", plan.model_dump_json())
        runtime.pause(job_id, request.conversation_id)
        runtime.resume(job_id, request.conversation_id)
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            record = runtime.get(job_id, request.conversation_id)
            if record.status == JobStatus.FAILED and record.active_lease_id is None:
                break
            time.sleep(0.01)
        else:
            raise AssertionError("synthetic caption failure was not settled")
        assert not record.resource_settlement_pending
        assert len(calls) == record.cost_used_units == 3
        assert record.generation == record.attempt_count == 1
        assert record.recovery_attempt_count == 0
        assert not list(runtime.service.store.job_dir(job_id).rglob("*.pptx"))
        for call in calls:
            assert "ShapeObject.text 是唯一权威可见标签" in call["system_prompt"]
            if line_caption:
                assert "shape_type=line 不能承载标签" in call["system_prompt"]
        for call in calls[1:]:
            directive = json.loads(call["user_prompt"])["retry_directive"]
            assert issue + "@$.objects[2]" in directive
            if line_caption:
                assert "shape_type=line 不能承载标签" in directive
                assert "TextObject" in directive
            else:
                assert "ShapeObject.text 是唯一权威可见标签" in directive
                assert "runs" in directive
            assert "合成私密标签" not in directive
            assert "synthetic-caption" not in directive
        feedback = runtime.service.store.checkpoint_path(job_id, "page-author-feedback-001.json")
        raw = feedback.read_text(encoding="utf-8")
        assert json.loads(raw)["schema_issues"] == [issue + "@$.objects[2]"]
        assert "合成私密标签" not in raw and "synthetic-caption" not in raw
        routes, original_deadline = record.events[-1].failure_routes, record.deadline_at
        history, original_request = record.events, record.request
    finally:
        runtime.shutdown()

    restarted = PresentationJobRuntime(tmp_path, page_author=UtilityPresentationPageAuthor("default"))
    try:
        restarted.retry_failed_after_review(
            job_id, request.conversation_id, reviewed_routes=routes,
            review_basis="synthetic-caption-feedback-review",
        )
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            record = restarted.get(job_id, request.conversation_id)
            if record.status == JobStatus.FAILED and record.active_lease_id is None:
                break
            time.sleep(0.01)
        else:
            raise AssertionError("synthetic caption recovery was not settled")
        assert len(calls) == record.cost_used_units == 6
        assert record.deadline_at == original_deadline
        assert record.request == original_request
        assert record.events[:len(history)] == history
        assert record.generation == 2 and record.recovery_attempt_count == 1
        directive = json.loads(calls[3]["user_prompt"])["retry_directive"]
        assert issue + "@$.objects[2]" in directive
        expected_guidance = (
            "shape_type=line 不能承载标签" if line_caption
            else "ShapeObject.text 是唯一权威可见标签"
        )
        assert expected_guidance in directive
        assert "合成私密标签" not in directive and "synthetic-caption" not in directive
        assert not list(restarted.service.store.job_dir(job_id).rglob("*.pptx"))
    finally:
        restarted.shutdown()


def test_public_author_rejects_whitespace_text_with_nonempty_caption_runs(monkeypatch):
    request, brief, plan, scaffold, scene = _case({
        "text": "   ", "runs": [{"text": "合成标签", "font_size": 24}],
    })

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps(scene), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    with pytest.raises(StructuredTaskValidationError) as captured:
        UtilityPresentationPageAuthor("default").author(
            request, brief, plan, plan.slides[0], scaffold,
        )
    assert "authored_shape_caption_text@$.objects[2].shape" in captured.value.validation_issues


def test_public_author_rejects_caption_on_a_line_that_cannot_render_text(monkeypatch):
    request, brief, plan, scaffold, scene = _case({
        "shape_type": "line", "text": "合成线条标签", "runs": [],
        "frame": {"x": 0.8, "y": 4, "w": 5, "h": 0},
    })

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps(scene), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    with pytest.raises(StructuredTaskValidationError) as captured:
        UtilityPresentationPageAuthor("default").author(
            request, brief, plan, plan.slides[0], scaffold,
        )
    assert "authored_shape_line_caption@$.objects[2].shape" in captured.value.validation_issues
    assert "合成线条标签" not in str(captured.value)


def _authored_deck(monkeypatch, shape):
    request, brief, plan, scaffold, scene = _case(shape)

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps(scene), model_label="synthetic/test")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    authored = UtilityPresentationPageAuthor("default").author(
        request, brief, plan, plan.slides[0], scaffold,
    )
    # Keep the public planner's fact registry alongside its authored slide.
    return build_deck_ir(request, plan).model_copy(update={"slides": [authored]})


def _captioned_deck(monkeypatch, shape_type):
    return _authored_deck(monkeypatch, {
        "shape_type": shape_type, "text": "合成可见标签",
        "runs": [
            {"text": "合成", "font_size": 24, "color": "102030"},
            {"text": "可见", "font_size": 24, "color": "204060", "bold": True},
            {"text": "标签", "font_size": 24, "color": "102030"},
        ],
    })


@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron", "line"])
@pytest.mark.parametrize("runs", [[], [{"text": "   ", "font_size": 24}]])
def test_legitimate_unlabelled_shapes_remain_authorable_and_native(
    tmp_path, monkeypatch, shape_type, runs,
):
    shape = {"shape_type": shape_type, "text": "", "runs": runs}
    if shape_type == "line":
        shape["frame"] = {"x": 0.8, "y": 4, "w": 5, "h": 0}
    deck = _authored_deck(monkeypatch, shape)
    output = tmp_path / "synthetic-unlabelled-shape.pptx"
    NativePptxCompiler().compile(deck, output)
    native = next(shape for shape in Presentation(str(output)).slides[0].shapes
                  if shape.name == "synthetic-caption")
    assert (native.text if native.has_text_frame else "") == ""
    assert native.shape_type.name == "AUTO_SHAPE"
    expected_preset = "roundRect" if shape_type == "round_rect" else shape_type
    assert native._element.xpath("./p:spPr/a:prstGeom")[0].get("prst") == expected_preset


@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron"])
def test_valid_authored_caption_exports_as_editable_native_rich_text(
    tmp_path, monkeypatch, shape_type,
):
    deck = _captioned_deck(monkeypatch, shape_type)
    output = tmp_path / "synthetic-caption.pptx"
    NativePptxCompiler().compile(deck, output)
    native = next(shape for shape in Presentation(str(output)).slides[0].shapes
                  if shape.name == "synthetic-caption")
    assert native.has_text_frame and native.text == "合成可见标签"
    runs = [run for paragraph in native.text_frame.paragraphs for run in paragraph.runs]
    assert [run.text for run in runs] == ["合成", "可见", "标签"]
    assert [str(run.font.color.rgb) for run in runs] == ["102030", "204060", "102030"]
    assert runs[1].font.bold is True
    assert native.shape_type.name == "AUTO_SHAPE"


@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron"])
@pytest.mark.parametrize("replacement", ["", "   ", "合成新标签"])
def test_native_label_edit_never_resurrects_the_immutable_old_caption_runs(
    tmp_path, monkeypatch, shape_type, replacement,
):
    deck = _captioned_deck(monkeypatch, shape_type)
    authored = deck.slides[0]
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, PresentationObjectEditPlan.model_validate({
        "slide_id": authored.id,
        "actions": [{"kind": "set_shape_text", "object_id": "synthetic-caption", "text": replacement}],
    }))
    old_shape = authored.objects[2]
    new_shape = changed.slides[0].objects[2]
    assert isinstance(old_shape, ShapeObject) and isinstance(new_shape, ShapeObject)
    assert ids == ["synthetic-caption"]
    assert new_shape.model_dump() == {**old_shape.model_dump(), "text": replacement}
    assert deck.model_dump_json() == before
    output = tmp_path / "synthetic-caption-edited.pptx"
    NativePptxCompiler().compile(changed, output)
    native = next(shape for shape in Presentation(str(output)).slides[0].shapes
                  if shape.name == "synthetic-caption")
    assert native.text == replacement
    assert "合成可见标签" not in native.text
