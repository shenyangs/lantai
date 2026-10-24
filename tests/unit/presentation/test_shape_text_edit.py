"""Explicit shape-label edits preserve Text-only and native-object boundaries."""

from __future__ import annotations

import json

import pytest
from pptx import Presentation
from pydantic import ValidationError

from lantai.app.structured_utility import StructuredUtilityResult
from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.intelligence import UtilityPresentationObjectEditPlanner
from lantai.presentation.models import (
    ChartObject, DeckIR, DeckSlide, ImageObject, MetricObject, ObjectFrame,
    ObjectGroup, ObjectRelationship, PresentationObjectEditPlan, RelationshipType,
    ReplaceTextAction, ShapeObject, TableObject, TextObject,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _deck(shape_type="round_rect", text="原标签"):
    return DeckIR(
        id="synthetic-shape-edit", title="形状文字边界",
        source_refs=["synthetic://shape-label"], required_facts=["保留限定"],
        immutable_data=["保留限定"],
        slides=[
            DeckSlide(
                id="target", layout_id="freeform", speaker_notes="原备注",
                objects=[
                    TextObject(id="title", slot="title", role="title", text="保留限定"),
                    ShapeObject(
                        id="label", shape_type=shape_type, text=text,
                        frame=ObjectFrame(x=2, y=2, w=3, h=1.5),
                        fill_color="ABCDEF", line_color="102030", line_width=2.5,
                        text_color="304050", font_size=24, rotate=7,
                    ),
                ],
                object_groups=[ObjectGroup(group_id="group", role="label", object_ids=["title", "label"])],
                relationships=[ObjectRelationship(
                    relationship_id="relation", relationship_type=RelationshipType.PROCESS,
                    source_object_ids=["title"], target_object_ids=["label"], label="保留关系",
                )],
            ),
            DeckSlide(id="selected", layout_id="freeform", objects=[
                ShapeObject(id="selected-label", text="附加页原标签", frame=ObjectFrame(x=1, y=2, w=3, h=1)),
            ]),
            DeckSlide(id="unselected", layout_id="freeform", objects=[
                ShapeObject(id="unselected-label", text="不可越界", frame=ObjectFrame(x=1, y=2, w=3, h=1)),
            ]),
        ],
    )


def _edit(*, object_id="label", text="新标签", additional=(), actions=None):
    return PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "additional_slide_ids": list(additional),
        "actions": actions if actions is not None else [
            {"kind": "set_shape_text", "object_id": object_id, "text": text},
        ],
    })


@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron"])
@pytest.mark.parametrize("text", ["新标签\n第二行", ""])
def test_explicit_shape_text_changes_only_text_and_keeps_native_identity(shape_type, text):
    deck = _deck(shape_type)
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, _edit(text=text))
    assert ids == ["label"]
    assert changed.slides[0].objects[1].model_dump() == {
        **deck.slides[0].objects[1].model_dump(), "text": text,
    }
    assert changed.slides[0].objects[0] == deck.slides[0].objects[0]
    assert changed.slides[0].object_groups == deck.slides[0].object_groups
    assert changed.slides[0].relationships == deck.slides[0].relationships
    assert changed.slides[0].speaker_notes == deck.slides[0].speaker_notes
    assert changed.slides[0].revision == deck.slides[0].revision + 1
    assert changed.slides[0].scene_hash
    assert changed.slides[1:] == deck.slides[1:]
    assert changed.source_refs == deck.source_refs
    assert changed.required_facts == deck.required_facts
    assert changed.immutable_data == deck.immutable_data
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("kind", ["text", "metric", "table", "chart", "image"])
def test_shape_action_rejects_other_object_kinds_atomically(kind):
    deck = _deck()
    unsupported = {
        "text": TextObject(id="wrong", role="body", text="原文本"),
        "metric": MetricObject(id="wrong", label="原指标", value="原值"),
        "table": TableObject(id="wrong", headers=["原表头"], rows=[["原格"]]),
        "chart": ChartObject(id="wrong", categories=["甲"], series={"合成": [1]}),
        "image": ImageObject(id="wrong", path="/private/tmp/synthetic-shape-edit.png", alt="合成图片"),
    }[kind]
    deck.slides[0].objects.append(unsupported)
    before = deck.model_dump_json()
    edit = _edit(actions=[
        {"kind": "set_shape_text", "object_id": "label", "text": "不得部分提交"},
        {"kind": "set_shape_text", "object_id": "wrong", "text": "非法目标"},
    ])
    with pytest.raises(ValueError, match="not an editable shape label"):
        apply_object_edit_plan(deck, edit)
    assert deck.model_dump_json() == before


def test_line_label_rejected_instead_of_silent_cross_compiler_difference():
    deck = _deck("line")
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="not an editable shape label"):
        apply_object_edit_plan(deck, _edit())
    assert deck.model_dump_json() == before


def test_existing_text_action_stays_text_only_for_shapes():
    deck = _deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="object is not editable text"):
        apply_object_edit_plan(deck, PresentationObjectEditPlan(
            slide_id="target", actions=[ReplaceTextAction(object_id="label", text="不应隐式接受")],
        ))
    assert deck.model_dump_json() == before


def test_same_shape_text_is_not_a_material_edit():
    deck = _deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="no material change"):
        apply_object_edit_plan(deck, _edit(text="原标签"))
    assert deck.model_dump_json() == before


def test_shape_action_cannot_target_unselected_page():
    deck = _deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="unknown object id"):
        apply_object_edit_plan(deck, _edit(object_id="unselected-label"))
    assert deck.model_dump_json() == before


def test_explicit_additional_shape_label_keeps_other_slides_untouched():
    deck = _deck()
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, _edit(
        object_id="selected-label", text="只改明确附加页", additional=["selected"],
    ))
    assert ids == ["selected-label"]
    assert changed.slides[0] == deck.slides[0]
    assert changed.slides[1].objects[0].text == "只改明确附加页"
    assert changed.slides[2] == deck.slides[2]
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("extra", ["frame", "shape_type", "fill_color"])
def test_shape_text_action_does_not_authorize_style_or_geometry(extra):
    with pytest.raises(ValidationError):
        _edit(actions=[{"kind": "set_shape_text", "object_id": "label", "text": "新标签", extra: None}])


@pytest.mark.parametrize("additional", [None, ["selected"]])
def test_planner_advertises_only_non_line_shapes_for_explicit_action(monkeypatch, additional):
    deck = _deck()
    deck.slides[0].objects.extend([
        ShapeObject(id="empty-label", text="", frame=ObjectFrame(x=5, y=2, w=2, h=1)),
        ShapeObject(id="connector", shape_type="line", text="非兼容文字", frame=ObjectFrame(x=1, y=4, w=3, h=0)),
        MetricObject(id="metric", label="不授权", value="原值"),
    ])
    calls = []

    async def utility(**kwargs):
        calls.append(kwargs)
        return StructuredUtilityResult(value=PresentationObjectEditPlan(
            slide_id="target", actions=[ReplaceTextAction(object_id="title", text="保留限定与新标题")],
        ), model_label="synthetic/synthetic", repaired=False)

    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", utility)
    UtilityPresentationObjectEditPlanner("synthetic-agent").plan(
        deck, "target", "仅明确指定范围", additional_slide_ids=additional,
    )
    payload = json.loads(calls[0]["user_prompt"])
    assert payload["editable_object_ids_by_action"]["set_shape_text"] == ["label", "empty-label"]
    assert payload["editable_object_ids_by_action"]["replace_text"] == ["title"]
    assert "set_shape_text" in calls[0]["system_prompt"]
    assert "unselected-label" not in calls[0]["user_prompt"]
    if additional:
        assert payload["additional_slides"][0]["editable_object_ids_by_action"]["set_shape_text"] == ["selected-label"]
    assert len(calls) == 1
    assert calls[0]["timeout_seconds"] == 90


@pytest.mark.parametrize("compiler", [PresentationCompiler, NativePptxCompiler])
def test_shape_label_reopens_as_native_text_with_same_shape_geometry(tmp_path, compiler):
    deck = _deck()
    original = tmp_path / "original.pptx"
    output = tmp_path / "changed.pptx"
    compiler().compile(deck, original)
    changed, _ = apply_object_edit_plan(deck, _edit(text="合成新标签"))
    compiler().compile(changed, output)
    original_native = Presentation(str(original))
    reopened = Presentation(str(output))
    old_shape = next(shape for shape in original_native.slides[0].shapes if shape.name == "label")
    new_shape = next(shape for shape in reopened.slides[0].shapes if shape.name == "label")
    assert (new_shape.left, new_shape.top, new_shape.width, new_shape.height, new_shape.rotation) == (
        old_shape.left, old_shape.top, old_shape.width, old_shape.height, old_shape.rotation,
    )
    visible = [shape.text for shape in reopened.slides[0].shapes if shape.has_text_frame]
    assert "合成新标签" in visible
    assert "原标签" not in visible
    assert [obj.kind for obj in changed.slides[0].objects] == ["text", "shape"]


@pytest.mark.parametrize("compiler", [PresentationCompiler, NativePptxCompiler])
@pytest.mark.parametrize("shape_type", ["rect", "round_rect", "ellipse", "chevron"])
@pytest.mark.parametrize("text", [
    "合成新标签\n第二行", "", "  合成标签  ", "   ", "合成标签\n",
])
def test_shape_label_survives_authoritative_reimport_edit_and_second_reimport(
    tmp_path, compiler, shape_type, text,
):
    deck = _deck(shape_type)
    # Freeform legacy pages require authored geometry for every native
    # object; an unframed title is not a valid round-trip fixture.
    deck.slides[0].objects[0].frame = ObjectFrame(x=0.8, y=0.4, w=10, h=1)
    original, output = tmp_path / "original.pptx", tmp_path / "edited.pptx"
    writer = compiler()
    writer.compile(deck, original)
    imported = reconcile_pptx(deck, original, assets_dir=tmp_path / "assets-first")
    label = next(obj for obj in imported.slides[0].objects if obj.id == "label")
    assert isinstance(label, ShapeObject)
    assert label.shape_type == shape_type
    assert label.text == "原标签"
    assert [obj.id for obj in imported.slides[0].objects] == ["title", "label"]

    changed, ids = apply_object_edit_plan(imported, _edit(text=text))
    assert ids == ["label"]
    assert changed.slides[0].objects[1].model_dump() == {**label.model_dump(), "text": text}
    writer.compile(changed, output)
    second = reconcile_pptx(changed, output, assets_dir=tmp_path / "assets-second")
    reimported = next(obj for obj in second.slides[0].objects if obj.id == "label")
    assert isinstance(reimported, ShapeObject)
    assert reimported.shape_type == shape_type
    # python-pptx exposes OOXML soft line breaks as vertical tabs, whereas
    # the native backend emits paragraph breaks. Both are visible newlines.
    assert reimported.text.replace("\v", "\n") == text
    assert [obj.id for obj in second.slides[0].objects] == ["title", "label"]
    assert second.slides[0].object_groups == imported.slides[0].object_groups
    assert second.slides[0].relationships == imported.slides[0].relationships
    assert second.required_facts == deck.required_facts
    assert second.immutable_data == deck.immutable_data
    assert second.source_refs == deck.source_refs
    original_shape = next(shape for shape in Presentation(str(original)).slides[0].shapes
                          if shape.name == "label")
    edited_shapes = [shape for shape in Presentation(str(output)).slides[0].shapes
                     if shape.name == "label"]
    assert len(edited_shapes) == 1
    native = edited_shapes[0]
    assert (native.left, native.top, native.width, native.height, native.rotation) == (
        original_shape.left, original_shape.top, original_shape.width,
        original_shape.height, original_shape.rotation,
    )
    assert native.text.replace("\v", "\n") == text
