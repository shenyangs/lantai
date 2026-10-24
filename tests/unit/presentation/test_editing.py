from __future__ import annotations

import hashlib

import pytest
from pptx import Presentation

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    DeckIR,
    DeckSlide,
    PresentationEditOperation,
    TextObject,
    TextReplacement,
)


def _deck(title: str = "原始标题") -> DeckIR:
    return DeckIR(
        id="editable-deck",
        title="事务编辑验证",
        slides=[
            DeckSlide(
                id="editable-slide",
                layout_id="cover",
                objects=[
                    TextObject(id="editable-title", role="title", text=title, slot="title"),
                    TextObject(id="untouched-subtitle", role="body", text="不应被修改", slot="subtitle"),
                ],
            )
        ],
    )


def _native_text(path, object_id: str) -> str:
    presentation = Presentation(str(path))
    shape = next(shape for shape in presentation.slides[0].shapes if shape.name == object_id)
    return shape.text


def _sha256(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _operation(operation_id: str = "edit-test") -> PresentationEditOperation:
    return PresentationEditOperation(
        operation_id=operation_id,
        deck_id="editable-deck",
        slide_ids=["editable-slide"],
        object_ids=["editable-title"],
        allowed_content_types={"text"},
        instruction="将标题改得更清楚",
    )


def test_preview_commit_and_rollback_are_source_isolated(tmp_path) -> None:
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(_deck(), source)
    original_hash = _sha256(source)
    service = PresentationEditTransactionService(tmp_path / "transactions")

    preview = service.prepare(
        source,
        _deck(),
        _operation(),
        [TextReplacement(object_id="editable-title", text="已验证的新标题")],
    )

    assert preview.verified is True
    assert _sha256(source) == original_hash
    assert _native_text(source, "editable-title") == "原始标题"
    assert _native_text(preview.candidate_path, "editable-title") == "已验证的新标题"
    assert _native_text(preview.candidate_path, "untouched-subtitle") == "不应被修改"

    receipt = service.commit(preview)
    assert receipt.verified is True
    assert _native_text(source, "editable-title") == "已验证的新标题"

    service.rollback(receipt)
    assert _sha256(source) == original_hash
    assert _native_text(source, "editable-title") == "原始标题"


def test_prepare_rejects_changes_outside_the_approved_scope(tmp_path) -> None:
    source = tmp_path / "source.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, source)
    service = PresentationEditTransactionService(tmp_path / "transactions")

    with pytest.raises(PermissionError, match="object scope"):
        service.prepare(
            source,
            deck,
            _operation(),
            [TextReplacement(object_id="untouched-subtitle", text="越界修改")],
        )


def test_commit_rejects_a_stale_preview(tmp_path) -> None:
    source = tmp_path / "source.pptx"
    deck = _deck()
    PresentationCompiler().compile(deck, source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    preview = service.prepare(
        source,
        deck,
        _operation(),
        [TextReplacement(object_id="editable-title", text="候选标题")],
    )

    PresentationCompiler().compile(_deck("用户同时修改的标题"), source)
    with pytest.raises(RuntimeError, match="source changed"):
        service.commit(preview)
    assert _native_text(source, "editable-title") == "用户同时修改的标题"


def test_edit_operation_requires_explicit_scope() -> None:
    with pytest.raises(ValueError, match="explicit slide or object scope"):
        PresentationEditOperation(
            deck_id="editable-deck",
            allowed_content_types={"text"},
            instruction="不能整份文档默认修改",
        )


def test_prepare_rejects_save_during_candidate_verification(tmp_path, monkeypatch) -> None:
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(_deck(), source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    verify = service._verify_text

    def concurrent_save(*args):
        verify(*args)
        PresentationCompiler().compile(_deck("人工保存"), source)

    monkeypatch.setattr(service, "_verify_text", concurrent_save)
    with pytest.raises(RuntimeError, match="source changed"):
        service.prepare(source, _deck(), _operation(), [
            TextReplacement(object_id="editable-title", text="候选标题"),
        ])
    assert _native_text(source, "editable-title") == "人工保存"
    assert not (tmp_path / "transactions/edit-test/preview.json").exists()


def test_redesign_rejects_save_during_compilation(tmp_path) -> None:
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(_deck(), source)
    service = PresentationEditTransactionService(tmp_path / "transactions")

    class ConcurrentCompiler:
        def compile(self, deck, candidate):
            PresentationCompiler().compile(deck, candidate)
            PresentationCompiler().compile(_deck("人工保存"), source)

    with pytest.raises(RuntimeError, match="source changed"):
        service.prepare_compiled_redesign(
            source, _deck(), _deck("重排标题"), _operation(), ConcurrentCompiler(),
        )
    assert _native_text(source, "editable-title") == "人工保存"
    assert not (tmp_path / "transactions/edit-test/preview.json").exists()


def test_redesign_compiles_one_native_slide_with_broader_review_scope(
    tmp_path,
) -> None:
    original = _deck().model_copy(
        update={
            "slides": [
                *_deck().slides,
                DeckSlide(
                    id="intent-context-slide",
                    layout_id="cover",
                    objects=[
                        TextObject(
                            id="intent-context-title",
                            role="title",
                            text="主判断上下文",
                            slot="title",
                        ),
                    ],
                ),
            ],
        },
    )
    changed = original.model_copy(
        update={
            "slides": [
                _deck("重排标题").slides[0],
                original.slides[1],
            ],
        },
    )
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(original, source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    operation = _operation().model_copy(
        update={
            "slide_ids": ["editable-slide", "intent-context-slide"],
        },
    )

    preview = service.prepare_compiled_redesign(
        source,
        original,
        changed,
        operation,
        PresentationCompiler(),
        compiled_slide_id="editable-slide",
    )

    assert _sha256(source) == preview.source_sha256
    assert _native_text(preview.candidate_path, "editable-title") == "重排标题"
    with pytest.raises(PermissionError, match="approved slide scope"):
        service.prepare_compiled_redesign(
            source,
            original,
            changed,
            operation.model_copy(update={"operation_id": "edit-outside"}),
            PresentationCompiler(),
            compiled_slide_id="outside-slide",
        )


def test_redesign_atomically_grafts_multiple_explicit_slides(tmp_path) -> None:
    original = _deck().model_copy(
        update={
            "slides": [
                *_deck().slides,
                DeckSlide(
                    id="second-slide",
                    layout_id="cover",
                    objects=[
                        TextObject(
                            id="second-title",
                            role="title",
                            text="第二页原文",
                            slot="title",
                        ),
                    ],
                ),
            ],
        },
    )
    changed = original.model_copy(
        update={
            "slides": [
                _deck("第一页已修").slides[0],
                original.slides[1].model_copy(
                    update={
                        "objects": [
                            original.slides[1].objects[0].model_copy(
                                update={"text": "第二页已修"},
                            ),
                        ],
                    },
                ),
            ],
        },
    )
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(original, source)
    source_hash = _sha256(source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    operation = PresentationEditOperation(
        operation_id="edit-multiple",
        deck_id=original.id,
        slide_ids=["editable-slide", "second-slide"],
        object_ids=["editable-title", "second-title"],
        allowed_content_types={"text"},
        instruction="一次修复两个互相依赖的页面",
    )

    preview = service.prepare_compiled_redesign(
        source,
        original,
        changed,
        operation,
        PresentationCompiler(),
        compiled_slide_ids=["editable-slide", "second-slide"],
    )

    assert _sha256(source) == source_hash
    presentation = Presentation(str(preview.candidate_path))
    by_name = {
        shape.name: shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    }
    assert by_name["editable-title"] == "第一页已修"
    assert by_name["second-title"] == "第二页已修"


def test_undo_preserves_newer_human_save(tmp_path) -> None:
    source = tmp_path / "source.pptx"
    PresentationCompiler().compile(_deck(), source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    preview = service.prepare(source, _deck(), _operation(), [
        TextReplacement(object_id="editable-title", text="已提交标题"),
    ])
    receipt = service.commit(preview)
    PresentationCompiler().compile(_deck("提交后人工修改"), source)
    saved_hash = _sha256(source)
    with pytest.raises(RuntimeError, match="source changed"):
        service.undo(receipt)
    assert _sha256(source) == saved_hash
    assert _native_text(source, "editable-title") == "提交后人工修改"
