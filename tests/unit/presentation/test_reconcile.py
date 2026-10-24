from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches

from lantai.presentation.models import (
    DeckIR,
    DeckSlide,
    ObjectFrame,
    ShapeObject,
    TextObject,
)
from lantai.presentation.reconcile import reconcile_pptx


def _deck() -> DeckIR:
    return DeckIR(
        id="reconcile-deck",
        title="人工保存重导入",
        slides=[
            DeckSlide(
                id="stable-slide",
                layout_id="freeform",
                objects=[
                    TextObject(
                        id="stable-title",
                        role="title",
                        text="保存前标题",
                        frame=ObjectFrame(x=0.8, y=0.7, w=5.5, h=0.8),
                    ),
                    ShapeObject(
                        id="deleted-shape",
                        frame=ObjectFrame(x=0.8, y=2, w=2, h=1),
                        fill_color="112233",
                    ),
                ],
            ),
        ],
    )


def test_reconcile_preserves_stable_ids_and_imports_user_objects(
    tmp_path: Path,
) -> None:
    source = tmp_path / "presentation.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    marker = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.01), Inches(0.01))
    marker.name = "meta:reconcile-deck:stable-slide"
    title = slide.shapes.add_textbox(Inches(1.25), Inches(0.85), Inches(6), Inches(0.8))
    title.name = "stable-title"
    title.text = "用户在澜台演示里改过的标题"
    extra = slide.shapes.add_textbox(Inches(8), Inches(4.8), Inches(4), Inches(0.9))
    extra.name = "用户新增批注"
    extra.text = "这是用户后来新增的内容"
    presentation.save(source)

    reconciled = reconcile_pptx(
        _deck(),
        source,
        assets_dir=tmp_path / "assets",
    )

    assert len(reconciled.slides) == 1
    assert reconciled.slides[0].id == "stable-slide"
    assert reconciled.slides[0].layout_id == "freeform"
    objects = {obj.id: obj for obj in reconciled.slides[0].objects}
    assert objects["stable-title"].text == "用户在澜台演示里改过的标题"
    assert objects["stable-title"].frame.x == pytest.approx(1.25)
    assert "deleted-shape" not in objects
    imported = [obj for obj in objects.values() if obj.id.startswith("imported-")]
    assert len(imported) == 1
    assert imported[0].text == "这是用户后来新增的内容"


def test_reconcile_falls_back_to_slide_position_for_legacy_pptx(
    tmp_path: Path,
) -> None:
    source = tmp_path / "legacy.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    title.name = "stable-title"
    title.text = "旧文件也能迁移"
    presentation.save(source)

    reconciled = reconcile_pptx(
        _deck(),
        source,
        assets_dir=tmp_path / "assets",
    )

    assert reconciled.slides[0].id == "stable-slide"
    assert reconciled.slides[0].objects[0].id == "stable-title"
    assert reconciled.slides[0].objects[0].text == "旧文件也能迁移"


@pytest.mark.parametrize("known", [True, False])
@pytest.mark.parametrize("text", ["  合成标签  ", "   ", "合成标签\n"])
def test_reconcile_keeps_shape_label_whitespace_without_changing_text_import(
    tmp_path: Path, known: bool, text: str,
) -> None:
    deck = _deck()
    if not known:
        deck.slides[0].objects = [deck.slides[0].objects[0]]
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(0.7), Inches(6), Inches(1))
    title.name = "stable-title"
    title.text = "  合成标题  "
    label = slide.shapes.add_shape(1, Inches(1), Inches(2), Inches(4), Inches(1))
    label.name = "deleted-shape"
    label.text = text
    source = tmp_path / "shape-label-whitespace.pptx"
    presentation.save(source)
    native_text = label.text
    source_bytes = source.read_bytes()
    original_deck = deck.model_dump_json()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")

    labels = [obj for obj in imported.slides[0].objects if isinstance(obj, ShapeObject)]
    assert len(labels) == 1
    assert labels[0].text == native_text == text
    assert (labels[0].id == "deleted-shape") == known
    title_object = next(obj for obj in imported.slides[0].objects if obj.id == "stable-title")
    assert title_object.text == "合成标题"
    assert source.read_bytes() == source_bytes
    assert deck.model_dump_json() == original_deck


def test_native_wide_reimport_does_not_rescale_textbox_width(tmp_path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = 12192000  # PptxGenJS LAYOUT_WIDE, 13 1/3 inches
    presentation.slide_height = Inches(7.5)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(6237580, 1993392, 5167732, 3383280)
    shape.name = "stable-title"
    shape.text = "在证据固定、口径确认与法务审查之后，采取必要法律措施。"
    source = tmp_path / "wide.pptx"
    presentation.save(source)
    result = reconcile_pptx(_deck(), source, assets_dir=tmp_path / "assets")
    frame = result.slides[0].objects[0].frame
    assert round(frame.x * 914400) == shape.left
    assert round(frame.w * 914400) == shape.width


def test_reconcile_does_not_restore_stale_rich_text(tmp_path: Path) -> None:
    from lantai.presentation.models import RichTextRun

    deck = _deck()
    deck.slides[0].objects[0].runs = [RichTextRun(text='保存前标题', bold=True)]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.name = 'stable-title'
    shape.text = '人工改过的标题'
    source = tmp_path / 'changed.pptx'
    presentation.save(source)
    reconciled = reconcile_pptx(deck, source, assets_dir=tmp_path / 'assets')
    title = reconciled.slides[0].objects[0]
    assert title.text == '人工改过的标题'
    assert ''.join(run.text for run in title.runs) == '人工改过的标题'
    assert deck.slides[0].objects[0].runs[0].text == '保存前标题'


def test_reconcile_rejects_group_without_dropping_children(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    shape.text = '分组内部的用户证据'
    source = tmp_path / 'grouped.pptx'
    presentation.save(source)
    original = source.read_bytes()
    assets = tmp_path / 'assets'
    with pytest.raises(ValueError, match='unsupported native group'):
        reconcile_pptx(_deck(), source, assets_dir=assets)
    assert source.read_bytes() == original
    assert not assets.exists()


@pytest.mark.parametrize('known', [False, True])
def test_reconcile_preserves_native_chart_type_and_values(tmp_path: Path, known: bool) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from lantai.presentation.models import ChartObject

    deck = _deck()
    if known:
        deck.slides[0].objects.append(ChartObject(
            id='editable-chart', chart_type='bar', categories=['旧数据'],
            series={'数量': [1]}, source_refs=['original-source'],
        ))
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['六月', '七月']
    data.add_series('数量', [4557, 4608])
    shape = slide.shapes.add_chart(
        XL_CHART_TYPE.LINE, Inches(1), Inches(1), Inches(8), Inches(4), data,
    )
    shape.name = 'editable-chart'
    source = tmp_path / 'chart.pptx'
    presentation.save(source)
    original = source.read_bytes()
    result = reconcile_pptx(deck, source, assets_dir=tmp_path / 'assets')
    chart = result.slides[0].objects[0]
    assert chart.chart_type == 'line'
    assert chart.categories == ['六月', '七月']
    assert chart.series == {'数量': [4557, 4608]}
    if known:
        assert chart.id == 'editable-chart'
        assert chart.source_refs == ['original-source']
    assert source.read_bytes() == original


def test_reconcile_rejects_stacked_chart_before_import(tmp_path: Path) -> None:
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['六月']
    data.add_series('数量', [4557])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_STACKED, Inches(1), Inches(1), Inches(8), Inches(4), data,
    )
    source = tmp_path / 'stacked.pptx'
    presentation.save(source)
    original = source.read_bytes()
    with pytest.raises(ValueError, match='unsupported native chart'):
        reconcile_pptx(_deck(), source, assets_dir=tmp_path / 'assets')
    assert source.read_bytes() == original
    assert not (tmp_path / 'assets').exists()


def test_inserted_and_duplicated_slides_do_not_steal_page_identity(tmp_path: Path) -> None:
    presentation = Presentation()
    for text, marked in [('新增页', False), ('原来的页', True), ('复制的页', True)]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        if marked:
            marker = slide.shapes.add_textbox(Inches(0), Inches(0), Inches(.01), Inches(.01))
            marker.name = 'meta:reconcile-deck:stable-slide'
        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        title.name = 'stable-title'
        title.text = text
    source = tmp_path / 'inserted.pptx'
    presentation.save(source)
    result = reconcile_pptx(_deck(), source, assets_dir=tmp_path / 'assets')
    assert len({slide.id for slide in result.slides}) == 3
    assert result.slides[1].id == 'stable-slide'
    assert result.slides[1].objects[0].id == 'stable-title'
    assert result.slides[0].id != 'stable-slide'
    assert result.slides[2].id != 'stable-slide'
    assert [slide.objects[0].text for slide in result.slides] == ['新增页', '原来的页', '复制的页']


def test_reconcile_reads_format_changes_even_when_text_is_unchanged(tmp_path: Path) -> None:
    from pptx.dml.color import RGBColor
    from pptx.util import Pt
    from lantai.presentation.models import RichTextRun

    deck = _deck()
    deck.slides[0].objects[0].runs = [RichTextRun(text='保存前标题', bold=True)]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    title.name = 'stable-title'
    title.text = '保存前标题'
    run = title.text_frame.paragraphs[0].runs[0]
    run.font.bold = False
    run.font.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string('CC3300')
    source = tmp_path / 'format.pptx'
    presentation.save(source)
    result = reconcile_pptx(deck, source, assets_dir=tmp_path / 'assets')
    imported = result.slides[0].objects[0].runs[0]
    assert imported.bold is False
    assert imported.italic is True
    assert imported.font_size == 12
    assert imported.color == 'CC3300'
    assert deck.slides[0].objects[0].runs[0].bold is True


def test_reconcile_does_not_turn_unsupported_shape_into_rectangle(tmp_path: Path) -> None:
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.STAR_5_POINT, Inches(1), Inches(1), Inches(2), Inches(2))
    source = tmp_path / 'star.pptx'
    presentation.save(source)
    original = source.read_bytes()
    with pytest.raises(ValueError, match='unsupported native shape'):
        reconcile_pptx(_deck(), source, assets_dir=tmp_path / 'assets')
    assert source.read_bytes() == original


def test_reconcile_preserves_line_preset_emitted_by_native_backend(tmp_path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(5), Inches(.01))
    shape.name = 'native-line'
    shape._element.xpath('./p:spPr/a:prstGeom')[0].set('prst', 'line')
    source = tmp_path / 'line.pptx'
    presentation.save(source)
    result = reconcile_pptx(_deck(), source, assets_dir=tmp_path / 'assets')
    assert result.slides[0].objects[0].shape_type == 'line'


def test_metric_reimport_does_not_duplicate_label_and_detail(tmp_path: Path) -> None:
    from lantai.presentation.models import MetricObject

    deck = _deck()
    deck.slides[0].objects = [MetricObject(id='metric', value='4557', label='讨论量', detail='六月')]
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    shape.name = 'metric'
    shape.text = '4608\n讨论量\n七月'
    source = tmp_path / 'metric.pptx'
    presentation.save(source)
    result = reconcile_pptx(deck, source, assets_dir=tmp_path / 'assets')
    imported = result.slides[0].objects[0]
    assert isinstance(imported, TextObject)
    assert imported.id == 'metric'
    assert imported.text == '4608\n讨论量\n七月'
    assert ''.join(run.text for run in imported.runs) == imported.text
    assert deck.slides[0].objects[0].value == '4557'


def test_reconcile_preserves_added_shape_cleared_text_and_edited_notes(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    cleared = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    cleared.name = "deleted-shape"
    cleared.text = ""
    cleared.fill.background()
    extra = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5), Inches(2), Inches(2), Inches(1))
    extra.text = "人工新增图形文字"
    extra.fill.solid()
    extra.fill.fore_color.rgb = RGBColor.from_string("123456")
    extra.rotation = 15
    slide.notes_slide.notes_text_frame.text = "人工修改的演讲备注"
    source = tmp_path / "human.pptx"
    presentation.save(source)
    deck = _deck()
    deck.slides[0].objects[1].text = "应当清空的旧文字"
    deck.slides[0].speaker_notes = "旧备注"
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].speaker_notes == "人工修改的演讲备注"
    known = next(o for o in imported.slides[0].objects if o.id == "deleted-shape")
    assert known.text == ""
    assert known.fill_color is None
    added = next(o for o in imported.slides[0].objects if o.id != "deleted-shape")
    assert isinstance(added, ShapeObject)
    assert added.shape_type == "ellipse"
    assert added.text == "人工新增图形文字"
    assert added.fill_color == "123456"
    assert added.rotate == 15
