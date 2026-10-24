# -*- coding: utf-8 -*-
# pylint: disable=protected-access
from __future__ import annotations

import threading
import hashlib
import json
from contextlib import contextmanager
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from lantai.app.structured_utility import StructuredTaskValidationError, StructuredUtilityResult
from lantai.exceptions import ProviderError
from lantai.presentation.design_system import prepare_deck_for_composition
from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.failure import failure_route_for_exception
from lantai.presentation.jobs import (
    PresentationExecutionLimits,
    PresentationJobInterruptedError,
    PresentationJobStore,
    cancel_active_job_for_runtime_run,
)
from lantai.presentation.intelligence import UtilityPresentationObjectEditPlanner, UtilityPresentationVisualCritic
from lantai.presentation.models import (
    CompositionEngine,
    DeckIR,
    DeckPlan,
    DeckPlanSlide,
    DeckSlide,
    FactRecord,
    JobStage,
    JobStatus,
    InputMaterial,
    ObjectFrame,
    ObjectGroup,
    ObjectRelationship,
    PlanningMode,
    PresentationObjectEditPlan,
    PresentationBrief,
    PresentationSourceSection,
    PresentationSourceCoverage,
    PresentationProviderCostStatus,
    PresentationJobRequest,
    PresentationStageMetric,
    QualityIssue,
    QualityReport,
    RelationshipType,
    ReplaceTextAction,
    RenderedSlideEvidence,
    SlideIntent,
    SlideIntentPatch,
    SourceCalculationProof,
    ShapeObject,
    TextObject,
    VisualCriticResult,
    VisualDirection,
    VisualIssue,
)
from lantai.presentation.repair import (
    RepairActionType,
    SlideRedesignProposal,
)
from lantai.presentation.quality import file_sha256, is_verified_delivery
from lantai.presentation.planner import build_brief
from lantai.presentation.service import PresentationService
from lantai.presentation.theme import THEMES
from lantai.presentation.runtime import PresentationJobRuntime
from lantai.presentation.render import RenderResult
from lantai.presentation.render import PresentationRenderUnavailable
from lantai.presentation.versioning import PresentationVersionStore
from tests.unit.presentation.visual_review_fixtures import synthetic_review_response


def _request(
    conversation_id: str = "conversation-a",
) -> PresentationJobRequest:
    return PresentationJobRequest(
        conversation_id=conversation_id,
        user_request="生成一份 6 页的产品汇报",
        slide_count=6,
        vision_critique_required=False,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        composition_engine=CompositionEngine.LEGACY_PYTHON_PPTX,
        materials=[
            InputMaterial(
                id="source-test-outline",
                name="outline.md",
                media_type="text/markdown",
                content=(
                    "# 产品汇报\n\n"
                    "## 第1页 封面\n- 产品汇报\n- 内部材料\n\n"
                    "## 第2页 背景\n- 用户需要可靠的演示文稿交付\n\n"
                    "## 第3页 问题\n- 交付前必须完成内容检查\n\n"
                    "## 第4页 方案\n- 输出保持原生可编辑\n\n"
                    "## 第5页 验证\n- 全量渲染后检查每一页\n\n"
                    "## 第6页 结论\n- 只有验证通过才可交付"
                ),
            ),
        ],
    )


def _render_for_edit_tests(
    source: str | Path,
    output_dir: str | Path,
    *,
    timeout: int = 120,
) -> RenderResult:
    _ = timeout
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"test pdf")
    count = len(Presentation(str(source)).slides)
    images: list[Path] = []
    for index in range(1, count + 1):
        image_path = destination / f"slide-{index}.png"
        image = Image.new("RGB", (640, 360), "white")
        image.paste("black", (0, 0, 280 + index, 360))
        image.save(image_path)
        images.append(image_path)
    return RenderResult(
        pdf_path=pdf_path,
        slide_images=tuple(images),
        renderer_id="test-renderer",
        renderer_version="1.0",
    )


def test_services_share_default_execution_capacity_across_jobs(
    tmp_path: Path,
) -> None:
    first = PresentationService(tmp_path / "first")
    second = PresentationService(tmp_path / "second")

    assert first.limits is second.limits


def test_stage_metric_never_serializes_unknown_provider_cost_as_zero() -> None:
    now = datetime.now(timezone.utc)
    metric = PresentationStageMetric(
        stage=JobStage.PLAN,
        operation="plan",
        repair_round=0,
        route_id="provider/model",
        outcome="passed",
        duration_ms=12.5,
        cost_units=10,
        provider_cost_status=PresentationProviderCostStatus.UNKNOWN,
        started_at=now,
        finished_at=now,
    )

    payload = metric.model_dump(mode="json")

    assert payload["schema_version"] == "1.1"
    assert payload["provider_cost_status"] == "unknown"
    assert payload["provider_cost_micro_cny"] is None
    assert payload["cost_units"] == 10


def test_stage_metric_requires_exact_status_for_explicit_free_provider_call(
) -> None:
    now = datetime.now(timezone.utc)
    common = {
        "stage": JobStage.QUALITY,
        "operation": "critic",
        "repair_round": 0,
        "route_id": "provider/model",
        "outcome": "passed",
        "duration_ms": 5.0,
        "cost_units": 1,
        "started_at": now,
        "finished_at": now,
    }

    exact_free = PresentationStageMetric(
        **common,
        provider_cost_status=PresentationProviderCostStatus.EXACT,
        provider_cost_micro_cny=0,
    )
    assert exact_free.provider_cost_micro_cny == 0
    with pytest.raises(ValueError, match="unknown provider cost cannot have"):
        PresentationStageMetric(
            **common,
            provider_cost_status=PresentationProviderCostStatus.UNKNOWN,
            provider_cost_micro_cny=0,
        )
    with pytest.raises(ValueError, match="exact provider cost requires"):
        PresentationStageMetric(
            **common,
            provider_cost_status=PresentationProviderCostStatus.EXACT,
        )


def test_stage_metric_dual_reads_v1_without_inventing_provider_cost() -> None:
    started_at = datetime.now(timezone.utc)
    metric = PresentationStageMetric.model_validate(
        {
            "schema_version": "1.0",
            "stage": "plan",
            "operation": "plan",
            "outcome": "passed",
            "duration_ms": 1.0,
            "cost_units": 10,
            "route_id": "historical-provider/model",
            "started_at": started_at.isoformat(),
            "finished_at": started_at.isoformat(),
        },
    )

    assert metric.schema_version == "1.0"
    assert (
        metric.provider_cost_status
        is PresentationProviderCostStatus.NOT_APPLICABLE
    )
    assert metric.provider_cost_micro_cny is None


def test_service_can_inject_execution_capacity_for_isolated_runs(
    tmp_path: Path,
) -> None:
    isolated = PresentationExecutionLimits(
        synthesis=1,
        rendering=1,
        critics=1,
    )

    service = PresentationService(
        tmp_path / "isolated",
        execution_limits=isolated,
    )

    assert service.limits is isolated


def test_artifact_query_does_not_create_an_absent_version_ledger(
    tmp_path: Path,
) -> None:
    runtime = PresentationJobRuntime(tmp_path)
    try:
        record = runtime.service.create_job(_request())
        job_dir = runtime.service.store.job_dir(record.job_id)
        (job_dir / "versions").rmdir()
        before = {
            path.relative_to(job_dir)
            for path in job_dir.rglob("*")
        }

        artifacts = runtime.artifacts(record.job_id, "conversation-a")

        after = {
            path.relative_to(job_dir)
            for path in job_dir.rglob("*")
        }
        assert "presentation_version" not in artifacts
        assert after == before
        assert not (job_dir / "versions").exists()
    finally:
        runtime.shutdown()


def test_shared_execution_capacity_serializes_critics_across_services(
    tmp_path: Path,
) -> None:
    shared = PresentationExecutionLimits(
        synthesis=1,
        rendering=1,
        critics=1,
    )
    first = PresentationService(
        tmp_path / "first-critic",
        execution_limits=shared,
    )
    second = PresentationService(
        tmp_path / "second-critic",
        execution_limits=shared,
    )
    first_entered = threading.Event()
    release_first = threading.Event()
    second_entered = threading.Event()

    def hold_first_slot() -> None:
        with first.limits.critic():
            first_entered.set()
            release_first.wait(timeout=2)

    def enter_second_slot() -> None:
        with second.limits.critic():
            second_entered.set()

    first_thread = threading.Thread(target=hold_first_slot)
    second_thread = threading.Thread(target=enter_second_slot)
    first_thread.start()
    assert first_entered.wait(timeout=1)
    second_thread.start()

    assert not second_entered.wait(timeout=0.05)
    release_first.set()
    first_thread.join(timeout=1)
    second_thread.join(timeout=1)

    assert not first_thread.is_alive()
    assert not second_thread.is_alive()
    assert second_entered.is_set()


def _completed_editable_runtime(
    tmp_path,
    monkeypatch,
    *,
    include_third_slide: bool = False,
    include_shape_labels: bool = False,
    page_count: int | None = None,
    composition_engine: CompositionEngine = CompositionEngine.LEGACY_PYTHON_PPTX,
) -> tuple[PresentationJobRuntime, str, DeckIR, Path]:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    runtime = PresentationJobRuntime(tmp_path)
    record = runtime.service.create_job(_request().model_copy(update={
        "composition_engine": composition_engine,
    }))
    deck = DeckIR(
        id="natural-edit-deck",
        title="自然语言编辑",
        slides=[
            DeckSlide(
                id="slide-a",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="title-a",
                        role="title",
                        text="原始标题",
                        slot="title",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-b",
                layout_id="closing",
                objects=[
                    TextObject(
                        id="title-b",
                        role="title",
                        text="保持不变",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    if include_third_slide:
        deck.slides.append(DeckSlide(
            id="slide-c", layout_id="closing",
            objects=[TextObject(id="title-c", role="title", text="未指定页必须保持")],
        ))
    if page_count is not None:
        assert page_count >= len(deck.slides)
        for index in range(len(deck.slides), page_count):
            deck.slides.append(DeckSlide(
                id=f"slide-extra-{index}", layout_id="closing",
                objects=[TextObject(id=f"title-extra-{index}", role="title", text=f"Original {index}")],
            ))
    if include_shape_labels:
        for slide in deck.slides:
            slide.layout_id = "freeform"
            slide.objects[0].frame = ObjectFrame(x=0.8, y=1, w=11.2, h=1.4)
            suffix = slide.id.removeprefix("slide-")
            label_id = f"label-{suffix}"
            slide.objects.append(ShapeObject(
                id=label_id, shape_type="round_rect", text=f"原始标签-{suffix}",
                frame=ObjectFrame(x=2, y=4.2, w=4, h=1),
                fill_color="ABCDEF", line_color="102030", text_color="304050",
                font_size=24, rotate=7,
            ))
            slide.object_groups = [ObjectGroup(
                group_id=f"group-{suffix}", role="label",
                object_ids=[f"title-{suffix}", label_id],
            )]
            slide.relationships = [ObjectRelationship(
                relationship_id=f"relation-{suffix}",
                relationship_type=RelationshipType.PROCESS,
                source_object_ids=[f"title-{suffix}"], target_object_ids=[label_id],
                label="保留原关系",
            )]
            slide.speaker_notes = "保留原备注"
    if composition_engine == CompositionEngine.NATIVE_PPTX:
        for slide in deck.slides:
            for obj in slide.objects:
                if isinstance(obj, TextObject):
                    obj.font_size = 48
    deck_path = runtime.service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    output = (
        runtime.service.store.job_dir(record.job_id)
        / "output"
        / "presentation.pptx"
    )
    runtime.service.compiler_for(composition_engine).compile(deck, output)
    report, _images = runtime.service.verify_artifact(
        deck,
        output,
        runtime.service.store.job_dir(record.job_id) / "renders" / "round-0",
        render_set_id="round-0",
    )
    if include_shape_labels:
        assert report.passed, [(issue.code, issue.slide_id) for issue in report.issues]
    quality_path = (
        runtime.service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    runtime.service.store.append_event(
        record.job_id,
        stage=JobStage.COMPLETED,
        status=JobStatus.COMPLETED,
        message="verified baseline",
        checkpoint=str(output),
    )
    assert deck_path.is_file()
    return runtime, record.job_id, deck, output


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("failure", [
    "cancel_after_review", "candidate_quality_write", "record_then_fail",
    "completed_event_write", "current_replace",
])
@pytest.mark.parametrize("edit_kind", ["replace_text", "set_shape_text"])
def test_redo_slide_commit_tail_keeps_file_and_version_ledger_atomic(
    tmp_path, monkeypatch, engine, failure, edit_kind,
):
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
        include_shape_labels=edit_kind == "set_shape_text",
    )
    store = runtime.service.store
    versions = PresentationVersionStore(store.job_dir(job_id))
    deck_path = store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    versions.bootstrap(output, deck, writer=engine.value, quality_report_path=quality_path,
                       render_dir=store.job_dir(job_id) / "renders/round-0")
    before = {path: path.read_bytes() for path in (
        output, deck_path, quality_path, versions.current_path, versions.history_path,
    )}
    base = versions.current()
    baseline_sequence = store.get(job_id).last_sequence

    class Planner:
        def plan(self, *args, **kwargs):
            target = "title" if edit_kind == "replace_text" else "label"
            return PresentationObjectEditPlan.model_validate({
                "slide_id": "slide-a", "additional_slide_ids": ["slide-b"],
                "actions": [
                    {"kind": edit_kind, "object_id": f"{target}-a", "text": "候选第一页"},
                    {"kind": edit_kind, "object_id": f"{target}-b", "text": "候选第二页"},
                ],
            })

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    if failure == "cancel_after_review":
        original_verify = runtime.service.verify_artifact

        def verify_then_cancel(*args, **kwargs):
            report, images = original_verify(*args, **kwargs)
            assert report.passed
            runtime.cancel(job_id, "conversation-a")
            return report, images

        monkeypatch.setattr(runtime.service, "verify_artifact", verify_then_cancel)
    elif failure == "candidate_quality_write":
        original_write_checkpoint = store.write_checkpoint

        def fail_candidate_diagnostic(job, name, content):
            if name.startswith("edit-quality-report-"):
                raise OSError("synthetic candidate diagnostic write failure")
            return original_write_checkpoint(job, name, content)

        monkeypatch.setattr(store, "write_checkpoint", fail_candidate_diagnostic)
    elif failure == "record_then_fail":
        original_record = PresentationVersionStore.record

        def append_then_fail(self, *args, **kwargs):
            result = original_record(self, *args, **kwargs)
            if kwargs.get("writer") == "lantai_ai_edit":
                raise OSError("synthetic error after version append")
            return result

        monkeypatch.setattr(PresentationVersionStore, "record", append_then_fail)
    else:
        failed_once = False
        if failure == "completed_event_write":
            original_write = store._write

            def fail_completed_write(record):
                nonlocal failed_once
                if (record.status == JobStatus.COMPLETED
                        and record.last_sequence > baseline_sequence and not failed_once):
                    failed_once = True
                    raise OSError("synthetic completed manifest error")
                return original_write(record)

            monkeypatch.setattr(store, "_write", fail_completed_write)
        else:
            import lantai.presentation.versioning as versioning_module
            original_replace = versioning_module.os.replace

            def fail_current_replace(source, destination):
                nonlocal failed_once
                if Path(destination) == versions.current_path and not failed_once:
                    failed_once = True
                    raise OSError("synthetic current version replace error")
                return original_replace(source, destination)

            monkeypatch.setattr(versioning_module.os, "replace", fail_current_replace)
    try:
        with pytest.raises((PresentationJobInterruptedError, OSError)):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "只改这两页",
                               file_sha256(output), additional_slide_ids=["slide-b"])
        assert {path: path.read_bytes() for path in before} == before
        assert versions.current() == base
        assert versions.assert_writer_base(output, deck) == base
        expected_status = JobStatus.CANCELLED if failure == "cancel_after_review" else JobStatus.COMPLETED
        current = store.get(job_id)
        assert current.status == expected_status
        if failure == "cancel_after_review":
            assert not any(event.status == JobStatus.COMPLETED for event in current.events
                           if event.sequence > baseline_sequence)
        if failure in {"cancel_after_review", "candidate_quality_write"}:
            assert not list((store.job_dir(job_id) / "checkpoints").glob("edit-quality-report-*.json"))
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
def test_redo_slide_shape_labels_publish_once_with_native_identity_and_scope(
    tmp_path, monkeypatch, engine,
):
    from zipfile import ZipFile

    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, include_shape_labels=True,
        composition_engine=engine,
    )
    calls = []
    effective_bases = []
    original_native = Presentation(str(output))
    old_frames = {
        shape.name: (shape.left, shape.top, shape.width, shape.height, shape.rotation)
        for slide in original_native.slides for shape in slide.shapes
        if shape.name in {"label-a", "label-b"}
    }
    with ZipFile(output) as package:
        unselected_bytes = package.read("ppt/slides/slide3.xml")

    class Planner:
        def plan(self, current, slide_id, instruction, *, deck_plan=None, additional_slide_ids=None):
            calls.append((slide_id, additional_slide_ids))
            effective_bases.append(current.model_copy(deep=True))
            return PresentationObjectEditPlan.model_validate({
                "slide_id": slide_id, "additional_slide_ids": ["slide-b"],
                "actions": [
                    {"kind": "set_shape_text", "object_id": "label-a", "text": "合成主页面标签"},
                    {"kind": "set_shape_text", "object_id": "label-b", "text": "合成附加页标签"},
                ],
            })

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    try:
        before = file_sha256(output)
        updated = runtime.redo_slide(
            job_id, "conversation-a", "slide-a", "只改明确两页标签", before,
            additional_slide_ids=["slide-b"],
        )
        assert updated.status == JobStatus.COMPLETED
        assert calls == [("slide-a", ["slide-b"])]
        # Native redo reads the authoritative PPTX first, importing run
        # formatting and a reconciliation revision. The edit must preserve
        # that effective base, not erase it to match the pre-export IR.
        effective_base = effective_bases[0]
        for original, imported in zip(deck.slides, effective_base.slides):
            assert [obj.id for obj in imported.objects] == [obj.id for obj in original.objects]
            assert imported.objects[0].text == original.objects[0].text
            assert imported.objects[1].text == original.objects[1].text
            assert imported.speaker_notes.strip() == original.speaker_notes
        changed = DeckIR.model_validate_json(
            runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_text(),
        )
        for index, text in enumerate(["合成主页面标签", "合成附加页标签"]):
            old, new = effective_base.slides[index], changed.slides[index]
            assert new.objects[0] == old.objects[0]
            assert new.objects[1].model_dump() == {**old.objects[1].model_dump(), "text": text}
            assert new.object_groups == old.object_groups
            assert new.relationships == old.relationships
            assert new.speaker_notes == old.speaker_notes
            assert new.revision == old.revision + 1
        assert changed.slides[2] == effective_base.slides[2]
        assert changed.source_refs == deck.source_refs
        assert changed.required_facts == deck.required_facts
        assert changed.immutable_data == deck.immutable_data
        with ZipFile(output) as package:
            assert package.read("ppt/slides/slide3.xml") == unselected_bytes
        reopened = Presentation(str(output))
        frames = {
            shape.name: (shape.left, shape.top, shape.width, shape.height, shape.rotation)
            for slide in reopened.slides for shape in slide.shapes
            if shape.name in old_frames
        }
        assert frames == old_frames
        visible = [shape.text for slide in reopened.slides for shape in slide.shapes
                   if shape.has_text_frame]
        assert "合成主页面标签" in visible and "合成附加页标签" in visible
        assert "原始标签-a" not in visible and "原始标签-b" not in visible
        assert "原始标签-c" in visible
        artifacts = runtime.artifacts(job_id, "conversation-a")
        assert len(artifacts["presentation_versions"]) == 2
        assert artifacts["presentation_version"]["writer"] == "lantai_ai_edit"
        assert artifacts["quality"]["output_sha256"] == file_sha256(output) != before
        assert artifacts["quality"]["visual_verified"] is True
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("text", ["合成第二轮\n仍可编辑", ""])
def test_redo_shape_label_reimports_previous_edit_then_replaces_or_clears(
    tmp_path, monkeypatch, engine, text,
):
    from zipfile import ZipFile

    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, include_shape_labels=True,
        composition_engine=engine,
    )
    with ZipFile(output) as package:
        unselected = {name: package.read(name) for name in (
            "ppt/slides/slide2.xml", "ppt/slides/slide3.xml",
        )}
    calls = []
    texts = ["合成第一轮", text]

    class Planner:
        def plan(self, current, slide_id, instruction, *, deck_plan=None):
            label = next(obj for obj in current.slides[0].objects if obj.id == "label-a")
            calls.append(label.text)
            return PresentationObjectEditPlan.model_validate({
                "slide_id": slide_id,
                "actions": [{"kind": "set_shape_text", "object_id": "label-a",
                             "text": texts[len(calls) - 1]}],
            })

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    try:
        for expected in texts:
            result = runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "连续编辑明确形状标签",
                file_sha256(output),
            )
            assert result.status == JobStatus.COMPLETED
            checkpoint = DeckIR.model_validate_json(
                runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_text(),
            )
            labels = [obj for obj in checkpoint.slides[0].objects if obj.id == "label-a"]
            assert len(labels) == 1 and isinstance(labels[0], ShapeObject)
            assert labels[0].text == expected
            native = [shape for shape in Presentation(str(output)).slides[0].shapes
                      if shape.name == "label-a"]
            # Legacy OOXML uses a soft line break (python-pptx's vertical
            # tab); compare the visible wording without changing the writer.
            assert len(native) == 1 and native[0].text.replace("\v", "\n") == expected
            with ZipFile(output) as package:
                assert {name: package.read(name) for name in unselected} == unselected
            artifacts = runtime.artifacts(job_id, "conversation-a")
            assert artifacts["quality"]["output_sha256"] == file_sha256(output)
            assert artifacts["quality"]["visual_verified"] is True
        assert calls == ["原始标签-a", "合成第一轮"]
        assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 3
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("invalid", ["wrong_kind", "line", "unselected"])
def test_redo_slide_shape_label_invalid_targets_leave_entire_base_unchanged(
    tmp_path, monkeypatch, engine, invalid,
):
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, include_shape_labels=True,
        composition_engine=engine,
    )
    store = runtime.service.store
    versions = PresentationVersionStore(store.job_dir(job_id))
    deck_path = store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    if invalid == "line":
        deck.slides[1].objects[1].shape_type = "line"
        deck.slides[1].objects[1].text = ""
        store.write_checkpoint(job_id, "deck-ir.json", deck.model_dump_json(indent=2))
        runtime.service.compiler_for(engine).compile(deck, output)
        report, _images = runtime.service.verify_artifact(
            deck, output, store.job_dir(job_id) / "renders/round-0", render_set_id="round-0",
        )
        assert report.passed, [(issue.code, issue.slide_id) for issue in report.issues]
        quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    versions.bootstrap(output, deck, writer=engine.value, quality_report_path=quality_path,
                       render_dir=store.job_dir(job_id) / "renders/round-0")
    before = {path: path.read_bytes() for path in (
        output, deck_path, quality_path, versions.current_path, versions.history_path,
    )}
    base = versions.current()
    second_target = {"wrong_kind": "title-b", "line": "label-b", "unselected": "label-c"}[invalid]

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan.model_validate({
                "slide_id": "slide-a", "additional_slide_ids": ["slide-b"],
                "actions": [
                    {"kind": "set_shape_text", "object_id": "label-a", "text": "不得部分提交"},
                    {"kind": "set_shape_text", "object_id": second_target, "text": "非法目标"},
                ],
            })

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    try:
        message = "unknown object id" if invalid == "unselected" else "not an editable shape label"
        with pytest.raises(ValueError, match=message):
            runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "拒绝越界与错误类型", file_sha256(output),
                additional_slide_ids=["slide-b"],
            )
        assert {path: path.read_bytes() for path in before} == before
        assert versions.current() == base
        assert versions.assert_writer_base(output, deck) == base
        assert store.get(job_id).status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


def _source_correction_runtime(tmp_path, monkeypatch):
    """A native baseline with a generated lock, not an approved user fact."""
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, composition_engine=CompositionEngine.NATIVE_PPTX,
    )
    old, new, ref = "原始标题", "材料倾向B。", "synthetic-source"
    body = new + "\n保持不变"
    digest = hashlib.sha256(body.encode()).hexdigest()
    coverage = PresentationSourceCoverage(
        material_id=ref, source_name="Synthetic", source_sha256=digest,
        method="inline", status="supported_content_extracted", extracted_chars=len(body),
    )
    deck.required_facts = [old, "保持不变"]
    deck.immutable_data = ["保持不变"]
    deck.source_refs = [ref]
    deck.source_coverage = [coverage]
    deck.facts = [FactRecord(fact_id="generated", text=old, source_refs=[ref]),
                  FactRecord(fact_id="fixed", text="保持不变", source_refs=[ref], immutable=True)]
    for slide, fact_id in zip(deck.slides, ["generated", "fixed"], strict=True):
        slide.intent_id = "intent-" + slide.id
        slide.source_fact_ids = [fact_id]
    plan = DeckPlan(
        title=deck.title, narrative_arc="核对来源",
        visual_direction=VisualDirection(
            visual_concept="核对", palette=["111111", "FFFFFF"], typography="清晰",
            image_language="无", chart_language="原生", shape_language="克制",
            grid="十二列", recurring_motif="细线",
        ),
        slides=[DeckPlanSlide(order=index + 1, title=slide.objects[0].text,
            content_blocks=[slide.objects[0].text], intent=SlideIntent(
                id=slide.intent_id, purpose="核对", key_message=slide.objects[0].text,
                required_facts=[slide.objects[0].text], evidence=[slide.objects[0].text],
                immutable_facts=["保持不变"] if index else [], source_refs=[ref],
            )) for index, slide in enumerate(deck.slides)],
    )
    brief = PresentationBrief(
        objective="核对来源", audience="团队", narrative="合成测试", source_refs=[ref],
        required_facts=["保持不变"], immutable_data=["保持不变"], source_coverage=[coverage],
        source_sections=[PresentationSourceSection(title="讨论", body=body, source_ref=ref)],
    )
    store = runtime.service.store
    for name, model in [("deck-ir.json", deck), ("deck-plan.json", plan), ("brief.json", brief)]:
        store.write_checkpoint(job_id, name, model.model_dump_json(indent=2))
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    report, _ = runtime.service.verify_artifact(deck, output, store.job_dir(job_id) / "renders/round-0",
                                               render_set_id="round-0")
    assert report.passed
    quality_path.write_text(report.model_dump_json(indent=2))
    versions = PresentationVersionStore(store.job_dir(job_id))
    versions.bootstrap(output, deck, writer="native_pptx", quality_report_path=quality_path,
                       render_dir=store.job_dir(job_id) / "renders/round-0")
    correction = dict(original_fact=old, corrected_fact=new, source_ref=ref,
                      source_sha256=digest, source_section_sha256=digest)
    return runtime, job_id, deck, output, versions, correction


@pytest.mark.parametrize("failure", [None, "review", "stop", "quality_commit", "record_then_fail"])
def test_source_fact_correction_is_native_atomic_and_whole_deck_reviewed(tmp_path, monkeypatch, failure):
    from zipfile import ZipFile
    runtime, job_id, deck, output, versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    store = runtime.service.store
    paths = [output, versions.current_path, versions.history_path,
             store.checkpoint_path(job_id, "deck-ir.json"), store.checkpoint_path(job_id, "deck-plan.json"),
             store.checkpoint_path(job_id, "brief.json"), store.job_dir(job_id) / "quality/quality-report.json"]
    before = {path: path.read_bytes() for path in paths}
    reviewed = []

    class Planner:
        def plan(self, current, slide_id, instruction, *, deck_plan, source_brief):
            assert current.required_facts == [correction["corrected_fact"], "保持不变"]
            assert deck_plan.slides[0].intent.required_facts == [correction["corrected_fact"]]
            assert source_brief.required_facts == ["保持不变"]
            return PresentationObjectEditPlan(slide_id=slide_id, actions=[
                ReplaceTextAction(object_id="title-a", text=correction["corrected_fact"]),
            ])

    class Critic:
        def review(self, slides):
            reviewed.extend(slide.slide_id for slide in slides)
            if failure == "stop":
                raise PresentationJobInterruptedError("synthetic user stop")
            return VisualCriticResult(critic_id="test-source-correction/1", issues=[
                VisualIssue(slide_id="slide-b", severity="warning", issue_type="content_accuracy",
                            description="Unselected page fails source review", proposed_action="Keep baseline"),
            ] if failure == "review" else [])

    runtime.object_edit_planner, runtime.service.visual_critic = Planner(), Critic()
    if failure == "quality_commit":
        original_write, failed = runtime._atomic_write_bytes, False

        def fail_quality(path, content):
            nonlocal failed
            if path == paths[-1] and not failed:
                failed = True
                raise OSError("synthetic quality commit failure")
            original_write(path, content)

        monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_quality)
    elif failure == "record_then_fail":
        original_record = PresentationVersionStore.record

        def record_then_fail(self, *args, **kwargs):
            result = original_record(self, *args, **kwargs)
            if kwargs.get("writer") == "lantai_ai_edit":
                raise OSError("synthetic error after version append")
            return result

        monkeypatch.setattr(PresentationVersionStore, "record", record_then_fail)
    with ZipFile(output) as package:
        untouched = package.read("ppt/slides/slide2.xml")
    try:
        def edit():
            return runtime.redo_slide(job_id, "conversation-a", "slide-a", "按原材料更正生成锁",
                                      file_sha256(output), source_fact_corrections=[correction])
        if failure is None:
            assert edit().status == JobStatus.COMPLETED
            changed = DeckIR.model_validate_json(paths[3].read_text())
            assert changed.required_facts == [correction["corrected_fact"], "保持不变"]
            assert len(changed.facts) == len(deck.facts)
            assert changed.immutable_data == deck.immutable_data
            assert len(changed.source_fact_corrections) == 1
            assert paths[5].read_bytes() == before[paths[5]]
            assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 2
        else:
            expected = {"review": ValueError, "stop": PresentationJobInterruptedError,
                        "quality_commit": OSError, "record_then_fail": OSError}
            with pytest.raises(expected[failure]):
                edit()
            assert {path: path.read_bytes() for path in paths} == before
        assert set(reviewed) == {"slide-a", "slide-b"} if failure != "stop" else bool(reviewed)
        with ZipFile(output) as package:
            assert package.read("ppt/slides/slide2.xml") == untouched
    finally:
        runtime.shutdown()


def test_source_fact_correction_invalid_pin_consumes_no_attempt_or_provider_call(tmp_path, monkeypatch):
    runtime, job_id, _deck, output, _versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    calls = []
    runtime.object_edit_planner = SimpleNamespace(plan=lambda *a, **k: calls.append(1))
    before = runtime.service.store.get(job_id).model_dump(mode="json")
    try:
        with pytest.raises(ValueError, match="source_fact_correction"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "更正来源", file_sha256(output),
                               source_fact_corrections=[{**correction, "source_sha256": "f" * 64}])
        assert calls == []
        assert runtime.service.store.get(job_id).model_dump(mode="json") == before
    finally:
        runtime.shutdown()


def test_source_fact_correction_requires_native_pixels_not_intent_only(tmp_path, monkeypatch):
    runtime, job_id, _deck, output, versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    store = runtime.service.store
    paths = [output, versions.current_path, versions.history_path,
             store.checkpoint_path(job_id, "deck-ir.json"), store.checkpoint_path(job_id, "deck-plan.json"),
             store.checkpoint_path(job_id, "brief.json"), store.job_dir(job_id) / "quality/quality-report.json"]
    before = {path: path.read_bytes() for path in paths}
    runtime.object_edit_planner = SimpleNamespace(plan=lambda *a, **k: PresentationObjectEditPlan(
        slide_id="slide-a", actions=[], intent_patch=SlideIntentPatch(
            slide_id="slide-a", key_message=correction["corrected_fact"],
        ),
    ))
    try:
        with pytest.raises(ValueError, match="source_fact_correction: native object actions"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "仅改判断不能更正成品",
                               file_sha256(output), source_fact_corrections=[correction])
        assert {path: path.read_bytes() for path in paths} == before
    finally:
        runtime.shutdown()


def test_source_fact_correction_without_object_planner_is_rejected_before_attempt(tmp_path, monkeypatch):
    runtime, job_id, _deck, output, _versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    before = runtime.service.store.get(job_id).model_dump(mode="json")
    try:
        with pytest.raises(ValueError, match="source_fact_correction: a native object edit planner"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "不能用旧文本替换绕过验证",
                               file_sha256(output), source_fact_corrections=[correction])
        assert runtime.service.store.get(job_id).model_dump(mode="json") == before
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("audit", ["valid", "missing_job", "broken_source", "bad_pin"])
def test_independent_artifact_review_replays_corrections_before_critic(tmp_path, monkeypatch, audit):
    from lantai.presentation.natural_edit import apply_source_fact_corrections

    runtime, job_id, deck, output, _versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    store = runtime.service.store
    _, brief = runtime.service._visual_source_context(job_id)
    changed, _ = apply_source_fact_corrections(deck, None, brief, [correction], allowed_slide_ids=["slide-a"])
    changed.slides[0].objects[0].text = correction["corrected_fact"]
    runtime.service.compiler_for(CompositionEngine.NATIVE_PPTX).compile(changed, output)
    if audit == "bad_pin":
        changed.source_fact_corrections[0].correction.source_sha256 = "e" * 64
    elif audit == "broken_source":
        store.checkpoint_path(job_id, "brief.json").write_text("{invalid source", encoding="utf-8")
    calls = []

    def critic(*args):
        calls.append(1)
        return [], [], "synthetic-audit/1", [], True

    monkeypatch.setattr(runtime.service, "_required_visual_review", critic)
    try:
        report, _ = runtime.service.verify_artifact(
            changed, output, store.job_dir(job_id) / "renders/source-audit",
            job_id=None if audit == "missing_job" else job_id,
        )
        issues = [issue for issue in report.issues if issue.code == "content.source_fact_correction_invalid"]
        if audit == "valid":
            assert issues == [] and calls == [1]
        else:
            assert len(issues) == 1 and issues[0].severity == "error" and not issues[0].repairable
            assert calls == [] and not report.passed
    finally:
        runtime.shutdown()


def test_later_native_edit_preserves_source_correction_and_reviews_all_pages(tmp_path, monkeypatch):
    runtime, job_id, _deck, output, _versions, correction = _source_correction_runtime(tmp_path, monkeypatch)
    reviewed = []

    class Planner:
        def plan(self, current, slide_id, instruction, **context):
            assert len(current.source_fact_corrections) == 1
            return PresentationObjectEditPlan(slide_id=slide_id, actions=[ReplaceTextAction(
                object_id="title-a" if slide_id == "slide-a" else "title-b",
                text=correction["corrected_fact"] if slide_id == "slide-a" else "保持不变 · 已核对",
            )])

    class Critic:
        def review(self, slides):
            reviewed.append({slide.slide_id for slide in slides})
            return VisualCriticResult(critic_id="synthetic-follow-up/1", issues=[])

    runtime.object_edit_planner, runtime.service.visual_critic = Planner(), Critic()
    try:
        runtime.redo_slide(job_id, "conversation-a", "slide-a", "按原材料更正",
                           file_sha256(output), source_fact_corrections=[correction])
        path = runtime.service.store.checkpoint_path(job_id, "deck-ir.json")
        original_record = DeckIR.model_validate_json(path.read_text()).source_fact_corrections[0]
        runtime.redo_slide(job_id, "conversation-a", "slide-b", "补充核对标签", file_sha256(output))
        changed = DeckIR.model_validate_json(path.read_text())
        assert changed.source_fact_corrections == [original_record]
        assert changed.required_facts == [correction["corrected_fact"], "保持不变"]
        assert reviewed == [{"slide-a", "slide-b"}, {"slide-a", "slide-b"}]
        assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 3
    finally:
        runtime.shutdown()


def test_redo_slide_atomically_commits_explicit_pages_and_preserves_unselected(tmp_path, monkeypatch):
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True,
    )
    calls = []

    class Planner:
        def plan(self, current, slide_id, instruction, *, deck_plan=None, additional_slide_ids=None):
            calls.append((slide_id, additional_slide_ids))
            return PresentationObjectEditPlan(
                slide_id=slide_id, additional_slide_ids=["slide-b"],
                actions=[
                    ReplaceTextAction(object_id="title-a", text="已修订主页面"),
                    ReplaceTextAction(object_id="title-b", text="已修订指定附加页"),
                ],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    try:
        before = file_sha256(output)
        runtime.redo_slide(job_id, "conversation-a", "slide-a", "同时改两页", before,
                           additional_slide_ids=["slide-b"])
        changed = DeckIR.model_validate_json(
            runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_text(),
        )
        assert calls == [("slide-a", ["slide-b"])]
        assert changed.slides[0].objects[0].text == "已修订主页面"
        assert changed.slides[1].objects[0].text == "已修订指定附加页"
        assert changed.slides[2] == deck.slides[2]
        native = {shape.name: shape.text for slide in Presentation(str(output)).slides
                  for shape in slide.shapes if getattr(shape, "has_text_frame", False)}
        assert native["title-a"] == "已修订主页面"
        assert native["title-b"] == "已修订指定附加页"
        assert native["title-c"] == "未指定页必须保持"
        artifacts = runtime.artifacts(job_id, "conversation-a")
        assert len(artifacts["presentation_versions"]) == 2
        assert artifacts["presentation_version"]["writer"] == "lantai_ai_edit"
        assert artifacts["quality"]["output_sha256"] == file_sha256(output) != before
        assert artifacts["quality"]["visual_verified"] is True
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("changed_page", ["slide-a", "slide-b"])
def test_redo_slide_permission_envelope_does_not_force_unchanged_page_rewrites(
    tmp_path, monkeypatch, engine, changed_page,
):
    from zipfile import ZipFile

    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    unchanged_part = "ppt/slides/slide2.xml" if changed_page == "slide-a" else "ppt/slides/slide1.xml"
    with ZipFile(output) as package:
        unchanged_bytes = package.read(unchanged_part)
        third_bytes = package.read("ppt/slides/slide3.xml")
    selected_object = "title-a" if changed_page == "slide-a" else "title-b"

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"],
                actions=[ReplaceTextAction(object_id=selected_object, text="实际修改的唯一页面")],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    prepare = PresentationEditTransactionService.prepare_compiled_redesign
    compiled_scopes = []

    def observe_prepare(self, *args, **kwargs):
        compiled_scopes.append(kwargs["compiled_slide_ids"])
        return prepare(self, *args, **kwargs)

    monkeypatch.setattr("lantai.presentation.editing.PresentationEditTransactionService.prepare_compiled_redesign", observe_prepare)
    try:
        updated = runtime.redo_slide(job_id, "conversation-a", "slide-a", "只修改需要修改的指定页",
                                     file_sha256(output), additional_slide_ids=["slide-b"])
        assert updated.status == JobStatus.COMPLETED
        assert compiled_scopes == [[changed_page]]
        with ZipFile(output) as package:
            assert package.read(unchanged_part) == unchanged_bytes
            assert package.read("ppt/slides/slide3.xml") == third_bytes
        assert runtime.artifacts(job_id, "conversation-a")["quality"]["visual_verified"] is True
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("additional", [["slide-a"], ["unknown"], ["slide-b", "slide-b"]])
def test_redo_slide_invalid_explicit_scope_never_calls_provider_or_mutates_base(tmp_path, monkeypatch, additional):
    runtime, job_id, _deck, output = _completed_editable_runtime(tmp_path, monkeypatch)
    calls = []

    class Planner:
        def plan(self, *args, **kwargs):
            calls.append(True)
            raise AssertionError("invalid scope must not call a provider")

    runtime.object_edit_planner = Planner()
    before = file_sha256(output)
    original_deck = runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_bytes()
    original_cost = runtime.service.store.get(job_id).cost_used_units
    try:
        with pytest.raises(ValueError, match="scope"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "有界修改", before,
                               additional_slide_ids=additional)
        assert calls == []
        assert runtime.service.store.get(job_id).cost_used_units == original_cost
        assert file_sha256(output) == before
        assert runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_bytes() == original_deck
        assert PresentationVersionStore(runtime.service.store.job_dir(job_id)).current() is None
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("overscope", ["additional", "intent"])
def test_redo_slide_rejects_unselected_model_scope_atomically(tmp_path, monkeypatch, overscope):
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True,
    )

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-c"] if overscope == "additional" else [],
                actions=[ReplaceTextAction(object_id="title-a", text="不得部分发布")],
                intent_patch=SlideIntentPatch(slide_id="slide-c", key_message="不得越权")
                if overscope == "intent" else None,
            )

    runtime.object_edit_planner = Planner()
    before = file_sha256(output)
    original_deck = runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_bytes()
    quality = runtime.service.store.job_dir(job_id) / "quality/quality-report.json"
    original_quality = quality.read_bytes()
    try:
        with pytest.raises((ValueError, PermissionError), match="scope"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "只改指定两页", before,
                               additional_slide_ids=["slide-b"])
        assert file_sha256(output) == before
        assert runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_bytes() == original_deck
        assert quality.read_bytes() == original_quality
        assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 1
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("failure", ["review", "quality_commit", "stop"])
def test_redo_slide_multi_page_failure_preserves_every_page_and_base_evidence(tmp_path, monkeypatch, failure):
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True,
    )

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"],
                actions=[ReplaceTextAction(object_id="title-a", text="候选主页面"),
                         ReplaceTextAction(object_id="title-b", text="候选附加页")],
            )

    class Critic:
        def review(self, slides):
            if failure == "stop":
                raise PresentationJobInterruptedError("synthetic user stop")
            return VisualCriticResult(
                critic_id="test-multi-page/1",
                issues=[VisualIssue(
                    slide_id="slide-b", object_id="title-b", severity="error",
                    issue_type="content_accuracy", description="指定附加页未通过事实验收",
                    proposed_action="核对来源后重试",
                )] if failure == "review" else [],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = Critic()
    before = file_sha256(output)
    deck_path = runtime.service.store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = runtime.service.store.job_dir(job_id) / "quality/quality-report.json"
    original_deck, original_quality = deck_path.read_bytes(), quality_path.read_bytes()
    if failure == "quality_commit":
        original_write = runtime._atomic_write_bytes
        failed_once = False

        def fail_first_quality(path, content):
            nonlocal failed_once
            if path == quality_path and not failed_once:
                failed_once = True
                raise OSError("synthetic quality commit failure")
            original_write(path, content)

        monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_first_quality)
    expected = {"review": ValueError, "quality_commit": OSError, "stop": PresentationJobInterruptedError}
    try:
        with pytest.raises(expected[failure]):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "原子修订两页", before,
                               additional_slide_ids=["slide-b"])
        assert file_sha256(output) == before
        assert deck_path.read_bytes() == original_deck
        assert quality_path.read_bytes() == original_quality
        versions = runtime.artifacts(job_id, "conversation-a")["presentation_versions"]
        assert len(versions) == 1
        assert versions[0]["pptx_sha256"] == before
    finally:
        runtime.shutdown()


def test_rejected_native_edit_keeps_candidate_quality_without_publishing_it(
    tmp_path, monkeypatch,
):
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True,
        composition_engine=CompositionEngine.NATIVE_PPTX,
    )

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a",
                actions=[ReplaceTextAction(object_id="title-a", text="未发布的候选标题")],
            )

    class Critic:
        def review(self, slides):
            return VisualCriticResult(
                critic_id="test-rejected-edit/1",
                issues=[VisualIssue(
                    slide_id="slide-b", object_ids=["slide-b::title-b"], severity="error",
                    issue_type="alignment", description="未修改页面存在图文相交",
                    proposed_action="保留原稿并检查候选渲染",
                )] if any(slide.slide_id == "slide-b" for slide in slides) else [],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = Critic()
    store = runtime.service.store
    job_dir = store.job_dir(job_id)
    before_hash = file_sha256(output)
    versions = PresentationVersionStore(job_dir)
    versions.bootstrap(
        output, _deck, writer=CompositionEngine.NATIVE_PPTX.value,
        quality_report_path=job_dir / "quality/quality-report.json",
        render_dir=job_dir / "renders/round-0",
    )
    before_versions = versions.history()
    protected_paths = [
        output, store.checkpoint_path(job_id, "deck-ir.json"),
        job_dir / "quality/quality-report.json",
        versions.current_path, versions.history_path,
        *sorted((job_dir / "versions").rglob("*.json")),
    ]
    before_bytes = {path: path.read_bytes() for path in protected_paths}
    try:
        with pytest.raises(ValueError, match="vision.alignment"):
            runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "只修订首页标题", before_hash,
            )

        reports = sorted((job_dir / "checkpoints").glob("edit-quality-report-*.json"))
        assert len(reports) == 1
        report = QualityReport.model_validate_json(reports[0].read_text(encoding="utf-8"))
        assert not report.passed
        assert not report.visual_verified
        assert report.visual_critic_id == "test-rejected-edit/1"
        assert any(
            issue.code == "vision.alignment" and issue.slide_id == "slide-b"
            for issue in report.issues
        )
        assert reports[0].name == f"edit-quality-report-{report.render_set_id}.json"
        candidate = next((job_dir / "edits").glob("*/candidate.pptx"))
        assert report.output_sha256 == file_sha256(candidate)
        assert report.output_sha256 != before_hash
        assert report.render_verified and report.rendered_slide_count == 3
        assert len(report.rendered_slides) == 3
        render_dir = job_dir / "renders" / report.render_set_id
        assert render_dir.is_dir()
        assert {evidence.sha256 for evidence in report.rendered_slides} == {
            file_sha256(path) for path in render_dir.glob("*.png")
        }
        assert store.get(job_id).status == JobStatus.COMPLETED
        assert all(path.read_bytes() == content for path, content in before_bytes.items())
        assert versions.history() == before_versions
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX])
@pytest.mark.parametrize("declare_extra", [False, True])
def test_redo_slide_required_pages_reject_partial_actions_before_compilation(
    tmp_path, monkeypatch, engine, declare_extra,
):
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    calls = []

    class Planner:
        def plan(self, *args, **kwargs):
            calls.append(kwargs)
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"] if declare_extra else [],
                actions=[ReplaceTextAction(object_id="title-a", text="只完成了主页面")],
            )

    runtime.object_edit_planner = Planner()
    monkeypatch.setattr(PresentationEditTransactionService, "prepare_compiled_redesign", lambda *a, **k: pytest.fail("partial mandatory edit must not compile"))
    monkeypatch.setattr(PresentationEditTransactionService, "prepare", lambda *a, **k: pytest.fail("partial mandatory edit must not compile"))
    runtime.service.visual_critic = SimpleNamespace(review=lambda *a, **k: pytest.fail("partial mandatory edit must not review"))
    job = runtime.service.store.job_dir(job_id)
    paths = [output, job / "checkpoints/deck-ir.json", job / "quality/quality-report.json"]
    before = {path: path.read_bytes() for path in paths}
    try:
        with pytest.raises(ValueError, match="edit.required_slide_unchanged"):
            runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "两页必须均完成修改", file_sha256(output),
                additional_slide_ids=["slide-b"], required_slide_ids=["slide-a", "slide-b"],
            )
        assert calls[0]["required_slide_ids"] == ["slide-a", "slide-b"]
        assert {path: path.read_bytes() for path in paths} == before
        assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 1
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("required", [["slide-c"], ["slide-a", "slide-a"], [""]])
@pytest.mark.parametrize("engine", [CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX])
def test_redo_slide_required_scope_is_validated_before_provider_or_budget(tmp_path, monkeypatch, required, engine):
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    calls = []
    runtime.object_edit_planner = SimpleNamespace(plan=lambda *a, **k: calls.append("provider"))
    store = runtime.service.store
    original_lease = store.writer_lease
    original_budget = store.assert_execution_budget

    @contextmanager
    def observed_lease(*args, **kwargs):
        calls.append("lease_enter")
        with original_lease(*args, **kwargs) as lease:
            yield lease

    def observed_budget(*args, **kwargs):
        calls.append("budget")
        return original_budget(*args, **kwargs)

    monkeypatch.setattr(store, "writer_lease", observed_lease)
    monkeypatch.setattr(store, "assert_execution_budget", observed_budget)
    monkeypatch.setattr(PresentationEditTransactionService, "prepare_compiled_redesign", lambda *a, **k: pytest.fail("invalid required scope must not compile"))
    monkeypatch.setattr(PresentationEditTransactionService, "prepare", lambda *a, **k: pytest.fail("invalid required scope must not compile"))
    monkeypatch.setattr(runtime.service, "verify_artifact", lambda *a, **k: pytest.fail("invalid required scope must not render"))
    runtime.service.visual_critic = SimpleNamespace(review=lambda *a, **k: pytest.fail("invalid required scope must not review"))
    before = file_sha256(output)
    cost = store.get(job_id).cost_used_units
    job = store.job_dir(job_id)
    before_files = {path.relative_to(job): path.read_bytes() for path in job.rglob("*") if path.is_file()}
    try:
        with pytest.raises(ValueError, match="required.*scope|scope.*required"):
            runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "不能扩大必改范围", before,
                additional_slide_ids=["slide-b"], required_slide_ids=required,
            )
        assert calls == []
        assert store.get(job_id).cost_used_units == cost
        assert file_sha256(output) == before
        assert {path.relative_to(job): path.read_bytes() for path in job.rglob("*") if path.is_file()} == before_files
        assert PresentationVersionStore(job).current() is None
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX])
@pytest.mark.parametrize("scope", [{"additional_slide_ids": ["slide-b"]}, {"required_slide_ids": ["slide-a"]}])
def test_redo_slide_scope_without_planner_preserves_attempt_and_ledger(tmp_path, monkeypatch, engine, scope):
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    store = runtime.service.store
    job = store.job_dir(job_id)
    versions = PresentationVersionStore(job)
    versions.bootstrap(output, deck, writer=engine.value, quality_report_path=job / "quality/quality-report.json",
                       render_dir=job / "renders/round-0")
    calls = []
    original_lease = store.writer_lease
    original_budget = store.assert_execution_budget

    @contextmanager
    def observed_lease(*args, **kwargs):
        calls.append("lease_enter")
        with original_lease(*args, **kwargs) as lease:
            yield lease

    def observed_budget(*args, **kwargs):
        calls.append("budget")
        return original_budget(*args, **kwargs)

    runtime.object_edit_planner = None
    monkeypatch.setattr(store, "writer_lease", observed_lease)
    monkeypatch.setattr(store, "assert_execution_budget", observed_budget)
    store.get(job_id)
    before_files = {path.relative_to(job): path.read_bytes() for path in job.rglob("*") if path.is_file()}
    try:
        with pytest.raises(ValueError, match="scope requires an object edit planner"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "范围不能绕过规划器", file_sha256(output), **scope)
        assert calls == []
        assert {path.relative_to(job): path.read_bytes() for path in job.rglob("*") if path.is_file()} == before_files
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX])
@pytest.mark.parametrize("removed_slide", ["slide-a", "slide-b"])
def test_redo_slide_rechecks_scope_after_lease_without_rewinding_new_checkpoint(
    tmp_path, monkeypatch, engine, removed_slide,
):
    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    store = runtime.service.store
    job = store.job_dir(job_id)
    deck_path = store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = job / "quality/quality-report.json"
    versions = PresentationVersionStore(job)
    versions.bootstrap(output, deck, writer=engine.value, quality_report_path=quality_path,
                       render_dir=job / "renders/round-0")
    retained = {path: path.read_bytes() for path in (output, quality_path, versions.current_path, versions.history_path)}
    original_lease = store.writer_lease
    changed_checkpoint = DeckIR.model_validate_json(deck_path.read_text(encoding="utf-8"))
    changed_checkpoint.slides = [slide for slide in changed_checkpoint.slides if slide.id != removed_slide]
    changed_text = changed_checkpoint.model_dump_json(indent=2)
    entered = []

    @contextmanager
    def changing_lease(*args, **kwargs):
        # A different operation replaces the checkpoint after preflight, before
        # this edit acquires its actual cross-process mutation authority.
        store.write_checkpoint(job_id, "deck-ir.json", changed_text)
        with original_lease(*args, **kwargs) as lease:
            entered.append(lease.lease_id)
            yield lease

    monkeypatch.setattr(store, "writer_lease", changing_lease)
    runtime.object_edit_planner = SimpleNamespace(plan=lambda *a, **k: pytest.fail("stale edit scope must not call the provider"))
    monkeypatch.setattr(PresentationEditTransactionService, "prepare_compiled_redesign", lambda *a, **k: pytest.fail("stale edit scope must not compile"))
    monkeypatch.setattr(PresentationEditTransactionService, "prepare", lambda *a, **k: pytest.fail("stale edit scope must not compile"))
    monkeypatch.setattr(runtime.service, "verify_artifact", lambda *a, **k: pytest.fail("stale edit scope must not render"))
    runtime.service.visual_critic = SimpleNamespace(review=lambda *a, **k: pytest.fail("stale edit scope must not review"))
    try:
        with pytest.raises(ValueError, match="requested.*scope.*unknown"):
            runtime.redo_slide(
                job_id, "conversation-a", "slide-a", "取得租约后仍须核对页面", file_sha256(output),
                additional_slide_ids=["slide-b"], required_slide_ids=["slide-a", "slide-b"],
            )
        assert len(entered) == 1
        assert deck_path.read_text(encoding="utf-8") == changed_text
        assert {path: path.read_bytes() for path in retained} == retained
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX])
def test_redo_slide_required_pages_commit_complete_object_actions_atomically(tmp_path, monkeypatch, engine):
    from zipfile import ZipFile

    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, include_third_slide=True, composition_engine=engine,
    )
    calls = []

    class Planner:
        def plan(self, *args, **kwargs):
            calls.append(kwargs)
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"],
                actions=[ReplaceTextAction(object_id="title-a", text="主页面已改"),
                         ReplaceTextAction(object_id="title-b", text="附加页已改")],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = _PassingVisualCritic()
    with ZipFile(output) as package:
        untouched = package.read("ppt/slides/slide3.xml")
    try:
        updated = runtime.redo_slide(
            job_id, "conversation-a", "slide-a", "两页必须完成修改", file_sha256(output),
            additional_slide_ids=["slide-b"], required_slide_ids=["slide-a", "slide-b"],
        )
        assert updated.status == JobStatus.COMPLETED
        assert calls[0]["required_slide_ids"] == ["slide-a", "slide-b"]
        changed = DeckIR.model_validate_json(runtime.service.store.checkpoint_path(job_id, "deck-ir.json").read_text())
        assert [slide.objects[0].text for slide in changed.slides[:2]] == ["主页面已改", "附加页已改"]
        with ZipFile(output) as package:
            assert package.read("ppt/slides/slide3.xml") == untouched
        assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 2
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("failure", [None, "review", "stop", "quality_commit"])
def test_twenty_page_native_edit_is_atomic_and_reviews_unselected_page(tmp_path, monkeypatch, failure):
    from zipfile import ZipFile
    from lantai.presentation.reconcile import reconcile_pptx

    runtime, job_id, deck, output = _completed_editable_runtime(
        tmp_path, monkeypatch, page_count=21, composition_engine=CompositionEngine.NATIVE_PPTX,
    )
    pages = [slide.id for slide in deck.slides[:20]]
    reviewed = []

    class Planner:
        def plan(self, *args, **kwargs):
            assert kwargs["additional_slide_ids"] == pages[1:]
            assert kwargs["required_slide_ids"] == pages
            return PresentationObjectEditPlan(
                slide_id=pages[0], additional_slide_ids=pages[1:],
                actions=[ReplaceTextAction(object_id=slide.objects[0].id, text=f"Revised {index}")
                         for index, slide in enumerate(deck.slides[:20])],
            )

    class Critic:
        def review(self, slides):
            reviewed.extend(slide.slide_id for slide in slides)
            if failure == "stop":
                raise PresentationJobInterruptedError("synthetic user stop")
            return VisualCriticResult(
                critic_id="test-twenty-page/1",
                issues=[VisualIssue(
                    slide_id=deck.slides[20].id, severity="error", issue_type="content_accuracy",
                    description="Unselected page still blocks whole-deck publication",
                    proposed_action="Preserve baseline until verified",
                )] if failure == "review" else [],
            )

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = Critic()
    before = file_sha256(output)
    deck_path = runtime.service.store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = runtime.service.store.job_dir(job_id) / "quality/quality-report.json"
    original_deck, original_quality = deck_path.read_bytes(), quality_path.read_bytes()
    with ZipFile(output) as package:
        untouched = package.read("ppt/slides/slide21.xml")
    # The compiled native file, not the pre-compile synthetic IR, owns
    # geometry/styles at edit time. Compare the untouched page against
    # that same normalized baseline while separately requiring exact XML.
    native_baseline = prepare_deck_for_composition(reconcile_pptx(
        deck, output, assets_dir=tmp_path / "baseline-native-assets",
    ))
    if failure == "quality_commit":
        original_write = runtime._atomic_write_bytes
        failed_once = False

        def fail_first_quality(path, content):
            nonlocal failed_once
            if path == quality_path and not failed_once:
                failed_once = True
                raise OSError("synthetic quality commit failure")
            original_write(path, content)

        monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_first_quality)
    try:
        def run_edit():
            return runtime.redo_slide(
                job_id, "conversation-a", pages[0], "Atomically revise twenty explicit pages",
                before, additional_slide_ids=pages[1:], required_slide_ids=pages,
            )

        if failure is None:
            assert run_edit().status == JobStatus.COMPLETED
            changed = DeckIR.model_validate_json(deck_path.read_text())
            assert [slide.objects[0].text for slide in changed.slides[:20]] == [f"Revised {i}" for i in range(20)]
            assert changed.slides[20].model_dump(exclude={"revision", "scene_hash"}) == native_baseline.slides[20].model_dump(exclude={"revision", "scene_hash"})
            assert file_sha256(output) != before
            assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 2
        else:
            expected = {"review": ValueError, "stop": PresentationJobInterruptedError, "quality_commit": OSError}
            with pytest.raises(expected[failure]):
                run_edit()
            assert file_sha256(output) == before
            assert deck_path.read_bytes() == original_deck
            assert quality_path.read_bytes() == original_quality
            assert len(runtime.artifacts(job_id, "conversation-a")["presentation_versions"]) == 1
        assert set(reviewed) == {slide.id for slide in deck.slides}
        with ZipFile(output) as package:
            assert package.read("ppt/slides/slide21.xml") == untouched
    finally:
        runtime.shutdown()


def test_redo_slide_commits_a_real_bounded_edit_with_new_evidence(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path,
        monkeypatch,
    )
    before_hash = file_sha256(output)
    try:
        updated = runtime.redo_slide(
            job_id,
            "conversation-a",
            "slide-a",
            "把本页标题改为：经过验证的新标题",
            before_hash,
        )
        deck_path = runtime.service.store.checkpoint_path(
            job_id,
            "deck-ir.json",
        )
        changed_deck = DeckIR.model_validate_json(
            deck_path.read_text(encoding="utf-8"),
        )
        quality = runtime.artifacts(job_id, "conversation-a")["quality"]
        presentation = Presentation(str(output))
        native_text = {
            shape.name: shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if getattr(shape, "has_text_frame", False)
        }

        assert updated.status == JobStatus.COMPLETED
        assert changed_deck.slides[0].objects[0].text == "经过验证的新标题"
        assert changed_deck.slides[1].objects[0].text == "保持不变"
        assert native_text["title-a"] == "经过验证的新标题"
        assert native_text["title-b"] == "保持不变"
        assert file_sha256(output) != before_hash
        assert quality["output_sha256"] == file_sha256(output)
        assert quality["render_set_id"].startswith("edit-")
        assert quality["render_verified"] is True
        assert quality["visual_verified"] is False
        assert quality["render_verified"] is True
        version = runtime.artifacts(job_id, "conversation-a")[
            "presentation_version"
        ]
        assert version["writer"] == "lantai_ai_edit"
        assert version["parent_version_id"] is not None
        assert updated.events[-1].message.startswith("单页修改已提交")

        versions = runtime.artifacts(job_id, "conversation-a")[
            "presentation_versions"
        ]
        assert len(versions) == 2
        edited_hash = file_sha256(output)
        restored = runtime.rollback_version(
            job_id,
            "conversation-a",
            target_version_id=versions[0]["version_id"],
            base_output_sha256=edited_hash,
        )
        restored_deck = DeckIR.model_validate_json(
            deck_path.read_text(encoding="utf-8"),
        )
        restored_artifacts = runtime.artifacts(job_id, "conversation-a")
        assert restored.writer == "version_rollback"
        assert restored.parent_version_id == versions[1]["version_id"]
        assert restored.generation > versions[1]["generation"]
        assert file_sha256(output) == before_hash
        assert restored_deck.slides[0].objects[0].text == "原始标题"
        assert restored_artifacts["quality"]["output_sha256"] == before_hash
        assert restored_artifacts["presentation_version"]["writer"] == (
            "version_rollback"
        )
        assert len(restored_artifacts["thumbnails"]) == 2
    finally:
        runtime.shutdown()


class _WholeSlideDesigner:
    def redesign(self, request):
        storyboard_title = next(
            obj
            for obj in request.slide.objects
            if isinstance(obj, TextObject) and obj.role == "title"
        )
        return SlideRedesignProposal(
            slide=DeckSlide(
                id=request.slide.id,
                intent_id=request.slide.intent_id,
                layout_id="quote",
                speaker_notes="通过重做强化唯一判断",
                objects=[
                    storyboard_title.model_copy(
                        update={"id": "redesign-storyboard-title"},
                    ),
                    TextObject(
                        id="redesign-title",
                        role="quote",
                        text="可信交付来自真实验证。",
                        slot="quote",
                    ),
                    TextObject(
                        id="redesign-quote",
                        role="body",
                        text="先看见最终页面，再决定是否交付",
                        slot="attribution",
                    ),
                ],
            ),
            rationale="用引语构图重建页面层级",
        )


class _PassingVisualCritic:
    def review(self, _slides):
        return VisualCriticResult(
            critic_id="test-vision/1",
            issues=[],
        )


class _IntentPatchPlanner:
    def plan(
        self,
        deck,
        slide_id,
        instruction,
        *,
        deck_plan=None,
    ):
        assert deck_plan is not None
        assert slide_id == "slide-a"
        assert instruction
        return PresentationObjectEditPlan(
            slide_id=slide_id,
            intent_patch=SlideIntentPatch(
                key_message="Q1 客户数为 18",
            ),
            rationale="按原始数据纠正作者主判断",
        )


def _completed_intent_edit_runtime(
    tmp_path,
    monkeypatch,
    *,
    composition_engine: CompositionEngine = CompositionEngine.LEGACY_PYTHON_PPTX,
) -> tuple[PresentationJobRuntime, str, DeckIR, DeckPlan, Path]:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    runtime = PresentationJobRuntime(
        tmp_path,
        object_edit_planner=_IntentPatchPlanner(),
        visual_critic=_PassingVisualCritic(),
    )
    record = runtime.service.create_job(_request().model_copy(update={
        "composition_engine": composition_engine,
    }))
    deck = DeckIR(
        id="intent-edit-deck",
        title="季度复盘",
        slides=[
            DeckSlide(
                id="slide-a",
                intent_id="intent-q1",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="title-a",
                        role="title",
                        text="Q1 客户数为 18",
                        slot="title",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-b",
                intent_id="intent-q2",
                layout_id="closing",
                objects=[
                    TextObject(
                        id="title-b",
                        role="title",
                        text="Q2 客户数为 21",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    direction = VisualDirection(
        visual_concept="克制经营复盘",
        palette=["#111111", "#FFFFFF"],
        typography="清晰",
        image_language="无装饰图",
        chart_language="原生图表",
        shape_language="少量分隔线",
        grid="十二列",
        recurring_motif="季度标记",
    )
    plan = DeckPlan(
        title="季度复盘",
        narrative_arc="按季度核对经营事实",
        visual_direction=direction,
        slides=[
            DeckPlanSlide(
                order=1,
                title="Q1",
                intent=SlideIntent(
                    id="intent-q1",
                    purpose="核对 Q1",
                    key_message="Q1 客户数为 21",
                ),
            ),
            DeckPlanSlide(
                order=2,
                title="Q2",
                intent=SlideIntent(
                    id="intent-q2",
                    purpose="核对 Q2",
                    key_message="Q2 客户数为 21",
                ),
            ),
        ],
    )
    if composition_engine == CompositionEngine.NATIVE_PPTX:
        for slide in deck.slides:
            for obj in slide.objects:
                obj.font_size = 48
    runtime.service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    runtime.service.store.write_checkpoint(
        record.job_id,
        "deck-plan.json",
        plan.model_dump_json(indent=2),
    )
    output = (
        runtime.service.store.job_dir(record.job_id)
        / "output"
        / "presentation.pptx"
    )
    runtime.service.compiler_for(composition_engine).compile(deck, output)
    report, _images = runtime.service.verify_artifact(
        deck,
        output,
        runtime.service.store.job_dir(record.job_id) / "renders" / "round-0",
        render_set_id="round-0",
    )
    quality_path = (
        runtime.service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    runtime.service.store.append_event(
        record.job_id,
        stage=JobStage.COMPLETED,
        status=JobStatus.COMPLETED,
        message="verified intent baseline",
        checkpoint=str(output),
    )
    return runtime, record.job_id, deck, plan, output


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("semantic_rejection", [False, True])
def test_redo_slide_carries_original_source_proof_through_native_edit_review(
    tmp_path, monkeypatch, engine, semantic_rejection,
):
    from zipfile import ZipFile

    runtime, job_id, deck, plan, output = _completed_intent_edit_runtime(
        tmp_path, monkeypatch, composition_engine=engine,
    )
    store = runtime.service.store
    source_ref = "material://quarter-data"
    brief = PresentationBrief(
        objective="季度新增客户数合计", audience="管理层", narrative="按来源合计",
        source_refs=[source_ref],
        source_sections=[PresentationSourceSection(
            title="季度新客户",
            body="Q1 客户数为 18 人；Q2 客户数为 21 人；客户数均指当季新增。口径：按季度合计新增客户数，单位为人。",
            source_ref=source_ref,
        )],
    )
    proof = SourceCalculationProof(
        result="39", unit="人", authorization_source_ref=source_ref,
        authorization_excerpt="按季度合计新增客户数，单位为人",
        terms=[
            {"source_ref": source_ref, "source_excerpt": "Q1 客户数为 18 人",
             "value": "18", "unit": "人", "coefficient": 1},
            {"source_ref": source_ref, "source_excerpt": "Q2 客户数为 21 人",
             "value": "21", "unit": "人", "coefficient": 1},
        ],
    )
    deck.source_refs = [source_ref]
    deck.required_facts = ["Q1 客户数为 18", "Q2 客户数为 21"]
    deck.facts = [
        FactRecord(fact_id=f"quarter-{index}", text=text, source_refs=[source_ref])
        for index, text in enumerate(deck.required_facts, start=1)
    ]
    for index, slide in enumerate(deck.slides, start=1):
        slide.source_fact_ids = [f"quarter-{index}"]
    store.write_checkpoint(job_id, "deck-ir.json", deck.model_dump_json(indent=2))
    store.write_checkpoint(job_id, "brief.json", brief.model_dump_json(indent=2))
    baseline_report, baseline_images = runtime.service.verify_artifact(
        deck, output, store.job_dir(job_id) / "renders/round-0", render_set_id="round-0",
    )
    assert is_verified_delivery(deck, baseline_report, output, baseline_images)
    (store.job_dir(job_id) / "quality/quality-report.json").write_text(
        baseline_report.model_dump_json(indent=2), encoding="utf-8",
    )
    plan_path = store.checkpoint_path(job_id, "deck-plan.json")
    source_bytes = store.checkpoint_path(job_id, "brief.json").read_bytes()
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    versions = PresentationVersionStore(store.job_dir(job_id))
    versions.bootstrap(output, deck, writer=engine.value, quality_report_path=quality_path,
                       render_dir=store.job_dir(job_id) / "renders/round-0")
    before = {path: path.read_bytes() for path in (
        output, plan_path, store.checkpoint_path(job_id, "deck-ir.json"),
        quality_path, versions.current_path, versions.history_path,
    )}
    with ZipFile(output) as package:
        untouched_page = package.read("ppt/slides/slide2.xml")
    planner_contexts, review_contexts = [], []
    metric = "期末存量客户数" if semantic_rejection else "新增客户数"
    changed_text = f"Q1 客户数为 18；{metric}合计39人"

    async def plan_from_sources(**kwargs):
        context = json.loads(kwargs["user_prompt"])
        planner_contexts.append(context)
        assert context["source_brief"] == brief.model_dump(mode="json")
        return StructuredUtilityResult(
            value=kwargs["result_model"].model_validate({
                "slide_id": "slide-a",
                "actions": [{"kind": "replace_text", "object_id": "title-a", "text": changed_text}],
                "intent_patch": {
                    "key_message": f"Q1与Q2{metric}合计39人",
                    "source_ref_additions": [source_ref],
                    "calculation_proof_additions": [proof.model_dump(mode="json")],
                },
            }), model_label="source-edit-fixture", repaired=False,
        )

    async def review_original_sources(**kwargs):
        context = json.loads(next(
            item["text"] for item in reversed(kwargs["user_content"]) if item["type"] == "text"
        ))
        review_contexts.append(context)
        if context.get("slide_id") == "slide-a":
            assert context["source_brief"] == brief.model_dump(mode="json")
            assert context["intent"]["calculation_proofs"] == [proof.model_dump(mode="json")]
        issues = [{
            "slide_id": "slide-a", "severity": "warning", "issue_type": "content_accuracy",
            "description": "合计新增人数不能表述为期末存量人数",
            "proposed_action": "按原材料保持新增口径",
        }] if semantic_rejection and context.get("slide_id") == "slide-a" else []
        return StructuredUtilityResult(
            value=kwargs["result_model"].model_validate(synthetic_review_response(kwargs, issues)),
            model_label="source-review-fixture", repaired=False,
        )

    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", plan_from_sources)
    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_content_task", review_original_sources)
    monkeypatch.setattr("lantai.presentation.intelligence._utility_route_cache_identity", lambda _agent_id: "fixture-route")
    runtime.object_edit_planner = UtilityPresentationObjectEditPlanner("fixture-agent")
    runtime.service.visual_critic = UtilityPresentationVisualCritic("fixture-agent")
    try:
        if semantic_rejection:
            with pytest.raises(ValueError, match="vision.content_accuracy"):
                runtime.redo_slide(job_id, "conversation-a", "slide-a", "按原来源合计季度新增客户数", file_sha256(output))
            assert {path: path.read_bytes() for path in before} == before
            assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
            assert store.checkpoint_path(job_id, "brief.json").read_bytes() == source_bytes
            assert len(planner_contexts) == 1 and len(review_contexts) == 3
            return
        updated = runtime.redo_slide(job_id, "conversation-a", "slide-a", "按原来源合计季度新增客户数", file_sha256(output))
        assert updated.status == JobStatus.COMPLETED
        changed_plan = DeckPlan.model_validate_json(plan_path.read_text())
        assert changed_plan.slides[0].intent.calculation_proofs == [proof]
        assert changed_plan.slides[0].intent.source_refs == [source_ref]
        assert changed_plan.slides[1] == plan.slides[1]
        assert store.checkpoint_path(job_id, "brief.json").read_bytes() == source_bytes
        assert len(planner_contexts) == 1
        assert len(review_contexts) == 3
        assert any(context.get("review_scope") == "whole_deck" for context in review_contexts)
        with ZipFile(output) as package:
            assert package.read("ppt/slides/slide2.xml") == untouched_page
        assert changed_text in [
            shape.text for shape in Presentation(str(output)).slides[0].shapes
            if shape.has_text_frame
        ]
        assert runtime.artifacts(job_id, "conversation-a")["quality"]["visual_verified"] is True
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("intent_page", ["slide-a", "slide-b"])
def test_redo_slide_atomically_reviews_intent_only_page_with_other_native_page(
    tmp_path, monkeypatch, engine, intent_page,
):
    from zipfile import ZipFile

    runtime, job_id, _deck, _plan, output = _completed_intent_edit_runtime(
        tmp_path, monkeypatch, composition_engine=engine,
    )
    object_page = "slide-b" if intent_page == "slide-a" else "slide-a"
    intent_part = "ppt/slides/slide1.xml" if intent_page == "slide-a" else "ppt/slides/slide2.xml"
    with ZipFile(output) as package:
        intent_bytes = package.read(intent_part)

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"],
                actions=[ReplaceTextAction(
                    object_id="title-b" if object_page == "slide-b" else "title-a",
                    text="Q2 客户数为 21（核对）" if object_page == "slide-b" else "Q1 客户数为 18（核对）",
                )],
                intent_patch=SlideIntentPatch(slide_id=intent_page,
                    key_message="Q1 客户数为 18" if intent_page == "slide-a" else "Q2 客户数为 21（复核）"),
            )

    runtime.object_edit_planner = Planner()
    original_verify = runtime.service.verify_artifact
    verified = []

    def observe_verify(deck, *args, plan=None, **kwargs):
        verified.append(([slide.id for slide in deck.slides], plan))
        return original_verify(deck, *args, plan=plan, **kwargs)

    monkeypatch.setattr(runtime.service, "verify_artifact", observe_verify)
    try:
        updated = runtime.redo_slide(job_id, "conversation-a", "slide-a", "在同一次整套审查中修订主判断和另一页对象",
                                     file_sha256(output), additional_slide_ids=["slide-b"])
        assert updated.status == JobStatus.COMPLETED
        assert len(verified) == 1 and verified[0][0] == ["slide-a", "slide-b"]
        changed_plan = DeckPlan.model_validate_json(runtime.service.store.checkpoint_path(job_id, "deck-plan.json").read_text())
        assert verified[0][1] == changed_plan
        expected = "Q1 客户数为 18" if intent_page == "slide-a" else "Q2 客户数为 21（复核）"
        assert changed_plan.slides[0 if intent_page == "slide-a" else 1].intent.key_message == expected
        with ZipFile(output) as package:
            assert package.read(intent_part) == intent_bytes
        assert runtime.artifacts(job_id, "conversation-a")["quality"]["visual_verified"] is True
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("mode", ["intent_only", "mixed", "review_rejected", "publication_failed"])
@pytest.mark.parametrize("required", [False, True])
def test_redo_slide_reviews_and_publishes_multiple_intents_atomically(
    tmp_path, monkeypatch, engine, mode, required,
):
    from zipfile import ZipFile

    runtime, job_id, _deck, _plan, output = _completed_intent_edit_runtime(
        tmp_path, monkeypatch, composition_engine=engine,
    )
    plan_path = runtime.service.store.checkpoint_path(job_id, "deck-plan.json")
    deck_path = runtime.service.store.checkpoint_path(job_id, "deck-ir.json")
    quality_path = runtime.service.store.job_dir(job_id) / "quality/quality-report.json"
    before = {path: path.read_bytes() for path in (output, plan_path, deck_path, quality_path)}
    with ZipFile(output) as package:
        unchanged_page_bytes = package.read("ppt/slides/slide2.xml")

    class Planner:
        def plan(self, *args, **kwargs):
            return PresentationObjectEditPlan(
                slide_id="slide-a", additional_slide_ids=["slide-b"],
                intent_patches=[SlideIntentPatch(key_message="Q1 客户数为 18"),
                                SlideIntentPatch(slide_id="slide-b", key_message="Q2 客户数为 21（复核）")],
                actions=[ReplaceTextAction(object_id="title-a", text="Q1 客户数为 18（复核）")]
                if mode == "mixed" else [],
            )

    runtime.object_edit_planner = Planner()
    verified = []
    original_verify = runtime.service.verify_artifact

    def observe_verify(deck, *args, plan=None, **kwargs):
        verified.append(([slide.id for slide in deck.slides], plan))
        report, images = original_verify(deck, *args, plan=plan, **kwargs)
        if mode == "review_rejected":
            report.passed = False
            report.issues.append(QualityIssue(code="test.blocking_review", severity="error",
                                            message="synthetic whole-deck rejection", repairable=False))
        return report, images

    monkeypatch.setattr(runtime.service, "verify_artifact", observe_verify)
    original_write = runtime._atomic_write_bytes
    failed_once = False

    def fail_quality_once(path, content):
        nonlocal failed_once
        if mode == "publication_failed" and path == quality_path and not failed_once:
            failed_once = True
            raise OSError("synthetic multi-intent publication failure")
        original_write(path, content)

    monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_quality_once)
    try:
        if mode in {"review_rejected", "publication_failed"}:
            with pytest.raises(ValueError if mode == "review_rejected" else OSError):
                runtime.redo_slide(job_id, "conversation-a", "slide-a", "同时核对两页作者判断",
                                   file_sha256(output), additional_slide_ids=["slide-b"],
                                   **({"required_slide_ids": ["slide-a", "slide-b"]} if required else {}))
            assert {path: path.read_bytes() for path in before} == before
            assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
            assert failed_once == (mode == "publication_failed")
        else:
            updated = runtime.redo_slide(job_id, "conversation-a", "slide-a", "同时核对两页作者判断",
                                         file_sha256(output), additional_slide_ids=["slide-b"],
                                         **({"required_slide_ids": ["slide-a", "slide-b"]} if required else {}))
            assert updated.status == JobStatus.COMPLETED
            changed_plan = DeckPlan.model_validate_json(plan_path.read_text())
            assert [slide.intent.key_message for slide in changed_plan.slides] == [
                "Q1 客户数为 18", "Q2 客户数为 21（复核）",
            ]
            assert verified[0][1] == changed_plan
            if mode == "intent_only":
                assert output.read_bytes() == before[output]
                assert deck_path.read_bytes() == before[deck_path]
            else:
                with ZipFile(output) as package:
                    assert package.read("ppt/slides/slide2.xml") == unchanged_page_bytes
        assert len(verified) == 1 and verified[0][0] == ["slide-a", "slide-b"]
        assert [slide.intent.key_message for slide in verified[0][1].slides] == [
            "Q1 客户数为 18", "Q2 客户数为 21（复核）",
        ]
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("engine", [
    CompositionEngine.LEGACY_PYTHON_PPTX, CompositionEngine.NATIVE_PPTX,
])
@pytest.mark.parametrize("legacy", [False, True])
@pytest.mark.parametrize("mixed", [False, True])
def test_redo_slide_rejects_shared_authored_identity_before_any_candidate_review(
    tmp_path, monkeypatch, engine, legacy, mixed,
):
    runtime, job_id, deck, plan, output = _completed_intent_edit_runtime(
        tmp_path, monkeypatch, composition_engine=engine,
    )
    deck.slides[1].intent_id = deck.slides[0].intent_id
    deck = DeckIR.model_validate_json(deck.model_dump_json())
    plan = DeckPlan.model_validate_json(plan.model_dump_json())
    store = runtime.service.store
    deck_path = store.checkpoint_path(job_id, "deck-ir.json")
    plan_path = store.checkpoint_path(job_id, "deck-plan.json")
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    store.write_checkpoint(job_id, "deck-ir.json", deck.model_dump_json(indent=2))
    report, images = runtime.service.verify_artifact(
        deck, output, store.job_dir(job_id) / "renders/round-0", plan=plan,
        vision_critique_required=True, render_set_id="round-0",
    )
    assert is_verified_delivery(deck, report, output, images)
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    before = {path: path.read_bytes() for path in (output, plan_path, deck_path, quality_path)}
    calls = {"planner": 0, "critic": 0}

    class Planner:
        def plan(self, *args, **kwargs):
            calls["planner"] += 1
            assert not kwargs.get("additional_slide_ids")
            patch = SlideIntentPatch(key_message="Q1 客户数为 18")
            return PresentationObjectEditPlan(
                slide_id="slide-a",
                actions=[ReplaceTextAction(object_id="title-a", text="Q1 客户数为 18（复核）")]
                if mixed else [],
                **({"intent_patch": patch} if legacy else {"intent_patches": [patch]}),
            )

    class Critic:
        def review(self, *args, **kwargs):
            calls["critic"] += 1
            pytest.fail("ambiguous intent must be rejected before candidate review")

    runtime.object_edit_planner = Planner()
    runtime.service.visual_critic = Critic()
    try:
        with pytest.raises(ValueError, match="ambiguous.*identity|identity.*ambiguous"):
            runtime.redo_slide(job_id, "conversation-a", "slide-a", "只修订明确授权的一页",
                               file_sha256(output))
        assert calls == {"planner": 1, "critic": 0}
        assert {path: path.read_bytes() for path in before} == before
        assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


@pytest.mark.parametrize("mode", ["passed", "review_rejected", "publication_failed", "stop"])
def test_notes_only_public_edit_preserves_native_bytes_and_rolls_back_all_authorities(
    tmp_path, monkeypatch, mode,
):
    from zipfile import ZipFile

    runtime, job_id, deck, _plan, output = _completed_intent_edit_runtime(
        tmp_path, monkeypatch, composition_engine=CompositionEngine.NATIVE_PPTX,
    )
    store = runtime.service.store
    deck_path = store.checkpoint_path(job_id, "deck-ir.json")
    plan_path = store.checkpoint_path(job_id, "deck-plan.json")
    quality_path = store.job_dir(job_id) / "quality/quality-report.json"
    versions = PresentationVersionStore(store.job_dir(job_id))
    versions.bootstrap(output, deck, writer="native_pptx", quality_report_path=quality_path,
                       render_dir=store.job_dir(job_id) / "renders/round-0")
    paths = [output, deck_path, plan_path, quality_path, versions.current_path, versions.history_path]
    before = {path: path.read_bytes() for path in paths}
    with ZipFile(output) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    runtime.object_edit_planner = SimpleNamespace(plan=lambda *a, **k: PresentationObjectEditPlan(
        slide_id="slide-a", speaker_notes_patches=[{"speaker_notes": "只解释本页客户口径，不增加新结论。"}],
    ))
    reviewed = []
    original_verify = runtime.service.verify_artifact
    def observe_verify(changed, *args, **kwargs):
        reviewed.append(([slide.id for slide in changed.slides], kwargs))
        if mode == "stop":
            raise PresentationJobInterruptedError("synthetic notes stop")
        report, images = original_verify(changed, *args, **kwargs)
        if mode == "review_rejected":
            report.passed = False
            report.issues.append(QualityIssue(code="test.notes_rejected", severity="error",
                                            message="synthetic whole-deck notes rejection", repairable=False))
        return report, images
    monkeypatch.setattr(runtime.service, "verify_artifact", observe_verify)
    original_write = runtime._atomic_write_bytes
    failed = False
    def fail_quality_once(path, content):
        nonlocal failed
        if mode == "publication_failed" and path == quality_path and not failed:
            failed = True
            raise OSError("synthetic notes publication failure")
        original_write(path, content)
    monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_quality_once)
    try:
        def edit():
            return runtime.redo_slide(job_id, "conversation-a", "slide-a", "只补本页讲解备注",
                                      file_sha256(output), required_slide_ids=["slide-a"])
        if mode == "passed":
            assert edit().status == JobStatus.COMPLETED
            changed = DeckIR.model_validate_json(deck_path.read_text())
            assert changed.slides[0].speaker_notes == "只解释本页客户口径，不增加新结论。"
            assert changed.slides[1].speaker_notes == deck.slides[1].speaker_notes
            assert plan_path.read_bytes() == before[plan_path]
            with ZipFile(output) as package:
                after = {name: package.read(name) for name in package.namelist()}
            assert parts.keys() == after.keys()
            assert [name for name in parts if parts[name] != after[name]] == ["ppt/notesSlides/notesSlide1.xml"]
            assert versions.current().writer == "lantai_ai_edit"
        else:
            exception = {"review_rejected": ValueError, "publication_failed": OSError,
                         "stop": PresentationJobInterruptedError}[mode]
            with pytest.raises(exception):
                edit()
            assert {path: path.read_bytes() for path in paths} == before
        assert len(reviewed) == 1 and reviewed[0][0] == ["slide-a", "slide-b"]
        assert reviewed[0][1]["visual_required"] is True
        assert reviewed[0][1]["vision_critique_required"] is True
        assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


def test_redo_slide_commits_intent_only_after_whole_deck_review(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, deck, plan, output = _completed_intent_edit_runtime(
        tmp_path,
        monkeypatch,
    )
    before_sha256 = file_sha256(output)
    deck_path = runtime.service.store.checkpoint_path(job_id, "deck-ir.json")
    plan_path = runtime.service.store.checkpoint_path(job_id, "deck-plan.json")
    before_deck_bytes = deck_path.read_bytes()
    try:
        updated = runtime.redo_slide(
            job_id,
            "conversation-a",
            "slide-a",
            "把 Q1 客户数主判断从 21 纠正为 18",
            before_sha256,
        )

        changed_plan = DeckPlan.model_validate_json(
            plan_path.read_text(encoding="utf-8"),
        )
        quality = runtime.artifacts(job_id, "conversation-a")["quality"]
        history = list(
            (
                runtime.service.store.job_dir(job_id)
                / "checkpoints"
                / "history"
            ).glob("intent-edit-*/deck-plan.json"),
        )

        assert updated.status == JobStatus.COMPLETED
        assert file_sha256(output) == before_sha256
        assert deck_path.read_bytes() == before_deck_bytes
        assert DeckIR.model_validate_json(
            deck_path.read_text(encoding="utf-8"),
        ) == deck
        assert changed_plan.slides[0].intent.key_message == "Q1 客户数为 18"
        assert changed_plan.slides[1].intent.key_message == "Q2 客户数为 21"
        assert len(history) == 1
        assert DeckPlan.model_validate_json(
            history[0].read_text(encoding="utf-8"),
        ) == plan
        assert quality["output_sha256"] == before_sha256
        assert quality["render_set_id"].startswith("edit-")
        assert quality["visual_verified"] is True
        assert quality["visual_critic_id"] == "test-vision/1"
        assert updated.events[-1].message.startswith("作者主判断修订已提交")
    finally:
        runtime.shutdown()


def test_redo_slide_rolls_back_intent_when_quality_commit_fails(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, _deck, plan, output = _completed_intent_edit_runtime(
        tmp_path,
        monkeypatch,
    )
    before_sha256 = file_sha256(output)
    plan_path = runtime.service.store.checkpoint_path(job_id, "deck-plan.json")
    quality_path = (
        runtime.service.store.job_dir(job_id)
        / "quality"
        / "quality-report.json"
    )
    before_quality_bytes = quality_path.read_bytes()
    original_atomic_write = runtime._atomic_write_bytes
    failed_once = False

    def fail_first_quality_commit(path, content):
        nonlocal failed_once
        if path == quality_path and not failed_once:
            failed_once = True
            raise OSError("simulated quality commit failure")
        original_atomic_write(path, content)

    monkeypatch.setattr(runtime, "_atomic_write_bytes", fail_first_quality_commit)
    try:
        with pytest.raises(OSError, match="quality commit failure"):
            runtime.redo_slide(
                job_id,
                "conversation-a",
                "slide-a",
                "把 Q1 客户数主判断从 21 纠正为 18",
                before_sha256,
            )

        assert failed_once is True
        assert file_sha256(output) == before_sha256
        assert DeckPlan.model_validate_json(
            plan_path.read_text(encoding="utf-8"),
        ) == plan
        assert quality_path.read_bytes() == before_quality_bytes
        assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


class _CheckpointVisualCritic:
    def __init__(self, *, deck_issue: bool = False) -> None:
        self.slide_calls: list[str] = []
        self.deck_calls = 0
        self.deck_issue = deck_issue
        self._lock = threading.Lock()

    def cache_identity(self) -> str:
        return "checkpoint-critic/v1"

    def review_slide(self, slide):
        with self._lock:
            self.slide_calls.append(slide.slide_id)
        return VisualCriticResult(
            critic_id="checkpoint-critic/model-a",
            issues=[],
        )

    def review_deck(self, slides):
        with self._lock:
            self.deck_calls += 1
        issues = []
        if self.deck_issue:
            issues.append(
                VisualIssue(
                    slide_id=slides[0].slide_id,
                    severity="error",
                    issue_type="repetition",
                    description="整套叙事重复，没有向前推进",
                    proposed_action="重写中段叙事节拍",
                ),
            )
        return VisualCriticResult(
            critic_id="checkpoint-critic/model-a",
            issues=issues,
        )

    def review(self, _slides):
        raise AssertionError("checkpoint critic must use page and deck scopes")


class _StopAwareBlockingVisualCritic:
    def __init__(self) -> None:
        self.started = threading.Event()
        self.release = threading.Event()
        self.cancelled = threading.Event()
        self.slide_calls: list[str] = []
        self._lock = threading.Lock()

    def cache_identity(self) -> str:
        return "stop-aware-critic/v1"

    def _record(self, slide_id: str) -> None:
        with self._lock:
            self.slide_calls.append(slide_id)

    def review_slide(self, slide):
        self._record(slide.slide_id)
        self.started.set()
        self.release.wait(timeout=2)
        return VisualCriticResult(critic_id="stop-aware-critic/v1", issues=[])

    def review_slide_cancellable(self, slide, ensure_active):
        self._record(slide.slide_id)
        self.started.set()
        try:
            while not self.release.wait(timeout=0.01):
                ensure_active()
        except PresentationJobInterruptedError:
            self.cancelled.set()
            raise
        return VisualCriticResult(critic_id="stop-aware-critic/v1", issues=[])

    def review_deck(self, _slides):
        return VisualCriticResult(critic_id="stop-aware-critic/v1", issues=[])

    def review(self, _slides):
        raise AssertionError("stop-aware critic must use scoped review")


class _ObservedPresentationExecutionLimits(PresentationExecutionLimits):
    def __init__(self) -> None:
        super().__init__(synthesis=1, rendering=1, critics=1)
        self.waiter_attempted = threading.Event()
        self._critic_attempts = 0
        self._attempt_lock = threading.Lock()

    @contextmanager
    def critic(self, *, ensure_active=None):
        with self._attempt_lock:
            self._critic_attempts += 1
            if self._critic_attempts >= 2:
                self.waiter_attempted.set()
        parent = (
            super().critic()
            if ensure_active is None
            else super().critic(ensure_active=ensure_active)
        )
        with parent:
            yield


def _two_slide_visual_deck() -> DeckIR:
    return DeckIR(
        id="visual-checkpoint-deck",
        title="视觉证据缓存",
        immutable_data=["两页都必须完成质检"],
        slides=[
            DeckSlide(
                id="visual-a",
                layout_id="cover",
                speaker_notes="先给出核心判断",
                objects=[
                    TextObject(
                        id="visual-a-title",
                        role="title",
                        text="核心判断",
                        slot="title",
                    ),
                ],
            ),
            DeckSlide(
                id="visual-b",
                layout_id="closing",
                speaker_notes="再给出下一步行动",
                objects=[
                    TextObject(
                        id="visual-b-title",
                        role="title",
                        text="下一步",
                        slot="title",
                    ),
                ],
            ),
        ],
    )


def test_visual_review_stop_cancels_inflight_call_without_more_page_cost(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    critic = _StopAwareBlockingVisualCritic()
    service = PresentationService(
        tmp_path,
        visual_critic=critic,
        execution_limits=PresentationExecutionLimits(
            synthesis=1,
            rendering=1,
            critics=1,
        ),
    )
    request = _request().model_copy(
        update={"cost_budget_units": 20},
    )
    record = service.create_job(request)
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    finished = threading.Event()
    outcome: dict[str, BaseException] = {}

    def verify() -> None:
        try:
            service.verify_artifact(
                deck,
                output,
                service.store.job_dir(record.job_id) / "renders" / "stopped",
                vision_critique_required=True,
                job_id=record.job_id,
            )
        except BaseException as exc:  # test thread must return the worker error
            outcome["error"] = exc
        finally:
            finished.set()

    worker = threading.Thread(target=verify)
    worker.start()
    assert critic.started.wait(timeout=1)
    service.store.cancel(record.job_id, record.conversation_id)
    stopped_promptly = finished.wait(timeout=0.5)
    critic.release.set()
    worker.join(timeout=2)

    assert stopped_promptly
    assert not worker.is_alive()
    assert isinstance(outcome.get("error"), PresentationJobInterruptedError)
    assert critic.cancelled.is_set()
    assert len(critic.slide_calls) == 1
    assert service.store.get(record.job_id).cost_used_units == 1
    checkpoints = (
        service.store.job_dir(record.job_id)
        / "quality"
        / "visual-checkpoints"
    )
    assert not list(checkpoints.glob("*.json"))


def test_visual_review_stop_leaves_waiting_critic_slots_uncharged(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    limits = _ObservedPresentationExecutionLimits()
    critic = _CheckpointVisualCritic()
    service = PresentationService(
        tmp_path,
        visual_critic=critic,
        execution_limits=limits,
    )
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    slot_held = threading.Event()
    release_slot = threading.Event()
    finished = threading.Event()
    outcome: dict[str, BaseException] = {}

    def occupy_slot() -> None:
        with limits.critic():
            slot_held.set()
            release_slot.wait(timeout=2)

    def verify() -> None:
        try:
            service.verify_artifact(
                deck,
                output,
                service.store.job_dir(record.job_id) / "renders" / "queued-stop",
                vision_critique_required=True,
                job_id=record.job_id,
            )
        except BaseException as exc:  # test thread must return the worker error
            outcome["error"] = exc
        finally:
            finished.set()

    holder = threading.Thread(target=occupy_slot)
    holder.start()
    assert slot_held.wait(timeout=1)
    worker = threading.Thread(target=verify)
    worker.start()
    assert limits.waiter_attempted.wait(timeout=1)
    service.store.cancel(record.job_id, record.conversation_id)
    stopped_promptly = finished.wait(timeout=0.5)
    release_slot.set()
    holder.join(timeout=2)
    worker.join(timeout=2)

    assert stopped_promptly
    assert not holder.is_alive()
    assert not worker.is_alive()
    assert isinstance(outcome.get("error"), PresentationJobInterruptedError)
    assert critic.slide_calls == []
    assert service.store.get(record.job_id).cost_used_units == 0


def test_visual_review_checkpoints_every_page_and_invalidates_dependencies(
    tmp_path,
    monkeypatch,
) -> None:
    render_revision = {"slide": None}

    def render(source, output_dir, *, timeout=120):
        result = _render_for_edit_tests(source, output_dir, timeout=timeout)
        if render_revision["slide"] is not None:
            image_path = result.slide_images[render_revision["slide"] - 1]
            with Image.open(image_path) as image:
                changed = image.copy()
            changed.putpixel((639, 359), (255, 0, 0))
            changed.save(image_path)
        return result

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    critic = _CheckpointVisualCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    first, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "visual-1",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert sorted(critic.slide_calls) == ["visual-a", "visual-b"]
    assert critic.deck_calls == 1
    assert len(first.visual_review_evidence) == 3
    assert not any(item.cache_hit for item in first.visual_review_evidence)
    assert service.store.get(record.job_id).cost_used_units == 3
    first_metrics = [
        item
        for item in service.store.get(record.job_id).stage_metrics
        if item.operation == "critic"
    ]
    assert {item.slide_id for item in first_metrics} == {
        "visual-a",
        "visual-b",
        None,
    }
    assert sum(item.cost_units for item in first_metrics) == 3
    assert all(
        item.provider_cost_status is PresentationProviderCostStatus.UNKNOWN
        for item in first_metrics
    )
    assert all(item.provider_cost_micro_cny is None for item in first_metrics)

    second, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "visual-2",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 2
    assert critic.deck_calls == 1
    assert all(item.cache_hit for item in second.visual_review_evidence)
    assert service.store.get(record.job_id).cost_used_units == 3
    cached_metrics = [
        item
        for item in service.store.get(record.job_id).stage_metrics
        if item.operation == "critic"
    ][len(first_metrics):]
    assert all(item.cost_units == 0 for item in cached_metrics)
    assert all(
        item.provider_cost_status
        is PresentationProviderCostStatus.NOT_APPLICABLE
        for item in cached_metrics
    )

    render_revision["slide"] = 2
    third, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "visual-3",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert critic.slide_calls.count("visual-a") == 1
    assert critic.slide_calls.count("visual-b") == 2
    assert critic.deck_calls == 2
    assert sum(not item.cache_hit for item in third.visual_review_evidence) == 2
    assert service.store.get(record.job_id).cost_used_units == 5

    shared_change = deck.model_copy(
        update={"immutable_data": ["共享事实变化必须失效全套缓存"]},
    )
    fourth, _ = service.verify_artifact(
        shared_change,
        output,
        service.store.job_dir(record.job_id) / "renders" / "visual-4",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert critic.slide_calls.count("visual-a") == 2
    assert critic.slide_calls.count("visual-b") == 3
    assert critic.deck_calls == 3
    assert not any(item.cache_hit for item in fourth.visual_review_evidence)
    assert service.store.get(record.job_id).cost_used_units == 8


def test_note_only_edit_invalidates_affected_page_and_deck_with_identical_pixels(
    tmp_path, monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)

    class NotesCritic(_CheckpointVisualCritic):
        def __init__(self) -> None:
            super().__init__()
            self.page_notes: dict[str, str | None] = {}
            self.deck_notes: list[list[str | None]] = []

        def review_slide(self, slide):
            self.page_notes[slide.slide_id] = getattr(slide, "speaker_notes", None)
            return super().review_slide(slide)

        def review_deck(self, slides):
            self.deck_notes.append([getattr(slide, "speaker_notes", None) for slide in slides])
            return super().review_deck(slides)

    critic = NotesCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request().model_copy(
        update={"composition_engine": CompositionEngine.NATIVE_PPTX},
    ))
    deck = _two_slide_visual_deck()
    for index, slide in enumerate(deck.slides, 1):
        slide.intent_id = f"notes-intent-{index}"
    plan = DeckPlan(
        title=deck.title, narrative_arc="先判断再行动",
        visual_direction=VisualDirection(
            visual_concept="editorial", palette=["#FFFFFF", "#111111"],
            typography="sans", image_language="documentary", chart_language="direct labels",
            shape_language="rules", grid="twelve columns", recurring_motif="folio",
        ), slides=[DeckPlanSlide(
            id=f"notes-plan-{index}", order=index, title="验证",
            intent=SlideIntent(id=slide.intent_id, purpose="explain", key_message="保留原计划判断"),
            content_blocks=["保留原计划判断"],
        ) for index, slide in enumerate(deck.slides, 1)],
    )
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    compiler = service.compiler_for(record.request.composition_engine)
    compiler.compile(deck, output)
    original_notes = [slide.speaker_notes for slide in deck.slides]
    original_deck = deck.model_dump(mode="json")
    original_plan = plan.model_dump(mode="json")

    def verify(candidate, name):
        return service.verify_artifact(
            candidate, output, service.store.job_dir(record.job_id) / "renders" / name,
            vision_critique_required=True, job_id=record.job_id, plan=plan,
        )

    first, first_images = verify(deck, "notes-original")
    assert sorted(critic.slide_calls) == ["visual-a", "visual-b"]
    assert critic.deck_calls == 1
    assert critic.page_notes == dict(zip([slide.id for slide in deck.slides], original_notes, strict=True))
    assert critic.deck_notes == [original_notes]
    assert not any(item.cache_hit for item in first.visual_review_evidence)
    assert service.store.get(record.job_id).cost_used_units == 3
    unchanged, _ = verify(deck, "notes-unchanged")
    assert all(item.cache_hit for item in unchanged.visual_review_evidence)
    assert len(critic.slide_calls) == 2 and critic.deck_calls == 1

    changed = deck.model_copy(deep=True)
    changed.slides[1].speaker_notes = "这次只修改第二页的原生备注，页面可见内容与意图保持原样"
    compiler.compile(changed, output)
    assert Presentation(str(output)).slides[1].notes_slide.notes_text_frame.text == changed.slides[1].speaker_notes
    revised, revised_images = verify(changed, "notes-revised")
    assert [file_sha256(path) for path in first_images] == [file_sha256(path) for path in revised_images]
    assert critic.slide_calls.count("visual-a") == 1
    assert critic.slide_calls.count("visual-b") == 2
    assert critic.deck_calls == 2
    assert critic.page_notes["visual-b"] == changed.slides[1].speaker_notes
    assert critic.deck_notes[-1] == [slide.speaker_notes for slide in changed.slides]
    assert [item.cache_hit for item in revised.visual_review_evidence] == [True, False, False]
    assert service.store.get(record.job_id).cost_used_units == 5
    original_fingerprints = {item.slide_id: item.dependency_fingerprint for item in first.visual_review_evidence}
    revised_fingerprints = {item.slide_id: item.dependency_fingerprint for item in revised.visual_review_evidence}
    assert original_fingerprints["visual-a"] == revised_fingerprints["visual-a"]
    assert original_fingerprints["visual-b"] != revised_fingerprints["visual-b"]
    assert original_fingerprints[None] != revised_fingerprints[None]
    metrics = [item for item in service.store.get(record.job_id).stage_metrics if item.operation == "critic"]
    assert sum(item.cost_units for item in metrics) == 5
    assert all(item.provider_cost_micro_cny is None for item in metrics if item.cost_units)
    assert deck.model_dump(mode="json") == original_deck
    assert plan.model_dump(mode="json") == original_plan


def test_visual_review_retries_only_failed_schema_scope_and_charges_both_calls(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )

    class SchemaRetryCritic(_CheckpointVisualCritic):
        def __init__(self) -> None:
            super().__init__()
            self.retry_directives: list[str | None] = []

        def review_deck(self, slides, *, retry_directive=None):
            self.deck_calls += 1
            self.retry_directives.append(retry_directive)
            if self.deck_calls == 1:
                raise StructuredTaskValidationError(
                    "structured utility task output failed schema validation",
                    validation_issues=(
                        "literal_error@$.issues[0].issue_type",
                    ),
                )
            return VisualCriticResult(
                critic_id="checkpoint-critic/model-a",
                issues=[],
            )

    critic = SchemaRetryCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    report, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "schema-retry",
        vision_critique_required=True,
        job_id=record.job_id,
    )

    assert report.visual_verified is True
    assert report.visual_evidence == "passed"
    assert sorted(critic.slide_calls) == ["visual-a", "visual-b"]
    assert critic.deck_calls == 2
    assert critic.retry_directives[0] is None
    assert "literal_error@$.issues[0].issue_type" in (
        critic.retry_directives[1] or ""
    )
    assert len(report.visual_review_evidence) == 3
    assert report.visual_review_evidence[-1].scope == "deck"
    assert report.visual_review_evidence[-1].cost_units == 2
    saved = service.store.get(record.job_id)
    assert saved.cost_used_units == 4
    deck_metrics = [
        item
        for item in saved.stage_metrics
        if item.operation == "critic" and item.slide_id is None
    ]
    assert [item.outcome for item in deck_metrics] == ["failed", "passed"]
    assert [item.cost_units for item in deck_metrics] == [1, 1]
    assert service.store.quality_checkpoint_path(
        record.job_id,
        "deck.json",
    ).is_file()


def test_visual_review_exhausted_schema_retry_keeps_typed_failure_and_metrics(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )

    class InvalidSchemaCritic(_CheckpointVisualCritic):
        def review_deck(self, _slides, *, retry_directive=None):
            self.deck_calls += 1
            raise StructuredTaskValidationError(
                "structured utility task output failed schema validation",
                validation_issues=(
                    "literal_error@$.issues[0].issue_type",
                ),
            )

    critic = InvalidSchemaCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    report, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "schema-exhausted",
        vision_critique_required=True,
        job_id=record.job_id,
    )

    assert report.visual_verified is False
    assert report.visual_evidence == "unavailable"
    assert critic.deck_calls == 2
    assert any(
        issue.code == "vision.critic_schema_invalid"
        for issue in report.issues
    )
    saved = service.store.get(record.job_id)
    assert saved.cost_used_units == 4
    deck_metrics = [
        item
        for item in saved.stage_metrics
        if item.operation == "critic" and item.slide_id is None
    ]
    assert [item.outcome for item in deck_metrics] == ["failed", "failed"]
    assert [item.cost_units for item in deck_metrics] == [1, 1]


@pytest.mark.parametrize("failures", [1, 3])
@pytest.mark.parametrize("private_location", ["root", "nested", "code", "index"])
def test_visual_schema_retry_redacts_model_owned_keys_and_messages(
    tmp_path, monkeypatch, failures, private_location,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    sentinel = "SYNTHETIC_PRIVATE_SCHEMA_KEY_4187"
    private_issue = {
        "root": f"extra_forbidden@$.{sentinel}",
        "nested": f"extra_forbidden@$.issues[0].{sentinel}",
        "code": f"{sentinel}@$",
        "index": f"literal_error@$.issues[{sentinel}].issue_type",
    }[private_location]

    class KeyCritic(_CheckpointVisualCritic):
        def __init__(self):
            super().__init__()
            self.directives = []

        def review_deck(self, slides, *, retry_directive=None):
            self.deck_calls += 1
            self.directives.append(retry_directive)
            if self.deck_calls <= failures:
                raise StructuredTaskValidationError(
                    sentinel,
                    validation_issues=(
                        private_issue,
                        "literal_error@$.issues[0].issue_type",
                    ),
                )
            return VisualCriticResult(critic_id="synthetic-key-critic", issues=[])

    critic = KeyCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    report, _ = service.verify_artifact(
        deck, output, service.store.job_dir(record.job_id) / "renders" / "schema-key",
        vision_critique_required=True, job_id=record.job_id,
    )
    assert critic.deck_calls == 2
    assert sentinel not in str(critic.directives)
    assert sentinel not in report.model_dump_json()
    assert "literal_error@$.issues[0].issue_type" in critic.directives[1]
    assert report.visual_verified == (failures == 1)
    assert service.store.get(record.job_id).cost_used_units == 4


def test_legacy_visual_schema_failure_does_not_echo_model_keys(tmp_path, monkeypatch):
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    sentinel = "SYNTHETIC_PRIVATE_SCHEMA_KEY_4187"

    class LegacyCritic:
        def review(self, slides):
            raise StructuredTaskValidationError(
                sentinel, validation_issues=(f"extra_forbidden@$.{sentinel}",),
            )

    service = PresentationService(tmp_path, visual_critic=LegacyCritic())
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    report, _ = service.verify_artifact(
        deck, output, service.store.job_dir(record.job_id) / "renders" / "legacy-schema-key",
        vision_critique_required=True, job_id=record.job_id,
    )
    assert not report.visual_verified
    assert any(i.code == "vision.critic_schema_invalid" for i in report.issues)
    assert sentinel not in report.model_dump_json()
    assert service.store.get(record.job_id).cost_used_units == 3


@pytest.mark.parametrize("invalid_reference", ["slide", "object"])
@pytest.mark.parametrize("failures", [1, 3])
def test_visual_review_bounds_reference_retries_and_never_caches_invalid_ids(
    tmp_path,
    monkeypatch,
    invalid_reference: str,
    failures: int,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx", _render_for_edit_tests,
    )
    rejected_id = "private-untrusted-reference-do-not-repeat"

    class ReferenceRetryCritic(_CheckpointVisualCritic):
        def __init__(self) -> None:
            super().__init__()
            self.retry_directives: list[str | None] = []

        def review_deck(self, slides, *, retry_directive=None):
            self.deck_calls += 1
            self.retry_directives.append(retry_directive)
            if self.deck_calls <= failures:
                issue = VisualIssue(
                    slide_id=(
                        rejected_id if invalid_reference == "slide"
                        else slides[0].slide_id
                    ),
                    severity="error",
                    issue_type="hierarchy",
                    object_ids=[rejected_id] if invalid_reference == "object" else [],
                    description="引用必须来自输入",
                    proposed_action="不得执行无效引用",
                )
                UtilityPresentationVisualCritic._validated_issues(
                    SimpleNamespace(value=SimpleNamespace(issues=[issue])), slides,
                )
            return VisualCriticResult(critic_id="checkpoint-critic/model-a", issues=[])

    critic = ReferenceRetryCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    original_sha256 = file_sha256(output)
    report, _ = service.verify_artifact(
        deck, output,
        service.store.job_dir(record.job_id) / "renders" / "reference-retry",
        vision_critique_required=True, job_id=record.job_id,
    )

    assert critic.deck_calls == 2
    assert sorted(critic.slide_calls) == ["visual-a", "visual-b"]
    assert critic.retry_directives[0] is None
    directive = critic.retry_directives[1] or ""
    field = "slide_id" if invalid_reference == "slide" else "object_ids"
    assert f"reference_error@$.issues[0].{field}" in directive
    assert rejected_id not in directive
    assert rejected_id not in report.model_dump_json()
    assert file_sha256(output) == original_sha256
    saved = service.store.get(record.job_id)
    assert saved.cost_used_units == 4
    metrics = [
        item for item in saved.stage_metrics
        if item.operation == "critic" and item.slide_id is None
    ]
    assert [item.cost_units for item in metrics] == [1, 1]
    assert [item.outcome for item in metrics] == (
        ["failed", "passed"] if failures == 1 else ["failed", "failed"]
    )
    checkpoint = service.store.quality_checkpoint_path(record.job_id, "deck.json")
    assert checkpoint.is_file() is (failures == 1)
    assert report.visual_verified is (failures == 1)
    assert report.visual_evidence == ("passed" if failures == 1 else "unavailable")
    if failures > 1:
        assert any(i.code == "vision.critic_schema_invalid" for i in report.issues)


def test_visual_reference_retry_honors_stop_before_second_call(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx", _render_for_edit_tests,
    )

    class StopDuringReferenceCritic(_CheckpointVisualCritic):
        def review_deck(self, slides, *, retry_directive=None):
            self.deck_calls += 1
            service.store.cancel(record.job_id, record.conversation_id)
            UtilityPresentationVisualCritic._validated_issues(
                SimpleNamespace(value=SimpleNamespace(issues=[VisualIssue(
                    slide_id="not-a-real-slide", severity="error",
                    issue_type="hierarchy", object_ids=[],
                    description="引用必须来自输入", proposed_action="不得执行",
                )])), slides,
            )
            raise AssertionError("invalid reference must be rejected")

    critic = StopDuringReferenceCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    original_sha256 = file_sha256(output)

    with pytest.raises(PresentationJobInterruptedError):
        service.verify_artifact(
            deck, output, service.store.job_dir(record.job_id) / "renders" / "reference-stop",
            vision_critique_required=True, job_id=record.job_id,
        )

    assert critic.deck_calls == 1
    assert service.store.get(record.job_id).cost_used_units == 3
    assert file_sha256(output) == original_sha256
    assert not service.store.quality_checkpoint_path(record.job_id, "deck.json").exists()


def test_deck_critic_ruleset_change_preserves_page_review_checkpoints(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    critic = _CheckpointVisualCritic()
    deck_identity = {"value": "checkpoint-deck/v1"}
    critic.slide_cache_identity = lambda: "checkpoint-slide/v1"
    critic.deck_cache_identity = lambda: deck_identity["value"]
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "split-identity-1",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    deck_identity["value"] = "checkpoint-deck/v2"
    report, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "split-identity-2",
        vision_critique_required=True,
        job_id=record.job_id,
    )

    assert sorted(critic.slide_calls) == ["visual-a", "visual-b"]
    assert critic.deck_calls == 2
    assert [item.cache_hit for item in report.visual_review_evidence] == [
        True,
        True,
        False,
    ]


def test_visual_review_cache_invalidates_for_shared_quality_dependencies(
    tmp_path,
    monkeypatch,
) -> None:
    render_identity = {"version": "1.0"}

    def render(source, output_dir, *, timeout=120):
        result = _render_for_edit_tests(source, output_dir, timeout=timeout)
        return RenderResult(
            pdf_path=result.pdf_path,
            slide_images=result.slide_images,
            renderer_id=result.renderer_id,
            renderer_version=render_identity["version"],
        )

    monkeypatch.setattr("lantai.presentation.service.render_pptx", render)
    critic = _CheckpointVisualCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "shared-1",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 2
    assert critic.deck_calls == 1

    themed = deck.model_copy(update={"theme_id": "lantai-warm"})
    service.compiler.compile(themed, output)
    theme_report, _ = service.verify_artifact(
        themed,
        output,
        service.store.job_dir(record.job_id) / "renders" / "shared-2",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 4
    assert critic.deck_calls == 2
    assert not any(item.cache_hit for item in theme_report.visual_review_evidence)

    monkeypatch.setitem(
        THEMES,
        "lantai-warm",
        replace(THEMES["lantai-warm"], body_font="Source Han Sans SC"),
    )
    service.compiler.compile(themed, output)
    font_report, _ = service.verify_artifact(
        themed,
        output,
        service.store.job_dir(record.job_id) / "renders" / "shared-font",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 6
    assert critic.deck_calls == 3
    assert not any(item.cache_hit for item in font_report.visual_review_evidence)

    render_identity["version"] = "2.0"
    renderer_report, _ = service.verify_artifact(
        themed,
        output,
        service.store.job_dir(record.job_id) / "renders" / "shared-3",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 8
    assert critic.deck_calls == 4
    assert not any(
        item.cache_hit for item in renderer_report.visual_review_evidence
    )

    monkeypatch.setattr(
        "lantai.presentation.service.VISUAL_REVIEW_RULESET_VERSION",
        "test-ruleset-vNext",
    )
    ruleset_report, _ = service.verify_artifact(
        themed,
        output,
        service.store.job_dir(record.job_id) / "renders" / "shared-4",
        vision_critique_required=True,
        job_id=record.job_id,
    )
    assert len(critic.slide_calls) == 10
    assert critic.deck_calls == 5
    assert not any(
        item.cache_hit for item in ruleset_report.visual_review_evidence
    )


def test_visual_review_binds_original_brief_and_invalidates_source_changes(
    tmp_path, monkeypatch,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    critic = _CheckpointVisualCritic()
    received = []
    original_review = critic.review_slide

    def review_slide(slide):
        received.append(slide)
        return original_review(slide)

    critic.review_slide = review_slide
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    brief = PresentationBrief(
        objective="保持来源语义", audience="内部评审", narrative="观察与局限",
        source_sections=[PresentationSourceSection(
            title="观察", body="有时停止，频率未知", source_ref="source-original",
        )],
        prohibited_changes=["不得夸大频率"],
    )
    brief_path = service.store.checkpoint_path(record.job_id, "brief.json")
    brief_path.write_text(brief.model_dump_json(), encoding="utf-8")

    def verify(round_id):
        return service.verify_artifact(
            deck, output, service.store.job_dir(record.job_id) / "renders" / round_id,
            vision_critique_required=True, job_id=record.job_id,
        )[0]

    verify("original-source")
    assert received[0].source_brief == brief
    assert received[0].original_user_request == record.request.user_request
    unchanged = verify("unchanged-source")
    assert len(received) == 2
    assert all(item.cache_hit for item in unchanged.visual_review_evidence)

    changed = brief.model_copy(deep=True)
    changed.source_sections[0].body = "尚无频率统计，不得推断大多数"
    brief_path.write_text(changed.model_dump_json(), encoding="utf-8")
    changed_report = verify("changed-source")
    assert len(received) == 4
    assert received[-1].source_brief == changed
    assert not any(item.cache_hit for item in changed_report.visual_review_evidence)


@pytest.mark.parametrize("failure_kind", ["json", "schema", "read", "utf8", "directory", "dangling_symlink"])
def test_visual_source_context_failure_is_private_and_cannot_reuse_cached_pass(
    tmp_path, monkeypatch, failure_kind,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    critic = _CheckpointVisualCritic()
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    original_hash = file_sha256(output)
    brief = PresentationBrief(
        objective="Synthetic privacy test", audience="test", narrative="test",
        source_sections=[PresentationSourceSection(
            title="test", body="valid synthetic source", source_ref="synthetic-source",
        )],
    )
    if failure_kind in {"directory", "dangling_symlink"}:
        # Make the historical fallback identical to the cached valid source;
        # otherwise the test would miss the stale-pass reuse edge case.
        brief = build_brief(record.request)
    brief_path = service.store.checkpoint_path(record.job_id, "brief.json")
    brief_path.write_text(brief.model_dump_json(), encoding="utf-8")

    def verify(name):
        return service.verify_artifact(
            deck, output, service.store.job_dir(record.job_id) / "renders" / name,
            vision_critique_required=True, job_id=record.job_id,
        )[0]

    assert verify("valid-source").visual_verified
    calls_before = (len(critic.slide_calls), critic.deck_calls)
    cost_before = service.store.get(record.job_id).cost_used_units
    sentinel = "SYNTHETIC-SOURCE-INPUT-MUST-NOT-ECHO-827"
    if failure_kind == "schema":
        payload = brief.model_dump(mode="json")
        payload["source_sections"][0]["body"] = [sentinel]
        brief_path.write_text(json.dumps(payload), encoding="utf-8")
    elif failure_kind == "json":
        brief_path.write_text('{"' + sentinel + '":', encoding="utf-8")
    elif failure_kind == "utf8":
        brief_path.write_bytes(sentinel.encode("utf-8") + b"\xff")
    elif failure_kind in {"directory", "dangling_symlink"}:
        brief_path.unlink()
        if failure_kind == "directory":
            brief_path.mkdir()
        else:
            brief_path.symlink_to(tmp_path / sentinel)
    else:
        original_read = Path.read_text

        def unreadable(path, *args, **kwargs):
            if path == brief_path:
                raise PermissionError(sentinel)
            return original_read(path, *args, **kwargs)

        monkeypatch.setattr(Path, "read_text", unreadable)

    report = verify("invalid-source")
    assert not report.passed
    assert not report.visual_verified
    assert report.visual_evidence == "unavailable"
    assert any(issue.code == "vision.source_context_unavailable" for issue in report.issues)
    assert not report.visual_review_evidence
    assert (len(critic.slide_calls), critic.deck_calls) == calls_before
    assert service.store.get(record.job_id).cost_used_units == cost_before
    assert file_sha256(output) == original_hash
    assert sentinel not in report.model_dump_json()
    failure = next(issue for issue in report.issues if issue.code == "vision.source_context_unavailable")
    assert "sha256=" in failure.message


@pytest.mark.parametrize("issue_type", ["content_accuracy", "source_coverage"])
def test_semantic_source_errors_remain_blocking_quality_issues(
    tmp_path, monkeypatch, issue_type,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)

    class SourceCritic:
        def review(self, slides):
            return VisualCriticResult(critic_id="source-critic", issues=[VisualIssue(
                slide_id=slides[0].slide_id, severity="error", issue_type=issue_type,
                object_ids=[], description="原始来源不支持该表述或要求遗漏",
                proposed_action="依据原始材料修订，保持品牌硬约束",
            )])

    service = PresentationService(tmp_path, visual_critic=SourceCritic())
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    report, _ = service.verify_artifact(
        deck, output, service.store.job_dir(record.job_id) / "renders" / "source-error",
        vision_critique_required=True, job_id=record.job_id,
    )
    assert not report.passed
    assert not report.visual_verified
    assert any(item.code == f"vision.{issue_type}" and item.severity == "error" for item in report.issues)


@pytest.mark.parametrize("issue_type", ["content_accuracy", "source_coverage"])
@pytest.mark.parametrize("reported_severity", ["info", "warning", "error"])
def test_semantic_source_findings_cannot_be_downgraded_to_deliverable(
    tmp_path, monkeypatch, issue_type, reported_severity,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    deck = _two_slide_visual_deck()
    finding = VisualIssue(
        slide_id=deck.slides[0].id, severity=reported_severity, issue_type=issue_type,
        object_ids=[],
        description="来源里的数量说明被用作指标，或反例图形与自己的结论矛盾",
        proposed_action="依原始材料纠正可见内容，不删除必需事实",
    )

    class SourceCritic:
        def review(self, slides):
            assert finding.slide_id in {slide.slide_id for slide in slides}
            return VisualCriticResult(critic_id="source-critic", issues=[finding])

    service = PresentationService(tmp_path, visual_critic=SourceCritic())
    record = service.create_job(_request())
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    artifact_hash = file_sha256(output)

    report, _ = service.verify_artifact(
        deck, output, service.store.job_dir(record.job_id) / "renders" / "source-finding",
        vision_critique_required=True, job_id=record.job_id,
    )

    assert not report.passed
    assert report.delivery_status == "failed"
    assert not report.visual_verified
    issue = next(item for item in report.issues if item.code == f"vision.{issue_type}")
    assert issue.severity == "error" and issue.repairable
    assert report.visual_issues[0].severity == "error"
    assert finding.severity == reported_severity
    assert file_sha256(output) == artifact_hash


@pytest.mark.parametrize("reported_severity", ["info", "warning"])
def test_non_source_visual_findings_keep_their_advisory_severity(
    tmp_path, monkeypatch, reported_severity,
) -> None:
    monkeypatch.setattr("lantai.presentation.service.render_pptx", _render_for_edit_tests)
    deck = _two_slide_visual_deck()
    finding = VisualIssue(
        slide_id=deck.slides[0].id, severity=reported_severity,
        issue_type="alignment", object_ids=[],
        description="可进一步优化留白对齐", proposed_action="小幅对齐",
    )

    class AdvisoryCritic:
        def review(self, slides):
            return VisualCriticResult(critic_id="advisory-critic", issues=[finding])

    service = PresentationService(tmp_path, visual_critic=AdvisoryCritic())
    record = service.create_job(_request())
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)
    artifact_hash = file_sha256(output)
    report, _ = service.verify_artifact(
        deck, output, service.store.job_dir(record.job_id) / "renders" / "advisory",
        vision_critique_required=True, job_id=record.job_id,
    )

    assert report.visual_verified
    issue = next(item for item in report.issues if item.code == "vision.alignment")
    assert issue.severity == reported_severity and not issue.repairable
    assert report.visual_issues[0].severity == reported_severity
    assert finding.severity == reported_severity
    assert file_sha256(output) == artifact_hash


def test_whole_deck_visual_review_is_a_blocking_quality_gate(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    critic = _CheckpointVisualCritic(deck_issue=True)
    service = PresentationService(tmp_path, visual_critic=critic)
    record = service.create_job(_request())
    deck = _two_slide_visual_deck()
    output = service.store.job_dir(record.job_id) / "candidate.pptx"
    service.compiler.compile(deck, output)

    report, _ = service.verify_artifact(
        deck,
        output,
        service.store.job_dir(record.job_id) / "renders" / "deck-gate",
        vision_critique_required=True,
        job_id=record.job_id,
    )

    assert critic.deck_calls == 1
    assert report.visual_verified is False
    assert report.visual_evidence == "failed"
    assert any(issue.code == "vision.repetition" for issue in report.issues)


def test_redo_slide_can_redesign_one_v3_slide_and_requires_visual_critic(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    runtime = PresentationJobRuntime(
        tmp_path,
        visual_critic=_PassingVisualCritic(),
        repair_designer=_WholeSlideDesigner(),
    )
    request = _request().model_copy(
        update={"composition_engine": CompositionEngine.PPTXGENJS_V3},
    )
    record = runtime.service.create_job(request)
    deck = DeckIR(
        id="semantic-redesign-deck",
        title="自然语言整页重做",
        slides=[
            DeckSlide(
                id="slide-a",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="old-title",
                        role="title",
                        text="原始标题",
                        slot="title",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-b",
                layout_id="closing",
                objects=[
                    TextObject(
                        id="untouched-title",
                        role="title",
                        text="第二页保持不变",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    runtime.service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    output = (
        runtime.service.store.job_dir(record.job_id)
        / "output"
        / "presentation.pptx"
    )
    runtime.service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
        deck,
        output,
    )
    report, _images = runtime.service.verify_artifact(
        deck,
        output,
        runtime.service.store.job_dir(record.job_id) / "renders" / "round-0",
        render_set_id="round-0",
    )
    quality_path = (
        runtime.service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    runtime.service.store.append_event(
        record.job_id,
        stage=JobStage.COMPLETED,
        status=JobStatus.COMPLETED,
        message="verified v3 baseline",
        checkpoint=str(output),
    )
    before_hash = file_sha256(output)

    try:
        updated = runtime.redo_slide(
            record.job_id,
            "conversation-a",
            "slide-a",
            "把这一页做得更高级，突出可信交付",
            before_hash,
        )
        changed = DeckIR.model_validate_json(
            runtime.service.store.checkpoint_path(
                record.job_id,
                "deck-ir.json",
            ).read_text(encoding="utf-8"),
        )
        quality = runtime.artifacts(record.job_id, "conversation-a")["quality"]

        assert updated.status == JobStatus.COMPLETED
        assert changed.slides[0].layout_id == "freeform"
        assert any(
            relationship.relationship_type == RelationshipType.QUOTE
            for relationship in changed.slides[0].relationships
        )
        assert any(
            obj.id == "redesign-title" for obj in changed.slides[0].objects
        )
        untouched_before = prepare_deck_for_composition(deck).slides[1]
        untouched_after = changed.slides[1]
        assert untouched_after.id == untouched_before.id
        assert [obj.id for obj in untouched_after.objects] == [
            obj.id for obj in untouched_before.objects
        ]
        assert [getattr(obj, "text", None) for obj in untouched_after.objects] == [
            getattr(obj, "text", None) for obj in untouched_before.objects
        ]
        assert all(obj.frame is not None for obj in untouched_after.objects)
        assert file_sha256(output) != before_hash
        assert quality["vision_critique_required"] is True
        assert quality["visual_verified"] is True
        assert quality["visual_critic_id"] == "test-vision/1"
        metered = runtime.service.store.get(record.job_id)
        assert metered.cost_used_units == 11
        critic_metrics = [
            item for item in metered.stage_metrics
            if item.operation == "critic"
        ]
        assert len(critic_metrics) == 1
        assert critic_metrics[0].cost_units == 3
        assert (
            critic_metrics[0].provider_cost_status
            is PresentationProviderCostStatus.UNKNOWN
        )
        assert updated.events[-1].message.startswith("整页重做已提交")
    finally:
        runtime.shutdown()


@pytest.mark.parametrize(
    ("instruction", "base_hash", "error"),
    [
        (
            "把本页标题改为：不会提交",
            "0" * 64,
            "文件已被其他操作修改",
        ),
        (
            "让这一页更高级",
            None,
            "暂未能把要求安全映射",
        ),
    ],
)
def test_redo_slide_rejects_stale_or_unmapped_edits_without_touching_source(
    tmp_path,
    monkeypatch,
    instruction,
    base_hash,
    error,
) -> None:
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path,
        monkeypatch,
    )
    before_hash = file_sha256(output)
    try:
        with pytest.raises((RuntimeError, ValueError), match=error):
            runtime.redo_slide(
                job_id,
                "conversation-a",
                "slide-a",
                instruction,
                base_hash or before_hash,
            )
        assert file_sha256(output) == before_hash
        assert (
            runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
        )
    finally:
        runtime.shutdown()


def test_redo_slide_failure_after_commit_preserves_newer_human_save(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path,
        monkeypatch,
    )
    before_hash = file_sha256(output)
    original_record = PresentationVersionStore.record

    def save_newer_human_version_then_fail(self, *args, **kwargs):
        if kwargs.get("writer") != "lantai_ai_edit":
            return original_record(self, *args, **kwargs)
        presentation = Presentation(str(output))
        title = next(
            shape
            for shape in presentation.slides[0].shapes
            if shape.name == "title-a"
        )
        title.text = "AI 提交后用户又保存的新标题"
        presentation.save(str(output))
        raise RuntimeError("simulated version persistence failure")

    monkeypatch.setattr(
        PresentationVersionStore,
        "record",
        save_newer_human_version_then_fail,
    )
    try:
        with pytest.raises(RuntimeError, match="version persistence failure"):
            runtime.redo_slide(
                job_id,
                "conversation-a",
                "slide-a",
                "把本页标题改为：AI 候选标题",
                before_hash,
            )

        presentation = Presentation(str(output))
        title = next(
            shape
            for shape in presentation.slides[0].shapes
            if shape.name == "title-a"
        )
        assert title.text == "AI 提交后用户又保存的新标题"
        assert file_sha256(output) != before_hash
        assert runtime.get(job_id, "conversation-a").status == JobStatus.FAILED
    finally:
        runtime.shutdown()


def test_redo_slide_does_not_promote_a_human_save_racing_version_record(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path,
        monkeypatch,
    )
    before_hash = file_sha256(output)
    original_record = PresentationVersionStore.record

    def save_newer_human_version_before_record(self, *args, **kwargs):
        if kwargs.get("writer") == "lantai_ai_edit":
            presentation = Presentation(str(output))
            title = next(
                shape
                for shape in presentation.slides[0].shapes
                if shape.name == "title-a"
            )
            title.text = "版本登记前用户保存的新标题"
            presentation.save(str(output))
        return original_record(self, *args, **kwargs)

    monkeypatch.setattr(
        PresentationVersionStore,
        "record",
        save_newer_human_version_before_record,
    )
    try:
        with pytest.raises(RuntimeError, match="changed before version record"):
            runtime.redo_slide(
                job_id,
                "conversation-a",
                "slide-a",
                "把本页标题改为：不能覆盖用户保存的 AI 候选",
                before_hash,
            )

        presentation = Presentation(str(output))
        title = next(
            shape
            for shape in presentation.slides[0].shapes
            if shape.name == "title-a"
        )
        assert title.text == "版本登记前用户保存的新标题"
        assert file_sha256(output) != before_hash
        versions = PresentationVersionStore(
            runtime.service.store.job_dir(job_id),
        ).history()
        assert len(versions) == 1
        assert versions[0].pptx_sha256 == before_hash
        assert runtime.get(job_id, "conversation-a").status == JobStatus.FAILED
    finally:
        runtime.shutdown()


def test_redo_slide_failure_rolls_back_its_unchanged_ai_commit(
    tmp_path,
    monkeypatch,
) -> None:
    runtime, job_id, _deck, output = _completed_editable_runtime(
        tmp_path,
        monkeypatch,
    )
    before_hash = file_sha256(output)
    original_record = PresentationVersionStore.record

    def fail_ai_version_record(self, *args, **kwargs):
        if kwargs.get("writer") != "lantai_ai_edit":
            return original_record(self, *args, **kwargs)
        raise RuntimeError("simulated version persistence failure")

    monkeypatch.setattr(
        PresentationVersionStore,
        "record",
        fail_ai_version_record,
    )
    try:
        with pytest.raises(RuntimeError, match="version persistence failure"):
            runtime.redo_slide(
                job_id,
                "conversation-a",
                "slide-a",
                "把本页标题改为：不会留存的 AI 候选",
                before_hash,
            )

        presentation = Presentation(str(output))
        title = next(
            shape
            for shape in presentation.slides[0].shapes
            if shape.name == "title-a"
        )
        assert title.text == "原始标题"
        assert file_sha256(output) == before_hash
        assert runtime.get(job_id, "conversation-a").status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


def test_job_events_are_monotonic_persistent_and_conversation_scoped(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    first = store.append_event(
        record.job_id,
        stage=JobStage.BRIEF,
        status=JobStatus.RUNNING,
        message="brief ready",
    )
    second = store.append_event(
        record.job_id,
        stage=JobStage.PLAN,
        status=JobStatus.RUNNING,
        message="plan ready",
    )

    reloaded = PresentationJobStore(tmp_path).get(record.job_id)
    assert (first.sequence, second.sequence, reloaded.last_sequence) == (
        1,
        2,
        2,
    )
    assert [event.sequence for event in reloaded.events] == [1, 2]
    assert [
        event.sequence
        for event in store.events_after(record.job_id, 1, "conversation-a")
    ] == [2]
    with pytest.raises(PermissionError):
        store.events_after(record.job_id, 0, "conversation-b")


def test_job_pause_resume_and_cancel_are_explicit_states(tmp_path) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    store.append_event(
        record.job_id,
        stage=JobStage.PLAN,
        status=JobStatus.RUNNING,
        message="planning",
    )
    assert (
        store.pause(record.job_id, "conversation-a").status == JobStatus.PAUSED
    )
    assert (
        store.resume(record.job_id, "conversation-a").status
        == JobStatus.RUNNING
    )
    assert (
        store.cancel(record.job_id, "conversation-a").stage
        == JobStage.CANCELLED
    )


def test_background_stage_event_cannot_revive_paused_or_cancelled_job(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    store.append_event(
        record.job_id,
        stage=JobStage.PLAN,
        status=JobStatus.RUNNING,
        message="planning",
    )
    store.pause(record.job_id, "conversation-a")
    with pytest.raises(PresentationJobInterruptedError):
        store.append_event(
            record.job_id,
            stage=JobStage.RENDER,
            status=JobStatus.RUNNING,
            message="late render event",
        )
    assert store.get(record.job_id).status == JobStatus.PAUSED

    store.resume(record.job_id, "conversation-a")
    store.cancel(record.job_id, "conversation-a")
    with pytest.raises(PresentationJobInterruptedError):
        store.append_event(
            record.job_id,
            stage=JobStage.COMPLETED,
            status=JobStatus.COMPLETED,
            message="late completion event",
        )
    cancelled = store.get(record.job_id)
    assert cancelled.status == JobStatus.CANCELLED
    assert all(event.status != JobStatus.COMPLETED for event in cancelled.events)


def test_verified_publication_cannot_run_twice_after_completion(tmp_path) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    store.append_event(
        record.job_id,
        stage=JobStage.PLAN,
        status=JobStatus.RUNNING,
        message="planning",
    )
    publications: list[str] = []
    store.complete_after_verified_publication(
        record.job_id,
        message="first delivery",
        checkpoint="/first.pptx",
        publish=lambda: publications.append("first"),
    )

    with pytest.raises(ValueError, match="completed"):
        store.complete_after_verified_publication(
            record.job_id,
            message="duplicate delivery",
            checkpoint="/second.pptx",
            publish=lambda: publications.append("second"),
        )

    assert publications == ["first"]
    assert store.get(record.job_id).checkpoints[JobStage.COMPLETED.value] == (
        "/first.pptx"
    )


def test_verified_edit_publication_uses_own_deadline_after_generation_expired(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    store.append_event(
        record.job_id,
        stage=JobStage.COMPLETED,
        status=JobStatus.COMPLETED,
        message="original verified delivery",
    )
    completed = store.get(record.job_id)
    original_deadline = datetime.now(timezone.utc) - timedelta(seconds=1)
    completed.deadline_at = original_deadline
    store._write(completed)
    publications: list[str] = []

    with store.writer_lease(record.job_id):
        store.append_event(
            record.job_id,
            stage=JobStage.REPAIR,
            status=JobStatus.RUNNING,
            message="explicit edit",
        )
        store.complete_after_verified_publication(
            record.job_id,
            message="verified edit delivery",
            checkpoint="/edited.pptx",
            publish=lambda: publications.append("edit"),
        )

    settled = store.get(record.job_id)
    assert publications == ["edit"]
    assert settled.status == JobStatus.COMPLETED
    assert settled.deadline_at == original_deadline
    assert settled.operation_history[-1].status == JobStatus.COMPLETED
    assert settled.operation_history[-1].deadline_at > original_deadline


def test_failed_job_requires_classified_retry_before_publication(tmp_path) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    store.append_event(
        record.job_id,
        stage=JobStage.RENDER,
        status=JobStatus.FAILED,
        message="render failed",
        recoverable=True,
        failure_routes=[
            failure_route_for_exception(
                ProviderError("render provider unavailable"),
                stage=JobStage.RENDER,
            )
        ],
    )
    publications: list[str] = []

    with pytest.raises(PresentationJobInterruptedError, match="failed"):
        store.complete_after_verified_publication(
            record.job_id,
            message="late delivery",
            checkpoint="/late.pptx",
            publish=lambda: publications.append("late"),
        )
    with pytest.raises(PresentationJobInterruptedError, match="failed"):
        store.append_event(
            record.job_id,
            stage=JobStage.RENDER,
            status=JobStatus.RUNNING,
            message="late worker revival",
        )
    with pytest.raises(PresentationJobInterruptedError, match="failed"):
        store.assert_execution_budget(record.job_id)

    assert publications == []
    assert store.get(record.job_id).status == JobStatus.FAILED
    with pytest.raises(ValueError, match="paused"):
        store.resume(record.job_id, "conversation-a")
    store.retry_failed(record.job_id, "conversation-a")
    store.complete_after_verified_publication(
        record.job_id,
        message="resumed delivery",
        checkpoint="/resumed.pptx",
        publish=lambda: publications.append("resumed"),
    )
    assert publications == ["resumed"]


def test_service_records_phase_timings_without_changing_job_state(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        title="阶段计时",
        slides=[
            DeckSlide(
                id="timed-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="timed-title",
                        role="title",
                        text="每一阶段都留下证据",
                        slot="title",
                    ),
                ],
            ),
        ],
    )

    service.compose(record.job_id, deck)
    completed = service.store.get(record.job_id)

    assert completed.status == JobStatus.COMPLETED
    assert {metric.operation for metric in completed.stage_metrics} >= {
        "budget_reservation",
        "compile",
        "render",
        "inspect",
        "promotion",
    }
    assert all(metric.duration_ms >= 0 for metric in completed.stage_metrics)
    assert sum(metric.cost_units for metric in completed.stage_metrics) == 8
    assert PresentationJobStore(tmp_path).get(record.job_id).stage_metrics == (
        completed.stage_metrics
    )


def test_user_cancel_during_blocking_render_prevents_promotion(
    tmp_path,
    monkeypatch,
) -> None:
    service = PresentationService(tmp_path)
    record = service.create_job(_request())

    def render_then_cancel(source, output_dir, *, timeout=120):
        rendered = _render_for_edit_tests(source, output_dir, timeout=timeout)
        service.store.cancel(record.job_id, "conversation-a")
        return rendered

    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        render_then_cancel,
    )
    deck = DeckIR(
        title="停止边界",
        slides=[
            DeckSlide(
                id="cancelled-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="cancelled-title",
                        role="title",
                        text="用户停止后不得发布",
                        slot="title",
                    ),
                ],
            ),
        ],
    )

    with pytest.raises(PresentationJobInterruptedError, match="cancelled"):
        service.compose(record.job_id, deck)

    cancelled = service.store.get(record.job_id)
    assert cancelled.status == JobStatus.CANCELLED
    assert cancelled.stage == JobStage.CANCELLED
    assert not (
        service.store.job_dir(record.job_id)
        / "output"
        / "presentation.pptx"
    ).exists()
    assert all(event.status != JobStatus.COMPLETED for event in cancelled.events)


def test_stop_racing_final_promotion_cannot_split_job_and_version_state(
    tmp_path,
    monkeypatch,
) -> None:
    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        _render_for_edit_tests,
    )
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        title="最终发布竞争",
        slides=[
            DeckSlide(
                id="race-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="race-title",
                        role="title",
                        text="发布与停止必须有单一顺序",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    entered_promotion = threading.Event()
    release_promotion = threading.Event()
    original_promote = service._promote_candidate
    outcomes: dict[str, object] = {}

    def controlled_promote(candidate, output):
        entered_promotion.set()
        assert release_promotion.wait(timeout=5)
        return original_promote(candidate, output)

    def compose():
        try:
            outcomes["compose"] = service.compose(record.job_id, deck)
        except Exception as exc:  # surfaced to the assertion below
            outcomes["compose_error"] = exc

    def cancel():
        try:
            outcomes["cancel"] = service.store.cancel(
                record.job_id,
                "conversation-a",
            )
        except Exception as exc:  # surfaced to the assertion below
            outcomes["cancel_error"] = exc

    monkeypatch.setattr(service, "_promote_candidate", controlled_promote)
    compose_thread = threading.Thread(target=compose)
    cancel_thread = threading.Thread(target=cancel)
    compose_thread.start()
    try:
        assert entered_promotion.wait(timeout=10)
        cancel_thread.start()
        # Without one publication boundary Stop can finish here, and the
        # worker later records a version for a cancelled job.
        cancel_thread.join(timeout=0.5)
        release_promotion.set()
        cancel_thread.join(timeout=10)
        compose_thread.join(timeout=10)
    finally:
        release_promotion.set()
        compose_thread.join(timeout=10)
        if cancel_thread.ident is not None:
            cancel_thread.join(timeout=10)

    assert not compose_thread.is_alive()
    assert not cancel_thread.is_alive()
    assert "compose_error" not in outcomes
    assert isinstance(outcomes.get("cancel_error"), ValueError)
    assert service.store.get(record.job_id).status == JobStatus.COMPLETED
    output = service.store.job_dir(record.job_id) / "output" / "presentation.pptx"
    assert PresentationVersionStore(service.store.job_dir(record.job_id)).current().pptx_sha256 == file_sha256(output)


def test_job_mutations_reject_cross_conversation_access(tmp_path) -> None:
    store = PresentationJobStore(tmp_path)
    record = store.create(_request())
    with pytest.raises(PermissionError):
        store.pause(record.job_id, "conversation-b")


def test_pending_choice_cancel_rejects_a_job_owned_by_another_run(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path / "presentation-jobs")
    record = store.create(
        _request(),
        runtime_run_id="run-ppt-other",
    )

    with pytest.raises(PermissionError, match="pending Runtime Run"):
        cancel_active_job_for_runtime_run(
            tmp_path,
            "run-ppt-source",
            record.job_id,
        )

    assert store.get(record.job_id).status == JobStatus.PENDING


def test_pending_choice_cancel_rejects_a_job_from_another_conversation(
    tmp_path,
) -> None:
    store = PresentationJobStore(tmp_path / "presentation-jobs")
    record = store.create(
        _request("conversation-private"),
        runtime_run_id="run-ppt-source",
    )

    with pytest.raises(PermissionError, match="different conversation"):
        cancel_active_job_for_runtime_run(
            tmp_path,
            "run-ppt-source",
            record.job_id,
            conversation_id="conversation-current",
        )

    assert store.get(record.job_id).status == JobStatus.PENDING


def test_service_composes_from_a_durable_deck_ir_checkpoint(tmp_path) -> None:
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        id="service-deck",
        title="可恢复生成",
        slides=[
            DeckSlide(
                id="service-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="service-title",
                        role="title",
                        text="可恢复生成",
                        slot="title",
                    ),
                ],
            ),
        ],
    )

    result, report = service.compose(record.job_id, deck)
    completed = service.store.get(record.job_id)

    assert result.output_path.is_file()
    assert report.passed is True
    assert completed.status == JobStatus.COMPLETED
    assert completed.last_sequence >= 5
    assert service.store.checkpoint_path(
        record.job_id,
        "deck-ir.json",
    ).is_file()


def test_service_fails_when_required_visual_verification_is_unavailable(
    tmp_path,
    monkeypatch,
) -> None:
    render_calls = 0

    def unavailable(*_args, **_kwargs):
        nonlocal render_calls
        render_calls += 1
        raise PresentationRenderUnavailable("LibreOffice is unavailable")

    monkeypatch.setattr("lantai.presentation.service.render_pptx", unavailable)
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        id="structural-only-deck",
        title="仅结构检查",
        slides=[
            DeckSlide(
                id="structural-only-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="structural-only-title",
                        role="title",
                        text="结构完整不等于视觉已验证",
                        slot="title",
                    ),
                ],
            ),
        ],
    )

    result, report = service.compose(record.job_id, deck)
    failed = service.store.get(record.job_id)
    authoritative = (
        service.store.job_dir(record.job_id) / "output" / "presentation.pptx"
    )

    assert result.output_path.is_file()
    assert result.output_path.parent.name == "round-0"
    assert authoritative.exists() is False
    assert report.passed is False
    assert report.visual_verified is False
    assert report.delivery_status == "failed"
    assert failed.status == JobStatus.FAILED
    assert any(issue.code == "render.unavailable" for issue in report.issues)
    assert "质量检查未通过" in failed.events[-1].message
    assert render_calls == 1


class _RejectedRepairDesigner:
    def redesign(self, _request):
        raise ValueError("object ids must be unique across the deck")


class _ServiceSplitDesigner:
    def __init__(self) -> None:
        self.requests = []

    def redesign(self, request):
        self.requests.append(request)
        assert request.planned_slide is not None
        original_plan = request.planned_slide
        second_intent = SlideIntent(
            id="intent-a-evidence",
            purpose="prove",
            key_message="证据边界单独成页",
        )
        return SlideRedesignProposal(
            slide=DeckSlide(
                id=request.slide.id,
                intent_id=request.slide.intent_id,
                layout_id="freeform",
                speaker_notes="先讲核心判断",
                objects=[
                    TextObject(
                        id="split-title",
                        role="title",
                        text="核心判断",
                    ),
                ],
            ),
            action=RepairActionType.SPLIT_SLIDE,
            rationale="密度问题需要拆页",
            additional_slides=(
                DeckSlide(
                    id="slide-a-evidence",
                    intent_id=second_intent.id,
                    layout_id="freeform",
                    speaker_notes="再讲证据边界",
                    objects=[
                        TextObject(
                            id="split-evidence",
                            role="body",
                            text="证据边界",
                        ),
                    ],
                ),
            ),
            plan_slides=(
                original_plan.model_copy(
                    update={"title": "核心判断"},
                ),
                DeckPlanSlide(
                    id="plan-a-evidence",
                    order=original_plan.order + 1,
                    title="证据边界",
                    intent=second_intent,
                ),
            ),
        )


class _MeteredMultiSlideDesigner:
    def __init__(self) -> None:
        self.calls: list[str] = []
        self.last_model_label: str | None = None

    def redesign(self, request):
        self.calls.append(request.slide.id)
        self.last_model_label = f"approved-provider/{request.slide.id}"
        replacement_objects = []
        changed_text = False
        for obj in request.slide.objects:
            if isinstance(obj, TextObject) and obj.role != "title":
                replacement_objects.append(
                    obj.model_copy(update={"text": f"{obj.text}（已修复）"}),
                )
                changed_text = True
            else:
                replacement_objects.append(obj)
        if not changed_text:
            replacement_objects.append(
                TextObject(
                    id=f"{request.slide.id}-repair-note",
                    role="body",
                    text="页面层级（已修复）",
                ),
            )
        return SlideRedesignProposal(
            slide=request.slide.model_copy(
                update={"objects": replacement_objects},
            ),
            action=RepairActionType.REDESIGN_SLIDE,
            rationale="每个失败页独立调用受控设计模型",
        )


class _PartiallyRejectedMultiSlideDesigner(_MeteredMultiSlideDesigner):
    def redesign(self, request):
        if request.slide.id == "repair-b":
            self.calls.append(request.slide.id)
            raise ValueError("second page proposal is invalid")
        return super().redesign(request)


def test_service_charges_each_repair_designer_call_to_parent_job(
    tmp_path,
    monkeypatch,
) -> None:
    designer = _MeteredMultiSlideDesigner()
    service = PresentationService(tmp_path, repair_designer=designer)
    request = _request().model_copy(
        update={"composition_engine": CompositionEngine.NATIVE_PPTX},
    )
    record = service.create_job(request)
    deck = DeckIR(
        id="multi-slide-repair-deck",
        title="逐调用计费",
        slides=[
            DeckSlide(
                id=slide_id,
                layout_id="freeform",
                objects=[
                    TextObject(
                        id=f"{slide_id}-title",
                        role="title",
                        text=f"{slide_id} 原稿",
                    ),
                ],
            )
            for slide_id in ("repair-a", "repair-b")
        ],
    )
    issues = [
        QualityIssue(
            code="vision.hierarchy",
            severity="error",
            slide_id=slide_id,
            message="页面层级不足",
            repairable=True,
        )
        for slide_id in ("repair-a", "repair-b")
    ]
    verification_calls = 0

    def verify(current_deck, output, render_dir, *_args, **_kwargs):
        nonlocal verification_calls
        verification_calls += 1
        Path(render_dir).mkdir(parents=True, exist_ok=True)
        passed = verification_calls == 2
        return (
            QualityReport(
                deck_id=current_deck.id,
                passed=passed,
                delivery_status="passed" if passed else "failed",
                issues=[] if passed else issues,
                output_sha256=file_sha256(output),
                render_set_id=f"round-{verification_calls - 1}",
            ),
            (),
        )

    monkeypatch.setattr(
        service,
        "compiler_for",
        lambda _engine: service.compiler,
    )
    monkeypatch.setattr(service, "verify_artifact", verify)
    monkeypatch.setattr(
        "lantai.presentation.service.is_verified_delivery",
        lambda *_args, **_kwargs: True,
    )

    _result, report = service.compose(record.job_id, deck)

    completed = service.store.get(record.job_id)
    repair_metrics = [
        metric
        for metric in completed.stage_metrics
        if metric.operation == "repair"
    ]
    assert report.passed is True
    assert designer.calls == ["repair-a", "repair-b"]
    assert [metric.slide_id for metric in repair_metrics] == [
        "repair-a",
        "repair-b",
    ]
    assert [metric.route_id for metric in repair_metrics] == [
        "approved-provider/repair-a",
        "approved-provider/repair-b",
    ]
    assert [metric.cost_units for metric in repair_metrics] == [1, 1]
    assert all(
        metric.provider_cost_status
        is PresentationProviderCostStatus.UNKNOWN
        for metric in repair_metrics
    )
    assert all(
        metric.provider_cost_micro_cny is None for metric in repair_metrics
    )


def test_service_checkpoints_each_accepted_page_before_later_repair_failure(
    tmp_path,
    monkeypatch,
) -> None:
    designer = _PartiallyRejectedMultiSlideDesigner()
    service = PresentationService(tmp_path, repair_designer=designer)
    request = _request().model_copy(
        update={"composition_engine": CompositionEngine.NATIVE_PPTX},
    )
    record = service.create_job(request)
    deck = DeckIR(
        id="partial-page-repair-deck",
        title="页级恢复",
        slides=[
            DeckSlide(
                id=slide_id,
                layout_id="freeform",
                objects=[
                    TextObject(
                        id=f"{slide_id}-title",
                        role="title",
                        text=f"{slide_id} 原稿",
                    ),
                ],
            )
            for slide_id in ("repair-a", "repair-b")
        ],
    )
    issues = [
        QualityIssue(
            code="vision.hierarchy",
            severity="error",
            slide_id=slide_id,
            message="页面层级不足",
            repairable=True,
        )
        for slide_id in ("repair-a", "repair-b")
    ]

    def verify(current_deck, output, render_dir, *_args, **_kwargs):
        Path(render_dir).mkdir(parents=True, exist_ok=True)
        return (
            QualityReport(
                deck_id=current_deck.id,
                passed=False,
                delivery_status="failed",
                issues=issues,
                output_sha256=file_sha256(output),
                render_set_id="round-0",
            ),
            (),
        )

    monkeypatch.setattr(
        service,
        "compiler_for",
        lambda _engine: service.compiler,
    )
    monkeypatch.setattr(service, "verify_artifact", verify)

    _result, report = service.compose(record.job_id, deck)

    checkpoint = service.store.checkpoint_path(
        record.job_id,
        "deck-ir-redesign-1-repair-a.json",
    )
    saved_deck = DeckIR.model_validate_json(
        service.store.checkpoint_path(
            record.job_id,
            "deck-ir.json",
        ).read_text(encoding="utf-8"),
    )
    saved_text = next(
        obj.text
        for obj in saved_deck.slides[0].objects
        if isinstance(obj, TextObject) and obj.role == "body"
    )
    assert report.passed is False
    assert checkpoint.is_file()
    assert saved_text.endswith("（已修复）")
    assert designer.calls == ["repair-a", "repair-b"]
    assert any(
        event.slide_id == "repair-a"
        and "已提交页级修复检查点" in event.message
        and event.recoverable is True
        for event in service.store.get(record.job_id).events
    )


def test_service_promotes_split_deck_and_plan_as_one_verified_repair(
    tmp_path,
    monkeypatch,
) -> None:
    designer = _ServiceSplitDesigner()
    service = PresentationService(tmp_path, repair_designer=designer)
    request = _request().model_copy(
        update={
            "composition_engine": CompositionEngine.NATIVE_PPTX,
            "slide_count": 3,
        },
    )
    record = service.create_job(request)
    intent_a = SlideIntent(
        id="intent-a",
        purpose="prove",
        key_message="核心判断和证据边界",
    )
    intent_b = SlideIntent(
        id="intent-b",
        purpose="decide",
        key_message="下一步",
    )
    plan = DeckPlan(
        title="结构修复",
        narrative_arc="判断到行动",
        visual_direction=VisualDirection(
            visual_concept="editorial evidence",
            palette=["#F7F4EE", "#171717"],
            typography="clear editorial hierarchy",
            image_language="documentary",
            chart_language="direct labels",
            shape_language="quiet rules",
            grid="twelve columns",
            recurring_motif="folio",
        ),
        slides=[
            DeckPlanSlide(
                id="plan-a",
                order=1,
                title="核心判断和证据边界",
                intent=intent_a,
            ),
            DeckPlanSlide(
                id="plan-b",
                order=2,
                title="下一步",
                intent=intent_b,
            ),
        ],
    )
    service.store.write_checkpoint(
        record.job_id,
        "deck-plan.json",
        plan.model_dump_json(indent=2),
    )
    deck = DeckIR(
        id="service-split-deck",
        title="结构修复",
        visual_direction=plan.visual_direction,
        slides=[
            DeckSlide(
                id="slide-a",
                intent_id="intent-a",
                layout_id="freeform",
                speaker_notes="原页面",
                objects=[
                    TextObject(
                        id="old-a",
                        role="body",
                        text="核心判断和证据边界",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-b",
                intent_id="intent-b",
                layout_id="freeform",
                speaker_notes="保持不变",
                objects=[TextObject(id="old-b", role="body", text="下一步")],
            ),
        ],
    )
    issue = QualityIssue(
        code="content.density",
        severity="error",
        slide_id="slide-a",
        message="判断和证据需要拆页",
        repairable=True,
    )
    verification_calls = []

    def verify(current_deck, output, render_dir, *_args, plan=None, **_kwargs):
        verification_calls.append((current_deck, plan))
        Path(render_dir).mkdir(parents=True, exist_ok=True)
        passed = len(verification_calls) == 2
        return (
            QualityReport(
                deck_id=current_deck.id,
                passed=passed,
                delivery_status="passed" if passed else "failed",
                issues=[] if passed else [issue],
                output_sha256=file_sha256(output),
                render_set_id=f"round-{len(verification_calls) - 1}",
            ),
            (),
        )

    monkeypatch.setattr(
        service,
        "compiler_for",
        lambda _engine: service.compiler,
    )
    monkeypatch.setattr(service, "verify_artifact", verify)
    monkeypatch.setattr(
        "lantai.presentation.service.is_verified_delivery",
        lambda *_args, **_kwargs: True,
    )

    _result, report = service.compose(record.job_id, deck)

    saved_plan = DeckPlan.model_validate_json(
        service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).read_text(encoding="utf-8"),
    )
    saved_deck = DeckIR.model_validate_json(
        service.store.checkpoint_path(
            record.job_id,
            "deck-ir.json",
        ).read_text(encoding="utf-8"),
    )
    assert report.passed is True
    assert service.store.get(record.job_id).status == JobStatus.COMPLETED
    assert [slide.id for slide in saved_deck.slides] == [
        "slide-a",
        "slide-a-evidence",
        "slide-b",
    ]
    assert [slide.id for slide in saved_plan.slides] == [
        "plan-a",
        "plan-a-evidence",
        "plan-b",
    ]
    assert len(verification_calls) == 2
    assert len(verification_calls[1][0].slides) == 3
    assert len(verification_calls[1][1].slides) == 3
    assert designer.requests[0].structural_repair_allowed is True
    assert designer.requests[0].slide_count_ceiling == 3
    assert service.store.checkpoint_path(
        record.job_id,
        "deck-plan-redesign-1.json",
    ).is_file()


def test_service_preserves_quality_report_when_repair_proposal_is_invalid(
    tmp_path,
    monkeypatch,
) -> None:
    service = PresentationService(
        tmp_path,
        repair_designer=_RejectedRepairDesigner(),
    )
    request = _request().model_copy(
        update={"composition_engine": CompositionEngine.NATIVE_PPTX},
    )
    record = service.create_job(request)
    deck = DeckIR(
        id="rejected-repair-deck",
        title="拒绝坏修复稿",
        slides=[
            DeckSlide(
                id="rejected-repair-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="rejected-repair-title",
                        role="title",
                        text="保留上一版诊断",
                        slot="title",
                    ),
                ],
            )
        ],
    )
    original_issue = QualityIssue(
        code="layout.capacity",
        severity="error",
        slide_id="rejected-repair-slide",
        message="slot 'point_1' is not defined by key_judgment",
        repairable=True,
    )

    def failing_report(*_args, **_kwargs):
        return (
            QualityReport(
                deck_id=deck.id,
                passed=False,
                issues=[original_issue],
                render_set_id="round-0",
            ),
            (),
        )

    monkeypatch.setattr(
        service,
        "compiler_for",
        lambda _engine: service.compiler,
    )
    monkeypatch.setattr(service, "verify_artifact", failing_report)

    result, report = service.compose(record.job_id, deck)

    failed = service.store.get(record.job_id)
    assert result.output_path.is_file()
    assert report.issues == [original_issue]
    assert failed.status == JobStatus.FAILED
    assert any(
        event.stage == JobStage.REPAIR
        and "自动修复提案被拒绝" in event.message
        and "object ids must be unique" in event.message
        for event in failed.events
    )
    assert (
        service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    ).is_file()
    assert not service.store.checkpoint_path(
        record.job_id,
        "deck-ir-redesign-1.json",
    ).exists()


def test_runtime_recovers_verified_output_without_recomposing(
    tmp_path,
    monkeypatch,
) -> None:
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        id="recoverable-deck",
        title="可恢复交付",
        slides=[
            DeckSlide(
                id="recoverable-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="recoverable-title",
                        role="title",
                        text="可恢复交付",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    deck_path = service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    output = (
        service.store.job_dir(record.job_id) / "output" / "presentation.pptx"
    )
    service.compiler.compile(deck, output)
    from lantai.presentation.quality import inspect_pptx

    render_dir = service.store.job_dir(record.job_id) / "renders" / "round-0"
    render_dir.mkdir(parents=True, exist_ok=True)
    render_path = render_dir / "slide-1.png"
    image = Image.new("RGB", (640, 360), "white")
    image.paste("black", (0, 0, 320, 360))
    image.save(render_path)
    report = inspect_pptx(deck, output, rendered_slide_count=1).model_copy(
        update={
            "renderer_id": "test-renderer",
            "renderer_version": "1.0",
            "font_manifest": ["Arial"],
            "rendered_slides": [
                RenderedSlideEvidence(
                    slide_id=deck.slides[0].id,
                    slide_index=1,
                    sha256=file_sha256(render_path),
                    width=640,
                    height=360,
                ),
            ],
        },
    )
    quality_path = (
        service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    service.store.append_event(
        record.job_id,
        stage=JobStage.QUALITY,
        status=JobStatus.RUNNING,
        message="quality checkpoint persisted before process exit",
        checkpoint=str(quality_path),
        recoverable=True,
    )
    service.store.append_event(
        record.job_id,
        stage=JobStage.DECK_IR,
        status=JobStatus.RUNNING,
        message="deck checkpoint",
        checkpoint=str(deck_path),
        recoverable=True,
    )

    compose_calls: list[str] = []

    def should_not_recompose(_self, job_id, _deck):
        compose_calls.append(job_id)
        raise AssertionError("verified output must not be recomposed")

    monkeypatch.setattr(PresentationService, "compose", should_not_recompose)
    runtime = PresentationJobRuntime(tmp_path)
    try:
        runtime._futures[record.job_id].result(timeout=10)
        recovered = runtime.get(record.job_id, "conversation-a")
    finally:
        runtime.shutdown()

    assert compose_calls == []
    assert recovered.status == JobStatus.COMPLETED
    assert recovered.events[-1].message == "已从经验证的交付检查点恢复"
    recovered_version = PresentationVersionStore(
        service.store.job_dir(record.job_id),
    ).current()
    assert recovered_version is not None
    assert recovered_version.pptx_sha256 == file_sha256(output)
    assert recovered_version.generation == recovered.generation
    assert recovered.runtime_run_id == runtime.runtime_run_id
    assert recovered.runtime_run_id != record.runtime_run_id


def test_runtime_rejects_checkpoint_without_bound_render_evidence(
    tmp_path,
) -> None:
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        id="legacy-unbound-deck",
        title="缺少证据绑定",
        slides=[
            DeckSlide(
                id="legacy-unbound-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="legacy-unbound-title",
                        role="title",
                        text="缺少证据绑定",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    output = (
        service.store.job_dir(record.job_id) / "output" / "presentation.pptx"
    )
    service.compiler.compile(deck, output)
    from lantai.presentation.quality import inspect_pptx

    report = inspect_pptx(deck, output, rendered_slide_count=1)
    quality_path = (
        service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")

    runtime = PresentationJobRuntime.__new__(PresentationJobRuntime)
    runtime.service = service

    assert runtime._complete_from_verified_checkpoint(record.job_id) is False


def test_runtime_rejects_a_verified_checkpoint_when_output_changed(
    tmp_path,
    monkeypatch,
) -> None:
    service = PresentationService(tmp_path)
    record = service.create_job(_request())
    deck = DeckIR(
        id="tampered-deck",
        title="被篡改的恢复点",
        slides=[
            DeckSlide(
                id="tampered-slide",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="tampered-title",
                        role="title",
                        text="原始内容",
                        slot="title",
                    ),
                ],
            ),
        ],
    )
    deck_path = service.store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    output = (
        service.store.job_dir(record.job_id) / "output" / "presentation.pptx"
    )
    service.compiler.compile(deck, output)
    from lantai.presentation.quality import inspect_pptx

    report = inspect_pptx(deck, output, rendered_slide_count=1)
    quality_path = (
        service.store.job_dir(record.job_id)
        / "quality"
        / "quality-report.json"
    )
    quality_path.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    output.write_bytes(output.read_bytes() + b"tampered")
    service.store.append_event(
        record.job_id,
        stage=JobStage.DECK_IR,
        status=JobStatus.RUNNING,
        message="deck checkpoint",
        checkpoint=str(deck_path),
        recoverable=True,
    )

    compose_calls: list[str] = []

    def record_recompose(_self, job_id, _deck, *, writer_lease=None):
        assert writer_lease is not None
        compose_calls.append(job_id)
        raise RuntimeError(
            "stop after proving recovery did not trust the artifact",
        )

    monkeypatch.setattr(PresentationService, "compose", record_recompose)
    runtime = PresentationJobRuntime(tmp_path)
    try:
        runtime._futures[record.job_id].result(timeout=10)
        current = runtime.get(record.job_id, "conversation-a")
    finally:
        runtime.shutdown()

    assert compose_calls == [record.job_id]
    assert current.status == JobStatus.FAILED


def test_runtime_completes_with_scoped_artifacts_and_monotonic_events(
    tmp_path,
    monkeypatch,
) -> None:
    def render_without_system_dependencies(
        _source: str | Path,
        output_dir: str | Path,
        *,
        timeout: int = 120,
    ) -> RenderResult:
        _ = timeout
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        pdf_path = destination / "presentation.pdf"
        pdf_path.write_bytes(b"test pdf")
        slide_images = []
        for index in range(1, 7):
            image_path = destination / f"slide-{index}.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste("black", (0, 0, 320, 360))
            image.save(image_path)
            slide_images.append(image_path)
        return RenderResult(
            pdf_path=pdf_path,
            slide_images=tuple(slide_images),
            renderer_id="test-renderer",
            renderer_version="1.0",
        )

    monkeypatch.setattr(
        "lantai.presentation.service.render_pptx",
        render_without_system_dependencies,
    )
    runtime = PresentationJobRuntime(tmp_path)
    try:
        created = runtime.create_and_start(_request())
        future = runtime._futures[created.job_id]
        future.result(timeout=15)

        completed = runtime.get(created.job_id, "conversation-a")
        artifacts = runtime.artifacts(created.job_id, "conversation-a")
        events = runtime.service.store.events_after(
            created.job_id,
            0,
            "conversation-a",
        )

        assert completed.status == JobStatus.COMPLETED
        assert runtime.output_path(created.job_id, "conversation-a").is_file()
        assert len(artifacts["deck_ir"]["slides"]) == 6
        assert artifacts["quality"]["passed"] is True
        assert len(artifacts["thumbnails"]) == 6
        assert (
            artifacts["editor_handoff"]["deck_id"]
            == artifacts["deck_ir"]["id"]
        )
        assert artifacts["editor_handoff"]["separate_edit_session"] is True
        assert [event.sequence for event in events] == list(
            range(1, len(events) + 1),
        )
        with pytest.raises(PermissionError):
            runtime.artifacts(created.job_id, "conversation-b")
    finally:
        runtime.shutdown()
