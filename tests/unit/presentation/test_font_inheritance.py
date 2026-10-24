"""Public import and QA locks for OOXML font inheritance, not guessed sizes."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.models import DeckIR, DeckSlide, ObjectFrame, ShapeObject, TextObject
from lantai.presentation.quality import inspect_pptx
from lantai.presentation.reconcile import reconcile_pptx


def _fixture(tmp_path: Path):
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shape.name = "body-1"
    shape.text = "合成字号继承测试"
    deck = DeckIR(id="font-deck", title="合成测试", slides=[DeckSlide(
        id="font-slide", layout_id="freeform", objects=[TextObject(
            id="body-1", role="body", text=shape.text,
            frame=ObjectFrame(x=1, y=1, w=6, h=1),
        )],
    )])
    return presentation, shape, deck, tmp_path / "input.pptx"


def test_import_and_quality_read_local_list_style_without_clamping(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    shape._element.find(qn("p:txBody")).find(qn("a:lstStyle")).append(parse_xml(
        f'<a:lvl1pPr {nsdecls("a")}><a:defRPr sz="2050"/></a:lvl1pPr>'
    ))
    presentation.save(source)
    before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    text = imported.slides[0].objects[0]
    assert text.font_size == 20.5
    assert text.runs[0].font_size == 20.5
    assert not [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    output = tmp_path / "exported.pptx"
    NativePptxCompiler().compile(imported, output)
    exported = reconcile_pptx(imported, output, assets_dir=tmp_path / "assets")
    assert exported.slides[0].objects[0].runs[0].font_size == 20.5
    assert source.read_bytes() == before


def test_ordinary_textbox_uses_presentation_default_not_master_other_style(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    defaults = presentation._element.find(qn("p:defaultTextStyle"))
    defaults.find(qn("a:lvl1pPr")).find(qn("a:defRPr")).set("sz", "1800")
    other = presentation.slide_master._element.find(qn("p:txStyles")).find(qn("p:otherStyle"))
    other.find(qn("a:lvl1pPr")).find(qn("a:defRPr")).set("sz", "1200")
    presentation.save(source)
    before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].font_size == 18
    assert imported.slides[0].objects[0].runs[0].font_size == 18
    assert not [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    output = tmp_path / "default-exported.pptx"
    NativePptxCompiler().compile(imported, output)
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "assets")
    assert second.slides[0].objects[0].runs[0].font_size == 18
    assert source.read_bytes() == before


@pytest.mark.parametrize("source_kind, expected", [
    ("layout_list", 23), ("layout_paragraph", 24),
    ("master_placeholder", 22), ("master_style", 21),
])
def test_body_placeholder_uses_nearest_matching_style(tmp_path, source_kind, expected):
    presentation, _shape, deck, source = _fixture(tmp_path)
    # Replace the fixture's ordinary slide with a real body placeholder chain.
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    shape = slide.placeholders[1]
    shape.name = "body-1"
    shape.text = "合成占位符继承测试"
    deck = deck.model_copy(update={"slides": [deck.slides[0].model_copy(update={"id": "placeholder-slide"})]})
    first_id = presentation.slides._sldIdLst[0]
    presentation.part.drop_rel(first_id.rId)
    presentation.slides._sldIdLst.remove(first_id)
    title = slide.shapes.title
    title._element.getparent().remove(title._element)
    layout = slide.slide_layout.placeholders[1]
    if source_kind == "layout_list":
        layout._element.find(qn("p:txBody")).find(qn("a:lstStyle")).append(parse_xml(
            f'<a:lvl1pPr {nsdecls("a")}><a:defRPr sz="2300"/></a:lvl1pPr>'
        ))
    elif source_kind == "layout_paragraph":
        layout.text_frame.paragraphs[0].font.size = Pt(24)
    elif source_kind == "master_placeholder":
        master = slide.slide_layout.slide_master.placeholders[1]
        master.text_frame.paragraphs[0].font.size = Pt(22)
    master_style = slide.slide_layout.slide_master._element.find(qn("p:txStyles")).find(qn("p:bodyStyle"))
    master_style.find(qn("a:lvl1pPr")).find(qn("a:defRPr")).set("sz", "2100")
    presentation._element.find(qn("p:defaultTextStyle")).find(qn("a:lvl1pPr")).find(qn("a:defRPr")).set("sz", "1800")
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].font_size == expected
    assert imported.slides[0].objects[0].runs[0].font_size == expected
    assert not [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]


def test_known_shape_uses_authoritative_inherited_size_on_export(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    shape._element.getparent().remove(shape._element)
    shape = presentation.slides[0].shapes.add_shape(1, Inches(1), Inches(1), Inches(6), Inches(1))
    shape.name = "body-1"
    shape.text = "合成形状内文本"
    shape._element.find(qn("p:txBody")).find(qn("a:lstStyle")).append(parse_xml(
        f'<a:lvl1pPr {nsdecls("a")}><a:defRPr sz="2350"/></a:lvl1pPr>'
    ))
    deck.slides[0].objects = [ShapeObject(id="body-1", shape_type="rect", text=shape.text,
                                          frame=ObjectFrame(x=1, y=1, w=6, h=1))]
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].font_size == 23.5
    output = tmp_path / "shape-exported.pptx"
    NativePptxCompiler().compile(imported, output)
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "assets")
    assert second.slides[0].objects[0].font_size == 23.5


def test_missing_font_authority_fails_import_instead_of_inventing_size(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    defaults = presentation._element.find(qn("p:defaultTextStyle"))
    presentation._element.remove(defaults)
    presentation.save(source)
    before = source.read_bytes()
    assert [i for i in inspect_pptx(deck, source).issues if i.code == "design.font_too_small"]
    with pytest.raises(ValueError, match="unresolved native font size"):
        reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert source.read_bytes() == before


@pytest.mark.parametrize("size, fails", [(17.8, False), (14.8, True)])
def test_explicit_run_overrides_paragraph_and_ending_cannot_hide_small_text(tmp_path, size, fails):
    presentation, shape, deck, source = _fixture(tmp_path)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.size = Pt(19.5)
    paragraph.runs[0].font.size = Pt(size)
    paragraph._p.append(parse_xml(f'<a:endParaRPr {nsdecls("a")} sz="5100"/>'))
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].runs[0].font_size == size
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert bool(issues) is fails


def test_autofit_scale_is_read_without_clamping_or_weakening_readability_floor(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    shape.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    properties = shape._element.find(qn("p:txBody")).find(qn("a:bodyPr"))
    for child in list(properties):
        if child.tag in {qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")}:
            properties.remove(child)
    properties.append(parse_xml(f'<a:normAutofit {nsdecls("a")} fontScale="80000"/>'))
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].runs[0].font_size == 16
    assert [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]


def test_paragraph_level_and_defppr_are_resolved_from_real_rules(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.level = 2
    style = shape._element.find(qn("p:txBody")).find(qn("a:lstStyle"))
    style.append(parse_xml(f'<a:defPPr {nsdecls("a")}><a:defRPr sz="2150"/></a:defPPr>'))
    style.append(parse_xml(f'<a:lvl1pPr {nsdecls("a")}><a:defRPr sz="3000"/></a:lvl1pPr>'))
    style.append(parse_xml(f'<a:lvl3pPr {nsdecls("a")}><a:defRPr sz="2250"/></a:lvl3pPr>'))
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].runs[0].font_size == 22.5
    style.remove(style.find(qn("a:lvl3pPr")))
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].runs[0].font_size == 21.5
