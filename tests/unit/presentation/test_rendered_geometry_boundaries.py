"""Deterministic boundaries around final-native/PDF glyph geometry evidence."""

from __future__ import annotations

from contextlib import contextmanager
from math import nan

import pytest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from lantai.presentation import rendered_geometry as geometry
from lantai.presentation.jobs import PresentationJobInterruptedError
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.service import PresentationService
from tests.unit.presentation.test_rendered_glyph_dividers import _deck


def _compile(tmp_path):
    deck = _deck()
    output = tmp_path / "presentation.pptx"
    NativePptxCompiler().compile(deck, output)
    return deck, output


def _fake_pdf(monkeypatch, tmp_path, *, glyphs=((235, 116, 242, 135),),
              size=(960, 540), rotation=0, pages=1, unicode=ord("M"), fail_charbox=False):
    """Test the parser boundary; these bytes are never rendered or delivered."""
    path = tmp_path / "parser-fixture.pdf"
    path.write_bytes(b"mock-only PDF parser fixture")

    class TextPage:
        closed = False
        loose_values = []

        def count_chars(self):
            return len(glyphs)

        def get_charbox(self, index, *, loose=False):
            self.loose_values.append(loose)
            if fail_charbox:
                raise ValueError("synthetic malformed glyph box")
            left, top, right, bottom = glyphs[index]
            return left, size[1] - bottom, right, size[1] - top

        def close(self):
            self.closed = True

    class Page:
        closed = False

        def get_size(self):
            return size

        def get_rotation(self):
            return rotation

        def get_textpage(self):
            return text_page

        def close(self):
            self.closed = True

    class Document:
        closed = False

        def __len__(self):
            return pages

        def __getitem__(self, index):
            return page

        def close(self):
            self.closed = True

    text_page, page, document = TextPage(), Page(), Document()
    monkeypatch.setattr(geometry.pdfium, "PdfDocument", lambda _: document)
    monkeypatch.setattr(geometry.pdfium.raw, "FPDFText_GetUnicode", lambda *_: unicode)
    return path, document, page, text_page


@pytest.mark.parametrize("case", [
    "no-fill", "transparent", "alpha-mod", "alpha-off", "rotated-rule",
    "rotated-text", "wide-rule", "short-rule", "background-rule",
    "rule-label", "round-rect", "outside-text",
])
def test_ambiguous_or_exempt_native_objects_do_not_require_glyph_proof(tmp_path, case):
    deck, output = _compile(tmp_path)
    native = Presentation(output)
    body = next(shape for shape in native.slides[0].shapes if shape.name == "body")
    rule = next(shape for shape in native.slides[0].shapes if shape.name == "divider")
    if case == "no-fill":
        rule.fill.background()
    elif case in {"transparent", "alpha-mod", "alpha-off"}:
        color = rule._element.xpath("./p:spPr/a:solidFill/*")[0]
        element = OxmlElement({"transparent": "a:alpha", "alpha-mod": "a:alphaMod",
                              "alpha-off": "a:alphaOff"}[case])
        element.set("val", "50000" if case == "transparent" else "100000")
        color.append(element)
    elif case == "rotated-rule":
        rule.rotation = 15
    elif case == "rotated-text":
        body.rotation = 15
    elif case == "wide-rule":
        rule.width = Inches(0.12)
    elif case == "short-rule":
        rule.height = Inches(0.2)
    elif case == "background-rule":
        tree = rule._element.getparent()
        tree.remove(rule._element)
        tree.insert(tree.index(body._element), rule._element)
    elif case == "rule-label":
        rule.text = "native lane label"
    elif case == "round-rect":
        rule._element.xpath("./p:spPr/a:prstGeom")[0].set("prst", "roundRect")
    elif case == "outside-text":
        rule.left = Inches(10)
    native.save(output)

    assert not geometry.has_glyph_divider_candidates(output)
    assert geometry.inspect_rendered_glyph_dividers(deck, output, None) == []


def test_tight_glyph_boxes_are_scaled_from_rendered_pdf_coordinates(tmp_path, monkeypatch):
    deck, output = _compile(tmp_path)
    path, document, page, text = _fake_pdf(
        monkeypatch, tmp_path, size=(480, 270), glyphs=((117.5, 58, 121, 67.5),),
    )
    issues = geometry.inspect_rendered_glyph_dividers(deck, output, path)

    assert len(issues) == 1 and issues[0].code == "layout.glyph_divider_collision"
    assert issues[0].object_id == "body"
    assert text.loose_values == [False]
    assert document.closed and page.closed and text.closed


def test_unknown_unicode_does_not_erase_visible_glyph_geometry(tmp_path, monkeypatch):
    deck, output = _compile(tmp_path)
    path, *_ = _fake_pdf(monkeypatch, tmp_path, unicode=0)

    assert geometry.inspect_rendered_glyph_dividers(deck, output, path)[0].code == "layout.glyph_divider_collision"


def test_final_native_rule_position_overrides_the_requested_ir(tmp_path, monkeypatch):
    deck, output = _compile(tmp_path)
    native = Presentation(output)
    rule = next(shape for shape in native.slides[0].shapes if shape.name == "divider")
    rule.left = Inches(4.2)
    native.save(output)
    path, *_ = _fake_pdf(monkeypatch, tmp_path, glyphs=((300, 116, 309, 135),))

    issue = geometry.inspect_rendered_glyph_dividers(deck, output, path)[0]
    assert issue.code == "layout.glyph_divider_collision" and issue.object_id == "body"
    assert deck.slides[0].objects[-1].frame.x == 3.3


@pytest.mark.parametrize("case", ["missing", "page-count", "aspect", "rotation", "non-finite",
                                  "no-glyphs", "whitespace", "malformed-glyph"])
def test_unavailable_or_ambiguous_pdf_geometry_is_blocking_not_layout_repair(tmp_path, monkeypatch, case):
    deck, output = _compile(tmp_path)
    options = {
        "page-count": {"pages": 2}, "aspect": {"size": (400, 400)},
        "rotation": {"rotation": 90}, "non-finite": {"size": (nan, 540)},
        "no-glyphs": {"glyphs": ()}, "whitespace": {"unicode": ord(" ")},
        "malformed-glyph": {"fail_charbox": True},
    }.get(case, {})
    path, document, page, text = _fake_pdf(monkeypatch, tmp_path, **options)
    issues = geometry.inspect_rendered_glyph_dividers(deck, output, None if case == "missing" else path)

    assert len(issues) == 1
    assert issues[0].code == "render.glyph_geometry_unavailable"
    assert issues[0].severity == "error" and not issues[0].repairable
    if case != "missing":
        assert document.closed
    if case in {"no-glyphs", "whitespace", "malformed-glyph"}:
        assert page.closed and text.closed


@pytest.mark.parametrize("error", [TimeoutError, PresentationJobInterruptedError])
def test_owned_deadline_or_cancel_is_not_swallowed_as_geometry_failure(tmp_path, monkeypatch, error):
    deck, output = _compile(tmp_path)
    path, document, page, text = _fake_pdf(monkeypatch, tmp_path)
    checks = 0

    def check():
        nonlocal checks
        checks += 1
        if checks >= 3:
            raise error("synthetic owned operation interruption")

    monkeypatch.setattr(geometry, "check_operation", check)
    with pytest.raises(error):
        geometry.inspect_rendered_glyph_dividers(deck, output, path)
    assert document.closed and page.closed and text.closed


def test_later_opaque_native_mask_prevents_claiming_the_rule_is_visible(tmp_path, monkeypatch):
    deck, output = _compile(tmp_path)
    native = Presentation(output)
    native.slides[0].shapes.add_shape(1, Inches(3.1), Inches(1.5), Inches(0.4), Inches(0.5)).fill.solid()
    native.save(output)
    path, *_ = _fake_pdf(monkeypatch, tmp_path)

    # This narrow check makes no overall quality claim for hidden text.
    assert geometry.inspect_rendered_glyph_dividers(deck, output, path) == []


def test_service_glyph_inspection_uses_existing_bounded_render_capacity(tmp_path, monkeypatch):
    deck, output = _compile(tmp_path)
    active = False

    class Limits:
        @contextmanager
        def rendering(self):
            nonlocal active
            active = True
            try:
                yield
            finally:
                active = False

    def inspect(*_):
        assert active, "glyph inspection must not bypass the render-capacity boundary"
        return []

    monkeypatch.setattr("lantai.presentation.service.inspect_rendered_glyph_dividers", inspect)
    service = PresentationService(tmp_path / "jobs", execution_limits=Limits())
    service._inspect_artifact(
        deck, output, (), None, None, None, None,
        tmp_path / "rendered", None, False, None, 0,
    )
    assert not active
