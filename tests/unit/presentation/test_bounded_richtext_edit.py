"""An unambiguous inline edit keeps every existing rich-text carrier."""

from copy import deepcopy

from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
import pytest

from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    CompositionAuthority, DeckIR, DeckSlide, ObjectFrame,
    PresentationEditOperation, PresentationObjectEditPlan, ReplaceTextAction,
    RichTextRun, TextObject, TextReplacement,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx
from lantai.presentation.text_editing import replace_native_text_object


def _text(separator=""):
    runs = [RichTextRun(text="保留的原话。", color="17242E", font_size=18, bold=True)]
    if separator:
        runs.append(RichTextRun(text=separator, paragraph_break=True))
    runs.append(RichTextRun(text="设计指出A减少陌生感。", color="5B6A74", font_size=18, italic=True))
    return TextObject(
        id="body", role="body", text="".join(run.text for run in runs), runs=runs,
        font_size=18, frame=ObjectFrame(x=1, y=1, w=8, h=2),
    )


@pytest.mark.parametrize("separator", ["", "\n", "\v"])
def test_unique_single_run_edit_keeps_styles_and_exact_neighbor_runs(separator):
    original = _text(separator)
    before = original.model_dump_json()
    new_text = "保留的原话。" + separator + "设计指出A可以减少陌生感。"
    changed = replace_native_text_object(original, new_text)
    assert changed.text == new_text
    assert changed.runs == [
        *original.runs[:-1],
        original.runs[-1].model_copy(update={"text": "设计指出A可以减少陌生感。"}),
    ]
    assert changed.native_metric_text is None
    assert original.model_dump_json() == before


@pytest.mark.parametrize("case", [
    "stale_runs", "whole_body", "two_runs", "new_break", "deleted_run",
    "invalid_break_flag", "embedded_break", "one_run", "boundary_insertion",
])
def test_unsupported_boundaries_keep_the_existing_unformatted_replacement(case):
    original = _text()
    new_text = "保留的原话。设计指出A可以减少陌生感。"
    if case == "stale_runs":
        original = original.model_copy(update={"text": original.text + "旧字"})
    elif case == "whole_body":
        new_text = "完全不同的合成正文。"
    elif case == "two_runs":
        new_text = "修改原话。设计指出A可以减少陌生感。"
    elif case == "new_break":
        new_text = "保留的原话。设计指出A可以\n减少陌生感。"
    elif case == "deleted_run":
        new_text = original.runs[0].text
    elif case == "invalid_break_flag":
        original = original.model_copy(update={"runs": [
            original.runs[0], original.runs[1].model_copy(update={"paragraph_break": True}),
        ]})
    elif case == "embedded_break":
        original = original.model_copy(update={"text": "保留的原话。\n设计指出A减少陌生感。", "runs": [
            original.runs[0].model_copy(update={"text": "保留的原话。\n"}), original.runs[1],
        ]})
        new_text = "保留的原话。\n设计指出A可以减少陌生感。"
    elif case == "one_run":
        original = original.model_copy(update={"runs": [RichTextRun(text=original.text)]})
    elif case == "boundary_insertion":
        new_text = original.runs[0].text + "新增" + original.runs[1].text
    changed = replace_native_text_object(original, new_text)
    assert changed.text == new_text
    assert changed.runs == [] and changed.native_metric_text is None


def _shape(path):
    return next(s for s in Presentation(path).slides[0].shapes if s.name == "body")


def _native_text_properties(shape):
    body = deepcopy(shape._element.find(qn("p:txBody")))
    for text in body.iter(qn("a:t")):
        text.text = ""
    return etree.tostring(body, method="c14n")


@pytest.mark.parametrize("mode", ["compiled", "transaction"])
@pytest.mark.parametrize("separator", ["", "\n"])
def test_public_edit_export_retains_all_native_text_properties_and_other_objects(tmp_path, mode, separator):
    deck = DeckIR(id="bounded-richtext", title="合成局部富文本编辑", slides=[
        DeckSlide(id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
                  objects=[_text(separator), TextObject(
                      id="untouched", role="body", text="不可顺手改写的对象", font_size=18,
                      frame=ObjectFrame(x=1, y=4, w=7, h=1),
                  )]),
        DeckSlide(id="untouched-page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
                  objects=[TextObject(id="other-page", role="body", text="不可修改的页面", font_size=18,
                                      frame=ObjectFrame(x=1, y=1, w=7, h=1))]),
    ])
    compiler = NativePptxCompiler()
    source = tmp_path / "source.pptx"
    compiler.compile(deck, source)
    original_bytes = source.read_bytes()
    before = Presentation(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    new_text = "保留的原话。" + separator + "设计指出A可以减少陌生感。"
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(
        slide_id="page", actions=[ReplaceTextAction(object_id="body", text=new_text)],
    ))
    service = PresentationEditTransactionService(tmp_path / "edits")
    operation = PresentationEditOperation(
        deck_id=deck.id, slide_ids=["page"], object_ids=ids,
        allowed_content_types={"text"}, instruction="只更正一个原生富文本分段",
    )
    if mode == "compiled":
        preview = service.prepare_compiled_redesign(source, imported, changed, operation, compiler)
    else:
        preview = service.prepare(source, imported, operation, [TextReplacement(object_id="body", text=new_text)])
    after = _shape(preview.candidate_path)
    assert after.text == new_text
    assert _native_text_properties(after) == _native_text_properties(_shape(source))
    assert after.shape_id == _shape(source).shape_id
    actual = Presentation(preview.candidate_path)
    assert next(s for s in actual.slides[0].shapes if s.name == "untouched")._element.xml == next(
        s for s in before.slides[0].shapes if s.name == "untouched")._element.xml
    assert actual.slides[1]._element.xml == before.slides[1]._element.xml
    assert source.read_bytes() == original_bytes
    reopened = reconcile_pptx(changed, preview.candidate_path, assets_dir=tmp_path / "reopened-assets")
    assert reopened.slides[0].objects[0].runs == changed.slides[0].objects[0].runs


def test_transaction_rejects_a_stale_native_run_carrier_instead_of_flattening(tmp_path):
    deck = DeckIR(id="stale-run", title="合成过期载体", slides=[DeckSlide(
        id="page", layout_id="freeform", composition_authority=CompositionAuthority.AUTHOR,
        objects=[_text()],
    )])
    source = tmp_path / "source.pptx"
    NativePptxCompiler().compile(deck, source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    native = Presentation(source)
    paragraph = _shape(source).text_frame.paragraphs[0]
    assert len(paragraph.runs) == 2
    shape = next(s for s in native.slides[0].shapes if s.name == "body")
    shape.text_frame.paragraphs[0].runs[0].text = shape.text
    shape.text_frame.paragraphs[0].runs[1].text = ""
    native.save(source)
    original_bytes = source.read_bytes()
    with pytest.raises(ValueError, match="native rich-text.*stale"):
        PresentationEditTransactionService(tmp_path / "edits").prepare(
            source, imported, PresentationEditOperation(deck_id=deck.id, slide_ids=["page"], object_ids=["body"],
                allowed_content_types={"text"}, instruction="只更正一个分段"),
            [TextReplacement(object_id="body", text="保留的原话。设计指出A可以减少陌生感。")],
        )
    assert source.read_bytes() == original_bytes
