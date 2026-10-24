from __future__ import annotations

import hashlib

import pytest
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from lantai.presentation.models import DeckIR, DeckSlide, ObjectFrame, TextObject
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx


def _native_fixture(tmp_path, *, text="测" * 46, height=0.68, role="title"):
    deck = DeckIR(
        id="synthetic-title-overflow",
        title="Synthetic native title geometry",
        theme_id="lantai-ink",
        slides=[
            DeckSlide(
                id="cover",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="native-title",
                        role=role,
                        slot="title",
                        text=text,
                        font_size=42,
                        frame=ObjectFrame(x=0.94, y=0.60, w=10.8, h=height),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "synthetic-overflow.pptx"
    NativePptxCompiler().compile(deck, output)
    native = Presentation(output)
    title = next(shape for shape in native.slides[0].shapes if shape.name == "native-title")
    title.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    return deck, native, title, output


@pytest.mark.parametrize("fit_mode", ["fixed", "norm-autofit-at-full-size"])
def test_native_quality_rejects_large_title_that_cannot_fit_its_frame(tmp_path, fit_mode) -> None:
    """A sparse character/area ratio cannot excuse a three-line 42pt title."""
    deck, native, title, output = _native_fixture(tmp_path)
    if fit_mode == "norm-autofit-at-full-size":
        title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        title._element.xpath("./p:txBody/a:bodyPr/a:normAutofit")[0].set("fontScale", "100000")
    native.save(output)
    before = hashlib.sha256(output.read_bytes()).hexdigest()

    report = inspect_pptx(deck, output)

    assert any(
        issue.code == "layout.text_overflow_risk"
        and issue.object_id == "native-title"
        and issue.severity == "error"
        for issue in report.issues
    )
    assert report.passed is False
    assert hashlib.sha256(output.read_bytes()).hexdigest() == before
    assert deck.slides[0].objects[0].font_size == 42
    assert deck.slides[0].objects[0].frame.h == 0.68


@pytest.mark.parametrize(
    "mode,text,height,role",
    [
        ("single-line", "合成标题", 0.68, "title"),
        ("roomy-multiline", "测" * 46, 2.4, "title"),
        ("shape-autofit", "测" * 46, 0.68, "title"),
        ("rotated", "测" * 46, 0.68, "title"),
        ("vertical", "测" * 46, 0.68, "title"),
        ("multicolumn", "测" * 46, 0.68, "title"),
        ("no-wrap", "测" * 46, 0.68, "title"),
        ("effective-font-scale", "测" * 46, 0.80, "body"),
        ("scaled-below-title-floor", "测" * 46, 0.80, "title"),
    ],
)
def test_native_quality_preserves_fitting_and_office_specific_text_frames(
    tmp_path, mode, text, height, role,
) -> None:
    deck, native, title, output = _native_fixture(tmp_path, text=text, height=height, role=role)
    body = title._element.xpath("./p:txBody/a:bodyPr")[0]
    if mode == "shape-autofit":
        title.text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    elif mode == "rotated":
        title.rotation = 90
    elif mode == "vertical":
        body.set("vert", "vert")
    elif mode == "multicolumn":
        body.set("numCol", "2")
    elif mode == "no-wrap":
        title.text_frame.word_wrap = False
    elif mode in {"effective-font-scale", "scaled-below-title-floor"}:
        title.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        body.xpath("./a:normAutofit")[0].set("fontScale", "50000")
    native.save(output)
    before = hashlib.sha256(output.read_bytes()).hexdigest()

    report = inspect_pptx(deck, output)

    assert not [issue for issue in report.issues if issue.code == "layout.text_overflow_risk"]
    assert bool([issue for issue in report.issues if issue.code == "design.font_too_small"]) is (
        mode == "scaled-below-title-floor"
    )
    assert hashlib.sha256(output.read_bytes()).hexdigest() == before
    assert deck.slides[0].objects[0].font_size == 42
    assert deck.slides[0].objects[0].frame.h == height


def test_condensed_native_character_spacing_requires_shaping_not_false_rejection(tmp_path) -> None:
    deck, native, title, output = _native_fixture(tmp_path)
    title.text_frame.paragraphs[0].runs[0].font._rPr.set("spc", "-2700")
    native.save(output)
    before = hashlib.sha256(output.read_bytes()).hexdigest()

    report = inspect_pptx(deck, output)

    assert not [issue for issue in report.issues if issue.code == "layout.text_overflow_risk"]
    assert hashlib.sha256(output.read_bytes()).hexdigest() == before
