"""Synthetic material locks at the actual production model boundary.

The provider resolver is replaced by a request capture; extraction, brief,
planner request construction and numeric validation remain production code.
"""
from __future__ import annotations

import hashlib
import json
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from lantai.app import utility_model
from lantai.app.utility_model import ResolvedUtilityModel
from lantai.config.config import UtilityIntelligenceConfig
from lantai.presentation.intelligence import UtilityPresentationSemanticPlanner
from lantai.presentation.material_extraction import PresentationMaterialError, read_material
from lantai.presentation.models import (
    ChartObject, DeckPlanSlide, InputMaterial, PresentationJobRequest, SlideIntent,
)
from lantai.presentation.page_authoring import (
    validate_planned_chart_sources, validate_planned_text_sources,
)
from lantai.presentation.planner import build_brief, build_deterministic_fixture_plan


class CapturedRequest(Exception):
    pass


@pytest.fixture
def capture(monkeypatch):
    calls = []

    async def model(messages, **kwargs):
        calls.append((messages, kwargs))
        raise CapturedRequest()

    monkeypatch.setattr(utility_model, "load_config", lambda: SimpleNamespace(
        utility_intelligence=UtilityIntelligenceConfig(mode="dedicated"),
    ))
    monkeypatch.setattr(utility_model, "resolve_utility_model", lambda *_: ResolvedUtilityModel(
        model=model, provider_id="isolated-material-capture", model_id="capture",
    ))
    return calls


def request(materials, count=2):
    return PresentationJobRequest(
        conversation_id="independent-synthetic-material-conservation",
        user_request="请根据所有材料做决策汇报，并保留来源与预算限制。",
        slide_count=count, materials=materials,
    )


def payload_for(req, brief, calls):
    with pytest.raises(CapturedRequest):
        UtilityPresentationSemanticPlanner("independent-synthetic-agent").plan(req, brief)
    assert len(calls) == 1
    messages, kwargs = calls[0]
    assert [message["role"] for message in messages] == ["system", "user"]
    assert kwargs["max_tokens"] > 0
    payload = json.loads(messages[1]["content"])
    assert payload["brief"] == brief.model_dump(mode="json")
    return payload


def scoped_sections(payload, source_ref):
    return [section for section in payload["brief"]["source_sections"]
            if source_ref in (section.get("source_refs") or [section["source_ref"]])]


@pytest.mark.parametrize("count", [1, 2, 5])
def test_every_material_h1_reaches_actual_model_in_its_own_source(count, capture):
    req = request([
        InputMaterial(id="alpha", name="same-name.md", content="# 甲组预算上限27万元\n## 实施\n甲组只做本地试点。"),
        InputMaterial(id="beta", name="same-name.md", content="# 乙组预算上限83万元\n## 实施\n乙组禁止上传外部云盘。"),
    ], count=count)
    brief = build_brief(req)
    payload = payload_for(req, brief, capture)
    assert payload["allowed_source_refs"] == ["alpha", "beta"]
    for ref, own, other in [
        ("alpha", "甲组预算上限27万元", "乙组预算上限83万元"),
        ("beta", "乙组预算上限83万元", "甲组预算上限27万元"),
    ]:
        scoped = json.dumps(scoped_sections(payload, ref), ensure_ascii=False)
        assert own in scoped, "document H1 absent from this source's model evidence"
        assert other not in scoped, "different document's H1 assigned to this source"


@pytest.mark.parametrize("preamble", ["", "范围仅限已授权客户。\n"])
def test_consecutive_h1_preserves_both_constraints_not_only_the_last(preamble, capture):
    req = request([InputMaterial(id="scope", name="scope.md", content=
        preamble + "# 预算上限27万元\n# 禁止上传外部云盘\n## 实施\n只做本地试点。")])
    payload = payload_for(req, build_brief(req), capture)
    scoped = json.dumps(scoped_sections(payload, "scope"), ensure_ascii=False)
    assert "预算上限27万元" in scoped
    assert "禁止上传外部云盘" in scoped
    assert "只做本地试点" in scoped


def test_empty_heading_between_heading_only_constraints_keeps_all_facts(capture):
    req = request([InputMaterial(id="scope", name="empty.md", content=
        "# 项目\n## 预算上限27万元\n##    \n## 禁止上传外部云盘\n## 实施\n只做本地试点。")])
    payload = payload_for(req, build_brief(req), capture)
    scoped = json.dumps(scoped_sections(payload, "scope"), ensure_ascii=False)
    assert "预算上限27万元" in scoped
    assert "禁止上传外部云盘" in scoped


@pytest.mark.parametrize("count", [1, 2])
def test_merging_few_pages_does_not_erase_late_short_heading_facts(count):
    materials = [InputMaterial(
        id=f"source-{index}", name=f"scope-{index}.md",
        content=f"## {'范围说明' * 6}{constraint}\n{body}",
    ) for index, constraint, body in [
        (1, "预算上限27万元", "预算执行依据。"),
        (2, "禁止上传外部云盘", "隐私授权依据。"),
        (3, "只做本地试点", "实施范围依据。"),
    ]]
    # Every individual title is under the legacy 48-character limit.
    assert all(len(material.content.splitlines()[0][3:]) < 48 for material in materials)
    req = request(materials, count=count)
    plan = build_deterministic_fixture_plan(req, build_brief(req))
    visible = "\n".join(slide.title + "\n" + "\n".join(slide.content_blocks) for slide in plan.slides)
    assert len(plan.slides) <= count
    for fact in ("预算上限27万元", "禁止上传外部云盘", "只做本地试点"):
        assert fact in visible, "merge discarded a fact present in an original short title"
    assert {ref for slide in plan.slides for ref in slide.intent.source_refs} == {
        "source-1", "source-2", "source-3",
    }


def test_h1_numeric_permissions_follow_exact_source_in_model_request(capture):
    req = request([
        InputMaterial(id="alpha", name="alpha.md", content="# 预算上限27万元\n## 实施\n本地试点。"),
        InputMaterial(id="beta", name="beta.md", content="# 预算上限83万元\n## 实施\n本地试点。"),
    ])
    payload = payload_for(req, build_brief(req), capture)
    allowed = payload["allowed_numeric_tokens_by_source"]
    assert "27" in allowed["alpha"]
    assert "83" not in allowed["alpha"]
    assert "83" in allowed["beta"]
    assert "27" not in allowed["beta"]


@pytest.mark.parametrize("surface", ["text", "chart"])
def test_h1_number_remains_usable_and_wrong_source_is_rejected(surface):
    req = request([
        InputMaterial(id="alpha", name="alpha.md", content="# 预算上限27万元\n## 实施\n本地试点。"),
        InputMaterial(id="beta", name="beta.md", content="# 预算上限83万元\n## 实施\n本地试点。"),
    ])
    brief = build_brief(req)

    def slide(ref):
        return DeckPlanSlide(order=1, title="预算", intent=SlideIntent(
            purpose="说明预算", key_message="预算上限27万元", source_refs=[ref],
        ), content_blocks=["预算上限27万元"], charts=[] if surface == "text" else [
            ChartObject(id="budget", chart_type="bar", categories=["上限"], series={"预算": [27]},
                        source_refs=[ref], unit="万元", derivation="材料给出的预算上限"),
        ])

    validator = validate_planned_text_sources if surface == "text" else validate_planned_chart_sources
    validator([slide("alpha")], brief)
    with pytest.raises(ValueError):
        validator([slide("beta")], brief)


def test_real_docx_h1_is_source_bound_with_other_inline_material(tmp_path, capture):
    path = tmp_path / "scope.docx"
    document = Document()
    document.add_heading("Word材料预算上限27万元", level=1)
    document.add_heading("实施", level=2)
    document.add_paragraph("Word材料禁止上传外部云盘。")
    document.save(path)
    req = request([
        InputMaterial(id="inline", name="scope.md", content="# 文本材料\n## 实施\n文本材料只做本地试点。"),
        InputMaterial(id="word", name=path.name, path=str(path)),
    ])
    payload = payload_for(req, build_brief(req), capture)
    word = json.dumps(scoped_sections(payload, "word"), ensure_ascii=False)
    assert "Word材料预算上限27万元" in word
    assert "Word材料禁止上传外部云盘" in word
    assert "文本材料只做本地试点" not in word
    coverage = {item["material_id"]: item for item in payload["brief"]["source_coverage"]}
    assert coverage["word"]["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert coverage["word"]["method"] == "docx_ooxml"


def write_real_pdf(path, pages):
    pdf = canvas.Canvas(str(path), pagesize=(612, 792), pageCompression=0)
    for lines in pages:
        text = pdf.beginText(36, 756)
        text.setFont("Helvetica", 8)
        text.setLeading(12)
        for line in lines:
            assert len(line) <= 80
            text.textLine(line)
        pdf.drawText(text)
        pdf.showPage()
    pdf.save()
    reader = PdfReader(path)
    assert len(reader.pages) == len(pages)
    assert "BUDGET_CEILING_27_USD" in reader.pages[-1].extract_text()


@pytest.mark.parametrize("page_count", [119, 120, 121])
def test_real_pdf_page_boundary_preserves_tail_or_explicitly_rejects_before_model(tmp_path, capture, page_count):
    path = tmp_path / f"pages-{page_count}.pdf"
    pages = [[f"Synthetic evidence page {index + 1}"] for index in range(page_count)]
    pages[-1] += ["BUDGET_CEILING_27_USD", "NO_EXTERNAL_CLOUD_STORAGE"]
    write_real_pdf(path, pages)
    req = request([InputMaterial(id="pdf", name=path.name, path=str(path))])
    try:
        brief = build_brief(req)
    except PresentationMaterialError:
        if page_count <= 120:
            raise
        assert capture == []
        return
    payload = payload_for(req, brief, capture)
    scoped = json.dumps(scoped_sections(payload, "pdf"), ensure_ascii=False)
    assert "BUDGET_CEILING_27_USD" in scoped, "actual PDF reader silently omitted the final page"
    assert "NO_EXTERNAL_CLOUD_STORAGE" in scoped
    assert payload["brief"]["source_coverage"][0]["source_sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_real_pdf_does_not_silently_clip_at_wiki_80000_character_limit(tmp_path, capture):
    path = tmp_path / "long-real-text.pdf"
    pages = [[f"row{row:02d} " + "a" * 74 for row in range(40)] for _ in range(27)]
    pages[-1] += ["BUDGET_CEILING_27_USD", "NO_EXTERNAL_CLOUD_STORAGE"]
    write_real_pdf(path, pages)
    assert len("\n".join(page.extract_text() for page in PdfReader(path).pages)) > 80_000
    req = request([InputMaterial(id="pdf", name=path.name, path=str(path))])
    try:
        brief = build_brief(req)
    except PresentationMaterialError:
        assert capture == []
        return
    payload = payload_for(req, brief, capture)
    scoped = json.dumps(scoped_sections(payload, "pdf"), ensure_ascii=False)
    assert "BUDGET_CEILING_27_USD" in scoped, "PDF extraction reported covered but discarded its tail"
    assert "NO_EXTERNAL_CLOUD_STORAGE" in scoped


@pytest.mark.parametrize("bad_first", [False, True])
def test_one_bad_declared_source_cannot_be_silently_dropped(tmp_path, capture, bad_first):
    bad = tmp_path / "bad.txt"
    bad.write_bytes(b"\xff\xfe")
    materials = [InputMaterial(id="good", name="good.md", content="## 有效事实\n预算上限27万元。"),
                 InputMaterial(id="bad", name="bad.txt", path=str(bad))]
    if bad_first:
        materials.reverse()
    with pytest.raises(PresentationMaterialError):
        build_brief(request(materials))
    assert capture == []


def test_literal_trailing_hash_is_not_silently_removed_from_heading_fact(capture):
    req = request([InputMaterial(
        id="scope", name="scope.md",
        content="# 项目\n## 技术路线必须采用C#\n仅安装本地SDK。",
    )])
    payload = payload_for(req, build_brief(req), capture)
    own_source = json.dumps(scoped_sections(payload, "scope"), ensure_ascii=False)
    assert "技术路线必须采用C#" in own_source

