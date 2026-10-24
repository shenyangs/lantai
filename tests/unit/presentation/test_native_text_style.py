"""Authoritative Office typography must not be replaced by semantic roles."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches
from pptx.dml.color import RGBColor

from lantai.presentation.models import DeckIR, DeckSlide, ObjectFrame, TextObject
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _fixture(tmp_path: Path, *, role="body", emphasis=False):
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shape.name = "native-text"
    shape.text = "合成原生字体保真测试"
    shape._element.find(qn("p:txBody")).find(qn("a:lstStyle")).append(parse_xml(
        f'<a:lvl1pPr {nsdecls("a")}><a:defRPr sz="2400"/></a:lvl1pPr>'
    ))
    deck = DeckIR(id="style-deck", title="合成测试", slides=[DeckSlide(
        id="style-slide", layout_id="freeform", objects=[TextObject(
            id=shape.name, role=role, emphasis=emphasis, text=shape.text,
            frame=ObjectFrame(x=1, y=1, w=6, h=1),
        )],
    )])
    return presentation, shape, deck, tmp_path / "source.pptx"


@pytest.mark.parametrize("role, emphasis", [
    ("title", False), ("heading", False), ("quote", False), ("body", True),
])
def test_plain_native_text_remains_plain_despite_stale_semantic_role(tmp_path, role, emphasis):
    presentation, _shape, deck, source = _fixture(tmp_path, role=role, emphasis=emphasis)
    presentation.save(source)
    before = source.read_bytes()
    old_deck = deck.model_dump_json()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    text = imported.slides[0].objects[0]
    assert text.runs and all(not run.bold and not run.italic for run in text.runs)
    output = tmp_path / "exported.pptx"
    NativePptxCompiler().compile(imported, output)
    native = next(s for s in Presentation(output).slides[0].shapes if s.name == "native-text")
    # PptxGenJS omits false attributes; absent flags are native plain text.
    assert all(run.font.bold is not True and run.font.italic is not True
               for paragraph in native.text_frame.paragraphs for run in paragraph.runs)
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "second-assets")
    assert all(not run.bold and not run.italic for run in second.slides[0].objects[0].runs)
    assert source.read_bytes() == before
    assert deck.model_dump_json() == old_deck


@pytest.mark.parametrize("authority", ["paragraph", "list_style", "presentation_default"])
def test_inherited_native_weight_slant_and_rgb_survive_export(tmp_path, authority):
    presentation, shape, deck, source = _fixture(tmp_path)
    if authority == "paragraph":
        properties = shape.text_frame.paragraphs[0]._p.get_or_add_pPr().get_or_add_defRPr()
    elif authority == "list_style":
        properties = shape._element.find(qn("p:txBody")).find(qn("a:lstStyle")).find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    else:
        properties = presentation._element.find(qn("p:defaultTextStyle")).find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    properties.set("b", "1")
    properties.set("i", "1")
    for fill in properties.findall(qn("a:solidFill")):
        properties.remove(fill)
    properties.append(parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="315F9A"/></a:solidFill>'))
    # A stale IR color may not replace the native inherited color.
    deck.slides[0].objects[0].color = "FFFFFF"
    presentation.save(source)
    before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    run = imported.slides[0].objects[0].runs[0]
    assert (run.bold, run.italic, run.color) == (True, True, "315F9A")
    output = tmp_path / "inherited-exported.pptx"
    NativePptxCompiler().compile(imported, output)
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "second-assets")
    run = second.slides[0].objects[0].runs[0]
    assert (run.bold, run.italic, run.color) == (True, True, "315F9A")
    assert source.read_bytes() == before


def test_explicit_plain_run_overrides_inherited_emphasis_and_color(tmp_path):
    presentation, shape, deck, source = _fixture(tmp_path, role="title", emphasis=True)
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.font.bold = True
    paragraph.font.italic = True
    paragraph.font.color.rgb = RGBColor.from_string("315F9A")
    run = paragraph.runs[0]
    run.font.bold = False
    run.font.italic = False
    run.font.color.rgb = RGBColor.from_string("17242E")
    presentation.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    result = imported.slides[0].objects[0].runs[0]
    assert (result.bold, result.italic, result.color) == (False, False, "17242E")
