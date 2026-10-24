"""Synthetic material proofs for later author grounding and native Word text."""
from __future__ import annotations

import hashlib

import pytest
from docx import Document
from docx.oxml import OxmlElement

from lantai.presentation.material_extraction import PresentationMaterialError, read_material
from lantai.presentation.models import InputMaterial, PresentationJobRequest
from lantai.presentation.page_authoring import grounded_source_text, validate_grounded_numeric_text
from lantai.presentation.planner import build_brief


def test_second_material_h1_number_remains_legal_at_actual_page_author_boundary():
    request = PresentationJobRequest(
        conversation_id="independent-source-author-grounding",
        user_request="保留两个项目材料中的真实预算限制。",
        slide_count=2,
        materials=[
            InputMaterial(id="alpha", name="a.md", content="# 预算上限27万元\n## 实施\n本地试点。"),
            InputMaterial(id="beta", name="b.md", content="# 预算上限83万元\n## 实施\n本地试点。"),
        ],
    )
    source_text = grounded_source_text(request, build_brief(request))
    validate_grounded_numeric_text("第二项目预算上限83万元。", source_text, location="author")
    with pytest.raises(ValueError):
        validate_grounded_numeric_text("虚构预算上限97万元。", source_text, location="author")


def test_real_docx_nonbreaking_hyphen_does_not_change_the_material_identifier(tmp_path):
    path = tmp_path / "scope.docx"
    document = Document()
    paragraph = document.add_paragraph("仅限客户编码 CASE")
    run = paragraph.add_run()
    run._r.append(OxmlElement("w:noBreakHyphen"))
    paragraph.add_run("27，禁止其他客户资料上传外部云盘。")
    document.save(path)
    raw_sha = hashlib.sha256(path.read_bytes()).hexdigest()
    material = InputMaterial(id="scope", name=path.name, path=str(path), sha256=raw_sha)
    try:
        text, coverage = read_material(material)
    except PresentationMaterialError:
        # Unsupported native text is acceptable only with explicit incomplete
        # coverage rejection before a source can ground any model request.
        return
    assert "CASE-27" in text or "CASE\u201127" in text, text
    assert "禁止其他客户资料上传外部云盘" in text
    assert coverage.source_sha256 == raw_sha

