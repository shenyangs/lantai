from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches

from lantai.presentation.slide_package import graft_compiled_slide, _resolve


def _presentation(path, image, title):
    presentation = Presentation()
    for text in ("Untouched", title):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = text
    slide.shapes.add_picture(str(image), 0, Inches(1), width=Inches(1))
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Revenue", [12, 34])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(1), Inches(4), Inches(3), data)
    slide.notes_slide.notes_text_frame.text = "Target notes"
    presentation.save(path)


def test_graft_preserves_original_parts_and_imports_target_dependencies(tmp_path):
    picture = tmp_path / "image.png"
    Image.new("RGB", (10, 10), "red").save(picture)
    source, candidate = tmp_path / "source.pptx", tmp_path / "candidate.pptx"
    _presentation(source, picture, "Old")
    saved = Presentation(source)
    saved.slides[0].shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, 0, Inches(2), Inches(2), Inches(1)).text = "Human edit"
    saved.save(source)
    original_bytes = source.read_bytes()
    Image.new("RGB", (10, 10), "blue").save(picture)
    _presentation(candidate, picture, "New")

    graft_compiled_slide(source, candidate, 1)

    assert source.read_bytes() == original_bytes
    with ZipFile(source) as before, ZipFile(candidate) as after:
        replaced = {"[Content_Types].xml", "ppt/slides/slide2.xml", "ppt/slides/_rels/slide2.xml.rels"}
        for name in before.namelist():
            if name not in replaced:
                assert before.read(name) == after.read(name), name
        assert len(after.namelist()) == len(set(after.namelist()))
        for name in after.namelist():
            if name.endswith(".rels"):
                owner = name.replace("/_rels/", "/")[:-5] if name != "_rels/.rels" else ""
                for rel in ET.fromstring(after.read(name)):
                    if rel.get("TargetMode") != "External":
                        assert _resolve(owner, rel.get("Target")) in after.namelist()
    reopened = Presentation(candidate)
    assert reopened.slides[0].shapes[-1].text == "Human edit"
    target = reopened.slides[1]
    assert target.shapes[0].text == "New"
    assert target.notes_slide.notes_text_frame.text == "Target notes"
    assert next(s for s in target.shapes if s.has_chart).chart.series[0].values == (12.0, 34.0)
    assert next(s for s in target.shapes if s.shape_type == 13).image.blob == picture.read_bytes()


def test_slide_count_failure_keeps_candidate_intact(tmp_path):
    image = tmp_path / "image.png"
    Image.new("RGB", (10, 10)).save(image)
    source, candidate = tmp_path / "source.pptx", tmp_path / "candidate.pptx"
    _presentation(source, image, "Old")
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(candidate)
    unchanged = candidate.read_bytes()
    with pytest.raises(ValueError, match="slide count"):
        graft_compiled_slide(source, candidate, 0)
    assert candidate.read_bytes() == unchanged
    assert not list(tmp_path.glob("*.tmp"))


def test_graft_retains_shared_compiled_image_and_implicit_notes(tmp_path):
    picture = tmp_path / "image.png"
    source, candidate = tmp_path / "source.pptx", tmp_path / "candidate.pptx"
    Image.new("RGB", (10, 10), "red").save(picture)
    _presentation(source, picture, "Old")
    original = Presentation(source)
    target = original.slides[1]
    original_picture = next(s for s in target.shapes if s.shape_type == 13)
    original_picture.name = "preserved-picture"
    shared = target.shapes.add_picture(str(picture), Inches(5), Inches(1), width=Inches(1))
    shared.name = "changed-picture"
    original.save(source)
    original_bytes = source.read_bytes()
    Image.new("RGB", (10, 10), "blue").save(picture)
    _presentation(candidate, picture, "New")
    compiled = Presentation(candidate)
    target = compiled.slides[1]
    next(s for s in target.shapes if s.shape_type == 13).name = "preserved-picture"
    shared = target.shapes.add_picture(str(picture), Inches(5), Inches(1), width=Inches(1))
    shared.name = "changed-picture"
    target.notes_slide.notes_text_frame.text = "New implicit notes"
    compiled.save(candidate)

    graft_compiled_slide(source, candidate, 1, preserved_object_ids={"preserved-picture"})

    target = Presentation(candidate).slides[1]
    actual = {s.name: s for s in target.shapes}
    preserved = next(s for s in Presentation(source).slides[1].shapes if s.name == "preserved-picture")
    assert actual["preserved-picture"]._element.xml == preserved._element.xml
    assert actual["preserved-picture"].image.blob == preserved.image.blob
    assert actual["changed-picture"].image.blob == picture.read_bytes()
    assert target.notes_slide.notes_text_frame.text == "New implicit notes"
    assert source.read_bytes() == original_bytes
