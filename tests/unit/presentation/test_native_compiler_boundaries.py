from __future__ import annotations

import json
from io import BytesIO
import os
import shutil
import signal
import subprocess
import sys
import time
from pathlib import Path
from xml.etree import ElementTree
import zipfile

import pytest

from lantai.presentation.models import (
    DeckIR, DeckSlide, ObjectFrame, TableObject, TextObject,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler


PROJECT_ROOT = Path(__file__).resolve().parents[3]
PRESENTATION_ROOT = PROJECT_ROOT / "src" / "lantai" / "presentation"


def test_runtime_can_resolve_from_configured_bundle(tmp_path, monkeypatch) -> None:
    monkeypatch.delenv("PRESENTATION_NODE_EXECUTABLE", raising=False)
    monkeypatch.delenv("PRESENTATION_NODE_MODULES", raising=False)
    runtime = tmp_path / "runtime"
    node = runtime / "node/bin/node"
    backend = runtime / "node/node_modules/pptxgenjs"
    node.parent.mkdir(parents=True)
    node.touch()
    backend.mkdir(parents=True)
    monkeypatch.setenv("LANTAI_PRESENTATION_RUNTIME", str(runtime))
    for name in ("LANTAI_NODE_BIN", "LANTAI_PPTXGENJS_MODULE", "LANTAI_NODE_MODULES"):
        monkeypatch.delenv(name, raising=False)
    compiler = NativePptxCompiler()
    assert compiler.node_binary == node
    assert compiler.backend_module_path == backend


def test_missing_runtime_does_not_use_developer_home(monkeypatch) -> None:
    monkeypatch.delenv("PRESENTATION_NODE_EXECUTABLE", raising=False)
    monkeypatch.delenv("PRESENTATION_NODE_MODULES", raising=False)
    from lantai.presentation.pptx_compiler import NativePptxCompilerUnavailable
    for name in ("LANTAI_PRESENTATION_RUNTIME", "LANTAI_NODE_BIN",
                 "LANTAI_PPTXGENJS_MODULE", "LANTAI_NODE_MODULES"):
        monkeypatch.delenv(name, raising=False)
    monkeypatch.setattr(NativePptxCompiler, "_runtime_roots", staticmethod(lambda: []))
    monkeypatch.setattr(shutil, "which", lambda _name: None)
    with pytest.raises(NativePptxCompilerUnavailable):
        NativePptxCompiler._resolve_node(None)
    with pytest.raises(NativePptxCompilerUnavailable):
        NativePptxCompiler._resolve_backend_module(None)


def _run_adapter(deck: dict[str, object]) -> subprocess.CompletedProcess[str]:
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node.js is unavailable")
    module_uri = (PRESENTATION_ROOT / "legacy_deckir_adapter.mjs").as_uri()
    script = f"""
      import {{ normalizeHistoricalDeckIR }} from {json.dumps(module_uri)};
      const deck = JSON.parse(process.argv[1]);
      const result = normalizeHistoricalDeckIR(deck);
      process.stdout.write(JSON.stringify(result));
    """
    return subprocess.run(
        [node, "--input-type=module", "-e", script, json.dumps(deck)],
        check=False,
        capture_output=True,
        text=True,
    )


def test_historical_adapter_never_composes_a_designed_deck() -> None:
    result = _run_adapter(
        {
            "id": "deck-designed",
            "design_context": {"profile": {"identity": "example"}},
            "slides": [
                {
                    "id": "slide-1",
                    "objects": [
                        {"id": "body-1", "kind": "text", "slot": "body"},
                    ],
                },
            ],
        },
    )

    assert result.returncode != 0
    assert "designed deck is missing an explicit valid frame" in result.stderr


def test_historical_adapter_is_explicit_and_reports_degradation() -> None:
    result = _run_adapter(
        {
            "id": "deck-historical",
            "slides": [
                {
                    "id": "slide-1",
                    "objects": [
                        {"id": "body-1", "kind": "text", "slot": "body"},
                    ],
                },
            ],
        },
    )

    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    assert payload["deck"]["slides"][0]["objects"][0]["frame"]
    assert payload["degradations"] == [
        "slide-1:body-1: historical DeckIR frame adapter used",
    ]


def test_historical_adapter_keeps_four_metric_slots_inside_the_slide() -> None:
    result = _run_adapter(
        {
            "id": "deck-four-metrics",
            "slides": [
                {
                    "id": "slide-metrics",
                    "objects": [
                        {"id": "title", "kind": "text", "slot": "title"},
                        *[
                            {
                                "id": f"metric-{index}",
                                "kind": "text",
                                "slot": f"metric_{index}",
                            }
                            for index in range(1, 5)
                        ],
                    ],
                },
            ],
        },
    )

    assert result.returncode == 0, result.stderr
    payload = json.loads(result.stdout)
    frames = [
        item["frame"]
        for item in payload["deck"]["slides"][0]["objects"]
        if item["id"].startswith("metric-")
    ]
    assert len(frames) == 4
    assert all(frame["y"] + frame["h"] <= 7.501 for frame in frames)


def test_backend_does_not_emit_invalid_color_warnings_for_transparent_shapes(
    tmp_path: Path,
) -> None:
    node = shutil.which("node")
    module_path = PROJECT_ROOT / "console" / "node_modules" / "pptxgenjs"
    if node is None or not module_path.exists():
        pytest.skip("PptxGenJS runtime is unavailable")
    backend_uri = (PRESENTATION_ROOT / "pptxgenjs_backend.mjs").as_uri()
    output = tmp_path / "transparent-shape.pptx"
    deck = {
        "id": "transparent-shape",
        "title": "Transparent shape",
        "slides": [
            {
                "id": "slide-1",
                "objects": [
                    {
                        "id": "shape-1",
                        "kind": "shape",
                        "shape_type": "rect",
                        "frame": {"x": 1, "y": 1, "w": 2, "h": 1},
                    },
                ],
            },
        ],
    }
    script = f"""
      import {{ compileWithPptxGenJS }} from {json.dumps(backend_uri)};
      const result = await compileWithPptxGenJS({{
        deck: JSON.parse(process.argv[1]),
        outputPath: process.argv[2],
        modulePath: process.argv[3],
      }});
      process.stdout.write(JSON.stringify(result));
    """

    result = subprocess.run(
        [
            node,
            "--input-type=module",
            "-e",
            script,
            json.dumps(deck),
            str(output),
            str(module_path),
        ],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert "not a valid scheme color" not in result.stderr


def test_backend_is_the_only_module_that_depends_on_pptxgenjs_api() -> None:
    backend = (PRESENTATION_ROOT / "pptxgenjs_backend.mjs").read_text(
        encoding="utf-8",
    )
    compiler = (PRESENTATION_ROOT / "pptx_compiler.mjs").read_text(
        encoding="utf-8",
    )
    adapter = (PRESENTATION_ROOT / "legacy_deckir_adapter.mjs").read_text(
        encoding="utf-8",
    )

    assert "createRequire" in backend
    assert "new PptxGenJS" in backend
    assert "pptx.ShapeType" in backend
    assert "createRequire" not in compiler
    assert "new PptxGenJS" not in compiler
    assert "pptx.ShapeType" not in compiler
    assert "createRequire" not in adapter


def test_backend_uses_one_color_for_a_single_non_pie_series() -> None:
    node = shutil.which("node")
    if node is None:
        pytest.skip("Node.js is unavailable")
    backend_uri = (PRESENTATION_ROOT / "pptxgenjs_backend.mjs").as_uri()
    script = f"""
      import {{ chartColorsFor }} from {json.dumps(backend_uri)};
      const palette = ["315F9A", "C16F3B", "3E8172", "7C8790"];
      process.stdout.write(JSON.stringify({{
        column: chartColorsFor({{ isPie: false, seriesCount: 1, palette }}),
        line: chartColorsFor({{ isPie: false, seriesCount: 2, palette }}),
        pie: chartColorsFor({{ isPie: true, seriesCount: 1, palette }}),
      }}));
    """

    result = subprocess.run(
        [node, "--input-type=module", "-e", script],
        check=False,
        capture_output=True,
        text=True,
    )

    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout) == {
        "column": ["315F9A"],
        "line": ["315F9A", "C16F3B", "3E8172", "7C8790"],
        "pie": ["315F9A", "C16F3B", "3E8172", "7C8790"],
    }


def test_compiler_and_backend_do_not_select_style_or_layout_policy() -> None:
    compiler = (PRESENTATION_ROOT / "pptx_compiler.mjs").read_text(
        encoding="utf-8",
    )
    backend = (PRESENTATION_ROOT / "pptxgenjs_backend.mjs").read_text(
        encoding="utf-8",
    )
    forbidden = (
        "design_context",
        "design_profile",
        "style_profile",
        "layout_id",
        "legacyFrame",
        "editorial-executive",
        "minimal-keynote",
        "product-launch",
        "data-analytical",
    )

    for source in (compiler, backend):
        assert not any(marker in source for marker in forbidden)


def test_native_compiler_repairs_dangling_content_type_overrides(tmp_path) -> None:
    deck = DeckIR(
        id="deck-package-integrity",
        title="Package integrity",
        slides=[
            DeckSlide(
                id=f"slide-{index}",
                layout_id="freeform",
                speaker_notes=f"Evidence note {index}",
                objects=[
                    TextObject(
                        id=f"title-{index}",
                        role="title",
                        text=f"Slide {index}",
                        frame=ObjectFrame(x=0.8, y=0.7, w=6.0, h=0.8),
                    ),
                ],
            )
            for index in range(1, 4)
        ],
    )
    output = tmp_path / "package-integrity.pptx"

    NativePptxCompiler().compile(deck, output)

    namespace = {
        "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
    }
    with zipfile.ZipFile(output) as package:
        members = set(package.namelist())
        content_types = ElementTree.fromstring(
            package.read("[Content_Types].xml"),
        )
        override_targets = {
            str(node.attrib["PartName"]).lstrip("/")
            for node in content_types.findall("ct:Override", namespace)
        }
        notes = package.read("ppt/notesSlides/notesSlide2.xml").decode("utf-8")

    assert override_targets <= members
    assert {
        target for target in override_targets if target.startswith("ppt/slideMasters/")
    } == {"ppt/slideMasters/slideMaster1.xml"}
    assert "Evidence note 2" in notes


def test_native_tables_use_standard_ooxml_vertical_anchor(tmp_path) -> None:
    """Native tables must reopen in strict editors without repairing XML."""

    deck = DeckIR(
        id="deck-native-table-anchor",
        title="Editable table",
        slides=[
            DeckSlide(
                id="slide-table",
                layout_id="freeform",
                objects=[
                    TableObject(
                        id="native-table",
                        headers=["Category", "Value"],
                        rows=[["A", "31"], ["B", "9"]],
                        frame=ObjectFrame(x=0.8, y=1.5, w=7.0, h=3.0),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "native-table-anchor.pptx"
    NativePptxCompiler().compile(deck, output)

    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(output) as package:
        slide = ElementTree.fromstring(package.read("ppt/slides/slide1.xml"))
    cells = slide.findall(".//a:tbl/a:tr/a:tc", namespace)
    anchors = [cell.find("a:tcPr", namespace).get("anchor") for cell in cells]

    assert len(cells) == 6
    assert anchors == ["ctr"] * 6
    assert [cell.find("a:txBody/a:p/a:r/a:t", namespace).text for cell in cells] == [
        "Category", "Value", "A", "31", "B", "9",
    ]


def test_native_compiler_sigkill_never_replaces_existing_candidate_with_partial_bytes(
    tmp_path: Path,
) -> None:
    """REC-01: the backend may die mid-write without corrupting a prior candidate."""

    fake_node = tmp_path / "fake-node"
    fake_node.write_text(
        """#!/usr/bin/env python3
import os
import sys
import time
from pathlib import Path

output = Path(sys.argv[3])
output.write_bytes(b"partial-pptx-package")
Path(os.environ["REC01_MARKER"]).write_text(str(output), encoding="utf-8")
time.sleep(30)
""",
        encoding="utf-8",
    )
    fake_node.chmod(0o755)
    backend = tmp_path / "fake-backend"
    backend.mkdir()
    output = tmp_path / "candidate.pptx"
    original = b"previous-valid-candidate"
    output.write_bytes(original)
    marker = tmp_path / "backend-started"
    worker_source = """
from pathlib import Path
import os
from lantai.presentation.models import DeckIR, DeckSlide
from lantai.presentation.pptx_compiler import NativePptxCompiler

compiler = NativePptxCompiler(
    node_binary=os.environ["REC01_NODE"],
    backend_module_path=os.environ["REC01_BACKEND"],
)
deck = DeckIR(
    id="rec01",
    title="REC-01",
    slides=[DeckSlide(id="slide-1", layout_id="freeform", objects=[])],
)
compiler.compile(deck, Path(os.environ["REC01_OUTPUT"]))
"""
    environment = os.environ.copy()
    environment.update(
        {
            "REC01_NODE": str(fake_node),
            "REC01_BACKEND": str(backend),
            "REC01_OUTPUT": str(output),
            "REC01_MARKER": str(marker),
        },
    )
    worker = subprocess.Popen(
        [sys.executable, "-c", worker_source],
        cwd=PROJECT_ROOT,
        env=environment,
        start_new_session=True,
    )
    deadline = time.monotonic() + 5
    while not marker.is_file() and worker.poll() is None:
        if time.monotonic() >= deadline:
            break
        time.sleep(0.02)
    try:
        assert marker.is_file(), "fake backend did not enter its write phase"
        os.killpg(worker.pid, signal.SIGKILL)
        worker.wait(timeout=5)
    finally:
        if worker.poll() is None:
            os.killpg(worker.pid, signal.SIGKILL)
            worker.wait(timeout=5)

    assert output.read_bytes() == original

    # Recovery reruns only the interrupted compile boundary and atomically
    # replaces the preserved candidate after a complete package is available.
    from pptx import Presentation

    valid_source = tmp_path / "valid-source.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(valid_source)
    fake_node.write_text(
        """#!/usr/bin/env python3
import json
import shutil
import sys
from pathlib import Path

output = Path(sys.argv[3])
shutil.copy2(Path(__VALID_SOURCE__), output)
print(json.dumps({"engine": "native_pptx", "editable_object_count": 0}))
""".replace("__VALID_SOURCE__", repr(str(valid_source))),
        encoding="utf-8",
    )
    fake_node.chmod(0o755)
    compiler = NativePptxCompiler(
        node_binary=fake_node,
        backend_module_path=backend,
    )
    result = compiler.compile(
        DeckIR(
            id="rec01",
            title="REC-01",
            slides=[DeckSlide(id="slide-1", layout_id="freeform", objects=[])],
        ),
        output,
    )

    assert result.output_path == output
    assert output.read_bytes() != original
    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        assert "[Content_Types].xml" in package.namelist()
    assert len(Presentation(str(output)).slides) == 1


def test_native_compiler_sigkill_after_atomic_commit_keeps_valid_candidate(
    tmp_path: Path,
) -> None:
    """REC-01: a kill after commit cannot expose a partial final package."""

    from pptx import Presentation

    valid_source = tmp_path / "valid-source.pptx"
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(valid_source)
    fake_node = tmp_path / "fake-node"
    fake_node.write_text(
        """#!/usr/bin/env python3
import json
import shutil
import sys
from pathlib import Path

output = Path(sys.argv[3])
shutil.copy2(Path(__VALID_SOURCE__), output)
print(json.dumps({"engine": "native_pptx", "editable_object_count": 0}))
""".replace("__VALID_SOURCE__", repr(str(valid_source))),
        encoding="utf-8",
    )
    fake_node.chmod(0o755)
    backend = tmp_path / "fake-backend"
    backend.mkdir()
    output = tmp_path / "candidate.pptx"
    output.write_bytes(b"previous-valid-candidate")
    marker = tmp_path / "compile-committed"
    worker_source = """
from pathlib import Path
import os
import time
from lantai.presentation.models import DeckIR, DeckSlide
from lantai.presentation.pptx_compiler import NativePptxCompiler

compiler = NativePptxCompiler(
    node_binary=os.environ["REC01_NODE"],
    backend_module_path=os.environ["REC01_BACKEND"],
)
deck = DeckIR(
    id="rec01",
    title="REC-01",
    slides=[DeckSlide(id="slide-1", layout_id="freeform", objects=[])],
)
compiler.compile(deck, Path(os.environ["REC01_OUTPUT"]))
Path(os.environ["REC01_MARKER"]).write_text("committed", encoding="utf-8")
time.sleep(30)
"""
    environment = os.environ.copy()
    environment.update(
        {
            "REC01_NODE": str(fake_node),
            "REC01_BACKEND": str(backend),
            "REC01_OUTPUT": str(output),
            "REC01_MARKER": str(marker),
        },
    )
    worker = subprocess.Popen(
        [sys.executable, "-c", worker_source],
        cwd=PROJECT_ROOT,
        env=environment,
        start_new_session=True,
    )
    deadline = time.monotonic() + 5
    while not marker.is_file() and worker.poll() is None:
        if time.monotonic() >= deadline:
            break
        time.sleep(0.02)
    try:
        assert marker.is_file(), "worker did not reach the committed checkpoint"
        os.killpg(worker.pid, signal.SIGKILL)
        worker.wait(timeout=5)
    finally:
        if worker.poll() is None:
            os.killpg(worker.pid, signal.SIGKILL)
            worker.wait(timeout=5)

    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        assert "[Content_Types].xml" in package.namelist()
    assert len(Presentation(str(output)).slides) == 1


def test_native_compiler_invalid_execution_evidence_preserves_existing_candidate(
    tmp_path: Path,
    monkeypatch,
) -> None:
    from pptx import Presentation

    node = tmp_path / "fake-node"
    node.touch()
    backend = tmp_path / "fake-backend"
    backend.mkdir()
    output = tmp_path / "candidate.pptx"
    original = b"previous-valid-candidate"
    output.write_bytes(original)

    def run(command, **_kwargs):
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(command[3])
        return subprocess.CompletedProcess(
            command,
            0,
            '{"engine":"unexpected","editable_object_count":0}',
            "",
        )

    monkeypatch.setattr(subprocess, "run", run)
    compiler = NativePptxCompiler(
        node_binary=node,
        backend_module_path=backend,
    )
    deck = DeckIR(
        id="rec01-invalid-evidence",
        title="REC-01 invalid evidence",
        slides=[DeckSlide(id="slide-1", layout_id="freeform", objects=[])],
    )

    with pytest.raises(RuntimeError, match="identity mismatch"):
        compiler.compile(deck, output)

    assert output.read_bytes() == original


def test_frozen_runtime_uses_electron_resources_without_duplicate_segment(tmp_path, monkeypatch) -> None:
    import sys

    resources = tmp_path / 'Lantai.app/Contents/Resources'
    executable = resources / 'lantai-server/lantai-server'
    monkeypatch.setattr(sys, 'frozen', True, raising=False)
    monkeypatch.setattr(sys, 'executable', str(executable))
    monkeypatch.delenv('LANTAI_PRESENTATION_RUNTIME', raising=False)
    roots = NativePptxCompiler._runtime_roots()
    assert resources / 'presentation-runtime' in roots
    assert resources / 'Resources/presentation-runtime' not in roots


def test_reimport_keeps_native_paragraph_boundaries_and_base_size(tmp_path) -> None:
    from pptx import Presentation
    from lantai.presentation.models import RichTextRun
    from lantai.presentation.reconcile import reconcile_pptx

    deck = DeckIR(id="paragraph-roundtrip", title="Paragraphs", slides=[
        DeckSlide(id="s1", layout_id="freeform", objects=[
            TextObject(id="body", role="body", text="标题\n第一段\n第二段",
                       font_size=31, frame=ObjectFrame(x=1, y=1, w=5, h=4),
                       runs=[RichTextRun(text="标题\n", font_size=32.24),
                             RichTextRun(text="第一段\n第二段", font_size=23)])
        ])
    ])
    compiler = NativePptxCompiler()
    first, second = tmp_path / "first.pptx", tmp_path / "second.pptx"
    compiler.compile(deck, first)
    reconciled = reconcile_pptx(deck, first, assets_dir=tmp_path / "assets")
    text = reconciled.slides[0].objects[0]
    assert text.font_size == 31
    assert any(run.paragraph_break for run in text.runs)
    compiler.compile(reconciled, second)
    def paragraphs(path):
        shape = next(s for s in Presentation(path).slides[0].shapes if s.name == "body")
        return [p.text for p in shape.text_frame.paragraphs]
    assert paragraphs(first) == paragraphs(second) == ["标题\n第一段", "第二段"]


def test_reimport_marks_fully_framed_text_page_as_author_composition(tmp_path) -> None:
    """A native text-only page must not be reflowed merely for lacking a Shape."""
    from lantai.presentation.models import CompositionAuthority
    from lantai.presentation.reconcile import reconcile_pptx

    deck = DeckIR(
        id="native-text-authority",
        title="Native text authority",
        slides=[
            DeckSlide(
                id="s1",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                objects=[
                    TextObject(
                        id="title",
                        role="title",
                        text="作者已确定坐标",
                        frame=ObjectFrame(x=0.75, y=0.45, w=11.83, h=1.2),
                    ),
                    TextObject(
                        id="body",
                        role="body",
                        text="纯文字页也必须保留作者构图权威。",
                        frame=ObjectFrame(x=0.75, y=2.0, w=7.5, h=1.0),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "native-text-authority.pptx"
    NativePptxCompiler().compile(deck, output)

    reconciled = reconcile_pptx(
        deck,
        output,
        assets_dir=tmp_path / "reconciled-assets",
    )

    slide = reconciled.slides[0]
    assert slide.composition_authority == CompositionAuthority.AUTHOR
    assert [obj.id for obj in slide.objects] == ["title", "body"]
    assert all(obj.frame is not None for obj in slide.objects)


def test_inline_regular_style_overrides_emphasized_textbox(tmp_path) -> None:
    from pptx import Presentation
    from lantai.presentation.models import RichTextRun

    deck = DeckIR(id="inline-weight", title="Inline weight", slides=[
        DeckSlide(id="s1", layout_id="freeform", objects=[
            TextObject(id="body", role="quote", emphasis=True, text="强调普通",
                       font_size=23, frame=ObjectFrame(x=1, y=1, w=5, h=2),
                       runs=[RichTextRun(text="强调", bold=True, italic=True),
                             RichTextRun(text="普通", bold=False, italic=False)])
        ])
    ])
    output = tmp_path / "inline-weight.pptx"
    NativePptxCompiler().compile(deck, output)
    shape = next(s for s in Presentation(output).slides[0].shapes if s.name == "body")
    runs = shape.text_frame.paragraphs[0].runs
    assert runs[0].font.bold and runs[0].font.italic
    assert not runs[1].font.bold and not runs[1].font.italic
    assert runs[1].font.size.pt == 23


@pytest.mark.parametrize("chart_type", ["bar", "column", "line"])
def test_native_two_dimensional_charts_reference_only_declared_axes(
    tmp_path, chart_type,
) -> None:
    from pptx import Presentation
    from lantai.presentation.models import ChartObject

    deck = DeckIR(
        id="native-chart-axis-integrity",
        title="Native chart axis integrity",
        slides=[DeckSlide(
            id="s1",
            layout_id="freeform",
            objects=[ChartObject(
                id="chart",
                chart_type=chart_type,
                categories=["Q1", "Q2"],
                series={"收入": [1, 2], "成本": [0.5, 1]},
                unit="万元",
                derivation="direct fixture values",
                source_refs=["fixture://native-chart-axis-integrity"],
                frame=ObjectFrame(x=1, y=1, w=8, h=4),
            )],
        )],
    )
    output = tmp_path / f"native-chart-{chart_type}.pptx"

    NativePptxCompiler().compile(deck, output)

    chart_namespace = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    with zipfile.ZipFile(output) as package:
        chart_xml = ElementTree.fromstring(package.read("ppt/charts/chart1.xml"))
    plot_area = chart_xml.find(f".//{{{chart_namespace}}}plotArea")
    assert plot_area is not None
    declared_axes = {
        axis.find(f"{{{chart_namespace}}}axId").attrib["val"]
        for axis in plot_area
        if axis.tag in {
            f"{{{chart_namespace}}}catAx", f"{{{chart_namespace}}}valAx",
            f"{{{chart_namespace}}}dateAx", f"{{{chart_namespace}}}serAx",
        }
    }
    chart_tag = "lineChart" if chart_type == "line" else "barChart"
    chart_element = plot_area.find(f"{{{chart_namespace}}}{chart_tag}")
    assert chart_element is not None
    referenced_axes = {
        axis.attrib["val"]
        for axis in chart_element.findall(f"{{{chart_namespace}}}axId")
    }
    assert referenced_axes == declared_axes
    assert len(referenced_axes) == 2
    native_chart = next(
        shape.chart for shape in Presentation(output).slides[0].shapes
        if shape.has_chart
    )
    assert [category.label for category in native_chart.plots[0].categories] == ["Q1", "Q2"]
    assert [tuple(series.values) for series in native_chart.series] == [(1.0, 2.0), (0.5, 1.0)]
    assert native_chart.part.chart_workbook.xlsx_part.blob


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie"])
def test_native_chart_workbooks_use_a_valid_table_range(tmp_path, chart_type) -> None:
    from pptx import Presentation
    from lantai.presentation.models import ChartObject

    deck = DeckIR(id="native-chart-workbook", title="Native chart workbook", slides=[
        DeckSlide(id="s1", layout_id="freeform", objects=[ChartObject(
            id="chart", chart_type=chart_type, categories=["Q1", "Q2"],
            series={"收入": [1, 2]}, unit="万元",
            derivation="direct fixture values",
            source_refs=["fixture://native-chart-workbook"],
            frame=ObjectFrame(x=1, y=1, w=8, h=4),
        )]),
    ])
    output = tmp_path / f"native-chart-workbook-{chart_type}.pptx"

    NativePptxCompiler().compile(deck, output)

    chart = next(shape.chart for shape in Presentation(output).slides[0].shapes
                 if shape.has_chart)
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
        table = ElementTree.fromstring(workbook.read("xl/tables/table1.xml"))
        worksheet = ElementTree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
    dimension = worksheet.find(f"{{{spreadsheet_ns}}}dimension")
    assert dimension is not None
    assert table.attrib["ref"] == dimension.attrib["ref"] == "A1:B3"
    cells = {
        cell.attrib["r"]: cell.find(f"{{{spreadsheet_ns}}}v").text
        for cell in worksheet.iter(f"{{{spreadsheet_ns}}}c")
    }
    assert {key: cells[key] for key in ("B2", "B3")} == {
        "B2": "1", "B3": "2",
    }
    assert [tuple(series.values) for series in chart.series] == [(1.0, 2.0)]


@pytest.mark.parametrize("chart_type", ["bar", "column", "line", "pie"])
def test_native_chart_workbooks_preserve_zero_as_a_number(tmp_path, chart_type) -> None:
    from pptx import Presentation
    from lantai.presentation.models import ChartObject

    deck = DeckIR(id="native-chart-zero", title="Native chart zero", slides=[
        DeckSlide(id="s1", layout_id="freeform", objects=[ChartObject(
            id="chart", chart_type=chart_type, categories=["Q1", "Q2", "Q3"],
            series={"收入": [0, 2, 0.5]}, unit="万元",
            derivation="direct fixture values",
            source_refs=["fixture://native-chart-zero"],
            frame=ObjectFrame(x=1, y=1, w=8, h=4),
        )]),
    ])
    output = tmp_path / f"native-chart-zero-{chart_type}.pptx"

    NativePptxCompiler().compile(deck, output)

    chart = next(shape.chart for shape in Presentation(output).slides[0].shapes
                 if shape.has_chart)
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with zipfile.ZipFile(BytesIO(chart.part.chart_workbook.xlsx_part.blob)) as workbook:
        worksheet = ElementTree.fromstring(workbook.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in worksheet.iter(f"{{{spreadsheet_ns}}}c")}
    assert cells["B2"].attrib.get("t", "n") == "n"
    assert cells["B2"].find(f"{{{spreadsheet_ns}}}v").text == "0"
    assert cells["B3"].find(f"{{{spreadsheet_ns}}}v").text == "2"
    assert cells["B4"].find(f"{{{spreadsheet_ns}}}v").text == "0.5"
    assert tuple(chart.series[0].values) == (0.0, 2.0, 0.5)


def test_native_pie_rejects_multiple_series_without_publishing_partial_facts(tmp_path) -> None:
    from lantai.presentation.models import ChartObject

    deck = DeckIR(id="multiple-pie-series", title="Multiple pie series", slides=[
        DeckSlide(id="s1", layout_id="freeform", objects=[ChartObject(
            id="chart", chart_type="pie", categories=["Q1", "Q2"],
            series={"收入": [1, 2], "成本": [0.5, 1]}, unit="万元",
            derivation="direct fixture values",
            source_refs=["fixture://multiple-pie-series"],
            frame=ObjectFrame(x=1, y=1, w=8, h=4),
        )]),
    ])
    output = tmp_path / "previous-candidate.pptx"
    output.write_bytes(b"previous candidate must remain unchanged")

    with pytest.raises(RuntimeError, match="pie chart requires exactly one data series"):
        NativePptxCompiler().compile(deck, output)

    assert output.read_bytes() == b"previous candidate must remain unchanged"
    assert sorted(path.name for path in tmp_path.iterdir()) == [output.name]


def test_native_compiler_honors_explicit_chart_legend_policy(tmp_path) -> None:
    from pptx import Presentation
    from lantai.presentation.models import ChartObject
    from lantai.presentation.reconcile import reconcile_pptx

    deck = DeckIR(
        id="native-chart-legend",
        title="Native chart legend",
        slides=[
            DeckSlide(
                id="s1",
                layout_id="freeform",
                objects=[
                    ChartObject(
                        id="chart",
                        chart_type="line",
                        categories=["Q1", "Q2"],
                        series={"收入": [1, 2], "成本": [0.5, 1]},
                        show_legend=False,
                        unit="万元",
                        derivation="direct fixture values",
                        source_refs=["fixture://native-chart-legend"],
                        frame=ObjectFrame(x=1, y=1, w=8, h=4),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "native-chart-legend.pptx"

    NativePptxCompiler().compile(deck, output)
    chart = next(
        shape.chart
        for shape in Presentation(output).slides[0].shapes
        if shape.has_chart
    )

    assert chart.has_legend is False
    reconciled = reconcile_pptx(
        deck,
        output,
        assets_dir=tmp_path / "reconciled-assets",
    )
    reconciled_chart = next(
        obj
        for obj in reconciled.slides[0].objects
        if isinstance(obj, ChartObject)
    )
    assert reconciled_chart.chart_type == "line"
    assert reconciled_chart.series == {"收入": [1.0, 2.0], "成本": [0.5, 1.0]}
    assert reconciled_chart.show_legend is False


def test_desktop_runtime_contract_and_electron_mode(tmp_path, monkeypatch):
    import os
    from pptx import Presentation
    for key in ("LANTAI_PRESENTATION_RUNTIME", "LANTAI_NODE_BIN", "LANTAI_NODE_MODULES", "LANTAI_PPTXGENJS_MODULE"):
        monkeypatch.delenv(key, raising=False)
    executable = tmp_path / "Electron"
    executable.touch()
    modules = tmp_path / "Resources/app.asar.unpacked/node_modules"
    (modules / "pptxgenjs").mkdir(parents=True)
    monkeypatch.setenv("PRESENTATION_NODE_EXECUTABLE", str(executable))
    monkeypatch.setenv("PRESENTATION_NODE_MODULES", str(modules))
    monkeypatch.setenv("PRESENTATION_NODE_MODE", "electron")
    monkeypatch.setattr(shutil, "which", lambda _: None)
    from lantai.presentation.models import PresentationJobRequest, PlanningMode, CompositionEngine
    from lantai.presentation.service import PresentationService
    request = PresentationJobRequest(conversation_id="desktop", user_request="Create presentation")
    assert request.planning_mode == PlanningMode.MODEL
    assert request.composition_engine == CompositionEngine.NATIVE_PPTX
    compiler = PresentationService(tmp_path / "jobs").compiler_for(request.composition_engine)
    assert compiler.node_binary == executable
    assert compiler.backend_module_path == modules / "pptxgenjs"
    deck = DeckIR(id="desktop", title="Desktop", slides=[DeckSlide(id="one", layout_id="cover", objects=[])])
    def run(command, **kwargs):
        assert command[0] == str(executable)
        assert kwargs["env"]["ELECTRON_RUN_AS_NODE"] == "1"
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(command[3])
        return subprocess.CompletedProcess(command, 0, '{"engine":"native_pptx","editable_object_count":0}', '')
    monkeypatch.setattr(subprocess, "run", run)
    compiler.compile(deck, tmp_path / "output.pptx")
    assert os.environ.get("ELECTRON_RUN_AS_NODE") != "1"
