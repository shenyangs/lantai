from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as XML
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches, Pt
from PIL import Image
import pytest

from lantai.app.structured_utility import StructuredUtilityResult
from lantai.presentation.intelligence import (
    UtilityPresentationObjectEditPlanner, UtilityPresentationRepairDesigner,
    UtilityPresentationVisualCritic, _RedesignResponse, _NonStructuralRedesignResponse,
)
from lantai.presentation.models import (
    AddObjectAction, DeckIR, DeckSlide, MoveResizeAction, ObjectFrame,
    PresentationObjectEditPlan, RemoveObjectAction, ReplaceImageAction,
    ReplaceTextAction, SetShapeTextAction, TemplatePackageBinding, TextObject,
    UpdateChartAction, VisualCriticSlideInput,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.layouts import check_capacity
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx
from lantai.presentation.reconcile import reconcile_pptx
from lantai.presentation.repair import SlideRedesignRequest


NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def _canonical(element) -> bytes:
    return XML.tostring(element, method="c14n", exclusive=True)


def _fixture(tmp_path: Path, *, complex_kind: str | None = None):
    deck = DeckIR(id="native-group-deck", title="组合对象往返", slides=[
        DeckSlide(id="native-group-slide", layout_id="freeform", objects=[
            TextObject(id="outside-title", role="title", slot="title",
                       text="组合以外的标题", frame=ObjectFrame(x=.8, y=.7, w=7, h=.8)),
        ]),
    ], required_facts=["组合内部的事实不丢失"])
    presentation = Presentation()
    presentation.slide_width = 12192000
    presentation.slide_height = 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.name = "用户原生组合"
    shape = group.shapes.add_shape(1, Inches(9), Inches(2), Inches(2), Inches(1.3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("683FA5")
    text = group.shapes.add_textbox(Inches(9.1), Inches(2.2), Inches(1.7), Inches(.8))
    text.text = "组合内部的事实不丢失"
    text.text_frame.paragraphs[0].runs[0].font.name = "Noto Sans SC"
    text.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    xfrm = group._element.find("p:grpSpPr/a:xfrm", NS)
    xfrm.find("a:ext", NS).set("cx", str(group.width + 2))
    xfrm.find("a:ext", NS).set("cy", str(group.height + 1))
    if complex_kind == "nested":
        nested = group.shapes.add_group_shape()
        nested.shapes.add_textbox(Inches(9), Inches(2), Inches(1), Inches(.5)).text = "不支持嵌套"
    elif complex_kind == "relationship":
        text._element.find("p:nvSpPr/p:cNvPr", NS).set(f"{{{NS['r']}}}id", "rId999")
    elif complex_kind == "foreign_extension":
        text._element.append(XML.Element("{https://invalid.example/native}payload"))
    elif complex_kind == "connector":
        group._element.append(XML.Element(f"{{{NS['p']}}}cxnSp"))
    marker = slide.shapes.add_textbox(0, 0, 100, 100)
    marker.name = "meta:native-group-deck:native-group-slide"
    title = slide.shapes.add_textbox(Inches(.8), Inches(.7), Inches(7), Inches(.8))
    title.name = "outside-title"
    title.text = "组合以外的标题"
    title.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    source = tmp_path / "user-saved.pptx"
    presentation.save(source)
    return deck, source, _canonical(group._element)


def test_native_group_survives_other_object_edit_export_and_second_import(tmp_path):
    deck, source, canonical = _fixture(tmp_path)
    original_source = source.read_bytes()
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    carriers = [obj for obj in imported.slides[0].objects if obj.kind == "native_group"]
    assert len(carriers) == 1
    assert not (tmp_path / "assets").exists()
    # Checkpoint serialization is part of the real edit/import contract.
    imported = DeckIR.model_validate_json(imported.model_dump_json())
    changed, changed_ids = apply_object_edit_plan(imported, PresentationObjectEditPlan(
        slide_id="native-group-slide",
        actions=[ReplaceTextAction(object_id="outside-title", text="AI 只改组合以外的标题")],
    ))
    assert changed_ids == ["outside-title"]
    output = tmp_path / "edited-export.pptx"
    result = NativePptxCompiler().compile(changed, output)
    assert result.degradations == []
    assert result.editable_object_count == 3  # two native children + edited title
    native = Presentation(output)
    groups = [shape for shape in native.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP]
    assert len(groups) == 1
    assert _canonical(groups[0]._element) == canonical
    assert len(groups[0].shapes) == 2
    ids = native.slides[0]._element.xpath(".//p:cNvPr/@id")
    assert len(ids) == len(set(ids))
    # The original group's z-order is retained, not appended after the title.
    assert native.slides[0].shapes[0].shape_type == MSO_SHAPE_TYPE.GROUP
    report = inspect_pptx(changed, output)
    assert not [i.code for i in report.issues if i.code in {
        "editability.object_missing", "editability.native_group_changed",
        "content.required_fact_missing", "editability.native_group_context_changed",
    }]
    second = reconcile_pptx(changed, output, assets_dir=tmp_path / "assets-2")
    assert [obj.id for obj in second.slides[0].objects if obj.kind == "native_group"] == [carriers[0].id]
    assert second.slides[0].objects[0].native_sha256 == carriers[0].native_sha256
    assert next(obj for obj in second.slides[0].objects if obj.id == "outside-title").text == "AI 只改组合以外的标题"
    assert source.read_bytes() == original_source


@pytest.mark.parametrize("complex_kind", ["nested", "relationship", "foreign_extension", "connector"])
def test_complex_native_groups_fail_before_importing_assets_or_touching_source(tmp_path, complex_kind):
    deck, source, _ = _fixture(tmp_path, complex_kind=complex_kind)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="unsupported native group"):
        reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert source.read_bytes() == original
    assert not (tmp_path / "assets").exists()


def _imported(tmp_path):
    deck, source, _ = _fixture(tmp_path)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    return imported, next(obj for obj in imported.slides[0].objects if obj.kind == "native_group")


def test_imported_pure_native_group_is_accepted_by_freeform_capacity(tmp_path):
    imported, _ = _imported(tmp_path)
    assert check_capacity(imported.slides[0]) == []


def test_native_group_children_still_consume_the_existing_freeform_object_limit(tmp_path):
    imported, _ = _imported(tmp_path)
    slide = imported.slides[0]
    title = next(obj for obj in slide.objects if obj.id == "outside-title")
    # The group has two native children: 78 outer objects reach the unchanged
    # 80-object limit, while 79 outer objects must still be rejected.
    slide.objects.extend(title.model_copy(update={"id": f"outer-{i}"}) for i in range(77))
    assert check_capacity(slide) == []
    slide.objects.append(title.model_copy(update={"id": "one-object-too-many"}))
    assert [violation.message for violation in check_capacity(slide)] == [
        "slot accepts at most 80 objects",
    ]


@pytest.mark.parametrize("kind", ["replace_text", "set_shape_text", "update_chart",
                                  "replace_image", "move_resize", "remove_object"])
def test_native_group_actions_reject_atomically_even_after_a_valid_action(tmp_path, kind):
    imported, carrier = _imported(tmp_path)
    action = {
        "replace_text": lambda: ReplaceTextAction(object_id=carrier.id, text="不授权"),
        "set_shape_text": lambda: SetShapeTextAction(object_id=carrier.id, text="不授权"),
        "update_chart": lambda: UpdateChartAction(object_id=carrier.id, categories=["甲"], series={"值": [1]}),
        "replace_image": lambda: ReplaceImageAction(object_id=carrier.id, asset_id="not-an-asset"),
        "move_resize": lambda: MoveResizeAction(object_id=carrier.id, frame=ObjectFrame(x=1, y=1, w=2, h=1)),
        "remove_object": lambda: RemoveObjectAction(object_id=carrier.id),
    }[kind]()
    before = imported.model_dump_json()
    with pytest.raises(ValueError, match="native group is read-only"):
        apply_object_edit_plan(imported, PresentationObjectEditPlan(
            slide_id=imported.slides[0].id,
            actions=[ReplaceTextAction(object_id="outside-title", text="不应部分落盘"), action],
        ))
    assert imported.model_dump_json() == before


@pytest.mark.parametrize("response_model", [_RedesignResponse, _NonStructuralRedesignResponse])
@pytest.mark.parametrize("typed", [False, True])
def test_model_authored_proposals_cannot_introduce_native_group_xml(tmp_path, response_model, typed):
    imported, carrier = _imported(tmp_path)
    with pytest.raises(ValueError):
        response_model.model_validate({"slide": (
            imported.slides[0] if typed else imported.slides[0].model_dump(mode="json")
        )})
    with pytest.raises(ValueError):
        AddObjectAction.model_validate({"object": carrier.model_dump(mode="json")})


def test_native_group_planner_capabilities_and_critic_are_read_only_without_xml(tmp_path, monkeypatch):
    imported, carrier = _imported(tmp_path)
    calls = []

    async def fake_utility(**kwargs):
        calls.append(kwargs)
        return StructuredUtilityResult(value=PresentationObjectEditPlan(
            slide_id=imported.slides[0].id,
            actions=[ReplaceTextAction(object_id="outside-title", text="授权外部标题")],
        ), model_label="fixture/fixture", repaired=False)

    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", fake_utility)
    UtilityPresentationObjectEditPlanner("fixture").plan(imported, imported.slides[0].id, "只改外部标题")
    payload = json.loads(calls[0]["user_prompt"])
    assert all(carrier.id not in ids for ids in payload["editable_object_ids_by_action"].values())
    group = next(obj for obj in payload["slide"]["objects"] if obj["kind"] == "native_group")
    assert group["read_only"] is True
    assert group["native_texts"] == ["组合内部的事实不丢失"]
    assert "native_xml" not in group and "style_context" not in group
    context = UtilityPresentationVisualCritic._slide_context(VisualCriticSlideInput(
        slide_id=imported.slides[0].id, slide_index=1, rendered_path="unused.png",
        rendered_sha256="a" * 64, key_message="不改组合", slide_objects=[carrier],
    ))
    assert context["slide_objects"] == [group]


def test_native_group_whole_page_redesign_fails_before_model_call(tmp_path, monkeypatch):
    imported, _ = _imported(tmp_path)
    async def forbidden_call(**kwargs):
        pytest.fail("native group redesign must not invoke the provider")
    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", forbidden_call)
    with pytest.raises(ValueError, match="native group is read-only"):
        UtilityPresentationRepairDesigner("fixture").redesign(SlideRedesignRequest(
            deck_id=imported.id, slide=imported.slides[0], intent=None, planned_slide=None,
            visual_direction=None, design_context=None, issues=(), immutable_facts=(), deck_context=(),
        ))


@pytest.mark.parametrize("tamper", ["xml", "hash", "frame", "style"])
def test_native_group_tampering_fails_before_overwriting_existing_output(tmp_path, tamper):
    imported, carrier = _imported(tmp_path)
    changes = {
        "xml": {"native_xml": carrier.native_xml.replace("组合内部的事实不丢失", "未经授权")},
        "hash": {"native_sha256": "a" * 64},
        "frame": {"frame": carrier.frame.model_copy(update={"w": carrier.frame.w + .01})},
        "style": {"style_context": carrier.style_context.model_copy(update={"sha256": "b" * 64})},
    }[tamper]
    bad = carrier.model_copy(update=changes)
    imported.slides[0].objects[0] = bad
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"existing user output must survive")
    with pytest.raises(ValueError, match="unsupported native group"):
        NativePptxCompiler().compile(imported, output)
    assert output.read_bytes() == b"existing user output must survive"
    assert not list(tmp_path.glob(".existing*"))


@pytest.mark.parametrize("blocker", ["template", "group_limit", "duplicate_group_name"])
def test_unsupported_native_group_deck_rejects_before_picture_extraction(tmp_path, blocker):
    deck, source, _ = _fixture(tmp_path)
    presentation = Presentation(source)
    slide = presentation.slides[0]
    image = BytesIO()
    Image.new("RGB", (2, 2), "white").save(image, format="PNG")
    image.seek(0)
    picture = slide.shapes.add_picture(image, Inches(1), Inches(3), Inches(1), Inches(1))
    # Ensure a picture would be processed before the group if validation were late.
    slide.shapes._spTree.remove(picture._element)
    slide.shapes._spTree.insert(2, picture._element)
    if blocker == "template":
        deck.template_binding = TemplatePackageBinding(source_path=str(source),
            source_sha256="a" * 64, source_slide_numbers=[1])
    else:
        original = next(shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP)
        for index in range(32 if blocker == "group_limit" else 1):
            group = deepcopy(original._element)
            for nv in group.findall(".//p:cNvPr", NS):
                nv.set("id", str(int(nv.get("id")) + (index + 1) * 100))
            if blocker == "group_limit":
                group.find("p:nvGrpSpPr/p:cNvPr", NS).set("name", f"extra-group-{index}")
            slide.shapes._spTree.append(group)
    presentation.save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="unsupported native group"):
        reconcile_pptx(deck, source, assets_dir=tmp_path / "must-not-write-assets")
    assert not (tmp_path / "must-not-write-assets").exists()
    assert source.read_bytes() == before


@pytest.mark.parametrize("tamper", ["native_xml", "inherited_style"])
def test_quality_gate_detects_changed_group_xml_or_inherited_style(tmp_path, tamper):
    imported, _ = _imported(tmp_path)
    output = tmp_path / "quality-tamper.pptx"
    NativePptxCompiler().compile(imported, output)
    with ZipFile(output) as package:
        entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
    part = "ppt/slides/slide1.xml" if tamper == "native_xml" else "ppt/theme/theme1.xml"
    with ZipFile(output, "w") as package:
        for entry, payload in entries:
            if entry.filename == part:
                if tamper == "native_xml":
                    payload = payload.replace("组合内部的事实不丢失".encode(), "组合内容被改动".encode())
                else:
                    payload = payload.replace(b"Calibri", b"NotTheInheritedFont")
            package.writestr(entry, payload)
    report = inspect_pptx(imported, output)
    expected = "editability.native_group_changed" if tamper == "native_xml" else "editability.native_group_context_changed"
    assert expected in {issue.code for issue in report.issues}
    assert report.passed is False


def test_office_theme_family_context_survives_native_group_roundtrip(tmp_path):
    """A native Office save adds relationship-free theme identity metadata."""
    deck, source, group_xml = _fixture(tmp_path)
    theme_ns = "http://schemas.microsoft.com/office/thememl/2012/main"
    with ZipFile(source) as package:
        entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
    with ZipFile(source, "w") as package:
        for entry, payload in entries:
            if entry.filename == "ppt/theme/theme1.xml":
                theme = XML.fromstring(payload)
                ext_list = XML.SubElement(theme, f"{{{NS['a']}}}extLst")
                extension = XML.SubElement(ext_list, f"{{{NS['a']}}}ext", {
                    "uri": "{05A4C25C-085E-4340-85A3-A5531E510DB2}",
                })
                XML.SubElement(extension, f"{{{theme_ns}}}themeFamily", {
                    "name": "Synthetic Office Theme",
                    "id": "{62F939B6-93AF-4DB8-9C6B-D6C7DFDC589F}",
                    "vid": "{4A3C46E8-61CC-4603-A589-7422A47A8E4A}",
                })
                payload = XML.tostring(theme, encoding="utf-8", xml_declaration=True)
                original_theme = _canonical(theme)
            package.writestr(entry, payload)
    source_bytes = source.read_bytes()
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    carrier = next(obj for obj in imported.slides[0].objects if obj.kind == "native_group")
    assert _canonical(XML.fromstring(carrier.style_context.theme_xml.encode())) == original_theme
    changed, _ = apply_object_edit_plan(imported, PresentationObjectEditPlan(
        slide_id="native-group-slide",
        actions=[ReplaceTextAction(object_id="outside-title", text="AI 只改外部标题，主题不变")],
    ))
    output = tmp_path / "theme-family-export.pptx"
    NativePptxCompiler().compile(changed, output)
    with ZipFile(output) as package:
        assert _canonical(XML.fromstring(package.read("ppt/theme/theme1.xml"))) == original_theme
    second = reconcile_pptx(changed, output, assets_dir=tmp_path / "assets-2")
    second_carrier = next(obj for obj in second.slides[0].objects if obj.kind == "native_group")
    assert second_carrier.native_sha256 == carrier.native_sha256
    assert second_carrier.style_context.sha256 == carrier.style_context.sha256
    assert _canonical(XML.fromstring(second_carrier.native_xml.encode())) == group_xml
    report = inspect_pptx(changed, output)
    assert not [issue.code for issue in report.issues if issue.code in {
        "editability.native_group_changed", "editability.native_group_context_changed",
        "editability.object_missing", "content.required_fact_missing",
    }]
    assert source.read_bytes() == source_bytes
