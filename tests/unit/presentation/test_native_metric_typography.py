"""Public native import and QA locks for authored metric detail captions."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
import pytest

from lantai.presentation.models import (
    CompositionAuthority, DeckIR, DeckSlide, DesignQualityFloor, MetricObject,
    NativeCompositionTokens, ObjectFrame, PresentationDesignProfile,
    ReferenceDeckStyleProfile, TextObject,
)
from lantai.presentation.native_fonts import NativeFontResolver
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx
from lantai.presentation.reconcile import reconcile_pptx


def test_native_metric_detail_retains_readability_identity_after_import(tmp_path):
    metric = MetricObject(
        id="metric-1", value="600", label="合成收入", detail="合成季度口径",
        frame=ObjectFrame(x=1, y=1, w=4, h=2),
    )
    deck = DeckIR(id="metric-deck", title="合成指标测试", composition_tokens=NativeCompositionTokens(
        canvas="FFFFFF", ink="171717", muted="666666", accent="D9583B", subtle="E5E5E0",
        title_font="Noto Sans SC", body_font="Noto Sans SC", title_size=42,
        body_size=18, caption_size=12,
    ), slides=[DeckSlide(
        id="metric-slide", layout_id="freeform",
        composition_authority=CompositionAuthority.AUTHOR, objects=[metric],
    )])
    source = tmp_path / "native-metric.pptx"
    NativePptxCompiler().compile(deck, source)
    before = source.read_bytes()
    shape = next(s for s in Presentation(source).slides[0].shapes if s.name == metric.id)
    sizes = NativeFontResolver(shape).visible_sizes()
    assert sizes == [54, 18, 12]

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    text = imported.slides[0].objects[0]
    assert isinstance(text, TextObject)
    assert text.text == "600\n合成收入\n合成季度口径"
    assert [run.font_size for run in text.runs if run.text.strip()] == sizes
    report = inspect_pptx(imported, source)
    assert not [issue for issue in report.issues if issue.code == "design.font_too_small"]
    assert source.read_bytes() == before
    assert metric.value == "600" and metric.detail == "合成季度口径"
    assert text.native_metric_text is not None
    restored = DeckIR.model_validate_json(imported.model_dump_json())
    output = tmp_path / "metric-export.pptx"
    NativePptxCompiler().compile(restored, output)
    second = reconcile_pptx(restored, output, assets_dir=tmp_path / "second-assets")
    assert second.slides[0].objects[0].native_metric_text == text.native_metric_text
    assert not [i for i in inspect_pptx(second, output).issues if i.code == "design.font_too_small"]
    assert NativeFontResolver(next(s for s in Presentation(output).slides[0].shapes
                                   if s.name == metric.id)).visible_sizes() == sizes
    assert source.read_bytes() == before


def _native_metric(tmp_path, paragraphs=None, sizes=(54, 18, 12)):
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = 12192000, 6858000
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(3))
    shape.name = "metric-1"
    shape.text = "\n".join(paragraphs or ["600", "合成收入", "合成季度口径"])
    for paragraph, size in zip(shape.text_frame.paragraphs, sizes):
        paragraph.font.size = Pt(size)
    metric = MetricObject(id=shape.name, value="600", label="合成收入", detail="合成季度口径")
    deck = DeckIR(id="metric-deck", title="合成指标", slides=[DeckSlide(
        id="metric-slide", layout_id="freeform", objects=[metric],
    )])
    source = tmp_path / "metric-native.pptx"
    presentation.save(source)
    return deck, source, presentation, shape


@pytest.mark.parametrize("sizes, floor", [((14, 18, 12), 17), ((54, 14, 12), 17), ((54, 18, 11), 12)])
def test_metric_value_label_and_detail_each_keep_their_floor(tmp_path, sizes, floor):
    deck, source, presentation, shape = _native_metric(tmp_path)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    for paragraph, size in zip(shape.text_frame.paragraphs, sizes):
        paragraph.font.size = Pt(size)
    presentation.save(source)
    before = source.read_bytes()
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert len(issues) == 1 and issues[0].severity == "error"
    assert f"{floor} pt" in issues[0].message
    assert NativeFontResolver(Presentation(source).slides[0].shapes[0]).visible_sizes() == list(sizes)
    assert source.read_bytes() == before


@pytest.mark.parametrize("paragraphs, sizes", [
    (["合成收入", "600", "合成季度口径"], (18, 54, 12)),
    (["600", "合成收入", "合成季度口径", "第二行合成口径"], (54, 18, 12, 12)),
])
def test_exact_known_metric_orders_and_multiline_detail_are_preserved(tmp_path, paragraphs, sizes):
    deck, source, _, _ = _native_metric(tmp_path, paragraphs, sizes)
    if len(paragraphs) == 4:
        deck.slides[0].objects[0].detail = "合成季度口径\n第二行合成口径"
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].native_metric_text is not None
    assert not [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]


@pytest.mark.parametrize("paragraphs", [
    ["601", "合成收入", "合成季度口径"],
    ["600", "新增正文", "合成季度口径"],
    ["600", "合成收入", "合成季度口径", "新增正文"],
    ["合成季度口径", "600", "合成收入"],
])
def test_unmatched_native_wording_cannot_reuse_caption_identity(tmp_path, paragraphs):
    deck, source, _, _ = _native_metric(tmp_path, paragraphs, (54, 18, 12, 12))
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].native_metric_text is None
    assert [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]


@pytest.mark.parametrize("role", ["body", "metric"])
def test_role_alone_never_exempts_unknown_small_body_text(tmp_path, role):
    deck, source, _, _ = _native_metric(tmp_path)
    deck.slides[0].objects = [TextObject(
        id="metric-1", role=role, text="600\n合成收入\n合成季度口径",
    )]
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    assert imported.slides[0].objects[0].native_metric_text is None
    assert [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]


def test_metric_caption_cannot_weaken_a_stricter_declared_quality_floor(tmp_path):
    deck, source, _, _ = _native_metric(tmp_path)
    deck.design_profile = PresentationDesignProfile(
        identity="synthetic-strict-caption", style_profile=ReferenceDeckStyleProfile(id="synthetic"),
        quality_floor=DesignQualityFloor(minimum_caption_font_size=14),
    )
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert len(issues) == 1 and "14 pt" in issues[0].message


@pytest.mark.parametrize("size, fails", [(12, False), (11.9, True)])
@pytest.mark.parametrize("field", [False, True])
def test_every_caption_run_and_field_is_checked_at_its_effective_size(tmp_path, size, fails, field):
    deck, source, presentation, shape = _native_metric(tmp_path)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    paragraph = shape.text_frame.paragraphs[2]
    paragraph.clear()
    if field:
        paragraph._p.append(parse_xml(
            f'<a:fld {nsdecls("a")} id="{{00000000-0000-0000-0000-000000000001}}" type="slidenum">'
            f'<a:rPr sz="{round(size * 100)}"/><a:t>合成季度口径</a:t></a:fld>'
        ))
    else:
        for text in ("合成", "季度口径"):
            run = paragraph.add_run()
            run.text, run.font.size = text, Pt(size)
    presentation.save(source)
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert bool(issues) is fails
    if fails:
        assert "12 pt" in issues[0].message


@pytest.mark.parametrize("scale, floor", [(80000, 17), (95000, 12)])
def test_native_autofit_cannot_hide_a_metric_readability_regression(tmp_path, scale, floor):
    deck, source, presentation, shape = _native_metric(tmp_path)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    body = shape._element.find(qn("p:txBody")).find(qn("a:bodyPr"))
    for child in list(body):
        if child.tag in {qn("a:noAutofit"), qn("a:normAutofit"), qn("a:spAutoFit")}:
            body.remove(child)
    body.append(parse_xml(f'<a:normAutofit {nsdecls("a")} fontScale="{scale}"/>'))
    presentation.save(source)
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert len(issues) == 1 and f"{floor} pt" in issues[0].message


def test_caption_metadata_is_discarded_after_native_wording_changes(tmp_path):
    deck, source, presentation, shape = _native_metric(tmp_path)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    original = imported.model_dump_json()
    shape.text_frame.paragraphs[2].text = "新加入的正文"
    presentation.save(source)
    before = source.read_bytes()
    assert [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    second = reconcile_pptx(imported, source, assets_dir=tmp_path / "second-assets")
    assert second.slides[0].objects[0].native_metric_text is None
    assert [i for i in inspect_pptx(second, source).issues if i.code == "design.font_too_small"]
    assert imported.model_dump_json() == original and source.read_bytes() == before


def test_declared_caption_floor_cannot_reduce_the_existing_12pt_minimum(tmp_path):
    deck, source, presentation, shape = _native_metric(tmp_path)
    deck.design_profile = PresentationDesignProfile(
        identity="synthetic-loose-caption", style_profile=ReferenceDeckStyleProfile(id="synthetic"),
        quality_floor=DesignQualityFloor(minimum_caption_font_size=9),
    )
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    shape.text_frame.paragraphs[2].font.size = Pt(11)
    presentation.save(source)
    issues = [i for i in inspect_pptx(imported, source).issues if i.code == "design.font_too_small"]
    assert len(issues) == 1 and "12 pt" in issues[0].message
