"""Native card containment cannot excuse a foreground hiding its content."""

from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.models import (
    ChartObject,
    DeckIR,
    DeckSlide,
    ObjectFrame,
    ShapeObject,
    TableObject,
    TextObject,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx


@pytest.fixture(params=[NativePptxCompiler, PresentationCompiler], ids=["native", "legacy"])
def compiler(request):
    return request.param()


def _content_deck(kind: str, *, background_last: bool) -> DeckIR:
    frame = ObjectFrame(x=2.0, y=2.0, w=8.5, h=3.5)
    if kind == "chart":
        content = ChartObject(
            id="editable-content", chart_type="column", categories=["A", "B"],
            series={"数量": [3, 5]}, unit="项", derivation="直接引用合成测试数据。",
            source_refs=["fixture://native-foreground-occlusion"], frame=frame,
        )
    elif kind == "table":
        content = TableObject(
            id="editable-content", headers=["类别", "数量"],
            rows=[["A", "3"], ["B", "5"]], frame=frame,
        )
    else:
        content = ShapeObject(
            id="editable-content", shape_type="rect", text="Category A",
            fill_color="1565C0", text_color="FFFFFF", font_size=24, frame=frame,
        )
    background = ShapeObject(
        id="content-background", shape_type="rect", fill_color="FFFFFF",
        frame=ObjectFrame(x=1.5, y=1.5, w=9.5, h=4.5),
    )
    title = TextObject(
        id="occlusion-title", role="title", slot="title", text="合成原生对象遮挡测试",
        font_size=32, frame=ObjectFrame(x=0.6, y=0.4, w=12, h=0.8),
    )
    objects = [content, background] if background_last else [background, content]
    return DeckIR(
        id="native-occlusion-deck", title=title.text, required_facts=[title.text],
        slides=[DeckSlide(id="native-occlusion-slide", layout_id="freeform", objects=[title, *objects])],
    )


def _native_content_shapes(presentation):
    return [shape for shape in presentation.slides[0].shapes
            if shape.name in {"editable-content", "content-background"}]


def _assert_occlusion_blocked(report) -> None:
    assert report.passed is False
    assert report.delivery_status == "failed"
    assert any(issue.code == "layout.complete_overlap" and issue.severity == "error"
               for issue in report.issues)


@pytest.mark.parametrize("kind", ["chart", "table", "shape"])
@pytest.mark.parametrize("shape_type", ["rect", "round_rect"])
def test_opaque_foreground_cannot_hide_native_content(tmp_path, compiler, kind, shape_type):
    deck = _content_deck(kind, background_last=True)
    deck.slides[0].objects[-1].shape_type = shape_type
    output = tmp_path / "covered-content.pptx"
    compiler.compile(deck, output)
    native_shapes = _native_content_shapes(Presentation(output))
    assert [shape.name for shape in native_shapes] == ["editable-content", "content-background"]
    assert native_shapes[0].width * native_shapes[0].height < native_shapes[1].width * native_shapes[1].height * 0.9
    if kind == "chart":
        assert native_shapes[0].has_chart
    elif kind == "table":
        assert native_shapes[0].has_table
    else:
        assert native_shapes[0].text == "Category A"

    _assert_occlusion_blocked(inspect_pptx(deck, output))


@pytest.mark.parametrize("kind", ["chart", "table", "shape"])
def test_native_content_remains_allowed_on_an_earlier_background(tmp_path, compiler, kind):
    deck = _content_deck(kind, background_last=False)
    output = tmp_path / "visible-content.pptx"
    compiler.compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is True, [issue.code for issue in report.issues]
    assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)


@pytest.mark.parametrize("kind", ["chart", "table", "shape"])
def test_occlusion_uses_the_reopened_native_order(tmp_path, compiler, kind):
    deck = _content_deck(kind, background_last=False)
    output = tmp_path / "reordered-content.pptx"
    compiler.compile(deck, output)
    native = Presentation(output)
    background, content = _native_content_shapes(native)
    background._element.getparent().remove(background._element)
    content._element.addnext(background._element)
    native.save(output)
    assert deck.slides[0].objects[1].id == "content-background"

    _assert_occlusion_blocked(inspect_pptx(deck, output))


@pytest.mark.parametrize("kind", ["chart", "table", "shape"])
@pytest.mark.parametrize("alpha", ["0", "50000", "100000"])
def test_foreground_native_opacity_is_not_guessed(tmp_path, compiler, kind, alpha):
    deck = _content_deck(kind, background_last=True)
    output = tmp_path / "alpha-content.pptx"
    compiler.compile(deck, output)
    native = Presentation(output)
    background = _native_content_shapes(native)[1]
    color = background._element.xpath("./p:spPr/a:solidFill/*")[0]
    opacity = OxmlElement("a:alpha")
    opacity.set("val", alpha)
    color.append(opacity)
    native.save(output)

    report = inspect_pptx(deck, output)

    if alpha == "100000":
        _assert_occlusion_blocked(report)
    else:
        assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)


@pytest.mark.parametrize("kind", ["chart", "table", "shape"])
@pytest.mark.parametrize("variation", [
    "alpha-mod", "alpha-offset", "no-fill", "rotated", "nonrectangular",
])
def test_native_occlusion_requires_proven_opaque_rectangular_geometry(
    tmp_path, compiler, kind, variation,
):
    deck = _content_deck(kind, background_last=True)
    output = tmp_path / "unproven-opaque-card.pptx"
    compiler.compile(deck, output)
    native = Presentation(output)
    background = _native_content_shapes(native)[1]
    if variation in {"alpha-mod", "alpha-offset"}:
        color = background._element.xpath("./p:spPr/a:solidFill/*")[0]
        transform = OxmlElement("a:alphaMod" if variation == "alpha-mod" else "a:alphaOff")
        transform.set("val", "50000" if variation == "alpha-mod" else "-50000")
        color.append(transform)
    elif variation == "no-fill":
        background.fill.background()
    elif variation == "rotated":
        background.rotation = 15
    else:
        geometry = background._element.xpath("./p:spPr/a:prstGeom")[0]
        geometry.set("prst", "ellipse")
    native.save(output)

    # These native states do not prove full rectangular opaque occlusion.
    # A separate rendered visual review can still reject the resulting slide.
    report = inspect_pptx(deck, output)

    assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)
