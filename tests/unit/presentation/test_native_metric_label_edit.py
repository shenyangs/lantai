"""Native metric label edits retain the imported paragraph typography."""

from pptx import Presentation
import pytest

from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    CompositionAuthority, DeckIR, DeckSlide, MetricObject, NativeCompositionTokens,
    ObjectFrame, PresentationEditOperation, PresentationObjectEditPlan,
    NativeMetricText, ReplaceTextAction, RichTextRun, TextObject,
)
from lantai.presentation.native_fonts import NativeFontResolver
from lantai.presentation.natural_edit import apply_object_edit_plan, plan_text_edit
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import inspect_pptx
from lantai.presentation.reconcile import reconcile_pptx


def _import_metric(tmp_path):
    deck = DeckIR(
        id="native-metric-edit", title="合成指标编辑", composition_tokens=NativeCompositionTokens(
            canvas="FFFFFF", ink="17242E", muted="5B6A74", accent="315F9A", subtle="E5E5E0",
            title_font="Noto Sans SC", body_font="Noto Sans SC", title_size=42,
            body_size=18, caption_size=12,
        ), slides=[
            DeckSlide(id="metric-slide", layout_id="freeform",
                      composition_authority=CompositionAuthority.AUTHOR, objects=[MetricObject(
                          id="metric-1", value="120", label="期初收入（万元）", detail="合成季度资料",
                          frame=ObjectFrame(x=.82, y=3.72, w=2.18, h=1.22),
                      )]),
            DeckSlide(id="untouched-slide", layout_id="freeform",
                      composition_authority=CompositionAuthority.AUTHOR, objects=[TextObject(
                          id="untouched", role="body", text="不在编辑范围内", font_size=18,
                          frame=ObjectFrame(x=1, y=1, w=5, h=1),
                      )]),
        ],
    )
    source = tmp_path / "source.pptx"
    NativePptxCompiler().compile(deck, source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    return imported, source


def _shape(path):
    return next(shape for shape in Presentation(path).slides[0].shapes if shape.name == "metric-1")


def _styles(shape):
    fonts = NativeFontResolver(shape)
    return [(fonts.run_size(paragraph, run), str(fonts.run_rgb(paragraph, run)),
             fonts.run_flag(paragraph, run, "b"), fonts.run_flag(paragraph, run, "i"))
            for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.text.strip()]


def test_compiled_label_edit_keeps_value_body_and_caption_styles(tmp_path):
    deck, source = _import_metric(tmp_path)
    original_bytes = source.read_bytes()
    before = _shape(source)
    assert NativeFontResolver(before).visible_sizes() == [34.16, 18, 12]
    untouched_xml = Presentation(source).slides[1]._element.xml
    new_text = "120\nQ1当季收入（万元）\n合成季度资料"
    changed, ids = apply_object_edit_plan(deck, PresentationObjectEditPlan(
        slide_id="metric-slide", actions=[ReplaceTextAction(object_id="metric-1", text=new_text)],
        rationale="只修订指标标签，保留原生格式",
    ))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, deck, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["metric-slide"], object_ids=ids,
            allowed_content_types={"text"}, instruction="只修订指标标签",
        ), NativePptxCompiler(), compiled_slide_id="metric-slide",
    )
    after = _shape(preview.candidate_path)
    assert after.text == new_text
    assert _styles(after) == _styles(before)
    assert changed.slides[0].objects[0].native_metric_text.label == "Q1当季收入（万元）"
    assert changed.slides[1] == deck.slides[1]
    assert Presentation(preview.candidate_path).slides[1]._element.xml == untouched_xml
    assert source.read_bytes() == original_bytes
    assert not [issue for issue in inspect_pptx(changed, preview.candidate_path).issues
                if issue.code in {"layout.text_overflow_risk", "design.font_too_small"}]


@pytest.mark.parametrize("new_text", [
    "121\nQ1当季收入（万元）\n合成季度资料",
    "120\nQ1当季收入（万元）\n新增正文",
    "120\nQ1当季收入（万元）\n合成季度资料\n新增正文",
    "120\nQ1当季收入\n（万元）\n合成季度资料",
    "120\n\n合成季度资料",
    "合成季度资料\n120\nQ1当季收入（万元）",
])
def test_other_replacements_cannot_inherit_the_small_caption_identity(new_text):
    text = TextObject(
        id="metric-1", role="metric", text="120\n期初收入（万元）\n合成季度资料",
        font_size=34.16, native_metric_text=NativeMetricText(
            value="120", label="期初收入（万元）", detail="合成季度资料",
        ), runs=[
            RichTextRun(text="120", font_size=34.16, bold=True),
            RichTextRun(text="\n", paragraph_break=True),
            RichTextRun(text="期初收入（万元）", font_size=18, bold=True),
            RichTextRun(text="\n", paragraph_break=True),
            RichTextRun(text="合成季度资料", font_size=12),
        ],
    )
    deck = DeckIR(id="metric-deck", title="合成指标", slides=[DeckSlide(
        id="metric-slide", layout_id="freeform", objects=[text],
    )])
    original = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, PresentationObjectEditPlan(
        slide_id="metric-slide", actions=[ReplaceTextAction(object_id="metric-1", text=new_text)],
    ))
    edited = changed.slides[0].objects[0]
    assert ids == ["metric-1"] and edited.text == new_text
    assert edited.native_metric_text is None and edited.runs == []
    assert deck.model_dump_json() == original


def test_identical_label_action_does_not_clear_styles_or_claim_a_change(tmp_path):
    deck, source = _import_metric(tmp_path)
    original_bytes, original_deck = source.read_bytes(), deck.model_dump_json()
    with pytest.raises(ValueError, match="no material change"):
        apply_object_edit_plan(deck, PresentationObjectEditPlan(
            slide_id="metric-slide", actions=[ReplaceTextAction(
                object_id="metric-1", text=deck.slides[0].objects[0].text,
            )],
        ))
    assert deck.model_dump_json() == original_deck and source.read_bytes() == original_bytes


def test_transactional_quoted_label_edit_keeps_native_paragraphs(tmp_path):
    deck, source = _import_metric(tmp_path)
    original_bytes = source.read_bytes()
    before = _shape(source)
    changed, replacements = plan_text_edit(deck, "metric-slide", "将“期初收入”改为“Q1当季收入”")
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare(
        source, deck, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["metric-slide"], object_ids=["metric-1"],
            allowed_content_types={"text"}, instruction="只修订指标标签",
        ), replacements,
    )
    after = _shape(preview.candidate_path)
    assert after.text == changed.slides[0].objects[0].text
    assert [paragraph.text for paragraph in after.text_frame.paragraphs] == [
        "120", "Q1当季收入（万元）", "合成季度资料",
    ]
    assert _styles(after) == _styles(before)
    assert source.read_bytes() == original_bytes
    assert not [issue for issue in inspect_pptx(changed, preview.candidate_path).issues
                if issue.code in {"layout.text_overflow_risk", "design.font_too_small"}]


@pytest.mark.parametrize("unsupported", [
    "body_role", "missing_runs", "stale_runs", "stale_metric",
    "invalid_break", "mixed_label_runs", "ambiguous_order",
])
def test_unproven_label_structures_cannot_keep_caption_identity(tmp_path, unsupported):
    deck, _ = _import_metric(tmp_path)
    obj = deck.slides[0].objects[0]
    if unsupported == "body_role":
        obj.role = "body"
    elif unsupported == "missing_runs":
        obj.runs = []
    elif unsupported == "stale_runs":
        obj.runs[0].text = "121"
    elif unsupported == "stale_metric":
        obj.native_metric_text.detail = "与原生内容不一致"
    elif unsupported == "invalid_break":
        obj.runs[1].text = "\v"
        obj.text = "".join(run.text for run in obj.runs)
    elif unsupported == "mixed_label_runs":
        label = obj.runs[2]
        obj.runs[2:3] = [
            label.model_copy(update={"text": "期初"}),
            label.model_copy(update={"text": "收入（万元）", "bold": False}),
        ]
    else:
        obj.native_metric_text.label = "120"
        obj.runs[2].text = "120"
        obj.text = "".join(run.text for run in obj.runs)
    original = deck.model_dump_json()
    old_label = obj.native_metric_text.label
    new_text = obj.text.replace(old_label, "Q1当季收入（万元）", 1)
    changed, _ = apply_object_edit_plan(deck, PresentationObjectEditPlan(
        slide_id="metric-slide", actions=[ReplaceTextAction(
            object_id="metric-1", text=new_text,
        )],
    ))
    edited = changed.slides[0].objects[0]
    assert edited.text == new_text
    assert edited.runs == [] and edited.native_metric_text is None
    assert deck.model_dump_json() == original
