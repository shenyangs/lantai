"""Existing native captions are not newly authored body typography."""

from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from lantai.presentation.models import DeckIR, DeckSlide, ShapeObject
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _native_caption(tmp_path: Path):
    native = Presentation()
    native.slide_width = 12192000
    native.slide_height = 6858000
    slide = native.slides.add_slide(native.slide_layouts[6])
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(6.9), Inches(5), Inches(.3),
    )
    shape.name = "existing-template-footer"
    shape.text = "Synthetic authorized footer"
    shape.text_frame.paragraphs[0].font.size = Pt(11)
    shape.text_frame.paragraphs[0].font.color.rgb = RGBColor.from_string("102030")
    for run in shape.text_frame.paragraphs[0].runs:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string("102030")
    source = tmp_path / "native-caption.pptx"
    native.save(source)
    deck = DeckIR(id="caption-deck", title="Synthetic native template", slides=[DeckSlide(
        id="caption-slide", layout_id="freeform", objects=[],
    )])
    return deck, source


def test_unknown_native_shape_caption_preserves_11pt_through_valid_ir_and_export(tmp_path):
    deck, source = _native_caption(tmp_path)
    source_before = source.read_bytes()
    deck_before = deck.model_dump_json()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    caption = imported.slides[0].objects[0]
    assert isinstance(caption, ShapeObject)
    assert caption.text == "Synthetic authorized footer"
    assert caption.runs and all(run.font_size == 11 for run in caption.runs)
    # A successful native import must remain valid when checkpointed, not only
    # survive an unchecked model_copy with an out-of-contract body font size.
    restored = DeckIR.model_validate_json(imported.model_dump_json())
    output = tmp_path / "exported-caption.pptx"
    NativePptxCompiler().compile(restored, output)
    native = next(shape for shape in Presentation(output).slides[0].shapes if shape.name == caption.id)
    assert native.text == caption.text
    assert all(run.font.size.pt == 11 for run in native.text_frame.paragraphs[0].runs)
    second = reconcile_pptx(restored, output, assets_dir=tmp_path / "second-assets")
    DeckIR.model_validate_json(second.model_dump_json())
    assert all(run.font_size == 11 for run in second.slides[0].objects[0].runs)
    assert source.read_bytes() == source_before
    assert deck.model_dump_json() == deck_before


def test_native_caption_layout_and_absent_paint_survive_public_roundtrip(tmp_path):
    deck, source = _native_caption(tmp_path)
    native = Presentation(source)
    shape = native.slides[0].shapes[0]
    shape.fill.background()
    shape.line.fill.background()
    shape.line.width = Pt(.75)
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    shape.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    for name, value in (("top", 3), ("right", 7), ("bottom", 5), ("left", 11)):
        setattr(shape.text_frame, "margin_" + name, Pt(value))
    native.save(source)
    source_before, deck_before = source.read_bytes(), deck.model_dump_json()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    restored = DeckIR.model_validate_json(imported.model_dump_json())
    output = tmp_path / "native-caption-layout.pptx"
    NativePptxCompiler().compile(restored, output)
    exported = next(s for s in Presentation(output).slides[0].shapes
                    if s.name == restored.slides[0].objects[0].id)
    assert exported.text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT
    assert exported.text_frame.vertical_anchor == MSO_ANCHOR.BOTTOM
    assert exported.line.width.pt == .75
    for name, value in (("top", 3), ("right", 7), ("bottom", 5), ("left", 11)):
        assert getattr(exported.text_frame, "margin_" + name).pt == value
    # Both noFill and a fully transparent native solid fill are invisible.
    for paint in (exported.fill, exported.line.fill):
        if paint.type is not None and paint.type != MSO_FILL.BACKGROUND:
            assert paint._xPr.xpath("./a:solidFill/*/a:alpha/@val") == ["0"]
    second = reconcile_pptx(restored, output, assets_dir=tmp_path / "second-assets")
    assert second.slides[0].objects[0].model_dump() == restored.slides[0].objects[0].model_dump()
    assert source.read_bytes() == source_before
    assert deck.model_dump_json() == deck_before


def test_new_authored_shape_keeps_existing_centered_caption_defaults(tmp_path):
    from lantai.presentation.models import ObjectFrame

    shape = ShapeObject(id="authored-shape", frame=ObjectFrame(x=1, y=1, w=3, h=1), text="New caption")
    deck = DeckIR(id="authored-deck", title="Synthetic author defaults", slides=[DeckSlide(
        id="authored-slide", layout_id="freeform", objects=[shape],
    )])
    output = tmp_path / "authored-defaults.pptx"
    NativePptxCompiler().compile(deck, output)
    native = next(s for s in Presentation(output).slides[0].shapes if s.name == shape.id)
    assert native.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER
    assert native.text_frame.vertical_anchor == MSO_ANCHOR.MIDDLE
    assert native.text_frame.margin_left.pt == .08
    assert native.line.fill.type is not None


@pytest.mark.parametrize("empty_line", [False, True])
def test_native_caption_without_line_paint_or_style_does_not_gain_an_outline(tmp_path, empty_line):
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn

    deck, source = _native_caption(tmp_path)
    native = Presentation(source)
    shape = native.slides[0].shapes[0]
    shape.fill.background()
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    line = shape._element.spPr.find(qn("a:ln"))
    if line is not None:
        shape._element.spPr.remove(line)
    if empty_line:
        shape._element.spPr.append(OxmlElement("a:ln"))
    native.save(source)
    source_before = source.read_bytes()

    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    caption = imported.slides[0].objects[0]
    assert caption.native_style.line_none is True
    assert caption.line_color is None
    output = tmp_path / "caption-without-outline.pptx"
    NativePptxCompiler().compile(imported, output)
    exported = next(s for s in Presentation(output).slides[0].shapes if s.name == caption.id)
    assert exported.line.fill._xPr.xpath("./a:solidFill/*/a:alpha/@val") == ["0"]
    second = reconcile_pptx(imported, output, assets_dir=tmp_path / "second-assets")
    assert second.slides[0].objects[0].model_dump() == caption.model_dump()
    assert source.read_bytes() == source_before


def test_native_caption_explicit_rgb_outline_remains_visible(tmp_path):
    deck, source = _native_caption(tmp_path)
    native = Presentation(source)
    native.slides[0].shapes[0].line.color.rgb = RGBColor.from_string("627384")
    native.slides[0].shapes[0].line.width = Pt(2.5)
    native.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    caption = imported.slides[0].objects[0]
    assert caption.native_style.line_none is False
    assert caption.line_color == "627384"
    output = tmp_path / "caption-with-outline.pptx"
    NativePptxCompiler().compile(imported, output)
    exported = next(s for s in Presentation(output).slides[0].shapes if s.name == caption.id)
    assert exported.line.fill.fore_color.rgb == RGBColor.from_string("627384")
    assert exported.line.width.pt == 2.5
    assert not exported.line.fill._xPr.xpath("./a:solidFill/*/a:alpha[@val='0']")


@pytest.mark.parametrize("margin", [-1, 145, float("nan"), float("inf")])
def test_native_caption_style_rejects_unbounded_point_margins(margin):
    from pydantic import ValidationError
    from lantai.presentation.models import NativeShapeStyle

    with pytest.raises(ValidationError):
        NativeShapeStyle(margin_top=margin, margin_right=0, margin_bottom=0, margin_left=0)


def test_native_caption_style_rejects_partial_margin_snapshot():
    from pydantic import ValidationError
    from lantai.presentation.models import NativeShapeStyle

    with pytest.raises(ValidationError):
        NativeShapeStyle(margin_left=5)


def test_native_caption_rejects_unsupported_alignment_before_extracting_assets(tmp_path):
    from io import BytesIO
    from PIL import Image

    deck, source = _native_caption(tmp_path)
    native = Presentation(source)
    shape = native.slides[0].shapes[0]
    shape.text_frame.paragraphs[0].alignment = PP_ALIGN.JUSTIFY
    stream = BytesIO()
    Image.new("RGB", (2, 2), "white").save(stream, format="PNG")
    stream.seek(0)
    picture = native.slides[0].shapes.add_picture(stream, Inches(1), Inches(1))
    native.slides[0].shapes._spTree.remove(picture._element)
    native.slides[0].shapes._spTree.insert(2, picture._element)
    native.save(source)
    source_before, deck_before = source.read_bytes(), deck.model_dump_json()
    assets = tmp_path / "assets"
    with pytest.raises(ValueError, match="unsupported native shape paragraph alignment"):
        reconcile_pptx(deck, source, assets_dir=assets)
    assert not assets.exists()
    assert source.read_bytes() == source_before
    assert deck.model_dump_json() == deck_before


def test_native_caption_text_only_edit_inherits_11pt_and_keeps_native_style(tmp_path):
    from lantai.presentation.models import PresentationObjectEditPlan
    from lantai.presentation.natural_edit import apply_object_edit_plan

    deck, source = _native_caption(tmp_path)
    native = Presentation(source)
    native.slides[0].shapes[0].line.fill.background()
    native.slides[0].shapes[0].text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    native.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    caption = imported.slides[0].objects[0]
    before = imported.model_dump_json()
    plan = PresentationObjectEditPlan.model_validate({
        "slide_id": imported.slides[0].id,
        "actions": [{"kind": "set_shape_text", "object_id": caption.id, "text": "Edited synthetic footer"}],
    })
    edited, changed = apply_object_edit_plan(imported, plan)
    assert changed == [caption.id]
    assert edited.slides[0].objects[0].model_dump() == {**caption.model_dump(), "text": "Edited synthetic footer"}
    output = tmp_path / "edited-native-caption.pptx"
    NativePptxCompiler().compile(edited, output)
    exported = next(s for s in Presentation(output).slides[0].shapes if s.name == caption.id)
    assert exported.text == "Edited synthetic footer"
    assert exported.text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT
    assert exported.line.fill._xPr.xpath("./a:solidFill/*/a:alpha/@val") == ["0"]
    assert all(run.font.size.pt == 11 for run in exported.text_frame.paragraphs[0].runs)
    assert imported.model_dump_json() == before
