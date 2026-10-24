from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.layouts import CORE_LAYOUTS
from lantai.presentation.models import AssetRecord, ChartObject, DeckIR, DeckSlide, ImageObject, MetricObject, ObjectFrame, ShapeObject, TableObject, TextObject
from lantai.presentation.quality import file_sha256


def _object(kind: str, slot: str, index: int, image_path: Path):
    if kind == "image":
        return ImageObject(
            asset_id="fixture-image",
            alt="可编辑图片占位",
            slot=slot,
        )
    if kind == "metric":
        return MetricObject(label="指标", value=f"{index + 1}00", slot=slot)
    if kind == "table":
        return TableObject(headers=["项目", "结论"], rows=[["A", "通过"]], slot=slot)
    if kind == "chart":
        return ChartObject(
            categories=["A", "B"],
            series={"指标": [1, 2]},
            unit="count",
            derivation="direct fixture values",
            source_refs=["fixture://all-layouts"],
            slot=slot,
        )
    role = "title" if slot == "title" else "body"
    return TextObject(role=role, text=f"{slot} 的可编辑内容", slot=slot)


def test_all_semantic_layouts_compile_to_native_editable_objects(tmp_path: Path) -> None:
    image_path = tmp_path / "reference-image.png"
    Image.new("RGB", (640, 360), "#204060").save(image_path)
    asset = AssetRecord(
        asset_id="fixture-image",
        kind="image",
        local_path=str(image_path),
        source_ref="fixture://all-layouts/image",
        rights_status="generated",
        sha256=file_sha256(image_path),
        dimensions=(640, 360),
        usage_role="explanation",
    )
    slides: list[DeckSlide] = []
    expected_ids: set[str] = set()
    for layout in CORE_LAYOUTS.values():
        objects = []
        if layout.id == "freeform":
            obj = ShapeObject(
                id="freeform-shape",
                frame=ObjectFrame(x=1, y=1.5, w=4, h=1.5),
                text="自由构图中的原生对象",
            )
            objects.append(obj)
            expected_ids.add(obj.id)
            slides.append(DeckSlide(layout_id=layout.id, objects=objects))
            continue
        for index, slot in enumerate(layout.slots):
            kind = sorted(slot.accepts)[0]
            obj = _object(kind, slot.id, index, image_path)
            objects.append(obj)
            expected_ids.add(obj.id)
        slides.append(DeckSlide(layout_id=layout.id, objects=objects))

    output = tmp_path / "all-layouts.pptx"
    result = PresentationCompiler().compile(
        DeckIR(title="17 类版式", slides=slides, assets=[asset]),
        output,
    )

    assert result.slide_count == 17
    assert result.degradations == []
    presentation = Presentation(output)
    actual_ids = {shape.name for slide in presentation.slides for shape in slide.shapes}
    assert expected_ids <= actual_ids
