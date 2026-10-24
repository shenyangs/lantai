# -*- coding: utf-8 -*-
"""Model-selected fact serialization through the production author boundary."""

from __future__ import annotations

import json
from copy import deepcopy
from itertools import product
from types import SimpleNamespace

import pytest
from jsonschema import Draft202012Validator
from pptx import Presentation

from lantai.app.structured_utility import StructuredTaskValidationError
from lantai.presentation.intelligence import UtilityPresentationPageAuthor
from lantai.presentation.models import InputMaterial, PresentationJobRequest, TextObject
from lantai.presentation.page_authoring import (
    PresentationPageAuthorContentError, grounded_source_text, validate_authored_page,
)
from lantai.presentation.planner import (
    build_brief,
    build_deck_ir,
    build_deterministic_fixture_plan,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler


def _case():
    request = PresentationJobRequest(
        conversation_id="author-fact-reference-transport-test",
        user_request="制作一页收入判断，不增添数据",
        slide_count=1,
        materials=[InputMaterial(
            name="合成事实.md",
            content="甲方收入为42万元。\n风险尚待核验。",
        )],
    )
    brief = build_brief(request)
    plan = build_deterministic_fixture_plan(request, brief)
    planned = plan.slides[0].model_copy(update={
        "title": "先核验证据，再判断收入",
        "charts": [],
        "intent": plan.slides[0].intent.model_copy(update={
            "required_facts": ["甲方收入为42万元。", "风险尚待核验。"],
            "immutable_facts": ["甲方收入为42万元。"],
        }),
    })
    plan = plan.model_copy(update={"slides": [planned]})
    scaffold = build_deck_ir(request, plan).slides[0]
    scene = {
        "objects": [
            {
                "kind": "text", "id": "title", "role": "title",
                "text_refs": ["storyboard.title"], "font_size": 36,
                "frame": {"x": 0.8, "y": 0.7, "w": 11.5, "h": 0.9},
            },
            {
                "kind": "text", "id": "evidence", "role": "body",
                "text_refs": [
                    "required_facts[1]", "required_facts[0]", "immutable_facts[0]",
                ],
                "font_size": 24,
                "frame": {"x": 0.8, "y": 2, "w": 11.5, "h": 3.8},
            },
        ],
        "speaker_notes": "",
        "object_groups": [{"role": "evidence", "object_ids": ["evidence"]}],
        "relationships": [],
    }
    return request, brief, plan, scaffold, scene


def _author(monkeypatch, case):
    request, brief, plan, scaffold, scene = case

    async def transport(**kwargs):
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="test/model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    return UtilityPresentationPageAuthor("agent-1").author(
        request, brief, plan, plan.slides[0], scaffold,
    )


def test_public_author_resolves_only_model_selected_same_page_facts(monkeypatch) -> None:
    request, brief, plan, scaffold, scene = _case()
    calls = []

    async def transport(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="test/model")

    # Replace only the external text transport, not schema parsing or authoring.
    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    authored = UtilityPresentationPageAuthor("agent-1").author(
        request, brief, plan, plan.slides[0], scaffold,
    )
    assert all(type(obj) is TextObject for obj in authored.objects)
    assert [obj.text for obj in authored.objects] == [
        plan.slides[0].title,
        "风险尚待核验。\n甲方收入为42万元。\n甲方收入为42万元。",
    ]
    for obj, decision in zip(authored.objects, scene["objects"], strict=True):
        assert obj.id == decision["id"]
        assert obj.role == decision["role"]
        assert obj.font_size == decision["font_size"]
        assert obj.frame.model_dump(mode="json") == decision["frame"]
        assert "text_refs" not in obj.model_dump(mode="json")
    assert authored.object_groups[0].model_dump(mode="json") == {
        "group_id": authored.object_groups[0].group_id,
        **scene["object_groups"][0],
    }
    accepted = validate_authored_page(
        authored, plan.slides[0], scaffold,
        source_text=grounded_source_text(request, brief), source_brief=brief,
    )
    assert accepted.scene_hash is not None
    contract = json.loads(calls[0]["user_prompt"])["fact_reference_transport"]
    assert contract["references"] == [
        {"ref": "storyboard.title", "text": plan.slides[0].title},
        {"ref": "required_facts[0]", "text": "甲方收入为42万元。"},
        {"ref": "required_facts[1]", "text": "风险尚待核验。"},
        {"ref": "immutable_facts[0]", "text": "甲方收入为42万元。"},
    ]


def _advertised_validator(monkeypatch):
    request, brief, plan, scaffold, scene = _case()
    calls = []

    async def transport(**kwargs):
        calls.append(kwargs)
        return SimpleNamespace(text=json.dumps(scene, ensure_ascii=False), model_label="test/model")

    monkeypatch.setattr("lantai.app.structured_utility.run_utility_text_task", transport)
    UtilityPresentationPageAuthor("agent-1").author(
        request, brief, plan, plan.slides[0], scaffold,
    )
    schema = calls[0]["structured_output_schema"]
    assert schema == json.loads(calls[0]["system_prompt"].split("JSON Schema:\n", 1)[1])
    Draft202012Validator.check_schema(schema)
    validator = Draft202012Validator(schema)
    assert validator.is_valid(scene)
    return validator, scene


def test_public_author_advertised_schema_rejects_two_text_representations(monkeypatch) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    ambiguous = deepcopy(scene)
    ambiguous["objects"][1]["text"] = "另一份正文"
    assert not validator.is_valid(ambiguous)


@pytest.mark.parametrize("representation", ["both", "neither"])
def test_public_author_reports_text_representation_rule_without_rejected_values(
    monkeypatch, representation,
) -> None:
    case = _case()
    private_marker = "synthetic-private-rejected-text-do-not-echo"
    obj = case[-1]["objects"][1]
    obj["id"] = private_marker
    if representation == "both":
        obj["text"] = private_marker
    else:
        obj.pop("text_refs")

    with pytest.raises(StructuredTaskValidationError) as captured:
        _author(monkeypatch, case)

    assert captured.value.validation_issues == (
        "authored_text_representation@$.objects[1].text",
    )
    assert private_marker not in str(captured.value)
    assert "甲方收入为42万元" not in str(captured.value)


def test_public_author_advertised_schema_rejects_rich_override_of_references(monkeypatch) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    overridden = deepcopy(scene)
    overridden["objects"][1]["runs"] = [{"text": "另一份正文", "font_size": 24}]
    assert not validator.is_valid(overridden)


def test_public_author_reports_rich_override_rule_without_rejected_values(monkeypatch) -> None:
    case = _case()
    private_marker = "synthetic-private-rejected-rich-text-do-not-echo"
    case[-1]["objects"][1].update({
        "id": private_marker,
        "runs": [{"text": private_marker, "font_size": 24}],
    })

    with pytest.raises(StructuredTaskValidationError) as captured:
        _author(monkeypatch, case)

    assert captured.value.validation_issues == (
        "authored_text_reference_runs@$.objects[1].text",
    )
    assert private_marker not in str(captured.value)
    assert "甲方收入为42万元" not in str(captured.value)


def test_public_author_advertised_schema_requires_object_kind(monkeypatch) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    unidentified = deepcopy(scene)
    unidentified["objects"][1].pop("kind")
    assert not validator.is_valid(unidentified)


def test_public_author_advertised_schema_rejects_partial_native_shape_margins(
    monkeypatch,
) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    scene["objects"].append({
        "kind": "shape", "id": "margin-shape", "shape_type": "rect",
        "frame": {"x": 0.8, "y": 6.2, "w": 2, "h": 0.4},
        "native_style": {"margin_top": 6},
    })
    case = (*_case()[:-1], scene)
    with pytest.raises(StructuredTaskValidationError):
        _author(monkeypatch, case)
    assert not validator.is_valid(scene)


@pytest.mark.parametrize("margin_states", product(("absent", "null", "number"), repeat=4))
def test_public_author_wire_and_parser_agree_on_atomic_point_margins(
    monkeypatch, margin_states,
) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    fields = ("margin_top", "margin_right", "margin_bottom", "margin_left")
    values = (0, 6.5, 144, 12)
    style = {"align": "center", "valign": "mid", "fill_none": True, "line_none": False}
    for field, state, value in zip(fields, margin_states, values, strict=True):
        if state != "absent":
            style[field] = value if state == "number" else None
    scene["objects"].append({
        "kind": "shape", "id": "atomic-margin-shape", "shape_type": "rect",
        "frame": {"x": 0.8, "y": 6.2, "w": 2, "h": 0.4},
        "native_style": style,
    })
    expected_valid = all(state == "number" for state in margin_states) or not any(
        state == "number" for state in margin_states
    )
    assert validator.is_valid(scene) is expected_valid
    case = (*_case()[:-1], scene)
    if not expected_valid:
        with pytest.raises(StructuredTaskValidationError):
            _author(monkeypatch, case)
        return
    authored_style = _author(monkeypatch, case).objects[-1].native_style
    for field, state, value in zip(fields, margin_states, values, strict=True):
        assert getattr(authored_style, field) == (value if state == "number" else None)
    for field in ("align", "valign", "fill_none", "line_none"):
        assert getattr(authored_style, field) == style[field]


@pytest.mark.parametrize("out_of_bounds", [-0.01, 144.01])
def test_public_author_atomic_margin_wire_keeps_original_point_bounds(
    monkeypatch, out_of_bounds,
) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    scene["objects"].append({
        "kind": "shape", "id": "bounded-margin-shape", "shape_type": "rect",
        "frame": {"x": 0.8, "y": 6.2, "w": 2, "h": 0.4},
        "native_style": {
            "margin_top": out_of_bounds, "margin_right": 6,
            "margin_bottom": 6, "margin_left": 6,
        },
    })
    assert not validator.is_valid(scene)
    with pytest.raises(StructuredTaskValidationError):
        _author(monkeypatch, (*_case()[:-1], scene))


def test_public_author_advertised_schema_preserves_title_font_floor(monkeypatch) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    unreadable = deepcopy(scene)
    unreadable["objects"][0]["font_size"] = 31
    assert not validator.is_valid(unreadable)


@pytest.mark.parametrize(("index", "minimum"), [(0, 32), (1, 17)])
def test_public_author_advertised_schema_preserves_rich_text_readability(
    monkeypatch, index, minimum,
) -> None:
    validator, scene = _advertised_validator(monkeypatch)
    literal = deepcopy(scene)
    obj = literal["objects"][index]
    obj.pop("text_refs")
    obj["text"] = "合成正文"
    obj["runs"] = [{"text": "合成正文", "font_size": minimum}]
    assert validator.is_valid(literal)
    obj["runs"][0]["font_size"] = minimum - 1
    assert not validator.is_valid(literal)


@pytest.mark.parametrize("reference", [
    "required_facts[99]", "deck.slides[1].required_facts[0]",
    "brief.required_facts[0]", "unknown-private-reference",
])
def test_unknown_or_cross_page_references_fail_closed_without_echoing_values(
    monkeypatch, reference,
) -> None:
    case = _case()
    case[-1]["objects"][1]["text_refs"] = [reference]
    with pytest.raises(PresentationPageAuthorContentError) as captured:
        _author(monkeypatch, case)
    assert captured.value.violation_code == "invalid_fact_reference"
    assert reference not in str(captured.value)


def test_omitted_reference_is_not_selected_or_filled_by_runtime(monkeypatch) -> None:
    case = _case()
    request, brief, plan, scaffold, scene = case
    scene["objects"][1]["text_refs"] = ["required_facts[0]"]
    authored = _author(monkeypatch, case)
    assert authored.objects[1].text == "甲方收入为42万元。"
    with pytest.raises(PresentationPageAuthorContentError) as captured:
        validate_authored_page(
            authored, plan.slides[0], scaffold,
            source_text=grounded_source_text(request, brief), source_brief=brief,
        )
    assert captured.value.violation_code == "required_visible_fact_omitted"
    assert captured.value.presentation_missing_fact_paths == ["required_facts[1]"]


@pytest.mark.parametrize("representation", [
    "both", "neither", "empty_refs", "rich_override", "nonstring_ref",
])
def test_reference_wire_rejects_ambiguous_or_invalid_representations(
    monkeypatch, representation,
) -> None:
    case = _case()
    obj = case[-1]["objects"][1]
    if representation == "both":
        obj["text"] = "另一份正文"
    elif representation == "neither":
        obj.pop("text_refs")
    elif representation == "empty_refs":
        obj["text_refs"] = []
    elif representation == "rich_override":
        obj["runs"] = [{"text": "另一份正文", "font_size": 24}]
    else:
        obj["text_refs"] = [42]
    with pytest.raises(StructuredTaskValidationError):
        _author(monkeypatch, case)


def test_existing_literal_and_rich_text_authoring_remains_supported(monkeypatch) -> None:
    case = _case()
    request, brief, plan, scaffold, scene = case
    texts = [plan.slides[0].title, "风险尚待核验。\n甲方收入为42万元。"]
    for obj, text in zip(scene["objects"], texts, strict=True):
        obj.pop("text_refs")
        obj["text"] = text
        obj["runs"] = [{"text": text, "font_size": obj["font_size"]}]
    authored = _author(monkeypatch, case)
    accepted = validate_authored_page(
        authored, plan.slides[0], scaffold,
        source_text=grounded_source_text(request, brief), source_brief=brief,
    )
    assert [obj.text for obj in accepted.objects] == texts
    assert [obj.runs[0].text for obj in accepted.objects] == texts


@pytest.mark.parametrize(("index", "font_size"), [(0, 31), (1, 16)])
def test_reference_text_still_obeys_authored_font_floor(monkeypatch, index, font_size) -> None:
    case = _case()
    case[-1]["objects"][index]["font_size"] = font_size
    with pytest.raises(StructuredTaskValidationError):
        _author(monkeypatch, case)


def test_selected_facts_compile_as_editable_native_pptx_text(monkeypatch, tmp_path) -> None:
    case = _case()
    request, brief, plan, scaffold, scene = case
    authored = _author(monkeypatch, case)
    output = tmp_path / "selected-facts-native.pptx"
    deck = build_deck_ir(request, plan).model_copy(update={"slides": [authored]})
    NativePptxCompiler().compile(deck, output)
    native = Presentation(output).slides[0]
    shapes = {shape.name: shape for shape in native.shapes}
    assert set(shapes) == {
        *(obj.id for obj in authored.objects),
        f"meta:{deck.id}:{authored.id}",
    }
    # The writer's established title/style floors apply to literal and
    # referenced text alike. Reference transport must not change those rules.
    literal_case = _case()
    for obj, selected in zip(literal_case[-1]["objects"], authored.objects, strict=True):
        obj.pop("text_refs")
        obj["text"] = selected.text
    literal = _author(monkeypatch, literal_case)
    baseline_output = tmp_path / "literal-native.pptx"
    NativePptxCompiler().compile(deck.model_copy(update={"slides": [literal]}), baseline_output)
    baseline_shapes = {shape.name: shape for shape in Presentation(baseline_output).slides[0].shapes}
    for obj in authored.objects:
        shape = shapes[obj.id]
        assert shape.has_text_frame
        assert shape.text == obj.text
        sizes = [run.font.size.pt
                 for paragraph in shape.text_frame.paragraphs for run in paragraph.runs]
        baseline_sizes = [run.font.size.pt
                          for paragraph in baseline_shapes[obj.id].text_frame.paragraphs
                          for run in paragraph.runs]
        assert sizes == baseline_sizes
        assert sizes and all((32 if obj.role == "title" else 17) <= size <= 72 for size in sizes)
        assert "required_facts[" not in shape.text


def test_invalid_fact_reference_retry_points_back_to_catalog_without_local_selection() -> None:
    from lantai.presentation.runtime import PresentationJobRuntime

    error = PresentationPageAuthorContentError(
        "page author selected a text reference outside its page catalog",
        violation_code="invalid_fact_reference",
    )
    directives = {}
    kind = PresentationJobRuntime._page_retry_description(None, error, directives, [], [])
    assert kind == "事实门禁未通过"
    assert "fact_reference_transport.references" in directives["fact_reference"]
    assert "text_refs" in directives["fact_reference"]
    assert "不会" in directives["fact_reference"]
