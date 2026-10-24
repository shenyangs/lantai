"""Notes-only edits are scoped, atomic and preserve every visible native byte."""

from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from lxml import etree as XML
from pptx import Presentation
from pydantic import ValidationError

from lantai.app.structured_utility import StructuredUtilityResult
from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.intelligence import UtilityPresentationObjectEditPlanner
from lantai.presentation.models import (
    DeckIR, DeckSlide, ObjectFrame, PresentationEditOperation,
    PresentationJobRequest, PresentationObjectEditPlan, TextObject,
)
from lantai.presentation.natural_edit import (
    apply_object_edit_plan, validate_requested_object_edit_scope,
    validate_required_object_edit_coverage,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.runtime import PresentationJobRuntime

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _deck():
    return DeckIR(id="synthetic-notes-edit", title="合成备注验证", slides=[
        DeckSlide(id="target", layout_id="freeform", objects=[
            TextObject(id="title", text="正文保持不变", role="title", font_size=32,
                       frame=ObjectFrame(x=0.6, y=0.5, w=12, h=0.8)),
        ]),
        DeckSlide(id="selected", layout_id="freeform", speaker_notes="保留原有备注", objects=[
            TextObject(id="second-title", text="保留附加页", role="title", font_size=32,
                       frame=ObjectFrame(x=0.6, y=0.5, w=12, h=0.8)),
        ]),
        DeckSlide(id="unselected", layout_id="freeform", speaker_notes="不可改动的原始引文", objects=[
            TextObject(id="third-title", text="保留未指定页", role="title", font_size=32,
                       frame=ObjectFrame(x=0.6, y=0.5, w=12, h=0.8)),
        ]),
    ])


def _plan(notes="只解释当前页的讲述边界。", *, target=None, additional=()):
    patch = {"speaker_notes": notes}
    if target is not None:
        patch["slide_id"] = target
    return PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "additional_slide_ids": list(additional),
        "speaker_notes_patches": [patch],
    })


def _operation(pages=("target",), content=("speaker_notes",)):
    return PresentationEditOperation(
        deck_id="synthetic-notes-edit", slide_ids=list(pages), object_ids=[],
        allowed_content_types=set(content), instruction="仅修订已指定页的讲解备注",
    )


def _changed(deck, plan=None):
    return apply_object_edit_plan(deck, plan or _plan())[0]


def _compile(tmp_path):
    source = tmp_path / "source.pptx"
    deck = _deck()
    NativePptxCompiler().compile(deck, source)
    return deck, source


def _parts(path):
    with ZipFile(path) as package:
        return {entry.filename: package.read(entry.filename) for entry in package.infolist()}


def _rewrite_part(path, name, transform):
    with ZipFile(path) as package:
        entries = [(entry, package.read(entry.filename)) for entry in package.infolist()]
    with ZipFile(path, "w") as package:
        for entry, payload in entries:
            package.writestr(entry, transform(payload) if entry.filename == name else payload)


def test_notes_only_plan_changes_no_object_or_source_fact():
    deck = _deck()
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, _plan())
    assert ids == []
    assert changed.slides[0].speaker_notes == "只解释当前页的讲述边界。"
    assert changed.slides[0].revision == deck.slides[0].revision + 1
    assert changed.slides[0].scene_hash
    assert changed.slides[0].model_dump(exclude={"speaker_notes", "revision", "scene_hash"}) == (
        deck.slides[0].model_dump(exclude={"speaker_notes", "revision", "scene_hash"})
    )
    assert changed.slides[1:] == deck.slides[1:]
    assert changed.model_dump(exclude={"slides"}) == deck.model_dump(exclude={"slides"})
    validate_required_object_edit_coverage(deck, changed, ["target"])
    assert deck.model_dump_json() == before


def test_notes_only_patch_to_explicit_additional_page_does_not_touch_primary():
    deck = _deck()
    plan = _plan("仅解释授权附加页。", target="selected", additional=["selected"])
    validate_requested_object_edit_scope(deck, "target", ["selected"], plan=plan)
    changed, ids = apply_object_edit_plan(deck, plan)
    assert ids == []
    assert changed.slides[0] == deck.slides[0]
    assert changed.slides[1].speaker_notes == "仅解释授权附加页。"
    assert changed.slides[2] == deck.slides[2]
    validate_required_object_edit_coverage(deck, changed, ["selected"])


def test_notes_target_cannot_escape_declared_or_requested_scope():
    deck = _deck()
    before = deck.model_dump_json()
    plan = _plan(target="unselected")
    with pytest.raises(PermissionError, match="notes scope"):
        validate_requested_object_edit_scope(deck, "target", plan=plan)
    with pytest.raises(PermissionError, match="notes scope"):
        apply_object_edit_plan(deck, plan)
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("patches", [
    [{"slide_id": "", "speaker_notes": "明确目标"}],
    [{"speaker_notes": "甲"}, {"slide_id": "target", "speaker_notes": "乙"}],
    [{"slide_id": f"page-{i}", "speaker_notes": "界限说明"} for i in range(21)],
])
def test_notes_patches_reject_empty_duplicate_or_over_cap_targets(patches):
    with pytest.raises(ValidationError):
        PresentationObjectEditPlan.model_validate({
            "slide_id": "target", "speaker_notes_patches": patches,
        })


@pytest.mark.parametrize("other", [
    {"actions": [{"kind": "replace_text", "object_id": "title", "text": "不得混合"}]},
    {"intent_patch": {"key_message": "不得混合"}},
    {"intent_patches": [{"key_message": "不得混合"}]},
])
def test_notes_only_v1_cannot_hide_native_or_intent_edits(other):
    with pytest.raises(ValidationError, match="notes-only"):
        PresentationObjectEditPlan.model_validate({
            "slide_id": "target", "speaker_notes_patches": [{"speaker_notes": "讲述边界"}],
            **other,
        })


def test_old_object_plan_schema_does_not_require_notes():
    plan = PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "actions": [
            {"kind": "replace_text", "object_id": "title", "text": "明确正文修订"},
        ],
    })
    assert plan.speaker_notes_patches == []
    assert plan.resolved_speaker_notes_patches() == []


def test_same_note_is_not_material_coverage():
    deck = _deck()
    with pytest.raises(ValueError, match="no material change"):
        apply_object_edit_plan(deck, _plan(""))
    with pytest.raises(ValueError, match="required_slide_unchanged"):
        validate_required_object_edit_coverage(deck, deck, ["target"])


@pytest.mark.parametrize("notes", ["杜撰 99% 的结论", "另有 42 项", "新增结果 3.5"])
def test_new_notes_preserve_existing_numeric_free_authoring_contract(notes):
    deck = _deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="numeric-free"):
        apply_object_edit_plan(deck, _plan(notes))
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("notes", ["讲述这页的证据边界。", "第一段。\n第二段。", ""])
def test_native_notes_transaction_changes_only_one_notes_part_and_can_undo(tmp_path, notes):
    deck, source = _compile(tmp_path)
    plan = _plan(notes, target="selected", additional=["selected"])
    changed = _changed(deck, plan)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    before = _parts(source)
    preview = service.prepare_speaker_notes(source, deck, changed, _operation(("target", "selected")))
    candidate = _parts(preview.candidate_path)
    assert before.keys() == candidate.keys()
    changed_parts = [name for name in before if before[name] != candidate[name]]
    assert len(changed_parts) == 1
    assert changed_parts[0].startswith("ppt/notesSlides/") and changed_parts[0].endswith(".xml")
    assert preview.changed_object_ids == [] and preview.verified
    assert _parts(source) == before
    reopened = Presentation(preview.candidate_path)
    assert [slide.notes_slide.notes_text_frame.text for slide in reopened.slides] == [
        "", notes, "不可改动的原始引文",
    ]
    receipt = service.commit(preview)
    assert _parts(source) == candidate
    service.rollback(receipt)
    assert _parts(source) == before


@pytest.mark.parametrize("field", ["title", "objects", "speaker_notes"])
def test_native_notes_transaction_rejects_other_deck_or_page_changes(tmp_path, field):
    deck, source = _compile(tmp_path)
    changed = _changed(deck).model_copy(deep=True)
    if field == "title":
        changed = changed.model_copy(update={"title": "越权整套变更"})
    elif field == "objects":
        changed.slides[0].objects[0].text = "越权正文变更"
    else:
        changed.slides[2].speaker_notes = "越权未指定页备注"
    before = source.read_bytes()
    service = PresentationEditTransactionService(tmp_path / "transactions")
    with pytest.raises(PermissionError):
        service.prepare_speaker_notes(source, deck, changed, _operation())
    assert source.read_bytes() == before


def test_native_notes_transaction_requires_notes_only_content_authority(tmp_path):
    deck, source = _compile(tmp_path)
    before = source.read_bytes()
    service = PresentationEditTransactionService(tmp_path / "transactions")
    with pytest.raises(PermissionError, match="notes-only"):
        service.prepare_speaker_notes(source, deck, _changed(deck), _operation(content=("text",)))
    assert source.read_bytes() == before


def test_native_notes_transaction_refuses_stale_ir_note_without_partial_commit(tmp_path):
    deck, source = _compile(tmp_path)
    deck.slides[0].speaker_notes = "不是当前原生文件的备注"
    before = source.read_bytes()
    service = PresentationEditTransactionService(tmp_path / "transactions")
    with pytest.raises(ValueError, match="stale"):
        service.prepare_speaker_notes(source, deck, _changed(deck), _operation())
    assert source.read_bytes() == before


def test_native_notes_batch_preserves_every_unselected_package_part(tmp_path):
    deck, source = _compile(tmp_path)
    plan = PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "additional_slide_ids": ["selected"],
        "speaker_notes_patches": [
            {"speaker_notes": "说明正文所列的证据边界。"},
            {"slide_id": "selected", "speaker_notes": "只解释附加页的讲述顺序。"},
        ],
    })
    changed = _changed(deck, plan)
    before = _parts(source)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    preview = service.prepare_speaker_notes(source, deck, changed, _operation(("target", "selected")))
    after = _parts(preview.candidate_path)
    assert before.keys() == after.keys()
    assert set(name for name in before if before[name] != after[name]) == {
        "ppt/notesSlides/notesSlide1.xml", "ppt/notesSlides/notesSlide2.xml",
    }
    assert _parts(source) == before
    assert [slide.notes_slide.notes_text_frame.text for slide in Presentation(preview.candidate_path).slides] == [
        "说明正文所列的证据边界。", "只解释附加页的讲述顺序。", "不可改动的原始引文",
    ]


def test_later_invalid_note_in_batch_leaves_no_partial_candidate(tmp_path):
    deck, source = _compile(tmp_path)
    stale = deck.model_copy(deep=True)
    stale.slides[1].speaker_notes = "与实际载体不一致的旧备注"
    plan = PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "additional_slide_ids": ["selected"],
        "speaker_notes_patches": [
            {"speaker_notes": "有效的主页面说明。"},
            {"slide_id": "selected", "speaker_notes": "附加页说明。"},
        ],
    })
    operation = _operation(("target", "selected"))
    before = source.read_bytes()
    service = PresentationEditTransactionService(tmp_path / "transactions")
    with pytest.raises(ValueError, match="stale"):
        service.prepare_speaker_notes(source, stale, _changed(stale, plan), operation)
    assert source.read_bytes() == before
    assert not (service.root / operation.operation_id / "candidate.pptx").exists()


@pytest.mark.parametrize("defect", ["missing", "duplicate", "external", "wrong_backlink", "mixed_format"])
def test_unsupported_native_note_carrier_fails_closed(tmp_path, defect):
    deck, source = _compile(tmp_path)
    name = "ppt/slides/_rels/slide1.xml.rels"
    def corrupt(payload):
        root = XML.fromstring(payload)
        notes = next(rel for rel in root if rel.get("Type", "").endswith("/notesSlide"))
        if defect == "missing":
            root.remove(notes)
        elif defect == "duplicate":
            added = deepcopy(notes)
            added.set("Id", "synthetic-duplicate")
            root.append(added)
        elif defect == "external":
            notes.set("TargetMode", "External")
        return XML.tostring(root)
    if defect in {"missing", "duplicate", "external"}:
        _rewrite_part(source, name, corrupt)
    elif defect == "wrong_backlink":
        _rewrite_part(source, "ppt/notesSlides/_rels/notesSlide1.xml.rels",
                      lambda payload: payload.replace(b"slide1.xml", b"slide2.xml"))
    else:
        def mixed(payload):
            root = XML.fromstring(payload)
            body = next(sp.find(f"{{{P}}}txBody") for sp in root.iter(f"{{{P}}}sp")
                        if sp.find(f"{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph").get("type") == "body")
            paragraph = body.find(f"{{{A}}}p")
            run = XML.SubElement(paragraph, f"{{{A}}}r")
            XML.SubElement(run, f"{{{A}}}rPr", b="1")
            XML.SubElement(run, f"{{{A}}}t").text = ""
            return XML.tostring(root)
        _rewrite_part(source, "ppt/notesSlides/notesSlide1.xml", mixed)
    before = source.read_bytes()
    service = PresentationEditTransactionService(tmp_path / "transactions")
    with pytest.raises(ValueError, match="notes"):
        service.prepare_speaker_notes(source, deck, _changed(deck), _operation())
    assert source.read_bytes() == before


def test_planner_advertises_explicit_notes_only_root_patch(monkeypatch):
    captured = {}
    expected = _plan()
    async def generate(**kwargs):
        captured.update(kwargs)
        return StructuredUtilityResult(value=expected, model_label="synthetic-notes/1", repaired=False)
    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", generate)
    assert UtilityPresentationObjectEditPlanner("synthetic-agent").plan(
        _deck(), "target", "仅补写当前页的讲解备注，不改正文",
    ) == expected
    payload = json.loads(captured["user_prompt"])
    assert payload["speaker_notes_edit_contract"]["notes_only"] is True
    assert payload["speaker_notes_edit_contract"]["visible_content_changes_forbidden"] is True
    assert payload["speaker_notes_edit_contract"]["new_notes_numeric_free"] is True
    assert "speaker_notes_patches" in captured["system_prompt"]
    assert payload["slide"]["speaker_notes"] == ""


def test_runtime_notes_candidate_requires_review_without_composing_visible_pages(tmp_path, monkeypatch):
    runtime = PresentationJobRuntime(tmp_path / "jobs", object_edit_planner=SimpleNamespace(
        plan=lambda *a, **k: _plan(),
    ))
    record = runtime.service.create_job(PresentationJobRequest(
        conversation_id="synthetic-notes-owner", user_request="只补讲解备注", slide_count=3,
    ))
    deck = _deck()
    def forbidden(*args, **kwargs):
        pytest.fail("notes-only edits must not normalize or compose visible pages")
    monkeypatch.setattr("lantai.presentation.runtime.compose_deck_ir_with_design_context", forbidden)
    monkeypatch.setattr("lantai.presentation.runtime.prepare_deck_for_composition", forbidden)
    try:
        result = runtime._prepare_edit_candidate(
            record.job_id, record, deck, deck, None, "target", "仅补写备注",
            required_slide_ids=["target"],
        )
        assert result[3] is True and result[4] is False
        assert result[7] == ["target"]
        assert result[8] == _changed(deck)
        assert result[9].allowed_content_types == {"speaker_notes"}
        assert result[9].object_ids == []
    finally:
        runtime.shutdown()


def test_runtime_notes_preview_uses_native_notes_transaction_not_compiler(tmp_path, monkeypatch):
    deck, source = _compile(tmp_path)
    runtime = PresentationJobRuntime(tmp_path / "jobs")
    record = runtime.service.create_job(PresentationJobRequest(
        conversation_id="synthetic-notes-owner", user_request="只补讲解备注", slide_count=3,
    ))
    def forbidden(*args, **kwargs):
        pytest.fail("notes-only preview must not compile visible slide graphs")
    monkeypatch.setattr(runtime.service, "compiler_for", forbidden)
    service = PresentationEditTransactionService(tmp_path / "transactions")
    try:
        preview = runtime._prepare_edit_preview(
            record, service, source, deck, _changed(deck), _operation(), True, ["target"], None,
        )
        assert preview.verified and preview.changed_object_ids == []
        assert _parts(source) != _parts(preview.candidate_path)
    finally:
        runtime.shutdown()
