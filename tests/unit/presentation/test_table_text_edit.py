"""Public table-text edits keep native identity and caller-owned page scope."""

from __future__ import annotations

import json
from copy import deepcopy

from lxml import etree as XML

import pytest
from pydantic import ValidationError
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

from lantai.app.structured_utility import StructuredUtilityResult
from lantai.presentation.intelligence import UtilityPresentationObjectEditPlanner
from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    DeckIR, DeckSlide, ObjectFrame, ObjectGroup, ObjectRelationship,
    PresentationEditOperation, PresentationObjectEditPlan, RelationshipType, RichTextRun, TableCellStyle,
    TableObject, TextObject,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _deck():
    table = TableObject(
        id="evidence-table", headers=["问题", "来源"],
        rows=[["原判断", "合成来源"], ["保留判断", "保留限定"]],
        frame=ObjectFrame(x=1, y=2, w=10, h=3), slot="evidence",
    )
    return DeckIR(
        id="synthetic-table-edit", title="原生表格文字编辑",
        source_refs=["synthetic://table-edit"], required_facts=["保留限定"],
        immutable_data=["保留限定"],
        slides=[
            DeckSlide(
                id="target", layout_id="freeform", speaker_notes="保留备注",
                objects=[TextObject(id="title", text="保留限定", role="title",
                                    frame=ObjectFrame(x=1, y=.4, w=10, h=1)), table],
                object_groups=[ObjectGroup(group_id="group", role="evidence",
                                           object_ids=["title", table.id])],
                relationships=[ObjectRelationship(
                    relationship_id="relation", relationship_type=RelationshipType.PROCESS,
                    source_object_ids=["title"], target_object_ids=[table.id], label="保留关系",
                )],
            ),
            DeckSlide(id="selected", layout_id="freeform", objects=[table.model_copy(update={"id": "selected-table"})]),
            DeckSlide(id="unselected", layout_id="freeform", objects=[table.model_copy(update={"id": "unselected-table"})]),
        ],
    )


def _edit(*, object_id="evidence-table", additional=(), headers=None, rows=None, actions=None):
    return PresentationObjectEditPlan.model_validate({
        "slide_id": "target", "additional_slide_ids": list(additional),
        "actions": actions if actions is not None else [{
            "kind": "update_table", "object_id": object_id,
            "headers": ["问题", "来源"] if headers is None else headers,
            "rows": [["新判断", "合成来源"], ["保留判断", "保留限定"]] if rows is None else rows,
        }],
    })


def test_explicit_table_text_edit_keeps_identity_geometry_evidence_and_other_pages():
    deck = _deck()
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, _edit())
    assert ids == ["evidence-table"]
    assert changed.slides[0].objects[1].model_dump() == {
        **deck.slides[0].objects[1].model_dump(),
        "rows": [["新判断", "合成来源"], ["保留判断", "保留限定"]],
    }
    assert changed.slides[0].objects[0] == deck.slides[0].objects[0]
    assert changed.slides[0].object_groups == deck.slides[0].object_groups
    assert changed.slides[0].relationships == deck.slides[0].relationships
    assert changed.slides[0].speaker_notes == deck.slides[0].speaker_notes
    assert changed.slides[0].revision == deck.slides[0].revision + 1
    assert changed.slides[0].scene_hash
    assert changed.slides[1:] == deck.slides[1:]
    assert changed.source_refs == deck.source_refs
    assert changed.required_facts == deck.required_facts
    assert changed.immutable_data == deck.immutable_data
    assert deck.model_dump_json() == before


def _styled_deck():
    deck = _deck()
    table = deck.slides[0].objects[1]
    table.cell_styles = [[TableCellStyle(
        align="left", valign="mid", font_size=18, fill_color="D4DDE1",
        runs=[RichTextRun(text=value, color="315F9A", font_size=18, bold=False, italic=False)],
    ) for value in row] for row in [table.headers, *table.rows]]
    return deck


def test_table_cell_text_updates_its_native_run_and_keeps_all_other_cell_styles():
    deck = _styled_deck()
    before = deck.model_dump_json()
    table = deck.slides[0].objects[1]
    changed, ids = apply_object_edit_plan(deck, _edit())
    updated = changed.slides[0].objects[1]
    assert ids == [table.id]
    assert updated.cell_styles[1][0].model_dump() == {
        **table.cell_styles[1][0].model_dump(),
        "runs": [{**table.cell_styles[1][0].runs[0].model_dump(), "text": "新判断"}],
    }
    assert [["".join(run.text for run in style.runs) for style in row]
            for row in updated.cell_styles] == [updated.headers, *updated.rows]
    for ri, row in enumerate(table.cell_styles):
        for ci, style in enumerate(row):
            if (ri, ci) != (1, 0):
                assert updated.cell_styles[ri][ci] == style
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("styled", [False, True])
@pytest.mark.parametrize("headers,rows", [
    (["问题"], [["新判断"], ["保留判断"]]),
    (["问题", "来源"], [["新判断", "合成来源"]]),
    (["问题", "来源"], [["新判断", "合成来源"], ["保留判断", "保留限定"], ["新增", "未经授权"]]),
])
def test_table_text_action_rejects_grid_changes_without_mutating_anything(styled, headers, rows):
    deck = _styled_deck() if styled else _deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="existing table grid"):
        apply_object_edit_plan(deck, _edit(headers=headers, rows=rows))
    assert deck.model_dump_json() == before


def test_planner_exposes_only_native_tables_in_explicit_primary_and_additional_scope(monkeypatch):
    calls = []

    async def utility(**kwargs):
        calls.append(kwargs)
        return StructuredUtilityResult(value=_edit(object_id="selected-table", additional=["selected"]),
                                       model_label="synthetic/synthetic", repaired=False)

    monkeypatch.setattr("lantai.presentation.intelligence.run_structured_utility_task", utility)
    UtilityPresentationObjectEditPlanner("synthetic-agent").plan(
        _deck(), "target", "仅修订选定附加页表格文字", additional_slide_ids=["selected"],
    )
    payload = json.loads(calls[0]["user_prompt"])
    assert payload["editable_object_ids_by_action"]["update_table"] == ["evidence-table"]
    assert payload["editable_object_ids_by_action"]["replace_text"] == ["title"]
    assert payload["additional_slides"][0]["editable_object_ids_by_action"]["update_table"] == ["selected-table"]
    assert "update_table" in calls[0]["system_prompt"]
    assert "unselected-table" not in calls[0]["user_prompt"]
    assert len(calls) == 1
    assert calls[0]["timeout_seconds"] == 90


def test_table_action_rejects_ragged_rows_in_its_typed_public_schema():
    with pytest.raises(ValidationError, match="every table row must match header width"):
        _edit(rows=[["少一个单元格"], ["保留判断", "保留限定"]])


def test_table_text_action_edits_only_the_explicit_additional_page():
    deck = _styled_deck()
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(
        deck, _edit(object_id="selected-table", additional=["selected"]),
    )
    assert ids == ["selected-table"]
    assert changed.slides[0] == deck.slides[0]
    assert changed.slides[2] == deck.slides[2]
    assert changed.slides[1].objects[0].rows[0][0] == "新判断"
    assert changed.slides[1].objects[0].id == "selected-table"
    assert changed.slides[1].revision == deck.slides[1].revision + 1
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("object_id,additional", [
    ("selected-table", []), ("unselected-table", ["selected"]), ("missing-table", []),
])
def test_table_text_action_cannot_expand_its_caller_owned_page_scope(object_id, additional):
    deck = _styled_deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="unknown object id"):
        apply_object_edit_plan(deck, _edit(object_id=object_id, additional=additional))
    assert deck.model_dump_json() == before


def test_wrong_kind_after_a_valid_cross_page_table_edit_rolls_back_every_page():
    deck = _styled_deck()
    deck.slides[1].objects.append(TextObject(id="selected-title", role="body", text="保留附加页正文"))
    before = deck.model_dump_json()
    actions = [*_edit().model_dump()["actions"],
               *_edit(object_id="selected-title").model_dump()["actions"]]
    with pytest.raises(ValueError, match="not an editable table"):
        apply_object_edit_plan(deck, _edit(additional=["selected"], actions=actions))
    assert deck.model_dump_json() == before


def test_unchanged_table_text_does_not_publish_a_fake_edit():
    deck = _styled_deck()
    before = deck.model_dump_json()
    table = deck.slides[0].objects[1]
    with pytest.raises(ValueError, match="no material change"):
        apply_object_edit_plan(deck, _edit(headers=table.headers, rows=table.rows))
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("extra", [
    {"frame": {"x": 2, "y": 2, "w": 3, "h": 3}},
    {"cell_styles": []}, {"slot": "other"}, {"source_refs": ["invented-source"]},
])
def test_table_text_schema_cannot_smuggle_geometry_style_or_evidence_changes(extra):
    action = _edit().model_dump()["actions"][0]
    with pytest.raises(ValidationError, match="Extra inputs are not permitted"):
        _edit(actions=[{**action, **extra}])


def _mixed_cell_deck():
    deck = _styled_deck()
    deck.slides[0].objects[1].cell_styles[1][0].runs = [
        RichTextRun(text="原", color="315F9A", font_size=18, bold=True),
        RichTextRun(text="判断", color="315F9A", font_size=18, bold=False),
    ]
    return deck


def test_changed_mixed_format_cell_fails_closed_instead_of_dropping_its_styles():
    deck = _mixed_cell_deck()
    before = deck.model_dump_json()
    with pytest.raises(ValueError, match="cannot preserve mixed native cell styles"):
        apply_object_edit_plan(deck, _edit())
    assert deck.model_dump_json() == before


def test_unchanged_mixed_format_cell_survives_a_different_cell_text_edit():
    deck = _mixed_cell_deck()
    before = deck.model_dump_json()
    changed, ids = apply_object_edit_plan(deck, _edit(
        rows=[["原判断", "新来源"], ["保留判断", "保留限定"]],
    ))
    assert ids == ["evidence-table"]
    assert changed.slides[0].objects[1].cell_styles[1][0] == deck.slides[0].objects[1].cell_styles[1][0]
    assert changed.slides[0].objects[1].cell_styles[1][1].runs[0].text == "新来源"
    assert deck.model_dump_json() == before


@pytest.mark.parametrize("text", ["新判断", "新判断\n第二行", "", "  新判断  "])
def test_table_text_survives_native_export_reimport_edit_and_second_reimport(tmp_path, text):
    deck = _styled_deck()
    writer = NativePptxCompiler()
    original, output = tmp_path / "original.pptx", tmp_path / "edited.pptx"
    writer.compile(deck, original)
    before = original.read_bytes()
    imported = reconcile_pptx(deck, original, assets_dir=tmp_path / "assets-first")
    table = next(obj for obj in imported.slides[0].objects if obj.kind == "table")
    changed, ids = apply_object_edit_plan(imported, _edit(
        rows=[[text, "合成来源"], ["保留判断", "保留限定"]],
    ))
    assert ids == [table.id]
    writer.compile(changed, output)
    second = reconcile_pptx(changed, output, assets_dir=tmp_path / "assets-second")
    reimported = next(obj for obj in second.slides[0].objects if obj.kind == "table")
    assert reimported.rows[0][0].replace("\v", "\n") == text
    assert [obj.id for obj in second.slides[0].objects] == [obj.id for obj in imported.slides[0].objects]
    assert second.slides[0].object_groups == imported.slides[0].object_groups
    assert second.slides[0].relationships == imported.slides[0].relationships
    assert second.slides[0].speaker_notes == imported.slides[0].speaker_notes
    assert second.required_facts == deck.required_facts
    assert second.immutable_data == deck.immutable_data
    assert second.source_refs == deck.source_refs
    old_shape = next(shape for shape in Presentation(str(original)).slides[0].shapes if shape.has_table)
    new_shapes = [shape for shape in Presentation(str(output)).slides[0].shapes if shape.has_table]
    assert len(new_shapes) == 1
    new_shape = new_shapes[0]
    assert (new_shape.name, new_shape.left, new_shape.top, new_shape.width, new_shape.height) == (
        old_shape.name, old_shape.left, old_shape.top, old_shape.width, old_shape.height,
    )
    assert [[cell.text.replace("\v", "\n") for cell in row.cells] for row in new_shape.table.rows] == [
        ["问题", "来源"], [text, "合成来源"], ["保留判断", "保留限定"],
    ]
    for row in new_shape.table.rows:
        for cell in row.cells:
            assert cell.vertical_anchor == MSO_ANCHOR.MIDDLE
            assert cell.fill.fore_color.rgb == old_shape.table.cell(0, 0).fill.fore_color.rgb
            for paragraph in cell.text_frame.paragraphs:
                assert paragraph.alignment == PP_ALIGN.LEFT
                assert all(run.font.size.pt == 18 and str(run.font.color.rgb) == "315F9A"
                           and run.font.bold is not True and run.font.italic is not True
                           for run in paragraph.runs)
    assert original.read_bytes() == before


@pytest.mark.parametrize("text", ["新判断", "新判断\n第二行", "", "  新判断  "])
def test_public_table_text_transaction_keeps_the_exact_original_cell_carrier(tmp_path, text):
    deck = _styled_deck()
    # This native fixture avoids the separately observed backend table-ID
    # collision, so that ambiguous IDs cannot mask a cell-style regression.
    deck.slides[0].objects.reverse()
    writer = NativePptxCompiler()
    source = tmp_path / "external.pptx"
    writer.compile(deck, source)
    external = Presentation(source)
    table = next(shape for shape in external.slides[0].shapes if shape.has_table)
    for row in table.table.rows:
        for cell in row.cells:
            cell._tc.tcPr.set("anchor", "mid")
            cell._tc.tcPr.set("marL", "12345")
            for paragraph in cell.text_frame.paragraphs:
                paragraph._p.get_or_add_pPr().attrib.pop("algn", None)
    external.save(source)
    source_bytes = source.read_bytes()
    original = next(shape for shape in Presentation(source).slides[0].shapes if shape.has_table)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, _edit(
        rows=[[text, "合成来源"], ["保留判断", "保留限定"]],
    ))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, imported, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["target"], object_ids=ids,
            allowed_content_types={"text", "table"}, instruction="只改表格文字，原生样式保持原样",
        ), writer,
    )
    actual = next(shape for shape in Presentation(preview.candidate_path).slides[0].shapes if shape.has_table)
    expected = deepcopy(original._element)
    expected.xpath(".//a:tr")[1].xpath("./a:tc")[0].xpath(".//a:t")[0].text = text
    assert XML.tostring(actual._element, method="c14n", exclusive=True) == XML.tostring(
        expected, method="c14n", exclusive=True,
    )
    assert source.read_bytes() == source_bytes
    assert actual.table.cell(1, 0).text == text
    reimported = reconcile_pptx(changed, preview.candidate_path, assets_dir=tmp_path / "reimport")
    assert next(obj for obj in reimported.slides[0].objects if obj.kind == "table").rows[0][0] == text


@pytest.mark.parametrize("mutation,message", [
    ("changed-text", "stale cell text"),
    ("removed-row", "stale grid"),
    ("removed-cell", "stale grid"),
    ("split-run", "ambiguous cell formatting"),
])
def test_public_table_transaction_rejects_stale_or_ambiguous_native_cells(tmp_path, mutation, message):
    deck = _styled_deck()
    deck.slides[0].objects.reverse()
    writer = NativePptxCompiler()
    source = tmp_path / "external.pptx"
    writer.compile(deck, source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, _edit())
    external = Presentation(source)
    table = next(shape for shape in external.slides[0].shapes if shape.has_table)
    if mutation == "changed-text":
        table.table.cell(1, 0).text = "另一个版本"
    elif mutation == "removed-row":
        table._element.xpath(".//a:tbl")[0].remove(table._element.xpath(".//a:tr")[-1])
    elif mutation == "removed-cell":
        row = table._element.xpath(".//a:tr")[1]
        row.remove(row.xpath("./a:tc")[-1])
    else:
        paragraph = table.table.cell(1, 0).text_frame.paragraphs[0]._p
        run = paragraph.xpath("./a:r")[0]
        extra = deepcopy(run)
        run.xpath("./a:t")[0].text = "原"
        extra.xpath("./a:t")[0].text = "判断"
        paragraph.append(extra)
    external.save(source)
    source_bytes = source.read_bytes()
    with pytest.raises(ValueError, match=message):
        PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
            source, imported, changed, PresentationEditOperation(
                deck_id=deck.id, slide_ids=["target"], object_ids=ids,
                allowed_content_types={"text", "table"}, instruction="只改选定表格文字",
            ), writer,
        )
    assert source.read_bytes() == source_bytes
    assert not list((tmp_path / "edits").rglob("preview.json"))
    assert not list((tmp_path / "edits").rglob("candidate.pptx"))


def test_public_multi_page_table_transaction_never_publishes_a_partial_native_edit(tmp_path):
    deck = _styled_deck()
    deck.slides[0].objects.reverse()
    deck.slides[1].objects.insert(0, TextObject(
        id="selected-title", text="保留附加页正文", role="title",
        frame=ObjectFrame(x=1, y=.4, w=10, h=1),
    ))
    writer = NativePptxCompiler()
    source = tmp_path / "external.pptx"
    writer.compile(deck, source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, _edit(
        additional=["selected"], actions=[
            *_edit().model_dump()["actions"],
            *_edit(object_id="selected-table").model_dump()["actions"],
        ],
    ))
    external = Presentation(source)
    next(shape for shape in external.slides[1].shapes if shape.has_table).table.cell(1, 0).text = "新版判断"
    external.save(source)
    source_bytes = source.read_bytes()
    with pytest.raises(ValueError, match="stale cell text"):
        PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
            source, imported, changed, PresentationEditOperation(
                deck_id=deck.id, slide_ids=["target", "selected"], object_ids=ids,
                allowed_content_types={"text", "table"}, instruction="两页表格必须原子编辑",
            ), writer, compiled_slide_ids=["target", "selected"],
        )
    assert source.read_bytes() == source_bytes
    assert not list((tmp_path / "edits").rglob("preview.json"))
    assert not list((tmp_path / "edits").rglob("candidate.pptx"))
