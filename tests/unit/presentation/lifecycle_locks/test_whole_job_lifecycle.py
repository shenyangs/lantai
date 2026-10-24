# -*- coding: utf-8 -*-
"""Independent LTA-007/020 SOURCE locks; synthetic material and no external
process.
"""
from __future__ import annotations

import asyncio
import threading
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest

from lantai.presentation import intelligence as intelligence_module
from lantai.presentation import jobs as jobs_module
from lantai.presentation import service as service_module
from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.intelligence import (
    UtilityPresentationPageAuthor,
    UtilityPresentationSemanticPlanner,
    UtilityPresentationVisualCritic,
)
from lantai.presentation.jobs import (
    PresentationExecutionLimits,
    PresentationJobInterruptedError,
    PresentationJobStore,
)
from lantai.presentation.models import (
    CompositionAuthority,
    CompositionEngine,
    DeckIR,
    DeckSlide,
    InputMaterial,
    JobStage,
    JobStatus,
    PlanningMode,
    PresentationJobRequest,
    TextObject,
    VisualCriticResult,
    VisualCriticSlideInput,
)
from lantai.presentation.planner import build_deterministic_fixture_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import file_sha256
from lantai.presentation.render import (
    PresentationRenderUnavailable,
    RenderResult,
)
from lantai.presentation.runtime import PresentationJobRuntime
from lantai.presentation.service import PresentationService
from lantai.presentation.versioning import PresentationVersionStore

WAIT = 1.5


def request(**updates):
    result = PresentationJobRequest(
        conversation_id="independent-lifecycle-synthetic",
        user_request="制作合成材料的三页复盘，保留原始事实",
        slide_count=3,
        timeout_seconds=30,
        cost_budget_units=1000,
        visual_verification_required=False,
        vision_critique_required=False,
        materials=[
            InputMaterial(
                name="synthetic.md",
                content=(
                    "# 合成项目复盘\n\n## 判断\n合成项目交付完成。\n\n"
                    "## 证据\n流程具有明确阶段。\n\n## 行动\n下一阶段继续核验。"
                ),
            ),
        ],
        planning_mode=PlanningMode.MODEL,
        composition_engine=CompositionEngine.LEGACY_PYTHON_PPTX,
    )
    return result.model_copy(update=updates)


def deck():
    return DeckIR(
        id="synthetic-deck",
        title="合成验证",
        slides=[
            DeckSlide(
                id="synthetic-slide",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id="synthetic-title",
                        role="title",
                        slot="title",
                        text="合成验证",
                    ),
                ],
            ),
        ],
    )


def authored(scaffold):
    return scaffold.model_copy(
        update={
            "composition_authority": CompositionAuthority.AUTHOR,
            "scene_hash": None,
            "speaker_notes": "",
            "objects": [
                obj
                for obj in scaffold.objects
                if obj.slot not in {"relationship-marker", "page-marker"}
            ],
            "object_groups": [],
            "relationships": [],
        },
    )


class FixtureAuthor:
    # Callback keeps the public request keyword in its signature.
    # pylint: disable-next=redefined-outer-name
    def author(self, request, brief, plan, planned, scaffold, **kwargs):
        del brief, kwargs, plan, planned, request
        return authored(scaffold)


class FixturePlanner:
    # Callback keeps the public request keyword in its signature.
    # pylint: disable-next=redefined-outer-name
    def plan(self, request, brief, **kwargs):
        del kwargs
        return build_deterministic_fixture_plan(request, brief)


class ControlledClock(datetime):
    current = datetime(2035, 1, 1, tzinfo=timezone.utc)

    @classmethod
    def now(cls, tz=None):
        return (
            cls.current if tz is not None else cls.current.replace(tzinfo=None)
        )


@pytest.fixture(name="clock")
def clock_fixture(monkeypatch):
    ControlledClock.current = datetime.now(timezone.utc)
    monkeypatch.setattr(jobs_module, "datetime", ControlledClock)
    return ControlledClock


class Worker:
    def __init__(self, action):
        self.errors = []
        self.finished = threading.Event()

        def run():
            try:
                action()
            except BaseException as exc:
                self.errors.append(exc)
            finally:
                self.finished.set()

        self.thread = threading.Thread(target=run, daemon=True)
        self.thread.start()

    def settle(self):
        assert self.finished.wait(
            5,
        ), "synthetic worker did not settle after cleanup"
        self.thread.join(1)


def interrupt(store, record, how, clock=None):
    if how == "cancel":
        store.cancel(record.job_id, record.conversation_id)
    elif how == "pause":
        store.pause(record.job_id, record.conversation_id)
    elif how == "deadline":
        assert clock is not None
        clock.current += timedelta(seconds=31)
    else:
        raise AssertionError(how)


def test_accepted_job_deadline_includes_time_before_first_worker(tmp_path):
    record = PresentationJobStore(tmp_path).create(request())
    assert (
        record.deadline_at is not None
    ), "accepted/queued job has no absolute deadline"
    remaining = (record.deadline_at - record.created_at).total_seconds()
    assert 0 < remaining <= record.timeout_seconds + 1


@pytest.mark.parametrize(
    "stale_owner",
    [False, True],
    ids=["resume", "restart"],
)
def test_original_deadline_and_cost_survive_new_writer(
    tmp_path,
    clock,
    stale_owner,
):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    with store.writer_lease(record.job_id) as first:
        original_deadline = first.deadline_at
        store.consume_cost(record.job_id, 7)
    if stale_owner:
        state = store.get(record.job_id)
        state.active_lease_id = "synthetic-dead-process-lease"
        # White-box lock exercises the actual private boundary.
        # pylint: disable-next=protected-access
        store._write(state)
    else:
        store.pause(record.job_id, record.conversation_id)
        store.resume(record.job_id, record.conversation_id)
    clock.current += timedelta(seconds=11)
    with store.writer_lease(record.job_id) as second:
        assert (
            second.deadline_at == original_deadline
        ), "same operation receives a fresh timeout"
        assert store.get(record.job_id).cost_used_units == 7


def test_expired_original_job_cannot_dispatch_under_new_lease(tmp_path, clock):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    with store.writer_lease(record.job_id):
        original = store.get(record.job_id).deadline_at
    clock.current = original + timedelta(seconds=1)
    dispatched = []
    try:
        with store.writer_lease(record.job_id):
            store.assert_execution_budget(record.job_id)
            dispatched.append("provider-or-compiler")
    except (TimeoutError, PresentationJobInterruptedError):
        pass
    assert (
        not dispatched
    ), "expired generation was given a new execution budget"


def test_cancelled_job_never_becomes_active_when_new_writer_is_acquired(
    tmp_path,
):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    store.cancel(record.job_id, record.conversation_id)
    with pytest.raises(PresentationJobInterruptedError):
        with store.writer_lease(record.job_id):
            store.assert_execution_budget(record.job_id)
    assert store.get(record.job_id).status == JobStatus.CANCELLED


def test_pause_resume_does_not_reauthorize_inflight_planner(
    tmp_path,
    monkeypatch,
):
    entered, release = threading.Event(), threading.Event()
    dispatched = []

    class BarrierPlanner:
        # Callback keeps the public request keyword in its signature.
        # pylint: disable-next=redefined-outer-name
        def plan(self, request, brief, **kwargs):
            del kwargs
            entered.set()
            assert release.wait(5)
            return build_deterministic_fixture_plan(request, brief)

    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=BarrierPlanner(),
        page_author=FixtureAuthor(),
    )
    record = runtime.service.create_job(request())
    monkeypatch.setattr(
        runtime.service,
        "compose",
        lambda *a, **k: dispatched.append("compose"),
    )

    def action():
        with runtime.service.store.writer_lease(record.job_id) as lease:
            # White-box lock exercises the actual private boundary.
            # pylint: disable-next=protected-access
            runtime._run_with_lease(record.job_id, lease)

    worker = Worker(action)
    try:
        assert entered.wait(5), worker.errors
        runtime.service.store.pause(record.job_id, record.conversation_id)
        runtime.service.store.resume(record.job_id, record.conversation_id)
        release.set()
        worker.settle()
        assert not runtime.service.store.checkpoint_path(
            record.job_id,
            "deck-plan.json",
        ).exists(), "retired planner published a plan after pause->resume"
        assert (
            not dispatched
        ), "old writer resumed composing under the new control state"
    finally:
        release.set()
        worker.settle()
        runtime.shutdown()


@pytest.mark.parametrize("stage", ["planner", "author"])
def test_real_async_model_wrapper_cancels_and_settles_on_stop(
    tmp_path,
    monkeypatch,
    stage,
):
    entered, released, coroutine_cancelled = (
        threading.Event(),
        threading.Event(),
        threading.Event(),
    )

    async def wait_for_release():
        entered.set()
        try:
            while not released.is_set():
                await asyncio.sleep(0.01)
        except asyncio.CancelledError:
            coroutine_cancelled.set()
            raise

    planner = (
        UtilityPresentationSemanticPlanner("synthetic-agent")
        if stage == "planner"
        else FixturePlanner()
    )
    author = (
        UtilityPresentationPageAuthor("synthetic-agent")
        if stage == "author"
        else FixtureAuthor()
    )
    if stage == "planner":

        async def delayed_plan(req, brief, context):
            del context
            await wait_for_release()
            return build_deterministic_fixture_plan(req, brief)

        monkeypatch.setattr(planner, "_plan_staged", delayed_plan)
    else:
        monkeypatch.setattr(
            author,
            "cache_identity",
            lambda: "synthetic-route-no-provider",
        )

        async def delayed_utility(**kwargs):
            del kwargs
            await wait_for_release()
            return SimpleNamespace(
                model_label="synthetic",
                value=SimpleNamespace(to_runtime=authored),
            )

        monkeypatch.setattr(
            intelligence_module,
            "run_structured_utility_task",
            delayed_utility,
        )
    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=planner,
        page_author=author,
    )
    record = runtime.service.create_job(request())
    monkeypatch.setattr(runtime.service, "compose", lambda *a, **k: None)

    def action():
        with runtime.service.store.writer_lease(record.job_id) as lease:
            # White-box lock exercises the actual private boundary.
            # pylint: disable-next=protected-access
            runtime._run_with_lease(record.job_id, lease)

    worker = Worker(action)
    try:
        assert entered.wait(5), worker.errors
        runtime.cancel(record.job_id, record.conversation_id)
        cancelled_before_release = coroutine_cancelled.wait(WAIT)
        settled_before_release = worker.finished.wait(0.1)
        owner_before_release = runtime.service.store.get(
            record.job_id,
        ).active_lease_id
        released.set()
        worker.settle()
        assert (
            cancelled_before_release
        ), f"{stage} async provider was not cancelled by Stop"
        assert (
            settled_before_release
        ), f"{stage} Stop remained blocked on uncancelled provider"
        assert (
            owner_before_release is None
        ), "settled producer retained stale writer"
        assert (
            runtime.service.store.get(record.job_id).status
            == JobStatus.CANCELLED
        )
    finally:
        released.set()
        worker.settle()
        runtime.shutdown()


@pytest.mark.parametrize("stage", ["synthesis", "rendering"])
@pytest.mark.parametrize("how", ["cancel", "deadline"])
def test_waiting_stage_reacts_before_capacity_is_released(
    tmp_path,
    monkeypatch,
    clock,
    stage,
    how,
):
    limits = PresentationExecutionLimits(synthesis=1, rendering=1, critics=1)
    service = PresentationService(tmp_path, execution_limits=limits)
    record = service.create_job(request())
    admitted, capacity_release = threading.Event(), threading.Event()
    gate = getattr(limits, stage)
    owner_started = threading.Event()

    def hold_capacity():
        with gate():
            owner_started.set()
            assert capacity_release.wait(5)

    holder = Worker(hold_capacity)
    assert owner_started.wait(2)
    if stage == "synthesis":

        class NoDispatchCompiler:
            def compile(self, *args, **kwargs):
                del args, kwargs
                admitted.set()
                raise AssertionError(
                    "compiler dispatch after lifecycle interruption",
                )

        monkeypatch.setattr(
            service,
            "compiler_for",
            lambda engine: NoDispatchCompiler(),
        )
        action = lambda lease: service.compose(
            record.job_id,
            deck(),
            writer_lease=lease,
        )
    else:
        output = tmp_path / "real-fixture.pptx"
        PresentationCompiler().compile(deck(), output)

        def fake_renderer(*args, **kwargs):
            del args, kwargs
            admitted.set()
            raise PresentationRenderUnavailable("synthetic renderer")

        monkeypatch.setattr(service_module, "render_pptx", fake_renderer)
        action = lambda lease: service.verify_artifact(
            deck(),
            output,
            tmp_path / "render",
            visual_required=False,
            job_id=record.job_id,
        )
    with service.store.writer_lease(record.job_id) as lease:
        blocked = threading.Event()
        semaphore = getattr(limits, "_" + stage)
        original_acquire = semaphore.acquire

        def acquire(*args, **kwargs):
            blocked.set()
            return original_acquire(*args, **kwargs)

        monkeypatch.setattr(semaphore, "acquire", acquire)
        worker = Worker(lambda: action(lease))
        try:
            assert blocked.wait(2), worker.errors
            assert not worker.finished.is_set(), worker.errors
            interrupt(service.store, record, how, clock)
            stopped_while_occupied = worker.finished.wait(WAIT)
            capacity_release.set()
            holder.settle()
            worker.settle()
            assert stopped_while_occupied, f"{stage} wait ignored {how}"
            assert (
                not admitted.is_set()
            ), f"{stage} dispatched work after {how}"
        finally:
            capacity_release.set()
            holder.settle()
            worker.settle()


def test_expired_publication_never_invokes_output_or_delivery_sink(
    tmp_path,
    clock,
):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    with store.writer_lease(record.job_id):
        store.append_event(
            record.job_id,
            stage=JobStage.QUALITY,
            status=JobStatus.RUNNING,
            message="synthetic verified candidate",
        )
        clock.current += timedelta(seconds=31)
        published = []
        try:
            store.complete_after_verified_publication(
                record.job_id,
                message="synthetic",
                checkpoint="synthetic-quality",
                publish=lambda: published.append("output+version"),
            )
        except (TimeoutError, PresentationJobInterruptedError):
            pass
        assert (
            not published
        ), "expired writer reached final publication under manifest lock"
        assert store.get(record.job_id).status != JobStatus.COMPLETED


def test_active_publication_preserves_existing_success_contract(tmp_path):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    with store.writer_lease(record.job_id):
        published = []
        store.complete_after_verified_publication(
            record.job_id,
            message="synthetic",
            checkpoint="synthetic-quality",
            publish=lambda: published.append("output+version"),
        )
    assert published == ["output+version"]
    assert store.get(record.job_id).status == JobStatus.COMPLETED


def test_cancelled_publication_never_invokes_output_sink(tmp_path):
    store = PresentationJobStore(tmp_path)
    record = store.create(request())
    with store.writer_lease(record.job_id):
        store.cancel(record.job_id, record.conversation_id)
        published = []
        with pytest.raises(PresentationJobInterruptedError):
            store.complete_after_verified_publication(
                record.job_id,
                message="synthetic",
                checkpoint="synthetic-quality",
                publish=lambda: published.append("output"),
            )
    assert not published


def test_real_native_compiler_uses_remaining_job_budget(
    tmp_path,
    monkeypatch,
    clock,
):
    service = PresentationService(tmp_path)
    record = service.create_job(
        request(composition_engine=CompositionEngine.NATIVE_PPTX),
    )
    # Real compiler code; only the process launch is replaced, so no binary
    # executes.
    compiler = NativePptxCompiler.__new__(NativePptxCompiler)
    compiler.node_binary = tmp_path / "synthetic-node"
    compiler.backend_module_path = tmp_path / "synthetic-backend"
    compiler.script_path = tmp_path / "synthetic-compiler.mjs"
    compiler.script_path.write_text("// synthetic", encoding="utf-8")
    observed = []

    class LaunchObserved(RuntimeError):
        pass

    def launch(*args, **kwargs):
        del args
        observed.append(kwargs.get("timeout"))
        raise LaunchObserved("synthetic launch observed")

    monkeypatch.setattr(service, "compiler_for", lambda engine: compiler)
    monkeypatch.setattr(
        "lantai.presentation.pptx_compiler.subprocess.run",
        launch,
    )
    with service.store.writer_lease(record.job_id) as lease:
        clock.current += timedelta(seconds=25)
        with pytest.raises(LaunchObserved):
            service.compose(record.job_id, deck(), writer_lease=lease)
    assert observed and observed[0] is not None
    assert (
        0 < observed[0] <= 5.1
    ), "native compiler still gets a fixed 180-second budget"


def test_render_receives_remaining_whole_job_budget(
    tmp_path,
    monkeypatch,
    clock,
):
    service = PresentationService(tmp_path)
    record = service.create_job(request())
    output = tmp_path / "real-fixture.pptx"
    PresentationCompiler().compile(deck(), output)
    observed = []

    def render(*args, **kwargs):
        del args
        observed.append(kwargs.get("timeout", 120))
        raise PresentationRenderUnavailable("synthetic absent renderer")

    monkeypatch.setattr(service_module, "render_pptx", render)
    with service.store.writer_lease(record.job_id):
        clock.current += timedelta(seconds=25)
        service.verify_artifact(
            deck(),
            output,
            tmp_path / "render",
            visual_required=False,
            job_id=record.job_id,
        )
    assert (
        observed and 0 < observed[0] <= 5.1
    ), "renderer gets a new fixed 120-second budget"


def test_expired_waiting_job_never_calls_model_when_executor_becomes_available(
    tmp_path,
    monkeypatch,
    clock,
):
    calls, entered, release = [], threading.Event(), threading.Event()

    class RecordingPlanner(FixturePlanner):
        # Callback keeps the public request keyword in its signature.
        # pylint: disable-next=redefined-outer-name
        def plan(self, request, brief, **kwargs):
            calls.append("planner")
            return super().plan(request, brief, **kwargs)

    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=RecordingPlanner(),
        page_author=FixtureAuthor(),
    )
    monkeypatch.setattr(runtime.service, "compose", lambda *a, **k: None)

    def occupy():
        entered.set()
        assert release.wait(5)

    # White-box lock exercises the actual private boundary.
    # pylint: disable-next=protected-access
    blocker = runtime._executor.submit(occupy)
    try:
        assert entered.wait(2)
        record = runtime.create_and_start(request())
        clock.current += timedelta(seconds=31)
        release.set()
        blocker.result(5)
        # White-box lock exercises the actual private boundary.
        # pylint: disable-next=protected-access
        runtime._futures[record.job_id].result(5)
        assert (
            not calls
        ), "expired queued job received a fresh deadline and called planner"
        assert (
            runtime.service.store.get(record.job_id).status
            != JobStatus.COMPLETED
        )
    finally:
        release.set()
        blocker.result(5)
        runtime.shutdown()


@pytest.mark.parametrize("resume_count", [1, 3])
def test_queued_resume_callbacks_cannot_reopen_completed_delivery(
    tmp_path,
    monkeypatch,
    resume_count,
):
    from PIL import Image
    from pptx import Presentation

    entered, release = threading.Event(), threading.Event()
    callbacks_settled = threading.Event()
    drains = []

    def render(source, output_dir, *, timeout=120):
        del timeout
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        pdf = destination / "presentation.pdf"
        pdf.write_bytes(b"synthetic-renderer-pdf")
        images = []
        for index in range(1, len(Presentation(str(source)).slides) + 1):
            path = destination / f"slide-{index}.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste("black", (0, 0, 280 + index, 360))
            image.save(path)
            images.append(path)
        return RenderResult(
            pdf_path=pdf,
            slide_images=tuple(images),
            renderer_id="synthetic-renderer",
            renderer_version="1.0",
        )

    monkeypatch.setattr(service_module, "render_pptx", render)
    runtime = PresentationJobRuntime(
        tmp_path,
        semantic_planner=FixturePlanner(),
        page_author=FixtureAuthor(),
    )

    def occupy():
        entered.set()
        assert release.wait(5)

    # Exercise the actual executor queue and all resume completion callbacks.
    # pylint: disable-next=protected-access
    blocker = runtime._executor.submit(occupy)
    try:
        assert entered.wait(2)
        record = runtime.create_and_start(request())
        original_deadline = record.deadline_at
        for _ in range(resume_count):
            runtime.pause(record.job_id, record.conversation_id)
            runtime.resume(record.job_id, record.conversation_id)
        # Register last: the FIFO drain runs after any stale callback's writer.
        # Future.result() alone does not wait for completion callbacks.
        # pylint: disable-next=protected-access
        original = runtime._futures[record.job_id]

        def drain_callbacks(_future):
            # pylint: disable-next=protected-access
            drains.append(runtime._executor.submit(lambda: None))
            callbacks_settled.set()

        original.add_done_callback(drain_callbacks)
        release.set()
        blocker.result(5)
        original.result(5)
        assert callbacks_settled.wait(5)
        drains[0].result(5)
        final = runtime.service.store.get(record.job_id)
        versions = PresentationVersionStore(
            runtime.service.store.job_dir(record.job_id),
        ).history()
        assert final.status == JobStatus.COMPLETED
        assert final.generation == final.attempt_count == 1
        assert (
            len(versions) == 1 and versions[0].generation == final.generation
        )
        assert final.active_operation is None
        assert final.deadline_at == original_deadline
        assert sum(e.status == JobStatus.COMPLETED for e in final.events) == 1
        assert final.active_lease_id is None
        assert not final.resource_settlement_pending
    finally:
        release.set()
        blocker.result(5)
        runtime.shutdown()


def test_completed_expired_job_keeps_explicit_editor_save_available(
    tmp_path,
    clock,
):
    runtime = PresentationJobRuntime(tmp_path)
    record = runtime.service.create_job(request())
    store = runtime.service.store
    output = store.job_dir(record.job_id) / "output" / "presentation.pptx"
    fixture = deck()
    PresentationCompiler().compile(fixture, output)
    versions = PresentationVersionStore(store.job_dir(record.job_id))
    initial = versions.bootstrap(output, fixture, writer="legacy_python_pptx")
    with store.writer_lease(record.job_id):
        store.append_event(
            record.job_id,
            stage=JobStage.COMPLETED,
            status=JobStatus.COMPLETED,
            message="synthetic completed",
        )
    state = store.get(record.job_id)
    state.deadline_at = clock.current - timedelta(days=1)
    state.cost_used_units = state.cost_budget_units
    # White-box lock exercises the actual private boundary.
    # pylint: disable-next=protected-access
    store._write(state)
    changed = fixture.model_copy(deep=True)
    changed.slides[0].objects[0].text = "用户显式编辑后的合成验证"
    PresentationCompiler().compile(changed, output)
    try:
        saved = runtime.register_editor_save(
            record.job_id,
            record.conversation_id,
            base_output_sha256=initial.pptx_sha256,
        )
        assert saved.writer == "genoffice_human" and not saved.reconciled
        assert saved.parent_version_id == initial.version_id
        assert saved.pptx_sha256 == file_sha256(output) != initial.pptx_sha256
        assert (
            versions.get(initial.version_id).pptx_sha256 == initial.pptx_sha256
        )
        assert store.get(record.job_id).status == JobStatus.COMPLETED
    finally:
        runtime.shutdown()


@pytest.mark.parametrize(
    "original_cost_exhausted",
    [False, True],
    ids=["available", "exhausted"],
)
def test_completed_verified_job_accepts_new_explicit_natural_edit(
    tmp_path,
    monkeypatch,
    clock,
    original_cost_exhausted,
):
    """A later authorized edit owns a new budget and retains the old
    artifact.
    """
    from PIL import Image
    from pptx import Presentation

    def render(source, output_dir, *, timeout=120):
        del timeout
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)
        pdf = destination / "presentation.pdf"
        pdf.write_bytes(b"synthetic-renderer-pdf")
        images = []
        for index in range(1, len(Presentation(str(source)).slides) + 1):
            path = destination / f"slide-{index}.png"
            image = Image.new("RGB", (640, 360), "white")
            image.paste("black", (0, 0, 280 + index, 360))
            image.save(path)
            images.append(path)
        return RenderResult(
            pdf_path=pdf,
            slide_images=tuple(images),
            renderer_id="synthetic-renderer",
            renderer_version="1.0",
        )

    monkeypatch.setattr(service_module, "render_pptx", render)
    runtime = PresentationJobRuntime(tmp_path)
    store = runtime.service.store
    record = runtime.service.create_job(request())
    fixture = DeckIR(
        id="synthetic-edit-deck",
        title="合成显式修改",
        slides=[
            DeckSlide(
                id="edit-first",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="edit-first-title",
                        role="title",
                        slot="title",
                        text="原始合成标题",
                    ),
                ],
            ),
            DeckSlide(
                id="edit-second",
                layout_id="closing",
                objects=[
                    TextObject(
                        id="edit-second-title",
                        role="title",
                        slot="title",
                        text="保持原页内容",
                    ),
                ],
            ),
        ],
    )
    store.write_checkpoint(
        record.job_id,
        "deck-ir.json",
        fixture.model_dump_json(indent=2),
    )
    output = store.job_dir(record.job_id) / "output" / "presentation.pptx"
    runtime.service.compiler.compile(fixture, output)
    report, images = runtime.service.verify_artifact(
        fixture,
        output,
        store.job_dir(record.job_id) / "renders" / "round-0",
        render_set_id="round-0",
    )
    from lantai.presentation.quality import is_verified_delivery

    assert is_verified_delivery(
        fixture,
        report,
        output,
        images,
    ), "synthetic baseline must satisfy real delivery checks"
    quality = store.job_dir(record.job_id) / "quality" / "quality-report.json"
    quality.parent.mkdir(parents=True, exist_ok=True)
    quality.write_text(report.model_dump_json(indent=2), encoding="utf-8")
    versions = PresentationVersionStore(store.job_dir(record.job_id))
    initial = versions.bootstrap(
        output,
        fixture,
        writer="legacy_python_pptx",
        quality_report_path=quality,
        render_dir=store.job_dir(record.job_id) / "renders" / "round-0",
    )
    store.append_event(
        record.job_id,
        stage=JobStage.COMPLETED,
        status=JobStatus.COMPLETED,
        message="synthetic verified complete",
        checkpoint=str(output),
    )
    state = store.get(record.job_id)
    state.deadline_at = clock.current - timedelta(days=1)
    if original_cost_exhausted:
        state.cost_used_units = state.cost_budget_units
    # White-box lock exercises the actual private boundary.
    # pylint: disable-next=protected-access
    store._write(state)
    try:
        updated = runtime.redo_slide(
            record.job_id,
            record.conversation_id,
            "edit-first",
            "把本页标题改为：显式修改后的合成标题",
            initial.pptx_sha256,
        )
        changed = DeckIR.model_validate_json(
            store.checkpoint_path(record.job_id, "deck-ir.json").read_text(),
        )
        assert updated.status == JobStatus.COMPLETED
        assert changed.slides[0].objects[0].text == "显式修改后的合成标题"
        assert changed.slides[1] == fixture.slides[1]
        assert (
            versions.get(initial.version_id).pptx_sha256 == initial.pptx_sha256
        )
        assert file_sha256(output) != initial.pptx_sha256
    finally:
        runtime.shutdown()


def test_parallel_real_critic_cleanup_retains_writer_until_all_workers_settle(
    tmp_path,
    monkeypatch,
):
    second_entered, release_cleanup, cancellation_seen = (
        threading.Event(),
        threading.Event(),
        threading.Event(),
    )
    second_settled = threading.Event()
    critic = UtilityPresentationVisualCritic("synthetic-agent")
    monkeypatch.setattr(
        critic,
        "slide_cache_identity",
        lambda: "synthetic-slide-critic",
    )
    monkeypatch.setattr(
        critic,
        "deck_cache_identity",
        lambda: "synthetic-deck-critic",
    )
    service = PresentationService(
        tmp_path,
        visual_critic=critic,
        execution_limits=PresentationExecutionLimits(critics=2),
    )
    record = service.create_job(request())

    async def review(slide, **kwargs):
        del kwargs
        if slide.slide_id == "first":
            while not second_entered.is_set():
                await asyncio.sleep(0.01)
            service.store.cancel(record.job_id, record.conversation_id)
            raise PresentationJobInterruptedError("synthetic cancelled")
        second_entered.set()
        try:
            try:
                while not release_cleanup.is_set():
                    await asyncio.sleep(0.01)
            except asyncio.CancelledError:
                cancellation_seen.set()
                while not release_cleanup.is_set():
                    await asyncio.sleep(0.01)
                raise
            return VisualCriticResult(critic_id="synthetic")
        finally:
            second_settled.set()

    monkeypatch.setattr(critic, "_review_slide", review)
    slides = [
        VisualCriticSlideInput(
            slide_id=s,
            slide_index=i,
            rendered_path="synthetic.png",
            rendered_sha256="a" * 64,
            key_message="合成视觉证据",
        )
        for i, s in enumerate(("first", "second"), 1)
    ]

    def action():
        with service.store.writer_lease(record.job_id):
            # White-box lock exercises the actual private boundary.
            # pylint: disable-next=protected-access
            service._review_visuals_with_checkpoints(
                slides,
                job_id=record.job_id,
                repair_rounds=0,
                shared_dependencies={"synthetic": True},
            )

    worker = Worker(action)
    try:
        assert second_entered.wait(3), worker.errors
        assert cancellation_seen.wait(3), worker.errors
        finished_before_cleanup = worker.finished.wait(0.2)
        owner_before_cleanup = service.store.get(record.job_id).active_lease_id
        release_cleanup.set()
        worker.settle()
        assert (
            not finished_before_cleanup
        ), "parent returned while actual sibling provider cleanup was pending"
        assert (
            owner_before_cleanup is not None
        ), "writer lease was cleared before sibling resource settled"
        assert service.store.get(record.job_id).active_lease_id is None
    finally:
        release_cleanup.set()
        worker.settle()
        assert second_settled.wait(
            3,
        ), "synthetic sibling cleanup did not settle"


def test_existing_critic_queue_stop_does_not_dispatch_model(tmp_path):
    limits = PresentationExecutionLimits(critics=1)
    service = PresentationService(tmp_path, execution_limits=limits)
    record = service.create_job(request())
    entered, release, calls = threading.Event(), threading.Event(), []

    def hold():
        with limits.critic():
            entered.set()
            assert release.wait(5)

    holder = Worker(hold)
    assert entered.wait(2)

    def action():
        with service.store.writer_lease(record.job_id):
            # White-box lock exercises the actual private boundary.
            # pylint: disable-next=protected-access
            service._review_visual_scope(
                job_id=record.job_id,
                repair_rounds=0,
                scope="slide",
                checkpoint_name="synthetic.json",
                fingerprint="a" * 64,
                critic_identity="synthetic",
                review=lambda *_: calls.append("provider")
                or VisualCriticResult(critic_id="synthetic"),
                slide_id="synthetic",
                rendered_sha256="b" * 64,
            )

    worker = Worker(action)
    try:
        service.store.cancel(record.job_id, record.conversation_id)
        assert worker.finished.wait(
            WAIT,
        ), "existing cancellable critic queue contract regressed"
        assert not calls
        assert isinstance(worker.errors[0], PresentationJobInterruptedError)
    finally:
        release.set()
        holder.settle()
        worker.settle()


def repair_case(tmp_path, monkeypatch, utility):
    from lantai.presentation.intelligence import (
        UtilityPresentationRepairDesigner,
    )
    from lantai.presentation.models import QualityIssue, QualityReport

    designer = UtilityPresentationRepairDesigner("synthetic-agent")
    service = PresentationService(tmp_path, repair_designer=designer)
    record = service.create_job(
        request(composition_engine=CompositionEngine.NATIVE_PPTX),
    )
    # A real synthetic PPTX and a deliberate eligible visual finding drive the
    # genuine repair dispatcher; no real renderer, Provider or child process.
    monkeypatch.setattr(
        service,
        "compiler_for",
        lambda engine: PresentationCompiler(),
    )
    monkeypatch.setattr(
        service,
        "verify_artifact",
        lambda candidate, *a, **k: (
            QualityReport(
                deck_id=candidate.id,
                passed=False,
                issues=[
                    QualityIssue(
                        code="vision.hierarchy",
                        severity="error",
                        slide_id=candidate.slides[0].id,
                        message="合成层级问题",
                        repairable=True,
                    ),
                ],
            ),
            (),
        ),
    )
    monkeypatch.setattr(
        intelligence_module,
        "run_structured_utility_task",
        utility,
    )
    monkeypatch.setattr(
        intelligence_module,
        "run_structured_utility_content_task",
        utility,
    )
    return service, record


def test_real_repair_model_coroutine_stops_and_settles(tmp_path, monkeypatch):
    entered, release, cancellation_seen = (
        threading.Event(),
        threading.Event(),
        threading.Event(),
    )

    async def utility(**kwargs):
        del kwargs
        entered.set()
        try:
            while not release.is_set():
                await asyncio.sleep(0.01)
        except asyncio.CancelledError:
            cancellation_seen.set()
            raise
        raise RuntimeError("synthetic model released without completion")

    service, record = repair_case(tmp_path, monkeypatch, utility)

    def action():
        with service.store.writer_lease(record.job_id) as lease:
            service.compose(record.job_id, deck(), writer_lease=lease)

    worker = Worker(action)
    try:
        assert entered.wait(3), worker.errors
        service.store.cancel(record.job_id, record.conversation_id)
        cancelled_before_release = cancellation_seen.wait(WAIT)
        settled_before_release = worker.finished.wait(0.1)
        release.set()
        worker.settle()
        assert (
            cancelled_before_release
        ), "actual repair async provider was not cancelled by Stop"
        assert (
            settled_before_release
        ), "repair producer remains blocked after Stop"
        assert service.store.get(record.job_id).status == JobStatus.CANCELLED
    finally:
        release.set()
        worker.settle()


@pytest.mark.parametrize("stage", ["planner", "author", "repair", "critic"])
def test_actual_model_request_uses_only_remaining_whole_job_budget(
    tmp_path,
    monkeypatch,
    clock,
    stage,
):
    observed = []

    class ObservedRequest(RuntimeError):
        pass

    async def utility(**kwargs):
        observed.append(
            {
                name: kwargs.get(name)
                for name in (
                    "timeout_seconds",
                    "first_output_timeout_seconds",
                    "idle_timeout_seconds",
                )
            },
        )
        raise ObservedRequest("synthetic utility dispatch observed")

    monkeypatch.setattr(
        intelligence_module,
        "run_structured_utility_task",
        utility,
    )
    monkeypatch.setattr(
        intelligence_module,
        "run_structured_utility_content_task",
        utility,
    )
    if stage == "repair":
        service, record = repair_case(tmp_path, monkeypatch, utility)
        with service.store.writer_lease(record.job_id) as lease:
            clock.current += timedelta(seconds=25)
            with pytest.raises(ObservedRequest):
                service.compose(record.job_id, deck(), writer_lease=lease)
    elif stage == "critic":
        critic = UtilityPresentationVisualCritic("synthetic-agent")
        monkeypatch.setattr(
            critic,
            "slide_cache_identity",
            lambda: "synthetic-slide",
        )
        monkeypatch.setattr(
            critic,
            "deck_cache_identity",
            lambda: "synthetic-deck",
        )
        service = PresentationService(tmp_path, visual_critic=critic)
        record = service.create_job(request())
        image = tmp_path / "synthetic.png"
        from PIL import Image

        Image.new("RGB", (100, 100), "white").save(image)
        slide = VisualCriticSlideInput(
            slide_id="first",
            slide_index=1,
            rendered_path=str(image),
            rendered_sha256=file_sha256(image),
            key_message="合成视觉核验",
        )
        with service.store.writer_lease(record.job_id):
            clock.current += timedelta(seconds=25)
            with pytest.raises(ObservedRequest):
                # White-box lock exercises the actual private boundary.
                # pylint: disable-next=protected-access
                service._review_visuals_with_checkpoints(
                    [slide],
                    job_id=record.job_id,
                    repair_rounds=0,
                    shared_dependencies={"synthetic": True},
                )
    else:
        planner = (
            UtilityPresentationSemanticPlanner("synthetic-agent")
            if stage == "planner"
            else FixturePlanner()
        )
        author = (
            UtilityPresentationPageAuthor("synthetic-agent")
            if stage == "author"
            else FixtureAuthor()
        )
        if stage == "author":
            monkeypatch.setattr(
                author,
                "cache_identity",
                lambda: "synthetic-author-no-provider",
            )
        runtime = PresentationJobRuntime(
            tmp_path,
            semantic_planner=planner,
            page_author=author,
        )
        record = runtime.service.create_job(request())
        try:
            with runtime.service.store.writer_lease(record.job_id) as lease:
                clock.current += timedelta(seconds=25)
                with pytest.raises(ObservedRequest):
                    # White-box lock exercises the actual private boundary.
                    # pylint: disable-next=protected-access
                    runtime._run_with_lease(record.job_id, lease)
        finally:
            runtime.shutdown()
    assert observed, "test fixture did not reach actual model dispatcher"
    assert (
        0 < observed[0]["timeout_seconds"] <= 5.1
    ), f"{stage} renewed its own model timeout"
    for name in ("first_output_timeout_seconds", "idle_timeout_seconds"):
        value = observed[0][name]
        assert (
            value is None or 0 < value <= 5.1
        ), f"{stage} local timeout exceeds remaining budget"
