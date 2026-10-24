from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.models import (
    AddObjectAction,
    ChartObject,
    DeckIR,
    DeckSlide,
    FactRecord,
    ObjectFrame,
    PresentationObjectEditPlan,
    TextObject,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.quality import inspect_pptx


SOURCE = "fixture://native-comparison"
FACT = "合成均分仅用于描述，不支持因果结论。"


def _edited_deck(*, chart_type="column", categories=None, series=None):
    original = DeckIR(
        id="native-comparison-deck",
        title="原生图表检查",
        source_refs=[SOURCE],
        required_facts=[FACT],
        facts=[FactRecord(fact_id="fact-boundary", text=FACT, source_refs=[SOURCE])],
        slides=[
            DeckSlide(
                id="native-comparison-slide",
                layout_id="freeform",
                semantic_layout="freeform",
                source_fact_ids=["fact-boundary"],
                objects=[
                    TextObject(
                        id="comparison-title", role="title", slot="title",
                        text="组均分：仅描述性比较，不能归因",
                        frame=ObjectFrame(x=0.6, y=0.4, w=12, h=0.8),
                    ),
                    TextObject(
                        id="comparison-boundary", role="body", slot="source",
                        text=FACT, frame=ObjectFrame(x=0.6, y=6.7, w=12, h=0.5),
                    ),
                ],
            ),
        ],
    )
    chart = ChartObject(
        id="native-comparison-chart", chart_type=chart_type,
        categories=categories if categories is not None else ["A班", "B班"],
        series=series if series is not None else {"前测": [61, 60], "后测": [76, 81]},
        unit="分", derivation="直接引用合成均分，不进行统计推断。",
        source_refs=[SOURCE], frame=ObjectFrame(x=0.8, y=1.6, w=11.5, h=4.6),
    )
    updated, changed = apply_object_edit_plan(
        original,
        PresentationObjectEditPlan(
            slide_id=original.slides[0].id, actions=[AddObjectAction(object=chart)],
        ),
    )
    assert changed == [chart.id]
    assert original.slides[0].relationships == []
    assert not any(isinstance(obj, ChartObject) for obj in original.slides[0].objects)
    assert updated.slides[0].layout_id == "freeform"
    assert updated.slides[0].relationships == []
    return updated


def _inspect(tmp_path, deck):
    output = tmp_path / "native-comparison.pptx"
    PresentationCompiler().compile(deck, output)
    return output, inspect_pptx(deck, output)


@pytest.mark.parametrize("chart_type", ["bar", "column"])
@pytest.mark.parametrize("compare_series", [False, True])
def test_real_native_comparison_survives_freeform_object_edit(
    tmp_path, chart_type, compare_series,
):
    deck = _edited_deck(
        chart_type=chart_type,
        categories=["A班"] if compare_series else None,
        series={"前测": [61], "后测": [76]} if compare_series else None,
    )
    output, report = _inspect(tmp_path, deck)
    assert any(shape.has_chart for shape in Presentation(output).slides[0].shapes)
    assert report.passed is True, [issue.code for issue in report.issues]
    assert report.editable is True
    assert report.visual_verified is False  # Structural success is not visual delivery.


@pytest.mark.parametrize("categories", [["A班"], ["A班", "A班"], ["", " "]])
def test_single_or_unlabelled_native_data_does_not_certify_comparison(tmp_path, categories):
    deck = _edited_deck(categories=categories, series={"前测": [61] * len(categories)})
    _, report = _inspect(tmp_path, deck)
    assert report.passed is False
    assert any(issue.code == "design.semantic_layout_mismatch" for issue in report.issues)


@pytest.mark.parametrize("fake_shape", [False, True])
def test_declared_chart_without_real_native_chart_does_not_certify_comparison(
    tmp_path, fake_shape,
):
    deck = _edited_deck()
    output, _ = _inspect(tmp_path, deck)
    presentation = Presentation(output)
    slide = presentation.slides[0]
    native = next(shape for shape in slide.shapes if shape.has_chart)
    name, left, top, width, height = native.name, native.left, native.top, native.width, native.height
    slide.shapes._spTree.remove(native._element)
    if fake_shape:
        slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height).name = name
    presentation.save(output)
    report = inspect_pptx(deck, output)
    assert report.passed is False
    assert any(issue.code == "design.semantic_layout_mismatch" for issue in report.issues)
    if not fake_shape:
        assert any(issue.code == "editability.object_missing" for issue in report.issues)


@pytest.mark.parametrize("wrong_type", [False, True])
def test_physical_chart_must_also_have_comparison_data_and_matching_type(tmp_path, wrong_type):
    deck = _edited_deck()
    output, _ = _inspect(tmp_path, deck)
    presentation = Presentation(output)
    slide = presentation.slides[0]
    native = next(shape for shape in slide.shapes if shape.has_chart)
    name, left, top, width, height = native.name, native.left, native.top, native.width, native.height
    slide.shapes._spTree.remove(native._element)
    data = ChartData()
    data.categories = ["A班", "B班"] if wrong_type else ["A班"]
    data.add_series("前测", [61, 60] if wrong_type else [61])
    slide.shapes.add_chart(
        XL_CHART_TYPE.LINE if wrong_type else XL_CHART_TYPE.COLUMN_CLUSTERED,
        left, top, width, height, data,
    ).name = name
    presentation.save(output)
    report = inspect_pptx(deck, output)
    assert report.passed is False
    assert any(issue.code == "design.semantic_layout_mismatch" for issue in report.issues)


def test_blank_series_labels_do_not_certify_comparison(tmp_path):
    deck = _edited_deck(categories=["A班"], series={"": [61], " ": [76]})
    _, report = _inspect(tmp_path, deck)
    assert report.passed is False
    assert any(issue.code == "design.semantic_layout_mismatch" for issue in report.issues)


@pytest.mark.parametrize(
    ("field", "value", "code"),
    [
        ("unit", "", "chart.unit_missing"),
        ("unit", "benchmark score", "chart.unit_placeholder"),
        ("derivation", "", "chart.derivation_missing"),
        ("source_refs", [], "chart.source_missing"),
        ("source_refs", ["fixture://invented"], "chart.source_invalid"),
    ],
)
def test_native_comparison_keeps_source_unit_and_derivation_gates(tmp_path, field, value, code):
    deck = _edited_deck()
    slide = deck.slides[0]
    changed = slide.model_copy(update={
        "objects": [obj.model_copy(update={field: value}) if isinstance(obj, ChartObject) else obj
                    for obj in slide.objects],
    })
    deck = deck.model_copy(update={"slides": [changed]})
    _, report = _inspect(tmp_path, deck)
    assert report.passed is False
    assert any(issue.code == code and issue.severity == "error" for issue in report.issues)


def test_native_comparison_chart_does_not_certify_process_composition(tmp_path):
    deck = _edited_deck()
    slide = deck.slides[0]
    changed = slide.model_copy(update={
        "objects": [obj.model_copy(update={"text": "研究流程"}) if obj.id == "comparison-title" else obj
                    for obj in slide.objects],
    })
    _, report = _inspect(tmp_path, deck.model_copy(update={"slides": [changed]}))
    assert report.passed is False
    assert any(issue.code == "design.semantic_layout_mismatch" for issue in report.issues)
