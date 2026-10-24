"""Native table export and scoped edits have unambiguous shape identities."""

from copy import deepcopy
from zipfile import ZipFile

from lxml import etree as XML
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches

from lantai.presentation.editing import PresentationEditTransactionService
from lantai.presentation.models import (
    DeckIR, DeckSlide, ObjectFrame, PresentationEditOperation,
    PresentationObjectEditPlan, TableObject, TextObject,
)
from lantai.presentation.natural_edit import apply_object_edit_plan
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.reconcile import reconcile_pptx


def _deck():
    return DeckIR(id="native-table-identity", title="原生表格身份", slides=[
        DeckSlide(id=slide_id, layout_id="freeform", objects=[
            TextObject(id=f"{slide_id}-title", text="保留标题", role="title",
                       frame=ObjectFrame(x=1, y=.4, w=10, h=1)),
            TableObject(id=f"{slide_id}-table", headers=["问题", "依据"],
                        rows=[["原判断", "保留来源"]],
                        frame=ObjectFrame(x=1, y=2, w=10, h=2)),
        ]) for slide_id in ("first", "second")
    ])


def test_native_table_export_has_unique_positive_shape_ids_on_every_page(tmp_path):
    output = tmp_path / "native-tables.pptx"
    result = NativePptxCompiler().compile(_deck(), output)
    assert result.degradations == [], "a lossless identity repair is not a content degradation"
    assert "p1: first-table: native duplicate table id 2 -> 5" in result.native_repairs
    for slide in Presentation(output).slides:
        ids = [int(node.get("id")) for node in slide._element.xpath(".//p:cNvPr")]
        assert all(value > 0 for value in ids)
        assert len(ids) == len(set(ids)), "native table export published duplicate shape IDs"
        assert len([shape for shape in slide.shapes if shape.has_table]) == 1


def _native_bytes(shape):
    return XML.tostring(shape._element, method="c14n", exclusive=True)


def test_public_text_edit_repairs_only_the_selected_legacy_table_identity(tmp_path):
    deck, writer = _deck(), NativePptxCompiler()
    source = tmp_path / "legacy.pptx"
    writer.compile(deck, source)
    legacy = Presentation(source)
    title = next(shape for shape in legacy.slides[0].shapes if shape.name == "first-title")
    table = next(shape for shape in legacy.slides[0].shapes if shape.has_table)
    table._element.xpath(".//p:cNvPr")[0].set("id", str(title.shape_id))
    legacy.save(source)
    source_bytes = source.read_bytes()
    original = Presentation(source)
    untouched = {shape.name: _native_bytes(shape) for shape in original.slides[0].shapes
                 if not shape.has_table}
    old_table = next(shape for shape in original.slides[0].shapes if shape.has_table)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan.model_validate({
        "slide_id": "first", "actions": [{
            "kind": "update_table", "object_id": "first-table", "headers": ["问题", "依据"],
            "rows": [["新判断", "保留来源"]],
        }],
    }))
    preview = PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
        source, imported, changed, PresentationEditOperation(
            deck_id=deck.id, slide_ids=["first"], object_ids=ids,
            allowed_content_types={"text", "table"}, instruction="只改第一张表格文字",
        ), writer,
    )
    result = Presentation(preview.candidate_path)
    actual = next(shape for shape in result.slides[0].shapes if shape.has_table)
    native_ids = [int(node.get("id")) for node in result.slides[0]._element.xpath(".//p:cNvPr")]
    assert len(native_ids) == len(set(native_ids))
    assert actual.shape_id != old_table.shape_id
    assert {shape.name: _native_bytes(shape) for shape in result.slides[0].shapes
            if not shape.has_table} == untouched
    expected = deepcopy(old_table._element)
    expected.xpath(".//p:cNvPr")[0].set("id", str(actual.shape_id))
    expected.xpath(".//a:tr")[1].xpath("./a:tc")[0].xpath(".//a:t")[0].text = "新判断"
    assert _native_bytes(actual) == XML.tostring(expected, method="c14n", exclusive=True)
    assert any("first-table: native duplicate table id" in item for item in preview.change_summary)
    with ZipFile(source) as before, ZipFile(preview.candidate_path) as after:
        assert after.read("ppt/slides/slide2.xml") == before.read("ppt/slides/slide2.xml")
    reimported = reconcile_pptx(changed, preview.candidate_path, assets_dir=tmp_path / "reimport")
    assert next(obj for obj in reimported.slides[0].objects if obj.id == "first-table").rows[0][0] == "新判断"
    assert source.read_bytes() == source_bytes


@pytest.mark.parametrize("mutation,message", [
    ("connector", "ambiguous reference"),
    ("animation", "ambiguous reference"),
    ("duplicate-name", "identity is missing or ambiguous"),
    ("three-way", "ambiguous shape ids"),
    ("root-collision", "identity is missing or ambiguous"),
    ("nonpositive", "ambiguous shape ids"),
    ("nonnumeric", "ambiguous shape ids"),
    ("unselected-table", "ambiguous shape ids"),
])
def test_public_edit_never_guesses_an_unapproved_or_ambiguous_native_identity(tmp_path, mutation, message):
    deck, writer = _deck(), NativePptxCompiler()
    source = tmp_path / "legacy.pptx"
    writer.compile(deck, source)
    external = Presentation(source)
    title = next(shape for shape in external.slides[0].shapes if shape.name == "first-title")
    table = next(shape for shape in external.slides[0].shapes if shape.has_table)
    table._element.xpath(".//p:cNvPr")[0].set("id", str(title.shape_id))
    external.save(source)
    imported = reconcile_pptx(deck, source, assets_dir=tmp_path / "assets")
    external = Presentation(source)
    slide = external.slides[0]
    title = next(shape for shape in slide.shapes if shape.name == "first-title")
    table = next(shape for shape in slide.shapes if shape.has_table)
    nv = table._element.xpath(".//p:cNvPr")[0]
    if mutation == "connector":
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(2), Inches(1), Inches(3), Inches(2),
        )
        connector.begin_connect(title, 0)
    elif mutation == "animation":
        namespace = "http://schemas.openxmlformats.org/presentationml/2006/main"
        timing = XML.SubElement(slide._element, f"{{{namespace}}}timing")
        target = XML.SubElement(timing, f"{{{namespace}}}spTgt")
        target.set("spid", str(title.shape_id))
    elif mutation == "duplicate-name":
        title.name = table.name
    elif mutation == "three-way":
        slide.shapes[-1]._element.xpath(".//p:cNvPr")[0].set("id", str(title.shape_id))
    elif mutation == "root-collision":
        nv.set("id", slide._element.xpath("./p:cSld/p:spTree/p:nvGrpSpPr/p:cNvPr")[0].get("id"))
    elif mutation in {"nonpositive", "nonnumeric"}:
        nv.set("id", "0" if mutation == "nonpositive" else "invalid")
    external.save(source)
    before = source.read_bytes()
    action = ({"kind": "replace_text", "object_id": "first-title", "text": "新标题"}
              if mutation == "unselected-table" else {
                  "kind": "update_table", "object_id": "first-table", "headers": ["问题", "依据"],
                  "rows": [["新判断", "保留来源"]],
              })
    changed, ids = apply_object_edit_plan(imported, PresentationObjectEditPlan.model_validate({
        "slide_id": "first", "actions": [action],
    }))
    with pytest.raises(ValueError, match=message):
        PresentationEditTransactionService(tmp_path / "edits").prepare_compiled_redesign(
            source, imported, changed, PresentationEditOperation(
                deck_id=deck.id, slide_ids=["first"], object_ids=ids,
                allowed_content_types={"text", "table"}, instruction="不猜测不扩展表格修复范围",
            ), writer,
        )
    assert source.read_bytes() == before
    assert not list((tmp_path / "edits").rglob("preview.json"))
    assert not list((tmp_path / "edits").rglob("candidate.pptx"))
    assert not list((tmp_path / "edits").rglob("candidate.stage-*.pptx"))
