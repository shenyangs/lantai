from __future__ import annotations

from pptx import Presentation
from pptx.util import Pt

from lantai.presentation.compiler import PresentationCompiler
from lantai.presentation.models import (
    ChartObject,
    DeckIR,
    DeckSlide,
    EvidenceState,
    FactRecord,
    MetricObject,
    ObjectRelationship,
    ObjectFrame,
    PresentationSourceCoverage,
    RelationshipType,
    ShapeObject,
    TableObject,
    TextObject,
)
from lantai.presentation.quality import inspect_pptx


def test_unverified_legacy_doc_conversion_cannot_pass_formal_quality(
    tmp_path,
) -> None:
    deck = _golden_deck().model_copy(
        update={
            "source_coverage": [
                PresentationSourceCoverage(
                    material_id="source-old-doc",
                    source_name="old.doc",
                    source_sha256="a" * 64,
                    converted_sha256="b" * 64,
                    method="libreoffice_doc_to_docx_ooxml",
                    status="converted_doc_unverified",
                    extracted_chars=42,
                )
            ]
        }
    )
    output = tmp_path / "unverified-doc.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    assert any(
        issue.code == "material.conversion_unverified"
        for issue in report.issues
    )


def _golden_deck() -> DeckIR:
    return DeckIR(
        id="deck-golden",
        title="澜台产品汇报",
        theme_id="lantai-ink",
        immutable_data=["95%"],
        slides=[
            DeckSlide(
                id="slide-cover",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="cover-title",
                        role="title",
                        text="让演示文稿真正可用",
                        slot="title",
                    ),
                    TextObject(
                        id="cover-subtitle",
                        role="body",
                        text="确定性排版与原生可编辑对象",
                        slot="subtitle",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-judgment",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id="judgment-title",
                        role="title",
                        text="核心判断",
                        slot="title",
                    ),
                    TextObject(
                        id="judgment-body",
                        role="body",
                        text="常规任务成功率目标为 95%",
                        slot="body",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-cards",
                layout_id="three_cards",
                objects=[
                    TextObject(
                        id="cards-title",
                        role="title",
                        text="三个工程支点",
                        slot="title",
                    ),
                    TextObject(
                        id="card-one",
                        role="heading",
                        text="DeckIR\n结构化语义",
                        slot="card_1",
                    ),
                    TextObject(
                        id="card-two",
                        role="heading",
                        text="Layout\n确定性约束",
                        slot="card_2",
                    ),
                    TextObject(
                        id="card-three",
                        role="heading",
                        text="QA\n逐页检查",
                        slot="card_3",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-process",
                layout_id="process",
                objects=[
                    TextObject(
                        id="process-title",
                        role="title",
                        text="生成流程",
                        slot="title",
                    ),
                    *[
                        TextObject(
                            id=f"step-{index}",
                            role="heading",
                            text=f"{index}\n阶段",
                            slot=f"step_{index}",
                        )
                        for index in range(1, 6)
                    ],
                ],
            ),
            DeckSlide(
                id="slide-data",
                layout_id="data_conclusion",
                objects=[
                    TextObject(
                        id="data-title",
                        role="title",
                        text="质量门槛",
                        slot="title",
                    ),
                    ChartObject(
                        id="native-chart",
                        chart_type="column",
                        categories=["基线", "V1"],
                        series={"可用率": [42.0, 95.0]},
                        unit="percent",
                        number_format="0.0%",
                        derivation="direct fixture values",
                        source_refs=["fixture://compiler-quality"],
                        slot="chart",
                    ),
                    MetricObject(
                        id="metric-one",
                        label="目标",
                        value="95%",
                        slot="metric_1",
                    ),
                    MetricObject(
                        id="metric-two",
                        label="可重开",
                        value="99.5%",
                        slot="metric_2",
                    ),
                    TextObject(
                        id="data-conclusion",
                        role="body",
                        text="先稳定，再智能。",
                        slot="conclusion",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-table",
                layout_id="table",
                objects=[
                    TextObject(
                        id="table-title",
                        role="title",
                        text="原生对象验证",
                        slot="title",
                    ),
                    TableObject(
                        id="native-table",
                        headers=["对象", "状态"],
                        rows=[["文字", "可编辑"], ["图表", "可编辑"]],
                        slot="table",
                    ),
                    TextObject(
                        id="table-conclusion",
                        role="body",
                        text="不使用整页截图。",
                        slot="conclusion",
                    ),
                ],
            ),
        ],
    )


def test_compiler_creates_reopenable_native_editable_pptx(tmp_path) -> None:
    deck = _golden_deck()
    output = tmp_path / "golden.pptx"

    result = PresentationCompiler().compile(deck, output)
    reopened = Presentation(output)
    native_names = {
        shape.name for slide in reopened.slides for shape in slide.shapes
    }

    assert result.slide_count == len(deck.slides)
    assert result.editable_object_count == sum(
        len(slide.objects) for slide in deck.slides
    )
    assert output.is_file()
    assert "native-chart" in native_names
    assert "native-table" in native_names
    assert any(
        shape.has_chart for slide in reopened.slides for shape in slide.shapes
    )
    assert any(
        shape.has_table for slide in reopened.slides for shape in slide.shapes
    )

    report = inspect_pptx(deck, output)
    assert report.passed is True, report.model_dump()
    assert report.editable is True
    assert report.rendered_slide_count == 0
    assert report.render_verified is False
    assert report.render_evidence == EvidenceState.NOT_RUN
    assert report.visual_verified is False


def test_summary_title_with_route_word_is_not_forced_into_process_layout(
    tmp_path,
) -> None:
    deck = DeckIR(
        id="deck-summary-route",
        title="路线总结",
        slides=[
            DeckSlide(
                id="slide-summary",
                layout_id="key_judgment",
                    relationships=[
                        ObjectRelationship(
                            source_object_ids=["summary-title"],
                            target_object_ids=["summary-body"],
                            relationship_type=RelationshipType.SUMMARY,
                        ),
                    ],
                objects=[
                    TextObject(
                        id="summary-title",
                        role="title",
                        text="一条路线已经形成",
                        slot="title",
                    ),
                    TextObject(
                        id="summary-body",
                        role="body",
                        text="这是一页收束判断，不是流程图。",
                        slot="body",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "summary-route.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert not any(
        issue.code == "design.semantic_layout_mismatch"
        for issue in report.issues
    )


def test_semantic_layout_mismatch_is_routed_to_controlled_redesign(
    tmp_path,
) -> None:
    deck = DeckIR(
        id="deck-comparison-mismatch",
        title="对比页",
        slides=[
            DeckSlide(
                id="slide-comparison",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id="comparison-title",
                        role="title",
                        text="收入与利润对比",
                        slot="title",
                    ),
                    TextObject(
                        id="comparison-body",
                        role="body",
                        text="当前页面尚未表达比较关系。",
                        slot="body",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "comparison-mismatch.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    mismatch = next(
        issue
        for issue in report.issues
        if issue.code == "design.semantic_layout_mismatch"
    )
    assert mismatch.severity == "error"
    assert mismatch.repairable is True


def test_quality_detects_small_inline_run_and_routes_repair(tmp_path) -> None:
    deck = DeckIR(
        id="deck-inline-readability",
        title="字号验证",
        slides=[
            DeckSlide(
                id="slide-inline-readability",
                layout_id="freeform",
                objects=[
                    TextObject(
                        id="body-with-small-run",
                        role="body",
                        text="正文注释",
                        font_size=20,
                        frame=ObjectFrame(x=1, y=1, w=5, h=1),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "inline-readability.pptx"
    PresentationCompiler().compile(deck, output)
    presentation = Presentation(output)
    shape = next(
        item
        for item in presentation.slides[0].shapes
        if item.name == "body-with-small-run"
    )
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.clear()
    paragraph.font.size = Pt(20)
    normal = paragraph.add_run()
    normal.text = "正文"
    normal.font.size = Pt(20)
    small = paragraph.add_run()
    small.text = "注释"
    small.font.size = Pt(12)
    presentation.save(output)

    report = inspect_pptx(deck, output)
    issue = next(
        item
        for item in report.issues
        if item.code == "design.font_too_small"
        and item.object_id == "body-with-small-run"
    )

    assert issue.severity == "error"
    assert issue.repairable is True


def test_native_connector_is_not_reported_as_object_overlap(tmp_path) -> None:
    deck = DeckIR(
        id="deck-connector",
        title="原生连接线",
        slides=[
            DeckSlide(
                id="slide-diagram",
                layout_id="freeform",
                objects=[
                    TextObject(
                        id="diagram-title",
                        role="title",
                        text="连接线可以穿过节点边界",
                        frame=ObjectFrame(x=0.8, y=0.5, w=7.0, h=0.8),
                    ),
                    ShapeObject(
                        id="diagram-node",
                        shape_type="rect",
                        frame=ObjectFrame(x=1.2, y=2.0, w=3.2, h=1.5),
                    ),
                    ShapeObject(
                        id="diagram-connector",
                        shape_type="line",
                        frame=ObjectFrame(x=1.2, y=2.7, w=3.2, h=0.08),
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "connector.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert not any(issue.code == "layout.complete_overlap" for issue in report.issues)


def test_compiler_supports_all_native_chart_types(tmp_path) -> None:
    for chart_type in ("bar", "column", "line", "pie"):
        deck = DeckIR(
            id=f"deck-{chart_type}",
            title=f"{chart_type} chart",
            slides=[
                DeckSlide(
                    id=f"slide-{chart_type}",
                    layout_id="data_conclusion",
                    objects=[
                        TextObject(
                            id=f"title-{chart_type}",
                            role="title",
                            text="原生图表",
                            slot="title",
                        ),
                        ChartObject(
                            id=f"chart-{chart_type}",
                            chart_type=chart_type,
                            categories=["第一项", "第二项"],
                            series={"数值": [1.0, 2.0]},
                            unit="count",
                            derivation="direct fixture values",
                            source_refs=["fixture://native-chart-types"],
                            slot="chart",
                        ),
                    ],
                )
            ],
        )
        output = tmp_path / f"{chart_type}.pptx"

        PresentationCompiler().compile(deck, output)
        reopened = Presentation(output)

        assert any(shape.has_chart for shape in reopened.slides[0].shapes)
        assert inspect_pptx(deck, output).passed is True


def test_compiler_honors_explicit_native_chart_legend_policy(tmp_path) -> None:
    deck = DeckIR(
        id="deck-chart-legend",
        title="chart legend policy",
        slides=[
            DeckSlide(
                id="slide-chart-legend",
                layout_id="data_conclusion",
                objects=[
                    TextObject(
                        id="title-chart-legend",
                        role="title",
                        text="显式图例策略",
                        slot="title",
                    ),
                    ChartObject(
                        id="chart-legend",
                        chart_type="line",
                        categories=["Q1", "Q2"],
                        series={"收入": [1.0, 2.0], "成本": [0.5, 1.0]},
                        show_legend=False,
                        unit="万元",
                        derivation="direct fixture values",
                        source_refs=["fixture://chart-legend-policy"],
                        slot="chart",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "legend-policy.pptx"

    PresentationCompiler().compile(deck, output)
    reopened = Presentation(output)
    chart = next(
        shape.chart
        for shape in reopened.slides[0].shapes
        if shape.has_chart
    )

    assert chart.has_legend is False


def test_quality_fails_when_source_required_fact_is_missing(tmp_path) -> None:
    deck = DeckIR(
        id="deck-source-grounding",
        title="来源覆盖测试",
        source_refs=["source-report"],
        required_facts=["核心样本数为 11 人"],
        slides=[
            DeckSlide(
                id="slide-source-grounding",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id="source-grounding-title",
                        role="title",
                        text="核心结论",
                        slot="title",
                    ),
                    TextObject(
                        id="source-grounding-body",
                        role="body",
                        text="这是一段没有覆盖材料事实的占位文字。",
                        slot="body",
                    ),
                ],
            )
        ],
    )
    output = tmp_path / "missing-source-fact.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    assert report.content_complete is False
    assert any(
        issue.code == "content.required_fact_missing"
        for issue in report.issues
    )


def test_quality_reads_required_facts_from_native_table_cells(tmp_path) -> None:
    fact = "表格事实 31415"
    deck = _golden_deck()
    table_slide = deck.slides[-1]
    table_objects = [
        obj.model_copy(
            update={"rows": [["文字", fact], ["图表", "可编辑"]]},
        )
        if isinstance(obj, TableObject)
        else obj
        for obj in table_slide.objects
    ]
    deck = deck.model_copy(
        update={
            "required_facts": [fact],
            "slides": [
                *deck.slides[:-1],
                table_slide.model_copy(update={"objects": table_objects}),
            ],
        },
    )
    output = tmp_path / "table-source-fact.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is True, report.model_dump()
    assert not any(
        issue.code == "content.required_fact_missing"
        for issue in report.issues
    )


def test_quality_matches_source_facts_across_powerpoint_line_breaks(
    tmp_path,
) -> None:
    fact = "第一条结论\n第二条结论"
    deck = DeckIR(
        id="deck-source-newlines",
        title="换行覆盖测试",
        source_refs=["source-report"],
        required_facts=[fact],
        facts=[
            FactRecord(
                fact_id="fact-source-newlines",
                text=fact,
                source_refs=["source-report"],
            ),
        ],
        slides=[
            DeckSlide(
                id="slide-source-newlines",
                layout_id="key_judgment",
                source_fact_ids=["fact-source-newlines"],
                objects=[
                    TextObject(
                        id="source-newlines-title",
                        role="title",
                        text="结论",
                        slot="title",
                    ),
                    TextObject(
                        id="source-newlines-body",
                        role="body",
                        text=fact,
                        slot="body",
                    ),
                ],
            )
        ],
    )
    output = tmp_path / "source-newlines.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is True
    assert report.content_complete is True


def test_compiler_enforces_presentation_font_floor_and_visual_surfaces(
    tmp_path,
) -> None:
    deck = DeckIR(
        id="deck-type-scale",
        title="演示字号标准",
        slides=[
            DeckSlide(
                id="slide-type-scale",
                layout_id="three_cards",
                objects=[
                    TextObject(
                        id="type-title",
                        role="title",
                        text="三个关键判断",
                        slot="title",
                    ),
                    *[
                        TextObject(
                            id=f"type-card-{index}",
                            role="body",
                            text=f"判断 {index}\n用一句话说清核心证据",
                            slot=f"card_{index}",
                        )
                        for index in range(1, 4)
                    ],
                ],
            )
        ],
    )
    output = tmp_path / "type-scale.pptx"

    PresentationCompiler().compile(deck, output)
    slide = Presentation(output).slides[0]
    shapes = {shape.name: shape for shape in slide.shapes}

    assert shapes["type-title"].text_frame.paragraphs[0].font.size.pt >= 32
    assert shapes["type-card-1"].text_frame.paragraphs[0].font.size.pt >= 17
    assert "type-card-1-surface" in shapes
    assert inspect_pptx(deck, output).passed is True


def test_quality_rejects_demo_chart_measurement_units(tmp_path) -> None:
    deck = DeckIR(
        id="deck-placeholder-unit",
        title="占位单位门禁",
        source_refs=["fixture://chart-semantics"],
        slides=[
            DeckSlide(
                id="slide-placeholder-unit",
                layout_id="data_conclusion",
                objects=[
                    TextObject(
                        id="placeholder-title",
                        role="title",
                        text="准时交付率持续改善",
                        slot="title",
                    ),
                    ChartObject(
                        id="placeholder-chart",
                        chart_type="line",
                        categories=["七月", "八月", "九月"],
                        series={"准时交付率": [88.0, 90.0, 92.0]},
                        unit="benchmark score",
                        derivation="直接引用月度运营数据",
                        source_refs=["fixture://chart-semantics"],
                        slot="chart",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "placeholder-unit.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    assert any(issue.code == "chart.unit_placeholder" for issue in report.issues)


def test_quality_rejects_temporal_metric_labels_on_composition_slide(
    tmp_path,
) -> None:
    deck = DeckIR(
        id="deck-composition-labels",
        title="结构占比语义门禁",
        source_refs=["fixture://chart-semantics"],
        slides=[
            DeckSlide(
                id="slide-composition-labels",
                layout_id="data_conclusion",
                objects=[
                    TextObject(
                        id="composition-title",
                        role="title",
                        text="核心客群贡献主要续费收入",
                        slot="title",
                    ),
                    ChartObject(
                        id="composition-chart",
                        chart_type="bar",
                        categories=["核心客群", "成长客群", "其他客群"],
                        series={"续费收入占比": [68.0, 22.0, 10.0]},
                        unit="续费收入占比（%）",
                        derivation="直接引用客户结构分析",
                        source_refs=["fixture://chart-semantics"],
                        slot="chart",
                    ),
                    MetricObject(
                        id="composition-start",
                        label="起点",
                        value="68%",
                        slot="metric_1",
                    ),
                ],
                relationships=[
                    ObjectRelationship(
                        relationship_type=RelationshipType.COMPOSITION,
                        source_object_ids=["composition-chart"],
                        target_object_ids=["composition-start"],
                        label="客户收入构成",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "composition-labels.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    assert any(
        issue.code == "chart.metric_relationship_mismatch"
        for issue in report.issues
    )


def test_quality_rejects_repetitive_text_dump_deck(tmp_path) -> None:
    slides = []
    for index in range(1, 5):
        slides.append(
            DeckSlide(
                id=f"slide-dump-{index}",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id=f"dump-title-{index}",
                        role="title",
                        text=f"第 {index} 部分",
                        slot="title",
                    ),
                    TextObject(
                        id=f"dump-body-{index}",
                        role="body",
                        text=(
                            "这是直接从文档复制进幻灯片的长段落，"
                            "没有信息分层、视觉叙事或图形化表达。" * 5
                        ),
                        slot="body",
                    ),
                ],
            )
        )
    deck = DeckIR(
        id="deck-text-dump",
        title="文档搬运反例",
        slides=slides,
    )
    output = tmp_path / "text-dump.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    assert any(
        issue.code == "design.repetitive_text_dump"
        for issue in report.issues
    )


def test_quality_rejects_duplicate_content_hidden_in_a_new_layout(tmp_path) -> None:
    repeated_points = [
        "历史状态会覆盖人工修改",
        "写入失败缺少可追溯版本",
        "编辑与生成边界不清晰",
        "版本继承必须能够追踪",
    ]
    deck = DeckIR(
        id="deck-repeated-content",
        title="换壳重复反例",
        slides=[
            DeckSlide(
                id="slide-judgment-copy",
                layout_id="key_judgment",
                objects=[
                    TextObject(
                        id="copy-title-1",
                        role="title",
                        text="稳定性的三个约束",
                        slot="title",
                    ),
                    TextObject(
                        id="copy-root-1",
                        role="body",
                        text="多写者竞争是稳定性的共同根因",
                        slot="body",
                    ),
                    *[
                        TextObject(
                            id=f"copy-proof-{index}",
                            role="body",
                            text=value,
                            slot="body",
                        )
                        for index, value in enumerate(repeated_points, start=1)
                    ],
                ],
            ),
            DeckSlide(
                id="slide-hierarchy-copy",
                layout_id="hierarchy",
                objects=[
                    TextObject(
                        id="copy-title-2",
                        role="title",
                        text="同一组信息换成关系图",
                        slot="title",
                    ),
                    *[
                        TextObject(
                            id=f"copy-node-{index}",
                            role="body",
                            text=value,
                            slot=f"node_{index}",
                        )
                        for index, value in enumerate(repeated_points, start=1)
                    ],
                ],
            ),
        ],
    )
    output = tmp_path / "repeated-content.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is False
    issue = next(
        issue for issue in report.issues if issue.code == "design.repetition"
    )
    assert issue.slide_id == "slide-hierarchy-copy"
    assert "100% of the shorter slide" in issue.message


def test_compiler_uses_dark_cover_and_wide_hierarchy_grid(tmp_path) -> None:
    deck = DeckIR(
        id="deck-cover-hierarchy",
        title="投研汇报",
        theme_id="lantai-ink",
        slides=[
            DeckSlide(
                id="slide-cover-dark",
                layout_id="cover",
                objects=[
                    TextObject(
                        id="cover-dark-title",
                        role="title",
                        text="深度求索投资者交流会纪要",
                        slot="title",
                    ),
                    TextObject(
                        id="cover-dark-subtitle",
                        role="body",
                        text="2026 年 7 月 23 日 · 内部资料",
                        slot="subtitle",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-hierarchy-grid",
                layout_id="hierarchy",
                objects=[
                    TextObject(
                        id="hierarchy-title",
                        role="title",
                        text="AGI 技术路线图",
                        slot="title",
                    ),
                    TextObject(
                        id="hierarchy-root",
                        role="body",
                        text="语言模型 → CoT → Agent → 持续学习",
                        slot="root",
                    ),
                    *[
                        TextObject(
                            id=f"hierarchy-node-{index}",
                            role="body",
                            text=f"关键阶段 {index}\n用一句话解释该阶段的主要判断",
                            slot=f"node_{index}",
                        )
                        for index in range(1, 5)
                    ],
                ],
            ),
        ],
    )
    output = tmp_path / "dark-cover-grid.pptx"

    PresentationCompiler().compile(deck, output)
    reopened = Presentation(output)
    cover = reopened.slides[0]
    hierarchy = {shape.name: shape for shape in reopened.slides[1].shapes}

    assert str(cover.background.fill.fore_color.rgb) == "203A5F"
    assert str(
        next(
            shape
            for shape in cover.shapes
            if shape.name == "cover-dark-title"
        ).text_frame.paragraphs[0].font.color.rgb
    ) == "FFFFFF"
    assert hierarchy["hierarchy-node-1"].width.inches > 5.5
    assert hierarchy["hierarchy-node-3"].top.inches > hierarchy[
        "hierarchy-node-1"
    ].top.inches
    assert inspect_pptx(deck, output).passed is True


def test_quality_flags_one_layout_dominating_content_deck_for_visual_review(
    tmp_path,
) -> None:
    slides = [
        DeckSlide(
            id=f"slide-repeat-{index}",
            layout_id="two_column",
            objects=[
                TextObject(
                    id=f"repeat-title-{index}",
                    role="title",
                    text=f"主题 {index}",
                    slot="title",
                ),
                TextObject(
                    id=f"repeat-left-{index}",
                    role="body",
                    text="左侧核心信息",
                    slot="left",
                ),
                TextObject(
                    id=f"repeat-right-{index}",
                    role="body",
                    text="右侧核心信息",
                    slot="right",
                ),
            ],
        )
        for index in range(1, 6)
    ]
    deck = DeckIR(
        id="deck-repetitive-layout",
        title="重复版式反例",
        slides=slides,
    )
    output = tmp_path / "repetitive-layout.pptx"
    PresentationCompiler().compile(deck, output)

    report = inspect_pptx(deck, output)

    assert report.passed is True
    issue = next(
        issue
        for issue in report.issues
        if issue.code == "design.repetitive_layout"
    )
    assert issue.severity == "warning"


def test_compiler_stacks_metric_cards_and_marks_quote_surfaces(tmp_path) -> None:
    deck = DeckIR(
        id="deck-metrics-quotes",
        title="数字与金句",
        slides=[
            DeckSlide(
                id="slide-metrics",
                layout_id="three_cards",
                objects=[
                    TextObject(
                        id="metrics-title",
                        role="title",
                        text="关键数字",
                        slot="title",
                    ),
                    MetricObject(
                        id="metric-top",
                        label="当前算力：",
                        value="约两万张",
                        detail="H 等效算力",
                        slot="card_1",
                    ),
                    MetricObject(
                        id="metric-bottom",
                        label="华为卡：",
                        value="约一万六千张",
                        detail="只够训当前一代",
                        slot="card_1",
                    ),
                ],
            ),
            DeckSlide(
                id="slide-quotes",
                layout_id="matrix",
                objects=[
                    TextObject(
                        id="quotes-title",
                        role="title",
                        text="金句摘录",
                        slot="title",
                    ),
                    TextObject(
                        id="quote-one",
                        role="quote",
                        text="「你越克制，你越有可能能够做成。」",
                        slot="cell_1",
                    ),
                ],
            ),
        ],
    )
    output = tmp_path / "metrics-quotes.pptx"

    PresentationCompiler().compile(deck, output)
    reopened = Presentation(output)
    metric_shapes = {
        shape.name: shape for shape in reopened.slides[0].shapes
    }
    quote_shapes = {shape.name for shape in reopened.slides[1].shapes}

    assert metric_shapes["metric-top"].top < metric_shapes["metric-bottom"].top
    assert metric_shapes["metric-top"].height < 3 * 914400
    assert "quote-one-quote-mark" in quote_shapes
    assert inspect_pptx(deck, output).passed is True
