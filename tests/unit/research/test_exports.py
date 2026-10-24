import asyncio
import hashlib
from pathlib import Path

import pytest

from lantai.research.exports import export_report, export_file
from lantai.research.projects import ResearchService, RevisionConflict
from lantai.research.library import LibraryService
from lantai.research.review import review_report


@pytest.fixture
def prepared(tmp_path):
    path = tmp_path / 'wiki/raw/sources/观察研究.txt'
    path.parent.mkdir(parents=True)
    path.write_text('样本量为 80。仅发现相关性，未证明因果关系。', encoding='utf-8')
    service = ResearchService(tmp_path)
    project = service.create_project('方法比较与适用边界', [str(path.relative_to(tmp_path))], idempotency_key='project')
    source = service.source(project['source_versions'][0])
    text = '观察研究发现相关性，不能据此推断因果关系。'
    report = service.save_report(project['project_id'], 1, text, [{
        'start': 0, 'end': len(text), 'evidence': [{'source_version_id': source['id'],
            'start': 0, 'end': len(source['text']), 'quote': source['text']}],
    }], idempotency_key='report')
    return service, project, report, source


def test_draft_export_has_versioned_citations_and_byte_receipt(prepared):
    service, project, report, source = prepared
    result = export_report(service, project['project_id'], 1, report['id'], 'markdown', 'gb7714', 'export')
    path = export_file(service, result['id'])
    text = path.read_text()
    assert '待核验' in text and '核验完成' not in text
    assert '[1]' in text and '观察研究' in text and source['text'] in text
    assert result['report_hash'] == report['text_hash']
    assert result['sources'][0]['raw_hash'] == source['raw_hash']
    assert result['artifact_hash'] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result['citation_provenance']['engine'] == 'citeproc-js'
    assert export_report(service, project['project_id'], 1, report['id'], 'markdown', 'gb7714', 'export')['id'] == result['id']
    other = ResearchService(service.workspace / 'other')
    with pytest.raises(FileNotFoundError):
        export_file(other, result['id'])
    path.write_text('tampered')
    with pytest.raises(ValueError, match='integrity'):
        export_file(service, result['id'])


def test_review_and_metadata_changes_are_visible_without_rewriting_old_exports(prepared):
    service, p, r, source = prepared
    library = LibraryService(service.workspace)
    first = library.import_items([{'title': 'Observation', 'DOI': '10.1000/test', 'author': [{'family': 'Li'}],
                                  'issued': {'date-parts': [[2024]]}}], origin={'provider': 'user'})['items'][0]
    library.link_source(first['work_id'], source['id'])
    service.read_source(p['project_id'], 1, source['id'], 0)

    async def reviewer(_):
        return {'claims': [{'start': 0, 'end': len(r['text']), 'assessment': 'supports',
                           'reason': 'The report preserves the observational boundary.', 'evidence_indices': [0]}], 'issues': []}
    asyncio.run(review_report(service, p['project_id'], 1, r['id'], reviewer,
                              reviewer_identity={'provider': 'test', 'model': 'independent'}))
    original = export_report(service, p['project_id'], 1, r['id'], 'markdown', 'numeric', 'original')
    original_bytes = export_file(service, original['id']).read_bytes()
    assert original['status'] == 'reviewed'
    revised = library.import_items([{'title': 'Corrected Observation', 'DOI': '10.1000/test'}], origin={'provider': 'user'})['items'][0]
    library.select_version(first['work_id'], revised['version_id'], first['version_id'])
    assert service.project(p['project_id'])['review_status'] == 'metadata_changed'
    changed = export_report(service, p['project_id'], 1, r['id'], 'markdown', 'numeric', 'changed')
    assert changed['status'] == 'metadata_changed'
    assert changed['bibliography_versions'][0]['version_id'] == first['version_id']
    assert export_file(service, original['id']).read_bytes() == original_bytes
    new_report = service.save_report(p['project_id'], 1, '新论点。', [], expected_report_id=r['id'], idempotency_key='edit')
    with pytest.raises(RevisionConflict):
        export_report(service, p['project_id'], 1, r['id'], 'markdown', 'numeric', 'stale')
    assert new_report['id'] != r['id']


def test_chinese_long_report_exports_as_editable_docx_and_paginated_pdf(prepared):
    service, p, r, _ = prepared
    text = '# 研究结果\n\n' + '\n\n'.join(f'第 {i} 项讨论：观察研究的样本、方法与局限需要分别核对。' for i in range(180))
    report = service.save_report(p['project_id'], 1, text, [], expected_report_id=r['id'], idempotency_key='long')
    from docx import Document
    from pypdf import PdfReader
    for format in ('docx', 'pdf'):
        value = export_report(service, p['project_id'], 1, report['id'], format, 'gb7714', format)
        path = export_file(service, value['id'])
        if format == 'docx':
            document = Document(path)
            full = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            assert '第 179 项讨论' in full and '待核验' in full
            assert any(paragraph.style.name.startswith('Heading') for paragraph in document.paragraphs)
        else:
            document = PdfReader(path)
            assert len(document.pages) > 3
            full = '\n'.join(page.extract_text() for page in document.pages)
            assert '第 179 项讨论' in full and '待核验' in full


def test_adopting_bibliography_keeps_history_and_requires_a_new_review(prepared):
    service, p, report, source = prepared
    library = LibraryService(service.workspace)
    first = library.import_items([{'title': 'Original Observation', 'DOI': '10.1000/adoption'}],
                                 origin={'provider': 'user'})['items'][0]
    library.link_source(first['work_id'], source['id'])
    service.read_source(p['project_id'], 1, source['id'], 0)
    original = export_report(service, p['project_id'], 1, report['id'], 'markdown', 'numeric', 'before-adoption')
    original_bytes = export_file(service, original['id']).read_bytes()
    selected = library.import_items([{'title': 'Corrected Observation', 'DOI': '10.1000/adoption'}],
                                    origin={'provider': 'user'})['items'][0]
    library.select_version(first['work_id'], selected['version_id'], first['version_id'])
    updated = service.adopt_bibliography(p['project_id'], 1, idempotency_key='adopt')
    assert updated['revision'] == 2 and updated['review_id'] is None
    assert updated['review_status'] == 'not_reviewed'
    current = service.project(p['project_id'])
    assert current['stale_bibliography'] == []
    assert current['reading']['read_blocks'] == 1
    revised_report = service.report(updated['report_id'])
    assert revised_report['text_hash'] == report['text_hash']
    assert revised_report['previous_report_id'] == report['id']
    assert revised_report['claims'][0]['evidence'][0]['source_version_id'] != source['id']
    result = export_report(service, p['project_id'], 2, revised_report['id'], 'markdown', 'numeric', 'after-adoption')
    assert result['bibliography_versions'][0]['version_id'] == selected['version_id']
    assert 'Corrected Observation' in export_file(service, result['id']).read_text()
    assert export_file(service, original['id']).read_bytes() == original_bytes
    assert service.source(source['id']) == source
    assert service.adopt_bibliography(p['project_id'], 1, idempotency_key='adopt')['id'] == updated['id']
    with pytest.raises(RevisionConflict):
        service.adopt_bibliography(p['project_id'], 1, idempotency_key='stale-adopt')


def test_chinese_exports_embed_their_font_without_viewer_font_dependencies(prepared):
    from io import BytesIO
    from zipfile import ZipFile
    from lxml import etree
    from pypdf import PdfReader
    service, p, r, _ = prepared
    value = export_report(service, p['project_id'], 1, r['id'], 'docx', 'gb7714', 'embedded-docx')
    with ZipFile(export_file(service, value['id'])) as archive:
        fonts = etree.fromstring(archive.read('word/fontTable.xml'))
        assert fonts.xpath('//*[local-name()="embedRegular"]')
        assert any(name.endswith('.odttf') for name in archive.namelist())
    value = export_report(service, p['project_id'], 1, r['id'], 'pdf', 'gb7714', 'embedded-pdf')
    pdf = PdfReader(BytesIO(export_file(service, value['id']).read_bytes()))
    descriptors = [font.get_object().get('/FontDescriptor')
                   for font in pdf.pages[0]['/Resources']['/Font'].values()]
    assert any(d and '/FontFile2' in d.get_object() for d in descriptors)


def test_word_citations_disclose_static_text_without_live_zotero_fields(prepared):
    from docx import Document
    service, p, report, _ = prepared
    value = export_report(service, p['project_id'], 1, report['id'], 'docx', 'gb7714', 'static')
    document = Document(export_file(service, value['id']))
    assert '引用为静态文本' in '\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert value['citation_mode'] == 'static_text'


def test_markdown_table_exports_real_cells_and_repeating_headers(prepared):
    from docx import Document
    from pypdf import PdfReader
    service, p, original, _ = prepared
    body = '## 结果比较\n\n| 编号 | 观察结果 | 限制 |\n| --- | --- | --- |\n'
    body += '\n'.join(f'| {i} | 观察到相关性 | 未排除混杂，不能推断因果 |' for i in range(100))
    report = service.save_report(p['project_id'], 1, body, [], expected_report_id=original['id'], idempotency_key='table')
    value = export_report(service, p['project_id'], 1, report['id'], 'docx', 'gb7714', 'table-docx')
    document = Document(export_file(service, value['id']))
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 101
    assert document.tables[0].rows[-1].cells[0].text == '99'
    assert document.tables[0].rows[0]._tr.xpath('./w:trPr/w:tblHeader')
    assert all(row._tr.xpath('./w:trPr/w:cantSplit') for row in document.tables[0].rows)
    value = export_report(service, p['project_id'], 1, report['id'], 'pdf', 'gb7714', 'table-pdf')
    pdf = PdfReader(export_file(service, value['id']))
    assert len(pdf.pages) >= 3
    table_pages = [page.extract_text() for page in pdf.pages if '未排除混杂' in page.extract_text()]
    assert len(table_pages) >= 3 and all('观察结果' in text for text in table_pages)


def test_table_export_preserves_escaped_pipes_code_blocks_and_malformed_rows():
    from io import BytesIO
    from docx import Document
    from lantai.research.exports import _docx
    text = ('```markdown\n| example | only |\n| --- | --- |\n```\n'
            '| 方法 | 限制 |\n| --- | --- |\n| A\\|B | C |\n'
            '| 有额外列 | 内容 | 不得丢失 |\n')
    document = Document(BytesIO(_docx(text)))
    assert len(document.tables) == 1
    assert document.tables[0].rows[1].cells[0].text == 'A|B'
    prose = '\n'.join(p.text for p in document.paragraphs)
    assert '| example | only |' in prose and '不得丢失' in prose


def test_single_table_row_longer_than_a_page_remains_readable():
    from io import BytesIO
    from pypdf import PdfReader
    from lantai.research.exports import _pdf
    body = '逐项描述证据和限制。' * 350 + '最终限制不可丢失。'
    pdf = PdfReader(BytesIO(_pdf('| 结果 | 限制 |\n| --- | --- |\n| 观察 | ' + body + ' |')))
    extracted = '\n'.join(p.extract_text() for p in pdf.pages)
    assert len(pdf.pages) > 1 and '最终限制不可丢失' in extracted
    assert all('结果' in p.extract_text() for p in pdf.pages)


def test_export_receipt_binds_citation_nodes_and_discloses_missing_fields(prepared):
    service, project, report, source = prepared
    value = export_report(service, project['project_id'], 1, report['id'], 'markdown', 'apa', 'audit')
    audit = value['citation_audit']
    assert audit['status'] == 'consistent'
    assert audit['nodes'][0]['report_end'] == len(report['text'])
    assert audit['nodes'][0]['bibliography_ids'] == [source['id']]
    assert audit['metadata_warnings'][0]['missing_fields'] == ['author', 'issued']
    text = export_file(service, value['id']).read_text()
    assert '引用元数据待补充' in text and '作者、出版日期' in text


def test_figure_markup_cannot_silently_claim_embedded_images(prepared):
    service, project, original, _ = prepared
    body = '结果见图。\n\n![图一](figure.png)'
    report = service.save_report(project['project_id'], 1, body, [], expected_report_id=original['id'], idempotency_key='figure')
    value = export_report(service, project['project_id'], 1, report['id'], 'docx', 'numeric', 'figure')
    from docx import Document
    text = '\n'.join(p.text for p in Document(export_file(service, value['id'])).paragraphs)
    assert 'Word/PDF 未嵌入图像' in text
    assert value['citation_audit']['figure_rendering'] == 'markup_only'
