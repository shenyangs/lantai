# -*- coding: utf-8 -*-
"""DSG-02 worker: keep a qualitative comparison as one native table."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from time import perf_counter
from xml.etree import ElementTree

from pypdf import PdfReader
from pptx import Presentation

from lantai.app.office_conversion import convert_office_document, find_soffice
from lantai.presentation.design_intelligence import (
    compose_deck_ir_with_design_context,
)
from lantai.presentation.models import (
    DeckIR,
    DeckSlide,
    FactRecord,
    PresentationBrief,
    PresentationJobRequest,
    RelationshipType,
    TableObject,
    TextObject,
)
from lantai.presentation.pdf_render import render_pdf_pages
from lantai.presentation.pptx_compiler import NativePptxCompiler


ORIGINAL_DECK = "dsg02-original-deck.json"
REQUEST_FILE = "dsg02-request.json"
BRIEF_FILE = "dsg02-brief.json"
COMPOSED_DECK = "dsg02-composed-deck.json"
COMPILED_PPTX = "dsg02-qualitative-comparison.pptx"
CONVERTED_PDF = "dsg02-qualitative-comparison.pdf"
TABLE_ID = "comparison-table"
REPEATED_HEADERS = ("截图能确认的", "澜台已确认的", "融合裁决")
HEADERS = ["维度", *REPEATED_HEADERS]
ROWS = [
    ["问题组织", "输入被展开为讨论题", "合同有目标和验收", "放第一层"],
    ["分工可见", "三个视角直接呈现", "成员有角色与依赖", "具体问题在前"],
    ["交叉讨论", "截图不足以确认", "有第二轮审阅", "保留真执行"],
    ["结论交付", "截屏未显示", "子结果与交付区分", "最终答案直接可读"],
    ["过程可控", "单张截图无法验证", "有停止暂停恢复", "不为换外观破坏保护"],
]


def _paths(state_root: Path) -> dict[str, Path]:
    return {
        "original_deck": state_root / ORIGINAL_DECK,
        "request": state_root / REQUEST_FILE,
        "brief": state_root / BRIEF_FILE,
        "composed_deck": state_root / COMPOSED_DECK,
        "pptx": state_root / "output" / COMPILED_PPTX,
        "pdf": state_root / "office" / CONVERTED_PDF,
        "rendered": state_root / "office" / "rendered",
    }


def _sha256(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _write_model(
    path: Path,
    model: DeckIR | PresentationJobRequest | PresentationBrief,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(model.model_dump_json(indent=2) + "\n", encoding="utf-8")


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _models() -> tuple[DeckIR, PresentationJobRequest, PresentationBrief]:
    facts = [
        FactRecord(
            fact_id=f"fact-comparison-{index:02d}",
            text="｜".join(row),
            source_refs=[f"source:comparison-row-{index:02d}"],
        )
        for index, row in enumerate(ROWS, start=1)
    ]
    deck = DeckIR(
        id="deck-dsg02-qualitative-table",
        title="截图与澜台：到底差在哪里",
        source_refs=[ref for fact in facts for ref in fact.source_refs],
        required_facts=[row[0] for row in ROWS],
        facts=facts,
        slides=[
            DeckSlide(
                id="slide-dsg02-comparison",
                layout_id="table",
                speaker_notes=(
                    "DSG02-NOTE：定性证据必须保持原始四列语义，"
                    "禁止重复列标题生成五组台账卡片。"
                ),
                source_fact_ids=[fact.fact_id for fact in facts],
                objects=[
                    TextObject(
                        id="comparison-table-title",
                        role="title",
                        text="截图能看到什么，澜台真正确认了什么",
                        slot="title",
                    ),
                    TableObject(id=TABLE_ID, headers=HEADERS, rows=ROWS),
                ],
            ),
        ],
    )
    request = PresentationJobRequest(
        conversation_id="dsg02-qualitative-table",
        user_request=(
            "为产品评审制作一页克制、清晰的定性对照页。"
            "保留四列比较关系与全部五条事实，原生对象必须可编辑。"
        ),
        audience="产品与工程评审人",
        slide_count=1,
        seed=17,
    )
    brief = PresentationBrief(
        objective="让评审人一眼区分可见现象、已确认事实与融合裁决",
        audience="产品与工程评审人",
        narrative="先给判断，再以一个定性矩阵呈现五维证据",
        required_facts=[row[0] for row in ROWS],
        source_refs=deck.source_refs,
        style_constraints=["克制", "高信息层级", "原生表格可编辑"],
        prohibited_changes=["不得把每行展开为重复列标题卡片"],
    )
    return deck, request, brief


def _phase_prepare(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    deck, request, brief = _models()
    _write_model(paths["original_deck"], deck)
    _write_model(paths["request"], request)
    _write_model(paths["brief"], brief)
    return {
        "phase": "prepare",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "slide_count": len(deck.slides),
        "headers": HEADERS,
        "rows": ROWS,
        "source_fact_ids": deck.slides[0].source_fact_ids,
        "object_ids": [obj.id for obj in deck.slides[0].objects],
    }


def _phase_compose(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    deck = DeckIR.model_validate_json(paths["original_deck"].read_text("utf-8"))
    request = PresentationJobRequest.model_validate_json(
        paths["request"].read_text("utf-8")
    )
    brief = PresentationBrief.model_validate_json(
        paths["brief"].read_text("utf-8")
    )
    composed = compose_deck_ir_with_design_context(deck, request, brief)
    slide = composed.slides[0]
    native_tables = [obj for obj in slide.objects if isinstance(obj, TableObject)]
    repeated_cards = [
        obj.id
        for obj in slide.objects
        if isinstance(obj, TextObject)
        and (obj.id == TABLE_ID or obj.id.startswith(f"{TABLE_ID}-row-"))
    ]
    comparison_relationships = [
        item.relationship_type.value
        for item in slide.relationships
        if item.relationship_type == RelationshipType.COMPARISON
    ]
    if len(native_tables) != 1:
        raise RuntimeError("DSG-02 composition did not retain exactly one table")
    table = native_tables[0]
    if table.headers != HEADERS or table.rows != ROWS:
        raise RuntimeError("DSG-02 composition changed qualitative table content")
    if table.frame is None:
        raise RuntimeError("DSG-02 composition did not measure the native table")
    if repeated_cards:
        raise RuntimeError(f"DSG-02 repeated ledger cards returned: {repeated_cards}")
    if not comparison_relationships:
        raise RuntimeError("DSG-02 comparison relationship was not retained")
    _write_model(paths["composed_deck"], composed)
    compilation = NativePptxCompiler().compile(composed, paths["pptx"])
    return {
        "phase": "compose",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "native_table_count": len(native_tables),
        "repeated_card_ids": repeated_cards,
        "comparison_relationships": comparison_relationships,
        "table_frame": table.frame.model_dump(mode="json"),
        "source_fact_ids": slide.source_fact_ids,
        "compiled_slide_count": compilation.slide_count,
        "editable_object_count": compilation.editable_object_count,
        "degradations": compilation.degradations,
        "pptx_sha256": _sha256(paths["pptx"]),
    }


def _table_matrix(table) -> list[list[str]]:
    return [[cell.text for cell in row.cells] for row in table.rows]


def _libreoffice_version() -> str:
    completed = subprocess.run(
        [str(find_soffice()), "--version"],
        capture_output=True,
        text=True,
        timeout=30,
        check=True,
    )
    return (completed.stdout or completed.stderr).strip()


def _phase_observe(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    original = DeckIR.model_validate_json(paths["original_deck"].read_text("utf-8"))
    composed = DeckIR.model_validate_json(paths["composed_deck"].read_text("utf-8"))
    presentation = Presentation(str(paths["pptx"]))
    if len(presentation.slides) != 1:
        raise RuntimeError("DSG-02 PPTX page count changed")
    slide = presentation.slides[0]
    native_table_shapes = [shape for shape in slide.shapes if shape.has_table]
    if len(native_table_shapes) != 1:
        raise RuntimeError("DSG-02 PPTX does not contain exactly one native table")
    table_shape = native_table_shapes[0]
    matrix = _table_matrix(table_shape.table)
    expected_matrix = [HEADERS, *ROWS]
    if matrix != expected_matrix:
        raise RuntimeError("DSG-02 PPTX native cells differ from the source matrix")
    non_table_texts = [
        shape.text
        for shape in slide.shapes
        if not shape.has_table and getattr(shape, "has_text_frame", False)
    ]
    all_text = "\n".join([*non_table_texts, *(cell for row in matrix for cell in row)])
    header_occurrences = {
        header: all_text.count(header) for header in REPEATED_HEADERS
    }
    if any(count != 1 for count in header_occurrences.values()):
        raise RuntimeError(
            f"DSG-02 repeated headers detected in PPTX: {header_occurrences}"
        )
    repeated_shape_names = [
        shape.name
        for shape in slide.shapes
        if shape.name.startswith(f"{TABLE_ID}-row-")
    ]
    if repeated_shape_names:
        raise RuntimeError(
            f"DSG-02 repeated row shapes detected: {repeated_shape_names}"
        )
    with zipfile.ZipFile(paths["pptx"]) as package:
        slide_xml = package.read("ppt/slides/slide1.xml")
    xml_root = ElementTree.fromstring(slide_xml)
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    serialized_table_count = len(xml_root.findall(f".//{{{drawing_ns}}}tbl"))
    if serialized_table_count != 1:
        raise RuntimeError("DSG-02 package contains an unexpected table count")

    converted = convert_office_document(
        paths["pptx"],
        "pdf",
        paths["pdf"].parent,
        output_name=paths["pdf"].name,
    )
    pdf_page_count = len(PdfReader(str(converted.output_path)).pages)
    rendered_pages = render_pdf_pages(
        converted.output_path,
        paths["rendered"],
        dpi=144,
        prefix="page-",
        zero_pad=2,
        image_format="PNG",
    )
    return {
        "phase": "observe",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "deck_page_count": len(composed.slides),
        "pptx_page_count": len(presentation.slides),
        "pdf_page_count": pdf_page_count,
        "rendered_page_count": len(rendered_pages),
        "native_table_shape_count": len(native_table_shapes),
        "native_table_shape_name": table_shape.name,
        "serialized_table_count": serialized_table_count,
        "cell_matrix": matrix,
        "header_occurrences": header_occurrences,
        "repeated_shape_names": repeated_shape_names,
        "source_fact_ids_before": original.slides[0].source_fact_ids,
        "source_fact_ids_after": composed.slides[0].source_fact_ids,
        "speaker_notes": slide.notes_slide.notes_text_frame.text,
        "shape_names": [shape.name for shape in slide.shapes],
        "pptx_sha256": _sha256(paths["pptx"]),
        "pdf_sha256": _sha256(paths["pdf"]),
        "pdf_converter_sha256": converted.sha256,
        "rendered_page_sha256": [_sha256(path) for path in rendered_pages],
        "libreoffice_version": _libreoffice_version(),
    }


def _run_phase(state_root: Path, phase: str) -> dict[str, object]:
    if phase == "prepare":
        return _phase_prepare(state_root)
    if phase == "compose":
        return _phase_compose(state_root)
    if phase == "observe":
        return _phase_observe(state_root)
    raise ValueError(f"unknown DSG-02 phase: {phase}")


def _invoke_phase(state_root: Path, phase: str) -> dict[str, object]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--state-root",
            str(state_root),
            "--phase",
            phase,
        ],
        env=environment,
        capture_output=True,
        text=True,
        timeout=240,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"DSG-02 {phase} failed: stdout={completed.stdout!r}; "
            f"stderr={completed.stderr!r}"
        )
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _copy_evidence(source: Path, destination: Path) -> dict[str, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return {"path": destination.name, "sha256": _sha256(destination)}


def _write_evidence_bundle(
    state_root: Path,
    result: dict[str, object],
    evidence_dir: Path,
) -> None:
    paths = _paths(state_root)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    sources = {
        "original_deck": paths["original_deck"],
        "request": paths["request"],
        "brief": paths["brief"],
        "composed_deck": paths["composed_deck"],
        "native_pptx": paths["pptx"],
        "libreoffice_pdf": paths["pdf"],
    }
    result["artifact_files"] = {
        label: _copy_evidence(source, evidence_dir / source.name)
        for label, source in sources.items()
    }
    result["rendered_page_files"] = [
        _copy_evidence(page, evidence_dir / f"dsg02-{page.name}")
        for page in sorted(paths["rendered"].glob("page-*.png"))
    ]
    _write_json(
        evidence_dir / "dsg02-qualitative-table-cross-boundary.json",
        result,
    )


def run(state_root: Path, evidence_dir: Path | None = None) -> dict[str, object]:
    state_root.mkdir(parents=True, exist_ok=True)
    phase_names = ("prepare", "compose", "observe")
    phases = {phase: _invoke_phase(state_root, phase) for phase in phase_names}
    result: dict[str, object] = {
        "matrix_id": "DSG-02",
        "minimum_evidence": "cross_boundary",
        "orchestrator_pid": os.getpid(),
        "child_pids": [int(phases[phase]["pid"]) for phase in phase_names],
        "phases": phases,
        "total_phase_seconds": round(
            sum(float(phases[phase]["duration_seconds"]) for phase in phase_names),
            6,
        ),
    }
    if evidence_dir is not None:
        _write_evidence_bundle(state_root, result, evidence_dir)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--phase", choices=("prepare", "compose", "observe"))
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    result = (
        _run_phase(args.state_root, args.phase)
        if args.phase is not None
        else run(args.state_root, args.evidence_dir)
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
