# -*- coding: utf-8 -*-
"""Independent QA-01 worker for incomplete 19-page review evidence."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

from lantai.presentation import service as service_module
from lantai.presentation.jobs import PresentationJobStore
from lantai.presentation.models import (
    CompositionAuthority,
    CompositionEngine,
    DeckIR,
    DeckSlide,
    EvidenceState,
    InputMaterial,
    ObjectFrame,
    PlanningMode,
    PresentationJobRequest,
    QualityReport,
    TextObject,
)
from lantai.presentation.quality import is_verified_delivery
from lantai.presentation.render import RenderResult
from lantai.presentation.service import PresentationService


CONVERSATION_ID = "conversation-qa01-full-page-coverage"
RUNTIME_RUN_ID = "native-run-qa01-full-page-coverage"


def _deck() -> DeckIR:
    return DeckIR(
        id="qa01-full-page-coverage-deck",
        title="QA-01 全页覆盖边界",
        slides=[
            DeckSlide(
                id=f"qa01-slide-{index:02d}",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=f"第 {index} 页的独立核验说明。",
                objects=[
                    TextObject(
                        id=f"qa01-title-{index:02d}",
                        role="title",
                        text=f"第 {index} 页判断",
                        slot="title",
                        frame=ObjectFrame(x=0.8, y=0.7, w=9.0, h=0.8),
                    ),
                    TextObject(
                        id=f"qa01-body-{index:02d}",
                        role="body",
                        text=f"这是只属于第 {index} 页的可编辑证据对象。",
                        slot="body",
                        frame=ObjectFrame(x=0.9, y=2.0, w=7.0, h=1.2),
                    ),
                ],
            )
            for index in range(1, 20)
        ],
    )


def _request() -> PresentationJobRequest:
    return PresentationJobRequest(
        conversation_id=CONVERSATION_ID,
        user_request="生成并逐页核验 19 页演示文稿",
        slide_count=19,
        vision_critique_required=True,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        composition_engine=CompositionEngine.PPTXGENJS_V3,
        materials=[
            InputMaterial(
                id="qa01-synthetic-source",
                name="qa01-synthetic-source.md",
                media_type="text/markdown",
                content="合成 QA 材料：19 页必须全部核验，不能抽查后宣称通过。",
            ),
        ],
    )


def _partial_three_page_render(
    source: str | Path,
    output_dir: str | Path,
    *,
    timeout: int = 120,
) -> RenderResult:
    _ = timeout
    source_slide_count = len(Presentation(str(source)).slides)
    if source_slide_count != 19:
        raise AssertionError(f"QA-01 expected 19 slides, got {source_slide_count}")
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"QA-01 deliberately incomplete render set\n")
    images: list[Path] = []
    for index in range(1, 4):
        image_path = destination / f"slide-{index}.png"
        image = Image.new("RGB", (640, 360), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, 180 + index * 20, 360), fill=(25, 25, 25))
        image.save(image_path)
        images.append(image_path)
    return RenderResult(
        pdf_path=pdf_path,
        slide_images=tuple(images),
        renderer_id="qa01-incomplete-renderer",
        renderer_version="1.0",
    )


class _ForbiddenCritic:
    def __init__(self) -> None:
        self.calls = 0

    def cache_identity(self) -> str:
        return "qa01-forbidden-critic/v1"

    def review(self, _slides):
        self.calls += 1
        raise AssertionError("critic must not run on an incomplete render set")


def _job_paths(state_root: Path, job_id: str) -> dict[str, Path]:
    job_dir = PresentationJobStore(state_root).job_dir(job_id)
    return {
        "job_dir": job_dir,
        "deck": job_dir / "checkpoints" / "authored-deck-ir.json",
        "report": job_dir / "checkpoints" / "qa01-partial-quality.json",
        "candidate": job_dir / "candidates" / "qa01-19-slide.pptx",
        "render_dir": job_dir / "renders" / "partial-three",
        "authoritative": job_dir / "output" / "presentation.pptx",
    }


def _execute(state_root: Path) -> dict[str, object]:
    critic = _ForbiddenCritic()
    service = PresentationService(state_root, visual_critic=critic)
    record = service.create_job(_request(), runtime_run_id=RUNTIME_RUN_ID)
    paths = _job_paths(state_root, record.job_id)
    deck = _deck()
    paths["candidate"].parent.mkdir(parents=True, exist_ok=True)
    service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
        deck,
        paths["candidate"],
    )
    service.store.write_checkpoint(
        record.job_id,
        "authored-deck-ir.json",
        deck.model_dump_json(indent=2),
    )
    original_render = service_module.render_pptx
    service_module.render_pptx = _partial_three_page_render
    try:
        report, rendered = service.verify_artifact(
            deck,
            paths["candidate"],
            paths["render_dir"],
            visual_required=True,
            vision_critique_required=True,
            render_set_id="qa01-partial-three",
            job_id=record.job_id,
        )
    finally:
        service_module.render_pptx = original_render
    service.store.write_checkpoint(
        record.job_id,
        "qa01-partial-quality.json",
        report.model_dump_json(indent=2),
    )
    return {
        "execute_pid": os.getpid(),
        "job_id": record.job_id,
        "slide_count": len(deck.slides),
        "rendered_image_count": len(rendered),
        "critic_calls": critic.calls,
        "report_passed": report.passed,
        "delivery_status": report.delivery_status,
        "render_verified": report.render_verified,
        "visual_verified": report.visual_verified,
        "render_evidence": report.render_evidence.value,
        "visual_evidence": report.visual_evidence.value,
        "blocking_gaps": report.blocking_gaps,
        "issue_codes": sorted({issue.code for issue in report.issues}),
        "actual_delivery_verified": is_verified_delivery(
            deck,
            report,
            paths["candidate"],
            list(rendered),
        ),
        "authoritative_output_exists": paths["authoritative"].exists(),
    }


def _observe(state_root: Path, job_id: str) -> dict[str, object]:
    store = PresentationJobStore(state_root)
    paths = _job_paths(state_root, job_id)
    deck = DeckIR.model_validate_json(paths["deck"].read_text(encoding="utf-8"))
    report = QualityReport.model_validate_json(
        paths["report"].read_text(encoding="utf-8"),
    )
    rendered = tuple(sorted(paths["render_dir"].glob("slide-*.png")))
    forged = report.model_copy(
        update={
            "passed": True,
            "delivery_status": "passed",
            "rendered_slide_count": len(deck.slides),
            "render_verified": True,
            "visual_verified": True,
            "issues": [],
            "content_evidence": EvidenceState.PASSED,
            "structure_evidence": EvidenceState.PASSED,
            "render_evidence": EvidenceState.PASSED,
            "visual_evidence": EvidenceState.PASSED,
            "editability_evidence": EvidenceState.PASSED,
            "provenance_evidence": EvidenceState.PASSED,
            "blocking_gaps": [],
        },
    )
    current = store.get(job_id)
    presentation = Presentation(str(paths["candidate"]))
    native_shape_count = sum(
        len(slide.shapes) for slide in presentation.slides
    )
    return {
        "observer_pid": os.getpid(),
        "job_id": job_id,
        "job_dir_count": len(list(state_root.glob("presentation-*"))),
        "manifest_count": len(
            list(state_root.glob("presentation-*/manifest.json")),
        ),
        "job_status": current.status.value,
        "completed_event_count": sum(
            event.status.value == "completed" for event in current.events
        ),
        "slide_count": len(deck.slides),
        "rendered_image_count": len(rendered),
        "persisted_render_evidence_count": len(report.rendered_slides),
        "actual_delivery_verified": is_verified_delivery(
            deck,
            report,
            paths["candidate"],
            list(rendered),
        ),
        "forged_green_flags_delivery_verified": is_verified_delivery(
            deck,
            forged,
            paths["candidate"],
            list(rendered),
        ),
        "candidate_exists": paths["candidate"].is_file(),
        "candidate_pptx_slide_count": len(presentation.slides),
        "candidate_native_shape_count": native_shape_count,
        "authoritative_output_exists": paths["authoritative"].exists(),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("execute", "observe"), required=True)
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--job-id")
    args = parser.parse_args()
    args.state_root.mkdir(parents=True, exist_ok=True)
    if args.mode == "execute":
        result = _execute(args.state_root)
    else:
        if not args.job_id:
            parser.error("--job-id is required for observe mode")
        result = _observe(args.state_root, args.job_id)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
