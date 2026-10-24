# -*- coding: utf-8 -*-
"""Independent QA-06 worker for shared-font visual cache invalidation."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import threading
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pptx import Presentation

from lantai.presentation import service as service_module
from lantai.presentation.jobs import PresentationJobStore
from lantai.presentation.models import (
    CompositionAuthority,
    CompositionEngine,
    DeckIR,
    DeckSlide,
    ObjectFrame,
    PlanningMode,
    PresentationJobRequest,
    TextObject,
    VisualCriticResult,
)
from lantai.presentation.quality import file_sha256
from lantai.presentation.render import RenderResult
from lantai.presentation.service import PresentationService
from lantai.presentation.theme import THEMES, get_theme


CONVERSATION_ID = "conversation-qa06-shared-font"
RUNTIME_RUN_ID = "native-run-qa06-shared-font"
CRITIC_IDENTITY = "qa06-controlled-critic/v1"
THEME_ID = "lantai-warm"
REPLACEMENT_BODY_FONT = "Source Han Sans SC"


def _deck() -> DeckIR:
    slides: list[DeckSlide] = []
    for index in range(1, 6):
        slides.append(
            DeckSlide(
                id=f"qa06-slide-{index:02d}",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=f"第 {index} 页来源说明保持稳定。",
                objects=[
                    TextObject(
                        id=f"qa06-title-{index:02d}",
                        role="title",
                        text=f"第 {index} 页共享字体检查",
                        slot="title",
                        frame=ObjectFrame(x=0.8, y=0.7, w=9.0, h=0.8),
                    ),
                    TextObject(
                        id=f"qa06-body-{index:02d}",
                        role="body",
                        text="页面内容与像素保持稳定，只替换主题字体依赖。",
                        slot="body",
                        frame=ObjectFrame(x=0.9, y=2.0, w=8.2, h=1.2),
                    ),
                ],
            ),
        )
    return DeckIR(
        id="qa06-shared-font-deck",
        title="QA-06 共享字体依赖",
        theme_id=THEME_ID,
        slides=slides,
    )


def _request() -> PresentationJobRequest:
    return PresentationJobRequest(
        conversation_id=CONVERSATION_ID,
        user_request="生成五页原生演示文稿，并验证共享主题字体变化会重审全部页面",
        slide_count=5,
        theme_id=THEME_ID,
        cost_budget_units=30,
        vision_critique_required=True,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        composition_engine=CompositionEngine.PPTXGENJS_V3,
    )


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _deterministic_render(
    source: str | Path,
    output_dir: str | Path,
    *,
    timeout: int = 120,
) -> RenderResult:
    """Render only from slide text so theme-token changes cannot fake a miss."""

    _ = timeout
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(str(source))
    images: list[Path] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = _slide_text(slide)
        digest = hashlib.sha256(text.encode("utf-8")).digest()
        image = Image.new("RGB", (640, 360), (digest[0], digest[1], digest[2]))
        draw = ImageDraw.Draw(image)
        draw.rectangle((40, 40, 600, 320), outline="white", width=4)
        draw.text((64, 64), f"slide {index}", fill="white")
        image_path = destination / f"slide-{index}.png"
        image.save(image_path)
        images.append(image_path)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"QA-06 deterministic render evidence\n")
    return RenderResult(
        pdf_path=pdf_path,
        slide_images=tuple(images),
        renderer_id="qa06-deterministic-renderer",
        renderer_version="1.0",
    )


class _RecordingCritic:
    def __init__(self) -> None:
        self.slide_calls: list[str] = []
        self.deck_calls = 0
        self._lock = threading.Lock()

    def cache_identity(self) -> str:
        return CRITIC_IDENTITY

    def review_slide(self, slide) -> VisualCriticResult:
        with self._lock:
            self.slide_calls.append(slide.slide_id)
        return VisualCriticResult(critic_id=CRITIC_IDENTITY, issues=[])

    def review_deck(self, _slides) -> VisualCriticResult:
        with self._lock:
            self.deck_calls += 1
        return VisualCriticResult(critic_id=CRITIC_IDENTITY, issues=[])


def _job_paths(state_root: Path, job_id: str) -> dict[str, Path]:
    job_dir = PresentationJobStore(state_root).job_dir(job_id)
    return {
        "job_dir": job_dir,
        "deck": job_dir / "checkpoints" / "qa06-deck.json",
        "candidate_seed": job_dir / "candidates" / "qa06-seed.pptx",
        "candidate_font": job_dir / "candidates" / "qa06-font.pptx",
        "render_seed": job_dir / "renders" / "qa06-seed",
        "render_font": job_dir / "renders" / "qa06-font",
        "quality": job_dir / "quality" / "visual-checkpoints",
        "seed_summary": job_dir / "checkpoints" / "qa06-seed-summary.json",
    }


def _checkpoint_hashes(path: Path) -> dict[str, str]:
    return {
        item.name: file_sha256(item)
        for item in sorted(path.glob("*.json"))
    }


def _slide_xml_hashes(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {
            f"slide-{index:02d}": hashlib.sha256(
                archive.read(f"ppt/slides/slide{index}.xml"),
            ).hexdigest()
            for index in range(1, 6)
        }


def _evidence_fingerprints(report) -> dict[str, str]:
    return {
        item.slide_id if item.scope == "slide" else "deck": (
            item.dependency_fingerprint
        )
        for item in report.visual_review_evidence
    }


def _verify(
    service: PresentationService,
    deck: DeckIR,
    candidate: Path,
    render_dir: Path,
    job_id: str,
):
    original_render = service_module.render_pptx
    service_module.render_pptx = _deterministic_render
    try:
        return service.verify_artifact(
            deck,
            candidate,
            render_dir,
            visual_required=True,
            vision_critique_required=True,
            job_id=job_id,
        )
    finally:
        service_module.render_pptx = original_render


def _seed(state_root: Path) -> dict[str, object]:
    critic = _RecordingCritic()
    service = PresentationService(state_root, visual_critic=critic)
    record = service.create_job(_request(), runtime_run_id=RUNTIME_RUN_ID)
    paths = _job_paths(state_root, record.job_id)
    deck = _deck()
    paths["candidate_seed"].parent.mkdir(parents=True, exist_ok=True)
    service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
        deck,
        paths["candidate_seed"],
    )
    service.store.write_checkpoint(
        record.job_id,
        "qa06-deck.json",
        deck.model_dump_json(indent=2),
    )
    report, rendered = _verify(
        service,
        deck,
        paths["candidate_seed"],
        paths["render_seed"],
        record.job_id,
    )
    summary = {
        "checkpoint_hashes": _checkpoint_hashes(paths["quality"]),
        "deck_ir_sha256": hashlib.sha256(
            deck.model_dump_json().encode("utf-8"),
        ).hexdigest(),
        "dependency_fingerprints": _evidence_fingerprints(report),
        "font_manifest": report.font_manifest,
        "render_hashes": [file_sha256(path) for path in rendered],
        "slide_xml_hashes": _slide_xml_hashes(paths["candidate_seed"]),
        "theme_body_font": get_theme(THEME_ID).body_font,
    }
    service.store.write_checkpoint(
        record.job_id,
        "qa06-seed-summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True),
    )
    return {
        "seed_pid": os.getpid(),
        "job_id": record.job_id,
        "slide_calls": sorted(critic.slide_calls),
        "deck_calls": critic.deck_calls,
        "report_passed": report.passed,
        "rendered_slide_count": len(rendered),
        "visual_evidence_count": len(report.visual_review_evidence),
        "font_manifest": report.font_manifest,
        "theme_body_font": get_theme(THEME_ID).body_font,
        "cost_used_units": service.store.get(record.job_id).cost_used_units,
    }


def _replace_font_and_verify(state_root: Path, job_id: str) -> dict[str, object]:
    original_theme = THEMES[THEME_ID]
    THEMES[THEME_ID] = replace(
        original_theme,
        body_font=REPLACEMENT_BODY_FONT,
    )
    try:
        critic = _RecordingCritic()
        service = PresentationService(state_root, visual_critic=critic)
        paths = _job_paths(state_root, job_id)
        seed = json.loads(paths["seed_summary"].read_text(encoding="utf-8"))
        persisted_deck = DeckIR.model_validate_json(
            paths["deck"].read_text(encoding="utf-8"),
        )
        current_deck = _deck()
        current_deck_sha256 = hashlib.sha256(
            current_deck.model_dump_json().encode("utf-8"),
        ).hexdigest()
        service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
            current_deck,
            paths["candidate_font"],
        )
        report, rendered = _verify(
            service,
            current_deck,
            paths["candidate_font"],
            paths["render_font"],
            job_id,
        )
        current_checkpoints = _checkpoint_hashes(paths["quality"])
        prior_checkpoints = dict(seed["checkpoint_hashes"])
        changed_checkpoints = sorted(
            name
            for name in current_checkpoints
            if current_checkpoints[name] != prior_checkpoints.get(name)
        )
        expected_checkpoints = sorted(
            [
                "slide-"
                + hashlib.sha256(slide.id.encode("utf-8")).hexdigest()
                + ".json"
                for slide in current_deck.slides
            ]
            + ["deck.json"],
        )
        current_slide_xml = _slide_xml_hashes(paths["candidate_font"])
        changed_slide_xml = sorted(
            name
            for name, digest in current_slide_xml.items()
            if digest != seed["slide_xml_hashes"].get(name)
        )
        current_render_hashes = [file_sha256(path) for path in rendered]
        changed_render_indexes = [
            index
            for index, (before, after) in enumerate(
                zip(seed["render_hashes"], current_render_hashes, strict=True),
                start=1,
            )
            if before != after
        ]
        fingerprints = _evidence_fingerprints(report)
        changed_fingerprint_scopes = sorted(
            scope
            for scope, fingerprint in fingerprints.items()
            if fingerprint != seed["dependency_fingerprints"].get(scope)
        )
        slide_evidence = {
            item.slide_id: item
            for item in report.visual_review_evidence
            if item.scope == "slide" and item.slide_id is not None
        }
        deck_evidence = [
            item for item in report.visual_review_evidence if item.scope == "deck"
        ]
        render_hash_by_slide_id = {
            slide.id: current_render_hashes[index]
            for index, slide in enumerate(current_deck.slides)
        }
        rebound_matches = {
            slide_id: evidence.rendered_sha256 == render_hash_by_slide_id[slide_id]
            for slide_id, evidence in slide_evidence.items()
        }
        record = service.store.get(job_id)
        return {
            "modify_pid": os.getpid(),
            "job_id": job_id,
            "deck_ir_unchanged": (
                current_deck_sha256 == seed["deck_ir_sha256"]
                and persisted_deck.model_dump(mode="json")
                == current_deck.model_dump(mode="json")
            ),
            "theme_body_font_before": seed["theme_body_font"],
            "theme_body_font_after": get_theme(THEME_ID).body_font,
            "font_manifest_before": seed["font_manifest"],
            "font_manifest_after": report.font_manifest,
            "changed_slide_xml": changed_slide_xml,
            "changed_render_indexes": changed_render_indexes,
            "slide_calls": sorted(critic.slide_calls),
            "deck_calls": critic.deck_calls,
            "cached_slide_ids": sorted(
                slide_id
                for slide_id, evidence in slide_evidence.items()
                if evidence.cache_hit
            ),
            "dirty_slide_ids": sorted(
                slide_id
                for slide_id, evidence in slide_evidence.items()
                if not evidence.cache_hit
            ),
            "deck_cache_hit": deck_evidence[0].cache_hit,
            "changed_fingerprint_scopes": changed_fingerprint_scopes,
            "changed_checkpoints": changed_checkpoints,
            "expected_changed_checkpoints": expected_checkpoints,
            "all_evidence_rebound_to_current_render": all(
                rebound_matches.values()
            ),
            "rebound_matches": rebound_matches,
            "report_passed": report.passed,
            "rendered_slide_count": len(rendered),
            "visual_evidence_count": len(report.visual_review_evidence),
            "cost_used_units": record.cost_used_units,
            "job_dir_count": len(list(state_root.glob("presentation-*"))),
            "manifest_count": len(
                list(state_root.glob("presentation-*/manifest.json")),
            ),
        }
    finally:
        THEMES[THEME_ID] = original_theme


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--mode",
        choices=("seed", "replace-font"),
        required=True,
    )
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--job-id")
    args = parser.parse_args()
    args.state_root.mkdir(parents=True, exist_ok=True)
    if args.mode == "seed":
        result = _seed(args.state_root)
    else:
        if not args.job_id:
            parser.error("--job-id is required for replace-font mode")
        result = _replace_font_and_verify(args.state_root, args.job_id)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
