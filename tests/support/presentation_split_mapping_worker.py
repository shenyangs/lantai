# -*- coding: utf-8 -*-
"""EDIT-05 worker: prove an approved split stays synchronized end to end."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from time import perf_counter

from pypdf import PdfReader
from pptx import Presentation

from lantai.app.office_conversion import convert_office_document
from lantai.presentation.models import (
    CompositionAuthority,
    DeckIR,
    DeckPlan,
    DeckPlanSlide,
    DeckSlide,
    FactRecord,
    ObjectFrame,
    QualityIssue,
    SlideIntent,
    TextObject,
    VisualDirection,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.pdf_render import render_pdf_pages
from lantai.presentation.repair import (
    RepairActionType,
    SlideRedesignProposal,
    repair_deck_with_designer,
)


ORIGINAL_DECK = "edit05-original-deck.json"
ORIGINAL_PLAN = "edit05-original-plan.json"
REPAIRED_DECK = "edit05-repaired-deck.json"
REPAIRED_PLAN = "edit05-repaired-plan.json"
COMPILED_PPTX = "edit05-approved-split.pptx"
CONVERTED_PDF = "edit05-approved-split.pdf"

JUDGMENT_NOTE = "EDIT05-NOTE-JUDGMENT：先讲清 95% 的核心判断。"
EVIDENCE_NOTE = "EDIT05-NOTE-EVIDENCE：再说明 240 份样本的证据边界。"
UNCHANGED_NOTE = "EDIT05-NOTE-ACTION：最后给出下一步行动。"


def _paths(state_root: Path) -> dict[str, Path]:
    return {
        "original_deck": state_root / ORIGINAL_DECK,
        "original_plan": state_root / ORIGINAL_PLAN,
        "repaired_deck": state_root / REPAIRED_DECK,
        "repaired_plan": state_root / REPAIRED_PLAN,
        "pptx": state_root / "output" / COMPILED_PPTX,
        "pdf": state_root / "office" / CONVERTED_PDF,
        "rendered": state_root / "office" / "rendered",
    }


def _sha256(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _write_model(path: Path, model: DeckIR | DeckPlan) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(model.model_dump_json(indent=2) + "\n", encoding="utf-8")


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _visual_direction() -> VisualDirection:
    return VisualDirection(
        visual_concept="editorial evidence split",
        palette=["#F7F4EE", "#171717", "#2F6B5F"],
        typography="clear editorial hierarchy",
        image_language="documentary",
        chart_language="direct labels",
        shape_language="quiet rules",
        grid="twelve columns",
        recurring_motif="evidence folio",
    )


def _text(
    object_id: str,
    role: str,
    text: str,
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    font_size: float,
) -> TextObject:
    return TextObject(
        id=object_id,
        role=role,
        text=text,
        font_size=font_size,
        frame=ObjectFrame(x=x, y=y, w=w, h=h),
    )


def _original_models() -> tuple[DeckIR, DeckPlan]:
    judgment_intent = SlideIntent(
        id="intent-judgment",
        purpose="prove",
        key_message="转化率判断必须与样本边界一起解释",
        source_refs=["source:judgment", "source:evidence"],
        required_facts=["95%", "240 份"],
        immutable_facts=["95%", "240"],
    )
    action_intent = SlideIntent(
        id="intent-action",
        purpose="decide",
        key_message="批准下一阶段验证",
        source_refs=["source:decision"],
        required_facts=["下一阶段"],
        immutable_facts=[],
    )
    plan = DeckPlan(
        title="EDIT-05 拆页映射",
        narrative_arc="判断—证据—行动",
        visual_direction=_visual_direction(),
        slides=[
            DeckPlanSlide(
                id="plan-judgment",
                order=1,
                title="核心判断与证据边界",
                intent=judgment_intent,
                content_blocks=["转化率达到 95%", "样本 240 份"],
                takeaway="结论成立，但必须保留样本边界。",
            ),
            DeckPlanSlide(
                id="plan-action",
                order=2,
                title="下一步行动",
                intent=action_intent,
                content_blocks=["批准下一阶段验证"],
                takeaway="进入下一阶段。",
            ),
        ],
    )
    deck = DeckIR(
        id="deck-edit05",
        title="EDIT-05 拆页映射",
        immutable_data=["95%", "240"],
        required_facts=["转化率达到 95%", "样本 240 份"],
        source_refs=["source:judgment", "source:evidence", "source:decision"],
        visual_direction=plan.visual_direction,
        facts=[
            FactRecord(
                fact_id="fact-judgment",
                text="转化率达到 95%",
                source_refs=["source:judgment"],
                immutable=True,
            ),
            FactRecord(
                fact_id="fact-evidence",
                text="样本 240 份",
                source_refs=["source:evidence"],
                immutable=True,
            ),
            FactRecord(
                fact_id="fact-decision",
                text="批准下一阶段验证",
                source_refs=["source:decision"],
            ),
        ],
        slides=[
            DeckSlide(
                id="slide-judgment",
                intent_id="intent-judgment",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes="原页同时解释判断与证据边界。",
                source_fact_ids=["fact-judgment", "fact-evidence"],
                objects=[
                    _text(
                        "original-title",
                        "title",
                        "核心判断与证据边界",
                        x=0.9,
                        y=0.7,
                        w=11.3,
                        h=0.8,
                        font_size=30,
                    ),
                    _text(
                        "original-body",
                        "body",
                        "转化率达到 95%，样本 240 份。",
                        x=0.9,
                        y=2.0,
                        w=10.8,
                        h=1.4,
                        font_size=24,
                    ),
                ],
            ),
            DeckSlide(
                id="slide-action",
                intent_id="intent-action",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=UNCHANGED_NOTE,
                source_fact_ids=["fact-decision"],
                objects=[
                    _text(
                        "action-title",
                        "title",
                        "下一步行动",
                        x=0.9,
                        y=0.7,
                        w=11.3,
                        h=0.8,
                        font_size=30,
                    ),
                    _text(
                        "action-body",
                        "body",
                        "批准下一阶段验证",
                        x=0.9,
                        y=2.0,
                        w=10.8,
                        h=1.4,
                        font_size=24,
                    ),
                ],
            ),
        ],
    )
    return deck, plan


class _ApprovedSplitDesigner:
    def redesign(self, request):
        if request.planned_slide is None:
            raise RuntimeError("EDIT-05 split requires its authoritative plan page")
        evidence_intent = SlideIntent(
            id="intent-evidence",
            purpose="prove",
            key_message="240 份样本界定证据边界",
            source_refs=["source:evidence"],
            required_facts=["240 份"],
            immutable_facts=["240"],
        )
        first_intent = request.planned_slide.intent.model_copy(
            update={
                "key_message": "转化率达到 95%",
                "source_refs": ["source:judgment"],
                "required_facts": ["95%"],
                "immutable_facts": ["95%"],
            }
        )
        first = DeckSlide(
            id=request.slide.id,
            intent_id=request.slide.intent_id,
            layout_id="freeform",
            composition_authority=CompositionAuthority.AUTHOR,
            speaker_notes=JUDGMENT_NOTE,
            source_fact_ids=["fact-judgment"],
            objects=[
                _text(
                    "judgment-title",
                    "title",
                    "核心判断",
                    x=0.9,
                    y=0.7,
                    w=11.3,
                    h=0.8,
                    font_size=30,
                ),
                _text(
                    "judgment-metric",
                    "body",
                    "转化率达到 95%",
                    x=0.9,
                    y=2.0,
                    w=10.8,
                    h=1.4,
                    font_size=34,
                ),
            ],
        )
        second = DeckSlide(
            id="slide-evidence",
            intent_id=evidence_intent.id,
            layout_id="freeform",
            composition_authority=CompositionAuthority.AUTHOR,
            speaker_notes=EVIDENCE_NOTE,
            source_fact_ids=["fact-evidence"],
            objects=[
                _text(
                    "evidence-title",
                    "title",
                    "证据边界",
                    x=0.9,
                    y=0.7,
                    w=11.3,
                    h=0.8,
                    font_size=30,
                ),
                _text(
                    "evidence-body",
                    "body",
                    "样本 240 份",
                    x=0.9,
                    y=2.0,
                    w=10.8,
                    h=1.4,
                    font_size=34,
                ),
            ],
        )
        return SlideRedesignProposal(
            slide=first,
            action=RepairActionType.SPLIT_SLIDE,
            rationale="用户批准把判断与证据拆成两页",
            additional_slides=(second,),
            plan_slides=(
                request.planned_slide.model_copy(
                    update={
                        "title": "核心判断",
                        "intent": first_intent,
                        "content_blocks": ["转化率达到 95%"],
                        "takeaway": "转化率达到 95%。",
                    }
                ),
                DeckPlanSlide(
                    id="plan-evidence",
                    order=request.planned_slide.order + 1,
                    title="证据边界",
                    intent=evidence_intent,
                    content_blocks=["样本 240 份"],
                    takeaway="样本边界必须与判断同时保留。",
                ),
            ),
        )


def _issue() -> QualityIssue:
    return QualityIssue(
        code="content.density",
        severity="error",
        slide_id="slide-judgment",
        message="判断与证据边界需要拆成两个可读页面",
        repairable=True,
    )


def _phase_prepare(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    deck, plan = _original_models()
    _write_model(paths["original_deck"], deck)
    _write_model(paths["original_plan"], plan)
    return {
        "phase": "prepare",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "deck_slide_ids": [slide.id for slide in deck.slides],
        "plan_slide_ids": [slide.id for slide in plan.slides],
        "page_count": len(deck.slides),
        "source_refs": plan.slides[0].intent.source_refs,
        "source_fact_ids": deck.slides[0].source_fact_ids,
    }


def _phase_split(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    deck = DeckIR.model_validate_json(paths["original_deck"].read_text("utf-8"))
    plan = DeckPlan.model_validate_json(paths["original_plan"].read_text("utf-8"))
    designer = _ApprovedSplitDesigner()
    rejection = ""
    try:
        repair_deck_with_designer(
            deck,
            plan,
            [_issue()],
            designer,
            allow_structural_repair=False,
            slide_count_ceiling=3,
        )
    except ValueError as exc:
        rejection = str(exc)
    if not rejection:
        raise RuntimeError("EDIT-05 split proceeded without explicit approval")

    repaired = repair_deck_with_designer(
        deck,
        plan,
        [_issue()],
        designer,
        allow_structural_repair=True,
        slide_count_ceiling=3,
    )
    if repaired.plan is None:
        raise RuntimeError("EDIT-05 approved split lost the DeckPlan")
    _write_model(paths["repaired_deck"], repaired.deck)
    _write_model(paths["repaired_plan"], repaired.plan)
    compilation = NativePptxCompiler().compile(repaired.deck, paths["pptx"])
    return {
        "phase": "split",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "unapproved_split_rejected": True,
        "rejection": rejection,
        "actions": list(repaired.actions),
        "deck_slide_ids": [slide.id for slide in repaired.deck.slides],
        "plan_slide_ids": [slide.id for slide in repaired.plan.slides],
        "intent_ids": [slide.intent.id for slide in repaired.plan.slides],
        "deck_page_count": len(repaired.deck.slides),
        "plan_page_count": len(repaired.plan.slides),
        "compiled_page_count": compilation.slide_count,
        "editable_object_count": compilation.editable_object_count,
        "pptx_sha256": _sha256(paths["pptx"]),
    }


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _slide_notes(slide) -> str:
    return slide.notes_slide.notes_text_frame.text


def _phase_observe(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    original_deck = DeckIR.model_validate_json(
        paths["original_deck"].read_text("utf-8")
    )
    original_plan = DeckPlan.model_validate_json(
        paths["original_plan"].read_text("utf-8")
    )
    repaired_deck = DeckIR.model_validate_json(
        paths["repaired_deck"].read_text("utf-8")
    )
    repaired_plan = DeckPlan.model_validate_json(
        paths["repaired_plan"].read_text("utf-8")
    )
    presentation = Presentation(str(paths["pptx"]))
    slide_texts = [_slide_text(slide) for slide in presentation.slides]
    slide_notes = [_slide_notes(slide) for slide in presentation.slides]
    converted = convert_office_document(
        paths["pptx"],
        "pdf",
        paths["pdf"].parent,
        output_name=paths["pdf"].name,
    )
    pdf_page_count = len(PdfReader(str(converted.output_path)).pages)
    rendered_pages = render_pdf_pages(
        converted.output_path,
        paths["rendered"],
        dpi=144,
        prefix="page-",
        zero_pad=2,
        image_format="PNG",
    )
    with zipfile.ZipFile(paths["pptx"]) as package:
        package_members = set(package.namelist())

    original_split_plan = original_plan.slides[0]
    split_plan = repaired_plan.slides[:2]
    original_split_deck = original_deck.slides[0]
    split_deck = repaired_deck.slides[:2]
    return {
        "phase": "observe",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "deck_slide_ids": [slide.id for slide in repaired_deck.slides],
        "plan_slide_ids": [slide.id for slide in repaired_plan.slides],
        "plan_orders": [slide.order for slide in repaired_plan.slides],
        "intent_ids": [slide.intent.id for slide in repaired_plan.slides],
        "deck_page_count": len(repaired_deck.slides),
        "plan_page_count": len(repaired_plan.slides),
        "pptx_page_count": len(presentation.slides),
        "pdf_page_count": pdf_page_count,
        "rendered_page_count": len(rendered_pages),
        "rendered_page_sha256": [_sha256(path) for path in rendered_pages],
        "slide_texts": slide_texts,
        "slide_notes": slide_notes,
        "source_refs_before": original_split_plan.intent.source_refs,
        "source_refs_after": [
            ref for slide in split_plan for ref in slide.intent.source_refs
        ],
        "required_facts_before": original_split_plan.intent.required_facts,
        "required_facts_after": [
            fact for slide in split_plan for fact in slide.intent.required_facts
        ],
        "source_fact_ids_before": original_split_deck.source_fact_ids,
        "source_fact_ids_after": [
            fact_id for slide in split_deck for fact_id in slide.source_fact_ids
        ],
        "notes_are_distinct": len(set(slide_notes[:2])) == 2,
        "unchanged_slide_preserved": (
            repaired_deck.slides[2] == original_deck.slides[1]
        ),
        "pptx_sha256": _sha256(paths["pptx"]),
        "pdf_sha256": _sha256(paths["pdf"]),
        "pdf_converter_sha256": converted.sha256,
        "has_slide_id_markers": all(
            f"ppt/slides/slide{index}.xml" in package_members
            for index in range(1, 4)
        ),
        "has_notes_parts": all(
            f"ppt/notesSlides/notesSlide{index}.xml" in package_members
            for index in range(1, 4)
        ),
    }


def _run_phase(state_root: Path, phase: str) -> dict[str, object]:
    if phase == "prepare":
        return _phase_prepare(state_root)
    if phase == "split":
        return _phase_split(state_root)
    if phase == "observe":
        return _phase_observe(state_root)
    raise ValueError(f"unknown EDIT-05 phase: {phase}")


def _invoke_phase(state_root: Path, phase: str) -> dict[str, object]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--state-root",
            str(state_root),
            "--phase",
            phase,
        ],
        env=environment,
        capture_output=True,
        text=True,
        timeout=240,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"EDIT-05 {phase} failed: stdout={completed.stdout!r}; "
            f"stderr={completed.stderr!r}"
        )
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _copy_evidence(source: Path, destination: Path) -> dict[str, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return {"path": destination.name, "sha256": _sha256(destination)}


def _write_evidence_bundle(
    state_root: Path,
    result: dict[str, object],
    evidence_dir: Path,
) -> None:
    paths = _paths(state_root)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    sources = {
        "original_deck": paths["original_deck"],
        "original_plan": paths["original_plan"],
        "repaired_deck": paths["repaired_deck"],
        "repaired_plan": paths["repaired_plan"],
        "approved_split_pptx": paths["pptx"],
        "libreoffice_pdf": paths["pdf"],
    }
    result["artifact_files"] = {
        label: _copy_evidence(source, evidence_dir / source.name)
        for label, source in sources.items()
    }
    result["rendered_page_files"] = [
        _copy_evidence(page, evidence_dir / f"edit05-{page.name}")
        for page in sorted(paths["rendered"].glob("page-*.png"))
    ]
    _write_json(evidence_dir / "edit05-split-mapping-cross-boundary.json", result)


def run(state_root: Path, evidence_dir: Path | None = None) -> dict[str, object]:
    state_root.mkdir(parents=True, exist_ok=True)
    phase_names = ("prepare", "split", "observe")
    phases = {phase: _invoke_phase(state_root, phase) for phase in phase_names}
    result: dict[str, object] = {
        "matrix_id": "EDIT-05",
        "minimum_evidence": "cross_boundary",
        "orchestrator_pid": os.getpid(),
        "child_pids": [int(phases[phase]["pid"]) for phase in phase_names],
        "phases": phases,
        "total_phase_seconds": round(
            sum(float(phases[phase]["duration_seconds"]) for phase in phase_names),
            6,
        ),
    }
    if evidence_dir is not None:
        _write_evidence_bundle(state_root, result, evidence_dir)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--phase", choices=("prepare", "split", "observe"))
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    result = (
        _run_phase(args.state_root, args.phase)
        if args.phase is not None
        else run(args.state_root, args.evidence_dir)
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
