# -*- coding: utf-8 -*-
"""EDIT-04 worker: reject a v1-based edit after v2 becomes authoritative."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from pptx import Presentation

from lantai.presentation.models import (
    CompositionAuthority,
    DeckIR,
    DeckSlide,
    ObjectFrame,
    TextObject,
)
from lantai.presentation.pptx_compiler import NativePptxCompiler
from lantai.presentation.quality import file_sha256
from lantai.presentation.versioning import (
    PresentationVersionStore,
    StalePresentationVersionError,
)


LEDGER_DIR_NAME = "presentation-edit04-stale-baseline"
OUTPUT_NAME = "presentation.pptx"
STALE_CANDIDATE_NAME = "prepared-from-v1.pptx"
CONTROL_NAME = "prepared-v1.json"


def _deck(title: str, *, revision: int) -> DeckIR:
    return DeckIR(
        id="edit04-versioned-deck",
        title=title,
        slides=[
            DeckSlide(
                id="edit04-slide-01",
                layout_id="freeform",
                revision=revision,
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes="EDIT-04 版本冲突验证；旧基准不得覆盖新版本。",
                objects=[
                    TextObject(
                        id="edit04-title",
                        role="title",
                        slot="title",
                        text=title,
                        frame=ObjectFrame(x=0.9, y=0.8, w=10.8, h=1.0),
                    ),
                    TextObject(
                        id="edit04-body",
                        role="body",
                        slot="body",
                        text=f"受管版本：v{revision}",
                        frame=ObjectFrame(x=0.9, y=2.2, w=5.2, h=0.8),
                    ),
                ],
            ),
        ],
    )


def _paths(state_root: Path) -> dict[str, Path]:
    ledger = state_root / LEDGER_DIR_NAME
    return {
        "ledger": ledger,
        "output": ledger / "output" / OUTPUT_NAME,
        "stale_candidate": ledger / "candidates" / STALE_CANDIDATE_NAME,
        "control": ledger / "checkpoints" / CONTROL_NAME,
    }


def _native_title(path: Path) -> str:
    presentation = Presentation(str(path))
    for shape in presentation.slides[0].shapes:
        if shape.name == "edit04-title" and getattr(shape, "has_text_frame", False):
            return shape.text
    raise RuntimeError(f"EDIT-04 title object missing from {path}")


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _read_control(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def _phase_prepare(state_root: Path) -> dict[str, object]:
    paths = _paths(state_root)
    paths["output"].parent.mkdir(parents=True, exist_ok=True)
    paths["stale_candidate"].parent.mkdir(parents=True, exist_ok=True)
    compiler = NativePptxCompiler()
    deck_v1 = _deck("权威版本 v1", revision=1)
    stale_deck = _deck("基于 v1 准备但被拒绝的修改", revision=2)
    compiler.compile(deck_v1, paths["output"])
    store = PresentationVersionStore(paths["ledger"])
    version_v1 = store.bootstrap(
        paths["output"],
        deck_v1,
        writer="pptxgenjs_v3",
    )
    compiler.compile(stale_deck, paths["stale_candidate"])
    control = {
        "base_version_id": version_v1.version_id,
        "base_pptx_sha256": version_v1.pptx_sha256,
        "stale_candidate_sha256": file_sha256(paths["stale_candidate"]),
    }
    _write_json(paths["control"], control)
    return {
        "phase": "prepare",
        "pid": os.getpid(),
        "version_id": version_v1.version_id,
        "output_sha256": version_v1.pptx_sha256,
        "stale_candidate_sha256": control["stale_candidate_sha256"],
        "output_title": _native_title(paths["output"]),
        "stale_candidate_title": _native_title(paths["stale_candidate"]),
        "history_count": len(store.history()),
    }


def _phase_advance(state_root: Path) -> dict[str, object]:
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    store = PresentationVersionStore(paths["ledger"])
    version_v1 = store.current()
    if version_v1 is None or version_v1.version_id != control["base_version_id"]:
        raise RuntimeError("EDIT-04 v1 base is not authoritative before advance")
    deck_v2 = _deck("权威版本 v2", revision=2)
    NativePptxCompiler().compile(deck_v2, paths["output"])
    version_v2 = store.record(
        paths["output"],
        deck_v2,
        writer="lantai_ai_edit",
        parent=version_v1,
        generation=version_v1.generation + 1,
        base_pptx_sha256=version_v1.pptx_sha256,
        expected_pptx_sha256=file_sha256(paths["output"]),
    )
    return {
        "phase": "advance",
        "pid": os.getpid(),
        "version_id": version_v2.version_id,
        "parent_version_id": version_v2.parent_version_id,
        "output_sha256": version_v2.pptx_sha256,
        "output_title": _native_title(paths["output"]),
        "history_count": len(store.history()),
    }


def _phase_stale_attempt(state_root: Path) -> dict[str, object]:
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    store = PresentationVersionStore(paths["ledger"])
    current_before = store.current()
    if current_before is None:
        raise RuntimeError("EDIT-04 authoritative v2 is missing")
    history_before = store.history()
    output_sha_before = file_sha256(paths["output"])
    output_bytes_before = paths["output"].read_bytes()
    artifact_bundles_before = sorted(
        path.name
        for path in (store.root / "artifacts").iterdir()
        if path.is_dir() and not path.name.startswith(".")
    )

    base_error = ""
    try:
        store.assert_writer_base(
            paths["output"],
            _deck("权威版本 v1", revision=1),
            expected_pptx_sha256=str(control["base_pptx_sha256"]),
        )
    except StalePresentationVersionError as exc:
        base_error = str(exc)
    if not base_error:
        raise AssertionError("EDIT-04 stale v1 base was accepted")

    parent_error = ""
    stale_parent = store.get(str(control["base_version_id"]))
    try:
        store.record(
            paths["stale_candidate"],
            _deck("基于 v1 准备但被拒绝的修改", revision=2),
            writer="lantai_ai_edit",
            parent=stale_parent,
            generation=current_before.generation + 1,
            base_pptx_sha256=str(control["base_pptx_sha256"]),
            expected_pptx_sha256=str(control["stale_candidate_sha256"]),
        )
    except StalePresentationVersionError as exc:
        parent_error = str(exc)
    if not parent_error:
        raise AssertionError("EDIT-04 stale parent was accepted")

    current_after = store.current()
    history_after = store.history()
    artifact_bundles_after = sorted(
        path.name
        for path in (store.root / "artifacts").iterdir()
        if path.is_dir() and not path.name.startswith(".")
    )
    staging_entries = sorted(
        path.name
        for path in (store.root / "artifacts").iterdir()
        if path.name.startswith(".")
    )
    return {
        "phase": "stale_attempt",
        "pid": os.getpid(),
        "base_error": base_error,
        "parent_error": parent_error,
        "authoritative_version_before": current_before.version_id,
        "authoritative_version_after": (
            current_after.version_id if current_after else None
        ),
        "authoritative_sha256_before": output_sha_before,
        "authoritative_sha256_after": file_sha256(paths["output"]),
        "authoritative_bytes_unchanged": (
            paths["output"].read_bytes() == output_bytes_before
        ),
        "authoritative_title": _native_title(paths["output"]),
        "stale_candidate_sha256": file_sha256(paths["stale_candidate"]),
        "stale_candidate_title": _native_title(paths["stale_candidate"]),
        "history_count_before": len(history_before),
        "history_count_after": len(history_after),
        "artifact_bundles_before": artifact_bundles_before,
        "artifact_bundles_after": artifact_bundles_after,
        "staging_entries": staging_entries,
    }


def _run_phase(state_root: Path, phase: str) -> dict[str, object]:
    if phase == "prepare":
        return _phase_prepare(state_root)
    if phase == "advance":
        return _phase_advance(state_root)
    if phase == "stale_attempt":
        return _phase_stale_attempt(state_root)
    raise ValueError(f"unknown EDIT-04 phase: {phase}")


def _invoke_phase(state_root: Path, phase: str) -> dict[str, object]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--state-root",
            str(state_root),
            "--phase",
            phase,
        ],
        env=environment,
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"EDIT-04 {phase} failed: stdout={completed.stdout!r}; "
            f"stderr={completed.stderr!r}"
        )
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _write_evidence_bundle(
    state_root: Path,
    result: dict[str, object],
    evidence_dir: Path,
) -> None:
    paths = _paths(state_root)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "authoritative_v2": (
            paths["output"],
            evidence_dir / "edit04-authoritative-v2.pptx",
        ),
        "rejected_stale_candidate": (
            paths["stale_candidate"],
            evidence_dir / "edit04-rejected-v1-edit.pptx",
        ),
        "version_current": (
            PresentationVersionStore(paths["ledger"]).current_path,
            evidence_dir / "edit04-version-current.json",
        ),
        "version_history": (
            PresentationVersionStore(paths["ledger"]).history_path,
            evidence_dir / "edit04-version-history.jsonl",
        ),
    }
    artifacts: dict[str, dict[str, str]] = {}
    for label, (source, destination) in files.items():
        shutil.copy2(source, destination)
        artifacts[label] = {
            "path": destination.name,
            "sha256": file_sha256(destination),
        }
    result["artifact_files"] = artifacts
    evidence_path = evidence_dir / "edit04-stale-baseline-cross-boundary.json"
    _write_json(evidence_path, result)


def run(state_root: Path, evidence_dir: Path | None = None) -> dict[str, object]:
    state_root.mkdir(parents=True, exist_ok=True)
    phases = {
        phase: _invoke_phase(state_root, phase)
        for phase in ("prepare", "advance", "stale_attempt")
    }
    result: dict[str, object] = {
        "matrix_id": "EDIT-04",
        "minimum_evidence": "cross_boundary",
        "orchestrator_pid": os.getpid(),
        "child_pids": [
            int(phases[phase]["pid"])
            for phase in ("prepare", "advance", "stale_attempt")
        ],
        "phases": phases,
        "ledger_count": len(list(state_root.glob(LEDGER_DIR_NAME))),
    }
    if evidence_dir is not None:
        _write_evidence_bundle(state_root, result, evidence_dir)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument(
        "--phase",
        choices=("prepare", "advance", "stale_attempt"),
    )
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    result = (
        _run_phase(args.state_root, args.phase)
        if args.phase is not None
        else run(args.state_root, args.evidence_dir)
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
