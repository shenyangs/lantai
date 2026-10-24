# -*- coding: utf-8 -*-
"""EDIT-06 worker: derive PDF from one exact immutable PPTX revision."""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from time import perf_counter

from pypdf import PdfReader
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from lantai.app.capabilities.artifacts import artifact_manifest_path
from lantai.app.office_conversion import find_soffice
from lantai.app.works import (
    add_work,
    create_work_edit_copy,
    derive_work_to_pdf,
    get_work,
    list_works,
    works_index_path,
)
from lantai.presentation.pdf_render import render_pdf_pages
from lantai.presentation.quality import file_sha256
from lantai.presentation.render import verify_rendered_text


CONTROL_NAME = "edit06-control.json"
V1_NAME = "edit06-specified-v1.pptx"
V1_ONLY_TOKEN = "EDIT06-SOURCE-V1-ONLY"
V2_ONLY_TOKEN = "EDIT06-SOURCE-V2-ONLY"


def _paths(state_root: Path) -> dict[str, Path]:
    return {
        "workspace": state_root / "workspace",
        "author": state_root / "author",
        "control": state_root / CONTROL_NAME,
        "v1_source": state_root / "author" / V1_NAME,
        "rendered": state_root / "observation" / "rendered-pdf",
    }


def _sha256(path: Path) -> str:
    return f"sha256:{file_sha256(path)}"


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _read_control(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def _add_text_slide(
    presentation: Presentation,
    *,
    title: str,
    body: str,
) -> None:
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)
    title_box = slide.shapes.add_textbox(
        Inches(0.85), Inches(0.72), Inches(11.65), Inches(0.9)
    )
    title_frame = title_box.text_frame
    title_frame.clear()
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.text = title
    title_paragraph.font.name = "PingFang SC"
    title_paragraph.font.size = Pt(30)
    title_paragraph.font.bold = True
    title_paragraph.alignment = PP_ALIGN.LEFT
    body_box = slide.shapes.add_textbox(
        Inches(0.9), Inches(2.05), Inches(10.8), Inches(1.25)
    )
    body_frame = body_box.text_frame
    body_frame.clear()
    body_paragraph = body_frame.paragraphs[0]
    body_paragraph.text = body
    body_paragraph.font.name = "PingFang SC"
    body_paragraph.font.size = Pt(20)


def _create_v1(path: Path) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    _add_text_slide(
        presentation,
        title="EDIT-06 指定版本 V1",
        body=f"来源校验令牌：{V1_ONLY_TOKEN}",
    )
    _add_text_slide(
        presentation,
        title="V1 第二页",
        body="本页必须出现在由指定旧版本派生的 PDF 中。",
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def _pptx_text(path: Path) -> str:
    presentation = Presentation(str(path))
    return "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _pdf_text(path: Path) -> tuple[int, str]:
    reader = PdfReader(str(path))
    return len(reader.pages), "\n".join(
        page.extract_text() or "" for page in reader.pages
    )


def _soffice_version() -> str:
    completed = subprocess.run(
        [find_soffice(), "--version"],
        capture_output=True,
        text=True,
        timeout=15,
        check=False,
    )
    return (completed.stdout or completed.stderr).strip() or "unknown"


def _phase_prepare(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    _create_v1(paths["v1_source"])
    version_one = add_work(
        paths["workspace"],
        source_path=paths["v1_source"],
        allowed_roots=[paths["author"]],
        title="EDIT-06 指定版本演示",
        source_chat_id="chat-edit06",
        source_run_id="run-edit06",
    )
    control: dict[str, object] = {
        "v1_work_id": version_one.id,
        "v1_artifact_id": version_one.artifact_id,
        "v1_version": version_one.version,
        "v1_digest": version_one.source_digest,
        "v1_managed_path": version_one.path,
    }
    _write_json(paths["control"], control)
    return {
        "phase": "prepare",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "work_id": version_one.id,
        "artifact_id": version_one.artifact_id,
        "version": version_one.version,
        "source_digest": version_one.source_digest,
        "pptx_sha256": _sha256(Path(version_one.path)),
        "page_count": len(Presentation(version_one.path).slides),
        "contains_v1_token": V1_ONLY_TOKEN in _pptx_text(Path(version_one.path)),
        "presentation_work_count": 1,
    }


def _phase_advance(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    version_one = get_work(paths["workspace"], str(control["v1_work_id"]))
    if version_one is None:
        raise RuntimeError("EDIT-06 v1 Work is missing before advance")
    edit_copy, parent_digest = create_work_edit_copy(
        paths["workspace"], version_one, version_one.path
    )
    presentation = Presentation(str(edit_copy))
    title_shape = next(
        shape
        for shape in presentation.slides[0].shapes
        if getattr(shape, "has_text_frame", False)
    )
    title_shape.text = "EDIT-06 当前版本 V2"
    _add_text_slide(
        presentation,
        title="V2 新增第三页",
        body=f"当前版本独有令牌：{V2_ONLY_TOKEN}",
    )
    presentation.save(edit_copy)
    version_two = add_work(
        paths["workspace"],
        source_path=edit_copy,
        allowed_roots=[edit_copy.parent],
    )
    control.update(
        {
            "v2_work_id": version_two.id,
            "v2_artifact_id": version_two.artifact_id,
            "v2_version": version_two.version,
            "v2_digest": version_two.source_digest,
            "v2_managed_path": version_two.path,
        }
    )
    _write_json(paths["control"], control)
    presentation_works = [
        item for item in list_works(paths["workspace"]) if item.kind == "presentation"
    ]
    return {
        "phase": "advance",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "work_id": version_two.id,
        "artifact_id": version_two.artifact_id,
        "version": version_two.version,
        "parent_work_id": version_two.parent_work_id,
        "parent_digest": version_two.parent_digest,
        "expected_parent_digest": parent_digest,
        "source_digest": version_two.source_digest,
        "pptx_sha256": _sha256(Path(version_two.path)),
        "page_count": len(Presentation(version_two.path).slides),
        "contains_v2_token": V2_ONLY_TOKEN in _pptx_text(Path(version_two.path)),
        "presentation_work_count": len(presentation_works),
    }


def _phase_derive(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    requested_work_id = str(control["v1_work_id"])
    before = list_works(paths["workspace"])
    derived = derive_work_to_pdf(paths["workspace"], requested_work_id)
    first_duration = perf_counter() - started
    repeated_started = perf_counter()
    repeated = derive_work_to_pdf(paths["workspace"], requested_work_id)
    repeat_duration = perf_counter() - repeated_started
    after = list_works(paths["workspace"])
    control.update(
        {
            "pdf_work_id": derived.id,
            "pdf_artifact_id": derived.artifact_id,
            "pdf_managed_path": derived.path,
            "pdf_digest": derived.source_digest,
        }
    )
    _write_json(paths["control"], control)
    relation = derived.derivation
    if relation is None:
        raise RuntimeError("EDIT-06 PDF derivation relation is missing")
    return {
        "phase": "derive",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "first_derivation_seconds": round(first_duration, 6),
        "repeat_lookup_seconds": round(repeat_duration, 6),
        "requested_work_id": requested_work_id,
        "current_v2_work_id": control["v2_work_id"],
        "derived_work_id": derived.id,
        "repeated_work_id": repeated.id,
        "idempotent_reuse": repeated.id == derived.id,
        "work_count_before": len(before),
        "work_count_after": len(after),
        "presentation_work_count_after": sum(
            item.kind == "presentation" for item in after
        ),
        "pdf_work_count_after": sum(item.kind == "pdf" for item in after),
        "derived_kind": derived.kind,
        "derived_version": derived.version,
        "derived_parent_work_id": derived.parent_work_id,
        "derived_lineage_root_id": derived.lineage_root_id,
        "derived_digest": derived.source_digest,
        "derived_sha256": _sha256(Path(derived.path)),
        "relation": relation.model_dump(mode="json"),
    }


def _phase_observe(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    works = list_works(paths["workspace"])
    by_id = {item.id: item for item in works}
    v1 = by_id[str(control["v1_work_id"])]
    v2 = by_id[str(control["v2_work_id"])]
    derived = by_id[str(control["pdf_work_id"])]
    v1_path = Path(v1.path)
    v2_path = Path(v2.path)
    pdf_path = Path(derived.path)
    page_count, pdf_text = _pdf_text(pdf_path)
    verify_rendered_text(v1_path, pdf_path)
    shutil.rmtree(paths["rendered"], ignore_errors=True)
    rendered_pages = render_pdf_pages(
        pdf_path,
        paths["rendered"],
        dpi=144,
        prefix="page-",
        zero_pad=2,
        image_format="PNG",
    )
    relation = derived.derivation
    if relation is None:
        raise RuntimeError("EDIT-06 observer found no derivation relation")
    index_payload = json.loads(
        works_index_path(paths["workspace"]).read_text(encoding="utf-8")
    )
    manifest_payload = json.loads(
        artifact_manifest_path(paths["workspace"]).read_text(encoding="utf-8")
    )
    return {
        "phase": "observe",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "soffice_version": _soffice_version(),
        "v1_page_count": len(Presentation(v1_path).slides),
        "v2_page_count": len(Presentation(v2_path).slides),
        "pdf_page_count": page_count,
        "rendered_page_count": len(rendered_pages),
        "rendered_page_sha256": [_sha256(path) for path in rendered_pages],
        "pdf_signature_valid": pdf_path.read_bytes().startswith(b"%PDF-"),
        "pdf_contains_v1_token": V1_ONLY_TOKEN in pdf_text,
        "pdf_excludes_v2_token": V2_ONLY_TOKEN not in pdf_text,
        "verified_rendered_text": True,
        "v1_sha256": _sha256(v1_path),
        "v2_sha256": _sha256(v2_path),
        "pdf_sha256": _sha256(pdf_path),
        "v1_digest": v1.source_digest,
        "v2_digest": v2.source_digest,
        "pdf_digest": derived.source_digest,
        "relation": relation.model_dump(mode="json"),
        "presentation_work_ids": [
            item.id for item in works if item.kind == "presentation"
        ],
        "pdf_work_ids": [item.id for item in works if item.kind == "pdf"],
        "works_index_count": len(index_payload["works"]),
        "artifact_manifest_count": len(manifest_payload["artifacts"]),
        "derived_is_revision": (
            derived.parent_work_id is not None or derived.version != 1
        ),
    }


def _run_phase(state_root: Path, phase: str) -> dict[str, object]:
    if phase == "prepare":
        return _phase_prepare(state_root)
    if phase == "advance":
        return _phase_advance(state_root)
    if phase == "derive":
        return _phase_derive(state_root)
    if phase == "observe":
        return _phase_observe(state_root)
    raise ValueError(f"unknown EDIT-06 phase: {phase}")


def _invoke_phase(state_root: Path, phase: str) -> dict[str, object]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--state-root",
            str(state_root),
            "--phase",
            phase,
        ],
        env=environment,
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"EDIT-06 {phase} failed: stdout={completed.stdout!r}; "
            f"stderr={completed.stderr!r}"
        )
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _copy_evidence(
    source: Path,
    destination: Path,
) -> dict[str, str]:
    shutil.copy2(source, destination)
    return {"path": destination.name, "sha256": _sha256(destination)}


def _write_evidence_bundle(
    state_root: Path,
    result: dict[str, object],
    evidence_dir: Path,
) -> None:
    paths = _paths(state_root)
    control = _read_control(paths["control"])
    evidence_dir.mkdir(parents=True, exist_ok=True)
    files = {
        "specified_v1_pptx": (
            Path(str(control["v1_managed_path"])),
            evidence_dir / "edit06-specified-v1.pptx",
        ),
        "current_v2_pptx": (
            Path(str(control["v2_managed_path"])),
            evidence_dir / "edit06-current-v2.pptx",
        ),
        "derived_from_v1_pdf": (
            Path(str(control["pdf_managed_path"])),
            evidence_dir / "edit06-derived-from-v1.pdf",
        ),
        "works_index": (
            works_index_path(paths["workspace"]),
            evidence_dir / "edit06-works-index.json",
        ),
        "artifact_manifest": (
            artifact_manifest_path(paths["workspace"]),
            evidence_dir / "edit06-artifact-manifest.json",
        ),
    }
    artifacts = {
        label: _copy_evidence(source, destination)
        for label, (source, destination) in files.items()
    }
    rendered_artifacts = []
    for page in sorted(paths["rendered"].glob("page-*.png")):
        rendered_artifacts.append(
            _copy_evidence(page, evidence_dir / f"edit06-{page.name}")
        )
    result["artifact_files"] = artifacts
    result["rendered_page_files"] = rendered_artifacts
    evidence_path = evidence_dir / "edit06-pdf-derivation-cross-boundary.json"
    _write_json(evidence_path, result)


def run(state_root: Path, evidence_dir: Path | None = None) -> dict[str, object]:
    state_root.mkdir(parents=True, exist_ok=True)
    phase_names = ("prepare", "advance", "derive", "observe")
    phases = {phase: _invoke_phase(state_root, phase) for phase in phase_names}
    result: dict[str, object] = {
        "matrix_id": "EDIT-06",
        "minimum_evidence": "cross_boundary",
        "orchestrator_pid": os.getpid(),
        "child_pids": [int(phases[phase]["pid"]) for phase in phase_names],
        "phases": phases,
        "total_phase_seconds": round(
            sum(float(phases[phase]["duration_seconds"]) for phase in phase_names),
            6,
        ),
    }
    if evidence_dir is not None:
        _write_evidence_bundle(state_root, result, evidence_dir)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument(
        "--phase",
        choices=("prepare", "advance", "derive", "observe"),
    )
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    result = (
        _run_phase(args.state_root, args.phase)
        if args.phase is not None
        else run(args.state_root, args.evidence_dir)
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
