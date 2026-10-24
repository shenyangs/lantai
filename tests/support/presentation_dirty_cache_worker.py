# -*- coding: utf-8 -*-
"""Independent QA-05 worker for page-local visual cache invalidation."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import threading
from pathlib import Path
from zipfile import ZipFile

from PIL import Image, ImageDraw
from pptx import Presentation

from lantai.presentation import service as service_module
from lantai.presentation.jobs import PresentationJobStore
from lantai.presentation.models import (
    CompositionAuthority,
    CompositionEngine,
    DeckIR,
    DeckSlide,
    ObjectFrame,
    PlanningMode,
    PresentationJobRequest,
    TextObject,
    VisualCriticResult,
)
from lantai.presentation.quality import file_sha256
from lantai.presentation.render import RenderResult
from lantai.presentation.service import PresentationService


CONVERSATION_ID = "conversation-qa05-dirty-cache"
RUNTIME_RUN_ID = "native-run-qa05-dirty-cache"
CRITIC_IDENTITY = "qa05-controlled-critic/v1"


def _deck(*, revision: int) -> DeckIR:
    slides: list[DeckSlide] = []
    for index in range(1, 6):
        body = f"第 {index} 页稳定正文：只允许目标页变化。"
        if index == 3 and revision == 2:
            body = "第 3 页修订正文：缓存必须只失效这一页。"
        slides.append(
            DeckSlide(
                id=f"qa05-slide-{index:02d}",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=f"第 {index} 页的来源说明保持稳定。",
                objects=[
                    TextObject(
                        id=f"qa05-title-{index:02d}",
                        role="title",
                        text=f"第 {index} 页判断",
                        slot="title",
                        frame=ObjectFrame(x=0.8, y=0.7, w=9.0, h=0.8),
                    ),
                    TextObject(
                        id=f"qa05-body-{index:02d}",
                        role="body",
                        text=body,
                        slot="body",
                        frame=ObjectFrame(x=0.9, y=2.0, w=7.6, h=1.2),
                    ),
                ],
            ),
        )
    return DeckIR(
        id="qa05-dirty-cache-deck",
        title="QA-05 Dirty 缓存边界",
        slides=slides,
    )


def _request() -> PresentationJobRequest:
    return PresentationJobRequest(
        conversation_id=CONVERSATION_ID,
        user_request="生成五页原生演示文稿，之后只修改第 3 页正文",
        slide_count=5,
        cost_budget_units=30,
        vision_critique_required=True,
        planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
        composition_engine=CompositionEngine.PPTXGENJS_V3,
    )


def _slide_text(slide) -> str:
    return "\n".join(
        shape.text
        for shape in slide.shapes
        if getattr(shape, "has_text_frame", False)
    )


def _deterministic_render(
    source: str | Path,
    output_dir: str | Path,
    *,
    timeout: int = 120,
) -> RenderResult:
    _ = timeout
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(str(source))
    images: list[Path] = []
    for index, slide in enumerate(presentation.slides, start=1):
        text = _slide_text(slide)
        digest = hashlib.sha256(text.encode("utf-8")).digest()
        image = Image.new("RGB", (640, 360), (digest[0], digest[1], digest[2]))
        draw = ImageDraw.Draw(image)
        draw.rectangle((40, 40, 600, 320), outline="white", width=4)
        draw.text((64, 64), f"slide {index}", fill="white")
        image_path = destination / f"slide-{index}.png"
        image.save(image_path)
        images.append(image_path)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"QA-05 deterministic render evidence\n")
    return RenderResult(
        pdf_path=pdf_path,
        slide_images=tuple(images),
        renderer_id="qa05-deterministic-renderer",
        renderer_version="1.0",
    )


class _RecordingCritic:
    def __init__(self) -> None:
        self.slide_calls: list[str] = []
        self.deck_calls = 0
        self._lock = threading.Lock()

    def cache_identity(self) -> str:
        return CRITIC_IDENTITY

    def review_slide(self, slide) -> VisualCriticResult:
        with self._lock:
            self.slide_calls.append(slide.slide_id)
        return VisualCriticResult(critic_id=CRITIC_IDENTITY, issues=[])

    def review_deck(self, _slides) -> VisualCriticResult:
        with self._lock:
            self.deck_calls += 1
        return VisualCriticResult(critic_id=CRITIC_IDENTITY, issues=[])


def _job_paths(state_root: Path, job_id: str) -> dict[str, Path]:
    job_dir = PresentationJobStore(state_root).job_dir(job_id)
    return {
        "job_dir": job_dir,
        "deck_v1": job_dir / "checkpoints" / "qa05-deck-v1.json",
        "candidate_v1": job_dir / "candidates" / "qa05-v1.pptx",
        "candidate_v2": job_dir / "candidates" / "qa05-v2.pptx",
        "render_v1": job_dir / "renders" / "qa05-v1",
        "render_v2": job_dir / "renders" / "qa05-v2",
        "quality": job_dir / "quality" / "visual-checkpoints",
        "seed_summary": job_dir / "checkpoints" / "qa05-seed-summary.json",
    }


def _checkpoint_hashes(path: Path) -> dict[str, str]:
    return {
        item.name: file_sha256(item)
        for item in sorted(path.glob("*.json"))
    }


def _slide_xml_hashes(path: Path) -> dict[str, str]:
    with ZipFile(path) as archive:
        return {
            f"slide-{index:02d}": hashlib.sha256(
                archive.read(f"ppt/slides/slide{index}.xml"),
            ).hexdigest()
            for index in range(1, 6)
        }


def _verify(
    service: PresentationService,
    deck: DeckIR,
    candidate: Path,
    render_dir: Path,
    job_id: str,
):
    original_render = service_module.render_pptx
    service_module.render_pptx = _deterministic_render
    try:
        return service.verify_artifact(
            deck,
            candidate,
            render_dir,
            visual_required=True,
            vision_critique_required=True,
            job_id=job_id,
        )
    finally:
        service_module.render_pptx = original_render


def _seed(state_root: Path) -> dict[str, object]:
    critic = _RecordingCritic()
    service = PresentationService(state_root, visual_critic=critic)
    record = service.create_job(_request(), runtime_run_id=RUNTIME_RUN_ID)
    paths = _job_paths(state_root, record.job_id)
    deck = _deck(revision=1)
    paths["candidate_v1"].parent.mkdir(parents=True, exist_ok=True)
    service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
        deck,
        paths["candidate_v1"],
    )
    service.store.write_checkpoint(
        record.job_id,
        "qa05-deck-v1.json",
        deck.model_dump_json(indent=2),
    )
    report, rendered = _verify(
        service,
        deck,
        paths["candidate_v1"],
        paths["render_v1"],
        record.job_id,
    )
    summary = {
        "checkpoint_hashes": _checkpoint_hashes(paths["quality"]),
        "slide_xml_hashes": _slide_xml_hashes(paths["candidate_v1"]),
        "render_hashes": [file_sha256(path) for path in rendered],
    }
    service.store.write_checkpoint(
        record.job_id,
        "qa05-seed-summary.json",
        json.dumps(summary, ensure_ascii=False, indent=2, sort_keys=True),
    )
    return {
        "seed_pid": os.getpid(),
        "job_id": record.job_id,
        "slide_calls": sorted(critic.slide_calls),
        "deck_calls": critic.deck_calls,
        "report_passed": report.passed,
        "rendered_slide_count": len(rendered),
        "visual_evidence_count": len(report.visual_review_evidence),
        "cost_used_units": service.store.get(record.job_id).cost_used_units,
    }


def _modify(state_root: Path, job_id: str) -> dict[str, object]:
    critic = _RecordingCritic()
    service = PresentationService(state_root, visual_critic=critic)
    paths = _job_paths(state_root, job_id)
    seed = json.loads(paths["seed_summary"].read_text(encoding="utf-8"))
    deck_v1 = DeckIR.model_validate_json(paths["deck_v1"].read_text(encoding="utf-8"))
    deck_v2 = _deck(revision=2)
    changed_ir_ids = [
        before.id
        for before, after in zip(deck_v1.slides, deck_v2.slides, strict=True)
        if before.model_dump(mode="json") != after.model_dump(mode="json")
    ]
    service.compiler_for(CompositionEngine.PPTXGENJS_V3).compile(
        deck_v2,
        paths["candidate_v2"],
    )
    report, rendered = _verify(
        service,
        deck_v2,
        paths["candidate_v2"],
        paths["render_v2"],
        job_id,
    )
    current_checkpoints = _checkpoint_hashes(paths["quality"])
    prior_checkpoints = dict(seed["checkpoint_hashes"])
    changed_checkpoints = sorted(
        name
        for name in current_checkpoints
        if current_checkpoints[name] != prior_checkpoints.get(name)
    )
    slide3_checkpoint = (
        "slide-"
        + hashlib.sha256("qa05-slide-03".encode("utf-8")).hexdigest()
        + ".json"
    )
    current_slide_xml = _slide_xml_hashes(paths["candidate_v2"])
    prior_slide_xml = dict(seed["slide_xml_hashes"])
    changed_slide_xml = sorted(
        name
        for name in current_slide_xml
        if current_slide_xml[name] != prior_slide_xml.get(name)
    )
    current_render_hashes = [file_sha256(path) for path in rendered]
    changed_render_indexes = [
        index
        for index, (before, after) in enumerate(
            zip(seed["render_hashes"], current_render_hashes, strict=True),
            start=1,
        )
        if before != after
    ]
    slide_evidence = {
        item.slide_id: item
        for item in report.visual_review_evidence
        if item.scope == "slide" and item.slide_id is not None
    }
    render_hash_by_slide_id = {
        slide.id: current_render_hashes[index]
        for index, slide in enumerate(deck_v2.slides)
    }
    rebound_matches = {
        slide_id: evidence.rendered_sha256 == render_hash_by_slide_id[slide_id]
        for slide_id, evidence in slide_evidence.items()
    }
    deck_evidence = [
        item for item in report.visual_review_evidence if item.scope == "deck"
    ]
    record = service.store.get(job_id)
    return {
        "modify_pid": os.getpid(),
        "job_id": job_id,
        "changed_ir_slide_ids": changed_ir_ids,
        "changed_slide_xml": changed_slide_xml,
        "changed_render_indexes": changed_render_indexes,
        "slide_calls": sorted(critic.slide_calls),
        "deck_calls": critic.deck_calls,
        "cached_slide_ids": sorted(
            slide_id
            for slide_id, evidence in slide_evidence.items()
            if evidence.cache_hit
        ),
        "dirty_slide_ids": sorted(
            slide_id
            for slide_id, evidence in slide_evidence.items()
            if not evidence.cache_hit
        ),
        "deck_cache_hit": deck_evidence[0].cache_hit,
        "changed_checkpoints": changed_checkpoints,
        "expected_changed_checkpoints": sorted([slide3_checkpoint, "deck.json"]),
        "all_evidence_rebound_to_current_render": all(rebound_matches.values()),
        "rebound_matches": rebound_matches,
        "report_passed": report.passed,
        "rendered_slide_count": len(rendered),
        "visual_evidence_count": len(report.visual_review_evidence),
        "cost_used_units": record.cost_used_units,
        "job_dir_count": len(list(state_root.glob("presentation-*"))),
        "manifest_count": len(list(state_root.glob("presentation-*/manifest.json"))),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("seed", "modify"), required=True)
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--job-id")
    args = parser.parse_args()
    args.state_root.mkdir(parents=True, exist_ok=True)
    if args.mode == "seed":
        result = _seed(args.state_root)
    else:
        if not args.job_id:
            parser.error("--job-id is required for modify mode")
        result = _modify(args.state_root, args.job_id)
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
