# -*- coding: utf-8 -*-
"""Independent RUN-08 worker for one-page repair and cache preservation."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import threading
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation

from lantai.presentation.design_system import prepare_deck_for_composition
from lantai.presentation.jobs import PresentationJobStore
from lantai.presentation.models import (
    CompositionAuthority,
    CompositionEngine,
    DeckIR,
    DeckSlide,
    JobStatus,
    ObjectFrame,
    PlanningMode,
    PresentationJobRequest,
    TextObject,
    VisualCriticResult,
    VisualIssue,
)
from lantai.presentation.quality import file_sha256
from lantai.presentation.render import RenderResult
from lantai.presentation.repair import (
    RepairActionType,
    SlideRedesignProposal,
)
from lantai.presentation.service import PresentationService


CONVERSATION_ID = "conversation-run08-local-repair"
RUNTIME_RUN_ID = "native-run-run08-local-repair"
CRITIC_IDENTITY = "run08-local-repair-critic/v1"
CRITIC_RESULT_ID = "run08-local-repair-critic/model-a"
TARGET_SLIDE_ID = "run08-slide-07"


def _canonical_hash(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        ).encode("utf-8"),
    ).hexdigest()


def _deck() -> DeckIR:
    slides: list[DeckSlide] = []
    for index in range(1, 20):
        slides.append(
            DeckSlide(
                id=f"run08-slide-{index:02d}",
                layout_id="freeform",
                composition_authority=CompositionAuthority.AUTHOR,
                speaker_notes=f"第 {index} 页只表达判断 {index}",
                objects=[
                    TextObject(
                        id=f"run08-title-{index:02d}",
                        role="title",
                        text=f"第 {index} 页判断",
                        slot="title",
                        frame=ObjectFrame(x=0.8, y=0.7, w=8.8, h=0.8),
                    ),
                    TextObject(
                        id=f"run08-body-{index:02d}",
                        role="body",
                        text=f"这是第 {index} 页唯一且必须保留的语义内容。",
                        slot="body",
                        frame=ObjectFrame(x=0.9, y=2.0, w=5.4, h=1.2),
                    ),
                ],
            ),
        )
    return DeckIR(
        id="run08-local-repair-deck",
        title="RUN-08 局部修复边界",
        slides=slides,
    )


def _append_jsonl(path: Path, payload: dict[str, object], lock: threading.Lock) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with lock:
        with path.open("a", encoding="utf-8") as handle:
            handle.write(
                json.dumps(payload, ensure_ascii=False, sort_keys=True) + "\n",
            )
            handle.flush()
            os.fsync(handle.fileno())


class _LocalFailureCritic:
    def __init__(self, ledger: Path) -> None:
        self.ledger = ledger
        self._lock = threading.Lock()

    def cache_identity(self) -> str:
        return CRITIC_IDENTITY

    def review_slide(self, slide) -> VisualCriticResult:
        render_round = Path(slide.rendered_path).parent.name
        _append_jsonl(
            self.ledger,
            {
                "scope": "slide",
                "slide_id": slide.slide_id,
                "render_round": render_round,
                "rendered_sha256": slide.rendered_sha256,
            },
            self._lock,
        )
        issues = []
        if slide.slide_id == TARGET_SLIDE_ID and render_round == "round-0":
            issues.append(
                VisualIssue(
                    slide_id=TARGET_SLIDE_ID,
                    severity="error",
                    issue_type="hierarchy",
                    object_ids=["run08-body-07"],
                    description="第 7 页正文层级不足，需要局部重构。",
                    proposed_action="只重构第 7 页的正文构图。",
                ),
            )
        return VisualCriticResult(
            critic_id=CRITIC_RESULT_ID,
            issues=issues,
        )

    def review_deck(self, slides) -> VisualCriticResult:
        render_round = Path(slides[0].rendered_path).parent.name
        _append_jsonl(
            self.ledger,
            {
                "scope": "deck",
                "slide_count": len(slides),
                "render_round": render_round,
            },
            self._lock,
        )
        return VisualCriticResult(critic_id=CRITIC_RESULT_ID, issues=[])

    def review(self, _slides):
        raise AssertionError("RUN-08 must use page and deck critic checkpoints")


class _SeventhSlideDesigner:
    def __init__(self, ledger: Path) -> None:
        self.ledger = ledger
        self._lock = threading.Lock()

    def redesign(self, request) -> SlideRedesignProposal:
        _append_jsonl(
            self.ledger,
            {
                "slide_id": request.slide.id,
                "issue_codes": [issue.code for issue in request.issues],
                "allowed_actions": [
                    action.value for action in request.allowed_actions
                ],
            },
            self._lock,
        )
        if request.slide.id != TARGET_SLIDE_ID:
            raise AssertionError("RUN-08 designer received a correct slide")
        objects = []
        for obj in request.slide.objects:
            if obj.id == "run08-body-07":
                objects.append(
                    obj.model_copy(
                        update={
                            "frame": ObjectFrame(
                                x=6.8,
                                y=2.0,
                                w=5.4,
                                h=1.2,
                            ),
                            "emphasis": True,
                        },
                    ),
                )
            else:
                objects.append(obj)
        return SlideRedesignProposal(
            slide=request.slide.model_copy(update={"objects": objects}),
            action=RepairActionType.REDESIGN_SLIDE,
            rationale="只修复第 7 页正文层级",
        )


def _render_pptx(
    source: str | Path,
    output_dir: str | Path,
    *,
    timeout: int = 120,
) -> RenderResult:
    _ = timeout
    destination = Path(output_dir)
    destination.mkdir(parents=True, exist_ok=True)
    pdf_path = destination / "presentation.pdf"
    pdf_path.write_bytes(b"RUN-08 deterministic render evidence\n")
    slide_count = len(Presentation(str(source)).slides)
    images: list[Path] = []
    for index in range(1, slide_count + 1):
        image_path = destination / f"slide-{index}.png"
        image = Image.new("RGB", (640, 360), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((0, 0, 220 + index, 360), fill=(25, 25, 25))
        if index == 7:
            color = (
                (190, 40, 40)
                if destination.name == "round-0"
                else (30, 130, 80)
            )
            draw.rectangle((460, 80, 610, 280), fill=color)
        image.save(image_path)
        images.append(image_path)
    return RenderResult(
        pdf_path=pdf_path,
        slide_images=tuple(images),
        renderer_id="run08-deterministic-renderer",
        renderer_version="1.0",
    )


def _slide_hashes(deck: DeckIR) -> dict[str, str]:
    return {
        slide.id: _canonical_hash(slide.model_dump(mode="json"))
        for slide in deck.slides
    }


def _pptx_slide_xml_hashes(path: Path) -> dict[str, str]:
    with zipfile.ZipFile(path, "r") as package:
        return {
            f"slide-{index:02d}": hashlib.sha256(
                package.read(f"ppt/slides/slide{index}.xml"),
            ).hexdigest()
            for index in range(1, 20)
        }


def _compose(state_root: Path) -> int:
    from lantai.presentation import service as service_module

    service_module.render_pptx = _render_pptx
    presentation_root = state_root / "presentation-jobs"
    critic_ledger = state_root / "critic-calls.jsonl"
    designer_ledger = state_root / "designer-calls.jsonl"
    service = PresentationService(
        presentation_root,
        visual_critic=_LocalFailureCritic(critic_ledger),
        repair_designer=_SeventhSlideDesigner(designer_ledger),
    )
    record = service.create_job(
        PresentationJobRequest(
            conversation_id=CONVERSATION_ID,
            user_request="RUN-08：19 页中第 7 页设计失败时只修必要页",
            slide_count=19,
            planning_mode=PlanningMode.DETERMINISTIC_FIXTURE,
            composition_engine=CompositionEngine.NATIVE_PPTX,
            cost_budget_units=200,
            visual_verification_required=True,
            vision_critique_required=True,
        ),
        runtime_run_id=RUNTIME_RUN_ID,
    )
    result, report = service.compose(
        record.job_id,
        _deck(),
        runtime_run_id=RUNTIME_RUN_ID,
    )
    final = service.store.get(record.job_id)
    print(
        json.dumps(
            {
                "compose_pid": os.getpid(),
                "job_id": record.job_id,
                "status": final.status.value,
                "report_passed": report.passed,
                "repair_rounds": report.repair_rounds,
                "slide_count": result.slide_count,
                "output_path": str(result.output_path),
            },
            ensure_ascii=False,
            sort_keys=True,
        ),
    )
    return 0 if final.status == JobStatus.COMPLETED and report.passed else 3


def _read_jsonl(path: Path) -> list[dict[str, object]]:
    return [
        json.loads(line)
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]


def _observe(state_root: Path, job_id: str) -> int:
    store = PresentationJobStore(
        state_root / "presentation-jobs",
        read_only=True,
    )
    record = store.get(job_id)
    job_dir = store.job_dir(job_id)
    original = prepare_deck_for_composition(_deck())
    repaired = DeckIR.model_validate_json(
        store.checkpoint_path(
            job_id,
            "deck-ir-redesign-1.json",
        ).read_text(encoding="utf-8"),
    )
    final = DeckIR.model_validate_json(
        store.checkpoint_path(job_id, "deck-ir.json").read_text(
            encoding="utf-8",
        ),
    )
    original_hashes = _slide_hashes(original)
    repaired_hashes = _slide_hashes(repaired)
    changed_slide_ids = [
        slide_id
        for slide_id, digest in original_hashes.items()
        if repaired_hashes[slide_id] != digest
    ]

    first_candidate = job_dir / "candidates" / "round-0" / "presentation.pptx"
    final_candidate = job_dir / "candidates" / "round-1" / "presentation.pptx"
    first_xml = _pptx_slide_xml_hashes(first_candidate)
    final_xml = _pptx_slide_xml_hashes(final_candidate)
    changed_pptx_pages = [
        page_id
        for page_id, digest in first_xml.items()
        if final_xml[page_id] != digest
    ]

    round_zero = json.loads(
        (job_dir / "quality" / "quality-report-round-0.json").read_text(
            encoding="utf-8",
        ),
    )
    round_one = json.loads(
        (job_dir / "quality" / "quality-report-round-1.json").read_text(
            encoding="utf-8",
        ),
    )
    second_slide_evidence = [
        item
        for item in round_one["visual_review_evidence"]
        if item["scope"] == "slide"
    ]
    second_deck_evidence = [
        item
        for item in round_one["visual_review_evidence"]
        if item["scope"] == "deck"
    ]
    critic_calls = _read_jsonl(state_root / "critic-calls.jsonl")
    designer_calls = _read_jsonl(state_root / "designer-calls.jsonl")
    output = job_dir / "output" / "presentation.pptx"
    summary = {
        "observer_pid": os.getpid(),
        "job_dir_count": len(
            list((state_root / "presentation-jobs").glob("presentation-*")),
        ),
        "manifest_count": len(
            list(
                (state_root / "presentation-jobs").glob(
                    "presentation-*/manifest.json",
                ),
            ),
        ),
        "job_id": job_id,
        "runtime_run_id": record.runtime_run_id,
        "status": record.status.value,
        "final_stage": record.stage.value,
        "repair_rounds": round_one["repair_rounds"],
        "round_zero_passed": round_zero["passed"],
        "round_one_passed": round_one["passed"],
        "round_zero_error_slide_ids": sorted(
            {
                issue.get("slide_id")
                for issue in round_zero["issues"]
                if issue["severity"] == "error" and issue.get("slide_id")
            },
        ),
        "changed_deck_slide_ids": changed_slide_ids,
        "unchanged_deck_slide_count": 19 - len(changed_slide_ids),
        "final_checkpoint_matches_repair": final == repaired,
        "changed_pptx_pages": changed_pptx_pages,
        "unchanged_pptx_page_count": 19 - len(changed_pptx_pages),
        "second_round_cached_slide_ids": sorted(
            item["slide_id"]
            for item in second_slide_evidence
            if item["cache_hit"]
        ),
        "second_round_uncached_slide_ids": sorted(
            item["slide_id"]
            for item in second_slide_evidence
            if not item["cache_hit"]
        ),
        "second_round_deck_cache_hit": second_deck_evidence[0]["cache_hit"],
        "critic_call_count": len(critic_calls),
        "second_round_critic_calls": [
            call for call in critic_calls if call["render_round"] == "round-1"
        ],
        "designer_calls": designer_calls,
        "cost_used_units": record.cost_used_units,
        "output_exists": output.is_file(),
        "output_sha256": file_sha256(output),
        "final_candidate_sha256": file_sha256(final_candidate),
        "final_pptx_slide_count": len(Presentation(str(output)).slides),
        "final_native_shape_count": sum(
            len(slide.shapes) for slide in Presentation(str(output)).slides
        ),
        "quality_metric_rounds": sorted(
            {
                metric.repair_round
                for metric in record.stage_metrics
                if metric.operation == "critic"
            },
        ),
    }
    print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    return 0


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=("compose", "observe"), required=True)
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--job-id")
    args = parser.parse_args()
    if args.mode == "observe":
        if not args.job_id:
            parser.error("--job-id is required in observe mode")
        return _observe(args.state_root, args.job_id)
    return _compose(args.state_root)


if __name__ == "__main__":
    raise SystemExit(main())
