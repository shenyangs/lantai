# -*- coding: utf-8 -*-
"""SRC-01 worker: exercise real binary DOC extraction across processes."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path
from time import perf_counter

import olefile
from docx import Document
from pypdf import PdfReader

from lantai.app.office_conversion import convert_office_document, find_soffice
from lantai.presentation.material_extraction import PresentationMaterialError
from lantai.presentation.models import (
    InputMaterial,
    PresentationBrief,
    PresentationJobRequest,
)
from lantai.presentation.pdf_render import render_pdf_pages
from lantai.presentation.planner import build_brief


SEED_DOCX = "src01-seed.docx"
LEGACY_DOC = "src01-real-word97.doc"
CORRUPT_DOC = "src01-corrupt.doc"
REQUEST_JSON = "src01-request.json"
BRIEF_JSON = "src01-brief.json"
EXTRACTED_TEXT = "src01-extracted.txt"
OBSERVED_DOCX = "src01-observed-conversion.docx"
OBSERVED_PDF = "src01-real-word97.pdf"

EXPECTED_PARAGRAPHS = [
    "SRC-01 旧稿事实边界",
    "正文事实：第三季度客户续约率为 92%。",
    "结论：先修复交付可靠性，再扩张作者策略。",
]
EXPECTED_TABLE = [
    ["指标", "本期", "口径"],
    ["续约率", "92%", "已签约客户"],
    ["样本量", "240家", "截至9月30日"],
]
EXPECTED_EXTRACTED_LINES = [
    EXPECTED_PARAGRAPHS[0],
    EXPECTED_PARAGRAPHS[1],
    *(" | ".join(row) for row in EXPECTED_TABLE),
    EXPECTED_PARAGRAPHS[2],
]


def _paths(state_root: Path) -> dict[str, Path]:
    return {
        "seed_docx": state_root / "source" / SEED_DOCX,
        "legacy_doc": state_root / "source" / LEGACY_DOC,
        "corrupt_doc": state_root / "source" / CORRUPT_DOC,
        "request": state_root / REQUEST_JSON,
        "brief": state_root / BRIEF_JSON,
        "extracted": state_root / EXTRACTED_TEXT,
        "observed_docx": state_root / "observation" / OBSERVED_DOCX,
        "observed_pdf": state_root / "observation" / OBSERVED_PDF,
        "rendered": state_root / "observation" / "rendered",
    }


def _sha256(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def _write_json(path: Path, payload: dict[str, object]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def _create_seed_docx(path: Path) -> None:
    document = Document()
    title = document.add_paragraph(EXPECTED_PARAGRAPHS[0])
    title.runs[0].bold = True
    document.add_paragraph(EXPECTED_PARAGRAPHS[1])
    table = document.add_table(rows=len(EXPECTED_TABLE), cols=3)
    for row, values in zip(table.rows, EXPECTED_TABLE, strict=True):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_paragraph(EXPECTED_PARAGRAPHS[2])
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _soffice_version() -> str:
    completed = subprocess.run(
        [find_soffice(), "--version"],
        capture_output=True,
        text=True,
        timeout=30,
        check=True,
    )
    return (completed.stdout or completed.stderr).strip()


def _phase_prepare(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    _create_seed_docx(paths["seed_docx"])
    legacy = convert_office_document(
        paths["seed_docx"],
        "doc",
        paths["legacy_doc"].parent,
        output_name=paths["legacy_doc"].name,
    )
    paths["corrupt_doc"].write_bytes(
        bytes.fromhex("d0cf11e0a1b11ae1")
        + b"SRC01-CORRUPT-NOT-A-WORD-DOCUMENT"
    )
    with olefile.OleFileIO(str(paths["legacy_doc"])) as compound:
        streams = sorted("/".join(parts) for parts in compound.listdir())
    request = PresentationJobRequest(
        conversation_id="src01-real-legacy-doc",
        user_request="根据旧版 Word 事实材料生成汇报",
        audience="产品与工程评审人",
        slide_count=3,
        materials=[
            InputMaterial(
                id="source-src01-real-doc",
                name=paths["legacy_doc"].name,
                path=str(paths["legacy_doc"]),
                media_type="application/msword",
                sha256=legacy.sha256,
            )
        ],
    )
    _write_json(paths["request"], request.model_dump(mode="json"))
    return {
        "phase": "prepare",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "fixture_class": "synthetic_content_real_binary_word_97_2003",
        "ole_signature": paths["legacy_doc"].read_bytes()[:8].hex(),
        "ole_is_valid": olefile.isOleFile(str(paths["legacy_doc"])),
        "ole_streams": streams,
        "seed_docx_sha256": _sha256(paths["seed_docx"]),
        "legacy_doc_sha256": _sha256(paths["legacy_doc"]),
        "corrupt_doc_sha256": _sha256(paths["corrupt_doc"]),
        "expected_paragraphs": EXPECTED_PARAGRAPHS,
        "expected_table": EXPECTED_TABLE,
        "expected_extracted_lines": EXPECTED_EXTRACTED_LINES,
        "libreoffice_version": _soffice_version(),
    }


def _phase_extract(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    request = PresentationJobRequest.model_validate_json(
        paths["request"].read_text(encoding="utf-8")
    )
    brief = build_brief(request)
    extracted = "\n".join(section.body for section in brief.source_sections)
    actual_lines = extracted.splitlines()
    if actual_lines != EXPECTED_EXTRACTED_LINES:
        raise RuntimeError(
            "SRC-01 production extraction changed body/table order: "
            f"{actual_lines!r}"
        )
    paths["brief"].write_text(
        brief.model_dump_json(indent=2) + "\n",
        encoding="utf-8",
    )
    paths["extracted"].write_text(extracted + "\n", encoding="utf-8")
    coverage = brief.source_coverage[0]
    return {
        "phase": "extract",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "source_section_count": len(brief.source_sections),
        "source_refs": brief.source_refs,
        "extracted_lines": actual_lines,
        "line_occurrences": {
            line: actual_lines.count(line) for line in EXPECTED_EXTRACTED_LINES
        },
        "coverage": coverage.model_dump(mode="json"),
        "brief_sha256": _sha256(paths["brief"]),
        "extracted_text_sha256": _sha256(paths["extracted"]),
    }


def _table_matrix(document: Document) -> list[list[str]]:
    if len(document.tables) != 1:
        raise RuntimeError("SRC-01 observed DOCX does not contain exactly one table")
    return [
        [cell.text for cell in row.cells]
        for row in document.tables[0].rows
    ]


def _phase_observe(state_root: Path) -> dict[str, object]:
    started = perf_counter()
    paths = _paths(state_root)
    brief = PresentationBrief.model_validate_json(
        paths["brief"].read_text(encoding="utf-8")
    )
    converted_docx = convert_office_document(
        paths["legacy_doc"],
        "docx",
        paths["observed_docx"].parent,
        output_name=paths["observed_docx"].name,
    )
    converted_pdf = convert_office_document(
        paths["legacy_doc"],
        "pdf",
        paths["observed_pdf"].parent,
        output_name=paths["observed_pdf"].name,
    )
    rendered_pages = render_pdf_pages(
        converted_pdf.output_path,
        paths["rendered"],
        dpi=144,
        prefix="page-",
        zero_pad=2,
        image_format="PNG",
    )
    observed = Document(str(converted_docx.output_path))
    observed_paragraphs = [
        paragraph.text for paragraph in observed.paragraphs if paragraph.text
    ]
    observed_table = _table_matrix(observed)
    corrupt_request = PresentationJobRequest(
        conversation_id="src01-corrupt-legacy-doc",
        user_request="根据损坏旧稿汇报",
        materials=[
            InputMaterial(
                id="source-src01-corrupt-doc",
                name=paths["corrupt_doc"].name,
                path=str(paths["corrupt_doc"]),
                media_type="application/msword",
            )
        ],
    )
    corrupt_error = ""
    try:
        build_brief(corrupt_request)
    except PresentationMaterialError as exc:
        corrupt_error = str(exc)
    if not corrupt_error:
        raise RuntimeError("SRC-01 corrupt DOC was reported as successfully read")
    if "not a valid OLE compound document" not in corrupt_error:
        raise RuntimeError(f"SRC-01 corrupt DOC failed ambiguously: {corrupt_error}")
    return {
        "phase": "observe",
        "pid": os.getpid(),
        "duration_seconds": round(perf_counter() - started, 6),
        "brief_source_sha256": brief.source_coverage[0].source_sha256,
        "legacy_doc_sha256": _sha256(paths["legacy_doc"]),
        "observed_docx_sha256": _sha256(paths["observed_docx"]),
        "observed_docx_converter_sha256": converted_docx.sha256,
        "observed_pdf_sha256": _sha256(paths["observed_pdf"]),
        "observed_pdf_converter_sha256": converted_pdf.sha256,
        "observed_paragraphs": observed_paragraphs,
        "observed_table": observed_table,
        "pdf_page_count": len(PdfReader(str(paths["observed_pdf"])).pages),
        "rendered_page_count": len(rendered_pages),
        "rendered_page_sha256": [_sha256(path) for path in rendered_pages],
        "corrupt_doc_rejected": True,
        "corrupt_error": corrupt_error,
        "libreoffice_version": _soffice_version(),
    }


def _run_phase(state_root: Path, phase: str) -> dict[str, object]:
    if phase == "prepare":
        return _phase_prepare(state_root)
    if phase == "extract":
        return _phase_extract(state_root)
    if phase == "observe":
        return _phase_observe(state_root)
    raise ValueError(f"unknown SRC-01 phase: {phase}")


def _invoke_phase(state_root: Path, phase: str) -> dict[str, object]:
    environment = os.environ.copy()
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    completed = subprocess.run(
        [
            sys.executable,
            str(Path(__file__).resolve()),
            "--state-root",
            str(state_root),
            "--phase",
            phase,
        ],
        env=environment,
        capture_output=True,
        text=True,
        timeout=240,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"SRC-01 {phase} failed: stdout={completed.stdout!r}; "
            f"stderr={completed.stderr!r}"
        )
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _copy_evidence(source: Path, destination: Path) -> dict[str, str]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return {"path": destination.name, "sha256": _sha256(destination)}


def _write_evidence_bundle(
    state_root: Path,
    result: dict[str, object],
    evidence_dir: Path,
) -> None:
    paths = _paths(state_root)
    evidence_dir.mkdir(parents=True, exist_ok=True)
    result["artifact_files"] = {
        label: _copy_evidence(source, evidence_dir / source.name)
        for label, source in {
            "seed_docx": paths["seed_docx"],
            "legacy_doc": paths["legacy_doc"],
            "corrupt_doc": paths["corrupt_doc"],
            "request": paths["request"],
            "brief": paths["brief"],
            "extracted_text": paths["extracted"],
            "observed_docx": paths["observed_docx"],
            "observed_pdf": paths["observed_pdf"],
        }.items()
    }
    result["rendered_page_files"] = [
        _copy_evidence(page, evidence_dir / f"src01-{page.name}")
        for page in sorted(paths["rendered"].glob("page-*.png"))
    ]
    _write_json(evidence_dir / "src01-real-doc-cross-boundary.json", result)


def run(state_root: Path, evidence_dir: Path | None = None) -> dict[str, object]:
    state_root.mkdir(parents=True, exist_ok=True)
    phases = {
        phase: _invoke_phase(state_root, phase)
        for phase in ("prepare", "extract", "observe")
    }
    result: dict[str, object] = {
        "matrix_id": "SRC-01",
        "minimum_evidence": "cross_boundary",
        "orchestrator_pid": os.getpid(),
        "child_pids": [int(phase["pid"]) for phase in phases.values()],
        "phases": phases,
        "total_phase_seconds": round(
            sum(float(phase["duration_seconds"]) for phase in phases.values()),
            6,
        ),
    }
    if evidence_dir is not None:
        _write_evidence_bundle(state_root, result, evidence_dir)
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--state-root", type=Path, required=True)
    parser.add_argument("--phase", choices=("prepare", "extract", "observe"))
    parser.add_argument("--evidence-dir", type=Path)
    args = parser.parse_args()
    result = (
        _run_phase(args.state_root, args.phase)
        if args.phase is not None
        else run(args.state_root, args.evidence_dir)
    )
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
