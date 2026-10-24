# -*- coding: utf-8 -*-
"""Editing one slide must not move unedited CJK pixels on another slide."""
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from pptx.util import Inches, Pt
import pytest

from lantai.app.office_conversion import OfficeConversionUnavailable, find_soffice
from lantai.presentation.render import render_pptx


def _subset_boundary_deck(destination: Path, *, edited: bool) -> None:
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    first = deck.slides.add_slide(deck.slide_layouts[6])
    text = ''.join(chr(0x4E00 + i) for i in range(255))
    if edited:
        text = text[:-1] + text[0]
    box = first.shapes.add_textbox(Inches(.75), Inches(1), Inches(11.8), Inches(5.5))
    box.text_frame.text = text
    box.text_frame.word_wrap = True
    second = deck.slides.add_slide(deck.slide_layouts[6])
    title = second.shapes.add_textbox(Inches(.75), Inches(.75), Inches(11.8), Inches(1))
    title.text_frame.text = '一零一'
    for slide, size in ((first, 12), (second, 33.987)):
        for shape in slide.shapes:
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Noto Sans SC'
                    run.font.size = Pt(size)
    deck.save(destination)


def test_cross_slide_cjk_subset_change_keeps_unedited_pixels_exact(tmp_path):
    try:
        find_soffice()
    except OfficeConversionUnavailable:
        pytest.skip('LibreOffice is unavailable')
    baseline = tmp_path / 'baseline.pptx'
    candidate = tmp_path / 'candidate.pptx'
    _subset_boundary_deck(baseline, edited=False)
    _subset_boundary_deck(candidate, edited=True)
    with ZipFile(baseline) as before, ZipFile(candidate) as after:
        assert before.read('ppt/slides/slide2.xml') == after.read('ppt/slides/slide2.xml')
    before_bytes = baseline.read_bytes()
    after_bytes = candidate.read_bytes()
    original = render_pptx(baseline, tmp_path / 'baseline-render')
    changed = render_pptx(candidate, tmp_path / 'candidate-render')
    assert len(original.slide_images) == len(changed.slide_images) == 2
    with Image.open(original.slide_images[0]) as first, Image.open(changed.slide_images[0]) as second:
        assert first.convert('RGB').tobytes() != second.convert('RGB').tobytes()
    with Image.open(original.slide_images[1]) as first, Image.open(changed.slide_images[1]) as second:
        assert first.size == second.size
        assert first.convert('RGB').tobytes() == second.convert('RGB').tobytes()
    assert baseline.read_bytes() == before_bytes
    assert candidate.read_bytes() == after_bytes


def test_all_slides_including_hidden_are_rendered_in_package_order(tmp_path):
    try:
        find_soffice()
    except OfficeConversionUnavailable:
        pytest.skip('LibreOffice is unavailable')
    deck = Presentation()
    for number in range(1, 12):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        title.text_frame.text = f'PAGE {number:02d}'
        if number == 6:
            slide._element.set('show', '0')
    # Page order is the relationship list, not numeric slide XML filenames.
    slide_ids = deck.slides._sldIdLst
    slide_ids.insert(0, slide_ids[-1])
    source = tmp_path / 'ordered-hidden.pptx'
    deck.save(source)
    original = source.read_bytes()

    rendered = render_pptx(source, tmp_path / 'render')

    expected = [11, *range(1, 11)]
    pages = PdfReader(rendered.pdf_path).pages
    assert len(pages) == len(rendered.slide_images) == 11
    assert [page.extract_text().strip() for page in pages] == [
        f'PAGE {number:02d}' for number in expected
    ]
    assert [path.name for path in rendered.slide_images] == [
        f'slide-{number}.png' for number in range(1, 12)
    ]
    assert source.read_bytes() == original
