from __future__ import annotations

import json
from pathlib import Path
import shutil
import subprocess
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from lantai.app.office_conversion import (
    convert_office_document,
    find_soffice,
    recalculate_xlsx,
)
from lantai.app.office_documents import preview_office_document
from lantai.app.routers import works as works_router
from lantai.app.work_previews import (
    get_or_create_work_preview,
    get_work_preview_page,
)
from lantai.app.works import add_work, create_work_edit_copy, list_works


def _require_soffice() -> str:
    try:
        return find_soffice()
    except RuntimeError:
        pytest.skip("LibreOffice is not installed in this test environment")


def test_real_office_read_convert_and_recalculate_pipeline(tmp_path: Path) -> None:
    """Exercise the real local Office boundary without mocks or global Python."""

    _require_soffice()
    source_dir = tmp_path / "source"
    output_dir = tmp_path / "rendered"
    source_dir.mkdir()

    docx_path = source_dir / "quarterly.docx"
    document = Document()
    document.add_heading("Quarterly report", 1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "42"
    document.save(docx_path)

    pptx_path = source_dir / "quarterly.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Quarterly presentation"
    slide.placeholders[1].text = "Revenue 42"
    presentation.save(pptx_path)

    xlsx_path = source_dir / "quarterly.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = 19
    worksheet["A2"] = 23
    worksheet["A3"] = "=SUM(A1:A2)"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(xlsx_path)

    docx_preview = preview_office_document(str(docx_path), mode="text")
    pptx_preview = preview_office_document(str(pptx_path), mode="text")
    xlsx_preview = preview_office_document(str(xlsx_path), mode="text")
    assert "Quarterly report" in docx_preview.content
    assert "Revenue" in docx_preview.content
    assert "Quarterly presentation" in pptx_preview.content
    assert "SUM(A1:A2)" in xlsx_preview.content

    docx_pdf = convert_office_document(docx_path, "pdf", output_dir)
    pptx_pdf = convert_office_document(pptx_path, "pdf", output_dir)
    assert docx_pdf.output_path.name == "quarterly.docx.pdf"
    assert pptx_pdf.output_path.name == "quarterly.pptx.pdf"
    assert docx_pdf.output_path != pptx_pdf.output_path
    assert docx_pdf.output_path.read_bytes().startswith(b"%PDF-")
    assert pptx_pdf.output_path.read_bytes().startswith(b"%PDF-")

    before = xlsx_path.read_bytes()
    result = recalculate_xlsx(xlsx_path)
    assert result.output_path == xlsx_path
    assert xlsx_path.read_bytes() != before
    recalculated = load_workbook(xlsx_path, data_only=True)
    assert recalculated.active["A3"].value == 42
    recalculated.close()


@pytest.mark.asyncio
async def test_real_script_to_ppt_preview_and_manual_version_lineage(
    tmp_path: Path,
    monkeypatch,
) -> None:
    """Exercise real PPTX rendering, revisioning, and the PDF derivation API."""

    _require_soffice()
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js is not installed in this test environment")
    repository = Path(__file__).resolve().parents[2]
    pptxgenjs = repository / "console" / "node_modules" / "pptxgenjs"
    if not pptxgenjs.exists():
        pytest.skip("PptxGenJS is not installed in this test environment")

    workspace = tmp_path / "workspace"
    task_output = workspace / "task-output"
    task_output.mkdir(parents=True)
    script = task_output / "build-deck.cjs"
    deck = task_output / "quarterly-review.pptx"
    script.write_text(
        "\n".join(
            [
                f"const pptxgen = require({json.dumps(str(pptxgenjs))});",
                "const deck = new pptxgen();",
                "deck.layout = 'LAYOUT_WIDE';",
                "for (let page = 1; page <= 3; page += 1) {",
                "  const slide = deck.addSlide();",
                "  slide.addText(`Quarterly review / page ${page}`, {",
                "    x: 0.8, y: 0.7, w: 10.8, h: 0.7, fontSize: 28",
                "  });",
                "  slide.addText(`Verified content ${page}`, {",
                "    x: 0.8, y: 1.8, w: 10.8, h: 0.5, fontSize: 18",
                "  });",
                "}",
                "deck.writeFile({ fileName: process.argv[2] });",
            ]
        ),
        encoding="utf-8",
    )
    subprocess.run(
        [node, str(script), str(deck)],
        cwd=repository / "console",
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert len(Presentation(deck).slides) == 3

    version_one = add_work(
        workspace,
        source_path=deck,
        allowed_roots=[workspace],
        title="Quarterly review",
        source_chat_id="chat-ap",
        source_run_id="run-ap",
    )
    preview_one = get_or_create_work_preview(workspace, version_one.id)
    page_three = get_work_preview_page(
        workspace,
        version_one.id,
        page=3,
        source_digest=version_one.source_digest or "",
    )

    assert preview_one.page_count == 3
    assert page_three.is_file()
    assert page_three.name == "slide-3.png"

    edit_copy, parent_digest = create_work_edit_copy(
        workspace,
        version_one,
        version_one.path,
    )
    manually_edited = Presentation(edit_copy)
    manually_edited.slides[2].shapes.add_textbox(0.8, 2.8, 10.8, 0.5).text = (
        "Manual editor save"
    )
    manually_edited.save(edit_copy)
    version_two = add_work(
        workspace,
        source_path=edit_copy,
        allowed_roots=[edit_copy.parent],
    )
    preview_two = get_or_create_work_preview(workspace, version_two.id)
    page_three_v2 = get_work_preview_page(
        workspace,
        version_two.id,
        page=3,
        source_digest=version_two.source_digest or "",
    )

    assert version_two.version == 2
    assert version_two.parent_work_id == version_one.id
    assert version_two.parent_digest == parent_digest
    assert version_two.lineage_root_id == version_one.id
    assert version_two.source_digest != version_one.source_digest
    assert preview_two.page_count == 3
    assert preview_two.source_digest == version_two.source_digest
    assert page_three_v2 != page_three
    assert page_three_v2.parent != page_three.parent
    with pytest.raises(ValueError, match="requested digest"):
        get_work_preview_page(
            workspace,
            version_two.id,
            page=3,
            source_digest=version_one.source_digest or "",
        )
    assert [(item.id, item.version) for item in list_works(workspace)] == [
        (version_two.id, 2),
        (version_one.id, 1),
    ]

    async def _agent_for_request(_request):
        return SimpleNamespace(workspace_dir=str(workspace))

    monkeypatch.setattr(
        works_router,
        "get_agent_for_request",
        _agent_for_request,
    )
    derived = await works_router.create_work_derivation(
        version_two.id,
        works_router.WorkDerivationRequest(chat_id="chat-ap"),
        SimpleNamespace(),
    )

    assert derived.kind == "pdf"
    assert Path(derived.path).read_bytes().startswith(b"%PDF-")
    assert derived.source_digest is not None
    assert derived.derivation is not None
    assert derived.derivation.source_work_id == version_two.id
    assert derived.derivation.source_artifact_id == version_two.artifact_id
    assert derived.derivation.source_version == version_two.version
    assert derived.derivation.source_digest == version_two.source_digest
    assert derived.derivation.output_digest == derived.source_digest
    assert derived.derivation.converter_output_digest.startswith("sha256:")
    assert "libreoffice" in derived.derivation.converter.lower()
