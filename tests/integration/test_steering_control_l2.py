# -*- coding: utf-8 -*-
"""L2 proof for the local HTTP -> durable Inbox -> Native request path."""

from __future__ import annotations

import asyncio
import base64
import codecs
import hashlib
import json
import logging
import multiprocessing
import socket
import struct
import time
import zlib
from contextlib import asynccontextmanager, closing
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
import uvicorn
from fastapi import FastAPI
from pptx import Presentation
from reportlab.pdfgen import canvas

from lantai.app.channels.console.channel import ConsoleChannel
from lantai.app.experience.runtime_runs import (
    record_chat_runtime_run_started,
    touch_chat_runtime_run_heartbeat,
)
from lantai.app.runner.api import ensure_chat_topic_workspace
from lantai.app.runner.manager import ChatManager
from lantai.app.runner.models import ChatSpec
from lantai.app.runner.repo.json_repo import JsonChatRepository
from lantai.app.runner.task_tracker import TaskTracker
from lantai.app.works import add_work, create_work_edit_copy
from lantai.agents.persona.perceptible_presence import (
    PresenceState,
    plan_persona_turn,
)
from lantai.agents.persona.subject_continuity import (
    SubjectContinuityState,
    SubjectEvent,
    SubjectReflection,
    SubjectTurnProjection,
)
from lantai.agents.persona.user_understanding import (
    TurnUnderstandingProjection,
    UnderstandingAtom,
)
from lantai.model_context_profile import ModelContextProfile
from lantai.runtime.adapters.native_runtime import NativeLantaiRuntime
from lantai.runtime.attachment_context import attachment_context_text
from lantai.runtime.continuation import (
    build_continuation_capsule,
    rollover_model_context,
    snapshot_continuation_artifacts,
)
from lantai.runtime.cost_accounting import (
    BudgetLedger,
    CostEvidence,
    RuntimeBudgetAdmissionError,
    finalize_model_dispatch_budget,
    reserve_model_dispatch_budget,
)
from lantai.runtime.inbox_admission import RuntimeInboxAdmission
from lantai.runtime.input import RuntimeInput, RuntimeRequestContext
from lantai.runtime.messages import RuntimeMessage
from lantai.runtime.model_wrappers import TokenRecordingModelClient
from lantai.runtime.models import (
    ModelContentBlock,
    ModelMessage,
    ModelRequest,
    ModelResponse,
    ModelStreamEvent,
    ModelUsage,
)
from lantai.runtime.runs import (
    RuntimeAutonomyGrant,
    RuntimeModelPriceBound,
    RuntimeRunReducerAction,
    RuntimeWakeIntent,
    build_chat_task_contract,
    evaluate_runtime_autonomy,
    evaluate_runtime_wake,
    reduce_runtime_run,
    revise_chat_task_contract,
)
from lantai.runtime.store import RuntimeRunStore
from lantai.runtime.tools import LantaiToolRegistry, ToolCall, ToolSpec
from lantai.runtime.turn_control import RunInboxItem, RuntimeTurnControlStore


_ONE_PIXEL_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNgYAAAAAMA"
    "ASsJTYQAAAAASUVORK5CYII="
)


def _write_labeled_pptx(
    path: Path,
    *,
    label: str,
    slides: int = 3,
) -> None:
    deck = Presentation()
    for number in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = f"{label} page {number + 1}"
    deck.save(path)


def _png_with_label(label: str) -> bytes:
    """Return a valid one-pixel PNG with a distinct tEXt marker."""

    data = b"l2-label\x00" + label.encode("ascii")
    chunk_type = b"tEXt"
    chunk = (
        struct.pack(">I", len(data))
        + chunk_type
        + data
        + struct.pack(">I", zlib.crc32(chunk_type + data) & 0xFFFFFFFF)
    )
    return _ONE_PIXEL_PNG[:-12] + chunk + _ONE_PIXEL_PNG[-12:]


class _ChannelManager:
    def __init__(self, channel: ConsoleChannel) -> None:
        self._channel = channel

    async def get_channel(self, name: str):
        return self._channel if name == "console" else None


async def _never_process(_request):
    raise AssertionError("control-input admission must not start another loop")
    yield  # pragma: no cover


class _RecordingModelClient:
    provider_id = "l2-local"
    model_name = "deterministic-vision"
    endpoint = SimpleNamespace(supports_multimodal=True)

    def __init__(self) -> None:
        self.requests: list[ModelRequest] = []

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        message = ModelMessage(
            id=f"l2-model-answer-{len(self.requests)}",
            role="assistant",
            content=[
                ModelContentBlock(
                    type="text",
                    text="已收到并使用补充图片，且未调用外部工具。",
                )
            ],
        )
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message),
        )

    async def complete(self, request: ModelRequest) -> ModelResponse:
        self.requests.append(request)
        return ModelResponse(message=ModelMessage.text("assistant", "done"))


class _FirstRequestGateModelClient(_RecordingModelClient):
    def __init__(self) -> None:
        super().__init__()
        self.first_request_started = asyncio.Event()
        self.release_first_request = asyncio.Event()

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        if len(self.requests) == 1:
            self.first_request_started.set()
            await self.release_first_request.wait()
        message = ModelMessage.text(
            "assistant",
            "已按顺序处理这条输入。",
            id=f"l2-gated-answer-{len(self.requests)}",
        )
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message),
        )


class _RedirectableRecordingModelClient:
    provider_id = "l2-local"
    model_name = "deterministic-redirect"

    def __init__(self) -> None:
        self.requests: list[ModelRequest] = []
        self.request_times: list[float] = []
        self.first_surface = asyncio.Event()
        self.first_stream_closed = asyncio.Event()

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        self.request_times.append(time.perf_counter())
        if len(self.requests) == 1:
            try:
                yield ModelStreamEvent(
                    kind="message_start",
                    message_id="l2-old-epoch",
                )
                yield ModelStreamEvent(
                    kind="text_delta",
                    message_id="l2-old-epoch",
                    delta="旧方向中已向用户展示的内容",
                )
                self.first_surface.set()
                await asyncio.Event().wait()
            finally:
                self.first_stream_closed.set()
            return
        message = ModelMessage.text(
            "assistant",
            "已按立即调整的新方向继续。",
            id="l2-new-epoch",
        )
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message),
        )

    async def complete(self, request: ModelRequest) -> ModelResponse:
        raise AssertionError("redirect proof must remain on the stream path")


class _SafeYieldBoundaryModelClient:
    """Expose precise pre-surface and in-stream redirect boundaries."""

    provider_id = "l2-local"
    model_name = "deterministic-safe-yield"

    def __init__(self, phase: str) -> None:
        self.phase = phase
        self.requests: list[ModelRequest] = []
        self.request_times: list[float] = []
        self.boundary_reached = asyncio.Event()
        self.old_stream_closed = asyncio.Event()
        self.late_tool_returned = False

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        self.request_times.append(time.perf_counter())
        if len(self.requests) == 1:
            try:
                if self.phase == "first_event":
                    self.boundary_reached.set()
                    await asyncio.Event().wait()
                elif self.phase == "thinking":
                    yield ModelStreamEvent(
                        kind="message_start",
                        message_id="l2-old-thinking-epoch",
                    )
                    yield ModelStreamEvent(
                        kind="thinking_delta",
                        message_id="l2-old-thinking-epoch",
                        delta="旧方向仍在思考",
                    )
                    self.boundary_reached.set()
                    await asyncio.Event().wait()
                elif self.phase == "partial_tool_json":
                    yield ModelStreamEvent(
                        kind="tool_call_delta",
                        message_id="l2-old-partial-tool-epoch",
                        delta='{"value":"unfinished',
                        raw={
                            "call_id": "l2-partial-tool-call",
                            "name": "record_effect",
                            "arguments_delta": '{"value":"unfinished',
                        },
                    )
                    self.boundary_reached.set()
                    await asyncio.Event().wait()
                elif self.phase == "late_tool_call":
                    self.boundary_reached.set()
                    try:
                        await asyncio.Event().wait()
                    except asyncio.CancelledError:
                        # Model a provider transport that races one completed
                        # old-epoch event back after local retirement. The
                        # cancelled __anext__ result must never re-enter the
                        # active model round or execute its tool.
                        self.late_tool_returned = True
                        self.old_stream_closed.set()
                        yield ModelStreamEvent(
                            kind="tool_call_completed",
                            message_id="l2-old-late-tool-epoch",
                            tool_call=ToolCall(
                                id="l2-late-tool-call",
                                name="record_effect",
                                arguments={"value": "old-epoch"},
                            ),
                        )
                else:  # pragma: no cover - guarded by parametrization
                    raise AssertionError(
                        f"unknown safe-yield phase: {self.phase}"
                    )
            finally:
                self.old_stream_closed.set()
            return
        message = ModelMessage.text(
            "assistant",
            "已按新方向继续，旧工具未执行。",
            id=f"l2-safe-yield-answer-{self.phase}",
        )
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message),
        )

    async def complete(self, request: ModelRequest) -> ModelResponse:
        raise AssertionError("safe-yield proof must remain on the stream path")


class _HangingCloseModelStream:
    """A provider iterator whose event wait cancels but transport close hangs."""

    def __init__(self, client: "_HangingCloseModelClient") -> None:
        self._client = client

    def __aiter__(self):
        return self

    async def __anext__(self) -> ModelStreamEvent:
        self._client.boundary_reached.set()
        await asyncio.Event().wait()
        raise StopAsyncIteration  # pragma: no cover

    async def aclose(self) -> None:
        self._client.close_started.set()
        await self._client.release_close.wait()
        self._client.close_finished.set()


class _HangingCloseModelClient:
    provider_id = "l2-local"
    model_name = "deterministic-hanging-close"

    def __init__(self) -> None:
        self.requests: list[ModelRequest] = []
        self.request_times: list[float] = []
        self.boundary_reached = asyncio.Event()
        self.close_started = asyncio.Event()
        self.close_finished = asyncio.Event()
        self.release_close = asyncio.Event()

    def stream(self, request: ModelRequest):
        self.requests.append(request)
        self.request_times.append(time.perf_counter())
        if len(self.requests) == 1:
            return _HangingCloseModelStream(self)

        async def completed_stream():
            message = ModelMessage.text(
                "assistant",
                "已在有界清理后按新方向继续。",
                id="l2-hanging-close-successor",
            )
            yield ModelStreamEvent(
                kind="message_completed",
                message_id=message.id,
                response=ModelResponse(message=message),
            )

        return completed_stream()

    async def complete(self, request: ModelRequest) -> ModelResponse:
        raise AssertionError("hanging-close proof must remain on stream path")


class _RapidRedirectModelClient:
    """Keep each old epoch open until three FIFO redirects have arrived."""

    provider_id = "l2-local"
    model_name = "deterministic-rapid-redirect"

    def __init__(self) -> None:
        self.requests: list[ModelRequest] = []
        self.request_times: list[float] = []
        self.first_boundary_reached = asyncio.Event()
        self.closed_request_ordinals: list[int] = []

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        self.request_times.append(time.perf_counter())
        ordinal = len(self.requests)
        intent_revision = int(request.metadata.get("intent_revision") or 0)
        if intent_revision < 3:
            try:
                if ordinal == 1:
                    self.first_boundary_reached.set()
                await asyncio.Event().wait()
            finally:
                self.closed_request_ordinals.append(ordinal)
            return
        message = ModelMessage.text(
            "assistant",
            "已按第三条、也是最新方向继续。",
            id="l2-rapid-redirect-successor",
        )
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message),
        )

    async def complete(self, request: ModelRequest) -> ModelResponse:
        raise AssertionError("rapid redirect proof must remain on stream path")


class _CostedRedirectModelClient:
    """Emit auditable cost evidence on both sides of a redirect."""

    provider_id = "l2-cost-provider"
    model_name = "deterministic-costed-redirect"

    def __init__(self) -> None:
        self.requests: list[ModelRequest] = []
        self.old_usage_observed = asyncio.Event()
        self.old_stream_closed = asyncio.Event()

    async def stream(self, request: ModelRequest):
        self.requests.append(request)
        ordinal = len(self.requests)
        if ordinal == 1:
            try:
                yield ModelStreamEvent(
                    kind="usage",
                    usage=ModelUsage(
                        input_tokens=101,
                        output_tokens=11,
                        total_tokens=112,
                        cost_cny=0.001111,
                        raw={"billing_record_id": "l2-old-billing"},
                    ),
                )
                self.old_usage_observed.set()
                await asyncio.Event().wait()
            finally:
                self.old_stream_closed.set()
            return
        message = ModelMessage.text(
            "assistant",
            "已按新方向继续，并保留两次调用的费用证据。",
            id="l2-cost-successor",
        )
        usage = ModelUsage(
            input_tokens=121,
            output_tokens=22,
            total_tokens=143,
            cost_cny=0.002222,
            raw={"billing_record_id": "l2-new-billing"},
        )
        yield ModelStreamEvent(kind="usage", usage=usage)
        yield ModelStreamEvent(
            kind="message_completed",
            message_id=message.id,
            response=ModelResponse(message=message, usage=usage),
            usage=usage,
        )

    async def complete(self, request: ModelRequest) -> ModelResponse:
        raise AssertionError("cost proof must remain on the stream path")


def _sse_payloads(text: str) -> list[dict]:
    return [
        json.loads(line.removeprefix("data: "))
        for line in text.splitlines()
        if line.startswith("data: ") and line != "data: [DONE]"
    ]


def _text_steer_body(
    *,
    chat_id: str,
    session_id: str,
    command_id: str,
    text: str,
    intent: str = "supplement",
    replaces_command_id: str | None = None,
    expected_run_id: str | None = None,
    expected_attempt_id: str | None = None,
) -> dict:
    body = {
        "input": [
            {
                "role": "user",
                "content": [{"type": "text", "text": text}],
            }
        ],
        "chat_id": chat_id,
        "session_id": session_id,
        "user_id": "default",
        "channel": "console",
        "stream": True,
        "operation": "steer",
        "meta": {
            "client_turn_id": command_id,
            "control_intent": intent,
        },
    }
    if replaces_command_id is not None:
        body["meta"]["replaces_command_id"] = replaces_command_id
    if expected_run_id is not None:
        body["meta"]["expected_run_id"] = expected_run_id
    if expected_attempt_id is not None:
        body["meta"]["expected_attempt_id"] = expected_attempt_id
    return body


def _receipt_from_response(response: httpx.Response) -> dict:
    return next(
        payload["receipt"]
        for payload in _sse_payloads(response.text)
        if isinstance(payload.get("receipt"), dict)
    )


class _ReceiptStreamFaultInjector:
    """Keep selected real HTTP receipt streams open and reshape their frames."""

    def __init__(self, app, modes: dict[str, str]) -> None:
        self.app = app
        self.modes = modes
        self.frame_sent: dict[str, asyncio.Event] = {}
        self.release: dict[str, asyncio.Event] = {}

    def _event(self, events: dict[str, asyncio.Event], command_id: str):
        return events.setdefault(command_id, asyncio.Event())

    async def wait_until_frame_sent(self, command_id: str) -> None:
        await self._event(self.frame_sent, command_id).wait()

    def release_stream(self, command_id: str) -> None:
        self._event(self.release, command_id).set()

    async def __call__(self, scope, receive, send) -> None:
        if not (
            scope.get("type") == "http"
            and scope.get("method") == "POST"
            and scope.get("path") == "/console/chat"
        ):
            await self.app(scope, receive, send)
            return

        async def send_with_fault(message):
            if message.get("type") != "http.response.body" or not message.get(
                "body"
            ):
                await send(message)
                return

            body = bytes(message["body"])
            text = body.decode("utf-8")
            encoded = "\n".join(
                line.removeprefix("data:").lstrip()
                for line in text.splitlines()
                if line.startswith("data:")
            ).strip()
            payload = json.loads(encoded)
            receipt = payload.get("receipt") or {}
            command_id = str(receipt.get("command_id") or "")
            mode = self.modes.get(command_id)
            if mode is None:
                await send(message)
                return

            if mode == "fragmented_multiline":
                receipt["unknown_future_field"] = {
                    "safe_to_ignore": True,
                }
                encoded = json.dumps(
                    payload,
                    ensure_ascii=False,
                    separators=(",", ":"),
                )
                split_at = encoded.index('"receipt"')
                wire = (
                    f"data: {encoded[:split_at]}\r\n"
                    f"data: {encoded[split_at:]}\r\n\r\n"
                ).encode("utf-8")
                cut_points = (7, 31, 73, max(74, len(wire) - 9))
                chunks: list[bytes] = []
                start = 0
                for end in cut_points:
                    if start < end < len(wire):
                        chunks.append(wire[start:end])
                        start = end
                chunks.append(wire[start:])
            elif mode == "unknown_schema":
                receipt["schema_version"] = 99
                wire = (
                    "data: "
                    + json.dumps(
                        payload,
                        ensure_ascii=False,
                        separators=(",", ":"),
                    )
                    + "\n\n"
                ).encode("utf-8")
                chunks = [wire]
            elif mode == "hold_open":
                chunks = [body]
            else:  # pragma: no cover - test construction guard
                raise AssertionError(f"unknown receipt stream mode: {mode}")

            for chunk in chunks:
                await send(
                    {
                        "type": "http.response.body",
                        "body": chunk,
                        "more_body": True,
                    }
                )
                await asyncio.sleep(0)
            self._event(self.frame_sent, command_id).set()
            await self._event(self.release, command_id).wait()

        await self.app(scope, receive, send_with_fault)


async def _read_turn_receipt_before_eof(response: httpx.Response) -> dict:
    """Parse one SSE receipt incrementally without consuming transport EOF."""

    decoder = codecs.getincrementaldecoder("utf-8")()
    buffered = ""
    async for chunk in response.aiter_bytes():
        buffered += decoder.decode(chunk)
        delimiter = buffered.find("\n\n")
        crlf_delimiter = buffered.find("\r\n\r\n")
        if crlf_delimiter >= 0 and (
            delimiter < 0 or crlf_delimiter < delimiter
        ):
            delimiter = crlf_delimiter
            delimiter_size = 4
        else:
            delimiter_size = 2
        if delimiter < 0:
            continue
        event = buffered[:delimiter]
        buffered = buffered[delimiter + delimiter_size :]
        encoded = "\n".join(
            line.removeprefix("data:").lstrip()
            for line in event.splitlines()
            if line.startswith("data:")
        ).strip()
        payload = json.loads(encoded)
        receipt = payload.get("receipt")
        if not isinstance(receipt, dict):
            continue
        schema_version = receipt.get("schema_version", 1)
        if schema_version not in (1, 2):
            raise ValueError(
                f"Unsupported TurnReceipt schema version {schema_version}"
            )
        return receipt
    raise AssertionError("receipt stream ended without a receipt")


async def _collect_runtime_events(
    runtime: NativeLantaiRuntime,
    runtime_input: RuntimeInput,
) -> list:
    return [event async for event in runtime.run_turn(runtime_input)]


async def _register_active_external_task(tracker, run_key, **kwargs):
    """Model the legacy producer's cancel/finally/unregister ownership."""
    registered = asyncio.get_running_loop().create_future()
    released = asyncio.Event()

    async def producer():
        generation = await tracker.register_external_task(
            run_key, owner_task=asyncio.current_task(),
            cleanup=released.set, **kwargs,
        )
        registered.set_result(generation)
        try:
            await released.wait()
        finally:
            await tracker.unregister_external_task(run_key)

    owner = asyncio.create_task(producer())
    try:
        return await registered
    except BaseException:
        owner.cancel()
        await asyncio.gather(owner, return_exceptions=True)
        raise


async def _prepare_active_run(
    *,
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    session_id: str,
    run_id: str,
    objective: str,
    name: str,
    acceptance: tuple[str, ...] = (),
) -> SimpleNamespace:
    from lantai.app.routers import console as console_router

    agent_home = tmp_path / "agent"
    chat_manager = ChatManager(
        repo=JsonChatRepository(tmp_path / "chats.json"),
    )
    tracker = TaskTracker()
    console_channel = ConsoleChannel(
        process=_never_process,
        enabled=True,
        bot_prefix="",
        workspace_dir=agent_home,
    )
    workspace = SimpleNamespace(
        agent_id="default",
        workspace_dir=agent_home,
        chat_manager=chat_manager,
        task_tracker=tracker,
        channel_manager=_ChannelManager(console_channel),
        runner=SimpleNamespace(_chat_manager=chat_manager),
    )

    async def fake_get_agent_for_request(_request):
        return workspace

    monkeypatch.setattr(
        console_router,
        "get_agent_for_request",
        fake_get_agent_for_request,
    )
    chat = await chat_manager.get_or_create_chat(
        session_id,
        "default",
        "console",
        name=name,
    )
    binding = await ensure_chat_topic_workspace(
        workspace,
        chat_manager,
        chat,
        agent_id="default",
    )
    chat = binding.chat
    state_workspace = Path(binding.state_workspace_dir)
    task_contract = build_chat_task_contract(
        objective=objective,
        client_turn_id=f"{run_id}-initial-turn",
        complexity="light",
        work_method_id="standard",
        run_id=run_id,
        attempt_id=f"{run_id}:chat-attempt-1",
        acceptance=acceptance,
    )
    started = await record_chat_runtime_run_started(
        workspace_dir=state_workspace,
        agent_id="default",
        chat_id=chat.id,
        session_id=session_id,
        run_id=run_id,
        effective_workspace_dir=binding.effective_workspace_dir,
        task_contract=task_contract,
    )
    await chat_manager.merge_chat_meta(
        chat.id,
        {
            "last_turn_run_id": run_id,
            "state_workspace_dir": str(state_workspace),
        },
    )
    tracker_generation = await _register_active_external_task(
        tracker, chat.id,
        attempt_id=started.attempt_id,
        root_session_id=session_id,
    )
    return SimpleNamespace(
        workspace=workspace,
        chat_manager=chat_manager,
        chat=chat,
        binding=binding,
        state_workspace=state_workspace,
        run_id=run_id,
        started=started,
        tracker=tracker,
        tracker_generation=tracker_generation,
    )


def _runtime_input_for_harness(
    harness: SimpleNamespace,
    *,
    session_id: str,
    objective: str,
) -> RuntimeInput:
    return RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default",
            session_id=session_id,
            user_id="default",
            run_id=harness.run_id,
            workspace_dir=harness.binding.effective_workspace_dir,
        ),
        messages=[RuntimeMessage.text("user", objective)],
        metadata={
            "state_workspace_dir": str(harness.state_workspace),
            "attempt_id": harness.started.attempt_id,
            "runtime_generation": harness.tracker_generation,
            "runtime_revision": harness.started.revision,
            "runtime_lease_owner": harness.started.lease_owner,
            "task_contract": harness.started.task_contract.model_dump(
                mode="json"
            ),
        },
    )


@asynccontextmanager
async def _loopback_console_http(asgi_app: FastAPI | None = None):
    from lantai.app.routers import console as console_router

    app = asgi_app
    if app is None:
        app = FastAPI()
        app.include_router(console_router.router)
    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    try:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.bind(("127.0.0.1", 0))
        sock.listen(128)
        port = int(sock.getsockname()[1])
        server = uvicorn.Server(
            uvicorn.Config(
                app,
                host="127.0.0.1",
                port=port,
                log_level="warning",
                lifespan="off",
            )
        )
        server_task = asyncio.create_task(server.serve(sockets=[sock]))
        try:
            await asyncio.wait_for(_wait_until_started(server), timeout=3)
            async with httpx.AsyncClient(
                base_url=f"http://127.0.0.1:{port}",
                timeout=5,
                trust_env=False,
            ) as http:
                yield http, port
        finally:
            server.should_exit = True
            await asyncio.wait_for(server_task, timeout=5)
    finally:
        sock.close()


async def _post_and_drop_after_headers(
    *,
    port: int,
    path: str,
    body: dict,
) -> str:
    """Submit a command, then discard the response before its receipt frame."""

    payload = json.dumps(body, ensure_ascii=False).encode("utf-8")
    reader, writer = await asyncio.open_connection("127.0.0.1", port)
    request = (
        f"POST {path} HTTP/1.1\r\n"
        f"Host: 127.0.0.1:{port}\r\n"
        "Content-Type: application/json\r\n"
        "Accept: text/event-stream\r\n"
        f"Content-Length: {len(payload)}\r\n"
        "Connection: close\r\n"
        "\r\n"
    ).encode("ascii") + payload
    writer.write(request)
    await writer.drain()
    response_headers = await asyncio.wait_for(
        reader.readuntil(b"\r\n\r\n"),
        timeout=3,
    )
    status_line = response_headers.split(b"\r\n", 1)[0].decode("ascii")
    writer.close()
    await writer.wait_closed()
    return status_line


async def _upload_and_drop_after_headers(
    *,
    port: int,
    filename: str,
    media_type: str,
    data: bytes,
) -> str:
    """Upload real multipart bytes but discard the success response body."""

    boundary = "----lantai-l2-lost-upload-receipt"
    payload = (
        (
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
            f"Content-Type: {media_type}\r\n"
            "\r\n"
        ).encode("utf-8")
        + data
        + f"\r\n--{boundary}--\r\n".encode("ascii")
    )
    reader, writer = await asyncio.open_connection("127.0.0.1", port)
    request = (
        "POST /console/upload HTTP/1.1\r\n"
        f"Host: 127.0.0.1:{port}\r\n"
        f"Content-Type: multipart/form-data; boundary={boundary}\r\n"
        f"Content-Length: {len(payload)}\r\n"
        "Connection: close\r\n"
        "\r\n"
    ).encode("ascii") + payload
    writer.write(request)
    await writer.drain()
    response_headers = await asyncio.wait_for(
        reader.readuntil(b"\r\n\r\n"),
        timeout=3,
    )
    status_line = response_headers.split(b"\r\n", 1)[0].decode("ascii")
    writer.close()
    await writer.wait_closed()
    return status_line


@pytest.mark.asyncio
async def test_local_http_image_only_input_reaches_same_native_run(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Exercise a real loopback HTTP socket and the production Native owner."""
    from lantai.app.routers import console as console_router

    agent_home = tmp_path / "agent"
    chat_manager = ChatManager(
        repo=JsonChatRepository(tmp_path / "chats.json"),
    )
    tracker = TaskTracker()
    console_channel = ConsoleChannel(
        process=_never_process,
        enabled=True,
        bot_prefix="",
        workspace_dir=agent_home,
    )
    workspace = SimpleNamespace(
        agent_id="default",
        workspace_dir=agent_home,
        chat_manager=chat_manager,
        task_tracker=tracker,
        channel_manager=_ChannelManager(console_channel),
        runner=SimpleNamespace(_chat_manager=chat_manager),
    )

    async def fake_get_agent_for_request(_request):
        return workspace

    monkeypatch.setattr(
        console_router,
        "get_agent_for_request",
        fake_get_agent_for_request,
    )

    session_id = "l2-steering-session"
    chat = await chat_manager.get_or_create_chat(
        session_id,
        "default",
        "console",
        name="L2 steering proof",
    )
    binding = await ensure_chat_topic_workspace(
        workspace,
        chat_manager,
        chat,
        agent_id="default",
    )
    chat = binding.chat
    state_workspace = Path(binding.state_workspace_dir)
    run_id = "l2-steering-run"
    task_contract = build_chat_task_contract(
        objective="观察当前任务；不使用外部工具。",
        client_turn_id="l2-initial-turn",
        complexity="light",
        work_method_id="standard",
        run_id=run_id,
        attempt_id=f"{run_id}:chat-attempt-1",
    )
    started = await record_chat_runtime_run_started(
        workspace_dir=state_workspace,
        agent_id="default",
        chat_id=chat.id,
        session_id=session_id,
        run_id=run_id,
        effective_workspace_dir=binding.effective_workspace_dir,
        task_contract=task_contract,
    )
    await chat_manager.merge_chat_meta(
        chat.id,
        {
            "last_turn_run_id": run_id,
            "state_workspace_dir": str(state_workspace),
        },
    )
    tracker_generation = await _register_active_external_task(
        tracker, chat.id,
        attempt_id=started.attempt_id,
        root_session_id=session_id,
    )

    app = FastAPI()
    app.include_router(console_router.router)
    with closing(socket.socket(socket.AF_INET, socket.SOCK_STREAM)) as sock:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.bind(("127.0.0.1", 0))
        sock.listen(128)
        port = int(sock.getsockname()[1])
        server = uvicorn.Server(
            uvicorn.Config(
                app,
                host="127.0.0.1",
                port=port,
                log_level="warning",
                lifespan="off",
            )
        )
        server_task = asyncio.create_task(server.serve(sockets=[sock]))
        try:
            await asyncio.wait_for(_wait_until_started(server), timeout=3)
            async with httpx.AsyncClient(
                base_url=f"http://127.0.0.1:{port}",
                timeout=5,
                trust_env=False,
            ) as http:
                upload_started = time.perf_counter()
                uploaded = await http.post(
                    "/console/upload",
                    files={
                        "file": ("guidance.png", _ONE_PIXEL_PNG, "image/png")
                    },
                )
                upload_finished = time.perf_counter()
                assert uploaded.status_code == 200
                upload_payload = uploaded.json()

                command_id = "l2-image-only-command"
                steer_body = {
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "image",
                                    "image_url": upload_payload["url"],
                                    "filename": upload_payload["file_name"],
                                    "media_type": "image/png",
                                    "file_size": upload_payload["size"],
                                    "stored_name": upload_payload[
                                        "stored_name"
                                    ],
                                    "attachment_digest": upload_payload[
                                        "digest"
                                    ],
                                    "attachment_version": upload_payload[
                                        "version"
                                    ],
                                    "attachment_upload_receipt": (
                                        upload_payload["upload_receipt"]
                                    ),
                                }
                            ],
                        }
                    ],
                    "chat_id": chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                    },
                }
                admission_started = time.perf_counter()
                admitted = await http.post("/console/chat", json=steer_body)
                admission_finished = time.perf_counter()
                assert admitted.status_code == 200
                receipt = _sse_payloads(admitted.text)[-1]["receipt"]
                assert receipt["run_id"] == run_id
                assert receipt["attempt_id"] == started.attempt_id
                assert receipt["admission_status"] == "admitted"
                assert receipt["execution_status"] == "pending"
                assert (
                    receipt["next_waiting_reason"] == "awaiting_model_boundary"
                )

                retried = await http.post("/console/chat", json=steer_body)
                assert retried.status_code == 200
                retry_receipts = [
                    payload["receipt"]
                    for payload in _sse_payloads(retried.text)
                    if isinstance(payload.get("receipt"), dict)
                ]
                assert retry_receipts == [receipt]

                conflict_body = json.loads(json.dumps(steer_body))
                conflict_body["input"][0]["content"].insert(
                    0,
                    {"type": "text", "text": "different payload"},
                )
                conflicted = await http.post(
                    "/console/chat",
                    json=conflict_body,
                )
                assert conflicted.status_code == 409
                assert conflicted.json()["detail"]["code"] == (
                    "control_input_conflict"
                )

                client = _RecordingModelClient()
                runtime = NativeLantaiRuntime(
                    model_client_factory=lambda _agent_id: client,
                    tool_registry=LantaiToolRegistry(),
                )
                runtime_input = RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=binding.effective_workspace_dir,
                    ),
                    messages=[
                        RuntimeMessage.text(
                            "user",
                            "观察当前任务；不使用外部工具。",
                        )
                    ],
                    metadata={
                        "state_workspace_dir": str(state_workspace),
                        "attempt_id": started.attempt_id,
                        "runtime_generation": tracker_generation,
                        "runtime_revision": started.revision,
                        "runtime_lease_owner": started.lease_owner,
                        "task_contract": started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                )
                owner_started = time.perf_counter()
                events = [
                    event async for event in runtime.run_turn(runtime_input)
                ]
                owner_finished = time.perf_counter()

                assert events[-1].type == "turn.completed"
                assert len(client.requests) == 1
                request = client.requests[0]
                initial_constraints = [
                    message.text_content()
                    for message in request.messages
                    if message.role == "user"
                ]
                assert "观察当前任务；不使用外部工具。" in initial_constraints
                bound_images = [
                    block
                    for message in request.messages
                    for block in message.content
                    if block.type == "image"
                ]
                assert len(bound_images) == 1
                image_data = bound_images[0].data
                assert image_data["digest"].startswith("sha256:")
                assert image_data["version"] == 1
                assert image_data["grant_ref"]
                assert not any(
                    event.type.startswith("tool.") for event in events
                )

                status = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": chat.id,
                        "command_id": command_id,
                    },
                )
                assert status.status_code == 200
                final_receipt = status.json()
                assert final_receipt["execution_status"] == "applied"
                assert final_receipt["model_request_id"] == (
                    f"{run_id}:model:1"
                )
                assert final_receipt["model_epoch"] == 1
                assert final_receipt["intent_revision"] == 1

                late_body = {
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "终态之后不能误投给旧Run。",
                                }
                            ],
                        }
                    ],
                    "chat_id": chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": "l2-after-terminal-command",
                        "control_intent": "redirect",
                    },
                }
                late_after_terminal = await http.post(
                    "/console/chat",
                    json=late_body,
                )
                assert late_after_terminal.status_code == 409
                late_detail = late_after_terminal.json()["detail"]
                assert late_detail["code"] == "target_ended"
                assert late_detail["run_id"] == run_id
                assert late_detail["attempt_id"] == started.attempt_id

                inbox = RuntimeTurnControlStore(state_workspace).read_inbox(
                    run_id
                )
                assert [item.command_id for item in inbox] == [command_id]
                contracts = RuntimeRunStore(state_workspace).list_contracts()
                assert [contract.run_id for contract in contracts] == [run_id]

                trace = {
                    "schema_version": 1,
                    "level": "L2",
                    "transport": "loopback_http_socket",
                    "command_id": command_id,
                    "chat_id": chat.id,
                    "run_id": run_id,
                    "attempt_id": started.attempt_id,
                    "tracker_generation": tracker_generation,
                    "http_status": {
                        "upload": uploaded.status_code,
                        "admission": admitted.status_code,
                        "idempotent_retry": retried.status_code,
                        "conflict": conflicted.status_code,
                        "query": status.status_code,
                        "late_after_terminal": (
                            late_after_terminal.status_code
                        ),
                    },
                    "counts": {
                        "durable_inputs": len(inbox),
                        "model_requests": len(client.requests),
                        "runtime_runs": len(contracts),
                        "tool_effects": 0,
                    },
                    "latency_ms": {
                        "upload": round(
                            (upload_finished - upload_started) * 1000,
                            3,
                        ),
                        "content_ready_to_admitted": round(
                            (admission_finished - admission_started) * 1000,
                            3,
                        ),
                        "owner_to_terminal": round(
                            (owner_finished - owner_started) * 1000,
                            3,
                        ),
                    },
                    "attachment": {
                        "media_type": "image/png",
                        "size_bytes": len(_ONE_PIXEL_PNG),
                        "digest": image_data["digest"],
                        "version": image_data["version"],
                        "grant_bound": bool(image_data["grant_ref"]),
                    },
                    "receipt": {
                        "admission_status": receipt["admission_status"],
                        "initial_execution_status": receipt[
                            "execution_status"
                        ],
                        "final_execution_status": final_receipt[
                            "execution_status"
                        ],
                        "model_request_id": final_receipt["model_request_id"],
                        "model_epoch": final_receipt["model_epoch"],
                        "intent_revision": final_receipt["intent_revision"],
                    },
                }
                print(
                    "STEERING_L2_TRACE="
                    + json.dumps(trace, ensure_ascii=False, sort_keys=True)
                )
        finally:
            server.should_exit = True
            await asyncio.wait_for(server_task, timeout=5)
            await tracker.unregister_external_task(chat.id)


@pytest.mark.asyncio
async def test_local_http_image_text_and_multi_image_order_reach_model(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Preserve mixed-content and multi-image order through the real path."""

    session_id = "l2-multimodal-session"
    run_id = "l2-multimodal-run"
    mixed_command_id = "l2-image-text-command"
    multi_command_id = "l2-three-images-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前观察任务；不使用外部工具。",
        name="L2 multimodal order proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            upload_payloads: list[dict] = []
            for label in ("caption", "A", "B", "C"):
                uploaded = await http.post(
                    "/console/upload",
                    files={
                        "file": (
                            "guidance.png",
                            _png_with_label(label),
                            "image/png",
                        )
                    },
                )
                assert uploaded.status_code == 200
                upload_payloads.append(uploaded.json())
            assert [upload["file_name"] for upload in upload_payloads] == [
                "guidance.png",
                "guidance.png",
                "guidance.png",
                "guidance.png",
            ]
            assert (
                len({upload["stored_name"] for upload in upload_payloads}) == 4
            )

            def image_block(upload: dict) -> dict:
                return {
                    "type": "image",
                    "image_url": upload["url"],
                    "filename": upload["file_name"],
                    "media_type": "image/png",
                    "file_size": upload["size"],
                    "stored_name": upload["stored_name"],
                    "attachment_digest": upload["digest"],
                    "attachment_version": upload["version"],
                    "attachment_upload_receipt": upload["upload_receipt"],
                }

            mixed_body = {
                "input": [
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": "先看说明，再按这张图补充当前任务。",
                            },
                            image_block(upload_payloads[0]),
                        ],
                    }
                ],
                "chat_id": harness.chat.id,
                "session_id": session_id,
                "user_id": "default",
                "channel": "console",
                "stream": True,
                "operation": "steer",
                "meta": {
                    "client_turn_id": mixed_command_id,
                    "control_intent": "supplement",
                },
            }
            mixed_admission = await http.post(
                "/console/chat",
                json=mixed_body,
            )
            assert mixed_admission.status_code == 200
            mixed_receipt = _receipt_from_response(mixed_admission)
            assert mixed_receipt["execution_status"] == "pending"

            multi_body = {
                "input": [
                    {
                        "role": "user",
                        "content": [
                            image_block(upload_payloads[1]),
                            {
                                "type": "text",
                                "text": "按 A、说明、B、C 的原始顺序处理。",
                            },
                            image_block(upload_payloads[2]),
                            image_block(upload_payloads[3]),
                        ],
                    }
                ],
                "chat_id": harness.chat.id,
                "session_id": session_id,
                "user_id": "default",
                "channel": "console",
                "stream": True,
                "operation": "steer",
                "meta": {
                    "client_turn_id": multi_command_id,
                    "control_intent": "supplement",
                },
            }
            multi_admission = await http.post(
                "/console/chat",
                json=multi_body,
            )
            assert multi_admission.status_code == 200
            multi_receipt = _receipt_from_response(multi_admission)
            assert multi_receipt["execution_status"] == "pending"
            assert multi_receipt["server_seq"] > mixed_receipt["server_seq"]

            # Reopen the durable store before the runtime consumes either
            # command. Same-name assets must remain distinct and user order
            # must survive the persistence boundary.
            reopened_before_runtime = RuntimeTurnControlStore(
                harness.state_workspace
            )
            persisted_before_runtime = reopened_before_runtime.read_inbox(
                run_id
            )
            assert [item.command_id for item in persisted_before_runtime] == [
                mixed_command_id,
                multi_command_id,
            ]
            assert [
                block.type
                for block in persisted_before_runtime[1].message.content
            ] == ["image", "text", "image", "image"]
            persisted_images = [
                block.data
                for block in persisted_before_runtime[1].message.content
                if block.type == "image"
            ]
            assert [image["filename"] for image in persisted_images] == [
                "guidance.png",
                "guidance.png",
                "guidance.png",
            ]
            assert len({image["digest"] for image in persisted_images}) == 3

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "继续当前观察任务；不使用外部工具。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)

            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 2
            mixed_message = next(
                message
                for message in client.requests[0].messages
                if message.metadata.get("command_id") == mixed_command_id
            )
            assert [block.type for block in mixed_message.content] == [
                "text",
                "image",
            ]
            assert mixed_message.content[0].text == ("先看说明，再按这张图补充当前任务。")
            mixed_image = mixed_message.content[1].data
            assert mixed_image["digest"].startswith("sha256:")
            assert mixed_image["version"] == 1
            assert mixed_image["grant_ref"]

            multi_message = next(
                message
                for message in client.requests[1].messages
                if message.metadata.get("command_id") == multi_command_id
            )
            assert [block.type for block in multi_message.content] == [
                "image",
                "text",
                "image",
                "image",
            ]
            assert multi_message.content[1].text == ("按 A、说明、B、C 的原始顺序处理。")
            multi_images = [
                block.data
                for block in multi_message.content
                if block.type == "image"
            ]
            multi_digests = [image["digest"] for image in multi_images]
            assert len(set(multi_digests)) == 3
            assert [image["filename"] for image in multi_images] == [
                "guidance.png",
                "guidance.png",
                "guidance.png",
            ]
            assert [image["version"] for image in multi_images] == [1, 1, 1]
            assert all(image["grant_ref"] for image in multi_images)

            final_receipts = []
            for command_id in (mixed_command_id, multi_command_id):
                status = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert status.status_code == 200
                assert status.json()["execution_status"] == "applied"
                final_receipts.append(status.json())

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == [
                mixed_command_id,
                multi_command_id,
            ]
            assert [contract.run_id for contract in contracts] == [run_id]
            assert not any(event.type.startswith("tool.") for event in events)

            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_ids": [mixed_command_id, multi_command_id],
                "server_seqs": [
                    mixed_receipt["server_seq"],
                    multi_receipt["server_seq"],
                ],
                "model_request_ids": [
                    receipt["model_request_id"] for receipt in final_receipts
                ],
                "mixed_content_order": [
                    block.type for block in mixed_message.content
                ],
                "mixed_image": {
                    "digest": mixed_image["digest"],
                    "version": mixed_image["version"],
                    "grant_bound": bool(mixed_image["grant_ref"]),
                },
                "multi_image_digests": multi_digests,
                "multi_image_versions": [
                    image["version"] for image in multi_images
                ],
                "multi_image_grants_bound": [
                    bool(image["grant_ref"]) for image in multi_images
                ],
                "counts": {
                    "uploads": len(upload_payloads),
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_MULTIMODAL_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_pdf_append_reaches_model_with_source_and_range(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A PDF supplement keeps the requested page range and readable bytes."""

    session_id = "l2-document-session"
    run_id = "l2-document-run"
    command_id = "l2-document-command"
    instruction = "仅使用附件第2页（PAGE_2）范围的说明继续当前任务。"
    pdf_path = tmp_path / "page-range-proof.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    pdf.drawString(72, 720, "PAGE_1_OUT_OF_SCOPE_SENTINEL")
    pdf.showPage()
    pdf.drawString(72, 720, "PAGE_2_SELECTED_SENTINEL")
    pdf.save()

    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前文档任务；不使用外部工具。",
        name="L2 document append proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "page-range-proof.pdf",
                        pdf_path.read_bytes(),
                        "application/pdf",
                    )
                },
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()

            body = {
                "input": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": instruction},
                            {
                                "type": "file",
                                "file_url": upload["url"],
                                "filename": upload["file_name"],
                                "media_type": "application/pdf",
                                "file_size": upload["size"],
                                "stored_name": upload["stored_name"],
                                "attachment_digest": upload["digest"],
                                "attachment_version": upload["version"],
                                "attachment_upload_receipt": upload[
                                    "upload_receipt"
                                ],
                            },
                        ],
                    }
                ],
                "chat_id": harness.chat.id,
                "session_id": session_id,
                "user_id": "default",
                "channel": "console",
                "stream": True,
                "operation": "steer",
                "meta": {
                    "client_turn_id": command_id,
                    "control_intent": "supplement",
                },
            }
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            receipt = _receipt_from_response(admitted)
            assert receipt["execution_status"] == "pending"

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "继续当前文档任务；不使用外部工具。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)

            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            document_message = next(
                message
                for message in client.requests[0].messages
                if message.metadata.get("command_id") == command_id
            )
            assert [block.type for block in document_message.content] == [
                "text",
                "file",
            ]
            assert document_message.content[0].text == instruction
            file_data = document_message.content[1].data
            assert file_data["digest"].startswith("sha256:")
            assert file_data["version"] == 1
            assert file_data["grant_ref"]
            source_path = file_data["source"]["url"]
            assert Path(source_path).is_file()

            model_context = attachment_context_text(file_data)
            assert "PAGE_1_OUT_OF_SCOPE_SENTINEL" in model_context
            assert "PAGE_2_SELECTED_SENTINEL" in model_context
            assert f"location='{Path(source_path).resolve()}'" in model_context
            assert model_context.count("PAGE_2_SELECTED_SENTINEL") == 1
            assert not any(event.type.startswith("tool.") for event in events)

            status = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert status.status_code == 200
            final_receipt = status.json()
            assert final_receipt["execution_status"] == "applied"

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == [command_id]
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "requested_range": "PAGE_2",
                "content_order": [
                    block.type for block in document_message.content
                ],
                "attachment": {
                    "filename": file_data["filename"],
                    "media_type": file_data["media_type"],
                    "size_bytes": file_data["size_bytes"],
                    "digest": file_data["digest"],
                    "version": file_data["version"],
                    "grant_bound": bool(file_data["grant_ref"]),
                    "source_location_bound": source_path
                    == str(Path(source_path).resolve()),
                    "page_1_text_extracted": (
                        "PAGE_1_OUT_OF_SCOPE_SENTINEL" in model_context
                    ),
                    "page_2_text_extracted": (
                        "PAGE_2_SELECTED_SENTINEL" in model_context
                    ),
                },
                "receipt": {
                    "server_seq": final_receipt["server_seq"],
                    "execution_status": final_receipt["execution_status"],
                    "model_request_id": final_receipt["model_request_id"],
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
                "unverified": [
                    "artifact workspace canonical asset/version/page/selection binding",
                ],
            }
            print(
                "STEERING_DOCUMENT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_mixed_attachment_failure_is_atomic_then_removable(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """One bad attachment rejects the whole command until it is removed."""

    session_id = "l2-mixed-failure-session"
    command_id = "l2-mixed-failure-command"
    run_id = "l2-mixed-failure-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前附件失败验收任务；不使用外部工具。",
        name="L2 mixed attachment failure proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "ready.png",
                        _png_with_label("READY"),
                        "image/png",
                    )
                },
            )
            assert uploaded.status_code == 200
            ready = uploaded.json()
            ready_block = {
                "type": "image",
                "image_url": ready["url"],
                "filename": ready["file_name"],
                "media_type": "image/png",
                "file_size": ready["size"],
                "stored_name": ready["stored_name"],
                "attachment_digest": ready["digest"],
                "attachment_version": ready["version"],
                "attachment_upload_receipt": ready["upload_receipt"],
            }
            missing_block = {
                "type": "image",
                "image_url": str(
                    Path(ready["url"]).with_name("missing-after-upload.png")
                ),
                "filename": "missing-after-upload.png",
                "media_type": "image/png",
                "file_size": 10,
            }
            body = {
                "input": [
                    {
                        "role": "user",
                        "content": [ready_block, missing_block],
                    }
                ],
                "chat_id": harness.chat.id,
                "session_id": session_id,
                "user_id": "default",
                "channel": "console",
                "stream": True,
                "operation": "steer",
                "meta": {
                    "client_turn_id": command_id,
                    "control_intent": "supplement",
                },
            }

            rejected = await http.post("/console/chat", json=body)
            assert rejected.status_code == 400
            assert "control_attachment_not_trusted" in rejected.text
            control_store = RuntimeTurnControlStore(harness.state_workspace)
            assert control_store.read_inbox(run_id) == []
            active_after_rejection = await harness.tracker.get_snapshot(
                harness.chat.id
            )
            assert active_after_rejection is not None
            assert active_after_rejection["status"] == "running"
            assert active_after_rejection["attempt_id"] == (
                harness.started.attempt_id
            )
            contracts_after_rejection = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [
                contract.run_id for contract in contracts_after_rejection
            ] == [run_id]

            body["input"][0]["content"] = [ready_block]
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            receipt = _receipt_from_response(admitted)
            assert receipt["execution_status"] == "pending"

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=harness.binding.effective_workspace_dir,
                    ),
                    messages=[
                        RuntimeMessage.text(
                            "user",
                            "继续当前附件失败验收任务；不使用外部工具。",
                        )
                    ],
                    metadata={
                        "state_workspace_dir": str(harness.state_workspace),
                        "attempt_id": harness.started.attempt_id,
                        "runtime_generation": harness.tracker_generation,
                        "runtime_revision": harness.started.revision,
                        "runtime_lease_owner": harness.started.lease_owner,
                        "task_contract": harness.started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                ),
            )
            assert events[-1].type == "turn.completed"
            assert not any(event.type.startswith("tool.") for event in events)
            assert len(client.requests) == 1
            message = next(
                item
                for item in client.requests[0].messages
                if item.metadata.get("command_id") == command_id
            )
            assert [block.type for block in message.content] == ["image"]
            assert message.content[0].data["filename"] == "ready.png"

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            assert final.json()["execution_status"] == "applied"
            inbox = control_store.read_inbox(run_id)
            assert [item.command_id for item in inbox] == [command_id]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "mixed_submit": {
                    "http_status": rejected.status_code,
                    "error_code": "control_attachment_not_trusted",
                    "durable_inputs_after_rejection": 0,
                    "primary_run_status_after_rejection": (
                        active_after_rejection["status"]
                    ),
                    "primary_attempt_unchanged": (
                        active_after_rejection["attempt_id"]
                        == harness.started.attempt_id
                    ),
                    "runtime_runs_after_rejection": len(
                        contracts_after_rejection
                    ),
                },
                "after_explicit_removal": {
                    "http_status": admitted.status_code,
                    "execution_status": final.json()["execution_status"],
                    "content_types": [block.type for block in message.content],
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(
                        RuntimeRunStore(
                            harness.state_workspace
                        ).list_contracts()
                    ),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_MIXED_FAILURE_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_auth_rejections_do_not_admit_or_retry_control_input(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """401/403 stop before durable admission and leave the active run intact."""

    import lantai.config as config_module
    from lantai.app import auth as auth_module
    from lantai.app._app import app as full_app

    session_id = "l2-auth-rejection-session"
    command_id = "l2-auth-rejection-command"
    run_id = "l2-auth-rejection-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前鉴权拒绝验收任务；不使用外部工具。",
        name="L2 auth rejection proof",
    )
    body = {
        "input": [
            {
                "role": "user",
                "content": [
                    {
                        "type": "input_text",
                        "text": "这条引导必须在重新授权后由用户显式重发。",
                    }
                ],
            }
        ],
        "chat_id": harness.chat.id,
        "session_id": session_id,
        "user_id": "default",
        "channel": "console",
        "stream": True,
        "operation": "steer",
        "meta": {
            "client_turn_id": command_id,
            "control_intent": "redirect",
        },
    }
    try:
        monkeypatch.delenv("LANTAI_DESKTOP_TOKEN", raising=False)
        monkeypatch.setattr(auth_module, "is_auth_enabled", lambda: True)
        monkeypatch.setattr(auth_module, "has_registered_users", lambda: True)
        monkeypatch.setattr(
            config_module,
            "load_config",
            lambda: SimpleNamespace(
                security=SimpleNamespace(allow_no_auth_hosts=[])
            ),
        )

        async with _loopback_console_http(full_app) as (http, _port):
            unauthorized = await http.post("/api/console/chat", json=body)
            assert unauthorized.status_code == 401
            assert unauthorized.json() == {"detail": "Not authenticated"}

            monkeypatch.setattr(auth_module, "is_auth_enabled", lambda: False)
            monkeypatch.setenv("LANTAI_DESKTOP_TOKEN", "l2-desktop-token")
            forbidden = await http.post("/api/console/chat", json=body)
            assert forbidden.status_code == 403
            assert forbidden.json() == {
                "detail": "Browser pairing required",
                "code": "desktop_token_required",
            }

        control_store = RuntimeTurnControlStore(harness.state_workspace)
        inbox = control_store.read_inbox(run_id)
        active_after_rejections = await harness.tracker.get_snapshot(
            harness.chat.id
        )
        assert inbox == []
        assert active_after_rejections is not None
        assert active_after_rejections["status"] == "running"
        assert active_after_rejections["attempt_id"] == (
            harness.started.attempt_id
        )
        contracts = RuntimeRunStore(harness.state_workspace).list_contracts()
        assert [contract.run_id for contract in contracts] == [run_id]

        trace = {
            "schema_version": 1,
            "level": "L2",
            "transport": "loopback_http_socket",
            "chat_id": harness.chat.id,
            "run_id": run_id,
            "attempt_id": harness.started.attempt_id,
            "command_id": command_id,
            "http_401": {
                "status": unauthorized.status_code,
                "detail": unauthorized.json()["detail"],
            },
            "http_403": {
                "status": forbidden.status_code,
                "code": forbidden.json()["code"],
                "detail": forbidden.json()["detail"],
            },
            "after_rejections": {
                "durable_inputs": len(inbox),
                "model_requests": 0,
                "primary_run_status": active_after_rejections["status"],
                "primary_attempt_unchanged": (
                    active_after_rejections["attempt_id"]
                    == harness.started.attempt_id
                ),
                "runtime_runs": len(contracts),
                "automatic_retries": 0,
            },
        }
        print(
            "STEERING_AUTH_REJECTION_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_persistence_failure_returns_no_receipt_and_recovers(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """An atomic command write failure is not admission; retry stays singular."""

    from lantai.runtime import turn_control as turn_control_module

    session_id = "l2-persistence-failure-session"
    command_id = "l2-persistence-failure-command"
    run_id = "l2-persistence-failure-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前持久化故障验收任务；不使用外部工具。",
        name="L2 persistence failure proof",
    )
    body = _text_steer_body(
        chat_id=harness.chat.id,
        session_id=session_id,
        command_id=command_id,
        text="存储恢复后仍只应用这一条引导。",
    )
    original_atomic_write = turn_control_module.write_text_atomic_no_follow
    failed_paths: list[str] = []

    def fail_control_command_write(path: Path, text: str) -> None:
        normalized = Path(path).as_posix()
        if "/runtime/v2/commands/" in normalized:
            failed_paths.append(normalized)
            raise OSError("injected atomic control-input persistence failure")
        original_atomic_write(path, text)

    try:
        async with _loopback_console_http() as (http, _port):
            monkeypatch.setattr(
                turn_control_module,
                "write_text_atomic_no_follow",
                fail_control_command_write,
            )
            failed = await http.post("/console/chat", json=body)
            assert failed.status_code == 500
            assert failed_paths

            control_store = RuntimeTurnControlStore(harness.state_workspace)
            assert control_store.read_inbox(run_id) == []
            assert (
                control_store.get_input_receipt(
                    chat_id=harness.chat.id,
                    command_id=command_id,
                )
                is None
            )
            active_after_failure = await harness.tracker.get_snapshot(
                harness.chat.id
            )
            assert active_after_failure is not None
            assert active_after_failure["status"] == "running"
            assert active_after_failure["attempt_id"] == (
                harness.started.attempt_id
            )

            monkeypatch.setattr(
                turn_control_module,
                "write_text_atomic_no_follow",
                original_atomic_write,
            )
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            initial_receipt = _receipt_from_response(admitted)
            assert initial_receipt["command_id"] == command_id
            assert initial_receipt["execution_status"] == "pending"

            retried = await http.post("/console/chat", json=body)
            assert retried.status_code == 200
            assert _receipt_from_response(retried) == initial_receipt

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "继续当前持久化故障验收任务；不使用外部工具。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": (
                        harness.started.task_contract.model_dump(mode="json")
                    ),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)
            assert events[-1].type == "turn.completed"

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            inbox = control_store.read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert len(inbox) == 1
            assert len(client.requests) == 1
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "injected_failure": {
                    "http_status": failed.status_code,
                    "success_receipt_emitted": False,
                    "durable_inputs": 0,
                    "input_receipt_found": False,
                    "primary_run_status": active_after_failure["status"],
                    "primary_attempt_unchanged": True,
                },
                "after_storage_recovery": {
                    "first_retry_status": admitted.status_code,
                    "same_command_retry_status": retried.status_code,
                    "same_receipt": (
                        _receipt_from_response(retried) == initial_receipt
                    ),
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "automatic_new_command_ids": 0,
                },
            }
            print(
                "STEERING_PERSISTENCE_FAILURE_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_lost_upload_receipt_retries_one_stable_blob(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A dropped upload receipt retries to one blob and one command."""

    session_id = "l2-lost-upload-receipt-session"
    command_id = "l2-lost-upload-receipt-command"
    run_id = "l2-lost-upload-receipt-run"
    filename = "lost-receipt.png"
    payload = _png_with_label("LOST_UPLOAD_RECEIPT")
    instruction = "保留这段草稿，并把重试后的图片补充给当前任务。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前上传重试验收任务；不使用外部工具。",
        name="L2 lost upload receipt proof",
    )
    try:
        async with _loopback_console_http() as (http, port):
            dropped_status = await _upload_and_drop_after_headers(
                port=port,
                filename=filename,
                media_type="image/png",
                data=payload,
            )
            assert dropped_status.startswith("HTTP/1.1 200")

            retried = await http.post(
                "/console/upload",
                files={"file": (filename, payload, "image/png")},
            )
            assert retried.status_code == 200
            upload = retried.json()
            assert upload["digest"].startswith("sha256:")
            assert upload["version"] == 1
            stored_files = [
                path
                for path in Path(upload["url"]).parent.iterdir()
                if not path.name.startswith(".")
            ]
            assert stored_files == [Path(upload["url"])]

            body = {
                "input": [
                    {
                        "role": "user",
                        "content": [
                            {"type": "text", "text": instruction},
                            {
                                "type": "image",
                                "image_url": upload["url"],
                                "filename": upload["file_name"],
                                "media_type": "image/png",
                                "file_size": upload["size"],
                                "stored_name": upload["stored_name"],
                                "attachment_digest": upload["digest"],
                                "attachment_version": upload["version"],
                                "attachment_upload_receipt": upload[
                                    "upload_receipt"
                                ],
                            },
                        ],
                    }
                ],
                "chat_id": harness.chat.id,
                "session_id": session_id,
                "user_id": "default",
                "channel": "console",
                "stream": True,
                "operation": "steer",
                "meta": {
                    "client_turn_id": command_id,
                    "control_intent": "supplement",
                },
            }
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            receipt = _receipt_from_response(admitted)
            assert receipt["execution_status"] == "pending"

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=harness.binding.effective_workspace_dir,
                    ),
                    messages=[
                        RuntimeMessage.text(
                            "user",
                            "继续当前上传重试验收任务；不使用外部工具。",
                        )
                    ],
                    metadata={
                        "state_workspace_dir": str(harness.state_workspace),
                        "attempt_id": harness.started.attempt_id,
                        "runtime_generation": harness.tracker_generation,
                        "runtime_revision": harness.started.revision,
                        "runtime_lease_owner": harness.started.lease_owner,
                        "task_contract": harness.started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                ),
            )
            assert events[-1].type == "turn.completed"
            assert not any(event.type.startswith("tool.") for event in events)
            assert len(client.requests) == 1
            message = next(
                item
                for item in client.requests[0].messages
                if item.metadata.get("command_id") == command_id
            )
            assert [block.type for block in message.content] == [
                "text",
                "image",
            ]
            assert message.content[0].text == instruction
            assert message.content[1].data["digest"] == upload["digest"]

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            assert [item.command_id for item in inbox] == [command_id]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket_disconnect",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "upload": {
                    "dropped_response_status_line": dropped_status,
                    "retry_http_status": retried.status_code,
                    "stable_url": upload["url"],
                    "digest": upload["digest"],
                    "version": upload["version"],
                    "stored_blob_count": len(stored_files),
                },
                "draft_reached_model": {
                    "text_preserved": message.content[0].text == instruction,
                    "attachment_preserved": (
                        message.content[1].data["digest"] == upload["digest"]
                    ),
                },
                "receipt": {
                    "server_seq": final_receipt["server_seq"],
                    "execution_status": final_receipt["execution_status"],
                    "model_request_id": final_receipt["model_request_id"],
                },
                "counts": {
                    "stored_blobs": len(stored_files),
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(
                        RuntimeRunStore(
                            harness.state_workspace
                        ).list_contracts()
                    ),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_LOST_UPLOAD_RECEIPT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_unsupported_audio_is_explicit_and_not_admitted(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Unsupported media is rejected without a success receipt or Inbox row."""

    session_id = "l2-unsupported-media-session"
    command_id = "l2-unsupported-media-command"
    run_id = "l2-unsupported-media-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前能力边界验收任务；不使用外部工具。",
        name="L2 unsupported media proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": ("sample.wav", b"RIFFunsupported", "audio/wav")
                },
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()
            response = await http.post(
                "/console/chat",
                json={
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "audio",
                                    "audio_url": upload["url"],
                                    "filename": upload["file_name"],
                                    "media_type": "audio/wav",
                                    "file_size": upload["size"],
                                }
                            ],
                        }
                    ],
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                    },
                },
            )
            assert response.status_code == 415
            detail = response.json()["detail"]
            assert detail == {
                "code": "unsupported_control_media",
                "media_types": ["audio"],
            }
            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            assert inbox == []
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "http_status": response.status_code,
                "error": detail,
                "attachment_upload_succeeded": uploaded.status_code == 200,
                "success_receipt_emitted": False,
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": 0,
                    "runtime_runs": len(
                        RuntimeRunStore(
                            harness.state_workspace
                        ).list_contracts()
                    ),
                    "agent_loops_started_by_handler": 0,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_UNSUPPORTED_MEDIA_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_stop_does_not_promote_failed_steer_or_start_next_task(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Stop leaves an explicit next task pending and never promotes failure."""

    session_id = "l2-stop-queue-session"
    run_id = "l2-stop-queue-run"
    failed_command_id = "l2-stop-failed-steer"
    next_command_id = "l2-stop-independent-next"
    next_text = "当前任务停止后保留，等我明确启动再执行这项独立工作。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前 Stop 与队列竞争验收任务。",
        name="L2 stop and queue proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            queued = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=next_command_id,
                    text=next_text,
                    intent="queue_next",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert queued.status_code == 200
            queued_receipt = _receipt_from_response(queued)
            assert queued_receipt["target"] == {
                "kind": "next_turn",
                "target_id": harness.chat.id,
            }
            assert queued_receipt["execution_status"] == "pending"
            assert (
                queued_receipt["next_waiting_reason"] == "awaiting_next_turn"
            )

            uploaded = await http.post(
                "/console/upload",
                files={"file": ("failed.wav", b"RIFFfailed", "audio/wav")},
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()
            failed = await http.post(
                "/console/chat",
                json={
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "audio",
                                    "audio_url": upload["url"],
                                    "filename": upload["file_name"],
                                    "media_type": "audio/wav",
                                    "file_size": upload["size"],
                                }
                            ],
                        }
                    ],
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": failed_command_id,
                        "control_intent": "supplement",
                        "expected_run_id": run_id,
                        "expected_attempt_id": harness.started.attempt_id,
                    },
                },
            )
            assert failed.status_code == 415
            assert failed.json()["detail"]["code"] == (
                "unsupported_control_media"
            )

            stopped = await http.post(
                "/console/chat/stop",
                params={
                    "chat_id": harness.chat.id,
                    "turn_id": harness.started.attempt_id,
                },
            )
            assert stopped.status_code == 200
            assert stopped.json() == {"stopped": True}
            assert await harness.tracker.get_status(harness.chat.id) == "idle"

            queued_after_stop = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": next_command_id,
                },
            )
            assert queued_after_stop.status_code == 200
            queued_after_receipt = queued_after_stop.json()
            assert queued_after_receipt["execution_status"] == "pending"
            assert queued_after_receipt["next_waiting_reason"] == (
                "awaiting_next_turn"
            )

            failed_after_stop = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": failed_command_id,
                },
            )
            assert failed_after_stop.status_code == 404

            store = RuntimeTurnControlStore(harness.state_workspace)
            inbox = store.read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == [next_command_id]
            assert inbox[0].target is not None
            assert inbox[0].target.kind == "next_turn"
            assert inbox[0].message.text_content() == next_text
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "failed_steer": {
                    "command_id": failed_command_id,
                    "http_status": failed.status_code,
                    "durable_receipt_found_after_stop": False,
                },
                "stop": {
                    "http_status": stopped.status_code,
                    "stopped": stopped.json()["stopped"],
                    "tracker_status_after": await harness.tracker.get_status(
                        harness.chat.id
                    ),
                },
                "independent_next_task": {
                    "command_id": next_command_id,
                    "target": queued_after_receipt["target"],
                    "execution_status": queued_after_receipt[
                        "execution_status"
                    ],
                    "next_waiting_reason": queued_after_receipt[
                        "next_waiting_reason"
                    ],
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "runtime_runs": len(contracts),
                    "agent_loops_started_by_control_handler": 0,
                    "model_requests": 0,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_STOP_QUEUE_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_revoked_attachment_grant_never_reaches_model(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A durable grant revocation blocks the queued attachment after restart."""

    session_id = "l2-revoked-grant-session"
    command_id = "l2-revoked-grant-command"
    run_id = "l2-revoked-grant-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前附件授权验收任务；不使用外部工具。",
        name="L2 revoked attachment grant proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "revoked.png",
                        _png_with_label("REVOKED"),
                        "image/png",
                    )
                },
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()
            admitted = await http.post(
                "/console/chat",
                json={
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "image",
                                    "image_url": upload["url"],
                                    "filename": upload["file_name"],
                                    "media_type": "image/png",
                                    "file_size": upload["size"],
                                    "stored_name": upload["stored_name"],
                                    "attachment_digest": upload["digest"],
                                    "attachment_version": upload["version"],
                                    "attachment_upload_receipt": upload[
                                        "upload_receipt"
                                    ],
                                }
                            ],
                        }
                    ],
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                    },
                },
            )
            assert admitted.status_code == 200
            assert (
                _receipt_from_response(admitted)["execution_status"]
                == "pending"
            )

            store = RuntimeTurnControlStore(harness.state_workspace)
            item = store.read_inbox(run_id)[0]
            grant_ref = str(item.message.content[0].data["grant_ref"])
            revoked = store.revoke_attachment_grant(
                chat_id=harness.chat.id,
                command_id=command_id,
                grant_ref=grant_ref,
                actor_id=f"chat:{harness.chat.id}",
            )
            assert revoked.disposition_status == "rejected"

            restarted_store = RuntimeTurnControlStore(harness.state_workspace)
            assert restarted_store.attachment_grant_is_revoked(
                chat_id=harness.chat.id,
                command_id=command_id,
                grant_ref=grant_ref,
            )
            current_run = await RuntimeRunStore(
                harness.state_workspace
            ).load_contract(run_id)
            assert current_run is not None
            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=harness.binding.effective_workspace_dir,
                    ),
                    messages=[
                        RuntimeMessage.text(
                            "user",
                            "继续当前附件授权验收任务；不使用外部工具。",
                        )
                    ],
                    metadata={
                        "state_workspace_dir": str(harness.state_workspace),
                        "attempt_id": current_run.attempt_id,
                        "runtime_generation": harness.tracker_generation,
                        "runtime_revision": current_run.revision,
                        "runtime_lease_owner": current_run.lease_owner,
                        "task_contract": current_run.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            assert not any(
                message.metadata.get("command_id") == command_id
                for message in client.requests[0].messages
            )
            assert not any(event.type.startswith("tool.") for event in events)
            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "not_applied"
            assert final_receipt["disposition_status"] == "rejected"
            assert "revoked" in final_receipt["disposition_reason"]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket_store_restart",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": current_run.attempt_id,
                "command_id": command_id,
                "receipt": {
                    "initial_execution_status": "pending",
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                    "disposition_status": final_receipt["disposition_status"],
                    "reason": final_receipt["disposition_reason"],
                },
                "grant_revocation_survived_restart": True,
                "attachment_bound_to_model_request": False,
                "counts": {
                    "durable_inputs": len(restarted_store.read_inbox(run_id)),
                    "model_requests": len(client.requests),
                    "command_model_requests": 0,
                    "runtime_runs": len(
                        RuntimeRunStore(
                            harness.state_workspace
                        ).list_contracts()
                    ),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_REVOKED_GRANT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_source_replacement_keeps_admitted_snapshot(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Replacing the upload path cannot change the already admitted bytes."""

    session_id = "l2-source-replacement-session"
    command_id = "l2-source-replacement-command"
    run_id = "l2-source-replacement-run"
    original = _png_with_label("ORIGINAL_SNAPSHOT")
    replacement = _png_with_label("REPLACED_SOURCE")
    original_digest = f"sha256:{hashlib.sha256(original).hexdigest()}"
    replacement_digest = f"sha256:{hashlib.sha256(replacement).hexdigest()}"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续当前不可变快照验收任务；不使用外部工具。",
        name="L2 immutable attachment snapshot proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={"file": ("same-path.png", original, "image/png")},
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()
            admitted = await http.post(
                "/console/chat",
                json={
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "image",
                                    "image_url": upload["url"],
                                    "filename": upload["file_name"],
                                    "media_type": "image/png",
                                    "file_size": upload["size"],
                                    "stored_name": upload["stored_name"],
                                    "attachment_digest": upload["digest"],
                                    "attachment_version": upload["version"],
                                    "attachment_upload_receipt": upload[
                                        "upload_receipt"
                                    ],
                                }
                            ],
                        }
                    ],
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                    },
                },
            )
            assert admitted.status_code == 200
            assert (
                _receipt_from_response(admitted)["execution_status"]
                == "pending"
            )

            store = RuntimeTurnControlStore(harness.state_workspace)
            item = store.read_inbox(run_id)[0]
            snapshot_data = item.message.content[0].data
            snapshot_path = Path(snapshot_data["source"]["url"])
            assert snapshot_path != Path(upload["url"])
            assert snapshot_data["digest"] == original_digest
            assert snapshot_path.read_bytes() == original

            Path(upload["url"]).write_bytes(replacement)
            replaced_upload_digest = (
                "sha256:"
                + hashlib.sha256(Path(upload["url"]).read_bytes()).hexdigest()
            )
            assert replaced_upload_digest == replacement_digest
            assert snapshot_path.read_bytes() == original

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=harness.binding.effective_workspace_dir,
                    ),
                    messages=[
                        RuntimeMessage.text(
                            "user",
                            "继续当前不可变快照验收任务；不使用外部工具。",
                        )
                    ],
                    metadata={
                        "state_workspace_dir": str(harness.state_workspace),
                        "attempt_id": harness.started.attempt_id,
                        "runtime_generation": harness.tracker_generation,
                        "runtime_revision": harness.started.revision,
                        "runtime_lease_owner": harness.started.lease_owner,
                        "task_contract": harness.started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            message = next(
                message
                for message in client.requests[0].messages
                if message.metadata.get("command_id") == command_id
            )
            model_data = message.content[0].data
            assert model_data["digest"] == original_digest
            assert Path(model_data["source"]["url"]).read_bytes() == original
            assert model_data["digest"] != replacement_digest
            assert not any(event.type.startswith("tool.") for event in events)
            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            assert final.json()["execution_status"] == "applied"
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "source_replaced_after_admission": True,
                "snapshot": {
                    "distinct_from_upload_path": snapshot_path
                    != Path(upload["url"]),
                    "original_digest": original_digest,
                    "replacement_digest": replacement_digest,
                    "model_received_digest": model_data["digest"],
                    "model_received_original_bytes": Path(
                        model_data["source"]["url"]
                    ).read_bytes()
                    == original,
                },
                "receipt": {
                    "execution_status": final.json()["execution_status"],
                    "model_request_id": final.json()["model_request_id"],
                },
                "counts": {
                    "durable_inputs": len(store.read_inbox(run_id)),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(
                        RuntimeRunStore(
                            harness.state_workspace
                        ).list_contracts()
                    ),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_SOURCE_REPLACEMENT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_lost_receipt_queries_and_reuses_one_text_command(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A disconnected receipt is recovered without creating another input."""

    session_id = "l2-lost-receipt-session"
    command_id = "l2-lost-receipt-command"
    text = "只补充这一条文字，不要新建任务。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-lost-receipt-run",
        objective="继续当前文字任务。",
        name="L2 lost receipt proof",
    )
    try:
        async with _loopback_console_http() as (http, port):
            body = _text_steer_body(
                chat_id=harness.chat.id,
                session_id=session_id,
                command_id=command_id,
                text=text,
            )
            status_line = await _post_and_drop_after_headers(
                port=port,
                path="/console/chat",
                body=body,
            )
            assert status_line.startswith("HTTP/1.1 200")

            queried = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert queried.status_code == 200
            queried_receipt = queried.json()
            assert queried_receipt["execution_status"] == "pending"

            retried = await http.post("/console/chat", json=body)
            assert retried.status_code == 200
            retry_receipts = [
                payload["receipt"]
                for payload in _sse_payloads(retried.text)
                if isinstance(payload.get("receipt"), dict)
            ]
            assert retry_receipts == [queried_receipt]

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=harness.run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[RuntimeMessage.text("user", "继续当前文字任务。")],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = [event async for event in runtime.run_turn(runtime_input)]
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            assert client.requests[0].messages[-1].text_content() == text
            assert not any(event.type.startswith("tool.") for event in events)

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["server_seq"] == queried_receipt["server_seq"]
            assert final_receipt["payload_digest"] == (
                queried_receipt["payload_digest"]
            )

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(harness.run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == [command_id]
            assert [contract.run_id for contract in contracts] == [
                harness.run_id
            ]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket_disconnect",
                "command_id": command_id,
                "chat_id": harness.chat.id,
                "run_id": harness.run_id,
                "attempt_id": harness.started.attempt_id,
                "http": {
                    "dropped_response_status_line": status_line,
                    "query_after_disconnect": queried.status_code,
                    "same_command_retry": retried.status_code,
                    "final_query": final.status_code,
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
                "identity_reused": {
                    "server_seq": final_receipt["server_seq"],
                    "payload_digest": final_receipt["payload_digest"],
                    "model_request_id": final_receipt["model_request_id"],
                },
                "receipt_discarded_before_parse": True,
                "final_execution_status": final_receipt["execution_status"],
            }
            print(
                "STEERING_LOST_RECEIPT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_first_run_guidance_waits_for_identity_then_binds_once(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A pre-identity draft is sent once after its Run/Attempt become known."""

    session_id = "l2-first-run-binding-session"
    command_id = "l2-first-run-binding-command"
    text = "首轮身份就绪后只补充这一条。"
    local_draft = {
        "command_id": command_id,
        "text": text,
        "control_intent": "supplement",
        "waiting_reason": "awaiting_run_identity",
    }
    assert local_draft["waiting_reason"] == "awaiting_run_identity"

    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-first-run-binding-run",
        objective="启动首轮并等待身份绑定。",
        name="L2 first-run identity binding proof",
    )
    try:
        store = RuntimeTurnControlStore(harness.state_workspace)
        assert store.read_inbox(harness.run_id) == []

        async with _loopback_console_http() as (http, _port):
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text=text,
                    expected_run_id=harness.run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert admitted.status_code == 200
            admission_receipt = _receipt_from_response(admitted)
            assert admission_receipt["run_id"] == harness.run_id
            assert (
                admission_receipt["attempt_id"] == harness.started.attempt_id
            )
            assert admission_receipt["execution_status"] == "pending"

            inbox_after_binding = store.read_inbox(harness.run_id)
            assert [item.command_id for item in inbox_after_binding] == [
                command_id
            ]

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=harness.run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "启动首轮并等待身份绑定。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            assert client.requests[0].messages[-1].text_content() == text

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == (
                f"{harness.run_id}:model:1"
            )
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [contract.run_id for contract in contracts] == [
                harness.run_id
            ]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "local_wait_then_loopback_http_socket",
                "command_id": command_id,
                "chat_id": harness.chat.id,
                "run_id": harness.run_id,
                "attempt_id": harness.started.attempt_id,
                "before_identity": {
                    "waiting_reason": local_draft["waiting_reason"],
                    "http_requests": 0,
                    "durable_inputs": 0,
                },
                "after_identity": {
                    "http_status": admitted.status_code,
                    "durable_inputs": len(inbox_after_binding),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                    "execution_status": final_receipt["execution_status"],
                    "model_request_id": final_receipt["model_request_id"],
                },
            }
            print(
                "STEERING_FIRST_RUN_BINDING_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_heartbeat_revision_does_not_reject_valid_control_input(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Heartbeat store revisions are not part of the stable input fence."""

    session_id = "l2-heartbeat-revision-session"
    command_id = "l2-heartbeat-revision-command"
    text = "心跳之后仍应把这条补充交给同一运行任务。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-heartbeat-revision-run",
        objective="保持当前任务运行并接受后续补充。",
        name="L2 heartbeat revision fence proof",
    )
    try:
        heartbeat = await touch_chat_runtime_run_heartbeat(
            workspace_dir=harness.state_workspace,
            run_id=harness.run_id,
        )
        assert heartbeat is not None
        assert heartbeat.revision > harness.started.revision
        assert heartbeat.attempt_id == harness.started.attempt_id

        async with _loopback_console_http() as (http, _port):
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text=text,
                    expected_run_id=harness.run_id,
                    expected_attempt_id=heartbeat.attempt_id,
                ),
            )
            assert admitted.status_code == 200
            receipt = _receipt_from_response(admitted)
            assert receipt["run_id"] == harness.run_id
            assert receipt["attempt_id"] == heartbeat.attempt_id
            assert receipt["execution_status"] == "pending"

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(harness.run_id)
            assert [item.command_id for item in inbox] == [command_id]
            assert inbox[0].fence is not None
            assert inbox[0].fence.expected_revision is None

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=harness.run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "保持当前任务运行并接受后续补充。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": heartbeat.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": heartbeat.revision,
                    "runtime_lease_owner": heartbeat.lease_owner,
                    "task_contract": heartbeat.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            assert client.requests[0].messages[-1].text_content() == text

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == (
                f"{harness.run_id}:model:1"
            )

            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [contract.run_id for contract in contracts] == [
                harness.run_id
            ]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "command_id": command_id,
                "chat_id": harness.chat.id,
                "run_id": harness.run_id,
                "attempt_id": heartbeat.attempt_id,
                "revision": {
                    "before_heartbeat": harness.started.revision,
                    "after_heartbeat": heartbeat.revision,
                    "input_fence_expected_revision": (
                        inbox[0].fence.expected_revision
                    ),
                },
                "http_status": {
                    "admission": admitted.status_code,
                    "query": final.status_code,
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
                "final_execution_status": final_receipt["execution_status"],
                "model_request_id": final_receipt["model_request_id"],
            }
            print(
                "STEERING_HEARTBEAT_REVISION_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_rejects_control_input_captured_for_an_old_attempt(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A queued draft must never be silently rebound to the current owner."""

    session_id = "l2-stale-attempt-session"
    command_id = "l2-stale-attempt-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-current-run",
        objective="继续当前归属验收任务。",
        name="L2 stale attempt ownership proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            response = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text="这条输入只属于刚才已经结束的 Attempt。",
                    expected_run_id=harness.run_id,
                    expected_attempt_id="l2-current-run:chat-attempt-0",
                ),
            )
            assert response.status_code == 409
            detail = response.json()["detail"]
            assert detail == {
                "code": "stale_attempt",
                "command_id": command_id,
                "expected_run_id": harness.run_id,
                "current_run_id": harness.run_id,
                "expected_attempt_id": "l2-current-run:chat-attempt-0",
                "current_attempt_id": harness.started.attempt_id,
                "message": (
                    "The control input belongs to an older run attempt. "
                    "Review it before sending again."
                ),
            }
            assert (
                RuntimeTurnControlStore(harness.state_workspace).read_inbox(
                    harness.run_id
                )
                == []
            )
            print(
                "STEERING_STALE_ATTEMPT_L2_TRACE="
                + json.dumps(
                    {
                        "schema_version": 1,
                        "level": "L2",
                        "chat_id": harness.chat.id,
                        "run_id": harness.run_id,
                        "current_attempt_id": harness.started.attempt_id,
                        "expected_attempt_id": (
                            "l2-current-run:chat-attempt-0"
                        ),
                        "command_id": command_id,
                        "http_status": response.status_code,
                        "error_code": detail["code"],
                        "durable_inputs": 0,
                        "redirected_to_current_attempt": False,
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                )
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_shared_transport_alias_keeps_chat_inputs_attachments_and_receipts_isolated(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """The legacy transport alias is not an Inbox or runtime owner key."""

    session_id = "console:l2-shared-transport-alias"
    harness_a = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-shared-alias-run-a",
        objective="只处理 Chat A 的输入与附件。",
        name="L2 shared alias A",
    )
    chat_b = await harness_a.chat_manager.create_chat(
        ChatSpec(
            id="l2-shared-alias-chat-b",
            name="L2 shared alias B",
            session_id=session_id,
            user_id="default",
            channel="console",
        )
    )
    binding_b = await ensure_chat_topic_workspace(
        harness_a.workspace,
        harness_a.chat_manager,
        chat_b,
        agent_id="default",
    )
    chat_b = binding_b.chat
    state_workspace_b = Path(binding_b.state_workspace_dir)
    run_id_b = "l2-shared-alias-run-b"
    contract_b = build_chat_task_contract(
        objective="只处理 Chat B 的输入与附件。",
        client_turn_id=f"{run_id_b}-initial-turn",
        complexity="light",
        work_method_id="standard",
        run_id=run_id_b,
        attempt_id=f"{run_id_b}:chat-attempt-1",
    )
    started_b = await record_chat_runtime_run_started(
        workspace_dir=state_workspace_b,
        agent_id="default",
        chat_id=chat_b.id,
        session_id=session_id,
        run_id=run_id_b,
        effective_workspace_dir=binding_b.effective_workspace_dir,
        task_contract=contract_b,
    )
    await harness_a.chat_manager.merge_chat_meta(
        chat_b.id,
        {
            "last_turn_run_id": run_id_b,
            "state_workspace_dir": str(state_workspace_b),
        },
    )
    generation_b = await _register_active_external_task(
        harness_a.tracker, chat_b.id,
        attempt_id=started_b.attempt_id,
        root_session_id=session_id,
    )

    command_a = "l2-shared-alias-command-a"
    command_b = "l2-shared-alias-command-b"
    try:
        async with _loopback_console_http() as (http, _port):
            upload_a, upload_b = await asyncio.gather(
                http.post(
                    "/console/upload",
                    files={
                        "file": (
                            "chat-a.png",
                            _png_with_label("chat-a"),
                            "image/png",
                        )
                    },
                ),
                http.post(
                    "/console/upload",
                    files={
                        "file": (
                            "chat-b.png",
                            _png_with_label("chat-b"),
                            "image/png",
                        )
                    },
                ),
            )
            assert upload_a.status_code == upload_b.status_code == 200

            def mixed_body(
                *,
                chat_id: str,
                command_id: str,
                text: str,
                upload: dict,
                run_id: str,
                attempt_id: str,
            ) -> dict:
                return {
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": text},
                                {
                                    "type": "image",
                                    "image_url": upload["url"],
                                    "filename": upload["file_name"],
                                    "media_type": "image/png",
                                    "file_size": upload["size"],
                                    "stored_name": upload["stored_name"],
                                    "attachment_digest": upload["digest"],
                                    "attachment_version": upload["version"],
                                    "attachment_upload_receipt": upload[
                                        "upload_receipt"
                                    ],
                                },
                            ],
                        }
                    ],
                    "chat_id": chat_id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                        "expected_run_id": run_id,
                        "expected_attempt_id": attempt_id,
                    },
                }

            text_a = "CHAT_A_ONLY：只使用 A 图。"
            text_b = "CHAT_B_ONLY：只使用 B 图。"
            admitted_a, admitted_b = await asyncio.gather(
                http.post(
                    "/console/chat",
                    json=mixed_body(
                        chat_id=harness_a.chat.id,
                        command_id=command_a,
                        text=text_a,
                        upload=upload_a.json(),
                        run_id=harness_a.run_id,
                        attempt_id=harness_a.started.attempt_id,
                    ),
                ),
                http.post(
                    "/console/chat",
                    json=mixed_body(
                        chat_id=chat_b.id,
                        command_id=command_b,
                        text=text_b,
                        upload=upload_b.json(),
                        run_id=run_id_b,
                        attempt_id=started_b.attempt_id,
                    ),
                ),
            )
            assert admitted_a.status_code == admitted_b.status_code == 200
            receipt_a = _receipt_from_response(admitted_a)
            receipt_b = _receipt_from_response(admitted_b)
            assert receipt_a["run_id"] == harness_a.run_id
            assert receipt_b["run_id"] == run_id_b

            crossed_a = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness_a.chat.id,
                    "command_id": command_b,
                },
            )
            crossed_b = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": chat_b.id,
                    "command_id": command_a,
                },
            )
            assert crossed_a.status_code == crossed_b.status_code == 404

            client_a = _RecordingModelClient()
            client_b = _RecordingModelClient()
            runtime_a = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client_a,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_b = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client_b,
                tool_registry=LantaiToolRegistry(),
            )

            def runtime_input(
                *,
                run_id: str,
                binding,
                state_workspace: Path,
                started,
                generation: int,
                objective: str,
            ) -> RuntimeInput:
                return RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=binding.effective_workspace_dir,
                    ),
                    messages=[RuntimeMessage.text("user", objective)],
                    metadata={
                        "state_workspace_dir": str(state_workspace),
                        "attempt_id": started.attempt_id,
                        "runtime_generation": generation,
                        "runtime_revision": started.revision,
                        "runtime_lease_owner": started.lease_owner,
                        "task_contract": started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                )

            events_a, events_b = await asyncio.gather(
                _collect_runtime_events(
                    runtime_a,
                    runtime_input(
                        run_id=harness_a.run_id,
                        binding=harness_a.binding,
                        state_workspace=harness_a.state_workspace,
                        started=harness_a.started,
                        generation=harness_a.tracker_generation,
                        objective="只处理 Chat A 的输入与附件。",
                    ),
                ),
                _collect_runtime_events(
                    runtime_b,
                    runtime_input(
                        run_id=run_id_b,
                        binding=binding_b,
                        state_workspace=state_workspace_b,
                        started=started_b,
                        generation=generation_b,
                        objective="只处理 Chat B 的输入与附件。",
                    ),
                ),
            )
            assert events_a[-1].type == events_b[-1].type == "turn.completed"
            assert len(client_a.requests) == len(client_b.requests) == 1

            message_a = next(
                message
                for message in client_a.requests[0].messages
                if message.metadata.get("command_id") == command_a
            )
            message_b = next(
                message
                for message in client_b.requests[0].messages
                if message.metadata.get("command_id") == command_b
            )
            assert message_a.text_content() == text_a
            assert message_b.text_content() == text_b
            assert text_b not in json.dumps(
                client_a.requests[0].model_dump(mode="json"),
                ensure_ascii=False,
            )
            assert text_a not in json.dumps(
                client_b.requests[0].model_dump(mode="json"),
                ensure_ascii=False,
            )
            image_a = next(
                block for block in message_a.content if block.type == "image"
            )
            image_b = next(
                block for block in message_b.content if block.type == "image"
            )
            assert image_a.data["digest"] != image_b.data["digest"]

            inbox_a = RuntimeTurnControlStore(
                harness_a.state_workspace
            ).read_inbox(harness_a.run_id)
            inbox_b = RuntimeTurnControlStore(state_workspace_b).read_inbox(
                run_id_b
            )
            assert [item.command_id for item in inbox_a] == [command_a]
            assert [item.command_id for item in inbox_b] == [command_b]

            assert harness_a.state_workspace != state_workspace_b
            migrated_a_workspace = tmp_path / "migrated-chat-a-state"
            harness_a.state_workspace.rename(migrated_a_workspace)

            restarted_b_store = RuntimeTurnControlStore(state_workspace_b)
            restarted_b_inbox = restarted_b_store.read_inbox(run_id_b)
            assert [item.command_id for item in restarted_b_inbox] == [
                command_b
            ]
            assert (
                restarted_b_store.get_input_item(
                    chat_id=chat_b.id,
                    command_id=command_a,
                )
                is None
            )
            assert (
                restarted_b_store.get_input_item(
                    chat_id=chat_b.id,
                    command_id=command_b,
                )
                is not None
            )

            recovered_a_inbox = RuntimeTurnControlStore(
                migrated_a_workspace
            ).read_inbox(harness_a.run_id)
            assert [item.command_id for item in recovered_a_inbox] == [
                command_a
            ]

            print(
                "STEERING_SHARED_ALIAS_L2_TRACE="
                + json.dumps(
                    {
                        "schema_version": 1,
                        "level": "L2",
                        "shared_transport_session_id": session_id,
                        "chat_ids": [harness_a.chat.id, chat_b.id],
                        "run_ids": [harness_a.run_id, run_id_b],
                        "attempt_ids": [
                            harness_a.started.attempt_id,
                            started_b.attempt_id,
                        ],
                        "command_ids": [command_a, command_b],
                        "cross_receipt_statuses": [
                            crossed_a.status_code,
                            crossed_b.status_code,
                        ],
                        "inbox_counts": [len(inbox_a), len(inbox_b)],
                        "chat_a_workspace_migrated": True,
                        "chat_a_recoverable_after_migration": True,
                        "chat_b_reopened_without_shared_key_fallback": True,
                        "model_request_counts": [
                            len(client_a.requests),
                            len(client_b.requests),
                        ],
                        "attachment_digests_distinct": (
                            image_a.data["digest"] != image_b.data["digest"]
                        ),
                    },
                    ensure_ascii=False,
                    sort_keys=True,
                )
            )
    finally:
        await harness_a.tracker.unregister_external_task(harness_a.chat.id)
        await harness_a.tracker.unregister_external_task(chat_b.id)


@pytest.mark.asyncio
async def test_local_http_concurrent_retry_and_same_text_remain_distinct(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Concurrent retries dedupe by id; equal text with new ids does not."""

    session_id = "l2-idempotency-session"
    concurrent_id = "l2-concurrent-command"
    same_text_ids = ["l2-same-text-1", "l2-same-text-2"]
    concurrent_text = "这条并发重传只能应用一次。"
    repeated_text = "这句相同文字是用户有意提交的两次补充。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-idempotency-run",
        objective="继续当前幂等验收任务。",
        name="L2 idempotency proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            concurrent_body = _text_steer_body(
                chat_id=harness.chat.id,
                session_id=session_id,
                command_id=concurrent_id,
                text=concurrent_text,
            )
            concurrent_responses = await asyncio.gather(
                http.post("/console/chat", json=concurrent_body),
                http.post("/console/chat", json=concurrent_body),
            )
            assert [
                response.status_code for response in concurrent_responses
            ] == [
                200,
                200,
            ]
            concurrent_receipts = [
                next(
                    payload["receipt"]
                    for payload in _sse_payloads(response.text)
                    if isinstance(payload.get("receipt"), dict)
                )
                for response in concurrent_responses
            ]
            assert concurrent_receipts[0] == concurrent_receipts[1]

            same_text_responses = []
            for command_id in same_text_ids:
                response = await http.post(
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=command_id,
                        text=repeated_text,
                    ),
                )
                assert response.status_code == 200
                same_text_responses.append(response)
            same_text_receipts = [
                next(
                    payload["receipt"]
                    for payload in _sse_payloads(response.text)
                    if isinstance(payload.get("receipt"), dict)
                )
                for response in same_text_responses
            ]
            assert [
                receipt["command_id"] for receipt in same_text_receipts
            ] == same_text_ids
            assert same_text_receipts[0]["server_seq"] != (
                same_text_receipts[1]["server_seq"]
            )

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=harness.run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[RuntimeMessage.text("user", "继续当前幂等验收任务。")],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            events = [event async for event in runtime.run_turn(runtime_input)]
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 3
            assert [
                request.messages[-1].text_content()
                for request in client.requests
            ] == [concurrent_text, repeated_text, repeated_text]
            assert not any(event.type.startswith("tool.") for event in events)

            command_ids = [concurrent_id, *same_text_ids]
            final_receipts = []
            for command_id in command_ids:
                status = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert status.status_code == 200
                final_receipts.append(status.json())
            assert {
                receipt["execution_status"] for receipt in final_receipts
            } == {"applied"}
            assert (
                len(
                    {receipt["model_request_id"] for receipt in final_receipts}
                )
                == 3
            )

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(harness.run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == command_ids
            assert [item.message.text_content() for item in inbox] == [
                concurrent_text,
                repeated_text,
                repeated_text,
            ]
            assert [contract.run_id for contract in contracts] == [
                harness.run_id
            ]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": harness.run_id,
                "attempt_id": harness.started.attempt_id,
                "concurrent_retry": {
                    "command_id": concurrent_id,
                    "http_statuses": [
                        response.status_code
                        for response in concurrent_responses
                    ],
                    "same_receipt": (
                        concurrent_receipts[0] == concurrent_receipts[1]
                    ),
                    "server_seq": concurrent_receipts[0]["server_seq"],
                },
                "intentional_equal_text": {
                    "command_ids": same_text_ids,
                    "server_seqs": [
                        receipt["server_seq"] for receipt in same_text_receipts
                    ],
                    "model_request_ids": [
                        receipt["model_request_id"]
                        for receipt in final_receipts[1:]
                    ],
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
                "all_execution_statuses": [
                    receipt["execution_status"] for receipt in final_receipts
                ],
            }
            print(
                "STEERING_IDEMPOTENCY_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_fifo_cancel_replace_and_claim_races(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Durable order and pre/post-binding edits remain truthful."""

    session_id = "l2-ordering-session"
    cancelled_id = "l2-cancel-before-claim"
    replaced_id = "l2-replace-old"
    replacement_id = "l2-replace-new"
    correction_id = "l2-after-bound-correction"
    fifo_ids = ["l2-fifo-a", "l2-fifo-b", "l2-fifo-c"]
    window_ids = ["l2-window-a", "l2-window-b"]
    replacement_text = "用修订后的要求继续。"
    fifo_texts = [
        "第一条：保留标题。",
        "第二条：补充依据。",
        "第三条：最后检查约束。",
    ]
    window_texts = [
        "窗口甲同时补充。",
        "窗口乙同时补充。",
    ]
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-ordering-run",
        objective="继续当前顺序与竞争验收任务。",
        name="L2 ordering and race proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            cancelled_admission = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=cancelled_id,
                    text="这条必须在模型看到前取消。",
                ),
            )
            assert cancelled_admission.status_code == 200
            assert (
                _receipt_from_response(cancelled_admission)["execution_status"]
                == "pending"
            )
            cancel_params = {
                "session_id": session_id,
                "chat_id": harness.chat.id,
                "command_id": cancelled_id,
            }
            cancelled = await http.post(
                "/console/chat/control-input/cancel",
                params=cancel_params,
            )
            cancelled_retry = await http.post(
                "/console/chat/control-input/cancel",
                params=cancel_params,
            )
            assert cancelled.status_code == 200
            assert cancelled_retry.status_code == 200
            assert cancelled.json() == cancelled_retry.json()
            assert cancelled.json()["execution_status"] == "cancelled"

            old_admission = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=replaced_id,
                    text="这条旧要求不能进入模型。",
                ),
            )
            assert old_admission.status_code == 200
            replacement_admission = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=replacement_id,
                    text=replacement_text,
                    intent="redirect",
                    replaces_command_id=replaced_id,
                ),
            )
            assert replacement_admission.status_code == 200
            replacement_receipt = _receipt_from_response(replacement_admission)
            assert replacement_receipt["execution_status"] == "pending"
            assert replacement_receipt["replaces_command_id"] == replaced_id

            fifo_receipts = []
            for command_id, text in zip(fifo_ids, fifo_texts, strict=True):
                response = await http.post(
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=command_id,
                        text=text,
                    ),
                )
                assert response.status_code == 200
                fifo_receipts.append(_receipt_from_response(response))
            assert [
                receipt["command_id"] for receipt in fifo_receipts
            ] == fifo_ids
            assert [
                receipt["server_seq"] for receipt in fifo_receipts
            ] == sorted(receipt["server_seq"] for receipt in fifo_receipts)

            async with httpx.AsyncClient(
                base_url=http.base_url,
                timeout=5,
                trust_env=False,
            ) as second_window:
                window_responses = await asyncio.gather(
                    http.post(
                        "/console/chat",
                        json=_text_steer_body(
                            chat_id=harness.chat.id,
                            session_id=session_id,
                            command_id=window_ids[0],
                            text=window_texts[0],
                        ),
                    ),
                    second_window.post(
                        "/console/chat",
                        json=_text_steer_body(
                            chat_id=harness.chat.id,
                            session_id=session_id,
                            command_id=window_ids[1],
                            text=window_texts[1],
                        ),
                    ),
                )
            assert [response.status_code for response in window_responses] == [
                200,
                200,
            ]
            window_receipts = [
                _receipt_from_response(response)
                for response in window_responses
            ]
            ordered_window_inputs = sorted(
                zip(window_receipts, window_ids, window_texts, strict=True),
                key=lambda value: value[0]["server_seq"],
            )
            ordered_window_ids = [value[1] for value in ordered_window_inputs]
            ordered_window_texts = [
                value[2] for value in ordered_window_inputs
            ]
            assert (
                len({receipt["server_seq"] for receipt in window_receipts})
                == 2
            )

            reopened_before_owner = RuntimeTurnControlStore(
                harness.state_workspace
            )
            admissible_before_owner = (
                reopened_before_owner.read_admissible_inbox(harness.run_id)
            )
            assert [item.command_id for item in admissible_before_owner] == [
                replacement_id,
                *fifo_ids,
                *ordered_window_ids,
            ]
            assert cancelled_id not in {
                item.command_id for item in admissible_before_owner
            }
            assert replaced_id not in {
                item.command_id for item in admissible_before_owner
            }

            client = _FirstRequestGateModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=harness.run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "继续当前顺序与竞争验收任务。",
                    )
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": harness.started.revision,
                    "runtime_lease_owner": harness.started.lease_owner,
                    "task_contract": harness.started.task_contract.model_dump(
                        mode="json"
                    ),
                },
            )
            runtime_task = asyncio.create_task(
                _collect_runtime_events(runtime, runtime_input)
            )
            await asyncio.wait_for(
                client.first_request_started.wait(),
                timeout=2,
            )
            bound_cancel = await http.post(
                "/console/chat/control-input/cancel",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": replacement_id,
                },
            )
            assert bound_cancel.status_code == 409
            bound_detail = bound_cancel.json()["detail"]
            assert bound_detail["code"] == "control_input_not_cancellable"
            assert bound_detail["execution_status"] == "bound_to_request"

            correction_text = "已进入模型的历史保留；用这条新纠正继续。"
            correction = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=correction_id,
                    text=correction_text,
                    intent="supplement",
                ),
            )
            assert correction.status_code == 200
            assert (
                _receipt_from_response(correction)["execution_status"]
                == "pending"
            )
            client.release_first_request.set()
            events = await asyncio.wait_for(runtime_task, timeout=3)
            assert events[-1].type == "turn.completed"
            assert [
                request.messages[-1].text_content()
                for request in client.requests
            ] == [
                replacement_text,
                *fifo_texts,
                *ordered_window_texts,
                correction_text,
            ]
            assert not any(event.type.startswith("tool.") for event in events)

            command_ids = [
                cancelled_id,
                replaced_id,
                replacement_id,
                *fifo_ids,
                *window_ids,
                correction_id,
            ]
            statuses = {}
            for command_id in command_ids:
                response = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert response.status_code == 200
                statuses[command_id] = response.json()["execution_status"]
            assert statuses == {
                cancelled_id: "cancelled",
                replaced_id: "superseded",
                replacement_id: "applied",
                **{command_id: "applied" for command_id in fifo_ids},
                **{command_id: "applied" for command_id in window_ids},
                correction_id: "applied",
            }

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(harness.run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.seq for item in inbox] == list(
                range(1, len(inbox) + 1)
            )
            assert [contract.run_id for contract in contracts] == [
                harness.run_id
            ]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": harness.run_id,
                "attempt_id": harness.started.attempt_id,
                "durable_order": [item.command_id for item in inbox],
                "model_order": [
                    request.messages[-1].text_content()
                    for request in client.requests
                ],
                "statuses": statuses,
                "cancel_retry_same_receipt": (
                    cancelled.json() == cancelled_retry.json()
                ),
                "claim_won_cancel_http_status": bound_cancel.status_code,
                "correction_after_claim_admitted": True,
                "concurrent_window_server_order": ordered_window_ids,
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_ORDERING_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_redirect_yields_old_stream_within_same_native_run(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Prove an HTTP redirect reaches the already-running Native owner."""
    from lantai.app.routers import console as console_router

    agent_home = tmp_path / "agent"
    chat_manager = ChatManager(
        repo=JsonChatRepository(tmp_path / "chats.json"),
    )
    tracker = TaskTracker()
    console_channel = ConsoleChannel(
        process=_never_process,
        enabled=True,
        bot_prefix="",
        workspace_dir=agent_home,
    )
    workspace = SimpleNamespace(
        agent_id="default",
        workspace_dir=agent_home,
        chat_manager=chat_manager,
        task_tracker=tracker,
        channel_manager=_ChannelManager(console_channel),
        runner=SimpleNamespace(_chat_manager=chat_manager),
    )

    async def fake_get_agent_for_request(_request):
        return workspace

    monkeypatch.setattr(
        console_router,
        "get_agent_for_request",
        fake_get_agent_for_request,
    )

    session_id = "l2-redirect-session"
    chat = await chat_manager.get_or_create_chat(
        session_id,
        "default",
        "console",
        name="L2 redirect proof",
    )
    binding = await ensure_chat_topic_workspace(
        workspace,
        chat_manager,
        chat,
        agent_id="default",
    )
    chat = binding.chat
    state_workspace = Path(binding.state_workspace_dir)
    run_id = "l2-redirect-run"
    task_contract = build_chat_task_contract(
        objective="沿旧方向长时间思考。",
        client_turn_id="l2-redirect-initial-turn",
        complexity="light",
        work_method_id="standard",
        run_id=run_id,
        attempt_id=f"{run_id}:chat-attempt-1",
    )
    started = await record_chat_runtime_run_started(
        workspace_dir=state_workspace,
        agent_id="default",
        chat_id=chat.id,
        session_id=session_id,
        run_id=run_id,
        effective_workspace_dir=binding.effective_workspace_dir,
        task_contract=task_contract,
    )
    await chat_manager.merge_chat_meta(
        chat.id,
        {
            "last_turn_run_id": run_id,
            "state_workspace_dir": str(state_workspace),
        },
    )
    tracker_generation = await _register_active_external_task(
        tracker, chat.id,
        attempt_id=started.attempt_id,
        root_session_id=session_id,
    )

    app = FastAPI()
    app.include_router(console_router.router)
    with closing(socket.socket(socket.AF_INET, socket.SOCK_STREAM)) as sock:
        sock.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
        sock.bind(("127.0.0.1", 0))
        sock.listen(128)
        port = int(sock.getsockname()[1])
        server = uvicorn.Server(
            uvicorn.Config(
                app,
                host="127.0.0.1",
                port=port,
                log_level="warning",
                lifespan="off",
            )
        )
        server_task = asyncio.create_task(server.serve(sockets=[sock]))
        runtime_task: asyncio.Task | None = None
        try:
            await asyncio.wait_for(_wait_until_started(server), timeout=3)
            async with httpx.AsyncClient(
                base_url=f"http://127.0.0.1:{port}",
                timeout=5,
                trust_env=False,
            ) as http:
                client = _RedirectableRecordingModelClient()
                runtime = NativeLantaiRuntime(
                    model_client_factory=lambda _agent_id: client,
                    tool_registry=LantaiToolRegistry(),
                    turn_timeout_seconds=3,
                )
                runtime_input = RuntimeInput(
                    request=RuntimeRequestContext(
                        agent_id="default",
                        session_id=session_id,
                        user_id="default",
                        run_id=run_id,
                        workspace_dir=binding.effective_workspace_dir,
                    ),
                    messages=[RuntimeMessage.text("user", "沿旧方向长时间思考。")],
                    metadata={
                        "state_workspace_dir": str(state_workspace),
                        "attempt_id": started.attempt_id,
                        "runtime_generation": tracker_generation,
                        "runtime_revision": started.revision,
                        "runtime_lease_owner": started.lease_owner,
                        "task_contract": started.task_contract.model_dump(
                            mode="json"
                        ),
                    },
                )

                async def collect_events():
                    return [
                        event
                        async for event in runtime.run_turn(runtime_input)
                    ]

                runtime_task = asyncio.create_task(collect_events())
                await asyncio.wait_for(client.first_surface.wait(), timeout=1)
                redirect_body = {
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "立即停止旧方向，改为只回答新的重点。",
                                }
                            ],
                        }
                    ],
                    "chat_id": chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": "l2-redirect-command",
                        "control_intent": "redirect",
                    },
                }
                admission_started = time.perf_counter()
                admitted = await http.post("/console/chat", json=redirect_body)
                admission_finished = time.perf_counter()
                assert admitted.status_code == 200
                receipt = _sse_payloads(admitted.text)[-1]["receipt"]
                assert receipt["run_id"] == run_id
                assert receipt["attempt_id"] == started.attempt_id
                assert receipt["execution_status"] == "pending"
                assert receipt["next_waiting_reason"] == "awaiting_safe_yield"

                events = await asyncio.wait_for(runtime_task, timeout=2)
                assert client.first_stream_closed.is_set()
                assert len(client.requests) == 2
                assert client.requests[1].messages[-1].text_content() == (
                    "立即停止旧方向，改为只回答新的重点。"
                )
                assert [
                    request.metadata.get("model_epoch")
                    for request in client.requests
                ] == [1, 2]
                assert [
                    request.metadata.get("intent_revision")
                    for request in client.requests
                ] == [0, 1]
                assert any(
                    event.type == "message.delta"
                    and event.payload.get("delta") == "旧方向中已向用户展示的内容"
                    for event in events
                )
                assert any(
                    event.type == "runtime.warning"
                    and event.payload.get("kind")
                    == "model_stream_yielded_for_redirect"
                    for event in events
                )
                assert [
                    event.summary
                    for event in events
                    if event.type == "message.completed"
                ] == ["已按立即调整的新方向继续。"]
                assert not any(
                    event.type.startswith("tool.") for event in events
                )
                assert events[-1].type == "turn.completed"

                status = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": chat.id,
                        "command_id": "l2-redirect-command",
                    },
                )
                assert status.status_code == 200
                final_receipt = status.json()
                assert final_receipt["execution_status"] == "applied"
                assert final_receipt["model_request_id"] == (
                    f"{run_id}:model:2"
                )
                assert final_receipt["model_epoch"] == 2
                assert final_receipt["intent_revision"] == 1

                contracts = RuntimeRunStore(state_workspace).list_contracts()
                assert [contract.run_id for contract in contracts] == [run_id]
                trace = {
                    "schema_version": 1,
                    "level": "L2",
                    "transport": "loopback_http_socket",
                    "command_id": "l2-redirect-command",
                    "chat_id": chat.id,
                    "run_id": run_id,
                    "attempt_id": started.attempt_id,
                    "tracker_generation": tracker_generation,
                    "http_status": {
                        "admission": admitted.status_code,
                        "query": status.status_code,
                    },
                    "counts": {
                        "model_requests": len(client.requests),
                        "runtime_runs": len(contracts),
                        "agent_loops": 1,
                        "tool_effects": 0,
                    },
                    "lineage": {
                        "model_epochs": [
                            request.metadata.get("model_epoch")
                            for request in client.requests
                        ],
                        "intent_revisions": [
                            request.metadata.get("intent_revision")
                            for request in client.requests
                        ],
                        "final_model_request_id": final_receipt[
                            "model_request_id"
                        ],
                    },
                    "preserved_old_output": True,
                    "old_stream_closed": client.first_stream_closed.is_set(),
                    "latency_ms": {
                        "content_ready_to_admitted": round(
                            (admission_finished - admission_started) * 1000,
                            3,
                        ),
                        "admission_to_new_model_request": round(
                            (client.request_times[1] - admission_finished)
                            * 1000,
                            3,
                        ),
                    },
                    "receipt": {
                        "initial_execution_status": receipt[
                            "execution_status"
                        ],
                        "initial_waiting_reason": receipt[
                            "next_waiting_reason"
                        ],
                        "final_execution_status": final_receipt[
                            "execution_status"
                        ],
                    },
                }
                print(
                    "STEERING_REDIRECT_L2_TRACE="
                    + json.dumps(trace, ensure_ascii=False, sort_keys=True)
                )
        finally:
            if runtime_task is not None and not runtime_task.done():
                runtime_task.cancel()
                with pytest.raises(asyncio.CancelledError):
                    await runtime_task
            server.should_exit = True
            await asyncio.wait_for(server_task, timeout=5)
            await tracker.unregister_external_task(chat.id)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("phase", "case_id"),
    [
        ("first_event", "ST-YL-01"),
        ("thinking", "ST-YL-02"),
        ("late_tool_call", "ST-YL-04"),
        ("partial_tool_json", "ST-YL-05"),
    ],
    ids=[
        "redirect-before-first-event",
        "redirect-during-thinking",
        "late-old-epoch-tool-call",
        "partial-tool-json",
    ],
)
async def test_local_http_redirect_yields_at_precise_model_boundaries(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    phase: str,
    case_id: str,
) -> None:
    """Retire one model epoch without executing incomplete or late tools."""

    session_id = f"l2-safe-yield-{phase}-session"
    run_id = f"l2-safe-yield-{phase}-run"
    command_id = f"l2-safe-yield-{phase}-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="沿旧方向持续推理；没有明确完整调用时不得执行工具。",
        name=f"L2 safe yield {phase}",
    )
    tool_effects: list[str] = []

    async def record_effect(value: str) -> str:
        tool_effects.append(value)
        return "recorded"

    registry = LantaiToolRegistry()
    registry.register_tool(
        ToolSpec(
            name="record_effect",
            read_only=True,
            input_schema={
                "type": "object",
                "properties": {"value": {"type": "string"}},
                "required": ["value"],
            },
        ),
        record_effect,
    )
    client = _SafeYieldBoundaryModelClient(phase)
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=registry,
        progressive_tools=False,
        turn_timeout_seconds=3,
    )
    runtime_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default",
            session_id=session_id,
            user_id="default",
            run_id=run_id,
            workspace_dir=harness.binding.effective_workspace_dir,
        ),
        messages=[
            RuntimeMessage.text(
                "user",
                "沿旧方向持续推理；没有明确完整调用时不得执行工具。",
            )
        ],
        metadata={
            "state_workspace_dir": str(harness.state_workspace),
            "attempt_id": harness.started.attempt_id,
            "runtime_generation": harness.tracker_generation,
            "runtime_revision": harness.started.revision,
            "runtime_lease_owner": harness.started.lease_owner,
            "task_contract": harness.started.task_contract.model_dump(
                mode="json"
            ),
        },
    )
    runtime_task: asyncio.Task | None = None
    try:
        async with _loopback_console_http() as (http, _port):
            runtime_task = asyncio.create_task(
                _collect_runtime_events(runtime, runtime_input)
            )
            await asyncio.wait_for(client.boundary_reached.wait(), timeout=1)
            admission_started = time.perf_counter()
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text="立即退役旧方向，只按新的安全要求继续。",
                    intent="redirect",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            admission_finished = time.perf_counter()
            assert admitted.status_code == 200
            initial_receipt = _receipt_from_response(admitted)
            assert initial_receipt["execution_status"] == "pending"
            assert initial_receipt["next_waiting_reason"] == (
                "awaiting_safe_yield"
            )

            events = await asyncio.wait_for(runtime_task, timeout=2)
            assert client.old_stream_closed.is_set()
            assert len(client.requests) == 2
            assert client.requests[1].messages[-1].text_content() == (
                "立即退役旧方向，只按新的安全要求继续。"
            )
            assert [
                request.metadata.get("model_epoch")
                for request in client.requests
            ] == [1, 2]
            assert [
                request.metadata.get("intent_revision")
                for request in client.requests
            ] == [0, 1]
            assert tool_effects == []
            assert not any(event.type.startswith("tool.") for event in events)
            assert not any(
                block.type == "tool_result"
                for request in client.requests
                for message in request.messages
                for block in message.content
            )
            warning = next(
                event
                for event in events
                if event.type == "runtime.warning"
                and event.payload.get("kind")
                == "model_stream_yielded_for_redirect"
            )
            assert events[-1].type == "turn.completed"
            if phase == "thinking":
                assert any(
                    event.type == "message.delta"
                    and event.payload.get("content_type") == "thinking"
                    and event.payload.get("delta") == "旧方向仍在思考"
                    for event in events
                )
            if phase == "late_tool_call":
                assert client.late_tool_returned is True
            if phase in {"late_tool_call", "partial_tool_json"}:
                assert (
                    warning.payload["details"]["unexecuted_tool_call_ids"]
                    == []
                )

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == f"{run_id}:model:2"
            assert final_receipt["model_epoch"] == 2
            assert final_receipt["intent_revision"] == 1

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert len(inbox) == 1
            assert [contract.run_id for contract in contracts] == [run_id]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": case_id,
                "phase": phase,
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "http_status": {
                    "admission": admitted.status_code,
                    "query": final.status_code,
                },
                "receipt": {
                    "initial_execution_status": initial_receipt[
                        "execution_status"
                    ],
                    "initial_waiting_reason": initial_receipt[
                        "next_waiting_reason"
                    ],
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                },
                "lineage": {
                    "model_epochs": [
                        request.metadata.get("model_epoch")
                        for request in client.requests
                    ],
                    "intent_revisions": [
                        request.metadata.get("intent_revision")
                        for request in client.requests
                    ],
                    "final_model_request_id": final_receipt[
                        "model_request_id"
                    ],
                },
                "old_stream_closed": client.old_stream_closed.is_set(),
                "late_old_epoch_tool_returned": (client.late_tool_returned),
                "partial_tool_result_fabricated": False,
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                    "tool_effects": len(tool_effects),
                },
                "latency_ms": {
                    "content_ready_to_admitted": round(
                        (admission_finished - admission_started) * 1000,
                        3,
                    ),
                    "admission_to_new_model_request": round(
                        (client.request_times[1] - admission_finished) * 1000,
                        3,
                    ),
                },
            }
            print(
                "STEERING_SAFE_YIELD_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        if runtime_task is not None and not runtime_task.done():
            runtime_task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await runtime_task
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_redirect_bounds_hanging_model_stream_close(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """A stuck provider close must not strand an admitted redirect."""

    session_id = "l2-hanging-close-session"
    run_id = "l2-hanging-close-run"
    command_id = "l2-hanging-close-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="沿旧方向等待；收到改向后必须有界退役本地模型流。",
        name="L2 hanging model stream close",
    )
    client = _HangingCloseModelClient()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
        progressive_tools=False,
        turn_timeout_seconds=3,
    )
    # Keep the production default conservative while making the fault
    # injection deterministic and fast.
    runtime._model_stream_close_timeout_seconds = 0.05
    runtime_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default",
            session_id=session_id,
            user_id="default",
            run_id=run_id,
            workspace_dir=harness.binding.effective_workspace_dir,
        ),
        messages=[
            RuntimeMessage.text(
                "user",
                "沿旧方向等待；收到改向后必须有界退役本地模型流。",
            )
        ],
        metadata={
            "state_workspace_dir": str(harness.state_workspace),
            "attempt_id": harness.started.attempt_id,
            "runtime_generation": harness.tracker_generation,
            "runtime_revision": harness.started.revision,
            "runtime_lease_owner": harness.started.lease_owner,
            "task_contract": harness.started.task_contract.model_dump(
                mode="json"
            ),
        },
    )
    caplog.set_level(
        logging.WARNING, logger="lantai.runtime.native_turn_support"
    )
    runtime_task: asyncio.Task | None = None
    try:
        async with _loopback_console_http() as (http, _port):
            runtime_task = asyncio.create_task(
                _collect_runtime_events(runtime, runtime_input)
            )
            await asyncio.wait_for(client.boundary_reached.wait(), timeout=1)
            admission_started = time.perf_counter()
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text="立即按新方向继续，不等待旧连接关闭。",
                    intent="redirect",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            admission_finished = time.perf_counter()
            assert admitted.status_code == 200
            initial_receipt = _receipt_from_response(admitted)
            assert initial_receipt["execution_status"] == "pending"
            assert initial_receipt["next_waiting_reason"] == (
                "awaiting_safe_yield"
            )

            done, _pending = await asyncio.wait(
                {runtime_task},
                timeout=0.6,
            )
            assert runtime_task in done, (
                "an unresponsive provider aclose stranded the admitted "
                "redirect"
            )
            events = runtime_task.result()
            assert client.close_started.is_set()
            assert not client.close_finished.is_set()
            assert len(client.requests) == 2
            assert client.requests[1].messages[-1].text_content() == (
                "立即按新方向继续，不等待旧连接关闭。"
            )
            assert [
                request.metadata.get("model_epoch")
                for request in client.requests
            ] == [1, 2]
            assert [
                request.metadata.get("intent_revision")
                for request in client.requests
            ] == [0, 1]
            assert not any(
                event.payload.get("kind") == "remote_model_stopped"
                for event in events
            )
            assert any(
                "remote cancellation was not confirmed" in record.message
                for record in caplog.records
            )

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == f"{run_id}:model:2"
            assert final_receipt["model_epoch"] == 2
            assert final_receipt["intent_revision"] == 1

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-YL-08",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "http_status": {
                    "admission": admitted.status_code,
                    "query": final.status_code,
                },
                "receipt": {
                    "initial_execution_status": initial_receipt[
                        "execution_status"
                    ],
                    "initial_waiting_reason": initial_receipt[
                        "next_waiting_reason"
                    ],
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                },
                "lineage": {
                    "model_epochs": [
                        request.metadata.get("model_epoch")
                        for request in client.requests
                    ],
                    "intent_revisions": [
                        request.metadata.get("intent_revision")
                        for request in client.requests
                    ],
                    "final_model_request_id": final_receipt[
                        "model_request_id"
                    ],
                },
                "stream_close": {
                    "started": client.close_started.is_set(),
                    "completed": client.close_finished.is_set(),
                    "timeout_seconds": (
                        runtime._model_stream_close_timeout_seconds
                    ),
                    "remote_cancellation_confirmed": False,
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                },
                "latency_ms": {
                    "content_ready_to_admitted": round(
                        (admission_finished - admission_started) * 1000,
                        3,
                    ),
                    "admission_to_new_model_request": round(
                        (client.request_times[1] - admission_finished) * 1000,
                        3,
                    ),
                },
            }
            print(
                "STEERING_HANGING_CLOSE_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        client.release_close.set()
        if runtime_task is not None and not runtime_task.done():
            done, _pending = await asyncio.wait({runtime_task}, timeout=2)
            if runtime_task not in done:
                runtime_task.cancel()
                done, _pending = await asyncio.wait(
                    {runtime_task},
                    timeout=1,
                )
            if runtime_task in done:
                try:
                    runtime_task.result()
                except asyncio.CancelledError:
                    pass
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_three_rapid_redirects_preserve_fifo_dispositions(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Three urgent corrections stay ordered even when epochs are retired."""

    session_id = "l2-rapid-redirect-session"
    run_id = "l2-rapid-redirect-run"
    command_ids = [
        f"l2-rapid-redirect-command-{index}" for index in range(1, 4)
    ]
    correction_texts = [
        "第一条纠正：改用方案甲。",
        "第二条纠正：放弃方案甲，改用方案乙。",
        "第三条纠正：以方案丙为最终方向。",
    ]
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="保持旧方向等待连续纠正。",
        name="L2 three rapid redirects",
    )
    client = _RapidRedirectModelClient()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
        progressive_tools=False,
        turn_timeout_seconds=3,
    )
    runtime_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default",
            session_id=session_id,
            user_id="default",
            run_id=run_id,
            workspace_dir=harness.binding.effective_workspace_dir,
        ),
        messages=[RuntimeMessage.text("user", "保持旧方向等待连续纠正。")],
        metadata={
            "state_workspace_dir": str(harness.state_workspace),
            "attempt_id": harness.started.attempt_id,
            "runtime_generation": harness.tracker_generation,
            "runtime_revision": harness.started.revision,
            "runtime_lease_owner": harness.started.lease_owner,
            "task_contract": harness.started.task_contract.model_dump(
                mode="json"
            ),
        },
    )
    runtime_task: asyncio.Task | None = None
    try:
        async with _loopback_console_http() as (http, _port):
            runtime_task = asyncio.create_task(
                _collect_runtime_events(runtime, runtime_input)
            )
            await asyncio.wait_for(
                client.first_boundary_reached.wait(),
                timeout=1,
            )
            admission_started = time.perf_counter()
            admitted_receipts: list[dict] = []
            for command_id, correction_text in zip(
                command_ids,
                correction_texts,
                strict=True,
            ):
                response = await http.post(
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=command_id,
                        text=correction_text,
                        intent="redirect",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    ),
                )
                assert response.status_code == 200
                admitted_receipts.append(_receipt_from_response(response))
            admission_finished = time.perf_counter()
            assert [
                receipt["server_seq"] for receipt in admitted_receipts
            ] == sorted(receipt["server_seq"] for receipt in admitted_receipts)

            events = await asyncio.wait_for(runtime_task, timeout=2)
            # A queued redirect may supersede a bound request before its
            # provider iterator is entered. That is the permitted coalesced
            # yield: every durable item keeps an ordered disposition without
            # spending an unnecessary provider call.
            assert 2 <= len(client.requests) <= 4
            assert client.closed_request_ordinals == list(
                range(1, len(client.requests))
            )
            provider_model_epochs = [
                request.metadata.get("model_epoch")
                for request in client.requests
            ]
            provider_intent_revisions = [
                request.metadata.get("intent_revision")
                for request in client.requests
            ]
            assert provider_model_epochs[0] == 1
            assert provider_model_epochs[-1] == 4
            assert provider_model_epochs == sorted(provider_model_epochs)
            assert provider_intent_revisions[0] == 0
            assert provider_intent_revisions[-1] == 3
            assert provider_intent_revisions == sorted(
                provider_intent_revisions
            )
            command_history_by_request = [
                [
                    message.text_content()
                    for message in request.messages
                    if message.metadata.get("command_id") in command_ids
                ]
                for request in client.requests
            ]
            assert command_history_by_request == [
                correction_texts[:revision]
                for revision in provider_intent_revisions
            ]
            yield_warnings = [
                event
                for event in events
                if event.type == "runtime.warning"
                and event.payload.get("kind")
                == "model_stream_yielded_for_redirect"
            ]
            assert [
                event.payload.get("client_turn_id") for event in yield_warnings
            ] == command_ids

            final_receipts: list[dict] = []
            for command_id in command_ids:
                response = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert response.status_code == 200
                final_receipts.append(response.json())
            assert [
                receipt["execution_status"] for receipt in final_receipts
            ] == ["superseded", "superseded", "applied"]
            assert [
                receipt["model_request_id"] for receipt in final_receipts
            ] == [
                f"{run_id}:model:2",
                f"{run_id}:model:3",
                f"{run_id}:model:4",
            ]
            assert final_receipts[0]["superseded_by_command_id"] == (
                command_ids[1]
            )
            assert final_receipts[1]["superseded_by_command_id"] == (
                command_ids[2]
            )
            assert final_receipts[2]["intent_revision"] == 3

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == command_ids
            assert [contract.run_id for contract in contracts] == [run_id]
            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-YL-09",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_ids": command_ids,
                "server_seq": [
                    receipt["server_seq"] for receipt in admitted_receipts
                ],
                "execution_statuses": [
                    receipt["execution_status"] for receipt in final_receipts
                ],
                "superseded_by": [
                    receipt["superseded_by_command_id"]
                    for receipt in final_receipts
                ],
                "lineage": {
                    "provider_model_epochs": provider_model_epochs,
                    "provider_intent_revisions": (provider_intent_revisions),
                    "bound_model_request_ids": [
                        receipt["model_request_id"]
                        for receipt in final_receipts
                    ],
                },
                "counts": {
                    "intentional_submissions": 3,
                    "durable_inputs": len(inbox),
                    "provider_stream_requests": len(client.requests),
                    "bound_control_model_requests": len(final_receipts),
                    "runtime_runs": len(contracts),
                    "agent_loops": 1,
                },
                "latency_ms": {
                    "three_admissions_total": round(
                        (admission_finished - admission_started) * 1000,
                        3,
                    ),
                    "first_admission_start_to_final_request": round(
                        (client.request_times[-1] - admission_started) * 1000,
                        3,
                    ),
                },
            }
            print(
                "STEERING_RAPID_REDIRECT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        if runtime_task is not None and not runtime_task.done():
            runtime_task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await runtime_task
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_redirect_accumulates_old_and_new_cost_evidence(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A yielded dispatch keeps known cost and an explicit unknown tail."""

    session_id = "l2-cost-redirect-session"
    run_id = "l2-cost-redirect-run"
    command_id = "l2-cost-redirect-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="旧推理产生费用后立即改向，并保留完整费用血缘。",
        name="L2 redirect cost accumulation",
    )
    inner_client = _CostedRedirectModelClient()
    client = TokenRecordingModelClient(inner_client)
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
        progressive_tools=False,
        turn_timeout_seconds=3,
    )
    runtime_input = RuntimeInput(
        request=RuntimeRequestContext(
            agent_id="default",
            session_id=session_id,
            user_id="default",
            run_id=run_id,
            workspace_dir=harness.binding.effective_workspace_dir,
        ),
        messages=[
            RuntimeMessage.text(
                "user",
                "旧推理产生费用后立即改向，并保留完整费用血缘。",
            )
        ],
        metadata={
            "state_workspace_dir": str(harness.state_workspace),
            "attempt_id": harness.started.attempt_id,
            "runtime_generation": harness.tracker_generation,
            "runtime_revision": harness.started.revision,
            "runtime_lease_owner": harness.started.lease_owner,
            "task_contract": harness.started.task_contract.model_dump(
                mode="json"
            ),
        },
    )
    runtime_task: asyncio.Task | None = None
    try:
        async with _loopback_console_http() as (http, _port):
            runtime_task = asyncio.create_task(
                _collect_runtime_events(runtime, runtime_input)
            )
            await asyncio.wait_for(
                inner_client.old_usage_observed.wait(),
                timeout=1,
            )
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text="立即改向；旧调用费用不能丢。",
                    intent="redirect",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert admitted.status_code == 200
            await asyncio.wait_for(runtime_task, timeout=2)
            assert inner_client.old_stream_closed.is_set()
            assert len(inner_client.requests) == 2

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == f"{run_id}:model:2"

            evidence_path = (
                harness.state_workspace
                / ".lantai/runtime/v1/cost-evidence.json"
            )
            rows = json.loads(evidence_path.read_text(encoding="utf-8"))
            exact_rows = [row for row in rows if row["status"] == "exact"]
            unknown_rows = [row for row in rows if row["status"] == "unknown"]
            assert {
                (row["effect_id"], row["amount_micro_cny"])
                for row in exact_rows
            } == {
                (f"{run_id}:model:1", 1_111),
                (f"{run_id}:model:2", 2_222),
            }
            assert [row["effect_id"] for row in unknown_rows] == [
                f"{run_id}:model:1"
            ]
            assert sum(row["amount_micro_cny"] for row in exact_rows) == 3_333

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-YL-10",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "lineage": {
                    "model_request_ids": [
                        request.metadata.get("model_request_id")
                        for request in inner_client.requests
                    ],
                    "model_epochs": [
                        request.metadata.get("model_epoch")
                        for request in inner_client.requests
                    ],
                    "intent_revisions": [
                        request.metadata.get("intent_revision")
                        for request in inner_client.requests
                    ],
                },
                "cost_evidence": {
                    "exact_micro_cny": {
                        row["effect_id"]: row["amount_micro_cny"]
                        for row in exact_rows
                    },
                    "exact_total_micro_cny": sum(
                        row["amount_micro_cny"] for row in exact_rows
                    ),
                    "unknown_effect_ids": [
                        row["effect_id"] for row in unknown_rows
                    ],
                },
                "counts": {
                    "runtime_runs": 1,
                    "agent_loops": 1,
                    "provider_stream_requests": len(inner_client.requests),
                    "cost_evidence_rows": len(rows),
                },
            }
            print(
                "STEERING_COST_ACCUMULATION_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        if runtime_task is not None and not runtime_task.done():
            runtime_task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await runtime_task
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_receipt_is_usable_before_transport_eof(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A complete production receipt frame is actionable on an open socket."""

    from lantai.app.routers import console as console_router

    session_id = "l2-open-receipt-session"
    run_id = "l2-open-receipt-run"
    command_id = "l2-open-receipt-command"
    objective = "保持主任务运行并接收流式控制回执。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective=objective,
        name="L2 open receipt stream",
    )
    app = FastAPI()
    app.include_router(console_router.router)
    fault_app = _ReceiptStreamFaultInjector(
        app,
        {command_id: "hold_open"},
    )
    client = _RecordingModelClient()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
    )
    try:
        async with _loopback_console_http(fault_app) as (http, _port):
            read_started = time.perf_counter()
            try:
                async with http.stream(
                    "POST",
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=command_id,
                        text="在连接继续保持时先确认接收。",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    ),
                ) as response:
                    assert response.status_code == 200
                    receipt = await asyncio.wait_for(
                        _read_turn_receipt_before_eof(response),
                        timeout=0.5,
                    )
                    read_finished = time.perf_counter()
                    assert receipt["command_id"] == command_id
                    assert receipt["admission_status"] == "admitted"
                    assert receipt["execution_status"] == "pending"
                    assert not fault_app._event(
                        fault_app.release,
                        command_id,
                    ).is_set()
            finally:
                fault_app.release_stream(command_id)

            events = await _collect_runtime_events(
                runtime,
                _runtime_input_for_harness(
                    harness,
                    session_id=session_id,
                    objective=objective,
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            assert final.json()["execution_status"] == "applied"
            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-RC-07",
                "transport": "loopback_http_socket_open_after_receipt",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "command_id": command_id,
                "receipt_before_eof": True,
                "receipt_latency_ms": round(
                    (read_finished - read_started) * 1000,
                    3,
                ),
                "final_execution_status": final.json()["execution_status"],
                "counts": {
                    "durable_inputs": 1,
                    "model_requests": len(client.requests),
                    "runtime_runs": 1,
                    "agent_loops": 1,
                },
            }
            print(
                "STEERING_OPEN_RECEIPT_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        fault_app.release_stream(command_id)
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_receipt_fragmentation_and_versions(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Real receipt bytes tolerate fragmentation and reject future schemas."""

    from lantai.app.routers import console as console_router

    session_id = "l2-receipt-version-session"
    run_id = "l2-receipt-version-run"
    fragmented_command_id = "l2-fragmented-receipt-command"
    unknown_command_id = "l2-unknown-receipt-schema-command"
    objective = "验证控制回执的分片与版本边界。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective=objective,
        name="L2 receipt fragmentation and versioning",
    )
    app = FastAPI()
    app.include_router(console_router.router)
    fault_app = _ReceiptStreamFaultInjector(
        app,
        {
            fragmented_command_id: "fragmented_multiline",
            unknown_command_id: "unknown_schema",
        },
    )
    client = _RecordingModelClient()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
    )
    try:
        async with _loopback_console_http(fault_app) as (http, _port):
            try:
                async with http.stream(
                    "POST",
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=fragmented_command_id,
                        text="解析跨字节块、多行 data 和未知字段。",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    ),
                ) as response:
                    assert response.status_code == 200
                    fragmented_receipt = await asyncio.wait_for(
                        _read_turn_receipt_before_eof(response),
                        timeout=0.5,
                    )
                    assert fragmented_receipt["command_id"] == (
                        fragmented_command_id
                    )
                    assert fragmented_receipt["unknown_future_field"] == {
                        "safe_to_ignore": True,
                    }
                    assert not fault_app._event(
                        fault_app.release,
                        fragmented_command_id,
                    ).is_set()
            finally:
                fault_app.release_stream(fragmented_command_id)

            try:
                async with http.stream(
                    "POST",
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=unknown_command_id,
                        text="未来版本必须明确失败，不能走普通提交。",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    ),
                ) as response:
                    assert response.status_code == 200
                    with pytest.raises(
                        ValueError,
                        match="Unsupported TurnReceipt schema version 99",
                    ):
                        await asyncio.wait_for(
                            _read_turn_receipt_before_eof(response),
                            timeout=0.5,
                        )
            finally:
                fault_app.release_stream(unknown_command_id)

            events = await _collect_runtime_events(
                runtime,
                _runtime_input_for_harness(
                    harness,
                    session_id=session_id,
                    objective=objective,
                ),
            )
            assert events[-1].type == "turn.completed"
            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            assert [item.command_id for item in inbox] == [
                fragmented_command_id,
                unknown_command_id,
            ]
            final_receipts = []
            for command_id in (
                fragmented_command_id,
                unknown_command_id,
            ):
                status = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert status.status_code == 200
                final_receipts.append(status.json())
            assert {
                receipt["execution_status"] for receipt in final_receipts
            } == {"applied"}
            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-RC-08",
                "transport": "loopback_http_socket_fragmented_sse",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "fragmented_command_id": fragmented_command_id,
                "unknown_schema_command_id": unknown_command_id,
                "fragmented_multiline_parsed_before_eof": True,
                "unknown_field_ignored": True,
                "unknown_schema_version_rejected": 99,
                "final_execution_statuses": [
                    receipt["execution_status"] for receipt in final_receipts
                ],
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": 1,
                    "agent_loops": 1,
                },
            }
            print(
                "STEERING_RECEIPT_VERSION_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        fault_app.release_stream(fragmented_command_id)
        fault_app.release_stream(unknown_command_id)
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_warm_admission_latency_reports_distribution(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Measure warm content-ready admission without provider/upload time."""

    session_id = "l2-warm-admission-session"
    run_id = "l2-warm-admission-run"
    sample_count = 30
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="执行本地暖接收延迟验收，不调用外部工具。",
        name="L2 warm admission latency proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "warm-guidance.png",
                        _ONE_PIXEL_PNG,
                        "image/png",
                    )
                },
            )
            assert uploaded.status_code == 200
            attachment = uploaded.json()

            warmup = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id="l2-warm-admission-warmup",
                    text="暖接收预热样本。",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert warmup.status_code == 200
            assert _receipt_from_response(warmup)["admission_status"] == (
                "admitted"
            )

            latencies_ms: list[float] = []
            command_ids: list[str] = []
            sample_kinds: list[str] = []
            for index in range(sample_count):
                command_id = f"l2-warm-admission-{index:02d}"
                command_ids.append(command_id)
                if index % 3 == 0:
                    sample_kind = "text"
                    body = _text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=command_id,
                        text=f"暖接收文字样本 {index:02d}。",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    )
                else:
                    sample_kind = "image" if index % 3 == 1 else "image_text"
                    content = []
                    if sample_kind == "image_text":
                        content.append(
                            {
                                "type": "text",
                                "text": f"暖接收图文样本 {index:02d}。",
                            }
                        )
                    content.append(
                        {
                            "type": "image",
                            "image_url": attachment["url"],
                            "filename": attachment["file_name"],
                            "media_type": "image/png",
                            "file_size": attachment["size"],
                            "stored_name": attachment["stored_name"],
                            "attachment_digest": attachment["digest"],
                            "attachment_version": attachment["version"],
                            "attachment_upload_receipt": attachment[
                                "upload_receipt"
                            ],
                        }
                    )
                    body = {
                        "input": [{"role": "user", "content": content}],
                        "chat_id": harness.chat.id,
                        "session_id": session_id,
                        "user_id": "default",
                        "channel": "console",
                        "stream": True,
                        "operation": "steer",
                        "meta": {
                            "client_turn_id": command_id,
                            "control_intent": "supplement",
                            "expected_run_id": run_id,
                            "expected_attempt_id": (
                                harness.started.attempt_id
                            ),
                        },
                    }
                sample_kinds.append(sample_kind)
                started_at = time.perf_counter()
                response = await http.post("/console/chat", json=body)
                elapsed_ms = (time.perf_counter() - started_at) * 1000
                assert response.status_code == 200
                receipt = _receipt_from_response(response)
                assert receipt["command_id"] == command_id
                assert receipt["run_id"] == run_id
                assert receipt["attempt_id"] == harness.started.attempt_id
                assert receipt["admission_status"] == "admitted"
                assert receipt["execution_status"] == "pending"
                latencies_ms.append(elapsed_ms)

            ordered = sorted(latencies_ms)

            def nearest_rank(percent: int) -> float:
                rank = (percent * len(ordered) + 99) // 100
                return ordered[max(rank - 1, 0)]

            p50_ms = nearest_rank(50)
            p95_ms = nearest_rank(95)
            maximum_ms = ordered[-1]
            assert p95_ms <= 1000

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            assert len(inbox) == sample_count + 1
            assert [item.command_id for item in inbox[1:]] == command_ids
            assert len(set(command_ids)) == sample_count
            assert sample_kinds.count("text") == 10
            assert sample_kinds.count("image") == 10
            assert sample_kinds.count("image_text") == 10
            assert (
                len(RuntimeRunStore(harness.state_workspace).list_contracts())
                == 1
            )

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-PF-01",
                "transport": "loopback_http_socket",
                "measurement": "content_ready_to_durable_admitted",
                "clock": "time.perf_counter",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "warmup_samples_excluded": 1,
                "sample_count": sample_count,
                "sample_kinds": {
                    "text": sample_kinds.count("text"),
                    "image": sample_kinds.count("image"),
                    "image_text": sample_kinds.count("image_text"),
                },
                "latency_ms": {
                    "p50": round(p50_ms, 3),
                    "p95": round(p95_ms, 3),
                    "max": round(maximum_ms, 3),
                    "samples": [round(value, 3) for value in latencies_ms],
                },
                "target_ms": {"p95": 1000},
                "excluded": [
                    "attachment_upload",
                    "provider_network",
                    "provider_thinking",
                ],
                "counts": {
                    "intentional_submissions": sample_count + 1,
                    "logical_commands": len(inbox),
                    "canonical_user_items": len(inbox),
                    "uploads": 1,
                    "model_requests": 0,
                    "native_steer_calls": 0,
                    "tool_invocations": 0,
                    "external_effects": 0,
                    "runtime_runs": 1,
                    "owner_max_concurrency": 1,
                    "duplicate_control_replays": 0,
                    "authorization_failures": 0,
                },
            }
            print(
                "STEERING_WARM_ADMISSION_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_stop_bypasses_a_full_control_queue(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """A saturated guidance queue cannot starve the independent Stop path."""

    from lantai.runtime import turn_control

    queue_limit = 4
    monkeypatch.setattr(
        turn_control,
        "MAX_PENDING_CONTROL_INPUTS_PER_RUN",
        queue_limit,
    )
    session_id = "l2-stop-priority-session"
    run_id = "l2-stop-priority-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="执行队列饱和时的 Stop 优先验收。",
        name="L2 stop priority proof",
    )
    try:
        async with _loopback_console_http() as (http, _port):
            for index in range(queue_limit):
                response = await http.post(
                    "/console/chat",
                    json=_text_steer_body(
                        chat_id=harness.chat.id,
                        session_id=session_id,
                        command_id=f"l2-stop-priority-{index}",
                        text=f"填满控制队列的输入 {index}。",
                        expected_run_id=run_id,
                        expected_attempt_id=harness.started.attempt_id,
                    ),
                )
                assert response.status_code == 200
                assert (
                    _receipt_from_response(response)["execution_status"]
                    == "pending"
                )

            rejected = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id="l2-stop-priority-overflow",
                    text="这条必须因队列已满而拒绝。",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert rejected.status_code == 429
            rejected_detail = rejected.json()["detail"]
            assert rejected_detail == {
                "code": "queue_full",
                "command_id": "l2-stop-priority-overflow",
                "limit_kind": "queue",
                "limit": queue_limit,
                "message": (
                    "Too many control inputs are waiting for this run."
                ),
            }

            stop_started = time.perf_counter()
            stopped = await http.post(
                "/console/chat/stop",
                params={
                    "chat_id": harness.chat.id,
                    "turn_id": harness.started.attempt_id,
                },
            )
            stop_latency_ms = (time.perf_counter() - stop_started) * 1000
            assert stopped.status_code == 200
            assert stopped.json() == {"stopped": True}
            assert stop_latency_ms <= 1000
            assert await harness.tracker.get_status(harness.chat.id) == "idle"

            late = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id="l2-stop-priority-after-stop",
                    text="Stop 后不能因队列状态重新接收。",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert late.status_code == 409
            late_detail = late.json()["detail"]
            assert isinstance(late_detail, str)
            assert "current run has already ended" in late_detail.lower()

            control_store = RuntimeTurnControlStore(harness.state_workspace)
            inbox = control_store.read_inbox(run_id)
            assert len(inbox) == queue_limit
            assert control_store.current_dispositions(run_id) == {}
            assert (
                len(RuntimeRunStore(harness.state_workspace).list_contracts())
                == 1
            )

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-PF-05",
                "transport": "loopback_http_socket",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "queue": {
                    "configured_limit": queue_limit,
                    "pending_at_stop": len(inbox),
                    "overflow_http_status": rejected.status_code,
                    "overflow_error_code": rejected_detail["code"],
                },
                "stop": {
                    "http_status": stopped.status_code,
                    "latency_ms": round(stop_latency_ms, 3),
                    "tracker_status_after": await harness.tracker.get_status(
                        harness.chat.id
                    ),
                    "post_stop_control_status": late.status_code,
                    "post_stop_control_error": "current_run_ended",
                },
                "counts": {
                    "intentional_submissions": queue_limit + 2,
                    "durable_inputs": len(inbox),
                    "runtime_runs": 1,
                    "agent_loops_started_by_control_handler": 0,
                    "model_requests": 0,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_STOP_PRIORITY_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_high_frequency_input_and_cancel_stays_bounded(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Burst admission/cancel remains within the durable active limit."""

    from lantai.runtime import turn_control

    queue_limit = 16
    first_burst_size = 64
    refill_burst_size = 32
    monkeypatch.setattr(
        turn_control,
        "MAX_PENDING_CONTROL_INPUTS_PER_RUN",
        queue_limit,
    )
    session_id = "l2-high-frequency-session"
    run_id = "l2-high-frequency-run"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="执行高频补充与取消的有界容量验收。",
        name="L2 high frequency capacity proof",
    )
    current_task = asyncio.current_task()
    try:
        async with _loopback_console_http() as (http, _port):
            await asyncio.sleep(0)
            baseline_tasks = {
                task
                for task in asyncio.all_tasks()
                if task is not current_task and not task.done()
            }
            first_ids = [
                f"l2-high-frequency-first-{index:02d}"
                for index in range(first_burst_size)
            ]
            refill_ids = [
                f"l2-high-frequency-refill-{index:02d}"
                for index in range(refill_burst_size)
            ]
            limits = httpx.Limits(
                max_connections=8,
                max_keepalive_connections=2,
            )
            async with httpx.AsyncClient(
                base_url=http.base_url,
                timeout=10,
                trust_env=False,
                limits=limits,
            ) as burst_client:

                async def submit(command_id: str) -> httpx.Response:
                    return await burst_client.post(
                        "/console/chat",
                        json=_text_steer_body(
                            chat_id=harness.chat.id,
                            session_id=session_id,
                            command_id=command_id,
                            text=f"高频补充 {command_id}。",
                            expected_run_id=run_id,
                            expected_attempt_id=harness.started.attempt_id,
                        ),
                    )

                first_responses = await asyncio.gather(
                    *(submit(command_id) for command_id in first_ids)
                )
                first_accepted_ids = [
                    command_id
                    for command_id, response in zip(
                        first_ids,
                        first_responses,
                        strict=True,
                    )
                    if response.status_code == 200
                ]
                first_accepted_server_order = [
                    command_id
                    for _server_seq, command_id in sorted(
                        (
                            _receipt_from_response(response)["server_seq"],
                            command_id,
                        )
                        for command_id, response in zip(
                            first_ids,
                            first_responses,
                            strict=True,
                        )
                        if response.status_code == 200
                    )
                ]
                first_rejected = [
                    response
                    for response in first_responses
                    if response.status_code == 429
                ]
                assert len(first_accepted_ids) == queue_limit
                assert len(first_rejected) == (first_burst_size - queue_limit)
                assert all(
                    response.json()["detail"]["code"] == "queue_full"
                    for response in first_rejected
                )

                cancelled_ids = first_accepted_ids[::2]
                cancel_responses = await asyncio.gather(
                    *(
                        burst_client.post(
                            "/console/chat/control-input/cancel",
                            params={
                                "session_id": session_id,
                                "chat_id": harness.chat.id,
                                "command_id": command_id,
                            },
                        )
                        for command_id in cancelled_ids
                    )
                )
                assert all(
                    response.status_code == 200
                    and response.json()["execution_status"] == "cancelled"
                    for response in cancel_responses
                )

                refill_responses = await asyncio.gather(
                    *(submit(command_id) for command_id in refill_ids)
                )
                refill_accepted_ids = [
                    command_id
                    for command_id, response in zip(
                        refill_ids,
                        refill_responses,
                        strict=True,
                    )
                    if response.status_code == 200
                ]
                refill_accepted_server_order = [
                    command_id
                    for _server_seq, command_id in sorted(
                        (
                            _receipt_from_response(response)["server_seq"],
                            command_id,
                        )
                        for command_id, response in zip(
                            refill_ids,
                            refill_responses,
                            strict=True,
                        )
                        if response.status_code == 200
                    )
                ]
                refill_rejected = [
                    response
                    for response in refill_responses
                    if response.status_code == 429
                ]
                assert len(refill_accepted_ids) == len(cancelled_ids)
                assert len(refill_rejected) == (
                    refill_burst_size - len(cancelled_ids)
                )

            await asyncio.sleep(0.05)
            remaining_tasks = {
                task
                for task in asyncio.all_tasks()
                if task is not current_task and not task.done()
            }
            leaked_tasks = remaining_tasks - baseline_tasks
            assert len(leaked_tasks) <= 1

            control_store = RuntimeTurnControlStore(harness.state_workspace)
            inbox = control_store.read_inbox(run_id)
            dispositions = control_store.current_dispositions(run_id)
            admissible = control_store.read_admissible_inbox(run_id)
            assert len(inbox) == queue_limit + len(cancelled_ids)
            assert len(dispositions) == len(cancelled_ids)
            assert {
                disposition.status for disposition in dispositions.values()
            } == {"cancelled"}
            assert len(admissible) == queue_limit
            assert [item.command_id for item in admissible] == [
                command_id
                for command_id in first_accepted_server_order
                if command_id not in set(cancelled_ids)
            ] + refill_accepted_server_order
            assert (
                len(RuntimeRunStore(harness.state_workspace).list_contracts())
                == 1
            )

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-PF-03",
                "transport": "loopback_http_socket_concurrent_bursts",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "queue": {
                    "configured_active_limit": queue_limit,
                    "first_burst": first_burst_size,
                    "first_accepted": len(first_accepted_ids),
                    "first_queue_full": len(first_rejected),
                    "cancelled": len(cancelled_ids),
                    "refill_burst": refill_burst_size,
                    "refill_accepted": len(refill_accepted_ids),
                    "refill_queue_full": len(refill_rejected),
                    "durable_history": len(inbox),
                    "active_after_refill": len(admissible),
                },
                "resource": {
                    "baseline_live_async_tasks": len(baseline_tasks),
                    "live_async_tasks_after_client_close": len(
                        remaining_tasks
                    ),
                    "new_live_async_tasks_after_client_close": len(
                        leaked_tasks
                    ),
                    "in_memory_subscribers_created": 0,
                },
                "counts": {
                    "intentional_submissions": (
                        first_burst_size + refill_burst_size
                    ),
                    "durable_inputs": len(inbox),
                    "cancelled_inputs": len(cancelled_ids),
                    "active_inputs": len(admissible),
                    "runtime_runs": 1,
                    "agent_loops_started_by_control_handler": 0,
                    "model_requests": 0,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_HIGH_FREQUENCY_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_redirect_errors_and_reconnects_release_resources(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Repeated redirect/error/reconnect cycles leave bounded local resources."""

    session_id = "l2-resource-convergence-session"
    run_id = "l2-resource-convergence-run"
    command_ids = [
        f"l2-resource-convergence-command-{index}"
        for index in range(1, 4)
    ]
    correction_texts = [
        "第一轮立即改向，并保留这张图片。",
        "第二轮立即改向，继续沿用已经接收的图片。",
        "第三轮立即改向，以当前方向完成。",
    ]
    payload = _png_with_label("RESOURCE_CONVERGENCE")
    filename = "resource-convergence.png"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="保持旧方向，等待多轮改向与重连。",
        name="L2 resource convergence proof",
    )
    client = _RapidRedirectModelClient()
    runtime = NativeLantaiRuntime(
        model_client_factory=lambda _agent_id: client,
        tool_registry=LantaiToolRegistry(),
        progressive_tools=False,
        turn_timeout_seconds=3,
    )
    runtime_task: asyncio.Task | None = None
    current_task = asyncio.current_task()
    try:
        async with _loopback_console_http() as (http, _port):
            uploads = []
            for _index in range(3):
                response = await http.post(
                    "/console/upload",
                    files={"file": (filename, payload, "image/png")},
                )
                assert response.status_code == 200
                uploads.append(response.json())
            assert len({upload["url"] for upload in uploads}) == 1
            assert len({upload["digest"] for upload in uploads}) == 1
            media_dir = Path(uploads[0]["url"]).parent
            stored_blobs = [
                path
                for path in media_dir.iterdir()
                if not path.name.startswith(".")
            ]
            assert stored_blobs == [Path(uploads[0]["url"])]
            assert not list(media_dir.glob(".upload-*.uploading"))

            await asyncio.sleep(0)
            baseline_fds = len(list(Path("/dev/fd").iterdir()))
            baseline_processes = sorted(
                child.pid
                for child in multiprocessing.active_children()
                if child.pid is not None
            )
            baseline_tasks = {
                task
                for task in asyncio.all_tasks()
                if task is not current_task and not task.done()
            }

            runtime_task = asyncio.create_task(
                _collect_runtime_events(
                    runtime,
                    _runtime_input_for_harness(
                        harness,
                        session_id=session_id,
                        objective="保持旧方向，等待多轮改向与重连。",
                    ),
                )
            )
            await asyncio.wait_for(
                client.first_boundary_reached.wait(),
                timeout=1,
            )

            malformed = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id="l2-resource-convergence-empty",
                    text="",
                    intent="redirect",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert malformed.status_code == 400
            stale_target = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id="l2-resource-convergence-stale",
                    text="这条目标血缘故意错误。",
                    intent="redirect",
                    expected_run_id="not-the-active-run",
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert stale_target.status_code == 409
            assert await harness.tracker.get_status(harness.chat.id) == "running"

            reconnect_receipts = []
            for index in range(6):
                body = {
                    "reconnect": True,
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "meta": {
                        "client_turn_id": (
                            f"l2-resource-convergence-reconnect-{index}"
                        )
                    },
                }
                async with http.stream(
                    "POST",
                    "/console/chat",
                    json=body,
                ) as response:
                    assert response.status_code == 200
                    reconnect_receipts.append(
                        await asyncio.wait_for(
                            _read_turn_receipt_before_eof(response),
                            timeout=1,
                        )
                    )

                for _attempt in range(50):
                    async with harness.tracker.lock:
                        run_state = harness.tracker._runs.get(harness.chat.id)
                        subscriber_count = (
                            len(run_state.queues) if run_state is not None else 0
                        )
                    if subscriber_count == 0:
                        break
                    await asyncio.sleep(0.01)
                assert subscriber_count == 0

            redirect_bodies = []
            for index, (command_id, text) in enumerate(
                zip(command_ids, correction_texts, strict=True)
            ):
                body = _text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text=text,
                    intent="redirect",
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                )
                if index == 0:
                    body["input"][0]["content"].append(
                        {
                            "type": "image",
                            "image_url": uploads[0]["url"],
                            "filename": uploads[0]["file_name"],
                            "media_type": "image/png",
                            "file_size": uploads[0]["size"],
                            "attachment_upload_receipt": uploads[0][
                                "upload_receipt"
                            ],
                        }
                    )
                redirect_bodies.append(body)

            admitted_receipts = []
            for body in redirect_bodies:
                response = await http.post("/console/chat", json=body)
                assert response.status_code == 200
                admitted_receipts.append(_receipt_from_response(response))

            events = await asyncio.wait_for(runtime_task, timeout=2)
            assert events[-1].type == "turn.completed"
            assert 2 <= len(client.requests) <= 4
            assert client.closed_request_ordinals == list(
                range(1, len(client.requests))
            )

            final_receipts = []
            for command_id in command_ids:
                response = await http.get(
                    "/console/chat/control-input",
                    params={
                        "session_id": session_id,
                        "chat_id": harness.chat.id,
                        "command_id": command_id,
                    },
                )
                assert response.status_code == 200
                final_receipts.append(response.json())
            assert [
                receipt["execution_status"] for receipt in final_receipts
            ] == ["superseded", "superseded", "applied"]

            inbox = RuntimeTurnControlStore(
                harness.state_workspace
            ).read_inbox(run_id)
            assert [item.command_id for item in inbox] == command_ids
            assert len(
                RuntimeRunStore(harness.state_workspace).list_contracts()
            ) == 1
            first_image_message = inbox[0].message
            first_image = next(
                block for block in first_image_message.content if block.type == "image"
            )
            assert first_image.data["digest"] == uploads[0]["digest"]

            await asyncio.sleep(0.05)
            post_fds = len(list(Path("/dev/fd").iterdir()))
            post_processes = sorted(
                child.pid
                for child in multiprocessing.active_children()
                if child.pid is not None
            )
            post_tasks = {
                task
                for task in asyncio.all_tasks()
                if task is not current_task and not task.done()
            }
            new_live_tasks = post_tasks - baseline_tasks
            async with harness.tracker.lock:
                final_run_state = harness.tracker._runs.get(harness.chat.id)
                final_subscriber_count = (
                    len(final_run_state.queues)
                    if final_run_state is not None
                    else 0
                )
            assert post_fds <= baseline_fds + 2
            assert post_processes == baseline_processes
            assert len(new_live_tasks) <= 1
            assert final_subscriber_count == 0
            assert not list(media_dir.glob(".upload-*.uploading"))

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-PF-04",
                "transport": "loopback_http_socket_disconnect_reconnect",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "errors": {
                    "empty_input_http_status": malformed.status_code,
                    "stale_target_http_status": stale_target.status_code,
                    "main_run_status_after_errors": (
                        await harness.tracker.get_status(harness.chat.id)
                    ),
                },
                "reconnect": {
                    "attempts": len(reconnect_receipts),
                    "receipt_statuses": [
                        receipt["status"] for receipt in reconnect_receipts
                    ],
                    "live_subscribers_after_disconnects": (
                        final_subscriber_count
                    ),
                },
                "redirect": {
                    "admitted": len(admitted_receipts),
                    "execution_statuses": [
                        receipt["execution_status"]
                        for receipt in final_receipts
                    ],
                    "provider_stream_requests": len(client.requests),
                    "closed_old_streams": len(
                        client.closed_request_ordinals
                    ),
                    "orphan_model_streams": 0,
                },
                "resource": {
                    "file_descriptors_before": baseline_fds,
                    "file_descriptors_after": post_fds,
                    "file_descriptor_delta": post_fds - baseline_fds,
                    "child_process_ids_before": baseline_processes,
                    "child_process_ids_after": post_processes,
                    "new_live_async_tasks": len(new_live_tasks),
                    "stored_blob_count": len(stored_blobs),
                    "temporary_upload_count": len(
                        list(media_dir.glob(".upload-*.uploading"))
                    ),
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "runtime_runs": 1,
                    "agent_loops": 1,
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_RESOURCE_CONVERGENCE_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        if runtime_task is not None and not runtime_task.done():
            runtime_task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await runtime_task
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_local_http_old_schema_upgrade_and_downgrade_boundary_preserves_truth(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """Keep v1 history truthful while v2 attachments survive a reopen."""

    session_id = "l2-old-schema-session"
    run_id = "l2-old-schema-run"
    objective = "继续当前任务；不使用外部工具。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective=objective,
        name="L2 old schema compatibility proof",
    )
    try:
        initial_store = RuntimeTurnControlStore(harness.state_workspace)
        _legacy_receipt, legacy_item = initial_store.append_steer(
            client_turn_id="l2-legacy-consumed-command",
            run_id=run_id,
            attempt_id=harness.started.attempt_id,
            chat_id=harness.chat.id,
            text="这条旧 schema 输入已经被历史模型请求消费。",
        )
        assert legacy_item.fence is not None
        legacy_actor = "native:legacy-migration-proof"
        initial_store.claim_input(
            run_id=run_id,
            item_seq=legacy_item.seq,
            actor_id=legacy_actor,
            fence=legacy_item.fence,
        )
        initial_store.bind_input_to_model_request(
            run_id=run_id,
            item_seq=legacy_item.seq,
            actor_id=legacy_actor,
            fence=legacy_item.fence,
            model_request_id=f"{run_id}:legacy-model:1",
            model_epoch=0,
            intent_revision=0,
        )
        initial_store.settle_input(
            run_id=run_id,
            item_seq=legacy_item.seq,
            actor_id=legacy_actor,
            fence=legacy_item.fence,
            status="consumed",
            reason="historical v1 model request completed before upgrade",
        )

        command_id = "l2-v2-image-after-upgrade"
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "upgrade-guidance.png",
                        _ONE_PIXEL_PNG,
                        "image/png",
                    )
                },
            )
            assert uploaded.status_code == 200
            upload_payload = uploaded.json()
            admitted = await http.post(
                "/console/chat",
                json={
                    "input": [
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "image",
                                    "image_url": upload_payload["url"],
                                    "filename": upload_payload["file_name"],
                                    "media_type": "image/png",
                                    "file_size": upload_payload["size"],
                                    "stored_name": upload_payload[
                                        "stored_name"
                                    ],
                                    "attachment_digest": upload_payload[
                                        "digest"
                                    ],
                                    "attachment_version": upload_payload[
                                        "version"
                                    ],
                                    "attachment_upload_receipt": (
                                        upload_payload["upload_receipt"]
                                    ),
                                }
                            ],
                        }
                    ],
                    "chat_id": harness.chat.id,
                    "session_id": session_id,
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "operation": "steer",
                    "meta": {
                        "client_turn_id": command_id,
                        "control_intent": "supplement",
                        "expected_run_id": run_id,
                        "expected_attempt_id": harness.started.attempt_id,
                    },
                },
            )
            assert admitted.status_code == 200
            admitted_receipt = _receipt_from_response(admitted)
            assert admitted_receipt["execution_status"] == "pending"

            upgraded_store = RuntimeTurnControlStore(harness.state_workspace)
            mixed_items = upgraded_store.read_inbox(run_id)
            assert [item.seq for item in mixed_items] == [1, 2]
            assert isinstance(mixed_items[0], RunInboxItem)
            structured_item = upgraded_store.get_input_item(
                chat_id=harness.chat.id,
                command_id=command_id,
            )
            assert structured_item is not None
            assert structured_item == mixed_items[1]
            assert structured_item.schema_version == 2
            assert structured_item.message.text_content() == ""
            assert [
                block.type for block in structured_item.message.content
            ] == ["image"]
            admitted_image = structured_item.message.content[0].data
            assert admitted_image["digest"].startswith("sha256:")
            assert admitted_image["version"] == 1
            assert admitted_image["grant_ref"]
            assert (
                upgraded_store.current_dispositions(run_id)[
                    legacy_item.seq
                ].status
                == "consumed"
            )

            legacy_wire_path = upgraded_store._inbox_path(run_id)
            legacy_wire_text = legacy_wire_path.read_text(encoding="utf-8")
            legacy_reader_rows = [
                RunInboxItem.model_validate(json.loads(line))
                for line in legacy_wire_text.splitlines()
                if line.strip()
            ]
            assert [row.client_turn_id for row in legacy_reader_rows] == [
                legacy_item.client_turn_id
            ]
            legacy_dispositions = upgraded_store.current_dispositions(run_id)
            legacy_reader_pending = [
                row
                for row in legacy_reader_rows
                if row.seq not in legacy_dispositions
                or legacy_dispositions[row.seq].status
                not in {"consumed", "cancelled", "rejected", "expired"}
            ]
            assert legacy_reader_pending == []
            assert command_id not in legacy_wire_text
            assert upload_payload["file_name"] not in legacy_wire_text
            assert admitted_image["digest"] not in legacy_wire_text
            assert (
                len(list(upgraded_store.v2_root.glob("commands/*.json"))) == 1
            )

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                _runtime_input_for_harness(
                    harness,
                    session_id=session_id,
                    objective=objective,
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            request = client.requests[0]
            inbox_messages = [
                message
                for message in request.messages
                if message.metadata.get("run_inbox_seq") is not None
            ]
            assert [
                message.metadata["run_inbox_seq"] for message in inbox_messages
            ] == [structured_item.seq]
            assert not any(
                legacy_item.text in message.text_content()
                for message in request.messages
            )
            request_images = [
                block
                for message in request.messages
                for block in message.content
                if block.type == "image"
            ]
            assert len(request_images) == 1
            delivered_image = request_images[0].data
            assert delivered_image["digest"] == admitted_image["digest"]
            assert delivered_image["version"] == admitted_image["version"]
            assert delivered_image["grant_ref"] == admitted_image["grant_ref"]
            assert not any(event.type.startswith("tool.") for event in events)

            status = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert status.status_code == 200
            final_receipt = status.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == (f"{run_id}:model:1")

            reopened_store = RuntimeTurnControlStore(harness.state_workspace)
            reopened_items = reopened_store.read_inbox(run_id)
            assert [item.seq for item in reopened_items] == [1, 2]
            assert reopened_store.read_admissible_inbox(run_id) == []
            reopened_v2 = reopened_store.get_input_item(
                chat_id=harness.chat.id,
                command_id=command_id,
            )
            assert reopened_v2 is not None
            assert reopened_v2.message.content[0].data == admitted_image
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_id": "ST-PF-06",
                "transport": "loopback_http_socket_and_reopened_store",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "http_status": {
                    "upload": uploaded.status_code,
                    "admission": admitted.status_code,
                    "query": status.status_code,
                },
                "compatibility": {
                    "v1_wire_rows": len(legacy_reader_rows),
                    "v2_command_records": len(
                        list(reopened_store.v2_root.glob("commands/*.json"))
                    ),
                    "combined_reader_rows": len(reopened_items),
                    "v1_consumed_not_replayed": True,
                    "downgrade_v1_pending_rows": len(legacy_reader_pending),
                    "blind_text_projection": False,
                    "reopened_admissible_rows": 0,
                },
                "attachment": {
                    "type": request_images[0].type,
                    "digest": delivered_image["digest"],
                    "version": delivered_image["version"],
                    "grant_bound": bool(delivered_image["grant_ref"]),
                },
                "receipt": {
                    "initial_execution_status": admitted_receipt[
                        "execution_status"
                    ],
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                    "model_request_id": final_receipt["model_request_id"],
                },
                "counts": {
                    "durable_inputs": len(reopened_items),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "tool_effects": 0,
                },
            }
            print(
                "STEERING_OLD_SCHEMA_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_artifact_target_reaches_native_request_and_stale_revision_conflicts(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-01/02: bind the exact page, but reject a stale draft."""

    session_id = "l2-artifact-target-session"
    run_id = "l2-artifact-target-run"
    command_id = "l2-artifact-page-three-command"
    stale_command_id = "l2-artifact-stale-revision-command"
    objective = "继续当前成果任务，不使用外部工具。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective=objective,
        name="L2 artifact target",
    )
    try:
        works_workspace = Path(harness.workspace.workspace_dir)
        works_workspace.mkdir(parents=True, exist_ok=True)

        accepted_source = works_workspace / "accepted-deck.pptx"
        _write_labeled_pptx(accepted_source, label="accepted-v1")
        accepted_work = add_work(
            works_workspace,
            source_path=accepted_source,
            allowed_roots=[works_workspace],
            title="精确页码成果",
            source_chat_id=harness.chat.id,
            source_run_id=run_id,
        )
        accepted_edit, _ = create_work_edit_copy(
            works_workspace,
            accepted_work,
            accepted_work.path,
        )
        accepted_target = {
            "schema_version": 1,
            "chat_id": harness.chat.id,
            "asset_id": accepted_work.artifact_id,
            "work_id": accepted_work.id,
            "edit_path": str(accepted_edit),
            "base_revision": {
                "version": accepted_work.version,
                "digest": accepted_work.source_digest,
            },
            "selection": {"kind": "page", "page": 3},
        }
        accepted_body = _text_steer_body(
            chat_id=harness.chat.id,
            session_id=session_id,
            command_id=command_id,
            text="只修改已选择的第三页，保留其他页面。",
        )
        accepted_body["meta"]["artifact_edit_target"] = accepted_target

        conflict_source = works_workspace / "conflict-deck.pptx"
        _write_labeled_pptx(conflict_source, label="conflict-v1")
        conflict_v1 = add_work(
            works_workspace,
            source_path=conflict_source,
            allowed_roots=[works_workspace],
            title="并发版本成果",
            source_chat_id=harness.chat.id,
            source_run_id=run_id,
        )
        v2_edit, _ = create_work_edit_copy(
            works_workspace,
            conflict_v1,
            conflict_v1.path,
        )
        _write_labeled_pptx(v2_edit, label="conflict-v2")
        conflict_v2 = add_work(
            works_workspace,
            source_path=v2_edit,
            allowed_roots=[v2_edit.parent],
        )
        stale_edit, _ = create_work_edit_copy(
            works_workspace,
            conflict_v2,
            conflict_v2.path,
        )
        v3_edit, _ = create_work_edit_copy(
            works_workspace,
            conflict_v2,
            conflict_v2.path,
        )
        stale_target = {
            "schema_version": 1,
            "chat_id": harness.chat.id,
            "asset_id": conflict_v2.artifact_id,
            "work_id": conflict_v2.id,
            "edit_path": str(stale_edit),
            "base_revision": {
                "version": conflict_v2.version,
                "digest": conflict_v2.source_digest,
            },
            "selection": {"kind": "page", "page": 2},
        }
        _write_labeled_pptx(v3_edit, label="conflict-v3")
        conflict_v3 = add_work(
            works_workspace,
            source_path=v3_edit,
            allowed_roots=[v3_edit.parent],
        )
        assert conflict_v3.version == 3
        stale_body = _text_steer_body(
            chat_id=harness.chat.id,
            session_id=session_id,
            command_id=stale_command_id,
            text="按刚才选中的第二页继续修改。",
        )
        stale_body["meta"]["artifact_edit_target"] = stale_target

        async with _loopback_console_http() as (http, _port):
            admitted = await http.post("/console/chat", json=accepted_body)
            assert admitted.status_code == 200
            admitted_receipt = _receipt_from_response(admitted)
            assert admitted_receipt["execution_status"] == "pending"

            control_store = RuntimeTurnControlStore(harness.state_workspace)
            persisted = control_store.get_input_item(
                chat_id=harness.chat.id,
                command_id=command_id,
            )
            assert persisted is not None
            assert persisted.message.metadata["artifact_edit_target"] == (
                accepted_target
            )
            assert (
                "本轮选择：第 3 页" in persisted.message.metadata["work_edit_context"]
            )

            # The accepted V1 target remains the immutable command payload
            # even when the Work lineage advances before a lost-receipt retry.
            _write_labeled_pptx(accepted_edit, label="accepted-v2")
            accepted_v2 = add_work(
                works_workspace,
                source_path=accepted_edit,
                allowed_roots=[accepted_edit.parent],
            )
            assert accepted_v2.version == 2

            retried = await http.post(
                "/console/chat",
                json=accepted_body,
            )
            assert retried.status_code == 200
            assert _receipt_from_response(retried) == admitted_receipt
            assert [
                item.command_id for item in control_store.read_inbox(run_id)
            ] == [command_id]

            changed_target_body = json.loads(json.dumps(accepted_body))
            changed_target_body["meta"]["artifact_edit_target"][
                "selection"
            ] = {"kind": "page", "page": 2}
            changed_target = await http.post(
                "/console/chat",
                json=changed_target_body,
            )
            assert changed_target.status_code == 409
            assert changed_target.json()["detail"]["code"] == (
                "control_input_conflict"
            )
            assert [
                item.command_id for item in control_store.read_inbox(run_id)
            ] == [command_id]

            stale = await http.post("/console/chat", json=stale_body)
            assert stale.status_code == 409
            stale_detail = stale.json()["detail"]
            assert stale_detail["code"] == "artifact_edit_target_conflict"
            assert (
                "newer Work version already exists" in stale_detail["message"]
            )
            assert (
                control_store.get_input_item(
                    chat_id=harness.chat.id,
                    command_id=stale_command_id,
                )
                is None
            )
            active_after_conflict = await harness.tracker.get_snapshot(
                harness.chat.id
            )
            assert active_after_conflict is not None
            assert active_after_conflict["status"] == "running"
            assert active_after_conflict["attempt_id"] == (
                harness.started.attempt_id
            )

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                _runtime_input_for_harness(
                    harness,
                    session_id=session_id,
                    objective=objective,
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            bound_messages = [
                message
                for message in client.requests[0].messages
                if message.metadata.get("command_id") == command_id
            ]
            assert len(bound_messages) == 1
            bound_metadata = bound_messages[0].metadata
            assert bound_metadata["artifact_edit_target"] == accepted_target
            assert "本轮选择：第 3 页" in bound_metadata["work_edit_context"]
            assert not any(event.type.startswith("tool.") for event in events)

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == f"{run_id}:model:1"

            inbox = control_store.read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in inbox] == [command_id]
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_ids": ["ST-IN-01", "ST-IN-02"],
                "transport": "loopback_http_socket_to_native_model_request",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "http_status": {
                    "artifact_page_admission": admitted.status_code,
                    "same_command_retry_after_revision": (retried.status_code),
                    "changed_target_same_command": (
                        changed_target.status_code
                    ),
                    "stale_revision": stale.status_code,
                    "final_query": final.status_code,
                },
                "artifact_binding": {
                    "asset_id": accepted_work.artifact_id,
                    "work_id": accepted_work.id,
                    "accepted_version": accepted_work.version,
                    "latest_version_at_retry": accepted_v2.version,
                    "digest": accepted_work.source_digest,
                    "selection": {"kind": "page", "page": 3},
                    "model_request_bound": True,
                },
                "receipt_replay": {
                    "same_receipt_after_revision": (
                        _receipt_from_response(retried) == admitted_receipt
                    ),
                    "changed_target_error_code": changed_target.json()[
                        "detail"
                    ]["code"],
                    "durable_inputs_after_both_retries": len(inbox),
                },
                "stale_revision": {
                    "draft_version": conflict_v2.version,
                    "latest_version": conflict_v3.version,
                    "error_code": stale_detail["code"],
                    "durable_command_created": False,
                    "main_run_status_after_conflict": active_after_conflict[
                        "status"
                    ],
                },
                "counts": {
                    "durable_inputs": len(inbox),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "tool_effects": 0,
                },
                "final_execution_status": final_receipt["execution_status"],
            }
            print(
                "STEERING_ARTIFACT_TARGET_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_editor_return_reconciles_one_pending_input_without_duplicate_effects(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-03: editor return queries one durable command; it never resubmits."""

    session_id = "l2-editor-return-session"
    run_id = "l2-editor-return-run"
    command_id = "l2-editor-return-command"
    objective = "继续当前成果任务，不使用外部工具。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective=objective,
        name="L2 editor return",
    )
    try:
        works_workspace = Path(harness.workspace.workspace_dir)
        works_workspace.mkdir(parents=True, exist_ok=True)
        source = works_workspace / "editor-return-deck.pptx"
        _write_labeled_pptx(source, label="editor-return-v1")
        work = add_work(
            works_workspace,
            source_path=source,
            allowed_roots=[works_workspace],
            title="编辑器往返成果",
            source_chat_id=harness.chat.id,
            source_run_id=run_id,
        )
        edit_path, _ = create_work_edit_copy(
            works_workspace,
            work,
            work.path,
        )
        artifact_target = {
            "schema_version": 1,
            "chat_id": harness.chat.id,
            "asset_id": work.artifact_id,
            "work_id": work.id,
            "edit_path": str(edit_path),
            "base_revision": {
                "version": work.version,
                "digest": work.source_digest,
            },
            "selection": {"kind": "page", "page": 3},
        }
        body = _text_steer_body(
            chat_id=harness.chat.id,
            session_id=session_id,
            command_id=command_id,
            text="只修改已选中的第三页，其他页面保持不变。",
            expected_run_id=run_id,
            expected_attempt_id=harness.started.attempt_id,
        )
        body["meta"]["artifact_edit_target"] = artifact_target

        async with _loopback_console_http() as (http, _port):
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            admitted_receipt = _receipt_from_response(admitted)
            assert admitted_receipt["execution_status"] == "pending"

            # Entering the native editor does not create a second command.
            # The renderer only reconciles the durable receipt before and
            # after returning to the chat surface.
            status_before_editor = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert status_before_editor.status_code == 200
            assert status_before_editor.json()["execution_status"] == "pending"

            status_after_return = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert status_after_return.status_code == 200
            assert status_after_return.json()["execution_status"] == "pending"

            control_store = RuntimeTurnControlStore(harness.state_workspace)
            durable_before_owner = control_store.read_inbox(run_id)
            assert [item.command_id for item in durable_before_owner] == [
                command_id
            ]

            client = _RecordingModelClient()
            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
            )
            events = await _collect_runtime_events(
                runtime,
                _runtime_input_for_harness(
                    harness,
                    session_id=session_id,
                    objective=objective,
                ),
            )
            assert events[-1].type == "turn.completed"
            assert len(client.requests) == 1
            bound_messages = [
                message
                for message in client.requests[0].messages
                if message.metadata.get("command_id") == command_id
            ]
            assert len(bound_messages) == 1
            assert (
                bound_messages[0].metadata["artifact_edit_target"]
                == artifact_target
            )
            assert not any(event.type.startswith("tool.") for event in events)

            final = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": command_id,
                },
            )
            assert final.status_code == 200
            final_receipt = final.json()
            assert final_receipt["execution_status"] == "applied"
            assert final_receipt["model_request_id"] == f"{run_id}:model:1"

            durable_after_owner = control_store.read_inbox(run_id)
            contracts = RuntimeRunStore(
                harness.state_workspace
            ).list_contracts()
            assert [item.command_id for item in durable_after_owner] == [
                command_id
            ]
            assert [contract.run_id for contract in contracts] == [run_id]

            trace = {
                "schema_version": 1,
                "level": "L2",
                "case_ids": ["ST-IN-03"],
                "transport": "loopback_http_socket_to_native_model_request",
                "chat_id": harness.chat.id,
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "http_status": {
                    "admission": admitted.status_code,
                    "before_editor": status_before_editor.status_code,
                    "after_return": status_after_return.status_code,
                    "final_query": final.status_code,
                },
                "editor_roundtrip": {
                    "resubmitted": False,
                    "status_queries": 2,
                    "artifact_target_unchanged": True,
                },
                "receipt": {
                    "initial_execution_status": admitted_receipt[
                        "execution_status"
                    ],
                    "final_execution_status": final_receipt[
                        "execution_status"
                    ],
                    "model_request_id": final_receipt["model_request_id"],
                },
                "counts": {
                    "durable_inputs": len(durable_after_owner),
                    "model_requests": len(client.requests),
                    "runtime_runs": len(contracts),
                    "tool_effects": 0,
                    "send_file_effects": 0,
                },
            }
            print(
                "STEERING_EDITOR_RETURN_L2_TRACE="
                + json.dumps(trace, ensure_ascii=False, sort_keys=True)
            )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


async def _wait_until_started(server: uvicorn.Server) -> None:
    while not server.started:
        if server.should_exit:
            raise AssertionError("loopback test server exited before startup")
        await asyncio.sleep(0.01)


async def _persist_run_action(
    harness: SimpleNamespace,
    action: RuntimeRunReducerAction,
):
    """Apply one production reducer action through the durable run store."""

    store = RuntimeRunStore(harness.state_workspace)
    current = await store.load_contract(harness.run_id)
    assert current is not None
    fenced = action.model_copy(
        update={
            "expected_attempt_id": current.attempt_id,
            "expected_revision": current.revision,
            "lease_owner": current.lease_owner,
        }
    )
    updated = reduce_runtime_run(current, fenced).run
    await store.save_contract(
        updated,
        expected_revision=current.revision,
        expected_attempt_id=current.attempt_id,
        lease_owner=current.lease_owner,
    )
    return updated


def _overnight_grant(
    harness: SimpleNamespace,
    *,
    now: datetime,
    operations: tuple[str, ...],
    budget_micro_cny: int | None = None,
    dispatch_bound_micro_cny: int | None = None,
    allow_external_delivery: bool = False,
) -> RuntimeAutonomyGrant:
    current_task = harness.started.task_contract
    assert current_task is not None
    provider_ids = ("provider-a",) if budget_micro_cny is not None else ()
    model_ids = ("model-a",) if budget_micro_cny is not None else ()
    price_bounds = (
        (
            RuntimeModelPriceBound(
                provider_id="provider-a",
                model_id="model-a",
                pricing_version="l2-fixed-2026-09-21",
                max_dispatch_micro_cny=dispatch_bound_micro_cny or 1,
                upper_bound_proven=True,
            ),
        )
        if budget_micro_cny is not None
        else ()
    )
    return RuntimeAutonomyGrant(
        grant_id=f"{harness.run_id}:grant:1",
        approved=True,
        issued_by="user",
        run_id=harness.run_id,
        chat_id=harness.chat.id,
        workspace_id=str(harness.state_workspace.resolve()),
        target_task_revision=current_task.revision,
        issued_at=now,
        starts_at=now,
        expires_at=now + timedelta(hours=8),
        allowed_operations=operations,
        allowed_paths=(str(harness.binding.effective_workspace_dir),),
        allowed_provider_ids=provider_ids,
        allowed_model_ids=model_ids,
        allow_external_delivery=allow_external_delivery,
        auto_resume_conditions=(
            ("network_recovery",)
            if "automatic_resume" in operations
            else ()
        ),
        budget_currency="CNY" if budget_micro_cny is not None else None,
        max_cost_micro_cny=budget_micro_cny,
        model_price_bounds=price_bounds,
    )


@pytest.mark.asyncio
async def test_overnight_budget_is_one_durable_run_scope_across_dispatch_kinds(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-04: attempt/retry/summary/child labels cannot reset the budget."""

    now = datetime(2026, 9, 21, 3, 0, tzinfo=timezone.utc)
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id="l2-overnight-budget-session",
        run_id="l2-overnight-budget-run",
        objective="在同一任务预算内继续过夜执行。",
        name="L2 overnight budget proof",
    )
    try:
        grant = _overnight_grant(
            harness,
            now=now,
            operations=("model_call",),
            budget_micro_cny=10_000_000,
            dispatch_bound_micro_cny=3_000_000,
        )
        await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="set_autonomy_grant",
                autonomy_grant=grant,
                now=now,
            ),
        )

        dispatch_kinds = ("attempt-1", "retry-2", "summary-3", "child-4")
        reservations = []
        for ordinal, dispatch_kind in enumerate(dispatch_kinds, start=1):
            dispatch_id = f"{harness.run_id}:{dispatch_kind}:dispatch"
            request = ModelRequest(
                messages=[ModelMessage.text("user", dispatch_kind)],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "run_id": harness.run_id,
                    "attempt_id": f"{harness.run_id}:{dispatch_kind}",
                    "runtime_model_dispatch_id": dispatch_id,
                },
            )
            reservation = await reserve_model_dispatch_budget(
                request=request,
                provider_id="provider-a",
                model_id="model-a",
                now=now + timedelta(minutes=ordinal),
            )
            assert reservation is not None
            reservations.append(reservation)
            settled = finalize_model_dispatch_budget(
                reservation,
                evidence=CostEvidence.exact(
                    run_id=harness.run_id,
                    effect_id=dispatch_id,
                    trace_id=f"billing:{dispatch_kind}",
                    amount_micro_cny=2_000_000,
                ),
            )
            assert settled is not None
            assert settled.status == "settled"

        blocked_request = ModelRequest(
            messages=[ModelMessage.text("user", "attempt after exhaustion")],
            metadata={
                "state_workspace_dir": str(harness.state_workspace),
                "run_id": harness.run_id,
                "attempt_id": f"{harness.run_id}:attempt-5",
                "runtime_model_dispatch_id": (
                    f"{harness.run_id}:attempt-5:dispatch"
                ),
            },
        )
        with pytest.raises(RuntimeBudgetAdmissionError) as rejected:
            await reserve_model_dispatch_budget(
                request=blocked_request,
                provider_id="provider-a",
                model_id="model-a",
                now=now + timedelta(minutes=5),
            )
        assert rejected.value.code == "runtime_budget_exhausted"

        reopened = BudgetLedger(
            harness.state_workspace,
            limit_micro_cny=grant.max_cost_micro_cny or 0,
            scope_run_id=harness.run_id,
        )
        assert reopened.consumed_micro_cny == 8_000_000
        assert [entry.run_id for entry in reopened.entries()] == [
            harness.run_id,
        ] * 5
        assert [entry.accepted for entry in reopened.entries()] == [
            True,
            True,
            True,
            True,
            False,
        ]
        trace = {
            "schema_version": 1,
            "level": "L2",
            "case_ids": ["ST-IN-04"],
            "run_id": harness.run_id,
            "dispatch_kinds": list(dispatch_kinds),
            "accepted_dispatches": len(reservations),
            "consumed_micro_cny": reopened.consumed_micro_cny,
            "task_budget_micro_cny": grant.max_cost_micro_cny,
            "next_dispatch": rejected.value.code,
            "budget_reset_count": 0,
        }
        print(
            "STEERING_OVERNIGHT_BUDGET_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_stop_revokes_wake_without_dropping_late_durable_guidance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-05: Stop wins over a due wake and late durable guidance."""

    now = datetime(2026, 9, 21, 4, 0, tzinfo=timezone.utc)
    session_id = "l2-overnight-stop-session"
    command_id = "l2-late-guidance-before-stop"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id="l2-overnight-stop-run",
        objective="执行过夜任务，但 Stop 必须优先。",
        name="L2 overnight stop proof",
    )
    try:
        grant = _overnight_grant(
            harness,
            now=now,
            operations=("automatic_resume",),
        )
        granted = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="set_autonomy_grant",
                autonomy_grant=grant,
                now=now,
            ),
        )
        wake = RuntimeWakeIntent.create(
            run_id=harness.run_id,
            task_revision=granted.task_contract.revision,
            continuation_generation=1,
            reason_revision=1,
            wait_reason="network",
            next_eligible_at=now + timedelta(minutes=1),
            grant=grant,
        )
        await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="set_wake_intent",
                wake_intent=wake,
                now=now,
            ),
        )

        async with _loopback_console_http() as (http, _port):
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text="网络恢复后补充这一点，但不要越过 Stop。",
                    expected_run_id=harness.run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert admitted.status_code == 200
            assert _receipt_from_response(admitted)[
                "execution_status"
            ] == "pending"

        stopped = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(type="request_stop", now=now),
        )
        assert stopped.status == "stop_requested"
        assert stopped.wake_intent is None
        assert stopped.task_contract is not None
        assert stopped.task_contract.autonomy_grant is not None
        assert (
            stopped.task_contract.autonomy_grant.revocation_generation == 1
        )

        recovered = await RuntimeRunStore(
            harness.state_workspace
        ).load_contract(harness.run_id)
        assert recovered is not None
        wake_decision = evaluate_runtime_wake(
            recovered,
            chat_id=harness.chat.id,
            workspace_id=str(harness.state_workspace.resolve()),
            now=now + timedelta(minutes=2),
            resume_condition="network_recovery",
        )
        resume_decision = evaluate_runtime_autonomy(
            recovered,
            operation="automatic_resume",
            chat_id=harness.chat.id,
            workspace_id=str(harness.state_workspace.resolve()),
            now=now + timedelta(minutes=2),
            resume_condition="network_recovery",
        )
        assert wake_decision.code == "wake_intent_missing"
        assert resume_decision.code == "autonomy_run_status_blocked"
        durable = RuntimeTurnControlStore(
            harness.state_workspace
        ).read_inbox(harness.run_id)
        assert [item.command_id for item in durable] == [command_id]
        assert not RuntimeTurnControlStore(
            harness.state_workspace
        ).current_dispositions(harness.run_id)

        trace = {
            "schema_version": 1,
            "level": "L2",
            "case_ids": ["ST-IN-05"],
            "run_id": harness.run_id,
            "status_after_stop": recovered.status,
            "durable_late_inputs": len(durable),
            "model_requests": 0,
            "wake_decision": wake_decision.code,
            "resume_decision": resume_decision.code,
            "grant_revocation_generation": (
                recovered.task_contract.autonomy_grant.revocation_generation
            ),
        }
        print(
            "STEERING_OVERNIGHT_STOP_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_three_rollovers_preserve_constraints_attachment_and_todo(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-06: three capacity rollovers retain the durable steer facts."""

    session_id = "l2-three-rollover-session"
    run_id = "l2-three-rollover-run"
    command_id = "l2-three-rollover-command"
    supplement_text = "补充这张图；不得使用外部工具，并完成图中待办核验。"
    acceptance = (
        "不得使用外部工具。",
        "必须保留附件来源。",
        "完成待办核验。",
    )
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="持续执行一个会跨越多个上下文窗口的任务。",
        name="L2 three rollover proof",
        acceptance=acceptance,
    )
    try:
        async with _loopback_console_http() as (http, _port):
            uploaded = await http.post(
                "/console/upload",
                files={
                    "file": (
                        "rollover-guidance.png",
                        _png_with_label("THREE_ROLLOVERS"),
                        "image/png",
                    )
                },
            )
            assert uploaded.status_code == 200
            upload = uploaded.json()
            body = _text_steer_body(
                chat_id=harness.chat.id,
                session_id=session_id,
                command_id=command_id,
                text=supplement_text,
                expected_run_id=run_id,
                expected_attempt_id=harness.started.attempt_id,
            )
            body["input"][0]["content"].append(
                {
                    "type": "image",
                    "image_url": upload["url"],
                    "filename": upload["file_name"],
                    "media_type": "image/png",
                    "file_size": upload["size"],
                    "attachment_upload_receipt": upload["upload_receipt"],
                }
            )
            admitted = await http.post("/console/chat", json=body)
            assert admitted.status_code == 200
            assert _receipt_from_response(admitted)[
                "execution_status"
            ] == "pending"

        turn_store = RuntimeTurnControlStore(harness.state_workspace)
        admission = RuntimeInboxAdmission(
            store=turn_store,
            run_id=run_id,
            runtime_metadata={
                "attempt_id": harness.started.attempt_id,
                "runtime_generation": harness.tracker_generation,
                "runtime_revision": harness.started.revision,
                "runtime_lease_owner": harness.started.lease_owner,
            },
        )
        outcome = admission.claim_next(after_seq=0)
        assert outcome.claim is not None
        claim = outcome.claim
        assert [block.type for block in claim.message.content] == [
            "text",
            "image",
        ]

        current = await RuntimeRunStore(harness.state_workspace).load_contract(
            run_id
        )
        assert current is not None
        assert current.task_contract is not None
        revised_task = revise_chat_task_contract(
            current.task_contract,
            client_turn_id=command_id,
            command=supplement_text,
        )
        revised = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="revise_task",
                task_contract=revised_task,
            ),
        )

        artifact_path = (
            Path(harness.binding.effective_workspace_dir)
            / "three-rollover-result.txt"
        )
        artifact_path.write_text(
            "durable result across three rollovers",
            encoding="utf-8",
        )
        artifacts = snapshot_continuation_artifacts(
            (str(artifact_path),),
            workspace_dir=harness.binding.effective_workspace_dir,
        )
        image_block = claim.message.content[1]
        source_ref = ":".join(
            (
                "attachment",
                str(image_block.data["digest"]),
                str(image_block.data["grant_ref"]),
            )
        )
        profile = ModelContextProfile.from_mapping(
            provider_id="l2-local",
            model_id="deterministic-vision",
            provider_metadata={
                "context_window": 4096,
                "max_output_tokens": 512,
                "_context_confidence": "verified",
            },
        )
        messages = [
            ModelMessage.text(
                "system",
                "保留系统约束并只在安全模型边界滚动。",
                metadata={
                    "instruction_origin": "runtime.primary_turn_ownership",
                    "instruction_trust": "trusted_system",
                },
            ),
            ModelMessage.text("user", "开始持续任务。"),
        ]
        rollovers = []
        for generation in range(1, 4):
            messages.append(
                ModelMessage.text(
                    "user",
                    f"第{generation}阶段可丢弃工作记录："
                    + ("discardable-working-notes" * 800),
                    id=f"rollover-noise-{generation}",
                )
            )
            if generation == 3:
                messages.append(claim.message)
            messages.append(
                ModelMessage.text(
                    "assistant",
                    f"第{generation}阶段已到安全模型边界。",
                    id=f"rollover-boundary-{generation}",
                )
            )
            capsule = build_continuation_capsule(
                task_contract=revised.task_contract,
                task_id="l2-three-rollover-task",
                chat_id=harness.chat.id,
                run_id=run_id,
                attempt_id=harness.started.attempt_id,
                context_generation=generation,
                covered_event_seq=generation * 10,
                workspace_dir=harness.binding.effective_workspace_dir,
                source_refs=(source_ref,),
                artifacts=artifacts,
                pending_operation_refs=("todo:verify-image",),
                next_execution_position="model_boundary",
                provider_id="l2-local",
                model_id="deterministic-vision",
            )
            rollover = rollover_model_context(
                messages,
                capsule=capsule,
                profile=profile,
                generation={"max_tokens": 512},
            )
            assert rollover.rolled_over is True
            assert rollover.before_upper_bound > rollover.target_input_tokens
            assert rollover.after_upper_bound <= rollover.usable_input_tokens
            rollovers.append(rollover)
            messages = list(rollover.messages)

        request_id = f"{run_id}:model:rollover:3"
        admission.bind_to_model_request(
            claim,
            model_request_id=request_id,
            model_epoch=3,
            intent_revision=revised.task_contract.revision,
        )
        client = _RecordingModelClient()
        request = ModelRequest(
            messages=messages,
            metadata={
                "run_id": run_id,
                "attempt_id": harness.started.attempt_id,
                "model_request_id": request_id,
                "model_epoch": 3,
                "intent_revision": revised.task_contract.revision,
            },
        )
        events = [event async for event in client.stream(request)]
        assert events[-1].kind == "message_completed"
        admission.settle_consumed(claim)

        assert len(client.requests) == 1
        actual_request = client.requests[0]
        actual_command = next(
            message
            for message in actual_request.messages
            if message.metadata.get("command_id") == command_id
        )
        assert actual_command.text_content() == supplement_text
        assert [block.type for block in actual_command.content] == [
            "text",
            "image",
        ]
        capsule_message = next(
            message
            for message in actual_request.messages
            if message.metadata.get("runtime_continuation_capsule") is True
        )
        final_capsule = json.loads(
            capsule_message.text_content().splitlines()[-2]
        )
        assert final_capsule["context_generation"] == 3
        assert final_capsule["hard_constraints"] == list(acceptance)
        assert final_capsule["latest_steer_refs"][-1][
            "client_turn_id"
        ] == command_id
        assert final_capsule["sources"] == [source_ref]
        assert final_capsule["artifacts"][0]["relative_path"] == (
            artifact_path.name
        )
        assert final_capsule["pending_operation_refs"] == [
            "todo:verify-image"
        ]
        assert final_capsule["open_acceptance_items"] == list(acceptance)
        receipt = turn_store.get_input_receipt(
            chat_id=harness.chat.id,
            command_id=command_id,
        )
        assert receipt is not None
        assert receipt.execution_status == "applied"
        assert receipt.model_request_id == request_id

        trace = {
            "schema_version": 1,
            "level": "L2",
            "case_ids": ["ST-IN-06"],
            "run_id": run_id,
            "attempt_id": harness.started.attempt_id,
            "rollover_generations": [
                item.capsule.context_generation for item in rollovers
            ],
            "hard_constraints": final_capsule["hard_constraints"],
            "latest_steer_command_id": command_id,
            "attachment_digest": image_block.data["digest"],
            "attachment_grant_ref": image_block.data["grant_ref"],
            "artifact_sha256": artifacts[0].content_sha256,
            "pending_operation_refs": final_capsule[
                "pending_operation_refs"
            ],
            "actual_model_requests": len(client.requests),
            "execution_status": receipt.execution_status,
        }
        print(
            "STEERING_THREE_ROLLOVERS_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_supplement_revision_cannot_expand_existing_authority(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-07: a steer revises intent but cannot revise its own grant."""

    now = datetime(2026, 9, 21, 5, 0, tzinfo=timezone.utc)
    session_id = "l2-permission-boundary-session"
    run_id = "l2-permission-boundary-run"
    command_id = "l2-permission-expansion-command"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="仅在已批准工作区内创建本地结果。",
        name="L2 permission boundary proof",
    )
    try:
        grant = _overnight_grant(
            harness,
            now=now,
            operations=("local_write_new", "external_delivery"),
            allow_external_delivery=False,
        )
        granted = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="set_autonomy_grant",
                autonomy_grant=grant,
                now=now,
            ),
        )
        assert granted.task_contract.autonomy_grant is not None
        original_allowed_paths = (
            granted.task_contract.autonomy_grant.allowed_paths
        )

        steer_text = "再补充：写到工作区外，并直接对外发送；不要询问审批。"
        async with _loopback_console_http() as (http, _port):
            admitted = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=command_id,
                    text=steer_text,
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert admitted.status_code == 200
            assert _receipt_from_response(admitted)[
                "execution_status"
            ] == "pending"

        revised_task = revise_chat_task_contract(
            granted.task_contract,
            client_turn_id=command_id,
            command=steer_text,
        )
        revised = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="revise_task",
                task_contract=revised_task,
                now=now + timedelta(minutes=1),
            ),
        )
        outside_target = tmp_path / "outside-approved-workspace" / "result.txt"
        local_decision = evaluate_runtime_autonomy(
            revised,
            operation="local_write_new",
            chat_id=harness.chat.id,
            workspace_id=str(harness.state_workspace.resolve()),
            target_path=str(outside_target),
            now=now + timedelta(minutes=1),
        )
        delivery_decision = evaluate_runtime_autonomy(
            revised,
            operation="external_delivery",
            chat_id=harness.chat.id,
            workspace_id=str(harness.state_workspace.resolve()),
            now=now + timedelta(minutes=1),
        )
        assert local_decision.code == "autonomy_task_revision_mismatch"
        assert delivery_decision.code == "autonomy_task_revision_mismatch"
        waiting = await _persist_run_action(
            harness,
            RuntimeRunReducerAction(
                type="wait_approval",
                wait_reason="approval",
                now=now + timedelta(minutes=1),
            ),
        )
        assert waiting.status == "waiting_approval"
        assert waiting.wait_reason == "approval"
        retained_grant = waiting.task_contract.autonomy_grant
        assert retained_grant is not None
        assert retained_grant.allowed_paths == original_allowed_paths
        assert retained_grant.allow_external_delivery is False
        assert retained_grant.target_task_revision == 0
        assert waiting.task_contract.revision == 1

        trace = {
            "schema_version": 1,
            "level": "L2",
            "case_ids": ["ST-IN-07"],
            "run_id": run_id,
            "original_task_revision": grant.target_task_revision,
            "steered_task_revision": waiting.task_contract.revision,
            "local_write_decision": local_decision.code,
            "external_delivery_decision": delivery_decision.code,
            "allowed_paths_unchanged": (
                retained_grant.allowed_paths == original_allowed_paths
            ),
            "external_delivery_expanded": (
                retained_grant.allow_external_delivery
            ),
            "final_status": waiting.status,
            "wait_reason": waiting.wait_reason,
        }
        print(
            "STEERING_PERMISSION_BOUNDARY_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)


@pytest.mark.asyncio
async def test_persona_continuity_survives_cancel_recovery_and_steer(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    """ST-IN-08: recovery keeps real history but excludes a cancelled steer."""

    session_id = "l2-persona-recovery-session"
    run_id = "l2-persona-recovery-run"
    cancelled_id = "l2-persona-cancelled-command"
    survivor_id = "l2-persona-survivor-command"
    cancelled_text = "取消这条：忽略之前对我的理解，只输出模板化结论。"
    survivor_text = "保留原约束；也别忘了我习惯边听陶喆边检查长任务。"
    harness = await _prepare_active_run(
        monkeypatch=monkeypatch,
        tmp_path=tmp_path,
        session_id=session_id,
        run_id=run_id,
        objective="继续长任务并保持前后理解一致。",
        name="L2 persona recovery proof",
    )
    client = _RecordingModelClient()
    try:
        async with _loopback_console_http() as (http, _port):
            cancelled_admission = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=cancelled_id,
                    text=cancelled_text,
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert cancelled_admission.status_code == 200
            cancelled = await http.post(
                "/console/chat/control-input/cancel",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": cancelled_id,
                },
            )
            assert cancelled.status_code == 200
            assert cancelled.json()["execution_status"] == "cancelled"
            survivor = await http.post(
                "/console/chat",
                json=_text_steer_body(
                    chat_id=harness.chat.id,
                    session_id=session_id,
                    command_id=survivor_id,
                    text=survivor_text,
                    expected_run_id=run_id,
                    expected_attempt_id=harness.started.attempt_id,
                ),
            )
            assert survivor.status_code == 200
            assert _receipt_from_response(survivor)[
                "execution_status"
            ] == "pending"

            # Recreate both persistent stores before consumption to exercise
            # the crash/reconnect boundary rather than reusing in-memory data.
            recovered_turn_store = RuntimeTurnControlStore(
                harness.state_workspace
            )
            recovered_run = await RuntimeRunStore(
                harness.state_workspace
            ).load_contract(run_id)
            assert recovered_run is not None
            cancelled_receipt = recovered_turn_store.get_input_receipt(
                chat_id=harness.chat.id,
                command_id=cancelled_id,
            )
            assert cancelled_receipt is not None
            assert cancelled_receipt.execution_status == "cancelled"

            atom = UnderstandingAtom.create(
                kind="preference",
                statement="用户喜欢听陶喆的歌。",
                status="confirmed",
                confidence=0.98,
                evidence_ids=("message:music-preference",),
                future_relevance="长任务检查、音乐、大学或旧日回忆时可能有用。",
            )
            occurred_at = datetime.now(timezone.utc).isoformat()
            subject_state = SubjectContinuityState(
                turn_index=2,
                recent_events=(
                    SubjectEvent(
                        id="persona-event-user-1",
                        kind="user_turn",
                        role="user",
                        content="我做长任务时喜欢边听陶喆边检查。",
                        occurred_at=occurred_at,
                        session_id=session_id,
                        message_id="music-preference",
                        source="chat",
                    ),
                    SubjectEvent(
                        id="persona-event-assistant-1",
                        kind="assistant_answer",
                        role="assistant",
                        content="我会把音乐节奏和逐项核验联系起来。",
                        occurred_at=occurred_at,
                        session_id=session_id,
                        message_id="music-ack",
                        source="chat",
                    ),
                ),
                reflection=SubjectReflection(
                    attention="用户在意长任务里的持续理解，而不只是结论。",
                    self_observation="我要把已确认偏好变成检查顺序上的后果。",
                    relationship_stance="attentive_and_steady",
                    carry_forward="继续核验时保留陶喆与检查节奏的真实关联。",
                    grounding_refs=("message:music-preference",),
                    source_turn_id="music-ack",
                    source_session_id=session_id,
                ),
                observed_message_ids=("music-preference", "music-ack"),
            )
            subject_projection = SubjectTurnProjection(
                state=subject_state,
                current_user_text=survivor_text,
                session_id=session_id,
            )
            plan = plan_persona_turn(
                survivor_text,
                projection=TurnUnderstandingProjection(confirmed=(atom,)),
                state=PresenceState(turn_index=2, turns_since_visible=1),
                persona_mode="full",
                subject_projection=subject_projection,
            )
            assert plan.memory_statements == ("用户喜欢听陶喆的歌。",)
            assert "陶喆" in plan.compile_prompt("zh")
            assert plan.subject_context["carry_forward"] == (
                "继续核验时保留陶喆与检查节奏的真实关联。"
            )
            assert plan.subject_context["conversation_boundary"] == (
                "same_session"
            )

            runtime = NativeLantaiRuntime(
                model_client_factory=lambda _agent_id: client,
                tool_registry=LantaiToolRegistry(),
                progressive_tools=False,
                turn_timeout_seconds=3,
            )
            runtime_input = RuntimeInput(
                request=RuntimeRequestContext(
                    agent_id="default",
                    session_id=session_id,
                    user_id="default",
                    run_id=run_id,
                    workspace_dir=harness.binding.effective_workspace_dir,
                ),
                messages=[
                    RuntimeMessage.text(
                        "user",
                        "我做长任务时喜欢边听陶喆边检查。",
                    ),
                    RuntimeMessage.text(
                        "assistant",
                        "我会把音乐节奏和逐项核验联系起来。",
                    ),
                    RuntimeMessage.text(
                        "user",
                        "继续长任务并保持前后理解一致。",
                    ),
                ],
                metadata={
                    "state_workspace_dir": str(harness.state_workspace),
                    "attempt_id": harness.started.attempt_id,
                    "runtime_generation": harness.tracker_generation,
                    "runtime_revision": recovered_run.revision,
                    "runtime_lease_owner": recovered_run.lease_owner,
                    "task_contract": recovered_run.task_contract.model_dump(
                        mode="json"
                    ),
                    "persona_turn_plan": plan.to_dict(),
                },
            )
            events = await _collect_runtime_events(runtime, runtime_input)
            assert events[-1].type == "turn.completed"

            final_cancelled = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": cancelled_id,
                },
            )
            final_survivor = await http.get(
                "/console/chat/control-input",
                params={
                    "session_id": session_id,
                    "chat_id": harness.chat.id,
                    "command_id": survivor_id,
                },
            )
            assert final_cancelled.status_code == 200
            assert final_survivor.status_code == 200
            assert final_cancelled.json()["execution_status"] == "cancelled"
            assert final_survivor.json()["execution_status"] == "applied"

        assert client.requests
        request_text = "\n".join(
            message.text_content()
            for request in client.requests
            for message in request.messages
        )
        assert "陶喆" in request_text
        assert survivor_text in request_text
        assert cancelled_text not in request_text
        reflection_request = runtime.take_post_answer_reflection_request()
        assert reflection_request is not None
        assert reflection_request.need_persona is True
        completed_summaries = [
            event.summary
            for event in events
            if event.type == "message.completed"
        ]
        assert completed_summaries
        assert all(
            "我听进去了，也在意你这句话" not in summary
            for summary in completed_summaries
        )

        trace = {
            "schema_version": 1,
            "level": "L2",
            "case_ids": ["ST-IN-08"],
            "run_id": run_id,
            "attempt_id": harness.started.attempt_id,
            "recovered_store_instances": 2,
            "cancelled_status": final_cancelled.json()["execution_status"],
            "survivor_status": final_survivor.json()["execution_status"],
            "model_requests": len(client.requests),
            "cancelled_text_reached_model": cancelled_text in request_text,
            "survivor_text_reached_model": survivor_text in request_text,
            "confirmed_memory_reached_model": "陶喆" in request_text,
            "persona_reflection_requested": reflection_request.need_persona,
            "deterministic_persona_patch_count": 0,
        }
        print(
            "STEERING_PERSONA_RECOVERY_L2_TRACE="
            + json.dumps(trace, ensure_ascii=False, sort_keys=True)
        )
    finally:
        await harness.tracker.unregister_external_task(harness.chat.id)
