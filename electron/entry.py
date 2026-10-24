#!/usr/bin/env python3
"""PyInstaller entry point for lantai server."""
import importlib
import io
import os
import runpy
import sys
import tempfile
from pathlib import Path

# Ensure PyInstaller's temp directory is in path so 'lantai' can be imported
if getattr(sys, "frozen", False):
    # Running in PyInstaller bundle
    bundle_dir = Path(sys._MEIPASS)
else:
    bundle_dir = Path(__file__).resolve().parent.parent / "src"

sys.path.insert(0, str(bundle_dir))

# The process supervisor runs before CLI, provider or keychain initialization.
# Its fixed entry accepts only inherited pipe/lock descriptors.
if sys.argv[1:2] == ["--owned-process-supervisor"]:
    from lantai.utils.posix_process_supervisor import main as run_process_supervisor

    raise SystemExit(run_process_supervisor(sys.argv[2:]))

if sys.argv[1:] == ["--desktop-process-probe"]:
    os.environ["LANTAI_METADATA_ONLY"] = "1"
    from lantai.security.desktop_process_probe import main as run_process_probe

    raise SystemExit(run_process_probe())

# A fixed-operation child uses this packaged executable, not a user-installed
# Python interpreter or a model-supplied script. Route before CLI/keychain init.
if sys.argv[1:] == ["--analysis-worker"]:
    os.environ["LANTAI_ANALYSIS_WORKER_CHILD"] = "1"
    from lantai.analysis.worker import main as run_analysis_worker

    raise SystemExit(run_analysis_worker())


def _run_skill_script_bridge() -> None:
    """Execute one runtime-authorized Skill script in this Python runtime."""

    raw_path = os.environ.pop("LANTAI_RUN_SKILL_SCRIPT_PATH", "").strip()
    if not raw_path:
        return
    script_path = Path(raw_path).expanduser().resolve()
    if not script_path.is_file() or script_path.suffix.lower() != ".py":
        raise SystemExit("Invalid Lantai Skill script bridge target")
    os.environ["LANTAI_SKILL_SCRIPT_CHILD"] = "1"
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    sys.dont_write_bytecode = True
    sys.path.insert(0, str(script_path.parent))
    sys.argv = [str(script_path), *sys.argv[1:]]
    runpy.run_path(str(script_path), run_name="__main__")
    raise SystemExit(0)


_run_skill_script_bridge()


def _run_fast_metadata_command() -> None:
    """Answer metadata-only commands without importing the backend runtime.

    Importing ``lantai.cli.main`` initializes services that may consult the macOS
    Keychain.  Release validation calls the frozen executable with ``--version``;
    that query must remain deterministic and side-effect free even when the
    Keychain is locked or unavailable.
    """

    if sys.argv[1:] != ["--version"]:
        return
    os.environ["LANTAI_METADATA_ONLY"] = "1"
    try:
        from lantai.__version__ import __version__
    finally:
        os.environ.pop("LANTAI_METADATA_ONLY", None)

    print(f"澜台, version {__version__}")
    raise SystemExit(0)


_run_fast_metadata_command()


def _run_smoke_imports() -> None:
    for module_name in (
        "lantai.agents.tools",
        "lantai.agents.tools.scheduled_task",
        "lantai.app._app",
    ):
        print(f"[backend-smoke] import {module_name}", flush=True)
        importlib.import_module(module_name)

    # Exercise the document libraries through real in-memory/file round trips.
    # Import-only smoke misses optional converter dependencies such as Mammoth.
    print("[backend-smoke] import document runtimes", flush=True)
    import pandas as pd
    import pdfplumber
    from docx import Document
    from markitdown import MarkItDown
    from openpyxl import Workbook, load_workbook
    from pptx import Presentation
    from pypdf import PdfReader
    from reportlab.pdfgen import canvas

    assert int(pd.DataFrame({"value": [1, 2]}).sum().iloc[0]) == 3

    workbook_buffer = io.BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "runtime"
    workbook.save(workbook_buffer)
    workbook_buffer.seek(0)
    assert load_workbook(workbook_buffer).active["A1"].value == "runtime"

    document_buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("Lantai DOCX runtime")
    document.save(document_buffer)
    document_buffer.seek(0)
    assert "Lantai DOCX runtime" in {
        paragraph.text for paragraph in Document(document_buffer).paragraphs
    }

    presentation_buffer = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[0])
    presentation.save(presentation_buffer)
    presentation_buffer.seek(0)
    assert len(Presentation(presentation_buffer).slides) == 1

    pdf_buffer = io.BytesIO()
    pdf = canvas.Canvas(pdf_buffer)
    pdf.drawString(72, 720, "Lantai PDF runtime")
    pdf.save()
    pdf_buffer.seek(0)
    assert len(PdfReader(pdf_buffer).pages) == 1
    pdf_buffer.seek(0)
    with pdfplumber.open(pdf_buffer) as opened_pdf:
        assert len(opened_pdf.pages) == 1

    with tempfile.TemporaryDirectory(prefix="lantai-office-smoke-") as temp_dir:
        path = Path(temp_dir) / "runtime.docx"
        converted_document = Document()
        converted_document.add_heading("Lantai MarkItDown runtime", level=1)
        converted_document.save(path)
        converted = MarkItDown().convert(path)
        assert "Lantai MarkItDown runtime" in converted.text_content
    print("[backend-smoke] complete", flush=True)


if os.environ.get("LANTAI_BACKEND_SMOKE_IMPORT") == "1":
    _run_smoke_imports()
    raise SystemExit(0)

from lantai.cli.main import cli

if __name__ == "__main__":
    cli()
