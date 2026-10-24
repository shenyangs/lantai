#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Prove edit -> undo -> re-edit -> save -> restart -> readback in packaged Office.

The script operates on copies of real DOCX/XLSX/PPTX/PDF fixtures. It drives
the shipped renderer through Chromium's debugging protocol and clicks the
product's own undo/save controls. Readback then happens in a fresh app process
and independently against the file format on disk.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import math
import os
import plistlib
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, cast

from websockets.sync.client import ClientConnection, connect

try:
    from packaged_office_support import (
        OwnedProcessCohort,
        app_metadata as _bound_app_metadata,
    )
except ModuleNotFoundError as import_error:
    if import_error.name != "packaged_office_support":
        raise
    from scripts.packaged_office_support import (
        OwnedProcessCohort,
        app_metadata as _bound_app_metadata,
    )


SUPPORTED = {".docx", ".xlsx", ".pptx", ".pdf"}
XLSX_NATIVE_NUMBER_WAIVER_ID = "genoffice.xlsx.native_number_input"
GLOBAL_BY_SUFFIX = {
    ".docx": "__lantaiDocs",
    ".xlsx": "desktopApi",
    ".pptx": "slidesApi",
    ".pdf": "pdfApi",
}
EDITOR_BY_SUFFIX = {
    ".docx": "docs",
    ".xlsx": "sheets",
    ".pptx": "slides",
    ".pdf": "pdf",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _app_executable(app: Path) -> Path:
    with (app / "Contents" / "Info.plist").open("rb") as handle:
        info = plistlib.load(handle)
    executable = app / "Contents" / "MacOS" / str(info["CFBundleExecutable"])
    if not executable.is_file():
        raise AssertionError(f"missing app executable: {executable}")
    return executable


def _app_metadata(app: Path) -> dict[str, Any]:
    return _bound_app_metadata(app)


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
        listener.bind(("127.0.0.1", 0))
        return int(listener.getsockname()[1])


def _all_targets(port: int) -> list[dict[str, Any]]:
    try:
        with urllib.request.urlopen(
            f"http://127.0.0.1:{port}/json/list",
            timeout=1,
        ) as response:
            value = json.load(response)
    except (OSError, ValueError):
        return []
    return [
        item
        for item in value
        if item.get("type") == "page" and item.get("webSocketDebuggerUrl")
    ]


def _targets(port: int) -> list[dict[str, Any]]:
    return [
        item
        for item in _all_targets(port)
        if str(item.get("url") or "").startswith("file://")
    ]


class Cdp:
    def __init__(self, url: str):
        self._socket: ClientConnection = connect(
            url,
            origin=None,
            open_timeout=2,
            close_timeout=2,
        )
        self._next_id = 1

    def close(self) -> None:
        self._socket.close()

    def command(
        self,
        method: str,
        params: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        request_id = self._next_id
        self._next_id += 1
        self._socket.send(
            json.dumps(
                {"id": request_id, "method": method, "params": params or {}},
            ),
        )
        while True:
            # PDF.js rasterization and large Office canvas paints can keep the
            # renderer main thread busy beyond two seconds. Keep the command
            # bounded, but allow a realistic frame budget before declaring
            # the packaged app unresponsive.
            reply = json.loads(self._socket.recv(timeout=10))
            if reply.get("id") != request_id:
                continue
            if "error" in reply:
                raise RuntimeError(f"CDP {method} failed: {reply['error']}")
            return dict(reply.get("result") or {})

    def evaluate(self, expression: str, *, await_promise: bool = False) -> Any:
        result = (
            self.command(
                "Runtime.evaluate",
                {
                    "expression": expression,
                    "awaitPromise": await_promise,
                    "returnByValue": True,
                    "userGesture": True,
                },
            ).get("result")
            or {}
        )
        if (
            result.get("subtype") == "error"
            or result.get("className") == "Error"
        ):
            raise RuntimeError(
                f"renderer evaluation failed: {result.get('description')}",
            )
        return result.get("value")

    def insert_text(self, text: str) -> None:
        self.command("Input.insertText", {"text": text})

    def key(
        self,
        key: str,
        *,
        code: str | None = None,
        virtual: int = 0,
    ) -> None:
        common = {
            "key": key,
            "code": code or key,
            "windowsVirtualKeyCode": virtual,
            "nativeVirtualKeyCode": virtual,
        }
        self.command("Input.dispatchKeyEvent", {"type": "keyDown", **common})
        self.command("Input.dispatchKeyEvent", {"type": "keyUp", **common})

    def drag(self, x1: float, y1: float, x2: float, y2: float) -> None:
        self.command(
            "Input.dispatchMouseEvent",
            {
                "type": "mousePressed",
                "x": x1,
                "y": y1,
                "button": "left",
                "clickCount": 1,
            },
        )
        for step in range(1, 6):
            ratio = step / 5
            self.command(
                "Input.dispatchMouseEvent",
                {
                    "type": "mouseMoved",
                    "x": x1 + (x2 - x1) * ratio,
                    "y": y1 + (y2 - y1) * ratio,
                    "button": "left",
                    "buttons": 1,
                },
            )
        self.command(
            "Input.dispatchMouseEvent",
            {
                "type": "mouseReleased",
                "x": x2,
                "y": y2,
                "button": "left",
                "clickCount": 1,
            },
        )

    def screenshot(self, target: Path) -> None:
        result = self.command(
            "Page.captureScreenshot",
            {"format": "png", "fromSurface": True},
        )
        encoded = str(result.get("data") or "")
        if not encoded:
            raise AssertionError("renderer returned an empty screenshot")
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(base64.b64decode(encoded))


@dataclass
class AppSession:
    process: subprocess.Popen[str]
    root: Path
    port: int
    cdp: Cdp
    log: Path
    workbench: dict[str, Any]
    ownership: OwnedProcessCohort | None = None
    process_settlement: dict[str, Any] = field(default_factory=dict)
    renderer_identity: dict[str, Any] = field(default_factory=dict)

    def __post_init__(self) -> None:
        if self.ownership is None:
            self.ownership = OwnedProcessCohort(self.process)

    def stop(self) -> None:
        close_error: BaseException | None = None
        try:
            self.cdp.close()
        except BaseException as exc:
            close_error = exc
        try:
            assert self.ownership is not None
            self.process_settlement = self.ownership.settle()
            _private_json(
                self.log.with_suffix(".settlement.json"),
                self.process_settlement,
            )
            if self.process_settlement["ok"]:
                shutil.rmtree(self.root)
            else:
                raise AssertionError(
                    "Office owned process settlement is incomplete; user data retained",
                )
        except BaseException as cleanup_error:
            if close_error is not None:
                raise close_error from cleanup_error
            raise
        if close_error is not None:
            raise close_error


def _private_json(path: Path, value: dict[str, Any]) -> None:
    fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    os.fchmod(fd, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as handle:
        handle.write(json.dumps(value, indent=2) + "\n")


def _stop_preserving_primary(session: AppSession) -> None:
    primary = sys.exc_info()[1]
    try:
        session.stop()
    except BaseException as cleanup_error:
        if primary is not None:
            raise primary.with_traceback(
                primary.__traceback__,
            ) from cleanup_error
        raise


def _find_office_cdp(port: int, global_name: str, timeout: float) -> Cdp:
    deadline = time.monotonic() + timeout
    errors: list[str] = []
    while time.monotonic() < deadline:
        for target in _targets(port):
            cdp: Cdp | None = None
            try:
                cdp = Cdp(str(target["webSocketDebuggerUrl"]))
                matched = bool(
                    cdp.evaluate(
                        f"typeof window[{json.dumps(global_name)}] !== 'undefined'",
                    ),
                )
                if matched:
                    selected = cdp
                    cdp = None
                    return selected
                cdp.close()
                cdp = None
            except Exception as exc:  # renderer is still navigating
                errors.append(str(exc))
            finally:
                if cdp is not None:
                    try:
                        cdp.close()
                    except Exception:
                        pass
        time.sleep(0.2)
    raise AssertionError(
        f"Office renderer {global_name} was unavailable: {errors[-3:]}",
    )


def _find_workbench_evidence(port: int, timeout: float) -> dict[str, Any]:
    """Prove the host window and its attached native Office view are usable."""

    deadline = time.monotonic() + timeout
    errors: list[str] = []
    while time.monotonic() < deadline:
        for target in _all_targets(port):
            if str(target.get("url") or "").startswith("file://"):
                continue
            cdp: Cdp | None = None
            try:
                cdp = Cdp(str(target["webSocketDebuggerUrl"]))
                evidence = cdp.evaluate(
                    "(async () => { "
                    "const api=window.electronAPI?.officeWorkbench; "
                    "if(!api?.getState) return null; "
                    "const state=await api.getState(); "
                    "return {host:{width:window.innerWidth,height:window.innerHeight},state}; "
                    "})()",
                    await_promise=True,
                )
                if not isinstance(evidence, dict):
                    continue
                host = evidence.get("host") or {}
                state = evidence.get("state") or {}
                bounds = state.get("bounds") or {}
                # Each term is a distinct required observation.
                # pylint: disable=too-many-boolean-expressions
                if (
                    float(host.get("width") or 0) >= 320
                    and float(host.get("height") or 0) >= 200
                    and state.get("active") is True
                    and state.get("visible") is True
                    and float(bounds.get("width") or 0) >= 320
                    and float(bounds.get("height") or 0) >= 200
                ):
                    return evidence
                # pylint: enable=too-many-boolean-expressions
            except Exception as exc:
                errors.append(str(exc))
            finally:
                if cdp is not None:
                    try:
                        cdp.close()
                    except Exception:
                        pass
        time.sleep(0.2)
    raise AssertionError(
        "Office host never exposed positive, visible native-view bounds: "
        f"{errors[-3:]}",
    )


def _office_environment(root: Path) -> dict[str, str]:
    """Give the child isolated stores without inheriting credentials or hooks."""
    for name in ("home", "tmp"):
        (root / name).mkdir(mode=0o700, exist_ok=True)
    env = {
        key: os.environ[key]
        for key in (
            "PATH",
            "USER",
            "LOGNAME",
            "LANG",
            "LC_ALL",
            "LANTAI_NODE_BIN",
            "LANTAI_PPTXGENJS_MODULE",
        )
        if key in os.environ
    }
    env.update(
        {
            "HOME": str(root / "home"),
            "TMPDIR": str(root / "tmp"),
            "LANTAI_DESKTOP_USER_DATA_DIR": str(root / "user-data"),
            "LANTAI_DESKTOP_LIFECYCLE_SMOKE_FILE": str(
                root / "lifecycle.json",
            ),
            "LANTAI_WORKING_DIR": str(root / "user-data" / "working"),
            "LANTAI_SECRET_DIR": str(root / "user-data" / "working.secret"),
            "LANTAI_AUTH_ENABLED": "0",
            "LANTAI_DISABLE_GITHUB_UPDATER": "1",
            "LANTAI_DISABLE_KEYRING": "1",
            "LANTAI_SKIP_KEYRING": "1",
            "LANTAI_REMINDER_DELIVERY_WORKER": "0",
            "PYTHONDONTWRITEBYTECODE": "1",
            "PYTHON_DOTENV_DISABLED": "1",
            "LANTAI_DISABLE_DOTENV": "1",
        },
    )
    return env


def _document_binding(
    cdp: Cdp,
    fixture: Path,
    workbench: dict[str, Any],
) -> dict[str, Any]:
    fixture = fixture.resolve(strict=True)
    suffix = fixture.suffix.lower()
    editor = EDITOR_BY_SUFFIX[suffix]
    state = workbench.get("state") or {}
    expected_id = hashlib.sha256(f"{editor}:{fixture}".encode()).hexdigest()[
        :20
    ]
    if (
        state.get("active") is not True
        or state.get("visible") is not True
        or state.get("kind") != editor
        or state.get("display_name") != fixture.name
        or state.get("workbench_id") != expected_id
    ):
        raise AssertionError(
            "Office workbench belongs to another document or editor",
        )
    renderer = cdp.evaluate(
        "(() => { const h=document.querySelector('.lt-office-file-title'); "
        "return {url:location.href,title:document.title,file_title:(h?.textContent||'').trim(),"
        "file_title_attribute:h?.getAttribute('title')||'',global_available:typeof window["
        + json.dumps(GLOBAL_BY_SUFFIX[suffix])
        + "] !== 'undefined'}; })()",
    )
    expected_title = (
        str(fixture) if suffix in {".docx", ".pdf"} else fixture.name
    )
    # Each term is a distinct required observation.
    # pylint: disable=too-many-boolean-expressions
    if (
        not isinstance(renderer, dict)
        or not str(renderer.get("url") or "").startswith("file://")
        or renderer.get("global_available") is not True
        or renderer.get("file_title") != fixture.name
        or renderer.get("file_title_attribute") != expected_title
    ):
        raise AssertionError(
            "Office renderer does not prove the selected document",
        )
    # pylint: enable=too-many-boolean-expressions
    return {
        **renderer,
        "editor": editor,
        "document_path": str(fixture),
        "workbench_id": expected_id,
        "global_name": GLOBAL_BY_SUFFIX[suffix],
    }


@dataclass
class _OfficeDocumentBindingProbe:
    cdp: Cdp
    renderer: dict[str, Any] = field(default_factory=dict)

    def evaluate(self, expression: str) -> Any:
        value = self.cdp.evaluate(expression)
        if isinstance(value, dict):
            self.renderer = {
                key: value.get(key)
                for key in (
                    "file_title",
                    "file_title_attribute",
                    "global_available",
                )
                if isinstance(value.get(key), (str, bool, int, float))
                or value.get(key) is None
            }
            # Observations exclude URL credentials, query values, and fragments.
            parsed = urllib.parse.urlsplit(str(value.get("url") or ""))
            host = parsed.hostname or ""
            self.renderer["url"] = urllib.parse.urlunsplit(
                (parsed.scheme, host, parsed.path, "", ""),
            )
        else:
            self.renderer = {"response_type": type(value).__name__}
        return value


def _wait_document_binding(
    cdp: Cdp,
    fixture: Path,
    port: int,
    *,
    deadline: float,
    process: subprocess.Popen[str],
    failure_log: Path,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Wait for the unchanged identity predicate within one launch deadline."""

    probe = _OfficeDocumentBindingProbe(cdp)
    workbench: dict[str, Any] = {}
    last_error = "not observed"
    returncode = None
    while time.monotonic() < deadline:
        returncode = process.poll()
        if returncode is not None:
            break
        try:
            workbench = _find_workbench_evidence(
                port,
                max(0.0, deadline - time.monotonic()),
            )
            identity = _document_binding(cast(Cdp, probe), fixture, workbench)
            returncode = process.poll()
            if returncode is not None:
                break
            if time.monotonic() < deadline:
                return workbench, identity
            last_error = "identity reply arrived after the launch deadline"
        except Exception as exc:
            last_error = type(exc).__name__
        remaining = deadline - time.monotonic()
        if remaining > 0:
            time.sleep(min(0.15, remaining))
    state = workbench.get("state") or {}
    _private_json(
        failure_log,
        {
            "accepted": False,
            "document_path": str(fixture),
            "process_returncode": returncode,
            "last_error": last_error,
            "renderer": probe.renderer,
            "workbench": {
                key: state.get(key)
                for key in (
                    "active",
                    "visible",
                    "kind",
                    "display_name",
                    "workbench_id",
                )
            },
        },
    )
    if returncode is not None:
        raise AssertionError(
            f"app exited while waiting for Office document binding: {returncode}",
        )
    raise AssertionError(
        "Office renderer did not prove the selected document within the launch "
        f"deadline: {last_error}",
    )


def _capture_launch_failure(
    root: Path,
    log: Path,
    port: int | None,
    stage: str,
    process: Any,
    primary: BaseException,
) -> None:
    """Keep bounded structural evidence before launch cleanup deletes the profile.

    Never copy document text, URL parameters, or raw backend log lines.
    Diagnostic errors cannot change the original failure or settlement.
    """
    evidence: dict[str, Any] = {
        "stage": stage,
        "error_type": type(primary).__name__,
        "process_started": process is not None,
        "targets": [],
        "logs": {},
    }
    if process is not None:
        evidence["process_returncode"] = process.poll()
    if port is not None:
        for target in _all_targets(port)[:16]:
            parsed = urllib.parse.urlsplit(str(target.get("url") or ""))
            evidence["targets"].append(
                {
                    "scheme": parsed.scheme,
                    "editor": next(
                        (
                            name
                            for name in ("docs", "sheets", "slides", "pdf")
                            if f"/{name}/" in parsed.path
                        ),
                        None,
                    ),
                    "type": str(target.get("type") or "")[:32],
                }
            )
    signals = {
        "backend_health_timeout": b"Backend health check timed out",
        "connection_refused": b"ERR_CONNECTION_REFUSED",
        "renderer_gone": b"render-process-gone",
        "uncaught_exception": b"uncaughtException",
    }
    for name, path in (
        ("electron", log),
        ("backend", root / "user-data" / "backend.log"),
    ):
        try:
            with path.open("rb") as handle:
                handle.seek(0, 2)
                size = handle.tell()
                handle.seek(max(0, size - 65536))
                tail = handle.read(65536)
            evidence["logs"][name] = {
                "exists": True,
                "size": size,
                "tail_truncated": size > len(tail),
                "signals": {
                    key: value in tail for key, value in signals.items()
                },
            }
        except OSError as error:
            evidence["logs"][name] = {
                "exists": path.exists(),
                "read_error": type(error).__name__,
            }
    _private_json(log.with_suffix(".launch-failure.json"), evidence)


def _launch_once(
    app: Path,
    fixture: Path,
    log: Path,
    timeout: float,
) -> AppSession:
    fixture = fixture.expanduser().resolve(strict=True)
    root = Path(tempfile.mkdtemp(prefix=f"lantai-edit-{fixture.suffix[1:]}-"))
    process: subprocess.Popen[str] | None = None
    ownership: OwnedProcessCohort | None = None
    cdp: Cdp | None = None
    port: int | None = None
    stage = "prepare"
    try:
        deadline = time.monotonic() + timeout
        port = _free_port()
        env = _office_environment(root)
        log.parent.mkdir(parents=True, exist_ok=True)
        fd = os.open(log, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
        os.fchmod(fd, 0o600)
        with os.fdopen(fd, "w", encoding="utf-8") as log_handle:
            # Ownership spans AppSession; stop() performs bounded settlement.
            process = subprocess.Popen(  # pylint: disable=consider-using-with
                [
                    str(_app_executable(app)),
                    f"--remote-debugging-port={port}",
                    str(fixture),
                ],
                cwd=str(app.parent),
                env=env,
                stdout=log_handle,
                stderr=subprocess.STDOUT,
                text=True,
                start_new_session=True,
            )
        stage = "ownership"
        ownership = OwnedProcessCohort(process)
        stage = "office-api"
        cdp = _find_office_cdp(
            port,
            GLOBAL_BY_SUFFIX[fixture.suffix.lower()],
            max(0.0, deadline - time.monotonic()),
        )
        stage = "body-content"
        _wait(
            lambda: bool(
                cdp.evaluate(
                    "document.body && document.body.innerText.length > 20",
                ),
            ),
            timeout=max(0.0, deadline - time.monotonic()),
            label=f"{fixture.suffix} body content",
            process=process,
        )
        stage = "viewport"
        _wait(
            lambda: bool(
                cdp.evaluate(
                    "window.innerWidth >= 300 && window.innerHeight >= 200",
                ),
            ),
            timeout=max(0.0, deadline - time.monotonic()),
            label=f"{fixture.suffix} interactive viewport",
            process=process,
        )
        stage = "document-binding"
        workbench, renderer_identity = _wait_document_binding(
            cdp,
            fixture,
            port,
            deadline=deadline,
            process=process,
            failure_log=log.with_suffix(".document-binding.json"),
        )
        return AppSession(
            process=process,
            root=root,
            port=port,
            cdp=cdp,
            log=log,
            workbench=workbench,
            ownership=ownership,
            renderer_identity=renderer_identity,
        )
    except BaseException as primary:
        try:
            _capture_launch_failure(root, log, port, stage, process, primary)
        except BaseException:
            pass  # Best-effort diagnostics must never prevent owner settlement.
        try:
            if cdp is not None:
                try:
                    cdp.close()
                except BaseException:
                    pass
            if process is None:
                settlement = {
                    "ok": True,
                    "live": [],
                    "unknown": [],
                    "signals": [],
                    "process_started": False,
                }
            else:
                settlement = (
                    ownership.settle()
                    if ownership is not None
                    else {
                        "ok": False,
                        "live": [],
                        "unknown": [
                            {
                                "pid": process.pid,
                                "reason": "initial Office process ownership could not be confirmed",
                            },
                        ],
                        "signals": [],
                    }
                )
            _private_json(log.with_suffix(".settlement.json"), settlement)
            if settlement["ok"]:
                shutil.rmtree(root)
            else:
                raise AssertionError(
                    "Office owned process settlement is incomplete; user data retained",
                ) from primary
        except BaseException as cleanup_error:
            raise primary.with_traceback(
                primary.__traceback__,
            ) from cleanup_error
        raise


def _launch(app: Path, fixture: Path, log: Path, timeout: float) -> AppSession:
    """Retry one clean cold start only when the Office renderer never appears."""

    first_error: AssertionError | None = None
    for attempt in (1, 2):
        attempt_log = (
            log
            if attempt == 1
            else log.with_name(f"{log.stem}-retry{log.suffix}")
        )
        try:
            return _launch_once(app, fixture, attempt_log, timeout)
        except AssertionError as exc:
            if (
                "Office renderer" not in str(exc)
                or exc.__cause__ is not None
                or attempt == 2
            ):
                if first_error is not None:
                    raise AssertionError(
                        f"{exc}; first cold-start failure: {first_error}",
                    ) from exc
                raise
            first_error = exc
            time.sleep(0.5)
    raise AssertionError("unreachable Office launch retry state")


def _wait(
    predicate: Callable[[], bool],
    *,
    timeout: float,
    label: str,
    process: subprocess.Popen[str] | None = None,
) -> None:
    deadline = time.monotonic() + timeout
    last_error = ""
    while time.monotonic() < deadline:
        if process is not None and process.poll() is not None:
            raise AssertionError(
                f"app exited while waiting for {label}: {process.returncode}",
            )
        try:
            if predicate():
                return
        except Exception as exc:
            last_error = str(exc)
        time.sleep(0.15)
    raise AssertionError(f"timed out waiting for {label}: {last_error}")


def _native_click_target(cdp: Cdp, expression: str) -> dict[str, Any] | None:
    """Return native-click geometry only for a visible, hit-testable control."""

    target = cdp.evaluate(
        "(() => { const element = "
        + expression
        + "; if (!(element instanceof HTMLElement) || element.disabled) return null; "
        "const style=getComputedStyle(element); const r=element.getBoundingClientRect(); "
        "if(r.width<=0||r.height<=0||style.display==='none"
        "'||style.visibility==='hidden'||Number(style.opac"
        "ity)===0) return null; "
        "const x=r.left+r.width/2,y=r.top+r.height/2; const hit=document.elementFromPoint(x,y); "
        "if(!(hit instanceof Element)||(hit!==element&&!el"
        "ement.contains(hit)&&!hit.contains(element))) "
        "return null; "
        "return {x,y,rect:{left:r.left,top:r.top,width:r.width,height:r.height},"
        "tag:element.tagName,title:element.title||'',aria:element.getAttribute('aria-label')||'',"
        "text:(element.innerText||'').trim(),hitTag:hit.ta"
        "gName,hitClass:String(hit.className||'')}; })()",
    )
    return target if isinstance(target, dict) else None


def _native_click(cdp: Cdp, expression: str, *, label: str) -> dict[str, Any]:
    """Hit-test and click a visible control through Chromium's native input path."""

    target = _native_click_target(cdp, expression)
    if not isinstance(target, dict):
        raise AssertionError(
            f"renderer control was not visible/hit-testable: {label}",
        )
    x, y = float(target["x"]), float(target["y"])
    cdp.command(
        "Input.dispatchMouseEvent",
        {
            "type": "mousePressed",
            "x": x,
            "y": y,
            "button": "left",
            "buttons": 1,
            "clickCount": 1,
        },
    )
    cdp.command(
        "Input.dispatchMouseEvent",
        {
            "type": "mouseReleased",
            "x": x,
            "y": y,
            "button": "left",
            "buttons": 0,
            "clickCount": 1,
        },
    )
    return {"label": label, **target, "input": "Input.dispatchMouseEvent"}


def _toolbar_button(title_part: str) -> str:
    needle = json.dumps(title_part, ensure_ascii=False)
    return (
        "[...document.querySelectorAll('button')].find((b) => "
        f"(b.title || b.getAttribute('aria-label') || '').includes({needle}))"
    )


def _text_button(text: str) -> str:
    needle = json.dumps(text, ensure_ascii=False)
    return (
        "[...document.querySelectorAll('button')].find((b) => "
        f"(b.innerText || '').trim() === {needle})"
    )


def _button_disabled(cdp: Cdp, title_part: str) -> bool:
    return bool(
        cdp.evaluate(f"Boolean(({_toolbar_button(title_part)})?.disabled)"),
    )


def _body_text(cdp: Cdp) -> str:
    return str(cdp.evaluate("document.body?.innerText || ''") or "")


def _docx_model_probe(cdp: Cdp) -> dict[str, Any] | None:
    """Read only the public editor state needed to prove a real DOCX edit."""

    probe = cdp.evaluate(
        "(() => { const e=window.__lantaiDocs?.editor; if(!e) return null; "
        "const s=e.state.selection; let idx=null; "
        "for(let d=s.$from.depth; d>=0; d--){ const v=s.$from.node(d)?.attrs?.docxIndex; "
        "if(v!==undefined && v!==null){ idx=v; break; } } "
        "const active=document.activeElement; const root=document.querySelector('.ProseMirror'); "
        "return {text:e.getText(),from:s.from,to:s.to,docxIndex:idx,"
        "activeEditor:Boolean(root && (active===root || root.contains(active))),"
        "activeTag:active?.tagName||'',activeClass:String(active?.className||'')}; })()",
    )
    return probe if isinstance(probe, dict) else None


def _focus_docx_existing_paragraph(cdp: Cdp) -> dict[str, Any]:
    """Native-click the final glyph of an existing, persistable DOCX block."""

    point = cdp.evaluate(
        "(() => { const blocks=[...document.querySelectorAll('.ProseMirror > [data-idx]')]"
        ".filter(e=>{const r=e.getBoundingClientRect(),s=getComputedStyle(e);"
        "return r.width>0&&r.height>0&&s.display!=='none'&&s.visibility!=='hidden'}); "
        "for(let i=blocks.length-1;i>=0;i--){ const block=blocks[i]; "
        "const walker=document.createTreeWalker(block,NodeFilter.SHOW_TEXT); let node,last=null; "
        "while((node=walker.nextNode())) if((node.nodeValue||'').trim()) last=node; "
        "if(!last) continue; const range=document.createRange(); const n=last.nodeValue.length; "
        "range.setStart(last,Math.max(0,n-1)); range.setEnd(last,n); "
        "const rect=[...range.getClientRects()].at(-1); "
        "if(!rect||rect.width<=0||rect.height<=0) "
        "continue; "
        "const x=Math.max(rect.left+1,rect.right-1),y=rect.top+rect.height/2; "
        "const hit=document.elementFromPoint(x,y),owner=hit?.closest?.('[data-idx]'); "
        "if(owner!==block) continue; return {x,y,docxIndex:Number(block.dataset.idx),"
        "blockText:(block.innerText||block.textContent||'').trim(),hitTag:hit.tagName,"
        "hitClass:String(hit.className||''),rect:{left:rec"
        "t.left,top:rect.top,width:rect.width,height:rect."
        "height}}; "
        "} return null; })()",
    )
    if not isinstance(point, dict):
        raise AssertionError(
            "DOCX existing paragraph had no visible native hit target",
        )
    for event_type, buttons in (("mousePressed", 1), ("mouseReleased", 0)):
        cdp.command(
            "Input.dispatchMouseEvent",
            {
                "type": event_type,
                "x": float(point["x"]),
                "y": float(point["y"]),
                "button": "left",
                "buttons": buttons,
                "clickCount": 1,
            },
        )
    cdp.key("End", code="End", virtual=35)
    probe = _docx_model_probe(cdp)
    if not probe or probe.get("activeEditor") is not True:
        raise AssertionError(
            f"DOCX native paragraph click did not focus the editor: {probe}",
        )
    if probe.get("docxIndex") != point.get("docxIndex"):
        raise AssertionError(
            f"DOCX native click selected block {probe.get('docxIndex')}, "
            f"expected {point.get('docxIndex')}",
        )
    return {**point, "selection": probe}


def _docx_insert_native_text(cdp: Cdp, value: str) -> None:
    """Enter DOCX text through Chromium's real printable-key/input path."""

    if not value:
        raise AssertionError("DOCX input must not be empty")
    first = value[0]
    code = "Space" if first == " " else "Unidentified"
    virtual = 32 if first == " " else 0
    cdp.command(
        "Input.dispatchKeyEvent",
        {
            "type": "keyDown",
            "key": first,
            "code": code,
            "text": first,
            "unmodifiedText": first,
            "windowsVirtualKeyCode": virtual,
            "nativeVirtualKeyCode": virtual,
        },
    )
    cdp.command(
        "Input.dispatchKeyEvent",
        {
            "type": "keyUp",
            "key": first,
            "code": code,
            "windowsVirtualKeyCode": virtual,
            "nativeVirtualKeyCode": virtual,
        },
    )
    if len(value) > 1:
        cdp.insert_text(value[1:])


def _xlsx_canvas_probe(cdp: Cdp) -> dict[str, Any]:
    # The Univer canvas includes its row/column headers. These points are cell
    # interiors in the shipped default grid, not offsets that land on a header.
    probe = cdp.evaluate(
        "(() => { "
        "const e=document.querySelector('canvas[id^=univer-sheet-main-canvas]'); "
        "if(!e) return null; const r=e.getBoundingClientRect(); "
        "const viewport=e.closest('section.univer-relative.univer-overflow-hidden'); "
        "const vr=viewport?.getBoundingClientRect(); "
        "if(r.width < 300 || r.height < 200 || !vr || "
        "vr.width < 300 || vr.height < 200) return null; "
        "const points={A1:{x:r.x+80,y:r.y+34},B2:{x:r.x+160,y:r.y+54},C2:{x:r.x+240,y:r.y+54}}; "
        "const hits={}; "
        "for(const [name,p] of Object.entries(points)){ "
        "if(p.x<r.left||p.x>=r.right||p.y<r.top||p.y>=r.bottom) return null; "
        "const hit=document.elementFromPoint(p.x,p.y); "
        "const expected=hit instanceof HTMLCanvasElement && "
        "((hit.id||'').startsWith('univer-sheet-') || Bool"
        "ean(hit.closest('section.univer-relative.univer-o"
        "verflow-hidden'))); "
        "if(!expected) return null; "
        "hits[name]={tag:hit.tagName,id:hit.id||'',className:String(hit.className||'')}; "
        "} "
        "return {canvas:{x:r.x,y:r.y,width:r.width,height:r.height},"
        "viewport:{x:vr.x,y:vr.y,width:vr.width,height:vr.height},points,hits}; })()",
    )
    if not isinstance(probe, dict):
        raise AssertionError(
            "XLSX canvas was not positively sized or native hit-testing missed its Univer layer",
        )
    return probe


def _xlsx_wait_for_canvas_probe(cdp: Cdp, timeout: float) -> dict[str, Any]:
    """Wait until Univer's canvas is sized and accepts native hit-testing."""

    probe: dict[str, Any] | None = None

    def ready() -> bool:
        nonlocal probe
        try:
            probe = _xlsx_canvas_probe(cdp)
        except AssertionError:
            return False
        return True

    _wait(ready, timeout=timeout, label="XLSX native canvas")
    if not isinstance(probe, dict):
        raise AssertionError(
            "XLSX native canvas wait completed without evidence",
        )
    return probe


def _xlsx_name_box(cdp: Cdp) -> str:
    return str(
        cdp.evaluate(
            "(() => { const candidates=[...document.querySelectorAll('input')].filter((e) => { "
            "const r=e.getBoundingClientRect(); return r.width>0 && r.height>0 && "
            "/^[A-Za-z]{1,3}[1-9][0-9]*$/.test((e.value||'').trim()); }); "
            "return candidates.sort((a,b) => b.getBoundingClie"
            "ntRect().y-a.getBoundingClientRect().y)[0]?.value"
            " || ''; })()",
        )
        or "",
    ).upper()


def _xlsx_select_cell(cdp: Cdp, address: str) -> None:
    # Focus must be transferred to the Office WebContents *before* the mouse
    # press. Otherwise Electron can update the canvas selection while leaving
    # keyboard routing on the parent workbench, and the user's next key never
    # reaches the sheet.
    cdp.command("Page.bringToFront", {})
    cdp.command("Emulation.setFocusEmulationEnabled", {"enabled": True})
    # Univer remounts its canvas after toolbar operations such as Undo.  The
    # workbook can already report a clean state while that replacement canvas
    # is still between layout and hit-testing.  Wait for the same positive
    # native probe used at editor startup before dispatching the next click.
    probe = _xlsx_wait_for_canvas_probe(cdp, timeout=10)
    point = (probe.get("points") or {}).get(address)
    if not isinstance(point, dict):
        raise AssertionError(
            f"XLSX {address} native coordinate was unavailable",
        )
    x, y = float(point["x"]), float(point["y"])
    for event_type in ("mousePressed", "mouseReleased"):
        cdp.command(
            "Input.dispatchMouseEvent",
            {
                "type": event_type,
                "x": x,
                "y": y,
                "button": "left",
                "clickCount": 1,
            },
        )


def _xlsx_prove_native_selection(cdp: Cdp, timeout: float) -> dict[str, Any]:
    """Click two real cells and require the visible Name Box to follow."""

    ready_probe = _xlsx_wait_for_canvas_probe(cdp, timeout)
    initial = _xlsx_name_box(cdp)
    _xlsx_select_cell(cdp, "B2")
    _wait(
        lambda: _xlsx_name_box(cdp) == "B2",
        timeout=timeout,
        label="XLSX B2 selection",
    )
    after_b2 = _xlsx_name_box(cdp)
    _xlsx_select_cell(cdp, "A1")
    _wait(
        lambda: _xlsx_name_box(cdp) == "A1",
        timeout=timeout,
        label="XLSX A1 selection",
    )
    return {
        "initial_name_box": initial,
        "after_b2_click": after_b2,
        "after_a1_click": _xlsx_name_box(cdp),
        "canvas_probe": _xlsx_canvas_probe(cdp),
        "ready_probe": ready_probe,
    }


def _xlsx_select_a1(cdp: Cdp) -> None:
    _xlsx_select_cell(cdp, "A1")


def _xlsx_visible_editor(cdp: Cdp) -> dict[str, Any] | None:
    editor = cdp.evaluate(
        "(() => { const e=document.querySelector('textarea[data-lantai-cell-editor=true]'); "
        "if(!(e instanceof HTMLTextAreaElement)) return null; const r=e.getBoundingClientRect(); "
        "const s=getComputedStyle(e); const x=r.left+r.width/2,y=r.top+r.height/2; "
        "const hit=document.elementFromPoint(x,y); "
        "return {address:e.dataset.cellAddress||'',value:e.value,active:document.activeElement===e,"
        "rect:{left:r.left,top:r.top,width:r.width,height:r.height},"
        "visible:r.width>0&&r.height>0&&s.display!=='none'"
        "&&s.visibility!=='hidden'&&Number(s.opacity)>0,"
        "hit:hit===e||Boolean(hit&&e.contains(hit))}; })()",
    )
    return editor if isinstance(editor, dict) else None


def _xlsx_replace_selected_cell(
    cdp: Cdp,
    address: str,
    value: str,
    timeout: float,
) -> dict[str, Any]:
    """Replace the selected cell through Univer's real keyboard edit path.

    ``Input.insertText`` alone does not open Univer's in-cell editor. A real
    printable keyDown is what switches the selected cell into replacement
    mode; the remainder can then use Chromium's text input path before Enter
    commits the edit. This deliberately avoids private renderer/test hooks.
    """

    if not value:
        raise AssertionError("XLSX replacement value must not be empty")
    # The first edit starts from a clean workbook, while later edits in the
    # same save transaction correctly start with Save already enabled. Only a
    # clean-to-dirty transition before Enter is suspicious; treating an
    # already-dirty workbook as a fresh failure makes every multi-cell edit
    # sequence fail deterministically on its second cell.
    save_was_disabled_before_edit = _button_disabled(cdp, "保存")
    # Keep focus emulation enabled through the first printable key. The actual
    # target transfer happened before selection in ``_xlsx_select_cell``.
    cdp.command("Page.bringToFront", {})
    cdp.command("Emulation.setFocusEmulationEnabled", {"enabled": True})
    first = value[0]
    upper = first.upper()
    is_ascii_letter = len(upper) == 1 and "A" <= upper <= "Z"
    virtual = ord(upper) if is_ascii_letter else 0
    code = f"Key{upper}" if is_ascii_letter else "Unidentified"
    cdp.command(
        "Input.dispatchKeyEvent",
        {
            "type": "keyDown",
            "key": first,
            "code": code,
            "text": first,
            "unmodifiedText": first,
            "windowsVirtualKeyCode": virtual,
            "nativeVirtualKeyCode": virtual,
        },
    )
    cdp.command(
        "Input.dispatchKeyEvent",
        {
            "type": "keyUp",
            "key": first,
            "code": code,
            "windowsVirtualKeyCode": virtual,
            "nativeVirtualKeyCode": virtual,
        },
    )
    _wait(
        lambda: bool(
            (probe := _xlsx_visible_editor(cdp))
            and probe.get("active") is True
            and probe.get("visible") is True
            and probe.get("hit") is True
            and probe.get("address") == address
            and probe.get("value") == first,
        ),
        timeout=min(timeout, 5),
        label="XLSX visible focused edit layer",
    )
    if save_was_disabled_before_edit and not _button_disabled(cdp, "保存"):
        raise AssertionError(
            "XLSX became dirty before the visible edit was committed",
        )
    if len(value) > 1:
        cdp.insert_text(value[1:])
    _wait(
        lambda: (_xlsx_visible_editor(cdp) or {}).get("value") == value,
        timeout=min(timeout, 5),
        label="XLSX marker visible before commit",
    )
    precommit = _xlsx_visible_editor(cdp)
    if not isinstance(precommit, dict) or not all(
        precommit.get(key) is True for key in ("active", "visible", "hit")
    ):
        raise AssertionError(
            "XLSX edit layer lost focus, visibility, or native hit-testing",
        )
    cdp.key("Enter", code="Enter", virtual=13)
    return {
        "address": address,
        "first_key": first,
        "marker_before_commit": value,
        "save_was_disabled_before_edit": save_was_disabled_before_edit,
        "editor": precommit,
    }


def _xlsx_copy_selected_cell(cdp: Cdp) -> str:
    # Univer renders cells into a canvas, so neither innerText nor the
    # accessibility tree contains their values. Exercise the same Copy path a
    # user does. Seed and inspect the macOS system clipboard from outside the
    # renderer so this remains independent of navigator.clipboard and any
    # renderer-local permission or test behavior.
    sentinel = f"__LANTAI_CLIPBOARD_PROBE_{time.time_ns()}__"
    subprocess.run(
        ["pbcopy"],
        input=sentinel,
        text=True,
        check=True,
        timeout=3,
    )
    cdp.command("Page.bringToFront", {})
    cdp.command("Emulation.setFocusEmulationEnabled", {"enabled": True})
    for event_type in ("keyDown", "keyUp"):
        cdp.command(
            "Input.dispatchKeyEvent",
            {
                "type": event_type,
                "key": "c",
                "code": "KeyC",
                "modifiers": 4,  # macOS Command modifier in CDP
                "windowsVirtualKeyCode": 67,
                "nativeVirtualKeyCode": 67,
            },
        )
    deadline = time.monotonic() + 2
    while time.monotonic() < deadline:
        value = subprocess.run(
            ["pbpaste"],
            capture_output=True,
            text=True,
            check=True,
            timeout=3,
        ).stdout.rstrip("\r\n")
        if value != sentinel:
            return value
        time.sleep(0.05)
    raise AssertionError("XLSX Copy did not reach the macOS system clipboard")


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode(
            "utf-8",
            errors="replace",
        )
    return "".join(
        re.findall(r"<w:t(?:\s[^>]*)?>(.*?)</w:t>", xml, flags=re.DOTALL),
    )


def _xlsx_cells(path: Path) -> dict[str, Any]:
    import openpyxl  # type: ignore[import-not-found]

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook.active
        return {
            "a1": str(sheet["A1"].value or ""),
            "b2": sheet["B2"].value,
            "b2_type": sheet["B2"].data_type,
            "c2": sheet["C2"].value,
            "c2_type": sheet["C2"].data_type,
        }
    finally:
        workbook.close()


def _pptx_slide_count(path: Path) -> int:
    with zipfile.ZipFile(path) as archive:
        return sum(
            bool(re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
            for name in archive.namelist()
        )


def _pdf_square_annotation_count(path: Path) -> int:
    from pypdf import PdfReader

    return sum(
        annotation.get_object().get("/Subtype") == "/Square"
        for page in PdfReader(path).pages
        for annotation in page.get("/Annots", [])
    )


def _pdf_annotation_readback(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader
    from pypdf.generic import ContentStream

    reader = PdfReader(path)
    annotations: list[dict[str, Any]] = []
    for page_index, page in enumerate(reader.pages):
        for reference in page.get("/Annots", []):
            annotation = reference.get_object()
            if annotation.get("/Subtype") != "/Square":
                continue
            appearance = (annotation.get("/AP") or {}).get("/N")
            if appearance is None or not hasattr(
                appearance.get_object(),
                "get_data",
            ):
                raise AssertionError(
                    "Saved PDF square has no normal appearance stream",
                )
            stream = appearance.get_object()
            matrix = list(
                map(float, stream.get("/Matrix", [1, 0, 0, 1, 0, 0])),
            )
            annotation_rect = list(map(float, annotation["/Rect"]))
            appearance_bbox = list(map(float, stream.get("/BBox", [])))
            if (
                int(annotation.get("/F", 0)) & (1 | 2 | 8 | 16 | 32)
                or matrix != [1, 0, 0, 1, 0, 0]
                or appearance_bbox != annotation_rect
            ):
                raise AssertionError(
                    "Saved PDF square appearance is hidden or has unsupported transforms",
                )
            operations = ContentStream(stream, reader).operations
            rectangles = [args for args, op in operations if op == b"re"]
            colors = [args for args, op in operations if op == b"RG"]
            widths = [args for args, op in operations if op == b"w"]
            # Each term is a distinct required observation.
            # pylint: disable=too-many-boolean-expressions
            if (
                len(rectangles) != 1
                or not colors
                or not widths
                or len(colors) != 1
                or len(widths) != 1
                or [op for _args, op in operations][-2:] != [b"re", b"S"]
                or any(
                    op not in {b"w", b"J", b"j", b"RG", b"re", b"S"}
                    for _args, op in operations
                )
            ):
                raise AssertionError(
                    "Saved PDF square appearance lacks a verifiable rectangle stroke",
                )
            # pylint: enable=too-many-boolean-expressions
            x, y, width, height = map(float, rectangles[0])
            color = list(map(float, colors[-1]))
            line_width = float(widths[-1][0])
            # Each term is a distinct required observation.
            # pylint: disable=too-many-boolean-expressions
            if (
                width <= 0
                or height <= 0
                or len(color) != 3
                or not all(0 <= value <= 1 for value in color)
                or not 0 < line_width <= 20
                or float(page.get("/UserUnit", 1)) != 1
            ):
                raise AssertionError(
                    "Saved PDF square appearance has unsupported geometry or color",
                )
            # pylint: enable=too-many-boolean-expressions
            annotations.append(
                {
                    "page_index": page_index,
                    "annotation_rect": annotation_rect,
                    "drawn_rect": [x, y, x + width, y + height],
                    "color": color,
                    "line_width": line_width,
                    "page_cropbox": list(map(float, page.cropbox)),
                    "page_rotation": int(page.get("/Rotate", 0)) % 360,
                    "appearance_sha256": hashlib.sha256(
                        stream.get_data(),
                    ).hexdigest(),
                },
            )
    return {
        "file_sha256": _sha256(path),
        "square_annotation_count": len(annotations),
        "annotations": annotations,
    }


def _pdf_rectangle_pixel_proof(
    image: Any,
    rect: list[float],
    color: list[int],
    line_width: float,
) -> dict[str, Any]:
    image = image.convert("RGB")
    if (
        len(rect) != 4
        or len(color) != 3
        or not all(math.isfinite(value) for value in rect)
        or not 0 < line_width <= 20
    ):
        raise AssertionError("Invalid PDF rectangle pixel proof input")
    left, top, right, bottom = rect
    radius = min(12, max(2, math.ceil(line_width / 2) + 1))
    margin = radius + 5
    if not (
        margin <= left < right <= image.width - margin
        and margin <= top < bottom <= image.height - margin
    ):
        raise AssertionError(
            "PDF rectangle is outside the visible pixel region",
        )
    pixels = image.load()
    ratios: dict[str, float] = {}
    for edge in ("top", "right", "bottom", "left"):
        matches = 0
        for index in range(16):
            fraction = 0.1 + index * 0.8 / 15
            horizontal = edge in {"top", "bottom"}
            x = (
                left + (right - left) * fraction
                if horizontal
                else (left if edge == "left" else right)
            )
            y = (
                (top if edge == "top" else bottom)
                if horizontal
                else top + (bottom - top) * fraction
            )
            matched = False
            for offset in range(-radius, radius + 1):
                sample_x = round(x if horizontal else x + offset)
                sample_y = round(y + offset if horizontal else y)
                pixel = pixels[sample_x, sample_y]
                if (
                    max(
                        abs(pixel[channel] - color[channel])
                        for channel in range(3)
                    )
                    > 32
                ):
                    continue
                direction = -1 if edge in {"top", "left"} else 1
                outside_x = round(
                    x if horizontal else x + direction * (radius + 4),
                )
                outside_y = round(
                    y + direction * (radius + 4) if horizontal else y,
                )
                outside = pixels[outside_x, outside_y]
                if (
                    max(
                        abs(pixel[channel] - outside[channel])
                        for channel in range(3)
                    )
                    >= 20
                ):
                    matched = True
                    break
            matches += matched
        ratios[edge] = matches / 16
    if min(ratios.values()) < 0.8:
        raise AssertionError(
            "Saved PDF rectangle is absent from the actual visible pixels",
        )
    return {
        "edge_match_ratios": list(ratios.values()),
        "edges": list(ratios),
        "samples_per_edge": 16,
        "image_size": [image.width, image.height],
        "rect": rect,
    }


def _pdf_normalized_rectangle(annotation: dict[str, Any]) -> list[float]:
    left, bottom, right, top = map(float, annotation["page_cropbox"])
    width, height = right - left, top - bottom
    rotation = int(annotation["page_rotation"])
    if width <= 0 or height <= 0 or rotation not in {0, 90, 180, 270}:
        raise AssertionError("Unsupported saved PDF rectangle page geometry")
    x1, y1, x2, y2 = map(float, annotation["drawn_rect"])
    points = []
    for x, y in ((x1 - left, y1 - bottom), (x2 - left, y2 - bottom)):
        if rotation == 0:
            points.append((x / width, (height - y) / height))
        elif rotation == 90:
            points.append((y / height, x / width))
        elif rotation == 180:
            points.append(((width - x) / width, y / height))
        else:
            points.append(((height - y) / height, (width - x) / width))
    result = [
        min(p[0] for p in points),
        min(p[1] for p in points),
        max(p[0] for p in points),
        max(p[1] for p in points),
    ]
    if not all(math.isfinite(v) and 0 <= v <= 1 for v in result):
        raise AssertionError(
            "Saved PDF rectangle is outside its page crop box",
        )
    return result


def _pdf_visible_rectangles(cdp: Cdp) -> list[dict[str, Any]]:
    result = cdp.evaluate(
        "(() => { const layer=document.querySelector('.pdf-row[data-idx=\"0\"] .pdf-draw-layer'); "
        "if(!layer) return null; const page=layer.getBoundingClientRect(); "
        "if(page.width<=0||page.height<=0) return null; "
        "return [...layer.querySelectorAll('.pdf-draw-shape rect')].flatMap(e=>{ "
        "const s=getComputedStyle(e),r=e.getBoundingClientRect(); "
        "const color=s.stroke.match(/^rgb\\(\\s*(\\d+)\\s*,\\s*(\\d+)\\s*,\\s*(\\d+)\\s*\\)$/); "
        "if(!color||s.display==='none'||s.visibility==='hidden'||Number(s.opacity)===0) return []; "
        "return [{page_index:0,normalized_rect:[(r.left-pa"
        "ge.left)/page.width,(r.top-page.top)/page.height,"
        "(r.right-page.left)/page.width,(r.bottom-page.top)/page.height],"
        "color:color.slice(1).map(Number),normalized_line_"
        "width:parseFloat(s.strokeWidth)/page.width}]; "
        "}); })()",
    )
    if not isinstance(result, list):
        raise AssertionError("PDF rectangle drawing layer is unavailable")
    for item in result:
        rect = item.get("normalized_rect") or []
        # Each term is a distinct required observation.
        # pylint: disable=too-many-boolean-expressions
        if (
            len(rect) != 4
            or not all(
                isinstance(v, (int, float))
                and math.isfinite(v)
                and 0 <= v <= 1
                for v in rect
            )
            or rect[0] >= rect[2]
            or rect[1] >= rect[3]
            or len(item.get("color") or []) != 3
            or not 0 < float(item.get("normalized_line_width") or 0) <= 0.1
        ):
            raise AssertionError(
                "PDF visible rectangle has invalid geometry or color",
            )
        # pylint: enable=too-many-boolean-expressions
    return result


def _pdf_saved_rectangle(
    before: dict[str, Any],
    after: dict[str, Any],
    interaction: dict[str, Any],
) -> dict[str, Any]:
    previous = {
        json.dumps(item, sort_keys=True) for item in before["annotations"]
    }
    added = [
        item
        for item in after["annotations"]
        if json.dumps(item, sort_keys=True) not in previous
    ]
    if (
        len(added) != 1
        or len(after["annotations"]) != len(before["annotations"]) + 1
    ):
        raise AssertionError(
            "PDF save did not retain exactly one new rectangle appearance",
        )
    annotation = added[0]
    visible = interaction["final_rectangles"][0]
    normalized = _pdf_normalized_rectangle(annotation)
    crop = annotation["page_cropbox"]
    displayed_width = (
        crop[3] - crop[1]
        if annotation["page_rotation"] % 180
        else crop[2] - crop[0]
    )
    normalized_line_width = annotation["line_width"] / displayed_width
    color = [round(value * 255) for value in annotation["color"]]
    if (
        annotation["page_index"] != visible["page_index"]
        or any(
            abs(a - b) > 0.005
            for a, b in zip(normalized, visible["normalized_rect"])
        )
        or color != visible["color"]
        or abs(normalized_line_width - visible["normalized_line_width"])
        > 0.001
    ):
        raise AssertionError(
            "Saved PDF rectangle differs from the final native drawing",
        )
    return {
        **annotation,
        "normalized_rect": normalized,
        "normalized_line_width": normalized_line_width,
    }


def _edit_docx(
    session: AppSession,
    marker: str,
    undo_marker: str,
    timeout: float,
) -> dict[str, Any]:
    cdp = session.cdp
    _wait(
        # Never ask CDP to serialize the TipTap editor graph; Chromium rejects
        # that cyclic/deep object before Python can coerce it to bool.
        lambda: bool(cdp.evaluate("Boolean(window.__lantaiDocs?.editor)")),
        timeout=timeout,
        label="DOCX editor",
        process=session.process,
    )
    temporary_target = _focus_docx_existing_paragraph(cdp)
    _docx_insert_native_text(cdp, f" {undo_marker}")
    _wait(
        lambda: undo_marker
        in ((_docx_model_probe(cdp) or {}).get("text") or ""),
        timeout=5,
        label="DOCX temporary edit in editor model",
    )
    temporary_model = _docx_model_probe(cdp)
    if not temporary_model or temporary_model.get(
        "docxIndex",
    ) != temporary_target.get("docxIndex"):
        raise AssertionError(
            f"DOCX temporary edit escaped its original block: {temporary_model}",
        )
    undo_clicks: list[dict[str, Any]] = []
    # Chromium reports the printable key and the following Input.insertText as
    # separate history items.  A single undo can therefore remove the marker
    # while leaving its leading space dirty.  Keep using the product's visible
    # undo button until the document is genuinely back at its clean baseline.
    for attempt in range(3):
        undo_clicks.append(
            _native_click(
                cdp,
                _toolbar_button("撤销"),
                label=f"DOCX undo {attempt + 1}",
            ),
        )
        if _button_disabled(cdp, "保存"):
            break
    _wait(
        lambda: undo_marker
        not in ((_docx_model_probe(cdp) or {}).get("text") or ""),
        timeout=5,
        label="DOCX undo in editor model",
    )
    _wait(
        lambda: _button_disabled(cdp, "保存"),
        timeout=5,
        label="DOCX undo restored clean state",
    )
    final_target = _focus_docx_existing_paragraph(cdp)
    _docx_insert_native_text(cdp, f" {marker}")
    _wait(
        lambda: marker in ((_docx_model_probe(cdp) or {}).get("text") or ""),
        timeout=5,
        label="DOCX final edit in editor model",
    )
    final_model = _docx_model_probe(cdp)
    if not final_model or final_model.get("docxIndex") != final_target.get(
        "docxIndex",
    ):
        raise AssertionError(
            f"DOCX final edit escaped its original block: {final_model}",
        )
    _wait(
        lambda: not _button_disabled(cdp, "保存"),
        timeout=8,
        label="DOCX final dirty state",
    )
    save_click = _native_click(cdp, _toolbar_button("保存"), label="DOCX save")
    _wait(
        lambda: _button_disabled(cdp, "保存"),
        timeout=min(timeout, 15),
        label=f"DOCX saved clean state; renderer text={_body_text(cdp)[-240:]}",
        process=session.process,
    )
    return {
        "temporary_visible": True,
        "undo_removed_temporary": True,
        "final_visible": True,
        "temporary_target": temporary_target,
        "temporary_model": temporary_model,
        "final_target": final_target,
        "final_model": final_model,
        "undo_clicks": undo_clicks,
        "save_click": save_click,
    }


def _edit_xlsx(
    session: AppSession,
    marker: str,
    undo_marker: str,
    original_a1: str,
    timeout: float,
) -> dict[str, Any]:
    cdp = session.cdp
    _wait(
        lambda: bool(_xlsx_canvas_probe(cdp)),
        timeout=timeout,
        label="XLSX workbook canvas",
    )

    def enter(address: str, value: str) -> dict[str, Any]:
        failures: list[str] = []
        for attempt in range(1, 4):
            _xlsx_select_cell(cdp, address)
            _wait(
                lambda: _xlsx_name_box(cdp) == address,
                timeout=min(timeout, 5),
                label=f"XLSX {address} selection before edit",
            )
            try:
                result = _xlsx_replace_selected_cell(
                    cdp,
                    address,
                    value,
                    timeout,
                )
                return {**result, "focus_attempt": attempt}
            except AssertionError as exc:
                failures.append(str(exc))
                # Electron can transiently return focus to the outer window
                # between canvas selection and the first printable key. Escape
                # cancels any uncommitted edit, after which the next attempt
                # repeats the same visible user path from a fresh selection.
                cdp.key("Escape", code="Escape", virtual=27)
        raise AssertionError(
            f"XLSX could not acquire the visible edit layer after 3 attempts: {failures}",
        )

    native_selection = _xlsx_prove_native_selection(cdp, min(timeout, 10))
    temporary_edit = enter("A1", undo_marker)
    _wait(
        lambda: not _button_disabled(cdp, "保存"),
        timeout=8,
        label="XLSX dirty edit",
    )
    undo_click = _native_click(cdp, _toolbar_button("撤销"), label="XLSX undo")
    try:
        _wait(
            lambda: _button_disabled(cdp, "保存"),
            timeout=8,
            label="XLSX undo restored clean state",
        )
    except AssertionError as exc:
        _xlsx_select_cell(cdp, "A1")
        _wait(
            lambda: _xlsx_name_box(cdp) == "A1",
            timeout=min(timeout, 5),
            label="XLSX A1 selection for undo failure evidence",
        )
        visible_value = _xlsx_copy_selected_cell(cdp)
        raise AssertionError(
            f"{exc}; A1 visible value after undo={visible_value!r}; "
            f"save_disabled={_button_disabled(cdp, '保存')}",
        ) from exc
    _xlsx_select_cell(cdp, "A1")
    _wait(
        lambda: _xlsx_name_box(cdp) == "A1",
        timeout=min(timeout, 5),
        label="XLSX A1 selection for post-undo evidence",
    )
    post_undo_visible_a1 = _xlsx_copy_selected_cell(cdp)
    if post_undo_visible_a1 != original_a1:
        raise AssertionError(
            "XLSX toolbar undo reported a clean workbook but did not restore "
            f"the visible A1 value: expected={original_a1!r}, "
            f"observed={post_undo_visible_a1!r}",
        )
    final_edit = enter("A1", marker)
    _wait(
        lambda: not _button_disabled(cdp, "保存"),
        timeout=8,
        label="XLSX final edit",
    )
    zoom_click = _native_click(
        cdp,
        (
            "[...document.querySelectorAll('button.zoom-btn')]"
            ".find((b) => (b.innerText||'').trim()==='+')"
        ),
        label="XLSX zoom in",
    )
    _wait(
        lambda: "110%" in _body_text(cdp),
        timeout=8,
        label="XLSX 110% zoom",
    )
    numeric_edit = enter("B2", "123.5")
    formula_edit = enter("C2", "=B2*2")
    _wait(
        lambda: not _button_disabled(cdp, "保存"),
        timeout=8,
        label="XLSX typed numeric and formula edits",
    )
    save_click = _native_click(cdp, _toolbar_button("保存"), label="XLSX save")
    _wait(
        lambda: _button_disabled(cdp, "保存"),
        timeout=min(timeout, 15),
        label="XLSX saved clean state",
        process=session.process,
    )
    return {
        "temporary_committed": True,
        "undo_clicked": True,
        "final_committed": True,
        "native_selection": native_selection,
        "temporary_edit": temporary_edit,
        "post_undo_visible_a1": post_undo_visible_a1,
        "post_undo_expected_a1": original_a1,
        "final_edit": final_edit,
        "zoom_click": zoom_click,
        "zoom_percent": 110,
        "numeric_edit": numeric_edit,
        "formula_edit": formula_edit,
        "undo_click": undo_click,
        "save_click": save_click,
    }


def _slide_status(cdp: Cdp) -> tuple[int, int] | None:
    match = re.search(
        r"(?:幻灯片|Slide)\s*(\d+)\s*/\s*(\d+)",
        _body_text(cdp),
        flags=re.I,
    )
    return (int(match.group(1)), int(match.group(2))) if match else None


def _pptx_add_slide_button() -> str:
    return (
        "[...document.querySelectorAll('button')].find((b) => "
        "(b.innerText || '').includes('新建幻灯片') || "
        "(b.title || '').includes('新建空白幻灯片'))"
    )


def _pptx_slide_menu_button() -> str:
    return (
        "[...document.querySelectorAll('.rb-drop-wrap > button')].find((b) => "
        "(b.innerText || '').trim() === '幻灯片' && "
        "(b.title || '').trim() === '幻灯片')"
    )


def _pptx_click_add_slide(cdp: Cdp, *, label: str) -> dict[str, Any]:
    """Use either the direct add command or the current two-step ribbon menu."""

    add = _pptx_add_slide_button()
    menu_click = None
    if _native_click_target(cdp, add) is None:
        menu_click = _native_click(
            cdp,
            _pptx_slide_menu_button(),
            label=f"{label} open slide menu",
        )
        _wait(
            lambda: _native_click_target(cdp, add) is not None,
            timeout=5,
            label=f"{label} add command",
        )
    add_click = _native_click(cdp, add, label=label)
    return {
        "menu_opened": menu_click is not None,
        "menu_click": menu_click,
        "add_click": add_click,
    }


def _edit_pptx(session: AppSession, timeout: float) -> dict[str, Any]:
    cdp = session.cdp
    _wait(
        lambda: _slide_status(cdp) is not None,
        timeout=timeout,
        label="PPTX page status",
    )
    initial = _slide_status(cdp)
    assert initial is not None
    initial_count = initial[1]
    add_temporary = _pptx_click_add_slide(
        cdp,
        label="PPTX add temporary slide",
    )
    _wait(
        lambda: (_slide_status(cdp) or (0, 0))[1] == initial_count + 1,
        timeout=8,
        label="PPTX temporary slide",
    )
    undo_click = _native_click(cdp, _toolbar_button("撤销"), label="PPTX undo")
    _wait(lambda: _slide_status(cdp) == initial, timeout=8, label="PPTX undo")
    add_final = _pptx_click_add_slide(cdp, label="PPTX add final slide")
    _wait(
        lambda: (_slide_status(cdp) or (0, 0))[1] == initial_count + 1,
        timeout=8,
        label="PPTX final slide",
    )
    save_click = _native_click(cdp, _toolbar_button("保存"), label="PPTX save")
    return {
        "initial_slides": initial_count,
        "temporary_slides": initial_count + 1,
        "undo_slides": initial_count,
        "final_slides": initial_count + 1,
        "add_temporary_click": add_temporary,
        "undo_click": undo_click,
        "add_final_click": add_final,
        "save_click": save_click,
    }


def _activate_pdf_rect(cdp: Cdp) -> dict[str, float]:
    _native_click(cdp, _text_button("注释"), label="PDF annotation tab")
    rect_button = (
        "[...document.querySelectorAll('button')].find((b) => "
        "(b.title || b.innerText || '').includes('矩形'))"
    )
    _wait(
        lambda: _native_click_target(cdp, rect_button) is not None,
        timeout=10,
        label="PDF rectangle tool",
    )
    if not cdp.evaluate(
        "Boolean(document.querySelector('.pdf-draw-layer.pdf-draw-active'))",
    ):
        _native_click(cdp, rect_button, label="PDF rectangle tool")
    geometry: object = None

    def drawing_layer_ready() -> bool:
        nonlocal geometry
        geometry = cdp.evaluate(
            "(() => { const e=document.querySelector('.pdf-draw-layer.pdf-draw-active'); "
            "if(!e) return null; const r=e.getBoundingClientRect(); "
            "return {x:r.x,y:r.y,width:r.width,height:r.height} })()",
        )
        return (
            isinstance(geometry, dict)
            and float(geometry.get("width") or 0) >= 100
        )

    # The PDF page and annotation controls become visible before the drawing
    # overlay has finished mounting.  Wait for a positively sized active layer
    # so a real native drag is not dispatched into the transient empty canvas.
    _wait(drawing_layer_ready, timeout=10, label="PDF drawing layer")
    assert isinstance(geometry, dict)
    return {key: float(value) for key, value in geometry.items()}


def _draw_pdf_rect(cdp: Cdp, *, offset: float) -> None:
    geometry = _activate_pdf_rect(cdp)
    x1 = geometry["x"] + min(120 + offset, geometry["width"] * 0.25)
    y1 = geometry["y"] + min(150 + offset, geometry["height"] * 0.25)
    x2 = min(x1 + 120, geometry["x"] + geometry["width"] - 30)
    y2 = min(y1 + 70, geometry["y"] + geometry["height"] - 30)
    cdp.drag(x1, y1, x2, y2)


def _edit_pdf(session: AppSession, _timeout: float) -> dict[str, Any]:
    cdp = session.cdp
    _draw_pdf_rect(cdp, offset=0)
    _wait(
        lambda: not _button_disabled(cdp, "撤销"),
        timeout=8,
        label="PDF temporary annotation",
    )
    _wait(
        lambda: len(_pdf_visible_rectangles(cdp)) == 1,
        timeout=8,
        label="PDF temporary rectangle count",
    )
    temporary_rectangles = _pdf_visible_rectangles(cdp)
    undo_click = _native_click(cdp, _toolbar_button("撤销"), label="PDF undo")
    _wait(lambda: _button_disabled(cdp, "撤销"), timeout=8, label="PDF undo")
    _wait(
        lambda: _pdf_visible_rectangles(cdp) == [],
        timeout=8,
        label="PDF undo rectangle count",
    )
    undo_rectangles = _pdf_visible_rectangles(cdp)
    _draw_pdf_rect(cdp, offset=45)
    _wait(
        lambda: not _button_disabled(cdp, "保存"),
        timeout=8,
        label="PDF final annotation",
    )
    _wait(
        lambda: len(_pdf_visible_rectangles(cdp)) == 1,
        timeout=8,
        label="PDF final rectangle count",
    )
    final_rectangles = _pdf_visible_rectangles(cdp)
    save_click = _native_click(cdp, _toolbar_button("保存"), label="PDF save")
    return {
        "temporary_annotation": True,
        "undo_cleared_stack": True,
        "final_annotation": True,
        "temporary_rectangles": temporary_rectangles,
        "undo_rectangles": undo_rectangles,
        "final_rectangles": final_rectangles,
        "undo_click": undo_click,
        "save_click": save_click,
    }


def _disk_state(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return {"text": _docx_text(path)}
    if suffix == ".xlsx":
        return _xlsx_cells(path)
    if suffix == ".pptx":
        return {"slide_count": _pptx_slide_count(path)}
    if suffix == ".pdf":
        return _pdf_annotation_readback(path)
    raise AssertionError(f"unsupported Office fixture: {path}")


def _assert_readback(
    *,
    suffix: str,
    before: dict[str, Any],
    after: dict[str, Any],
    marker: str,
    undo_marker: str,
    xlsx_native_number_waiver_reason: str | None = None,
) -> list[dict[str, Any]]:
    waivers: list[dict[str, Any]] = []
    if suffix == ".docx":
        if marker not in after["text"] or undo_marker in after["text"]:
            raise AssertionError(
                "DOCX disk readback did not preserve only the final edit",
            )
    elif suffix == ".xlsx":
        if after["a1"] != marker or after["a1"] == undo_marker:
            raise AssertionError(f"XLSX A1 readback mismatch: {after['a1']!r}")
        if after["b2"] != 123.5 or after["b2_type"] != "n":
            if (
                xlsx_native_number_waiver_reason
                and str(after["b2"]) == "123.5"
                and after["b2_type"] == "s"
            ):
                waivers.append(
                    {
                        "id": XLSX_NATIVE_NUMBER_WAIVER_ID,
                        "reason": xlsx_native_number_waiver_reason,
                        "observed_value": after["b2"],
                        "observed_type": after["b2_type"],
                    },
                )
            else:
                raise AssertionError(
                    "XLSX numeric input lost native semantics: "
                    f"{after['b2']!r}/{after['b2_type']!r}",
                )
        if after["c2"] != "=B2*2" or after["c2_type"] != "f":
            raise AssertionError(
                f"XLSX formula input lost native semantics: {after['c2']!r}/{after['c2_type']!r}",
            )
    elif suffix == ".pptx":
        if after["slide_count"] != before["slide_count"] + 1:
            raise AssertionError(
                f"PPTX slide readback mismatch: {before} -> {after}",
            )
    elif suffix == ".pdf":
        if (
            after["square_annotation_count"]
            != before["square_annotation_count"] + 1
        ):
            raise AssertionError(
                f"PDF annotation readback mismatch: {before} -> {after}",
            )
    return waivers


def _pdf_fresh_pixel_readback(
    cdp: Cdp,
    fixture: Path,
    annotation: dict[str, Any],
    output: Path,
    timeout: float,
) -> dict[str, Any]:
    from PIL import Image

    expected_hash = _sha256(fixture)
    renderer_hash = cdp.evaluate(
        "(async () => { const data=await window.pdfApi.readFile("
        + json.dumps(str(fixture))
        + "); const hash=await crypto.subtle.digest('SHA-256',new Uint8Array(data)); "
        "return [...new Uint8Array(hash)].map(v=>v.toString(16).padStart(2,'0')).join(''); })()",
        await_promise=True,
    )
    if renderer_hash != expected_hash:
        raise AssertionError(
            "Fresh PDF renderer read different saved file bytes",
        )
    page_index = int(annotation["page_index"])
    canvas: dict[str, Any] | None = None

    def painted_canvas() -> bool:
        nonlocal canvas
        value = cdp.evaluate(
            "(() => { const e=document.querySelector("
            + json.dumps(
                f'.pdf-row[data-idx="{page_index}"] .pdf-page-content > canvas',
            )
            + "); if(!(e instanceof HTMLCanvasElement)||e.width<=0||e.height<=0) return null; "
            "const r=e.getBoundingClientRect(),s=getComputedStyle(e); "
            "if(r.width<=0||r.height<=0||s.display==='none'||s.visibility==='hidden') return null; "
            "return {data:e.toDataURL('image/png'),width:e.width,height:e.height,"
            "rect:{left:r.left,top:r.top,width:r.width,height:r.height},"
            "viewport:[window.innerWidth,window.innerHeight]}; })()",
        )
        canvas = value if isinstance(value, dict) else None
        return canvas is not None

    _wait(painted_canvas, timeout=timeout, label="PDF saved-page canvas")
    assert canvas is not None
    canvas_data: dict[str, Any] = dict(canvas)
    encoded = str(canvas_data.get("data") or "")
    if not encoded.startswith("data:image/png;base64,"):
        raise AssertionError(
            "Fresh PDF saved-page canvas did not return PNG pixels",
        )
    canvas_bytes = base64.b64decode(encoded.partition(",")[2], validate=True)
    with Image.open(io.BytesIO(canvas_bytes)) as source:
        image = source.convert("RGB")
    if list(image.size) != [canvas_data["width"], canvas_data["height"]]:
        raise AssertionError("Fresh PDF canvas dimensions differ from its PNG")
    normalized = annotation["normalized_rect"]
    color = [round(value * 255) for value in annotation["color"]]
    canvas_rect = [
        normalized[0] * image.width,
        normalized[1] * image.height,
        normalized[2] * image.width,
        normalized[3] * image.height,
    ]
    canvas_proof = _pdf_rectangle_pixel_proof(
        image,
        canvas_rect,
        color,
        annotation["normalized_line_width"] * image.width,
    )
    canvas_path = output.with_name(f"{output.stem}-saved-page.png")
    fd = os.open(canvas_path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    os.fchmod(fd, 0o600)
    with os.fdopen(fd, "wb") as handle:
        handle.write(canvas_bytes)
    cdp.screenshot(output)
    with Image.open(output) as source:
        screenshot = source.convert("RGB")
    viewport_width, viewport_height = map(float, canvas_data["viewport"])
    if viewport_width <= 0 or viewport_height <= 0:
        raise AssertionError("Fresh PDF screenshot viewport is invalid")
    scale_x, scale_y = (
        screenshot.width / viewport_width,
        screenshot.height / viewport_height,
    )
    if abs(scale_x - scale_y) > 0.02:
        raise AssertionError(
            "Fresh PDF screenshot has unsupported viewport scaling",
        )
    bounds = canvas_data["rect"]
    screenshot_rect = [
        (bounds["left"] + normalized[0] * bounds["width"]) * scale_x,
        (bounds["top"] + normalized[1] * bounds["height"]) * scale_y,
        (bounds["left"] + normalized[2] * bounds["width"]) * scale_x,
        (bounds["top"] + normalized[3] * bounds["height"]) * scale_y,
    ]
    screenshot_proof = _pdf_rectangle_pixel_proof(
        screenshot,
        screenshot_rect,
        color,
        annotation["normalized_line_width"] * bounds["width"] * scale_x,
    )
    if _sha256(fixture) != expected_hash:
        raise AssertionError(
            "Saved PDF changed during fresh renderer pixel readback",
        )
    return {
        "file_sha256": expected_hash,
        "renderer_file_sha256": renderer_hash,
        "square_annotation_count": _pdf_square_annotation_count(fixture),
        "visible_square_count": 1,
        "annotation": annotation,
        "canvas": str(canvas_path),
        "canvas_sha256": _sha256(canvas_path),
        "canvas_geometry": {
            "rect": bounds,
            "viewport": canvas_data["viewport"],
        },
        "canvas_pixel_proof": canvas_proof,
        "screenshot_pixel_proof": screenshot_proof,
    }


def _fresh_renderer_readback(
    *,
    app: Path,
    fixture: Path,
    marker: str,
    output: Path,
    timeout: float,
    pdf_annotation: dict[str, Any] | None = None,
) -> dict[str, Any]:
    session = _launch(app, fixture, output.with_suffix(".log"), timeout)
    try:
        process_pid = session.process.pid
        user_data_dir = str(session.root / "user-data")
        text = _body_text(session.cdp)
        if fixture.suffix.lower() == ".docx" and marker not in text:
            raise AssertionError(
                "fresh DOCX renderer did not show the saved marker",
            )
        pptx_slide_count = None
        if fixture.suffix.lower() == ".pptx":
            status = _slide_status(session.cdp)
            if status is None or status[1] != _pptx_slide_count(fixture):
                raise AssertionError(
                    "fresh PPTX renderer did not show the added slide",
                )
            pptx_slide_count = status[1]
        pdf_annotations = None
        if fixture.suffix.lower() == ".pdf":
            if pdf_annotation is None:
                raise AssertionError(
                    "Fresh PDF readback requires the saved final rectangle appearance",
                )
            pdf_annotations = _pdf_fresh_pixel_readback(
                session.cdp,
                fixture,
                pdf_annotation,
                output,
                timeout,
            )
        if fixture.suffix.lower() == ".xlsx":
            native_selection = _xlsx_prove_native_selection(
                session.cdp,
                min(timeout, 10),
            )
            value = _xlsx_copy_selected_cell(session.cdp)
            if value != marker:
                raise AssertionError(
                    f"fresh XLSX renderer A1 mismatch: {value!r}",
                )
        else:
            native_selection = None
        if pdf_annotations is None:
            session.cdp.screenshot(output)
        result = {
            "title": str(session.cdp.evaluate("document.title") or ""),
            "screenshot": str(output),
            "screenshot_sha256": _sha256(output),
            "workbench": session.workbench,
            "renderer_identity": session.renderer_identity,
            "docx_marker_visible": marker in text
            if fixture.suffix.lower() == ".docx"
            else None,
            "pptx_slide_count": pptx_slide_count,
            "pdf_annotations": pdf_annotations,
            "process_pid": process_pid,
            "user_data_dir": user_data_dir,
            "native_selection": native_selection,
            "xlsx_clipboard_readback": value
            if fixture.suffix.lower() == ".xlsx"
            else None,
        }
    finally:
        _stop_preserving_primary(session)
    result["process_settlement"] = session.process_settlement
    return result


def _run_one(
    *,
    app: Path,
    source: Path,
    work_dir: Path,
    case_id: str,
    timeout: float,
    xlsx_native_number_waiver_reason: str | None = None,
) -> dict[str, Any]:
    source = source.expanduser().resolve(strict=True)
    work_dir = work_dir.expanduser().resolve()
    suffix = source.suffix.lower()
    if suffix not in SUPPORTED:
        raise AssertionError(f"unsupported Office fixture: {source}")
    work_dir.mkdir(parents=True, exist_ok=True)
    fixture = work_dir / f"roundtrip{suffix}"
    shutil.copy2(source, fixture)
    marker = f"LANTAI_2_0_FINAL_{case_id}_{suffix[1:].upper()}"
    undo_marker = f"LANTAI_2_0_UNDO_{case_id}_{suffix[1:].upper()}"
    before_hash = _sha256(fixture)
    before_state = _disk_state(fixture)
    session = _launch(
        app,
        fixture,
        work_dir / f"edit-{suffix[1:]}.log",
        timeout,
    )
    edit_pid = session.process.pid
    edit_user_data = str(session.root / "user-data")
    try:
        if suffix == ".docx":
            interaction = _edit_docx(session, marker, undo_marker, timeout)
        elif suffix == ".xlsx":
            interaction = _edit_xlsx(
                session,
                marker,
                undo_marker,
                str(before_state["a1"]),
                timeout,
            )
        elif suffix == ".pptx":
            interaction = _edit_pptx(session, timeout)
        else:
            interaction = _edit_pdf(session, timeout)
        _wait(
            lambda: _sha256(fixture) != before_hash,
            timeout=timeout,
            label=f"{suffix} file hash change",
            process=session.process,
        )
        session.cdp.screenshot(work_dir / f"saved-{suffix[1:]}.png")
    finally:
        _stop_preserving_primary(session)
    if session.process.poll() is None:
        raise AssertionError(
            "edit process remained alive after the stop boundary",
        )
    edit_exit_code = session.process.returncode
    after_hash = _sha256(fixture)
    after_state = _disk_state(fixture)
    waivers = _assert_readback(
        suffix=suffix,
        before=before_state,
        after=after_state,
        marker=marker,
        undo_marker=undo_marker,
        xlsx_native_number_waiver_reason=xlsx_native_number_waiver_reason,
    )
    pdf_annotation = (
        _pdf_saved_rectangle(before_state, after_state, interaction)
        if suffix == ".pdf"
        else None
    )
    renderer = _fresh_renderer_readback(
        app=app,
        fixture=fixture,
        marker=marker,
        output=work_dir / f"reopened-{suffix[1:]}.png",
        timeout=timeout,
        pdf_annotation=pdf_annotation,
    )
    if _sha256(fixture) != after_hash:
        raise AssertionError(
            "Office fixture changed during the fresh renderer readback",
        )
    if int(renderer["process_pid"]) == edit_pid:
        raise AssertionError("fresh renderer reused the edit process PID")
    if str(renderer["user_data_dir"]) == edit_user_data:
        raise AssertionError(
            "fresh renderer reused the edit process user-data directory",
        )
    return {
        "source": str(source),
        "fixture_copy": str(fixture),
        "suffix": suffix,
        "marker": marker,
        "undo_marker": undo_marker,
        "before_sha256": before_hash,
        "after_sha256": after_hash,
        "hash_changed": after_hash != before_hash,
        "before_state": before_state,
        "after_state": after_state,
        "interaction": interaction,
        "workbench": session.workbench,
        "renderer_identity": session.renderer_identity,
        "edit_process_pid": edit_pid,
        "edit_process_exit_code": edit_exit_code,
        "edit_process_settlement": session.process_settlement,
        "edit_user_data_dir": edit_user_data,
        "fresh_process_is_distinct": True,
        "fresh_user_data_is_distinct": True,
        "fresh_renderer": renderer,
        "waivers": waivers,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--app", type=Path, required=True)
    parser.add_argument("--fixture", type=Path, action="append", required=True)
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--case-id", default="case-01")
    parser.add_argument("--timeout", type=float, default=60)
    parser.add_argument("--evidence-output", type=Path, required=True)
    parser.add_argument(
        "--waive-xlsx-native-number-reason",
        help=(
            "Explicitly waive only GenOffice's string-typed XLSX numeric input; "
            "the persisted value must still be exactly 123.5 and every other check remains strict"
        ),
    )
    args = parser.parse_args(argv)
    app = args.app.expanduser().resolve(strict=True)
    fixtures = [
        path.expanduser().resolve(strict=True) for path in args.fixture
    ]
    suffixes = [path.suffix.lower() for path in fixtures]
    if len(fixtures) != 4 or set(suffixes) != SUPPORTED:
        raise SystemExit(
            "Exactly one DOCX, XLSX, PPTX, and PDF fixture is required",
        )
    args.work_dir = args.work_dir.expanduser().resolve()
    args.work_dir.mkdir(parents=True, exist_ok=True)
    started = time.time()
    metadata = _app_metadata(app)
    results = [
        _run_one(
            app=app,
            source=fixture,
            work_dir=args.work_dir / fixture.suffix[1:].lower(),
            case_id=args.case_id,
            timeout=args.timeout,
            xlsx_native_number_waiver_reason=args.waive_xlsx_native_number_reason,
        )
        for fixture in fixtures
    ]
    waivers = [waiver for result in results for waiver in result["waivers"]]
    if _app_metadata(app) != metadata:
        raise AssertionError(
            "Office app component artifacts changed during the roundtrip",
        )
    payload = {
        "schema": "com.lantai.packaged-office-edit-roundtrip.v1",
        "ok": True,
        "status": "passed_with_waiver" if waivers else "passed",
        "started_at": started,
        "finished_at": time.time(),
        "app": metadata,
        "case_id": args.case_id,
        "results": results,
        "waivers": waivers,
    }
    args.evidence_output.parent.mkdir(parents=True, exist_ok=True)
    args.evidence_output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({"ok": True, "evidence": str(args.evidence_output)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
