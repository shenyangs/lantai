#!/usr/bin/env python3
"""Verify that a packaged Lantai app can execute a bundled Skill script."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import plistlib
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from PIL import Image
from pptx import Presentation
from reportlab.pdfgen import canvas

try:
    from scripts.packaged_runtime_evidence import SCHEMA as EVIDENCE_SCHEMA
except ModuleNotFoundError:  # direct script execution from scripts/
    from packaged_runtime_evidence import SCHEMA as EVIDENCE_SCHEMA

ENGINE_VERSION = "4.0.1"
_HOST_ENV_KEYS = (
    "CONDA_PREFIX",
    "DYLD_FALLBACK_LIBRARY_PATH",
    "DYLD_LIBRARY_PATH",
    "LD_LIBRARY_PATH",
    "NODE_PATH",
    "NPM_CONFIG_PREFIX",
    "PYTHONHOME",
    "PYTHONPATH",
    "PYTHONUSERBASE",
    "VIRTUAL_ENV",
)


def _required_file(path: Path, label: str) -> Path:
    if not path.is_file():
        raise RuntimeError(f"Packaged {label} is missing: {path}")
    return path


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sanitized_env(
    extra: dict[str, str] | None = None,
    *,
    clean_path: bool = True,
) -> dict[str, str]:
    env = dict(os.environ)
    for key in _HOST_ENV_KEYS:
        env.pop(key, None)
    env["PYTHONNOUSERSITE"] = "1"
    env["PYTHONDONTWRITEBYTECODE"] = "1"
    if clean_path:
        env["PATH"] = "/usr/bin:/bin"
    if extra:
        env.update(extra)
    return env


def _run_isolated(
    command: list[str],
    *,
    env: dict[str, str],
    timeout: int,
) -> subprocess.CompletedProcess[str]:
    with tempfile.TemporaryDirectory(prefix="lantai-packaged-probe-cwd-") as raw:
        probe_root = Path(raw)
        isolated_env = dict(env)
        isolated_env.update(
            {
                "LANTAI_WORKING_DIR": str(probe_root / "working"),
                "LANTAI_SECRET_DIR": str(probe_root / "working.secret"),
                "LANTAI_DESKTOP_USER_DATA_DIR": str(probe_root / "desktop-data"),
                "LANTAI_DISABLE_KEYRING": "1",
            }
        )
        return subprocess.run(
            command,
            cwd=probe_root,
            env=isolated_env,
            capture_output=True,
            text=True,
            timeout=timeout,
            check=False,
        )


def _skill_resource_manifest(
    root: Path,
    *,
    ignore_runtime_caches: bool,
) -> dict[str, str]:
    manifest: dict[str, str] = {}
    for path in sorted(root.rglob("*")):
        if not path.is_file():
            continue
        relative = path.relative_to(root)
        if ignore_runtime_caches and (
            "__pycache__" in relative.parts
            or path.suffix.lower() == ".pyc"
            or path.name == ".DS_Store"
        ):
            continue
        manifest[relative.as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    return manifest


def _compare_skill_resources(source: Path, packaged: Path) -> dict[str, object]:
    source_manifest = _skill_resource_manifest(
        source,
        ignore_runtime_caches=True,
    )
    packaged_manifest = _skill_resource_manifest(
        packaged,
        ignore_runtime_caches=False,
    )
    source_paths = set(source_manifest)
    packaged_paths = set(packaged_manifest)
    missing = sorted(source_paths - packaged_paths)
    unexpected = sorted(packaged_paths - source_paths)
    mismatched = sorted(
        path
        for path in source_paths & packaged_paths
        if source_manifest[path] != packaged_manifest[path]
    )
    source_manifest_sha256 = hashlib.sha256(
        json.dumps(
            source_manifest,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()
    packaged_manifest_sha256 = hashlib.sha256(
        json.dumps(
            packaged_manifest,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        ).encode("utf-8")
    ).hexdigest()
    return {
        "status": "failed" if missing or unexpected or mismatched else "ok",
        "source_count": len(source_manifest),
        "packaged_count": len(packaged_manifest),
        "missing": missing,
        "unexpected": unexpected,
        "mismatched": mismatched,
        "source_manifest_sha256": source_manifest_sha256,
        "packaged_manifest_sha256": packaged_manifest_sha256,
    }


def _packaged_identity(
    app: Path,
    *,
    resources: Path,
    backend: Path,
) -> dict[str, object]:
    provenance_path = _required_file(
        resources / "console" / "dist" / "release-provenance.json",
        "release provenance",
    )
    provenance = json.loads(provenance_path.read_text(encoding="utf-8"))
    info_path = _required_file(app / "Contents" / "Info.plist", "Info.plist")
    with info_path.open("rb") as handle:
        info = plistlib.load(handle)
    executable_name = str(info.get("CFBundleExecutable") or "").strip()
    if not executable_name:
        raise RuntimeError("Packaged Info.plist has no CFBundleExecutable")
    executable = _required_file(
        app / "Contents" / "MacOS" / executable_name,
        "app executable",
    )
    app_asar = _required_file(resources / "app.asar", "app.asar")
    return {
        "source": {
            "git_sha": provenance.get("gitSha"),
            "source_branch": provenance.get("sourceBranch"),
            "build_time": provenance.get("buildTime"),
            "canonical_main": provenance.get("canonicalMain"),
        },
        "app": {
            "name": app.name,
            "path": str(app),
            "version": info.get("CFBundleShortVersionString"),
            "bundle_id": info.get("CFBundleIdentifier"),
        },
        "artifacts": {
            "app_executable": {
                "path": f"Contents/MacOS/{executable_name}",
                "sha256": _sha256_file(executable),
            },
            "app_asar": {
                "path": "Contents/Resources/app.asar",
                "sha256": _sha256_file(app_asar),
            },
            "backend_executable": {
                "path": (
                    "Contents/Resources/lantai-server/Lantai Backend.app/"
                    "Contents/MacOS/lantai-server"
                ),
                "sha256": _sha256_file(backend),
            },
        },
    }


def _packaged_skill_runner(
    backend_app: Path,
    skill_name: str,
    script_name: str,
) -> Path:
    """Resolve PyInstaller BUNDLE data under the helper's Resources root."""

    return (
        backend_app
        / "Contents"
        / "Resources"
        / "lantai"
        / "agents"
        / "skills"
        / skill_name
        / "scripts"
        / script_name
    )


def run_smoke(app: Path, *, timeout: int) -> dict[str, object]:
    app = app.expanduser().resolve()
    resources = app / "Contents" / "Resources"
    backend_root = resources / "lantai-server"
    backend_app = backend_root / "Lantai Backend.app"
    backend = _required_file(
        backend_app / "Contents" / "MacOS" / "lantai-server",
        "backend",
    )
    packaged_identity = _packaged_identity(
        app,
        resources=resources,
        backend=backend,
    )
    officecli = _required_file(backend_root / "bin" / "officecli", "OfficeCLI")
    electron = _required_file(
        app / "Contents" / "MacOS" / app.stem,
        "Electron executable",
    )
    modules = resources / "app.asar.unpacked" / "node_modules"
    manifest = _required_file(
        modules / "pptxgenjs" / "package.json",
        "presentation engine manifest",
    )
    engine_version = str(json.loads(manifest.read_text(encoding="utf-8"))["version"])
    if engine_version != ENGINE_VERSION:
        raise RuntimeError(
            "Packaged presentation engine version mismatch: "
            f"expected {ENGINE_VERSION}, got {engine_version}",
        )

    import_probe_script = _required_file(
        Path(__file__).with_name("packaged_skill_import_probe.py"),
        "Skill import probe",
    )
    packaged_skills_root = (
        backend_app / "Contents" / "Resources" / "lantai" / "agents" / "skills"
    )
    source_skills_root = (
        Path(__file__).resolve().parents[1] / "src" / "lantai" / "agents" / "skills"
    )
    resource_parity = _compare_skill_resources(
        source_skills_root,
        packaged_skills_root,
    )
    if resource_parity["status"] != "ok":
        raise RuntimeError(
            "Packaged Skill resources differ from the release source: "
            + json.dumps(resource_parity, ensure_ascii=False, sort_keys=True),
        )
    probe_timeout = max(30, min(timeout, 120))
    per_import_timeout = max(10, min(60, probe_timeout - 5))
    import_probe = _run_isolated(
        [
            str(backend),
            "--skills-root",
            str(packaged_skills_root),
            "--per-import-timeout",
            str(per_import_timeout),
        ],
        env=_sanitized_env(
            {"LANTAI_RUN_SKILL_SCRIPT_PATH": str(import_probe_script)},
        ),
        timeout=probe_timeout,
    )
    if import_probe.returncode != 0:
        detail = (import_probe.stderr or import_probe.stdout or "no output").strip()
        raise RuntimeError(
            "Packaged Skill import probe failed with exit code "
            f"{import_probe.returncode}: {detail[-4000:]}",
        )
    import_report = json.loads(import_probe.stdout)

    runtime_probe_script = _required_file(
        Path(__file__).with_name("packaged_runtime_import_probe.py"),
        "runtime import probe",
    )
    source_root = Path(__file__).resolve().parents[1] / "src" / "lantai"
    runtime_probe = _run_isolated(
        [
            str(backend),
            "--source-root",
            str(source_root),
            "--per-import-timeout",
            str(per_import_timeout),
        ],
        env=_sanitized_env(
            {"LANTAI_RUN_SKILL_SCRIPT_PATH": str(runtime_probe_script)},
        ),
        timeout=probe_timeout,
    )
    if runtime_probe.returncode != 0:
        detail = (runtime_probe.stderr or runtime_probe.stdout or "no output").strip()
        raise RuntimeError(
            "Packaged runtime import probe failed with exit code "
            f"{runtime_probe.returncode}: {detail[-4000:]}",
        )
    runtime_report = json.loads(runtime_probe.stdout)
    research_probe_script = _required_file(
        Path(__file__).with_name("packaged_research_export_probe.py"),
        "research export probe",
    )
    research_probe = _run_isolated(
        [str(backend)],
        env=_sanitized_env(
            {"LANTAI_RUN_SKILL_SCRIPT_PATH": str(research_probe_script)},
        ),
        timeout=probe_timeout,
    )
    if research_probe.returncode != 0:
        detail = (research_probe.stderr or research_probe.stdout or "no output").strip()
        raise RuntimeError(
            "Packaged research export probe failed with exit code "
            f"{research_probe.returncode}: {detail[-4000:]}",
        )
    research_report = json.loads(research_probe.stdout)
    cli_help_failures: list[dict[str, object]] = []
    for command in sorted(runtime_report["lazy_cli_commands"]):
        cli_help = _run_isolated(
            [str(backend), command, "--help"],
            env=_sanitized_env(),
            timeout=probe_timeout,
        )
        if cli_help.returncode != 0:
            cli_help_failures.append(
                {
                    "command": command,
                    "exit_code": cli_help.returncode,
                    "output": (
                        cli_help.stderr or cli_help.stdout or "no output"
                    ).strip()[-1000:],
                }
            )
    if cli_help_failures:
        raise RuntimeError(
            "Packaged lazy CLI surface failed: "
            + json.dumps(cli_help_failures, ensure_ascii=False, sort_keys=True),
        )

    node_probe_script = _required_file(
        Path(__file__).with_name("packaged_skill_node_probe.cjs"),
        "Skill Node import probe",
    )
    node_probe = _run_isolated(
        [
            str(electron),
            str(node_probe_script),
            "--skills-root",
            str(packaged_skills_root),
            "--node-modules",
            str(modules),
        ],
        env=_sanitized_env({"ELECTRON_RUN_AS_NODE": "1"}),
        timeout=probe_timeout,
    )
    if node_probe.returncode != 0:
        detail = (node_probe.stderr or node_probe.stdout or "no output").strip()
        raise RuntimeError(
            "Packaged Skill Node import probe failed with exit code "
            f"{node_probe.returncode}: {detail[-4000:]}",
        )
    node_report = json.loads(node_probe.stdout)

    shell_scripts = sorted(packaged_skills_root.glob("*/scripts/**/*.sh"))
    shell_failures: list[dict[str, object]] = []
    for shell_script in shell_scripts:
        checked = _run_isolated(
            ["/bin/sh", "-n", str(shell_script)],
            env=_sanitized_env(),
            timeout=probe_timeout,
        )
        if checked.returncode != 0:
            shell_failures.append(
                {
                    "script": shell_script.relative_to(
                        packaged_skills_root,
                    ).as_posix(),
                    "exit_code": checked.returncode,
                    "output": (checked.stderr or checked.stdout or "no output").strip()[
                        -1000:
                    ],
                }
            )
    if shell_failures:
        raise RuntimeError(
            "Packaged Skill shell syntax probe failed: "
            + json.dumps(shell_failures, ensure_ascii=False, sort_keys=True),
        )

    office_env = _sanitized_env(
        {
            "PATH": os.pathsep.join(
                [str(officecli.parent), "/usr/bin", "/bin"],
            ),
            "OFFICECLI_PATH": str(officecli),
            "OFFICECLI_SKIP_UPDATE": "1",
        },
    )
    office_result = _run_isolated(
        ["/bin/sh", "-c", "command -v officecli && officecli --version"],
        env=office_env,
        timeout=probe_timeout,
    )
    if office_result.returncode != 0:
        detail = (office_result.stderr or office_result.stdout or "no output").strip()
        raise RuntimeError(
            "Packaged OfficeCLI shell resolution failed with exit code "
            f"{office_result.returncode}: {detail[-2000:]}",
        )

    with tempfile.TemporaryDirectory(prefix="lantai-packaged-skill-smoke-") as raw:
        workspace = Path(raw)
        stable_presentations: dict[str, dict[str, object]] = {}
        for locale, language, user_request, fixture_material in (
            (
                "pptx-zh",
                "zh-CN",
                "验证打包后的澜台稳定演示文稿生成器",
                "# 打包态验证\n\n## 结论\n澜台已打包的演示文稿运行时可以根据明确材料生成一页可编辑演示文稿。",
            ),
            (
                "pptx-en",
                "en-US",
                "Verify packaged stable presentation generation",
                "# Packaged runtime verification\n\n## Conclusion\nThe packaged Lantai presentation runtime can generate one editable slide from explicit source material.",
            ),
        ):
            stable_runner = _required_file(
                _packaged_skill_runner(
                    backend_app,
                    locale,
                    "run_presentation_v3.py",
                ),
                f"{locale} stable PPT Skill runner",
            )
            stable_request = workspace / f"request-{locale}.json"
            stable_output = workspace / f"stable-result-{locale}.pptx"
            stable_request.write_text(
                json.dumps(
                    {
                        "schema_version": "1.0",
                        "conversation_id": f"packaged-presentation-smoke-{locale}",
                        "user_request": user_request,
                        "audience": "release verification",
                        "language": language,
                        "slide_count": 1,
                        "mode": "stable",
                        "theme_id": "lantai-ink",
                        "planning_mode": "deterministic_fixture",
                        "visual_verification_required": True,
                        "vision_critique_required": False,
                        "materials": [
                            {
                                "id": f"packaged-smoke-material-{locale}",
                                "name": f"packaged-smoke-{locale}.md",
                                "content": fixture_material,
                                "media_type": "text/markdown",
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            stable_completed = subprocess.run(
                [
                    str(backend),
                    "--request",
                    stable_request.name,
                    "--job-root",
                    f".presentation-jobs-{locale}",
                    "--output",
                    stable_output.name,
                ],
                cwd=workspace,
                env=_sanitized_env(
                    {
                        "ELECTRON_RUN_AS_NODE": "1",
                        "LANTAI_NODE_BIN": str(electron),
                        "LANTAI_NODE_MODULES": str(modules),
                        "LANTAI_RUN_SKILL_SCRIPT_PATH": str(stable_runner),
                    },
                ),
                capture_output=True,
                text=True,
                timeout=max(30, timeout),
                check=False,
            )
            if stable_completed.returncode != 0:
                detail = (
                    stable_completed.stderr or stable_completed.stdout or "no output"
                ).strip()
                raise RuntimeError(
                    f"Packaged {locale} stable presentation runtime failed "
                    f"with exit code {stable_completed.returncode}: "
                    f"{detail[-2000:]}",
                )
            _required_file(
                stable_output,
                f"{locale} stable Skill-generated presentation",
            )
            stable_presentation = Presentation(stable_output)
            if len(stable_presentation.slides) != 1:
                raise RuntimeError(
                    f"Packaged {locale} stable presentation has an invalid slide count",
                )
            stable_presentations[locale] = {
                "slide_count": len(stable_presentation.slides),
                "sha256": _sha256_file(stable_output),
            }

        pdf_source = workspace / "one-page.pdf"
        pdf = canvas.Canvas(str(pdf_source))
        pdf.drawString(72, 720, "Lantai packaged PDFium verification")
        pdf.save()
        pdf_renderers: dict[str, dict[str, object]] = {}
        for locale in ("pdf-zh", "pdf-en"):
            converter = _required_file(
                _packaged_skill_runner(
                    backend_app,
                    locale,
                    "render_pages.py",
                ),
                f"{locale} PDF renderer",
            )
            render_dir = workspace / f"render-{locale}"
            rendered = subprocess.run(
                [
                    str(backend),
                    "--path",
                    str(pdf_source),
                    "--output-dir",
                    str(render_dir),
                ],
                cwd=workspace,
                env=_sanitized_env(
                    {"LANTAI_RUN_SKILL_SCRIPT_PATH": str(converter)},
                ),
                capture_output=True,
                text=True,
                timeout=probe_timeout,
                check=False,
            )
            if rendered.returncode != 0:
                detail = (rendered.stderr or rendered.stdout or "no output").strip()
                raise RuntimeError(
                    f"Packaged {locale} PDFium render failed with exit code "
                    f"{rendered.returncode}: {detail[-2000:]}",
                )
            image_path = _required_file(
                render_dir / "page-001.png",
                f"{locale} PDFium output",
            )
            with Image.open(image_path) as image:
                if image.format != "PNG" or image.width <= 0 or image.height <= 0:
                    raise RuntimeError(
                        f"Packaged {locale} PDFium output is invalid",
                    )
                pdf_renderers[locale] = {
                    "page_count": 1,
                    "width": image.width,
                    "height": image.height,
                    "sha256": _sha256_file(image_path),
                }

        post_probe_parity = _compare_skill_resources(
            source_skills_root,
            packaged_skills_root,
        )
        if (
            post_probe_parity["status"] != "ok"
            or post_probe_parity["packaged_manifest_sha256"]
            != resource_parity["packaged_manifest_sha256"]
        ):
            raise RuntimeError(
                "Packaged Skill probes mutated the signed bundle: "
                + json.dumps(
                    post_probe_parity,
                    ensure_ascii=False,
                    sort_keys=True,
                ),
            )
        return {
            "schema": EVIDENCE_SCHEMA,
            "status": "ok",
            "generated_at": datetime.now(timezone.utc)
            .isoformat()
            .replace(
                "+00:00",
                "Z",
            ),
            **packaged_identity,
            "skill": "pptx-zh",
            "script": "scripts/run_presentation_v3.py",
            "engine_version": engine_version,
            "runtime_import_probe": {
                "module_count": runtime_report["module_count"],
                "failed_count": runtime_report["failed_count"],
                "cli_help_count": len(runtime_report["lazy_cli_commands"]),
                "per_import_timeout_seconds": per_import_timeout,
            },
            "research_export_probe": research_report,
            "node_import_probe": {
                "script_count": node_report["scriptCount"],
                "dependency_count": node_report["dependencyCount"],
                "failed_count": node_report["failedCount"],
            },
            "shell_syntax_count": len(shell_scripts),
            "skill_resource_parity": {
                "source_count": resource_parity["source_count"],
                "packaged_count": resource_parity["packaged_count"],
                "source_manifest_sha256": resource_parity["source_manifest_sha256"],
                "packaged_manifest_sha256": resource_parity["packaged_manifest_sha256"],
                "bundle_unchanged_after_probes": True,
            },
            "skill_import_probe": {
                "script_count": import_report["script_count"],
                "import_count": import_report["import_count"],
                "failed_count": import_report["failed_count"],
            },
            "stable_presentations": stable_presentations,
            "pdf_renderers": pdf_renderers,
            "sanitized_environment": {
                "clean_path": True,
                "dont_write_bytecode": True,
                "removed": list(_HOST_ENV_KEYS),
                "isolated_probe_cwd": True,
                "isolated_runtime_roots": True,
                "keyring_disabled": True,
            },
            "slide_count": stable_presentations["pptx-zh"]["slide_count"],
            "officecli": (office_result.stdout or "").strip().splitlines()[-1],
        }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--app", required=True, type=Path)
    parser.add_argument("--timeout", type=int, default=120)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    rendered = (
        json.dumps(
            run_smoke(args.app, timeout=args.timeout),
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n"
    )
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered, encoding="utf-8")
    print(rendered, end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
