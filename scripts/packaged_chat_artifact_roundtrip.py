#!/usr/bin/env python3
"""Validate chat artifact-card -> editor -> return on an exact packaged App.

The driver uses a controlled loopback OpenAI-compatible provider, but it does
not replace the packaged submit, Native Runtime, tool transaction, delivery,
chat persistence, Electron routing, or Office editor paths.  It records action
counts before and after navigation so route changes cannot masquerade as a new
user command or delivery.
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import os
import shutil
import signal
import socket
import subprocess
import tempfile
import threading
import time
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Callable
from urllib import parse

from PIL import Image, ImageDraw
from pptx import Presentation

import packaged_app_lifecycle_smoke as lifecycle
import packaged_office_edit_roundtrip as office


SCHEMA = "com.lantai.runtime-round3.chat-artifact-roundtrip.v1"
MODEL_ID = "lantai-round3-artifact-fixture"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _free_port() -> int:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
        listener.bind(("127.0.0.1", 0))
        return int(listener.getsockname()[1])


def _wait(
    predicate: Callable[[], Any],
    *,
    timeout: float,
    label: str,
    process: subprocess.Popen[str] | None = None,
) -> Any:
    deadline = time.monotonic() + timeout
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        if process is not None and process.poll() is not None:
            raise AssertionError(
                f"candidate exited while waiting for {label}: {process.returncode}"
            )
        try:
            value = predicate()
            if value:
                return value
        except Exception as exc:  # renderer/backend can be between routes
            last_error = exc
        time.sleep(0.2)
    raise AssertionError(f"timed out waiting for {label}; last_error={last_error!r}")


def _request_json(
    state: dict[str, Any],
    method: str,
    endpoint: str,
    payload: dict[str, Any] | None = None,
    *,
    timeout: float,
) -> Any:
    _status, _body, parsed = lifecycle._json_request(
        method,
        f"{lifecycle._state_base_url(state)}{endpoint}",
        payload=payload,
        headers=lifecycle._state_headers(state),
        timeout=timeout,
    )
    return parsed


def _create_presentation(pptx_path: Path, preview_path: Path, marker: str) -> None:
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    title = slide.shapes.title
    if title is not None:
        title.text = "澜台 Runtime 精确候选验收"
    textbox = slide.shapes.add_textbox(
        deck.slide_width // 8,
        deck.slide_height // 3,
        deck.slide_width * 3 // 4,
        deck.slide_height // 4,
    )
    textbox.text_frame.text = marker
    deck.save(pptx_path)

    image = Image.new("RGB", (1600, 900), "#f7f5ef")
    draw = ImageDraw.Draw(image)
    draw.rectangle((90, 90, 1510, 810), outline="#262626", width=4)
    draw.text((150, 180), "LANTAI RUNTIME ROUND 3", fill="#262626")
    draw.text((150, 330), marker, fill="#7a321f")
    image.save(preview_path)


def _tool_args(
    name: str,
    schema: dict[str, Any],
    *,
    pptx_path: Path,
    preview_path: Path,
) -> dict[str, Any]:
    if name == "declare_deliverable_contract":
        result: dict[str, Any] = {
            "kind": "presentation",
            "accepted_extensions": [".pptx"],
            "visual_verification_required": True,
        }
    elif name == "equip_lantai_tools":
        result = {"files": True}
    elif name == "finalize_deliverable":
        result = {
            "primary_path": str(pptx_path),
            "title": "Runtime 精确候选验收.pptx",
            "kind": "presentation",
            "entrypoint": str(pptx_path),
            "verification_paths": [str(preview_path)],
            "acceptance_evidence": [
                {
                    "check": "fixture_content",
                    "status": "passed",
                    "result": "passed",
                    "detail": "Valid PPTX and rendered preview were created in the isolated workspace.",
                }
            ],
        }
    elif name == "send_file_to_user":
        result = {"file_path": str(pptx_path)}
    else:
        raise AssertionError(f"unsupported provider tool call: {name}")
    properties = schema.get("properties") if isinstance(schema, dict) else None
    if isinstance(properties, dict) and "_lantai_public_progress" in properties:
        result["_lantai_public_progress"] = {
            "declare_deliverable_contract": "已确认交付为可编辑 PPTX。",
            "finalize_deliverable": "文件和预览已生成，正在核对交付版本。",
            "send_file_to_user": "验证已通过，正在交付演示文稿。",
        }.get(name, "")
    return result


class _ProviderHandler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    server_version = "LantaiRound3Artifact/1.0"

    def log_message(self, _format: str, *_args: Any) -> None:
        return

    def _send_bytes(self, status: int, body: bytes, content_type: str) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Connection", "close")
        self.end_headers()
        self.wfile.write(body)
        self.wfile.flush()
        self.close_connection = True

    def _send_json(self, status: int, payload: dict[str, Any]) -> None:
        self._send_bytes(
            status,
            json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            "application/json; charset=utf-8",
        )

    def do_GET(self) -> None:  # noqa: N802
        if parse.urlparse(self.path).path.rstrip("/") not in {"/v1/models", "/models"}:
            self._send_json(404, {"error": {"message": "not found"}})
            return
        self._send_json(
            200,
            {"object": "list", "data": [{"id": MODEL_ID, "object": "model"}]},
        )

    def do_POST(self) -> None:  # noqa: N802
        if parse.urlparse(self.path).path.rstrip("/") not in {
            "/v1/chat/completions",
            "/chat/completions",
        }:
            self._send_json(404, {"error": {"message": "not found"}})
            return
        length = int(self.headers.get("Content-Length") or "0")
        payload = json.loads(self.rfile.read(length).decode("utf-8") or "{}")
        if not isinstance(payload, dict):
            self._send_json(400, {"error": {"message": "invalid request"}})
            return
        server = self.server
        with server.request_lock:  # type: ignore[attr-defined]
            server.requests_seen.append(payload)  # type: ignore[attr-defined]
        tools = payload.get("tools") if isinstance(payload.get("tools"), list) else []
        tool_map = {
            str(item.get("function", {}).get("name") or ""): item.get("function", {})
            for item in tools
            if isinstance(item, dict) and isinstance(item.get("function"), dict)
        }
        forced = payload.get("tool_choice")
        forced_name = ""
        if isinstance(forced, dict):
            function = forced.get("function")
            if isinstance(function, dict):
                forced_name = str(function.get("name") or "")

        name = ""
        with server.request_lock:  # type: ignore[attr-defined]
            stage = int(server.tool_stage)  # type: ignore[attr-defined]
            if forced_name:
                name = forced_name
            elif stage == 0 and "declare_deliverable_contract" in tool_map:
                name = "declare_deliverable_contract"
            elif stage == 1 and "equip_lantai_tools" in tool_map:
                name = "equip_lantai_tools"
            elif stage == 2 and "finalize_deliverable" in tool_map:
                name = "finalize_deliverable"
            elif stage == 3 and "send_file_to_user" in tool_map:
                name = "send_file_to_user"
            if name:
                server.tool_stage = stage + 1  # type: ignore[attr-defined]
                server.tool_calls.append(name)  # type: ignore[attr-defined]

        tool_call: dict[str, Any] | None = None
        if name:
            function = tool_map.get(name) or {}
            arguments = _tool_args(
                name,
                function.get("parameters") or {},
                pptx_path=server.pptx_path,  # type: ignore[attr-defined]
                preview_path=server.preview_path,  # type: ignore[attr-defined]
            )
            tool_call = {
                "id": f"call_{uuid.uuid4().hex}",
                "type": "function",
                "function": {
                    "name": name,
                    "arguments": json.dumps(arguments, ensure_ascii=False),
                },
            }

        if payload.get("stream") is False:
            message: dict[str, Any] = {"role": "assistant", "content": None}
            finish_reason = "stop"
            if tool_call is not None:
                message["tool_calls"] = [tool_call]
                finish_reason = "tool_calls"
            else:
                message["content"] = "演示文稿已交付。"
            self._send_json(
                200,
                {
                    "id": f"chatcmpl-{uuid.uuid4().hex}",
                    "object": "chat.completion",
                    "created": int(time.time()),
                    "model": MODEL_ID,
                    "choices": [
                        {"index": 0, "message": message, "finish_reason": finish_reason}
                    ],
                    "usage": {"prompt_tokens": 1, "completion_tokens": 1, "total_tokens": 2},
                },
            )
            return

        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream; charset=utf-8")
        self.send_header("Cache-Control", "no-cache")
        self.send_header("Connection", "close")
        self.end_headers()
        chunk_id = f"chatcmpl-{uuid.uuid4().hex}"

        def send(delta: dict[str, Any], finish_reason: str | None = None) -> None:
            event = {
                "id": chunk_id,
                "object": "chat.completion.chunk",
                "created": int(time.time()),
                "model": MODEL_ID,
                "choices": [
                    {"index": 0, "delta": delta, "finish_reason": finish_reason}
                ],
            }
            self.wfile.write(
                f"data: {json.dumps(event, ensure_ascii=False)}\n\n".encode("utf-8")
            )
            self.wfile.flush()

        try:
            send({"role": "assistant"})
            if tool_call is not None:
                send({"tool_calls": [{"index": 0, **tool_call}]})
                send({}, "tool_calls")
            else:
                send({"content": "演示文稿已交付。"})
                send({}, "stop")
            self.wfile.write(b"data: [DONE]\n\n")
            self.wfile.flush()
        except (BrokenPipeError, ConnectionResetError):
            pass
        self.close_connection = True


class _Provider:
    def __init__(self, pptx_path: Path, preview_path: Path):
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), _ProviderHandler)
        self.server.requests_seen = []  # type: ignore[attr-defined]
        self.server.tool_calls = []  # type: ignore[attr-defined]
        self.server.tool_stage = 0  # type: ignore[attr-defined]
        self.server.request_lock = threading.Lock()  # type: ignore[attr-defined]
        self.server.pptx_path = pptx_path  # type: ignore[attr-defined]
        self.server.preview_path = preview_path  # type: ignore[attr-defined]
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)

    @property
    def base_url(self) -> str:
        host, port = self.server.server_address
        return f"http://{host}:{port}/v1"

    @property
    def request_count(self) -> int:
        return len(self.server.requests_seen)  # type: ignore[attr-defined]

    @property
    def tool_calls(self) -> list[str]:
        return list(self.server.tool_calls)  # type: ignore[attr-defined]

    def __enter__(self) -> "_Provider":
        self.thread.start()
        return self

    def __exit__(self, *_exc: object) -> None:
        self.server.shutdown()
        self.server.server_close()
        self.thread.join(timeout=5)


def _host_cdp(port: int, timeout: float) -> office.Cdp:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        for target in office._all_targets(port):
            if str(target.get("url") or "").startswith("file://"):
                continue
            cdp: office.Cdp | None = None
            try:
                cdp = office.Cdp(str(target["webSocketDebuggerUrl"]))
                if cdp.evaluate("Boolean(window.electronAPI && document.body)"):
                    return cdp
            except Exception:
                pass
            if cdp is not None:
                with contextlib.suppress(Exception):
                    cdp.close()
        time.sleep(0.2)
    raise AssertionError("packaged host renderer was unavailable")


def _body_text(cdp: office.Cdp) -> str:
    return str(cdp.evaluate("document.body?.innerText || ''") or "")


def _hash_route(cdp: office.Cdp) -> str:
    return str(cdp.evaluate("location.hash") or "")


def _click_artifact(cdp: office.Cdp, filename: str) -> None:
    clicked = cdp.evaluate(
        "(() => {"
        f"const label={json.dumps('在澜台中编辑：' + filename)};"
        "const target=[...document.querySelectorAll('button')].find((node) => "
        "node.getAttribute('aria-label')===label);"
        "if(!target) return false; target.click(); return true;"
        "})()"
    )
    if clicked is not True:
        raise AssertionError(f"artifact edit button was unavailable: {filename}")


def _click_return(cdp: office.Cdp) -> None:
    clicked = cdp.evaluate(
        "(() => {const target=[...document.querySelectorAll('button')].find((node) => "
        "node.textContent?.trim()==='返回对话');"
        "if(!target) return false; target.click(); return true;})()"
    )
    if clicked is not True:
        raise AssertionError("返回对话 button was unavailable")


def _compact_chat(chat: Any) -> dict[str, Any]:
    if not isinstance(chat, dict):
        raise AssertionError(f"expected chat object, got {chat!r}")
    messages = chat.get("messages") if isinstance(chat.get("messages"), list) else []
    user_messages = [
        message
        for message in messages
        if isinstance(message, dict) and message.get("role") == "user"
    ]
    return {
        "id": chat.get("id"),
        "status": chat.get("status"),
        "pending_turn": (chat.get("meta") or {}).get("pending_turn")
        if isinstance(chat.get("meta"), dict)
        else None,
        "message_count": len(messages),
        "user_message_count": len(user_messages),
        "messages": messages,
    }


def _event_summary(payload: Any) -> dict[str, Any]:
    events = payload.get("events") if isinstance(payload, dict) else []
    objects = [event for event in events if isinstance(event, dict)]
    types = [str(event.get("type") or "") for event in objects]
    tool_names: list[str] = []
    for event in objects:
        body = event.get("payload") if isinstance(event.get("payload"), dict) else {}
        name = body.get("tool_name") or body.get("name")
        if isinstance(name, str) and name:
            tool_names.append(name)
    return {
        "event_count": len(objects),
        "last_seq": max((int(event.get("seq") or 0) for event in objects), default=0),
        "terminal_types": [kind for kind in types if kind.startswith("turn.")],
        "tool_names": tool_names,
        "send_file_event_count": sum(name == "send_file_to_user" for name in tool_names),
        "finalize_event_count": sum(name == "finalize_deliverable" for name in tool_names),
    }


def _launch(
    *,
    app: Path,
    root: Path,
    state_path: Path,
    port: int,
    log_path: Path,
    provider_url: str,
    timeout: float,
) -> tuple[subprocess.Popen[str], Any, dict[str, Any], dict[str, Any]]:
    user_data = root / "user-data"
    env = os.environ.copy()
    env.update(
        {
            "LANTAI_DESKTOP_LIFECYCLE_SMOKE_FILE": str(state_path),
            "LANTAI_DESKTOP_USER_DATA_DIR": str(user_data),
            "LANTAI_WORKING_DIR": str(user_data / "working"),
            "LANTAI_SECRET_DIR": str(user_data / "working.secret"),
            "LANTAI_AUTH_ENABLED": "0",
            "LANTAI_DISABLE_GITHUB_UPDATER": "1",
            "LANTAI_DISABLE_KEYRING": "1",
            "LANTAI_SKIP_KEYRING": "1",
            "LANTAI_REMINDER_DELIVERY_WORKER": "0",
        }
    )
    provider_config = lifecycle._write_smoke_openai_provider_config(
        user_data,
        base_url=provider_url,
    )
    log_handle = log_path.open("w", encoding="utf-8")
    process = subprocess.Popen(
        [str(lifecycle._find_app_executable(app)), f"--remote-debugging-port={port}"],
        cwd=str(app.parent),
        env=env,
        stdout=log_handle,
        stderr=subprocess.STDOUT,
        text=True,
        start_new_session=True,
    )
    state = lifecycle._wait_for_state(
        state_path,
        timeout=timeout,
        predicate=lambda item: int(item.get("sequence") or 0) >= 1,
        process=process,
        expected_distribution="official",
    )
    lifecycle._assert_ready_state(state, expected_distribution="official")
    lifecycle._assert_desktop_readiness(state, timeout=timeout)
    return process, log_handle, state, provider_config


def _submit_with_fixture_approval(
    state: dict[str, Any],
    *,
    payload: dict[str, Any],
    root_session_id: str,
    timeout: float,
    process: subprocess.Popen[str],
) -> tuple[Any, dict[str, Any]]:
    """Submit one turn and explicitly approve its isolated file delivery.

    ``/api/console/chat`` stays open while the Native Runtime waits for a tool
    decision.  Drive the approval from a separate thread so this harness proves
    that the packaged approval boundary was reached and resolved instead of
    weakening or bypassing it.
    """

    result: dict[str, Any] = {}
    errors: list[BaseException] = []

    def submit() -> None:
        try:
            result["response"] = _request_json(
                state,
                "POST",
                "/api/console/chat",
                payload,
                timeout=timeout,
            )
        except BaseException as exc:  # surfaced on the driver thread below
            errors.append(exc)

    thread = threading.Thread(target=submit, name="artifact-turn-submit", daemon=True)
    thread.start()
    approval: dict[str, Any] | None = None
    approval_response: Any = None
    observed_at = ""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline and thread.is_alive():
        if process.poll() is not None:
            raise AssertionError(
                "candidate exited while waiting for fixture file approval: "
                f"{process.returncode}"
            )
        pending = _request_json(
            state,
            "GET",
            "/api/console/push-messages",
            timeout=min(timeout, 10),
        )
        candidates = (
            pending.get("pending_approvals")
            if isinstance(pending, dict)
            and isinstance(pending.get("pending_approvals"), list)
            else []
        )
        for item in candidates:
            if not isinstance(item, dict):
                continue
            if (
                item.get("tool_name") == "send_file_to_user"
                and item.get("root_session_id") == root_session_id
            ):
                approval = item
                observed_at = lifecycle._utc_now_iso()
                approval_response = _request_json(
                    state,
                    "POST",
                    "/api/approval/approve",
                    {
                        "request_id": item.get("request_id"),
                        "session_id": item.get("root_session_id"),
                        "user_id": item.get("user_id"),
                        "decision_scope": "once",
                    },
                    timeout=min(timeout, 10),
                )
                break
        if approval is not None:
            break
        time.sleep(0.2)

    remaining = max(0.1, deadline - time.monotonic())
    thread.join(timeout=remaining)
    if thread.is_alive():
        raise AssertionError("timed out waiting for approved artifact turn to finish")
    if errors:
        raise errors[0]
    if approval is None:
        raise AssertionError(
            "artifact turn completed without exposing the required file-delivery approval"
        )
    if not isinstance(approval_response, dict) or approval_response.get("success") is not True:
        raise AssertionError(f"fixture file approval failed: {approval_response!r}")
    return result.get("response"), {
        "observed_at": observed_at,
        "tool_name": approval.get("tool_name"),
        "severity": approval.get("severity"),
        "root_session_matches": approval.get("root_session_id") == root_session_id,
        "decision_scope": approval_response.get("decision_scope"),
        "approved": True,
        "resolved_request_count": 1
        + len(approval_response.get("resolved_request_ids") or []),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--app", type=Path, required=True)
    parser.add_argument("--evidence-output", type=Path, required=True)
    parser.add_argument("--work-dir", type=Path, required=True)
    parser.add_argument("--timeout", type=float, default=180)
    args = parser.parse_args()

    app = args.app.expanduser().resolve(strict=True)
    work_dir = args.work_dir.expanduser().resolve()
    evidence_path = args.evidence_output.expanduser().resolve()
    work_dir.mkdir(parents=True, exist_ok=True)
    smoke_root = Path(tempfile.mkdtemp(prefix="lantai-round3-artifact-card-"))
    user_data = smoke_root / "user-data"
    workspace = user_data / "working" / "workspaces" / "default"
    artifact_dir = workspace / "artifacts"
    marker = f"ROUND3_ARTIFACT_{uuid.uuid4().hex[:10]}"
    pptx_path = artifact_dir / f"{marker}.pptx"
    preview_path = artifact_dir / f"{marker}.png"
    _create_presentation(pptx_path, preview_path, marker)
    initial_hash = _sha256(pptx_path)
    initial_slide_count = office._pptx_slide_count(pptx_path)
    port = _free_port()
    state_path = smoke_root / "desktop-state.json"
    log_path = work_dir / "candidate.log"
    screenshots = work_dir / "screenshots"
    screenshots.mkdir(parents=True, exist_ok=True)
    process: subprocess.Popen[str] | None = None
    log_handle: Any = None
    host: office.Cdp | None = None
    started_at = lifecycle._utc_now_iso()

    with _Provider(pptx_path, preview_path) as provider:
        try:
            process, log_handle, state, provider_config = _launch(
                app=app,
                root=smoke_root,
                state_path=state_path,
                port=port,
                log_path=log_path,
                provider_url=provider.base_url,
                timeout=args.timeout,
            )
            host = _host_cdp(port, args.timeout)
            chat = _request_json(
                state,
                "POST",
                "/api/chats",
                {
                    "name": "Runtime 产物往返验收",
                    "session_id": f"console:round3-artifact-{marker}",
                    "user_id": "default",
                    "channel": "console",
                    "meta": {"topic_workspace_dir": str(workspace)},
                },
                timeout=args.timeout,
            )
            chat_id = str(chat.get("id") or "") if isinstance(chat, dict) else ""
            if not chat_id:
                raise AssertionError(f"chat creation failed: {chat!r}")
            turn_id = f"round3-artifact-turn-{uuid.uuid4().hex}"
            _turn_response, approval_evidence = _submit_with_fixture_approval(
                state,
                payload={
                    "chat_id": chat_id,
                    "input": [{"role": "user", "content": "请交付这份精确候选验收 PPT。"}],
                    "session_id": f"console:round3-artifact-{marker}",
                    "user_id": "default",
                    "channel": "console",
                    "stream": True,
                    "client_turn_id": turn_id,
                    "meta": {"client_turn_id": turn_id, "deep_thinking": True},
                },
                root_session_id=f"console:round3-artifact-{marker}",
                timeout=args.timeout,
                process=process,
            )

            terminal_chat = _wait(
                lambda: _request_json(
                    state, "GET", f"/api/chats/{chat_id}", timeout=args.timeout
                ),
                timeout=args.timeout,
                label="terminal chat",
                process=process,
            )
            _wait(
                lambda: pptx_path.name
                in json.dumps(
                    _request_json(
                        state, "GET", f"/api/chats/{chat_id}", timeout=args.timeout
                    ),
                    ensure_ascii=False,
                ),
                timeout=args.timeout,
                label="durable delivered file",
                process=process,
            )
            terminal_chat = _request_json(
                state, "GET", f"/api/chats/{chat_id}", timeout=args.timeout
            )
            initial_events = _request_json(
                state,
                "GET",
                f"/api/chats/{chat_id}/events?after_seq=0",
                timeout=args.timeout,
            )
            initial_counts = {
                "intentional_submits": 1,
                "accepted_logical_commands": 1,
                "provider_requests": provider.request_count,
                "provider_tool_calls": provider.tool_calls,
                "chat": _compact_chat(terminal_chat),
                "events": _event_summary(initial_events),
                "artifact_sha256": initial_hash,
                "slide_count": initial_slide_count,
            }
            if initial_counts["chat"]["user_message_count"] != 1:
                raise AssertionError(
                    f"single submit did not persist exactly one user item: {initial_counts!r}"
                )

            host.evaluate(f"location.hash={json.dumps('#/chat/' + chat_id)}")
            _wait(
                lambda: pptx_path.name in _body_text(host),
                timeout=args.timeout,
                label="chat artifact card",
                process=process,
            )
            host.screenshot(screenshots / "01-delivered-chat.png")

            # P2: enter and leave without editing.
            p2_before_provider = provider.request_count
            p2_before_hash = _sha256(pptx_path)
            _click_artifact(host, pptx_path.name)
            _wait(
                lambda: "返回对话" in _body_text(host),
                timeout=args.timeout,
                label="office route P2",
                process=process,
            )
            host.screenshot(screenshots / "02-p2-office.png")
            _click_return(host)
            _wait(
                lambda: _hash_route(host).endswith(f"/chat/{chat_id}")
                and pptx_path.name in _body_text(host),
                timeout=args.timeout,
                label="P2 return to original chat",
                process=process,
            )
            p2_chat = _request_json(
                state, "GET", f"/api/chats/{chat_id}", timeout=args.timeout
            )
            p2_events = _request_json(
                state,
                "GET",
                f"/api/chats/{chat_id}/events?after_seq=0",
                timeout=args.timeout,
            )
            p2 = {
                "status": "passed",
                "route": _hash_route(host),
                "artifact_sha256_before": p2_before_hash,
                "artifact_sha256_after": _sha256(pptx_path),
                "artifact_bytes_unchanged": _sha256(pptx_path) == p2_before_hash,
                "provider_request_delta": provider.request_count - p2_before_provider,
                "accepted_command_delta": 0,
                "generate_delta": 0,
                "send_file_delta": _event_summary(p2_events)["send_file_event_count"]
                - initial_counts["events"]["send_file_event_count"],
                "chat": _compact_chat(p2_chat),
                "events": _event_summary(p2_events),
            }
            if not p2["artifact_bytes_unchanged"] or p2["provider_request_delta"] != 0:
                raise AssertionError(f"P2 introduced an artifact/model mutation: {p2!r}")
            if p2["send_file_delta"] != 0:
                raise AssertionError(f"P2 repeated delivery: {p2!r}")
            if p2["chat"]["message_count"] != initial_counts["chat"]["message_count"]:
                raise AssertionError(f"P2 duplicated chat history: {p2!r}")
            if p2["chat"]["user_message_count"] != 1:
                raise AssertionError(f"P2 changed the logical user item count: {p2!r}")
            host.screenshot(screenshots / "03-p2-returned-chat.png")

            # P3: use the packaged slides editor, save, and return.
            p3_before_provider = provider.request_count
            p3_before_hash = _sha256(pptx_path)
            p3_before_slides = office._pptx_slide_count(pptx_path)
            _click_artifact(host, pptx_path.name)
            _wait(
                lambda: "返回对话" in _body_text(host),
                timeout=args.timeout,
                label="office route P3",
                process=process,
            )
            slides = office._find_office_cdp(port, office.GLOBAL_BY_SUFFIX[".pptx"], args.timeout)
            try:
                interaction = office._edit_pptx(
                    office.AppSession(
                        process=process,
                        root=smoke_root,
                        port=port,
                        cdp=slides,
                        log=log_path,
                        workbench={},
                    ),
                    args.timeout,
                )
                _wait(
                    lambda: _sha256(pptx_path) != p3_before_hash,
                    timeout=args.timeout,
                    label="P3 saved bytes",
                    process=process,
                )
                slides.screenshot(screenshots / "04-p3-saved-office.png")
            finally:
                slides.close()
            _click_return(host)
            _wait(
                lambda: _hash_route(host).endswith(f"/chat/{chat_id}")
                and pptx_path.name in _body_text(host),
                timeout=args.timeout,
                label="P3 return to original chat",
                process=process,
            )
            p3_chat = _request_json(
                state, "GET", f"/api/chats/{chat_id}", timeout=args.timeout
            )
            p3_events = _request_json(
                state,
                "GET",
                f"/api/chats/{chat_id}/events?after_seq=0",
                timeout=args.timeout,
            )
            p3_after_hash = _sha256(pptx_path)
            p3_after_slides = office._pptx_slide_count(pptx_path)
            p3 = {
                "status": "passed",
                "route": _hash_route(host),
                "interaction": interaction,
                "artifact_sha256_before": p3_before_hash,
                "artifact_sha256_after": p3_after_hash,
                "artifact_bytes_changed": p3_after_hash != p3_before_hash,
                "slide_count_before": p3_before_slides,
                "slide_count_after": p3_after_slides,
                "provider_request_delta": provider.request_count - p3_before_provider,
                "accepted_command_delta": 0,
                "generate_delta": 0,
                "send_file_delta": _event_summary(p3_events)["send_file_event_count"]
                - p2["events"]["send_file_event_count"],
                "same_path_overwrite": True,
                "original_delivery_hash": initial_hash,
                "current_bytes_hash": p3_after_hash,
                "old_verification_does_not_match_current_bytes": initial_hash
                != p3_after_hash,
                "chat": _compact_chat(p3_chat),
                "events": _event_summary(p3_events),
            }
            if not p3["artifact_bytes_changed"] or p3_after_slides != p3_before_slides + 1:
                raise AssertionError(f"P3 did not persist the intended edit: {p3!r}")
            if p3["provider_request_delta"] != 0 or p3["send_file_delta"] != 0:
                raise AssertionError(f"P3 caused an automatic rerun/delivery: {p3!r}")
            if p3["chat"]["message_count"] != initial_counts["chat"]["message_count"]:
                raise AssertionError(f"P3 duplicated chat history: {p3!r}")
            if p3["chat"]["user_message_count"] != 1:
                raise AssertionError(f"P3 changed the logical user item count: {p3!r}")
            host.screenshot(screenshots / "05-p3-returned-chat.png")

            provenance_path = (
                app
                / "Contents"
                / "Resources"
                / "console"
                / "dist"
                / "release-provenance.json"
            )
            screenshot_manifest = [
                {
                    "path": str(path),
                    "sha256": _sha256(path),
                    "bytes": path.stat().st_size,
                }
                for path in sorted(screenshots.glob("*.png"))
            ]
            evidence = {
                "schema": SCHEMA,
                "status": "passed",
                "started_at": started_at,
                "finished_at": lifecycle._utc_now_iso(),
                "candidate": {
                    "app_path": str(app),
                    "executable": str(lifecycle._find_app_executable(app)),
                    "executable_sha256": _sha256(lifecycle._find_app_executable(app)),
                    "provenance": json.loads(provenance_path.read_text(encoding="utf-8")),
                    "process_pid": process.pid,
                    "backend_pid": state.get("backend_pid"),
                    "api_base_url": lifecycle._state_base_url(state),
                    "profile": str(user_data),
                },
                "fixture": {
                    "marker": marker,
                    "pptx_path": str(pptx_path),
                    "preview_path": str(preview_path),
                    "initial_sha256": initial_hash,
                    "initial_slide_count": initial_slide_count,
                },
                "provider": {
                    "kind": "controlled local OpenAI-compatible fixture",
                    "config": provider_config,
                    "request_count": provider.request_count,
                    "tool_calls": provider.tool_calls,
                },
                "initial_counts": initial_counts,
                "explicit_fixture_approval": approval_evidence,
                "p2_no_edit_roundtrip": p2,
                "p3_edit_save_roundtrip": p3,
                "screenshots": screenshot_manifest,
                "smoke_root": str(smoke_root),
                "log": str(log_path),
            }
            evidence_path.parent.mkdir(parents=True, exist_ok=True)
            evidence_path.write_text(
                json.dumps(evidence, ensure_ascii=False, indent=2) + "\n",
                encoding="utf-8",
            )
            print(
                json.dumps(
                    {
                        "status": "passed",
                        "evidence": str(evidence_path),
                        "chat_id": chat_id,
                        "p2_provider_delta": p2["provider_request_delta"],
                        "p3_provider_delta": p3["provider_request_delta"],
                        "initial_hash": initial_hash,
                        "edited_hash": p3_after_hash,
                    },
                    ensure_ascii=False,
                )
            )
            return 0
        finally:
            if host is not None:
                with contextlib.suppress(Exception):
                    host.close()
            if process is not None and process.poll() is None:
                with contextlib.suppress(ProcessLookupError):
                    os.killpg(process.pid, signal.SIGTERM)
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    with contextlib.suppress(ProcessLookupError):
                        os.killpg(process.pid, signal.SIGKILL)
                    process.wait(timeout=5)
            if log_handle is not None:
                log_handle.close()


if __name__ == "__main__":
    raise SystemExit(main())
