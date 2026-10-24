#!/usr/bin/env python3
"""Opt-in real-provider analysis journey, isolated from the desktop's data.

Scenario text is supplied over stdin, never checked into the repository or
retained as evidence. Only encrypted provider stores are copied. Backend logs,
conversation text and model output are discarded when the temporary run ends.
The JSON result contains checks, counts and hashes, not transcripts.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
import uuid
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from decimal import Decimal
from pathlib import Path
from typing import Any, Callable
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
DELIVERY_TOOLS = frozenset({
    "create_analysis_chart", "create_analysis_report",
    "create_analysis_presentation", "export_analysis_evidence",
    # A long-running analysis presentation can outlive the foreground tool
    # call.  This scoped probe resumes the already-authorized durable job; it
    # cannot create another deck and still passes the stored analysis binding
    # and final-file verification before delivery.
    "get_presentation_generation_status",
})
ANALYSIS_TOOLS = DELIVERY_TOOLS | frozenset({
    "equip_lantai_tools", "import_data_table", "analyze_data_table",
    "get_analysis_result", "refresh_analysis_result",
})
DIAGNOSTIC_TOOLS = ANALYSIS_TOOLS | frozenset({
    "execute_shell_command", "execute_python_code", "write_file", "edit_file",
    "send_file_to_user", "create_presentation", "get_presentation_generation_status",
})


class SmokeFailure(RuntimeError):
    """A fixed diagnostic code; never a raw provider/server exception."""


def require(condition: bool, code: str) -> None:
    if not condition:
        raise SmokeFailure(code)


def digest(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()


def source_presentation_runtime_env() -> dict[str, str]:
    """Bind source smoke to the repository's installed native PPTX runtime."""
    module = ROOT / "console" / "node_modules" / "pptxgenjs"
    require(module.is_dir(), "source_presentation_backend_missing")
    return {"LANTAI_PPTXGENJS_MODULE": str(module)}


def validate_approval(item: dict, session: str) -> None:
    """Fail closed with metadata-only codes; never include approval arguments."""
    require(item.get("root_session_id") == session, "approval_root_session_mismatch")
    require(item.get("user_id") == "default", "approval_user_mismatch")
    require(item.get("agent_id") == "default", "approval_agent_mismatch")
    tool = item.get("tool_name")
    if tool == "send_file_to_user":
        policy = item.get("policy_decision") or {}
        require(not policy.get("workspace_escape") and not policy.get("hard_denied"),
                "file_delivery_workspace_escape")
        require(bool(re.fullmatch(r"[0-9a-f]{64}", str(policy.get("analysis_smoke_final_file_sha256", "")))),
                "file_delivery_attestation_missing")
        return
    diagnostic = tool if tool in DIAGNOSTIC_TOOLS else "other"
    require(tool in DELIVERY_TOOLS, f"unexpected_tool_approval_{diagnostic}")


def sealed_record(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_bytes())
    checksum = value.pop("record_sha256", None)
    canonical = json.dumps(value, ensure_ascii=False, sort_keys=True,
                           separators=(",", ":"), allow_nan=False).encode()
    require(checksum == digest(canonical), "record_checksum_mismatch")
    value["record_sha256"] = checksum
    return value


def sealed_file_approval_hash(state_root: Path, session: str, workspace: Path, file_path: str) -> str:
    """Inspect the exact requested file, not a workspace-wide permission rule."""
    candidate = Path(file_path).expanduser()
    candidate = candidate if candidate.is_absolute() else workspace / candidate
    require(not candidate.is_symlink() and candidate.resolve().is_relative_to(workspace.resolve()),
            "file_delivery_outside_workspace")
    receipts = state_root / "analysis" / digest(session.encode()) / "deliveries"
    for path in receipts.glob("*.json"):
        receipt = sealed_record(path)
        if Path(receipt["artifact_path"]).resolve() != candidate.resolve():
            continue
        require(receipt["kind"] in {"evidence", "svg", "docx", "pptx"}
                and bool(receipt.get("bound_elements"))
                and receipt["verification"]["delivery_consistency"]["status"] == "passed",
                "file_delivery_not_verified")
        require(candidate.is_file() and candidate.stat().st_size <= 100 * 1024 * 1024,
                "file_delivery_missing_or_oversized")
        require(digest(candidate.read_bytes()) == receipt["content_sha256"], "file_delivery_hash_mismatch")
        return receipt["content_sha256"]
    raise SmokeFailure("file_delivery_receipt_missing")


def observed_backend() -> None:
    """Test-only metadata observation; retain real policy and once-only approval.

    The production approval API intentionally excludes tool arguments. Inside
    this disposable backend we observe (never publish) the requested file path
    and add only its verified sealed hash to the pending policy metadata. No
    tool implementation, policy decision, model response or approval result is
    substituted. Shell, deletion and unbound-file requests remain fail-closed.
    """
    import asyncio

    from lantai.app.approvals.service import ApprovalService
    from lantai.cli.main import cli
    from lantai.runtime.adapters.native_runtime import NativeLantaiRuntime
    from lantai.runtime.assistant_artifact_projection import (
        AssistantArtifactProjectionAccumulator,
    )

    original = ApprovalService.create_pending
    original_projection_observe = AssistantArtifactProjectionAccumulator.observe
    original_projection_project = AssistantArtifactProjectionAccumulator.project
    original_projection_pending = (
        AssistantArtifactProjectionAccumulator.pending_completion
    )
    original_projection_has_pending = (
        AssistantArtifactProjectionAccumulator.has_pending
    )
    original_run_turn = NativeLantaiRuntime.run_turn
    projection_observation_path = (
        Path(os.environ["LANTAI_WORKING_DIR"])
        / "analysis-smoke-projection-observations.jsonl"
    )

    async def observe(service, *args, **kwargs):
        pending = await original(service, *args, **kwargs)
        if pending.tool_name == "send_file_to_user":
            extra = kwargs.get("extra") or {}
            try:
                scope = json.loads((Path(os.environ["LANTAI_WORKING_DIR"]) / "analysis-smoke-scope.json").read_bytes())
                state = Path(scope["state_root"])
                workspace = Path(extra["workspace_dir"])
                secret_root = Path(os.environ["LANTAI_SECRET_DIR"]).resolve()
                test_root = secret_root.parent
                require(pending.session_id == scope["session"] == pending.root_session_id
                        and pending.user_id == "default" and pending.agent_id == "default"
                        and workspace.resolve() == Path(scope["artifact_root"]).resolve()
                        and state.resolve().is_relative_to(test_root)
                        and workspace.resolve().is_relative_to(test_root)
                        and not workspace.resolve().is_relative_to(secret_root)
                        and not state.resolve().is_relative_to(secret_root), "observer_not_isolated")
                checksum = sealed_file_approval_hash(
                    state, pending.session_id, workspace, extra["tool_call"]["input"]["file_path"])
            except Exception:
                checksum = None
            if checksum:
                pending.policy_decision = {**(pending.policy_decision or {}),
                                           "analysis_smoke_final_file_sha256": checksum}
        return pending

    def write_projection_observation(**values: bool) -> None:
        observation = {
            key: values.get(key) is True
            for key in (
                "observe_called",
                "spec_present",
                "spec_name_matches",
                "spec_origin_builtin",
                "projection_marker_present",
                "result_ok",
                "reference_shape_present",
                "observe_has_pending_after",
                "project_called",
                "project_event_eligible",
                "project_had_pending_before",
                "project_has_pending_after",
                "project_added_chart_fence",
                "pending_called",
                "pending_had_pending_before",
                "pending_produced_event",
                "pending_event_has_chart_fence",
                "run_turn_finalized",
                "run_turn_exhausted",
                "run_turn_terminal_seen",
                "run_turn_cancelled",
            )
        }
        with projection_observation_path.open("a", encoding="utf-8") as stream:
            stream.write(json.dumps(observation, sort_keys=True) + "\n")

    def observe_projection(accumulator, spec, result):
        """Persist only structural projection diagnostics in the test clone."""

        is_chart_result = (
            str(getattr(result, "name", "") or "")
            == "create_analysis_chart"
        )
        diagnostic: dict[str, bool] | None = None
        if is_chart_result:
            metadata = result.metadata if isinstance(result.metadata, dict) else {}
            reference = metadata.get("chat_chart_reference")
            diagnostic = dict(
                observe_called=True,
                spec_present=spec is not None,
                spec_name_matches=bool(
                    spec is not None and spec.name == result.name
                ),
                spec_origin_builtin=bool(
                    spec is not None and spec.origin == "builtin"
                ),
                projection_marker_present=bool(
                    spec is not None
                    and spec.metadata.get("assistant_artifact_projection")
                    == "analysis_chart_reference"
                ),
                result_ok=bool(result.ok),
                reference_shape_present=bool(
                    isinstance(reference, dict)
                    and set(reference)
                    == {"analysis_id", "x_column", "y_column", "type"}
                ),
            )
        observed = original_projection_observe(accumulator, spec, result)
        if diagnostic is not None:
            write_projection_observation(
                **diagnostic,
                observe_has_pending_after=original_projection_has_pending(
                    accumulator
                ),
            )
        return observed

    def project_projection(accumulator, event):
        pending_before = original_projection_has_pending(accumulator)
        projected = original_projection_project(accumulator, event)
        before_text = str(getattr(event, "summary", "") or "")
        after_text = str(getattr(projected, "summary", "") or "")
        write_projection_observation(
            project_called=True,
            project_event_eligible=bool(
                event.type == "message.completed"
                and event.status == "completed"
                and event.source == "assistant"
            ),
            project_had_pending_before=pending_before,
            project_has_pending_after=original_projection_has_pending(
                accumulator
            ),
            project_added_chart_fence=bool(
                "```analysis-chart" not in before_text
                and "```analysis-chart" in after_text
            ),
        )
        return projected

    def pending_projection(accumulator, **kwargs):
        pending_before = original_projection_has_pending(accumulator)
        event = original_projection_pending(accumulator, **kwargs)
        write_projection_observation(
            pending_called=True,
            pending_had_pending_before=pending_before,
            pending_produced_event=event is not None,
            pending_event_has_chart_fence=bool(
                event is not None
                and "```analysis-chart" in str(event.summary or "")
            ),
        )
        return event

    async def observe_run_turn(
        runtime,
        runtime_input,
        *,
        request_policy_observation_invocation=None,
    ):
        """Record only whether the real Native stream reached its own terminal."""

        exhausted = False
        terminal_seen = False
        cancelled = False
        try:
            async for event in original_run_turn(
                runtime,
                runtime_input,
                request_policy_observation_invocation=(
                    request_policy_observation_invocation
                ),
            ):
                if event.type in {
                    "turn.completed",
                    "turn.failed",
                    "turn.interrupted",
                    "turn.stalled",
                }:
                    terminal_seen = True
                yield event
            exhausted = True
        except asyncio.CancelledError:
            cancelled = True
            raise
        finally:
            write_projection_observation(
                run_turn_finalized=True,
                run_turn_exhausted=exhausted,
                run_turn_terminal_seen=terminal_seen,
                run_turn_cancelled=cancelled,
            )

    ApprovalService.create_pending = observe
    AssistantArtifactProjectionAccumulator.observe = observe_projection
    AssistantArtifactProjectionAccumulator.project = project_projection
    AssistantArtifactProjectionAccumulator.pending_completion = pending_projection
    NativeLantaiRuntime.run_turn = observe_run_turn
    cli()


def projection_observation_snapshot(path: Path) -> int:
    """Return the current line count for structural smoke-only observations."""

    if not path.is_file():
        return 0
    return len(path.read_text(encoding="utf-8").splitlines())


def projection_observations_since(path: Path, baseline: int) -> list[dict[str, bool]]:
    """Read bounded boolean-only observations written by the disposable backend."""

    if not path.is_file():
        return []
    observations: list[dict[str, bool]] = []
    for line in path.read_text(encoding="utf-8").splitlines()[baseline:]:
        try:
            value = json.loads(line)
        except ValueError:
            continue
        if not isinstance(value, dict):
            continue
        observations.append(
            {
                key: value.get(key) is True
                for key in (
                    "observe_called",
                    "spec_present",
                    "spec_name_matches",
                    "spec_origin_builtin",
                    "projection_marker_present",
                    "result_ok",
                    "reference_shape_present",
                    "observe_has_pending_after",
                    "project_called",
                    "project_event_eligible",
                    "project_had_pending_before",
                    "project_has_pending_after",
                    "project_added_chart_fence",
                    "pending_called",
                    "pending_had_pending_before",
                    "pending_produced_event",
                    "pending_event_has_chart_fence",
                    "run_turn_finalized",
                    "run_turn_exhausted",
                    "run_turn_terminal_seen",
                    "run_turn_cancelled",
                )
            }
        )
    return observations


def tool_counts(frames: list[tuple[int, dict]]) -> dict[str, int]:
    """Count unique public calls; do not retain arguments or arbitrary names."""
    calls: dict[str, set[str]] = {}

    def visit(value):
        if isinstance(value, dict):
            call = value.get("data") if isinstance(value.get("data"), dict) else value
            name = call.get("name") or call.get("tool_name")
            if value.get("type") in {"tool_call", "function_call", "tool_use"} and name in DIAGNOSTIC_TOOLS:
                calls.setdefault(name, set()).add(str(call.get("call_id") or call.get("id") or name))
            for nested in value.values():
                if isinstance(nested, (dict, list)):
                    visit(nested)
        elif isinstance(value, list):
            for nested in value:
                visit(nested)

    for _at, frame in frames:
        visit(frame)
    return {name: len(calls[name]) for name in sorted(calls)}


def _event_ledger_tool_start(event: Any) -> tuple[str, str, str] | None:
    """Return allow-listed start metadata without retaining arguments or output."""
    start_types = frozenset({
        "chat.tool_call.started", "tool.call.started", "tool.started",
    })
    event_type = str(event.type or "")
    event_kind = str(event.kind or "")
    if (
        (event_type and event_type not in start_types)
        or (event_kind and event_kind not in start_types)
        or not (event_type or event_kind)
    ):
        return None
    payload = event.payload if isinstance(event.payload, dict) else {}
    nested = payload.get("tool_call")
    nested = nested if isinstance(nested, dict) else {}
    name = str(
        payload.get("tool_name")
        or payload.get("name")
        or nested.get("name")
        or ""
    )
    if name not in DIAGNOSTIC_TOOLS:
        return None
    call_id = str(
        payload.get("tool_call_id")
        or payload.get("call_id")
        or payload.get("id")
        or nested.get("tool_call_id")
        or nested.get("call_id")
        or nested.get("id")
        or event.item_id
        or event.id
    )
    scope = str(
        event.run_id
        or f"chat:{event.chat_id or ''}:session:{event.session_id or ''}"
    )
    return name, call_id, scope


def event_ledger_snapshot(
    events: list[Any],
) -> tuple[frozenset[str], int, frozenset[str]]:
    """Capture event and stable call identities at the boundary before one turn."""
    return (
        frozenset(str(event.id) for event in events),
        max((int(event.seq) for event in events), default=0),
        frozenset(
            metadata[1]
            for event in events
            if (metadata := _event_ledger_tool_start(event)) is not None
        ),
    )


def event_ledger_assistant_projection_shape(
    events: list[Any],
    snapshot: tuple[frozenset[str], int, frozenset[str]],
) -> dict[str, int]:
    """Count new assistant completion structure without retaining text."""

    baseline_ids, baseline_seq, _baseline_call_ids = snapshot
    completions = 0
    with_text = 0
    with_chart_fence = 0
    for event in sorted(events, key=lambda item: int(item.seq)):
        if str(event.id) in baseline_ids or int(event.seq) <= baseline_seq:
            continue
        if (
            str(getattr(event, "type", "") or "") != "message.completed"
            or str(getattr(event, "status", "") or "") != "completed"
            or str(getattr(event, "source", "") or "") != "assistant"
        ):
            continue
        payload = event.payload if isinstance(event.payload, dict) else {}
        message = payload.get("message")
        if not isinstance(message, dict) or message.get("role") != "assistant":
            continue
        completions += 1
        content = message.get("content")
        texts = [
            str(block.get("text") or "")
            for block in content
            if isinstance(block, dict)
            and block.get("type") == "text"
            and isinstance(block.get("text"), str)
        ] if isinstance(content, list) else []
        visible = "".join(texts)
        if visible:
            with_text += 1
        if "```analysis-chart" in visible:
            with_chart_fence += 1
    return {
        "assistant_completions": completions,
        "assistant_completions_with_text": with_text,
        "assistant_completions_with_chart_fence": with_chart_fence,
    }


def event_ledger_terminal_shape(
    events: list[Any],
    snapshot: tuple[frozenset[str], int, frozenset[str]],
) -> list[dict[str, str]]:
    """Return new runtime terminal types and bounded error codes only."""

    baseline_ids, baseline_seq, _baseline_call_ids = snapshot
    terminals: list[dict[str, str]] = []
    allowed_codes = {
        "missing_visible_answer",
        "model_first_event_timeout",
        "model_output_invalid",
        "model_output_truncated",
        "model_request_timeout",
        "model_stream_idle_timeout",
        "provider_connection_interrupted",
        "repeated_tool_failure",
        "required_capability_failed",
        "runtime_durability_blocked",
        "task_budget_exceeded",
        "tool_argument_repair_exhausted",
    }
    for event in sorted(events, key=lambda item: int(item.seq)):
        if str(event.id) in baseline_ids or int(event.seq) <= baseline_seq:
            continue
        event_type = str(getattr(event, "type", "") or "")
        if event_type not in {
            "turn.completed",
            "turn.failed",
            "turn.interrupted",
            "turn.stalled",
        }:
            continue
        payload = event.payload if isinstance(event.payload, dict) else {}
        raw_code = str(payload.get("error_code") or "")
        terminal = {
            "type": event_type,
            "status": str(getattr(event, "status", "") or ""),
        }
        if raw_code:
            terminal["error_code"] = (
                raw_code if raw_code in allowed_codes else "other"
            )
            terminal["error_code_sha256"] = digest(raw_code.encode())
        terminals.append(terminal)
    return terminals


def event_ledger_tool_counts(
    events: list[Any],
    snapshot: tuple[frozenset[str], int, frozenset[str]],
) -> dict[str, int]:
    """Count new allow-listed starts without retaining arguments or outputs."""
    baseline_ids, baseline_seq, baseline_call_ids = snapshot
    calls: dict[str, set[tuple[str, str]]] = {}
    seen: set[tuple[str, str]] = set()
    for event in sorted(events, key=lambda item: int(item.seq)):
        if str(event.id) in baseline_ids or int(event.seq) <= baseline_seq:
            continue
        metadata = _event_ledger_tool_start(event)
        if metadata is None:
            continue
        name, call_id, scope = metadata
        if call_id in baseline_call_ids:
            continue
        identity = (scope, call_id)
        if identity in seen:
            continue
        seen.add(identity)
        calls.setdefault(name, set()).add(identity)
    return {name: len(calls[name]) for name in sorted(calls)}


def _presentation_error_class(message: str) -> str:
    """Classify expected PPT failures without publishing tool output."""
    lowered = message.lower()
    classifiers = (
        ("already exists", "presentation_filename_exists"),
        ("unbound numeric fact", "presentation_unbound_numeric_fact"),
        ("fact reference", "presentation_fact_reference_invalid"),
        ("fact row", "presentation_fact_row_invalid"),
        ("fact column", "presentation_fact_column_invalid"),
        ("missing value", "presentation_fact_value_missing"),
        ("analysis binding", "presentation_binding_invalid"),
        ("must be bound", "presentation_binding_missing"),
        ("needs a result fact reference", "presentation_slide_fact_missing"),
        ("slide count", "presentation_slide_count_invalid"),
        ("final-file consistency check", "presentation_final_file_invalid"),
        ("delivery path", "presentation_delivery_path_invalid"),
        ("presentation_generation_pending", "presentation_generation_pending"),
        ("presentation_generation_failed", "presentation_generation_failed"),
    )
    return next((code for marker, code in classifiers if marker in lowered), "other")


def _event_ledger_tool_completion(event: Any) -> tuple[str, str, str, str, str, str] | None:
    """Return metadata-only completion evidence for an allow-listed tool."""
    completion_types = frozenset({
        "chat.tool_call.finished", "chat.tool_result",
        "tool.call.completed", "tool.completed",
    })
    event_type = str(event.type or "")
    event_kind = str(event.kind or "")
    if (
        (event_type and event_type not in completion_types)
        or (event_kind and event_kind not in completion_types)
        or not (event_type or event_kind)
    ):
        return None
    payload = event.payload if isinstance(event.payload, dict) else {}
    nested = payload.get("tool_call")
    nested = nested if isinstance(nested, dict) else {}
    name = str(
        payload.get("tool_name")
        or payload.get("name")
        or nested.get("name")
        or ""
    )
    if name not in DIAGNOSTIC_TOOLS:
        return None
    call_id = str(
        payload.get("tool_call_id")
        or payload.get("call_id")
        or payload.get("id")
        or nested.get("tool_call_id")
        or nested.get("call_id")
        or nested.get("id")
        or event.item_id
        or event.id
    )
    scope = str(
        event.run_id
        or f"chat:{event.chat_id or ''}:session:{event.session_id or ''}"
    )
    succeeded = payload.get("ok") is True or str(getattr(event, "status", "")) == "completed"
    error = "" if succeeded else str(payload.get("error") or payload.get("error_code") or "")
    error_code = "" if succeeded else str(payload.get("error_code") or "")
    return name, call_id, scope, "completed" if succeeded else "failed", error, error_code


def event_ledger_tool_outcomes(
    events: list[Any],
    snapshot: tuple[frozenset[str], int, frozenset[str]],
) -> list[dict[str, str]]:
    """Return safe completion diagnostics, never arguments, output or error text."""
    baseline_ids, baseline_seq, baseline_call_ids = snapshot
    outcomes: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    for event in sorted(events, key=lambda item: int(item.seq)):
        if str(event.id) in baseline_ids or int(event.seq) <= baseline_seq:
            continue
        metadata = _event_ledger_tool_completion(event)
        if metadata is None:
            continue
        name, call_id, scope, status, error, error_code = metadata
        if call_id in baseline_call_ids:
            continue
        identity = (scope, call_id)
        if identity in seen:
            continue
        seen.add(identity)
        outcome = {"tool": name, "status": status}
        if status == "failed":
            known_presentation_codes = {
                "analysis_presentation_consistency_failed",
                "analysis_presentation_validation_failed",
                "invalid_presentation_arguments",
                "presentation_delivery_verification_failed",
                "presentation_generation_failed",
                "presentation_generation_pending",
                "presentation_quality_failed",
                "presentation_verification_attestation_failed",
                "presentation_verification_failed",
            }
            outcome.update(
                error_class=(
                    error_code
                    if name == "create_analysis_presentation"
                    and error_code in known_presentation_codes
                    else _presentation_error_class(error)
                    if name == "create_analysis_presentation"
                    else "other"
                ),
                error_sha256=digest(error.encode()),
            )
        outcomes.append(outcome)
    return outcomes


def failure_metadata(frames: list[tuple[int, dict]]) -> list[dict[str, str]]:
    """Allow-listed terminal codes only; never retain an error message."""
    known = {
        "completion_gate_failed", "console_stream_exception", "console_stream_failed",
        "empty_backend_stream", "missing_terminal_response", "rate_limit_exceeded",
        "context_length_exceeded", "invalid_api_key", "invalid_request_error",
        "authentication_error", "insufficient_quota", "timeout", "cancelled",
    }
    failures = []
    for _at, frame in frames:
        if frame.get("object") != "response" or frame.get("status") != "failed":
            continue
        error = frame.get("error")
        code = str(error.get("code") or "") if isinstance(error, dict) else ""
        failures.append({"code": code if code in known else "other",
                         "code_sha256": digest(code.encode())})
    return failures


def assistant_text(frames: list[tuple[int, dict]]) -> str:
    """Only assistant messages, never tool output or an echoed input."""
    messages: dict[str, str] = {}
    for index, (_at, frame) in enumerate(frames):
        candidates = [frame]
        if frame.get("object") == "response":
            candidates = frame.get("output") or []
        for item in candidates:
            if not isinstance(item, dict) or item.get("role") != "assistant":
                continue
            if item.get("type") not in {None, "message"}:
                continue
            content = item.get("content")
            if not isinstance(content, list):
                continue
            text = "".join(block.get("text", "") for block in content
                           if isinstance(block, dict) and block.get("type") == "text"
                           and isinstance(block.get("text"), str))
            if text:
                messages[str(item.get("id") or index)] = text
    return "\n".join(messages.values())


def assistant_surface_shape(frames: list[tuple[int, dict]]) -> dict:
    """Return only allow-listed assistant surface structure, never text."""

    message_count = 0
    messages_with_text = 0
    messages_with_chart_fence = 0
    content_types: Counter[str] = Counter()
    allowed_content_types = {
        "text", "output_text", "file", "image", "audio", "tool_call",
    }
    for _index, (_at, frame) in enumerate(frames):
        candidates = [frame]
        if frame.get("object") == "response":
            candidates = frame.get("output") or []
        for item in candidates:
            if not isinstance(item, dict) or item.get("role") != "assistant":
                continue
            if item.get("type") not in {None, "message"}:
                continue
            message_count += 1
            content = item.get("content")
            if not isinstance(content, list):
                continue
            texts = []
            for block in content:
                if not isinstance(block, dict):
                    continue
                raw_type = str(block.get("type") or "other")
                content_types[
                    raw_type if raw_type in allowed_content_types else "other"
                ] += 1
                if raw_type in {"text", "output_text"} and isinstance(
                    block.get("text"), str
                ):
                    texts.append(block["text"])
            visible = "".join(texts)
            if visible:
                messages_with_text += 1
            if "```analysis-chart" in visible:
                messages_with_chart_fence += 1
    return {
        "assistant_messages": message_count,
        "assistant_messages_with_text": messages_with_text,
        "assistant_messages_with_chart_fence": messages_with_chart_fence,
        "assistant_content_types": dict(sorted(content_types.items())),
    }


def chart_references(text: str) -> list[dict]:
    references = []
    for match in re.finditer(r"```analysis-chart\s*\n(.*?)\n```", text, re.S):
        try:
            reference = json.loads(match.group(1))
        except ValueError:
            continue
        if (isinstance(reference, dict)
                and set(reference) == {"analysis_id", "x_column", "y_column", "type"}
                and re.fullmatch(r"[0-9a-f]{64}", str(reference["analysis_id"]))
                and reference["type"] in {"bar", "line"}):
            references.append(reference)
    return references


def matching_result(state_root: Path, source_hash: str, scenario: dict,
                    expected: dict[str, str]) -> tuple[dict, str]:
    matches = []
    for path in state_root.glob("**/analysis/*/results/*.json"):
        record = sealed_record(path)
        if record.get("evidence", {}).get("source_sha256") != source_hash:
            continue
        plan = record.get("plan", {})
        if plan.get("group_by") != [scenario["group_column"]]:
            continue
        if plan.get("filters") or plan.get("deduplicate_by"):
            continue
        for metric in plan.get("metrics", []):
            if metric.get("op") != "sum" or metric.get("column") != scenario["value_column"]:
                continue
            alias = metric.get("as")
            try:
                values = {str(row[scenario["group_column"]]): Decimal(str(row[alias]))
                          for row in record["rows"]}
                gold = {key: Decimal(value) for key, value in expected.items()}
            except (KeyError, TypeError, ValueError, ArithmeticError):
                continue
            if values == gold and len(record["rows"]) == len(gold):
                matches.append((record, alias))
    require(bool(matches), "expected_source_backed_analysis_missing")
    return matches[-1]


def chart_category_labels(chart_roots: list[ElementTree.Element]) -> set[str]:
    """Read category points across the legal OOXML cache representations."""

    namespace = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    return {
        str(node.text)
        for chart in chart_roots
        for node in chart.findall(".//c:cat//c:pt/c:v", namespace)
        if node.text is not None
    }


def verify_deliveries(state_root: Path, result: dict, required: set[str],
                      artifact_root: Path) -> dict[str, str]:
    """Independently reopen final files and verify their sealed lineage."""
    verified = {}
    for path in state_root.glob("**/analysis/*/deliveries/*.json"):
        receipt = sealed_record(path)
        if receipt.get("analysis_id") != result["analysis_id"]:
            continue
        kind = receipt.get("kind")
        artifact = Path(receipt["artifact_path"]).resolve()
        require(artifact.is_relative_to(artifact_root.resolve()), "delivery_outside_test_workspace")
        blob = artifact.read_bytes()
        require(digest(blob) == receipt["content_sha256"], "final_file_hash_mismatch")
        require(bool(receipt["bound_elements"]), "delivery_bindings_missing")
        verification = receipt["verification"]
        require(verification["delivery_consistency"]["status"] == "passed", "delivery_not_verified")
        require(verification["semantic_review"]["status"] == "pending", "semantic_review_falsely_passed")
        if kind == "svg":
            ElementTree.fromstring(blob)
            require(result["evidence"]["source_sha256"] in blob.decode(), "svg_source_hash_missing")
        elif kind == "docx":
            from docx import Document
            document = Document(io.BytesIO(blob))
            text = "\n".join([paragraph.text for paragraph in document.paragraphs] +
                             [cell.text for table in document.tables for row in table.rows for cell in row.cells])
            require(result["analysis_id"] in text, "docx_analysis_id_missing")
            for row in result["rows"]:
                for value in row.values():
                    require(str(value) in text, "docx_bound_value_missing")
        elif kind == "pptx":
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(blob))
            require(bool(presentation.slides), "pptx_empty")
            with zipfile.ZipFile(io.BytesIO(blob)) as archive:
                charts = [name for name in archive.namelist()
                          if re.fullmatch(r"ppt/charts/chart\d+\.xml", name)]
                require(bool(charts), "pptx_native_editable_chart_missing")
                chart_roots = [ElementTree.fromstring(archive.read(name)) for name in charts]
                namespace = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                numeric_values = {Decimal(node.text) for chart in chart_roots
                                  for node in chart.findall(".//c:numCache/c:pt/c:v", namespace)}
                labels = chart_category_labels(chart_roots)
                require(labels.issuperset(str(row[result["plan"]["group_by"][0]]) for row in result["rows"]),
                        "pptx_chart_bound_label_missing")
                for metric in result["plan"]["metrics"]:
                    if metric.get("op") == "sum":
                        require(numeric_values.issuperset(Decimal(str(row[metric["as"]])) for row in result["rows"]),
                                "pptx_chart_bound_value_missing")
                require(any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx")
                            for name in archive.namelist()), "pptx_embedded_chart_workbook_missing")
        elif kind != "evidence":
            raise SmokeFailure("unexpected_delivery_kind")
        verified[kind] = receipt["content_sha256"]
    allowed_kinds = {"evidence", "svg", "docx", "pptx"}
    require(required.issubset(allowed_kinds), "requested_verified_delivery_kind_invalid")
    missing = sorted(required.difference(verified))
    if missing:
        raise SmokeFailure("requested_verified_delivery_missing_" + "_".join(missing))
    return verified


def stage_utility_config(source_working_dir: Path, working: Path,
                         provider_id: str | None, model: str | None) -> dict:
    """Copy configuration metadata only; overrides affect the test clone."""
    configured = json.loads((source_working_dir / "config.json").read_bytes()).get("utility_intelligence", {})
    utility = {key: configured[key] for key in ("mode", "timeout_seconds", "background_processing_enabled")
               if key in configured}
    slot = configured.get("model", {})
    if isinstance(slot, dict):
        utility["model"] = {key: slot[key] for key in ("provider_id", "model") if key in slot}
    require(bool(provider_id) == bool(model), "utility_override_incomplete")
    if provider_id:
        utility.update(mode="dedicated", model={"provider_id": provider_id, "model": model},
                       background_processing_enabled=True)
    path = working / "config.json"
    config = json.loads(path.read_bytes())
    config["utility_intelligence"] = utility
    path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    return {"overridden_in_test_copy": bool(provider_id), "mode": utility.get("mode", "follow_active"),
            **utility.get("model", {})}


def run(*, source_working_dir: Path, source_secret_dir: Path, scenario: dict,
        provider_id: str | None = None, model: str | None = None,
        utility_provider_id: str | None = None, utility_model: str | None = None,
        timeout: float = 420, progress: Callable[[dict], None] | None = None) -> dict:
    spec = importlib.util.spec_from_file_location(
        "analysis_provider_helpers", ROOT / "scripts/source_configured_provider_stream_smoke.py")
    helpers = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(helpers)
    smoke_root = Path(tempfile.mkdtemp(prefix="lantai-analysis-provider-smoke-"))
    process = None
    state_root: Path | None = None
    ledger_snapshot: tuple[frozenset[str], int, frozenset[str]] | None = None
    pool = ThreadPoolExecutor(max_workers=1)
    evidence = {"ok": False, "same_conversation": True, "turns": [], "checks": {}}
    try:
        require(isinstance(scenario.get("turns"), list) and 2 <= len(scenario["turns"]) <= 12,
                "scenario_turn_count_invalid")
        filename = scenario.get("filename")
        require(isinstance(filename, str) and Path(filename).name == filename
                and filename.endswith(".csv"), "scenario_filename_invalid")
        working, secret, selected = helpers._stage_profile(
            source_working_dir=source_working_dir, source_secret_dir=source_secret_dir,
            agent_id="default", target_root=smoke_root)
        evidence["utility"] = stage_utility_config(source_working_dir, working, utility_provider_id, utility_model)
        if provider_id or model:
            require(bool(provider_id and model), "provider_override_incomplete")
            agent_path = Path(selected["agent_path"])
            agent = json.loads(agent_path.read_bytes())
            agent["active_model"] = {"provider_id": provider_id, "model": model}
            agent_path.write_text(json.dumps(agent, ensure_ascii=False), encoding="utf-8")
            (secret / "providers/active_model.json").write_text(
                json.dumps(agent["active_model"]), encoding="utf-8")
            selected.update(provider_id=provider_id, model=model)
        evidence.update(provider_id=selected["provider_id"], model=selected["model"])
        env = os.environ.copy()
        for key in ("LANTAI_RUNTIME_ENGINE", "LANTAI_RUNTIME_ENGINE_DEFAULT",
                    "COPAW_RUNTIME_ENGINE", "COPAW_RUNTIME_ENGINE_DEFAULT"):
            env.pop(key, None)
        desktop_token = uuid.uuid4().hex
        env.update(LANTAI_WORKING_DIR=str(working), LANTAI_SECRET_DIR=str(secret),
                   LANTAI_DESKTOP_TOKEN=desktop_token,
                   LANTAI_AUTH_ENABLED="0", LANTAI_DISABLE_KEYRING="1",
                   LANTAI_SKIP_KEYRING="1", LANTAI_REMINDER_DELIVERY_WORKER="0",
                   PYTHONPATH=os.pathsep.join([str(ROOT / "src"), str(ROOT)]))
        env.update(source_presentation_runtime_env())
        base_url = f"http://127.0.0.1:{helpers._available_port()}"
        process = subprocess.Popen(
            [str(ROOT / ".venv/bin/python"), "-c",
             "from scripts.analysis_configured_provider_smoke import observed_backend; observed_backend()",
             "app", "--host", "127.0.0.1",
             "--port", base_url.rsplit(":", 1)[-1], "--log-level", "warning"],
            cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        helpers._wait_for_health(base_url, process, min(timeout, 60))

        def request(method: str, path: str, payload: dict | None = None) -> Any:
            data = json.dumps(payload).encode() if payload is not None else None
            try:
                with urlopen(Request(base_url + path, data=data, method=method,
                                     headers={"Content-Type": "application/json", "X-Agent-Id": "default",
                                              "X-Lantai-Desktop-Token": desktop_token}),
                             timeout=10) as response:
                    return json.loads(response.read())
            except HTTPError as error:
                raise SmokeFailure(f"http_{error.code}") from None

        session = "console:analysis-provider-smoke-" + uuid.uuid4().hex
        workspace = request("POST", "/api/chats/topic-workspace", {
            "session_id": session, "user_id": "default", "channel": "console", "name": "Analysis smoke"})
        chat = workspace["chat"]
        session = chat["session_id"]
        artifact_root = Path(workspace["effective_workspace_dir"])
        state_root = Path(workspace["state_workspace_dir"])
        require(artifact_root.resolve().is_relative_to(smoke_root.resolve()), "test_workspace_not_isolated")
        (working / "analysis-smoke-scope.json").write_text(json.dumps({
            "session": session, "artifact_root": str(artifact_root), "state_root": str(state_root),
        }), encoding="utf-8")
        projection_observation_path = (
            working / "analysis-smoke-projection-observations.jsonl"
        )
        source = artifact_root / filename
        source.write_text(scenario["source_data"], encoding="utf-8")
        original_files: dict[Path, str] = {}
        original_result = None
        required_kinds: set[str] = set()

        def read_turn(prompt: str) -> list[tuple[int, dict]]:
            payload = helpers._configured_provider_request_payload(uuid.uuid4().hex, session_id=session, prompt=prompt)
            payload["chat_id"] = chat["id"]
            frames, lines = [], []
            started = time.monotonic()
            try:
                with urlopen(Request(base_url + "/api/console/chat", data=json.dumps(payload).encode(),
                                     headers={"Content-Type": "application/json", "X-Agent-Id": "default",
                                              "X-Lantai-Desktop-Token": desktop_token}, method="POST"),
                             timeout=timeout) as response:
                    while line := response.readline():
                        text = line.decode().rstrip("\r\n")
                        if text.startswith("data:"):
                            lines.append(text[5:].lstrip())
                        elif not text and lines:
                            frames.append((round((time.monotonic() - started) * 1000), json.loads("\n".join(lines))))
                            lines = []
            except HTTPError as error:
                raise SmokeFailure(f"chat_http_{error.code}") from None
            return frames

        for number, turn in enumerate(scenario["turns"], 1):
            if "source_data" in turn:
                for path, checksum in original_files.items():
                    require(digest(path.read_bytes()) == checksum, "old_artifact_changed_before_refresh")
                source.write_text(turn["source_data"], encoding="utf-8")
            from lantai.app.experience.events import collect_workspace_events

            ledger_snapshot = event_ledger_snapshot(
                collect_workspace_events(state_root),
            )
            projection_observation_baseline = projection_observation_snapshot(
                projection_observation_path,
            )
            started = time.monotonic()
            future = pool.submit(read_turn, turn["prompt"])
            approvals = 0
            while not future.done():
                require(time.monotonic() - started < timeout, "turn_deadline_exceeded")
                pending = request("GET", "/api/approval/list?" + urlencode({"session_id": session}))
                for item in pending.get("pending_approvals", []):
                    validate_approval(item, session)
                    reply = request("POST", "/api/approval/approve", {
                        "request_id": item["request_id"], "session_id": session,
                        "user_id": "default", "decision_scope": "once"})
                    require(reply.get("success") is True, "scoped_approval_failed")
                    approvals += 1
                time.sleep(0.2)
            frames = future.result()
            trace = {"turn": number, "frames": len(frames), "approvals": approvals,
                     "elapsed_ms": round((time.monotonic() - started) * 1000),
                     "reasoning_frames": sum(frame.get("type") == "reasoning" for _, frame in frames),
                     "completed": any(frame.get("object") == "response" and frame.get("status") == "completed"
                                      for _, frame in frames)}
            trace["analysis_records"] = len(list(state_root.glob("**/analysis/*/results/*.json")))
            trace["dataset_records"] = len(list(state_root.glob("**/analysis/*/datasets/*.json")))
            ledger_events = collect_workspace_events(state_root)
            trace["analysis_tool_calls"] = event_ledger_tool_counts(
                ledger_events, ledger_snapshot,
            )
            trace["analysis_tool_outcomes"] = event_ledger_tool_outcomes(
                ledger_events, ledger_snapshot,
            )
            trace["assistant_projection_ledger_shape"] = (
                event_ledger_assistant_projection_shape(
                    ledger_events,
                    ledger_snapshot,
                )
            )
            trace["runtime_terminal_shape"] = event_ledger_terminal_shape(
                ledger_events,
                ledger_snapshot,
            )
            trace["analysis_tool_count_source"] = "workspace_event_ledger"
            trace["assistant_surface_shape"] = assistant_surface_shape(frames)
            trace["assistant_projection_observations"] = (
                projection_observations_since(
                    projection_observation_path,
                    projection_observation_baseline,
                )
            )
            trace["failures"] = failure_metadata(frames)
            evidence["turns"].append(trace)
            if progress:
                progress(trace)
            require(trace["completed"], "chat_turn_not_completed")
            result, alias = matching_result(state_root, digest(source.read_bytes()), scenario, turn["expected"])
            require(result["verification"]["semantic_review"]["status"] == "pending", "semantic_review_falsely_passed")
            require(result["evidence"]["source_row_count"] == len(list(csv.DictReader(io.StringIO(source.read_text())))),
                    "source_row_count_mismatch")
            if original_result is None:
                original_result = result
            elif "source_data" not in turn:
                require(result["analysis_id"] == original_result["analysis_id"], "followup_replaced_analysis")
            else:
                require(result["analysis_id"] != original_result["analysis_id"] and
                        result["evidence"]["source_sha256"] != original_result["evidence"]["source_sha256"],
                        "refresh_did_not_create_new_version")
                require(result["plan"] == original_result["plan"], "refresh_changed_measurement_plan")
                for path, checksum in original_files.items():
                    require(digest(path.read_bytes()) == checksum, "refresh_overwrote_old_artifact")
                evidence["checks"]["refresh_preserves_old_deliveries"] = True
            required_kinds.update(turn.get("delivery_kinds", []))
            if "source_data" not in turn:
                verified = verify_deliveries(state_root, result, required_kinds, artifact_root)
                evidence["checks"]["delivery_hashes"] = verified
                for path in state_root.glob("**/analysis/*/deliveries/*.json"):
                    receipt = sealed_record(path)
                    original_files[Path(receipt["artifact_path"])] = receipt["content_sha256"]
            if turn.get("chat_chart"):
                references = chart_references(assistant_text(frames))
                reference = next((item for item in references if item["analysis_id"] == result["analysis_id"]
                                  and item["x_column"] == scenario["group_column"] and item["y_column"] == alias), None)
                require(reference is not None, "assistant_chat_chart_reference_missing")
                path = f"/api/chats/{chat['id']}/analysis/{result['analysis_id']}/chart?"
                chart = request("GET", path + urlencode({"x_column": scenario["group_column"], "y_column": alias}))
                require({item["label"]: Decimal(item["value"]) for item in chart["points"]} ==
                        {key: Decimal(value) for key, value in turn["expected"].items()}, "chat_chart_values_mismatch")
                evidence["checks"]["assistant_chart_uses_sealed_server_data"] = True
            evidence["checks"]["analysis_id"] = result["analysis_id"]
            evidence["checks"]["source_sha256"] = result["evidence"]["source_sha256"]
        request("DELETE", f"/api/chats/{chat['id']}")
        try:
            request("GET", f"/api/chats/{chat['id']}/analysis/{result['analysis_id']}/chart?" +
                    urlencode({"x_column": scenario["group_column"], "y_column": alias}))
        except SmokeFailure as failure:
            require(str(failure) == "http_404", "deleted_chat_unexpected_status")
        else:
            raise SmokeFailure("deleted_chat_still_exposes_chart")
        evidence["checks"]["deleted_chat_chart_unavailable"] = True
        evidence["ok"] = True
    except Exception as error:
        evidence["failure_code"] = str(error) if isinstance(error, SmokeFailure) else type(error).__name__
        if state_root is not None and ledger_snapshot is not None:
            try:
                from lantai.app.experience.events import collect_workspace_events

                failure_events = collect_workspace_events(state_root)
                evidence["checks"]["failure_tool_calls"] = (
                    event_ledger_tool_counts(failure_events, ledger_snapshot)
                )
                evidence["checks"]["failure_tool_outcomes"] = (
                    event_ledger_tool_outcomes(failure_events, ledger_snapshot)
                )
                evidence["checks"]["failure_assistant_projection_ledger_shape"] = (
                    event_ledger_assistant_projection_shape(
                        failure_events,
                        ledger_snapshot,
                    )
                )
            except (OSError, ValueError):
                evidence["checks"]["failure_event_ledger_unavailable"] = True
    finally:
        helpers._terminate(process)
        pool.shutdown(wait=True, cancel_futures=True)
        shutil.rmtree(smoke_root)
        evidence["temporary_data_removed"] = not smoke_root.exists()
    return evidence


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-working-dir", required=True, type=Path)
    parser.add_argument("--source-secret-dir", required=True, type=Path)
    parser.add_argument("--provider-id")
    parser.add_argument("--model")
    parser.add_argument("--utility-provider-id")
    parser.add_argument("--utility-model")
    parser.add_argument("--timeout", type=float, default=420)
    args = parser.parse_args()
    scenario = json.load(sys.stdin)
    result = run(**vars(args), scenario=scenario,
                 progress=lambda value: print(json.dumps({"progress": value}), flush=True))
    print(json.dumps(result, ensure_ascii=False), flush=True)
    raise SystemExit(0 if result["ok"] else 1)


if __name__ == "__main__":
    main()
