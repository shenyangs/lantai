#!/usr/bin/env python3
"""Build the synthetic D06 strict-template development deck from its PPTX fixture.

This is a bounded development author, not a second Lantai production runtime.
The fixture's five visible HARD shapes are retained on every slide. The brief
mentions a master logo, but the supplied fixture has no master/layout shapes;
that discrepancy is reported rather than represented as a passed preservation.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


PURPLE = RGBColor(0x68, 0x3F, 0xA5)
INK = RGBColor(0x24, 0x26, 0x32)
MUTED = RGBColor(0x66, 0x69, 0x76)
PALE = RGBColor(0xF4, 0xF1, 0xF8)
PALE_GREEN = RGBColor(0xE9, 0xF2, 0xEC)
GREEN = RGBColor(0x3D, 0x75, 0x5C)
RULE = RGBColor(0xDC, 0xD9, 0xE3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "PingFang SC"
SOURCE = "SRC-B01（合成开发材料）"


def text(slide, value, x, y, w, h, *, size=18, color=INK, bold=False,
         align=PP_ALIGN.LEFT, margin=0, valign=MSO_ANCHOR.MIDDLE):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.15
        for run in paragraph.runs:
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def rect(slide, x, y, w, h, fill, *, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def note(slide, detail):
    slide.notes_slide.notes_text_frame.text = (
        f"来源：{SOURCE}；附件 D06_template_brief.md 与 strict-template-fixture.pptx。"
        f"本页为计划/建议，不是已完成事实。{detail}"
    )


def set_template_text(shape, value):
    frame = shape.text_frame
    if frame.paragraphs and frame.paragraphs[0].runs:
        frame.paragraphs[0].runs[0].text = value
    else:
        frame.text = value


def prepare_slide(prs, index, title, source_shapes):
    if index <= 3:
        slide = prs.slides[index - 1]
    else:
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        for source_shape in source_shapes:
            slide.shapes._spTree.insert_element_before(deepcopy(source_shape), "p:extLst")
    expected = ["Shape 0", "Text 1", "Shape 2", "Text 3", "Text 4", "Text 5"]
    shapes = list(slide.shapes)
    if [shape.name for shape in shapes[:6]] != expected:
        raise ValueError(f"Unexpected fixture shapes on slide {index}")
    if shapes[3].text != "青禾实验室 · 内部讨论":
        raise ValueError("HARD footer differs from fixture")
    placeholder = shapes[5]
    placeholder._element.getparent().remove(placeholder._element)
    set_template_text(shapes[1], title)
    set_template_text(shapes[4], str(index))
    # The actual fixture has no master logo. This in-content identifier is a
    # visible supplement, never counted as preservation of a nonexistent logo.
    text(slide, "青禾实验室", 11.35, 0.62, 1.08, 0.35,
         size=12, color=PURPLE, bold=True, align=PP_ALIGN.RIGHT)
    return slide


def eyebrow(slide, label, *, y=1.50):
    text(slide, label, 0.94, y, 11.3, 0.34, size=13, color=PURPLE, bold=True)


def accent_number(slide, number, label, x, y, *, width=3.7):
    text(slide, number, x, y, width, 0.68, size=34, color=PURPLE, bold=True)
    text(slide, label, x, y + 0.75, width, 0.62, size=17, color=INK)


def panel(slide, x, y, w, h, *, fill=PALE, stripe=PURPLE):
    rect(slide, x, y, w, h, fill, radius=True)
    rect(slide, x, y, 0.055, h, stripe)


def make_deck(fixture: Path, output: Path, revision: str):
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite {output}")
    prs = Presentation(fixture)
    if len(prs.slides) != 3 or len(prs.slide_masters) != 1:
        raise ValueError("Expected exact three-slide synthetic fixture")
    if prs.slide_masters[0].shapes or prs.slide_layouts[0].shapes:
        raise ValueError("Fixture changed: inspect master/layout logo before authoring")
    source_shapes = [deepcopy(shape._element) for shape in prs.slides[1].shapes]

    # 01 — background. A need, not a fabricated claim about current failures.
    s = prepare_slide(prs, 1, "四周，把资料与结果放到同一条链上", source_shapes)
    eyebrow(s, "项目启动｜背景")
    text(s, "资料与结果需要统一。", 0.94, 2.02, 7.2, 0.85,
         size=34, bold=True)
    text(s, "这是一份拟议启动方案，不是进展汇报。", 0.97, 2.94, 6.5, 0.47,
         size=19, color=MUTED)
    rect(s, 0.95, 3.74, 7.15, 0.025, RULE)
    text(s, "盘点输入 → 可编辑初稿 → 双人评审 → 内部候选", 0.96, 4.05, 7.7, 0.77,
         size=24, color=PURPLE, bold=True)
    panel(s, 9.18, 2.07, 2.56, 3.55, fill=PALE_GREEN, stripe=GREEN)
    accent_number(s, "04", "周；每阶段一周", 9.48, 2.47, width=1.9)
    accent_number(s, "04", "位拟定团队成员", 9.48, 4.1, width=1.9)
    text(s, "预算尚未确定", 0.97, 5.72, 4.4, 0.38, size=16, color=MUTED)
    note(s, "四周及团队人数来自原 brief；预算明确为未定。")

    # 02 — goal. Separate project target from proof of success.
    s = prepare_slide(prs, 2, "目标：形成可审、可改、可内部使用的候选", source_shapes)
    eyebrow(s, "目标｜三层结果")
    for n, heading, body, x in [
        ("01", "统一入口", "先盘点输入与结果，建立同一份工作依据。", 0.95),
        ("02", "原生可改", "制作可编辑初稿，供内容与视觉继续修订。", 5.09),
        ("03", "双审再发", "两位评审查事实与视觉，再发布内部候选。", 9.23),
    ]:
        panel(s, x, 2.18, 3.57, 3.10, fill=PALE if x != 5.09 else PALE_GREEN,
              stripe=PURPLE if x != 5.09 else GREEN)
        text(s, n, x + 0.27, 2.46, 0.85, 0.57, size=28, color=PURPLE, bold=True)
        text(s, heading, x + 0.27, 3.14, 2.9, 0.47, size=24, bold=True)
        text(s, body, x + 0.27, 3.82, 2.91, 0.90, size=16, color=MUTED)
    text(s, "以上均为拟定目标；本稿不宣称候选已获批准。", 0.96, 5.85, 10.5, 0.42,
         size=16, color=MUTED)
    note(s, "目标从原 brief 的四阶段计划提炼；不把计划表述为完成事实。")

    # 03 — scope. The unknowns remain explicit.
    s = prepare_slide(prs, 3, "范围：先做内部候选，不越过未定义边界", source_shapes)
    eyebrow(s, "范围｜已给定与待定")
    panel(s, 0.95, 2.08, 5.48, 3.72, fill=PALE_GREEN, stripe=GREEN)
    text(s, "本阶段明确包含", 1.24, 2.44, 4.7, 0.48, size=23, bold=True)
    text(s, "输入盘点\n可编辑初稿\n事实与视觉双审\n内部候选", 1.25, 3.16, 4.66, 2.03,
         size=21, color=INK, valign=MSO_ANCHOR.TOP)
    panel(s, 6.79, 2.08, 5.48, 3.72, fill=PALE, stripe=PURPLE)
    text(s, "材料未定义", 7.08, 2.44, 4.7, 0.48, size=23, bold=True)
    text(s, "预算金额\n对外发布\n量化效果指标\n后续扩展范围", 7.09, 3.16, 4.63, 2.03,
         size=21, color=MUTED, valign=MSO_ANCHOR.TOP)
    note(s, "右栏表示 brief 中没有给出这些事项，不代表项目明确禁止。")

    # 04 — deliverables. Native editable ownership is a proposed acceptance.
    s = prepare_slide(prs, 4, "交付物：每周留下可接续的真实对象", source_shapes)
    eyebrow(s, "交付｜四件计划产物")
    rows = [
        ("W1", "输入盘点", "资料与结果的清单", "进入初稿前校对来源"),
        ("W2", "可编辑初稿", "文字、表格、图解可继续修改", "留给评审真实改动空间"),
        ("W3", "评审记录", "事实与视觉两线意见", "指向具体页与对象"),
        ("W4", "内部候选", "修订后版本与未决边界", "供内部讨论，不等于正式发布"),
    ]
    for i, (week, name, thing, meaning) in enumerate(rows):
        y = 2.0 + i * 0.98
        rect(s, 0.95, y, 11.55, 0.82, WHITE if i % 2 else PALE)
        rect(s, 0.95, y, 0.08, 0.82, PURPLE if i % 2 else GREEN)
        text(s, week, 1.22, y + 0.10, 0.77, 0.55, size=18, color=PURPLE, bold=True)
        text(s, name, 2.11, y + 0.10, 2.05, 0.55, size=20, bold=True)
        text(s, thing, 4.29, y + 0.10, 4.05, 0.55, size=17)
        text(s, meaning, 8.56, y + 0.10, 3.58, 0.55, size=15, color=MUTED)
    note(s, "交付物是基于四阶段的作者建议；原 brief 没有给定文件命名、数量或完成状态。")

    # 05 — roles. Proposed allocation explicitly distinguished from facts.
    s = prepare_slide(prs, 5, "分工：四人团队，两位评审需预留席位", source_shapes)
    eyebrow(s, "分工｜人数为材料事实；职责为启动建议")
    rect(s, 0.96, 2.08, 11.47, 0.035, RULE)
    groups = [
        ("2", "工程", "搭建可编辑初稿\n处理文件与版本", 0.95, PURPLE),
        ("1", "设计", "建立视觉规则\n修订重点页", 5.07, GREEN),
        ("1", "内容编辑", "盘点材料与主张\n记录来源边界", 9.18, PURPLE),
    ]
    for n, role, duty, x, color in groups:
        panel(s, x, 2.45, 3.2, 2.81, fill=PALE_GREEN if color == GREEN else PALE,
              stripe=color)
        text(s, n, x + 0.26, 2.75, 0.78, 0.71, size=37, bold=True, color=color)
        text(s, role, x + 1.11, 2.90, 1.76, 0.47, size=22, bold=True)
        text(s, duty, x + 0.26, 3.83, 2.57, 1.05, size=17, color=MUTED)
    text(s, "评审：原材料只写“两位评审”；具体人选与授权尚未确定。", 0.97, 5.76, 10.93, 0.54,
         size=17, color=INK)
    note(s, "2 工程 + 1 设计 + 1 内容编辑为给定团队规模。职责映射是建议，不是既定人事安排。")

    # 06 — workflow. Four visually distinct gates, no fake progress.
    s = prepare_slide(prs, 6, "流程：四道门，先证据后候选", source_shapes)
    eyebrow(s, "流程｜按一周一阶段规划")
    steps = [
        ("01", "盘点输入", "资料与结果\n先归一", PURPLE),
        ("02", "制作初稿", "关键对象\n保持可编辑", GREEN),
        ("03", "双人评审", "事实与视觉\n分别核查", PURPLE),
        ("04", "内部候选", "修订后提交\n供内部讨论", GREEN),
    ]
    for i, (n, title, body, color) in enumerate(steps):
        x = 0.95 + i * 3.06
        panel(s, x, 2.35, 2.68, 3.16, fill=PALE if i % 2 == 0 else PALE_GREEN,
              stripe=color)
        text(s, n, x + 0.22, 2.64, 0.73, 0.55, size=28, color=color, bold=True)
        text(s, title, x + 0.22, 3.33, 2.20, 0.46, size=21, bold=True)
        text(s, body, x + 0.22, 4.02, 2.23, 0.91, size=17, color=MUTED)
        if i < 3:
            text(s, "→", x + 2.67, 3.59, 0.39, 0.61, size=27,
                 color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    text(s, "阶段顺序为计划；本页不表示进度已推进到任一门。", 0.96, 5.93, 10.91, 0.41,
         size=16, color=MUTED)
    note(s, "四阶段名称及每阶段一周来自原始 brief。")

    # 07 — acceptance. Proposed criteria, not passed checkmarks.
    s = prepare_slide(prs, 7, "验收：把“能打开”与“能接手”分开", source_shapes)
    eyebrow(s, "验收｜建议检查项；当前均未执行")
    checks = [
        ("来源", "输入与结果可追溯；未定事项不补写成事实"),
        ("编辑", "核心文字、表格与图解可在 PPT 中继续改"),
        ("双审", "两位评审分别检查事实与视觉并留下修改记录"),
        ("场景", "内部候选能够用于讨论，未决边界清楚可见"),
    ]
    for i, (label, body) in enumerate(checks):
        y = 2.08 + i * 0.93
        rect(s, 0.95, y, 11.45, 0.75, PALE if i % 2 == 0 else WHITE)
        rect(s, 0.95, y, 0.07, 0.75, PURPLE)
        text(s, label, 1.28, y + 0.12, 1.12, 0.45, size=19, bold=True, color=PURPLE)
        text(s, body, 2.6, y + 0.11, 9.38, 0.47, size=17)
    text(s, "状态：待检查。作者稿、渲染图或自评不能替代评审与实际编辑保存。",
         0.97, 5.96, 11.3, 0.48, size=16, color=MUTED)
    note(s, "验收条款由作者根据 brief 拟定；没有已通过的事实、视觉或编辑性结论。")

    # 08 — risks. Conditions rather than invented incidents.
    s = prepare_slide(prs, 8, "风险：先暴露未决项，不给它们填假答案", source_shapes)
    eyebrow(s, "风险｜潜在项与预案")
    risks = [
        ("预算未定", "首周确定额度与审批边界", "已知未决"),
        ("输入缺口", "盘点时记录缺失来源，暂不编造", "需检查"),
        ("评审拥堵", "提前确认两位评审的时间", "需确认"),
        ("编辑性不足", "在初稿阶段试改核心对象", "需测试"),
    ]
    for i, (risk, action, status) in enumerate(risks):
        y = 2.02 + i * 0.98
        panel(s, 0.95, y, 11.53, 0.81, fill=PALE if i % 2 == 0 else WHITE,
              stripe=PURPLE if i == 0 else GREEN)
        text(s, risk, 1.25, y + 0.12, 2.16, 0.47, size=19, bold=True)
        text(s, action, 3.57, y + 0.12, 6.25, 0.47, size=17)
        text(s, status, 10.47, y + 0.13, 1.54, 0.45, size=15, color=MUTED,
             align=PP_ALIGN.RIGHT)
    note(s, "预算未定为材料事实；其余为条件性风险推断，不表示这些问题已经发生。")

    # 09 — time. One native text-and-shape timeline, four weeks exact.
    s = prepare_slide(prs, 9, "时间：四周四步，不压缩评审", source_shapes)
    eyebrow(s, "时间｜计划，不是完成进度")
    weeks = [
        ("第 1 周", "盘点输入", "形成统一资料清单"),
        ("第 2 周", "可编辑初稿", "内容与图解可继续修改"),
        ("第 3 周", "双人评审", "事实 + 视觉"),
        ("第 4 周", "内部候选", "修订并提交讨论"),
    ]
    rect(s, 1.47, 3.18, 10.45, 0.035, RULE)
    for i, (week, name, output_hint) in enumerate(weeks):
        x = 0.95 + i * 3.05
        rect(s, x, 2.79, 0.82, 0.82, PURPLE if i % 2 == 0 else GREEN, radius=True)
        text(s, str(i + 1), x + 0.22, 2.90, 0.38, 0.55,
             size=23, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        text(s, week, x, 2.18, 2.43, 0.45, size=18, color=PURPLE, bold=True)
        text(s, name, x, 3.96, 2.67, 0.49, size=21, bold=True)
        text(s, output_hint, x, 4.69, 2.61, 0.84, size=16, color=MUTED)
    text(s, "每阶段一周；若输入、评审或预算未决，需重新评估，而非暗中跳步。",
         0.96, 5.9, 11.25, 0.46, size=16, color=MUTED)
    note(s, "四周、每阶段一周是原 brief 的计划值；没有日期或已完成比例。")

    # 10 — decisions. Keep the meeting actionable without invented approvals.
    s = prepare_slide(prs, 10, "今天需要定下三件事", source_shapes)
    eyebrow(s, "决策｜请启动会确认")
    items = [
        ("1", "确认输入边界", "哪些资料与结果进入第 1 周盘点？"),
        ("2", "确定评审人选", "两位评审如何分担事实与视觉检查？"),
        ("3", "明确预算路径", "金额尚未确定，谁负责审批与何时答复？"),
    ]
    for i, (number, headline, question) in enumerate(items):
        y = 2.04 + i * 1.24
        panel(s, 0.95, y, 11.47, 1.07,
              fill=PALE_GREEN if i == 1 else PALE,
              stripe=GREEN if i == 1 else PURPLE)
        text(s, number, 1.27, y + 0.19, 0.62, 0.62, size=30,
             color=PURPLE, bold=True)
        text(s, headline, 2.15, y + 0.10, 3.01, 0.43, size=20, bold=True)
        text(s, question, 2.16, y + 0.54, 9.54, 0.43, size=17, color=MUTED)
    text(s, "以上均待决策；本稿不代替会议纪要或授权。", 0.97, 6.07, 11.3, 0.40,
         size=16, color=MUTED)
    note(s, "三项为依照原材料的启动会建议，不是已有决定。")

    if len(prs.slides) != 10:
        raise AssertionError("D06 must contain ten complete slides")
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--fixture", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--revision", default="r1")
    args = parser.parse_args()
    make_deck(args.fixture, args.output, args.revision)
    print(args.output)


if __name__ == "__main__":
    main()
